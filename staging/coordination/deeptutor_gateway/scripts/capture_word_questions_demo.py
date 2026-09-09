"""Capture only owned Qt widgets with synthetic teaching material, no settings."""

from __future__ import annotations

import argparse
import json
import os
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(ROOT))


def build_demo(folder):
    from docx import Document

    source = folder / "演示讲义 电解质与电离.docx"
    if source.exists():
        raise RuntimeError("Use a fresh demo folder")
    folder.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("电解质与电离", 1)
    doc.add_paragraph("知识点01 电解质的判断")
    doc.add_paragraph("【即学即练1】下列物质中属于电解质的是（　）")
    doc.add_paragraph("A．铜　　B．氯化钠　　C．乙醇　　D．蔗糖")
    doc.add_paragraph("【答案】B")
    doc.add_paragraph(
        "【解析】氯化钠在水溶液或熔融状态下能导电，是电解质。铜是单质，不属于电解质。"
    )
    doc.add_paragraph("知识点02 电离方程式")
    doc.add_paragraph("【典例2】写出氯化钠在水中的电离方程式。")
    p = doc.add_paragraph("【答案】NaCl = Na")
    p.add_run("+").font.superscript = True
    p.add_run(" + Cl")
    p.add_run("−").font.superscript = True
    doc.add_paragraph("【解析】分别检查元素种类与数目、电荷总数是否守恒。")
    doc.add_paragraph(
        "【变式2-1】氯化钠固体不导电，能否据此判断它不是电解质？说明理由。"
    )
    doc.add_paragraph("【答案】不能。固体中离子不能自由移动；溶于水或熔融后可导电。")
    doc.save(source)
    return source


def capture(folder, output):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtCore import Qt
    from PySide6.QtGui import QFont, QFontDatabase
    from PySide6.QtTest import QTest

    from integrations.deeptutor_shchem_v1.desktop_facade import DesktopWorkbenchFacade
    from integrations.deeptutor_shchem_v1.desktop_paths import DesktopPaths
    from integrations.deeptutor_shchem_v1.desktop_workbench.app import (
        create_application,
    )
    from integrations.deeptutor_shchem_v1.desktop_workbench.dialogs import ImportDialog
    from integrations.deeptutor_shchem_v1.desktop_workbench.tasks import (
        DesktopTaskBridge,
    )
    from integrations.deeptutor_shchem_v1.desktop_workbench.word_question_dialog import (
        WordQuestionDialog,
    )

    app = create_application([])
    for font in ("msyh.ttc", "msyhbd.ttc", "msyhl.ttc", "arial.ttf", "segoeui.ttf"):
        QFontDatabase.addApplicationFont("C:/Windows/Fonts/" + font)
    app.setFont(QFont("Microsoft YaHei", 10))
    state = folder / "isolated-state"

    class DemoProviders:
        def list_metadata(self):
            return ()

        def __getattr__(self, name):
            raise RuntimeError("Demo must not access provider configuration: " + name)

    facade = DesktopWorkbenchFacade(
        DesktopPaths.from_workspace(ROOT, state_root=state),
        provider_store=DemoProviders(),
    )
    facade.word_question_save_selection([])
    source = folder / "演示讲义 电解质与电离.docx"
    receipt = facade.save_visual_import_batch(
        handout_files=(source,), source_type="教师讲义"
    )
    tasks = DesktopTaskBridge()
    output.mkdir(parents=True, exist_ok=True)

    def settle(condition, seconds=15):
        deadline = time.monotonic() + seconds
        while time.monotonic() < deadline:
            app.processEvents()
            if condition():
                for _ in range(5):
                    app.processEvents()
                return
            time.sleep(0.015)
        raise RuntimeError("UI did not settle")

    dialog = WordQuestionDialog(facade, tasks, batch_id=receipt.batch_id)
    dialog.resize(1200, 850)
    dialog.show()
    settle(lambda: not dialog._catalog_busy and dialog.question_list.count() == 3)
    dialog.question_list.item(0).setCheckState(Qt.CheckState.Checked)
    dialog.question_list.item(1).setCheckState(Qt.CheckState.Checked)
    dialog.question_list.setCurrentRow(0)
    settle(
        lambda: not dialog._selection_save_busy and not dialog._save_timer.isActive()
    )
    dialog.grab().save(str(output / "word-questions.png"))
    dialog.tabs.setCurrentIndex(1)
    QTest.qWait(250)
    dialog.grab().save(str(output / "word-answers.png"))
    dialog.preview_button.click()
    settle(lambda: not dialog._reference_busy and dialog.import_button.isEnabled())
    dialog.grab().save(str(folder / "selected-reference.png"))
    dialog.resize(700, 1020)
    QTest.qWait(250)
    dialog.grab().save(str(folder / "narrow.png"))
    selected = dialog.selections
    dialog.reject()
    settle(lambda: not dialog.isVisible())
    import_dialog = ImportDialog(facade, tasks)
    import_dialog.resize(840, 820)
    import_dialog.show()
    app.processEvents()
    import_dialog.scroll.ensureWidgetVisible(import_dialog.word_questions_button)
    app.processEvents()
    import_dialog.grab().save(str(folder / "import-entry.png"))
    import_dialog.reject()
    tasks.shutdown()
    report = {
        "synthetic_only": True,
        "question_count": 3,
        "selected_count": len(selected),
        "question_reference_confirmable": True,
        "provider_called": False,
        "screenshots": ["word-questions.png", "word-answers.png"],
    }
    (folder / "ui-check.json").write_text(
        json.dumps(report, indent=2), encoding="utf-8"
    )
    print(json.dumps(report))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--folder", type=Path, required=True)
    parser.add_argument("--build-demo", action="store_true")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    if args.build_demo:
        print(build_demo(args.folder))
    else:
        capture(args.folder, args.output or args.folder)
