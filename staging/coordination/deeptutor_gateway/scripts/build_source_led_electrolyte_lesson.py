"""Build a source-led lesson exemplar; no provider, app-state or source writes.

This is an explicit teacher-review candidate, not the generic desktop generator.
Only the inspected textbook figure is raster; all slide text is editable.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).resolve().parents[4]
BOOK = ROOT / "课本/沪科技化学必修第一册【高清教材】.pdf"
WORD = ROOT / (
    "sh-chem-db/.intake/2026-07-30-user-teaching-pack/expanded/PKG-032/"
    "第04讲 离子反应和离子方程式（复习讲义）（上海专用）（解析版）.docx"
)
EXPECTED = {
    "book": "a565f0a15ffd10c704f4be42bfe7200c125b68959d11ef582acdc45dde2ccf22",
    "word": "d60317b8e533b957943e98b481305b85557d030d3056bf2eb0e9273f2811162d",
}
INK, NAVY, TEAL, MUTED = "172B3A", "17324D", "138A86", "617583"
PAPER, LIGHT, WHITE = "F7F9FB", "DDF2F0", "FFFFFF"
FONT = "Microsoft YaHei"
W, H = 13.333333, 7.5
T56 = "教材印刷第 56 页（PDF 第 61 页）"
T57 = "教材印刷第 57 页（PDF 第 62 页）"
T58 = "教材印刷第 58 页（PDF 第 63 页）"
W34 = "复习讲义 Word 第 3—4 页（本机导出页码）"
W45 = "复习讲义 Word 第 4—5 页（本机导出页码）"


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class Lesson:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.prs.core_properties.title = "电解质与电离方程式"
        self.prs.core_properties.subject = "依据复习讲义与沪科技教材的四十分钟复习课"
        self.records = []
        self.minute = 0
        self.boxes = []

    def text(self, slide, text, x, y, w, h, size=25, color=INK, bold=False):
        if min(x, y, w, h) < 0 or x + w > W + 0.01 or y + h > H + 0.01:
            raise ValueError("Text box outside slide")
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for index, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
            p.space_after = Pt(8)
            p.line_spacing = 1.16
            for token in re.split(r"([_^]\{[^}]+\})", line):
                if not token:
                    continue
                run = p.add_run()
                scripted = bool(re.fullmatch(r"[_^]\{[^}]+\}", token))
                run.text = token[2:-1] if scripted else token
                run.font.name = FONT
                run.font.size = Pt(size * 0.72 if scripted else size)
                run.font.bold = bold
                run.font.color.rgb = RGBColor.from_string(color)
                prop = run._r.get_or_add_rPr()
                ea = OxmlElement("a:ea")
                ea.set("typeface", FONT)
                prop.append(ea)
                if scripted:
                    prop.set("baseline", "30000" if token[0] == "^" else "-20000")
        self.boxes.append(
            {
                "slide": len(self.records),
                "text": text,
                "font_pt": size,
                "bbox_inches": [x, y, w, h],
            }
        )
        return box

    def rule(self, slide, x, y, w, color=TEAL, height=0.035):
        shape = slide.shapes.add_shape(
            1, Inches(x), Inches(y), Inches(w), Inches(height)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
        shape._element.spPr.append(OxmlElement("a:effectLst"))

    def page(self, title, stage, minutes, source, notes, *, note=None, pair=None):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(PAPER)
        number = len(self.records) + 1
        interval = f"{self.minute}—{self.minute + minutes} 分钟"
        self.minute += minutes
        self.records.append(
            {
                "page": number,
                "title": title,
                "stage": stage,
                "minutes": minutes,
                "interval": interval,
                "source": source,
                "teacher_notes": notes,
                "notebook_section": note,
                "response_pair": pair,
            }
        )
        self.text(slide, stage, 0.65, 0.32, 9.9, 0.4, 15, TEAL, True)
        self.text(slide, f"{number:02d}", 11.65, 0.3, 1.0, 0.4, 17, MUTED)
        self.text(slide, title, 0.65, 0.9, 12.0, 0.68, 32, NAVY, True)
        self.rule(slide, 0.65, 1.72, 0.6)
        self.text(slide, "依据：" + source, 0.65, 7.02, 11.95, 0.32, 11, MUTED)
        slide.notes_slide.notes_text_frame.text = (
            f"第 {number} 页 | 建议时间 {interval} | 本页 {minutes} 分钟\n"
            f"资料依据：{source}\n学生笔记：{note or '无需整页抄写'}\n\n{notes}\n\n"
            "使用范围：教师复核候选。课堂安排与反馈为本课设计，不是原始试卷官方答案；"
            "尚无真实班级试教数据。"
        )
        return slide

    def lines(self, slide, lines, *, y=2.15, gap=1.05, size=26, x=0.8, w=11.7):
        for i, line in enumerate(lines):
            self.text(slide, line, x, y + i * gap, w, 0.85, size)

    def cue(self, slide, text):
        self.text(slide, text, 0.8, 6.28, 11.7, 0.52, 21, TEAL, True)

    def table(self, slide, headers, rows, *, widths, y=2.08, row_h=0.88, size=24):
        x = 0.8
        for j, heading in enumerate(headers):
            self.text(
                slide,
                heading,
                x + 0.12,
                y + 0.12,
                widths[j] - 0.25,
                0.62,
                size,
                NAVY,
                True,
            )
            x += widths[j]
        self.rule(slide, 0.8, y + 0.77, sum(widths), MUTED, 0.012)
        for i, row in enumerate(rows):
            x = 0.8
            top = y + 0.85 + i * row_h
            for j, value in enumerate(row):
                self.text(
                    slide,
                    value,
                    x + 0.12,
                    top + 0.06,
                    widths[j] - 0.25,
                    row_h - 0.12,
                    size,
                    INK,
                    j == 0,
                )
                x += widths[j]
            self.rule(slide, 0.8, top + row_h - 0.02, sum(widths), "CCD7DE", 0.008)


def build_deck(figure):
    d = Lesson()
    s = d.page(
        "电解质与电离方程式",
        "本课任务",
        2,
        W34 + "；教材第 56—58 页",
        "按高二复习、40分钟试排。对应讲义考点一，不在本课挤入离子共存与检验。"
        "先让学生在笔记页写姓名，浏览四段标题。三个任务用于下课回看，不预先讲出答案。",
    )
    d.lines(
        s,
        [
            "判类别：物质属于哪一类，依据是什么？",
            "说原因：同一种物质，为什么导电情况会不同？",
            "作表达：怎样用电离方程式表示微观过程？",
        ],
        gap=1.12,
        size=28,
    )
    d.cue(s, "准备笔记页：概念边界 / 微观解释 / 强弱比较 / 符号表达")

    s = d.page(
        "氯化钠不导电时还是电解质吗",
        "先思考 · 暂不抄写",
        1,
        T56 + "；" + W34,
        "沿用教材思考中的固体与溶液比较，提问而不虚构本班实验。让学生先选观点并说一句依据。"
        "记录两种不同理由，暂不判定；第7页用粒子状态回扣。",
    )
    d.table(
        s,
        ["比较对象", "教材给出的导电情况"],
        [["氯化钠固体", "不能导电"], ["氯化钠水溶液", "能够导电"]],
        widths=[4.0, 7.7],
        row_h=1.08,
        size=27,
    )
    d.text(
        s, "物质类别是否随着这两种状态一起改变？", 0.9, 5.55, 11.5, 0.65, 29, NAVY, True
    )

    s = d.page(
        "判断电解质要保留哪些条件",
        "一 · 概念边界",
        2,
        T56 + "；" + W34,
        "先让学生找定义中的研究对象、状态、导电条件。解释‘或’是至少一种情况满足。"
        "定义中的导电由化合物自身电离说明。暂不展开CO2等与水反应例外。",
    )
    d.text(
        s,
        "电解质是在水溶液中或熔融状态下\n能够导电的化合物。",
        0.85,
        2.15,
        11.65,
        1.4,
        31,
        NAVY,
        True,
    )
    d.lines(
        s,
        [
            "共同对象：化合物，单质和混合物不在此分类范围。",
            "非电解质：在上述两种条件下均不能导电的化合物。",
            "注意“或”与“和”；依据化合物是否能自身电离。",
        ],
        y=3.9,
        gap=0.71,
        size=25,
    )
    d.cue(s, "先圈关键词，稍后在笔记一中合成完整表述。")

    s = d.page(
        "先判研究对象再判断类别",
        "独立判断 · 说出依据",
        2,
        W34 + "；" + T56,
        "这是依据讲义概念改编的课堂判断，不标原题题号。学生先独立写再同桌说理由。"
        "接受‘不在该分类范围’表述。下一页才反馈，避免答案提前暴露。",
        pair="boundary-question",
    )
    d.table(
        s,
        ["对象", "你的类别判断", "判定依据"],
        [
            ["铜", "？", "？"],
            ["氯化钠", "？", "？"],
            ["氯化钠水溶液", "？", "？"],
            ["乙醇", "？", "？"],
        ],
        widths=[3.6, 3.3, 4.8],
        row_h=0.73,
        size=25,
    )
    d.cue(s, "类别可选：电解质 / 非电解质 / 不属于二者的分类范围")

    s = d.page(
        "整理笔记一 概念的边界",
        "核对后记录 · 留 60 秒",
        2,
        W34 + "；" + T56,
        "先快速核对：铜为单质，氯化钠溶液为混合物；二者不属于电解质或非电解质。"
        "氯化钠是电解质，乙醇是非电解质。用60秒补记定义和一个边界例，不能只给分类名单。",
        note="一 概念边界",
        pair="boundary-answer",
    )
    d.table(
        s,
        ["对象", "判断与理由"],
        [
            ["铜 / 氯化钠水溶液", "单质 / 混合物，均不在此分类范围"],
            ["氯化钠", "电解质；水溶液或熔融状态可导电"],
            ["乙醇", "非电解质；两种条件下均不能自身电离"],
        ],
        widths=[4.05, 7.65],
        row_h=0.88,
        size=24,
    )
    d.cue(s, "笔记保留：化合物 + 状态条件 + 自身电离；边界例任选一个。")

    s = d.page(
        "看图解释氯化钠的电离",
        "二 · 先观察教材图 2.14",
        3,
        T57 + " 图 2.14",
        "先用45秒观察全图；学生在图左标晶体，在两条路径旁标条件，口头区分粒子存在与自由移动。"
        "完整保留水分子、两条路径、离子图例和原图题，不把水合离子误画成裸离子。"
        "图示不是电离实测录像；本课也不安排熔盐演示。",
    )
    with Image.open(figure) as im:
        ratio = im.height / im.width
    s.shapes.add_picture(
        str(figure),
        Inches(0.7),
        Inches(1.93),
        width=Inches(8.05),
        height=Inches(8.05 * ratio),
    )
    d.text(s, "观察两件事", 9.35, 2.15, 3.2, 0.6, 25, NAVY, True)
    d.text(
        s,
        "1 晶体中的离子\n能自由移动吗？\n\n2 溶于水和熔融\n两条路径的\n共同结果是什么？",
        9.35,
        3.02,
        3.2,
        3.25,
        24,
    )

    s = d.page(
        "电离使离子能够自由移动",
        "用微观解释回到开场问题",
        2,
        T57 + "；" + W34,
        "请学生先根据图完成三行解释，再用表核对。晶体中已有离子，不应说从没有离子变成有离子。"
        "回扣第2页：氯化钠固体是电解质，只是该状态下离子不能自由移动。"
        "导电涉及外加电场下定向移动，电离本身无需通电。",
    )
    d.table(
        s,
        ["状态", "离子状态", "导电解释"],
        [
            ["氯化钠晶体", "离子不能自由移动", "该状态不能导电"],
            ["溶于水", "形成自由移动的水合离子", "外接电源、闭合通路时导电"],
            ["熔融", "离子能够自由移动", "外接电源、闭合通路时导电"],
        ],
        widths=[2.65, 4.55, 4.5],
        row_h=1.0,
        size=24,
    )
    d.cue(s, "开场问题：氯化钠固体不能导电，仍然是电解质。")

    s = d.page(
        "整理笔记二 电离与导电",
        "讲解后记录 · 留 60 秒",
        2,
        T57,
        "先用30秒请学生说出电离定义，然后留60秒填笔记二，最后抽问电离是否需要通电。"
        "板书只留形成自由移动离子和外电场下定向移动两层，不要求抄讲稿。",
        note="二 微观解释",
    )
    d.text(s, "电离", 0.9, 2.2, 2.0, 0.6, 29, TEAL, True)
    d.text(
        s,
        "电解质在水溶液中或熔融状态下，\n形成可以自由移动离子的过程。",
        3.1,
        2.2,
        9.3,
        1.3,
        28,
    )
    d.text(s, "导电", 0.9, 4.0, 2.0, 0.6, 29, TEAL, True)
    d.text(
        s,
        "自由移动的离子在外加电场下定向移动。\n电离本身不以通电为前提。",
        3.1,
        4.0,
        9.3,
        1.3,
        28,
    )
    d.cue(s, "电离：形成自由移动的离子。导电：离子在外加电场下定向移动。")

    s = d.page(
        "强弱电解质按什么区分",
        "三 · 从电离程度比较",
        3,
        T58 + "；" + W45,
        "限定为水溶液中的电离程度。先解释全部与部分，再引出强弱名称。"
        "举教材NaCl与醋酸，不引入导电亮度比较或未经控制的浓度条件。"
        "弱电解质溶液有未电离分子；离子与分子共存，水不计入这两类对比对象。",
    )
    d.text(
        s,
        "共同前提：比较电解质在水溶液中的电离程度。",
        0.85,
        2.12,
        11.7,
        0.8,
        27,
        NAVY,
        True,
    )
    d.text(s, "强电解质", 0.95, 3.3, 5.4, 0.6, 29, TEAL, True)
    d.text(s, "全部电离\n如氯化钠、氯化氢、氢氧化钠", 0.95, 4.2, 5.25, 1.45, 25)
    d.text(s, "弱电解质", 7.05, 3.3, 5.4, 0.6, 29, TEAL, True)
    d.text(s, "部分电离\n如醋酸、一水合氨、水", 7.05, 4.2, 5.25, 1.45, 25)
    d.cue(s, "强、弱说的是电离程度，不能直接换成“灯泡亮、暗”。")

    s = d.page(
        "整理笔记三 用同一标准比较",
        "归纳后记录 · 留 60 秒",
        2,
        T58 + "；" + W45,
        "学生用60秒补齐笔记页比较表。强调先读比较维度，再读强弱两列。"
        "表中微粒只讨论电解质电离情况，不展开水的电离及其他次级过程。",
        note="三 强弱比较",
    )
    d.table(
        s,
        ["比较维度", "强电解质", "弱电解质"],
        [
            ["水溶液中电离程度", "全部电离", "部分电离"],
            ["未电离的该电解质分子", "无", "有"],
            ["电离方程式的符号", "等号  =", "可逆符号  ⇌"],
            ["本课典型例子", "氯化钠", "醋酸"],
        ],
        widths=[4.6, 3.55, 3.55],
        row_h=0.73,
        size=24,
    )
    d.cue(s, "另记一句：难溶不等于弱电解质，不能把溶解性当成电离程度。")

    s = d.page(
        "这两句话的依据够吗",
        "短检查 · 先判断再说明",
        1,
        W45 + "；" + T58,
        "依据讲义易混点改编的课堂检查。先不显示答案。学生用‘判据’一词纠正两个说法。"
        "不要求讨论具体导电率公式。",
        pair="strength-question",
    )
    d.lines(
        s,
        [
            "甲：某物质难溶于水，所以它是弱电解质。",
            "乙：某溶液导电较强，所以溶质一定是强电解质。",
        ],
        gap=1.4,
        size=28,
    )
    d.cue(s, "在笔记三下方写一句纠正；重点是补上正确判据。")

    s = d.page(
        "回到电离程度判断",
        "核对 · 修正自己的表述",
        1,
        W45 + "；" + T58,
        "甲乙的推断依据均不足。用讲义BaSO4例说明难溶与强弱无必然对应，但不展开溶解平衡。"
        "乙只强调灯泡亮暗受溶液条件影响，不能单凭现象判断电离程度。",
        pair="strength-answer",
    )
    d.lines(
        s,
        [
            "甲：溶解性不能代替电离程度。\n讲义中的硫酸钡难溶，但属于强电解质。",
            "乙：不能仅凭溶液导电强弱作出该判断。\n仍须依据溶质在水溶液中的电离程度。",
        ],
        gap=1.68,
        size=27,
    )

    s = d.page(
        "把微观过程写成电离方程式",
        "四 · 先看一个完整示例",
        2,
        T57 + "；" + W45,
        "先读物质和离子名称。解释离子符号简写水合离子，不能把图中的水合过程理解为改变元素。"
        "示例是教材已有式；依次讲电解质、电离程度、离子、守恒。",
    )
    d.text(s, "氯化钠在水溶液中的电离", 0.9, 2.15, 11.5, 0.65, 27, NAVY, True)
    d.text(s, "NaCl = Na^{+} + Cl^{-}", 1.1, 3.25, 11.0, 1.0, 42, NAVY, True)
    d.lines(
        s,
        [
            "氯化钠 → 钠离子、氯离子（名称对应）",
            "原子种类与数目不变；右侧电荷总和为 0。",
        ],
        y=4.65,
        gap=0.8,
        size=26,
    )
    d.cue(s, "按教材习惯，用简单离子符号表示溶液中的水合离子。")

    s = d.page(
        "弱电解质要表达部分电离",
        "从判据选择表示符号",
        1,
        T58,
        "醋酸为教材已有示例，说明一个酸性氢电离，不能按化学式把所有氢拆出。"
        "本课只讨论醋酸这一基础表达；多元弱酸、酸式盐和特殊物质放到后续课时。",
    )
    d.text(s, "醋酸在水溶液中的电离", 0.9, 2.15, 11.5, 0.7, 27, NAVY, True)
    d.text(
        s, "CH_{3}COOH ⇌ H^{+} + CH_{3}COO^{-}", 0.95, 3.4, 11.6, 1.05, 38, NAVY, True
    )
    d.lines(
        s,
        [
            "产物名称：氢离子和醋酸根离子。",
            "选择可逆符号：因为醋酸在水溶液中部分电离。",
        ],
        y=4.8,
        gap=0.76,
        size=26,
    )

    s = d.page(
        "独立完成教材书写任务",
        "先作答 · 暂不展示答案",
        2,
        T57 + " 书写表达",
        "直接选用已阅读教材中的三个书写对象，不虚构试卷年份题号。明确在水溶液中。"
        "学生在笔记页背面练习区作答；巡视先看OH与SO4是否整体保留，再看系数与电荷。",
        pair="equation-question",
    )
    d.text(
        s,
        "写出下列电解质在水溶液中的电离方程式。",
        0.85,
        2.1,
        11.7,
        0.8,
        27,
        NAVY,
        True,
    )
    d.lines(
        s,
        ["氢氧化钡  Ba(OH)_{2}", "硫酸钠  Na_{2}SO_{4}", "氯化钡  BaCl_{2}"],
        y=3.15,
        gap=0.92,
        size=31,
    )
    d.cue(s, "完成后自行检查：粒子种类、系数、离子电荷、表示符号。")

    s = d.page(
        "核对时说出每个系数的依据",
        "教材任务反馈",
        2,
        T57 + " 任务；本课推导反馈",
        "以下为本课依概念推导的反馈，不称官方评分细则。逐式检查原子与电荷守恒。"
        "2OH−中的2是系数，OH−为整体；Na2SO4右侧硫酸根电荷是2−。"
        "让学生用不同颜色改一处，并写出是系数、电荷还是整体离子错误。",
        pair="equation-answer",
    )
    d.lines(
        s,
        [
            "氢氧化钡  Ba(OH)_{2} = Ba^{2+} + 2OH^{-}",
            "硫酸钠  Na_{2}SO_{4} = 2Na^{+} + SO_{4}^{2-}",
            "氯化钡  BaCl_{2} = Ba^{2+} + 2Cl^{-}",
        ],
        y=2.25,
        gap=1.12,
        size=35,
    )
    d.cue(s, "硫酸根、氢氧根作为整体书写；系数与离子电荷的位置不同。")

    s = d.page(
        "整理笔记四 书写和检查",
        "记录步骤 · 留 60 秒",
        2,
        T57 + "；" + T58 + "；" + W45,
        "学生保留一个完整式、离子名称和四个检查点，不抄全部讲稿。"
        "弱电解质式在旁注明可逆符号的依据。核对三个任务后，留60秒集中记录。",
        note="四 符号表达",
    )
    d.lines(
        s,
        [
            "先判断：物质在水溶液中的电离程度。",
            "再表示：写出离子，并选择等号或可逆符号。",
            "后检查：原子数守恒，电荷总和相等。",
        ],
        gap=1.02,
        size=28,
    )
    d.cue(s, "笔记中保留一个完整示例，并圈出最容易错的系数或电荷。")

    s = d.page(
        "合上讲义重建自己的知识框架",
        "独立回忆 · 先不翻笔记",
        3,
        "本课依据讲义与教材设计的整理任务",
        "前90秒保持本页只显示问题，不翻到答案页；让学生在背面独立重建区作答。"
        "后90秒可同桌互讲。基础班可用标题提示，但不要一开始展示填好的框架。"
        "巡视记录缺失的是条件、判据还是符号，下一页只补真正的缺口。",
        pair="recall-question",
    )
    d.lines(
        s,
        [
            "判类别时，必须先看什么、再看什么？",
            "怎样用离子状态解释氯化钠的导电情况？",
            "电离程度怎样影响电离方程式的表示？",
        ],
        gap=1.08,
        size=28,
    )
    d.cue(s, "用关键词、箭头和一个完整方程式即可；不要求复述讲稿。")

    s = d.page(
        "核对框架只补缺失的关系",
        "回看笔记 · 整理 2 分钟",
        2,
        W34 + "；" + W45 + "；教材第 56—58 页",
        "此页是回忆后的简洁核对，不代替前三分钟独立重建。学生只补缺口并圈一处易混。"
        "讲义考点二离子反应留待下一课，由电离出的离子继续衔接，今天不另开新主题。",
        note="整合与自查",
        pair="recall-answer",
    )
    d.table(
        s,
        ["笔记部分", "核对的核心关系"],
        [
            ["一 概念边界", "化合物范围 + 状态条件 + 自身电离"],
            ["二 微观解释", "形成自由移动离子 → 解释该状态的导电"],
            ["三 强弱比较", "水溶液中全部 / 部分电离 → 强 / 弱"],
            ["四 符号表达", "离子种类 + 系数 + 电荷 + 等号 / 可逆符号"],
        ],
        widths=[3.2, 8.5],
        row_h=0.76,
        size=25,
    )
    d.cue(s, "圈出仍混淆的一处：对象 / 状态 / 电离程度 / 书写。")

    s = d.page(
        "离堂前完成两个小任务",
        "独立反馈 · 2 分钟",
        2,
        "依据本课教材与讲义概念改编",
        "与开场三个任务对应：第一问检查分类和微观原因，第二问检查符号。"
        "收取或拍下匿名错误类型即可；两分钟抽样不能证明全班达成。下一页用于即时反馈。",
        pair="exit-question",
    )
    d.text(
        s,
        "1 氯化钠固体不能导电。\n   它是不是电解质？请用粒子状态说明理由。",
        0.9,
        2.2,
        11.5,
        1.5,
        28,
    )
    d.text(
        s,
        "2 写出氢氧化钠在水溶液中的电离方程式，\n   并说明表示符号的选择依据。",
        0.9,
        4.2,
        11.5,
        1.5,
        28,
    )

    s = d.page(
        "带着一个明确的问题复习",
        "离堂核对与课后任务",
        1,
        T57 + "；复习讲义考点一",
        "简要反馈两问：NaCl水溶液或熔融状态能导电，所以是电解质；固体离子不能自由移动，故固体不导电。NaOH完全电离，用等号。"
        "课后只补本课笔记中的一处缺口，按讲义知识点1—3回读，基础例题待教师选定。"
        "未核验的讲义后续题号不写入作业。若多数学生错电荷，下一课先用2分钟纠错。",
        pair="exit-answer",
    )
    d.lines(
        s,
        [
            "氯化钠是电解质：水溶液或熔融状态能够导电。",
            "固体不导电：晶体中的离子不能自由移动。",
            "氢氧化钠  NaOH = Na^{+} + OH^{-}；全部电离，用等号。",
        ],
        y=2.1,
        gap=0.82,
        size=26,
    )
    d.text(
        s,
        "课后：回读讲义考点一知识点 1—3，\n补好自己圈出的一个缺口，再独立写一遍今天改错的式子。",
        0.9,
        4.95,
        11.5,
        1.1,
        25,
        NAVY,
        True,
    )
    assert d.minute == 40, d.minute
    assert len(d.records) == 21
    return d


def build_notebook(path):
    """Two actual writable pages, not a printed full-answer worksheet."""
    pdfmetrics.registerFont(
        TTFont("Notebook", "C:/Windows/Fonts/msyh.ttc", subfontIndex=0)
    )
    c = Canvas(str(path), pagesize=A4)
    c.setTitle("电解质与电离方程式 学生课堂笔记")
    _, height = A4

    def text(x, y, content, size=11):
        c.setFillColorRGB(0.09, 0.13, 0.17)
        c.setFont("Notebook", size)
        c.drawString(x, height - y, content)

    def line(y, x=43, end=552):
        c.setStrokeColorRGB(0.72, 0.76, 0.78)
        c.setLineWidth(0.45)
        c.line(x, height - y, end, height - y)

    def lines(start, count, gap=22):
        for i in range(count):
            line(start + gap * i)

    text(43, 46, "电解质与电离方程式", 19)
    text(43, 72, "课堂笔记 1 / 2     姓名____________  日期____________", 10)
    text(43, 98, "听讲时先看图、说理由；到整理页再补记关键词，不抄整页讲稿。", 10)
    text(43, 135, "一 概念边界    对应 PPT 第 3—5 页", 13)
    text(43, 160, "用完整的一句话写出电解质定义，圈出对象和条件。")
    lines(185, 2)
    text(43, 235, "一个不在电解质或非电解质分类范围内的例子及理由")
    lines(260, 2)
    text(43, 320, "二 微观解释    对应 PPT 第 6—8 页", 13)
    text(43, 345, "解释氯化钠晶体、溶于水、熔融三种情况下离子的状态。")
    lines(370, 3)
    text(43, 442, "电离与导电的区别    比较自由移动离子的形成与定向移动")
    lines(468, 2)
    text(43, 535, "三 强弱比较    对应 PPT 第 9—12 页", 13)
    # Real column labels match row objects and response dimensions.
    xs = [43, 257, 401, 552]
    top, row_height = 555, 39
    for row in range(5):
        yy = top + row * row_height
        line(yy)
    for x in xs:
        c.line(x, height - top, x, height - (top + 4 * row_height))
    text(52, 580, "比较维度", 10)
    text(267, 580, "强电解质", 10)
    text(411, 580, "弱电解质", 10)
    for yy, label in [
        (619, "水溶液中的电离程度"),
        (658, "电离方程式的表示符号"),
        (697, "一个典型例子"),
    ]:
        text(52, yy, label, 10)
    text(43, 737, "一句易混点提醒")
    line(762)
    text(43, 812, "内容依据 复习讲义考点一；沪科技教材必修第一册印刷第56—58页", 8)
    c.showPage()
    text(43, 46, "电解质与电离方程式", 19)
    text(43, 72, "课堂笔记 2 / 2     先独立作答，再用不同颜色订正", 10)
    text(43, 112, "四 符号表达    对应 PPT 第 13—17 页", 13)
    text(43, 139, "先写自己的书写顺序和检查点，再保留一个完整示例。")
    lines(165, 2)
    text(43, 222, "教材任务    写出水溶液中的电离方程式", 11)
    text(43, 252, "氢氧化钡")
    line(267, 112)
    text(43, 296, "硫酸钠")
    line(311, 112)
    text(43, 340, "氯化钡")
    line(355, 112)
    text(43, 389, "我的一处订正及原因")
    line(414)
    text(43, 453, "独立重建    对应 PPT 第 18—19 页", 13)
    text(43, 478, "先不翻笔记，用关键词、箭头和一个方程式连接今天的内容。", 10)
    lines(509, 4, 25)
    text(43, 620, "离堂反馈    对应 PPT 第 20 页", 13)
    text(43, 647, "1 氯化钠固体是不是电解质？用粒子状态说明理由。", 10)
    lines(671, 2)
    text(43, 727, "2 写出氢氧化钠的电离方程式，并说明符号选择依据。", 10)
    lines(753, 2)
    text(43, 812, "完整任务见课件；课堂笔记与反馈为本课设计，不是教材原版练习纸。", 8)
    c.save()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--pdftoppm", type=Path, required=True)
    args = parser.parse_args()
    before = {"book": sha(BOOK), "word": sha(WORD)}
    if before != EXPECTED:
        raise SystemExit(
            "Source hash mismatch; reread the changed source before building."
        )
    output = args.output_dir.resolve()
    if output.exists():
        raise SystemExit(
            "Use a new output directory; existing artifacts are never overwritten."
        )
    output.mkdir(parents=True)
    figure = output / "textbook-figure-2-14.png"
    # This is extraction of an inspected source region, not an altered/redrawn figure.
    subprocess.run(
        [
            str(args.pdftoppm),
            "-f",
            "62",
            "-l",
            "62",
            "-r",
            "240",
            "-x",
            "320",
            "-y",
            "880",
            "-W",
            "1400",
            "-H",
            "864",
            "-singlefile",
            "-png",
            str(BOOK),
            str(figure.with_suffix("")),
        ],
        check=True,
    )
    lesson = build_deck(figure)
    pptx = output / "电解质与电离方程式_课堂重构版.pptx"
    notebook = output / "电解质与电离方程式_学生笔记.pdf"
    lesson.prs.save(pptx)
    build_notebook(notebook)
    after = {"book": sha(BOOK), "word": sha(WORD)}
    if before != after:
        raise RuntimeError("Source unexpectedly changed")
    manifest = {
        "schema_version": "shchem.source-led-lesson-exemplar.v1",
        "candidate_only": True,
        "teacher_review_required": True,
        "publication_allowed": False,
        "generic_desktop_generator_integrated": False,
        "provider_calls": 0,
        "source_hashes": before,
        "sources_unchanged": True,
        "source_scope": {
            "word_export_pages": [3, 4, 5],
            "textbook_pdf_pages": [61, 62, 63],
            "textbook_printed_pages": [56, 57, 58],
            "concepts": ["TB-M1-C2-S22-C05", "TB-M1-C2-S22-C06", "TB-M1-C2-S22-C07"],
        },
        "figure": {
            "path": figure.name,
            "sha256": sha(figure),
            "pdf_page": 62,
            "printed_page": 57,
            "figure_id": "2.14",
            "dpi": 240,
            "pixel_clip": [320, 880, 1400, 864],
            "redrawn": False,
        },
        "slides": lesson.records,
        "minutes": lesson.minute,
        "layout_boxes": lesson.boxes,
        "artifacts": {p.name: sha(p) for p in [pptx, notebook]},
        "visual_review": "pending",
        "classroom_trial": "not_performed",
    }
    (output / "lesson-manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    guide = [
        "# 电解质与电离方程式 教师讲解说明",
        "",
        "本课按高二复习40分钟试排，以复习讲义考点一和教材印刷56—58页为依据。",
        "PPT正文面向学生，逐页备注包含讲解、提问、记录时间和反馈建议。尚需任课教师审核与班级试教。",
        "学生笔记PDF为A4双面两页；第18页先独立回忆，第19页再投影核对。",
        "图2.14取自用户提供教材母本，原图仅用于本地备课；对外传播前需另核使用权。",
        "",
    ]
    for record in lesson.records:
        guide += [
            f"## 第{record['page']}页 {record['title']}",
            "",
            f"时间 {record['interval']}。依据 {record['source']}。",
            "",
            record["teacher_notes"],
            "",
        ]
    (output / "教师讲解说明.md").write_text("\n".join(guide), encoding="utf-8")
    print(
        json.dumps(
            {
                "output": str(output),
                "slides": len(lesson.records),
                "minutes": lesson.minute,
                "source_unchanged": True,
                "visual_review": "pending",
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
