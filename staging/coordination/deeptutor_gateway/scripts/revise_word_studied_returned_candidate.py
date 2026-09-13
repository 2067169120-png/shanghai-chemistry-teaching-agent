"""Explicit offline editorial revision of one retained response; no provider access.

Retain the raw response and mothers. Emit a new assistant-reviewed candidate,
not a teacher-approved artifact or a silently recovered native task.
"""

from __future__ import annotations

import copy
import hashlib
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(ROOT))

from verify_real_blueprint_preparation import BundledArtifactRenderer

from integrations.deeptutor_shchem_v1.desktop_preparation import (
    normalize_preparation_candidate,
    normalize_preparation_payload,
)

LIVE = ROOT / "runtime/deeptutor_shchem/qa/word-studied-live-generation-20260909"
BRIEF = ROOT / "outputs/备课/2026-09-09-电解质的电离-讲义精读生成材料"
OUT = ROOT / "outputs/备课/2026-09-09-电解质的电离-讲义精读讲练版-r9"
RAW_HASH = "99675c2e109b299b81523525604e903e72e3a30681610865fd8af2c77ac53c76"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def table(columns, rows, dimension="比较维度"):
    return {
        "kind": "comparison",
        "comparison": {
            "dimension_label": dimension,
            "columns": columns,
            "rows": [{"label": label, "values": values} for label, values in rows],
        },
        "steps": [],
    }


def worksheet_table(heading, prompt, rows, columns):
    return {
        "heading": heading,
        "prompt": prompt,
        "response_kind": "table",
        "response_lines": 0,
        "columns": columns,
        "row_labels": rows,
    }


def revise(raw):
    c = copy.deepcopy(raw)
    c["title"] = "电解质的电离｜两课时讲练"
    old = {i: s for i, s in enumerate(c["slides"], 1)}
    old[5]["content"] = [
        "教材摘录：这类能在水溶液中或熔融状态下导电的化合物叫做电解质。",
        "教材摘录：另一类化合物，如蔗糖、酒精等，在水溶液中和熔融状态下均不能导电，这类化合物叫做非电解质。",
        "必记：对象是化合物；电解质的条件用“或”，非电解质的条件用“和”。",
    ]
    old[11]["content"] = [
        "教材摘录：电解质在水溶液中或熔融状态下，形成可以自由移动离子的过程称为电离。",
        "必记：电离不需要通电；电解质导电还需外接电源与闭合通路，离子定向移动。",
    ]
    old[16]["content"] = [
        "教材摘录：像氯化钠、氯化氢、氢氧化钠等在水溶液中能够全部电离为自由移动离子的电解质称为强电解质。",
        "教材摘录：像醋酸、一水合氨（NH₃·H₂O）等在水溶液中仅有部分分子能电离出自由移动离子的电解质称为弱电解质。",
        "必记：根本判据是在水溶液中的电离程度。",
    ]
    old[8]["content"][1] = "铜既不是电解质，也不是非电解质——铜是单质，不归入这两类。"
    old[8]["content"][2] = (
        "CO₂是非电解质——其水溶液导电与生成的H₂CO₃电离有关，不是CO₂自身电离。"
    )
    old[8]["teacher_notes"] = (
        "核对三个边界反例，强调判断的研究对象为化合物，不把导电溶液这一混合物称为电解质。"
    )
    old[12]["teacher_notes"] = (
        "观察晶体中的离子及其位置受约束的状态，再比较水中离子的移动能力。结合水分子说明水合离子；下一页整理三状态，不将溶液画成只有裸离子。"
    )
    # A five-row comparison is split explicitly, never truncated or widened globally.
    old[17]["title"] = "K3 必记表：强弱判据与表达"
    full_rows = copy.deepcopy(old[17]["visual"]["comparison"]["rows"])
    old[17]["visual"]["comparison"]["rows"] = full_rows[:3]
    old[17]["content"] = ["强弱只按电离程度区分；不能用当前导电性或溶解度替代。"]
    old[17]["teacher_notes"] = (
        "先填学习单3前3行，再投影核对。导电能力还受浓度等条件影响；不将灯泡亮暗直接等同于电解质强弱。"
    )
    extra17 = copy.deepcopy(old[17])
    extra17["title"] = "K3 必记表：类型与溶质存在形式"
    extra17["visual"]["comparison"]["rows"] = full_rows[3:]
    extra17["content"] = [
        "这里仅比较目标溶质的电离，不是列出整个水溶液的全部微粒。",
        "难溶不等于弱电解质；导电能力不能在未控制浓度等条件时替代电离程度。",
    ]
    extra17["teacher_notes"] = (
        "完成学习单3后两行，用NaCl和CH₃COOH举例。此时不逐一列出后面Q1的干扰项答案；Q1讲评后补上BaSO₄、Q2讲评后补上NH₄Cl。"
    )
    for n in (9, 14, 18, 20, 31, 33):
        old[n]["teacher_notes"] = (
            "先让学生独立作答，按本页预算留出纸笔时间；巡视记录错误，不口头提示选项的正误或直接给出理由。下一页再核对，学生订正并补写依据。"
        )
    old[21]["content"] = [
        "讲义参考答案：A（强电解质）。",
        "依据：氯化铵 NH₄Cl＝NH₄⁺＋Cl⁻（水溶液）。",
        "盐与弱碱有关，不等于盐是弱电解质；仍按盐自身的电离程度判断。",
    ]
    old[21]["teacher_notes"] += (
        " 若已学水解，可补充：水解是电离后离子的反应，不等于NH₄Cl只部分电离。未学班级不展开。"
    )
    old[24]["visual"]["steps"] = [
        {"label": "①判类别＋②定条件", "detail": "先判强弱，再看水溶液或熔融等条件。"},
        {"label": "③确定离子", "detail": "辨清离子种类；NH₄⁺、SO₄²⁻、ClO⁻等保持整体。"},
        {"label": "④选用符号", "detail": "本课强电解质用＝，弱电解质用⇌。"},
        {"label": "⑤核对", "detail": "原子数、电荷、系数，以及是否需要分步。"},
    ]
    full28 = copy.deepcopy(old[28]["visual"]["comparison"]["rows"])
    old[28]["title"] = "K5 必记例式：醋酸与一水合氨"
    old[28]["visual"]["comparison"]["rows"] = full28[:2]
    old[28]["content"] = ["水溶液中的弱电解质仅部分电离，用可逆符号。"]
    extra28 = copy.deepcopy(old[28])
    extra28["title"] = "K5 必记例式：多元弱酸与水"
    extra28["visual"]["comparison"]["rows"] = full28[2:]
    extra28["content"] = ["多元弱酸分步电离，以第一步为主；不能写成完全电离的总式。"]
    old[29]["content"][1] = (
        "微观表达：2H₂O ⇌ H₃O⁺＋OH⁻；本课简写：H₂O ⇌ H⁺＋OH⁻。H⁺是水溶液中氢离子的简写，不表示裸质子独立存在。"
    )
    old[30]["content"] = [
        "必记：NaHSO₄按状态区分；NaHCO₃的盐解离与HCO₃⁻部分电离分开写。",
        "不能仅由下面的电离式判断NaHCO₃溶液酸碱性，还涉及其他平衡。",
    ]
    old[32]["content"] = [
        "A：电离不需通电。硝酸钾 KNO₃＝K⁺＋NO₃⁻。",
        "B：弱酸分步。硫化氢 H₂S ⇌ H⁺＋HS⁻；HS⁻ ⇌ H⁺＋S²⁻。",
        "C：原式正确。一水合氨 NH₃·H₂O ⇌ NH₄⁺＋OH⁻。",
        "D：ClO⁻保持整体。次氯酸钠 NaClO＝Na⁺＋ClO⁻。",
    ]
    old[34]["content"][1] = (
        "BaSO₄（硫酸钡）是强电解质；中学表述中，溶入水的部分全部电离。难溶只说明溶解量小，不说明已溶部分仅部分电离。"
    )
    old[34]["teacher_notes"] = (
        "区分溶解量与电离程度；溶解平衡式不是弱电解质证据。按本课范围解释，不延伸到熔盐操作或没有来源的实测现象。"
    )
    old[35]["image"]["observation_prompt"] = (
        "观察a、b、c中的粒子种类、排列和水分子标示，独立判断四个选项，并写出图中证据。"
    )
    old[35]["teacher_notes"] = (
        "先读图再独立作答，不预先指定各图状态；本课不要求电极反应。学生在学习单所附同一题图旁写依据，下一页再讲评。"
    )
    summary = []
    summaries = [
        (
            "常考知识总结① 分类与导电",
            table(
                ["必须记住的结论", "易错边界"],
                [
                    (
                        "电解质",
                        [
                            "化合物；水溶液中或熔融状态下能导电",
                            "按自身电离判断，不凭当前样品导电性",
                        ],
                    ),
                    (
                        "非电解质",
                        [
                            "化合物；水溶液中和熔融状态下均不能导电",
                            "铜等单质、溶液等混合物不归入两类",
                        ],
                    ),
                    (
                        "电离",
                        [
                            "在水溶液或熔融状态下形成自由移动离子的过程",
                            "不需要通电；NaCl晶体本来就有离子",
                        ],
                    ),
                    (
                        "电解质导电",
                        [
                            "自由移动离子在外接电源、闭合通路中定向移动",
                            "NaCl固体离子不能自由移动，故不导电",
                        ],
                    ),
                ],
            ),
        ),
        (
            "常考知识总结② 强弱与典型干扰",
            table(
                ["判断依据", "例子与提醒"],
                [
                    (
                        "强电解质",
                        [
                            "在水溶液中全部电离；本课用＝",
                            "强酸、强碱、大多数盐；NH₄Cl、BaSO₄均属强",
                        ],
                    ),
                    (
                        "弱电解质",
                        [
                            "在水溶液中部分电离；用⇌",
                            "弱酸、弱碱、水；CH₃COOH、NH₃·H₂O、H₂S",
                        ],
                    ),
                    (
                        "排除干扰",
                        [
                            "溶解度、当前导电性、盐与哪种酸碱有关，均不能替代电离程度",
                            "BaSO₄难溶不等于弱；溶解平衡的⇌不等于弱电离",
                        ],
                    ),
                ],
            ),
        ),
        (
            "常考知识总结③ 书写与条件",
            table(
                ["规则", "核对例式"],
                [
                    (
                        "基本步骤",
                        [
                            "判类别→定条件→写离子→选符号→查原子、电荷与分步",
                            "硫酸铵：(NH₄)₂SO₄＝2NH₄⁺＋SO₄²⁻",
                        ],
                    ),
                    (
                        "多元弱酸",
                        ["分步且以第一步为主", "H₂S ⇌ H⁺＋HS⁻；HS⁻ ⇌ H⁺＋S²⁻"],
                    ),
                    (
                        "硫酸氢钠",
                        [
                            "水溶液与熔融条件不能混用",
                            "水中：NaHSO₄＝Na⁺＋H⁺＋SO₄²⁻；熔融：NaHSO₄＝Na⁺＋HSO₄⁻",
                        ],
                    ),
                    (
                        "碳酸氢钠",
                        [
                            "盐完全解离与酸式酸根部分电离分开",
                            "NaHCO₃＝Na⁺＋HCO₃⁻；HCO₃⁻ ⇌ H⁺＋CO₃²⁻",
                        ],
                    ),
                ],
            ),
        ),
    ]
    for title, visual in summaries:
        s = copy.deepcopy(old[40])
        s.update(
            title=title,
            visual=visual,
            minutes=1,
            content=["据教材与讲义考向整理；对照前面笔记查漏补缺，不在此重新抄写。"],
            teacher_notes="总结用于快速定位已有笔记，不要求在1分钟内重抄整表。常考指讲义列出的考向，不是考试频率统计。课后订正Q1—Q8，Q9选做。Q9教师参考：2H₂O₂ ⇌ H₃O₂⁺＋HO₂⁻。",
        )
        summary.append(s)
    times = {
        1: 1,
        2: 1,
        3: 1,
        4: 1,
        5: 2,
        6: 2,
        7: 1,
        8: 1,
        9: 2,
        10: 2,
        11: 2,
        12: 2,
        13: 2,
        14: 1,
        15: 1,
        16: 3,
        17: 2,
        18: 2,
        19: 2,
        20: 1,
        21: 1,
        22: 4,
        23: 3,
        24: 2,
        25: 2,
        26: 2,
        27: 2,
        28: 2,
        29: 2,
        30: 4,
        31: 2,
        32: 2,
        33: 2,
        34: 2,
        35: 2,
        36: 1,
        37: 2,
        38: 2,
        39: 1,
    }
    c["slides"] = []
    for i in range(1, 40):
        old[i]["minutes"] = times[i]
        c["slides"].append(old[i])
        if i == 17:
            extra17["minutes"] = 3
            c["slides"].append(extra17)
        if i == 28:
            extra28["minutes"] = 2
            c["slides"].append(extra28)
    c["slides"].extend(summary)
    for s in c["slides"]:
        if s["image"]:
            s["image"]["observation_prompt"] = {
                "K1 电解质与非电解质定义": "读原句，圈出对象与条件。",
                "K3 强弱电解质的定义": "读原句，圈出“全部”与“部分”。",
                "图2.15 水分子电离过程示意图": "观察微粒，比较两种表达。",
                "Q7 NaCl三种状态图像综合": "据图判断，写出证据。",
            }.get(s["title"], s["image"]["observation_prompt"])
    old[39]["teacher_notes"] = (
        "独立口答回扣四个问题，1分钟内快速定位薄弱项；随后3页查阅已有笔记，不重抄整表。"
    )
    # Per-page source attribution, including question adaptation status.
    for s in c["slides"]:
        t = s["title"]
        source = "教材印刷56—58页；讲义阅读件3—8页；助手整理，待教师审核。"
        if "Q1 " in t or "Q2 " in t or "Q3 " in t:
            source = "讲义阅读件第6页，讲义参考题及答案；原卷标签未独立回查。"
        elif "Q4 " in t:
            source = "讲义阅读件第6—7页，讲义参考题及答案。"
        elif "Q5 " in t:
            source = "据讲义第7页考向2例2整理为纠错题；非残缺原题原样复制。"
        elif "Q6 " in t:
            source = "据讲义第8页难溶盐考点改写；整理题参考。"
        elif "Q7 " in t:
            source = "讲义阅读件第8页NaCl三状态题，讲义参考答案。"
        elif "Q8 " in t:
            source = "教材印刷57页前3道书写任务；58页两例式整理及水的书写任务。"
        s["teacher_notes"] += "\n来源：" + source
        if t.startswith("Q"):
            s["content"].append("来源：" + source)
    c["activities"][0]["minutes"] = 14
    c["activities"][2]["worksheet"]["instructions"] = [
        "按教材58页填写强弱比较表，再完成Q1、Q2，补写判断依据。"
    ]
    c["activities"][2]["worksheet"]["sections"][1]["response_lines"] = 3
    c["activities"][2]["worksheet"]["sections"][2]["response_lines"] = 2
    w4 = c["activities"][4]["worksheet"]
    w4["sections"][0].update(
        heading="K4 书写五步",
        prompt="记录判类别、定条件、写离子、选符号、核对五步，并标出容易遗漏的检查项。",
    )
    w4["sections"][1:1] = [
        worksheet_table(
            "K5 弱电解质与分步",
            "填写完整电离式，并在式旁标记可逆符号与分步。",
            [
                "醋酸（水溶液）",
                "一水合氨（水溶液）",
                "硫化氢第一步",
                "硫化氢第二步",
                "水（简单离子符号）",
            ],
            ["物质/步骤", "完整电离式（学生填写）"],
        ),
        worksheet_table(
            "K6 酸式盐与条件",
            "按所列状态填写，不将盐解离与酸式酸根电离合并。",
            [
                "NaHSO₄（水溶液）",
                "NaHSO₄（熔融）",
                "NaHCO₃（水中盐解离）",
                "HCO₃⁻（部分电离）",
            ],
            ["物质/条件", "完整电离式（学生填写）"],
        ),
    ]
    w4["sections"][0]["response_lines"] = 3
    w4["sections"][3]["response_lines"] = 6
    w4["sections"] = [w4["sections"][i] for i in [0, 4, 1, 2, 3]]
    q9 = "课后选做Q9：已知水的自偶电离为2H₂O ⇌ H₃O⁺＋OH⁻，H₂O₂也存在极微弱的自偶电离。类比写出其表达式，并检查原子数和电荷。"
    c["homework"]["tasks"][1]["instruction"] = q9
    c["activities"][5]["worksheet"]["sections"].append(
        {
            "heading": "Q9 课后选做",
            "prompt": q9,
            "response_kind": "lines",
            "response_lines": 3,
            "columns": [],
            "row_labels": [],
        }
    )
    # Assign one explicit timed activity per page, including the orientation.
    for start, end, owner in [
        (0, 10, 1),
        (10, 15, 2),
        (15, 22, 3),
        (22, 23, 4),
        (23, 24, 6),
        (24, 34, 5),
        (34, 44, 6),
    ]:
        for s in c["slides"][start:end]:
            s["activity_numbers"] = [owner]
    # Group the plan using the exact ordered slide segments; no second clock.
    groups = [
        ("第一课时｜导入与教材读图", 0, 4, 1),
        ("第一课时｜K1分类与Q3", 4, 10, 1),
        ("第一课时｜K2微观解释与Q4", 10, 15, 2),
        ("第一课时｜K3强弱与Q1、Q2", 15, 22, 3),
        ("第一课时｜笔记与独立回顾", 22, 23, 4),
        ("第二课时｜回顾", 23, 24, 6),
        ("第二课时｜K4方法与Q8前3题", 24, 28, 5),
        ("第二课时｜K5、K6完整例式", 28, 32, 5),
        ("第二课时｜Q5纠错", 32, 34, 5),
        ("第二课时｜Q6辨析", 34, 36, 6),
        ("第二课时｜Q7图像与Q8后3题", 36, 40, 6),
        ("第二课时｜独立回扣与常考总结", 40, 44, 6),
    ]
    stages = []
    for title, start, end, activity in groups:
        slides = c["slides"][start:end]
        stages.append(
            {
                "title": title,
                "objective_numbers": sorted(
                    {n for s in slides for n in s["objective_numbers"]}
                ),
                "activity_numbers": sorted(
                    {n for s in slides for n in s["activity_numbers"]}
                )
                or [activity],
                "assessment_numbers": sorted(
                    {n for s in slides for n in s["assessment_numbers"]}
                ),
                "minutes": sum(s["minutes"] for s in slides),
                "teacher_action": f"按PPT第{start + 1}—{end}页推进："
                + "；".join(s["title"] for s in slides)
                + "。题页先独立作答，讲评页再核对；具体提醒与参考答案见对应页教师备注。",
                "student_action": "依次观察/阅读、独立作答、核对订正；知识表在首次讲解时填写，课末只补缺。作答与记录时间已含本环节预算。",
                "materials": [
                    f"PPT第{start + 1}—{end}页",
                    "教材印刷56—58页；讲义阅读件3—8页；对应学习单",
                ],
                "assessment": "检查答案和依据是否同时完整，记录具体漏项；未掌握的题号带入课后订正，不把课堂预测写成本班实测。",
            }
        )
    c["lesson_stages"] = stages
    c["uncertainties"].append(
        {
            "field": "离线修订与使用边界",
            "description": "本版由助手对一次模型返回稿进行明确拆页、计时、题解和笔记修订；没有追加模型请求，也未通过教师或真实课堂验收。",
            "teacher_action": "请核对教材摘录、讲义整理题与本班节奏。先试讲确认投影与抄写时间，不能仅以文件检查通过视为教学就绪。",
        }
    )
    return c


def validate(c):
    assert len(c["slides"]) == 44
    assert sum(s["minutes"] for s in c["slides"][:23]) == 40
    assert sum(s["minutes"] for s in c["slides"][23:]) == 40
    assert sum(s["minutes"] for s in c["lesson_stages"]) == 80
    assert sum(s["minutes"] for s in c["activities"]) == 80
    assert len({s["image"]["asset_id"] for s in c["slides"] if s["image"]}) == 8
    assert "H₃O₂⁺＋HO₂⁻" not in c["homework"]["tasks"][1]["instruction"]
    for number in range(1, 9):
        question = next(
            i for i, s in enumerate(c["slides"]) if s["title"].startswith(f"Q{number} ")
        )
        answer = next(
            i
            for i, s in enumerate(c["slides"])
            if s["title"].startswith(f"Q{number} ") and "核对" in s["title"]
        )
        assert question < answer


def main():
    assert digest(LIVE / "returned-candidate.json") == RAW_HASH
    if OUT.exists() and {p.name for p in OUT.iterdir()} != {
        "edited-raw-candidate.json"
    }:
        raise SystemExit("Output exists; refusing replacement")
    raw = json.loads((LIVE / "returned-candidate.json").read_text("utf-8"))
    revised = revise(raw)
    validate(revised)
    payload = json.loads((LIVE / "teacher-brief.json").read_text("utf-8"))
    # Cover uses the payload topic, not the saved draft's organizational label.
    payload["topic"] = "电解质的电离"
    canonical = normalize_preparation_candidate(
        revised, normalize_preparation_payload(payload)
    )
    image_data = {}
    for p in (BRIEF / "assets").glob("*.image"):
        image_data["IMG-" + digest(p)] = p.read_bytes()
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "teacher-brief.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), "utf-8"
    )
    edited_path = OUT / "edited-raw-candidate.json"
    if edited_path.exists():
        assert json.loads(edited_path.read_text("utf-8")) == revised
    else:
        edited_path.write_text(
            json.dumps(revised, ensure_ascii=False, indent=2), "utf-8"
        )
    result = BundledArtifactRenderer(
        Path(
            "C:/Users/20671/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/python.exe"
        )
    ).render(
        canonical,
        output_kind="joint",
        output_dir=OUT,
        image_data=image_data,
        report_progress=None,
        is_cancelled=lambda: False,
    )
    # Insert the same unedited source image before the worksheet's Q7 prompt.
    import io

    from docx import Document
    from docx.enum.table import WD_ROW_HEIGHT_RULE
    from docx.shared import Inches

    docpath = OUT / "student_worksheet.docx"
    doc = Document(docpath)
    for row in doc.tables[2].rows[1:]:
        row.height = Inches(0.32)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    # Keep a short writing table together, preserving all cells and answer space.
    for tab in doc.tables:
        for row in tab.rows[:-1]:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
    q7 = next(p for p in doc.paragraphs if p.text == "Q7 图像综合")
    image_p = q7.insert_paragraph_before()
    image_p.add_run().add_picture(
        io.BytesIO(
            image_data[
                "IMG-d4b7868f7cd0721958e3806018d8d3130cc2fac8fb221780533d117158ecb363"
            ]
        ),
        width=Inches(5.8),
    )
    image_p.paragraph_format.keep_with_next = True
    q7.insert_paragraph_before("题图来源：讲义阅读件第8页；与PPT Q7为同一原图。")
    doc.save(docpath)
    receipt = {
        "raw_candidate_sha256": RAW_HASH,
        "api_calls_during_revision": 0,
        "revision_kind": "assistant_offline_editorial_candidate",
        "slides": 44,
        "lesson_minutes": [40, 40],
        "source_images": 8,
        "renderer_result": result,
        "worksheet_q7_source_image_added": True,
        "office_visual_review": "pending",
        "teacher_review_required": True,
        "publication_allowed": False,
        "artifacts": {
            p.name: digest(p) for p in OUT.iterdir() if p.suffix in {".pptx", ".docx"}
        },
    }
    (OUT / "revision-verification.json").write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2), "utf-8"
    )
    print(
        json.dumps(
            {k: v for k, v in receipt.items() if k != "renderer_result"},
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
