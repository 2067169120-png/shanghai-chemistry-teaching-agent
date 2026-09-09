"""Task-local DOCX pagination fix; preserves text and all previous artifacts."""
from pathlib import Path
from copy import deepcopy
import hashlib
import json
import re
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[4]
source = ROOT / "outputs/备课/2026-09-09-电解质的电离-讲义主线两课时-r2"
target = ROOT / "outputs/备课/2026-09-09-电解质的电离-讲义主线两课时-r4"
if target.exists():
    raise RuntimeError("Use a new target; previous artifacts are retained")
shutil.copytree(source, target)
path = target / "lesson_plan.docx"
doc = Document(path)
before = "\n".join(p.text for p in doc.paragraphs)
changed = []
header = re.compile(r"^(?:教师活动：\s*)?PPT\s*第\s*\d+")
for p in list(doc.paragraphs):
    if not header.match(p.text):
        continue
    # These are eight long activity paragraphs with soft breaks. Split only
    # slide-heading/body groups, retaining every formatted run and charge.
    lines = [[]]
    for run in p._p.findall(qn("w:r")):
        for child in run:
            if child.tag == qn("w:rPr"):
                continue
            if child.tag == qn("w:br"):
                lines.append([])
            else:
                r = OxmlElement("w:r")
                if run.rPr is not None:
                    r.append(deepcopy(run.rPr))
                r.append(deepcopy(child))
                lines[-1].append(r)
    groups = []
    body = []
    for line in lines:
        text = "".join(t.text or "" for r in line for t in r.iter(qn("w:t")))
        if header.match(text):
            if body:
                groups.append((False, body)); body = []
            groups.append((True, [line])); changed.append(text)
        else:
            body.append(line)
    if body:
        groups.append((False, body))
    for is_header, group in groups:
        # Empty separating lines become paragraph spacing, not blank pages.
        while group and not group[0]: group.pop(0)
        while group and not group[-1]: group.pop()
        if not group: continue
        node = OxmlElement("w:p")
        if p._p.pPr is not None: node.append(deepcopy(p._p.pPr))
        for line_index, line in enumerate(group):
            if line_index:
                r = OxmlElement("w:r"); r.append(OxmlElement("w:br")); node.append(r)
            for r in line: node.append(r)
        p._p.addprevious(node)
        Paragraph(node, p._parent).paragraph_format.keep_with_next = is_header
    p._p.getparent().remove(p._p)
assert len(changed) == 44, len(changed)
assert re.sub(r"\s+", "", "\n".join(p.text for p in doc.paragraphs)) == re.sub(r"\s+", "", before)
doc.save(path)
assert re.sub(r"\s+", "", "\n".join(p.text for p in Document(path).paragraphs)) == re.sub(r"\s+", "", before)
report_path = target / "verification.json"
report = json.loads(report_path.read_text("utf-8"))
report["artifacts"] = {
    p.name: hashlib.sha256(p.read_bytes()).hexdigest()
    for p in target.iterdir() if p.suffix in (".docx", ".pptx")
}
report["pagination_fix"] = {"keep_slide_heading_with_next":44, "document_text_unchanged":True}
report["visual_review"] = "pending final Office export; PPT and worksheet unchanged from r2"
report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(report, ensure_ascii=False))
