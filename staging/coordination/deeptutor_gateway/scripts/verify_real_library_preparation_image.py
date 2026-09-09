"""Offline real-source image handoff and native UI inspection evidence.

Uses isolated state, no credentials/provider, and an automated QA confirmation.
It does not represent human teaching approval or a model understanding pixels.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import time
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
ROOT = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(ROOT))


def wait(app, predicate, seconds=45):
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        app.processEvents()
        if predicate():
            return
        time.sleep(0.01)
    raise TimeoutError("Local UI did not finish the verified image operation")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    output = args.output.resolve()
    output.relative_to(ROOT / "runtime/deeptutor_shchem/qa")
    output.mkdir(parents=True, exist_ok=False)

    from PySide6.QtCore import QTimer
    from PySide6.QtWidgets import QApplication, QScrollArea

    from integrations.deeptutor_shchem_v1.desktop_facade import DesktopWorkbenchFacade
    from integrations.deeptutor_shchem_v1.desktop_paths import DesktopPaths
    from integrations.deeptutor_shchem_v1.desktop_preparation_images import (
        PreparationImageStore,
    )
    from integrations.deeptutor_shchem_v1.desktop_workbench.library_detail import (
        LibraryDetailDialog,
    )
    from integrations.deeptutor_shchem_v1.desktop_workbench.main_window import (
        WORKBENCH_STYLE,
        install_font_fallbacks,
    )
    from integrations.deeptutor_shchem_v1.desktop_workbench.preparation_images_widget import (
        PreparationImageMetadataDialog,
    )
    from integrations.deeptutor_shchem_v1.desktop_workbench.tasks import (
        DesktopTaskBridge,
    )
    from integrations.deeptutor_shchem_v1.desktop_workbench.workflow_pages import (
        PreparationPage,
    )

    app = QApplication.instance() or QApplication([])
    app.setStyleSheet(WORKBENCH_STYLE)
    install_font_fallbacks()
    paths = DesktopPaths.from_workspace(
        ROOT, state_root=output / "isolated-personal-state"
    )
    facade = DesktopWorkbenchFacade(paths, provider_store=object())
    tasks = DesktopTaskBridge()
    try:
        card = facade.search_themes(scope="master", query="银镜", limit=1).cards[0]
        detail = facade.library_theme_detail(card)
        descriptors = [
            *detail.shared_images,
            *(image for part in detail.parts for image in part.question_images),
        ]
        if not descriptors:
            raise RuntimeError("No real source image; no fixture fallback")
        descriptor = descriptors[0]
        original = facade.library_image(descriptor)
        page = PreparationPage(facade, tasks)
        page._availability_timer.stop()
        page.topic.setText("原题图片接入验证")
        page.audience.setText("本地QA，不是授课样课")
        page.objective.setPlainText("核对原图、来源说明与本地备课图片传递。")
        page.materials.setPlainText(
            "保留教材知识主线，原题图作为辅助材料。此为本地接入验证。"
        )
        before = page._payload()
        dialog = LibraryDetailDialog(detail, tasks, facade.library_image)
        dialog.preparation_image_requested.connect(page.import_library_image)
        dialog.show()
        wait(app, lambda: all(w._task_id is None for w in dialog._image_widgets))
        assert dialog._image_widgets[0].has_image
        screenshots = []
        for width in (420, 900):
            dialog.resize(width, 800)
            for _ in range(5):
                app.processEvents()
            assert dialog.minimumSizeHint().width() <= width
            scroll = dialog.tabs.widget(0)
            assert scroll.horizontalScrollBar().maximum() == 0
            path = output / f"library-{width}.png"
            assert dialog.grab().save(str(path))
            screenshots.append(str(path))

        # Exercise the actual button and metadata dialog. This is a scripted QA
        # click in isolated state, not a claim of a real teacher's approval.
        def confirm():
            metadata = app.activeModalWidget()
            if not isinstance(metadata, PreparationImageMetadataDialog):
                raise TypeError("Expected the native image metadata dialog")
            metadata.resize(420, 430)
            app.processEvents()
            metadata.grab().save(str(output / "metadata-420.png"))
            screenshots.append(str(output / "metadata-420.png"))
            metadata.confirm_button.click()

        QTimer.singleShot(0, confirm)
        dialog._preparation_buttons[0].click()
        wait(app, lambda: page._library_image_task_id is None)
        assets = page.image_assets_widget.assets()
        assert len(assets) == 1
        asset = assets[0]
        store = PreparationImageStore(paths.task_root / "preparation-v1/images")
        assert store.load(asset) == original == facade.library_image(descriptor)
        assert asset["sha256"] == descriptor.sha256
        after = page._payload()
        for key, value in before.items():
            assert after[key] == value
        assert "view_id" not in json.dumps(after) and "crop_id" not in json.dumps(after)
        # A local renderer probe, not a model-generated lesson: only a labelled
        # cover and the selected source image. No DOCX or teaching claims.
        from pptx import Presentation

        from integrations.deeptutor_shchem_v1.desktop_preparation_renderer import (
            NativePreparationRenderer,
        )

        probe = {
            "title": "原题图片接入验证",
            "topic": "原题图片接入验证",
            "candidate_only": True,
            "teacher_review_required": True,
            "publication_allowed": False,
            "official_claim_allowed": False,
            "image_assets": [asset],
            "objectives": [],
            "activities": [],
            "assessments": [],
            "lesson_stages": [],
            "slides": [
                {
                    "id": "S01",
                    "title": "原题图片接入验证",
                    "purpose": "本地排版验证，不是授课样课",
                    "content": ["仅核对原图、来源和排版；未调用模型。"],
                    "minutes": 1,
                },
                {
                    "id": "S02",
                    "title": "题库原图",
                    "purpose": "核对原图",
                    "content": ["本页为原图接入测试，不提供题目解答。"],
                    "minutes": 1,
                    "image": {
                        "asset_id": asset["asset_id"],
                        "observation_prompt": "核对图中的符号和原题条件。",
                    },
                },
            ],
        }
        NativePreparationRenderer().render(
            probe,
            output_kind="ppt",
            output_dir=output / "renderer-probe",
            image_data={asset["asset_id"]: original},
        )
        pptx = output / "renderer-probe/lesson_presentation.pptx"
        picture = next(
            shape
            for slide in Presentation(pptx).slides
            for shape in slide.shapes
            if shape.shape_type == 13
        )
        assert picture.image.blob == original
        assert (
            abs(
                (picture.width / picture.height) / (asset["width"] / asset["height"])
                - 1
            )
            < 0.001
        )
        for width in (420, 900):
            page.resize(width, 900)
            page.show()
            app.processEvents()
            scroll = page.findChild(QScrollArea)
            scroll.ensureWidgetVisible(page.image_assets_widget, 0, 12)
            app.processEvents()
            assert scroll.horizontalScrollBar().maximum() == 0
            path = output / f"preparation-{width}.png"
            assert page.grab().save(str(path))
            screenshots.append(str(path))
        report = {
            "source_title": detail.title_zh,
            "source_paper": detail.paper_title_zh,
            "source_image_caption": descriptor.caption_zh,
            "source_sha256": hashlib.sha256(original).hexdigest(),
            "source_bytes": len(original),
            "asset": asset,
            "exact_source_copy": True,
            "existing_form_preserved": True,
            "renderer_probe_pptx": str(pptx),
            "pptx_embeds_exact_original_bytes": True,
            "pptx_picture_aspect_ratio_preserved": True,
            "model_calls": 0,
            "credential_store_accessed": False,
            "confirmation": "automated_qa_not_teacher_approval",
            "screenshots": screenshots,
            "visual_review": "pending",
        }
        (output / "verification.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        print(json.dumps(report, ensure_ascii=False))
        page.close()
    finally:
        tasks.shutdown(5000)
        facade.shutdown()


if __name__ == "__main__":
    main()
