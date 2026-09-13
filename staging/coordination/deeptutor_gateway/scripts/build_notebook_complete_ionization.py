"""Revise the frozen Word-led lesson into self-contained notebook pages, offline."""

from __future__ import annotations

import argparse
import json
import re
import subprocess
from copy import deepcopy
from pathlib import Path

from build_word_led_ionization import (
    B57,
    B58,
    EXPECTED,
    QUOTE,
    ROOT,
    W5,
    BundledArtifactRenderer,
    _canonical_candidate_digest,
    assets,
    comparison,
    normalize_preparation_candidate,
    sha,
)
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

BASE = ROOT / "outputs/备课/2026-09-09-电解质的电离-讲义主线两课时-r4"
CROPS = ROOT / "runtime/deeptutor_shchem/qa/notebook-definition-assets-20260909"


def revised_lesson():
    raw = json.loads((BASE / "word-led-raw-candidate.json").read_text("utf-8"))
    brief = json.loads((BASE / "teacher-brief.json").read_text("utf-8"))
    coverage = json.loads((BASE / "source-coverage.json").read_text("utf-8"))
    meta, payload, _ = assets()
    crop_ids = []
    for item in json.loads((CROPS / "manifest.json").read_text("utf-8")):
        path = CROPS / item["name"]
        assert sha(path) == item["sha256"]
        assert sha(Path(item["source_page"])) == item["source_sha256"]
        asset_id = "IMG-" + item["sha256"]
        crop_ids.append(asset_id)
        meta.append(
            {
                "asset_id": asset_id,
                "sha256": item["sha256"],
                "caption": item["source"],
                "source": item["source"],
                "purpose": "对照原句与可编辑笔记，圈出定义条件。",
                "width": item["width"],
                "height": item["height"],
                "content_type": "image/png",
            }
        )
        payload[asset_id] = path.read_bytes()
    brief["image_assets"] = meta
    brief["advanced"]["template_and_delivery"] += (
        " 原句截图与可编辑文字对应；总结保留完整例式。"
    )
    raw["slides"][13]["image"] = {
        "asset_id": crop_ids[0],
        "observation_prompt": "找出状态、自由移动与过程三个条件。",
    }
    raw["slides"][13]["content"] = [
        "教材第57页原句：\n“" + QUOTE + "”",
        "笔记：圈出状态与自由移动。电离不需要通电；导电条件见下一页。",
    ]
    raw["slides"][13]["teacher_notes"] = (
        "先读原图，再摘录右侧完整原句，至少留45秒。定义与导电条件分开，下一页对照。来源："
        + B57
    )
    raw["slides"][17]["image"] = {
        "asset_id": crop_ids[1],
        "observation_prompt": "对照两句定义，圈出水溶液中、全部、仅有部分。",
    }
    raw["slides"][17]["content"] = [
        "教材第58页原句：\n“像氯化钠、氯化氢、氢氧化钠等在水溶液中能够全部电离为自由移动离子的电解质称为强电解质。”",
        "“像醋酸、一水合氨（NH₃·H₂O）等在水溶液中仅有部分分子能电离出自由移动离子的电解质称为弱电解质。”",
    ]
    raw["slides"][17]["teacher_notes"] = (
        "原图与右侧两句逐一对应。先读、圈条件，随后用下一页完整表记录强弱类别与易错点，不用灯泡亮度判强弱。来源："
        + B58
    )
    original = deepcopy(raw["slides"][39])
    first = deepcopy(original)
    first.update(
        title="常考知识总结五 分步电离",
        minutes=1,
        content=["本表列出完整笔记：保留每步离子、符号与判断依据。"],
        teacher_notes="先对照学习单6订正，不重新抄题。H₂S每步释放一个H⁺；NaHCO₃先完全电离，HCO₃⁻再部分电离。来源："
        + W5,
        visual=comparison(
            ["完整电离式", "判断与易错"],
            [
                (
                    "硫化氢 H₂S",
                    [
                        "H₂S ⇌ H⁺ + HS⁻\nHS⁻ ⇌ H⁺ + S²⁻",
                        "多元弱酸分步电离；\n以第一步为主，不能一步写到底。",
                    ],
                ),
                (
                    "碳酸氢钠 NaHCO₃",
                    [
                        "NaHCO₃ = Na⁺ + HCO₃⁻\nHCO₃⁻ ⇌ H⁺ + CO₃²⁻",
                        "盐的完全电离与酸根的部分电离\n是两步，不能合写成全部电离。",
                    ],
                ),
            ],
        ),
    )
    second = deepcopy(original)
    second.update(
        title="常考知识总结六 酸式盐看状态",
        minutes=1,
        content=["硫酸氢钠 NaHSO₄：状态不同，电离式不同（按本讲义的中学简化表达）。"],
        teacher_notes="对照两种状态核对学习单6，不能只记“分别处理”。难溶盐的电离程度与溶解平衡辨析见第35—36页。来源："
        + W5,
        visual=comparison(
            ["完整电离式", "判断与易错"],
            [
                (
                    "水溶液",
                    [
                        "NaHSO₄ = Na⁺ + H⁺ + SO₄²⁻",
                        "本讲义的水溶液简化表达\n写出Na⁺、H⁺、SO₄²⁻。",
                    ],
                ),
                (
                    "熔融状态",
                    [
                        "NaHSO₄ = Na⁺ + HSO₄⁻",
                        "保留HSO₄⁻整体；\n不能照搬水溶液中的写法。",
                    ],
                ),
            ],
        ),
    )
    raw["slides"][39:40] = [first, second]
    coverage[39:40] = [
        {"slide": 40, "title": first["title"], "source": W5, "role": "笔记总结"},
        {"slide": 41, "title": second["title"], "source": W5, "role": "笔记总结"},
    ]
    for n, entry in enumerate(coverage, 1):
        entry["slide"] = n

    # Only post-insertion literal page references need adjustment; no question text changes.
    def remap(value):
        if isinstance(value, str):
            return re.sub(
                r"PPT第(4[1-4])页", lambda m: "PPT第" + str(int(m[1]) + 1) + "页", value
            )
        if isinstance(value, list):
            return [remap(v) for v in value]
        if isinstance(value, dict):
            return {k: remap(v) for k, v in value.items()}
        return value

    raw = remap(raw)
    # Keep all original worksheet response spaces and tasks; synchronize page guidance.
    raw["activities"][5]["worksheet"]["instructions"].append(
        "第40—41页提供完整订正表：每式保留状态与判断依据。"
    )
    for activity, stage in zip(raw["activities"], raw["lesson_stages"]):
        number = raw["activities"].index(activity) + 1
        selected = [
            (n, slide)
            for n, slide in enumerate(raw["slides"], 1)
            if slide["activity_numbers"] == [number]
        ]
        sections = []
        for n, slide in selected:
            body = "\n".join(slide["content"])
            # Knowledge tables are actual teaching content, not just the slide title.
            if n in (40, 41):
                body += "\n" + "\n".join(
                    row["label"] + "：" + "；".join(row["values"])
                    for row in slide["visual"]["comparison"]["rows"]
                )
            sections.append(
                f"PPT第{n}页 {slide['title']} {slide['minutes']}分钟\n{body}\n{slide['teacher_notes']}"
            )
        action = "\n\n".join(sections)
        for target in (activity, stage):
            target["teacher_action"] = action
            target["minutes"] = sum(slide["minutes"] for _, slide in selected)
            target["materials"][0] = f"PPT第{selected[0][0]}—{selected[-1][0]}页"
    assert len(raw["slides"]) == 45
    assert [
        sum(s["minutes"] for s in raw["slides"][:24]),
        sum(s["minutes"] for s in raw["slides"][24:]),
    ] == [40, 40]
    assert len(meta) == 8
    return raw, brief, coverage, payload


def split_plan_headings(path):
    """Retain all text/run formatting while keeping each slide heading with its body."""
    doc = Document(path)
    text_before = re.sub(r"\s+", "", "\n".join(p.text for p in doc.paragraphs))
    header = re.compile(r"^(?:教师活动：\s*)?PPT\s*第\s*\d+")
    count = 0
    for p in list(doc.paragraphs):
        if not header.match(p.text):
            continue
        lines = [[]]
        for run in p._p.findall(qn("w:r")):
            for child in run:
                if child.tag == qn("w:rPr"):
                    continue
                if child.tag == qn("w:br"):
                    lines.append([])
                else:
                    node = OxmlElement("w:r")
                    if run.rPr is not None:
                        node.append(deepcopy(run.rPr))
                    node.append(deepcopy(child))
                    lines[-1].append(node)
        groups, body = [], []
        for line in lines:
            text = "".join(t.text or "" for r in line for t in r.iter(qn("w:t")))
            if header.match(text):
                if body:
                    groups.append((False, body))
                    body = []
                groups.append((True, [line]))
                count += 1
            else:
                body.append(line)
        if body:
            groups.append((False, body))
        for is_header, group in groups:
            while group and not group[0]:
                group.pop(0)
            while group and not group[-1]:
                group.pop()
            if not group:
                continue
            node = OxmlElement("w:p")
            if p._p.pPr is not None:
                node.append(deepcopy(p._p.pPr))
            for index, line in enumerate(group):
                if index:
                    run = OxmlElement("w:r")
                    run.append(OxmlElement("w:br"))
                    node.append(run)
                for run in line:
                    node.append(run)
            p._p.addprevious(node)
            Paragraph(node, p._parent).paragraph_format.keep_with_next = is_header
        p._p.getparent().remove(p._p)
    assert count == 45, count
    assert text_before == re.sub(r"\s+", "", "\n".join(p.text for p in doc.paragraphs))
    doc.save(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--artifact-python", type=Path, required=True)
    args = parser.parse_args()
    output = args.output.resolve()
    output.relative_to(ROOT)
    if output.exists():
        raise RuntimeError("Retain existing output; choose a new directory")
    assert all(sha(path) == digest for path, digest in EXPECTED.items())
    raw, brief, coverage, payload = revised_lesson()
    candidate = normalize_preparation_candidate(raw, brief)
    candidate["source_basis"]["statement_zh"] = (
        "依据Word讲义第3—8页与教材印刷56—58页；原句、讲义归纳和整理题分别标明。"
    )
    for issue in candidate["uncertainties"]:
        if issue["field"] == "source_basis":
            issue["description"] = (
                "已按母本核对8幅来源图片：4幅教材图、3处教材原段、1幅讲义题图；本轮未调用模型API。"
            )
    candidate["candidate_id"] = (
        "PREPCAND-" + _canonical_candidate_digest(candidate)[:32]
    )
    assert all(
        sum(s["minutes"] for s in candidate[key]) == 80
        for key in ("slides", "activities", "lesson_stages")
    )
    BundledArtifactRenderer(args.artifact_python).render(
        candidate,
        output_kind="linked_bundle",
        output_dir=output,
        report_progress=lambda *_: None,
        is_cancelled=lambda: False,
        image_data=payload,
    )
    # Orchestration uses the project interpreter; all DOCX authoring uses the
    # authoritative bundled artifact runtime, including this pagination edit.
    pagination_code = (
        "import sys,re; from pathlib import Path; from copy import deepcopy; "
        "from docx import Document; from docx.oxml import OxmlElement; "
        "from docx.oxml.ns import qn; from docx.text.paragraph import Paragraph; "
        "source=Path(sys.argv[1]).read_text('utf-8'); "
        "exec(source[source.index('def split_plan_headings('):source.index('\\ndef main():')]); "
        "split_plan_headings(Path(sys.argv[2]))"
    )
    subprocess.run(
        [
            str(args.artifact_python),
            "-X",
            "utf8",
            "-B",
            "-c",
            pagination_code,
            str(Path(__file__).resolve()),
            str(output / "lesson_plan.docx"),
        ],
        check=True,
    )
    for filename, value in [
        ("word-led-raw-candidate.json", raw),
        ("teacher-brief.json", brief),
        ("source-coverage.json", coverage),
    ]:
        (output / filename).write_text(
            json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    report = {
        "slide_count": 45,
        "period_minutes": [40, 40],
        "source_images": 8,
        "worksheet_units": 8,
        "source_files_unchanged": all(
            sha(path) == digest for path, digest in EXPECTED.items()
        ),
        "new_model_calls": 0,
        "teacher_approval": False,
        "classroom_validation": False,
        "visual_review": "pending",
        "artifacts": {
            path.name: sha(path)
            for path in output.iterdir()
            if path.suffix in (".pptx", ".docx")
        },
    }
    (output / "verification.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
