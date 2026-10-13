"""Native selection, stale/cancelled tasks, and real preparation handoff."""
from copy import deepcopy
import math

import pytest
from PySide6.QtCore import Qt, QPointF
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QMessageBox

from test_textbook_section_reader import section_fixture as section_fixture
from test_textbook_asset_catalog import asset_fixture as asset_fixture
from test_textbook_bookshelf import shelf_fixture as shelf_fixture, disk_bytes
from test_textbook_bookshelf_ui import app as app, page as page
from test_textbook_preparation_crops import crop_fixture as crop_fixture, request
from test_desktop_library_ui import _ManualBridge
from test_desktop_studio_ui import window as window, settle
from integrations.deeptutor_shchem_v1.desktop_textbook_preparation_crops import TextbookPreparationCropService
from integrations.deeptutor_shchem_v1.desktop_workbench.textbook_preparation_crop_dialog import TextbookPreparationCropDialog
from integrations.deeptutor_shchem_v1.desktop_workbench.tasks import DesktopTaskBridge


@pytest.fixture
def dialog(app,crop_fixture):
    f,bridge=crop_fixture,_ManualBridge()
    value=TextbookPreparationCropDialog(f.facade,bridge,f.selection)
    value.show();app.processEvents()
    bridge.succeed(bridge.take('读取教材框选原页'),run=True)
    app.processEvents()
    yield value,bridge,f
    value.close();value.deleteLater();app.processEvents()


def select(value,bridge):
    value._canvas_changed((100,100,610,630));value.purpose.setText('观察合成色块与边缘')
    value.caption.setText('合成图表');value.preview_button.click()
    bridge.succeed(bridge.take('核对教材裁片'),run=True)


def test_dialog_has_no_implicit_selection_and_confirmation_only_returns_request(dialog):
    value,bridge,f=dialog
    before=disk_bytes(f.root)
    assert value._chosen_bounds is None and not value.preview_button.isEnabled()
    assert not value.confirm.isEnabled() and not value.checked.isChecked()
    select(value,bridge)
    assert value.canvas._selection == (100,100,610,630)
    assert value.crop_preview.has_image and not value.confirm.isEnabled()
    value.checked.setChecked(True);assert value.confirm.isEnabled()
    value.confirm.click()
    assert value.selection_request==request(f)
    assert not f.task_root.exists() and disk_bytes(f.root)==before


@pytest.mark.parametrize('field',['caption','purpose','bounds'])
def test_editing_current_crop_invalidates_previous_preview(dialog,field):
    value,bridge,_=dialog;select(value,bridge);value.checked.setChecked(True)
    if field=='bounds':value.edges['left'].setValue(101)
    else:getattr(value,field).setText('新内容')
    assert value._preview is None and not value.checked.isChecked() and not value.confirm.isEnabled()
    assert not value.crop_preview.has_image


def test_whole_page_is_explicit_and_reloaded_larger_page_has_correct_edges(dialog,app):
    from PIL import Image
    import io
    value,bridge,f=dialog
    value.whole_page.click()
    assert value._chosen_bounds=={'left':0,'top':0,'right':f.page['width'],'bottom':f.page['height']}
    smaller=deepcopy(f.page);buffer=io.BytesIO();Image.new('RGB',(100,150),'blue').save(buffer,format='PNG')
    smaller.update(page_data=buffer.getvalue(),width=100,height=150)
    value.reload.click();bridge.succeed(bridge.take('读取教材框选原页'),smaller)
    assert value.edges['right'].value()==100 and value._chosen_bounds is None
    value.reload.click();bridge.succeed(bridge.take('读取教材框选原页'),f.page);app.processEvents()
    assert value.edges['right'].value()==f.page['width'] and value.edges['bottom'].value()==f.page['height']
    assert value._chosen_bounds is None


def test_mouse_drag_is_in_original_pixel_coordinates_and_source_zoom_preserves_box(dialog,app):
    value,_,_=dialog;app.processEvents()
    start=value.canvas.mapFromScene(QPointF(200,250));end=value.canvas.mapFromScene(QPointF(600,700))
    QTest.mousePress(value.canvas.viewport(),Qt.MouseButton.LeftButton,pos=start)
    QTest.mouseMove(value.canvas.viewport(),end)
    QTest.mouseRelease(value.canvas.viewport(),Qt.MouseButton.LeftButton,pos=end)
    bounds=value._chosen_bounds;assert bounds
    tolerance=math.ceil(1/value.canvas.transform().m11())+2
    for edge,expected in zip(bounds.values(),(200,250,600,700),strict=True):assert abs(edge-expected)<=tolerance
    value.actual_source.click();assert value.canvas.transform().m11()==1
    assert value._chosen_bounds==bounds
    value.fit_source.click();assert value._chosen_bounds==bounds


def test_cancelled_read_recovers_controls_and_closing_cancels_own_task(dialog):
    value,bridge,_=dialog
    value.reload.click();task=bridge.take('读取教材框选原页')
    value._task_cancelled(task.task_id,task.label)
    assert not value._busy and value.reload.isEnabled() and not value.confirm.isEnabled()
    value.reload.click();later=bridge.take('读取教材框选原页');value.reject()
    assert later.task_id in bridge.cancelled
    bridge.succeed(later,run=True)
    assert value._closed and value._page is None and value.selection_request is None


def test_stale_preview_callback_cannot_confirm_new_fields(dialog):
    value,bridge,_=dialog
    value.whole_page.click();value.purpose.setText('用途');value.preview_button.click()
    task=bridge.take('核对教材裁片');value.caption.setText('另一图题')
    bridge.succeed(task,run=True);value._task_finished(task.task_id)
    assert not value._busy and value._preview is None and not value.confirm.isEnabled()


def test_native_source_and_crop_paint_at_narrow_width(dialog,app):
    value,bridge,_=dialog;value.resize(420,820);app.processEvents();select(value,bridge);app.processEvents()
    assert value.splitter.orientation()==Qt.Orientation.Vertical
    assert value.width()==420 and value.height()==820
    for widget in (value.status,value.preview_button,value.confirm,value.cancel):
        rectangle=widget.rect();top=widget.mapTo(value,rectangle.topLeft());bottom=widget.mapTo(value,rectangle.bottomRight())
        assert top.x()>=0 and bottom.x()<value.width() and top.y()>=0 and bottom.y()<value.height()
    assert value.edges['right'].width()>=95
    for view in (value.canvas,value.crop_preview.image):
        value.scroll.ensureWidgetVisible(view);app.processEvents()
        image=view.viewport().grab().toImage()
        colored=sum(max(image.pixelColor(x,y).red(),image.pixelColor(x,y).blue())-image.pixelColor(x,y).green()>70
                    for x in range(0,image.width(),10) for y in range(0,image.height(),10))
        assert colored>15  # Verify actual PDF/crop paint, not just metadata widgets.


def test_real_worker_renders_page_and_cancelled_preview_is_recoverable(app,crop_fixture,monkeypatch):
    from integrations.deeptutor_shchem_v1.reader_cancellation import ReadCancelled
    f=crop_fixture;bridge=DesktopTaskBridge()
    value=TextbookPreparationCropDialog(f.facade,bridge,f.selection);value.show()
    settle(app,lambda:not value._busy)
    assert value._page and value.canvas.scene().items()
    def stopped(*_):raise ReadCancelled()
    monkeypatch.setattr(f.facade,'preparation_textbook_crop_preview',stopped)
    value.whole_page.click();value.purpose.setText('合成用途');value.preview_button.click()
    settle(app,lambda:not value._busy)
    assert value.preview_button.isEnabled() and not value.confirm.isEnabled()
    assert '取消' in value.status.text()
    value.close();bridge.shutdown();value.deleteLater();app.processEvents()


def test_bookshelf_opens_crop_and_emits_explicit_request(page,monkeypatch):
    from integrations.deeptutor_shchem_v1.desktop_workbench import textbook_preparation_crop_dialog as module
    value,_,_=page;value.tabs.setCurrentIndex(value.asset_tab)
    chosen={'expected_image_sha256':'a'*64};observed=[]
    class Accepted:
        selection_request=chosen
        class DialogCode:Accepted=1
        def __init__(self,facade,tasks,selection,parent,*,allow_composition_handoff=False):
            assert allow_composition_handoff; observed.append(selection)
        def exec(self):return 1
        def deleteLater(self):pass
    monkeypatch.setattr(module,'TextbookPreparationCropDialog',Accepted)
    value.preparation_image_requested.connect(lambda item:observed.append(item))
    value.crop_button.click()
    assert len(observed)==2 and observed[1]==chosen
    assert set(observed[0])=={'visual_asset_id','revision'}


def bind_import(win,f,monkeypatch):
    bridge=_ManualBridge();prep=win.preparation_page;prep.tasks=bridge
    service=TextbookPreparationCropService(f.root,win.facade.paths.task_root)
    monkeypatch.setattr(win.facade,'import_textbook_preparation_crop',lambda choice:service.confirm(**choice))
    return prep,bridge


def test_main_window_image_handoff_preserves_form_and_blocks_new_lesson(window,crop_fixture,monkeypatch):
    win,app=window;f=crop_fixture;prep,bridge=bind_import(win,f,monkeypatch)
    prep.topic.setText('原有课题');prep.materials.setPlainText('原有公共材料');prep.objective.setPlainText('原有目标')
    prep.image_assets_widget.set_image_input_mode('local_only')
    before=deepcopy(prep._payload());choice=request(f)
    win.textbook_page.preparation_image_requested.emit(choice)
    assert win.stack.currentWidget() is prep and prep.studio_busy() and not prep.isEnabled()
    win.new_preparation();assert prep._payload()==before
    assert not prep.apply_studio_template('unit')
    bridge.succeed(bridge.take('添加教材备课图表'),run=True);app.processEvents()
    after=prep._payload();assert len(after['image_assets'])==1
    after.pop('image_assets');assert after==before
    assert prep.image_assets_widget.image_input_mode()=='local_only' and not prep.studio_busy() and prep.isEnabled()
    monkeypatch.setattr(QMessageBox,'information',lambda *_:None)
    assert not prep.import_textbook_image(choice) and not bridge.pending


def test_imported_crop_can_be_bound_in_native_lesson_dialog_then_exported(window,crop_fixture,monkeypatch):
    from integrations.deeptutor_shchem_v1.desktop_workbench.lesson_design_dialog import LessonDesignDialog
    from integrations.deeptutor_shchem_v1.desktop_lesson_output import export_design, checked_file
    from pptx import Presentation
    win,app=window;f=crop_fixture;prep,bridge=bind_import(win,f,monkeypatch)
    prep.topic.setText('观察合成图表');prep.objective.setPlainText('识别色块与边缘')
    prep.materials.setPlainText('合成课堂资料，不是真实教材题目')
    assert prep.import_textbook_image(request(f));bridge.succeed(bridge.take('添加教材备课图表'),run=True)
    value=LessonDesignDialog(prep);value.show();value.add_node();value.bind.click()
    asset=prep.image_assets_widget.assets()[0];node=value.history.value['nodes'][0]
    assert node['image_ids']==[asset['asset_id']] and node['locked'] and not node['student_material']
    value.locked.setChecked(False);value.student_material.setChecked(True)
    value.fields['student_task'].setPlainText('观察色块与页边，描述可见差异。')
    value.fields['teacher_answer'].setPlainText('教师答案-ONLY')
    app.processEvents();payload=prep._payload()
    assert payload['lesson_design']['nodes'][0]['student_material']
    output=export_design(win.facade,payload);payload['lesson_design']['exports'].append(output)
    ppt=Presentation(checked_file(win.facade,payload['lesson_design'],output['id'],'lesson_presentation.pptx'))
    pictures=[shape for slide in ppt.slides for shape in slide.shapes if shape.shape_type==13]
    assert len(pictures)==1 and pictures[0].image.blob==win.facade.preparation_image_bytes(asset)
    assert '教师答案-ONLY' not in '\n'.join(shape.text for slide in ppt.slides for shape in slide.shapes if shape.has_text_frame)
    value.close();value.deleteLater();app.processEvents()


@pytest.mark.parametrize('result',['changed-form','failure','cancelled','wrong-image','closed-page'])
def test_image_import_rejects_old_context_and_recovers_page(window,crop_fixture,monkeypatch,result):
    win,_=window;f=crop_fixture;prep,bridge=bind_import(win,f,monkeypatch)
    prep.topic.setText('原课题');before=deepcopy(prep._payload())
    assert prep.import_textbook_image(request(f));task=bridge.take('添加教材备课图表')
    if result=='changed-form':
        prep.topic.setText('另一次备课');bridge.succeed(task,run=True)
    elif result=='failure':bridge.fail(task,'原页发生变化')
    elif result=='cancelled':
        prep._qt_task_cancelled(task.task_id,task.label);bridge.succeed(task,run=True)
    elif result=='closed-page':
        prep.close();assert task.task_id in bridge.cancelled;bridge.succeed(task,run=True)
    else:bridge.succeed(task,{'sha256':'b'*64})
    assert prep.isEnabled() and not prep.studio_busy() and not prep.image_assets_widget.assets()
    if result=='changed-form':assert prep.topic.text()=='另一次备课'
    else:assert prep._payload()==before
