from io import BytesIO
from hashlib import sha256
from copy import deepcopy
from zipfile import ZipFile
import pytest
from PIL import Image, ImageDraw, ImageChops
from docx import Document
from docx.shared import Inches
from integrations.deeptutor_shchem_v1.desktop_raster_numbers import (
    image_catalog, validate_edits, render_numbered_image, apply_to_documents, RasterNumberError,
)


def fixture_documents(fmt='PNG'):
    picture = Image.new('RGB',(640,260),'white')
    draw=ImageDraw.Draw(picture)
    draw.text((15,15),'27. Original raster question',fill='black')
    draw.line((15,150,550,150),fill='black',width=3)
    out=BytesIO(); picture.save(out,format=fmt); image=out.getvalue()
    docs={}
    for role in ('student','teacher'):
        doc=Document();doc.add_paragraph('1. Test chemistry 12.5')
        p=doc.add_paragraph(); p.add_run('Fe').font.bold=True
        p.add_run('3+').font.superscript=True
        doc.add_picture(BytesIO(image),width=Inches(4.5))
        if role=='teacher': doc.add_paragraph('Answer for question 1')
        out=BytesIO();doc.save(out);docs[role+'_bytes']=out.getvalue()
    return docs,image


def edit_for(image,number=1,box=None):
    return {'image_sha256':sha256(image).hexdigest(),'box':box or [10,10,42,37], 'number':number,'punctuation':'.'}


def test_shared_image_index_keeps_role_scope():
    docs,image=fixture_documents()
    cat=image_catalog(docs)
    assert len(cat['images'])==1
    assert cat['images'][0]['audiences']==['student','teacher']
    assert cat['images'][0]['data']==image


@pytest.mark.parametrize('fmt',['PNG','JPEG','BMP'])
def test_only_selected_pixels_change_and_result_is_lossless(fmt):
    _,image=fixture_documents(fmt); e=edit_for(image)
    output=render_numbered_image(image,[e])
    a=Image.open(BytesIO(image)).convert('RGBA'); b=Image.open(BytesIO(output)).convert('RGBA')
    assert output.startswith(b'\x89PNG') and a.size==b.size
    difference=ImageChops.difference(a.convert('RGB'),b.convert('RGB'))
    assert difference.getbbox() is not None
    difference.paste((0,0,0),tuple(e['box']))
    assert difference.getbbox() is None


@pytest.mark.parametrize('fmt',['PNG','JPEG'])
def test_both_copies_change_but_ooxml_formulas_positions_and_original_stay_identical(fmt):
    docs,image=fixture_documents(fmt); before=deepcopy(docs)
    result=apply_to_documents(docs,[edit_for(image)])
    assert docs==before
    shared=[]
    for role in ('student','teacher'):
        source=ZipFile(BytesIO(docs[role+'_bytes'])); new=ZipFile(BytesIO(result[role+'_bytes']))
        assert source.namelist()==new.namelist()
        for name in source.namelist():
            if name.startswith('word/media/'):
                assert new.read(name)!=source.read(name);shared.append(new.read(name))
            elif name!='[Content_Types].xml':
                assert new.read(name)==source.read(name)
        loaded=Document(BytesIO(result[role+'_bytes']))
        assert loaded.inline_shapes[0].width==Inches(4.5)
        assert '12.5' in loaded.paragraphs[0].text
        assert loaded.paragraphs[1].runs[1].font.superscript
    assert shared[0]==shared[1]


def test_teacher_only_image_never_enters_student_file():
    docs,image=fixture_documents()
    teacher=Document(BytesIO(docs['teacher_bytes']))
    out=BytesIO();Image.new('RGB',(150,50),'gray').save(out,format='PNG'); answer=out.getvalue()
    teacher.add_picture(BytesIO(answer));out=BytesIO();teacher.save(out);docs['teacher_bytes']=out.getvalue()
    result=apply_to_documents(docs,[edit_for(answer)])
    assert result['student_bytes']==docs['student_bytes']
    assert len(Document(BytesIO(result['teacher_bytes'])).inline_shapes)==2


@pytest.mark.parametrize('field,value',[
    ('number',0),('number',True),('number',1000),('number','2'),('punctuation','Fe2+'),
    ('box',[-1,0,30,30]),('box',[0,0,9999,30]),('box',[2,2,1,1]),('box',[0,0,1,2]),
    ('box',[0,0,10.2,20]),('image_sha256','f'*64),
])
def test_invalid_or_stale_region_is_not_silently_applied(field,value):
    docs,image=fixture_documents(); edit=edit_for(image);edit[field]=value
    with pytest.raises(RasterNumberError):apply_to_documents(docs,[edit])


def test_overlap_rejected_but_adjacent_regions_allowed():
    docs,image=fixture_documents();cat=image_catalog(docs)['images']
    e=edit_for(image)
    with pytest.raises(RasterNumberError):validate_edits([e,edit_for(image,2,[20,20,50,45])],cat)
    assert len(validate_edits([e,edit_for(image,2,[42,10,70,37])],cat))==2


def test_no_changes_is_exact_original_and_second_number_renders_from_original():
    docs,image=fixture_documents()
    assert apply_to_documents(docs,[])==docs
    assert render_numbered_image(image,[edit_for(image,1)])!=render_numbered_image(image,[edit_for(image,2)])
    assert image_catalog(docs)['images'][0]['data']==image


def test_app_inserted_old_notice_is_not_kept_after_number_changes():
    from integrations.deeptutor_shchem_v1.desktop_paper_numbering import RASTER_NOTICE
    documents, _ = fixture_documents()
    for key in ('student_bytes', 'teacher_bytes'):
        doc = Document(BytesIO(documents[key]))
        doc.add_paragraph(RASTER_NOTICE)
        out = BytesIO(); doc.save(out); documents[key] = out.getvalue()
    catalog = image_catalog(documents)
    key = catalog['images'][0]['image_sha256']
    edit = {'image_sha256':key,'box':[0,0,40,30],'number':2,'punctuation':'.'}
    result = apply_to_documents(documents,[edit])
    for audience in ('student_bytes','teacher_bytes'):
        text = '\n'.join(p.text for p in Document(BytesIO(result[audience])).paragraphs)
        assert RASTER_NOTICE not in text
        assert '未框选的旧号' in text
