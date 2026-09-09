from __future__ import annotations

import copy
import hashlib
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Twips
from jsonschema import Draft202012Validator

from integrations.deeptutor_shchem_v1.desktop_preparation import (
    DesktopPreparationError,
    DesktopPreparationManager,
    _validate_canonical_candidate,
    normalize_preparation_candidate,
    preparation_candidate_schema,
)
from integrations.deeptutor_shchem_v1.desktop_preparation_renderer import (
    NativePreparationRenderer,
)
from integrations.deeptutor_shchem_v1.desktop_preparation_worksheet import (
    build_worksheet_document,
    normalize_worksheet,
)
from staging.coordination.deeptutor_gateway.scripts.verify_preparation_worksheets import (
    worksheet_sample,
)
from staging.coordination.deeptutor_gateway.tests.test_desktop_preparation import (
    RecordingProvider,
    _payload,
    _raw_candidate,
)


def test_optional_worksheet_preserves_legacy_candidate_and_immutable_input() -> None:
    raw = _raw_candidate()
    before = copy.deepcopy(raw)
    legacy = normalize_preparation_candidate(raw, _payload())
    assert all("worksheet" not in activity for activity in legacy["activities"])
    assert raw == before
    for activity in raw["activities"]:
        activity["worksheet"] = None
    assert normalize_preparation_candidate(raw, _payload()) == legacy
    assert _validate_canonical_candidate(legacy, _payload()) == legacy
    assert normalize_worksheet(None) is None


def test_worksheet_print_uses_a4_and_table_width_matches_usable_page_width() -> None:
    raw = _raw_candidate()
    raw["activities"][0]["worksheet"] = worksheet_sample()
    document = build_worksheet_document(raw)
    section = document.sections[0]
    assert section.page_width == Twips(11906)  # A4 width, quantized to twips
    assert section.page_height == Twips(16838)  # A4 height, quantized to twips
    usable_width = section.page_width - section.left_margin - section.right_margin
    assert usable_width > 0
    assert document.tables
    for table in document.tables:
        grid_widths = [int(column.width) for column in table.columns]
        assert sum(grid_widths) == usable_width
        for row in table.rows:
            cell_widths = [int(cell.width) for cell in row.cells]
            assert sum(cell_widths) == usable_width
            assert cell_widths == grid_widths


def test_multiple_worksheets_break_before_title_without_standalone_page_break() -> None:
    raw = _raw_candidate()
    first = worksheet_sample()
    second = copy.deepcopy(first)
    second["title"] = "第二份学习单"
    raw["activities"][0]["worksheet"] = first
    raw["activities"][1]["worksheet"] = second

    document = build_worksheet_document(raw)
    title_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text in {first["title"], second["title"]}
    ]
    assert len(title_paragraphs) == 2
    assert title_paragraphs[0].paragraph_format.page_break_before is None
    assert title_paragraphs[1].paragraph_format.page_break_before is True
    assert document._element.body.xpath('.//w:br[@w:type="page"]') == []


def test_current_schema_and_worksheet_normalization() -> None:
    raw = _raw_candidate()
    for activity in raw["activities"]:
        activity["worksheet"] = None
    for slide in raw["slides"]:
        slide["visual"] = None
        slide["image"] = None
    sheet = worksheet_sample()
    sheet["title"] = "  整理记录  "
    raw["activities"][0]["worksheet"] = sheet
    before = copy.deepcopy(raw)
    Draft202012Validator(preparation_candidate_schema()).validate(raw)
    result = normalize_preparation_candidate(raw, _payload())
    assert result["activities"][0]["worksheet"]["title"] == "整理记录"
    assert result["activities"][0]["id"] == "A01"
    assert _validate_canonical_candidate(result, _payload()) == result
    assert before == raw


@pytest.mark.parametrize(
    "field,value",
    [
        ("response_kind", "image"),
        ("response_lines", True),
        ("response_lines", 2),
        ("columns", ["一列"]),
        ("columns", ["1", "2", "3", "4", "5"]),
        ("row_labels", []),
        ("row_labels", ["x"] * 7),
        ("row_labels", [""]),
        ("prompt", ""),
        ("prompt", "x" * 501),
        ("heading", "x" * 81),
        ("answer", "禁止额外答案字段"),
    ],
)
def test_table_worksheet_rejects_bad_fields(field: str, value: object) -> None:
    sheet = worksheet_sample()
    sheet["sections"][0][field] = value
    with pytest.raises(ValueError):
        normalize_worksheet(sheet)


@pytest.mark.parametrize(
    "field,value",
    [
        ("response_lines", 0),
        ("response_lines", 11),
        ("response_lines", 3.0),
        ("columns", ["不能混用"]),
        ("row_labels", ["不能混用"]),
    ],
)
def test_writing_worksheet_rejects_bad_space(field: str, value: object) -> None:
    sheet = worksheet_sample()
    sheet["sections"][1][field] = value
    with pytest.raises(ValueError):
        normalize_worksheet(sheet)


@pytest.mark.parametrize(
    "sheet", [[], {}, "bad", {"title": "empty", "instructions": [], "sections": []}]
)
def test_worksheet_rejects_empty_or_nonobject(sheet: object) -> None:
    with pytest.raises(ValueError):
        normalize_worksheet(sheet)


@pytest.mark.parametrize("output_kind", ["ppt", "lesson_plan", "joint"])
def test_manager_registers_real_worksheet_and_keeps_teacher_text_out(
    tmp_path: Path, output_kind: str
) -> None:
    raw = _raw_candidate()
    raw["activities"][0]["worksheet"] = worksheet_sample()
    sentinel = "仅教师可见的评分提示标记"
    raw["activities"][0]["teacher_action"] = sentinel
    provider = RecordingProvider(raw)
    manager = DesktopPreparationManager(tmp_path / "tasks", NativePreparationRenderer())
    prepared = manager.prepare(
        _payload(output_kind=output_kind), "OFFLINE-PROFILE", "REV-1"
    )
    completed = manager.run(prepared["task_id"], provider, None, lambda: False)
    assert completed["status"] == "completed"
    assert len(provider.calls) == 1
    path, content_type = manager.artifact_path(
        prepared["task_id"], "student_worksheet_docx"
    )
    assert content_type.endswith("wordprocessingml.document")
    receipt = next(
        item
        for item in completed["artifacts"]
        if item["artifact_id"] == "student_worksheet_docx"
    )
    assert receipt["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    document = Document(path)
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert worksheet_sample()["title"] in texts
    assert sentinel not in "\n".join(texts)
    assert document.paragraphs[0].style.name == "Title"
    assert str(document.styles["Title"].font.color.rgb) == "000000"
    assert len(document.tables) == 1
    table = document.tables[0]
    assert (
        table.rows[0]._tr.find("./" + qn("w:trPr") + "/" + qn("w:tblHeader"))
        is not None
    )
    for row, label in zip(
        table.rows[1:], worksheet_sample()["sections"][0]["row_labels"], strict=True
    ):
        assert row.cells[0].text == label
        assert all(not cell.text for cell in row.cells[1:])
        height = row._tr.find("./" + qn("w:trPr") + "/" + qn("w:trHeight"))
        assert height is not None and height.get(qn("w:hRule")) == "atLeast"
    if output_kind != "ppt":
        lesson, _ = manager.artifact_path(prepared["task_id"], "lesson_plan_docx")
        lesson_document = Document(lesson)
        assert "配套学习单：" in "\n".join(p.text for p in lesson_document.paragraphs)
        section = lesson_document.sections[0]
        usable_width = section.page_width - section.left_margin - section.right_margin
        for lesson_table in lesson_document.tables:
            grid_widths = [column.width for column in lesson_table.columns]
            assert sum(grid_widths) <= usable_width
            # Word uses the header's preferred cell widths as well as tblGrid.
            # A default equal-width header must not contradict the body grid.
            for row in lesson_table.rows:
                assert [cell.width for cell in row.cells] == grid_widths


def test_bad_worksheet_fails_before_provider_candidate_can_be_saved() -> None:
    raw = _raw_candidate()
    raw["activities"][0]["worksheet"] = worksheet_sample()
    raw["activities"][0]["worksheet"]["sections"][0]["row_labels"] = []
    with pytest.raises(DesktopPreparationError) as caught:
        normalize_preparation_candidate(raw, _payload())
    assert caught.value.code == "preparation_candidate_worksheet_invalid"
