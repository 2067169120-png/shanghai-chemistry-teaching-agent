"""Explicit batch-local labels use synthetic DOCX and temporary personal state."""

import hashlib
import io
import json
import sqlite3
from copy import deepcopy
from types import SimpleNamespace

import pytest
from docx import Document

from integrations.deeptutor_shchem_v1 import desktop_word_questions as module
from integrations.deeptutor_shchem_v1.desktop_facade import DesktopFacadeError
from integrations.deeptutor_shchem_v1.desktop_state import DesktopStateStore
from integrations.deeptutor_shchem_v1.desktop_word_questions import (
    RANGES_DRAFT,
    WordQuestionError,
    WordQuestionService,
)


def _docx(*texts):
    document = Document()
    document.add_heading("合成课堂练习", 1)
    for number, text in enumerate(texts, 1):
        document.add_paragraph(f"【即学即练{number}】{text}")
        document.add_paragraph("本题补充条件。")
        document.add_paragraph("【答案】合成参考答案。")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


class Facade:
    """Only source restoration and a temp state exist; provider access fails."""

    def __init__(self, root):
        self.paths = SimpleNamespace(
            workspace_root=root / "workspace", state_root=root / "state"
        )
        self.state_store = DesktopStateStore(self.paths.state_root)
        self.batches = {}
        self.reads = []

    def add(self, batch_id, content, name="合成讲义.docx", source_id="SOURCE1"):
        path = self.paths.workspace_root / batch_id / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)
        row = {
            "source_id": source_id,
            "source_name": name,
            "sha256": hashlib.sha256(content).hexdigest(),
            "path": path,
        }
        self.batches.setdefault(batch_id, []).append(row)
        return path

    def _saved_visual_import_batch(self, batch_id):
        if batch_id not in self.batches:
            raise DesktopFacadeError("missing", "未找到已保存批次。")
        return {"batch_id": batch_id}

    def imported_word_sources(self, batch_id):
        self._saved_visual_import_batch(batch_id)
        return deepcopy(self.batches[batch_id])

    def _restore_visual_import_sources(self, descriptor, *, source_ids=None):
        batch_id = descriptor["batch_id"]
        result = []
        for row in self.batches[batch_id]:
            if source_ids is not None and row["source_id"] not in source_ids:
                continue
            self.reads.append((batch_id, row["source_id"]))
            raw = row["path"].read_bytes()
            if hashlib.sha256(raw).hexdigest() != row["sha256"]:
                raise DesktopFacadeError("source_changed", "原文件归档已变化。")
            result.append(
                SimpleNamespace(
                    content=raw,
                    filename=row["source_name"],
                    source_sha256=row["sha256"],
                    effective_source_file_id=row["source_id"],
                )
            )
        return result

    def list_imported_word_batches(self):
        return tuple(SimpleNamespace(batch_id=key) for key in self.batches)


@pytest.fixture
def setup(tmp_path, monkeypatch):
    facade = Facade(tmp_path)
    taxonomy = facade.paths.workspace_root / "sh-chem-db/kb/knowledge_taxonomy.json"
    directory = (
        facade.paths.workspace_root
        / "sh-chem-db/kb/classification/supplemental_wechat_textbook_tagging_v1_2026-08-27/textbook_directory_nodes.json"
    )
    taxonomy.parent.mkdir(parents=True)
    directory.parent.mkdir(parents=True)
    taxonomy.write_text(
        json.dumps(
            {"dimensions": {"knowledge_points": [{"id": "K11", "name": "氧化还原"}]}}
        ),
        encoding="utf-8",
    )
    directory.write_text(
        json.dumps(
            {
                "nodes": [
                    {
                        "node_key": "S1",
                        "chapter_id": "C1",
                        "volume_id": "V1",
                        "section_title": "氧化还原反应",
                        "volume_title": "合成教材",
                        "chapter_title": "合成章节",
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(module, "source_quality_notes", lambda _root: {})
    source = facade.add(
        "B1", _docx("（2025·上海·高三二模）解释氧化还原反应。", "说明本题现象。")
    )
    service = WordQuestionService(facade)
    return service, facade, source, directory


def _rows(service):
    return service._annotation_snapshot("B1", {})[0]


def _range(service, row):
    service.state.save_draft(
        RANGES_DRAFT,
        {
            "ranges": {
                row["key"]: {
                    "source_revision": row["source_revision"],
                    "range": {
                        "block_start": row["block_start"],
                        "question_end": row["question_blocks"][0]["index"],
                        "answer_start": None,
                        "block_end": row["question_blocks"][0]["index"],
                    },
                }
            }
        },
    )


def test_explicit_batch_annotates_current_rows_and_keeps_unknown(setup):
    service, facade, source, _ = setup
    before = source.read_bytes()
    result = service.annotate_imported_batch("B1")
    assert result["method"] == "local_rules" and result["model_invoked"] is False
    assert result["batch_count"] == result["source_count"] == 1
    assert (
        result["question_count"]
        == result["saved_question_count"]
        == result["changed_question_count"]
        == 2
    )
    assert result["valid_label_counts"] == {
        "source_bound_questions": 2,
        "primary_knowledge": 1,
        "curriculum_mapped": 1,
        "applicable_grade": 0,
        "original_exam_type": 1,
    }
    rows = _rows(service)
    saved = service.attribute_store.get_many([row["key"] for row in rows])
    assert saved[rows[1]["key"]]["primary_knowledge"]["id"] == "unknown"
    assert saved[rows[1]["key"]]["original_source"]["exam_type"]["value"] == "unknown"
    assert source.read_bytes() == before
    assert not service.preview_cache.root.exists()
    assert not facade.state_store.path.exists()


def test_retry_is_idempotent_and_warm_cache_unchanged(setup):
    service, _, source, _ = setup
    preview = service.reader.word_preview_bytes(source.read_bytes(), source.name)
    service.preview_cache.save(source.read_bytes(), source.name, preview)
    frozen = {p.name: p.read_bytes() for p in service.preview_cache.root.iterdir()}
    service.annotate_imported_batch("B1")
    rows = _rows(service)
    history = {row["key"]: service.attribute_store.history(row["key"]) for row in rows}
    result = service.annotate_imported_batch("B1")
    assert result["changed_question_count"] == 0
    assert history == {
        row["key"]: service.attribute_store.history(row["key"]) for row in rows
    }
    assert frozen == {
        p.name: p.read_bytes() for p in service.preview_cache.root.iterdir()
    }


@pytest.mark.parametrize("stale", [False, True])
def test_teacher_modified_rows_and_history_are_preserved(setup, stale):
    service, _, _, _ = setup
    service.annotate_imported_batch("B1")
    row = _rows(service)[0]
    original = service.attribute_store.get(row["key"])
    edited = service.attribute_store.save_teacher_edit(
        row["key"],
        {"teacher_note": "教师修改须保留"},
        expected_revision=original["revision"],
    )
    history = service.attribute_store.history(row["key"])
    if stale:
        _range(service, row)
    result = service.annotate_imported_batch("B1")
    assert result["preserved_teacher_count"] == 1
    assert result["stale_teacher_count"] == int(stale)
    assert result["valid_label_counts"]["source_bound_questions"] == 2 - int(stale)
    assert result["valid_label_counts"]["primary_knowledge"] == 1 - int(stale)
    assert service.attribute_store.get(row["key"]) == edited
    assert service.attribute_store.history(row["key"]) == history
    if stale:
        assert any("旧范围" in warning for warning in result["warnings"])


def test_existing_source_metadata_survives_auto_refresh(setup):
    service, _, _, _ = setup
    row = _rows(service)[0]
    catalog = module.load_attribute_catalog(service.facade.paths.workspace_root)
    original = module.suggest_attributes(
        row,
        {
            "source_name": row["source_name"],
            "usage_context": "高三课堂",
            "package_id": "SYNTHETIC",
        },
        catalog,
    )
    service.attribute_store.save_many([original])
    _range(service, row)
    result = service.annotate_imported_batch("B1")
    saved = service.attribute_store.get(row["key"])
    assert saved["source"]["package_id"] == "SYNTHETIC"
    assert saved["question_revision"] == _rows(service)[0]["revision"]
    assert result["valid_label_counts"]["applicable_grade"] == 1


@pytest.mark.parametrize("change", ["source", "range", "directory"])
def test_changes_during_suggestion_fail_before_attribute_transaction(
    setup, monkeypatch, change
):
    service, _, source, directory = setup
    row = _rows(service)[0]
    original = module.suggest_attributes
    changed = False

    def mutate(*args, **kwargs):
        nonlocal changed
        result = original(*args, **kwargs)
        if not changed:
            changed = True
            if change == "source":
                source.write_bytes(_docx("新版本合成题。"))
            elif change == "range":
                _range(service, row)
            else:
                data = json.loads(directory.read_text(encoding="utf-8"))
                data["nodes"][0]["chapter_title"] = "目录新修订"
                directory.write_text(json.dumps(data), encoding="utf-8")
        return result

    monkeypatch.setattr(module, "suggest_attributes", mutate)
    with pytest.raises(WordQuestionError, match="未提交"):
        service.annotate_imported_batch("B1")
    assert service.attribute_store.get(row["key"]) is None
    assert service.attribute_store.history(row["key"]) == []


@pytest.mark.parametrize("failure", ["missing", "invalid", "range"])
def test_failed_catalog_or_stale_range_cannot_save_labels(setup, failure):
    service, _, source, directory = setup
    row = _rows(service)[0]
    before = source.read_bytes()
    if failure == "missing":
        directory.unlink()
    elif failure == "invalid":
        directory.write_text("not json", encoding="utf-8")
    else:
        service.state.save_draft(
            RANGES_DRAFT,
            {"ranges": {row["key"]: {"source_revision": "stale", "range": {}}}},
        )
    with pytest.raises(WordQuestionError, match="未提交"):
        service.annotate_imported_batch("B1")
    assert service.attribute_store.get(row["key"]) is None
    assert source.read_bytes() == before


def test_no_questions_is_explicit_without_creating_attribute_store(setup):
    service, facade, _, _ = setup
    facade.add("EMPTY", _docx())
    result = service.annotate_imported_batch("EMPTY")
    assert result["status"] == "no_questions"
    assert result["source_count"] == 1
    assert result["question_count"] == result["saved_question_count"] == 0
    assert not facade.paths.state_root.exists()


def test_batch_isolation_and_unknown_batch(setup):
    service, facade, _, _ = setup
    other = facade.add("B2", _docx("另一个批次的烷烃题。"))
    other.write_bytes(b"unreadable unrelated source")
    service.annotate_imported_batch("B1")
    assert {batch for batch, _ in facade.reads} == {"B1"}
    with pytest.raises(WordQuestionError, match="未找到"):
        service.annotate_imported_batch("MISSING")
    with pytest.raises(WordQuestionError, match="请选择"):
        service.annotate_imported_batch("")


def test_public_catalog_supplies_read_only_directory_and_survives_missing_file(setup):
    service, _, _, directory = setup
    result = service.catalog()
    assert len(result["items"]) == 2
    assert result["attribute_catalog"]["nodes"][0]["chapter_title"] == "合成章节"
    assert service.attribute_store.get(result["items"][0]["key"]) is None
    directory.unlink()
    result = service.catalog()
    assert len(result["items"]) == 2 and result["attribute_catalog"] is None
    assert any("教材筛选不可用" in warning for warning in result["warnings"])


def test_partial_sql_write_rolls_back_the_whole_batch(setup, monkeypatch):
    service, _, source, _ = setup
    rows = _rows(service)
    frozen = source.read_bytes()
    original = service.attribute_store._write
    calls = []

    def fail_second(connection, row):
        calls.append(row["key"])
        if len(calls) == 2:
            raise sqlite3.OperationalError("synthetic disk failure")
        return original(connection, row)

    monkeypatch.setattr(service.attribute_store, "_write", fail_second)
    with pytest.raises(WordQuestionError, match="未提交"):
        service.annotate_imported_batch("B1")
    assert len(calls) == 2
    assert service.attribute_store.get_many([row["key"] for row in rows]) == {}
    assert all(service.attribute_store.history(row["key"]) == [] for row in rows)
    assert source.read_bytes() == frozen


def test_multi_source_counts_deduplicate_exact_source_bytes(setup):
    service, facade, source, _ = setup
    facade.add("B1", source.read_bytes(), "同一来源副本.docx", "COPY")
    facade.add("B1", _docx("分析另一份合成资料。"), "不同来源.docx", "SOURCE2")
    result = service.annotate_imported_batch("B1")
    assert result["source_count"] == 2
    assert result["question_count"] == result["saved_question_count"] == 3
    assert result["valid_label_counts"]["source_bound_questions"] == 3


def test_malformed_directory_preserves_browser_but_stops_annotation(setup):
    service, _, _, directory = setup
    directory.write_text(json.dumps({"nodes": "invalid"}), encoding="utf-8")
    result = service.catalog()
    assert len(result["items"]) == 2 and result["attribute_catalog"] is None
    with pytest.raises(WordQuestionError, match="未提交"):
        service.annotate_imported_batch("B1")
    assert service.attribute_store.get(result["items"][0]["key"]) is None
