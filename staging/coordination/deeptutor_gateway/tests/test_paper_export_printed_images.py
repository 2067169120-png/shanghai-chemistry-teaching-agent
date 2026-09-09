from copy import deepcopy

import pytest
from docx import Document
from PIL import Image

from integrations.deeptutor_shchem_v1 import paper_export_renderer as renderer


def test_projection_preserves_text_distinct_assets_and_same_part_repeats():
    image = {"block_type": "image", "asset_ref": "q.png"}
    text = {"block_type": "paragraph", "text_zh": "仍须回答的文字"}
    other = {"block_type": "apparatus", "asset_ref": "other.png"}
    parts = [
        {"question_blocks": [image, image]},
        {"question_blocks": [image, text, other]},
    ]
    original = deepcopy(parts)
    blocks, suppressed = renderer._printed_question_blocks(parts)
    assert suppressed
    assert blocks == [[image, image], [text, other]]
    assert parts == original


@pytest.mark.parametrize("audience", ["student", "teacher"])
def test_one_image_per_printed_preserves_all_atomic_scores_and_answers(
    tmp_path, monkeypatch, audience
):
    bundle = renderer.build_synthetic_demo_bundle()
    section = deepcopy(bundle["teacher_plan"]["visible"]["theme_sections"][0])
    section["shared_materials"] = []
    template = section["printed_questions"][0]["atomic_parts"][0]
    printed = []
    for question_number in (1, 2):
        parts = []
        for index in (1, 2):
            part = deepcopy(template)
            part["question_blocks"] = [
                {"block_type": "image", "asset_ref": "q.png", "alt_text_zh": "测试题图"}
            ]
            part["score"] = index
            part["answer_space"]["lines"] = index
            part["teacher_notes"]["explanation_zh"] = f"解析{question_number}-{index}"
            parts.append(part)
        printed.append({"question_number": question_number, "atomic_parts": parts})
    section["printed_questions"] = printed
    visible = {"theme_sections": [section]}
    original = deepcopy(visible)
    Image.new("RGB", (300, 80), "white").save(tmp_path / "q.png")
    doc = Document()
    renderer._style_document(doc, bundle["preset"])
    answer_lines = []
    add_answer_space = renderer._add_answer_space

    def capture_space(document, lines, **kwargs):
        answer_lines.append(lines)
        add_answer_space(document, lines, **kwargs)

    monkeypatch.setattr(renderer, "_add_answer_space", capture_space)
    renderer._add_theme_sections(
        doc, visible, audience=audience, asset_root=tmp_path,
        content_width_dxa=9000, body_size_pt=11,
    )
    assert len(doc.inline_shapes) == 2  # No dedup across printed-question boundaries.
    paragraphs = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    notes_text = "\n".join(
        paragraph.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    assert paragraphs.count("作答单元1 （1 分）") == 2
    assert paragraphs.count("作答单元2 （2 分）") == 2
    assert answer_lines == ([1, 2, 1, 2] if audience == "student" else [0] * 4)
    for q in (1, 2):
        for unit in (1, 2):
            assert (f"解析{q}-{unit}" in notes_text) is (audience == "teacher")
    assert visible == original
