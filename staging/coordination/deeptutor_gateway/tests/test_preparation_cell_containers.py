"""Word cell containers must preserve ordered native text or explicit gaps."""

from __future__ import annotations

import pytest
from docx import Document
from docx.oxml import OxmlElement

from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
    PreparationSourcesService,
    _block_text,
    _table_rows,
)


def _paragraph(text):
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    paragraph.append(run)
    return paragraph


def _control(*children):
    wrapper = OxmlElement("w:sdt")
    wrapper.append(OxmlElement("w:sdtPr"))
    content = OxmlElement("w:sdtContent")
    for child in children:
        content.append(child)
    wrapper.append(content)
    wrapper.append(OxmlElement("w:sdtEndPr"))
    return wrapper


def _table():
    document = Document()
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)._tc
    for paragraph in list(cell.p_lst):
        cell.remove(paragraph)
    return document, table, cell


def _cell_text(document, table):
    warnings = set()
    rows = _table_rows(table._tbl, document, warnings)
    return rows[0]["cells"][0]["text"], warnings


def test_cell_controls_preserve_summary_example_answer_order_once(tmp_path):
    document, table, cell = _table()
    cell.append(_paragraph("导语"))
    cell.append(_control(_paragraph("知识总结：保留条件")))
    custom = OxmlElement("w:customXml")
    custom.append(OxmlElement("w:customXmlPr"))
    custom.append(_control(_paragraph("例题：完整题面"), _paragraph("答案：B")))
    cell.append(custom)
    cell.append(_paragraph("变式训练"))
    expected = "导语\n知识总结：保留条件\n例题：完整题面\n答案：B\n变式训练"
    text, warnings = _cell_text(document, table)
    assert text == expected
    assert warnings == set()
    source = tmp_path / "cell-controlled-handout.docx"
    document.save(source)
    before = source.read_bytes()
    service = PreparationSourcesService(tmp_path)
    preview = service.word_preview(source)
    assert len(preview["blocks"]) == 1
    reference = service.reference(source, preview["source_sha256"], 1, 1, [])
    for line in expected.splitlines():
        assert reference["materials"].count(line) == 1
    assert reference["warnings"] == []
    assert source.read_bytes() == before


def test_nested_table_in_control_retains_structure_and_cell_order():
    document, table, cell = _table()
    nested_document = Document()
    nested = nested_document.add_table(rows=2, cols=2)
    nested.cell(0, 0).merge(nested.cell(0, 1)).text = "比较维度"
    nested.cell(1, 0).text = "判断条件"
    nested.cell(1, 1).text = "对应例子"
    cell.append(_control(_paragraph("表前"), nested._tbl, _paragraph("表后")))
    text, warnings = _cell_text(document, table)
    assert text == (
        "表前\n【表格开始：按原行列顺序；未标合并即未合并】\n"
        "〔第1行·第1—2列；横向合并〕\n比较维度\n"
        "〔第2行·第1列〕\n判断条件\n"
        "〔第2行·第2列〕\n对应例子\n【表格结束】\n表后"
    )
    outer = _block_text(table._tbl, document, set())
    assert outer.count("【表格开始：") == outer.count("【表格结束】") == 2
    assert text in outer
    assert warnings == set()


def test_unknown_cell_container_stays_at_original_position_as_explicit_gap():
    document, table, cell = _table()
    unknown = OxmlElement("w:smartTag")
    unknown.append(_paragraph("未经支持的容器内容"))
    cell.append(_control(_paragraph("之前"), unknown, _paragraph("之后")))
    text, warnings = _cell_text(document, table)
    assert text == "之前\n【待查看原文：特殊正文容器】\n之后"
    assert warnings == {"特殊正文容器未提取；需查看原文件。"}


@pytest.mark.parametrize(
    "tag,label", [("m:oMath", "数学公式"), ("w:object", "嵌入对象或旧公式")]
)
def test_wrapped_formula_remains_a_gap_instead_of_guessed_text(tag, label):
    document, table, cell = _table()
    paragraph = _paragraph("例式：")
    unknown = OxmlElement(tag)
    unknown.append(_paragraph("不得当作已识别公式"))
    paragraph.append(unknown)
    cell.append(_control(paragraph, _paragraph("分析条件")))
    text, warnings = _cell_text(document, table)
    assert text == f"例式：【待查看原文：{label}】\n分析条件"
    assert "不得当作已识别公式" not in text
    assert f"{label}未提取；需查看原文件，不得根据残句补写。" in warnings


def test_control_properties_and_bookmarks_do_not_create_fake_body_or_warnings():
    document, table, cell = _table()
    control = _control(_paragraph("唯一正文"))
    properties = control.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sdtPr"
    )
    properties.append(_paragraph("属性内不是正文"))
    cell.append(OxmlElement("w:bookmarkStart"))
    cell.append(control)
    cell.append(OxmlElement("w:bookmarkEnd"))
    cell.append(OxmlElement("w:proofErr"))
    text, warnings = _cell_text(document, table)
    assert text == "唯一正文"
    assert warnings == set()
