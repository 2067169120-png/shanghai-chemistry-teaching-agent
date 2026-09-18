from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from hashlib import sha256
import json
import pytest
from integrations.deeptutor_shchem_v1.desktop_lesson_design import *
from integrations.deeptutor_shchem_v1.desktop_lesson_output import export_design, checked_file, write_documents
from integrations.deeptutor_shchem_v1.desktop_editor_recovery import PreparationRecoveryStore
from integrations.deeptutor_shchem_v1.desktop_preparation import normalize_preparation_payload
from integrations.deeptutor_shchem_v1.desktop_preparation_drafts import PreparationDraftService


def sample():
    p = {"output_kind":"joint", "topic":"动态平衡的证据与表达", "audience":"高二·合成验收",
         "lesson_route":"复习", "lesson_timing":"1课时×40分钟",
         "objective":"根据速率证据判断平衡\n写出平衡常数表达式",
         "materials":"合成流程测试资料，不是真实试卷",
         "advanced":dict.fromkeys(("learning_and_experiment", "template_and_delivery", "homework_and_strategy"),"")}
    d = new_design(p)
    n = new_node("课前诊断")
    n.update(title="依据速率判断", objective_ids=[d["objectives"][0]["id"]],
             student_task="根据正逆速率说明判断依据。", expected_output="判断及证据。",
             criteria="判断与条件一致。", teacher_action="收集解释，追问证据。",
             teacher_answer="仅教师答案-SECRET", material_text="合成反应A⇌B；12.5 mol·L⁻¹。", student_material=True,
             minutes=10, confirmed=True)
    d["nodes"].append(n); p["lesson_design"]=d
    return p


def test_second_goal_is_not_covered_by_first_goal_task():
    p=sample(); d=p["lesson_design"]
    assert [g["status"] for g in coverage(d)["objectives"]] == ["已明确关联","未安排"]


def test_adding_specific_task_only_changes_its_goal():
    d=sample()["lesson_design"]
    n=new_node(); n.update(objective_ids=[d["objectives"][1]["id"]],
        student_task="写表达式",expected_output="表达式",criteria="结构符合反应式",confirmed=True)
    d["nodes"].append(n)
    assert [g["status"] for g in coverage(d)["objectives"]] == ["已明确关联","已明确关联"]


def test_link_without_evidence_stays_pending():
    d=sample()["lesson_design"]; d["nodes"][0]["criteria"]=""
    assert coverage(d)["objectives"][0]["status"]=="待核对"


def test_goal_edit_invalidates_only_its_links():
    d=sample()["lesson_design"]
    out=update_goal(d,d["objectives"][0]["id"],"新目标")
    assert not out["nodes"][0]["confirmed"] and d["nodes"][0]["confirmed"]


def test_task_edit_clears_previous_confirmation():
    d=sample()["lesson_design"]
    out=update_node(d,d["nodes"][0]["id"],{"student_task":"改后的任务","confirmed":True})
    assert not out["nodes"][0]["confirmed"]


def test_notes_and_minutes_do_not_change_confirmed_links():
    d=sample()["lesson_design"];n=d["nodes"][0]
    out=update_node(d,n["id"],{"notes":"仅教师备注","minutes":12})
    assert out["nodes"][0]["confirmed"]


@pytest.mark.parametrize("field,value",[("material_text","changed"),("teacher_answer","changed"),
    ("image_ids",["IMG-"+"a"*64]),("student_material",False)])
def test_lock_cannot_be_bypassed_in_same_patch(field,value):
    d=sample()["lesson_design"];n=d["nodes"][0];n["locked"]=True
    with pytest.raises(ValueError): update_node(d,n["id"],{field:value,"locked":False})


def test_unlocked_change_and_revision_check():
    d=sample()["lesson_design"];n=d["nodes"][0]
    with pytest.raises(ValueError):update_node(d,n["id"],{"notes":"x"},expected="old")
    out=update_node(d,n["id"],{"teacher_action":"new"},expected=digest(d))
    assert out["nodes"][0]["teacher_answer"]==n["teacher_answer"]


def test_material_binding_defaults_to_teacher_only_and_keeps_exact_text():
    d=sample()["lesson_design"];n=d["nodes"][0]
    text="完整材料\n\n原题① Fe³⁺⇌Fe²⁺；12.5\n答案：不能给学生"
    out=bind_material(d,n["id"],text)
    assert out["nodes"][0]["material_text"]==text
    assert out["nodes"][0]["locked"] and not out["nodes"][0]["student_material"]


def test_unknown_time_is_not_zero_or_allocated_from_total():
    d=sample()["lesson_design"];d["nodes"][0]["minutes"]=None
    c=coverage(d,40)
    assert c["known_minutes"]==0 and c["unestimated"]==1
    d["nodes"][0]["minutes"]=80;assert coverage(d,40)["over_budget"]


def test_undo_and_redo_preserve_node_identity():
    d=sample()["lesson_design"];h=DesignHistory(d)
    v=deepcopy(d);v["nodes"].append(new_node());h.put(v)
    assert h.travel()==d
    assert h.travel(True)==v


def test_incomplete_design_roundtrips_in_original_recovery(tmp_path):
    p=sample();p["lesson_design"]["nodes"][0]["student_task"]=""
    store=PreparationRecoveryStore(tmp_path)
    store.save(p,dirty=True,sequence=0)
    assert PreparationRecoveryStore(tmp_path).load()["payload"]==p


def test_original_model_request_contract_unchanged():
    p=sample();old=deepcopy(p);old.pop("lesson_design")
    assert normalize_preparation_payload(p)==normalize_preparation_payload(old)


def test_saved_draft_reader_retains_design():
    p=sample()
    row={"kind":"preparation","contract_version":"lesson-blueprint/2.0.0","status":"draft",
         "output_kind":"joint","core_fields":{k:p[k] for k in ("topic","audience","lesson_route","lesson_timing","objective","materials")},
         "advanced":p["advanced"],"created_at":"2026-09-18T01:00:00Z","lesson_design":p["lesson_design"]}
    assert PreparationDraftService._payload(row)["lesson_design"]==p["lesson_design"]


def test_future_schema_not_silently_discarded():
    p=sample();p["lesson_design"]["schema_version"]="future"
    with pytest.raises(Exception):normalize_preparation_payload(p)


def test_exports_share_nodes_without_student_answer_or_extra_lines(tmp_path):
    from docx import Document
    from pptx import Presentation
    p=sample();original=deepcopy(p)
    f=SimpleNamespace(paths=SimpleNamespace(task_root=tmp_path))
    output=export_design(f,p)
    p["lesson_design"]["exports"].append(output)
    assert content_fingerprint(p)==output["fingerprint"]
    doc=Document(checked_file(f,p["lesson_design"],output["id"],"lesson_plan.docx"))
    student=Document(checked_file(f,p["lesson_design"],output["id"],"student_worksheet.docx"))
    ppt=Presentation(checked_file(f,p["lesson_design"],output["id"],"lesson_presentation.pptx"))
    teacher_text="\n".join(x.text for x in doc.paragraphs)
    student_text="\n".join(x.text for x in student.paragraphs)
    screen="\n".join(s.text for sl in ppt.slides for s in sl.shapes if s.has_text_frame)
    notes="\n".join(sl.notes_slide.notes_text_frame.text for sl in ppt.slides)
    assert "仅教师答案-SECRET" in teacher_text and "仅教师答案-SECRET" in notes
    assert "仅教师答案-SECRET" not in screen and "仅教师答案-SECRET" not in student_text
    assert original["lesson_design"]["nodes"][0]["student_task"] in student_text
    assert "___" not in student_text and not student.tables
    p["lesson_design"]["nodes"][0]["student_task"]="new task"
    assert content_fingerprint(p)!=output["fingerprint"]
    # Old output remains intact and accessible as its old version.
    assert checked_file(f,p["lesson_design"],output["id"],"lesson_plan.docx").is_file()
    again=export_design(f,p)
    assert again["id"] != output["id"]


def test_missing_picture_and_overlong_student_text_have_specific_error():
    p=sample();n=p["lesson_design"]["nodes"][0];n["image_ids"]=["IMG-"+"b"*64]
    with pytest.raises(ValueError,match="图片"):candidate_from_design(p)
    n["image_ids"]=[];n["material_text"]="长"*551
    with pytest.raises(ValueError,match="过长"):candidate_from_design(p)


def test_explicit_teacher_only_source_not_in_student_outputs(tmp_path):
    from docx import Document
    p=sample();p["lesson_design"]["nodes"][0]["student_material"]=False
    write_documents(p,tmp_path,{})
    text="\n".join(x.text for x in Document(tmp_path/"student_worksheet.docx").paragraphs)
    assert "12.5" not in text
    assert "12.5" not in str([s["content"] for s in candidate_from_design(p)["slides"]])
