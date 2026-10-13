"""Static EQ evidence survives whole-section import, native layout and Office export."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import gzip
import json
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Mm, Pt
from PIL import Image, ImageDraw
from pptx import Presentation
import pytest

from test_composition_render import app as app
from test_handout_knowledge import make_facade, request_for, word_bytes
from test_lesson_design import sample
from test_static_eq_boundaries import _field, _run, A, R, W
from integrations.deeptutor_shchem_v1 import desktop_handout_knowledge as handout
from integrations.deeptutor_shchem_v1 import desktop_preparation_sources as sources
from integrations.deeptutor_shchem_v1.desktop_composition_inline import NativeParagraphDisplay, display_tokens
from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition, layout_composition, write_pptx
from integrations.deeptutor_shchem_v1.desktop_composition_rich_render import flatten_elements, rich_lines
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    add_resource, block, digest, edit_resource, enable_composition, new_composition, validate_resource,
)
from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
from integrations.deeptutor_shchem_v1.desktop_word_preview_cache import WordPreviewCache
from integrations.deeptutor_shchem_v1.offline_word_image_label_review import _xml_sha
from integrations.deeptutor_shchem_v1.word_native_text import NATIVE_WORD_TEXT_REVISION
from integrations.deeptutor_shchem_v1.word_static_eq import REVISION as STATIC_EQ_REVISION, layout_eq, png_bytes


ARTIFACTS = Path("E:/项目/上海高中化学等级考/outputs/2026-10-08-polymer-field-workbench/eq-worker/projection-tests")
EQ_TEXT = "――→〔上移7pt：H^{＋}；下移5pt：△〕"
CROP = {"l": 20000, "t": 10000, "r": 20000, "b": 10000}


def _native_png():
    image = Image.new("RGB", (100, 50), "#20bb50")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 19, 49), fill="#ff0000")
    draw.rectangle((80, 0, 99, 49), fill="#0000ff")
    draw.line((25, 25, 75, 25), fill="black", width=2)
    draw.polygon([(75, 25), (65, 18), (65, 32)], fill="black")
    output = BytesIO()
    image.save(output, "PNG")
    return output.getvalue()


def _add_text_eq(paragraph):
    paragraph._p.extend(_field(
        _run(instruction=r" EQ \o(――→,\s\up7(H", font_pt=12),
        _run(instruction="＋", script="superscript", font_pt=12),
        _run(instruction=r"),\s\do5(△)) ", font_pt=12),
    ))


def _add_image_eq(paragraph, raw):
    paragraph._p.extend([_run(boundary="begin"), _run(instruction=r" EQ \o(", font_pt=12)])
    picture_run = paragraph.add_run()
    picture_run.font.size = Pt(12)
    picture = picture_run.add_picture(BytesIO(raw), width=Mm(24), height=Mm(4))
    fill = picture._inline.find(".//{http://schemas.openxmlformats.org/drawingml/2006/picture}blipFill")
    crop = OxmlElement("a:srcRect")
    for key, value in CROP.items():
        crop.set(key, str(value))
    fill.insert(1, crop)
    paragraph._p.extend([
        _run(instruction=r",\s\up7(催化剂)) ", font_pt=12),
        _run(boundary="separate"), _run(boundary="end"),
    ])


def _rows(item, kind="text", *, original=False):
    rows = item["original"]["blocks"] if original else item["blocks"]
    return [row for row in rows if row.get("presentation", {}).get("eqs")
            and row["presentation"]["eqs"][0]["field"]["base"]["kind"] == kind]


def _tokens(item, row):
    original = next(value["value"] for value in item["original"]["blocks"] if value["id"] == row["id"])
    return display_tokens(row, {value["id"]: value for value in item["blocks"]}, original)


def _transcription(tokens):
    parts = []
    for token in tokens:
        if token["kind"] == "eq":
            parts.append(token["field"]["transcription"])
        elif token["kind"] == "text":
            text = token["text"]
            parts.append(text if token["script"] == "baseline" else
                         ("_{" if token["script"] == "subscript" else "^{") + text + "}")
        else:
            parts.append(token["placeholder"])
    return "".join(parts)


def _composition(item):
    value = new_composition()
    return add_resource(value, item, {"page_id": value["pages"][0], "x": 12, "y": 12, "width": 186, "height": 273})


def _eq_token(material, kind="text"):
    return next(token for token in _tokens(material.item, _rows(material.item, kind)[0]) if token["kind"] == "eq")


def _flat_picture(picture):
    return list(flatten_elements([{**picture, "x": 0, "y": 0}]))


@pytest.fixture
def eq_material(tmp_path):
    document = Document()
    document.add_heading("知识点1 聚合反应条件", 3)
    paragraph = document.add_paragraph("水H")
    paragraph.add_run("2").font.subscript = True
    paragraph.add_run("O；")
    _add_text_eq(paragraph)
    paragraph.add_run("；生成H")
    paragraph.add_run("2").font.subscript = True
    paragraph.add_run("O。")
    native_png = _native_png()
    paragraph = document.add_paragraph("原图条件：")
    _add_image_eq(paragraph, native_png)
    paragraph.add_run("；条件说明完整。")
    document.add_paragraph("同一原图的普通出现：").add_run().add_picture(BytesIO(native_png), width=Mm(24), height=Mm(12))
    document.add_paragraph("完整栏目尾部说明。")
    document.add_heading("知识点2 下一栏目", 3)
    document.add_paragraph("此段不属于选中的完整栏目。")
    raw = word_bytes(document)
    facade = make_facade(raw, tmp_path / "source")
    request = request_for(facade, "知识点1 聚合反应条件")
    result = handout.import_handout_knowledge(facade, [request], [], [])
    store = PreparationImageStore(facade.paths.task_root / "preparation-v1/images")
    image_data = {asset["asset_id"]: store.load(asset) for asset in result["image_assets"]}
    return SimpleNamespace(raw=raw, facade=facade, request=request, result=result,
                           item=result["resources"][0], assets=result["image_assets"],
                           image_data=image_data, native_png=native_png)


@pytest.fixture
def artifact_dir():
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    return ARTIFACTS


def test_closed_uncached_eq_imports_the_entire_section_and_keeps_source_evidence(eq_material):
    material = eq_material
    item = material.item
    snapshot = item["original"]["evidence"]["section_snapshot"]
    assert [row["index"] for row in snapshot["blocks"]] == [1, 2, 3, 4, 5]
    assert snapshot["blocks"][-1]["text"] == "完整栏目尾部说明。"
    assert all("此段不属于" not in row["text"] for row in snapshot["blocks"])
    assert len(_rows(item)) == 1 and len(_rows(item, "image")) == 1
    assert sum(len(row.get("static_eqs", [])) for row in snapshot["blocks"]) == 2
    for row in _rows(item, original=True) + _rows(item, "image", original=True):
        source = next(source for source in snapshot["blocks"] if source["index"] == row["source_index"])
        assert source["static_eqs"] == row["presentation"]["eqs"]
    assert len(material.assets) == 1
    assert len([row for row in item["blocks"] if row["kind"] == "image"]) == 2
    assert validate_resource(item) == item
    assert material.facade.original.read_bytes() == material.raw


@pytest.mark.parametrize("unknown", [" PAGE ", r' TOC \o "1-3" ', r' QUOTE "条件" '])
def test_unknown_field_rejects_a_mixed_import_before_writing_any_picture(tmp_path, unknown):
    document = Document()
    document.add_heading("知识点1 可读栏目", 3)
    _add_text_eq(document.add_paragraph("反应条件："))
    document.add_paragraph().add_run().add_picture(BytesIO(_native_png()))
    document.add_heading("知识点2 未支持字段", 3)
    document.add_paragraph("原条件：")._p.extend(_field(_run(instruction=unknown)))
    raw = word_bytes(document)
    facade = make_facade(raw, tmp_path / "atomic-source")
    requests = [request_for(facade, "知识点1 可读栏目"), request_for(facade, "知识点2 未支持字段")]
    before = deepcopy(requests)
    with pytest.raises(ValueError, match="Word公式域或字段"):
        handout.import_handout_knowledge(facade, requests, [], [])
    assert requests == before
    assert not facade.paths.task_root.exists()
    assert facade.original.read_bytes() == raw


def test_table_with_unprojected_eq_keeps_the_complete_section_guard(tmp_path):
    document = Document()
    document.add_heading("知识点1 表格条件", 3)
    document.add_paragraph("完整表格如下。")
    table = document.add_table(rows=1, cols=1)
    _add_text_eq(table.cell(0, 0).paragraphs[0])
    raw = word_bytes(document)
    facade = make_facade(raw, tmp_path / "table-source")
    request = request_for(facade, "知识点1 表格条件")
    preview = sources.PreparationSourcesService(tmp_path).word_preview_bytes(raw, "table-eq.docx")
    assert preview["extraction_revision"] == NATIVE_WORD_TEXT_REVISION + "+" + STATIC_EQ_REVISION
    assert EQ_TEXT in preview["blocks"][-1]["text"]
    assert all("static_eqs" not in row for row in preview["blocks"])
    with pytest.raises(ValueError, match="Word公式域或字段"):
        handout.import_handout_knowledge(facade, [request], [], [])
    assert not facade.paths.task_root.exists()
    assert facade.original.read_bytes() == raw


def test_native_bind_keeps_scripts_on_both_sides_of_the_eq_span(eq_material):
    display = NativeParagraphDisplay(eq_material.raw)
    native = display.reader.read(display.elements[1])
    row = block("text", "knowledge", native.text, 2)
    before = _xml_sha(display.elements[1])
    display.bind(row, [], [])
    assert row["value"] == "水H_{2}O；" + EQ_TEXT + "；生成H_{2}O。"
    tokens = display_tokens(row, {})
    assert [token["text"] for token in tokens if token["kind"] == "text" and token["script"] == "subscript"] == ["2", "2"]
    assert _transcription(tokens) == row["value"]
    eq = next(token for token in tokens if token["kind"] == "eq")
    assert eq["field"]["source_font_pt"] == 12
    assert eq["field"]["overlays"][0]["runs"][-1] == {"text": "＋", "script": "superscript"}
    assert _xml_sha(display.elements[1]) == before


def test_work_copy_prefix_retains_the_whole_field_but_edited_label_becomes_visible_raw_text(eq_material):
    item = eq_material.item
    row = _rows(item)[0]
    value = edit_resource(_composition(item), item["id"], locked=False)
    prefix = "新增讲解。"
    prefixed = edit_resource(value, item["id"], values={row["id"]: prefix + row["value"]})["resources"][0]
    prefixed_row = next(part for part in prefixed["blocks"] if part["id"] == row["id"])
    assert _transcription(_tokens(prefixed, prefixed_row)) == prefix + row["value"]
    assert next(token for token in _tokens(prefixed, prefixed_row) if token["kind"] == "eq")["field"] == row["presentation"]["eqs"][0]["field"]
    edited_text = row["value"].replace("H^{＋}", "Ca^{2＋}")
    edited = edit_resource(value, item["id"], values={row["id"]: edited_text})["resources"][0]
    edited_row = next(part for part in edited["blocks"] if part["id"] == row["id"])
    tokens = _tokens(edited, edited_row)
    assert not any(token["kind"] == "eq" for token in tokens)
    assert _transcription(tokens) == edited_text
    assert any("――→〔上移7pt：Ca^{2＋}；下移5pt：△〕" in token.get("text", "") and token["script"] == "baseline" for token in tokens)
    assert edited["original"] == prefixed["original"] == item["original"]
    assert validate_resource(edited) == edited


@pytest.mark.parametrize("property_name,replacement", [("source_font_pt", 16), ("field_sha256", "0" * 64)])
def test_work_copy_presentation_cannot_be_silently_replaced(eq_material, property_name, replacement):
    altered = deepcopy(eq_material.item)
    _rows(altered)[0]["presentation"]["eqs"][0]["field"][property_name] = replacement
    with pytest.raises(ValueError, match="不能被静默替换"):
        validate_resource(altered)


def test_resealed_source_field_sha_must_match_the_whole_section_snapshot(eq_material):
    altered = deepcopy(eq_material.item)
    snapshot = altered["original"]["evidence"]["section_snapshot"]["blocks"]
    next(row for row in snapshot if row.get("static_eqs"))["static_eqs"][0]["field"]["field_sha256"] = "0" * 64
    altered["original_sha256"] = digest(altered["original"])
    with pytest.raises(ValueError, match="静态EQ.*快照失配"):
        validate_resource(altered)


def test_eq_revision_changes_only_sources_with_an_accepted_eq(eq_material, tmp_path, monkeypatch):
    service = sources.PreparationSourcesService(tmp_path)
    ordinary = Document()
    ordinary.add_paragraph("普通正文H").add_run("2").font.subscript = True
    raw = word_bytes(ordinary)
    plain_before = service.word_preview_bytes(raw, "ordinary.docx")
    eq_before = service.word_preview_bytes(eq_material.raw, "eq.docx")
    assert plain_before["extraction_revision"] == NATIVE_WORD_TEXT_REVISION
    assert eq_before["extraction_revision"] == NATIVE_WORD_TEXT_REVISION + "+" + STATIC_EQ_REVISION
    monkeypatch.setattr(sources, "STATIC_EQ_REVISION", "future-static-eq-revision")
    assert service.word_preview_bytes(raw, "ordinary.docx") == plain_before
    eq_after = service.word_preview_bytes(eq_material.raw, "eq.docx")
    assert eq_after["source_sha256"] == eq_before["source_sha256"]
    assert eq_after["revision"] != eq_before["revision"]


@pytest.mark.parametrize("stale_part", ["record", "extraction"])
def test_preview_cache_rejects_a_stale_static_eq_revision(eq_material, tmp_path, stale_part):
    preview = sources.PreparationSourcesService(tmp_path).word_preview_bytes(eq_material.raw, "eq.docx")
    cache = WordPreviewCache(tmp_path / "cache")
    assert cache.save(eq_material.raw, "eq.docx", preview)
    assert cache.load(eq_material.raw, "eq.docx") == preview
    path = next(cache.root.glob("*.json.gz"))
    record = json.loads(gzip.decompress(path.read_bytes()))
    if stale_part == "record":
        record["static_eq_revision"] = "old-static-eq-revision"
    else:
        record["preview"]["extraction_revision"] = NATIVE_WORD_TEXT_REVISION + "+old-static-eq-revision"
        record["preview"]["revision"] = sources._digest({key: value for key, value in record["preview"].items() if key != "revision"})
    path.write_bytes(gzip.compress(json.dumps(record, ensure_ascii=False).encode("utf-8")))
    assert cache.load(eq_material.raw, "eq.docx") is None


def test_qt_eq_primitives_have_real_charge_and_upper_lower_geometry(eq_material, app, artifact_dir):
    token = _eq_token(eq_material)
    picture = layout_eq(token, {}, {}, font_px=16)
    leaves = _flat_picture(picture)
    by_text = {leaf.get("text"): leaf for leaf in leaves if leaf["kind"] == "text"}
    assert set(by_text) == {"――→", "H", "＋", "△"}
    base, hydrogen, charge, lower = [by_text[text] for text in ("――→", "H", "＋", "△")]
    assert charge["y"] < hydrogen["y"] < base["y"] < lower["y"]
    assert charge["font_px"] < hydrogen["font_px"] == base["font_px"]
    assert "SuperScript" in charge["script"]
    assert base["y"] - hydrogen["y"] == pytest.approx(7 * 25.4 / 72, abs=0.4)
    assert lower["y"] - base["y"] == pytest.approx(5 * 25.4 / 72, abs=0.4)
    assert picture["descent_mm"] > 0
    assert 0 < picture["baseline_mm"] < picture["height"]
    for leaf in leaves:
        assert -0.01 <= leaf["x"] and leaf["x"] + leaf["width"] <= picture["width"] + 0.25
        assert -0.01 <= leaf["y"] and leaf["y"] + leaf["height"] <= picture["height"] + 0.01
    (artifact_dir / "text-eq.png").write_bytes(png_bytes(picture["_image"]))


def test_actual_qt_linebreak_without_glyphs_keeps_text_and_uses_real_font(app):
    row = {"id": "source-linebreak", "value": "\u2028"}
    tokens = [{"kind": "text", "text": "\u2028", "script": "baseline"}]
    original_row, original_tokens = deepcopy(row), deepcopy(tokens)
    chunks = rich_lines(row, {}, 80, {}, {}, 120, tokens=tokens)
    positioned = []
    offset = 0
    for chunk in chunks:
        positioned.append({**chunk, "x": 0, "y": offset})
        offset += chunk["height"]
    leaves = [leaf for leaf in flatten_elements(positioned) if leaf["kind"] == "text"]
    assert leaves and "".join(leaf["text"] for leaf in leaves) == "\n"
    assert all(isinstance(leaf["font_family"], str) and leaf["font_family"].strip() for leaf in leaves)
    assert all(leaf["font_px"] > 0 and leaf["height"] > 0 for leaf in leaves)
    assert row == original_row and tokens == original_tokens


def test_qt_cropped_image_uses_bound_original_and_keeps_source_bytes_and_evidence(eq_material, app, artifact_dir):
    material = eq_material
    before = deepcopy(material.item)
    raw_before = deepcopy(material.image_data)
    token = _eq_token(material, "image")
    picture = layout_eq(token, {asset["asset_id"]: asset for asset in material.assets}, material.image_data)
    image = next(leaf for leaf in _flat_picture(picture) if leaf["kind"] == "image")
    assert image["asset_id"] == token["asset_id"]
    assert (image["width"], image["height"]) == pytest.approx((24, 4))
    assert image["crop"] == CROP
    assert image["_image"].pixelColor(5, 5).getRgb()[:3] == (32, 187, 80)
    assert Image.open(BytesIO(material.native_png)).getpixel((5, 5)) == (255, 0, 0)
    source = NativeParagraphDisplay(material.raw)
    drawing = source.elements[2].find(".//" + W + "drawing")
    assert token["field"]["base"]["container_xml_sha256"] == _xml_sha(drawing)
    assert sha256(material.image_data[token["asset_id"]]).hexdigest() == material.assets[0]["sha256"]
    assert material.image_data == raw_before and material.item == before
    assert material.facade.original.read_bytes() == material.raw
    (artifact_dir / "native-source.png").write_bytes(material.native_png)
    (artifact_dir / "image-eq.png").write_bytes(png_bytes(picture["_image"]))


def test_full_layout_keeps_eq_base_with_outer_baseline_and_does_not_repeat_bound_image(eq_material, app):
    value = arrange_composition(_composition(eq_material.item), eq_material.assets, eq_material.image_data)
    layout = layout_composition(value, eq_material.assets, eq_material.image_data, student=True)
    assert not layout["errors"]
    leaves = [leaf for page in layout["pages"].values() for leaf in flatten_elements(page)]
    text_row = _rows(eq_material.item)[0]
    field_sha = text_row["presentation"]["eqs"][0]["field"]["field_sha256"]
    eq_parts = [leaf for leaf in leaves if leaf.get("eq_field_sha256") == field_sha]
    assert eq_parts and all(leaf["block_id"] == text_row["id"] for leaf in eq_parts)
    base = next(leaf for leaf in leaves if leaf.get("eq_field_sha256") == field_sha and leaf.get("text") == "――→")
    outer = next(leaf for leaf in leaves if leaf.get("block_id") == text_row["id"] and leaf.get("text") == "水H")
    assert abs(base["y"] - outer["y"]) < 0.7
    assert sum(leaf["kind"] == "image" for leaf in leaves) == 2
    assert sum(leaf.get("crop") == CROP for leaf in leaves) == 1
    assert eq_material.item["blocks"] == value["resources"][0]["blocks"]


def test_actual_pptx_contains_editable_eq_text_and_original_image_with_exact_crop(eq_material, app, artifact_dir):
    value = arrange_composition(_composition(eq_material.item), eq_material.assets, eq_material.image_data)
    layout = layout_composition(value, eq_material.assets, eq_material.image_data, student=True)
    assert not layout["errors"]
    path = artifact_dir / "native-primitives.pptx"
    write_pptx(layout, path, eq_material.image_data)
    presentation = Presentation(path)
    shapes = [shape for slide in presentation.slides for shape in slide.shapes]
    text_shapes = {shape.text: shape for shape in shapes if shape.has_text_frame}
    assert {"――→", "H", "＋", "△", "催化剂"} <= set(text_shapes)
    assert text_shapes["＋"].top < text_shapes["H"].top < text_shapes["――→"].top < text_shapes["△"].top
    assert text_shapes["＋"].text_frame.paragraphs[0].runs[0].font.size < text_shapes["H"].text_frame.paragraphs[0].runs[0].font.size
    pictures = [shape for shape in shapes if shape.shape_type == 13]
    assert len(pictures) == 2
    assert all(picture.image.blob == eq_material.native_png for picture in pictures)
    cropped = next(picture for picture in pictures if picture.crop_left > 0)
    assert (cropped.crop_left, cropped.crop_top, cropped.crop_right, cropped.crop_bottom) == pytest.approx((0.2, 0.1, 0.2, 0.1))
    plain = next(picture for picture in pictures if picture is not cropped)
    assert (plain.crop_left, plain.crop_top, plain.crop_right, plain.crop_bottom) == (0, 0, 0, 0)


def test_actual_docx_uses_exact_eq_visual_bytes_and_keeps_ordinary_native_image(eq_material, app, artifact_dir):
    material = eq_material
    assets = {asset["asset_id"]: asset for asset in material.assets}
    pictures = {kind: layout_eq(_eq_token(material, kind), assets, material.image_data) for kind in ("text", "image")}
    expected = {kind: png_bytes(picture["_image"]) for kind, picture in pictures.items()}
    payload = sample()
    payload["lesson_design"] = enable_composition(payload["lesson_design"])
    payload["lesson_design"]["composition"] = arrange_composition(_composition(material.item), material.assets, material.image_data)
    payload["image_assets"] = material.assets
    before = deepcopy(payload)
    directory = artifact_dir / "documents"
    directory.mkdir(parents=True, exist_ok=True)
    write_documents(payload, directory, material.image_data)
    for name in ("lesson_plan.docx", "student_worksheet.docx"):
        path = directory / name
        document = Document(path)
        with ZipFile(path) as package:
            media = {package.read(member) for member in package.namelist() if member.startswith("word/media/")}
        assert media == {material.native_png, *expected.values()}
        assert len(document.inline_shapes) == 3
        paragraph = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("水H"))
        assert paragraph.text == "水H2O；；生成H2O。"
        assert [run.text for run in paragraph.runs if run.font.subscript] == ["2", "2"]
        for run in paragraph.runs:
            drawing = run._r.find(W + "drawing")
            if drawing is None:
                continue
            relationship = drawing.find(".//" + A + "blip").get(R + "embed")
            assert document.part.related_parts[relationship].blob == expected["text"]
            position = run._r.find(W + "rPr/" + W + "position")
            assert position is not None and int(position.get(W + "val")) < 0
    assert payload == before
    assert material.facade.original.read_bytes() == material.raw
    assert set(material.image_data.values()) == {material.native_png}
