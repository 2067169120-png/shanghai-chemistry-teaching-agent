from __future__ import annotations

import hashlib
import json

import pytest
from docx import Document
from docx.oxml import OxmlElement

from integrations.deeptutor_shchem_v1.desktop_preparation_limits import MAX_MATERIALS
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
    CONCEPTS,
    PreparationSourceError,
    PreparationSourcesService,
)


@pytest.fixture
def source_service(tmp_path):
    book = tmp_path / "books/textbook.pdf"
    book.parent.mkdir()
    book.write_bytes(b"fixture-source-book")
    row = {
        "concept_id": "C01",
        "title": "电解质",
        "statement": "概念候选及条件",
        "volume_id": "TB-M1",
        "source_path": "books/textbook.pdf",
        "source_sha256": hashlib.sha256(book.read_bytes()).hexdigest(),
        "pdf_pages": [61],
        "candidate_only": True,
        "human_reviewed": False,
        "teaching_use_allowed": False,
        "generation_allowed": False,
        "publication_allowed": False,
    }
    catalog = tmp_path / CONCEPTS
    catalog.parent.mkdir(parents=True)
    catalog.write_text(json.dumps(row, ensure_ascii=False) + "\n", encoding="utf-8")
    return PreparationSourcesService(tmp_path)


def _word(tmp_path):
    document = Document()
    document.add_heading("知识点一", 1)
    paragraph = document.add_paragraph("SO")
    paragraph.add_run("4").font.subscript = True
    paragraph.add_run("2−").font.superscript = True
    paragraph.add_run("保留条件")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "概念比较"
    table.cell(1, 0).text = "状态"
    table.cell(1, 1).text = "溶液"
    paragraph = document.add_paragraph("旧公式在此：")
    paragraph.add_run()._r.append(OxmlElement("w:object"))
    paragraph.add_run("后续文字")
    path = tmp_path / "teacher-handout.docx"
    document.save(path)
    return path


def test_word_native_order_scripts_tables_and_object_gap(source_service, tmp_path):
    path = _word(tmp_path)
    before = path.read_bytes()
    preview = source_service.word_preview(path)
    blocks = preview["blocks"]
    assert [b["index"] for b in blocks] == [1, 2, 3, 4]
    assert "SO_{4}^{2−}保留条件" == blocks[1]["text"]
    assert "〔第1行·第1—2列；横向合并〕\n概念比较" in blocks[2]["text"]
    assert "状态" in blocks[2]["text"] and "溶液" in blocks[2]["text"]
    assert "旧公式在此：【待查看原文：嵌入对象或旧公式】后续文字" == blocks[3]["text"]
    assert blocks[3]["warnings"]
    assert path.read_bytes() == before
    assert str(path) not in json.dumps(preview)


def test_word_sections_exclude_toc_and_include_nested_content(source_service, tmp_path):
    from docx.enum.style import WD_STYLE_TYPE

    document = Document()
    document.styles.add_style("TOC 2", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("考点一 电解质", style="TOC 2")
    document.add_heading("第三章", 1)
    document.add_heading("考点一 电解质", 2)
    document.add_heading("知识点1", 3)
    document.add_paragraph("完整知识总结")
    document.add_paragraph("例1 题干与条件")
    document.add_paragraph("【解析】完整答案与依据")
    document.add_heading("考点二 离子反应", 2)
    document.add_paragraph("下一节正文")
    document.add_heading("第四章", 1)
    path = tmp_path / "sections.docx"
    document.save(path)
    preview = source_service.word_preview(path)
    assert preview["sections"] == [
        {"start": 2, "end": 9, "level": 1, "title": "第三章"},
        {"start": 3, "end": 7, "level": 2, "title": "考点一 电解质"},
        {"start": 8, "end": 9, "level": 2, "title": "考点二 离子反应"},
        {"start": 10, "end": 10, "level": 1, "title": "第四章"},
    ]
    selected = source_service.reference(str(path), preview["source_sha256"], 3, 7, [])
    assert all(
        text in selected["materials"]
        for text in ("完整知识总结", "例1 题干与条件", "【解析】完整答案与依据")
    )
    assert "下一节正文" not in selected["materials"]


def test_unstyled_word_does_not_guess_chapters_from_topic_words(
    source_service, tmp_path
):
    document = Document()
    document.add_paragraph("考点一 电解质")
    document.add_paragraph("知识点、例题、总结只是普通文本")
    path = tmp_path / "no-headings.docx"
    document.save(path)
    assert source_service.word_preview(path)["sections"] == []


def test_directory_heading_is_not_offered_as_whole_document_section(
    source_service, tmp_path
):
    document = Document()
    document.add_heading("目 录", 1)
    document.add_paragraph("章节列表")
    document.add_heading("考点一 电解质", 2)
    document.add_paragraph("正文")
    path = tmp_path / "directory.docx"
    document.save(path)
    assert source_service.word_preview(path)["sections"] == [
        {"start": 3, "end": 4, "level": 2, "title": "考点一 电解质"}
    ]


def test_word_and_concepts_compile_exact_selected_snapshot(source_service, tmp_path):
    path = _word(tmp_path)
    preview = source_service.word_preview(path)
    option = source_service.concept_options("电解质")[0]
    reference = source_service.reference(
        str(path),
        preview["source_sha256"],
        2,
        4,
        [{k: option[k] for k in ("concept_id", "revision")}],
    )
    text = reference["materials"]
    assert "SO_{4}^{2−}" in text
    assert "知识点一" not in text  # Outside selected range.
    assert "概念候选及条件" in text and "PDF文件页序：[61]" in text
    assert '"generation_allowed": false' in text
    assert preview["source_sha256"] in text and option["revision"] in text
    assert reference["warnings"] and "不得根据残句补写" in text
    assert str(tmp_path) not in text


def test_concept_search_and_concept_only(source_service):
    assert source_service.concept_options("not-present") == []
    option = source_service.concept_options("C01 条件")[0]
    chosen = {k: option[k] for k in ("concept_id", "revision")}
    text = source_service.reference(None, None, 1, 1, [chosen])["materials"]
    assert "概念候选及条件" in text and "Word区块" not in text


@pytest.mark.parametrize(
    "case",
    [
        "word_changed",
        "concept_changed",
        "book_changed",
        "duplicate",
        "bad_range",
        "empty",
    ],
)
def test_stale_or_invalid_selection_is_not_compiled(source_service, tmp_path, case):
    path = _word(tmp_path)
    preview = source_service.word_preview(path)
    option = source_service.concept_options()[0]
    chosen = {k: option[k] for k in ("concept_id", "revision")}
    concepts = [chosen]
    start, end = 1, 2
    word_sha = preview["source_sha256"]
    if case == "word_changed":
        word_sha = "0" * 64
    elif case == "concept_changed":
        chosen["revision"] = "0" * 64
    elif case == "book_changed":
        (tmp_path / "books/textbook.pdf").write_bytes(b"changed-fixture")
    elif case == "duplicate":
        concepts *= 2
    elif case == "bad_range":
        end = 999
    else:
        path, concepts = None, []
    with pytest.raises(PreparationSourceError):
        source_service.reference(path, word_sha, start, end, concepts)


def test_long_material_fails_without_truncation(source_service, tmp_path):
    document = Document()
    unit = "长资料"
    document.add_paragraph(unit * (MAX_MATERIALS // len(unit) + 1))
    path = tmp_path / "long.docx"
    document.save(path)
    preview = source_service.word_preview(path)
    with pytest.raises(PreparationSourceError, match="未截断"):
        source_service.reference(path, preview["source_sha256"], 1, 1, [])


def test_hyperlink_math_and_hidden_text_are_not_silently_lost(source_service, tmp_path):
    document = Document()
    paragraph = document.add_paragraph("前")
    link = OxmlElement("w:hyperlink")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "链接显示文字"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    paragraph._p.append(OxmlElement("m:oMath"))
    paragraph.add_run("隐藏内容").font.hidden = True
    path = tmp_path / "special.docx"
    document.save(path)
    block = source_service.word_preview(path)["blocks"][0]
    assert "链接显示文字" in block["text"]
    assert "【待查看原文：数学公式】" in block["text"]
    assert "隐藏内容" not in block["text"]
    assert len(block["warnings"]) == 3


def test_facade_reference_draft_and_confirmed_fake_provider_roundtrip(
    source_service, tmp_path
):
    from test_desktop_preparation import FileRenderer
    from test_desktop_preparation_facade import (
        _facade,
        _payload,
        _ProviderStore,
        _Transport,
    )

    from integrations.deeptutor_shchem_v1.desktop_paths import DesktopPaths
    from integrations.deeptutor_shchem_v1.desktop_preparation import (
        DesktopPreparationManager,
    )

    paths = DesktopPaths.from_workspace(
        tmp_path, state_root=tmp_path / "isolated-state"
    )
    provider_store, transport = _ProviderStore(), _Transport()
    facade = _facade(
        paths,
        provider_store,
        transport,
        manager=DesktopPreparationManager(
            paths.task_root / "preparation-v1", FileRenderer()
        ),
    )
    path = _word(tmp_path)
    preview = facade.preparation_word_preview(str(path))
    option = facade.preparation_concept_options("电解质")[0]
    ref = facade.preparation_source_reference(
        str(path),
        preview["source_sha256"],
        1,
        4,
        [{k: option[k] for k in ("concept_id", "revision")}],
    )
    payload = {**_payload(output_kind="ppt"), "materials": ref["materials"]}
    receipt = facade.create_preparation_draft(payload)
    draft = next(
        row
        for row in facade.preparation_draft_options()
        if row["draft_id"] == receipt.draft_id
    )
    loaded = facade.load_preparation_draft(receipt.draft_id, draft["revision"])
    assert loaded["payload"]["materials"] == ref["materials"]
    assert transport.calls == 0 and provider_store.borrow_calls == []
    prepared = facade.prepare_preparation(
        loaded["payload"], "teacher-text", provider_store.revision
    )
    result = facade.generate_preparation(
        prepared.task_id, teacher_confirmed=True, should_cancel=lambda: False
    )
    assert result.status == "completed"
    assert transport.calls == 1
    body = json.loads(transport.request_bodies[0])
    prompt = body["input"][0]["content"][0]["text"]
    embedded, _ = json.JSONDecoder().raw_decode(
        prompt.split("教师备课简报 JSON：\n", 1)[1]
    )
    assert embedded["materials"] == ref["materials"]
    assert "〔第1行·第1—2列；横向合并〕\n概念比较" in embedded["materials"]
    for text in (
        "SO_{4}^{2−}",
        "概念候选及条件",
        preview["source_sha256"],
        "待查看原文",
        "区块不是页码",
    ):
        assert text in prompt
    assert str(tmp_path) not in prompt
