from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt
from PIL import Image

from integrations.deeptutor_shchem_v1.desktop_word_question_export import (
    O,
    R,
    V,
    WordQuestionExportError,
    _body_blocks,
    export_word_questions,
)


def data(document):
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def item(document, question, answer=(), context=(), key="one", points=2):
    content = data(document)
    return {
        "source_bytes": content, "source_sha256": sha256(content).hexdigest(),
        "source_name": "演示讲义.docx", "key": key, "revision": "test-revision",
        "points": points, "export_ready": True,
        "question_blocks": [{"index": n} for n in question],
        "answer_blocks": [{"index": n} for n in answer],
        "context_blocks": [{"index": n} for n in context],
    }


def output_doc(result, role="student"):
    return Document(BytesIO(result[role + "_bytes"]))


def text(document):
    return "\n".join(document.element.body.itertext())


def picture(document, color):
    stream = BytesIO()
    Image.new("RGB", (32, 16), color).save(stream, format="PNG")
    stream.seek(0)
    document.add_paragraph().add_run().add_picture(stream, width=Inches(1))


def test_only_selected_question_and_teacher_answer_scores():
    source = Document()
    for value in ("第一题题干", "【答案】第一答案", "第二题未选", "【答案】第二答案"):
        source.add_paragraph(value)
    result = export_word_questions("化学课堂练习", [item(source, [1], [2], points=4)])
    student, teacher = output_doc(result), output_doc(result, "teacher")
    assert "第一题题干" in text(student)
    assert "第一答案" not in text(student)
    assert "第二题未选" not in text(student)
    assert "第二答案" not in text(teacher)
    assert "第一答案" in text(teacher)
    assert any(p.text == "练习 1  4 分" for p in teacher.paragraphs)
    assert any(p.text == "练习 1" for p in student.paragraphs)
    assert not any("___" in p.text for p in student.paragraphs)
    shown = export_word_questions("化学课堂练习", [item(source, [1], [2], points=4)], show_student_scores=True)
    assert any(p.text == "练习 1  4 分" for p in output_doc(shown).paragraphs)


def test_source_printed_score_and_existing_answer_space_not_erased_or_duplicated():
    source = Document()
    source.add_paragraph("1 （3分）填写化学式 __________")
    result = export_word_questions("课堂练习", [item(source, [1], points=5)])
    student = output_doc(result)
    assert sum(p.text.count("__________") for p in student.paragraphs) == 1
    assert "（3分）" in text(student)
    teacher_text = text(output_doc(result, "teacher"))
    assert "当前选定范围未识别到本题答案" in teacher_text
    assert "这不表示原文没有答案" in teacher_text
    assert "原文未提供本题答案。" not in teacher_text


def test_source_images_with_colliding_relationship_ids_are_rebound():
    first, second = Document(), Document()
    picture(first, "red")
    picture(second, "blue")
    result = export_word_questions("图题练习", [item(first, [1], key="red"), item(second, [1], key="blue")])
    student = output_doc(result)
    relationships = [r for r in student.part.rels.values() if r.reltype == RT.IMAGE]
    assert len(relationships) == 2
    assert len({sha256(r.target_part.blob).hexdigest() for r in relationships}) == 2
    refs = student.element.xpath(".//a:blip/@r:embed")
    assert len(refs) == 2 and len(set(refs)) == 2
    assert all(student.part.rels[r].reltype == RT.IMAGE for r in refs)
    assert len(set(student.element.xpath(".//wp:docPr/@id"))) == 2


def test_native_scripts_omml_table_and_source_styles_are_retained():
    source = Document()
    style = source.styles.add_style("Chemistry", WD_STYLE_TYPE.PARAGRAPH)
    style.font.size = Pt(17)
    paragraph = source.add_paragraph(style="Chemistry")
    paragraph.add_run("SO")
    paragraph.add_run("4").font.subscript = True
    paragraph.add_run("2−").font.superscript = True
    paragraph._p.append(parse_xml(f'<m:oMath {nsdecls("m")}><m:f><m:num><m:r><m:t>n</m:t></m:r></m:num><m:den><m:r><m:t>V</m:t></m:r></m:den></m:f></m:oMath>'))
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "浓度", "mol/L"
    result = export_word_questions("化学表达练习", [item(source, [1, 2])])
    student = output_doc(result)
    assert student.element.xpath(".//m:oMath/m:f/m:den/m:r/m:t")[0].text == "V"
    assert set(student.element.xpath(".//w:vertAlign/@w:val")) == {"subscript", "superscript"}
    assert student.tables[0].cell(0, 0).text == "浓度"
    chemistry = next(p for p in student.paragraphs if "SO" in p.text)
    assert chemistry.style.font.size == Pt(17)
    assert chemistry.style.style_id != "Chemistry"


def test_same_style_names_from_different_sources_keep_distinct_definitions():
    sources = []
    for i, size in enumerate((11, 18)):
        doc = Document()
        style = doc.styles.add_style("SameName", WD_STYLE_TYPE.PARAGRAPH)
        style.font.size = Pt(size)
        doc.add_paragraph(f"选题{i}", "SameName")
        sources.append(item(doc, [1], key=str(i)))
    result = output_doc(export_word_questions("练习", sources))
    paragraphs = [p for p in result.paragraphs if p.text in ("1. 选题0", "2. 选题1")]
    assert [p.style.font.size for p in paragraphs] == [Pt(11), Pt(18)]
    assert len({p.style.style_id for p in paragraphs}) == 2


def test_numbering_definitions_are_rebound_between_sources():
    documents = []
    for label in ("甲题", "乙题"):
        source = Document()
        source.add_paragraph(label, "List Number")
        documents.append(item(source, [1], key=label))
    student = output_doc(export_word_questions("编号练习", documents))
    nums = []
    for paragraph in [p for p in student.paragraphs if p.text.endswith("题")]:
        refs = paragraph.style.element.xpath("./w:pPr/w:numPr/w:numId/@w:val")
        assert len(refs) == 1
        nums.extend(refs)
    assert len(set(nums)) == 2
    numbering = student.part.numbering_part.element
    known = {n.get(qn("w:numId")) for n in numbering if n.tag == qn("w:num")}
    assert set(nums) <= known


def test_shared_material_is_once_per_source_block_in_order():
    source = Document()
    for value in ("公共实验材料", "第一题", "【答案】甲", "第二题", "【答案】乙"):
        source.add_paragraph(value)
    one = item(source, [2], [3], [1], key="one")
    two = {**one, "key": "two", "question_blocks": [{"index": 4}], "answer_blocks": [{"index": 5}]}
    result = export_word_questions("材料练习", [one, two])
    for role in ("student", "teacher"):
        paragraphs = [p.text for p in output_doc(result, role).paragraphs]
        assert paragraphs.count("公共实验材料") == 1
        assert paragraphs.index("公共实验材料") < paragraphs.index("1. 第一题") < paragraphs.index("2. 第二题")


@pytest.mark.parametrize("mutation,match", [
    ({"source_sha256": "bad"}, "原文件与预览"),
    ({"export_ready": False}, "边界"),
    ({"question_blocks": [{"index": 0}]}, "范围"),
    ({"question_blocks": [{"index": 1, "split_display": True}]}, "同一"),
    ({"answer_blocks": [{"index": 1}]}, "重叠"),
    ({"points": float("nan")}, "分值"),
    ({"points": 0}, "分值"),
])
def test_invalid_selection_is_rejected(mutation, match):
    source = Document()
    source.add_paragraph("题干")
    selection = item(source, [1])
    selection.update(mutation)
    with pytest.raises(WordQuestionExportError, match=match):
        export_word_questions("练习", [selection])


def test_inline_answer_and_cross_question_answer_leak_are_rejected():
    source = Document()
    source.add_paragraph("题干【答案】不能出现在学生卷")
    with pytest.raises(WordQuestionExportError, match="答案标记"):
        export_word_questions("练习", [item(source, [1])])
    source = Document()
    source.add_paragraph("题干一")
    source.add_paragraph("无标签的答案文本")
    source.add_paragraph("题干二")
    one = item(source, [1], [2], key="one")
    two = {**one, "key": "two", "question_blocks": [{"index": 3}], "answer_blocks": [], "context_blocks": [{"index": 2}]}
    with pytest.raises(WordQuestionExportError, match="另一道题的答案"):
        export_word_questions("练习", [one, two])


@pytest.mark.parametrize("kind", ["hyperlink", "external_image", "field", "hidden", "hidden_style", "comment"])
def test_unsupported_or_nonvisible_source_dependencies_fail(kind):
    source = Document()
    paragraph = source.add_paragraph("题目")
    if kind == "hyperlink":
        relation = source.part.relate_to("https://example.com", RT.HYPERLINK, is_external=True)
        node = OxmlElement("w:hyperlink")
        node.set(qn("r:id"), relation)
        paragraph._p.append(node)
    elif kind == "external_image":
        relation = source.part.relate_to("https://example.com/image.png", RT.IMAGE, is_external=True)
        paragraph._p.append(parse_xml(f'<w:r {nsdecls("w", "r", "a")}><w:drawing><a:blip r:link="{relation}"/></w:drawing></w:r>'))
    elif kind == "field":
        paragraph._p.append(OxmlElement("w:fldSimple"))
    elif kind == "hidden":
        paragraph.add_run("隐藏答案").font.hidden = True
    elif kind == "hidden_style":
        style = source.styles.add_style("HiddenAnswer", WD_STYLE_TYPE.CHARACTER)
        style.font.hidden = True
        paragraph.add_run("隐藏答案", "HiddenAnswer")
    else:
        paragraph._p.append(OxmlElement("w:commentReference"))
    with pytest.raises(WordQuestionExportError):
        export_word_questions("练习", [item(source, [1])])


def vml_source(*, ole=False, draw_aspect="Content", fallback=True):
    source = Document()
    picture(source, "green")
    paragraph = source.paragraphs[0]
    rid = paragraph._p.xpath(".//a:blip/@r:embed")[0]
    paragraph.clear()
    tag = "object" if ole else "pict"
    image = f'<v:imagedata r:id="{rid}"/>' if fallback else ""
    embedded = f'<o:OLEObject Type="Embed" ProgID="Equation.3" DrawAspect="{draw_aspect}" ShapeID="_oldshape" r:id="unresolvedOLE"/>' if ole else ""
    paragraph._p.append(parse_xml(f'<w:r {nsdecls("w", "r")} xmlns:v="{V}" xmlns:o="{O}"><w:{tag}><v:shapetype id="_x0000_t75" coordsize="21600,21600"/><v:shape id="_oldshape" type="#_x0000_t75" style="width:72pt;height:36pt">{image}</v:shape>{embedded}</w:{tag}></w:r>'))
    return source


def test_static_vml_image_type_and_relationship_are_retained():
    source = vml_source()
    result = export_word_questions("图题", [item(source, [1])])
    student = output_doc(result)
    shapes = list(student.element.iter(f"{{{V}}}shape"))
    types = list(student.element.iter(f"{{{V}}}shapetype"))
    assert len(shapes) == len(types) == 1
    assert shapes[0].get("type") == "#" + types[0].get("id")
    image = next(student.element.iter(f"{{{V}}}imagedata"))
    assert student.part.rels[image.get(f"{{{R}}}id")].reltype == RT.IMAGE


def test_standard_vml_picture_without_local_type_keeps_image_and_geometry():
    source = vml_source()
    definition = next(source.element.iter(f"{{{V}}}shapetype"))
    definition.getparent().remove(definition)
    shape = next(source.element.iter(f"{{{V}}}shape"))
    shape.set(f"{{{O}}}spt", "75")
    shape.set("style", "height:106.8pt;width:163.2pt;rotation:10")
    original_image = next(shape.iter(f"{{{V}}}imagedata"))
    original_image.set("cropleft", "0.1")
    original_blob = source.part.rels[original_image.get(f"{{{R}}}id")].target_part.blob
    student = output_doc(export_word_questions("图片练习", [item(source, [1])]))
    rect = next(student.element.iter(f"{{{V}}}rect"))
    assert rect.get("style") == shape.get("style")
    assert rect.get("stroked") == "f"
    assert not list(student.element.iter(f"{{{V}}}shapetype"))
    image = next(rect.iter(f"{{{V}}}imagedata"))
    assert image.get("cropleft") == "0.1"
    assert student.part.rels[image.get(f"{{{R}}}id")].target_part.blob == original_blob


@pytest.mark.parametrize("custom", ["unknown_type", "custom_geometry"])
def test_unknown_missing_vml_definition_still_rejected(custom):
    source = vml_source()
    definition = next(source.element.iter(f"{{{V}}}shapetype"))
    definition.getparent().remove(definition)
    shape = next(source.element.iter(f"{{{V}}}shape"))
    if custom == "unknown_type":
        shape.set("type", "#unknown")
    else:
        shape.set("path", "m1,1l10,10xe")
    with pytest.raises(WordQuestionExportError, match="形状定义缺失"):
        export_word_questions("图片练习", [item(source, [1])])


def test_ole_complete_content_fallback_becomes_static_image_with_warning():
    result = export_word_questions("旧版公式", [item(vml_source(ole=True), [1])])
    student = output_doc(result)
    assert not list(student.element.iter(f"{{{O}}}OLEObject"))
    assert not list(student.element.iter(qn("w:object")))
    assert list(student.element.iter(f"{{{V}}}imagedata"))
    assert any("静态预览图" in warning for warning in result["warnings"])
    with ZipFile(BytesIO(result["student_bytes"])) as archive:
        assert not any("embeddings" in name for name in archive.namelist())


@pytest.mark.parametrize("kwargs", [{"draw_aspect": "Icon"}, {"fallback": False}])
def test_ole_without_complete_content_fallback_is_rejected(kwargs):
    with pytest.raises(WordQuestionExportError, match="预览不完整"):
        export_word_questions("图题", [item(vml_source(ole=True, **kwargs), [1])])


def test_body_index_walker_matches_preparation_preview_in_wrappers():
    from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
        _body_blocks as preview_blocks,
    )
    source = Document()
    source.add_paragraph("前")
    source.add_paragraph("容器中")
    original = source.paragraphs[1]._p
    wrapper = OxmlElement("w:sdt")
    wrapper.append(OxmlElement("w:sdtPr"))
    contents = OxmlElement("w:sdtContent")
    contents.append(deepcopy(original))
    wrapper.append(contents)
    original.getparent().replace(original, wrapper)
    source.add_table(rows=1, cols=1).cell(0, 0).text = "表格"
    assert list(_body_blocks(source.element.body)) == list(preview_blocks(source.element.body))
    # The current preview also indexes the unrecognized sdtPr metadata as a
    # visible gap.  Reuse those exact indices; never silently renumber blocks.
    result = output_doc(export_word_questions("练习", [item(source, [3, 4])]))
    assert "容器中" in text(result)
    assert "前" not in text(result)
    assert result.tables[0].cell(0, 0).text == "表格"


def test_generated_title_rules_are_removed_without_changing_source_decoration(monkeypatch):
    from integrations.deeptutor_shchem_v1 import (
        desktop_word_question_export as exporter,
    )

    real_document = exporter.Document

    def decorated_template(*args, **kwargs):
        document = real_document(*args, **kwargs)
        if not args and not kwargs:
            for name in ("Title", "Subtitle"):
                border = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:color"), "4F81BD")
                border.append(bottom)
                document.styles[name].element.get_or_add_pPr().append(border)
                document.styles[name].font.underline = True
        return document

    source = Document()
    paragraph = source.add_paragraph("原题蓝色题号")
    paragraph.runs[0]._r.get_or_add_rPr().append(parse_xml(f'<w:color {nsdecls("w")} w:val="0000FF"/>'))
    selection = item(source, [1])
    monkeypatch.setattr(exporter, "Document", decorated_template)
    result = output_doc(export_word_questions("课堂练习", [selection]))
    for name in ("Title", "Subtitle"):
        assert not list(result.styles[name].element.iter(qn("w:pBdr")))
        assert result.styles[name].font.underline is False
    for paragraph in result.paragraphs[:2]:
        assert not list(paragraph._p.iter(qn("w:pBdr")))
        assert all(run.font.underline is False for run in paragraph.runs)
    original = next(p for p in result.paragraphs if p.text == "1. 原题蓝色题号")
    assert original._p.xpath(".//w:color/@w:val") == ["0000FF"]


def test_question_chain_keeps_options_together_but_releases_before_answer():
    source = Document()
    for value in ("题干", "装置说明", "A 选项", "B 选项", "【答案】参考内容"):
        source.add_paragraph(value)
    result = output_doc(export_word_questions("课堂练习", [item(source, [1, 2, 3, 4], [5])]), "teacher")
    paragraphs = {p.text: p for p in result.paragraphs}
    for value in ("1. 题干", "装置说明", "A 选项"):
        assert paragraphs[value].paragraph_format.keep_with_next is True
        assert paragraphs[value].paragraph_format.keep_together is True
    assert paragraphs["B 选项"].paragraph_format.keep_with_next is False
    assert paragraphs["【答案】参考内容"].paragraph_format.keep_with_next is None
    assert not result.element.xpath(".//w:br[@w:type='page']")


def test_terminal_question_table_releases_last_cell_paragraphs_only():
    source = Document()
    source.add_paragraph("题干")
    table = source.add_table(rows=2, cols=2)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.text = f"选项{row_index}{col_index}"
            cell.add_paragraph("说明")
    source.add_paragraph("【答案】参考")
    student = output_doc(export_word_questions("课堂练习", [item(source, [1, 2], [3])]))
    assert next(p for p in student.paragraphs if p.text == "1. 题干").paragraph_format.keep_with_next is True
    exported = student.tables[0]
    for cell in exported.rows[0].cells:
        assert all(p.paragraph_format.keep_with_next is True for p in cell.paragraphs)
    for cell in exported.rows[-1].cells:
        assert cell.paragraphs[0].paragraph_format.keep_with_next is True
        assert cell.paragraphs[-1].paragraph_format.keep_with_next is False
    assert not student.element.xpath(".//w:cantSplit")


def test_question_table_between_paragraphs_keeps_following_content():
    source = Document()
    source.add_paragraph("题干")
    source.add_table(rows=1, cols=2).cell(0, 0).text = "选项"
    source.add_paragraph("最后一项要求")
    student = output_doc(export_word_questions("课堂练习", [item(source, [1, 2, 3])]))
    assert all(p.paragraph_format.keep_with_next is True for cell in student.tables[0].rows[0].cells for p in cell.paragraphs)
    assert next(p for p in student.paragraphs if p.text == "最后一项要求").paragraph_format.keep_with_next is False
