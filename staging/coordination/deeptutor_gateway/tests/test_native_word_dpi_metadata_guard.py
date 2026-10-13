"""Exact authored DPI metadata must not hide a complete embedded source image."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO

import pytest
from docx import Document
from docx.shared import Inches
from lxml import etree
from PIL import Image

from integrations.deeptutor_shchem_v1.desktop_preparation_sources import _body_blocks, _word_images
from integrations.deeptutor_shchem_v1.offline_word_image_label_review import (
    OfflineWordLabelReviewError, _LOCAL_DPI_URI, _objects, _q,
)


def source(value="0"):
    raw = BytesIO()
    Image.new("RGB", (24, 18), "white").save(raw, "PNG")
    document = Document()
    paragraph = document.add_paragraph("H₂N—CH₂—COOH ")
    paragraph.add_run().add_picture(BytesIO(raw.getvalue()), width=Inches(0.3))
    element = list(_body_blocks(document._element.body))[0]
    blip = next(element.iter(_q("a:blip")))
    extensions = etree.SubElement(blip, _q("a:extLst"))
    extension = etree.SubElement(extensions, _q("a:ext"), uri=_LOCAL_DPI_URI)
    dpi = etree.SubElement(extension, _q("a14:useLocalDpi"), **({"val": value} if value is not None else {}))
    return document, element, blip, extensions, extension, dpi, raw.getvalue()


@pytest.mark.parametrize("value", ["0", "1", "false", "true", None])
def test_supported_boolean_dpi_extension_preserves_xml_and_original_bytes(value):
    document, element, _blip, _extensions, _extension, _dpi, original = source(value)
    assets = _word_images([element], document)
    before = etree.tostring(element)
    objects = _objects(element, document, 1, assets)
    assert etree.tostring(element) == before
    assert len(objects) == 1
    assert objects[0]["image"]["sha256"] == sha256(original).hexdigest()
    assert objects[0]["image"]["relationship_id"] == _blip.get(_q("r:embed"))
    assert objects[0]["coverage"] == "embedded_image"


@pytest.mark.parametrize("mutation", [
    "unknown_uri", "unknown_child", "extra_extension",
    "invalid_boolean", "extra_dpi_attribute", "extra_extension_attribute",
    "extra_list_attribute", "nested_dpi_content", "nonempty_text",
    "misplaced_list", "misplaced_dpi", "misplaced_extension", "brightness_effect", "floating_anchor",
    "cropped_image", "rotated_image", "flipped_image", "external_link",
])
def test_unknown_or_transformed_image_metadata_remains_blocked(mutation):
    document, element, blip, extensions, extension, dpi, _raw = source()
    if mutation == "unknown_uri": extension.set("uri", "{UNKNOWN}")
    elif mutation == "unknown_child": etree.SubElement(extension, _q("a:lum"))
    elif mutation == "extra_extension": extensions.append(deepcopy(extension))
    elif mutation == "invalid_boolean": dpi.set("val", "on")
    elif mutation == "extra_dpi_attribute": dpi.set("ignored", "1")
    elif mutation == "extra_extension_attribute": extension.set("ignored", "1")
    elif mutation == "extra_list_attribute": extensions.set("ignored", "1")
    elif mutation == "nested_dpi_content": etree.SubElement(dpi, _q("a:ext"))
    elif mutation == "nonempty_text": dpi.text = "not metadata"
    elif mutation == "misplaced_list":
        blip.remove(extensions)
        next(element.iter(_q("pic:spPr"))).append(extensions)
    elif mutation == "misplaced_dpi":
        extension.remove(dpi)
        next(element.iter(_q("pic:spPr"))).append(dpi)
    elif mutation == "misplaced_extension":
        blip.remove(extensions)
        extensions.remove(extension)
        next(element.iter(_q("pic:spPr"))).append(extension)
    elif mutation == "brightness_effect": etree.SubElement(blip, _q("a:lum"), bright="50000")
    elif mutation == "floating_anchor": next(element.iter(_q("wp:inline"))).tag = _q("wp:anchor")
    elif mutation == "cropped_image":
        etree.SubElement(next(element.iter(_q("pic:blipFill"))), _q("a:srcRect"), l="1")
    elif mutation == "rotated_image": next(element.iter(_q("a:xfrm"))).set("rot", "60000")
    elif mutation == "flipped_image": next(element.iter(_q("a:xfrm"))).set("flipH", "1")
    elif mutation == "external_link": blip.set(_q("r:link"), "rIdExternal")
    assets = _word_images([element], document)
    before = etree.tostring(element)
    with pytest.raises(OfflineWordLabelReviewError):
        _objects(element, document, 1, assets)
    assert etree.tostring(element) == before


def test_plain_original_picture_keeps_existing_admission():
    document, element, blip, extensions, _extension, _dpi, _raw = source()
    blip.remove(extensions)
    assert len(_objects(element, document, 1, _word_images([element], document))) == 1
