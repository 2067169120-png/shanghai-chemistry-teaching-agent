"""Result equivalence and work-count tests, never timing thresholds on shared CI."""
from copy import deepcopy
from concurrent.futures import ThreadPoolExecutor
import random
from types import SimpleNamespace
import pytest
from test_word_question_filters import _row, _catalog
from test_desktop_visual_import_facade import desktop_paths, _facade, FakeProviderStore
from integrations.deeptutor_shchem_v1.desktop_question_explorer import personal_results
from integrations.deeptutor_shchem_v1.desktop_explorer_index import PersonalSearchIndex, SearchCancelled
from integrations.deeptutor_shchem_v1.desktop_word_image_batch import load_word_image_batch
from integrations.deeptutor_shchem_v1.desktop_word_questions import WordQuestionError


def source_catalog(size=200):
    rows = [_row(key=f'q{i}', source=f's{i%5}', knowledge=(f'K{1+i%19:02d}',),
                 mappings=(('S1','S3') if i%3==0 else ('S2',)),
                 exam=('first_mock','second_mock','school_exam')[i%3]) for i in range(size)]
    for i in range(0, size, 11):
        rows[i]['revision'] = 'changed-source-range'
    return {'items':rows, 'attribute_catalog':_catalog(), 'warnings':[]}


def comparable(result):
    return {k:v for k,v in result.items() if k != 'performance'}


def test_two_hundred_fixed_random_requests_match_the_original():
    catalog = source_catalog(); before = deepcopy(catalog)
    index = PersonalSearchIndex(catalog,'word_native')
    rng = random.Random(195)
    for _ in range(200):
        selection = {g:rng.sample(v,rng.randrange(len(v)+1)) for g,v in {
            'source':['s1','s3','unknown'], 'knowledge':['K01','K09','unknown'],
            'book':['V1','V2','unknown'], 'exam':['first_mock','second_mock','unknown']}.items()}
        selection['knowledge_mode'] = rng.choice(['any','all'])
        query = rng.choice(['','CONTEXT','OMEGA','第一册','原题','ABSENT'])
        page = rng.randrange(3)
        assert comparable(index.search(selection,query,page,8)) == personal_results(catalog,'word_native',selection,query,page,8)
    assert catalog == before


def test_paging_reuses_matching_ids_but_has_all_facets():
    catalog=source_catalog();index=PersonalSearchIndex(catalog,'word_native')
    first=index.search({},'',0,8);second=index.search({},'',1,8)
    assert not first['performance']['query_cache_hit']
    assert second['performance']['query_cache_hit']
    assert second['entries'][0]['key']=='q8'
    assert second['facets']==first['facets']
    assert second['entries'][0]['payload'] is catalog['items'][8]


def test_cancellation_does_not_publish_partial_query_cache():
    index=PersonalSearchIndex(source_catalog(),'word_native');calls=[]
    def cancelled():
        calls.append(1);return len(calls)>=3
    with pytest.raises(SearchCancelled):index.search({},'CONTEXT',cancelled=cancelled)
    assert not index._matches
    assert index.search({},'CONTEXT')['total']==200


def test_cancelled_index_build_never_returns_a_partial_index():
    with pytest.raises(SearchCancelled):PersonalSearchIndex(source_catalog(),'word_native',cancelled=lambda:True)


def test_result_cache_is_bounded_and_separate_per_catalogue():
    index=PersonalSearchIndex(source_catalog(),'word_native')
    for i in range(10):index.search({},str(i))
    assert len(index._matches)==4
    changed=source_catalog(1);changed['items'][0]['question_blocks'][0]['text']='new exact text'
    second=PersonalSearchIndex(changed,'word_native')
    assert second.search({},'new exact text')['total']==1
    assert index.search({},'new exact text')['total']==0


def test_stale_labels_and_cross_branch_curriculum_match_old_rules():
    c=source_catalog(3);index=PersonalSearchIndex(c,'word_native')
    for selection in ({'knowledge':['K01']},{'book':['V2'],'section':['section:["V1","C1","S1"]']},
                      {'knowledge':7},{'knowledge_mode':'invalid'}):
        assert comparable(index.search(selection,''))==personal_results(c,'word_native',selection,'')


def test_answer_text_not_in_projection_or_results():
    index=PersonalSearchIndex(source_catalog(),'word_native')
    assert all('OMEGA' not in r[3] and 'OMEGA' not in r[4] for r in index._rows)
    assert index.search({},'OMEGA')['total']==0


def test_parallel_filters_keep_their_own_query():
    c=source_catalog();index=PersonalSearchIndex(c,'word_native')
    queries=['ALPHA','CONTEXT','OMEGA','ABSENT']*3
    with ThreadPoolExecutor(4) as pool:
        results=list(pool.map(lambda q: index.search({},q),queries))
    assert [r['total'] for r in results]==[personal_results(c,'word_native',{},q)['total'] for q in queries]


def test_visual_catalogue_retains_old_matching_and_answer_exclusion():
    c={'items':[{'key':'v1','source_name':'合成图','title':'题目', 'question_text':'stem',
                  'shared_text':'context','answer_text':'OMEGA'}], 'filter_options':{},'warnings':[]}
    index=PersonalSearchIndex(c,'visual_native')
    for query in ('','stem','context','OMEGA'):
        assert comparable(index.search({},query))==personal_results(c,'visual_native',{},query)


def make_images_fixture(paths, tmp_path, count=6, questions=12):
    from io import BytesIO
    from docx import Document
    from PIL import Image
    doc=Document()
    for n in range(questions):
        doc.add_paragraph(f'【例{n+1}】合成多图题，观察图示并作答。')
        for i in range(count if n==0 else 1):
            data=BytesIO();Image.new('RGB',(240,160),(n%255,i*30%255,80)).save(data,format='PNG')
            doc.add_picture(BytesIO(data.getvalue()))
        doc.add_paragraph('【答案】合成答案，不进入题面关键词。')
    path=tmp_path/'合成多图讲义.docx';doc.save(path)
    facade=_facade(paths,FakeProviderStore(configured=False))
    facade.save_visual_import_batch(handout_files=(path,),source_type='教师讲义')
    row=facade.word_question_catalog()['items'][0]
    ids=[a['asset_id'] for b in row['question_blocks'] for a in b.get('assets',[])]
    assert len(ids)==count
    return facade,row,ids


def test_batch_images_match_old_bytes_with_one_resolve(desktop_paths,tmp_path,monkeypatch):
    facade,row,ids=make_images_fixture(desktop_paths,tmp_path)
    expected={i:facade.word_question_image(row['key'],row['revision'],i) for i in ids}
    service=facade._word_questions();original=service._resolve;calls=[]
    def resolve(*args,**kwargs):calls.append(1);return original(*args,**kwargs)
    monkeypatch.setattr(service,'_resolve',resolve)
    outputs=[];stats=load_word_image_batch(facade,row['key'],row['revision'],ids,progress=outputs.append)
    assert len(calls)==1 and stats['completed']==len(ids) and stats['failed']==0
    assert {v['asset_id']:v['result'] for v in outputs}==expected
    facade.shutdown()


def test_batch_rechecks_revision_before_reading_images(desktop_paths,tmp_path):
    facade,row,ids=make_images_fixture(desktop_paths,tmp_path)
    outputs=[]
    with pytest.raises(WordQuestionError):load_word_image_batch(facade,row['key'],'stale',ids,progress=outputs.append)
    assert outputs==[]
    facade.shutdown()


def test_batch_rejects_nonmember_before_emitting_any_image(desktop_paths,tmp_path):
    facade,row,ids=make_images_fixture(desktop_paths,tmp_path)
    outputs=[]
    with pytest.raises(WordQuestionError):load_word_image_batch(facade,row['key'],row['revision'],ids+['bad-id'],progress=outputs.append)
    assert outputs==[]
    facade.shutdown()


def test_cancelled_batch_does_not_read_sources():
    facade=SimpleNamespace(_word_questions=lambda: pytest.fail('must not read'))
    with pytest.raises(SearchCancelled):load_word_image_batch(facade,'q','v',['a'],cancelled=lambda:True)


def test_fallback_batch_deduplicates_and_reports_per_image_errors():
    calls=[]
    def read(k,r,a):
        calls.append(a)
        if a=='bad':raise OSError('private path and key must not leak')
        return {'bytes':b'ok'}
    outputs=[]
    stats=load_word_image_batch(SimpleNamespace(word_question_image=read),'q','v',['a','a','bad'],progress=outputs.append)
    assert calls==['a','bad'] and stats['failed']==1
    assert outputs[-1]=={'asset_id':'bad','failed':True}
    assert 'private' not in str(stats)


def test_cancelled_cold_read_is_reused_not_repeated():
    from integrations.deeptutor_shchem_v1.desktop_explorer_index import PersonalCatalogSession
    session=PersonalCatalogSession('word_native');calls=[];stopped=[False]
    def load():calls.append(1);stopped[0]=True;return source_catalog(5)
    with pytest.raises(SearchCancelled):session.get(load,cancelled=lambda:stopped[0])
    assert session.catalog is not None and session.index is None
    index,metrics=session.get(load)
    assert calls==[1] and index.search({},'')['total']==5 and metrics['catalogue_reused']


def test_parallel_first_queries_share_a_single_catalogue_read():
    from integrations.deeptutor_shchem_v1.desktop_explorer_index import PersonalCatalogSession
    import time
    session=PersonalCatalogSession('word_native');calls=[]
    def load():calls.append(1);time.sleep(.03);return source_catalog(5)
    with ThreadPoolExecutor(4) as pool:
        results=list(pool.map(lambda _:session.get(load)[0],range(4)))
    assert calls==[1] and all(value is results[0] for value in results)
