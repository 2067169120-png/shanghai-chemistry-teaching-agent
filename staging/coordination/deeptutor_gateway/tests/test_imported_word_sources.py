import hashlib
import io
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.shared import Inches
from test_desktop_visual_import_facade import (
    FakeProviderStore,
    FakeVisualTransport,
    _facade,
    _png,
)
from test_desktop_visual_import_facade import (
    desktop_paths as desktop_paths,  # noqa: PLC0414 - pytest fixture registration
)

from integrations.deeptutor_shchem_v1.desktop_facade import DesktopFacadeError
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
    PreparationSourceError,
)


def word_bytes(*, hybrid=False):
    doc = Document()
    doc.add_heading("电离与离子方程式", 1)
    doc.add_paragraph("1．写出对应的化学式与电荷。")
    p = doc.add_paragraph("SO")
    style = doc.styles.add_style("ChemicalSub", WD_STYLE_TYPE.CHARACTER)
    style.font.subscript = True
    p.add_run("4", style=style)
    p.add_run("2−").font.superscript = True
    p = doc.add_paragraph("原生方程式：")
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "H₂O ⇌ H⁺ + OH⁻"
    run.append(text)
    math.append(run)
    p._p.append(math)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "条件"
    table.cell(0, 1).text = "水溶液"
    if hybrid:
        p = doc.add_paragraph("原图条件不可省略：")
        p.add_run().add_picture(io.BytesIO(_png("green")), width=Inches(0.5))
        doc.add_paragraph("旧公式：").add_run()._r.append(OxmlElement("w:object"))
    result = io.BytesIO()
    doc.save(result)
    return result.getvalue()


@pytest.mark.parametrize("hybrid", [False, True])
def test_archived_word_to_preparation_preserves_original_name_content_and_gaps(
    desktop_paths, tmp_path, hybrid,
):
    source_path = tmp_path / "电离讲义（教师版）.docx"
    original = word_bytes(hybrid=hybrid)
    source_path.write_bytes(original)
    provider = FakeProviderStore(configured=False)
    transport = FakeVisualTransport()
    facade = _facade(desktop_paths, provider, transport=transport)
    saved = facade.save_visual_import_batch(handout_files=(source_path,), source_type="教师讲义")
    assert saved.native_quick_count == (0 if hybrid else 1)
    assert saved.visual_queue_count == (1 if hybrid else 0)
    # This is our temporary fixture, not a user's file. Prove source moves do
    # not break the stable import -> reference path.
    source_path.unlink()
    facade = _facade(desktop_paths, provider, transport=transport)
    assert saved.batch_id in {item.batch_id for item in facade.list_imported_word_batches()}
    source = facade.imported_word_sources(saved.batch_id)[0]
    preview = facade.imported_word_preview(saved.batch_id, source["source_id"])
    combined = "\n".join(block["text"] for block in preview["blocks"])
    assert "SO_{4}^{2−}" in combined
    assert "H₂O ⇌ H⁺ + OH⁻" in combined
    assert "条件" in combined and "水溶液" in combined
    assert preview["source_name"] == source_path.name
    assert preview["source_sha256"] == hashlib.sha256(original).hexdigest()
    ref = facade.imported_word_reference(
        saved.batch_id, source["source_id"], preview["source_sha256"],
        1, len(preview["blocks"]), preview["revision"],
    )
    assert source_path.name in ref["materials"]
    assert "SO_{4}^{2−}" in ref["materials"]
    assert "H₂O ⇌ H⁺ + OH⁻" in ref["materials"]
    assert str(desktop_paths.state_root) not in ref["materials"]
    if hybrid:
        assert "【待查看原文：嵌入对象或旧公式】" in ref["materials"]
        assert ref["warnings"]
        assert len(preview["assets"]) == 1
        image = facade.imported_word_asset(saved.batch_id, source["source_id"], preview["assets"][0]["asset_id"])
        assert image["bytes"] == _png("green")
        assert preview["assets"][0]["block_index"] == 6
    else:
        assert not ref["warnings"]
        records = list((desktop_paths.state_root / "word-handout-import" / "documents").glob("*.json"))
        text = records[0].read_text(encoding="utf-8")
        assert "SO_{4}^{2−}" in text
    assert provider.borrow_calls == transport.calls == 0


def test_archived_source_hash_range_and_preview_revision_are_checked(desktop_paths, tmp_path):
    source_path = tmp_path / "source.docx"
    source_path.write_bytes(word_bytes())
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    saved = facade.save_visual_import_batch(handout_files=(source_path,), source_type="讲义")
    source_id = facade.imported_word_sources(saved.batch_id)[0]["source_id"]
    preview = facade.imported_word_preview(saved.batch_id, source_id)
    with pytest.raises(DesktopFacadeError, match="变化"):
        facade.imported_word_reference(saved.batch_id, source_id, preview["source_sha256"], 1, 1, "stale")
    with pytest.raises(PreparationSourceError, match="范围"):
        facade.imported_word_reference(saved.batch_id, source_id, preview["source_sha256"], 1, 999, preview["revision"])
    with pytest.raises(DesktopFacadeError, match="未找到"):
        facade.imported_word_preview(saved.batch_id, "different-source")
    archive = desktop_paths.state_root / "visual-import-v2" / "sources" / f"{preview['source_sha256']}.docx"
    archive.write_bytes(b"changed fixture")
    with pytest.raises(DesktopFacadeError, match="变化"):
        facade.imported_word_preview(saved.batch_id, source_id)


def test_large_reference_is_not_silently_truncated(desktop_paths, tmp_path):
    doc = Document()
    doc.add_paragraph("条件不可丢失" * 4100)
    source = tmp_path / "large.docx"
    doc.save(source)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    saved = facade.save_visual_import_batch(handout_files=(source,), source_type="讲义")
    sid = facade.imported_word_sources(saved.batch_id)[0]["source_id"]
    preview = facade.imported_word_preview(saved.batch_id, sid)
    with pytest.raises(PreparationSourceError, match="未截断"):
        facade.imported_word_reference(saved.batch_id, sid, preview["source_sha256"], 1, 1, preview["revision"])


def test_metadata_list_is_lazy_but_opening_checks_selected_archive(desktop_paths, tmp_path, monkeypatch):
    first, second = tmp_path / "first.docx", tmp_path / "second.docx"
    first.write_bytes(word_bytes())
    second.write_bytes(word_bytes(hybrid=True))
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    saved = facade.save_visual_import_batch(handout_files=(first, second), source_type="完整教案")
    original_restore = facade._restore_visual_import_sources

    def refuse_restore(*args, **kwargs):
        raise AssertionError("Listing must not load or parse all original documents")

    monkeypatch.setattr(facade, "_restore_visual_import_sources", refuse_restore)
    rows = facade.imported_word_sources(saved.batch_id)
    assert len(rows) == 2
    monkeypatch.setattr(facade, "_restore_visual_import_sources", original_restore)
    first_preview = facade.imported_word_preview(saved.batch_id, rows[0]["source_id"])
    second_path = Path(facade.imported_word_path(saved.batch_id, rows[1]["source_id"]))
    second_path.write_bytes(b"bad source")
    assert facade.imported_word_preview(saved.batch_id, rows[0]["source_id"]) == first_preview
    with pytest.raises(DesktopFacadeError, match="变化"):
        facade.imported_word_path(saved.batch_id, rows[1]["source_id"])
