"""Actual Word object order, answer isolation and selectable native Qt flow."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO

import pytest
from docx import Document
from docx.shared import Mm
from PIL import Image
from test_desktop_visual_import_facade import FakeProviderStore, _facade, desktop_paths as desktop_paths
from test_word_question_dialog import _isolated_qt_app, _Tasks

from integrations.deeptutor_shchem_v1.desktop_word_questions import WordQuestionError


def _png(color):
    output = BytesIO(); Image.new('RGB', (160, 80), color).save(output, format='PNG')
    return output.getvalue()


@pytest.fixture
def native(desktop_paths, tmp_path):
    doc = Document()
    p = doc.add_paragraph('1．选择正确的图。NH'); p.add_run('3').font.subscript = True
    p.add_run('与文字 <img src="https://example.invalid/test"> 保留。')
    p = doc.add_paragraph('A．'); p.add_run().add_picture(BytesIO(_png('blue')), width=Mm(25))
    p.add_run(' B．'); p.add_run().add_picture(BytesIO(_png('red')), width=Mm(30))
    doc.add_paragraph('【答案】B')
    p = doc.add_paragraph('【解析】仅教师答案'); p.add_run().add_picture(BytesIO(_png('blue')), width=Mm(18))
    p.add_run('Na'); p.add_run('+').font.superscript = True
    path = tmp_path / '原生图文顺序.docx'; doc.save(path)
    provider = FakeProviderStore(configured=False); facade = _facade(desktop_paths, provider)
    facade.save_visual_import_batch(handout_files=(path,), source_type='教师讲义')
    catalog = facade.word_question_catalog()
    row = catalog['items'][0]
    assert len(catalog['items']) == 1
    return facade, row, provider


def test_actual_source_inline_tokens_keep_option_and_script_order_without_writes(native):
    facade, row, provider = native
    before = facade.state_store.snapshot()
    result = facade.word_question_reading_layout(row['key'], row['revision'])
    first, options = result['paragraphs']
    assert [(t['text'], t['script']) for t in first['tokens'] if t['kind'] == 'text' and t['script'] != 'baseline'] == [('3', 'subscript')]
    assert [t['kind'] for t in options['tokens']] == ['text', 'image', 'text', 'image']
    assert [t['text'] for t in options['tokens'] if t['kind'] == 'text'] == ['A．', ' B．']
    assert [round(t['width_mm']) for t in options['tokens'] if t['kind'] == 'image'] == [25, 30]
    assert result['scope'] == 'question' and len(result['images']) == 2
    assert '仅教师答案' not in str(result['paragraphs'])
    assert facade.state_store.snapshot() == before and provider.borrow_calls == 0
    assert not (facade.paths.task_root / 'preparation-v1/images').exists()


def test_answer_only_returns_its_own_occurrences_even_when_pixels_reused(native):
    facade, row, _provider = native
    question = facade.word_question_reading_layout(row['key'], row['revision'])
    answer = facade.word_question_reading_layout(row['key'], row['revision'], scope='answer')
    assert len(answer['images']) == 1 and set(answer['images']).isdisjoint(question['images'])
    assert '仅教师答案' in str(answer['paragraphs'])
    assert [t['text'] for p in answer['paragraphs'] for t in p['tokens'] if t.get('script') == 'superscript'] == ['+']
    assert next(iter(answer['images'].values()))['bytes'] == next(iter(question['images'].values()))['bytes']


@pytest.mark.parametrize('scope', ['student', '', None])
def test_unknown_scope_cannot_expose_a_combined_question_and_answer(native, scope):
    facade, row, _provider = native
    with pytest.raises(WordQuestionError): facade.word_question_reading_layout(row['key'], row['revision'], scope=scope)


def test_stale_revision_is_rejected_before_pixels_are_returned(native):
    facade, row, _provider = native
    with pytest.raises(WordQuestionError): facade.word_question_reading_layout(row['key'], '0' * 64)


@pytest.mark.parametrize('width', [900, 320])
def test_qt_reader_escapes_source_html_retains_pixels_and_fits_width(native, width):
    from PySide6.QtCore import QUrl
    from PySide6.QtGui import QTextDocument
    from integrations.deeptutor_shchem_v1.desktop_workbench.word_inline_reader import WordInlineReader
    facade, row, _provider = native
    value = facade.word_question_reading_layout(row['key'], row['revision'])
    frozen = deepcopy(value)
    with _isolated_qt_app() as app:
        reader = WordInlineReader(value); reader.resize(width, 300); reader.show()
        for _ in range(5): app.processEvents()
        plain = reader.toPlainText()
        assert 'NH3' in plain and '<img src="https://example.invalid/test">' in plain
        assert '待查看原文' not in plain and '仅教师答案' not in plain
        assert reader.height() >= reader.document().size().height()
        cursor = reader.document().begin(); pictures = []
        while cursor.isValid():
            iterator = cursor.begin()
            while not iterator.atEnd():
                fragment = iterator.fragment()
                if fragment.charFormat().isImageFormat():
                    pictures.append(fragment.charFormat().toImageFormat())
                iterator += 1
            cursor = cursor.next()
        assert len(pictures) == 2
        assert all(p.width() <= reader.viewport().width() for p in pictures)
        for asset_id, asset in value['images'].items():
            image = reader.document().resource(QTextDocument.ResourceType.ImageResource, QUrl('word-source-image:' + asset_id))
            assert image is not None and not image.isNull() and image.pixelColor(0, 0).name() in {'#0000ff', '#ff0000'}
        assert not reader.openLinks() and not reader.openExternalLinks()
        reader.close()
    assert value == frozen


def test_tampered_pixel_bytes_do_not_become_a_successful_reader(native):
    from integrations.deeptutor_shchem_v1.desktop_workbench.word_inline_reader import WordInlineReader
    facade, row, _provider = native
    value = facade.word_question_reading_layout(row['key'], row['revision'])
    next(iter(value['images'].values()))['bytes'] = _png('black')
    with _isolated_qt_app():
        with pytest.raises(ValueError, match='来源位置'): WordInlineReader(value)


def test_native_dialog_displays_inline_question_then_separate_answer(native):
    from integrations.deeptutor_shchem_v1.desktop_workbench.word_question_dialog import WordQuestionDialog
    from integrations.deeptutor_shchem_v1.desktop_workbench.word_inline_reader import WordInlineReader
    facade, row, _provider = native
    with _isolated_qt_app() as app:
        tasks = _Tasks(); dialog = WordQuestionDialog(facade, tasks)
        dialog.show(); tasks.flush(); app.processEvents(); tasks.flush(); app.processEvents()
        reader = dialog._panels[0].findChild(WordInlineReader)
        assert reader is not None and '仅教师答案' not in reader.toPlainText()
        assert dialog.current_question_readiness()[0]
        dialog.tabs.setCurrentIndex(1); tasks.flush(); app.processEvents()
        answer = dialog._panels[1].findChild(WordInlineReader)
        assert answer is not None and '仅教师答案' in answer.toPlainText()
        dialog.reject()


def test_incorrect_source_binding_falls_back_to_full_original_preview(native):
    from integrations.deeptutor_shchem_v1.desktop_workbench.word_question_dialog import WordQuestionDialog
    from integrations.deeptutor_shchem_v1.desktop_workbench.word_inline_reader import WordInlineReader
    facade, row, _provider = native
    original = facade.word_question_reading_layout
    def changed(*args, **kwargs):
        value = original(*args, **kwargs); value['source_sha256'] = '0' * 64; return value
    facade.word_question_reading_layout = changed
    with _isolated_qt_app() as app:
        tasks = _Tasks(); dialog = WordQuestionDialog(facade, tasks)
        dialog.show(); tasks.flush(); app.processEvents(); tasks.flush(); app.processEvents()
        assert dialog._panels[0].findChild(WordInlineReader) is None
        labels = [w.text() for w in dialog._panels[0].findChildren(__import__('PySide6.QtWidgets', fromlist=['QLabel']).QLabel)]
        assert any('不能按原文位置排版' in text for text in labels)
        assert any('A．' in text and 'B．' in text for text in labels)
        dialog.reject()
