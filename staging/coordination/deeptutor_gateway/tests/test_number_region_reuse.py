from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image
from test_raster_numbers import fixture_documents, edit_for
from test_word_question_export import item, data
from integrations.deeptutor_shchem_v1.desktop_state import DesktopStateStore, DesktopStateError
from integrations.deeptutor_shchem_v1.desktop_raster_numbers import image_catalog, RasterNumberError
from integrations.deeptutor_shchem_v1.desktop_number_regions import NumberRegionStore, suggest_number
from integrations.deeptutor_shchem_v1.desktop_mixed_paper_export import build_mixed_paper_docx


def store_image(tmp_path):
    state = DesktopStateStore(tmp_path)
    docs, raw = fixture_documents()
    image = image_catalog(docs)['images'][0]
    return state, NumberRegionStore(state), image, raw


def test_reading_empty_library_does_not_create_a_file(tmp_path):
    state, store, image, _ = store_image(tmp_path)
    assert store.load(image) == []
    assert not state.path.exists()


def test_remember_survives_reopen_and_never_saves_a_paper_number(tmp_path):
    state, store, image, raw = store_image(tmp_path)
    state.save_draft('d', {'materials': '教师原材料'})
    state.add_to_basket({'key': 'q', 'title': '原题'})
    before = state.snapshot()
    store.remember(image, [edit_for(raw, 83)])
    reopened = NumberRegionStore(DesktopStateStore(tmp_path))
    rows = reopened.load(image)
    assert rows == [{'box': [10,10,42,37], 'punctuation': '.'}]
    after = state.snapshot()
    assert after['basket'] == before['basket'] and after['drafts'] == before['drafts']
    assert 'number' not in after['scan_number_regions']['images'][image['image_sha256']]['regions'][0]
    assert 'data' not in repr(after['scan_number_regions'])
    rows[0]['box'][0] = 99
    assert reopened.load(image)[0]['box'][0] == 10


def test_changed_image_does_not_match_old_locations(tmp_path):
    state, store, image, raw = store_image(tmp_path)
    store.remember(image, [edit_for(raw)])
    changed = {**image, 'image_sha256': 'a'*64}
    assert store.load(changed) == []
    with pytest.raises(RasterNumberError, match='尺寸'):
        store.load({**image, 'width': 1000})


@pytest.mark.parametrize('bad', [None, [], {'schema_version': 'future', 'images': {}}, {'images': []}])
def test_invalid_library_does_not_get_silently_overwritten(tmp_path, bad):
    state, store, image, raw = store_image(tmp_path)
    state._update(lambda value: value.__setitem__('scan_number_regions', bad))
    before = state.path.read_bytes()
    with pytest.raises(RasterNumberError): store.remember(image, [edit_for(raw)])
    assert state.path.read_bytes() == before


def test_outside_or_overlapping_boxes_never_saved(tmp_path):
    state, store, image, raw = store_image(tmp_path)
    with pytest.raises(RasterNumberError): store.remember(image, [edit_for(raw, box=[-1,0,20,40])])
    with pytest.raises(RasterNumberError): store.remember(image, [edit_for(raw), edit_for(raw)])
    assert not state.path.exists()


def test_failed_write_keeps_old_locations_and_other_work(tmp_path, monkeypatch):
    state, store, image, raw = store_image(tmp_path)
    store.remember(image, [edit_for(raw)])
    old = state.path.read_bytes()
    def fail(*_): raise OSError('simulated disk error')
    monkeypatch.setattr('integrations.deeptutor_shchem_v1.desktop_state.os.replace', fail)
    with pytest.raises(DesktopStateError): store.remember(image, [edit_for(raw, box=[12,12,50,50])])
    assert state.path.read_bytes() == old


def test_forget_only_one_image_and_does_not_mutate_current_edits(tmp_path):
    state, store, image, raw = store_image(tmp_path)
    edits = [edit_for(raw)]
    before = deepcopy(edits)
    other = {**image, 'image_sha256': 'b'*64}
    store.remember(image, edits)
    store.remember(other, [{**edits[0], 'image_sha256': 'b'*64}])
    store.forget(image)
    assert store.load(image) == [] and store.load(other)
    assert edits == before


def two_word_items(same_image=False, shared=False):
    doc = Document()
    raws = []
    for color in ('white', 'white' if same_image else 'gray'):
        stream = BytesIO(); Image.new('RGB', (80,40), color).save(stream, 'PNG'); raws.append(stream.getvalue())
    if shared: doc.add_picture(BytesIO(raws[0]), width=Inches(1))
    ranges = []
    for index, number in enumerate((27,38)):
        start = len(doc.element.body)
        doc.add_paragraph(f'【例{number}】合成题12.5')
        doc.add_picture(BytesIO(raws[index]), width=Inches(1))
        doc.add_paragraph(f'【答案】{number}. 说明')
        doc.add_picture(BytesIO(raws[index]), width=Inches(1))
        ranges.append((start, start+1, start+2, start+3))
    content = data(doc)
    results = []
    for index, (a,b,c,d) in enumerate(ranges):
        row = item(doc, [a,b], [c,d], context=[1] if shared else [], key=str(index))
        row['source_bytes'] = content; row['source_sha256'] = sha256(content).hexdigest()
        results.append(row)
    return results, raws


def test_suggestions_follow_current_order_without_changing_source_or_result_contract():
    rows, raws = two_word_items()
    before = deepcopy(rows)
    for choices, wanted in ((rows,[1,2]),(rows[::-1],[2,1])):
        hints = {}
        result = build_mixed_paper_docx('合成编号练习', [{'kind':'word_question', 'question':r} for r in choices], image_number_hints=hints)
        assert set(result) == {'student_bytes','teacher_bytes','warnings'}
        images = {r['image_sha256']:r for r in image_catalog(result)['images']}
        for raw, number in zip(raws,wanted):
            assert suggest_number(images[sha256(raw).hexdigest()],hints) == number
    assert rows == before


def test_same_image_in_two_questions_never_gets_one_automatic_number():
    rows, raws = two_word_items(same_image=True)
    hints={}
    result=build_mixed_paper_docx('合成', [{'kind':'word_question','question':r} for r in rows], image_number_hints=hints)
    assert suggest_number(image_catalog(result)['images'][0], hints) is None


def test_shared_material_image_vetoes_same_image_question_hint():
    rows, raws = two_word_items(shared=True)
    hints={}
    result=build_mixed_paper_docx('合成', [{'kind':'word_question','question':r} for r in rows[:1]], image_number_hints=hints)
    assert suggest_number(image_catalog(result)['images'][0], hints) is None


def test_number_hints_never_change_any_exported_document_xml():
    rows,_=two_word_items()
    args=[{'kind':'word_question','question':rows[0]}]
    ordinary=build_mixed_paper_docx('合成',args)
    hinted=build_mixed_paper_docx('合成',args,image_number_hints={})
    for role in ('student','teacher'):
        with ZipFile(BytesIO(ordinary[role+'_bytes'])) as a, ZipFile(BytesIO(hinted[role+'_bytes'])) as b:
            assert a.namelist()==b.namelist()
            for name in a.namelist():
                assert a.read(name)==b.read(name)
