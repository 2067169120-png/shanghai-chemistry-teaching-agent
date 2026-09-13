"""Only synthetic DOCX/state; never execute the real pinned --apply operation."""

from __future__ import annotations

import hashlib
import importlib.util
import io
import json
from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document
from test_desktop_visual_import_facade import _facade, _png
from test_desktop_visual_import_facade import (
    desktop_paths as desktop_paths,  # noqa: PLC0414
)

from integrations.deeptutor_shchem_v1.desktop_word_question_attributes import (
    suggest_attributes,
)
from integrations.deeptutor_shchem_v1.desktop_word_question_index import (
    index_word_questions,
)
from integrations.deeptutor_shchem_v1.desktop_word_questions import RANGES_DRAFT

SCRIPT = Path(__file__).parents[1] / "scripts" / "apply_word_boundary_repairs.py"
spec = importlib.util.spec_from_file_location(
    "apply_word_boundary_repairs_test_module", SCRIPT
)
repair = importlib.util.module_from_spec(spec)
spec.loader.exec_module(repair)
TAXONOMY = {
    "knowledge_points": [
        {"id": "K09", "name": "反应方向、限度和速率进阶"},
        {"id": "K11", "name": "氧化还原与电化学"},
        {"id": "K12", "name": "原子结构与元素性质进阶"},
        {"id": "K15", "name": "有机结构、同分异构与命名"},
    ],
    "nodes": [
        {
            "node_key": "synthetic-organic-section",
            "section_title": "有机化合物的结构",
            "chapter_id": "synthetic-chapter",
            "volume_id": "synthetic-volume",
        }
    ],
}


@pytest.fixture
def corpus(desktop_paths, tmp_path, monkeypatch):
    pytest.importorskip("PySide6")
    contents = {
        "PKG-022": {
            316: "（单选）【例1】请写出合成条件甲。",
            317: "【答案】合成答案甲。",
            340: "（多选）【例2】请判断本合成机理题的选项。",
            344: "B．这里是含决速步骤词语的合成选项，不是真实题目。",
            345: "C．这里是含活化能词语的合成选项，不是真实题目。",
            346: "D．反应1的机理仅作合成测试，符号 Fe^{2+、Fe^{3+、H_{2}O_{2 为测试锚词。",
            347: "【答案】合成答案乙。",
            348: "知识点02 后续知识",
        },
        "PKG-040": {
            265: "【例1】(2025·上海卷合成来源定位)请选择合成电负性比较题的选项（ ）",
            266: "A．C\nB．H\nC．O",
            267: "(1)B",
            268: "(1)这是一段足够完整而且没有提问句的合成解析。",
            269: "知识点02 后续知识",
        },
        "PKG-069": {
            293: "【例1】广义的水解观认为：这里是合成定义。根据上述观点，请回答合成问题。",
            294: "本题给定条件甲。",
            295: "本题给定条件乙。",
            296: "【答案】合成答案。",
            297: "【解析】合成解析。",
            298: "知识点02 后续知识",
        },
        "PKG-092": {
            79: "【变式训练1·变题型请写出合成答案。",
            80: "本题条件。",
            84: "本题末尾条件。",
            85: "【答案】合成答案。",
            86: "【解析】合成解析。",
            87: "合成解析末尾。",
            88: "知识点02 后续知识",
        },
    }
    files, pins = [], []
    for package, _, origins in repair.SOURCES:
        document = Document()
        for index in range(1, max(contents[package]) + 1):
            paragraph = document.add_paragraph(
                "知识点01 有机化合物的结构"
                if index == 1
                else contents[package].get(index, "合成讲义说明。")
            )
            if package == "PKG-069" and index == 295:
                paragraph.add_run().add_picture(io.BytesIO(_png("blue")))
        path = tmp_path / f"{package}-synthetic.docx"
        document.save(path)
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        files.append(path)
        pins.append((package, digest, origins))
    facade = _facade(desktop_paths, repair.NoProviders())
    receipt = facade.save_visual_import_batch(
        handout_files=files, source_type="synthetic local repair fixture"
    )
    monkeypatch.setattr(repair, "BATCH_ID", receipt.batch_id)
    monkeypatch.setattr(repair, "SOURCES", tuple(pins))
    descriptor = facade._saved_visual_import_batch(receipt.batch_id)
    sources = facade._restore_visual_import_sources(descriptor)
    ids = {source.source_sha256: source.effective_source_file_id for source in sources}
    monkeypatch.setattr(repair, "_source_id", lambda package, digest: ids[digest])
    monkeypatch.setattr(
        repair, "load_attribute_catalog", lambda root: deepcopy(TAXONOMY)
    )
    service = facade._word_questions()
    old_rows = []
    for source in sources:
        preview = service.reader.word_preview_bytes(source.content, source.filename)
        assert service.preview_cache.save(source.content, source.filename, preview)
        package = next(
            package for package, digest, _ in pins if digest == source.source_sha256
        )
        for row in index_word_questions(preview):
            if package == "PKG-022" and row["origin_block_start"] == 340:
                continue
            # Simulate the prior parser's changed binding on the first exercise.
            row = deepcopy(row)
            if package == "PKG-022":
                row["revision"] = "old-source-bound-question-revision"
            old_rows.append(
                suggest_attributes(
                    row,
                    {
                        "package_id": package,
                        "source_name": source.filename,
                        "collection_name": "Synthetic collection",
                        "lecture_topic": "合成课题",
                        "document_role": "解析版",
                        "usage_context": "高三一轮复习",
                    },
                    [],
                )
            )
    # Re-seal through the public suggestion builder, not a forged attribute row.
    sample = index_word_questions(
        service.preview_cache.load(sources[0].content, sources[0].filename)
    )[0]
    sample = {**sample, "key": "unrelated-key"}
    old_rows.append(suggest_attributes(sample, {"source_name": "unrelated.docx"}, []))
    service.attribute_store.save_many(old_rows)
    facade.state_store.save_draft("unrelated-draft", {"keep": [1, 2, 3]})
    # The repair must use warm previews, not renderer/extractor fallbacks.
    monkeypatch.setattr(
        service.reader,
        "word_preview_bytes",
        lambda *a, **kw: pytest.fail("reparsed source"),
    )
    monkeypatch.setattr(
        service.preview_cache, "save", lambda *a, **kw: pytest.fail("rewrote cache")
    )
    restore = facade._restore_visual_import_sources
    calls = []

    def selected_only(descriptor, *, source_ids=None):
        assert source_ids is not None
        assert source_ids <= set(ids.values())
        calls.append(set(source_ids))
        return restore(descriptor, source_ids=source_ids)

    monkeypatch.setattr(facade, "_restore_visual_import_sources", selected_only)
    return facade, calls


def run(corpus, tmp_path, *, apply=False, name="report.json"):
    facade, _ = corpus
    return repair.run_repair(
        state_root=facade.paths.state_root,
        report_path=tmp_path / name,
        apply=apply,
        facade_factory=lambda root: facade,
    )


def test_dry_run_is_default_and_has_five_validated_targets_without_writes(
    corpus, tmp_path
):
    facade, calls = corpus
    before = facade.state_store.path.read_bytes()
    report = run(corpus, tmp_path)
    assert report["status"] == "dry_run_validated", report
    assert len(report["targets"]) == 5
    assert not report["write_started"] and report["range_writes"] == []
    assert facade.state_store.path.read_bytes() == before
    assert len(report["preview_revisions"]) == 4
    assert all(
        "automatic_proposal" in target and "source_quality" in target
        for target in report["targets"]
    )
    assert (
        report["provider_calls"] == 0 and report["student_manager_initialized"] is False
    )
    assert (
        report["integrity"]["database_counts_before"]
        == report["integrity"]["database_counts_after"]
    )
    assert calls and all(len(call) <= 4 for call in calls)
    assert "合成答案" not in json.dumps(report, ensure_ascii=False)
    assert not (facade.paths.state_root / ".desktop-workbench.lock").exists()


def test_apply_updates_four_adds_one_and_replay_is_idempotent(corpus, tmp_path):
    report = run(corpus, tmp_path, apply=True)
    assert report["status"] == "applied", report
    assert len(report["range_writes"]) == 3
    assert sorted(row["action"] for row in report["attribute_results"]) == [
        "added",
        "updated",
        "updated",
        "updated",
        "updated",
    ]
    assert all(row["current_question_binding"] for row in report["attribute_results"])
    integrity = report["integrity"]
    assert (
        integrity["database_counts_after"]["attributes"]
        == integrity["database_counts_before"]["attributes"] + 1
    )
    assert (
        integrity["database_counts_after"]["history"]
        == integrity["database_counts_before"]["history"] + 5
    )
    again = run(corpus, tmp_path, apply=True, name="replay.json")
    assert again["status"] == "applied", again
    assert not again["range_writes"]
    assert {row["action"] for row in again["attribute_results"]} == {"unchanged"}
    assert (
        again["integrity"]["database_counts_before"]
        == again["integrity"]["database_counts_after"]
    )


def test_teacher_modified_and_history_are_not_rebound(corpus, tmp_path):
    facade, _ = corpus
    store = facade._word_questions().attribute_store
    plan = repair.build_plan(facade)
    key = plan["targets"][2]["planned"]["key"]
    old = store.get(key)
    teacher = store.save_teacher_edit(
        key,
        {"teacher_note": "合成教师记录，必须保留"},
        expected_revision=old["revision"],
    )
    history = store.history(key)
    report = run(corpus, tmp_path, apply=True)
    assert report["status"] == "applied", report
    assert store.get(key) == teacher and store.history(key) == history
    result = next(row for row in report["attribute_results"] if row["key"] == key)
    assert result["action"] == "teacher_modified_preserved"
    assert not result["current_question_binding"]


def test_different_teacher_range_stops_before_any_apply(corpus, tmp_path):
    facade, _ = corpus
    plan = repair.build_plan(facade)
    item = plan["manual"][0]
    boundaries = {**item["boundaries"], "reviewed_issues": []}
    facade.word_question_update_range(item["key"], item["revision"], **boundaries)
    before = facade.state_store.path.read_bytes()
    report = run(corpus, tmp_path, apply=True)
    assert report["status"] == "blocked_no_apply"
    assert report["error"] == "different_teacher_range_exists_stop"
    assert facade.state_store.path.read_bytes() == before


def test_missing_cache_stops_without_reparse_or_state_write(corpus, tmp_path):
    facade, _ = corpus
    service = facade._word_questions()
    package, digest, _ = repair.SOURCES[0]
    path = service.preview_cache._path(digest, f"{package}-synthetic.docx")
    path.unlink()  # Only the disposable synthetic test cache.
    before = facade.state_store.path.read_bytes()
    report = run(corpus, tmp_path)
    assert report["status"] == "blocked_no_apply"
    assert report["error"] == "verified_warm_preview_required_no_reparse"
    assert facade.state_store.path.read_bytes() == before and not path.exists()


def test_live_desktop_lock_is_not_deleted_or_bypassed(corpus, tmp_path):
    from PySide6.QtCore import QLockFile

    facade, _ = corpus
    path = facade.paths.state_root / ".desktop-workbench.lock"
    owner = QLockFile(str(path))
    assert owner.tryLock(0)
    before = path.read_bytes()
    try:
        report = run(corpus, tmp_path, apply=True)
        assert report["status"] == "blocked_no_apply"
        assert report["error"] == "desktop_lock_exists_stop"
        assert path.read_bytes() == before and owner.isLocked()
    finally:
        owner.unlock()


def test_existing_report_is_not_overwritten(corpus, tmp_path):
    path = tmp_path / "report.json"
    path.write_bytes(b"preserve original report")
    with pytest.raises(repair.RepairStop, match="fresh_report_path"):
        run(corpus, tmp_path, apply=True)
    assert path.read_bytes() == b"preserve original report"


def test_source_damage_stops_before_range_writes(corpus, tmp_path):
    facade, _ = corpus
    digest = repair.SOURCES[0][1]
    archive = facade._visual_import_root / "sources" / f"{digest}.docx"
    archive.write_bytes(b"damaged synthetic source")
    report = run(corpus, tmp_path, apply=True)
    assert report["status"] == "blocked_no_apply"
    assert not report["write_started"]


def test_interrupted_attribute_write_reports_partial_ranges_not_success(
    corpus, tmp_path, monkeypatch
):
    facade, _ = corpus
    store = facade._word_questions().attribute_store

    def fail(_rows):
        raise OSError("synthetic disk failure")

    monkeypatch.setattr(store, "save_many", fail)
    report = run(corpus, tmp_path, apply=True)
    assert report["status"] == "failed_partial_apply"
    assert len(report["range_writes"]) == 3 and not report["attribute_results"]
    assert (
        report["integrity"]["database_counts_before"]
        == report["integrity"]["database_counts_after"]
    )
    assert len(facade.state_store.snapshot()["drafts"][RANGES_DRAFT]["ranges"]) == 3


def test_non_target_state_change_is_detected_before_mutation(corpus, tmp_path):
    facade, _ = corpus
    plan = repair.build_plan(facade)
    facade.state_store.save_draft("unrelated-draft", {"changed": True})
    with pytest.raises(repair.RepairStop, match="non_target_state_changed"):
        repair.apply_plan(facade, plan, {"write_started": False}, lambda: None)


def test_provider_guard_never_resolves_a_profile():
    providers = repair.NoProviders()
    assert providers.list_metadata() == ()
    with pytest.raises(repair.RepairStop, match="provider_access_forbidden"):
        _ = providers.borrow_invocation_context


def test_unfixed_prefix_parser_cannot_apply(corpus, tmp_path, monkeypatch):
    original = repair.index_word_questions
    monkeypatch.setattr(
        repair,
        "index_word_questions",
        lambda preview: [
            row for row in original(preview) if row["origin_block_start"] != 340
        ],
    )
    report = run(corpus, tmp_path, apply=True)
    assert report["status"] == "blocked_no_apply"
    assert report["error"] == "exact_target_missing_or_duplicated"
    assert report["range_writes"] == []


def test_known_source_quality_issue_is_not_silently_promoted(corpus, tmp_path):
    facade, _ = corpus
    plan = repair.build_plan(facade)
    first = plan["targets"][0]["planned"]
    notes_path = (
        facade.paths.workspace_root / "knowledge/lectures/source-quality-notes.json"
    )
    notes_path.parent.mkdir(parents=True)
    notes_path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "sources": [
                    {
                        "source_sha256": first["source_sha256"],
                        "source_revision": first["source_revision"],
                        "issues": [
                            {
                                "id": "synthetic-issue",
                                "block_indices": [316],
                                "summary": "合成待核对问题",
                                "suggested_correction": "合成建议",
                            }
                        ],
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    report = run(corpus, tmp_path, apply=True)
    assert report["status"] == "blocked_no_apply"
    assert report["error"] == "planned_target_not_complete_or_exportable"
    assert not report["write_started"]


def test_actual_provider_attempt_count_is_reported_on_failure(
    corpus, tmp_path, monkeypatch
):
    facade, _ = corpus

    def forbidden(_facade):
        _ = facade._providers.borrow_invocation_context

    monkeypatch.setattr(repair, "build_plan", forbidden)
    report = run(corpus, tmp_path)
    assert report["status"] == "blocked_no_apply"
    assert report["provider_calls"] == 1
    assert report["error"] == "unexpected_provider_access"


def test_main_does_not_apply_without_explicit_flag(monkeypatch, tmp_path):
    calls = []

    def fake(**kwargs):
        calls.append(kwargs)
        return {
            "status": "dry_run_validated",
            "range_writes": [],
            "attribute_results": [],
        }

    monkeypatch.setattr(repair, "run_repair", fake)
    assert (
        repair.main(
            ["--state-root", str(tmp_path), "--report", str(tmp_path / "dry.json")]
        )
        == 0
    )
    assert calls[0]["apply"] is False


def test_default_factory_does_not_create_settings_or_student_files(
    desktop_paths, monkeypatch
):
    desktop_paths.state_root.mkdir()
    monkeypatch.setattr(repair, "ROOT", desktop_paths.workspace_root)
    before = list(desktop_paths.state_root.rglob("*"))
    facade = repair._make_facade(desktop_paths.state_root)
    assert isinstance(facade._providers, repair.NoProviders)
    assert facade._providers.calls == 0 and facade._student_analysis_manager is None
    assert list(desktop_paths.state_root.rglob("*")) == before


def pinned_case(package):
    origin, text_by_index = {
        "PKG-022": (
            340,
            {
                340: "这是一个合成机理案例。",
                344: "合成选项：决速步骤信息。",
                345: "合成选项：活化能信息。",
                346: "合成反应1的机理，含 Fe^{2+、Fe^{3+ 和 H_{2}O_{2 符号。",
            },
        ),
        "PKG-040": (
            265,
            {
                265: "(2025·上海卷合成定位)这是电负性合成题。",
                266: "合成备选 C / H / O。",
            },
        ),
        "PKG-069": (
            293,
            {
                293: "广义的水解观认为：合成前提。根据上述观点判断合成产物。",
            },
        ),
    }[package]
    digest = next(sha for item, sha, _ in repair.SOURCES if item == package)
    question = {
        "key": repair._key(digest, origin),
        "source_sha256": digest,
        "origin_block_start": origin,
        "revision": "synthetic-question-v1",
        "source_revision": "synthetic-preview-v1",
        "index_revision": "synthetic-index",
        "extraction_revision": "synthetic-extractor",
        "chapter": "有机化合物的结构",
        "question_blocks": [
            {"index": index, "text": text, "assets": []}
            for index, text in text_by_index.items()
        ],
        "answer_blocks": [],
        "context_blocks": [],
        "warnings": [],
    }
    metadata = {"package_id": package, "source_name": f"{package}-synthetic.docx"}
    return question, suggest_attributes(question, metadata, TAXONOMY)


@pytest.mark.parametrize(
    ("package", "primary", "supporting"),
    [
        ("PKG-022", "K09", ["K11"]),
        ("PKG-040", "K12", []),
        ("PKG-069", "unknown", []),
    ],
)
def test_pinned_labels_use_question_evidence_not_organic_chapter(
    package, primary, supporting
):
    question, automatic = pinned_case(package)
    assert automatic["primary_knowledge"]["id"] == "K15"
    assert automatic["curriculum_candidates"]
    before = deepcopy(automatic)
    corrected = repair.source_pinned_suggestion(question, automatic, TAXONOMY)
    assert automatic == before
    assert corrected["primary_knowledge"]["id"] == primary
    assert [entry["id"] for entry in corrected["supporting_knowledge"]] == supporting
    assert corrected["source"] == automatic["source"]
    assert (
        corrected["annotation_source"] == "auto_suggested"
        and corrected["teacher_note"] == ""
    )
    assert (
        corrected["curriculum_candidates"] == []
        and corrected["curriculum_status"] == "pending_mapping"
    )
    assert corrected["rule_revision"] == repair.PINNED_ATTRIBUTE_REVISION
    for field in (
        "answer_status",
        "material_status",
        "applicable_grades",
        "response_forms",
    ):
        assert corrected[field] == automatic[field]
    blocks = {block["index"]: block["text"] for block in question["question_blocks"]}
    for label in [corrected["primary_knowledge"], *corrected["supporting_knowledge"]]:
        assert label["status"] == "auto_suggested"
        assert all(
            entry["kind"] == "question_text"
            and entry["quote"] in blocks[entry["block_index"]]
            for entry in label["evidence"]
        )
    if package == "PKG-069":
        assert corrected["primary_knowledge"]["label"] == "广义水解观点下的产物判断"
    if package != "PKG-040":
        assert corrected["original_source"] == automatic["original_source"]


def test_shanghai_attribution_only_observes_year_and_region():
    question, automatic = pinned_case("PKG-040")
    corrected = repair.source_pinned_suggestion(question, automatic, TAXONOMY)
    original = corrected["original_source"]
    assert original["year"]["value"] == "2025" and original["region"]["value"] == "上海"
    assert (
        original["year"]["status"] == original["region"]["status"] == "source_observed"
    )
    assert original["grade"] == automatic["original_source"]["grade"]
    assert original["exam_type"] == automatic["original_source"]["exam_type"]
    assert original["grade"]["value"] == original["exam_type"]["value"] == "unknown"
    assert original["citation_quotes"][0]["block_index"] == 265
    assert (
        original["citation_quotes"][0]["quote"]
        in question["question_blocks"][0]["text"]
    )


@pytest.mark.parametrize("change", ["sha", "origin"])
def test_identical_words_on_another_source_or_question_do_not_trigger(change):
    question, _ = pinned_case("PKG-022")
    if change == "sha":
        question["source_sha256"] = "a" * 64
    else:
        question["origin_block_start"] = 316
    question["key"] = repair._key(
        question["source_sha256"], question["origin_block_start"]
    )
    automatic = suggest_attributes(question, {"package_id": "PKG-022"}, TAXONOMY)
    assert repair.source_pinned_suggestion(question, automatic, TAXONOMY) == automatic


@pytest.mark.parametrize("package", ["PKG-022", "PKG-040", "PKG-069"])
def test_pinned_missing_or_relocated_words_fail_closed(package):
    question, automatic = pinned_case(package)
    question["question_blocks"][-1]["text"] = "不含所需词语的合成区块。"
    with pytest.raises(repair.RepairStop, match="source_anchor_missing"):
        repair.source_pinned_suggestion(question, automatic, TAXONOMY)


def test_pinned_taxonomy_mismatch_cannot_guess_a_replacement():
    question, automatic = pinned_case("PKG-022")
    with pytest.raises(repair.RepairStop, match="taxonomy_mismatch"):
        repair.source_pinned_suggestion(
            question, automatic, {"knowledge_points": [], "nodes": []}
        )


def test_pinned_override_cannot_promote_or_replace_teacher_attributes():
    question, automatic = pinned_case("PKG-022")
    teacher = repair._seal({**automatic, "annotation_source": "teacher_modified"})
    with pytest.raises(repair.RepairStop, match="requires_automatic_proposal"):
        repair.source_pinned_suggestion(question, teacher, TAXONOMY)


def test_pinned_exact_key_and_revision_bindings_are_required():
    question, automatic = pinned_case("PKG-022")
    question["key"] = "wrong-key"
    with pytest.raises(repair.RepairStop, match="question_identity_mismatch"):
        repair.source_pinned_suggestion(question, automatic, TAXONOMY)
    question, automatic = pinned_case("PKG-022")
    question["revision"] = "changed-current-question"
    with pytest.raises(repair.RepairStop, match="source_binding_mismatch"):
        repair.source_pinned_suggestion(question, automatic, TAXONOMY)
