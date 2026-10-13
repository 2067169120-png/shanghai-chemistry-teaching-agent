"""Short source stems remain with a first diagram; long text flows normally."""
from io import BytesIO
from docx import Document
from docx.shared import Mm
from PIL import Image
from copy import deepcopy
import zipfile
import pytest
from integrations.deeptutor_shchem_v1.desktop_lesson_output import _keep_question_opening
from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    add_resource, block, new_composition, resource)
from test_composition_student_pages import payload
from test_composition_inline_personal import inline
from test_composition_render import app


def picture(paragraph,width,height):
    raw=BytesIO();Image.new('RGB',(120,90),'#547691').save(raw,format='PNG')
    paragraph.add_run().add_picture(BytesIO(raw.getvalue()),width=Mm(width),height=Mm(height))


def test_short_title_and_stem_follow_whole_diagram_without_chaining_options():
    doc=Document();title=doc.add_heading('气体装置原题',1);stem=doc.add_paragraph('依据题给条件比较制备与检验。')
    image=doc.add_paragraph();picture(image,100,50);option=doc.add_paragraph('A．原题选项保持。')
    _keep_question_opening(doc,[title,stem,image,option])
    assert title.paragraph_format.keep_with_next and stem.paragraph_format.keep_with_next
    assert image.paragraph_format.keep_together and not image.paragraph_format.keep_with_next
    assert option.paragraph_format.keep_together is None


def test_small_formula_and_badge_do_not_stop_the_group_before_diagram():
    doc=Document();title=doc.add_heading('来源公式与装置',1)
    intro=doc.add_paragraph('短导语');picture(intro,8,5)
    badge=doc.add_paragraph();picture(badge,20,8)
    stem=doc.add_paragraph('原装置图如下。');image=doc.add_paragraph();picture(image,90,45)
    _keep_question_opening(doc,[title,intro,badge,stem,image])
    assert all(p.paragraph_format.keep_with_next for p in (title,intro,badge,stem))
    assert not image.paragraph_format.keep_with_next


def test_long_stem_and_tall_diagram_keep_normal_flow_and_complete_contents():
    for long_stem,tall in ((True,False),(False,True)):
        doc=Document();title=doc.add_heading('长题完整保存',1)
        text='条件、单位和共享材料。'*100 if long_stem else '短导语'
        stem=doc.add_paragraph(text);image=doc.add_paragraph();picture(image,110,230 if tall else 45)
        _keep_question_opening(doc,[title,stem,image])
        assert stem.text==text and stem.paragraph_format.keep_with_next is None
        assert image.paragraph_format.keep_together is None


def output_payload(blocks):
    item=resource('complete_question','完整题面',blocks,
        {'source_path':'合成原题.docx','source_sha256':'a'*64,
         'source_revision':'b'*64,'question_key':'synthetic-word-question',
         'scope_review_sha256':'c'*64,'scope_status':'complete_native_scope',
         'hierarchy':dict.fromkeys(('paper_id','theme_big_question_id',
             'printed_question_id','atomic_part_id'),'unknown')})
    value=new_composition()
    value=add_resource(value,item,{'page_id':value['pages'][0],
        'x':12,'y':12,'width':186,'height':273})
    return payload(value)


@pytest.mark.parametrize('student',[False,True])
@pytest.mark.parametrize('shared',[False,True])
def test_short_printed_question_keeps_all_options_and_leaves_teacher_answer_separate(tmp_path,student,shared):
    blocks=([block('text','shared_material','共同材料的单位和适用条件。',1)] if shared else [])
    blocks += [block('text','question','12．请依据题给条件选择。',2),
        *[block('text','question',letter+'．原选项与全部条件须保留。',i)
          for i,letter in enumerate('ABCD',3)],
        block('text','answer','仅教师答案与解析。',7)]
    data=output_payload(blocks);before=deepcopy(data);write_documents(data,tmp_path,{})
    doc=Document(tmp_path/('student_worksheet.docx' if student else 'lesson_plan.docx'))
    start=next(i for i,p in enumerate(doc.paragraphs) if p.text=='完整题面')
    group=doc.paragraphs[start:start+len(blocks)]
    assert all(p.paragraph_format.keep_together and p.paragraph_format.keep_with_next for p in group[:-1])
    assert group[-1].paragraph_format.keep_together and group[-1].paragraph_format.keep_with_next is False
    text=''.join(p.text for p in doc.paragraphs)
    assert ('仅教师答案与解析。' in text)==(not student)
    # Classroom and teacher products contain teaching content. Provenance is
    # retained in the sealed input, rather than appended to the lesson body.
    evidence=data['lesson_design']['composition']['resources'][0]['original']['evidence']
    assert not any(p.text.startswith('来源与内容：') for p in doc.paragraphs)
    assert evidence['source_path'] not in text and evidence['source_sha256'] not in text
    assert evidence==before['lesson_design']['composition']['resources'][0]['original']['evidence']
    if not student:
        answer=next(p for p in doc.paragraphs if p.text=='仅教师答案与解析。')
        assert not answer.paragraph_format.keep_with_next
    assert data==before


@pytest.mark.parametrize('kind',['long_text','two_tall_images'])
def test_printed_question_larger_than_a_page_retains_normal_flow_and_every_source_byte(tmp_path,kind):
    blocks=[block('text','question','长题题干和单位。',1)]
    assets=[];raw={}
    if kind=='long_text':
        blocks.append(block('text','question','不能省略的原始条件。'*400,2))
    else:
        from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
        buf=BytesIO();Image.new('RGB',(120,600),'#315b78').save(buf,format='PNG')
        data=buf.getvalue();store=PreparationImageStore(tmp_path/'images')
        for i in (2,3):
            asset=store.import_bytes(data,'完整原图','测试作者','长题分页')
            assets.append(asset);raw[asset['asset_id']]=data
            blocks.append(block('image','question',asset['asset_id'],i))
    p=output_payload(blocks);p['image_assets']=assets;before=deepcopy(p)
    write_documents(p,tmp_path,raw)
    for name in ('lesson_plan.docx','student_worksheet.docx'):
        doc=Document(tmp_path/name)
        stem=next(row for row in doc.paragraphs if row.text=='长题题干和单位。')
        assert not stem.paragraph_format.keep_with_next
        if kind=='long_text':
            assert blocks[1]['value'] in ''.join(row.text for row in doc.paragraphs)
        else:
            assert len(doc.inline_shapes)==2
            with zipfile.ZipFile(tmp_path/name) as z:
                assert data in [z.read(n) for n in z.namelist() if '/media/' in n]
    assert p==before


@pytest.mark.parametrize('student',[False,True])
def test_interleaved_teacher_answer_does_not_join_question_keep_groups(tmp_path,student):
    p=output_payload([block('text','question','前一作答任务。',1),
        block('text','answer','仅教师中间解析。',2),
        block('text','question','后一作答任务。',3)])
    write_documents(p,tmp_path,{})
    doc=Document(tmp_path/('student_worksheet.docx' if student else 'lesson_plan.docx'))
    first=next(row for row in doc.paragraphs if row.text=='前一作答任务。')
    last=next(row for row in doc.paragraphs if row.text=='后一作答任务。')
    assert first.paragraph_format.keep_with_next is student
    assert last.paragraph_format.keep_together and last.paragraph_format.keep_with_next is False
    if not student:
        answer=next(row for row in doc.paragraphs if row.text=='仅教师中间解析。')
        assert not answer.paragraph_format.keep_with_next


@pytest.mark.parametrize('student',[False,True])
def test_short_inline_question_groups_displayed_formula_and_one_original_picture(inline,tmp_path,student):
    composition,assets,raw,_facade,native=inline
    p=payload(composition);p['image_assets']=assets;before=deepcopy(p)
    write_documents(p,tmp_path,raw)
    name='student_worksheet.docx' if student else 'lesson_plan.docx'
    doc=Document(tmp_path/name)
    question_p=next(row for row in doc.paragraphs if '图前H' in row.text)
    assert question_p.paragraph_format.keep_together and question_p.paragraph_format.keep_with_next is False
    assert '图前H2O' in question_p.text and '图后Na+尾部条件。' in question_p.text
    assert any(r.font.subscript for r in question_p.runs) and any(r.font.superscript for r in question_p.runs)
    assert len(doc.inline_shapes)==1
    with zipfile.ZipFile(tmp_path/name) as z:
        assert next(iter(raw.values())) in [z.read(n) for n in z.namelist() if '/media/' in n]
    assert ('仅教师答案H2O' in ''.join(row.text for row in doc.paragraphs))==(not student)
    assert p==before and composition['resources'][0]['original']['blocks'][0]['value']==native


def test_word_flow_revision_creates_a_fresh_output_fingerprint(monkeypatch):
    from integrations.deeptutor_shchem_v1.desktop_lesson_design import content_fingerprint
    from integrations.deeptutor_shchem_v1 import desktop_lesson_composition as composition
    p=output_payload([block('text','question','完整题干。',1)])
    current=content_fingerprint(p)
    monkeypatch.setattr(composition,'RENDER_REVISION','shchem.composition-student-arrangement.v16')
    assert content_fingerprint(p)!=current
