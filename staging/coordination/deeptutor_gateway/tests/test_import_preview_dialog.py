from __future__ import annotations

import hashlib
import io
import os
from copy import deepcopy

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from docx import Document
from PySide6.QtCore import QBuffer, QByteArray, QIODevice, Qt
from PySide6.QtGui import QImage, QPainter, QPdfWriter
from PySide6.QtPdf import QPdfDocument
from PySide6.QtWidgets import QApplication, QDialog
from test_word_question_dialog import _Tasks

from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
    PreparationSourcesService,
)
from integrations.deeptutor_shchem_v1.desktop_word_question_index import (
    index_word_questions,
)
from integrations.deeptutor_shchem_v1.desktop_workbench.import_preview_dialog import (
    ImportPreviewDialog,
)


@pytest.fixture
def qt_app():
    return QApplication.instance() or QApplication([])


def image_bytes():
    image = QImage(260, 130, QImage.Format.Format_RGB32)
    image.fill(0xFF4488AA)
    data = QByteArray()
    buffer = QBuffer(data)
    buffer.open(QIODevice.OpenModeFlag.WriteOnly)
    image.save(buffer, "PNG")
    return bytes(data)


def pdf_bytes():
    data = QByteArray()
    buffer = QBuffer(data)
    buffer.open(QIODevice.OpenModeFlag.WriteOnly)
    writer = QPdfWriter(buffer)
    painter = QPainter(writer)
    painter.drawText(300, 300, "Preview page one")
    writer.newPage()
    painter.drawText(300, 300, "Preview page two")
    painter.end()
    del writer
    return bytes(data)


class Facade:
    def __init__(self, tmp_path):
        self.calls = []
        self.asset_fails = False
        self.source_fails = False
        self.png = image_bytes()
        document = Document()
        document.add_paragraph("知识点：电解质的概念与分类")
        document.add_paragraph("【例1】下列属于电解质的是哪一种？")
        document.add_picture(io.BytesIO(self.png))
        document.add_paragraph("A. 氯化钠 B. 铜 C. 葡萄糖 D. 酒精")
        document.add_paragraph("【答案】A")
        data = io.BytesIO()
        document.save(data)
        self.docx = data.getvalue()
        self.reader = PreparationSourcesService(tmp_path)
        self.preview = self.reader.word_preview_bytes(self.docx, "演示解析版.docx")
        self.pdf = pdf_bytes()
        self.sources = [
            self.source("word", "handout", "docx", "演示解析版.docx", self.docx),
            self.source("picture", "question", "image", "演示题图.png", self.png),
            self.source("answer", "answer", "image", "演示答案.png", self.png),
            self.source("pdf", "question", "pdf", "演示两页.pdf", self.pdf),
        ]
        self.catalog = {
            "preview_id": "P1",
            "revision": "R1",
            "sources": self.sources,
            "warnings": [],
        }

    @staticmethod
    def source(source_id, role, kind, name, data):
        return {
            "source_id": source_id,
            "role": role,
            "kind": kind,
            "source_name": name,
            "source_sha256": hashlib.sha256(data).hexdigest(),
            "order_index": 1,
        }

    def preview_import_source(self, preview_id, revision, source_id):
        self.calls.append(("source", preview_id, revision, source_id))
        if self.source_fails:
            raise ValueError("source changed")
        source = next(row for row in self.sources if row["source_id"] == source_id)
        if source["kind"] == "docx":
            return {
                **source,
                "preview": deepcopy(self.preview),
                "questions": index_word_questions(self.preview),
                "warnings": [],
            }
        return {
            **source,
            "bytes": self.pdf if source["kind"] == "pdf" else self.png,
            "mime_type": "application/pdf" if source["kind"] == "pdf" else "image/png",
            "warnings": [],
        }

    def preview_import_asset(self, preview_id, revision, source_id, asset_id):
        self.calls.append(("asset", preview_id, revision, source_id, asset_id))
        if self.asset_fails:
            raise ValueError("private-source-error")
        return self.reader.word_asset_bytes(self.docx, asset_id)


def view(tmp_path):
    facade, tasks = Facade(tmp_path), _Tasks()
    dialog = ImportPreviewDialog(facade, tasks, facade.catalog)
    return dialog, facade, tasks


def close(dialog, tasks):
    dialog.reject()
    tasks.flush()
    dialog.deleteLater()


def test_initial_manifest_is_not_saved_or_marked_reviewed_and_word_is_lazy(
    qt_app, tmp_path
):
    dialog, facade, tasks = view(tmp_path)
    assert not facade.calls and not dialog.selected_source_ids and not dialog._viewed
    assert "尚未打开" in dialog.file_list.item(0).text()
    assert not dialog.confirm_button.isEnabled()
    tasks.finish("预览导入来源")
    assert "知识点：电解质" in dialog.texts[0].toPlainText()
    assert dialog._viewed == {"word"}
    assert dialog.question_combo.count() == 1
    assert "下列属于" in dialog.texts[1].toPlainText()
    assert "【答案】" not in dialog.texts[1].toPlainText()
    assert "A" in dialog.texts[2].toPlainText()
    tasks.finish("预览导入原图")
    assert not dialog.image_labels[0].pixmap().isNull()
    assert all(call[0] in {"source", "asset"} for call in facade.calls)
    close(dialog, tasks)


def test_selection_is_whole_files_and_unopened_count_is_explicit(qt_app, tmp_path):
    dialog, _, tasks = view(tmp_path)
    tasks.flush()
    dialog.select_all.click()
    assert "已打开 1 份，尚未打开 3 份" in dialog.selection_note.text()
    assert "整份文件及其全部候选" in dialog.selection_note.text()
    dialog.file_list.item(1).setCheckState(Qt.CheckState.Unchecked)
    dialog.file_list.item(2).setCheckState(Qt.CheckState.Unchecked)
    dialog.confirm_button.click()
    assert dialog.result() == QDialog.DialogCode.Accepted
    assert dialog.selected_source_ids == ["word", "pdf"]
    tasks.flush()
    dialog.deleteLater()


def test_answer_only_cannot_confirm(qt_app, tmp_path):
    dialog, _, tasks = view(tmp_path)
    tasks.flush()
    dialog.file_list.item(2).setCheckState(Qt.CheckState.Checked)
    assert not dialog.confirm_button.isEnabled()
    assert "参考答案不能单独导入" in dialog.status.text()
    close(dialog, tasks)


def test_source_image_and_pdf_show_real_pixels_and_pages(qt_app, tmp_path):
    dialog, _, tasks = view(tmp_path)
    tasks.flush()
    dialog.file_list.setCurrentRow(1)
    tasks.finish("预览导入来源")
    assert not dialog.image_labels[0].pixmap().isNull()
    dialog.file_list.setCurrentRow(3)
    tasks.finish("预览导入来源")
    for _ in range(5):
        qt_app.processEvents()
    assert dialog.pdf_document.status() == QPdfDocument.Status.Ready
    assert dialog.pdf_document.pageCount() == 2
    assert not dialog.pdf_document.render(0, dialog.pdf_view.size()).isNull()
    dialog.pdf_page.setValue(2)
    assert dialog.pdf_view.pageNavigator().currentPage() == 1
    assert "pdf" in dialog._viewed
    close(dialog, tasks)


def test_old_image_result_cannot_replace_other_source_and_failure_clears_pixels(
    qt_app, tmp_path
):
    dialog, facade, tasks = view(tmp_path)
    tasks.finish("预览导入来源")
    dialog.file_list.setCurrentRow(1)
    tasks.finish("预览导入原图", allow_cancelled=True)
    assert dialog._pixmaps[0] is None
    tasks.finish("预览导入来源")
    assert dialog._pixmaps[0] is not None
    facade.source_fails = True
    dialog.file_list.setCurrentRow(0)
    tasks.finish("预览导入来源")
    assert dialog._pixmaps[0] is None
    assert "暂不能完整预览" in dialog.source_note.text()
    assert "private" not in dialog.source_note.text()
    close(dialog, tasks)


def test_failed_word_asset_shows_explicit_gap_instead_of_old_image(qt_app, tmp_path):
    dialog, facade, tasks = view(tmp_path)
    facade.asset_fails = True
    tasks.flush()
    assert dialog._pixmaps[0] is None
    assert "未能预览" in dialog.image_labels[0].text()
    close(dialog, tasks)


def test_cancel_clears_selection_and_ignores_late_source_callbacks(qt_app, tmp_path):
    dialog, _, tasks = view(tmp_path)
    dialog.select_all.click()
    dialog.reject()
    tasks.finish("预览导入来源", allow_cancelled=True)
    assert (
        dialog.selected_source_ids == []
        and dialog.result() == QDialog.DialogCode.Rejected
    )
    assert not dialog.texts[0].toPlainText()
    dialog.deleteLater()


@pytest.mark.parametrize("width", [420, 760, 1200])
def test_preimport_dialog_narrow_and_wide_layouts(qt_app, tmp_path, width):
    dialog, _, tasks = view(tmp_path)
    tasks.flush()
    dialog.resize(width, 800)
    dialog.show()
    for _ in range(3):
        qt_app.processEvents()
    assert dialog.width() == width
    assert dialog.splitter.orientation() == (
        Qt.Orientation.Vertical if width < 760 else Qt.Orientation.Horizontal
    )
    assert dialog.confirm_button.isVisible()
    close(dialog, tasks)
