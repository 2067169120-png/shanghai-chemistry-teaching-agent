from __future__ import annotations

import hashlib
import io
import json
from copy import deepcopy

import pytest
from PIL import Image
from pptx import Presentation
from test_desktop_preparation import (
    FileRenderer,
    RecordingProvider,
    _payload,
    _raw_candidate,
)
from test_desktop_preparation_provider import _candidate, _context, _Transport

from integrations.deeptutor_shchem_v1.desktop_preparation import (
    DesktopPreparationError,
    DesktopPreparationManager,
    _validate_canonical_candidate,
    normalize_preparation_candidate,
    normalize_preparation_payload,
)
from integrations.deeptutor_shchem_v1.desktop_preparation_images import (
    MAX_IMAGES,
    PreparationImageError,
    PreparationImageStore,
    image_info,
    normalize_image_assets,
)
from integrations.deeptutor_shchem_v1.desktop_preparation_provider import (
    DesktopPreparationProviderError,
    StructuredPreparationProvider,
)
from integrations.deeptutor_shchem_v1.desktop_preparation_renderer import (
    NativePreparationRenderer,
)


def _import(tmp_path, store, fmt="PNG"):
    path = tmp_path / ("private-original." + fmt.lower())
    Image.new("RGB", (640, 320), "#268090").save(path, format=fmt)
    asset = store.import_image(
        path, "课堂观察图", "教师本地测试图", "比较两个区域并写下证据"
    )
    return path, asset


def _image_candidate(asset):
    raw = _raw_candidate()
    raw["slides"][1]["image"] = {
        "asset_id": asset["asset_id"],
        "observation_prompt": "观察哪些信息能支持判断？",
    }
    return raw


@pytest.mark.parametrize("fmt", ["PNG", "JPEG", "WEBP"])
def test_image_copy_metadata_and_load_preserve_source(tmp_path, fmt):
    store = PreparationImageStore(tmp_path / "images")
    path, asset = _import(tmp_path, store, fmt)
    original = path.read_bytes()
    assert store.load(asset) == original
    assert asset["sha256"] == hashlib.sha256(original).hexdigest()
    assert asset["width"] == 640 and asset["height"] == 320
    assert (
        store.import_image(path, asset["caption"], asset["source"], asset["purpose"])
        == asset
    )
    assert path.read_bytes() == original
    assert str(path) not in json.dumps(asset)


@pytest.mark.parametrize(
    "case", ["path", "duplicate", "missing", "hash", "dimension", "blank", "too_many"]
)
def test_reject_invalid_metadata(tmp_path, case):
    _, asset = _import(tmp_path, PreparationImageStore(tmp_path / "images"))
    rows = [asset]
    if case == "path":
        asset["local_path"] = "C:/private/file.png"
    elif case == "duplicate":
        rows.append(deepcopy(asset))
    elif case == "missing":
        del asset["purpose"]
    elif case == "hash":
        asset["sha256"] = "0" * 64
    elif case == "dimension":
        asset["width"] = True
    elif case == "blank":
        asset["source"] = " "
    else:
        rows *= MAX_IMAGES + 1
    with pytest.raises(PreparationImageError):
        normalize_image_assets(rows)


def test_twenty_eight_distinct_images_survive_draft_provider_and_real_pptx(tmp_path):
    manager = DesktopPreparationManager(
        tmp_path / "manager", NativePreparationRenderer()
    )
    assets, originals = [], []
    image_count = 28
    for number in range(image_count):
        stream = io.BytesIO()
        Image.new(
            "RGB",
            (160 + number, 90),
            ((number * 7) % 256, (40 + number * 11) % 256, (80 + number * 13) % 256),
        ).save(stream, format="PNG")
        originals.append(stream.getvalue())
        assets.append(
            manager.image_store.import_bytes(
                originals[-1],
                f"测试图{number + 1}",
                "合成测试图，不是教材",
                "检查原图嵌入顺序",
            )
        )
    assert MAX_IMAGES == 48
    assert normalize_image_assets(assets) == assets
    too_many = list(assets)
    for number in range(image_count, MAX_IMAGES + 1):
        extra = deepcopy(assets[0])
        extra["sha256"] = f"{number + 100000:064x}"
        extra["asset_id"] = "IMG-" + extra["sha256"]
        too_many.append(extra)
    with pytest.raises(PreparationImageError, match=rf"最多选择{MAX_IMAGES}张"):
        normalize_image_assets(too_many)
    payload = {**_payload(output_kind="ppt"), "image_assets": assets}
    assert normalize_preparation_payload(payload)["image_assets"] == assets
    raw = _raw_candidate()
    template = raw["slides"][1]
    raw["slides"] = []
    for number, asset in enumerate(assets, 1):
        slide = deepcopy(template)
        slide.update(
            title=f"图片测试{number}",
            image={
                "asset_id": asset["asset_id"],
                "observation_prompt": "检查测试图是否完整显示。",
            },
        )
        raw["slides"].append(slide)
    provider = RecordingProvider(raw)
    prepared = manager.prepare(payload, "fixture", "r1")
    result = manager.run(prepared["task_id"], provider)
    assert result["status"] == "completed", result.get("error")
    assert len(provider.calls) == 1
    pptx, _ = manager.artifact_path(result["task_id"], "pptx")
    deck = Presentation(pptx)
    assert len(deck.slides) == image_count
    for slide, expected in zip(deck.slides, originals, strict=True):
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        assert len(pictures) == 1
        assert pictures[0].image.blob == expected


def test_preparation_materials_use_extended_40000_character_limit():
    from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
        MAX_MATERIALS,
    )

    assert MAX_MATERIALS == 40_000
    materials = "资料" * (MAX_MATERIALS // 2)
    normalized = normalize_preparation_payload(
        {**_payload(), "materials": materials}
    )
    assert normalized["materials"] == materials
    with pytest.raises(DesktopPreparationError, match=rf"超过{MAX_MATERIALS}字"):
        normalize_preparation_payload(
            {**_payload(), "materials": materials + "多"}
        )


def test_invalid_animation_and_orientation_rejected():
    with pytest.raises(PreparationImageError):
        image_info(b"not-an-image")
    image = Image.new("RGB", (12, 12), "red")
    second = Image.new("RGB", (12, 12), "blue")
    stream = io.BytesIO()
    image.save(stream, format="PNG", save_all=True, append_images=[second])
    with pytest.raises(PreparationImageError, match="静态"):
        image_info(stream.getvalue())
    stream = io.BytesIO()
    exif = Image.Exif()
    exif[274] = 6
    image.save(stream, format="JPEG", exif=exif)
    with pytest.raises(PreparationImageError, match="旋转"):
        image_info(stream.getvalue())


def test_candidate_binds_selected_image_and_legacy_remains_unchanged(tmp_path):
    _, asset = _import(tmp_path, PreparationImageStore(tmp_path / "images"))
    legacy = normalize_preparation_candidate(_raw_candidate(), _payload())
    empty = normalize_preparation_candidate(
        _raw_candidate(), {**_payload(), "image_assets": []}
    )
    assert legacy == empty
    assert "image_assets" not in legacy and "image" not in legacy["slides"][1]
    payload = {**_payload(), "image_assets": [asset]}
    result = normalize_preparation_candidate(_image_candidate(asset), payload)
    assert result["image_assets"] == [asset]
    assert result["source_basis"]["mode"] == "teacher_text_with_local_images"
    assert _validate_canonical_candidate(result, payload) == result
    raw = _image_candidate({"asset_id": "IMG-" + "f" * 64})
    with pytest.raises(DesktopPreparationError, match="未选择"):
        normalize_preparation_candidate(raw, payload)
    raw = _image_candidate(asset)
    raw["slides"][1]["visual"] = {
        "kind": "process",
        "comparison": None,
        "steps": [
            {"label": "观察", "detail": "描述"},
            {"label": "解释", "detail": "核对"},
        ],
    }
    with pytest.raises(DesktopPreparationError):
        normalize_preparation_candidate(raw, payload)


@pytest.mark.parametrize("fmt", ["PNG", "JPEG", "WEBP"])
def test_manager_renders_real_embedded_image_and_shared_preview(tmp_path, fmt):
    manager = DesktopPreparationManager(
        tmp_path / "manager", NativePreparationRenderer()
    )
    path, asset = _import(tmp_path, manager.image_store, fmt)
    provider = RecordingProvider(_image_candidate(asset))
    prepared = manager.prepare(
        {**_payload(output_kind="ppt"), "image_assets": [asset]}, "fixture", "r1"
    )
    result = manager.run(prepared["task_id"], provider)
    assert result["status"] == "completed", result.get("error")
    pptx, _ = manager.artifact_path(result["task_id"], "pptx")
    picture = next(s for s in Presentation(pptx).slides[1].shapes if s.shape_type == 13)
    if fmt != "WEBP":
        assert picture.image.blob == path.read_bytes()
    else:
        assert image_info(picture.image.blob)["content_type"] == "image/png"
    assert abs(picture.width / picture.height - 2) < 0.01
    rendered = Presentation(pptx).slides[1]
    assert asset["purpose"] in rendered.notes_slide.notes_text_frame.text
    assert asset["source"] in "\n".join(
        s.text for s in rendered.shapes if s.has_text_frame
    )
    qa = json.loads((pptx.parent / "qa_report.json").read_text(encoding="utf-8"))
    assert all(c["passed"] for c in qa["checks"]), qa["checks"]
    assert len(provider.calls) == 1
    assert str(path) not in json.dumps(provider.calls)


def test_missing_and_changed_image_fail_before_provider(tmp_path):
    manager = DesktopPreparationManager(
        tmp_path / "manager", NativePreparationRenderer()
    )
    _, asset = _import(tmp_path, manager.image_store)
    payload = {**_payload(), "image_assets": [asset]}
    task = manager.prepare(payload, "fixture", "r1")
    stored = manager.image_store.root / (asset["sha256"] + ".image")
    stored.write_bytes(b"corrupt-test-fixture")
    provider = RecordingProvider(_image_candidate(asset))
    result = manager.run(task["task_id"], provider)
    assert result["status"] == "failed"
    assert provider.calls == []
    stored.unlink()
    with pytest.raises(DesktopPreparationError, match="丢失"):
        manager.prepare(payload, "fixture", "r2")


def test_legacy_renderer_cannot_silently_discard_image_and_retry_uses_frozen_candidate(
    tmp_path,
):
    manager = DesktopPreparationManager(tmp_path / "manager", FileRenderer())
    _, asset = _import(tmp_path, manager.image_store)
    provider = RecordingProvider(_image_candidate(asset))
    task = manager.prepare(
        {**_payload(output_kind="ppt"), "image_assets": [asset]}, "fixture", "r1"
    )
    first = manager.run(task["task_id"], provider)
    assert first["status"] == "failed"
    assert first["candidate_available"]
    manager.retry(task["task_id"])
    second = manager.run(task["task_id"], provider)
    assert second["status"] == "failed"
    assert len(provider.calls) == 1


@pytest.mark.parametrize("style", ["responses", "chat_completions"])
def test_provider_sends_metadata_without_pixels_or_paths(tmp_path, style):
    path, asset = _import(tmp_path, PreparationImageStore(tmp_path / "images"))
    transport = _Transport(_candidate())
    provider = StructuredPreparationProvider(
        _context(api_style=style), transport=transport
    )
    provider.generate(
        normalize_preparation_payload({**_payload(), "image_assets": [asset]})
    )
    body = transport.requests[0].body.decode("utf-8")
    assert asset["asset_id"] in body
    assert "input_image" not in body and "data:image" not in body
    assert path.name not in body and str(path) not in body
    assert "像素" in body


def test_provider_rejects_private_metadata_fields_before_transport(tmp_path):
    path, asset = _import(tmp_path, PreparationImageStore(tmp_path / "images"))
    transport = _Transport(_candidate())
    provider = StructuredPreparationProvider(_context(), transport=transport)
    asset["local_path"] = str(path)
    with pytest.raises(DesktopPreparationProviderError):
        provider.generate({**_payload(), "image_assets": [asset]})
    assert transport.requests == []
