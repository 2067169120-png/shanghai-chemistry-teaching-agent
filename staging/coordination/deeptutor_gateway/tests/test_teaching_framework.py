"""Review-free teaching projections retain evidence, questions and saved edits."""
from copy import deepcopy
import json
from pathlib import Path
import pytest
from docx import Document
from pptx import Presentation
from test_lesson_design import sample
from test_lesson_design_ui import desk
from test_question_explorer_ui import window
from test_desktop_preparation_renderer import _candidate, _artifact_map
from integrations.deeptutor_shchem_v1.desktop_teaching_framework import *
from integrations.deeptutor_shchem_v1.desktop_lesson_design import validate_design, content_fingerprint
from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import resource, block, visible_blocks
from integrations.deeptutor_shchem_v1.desktop_preparation_renderer import NativePreparationRenderer


def full_framework():
    return {'schema_version': SCHEMA, **{key: label + '具体内容' for key, label in FIELDS}}


def test_saved_framework_is_validated_and_invalidates_old_output_without_changing_nodes():
    payload = sample(); before = deepcopy(payload); old = content_fingerprint(payload)
    payload['lesson_design']['teaching_framework'] = full_framework()
    assert validate_design(payload['lesson_design']) == payload['lesson_design']
    assert content_fingerprint(payload) != old
    assert payload['lesson_design']['nodes'] == before['lesson_design']['nodes']
    assert payload['lesson_design']['objectives'] == before['lesson_design']['objectives']


@pytest.mark.parametrize('change', [{'extra': 'x'}, {'methods': 123}, {'methods': 'x' * 12001}, {'schema_version': 'bad'}])
def test_bad_framework_is_rejected_without_silently_dropping_data(change):
    payload = sample(); payload['lesson_design']['teaching_framework'] = {**full_framework(), **change}
    with pytest.raises(ValueError): validate_design(payload['lesson_design'])


def test_source_conditions_are_kept_but_unverified_numbers_are_not_promoted():
    text = '本地整理候选，待教师核对。\n组成和结构相似时进行比较，温度单位为℃。\nHF沸点为50℃（待核实）。'
    result = classroom_text(text)
    assert result == '组成和结构相似时进行比较，温度单位为℃。'
    assert '50℃' not in result


def note(kind='knowledge_note'):
    return resource(kind, '知识卡', [block('text', 'knowledge' if kind == 'knowledge_note' else 'question',
        '本地整理候选，待教师核对。\n条件相同，数据保持12.5 mol·L⁻¹。'),
        block('text', 'teacher_note', '原页依据（待核对）：原书完整文字。')],
        {'source_path': '原书.pdf', 'source_sha256': 'a'*64, 'source_revision': 'b'*64,
         'record_id': 'NOTE-1', 'pdf_pages': [54], 'printed_pages': 'unknown'})


def test_knowledge_display_drops_raw_quote_and_reminders_keeps_exact_original_and_flags():
    item = note(); before = deepcopy(item)
    for student in (True, False):
        rows = visible_blocks(item, student=student)
        assert len(rows) == 1 and rows[0]['value'] == '条件相同，数据保持12.5 mol·L⁻¹。'
    assert item == before and item['original_sha256'] == before['original_sha256']
    assert item['human_reviewed'] is False and item['publication_allowed'] is False


def test_complete_source_question_is_never_cleaned_as_generation_boilerplate():
    # Use the source-card API with a complete-question evidence contract.
    item = resource('complete_question', '原题', [block('text', 'question', '题面：待核验是本题给出的原文条件。', 2)],
        {'source_path':'原卷.docx','source_sha256':'a'*64,'source_revision':'b'*64,
         'question_key':'source-q','scope_review_sha256':'c'*64,'scope_status':'complete_native_scope',
         'hierarchy':dict.fromkeys(('paper_id','theme_big_question_id','printed_question_id','atomic_part_id'),'unknown')})
    assert visible_blocks(item)[0]['value'] == item['original']['blocks'][0]['value']


def test_all_five_word_sections_and_teacher_student_isolation(tmp_path):
    payload = sample(); payload['lesson_design']['teaching_framework'] = full_framework()
    payload['lesson_design']['nodes'][0]['teacher_action'] += '\n待教师核对。'
    write_documents(payload, tmp_path, {})
    teacher = '\n'.join(p.text for p in Document(tmp_path/'lesson_plan.docx').paragraphs)
    student = '\n'.join(p.text for p in Document(tmp_path/'student_worksheet.docx').paragraphs)
    assert [teacher.index(h) for h in HEADINGS] == sorted(teacher.index(h) for h in HEADINGS)
    assert all(label+'具体内容' in teacher for _, label in FIELDS)
    assert '待教师' not in teacher and '仅教师答案-SECRET' in teacher
    assert '仅教师答案-SECRET' not in student and '板书设计具体内容' not in student


def test_framework_native_edits_survive_recovery_and_undo(desk):
    dialog, app, page = desk
    before = deepcopy(dialog.history.value)
    dialog.framework_fields['difficulty'].setPlainText('用粒子间作用解释物质性质。')
    dialog.framework_fields['board_design'].setPlainText('结构 → 作用 → 性质')
    dialog.apply_framework_button.click()
    assert page._payload()['lesson_design']['teaching_framework']['board_design'] == '结构 → 作用 → 性质'
    assert page.recovery.flush()
    assert page.recovery.store.load()['payload']['lesson_design'] == dialog.history.value
    dialog.travel(False)
    assert dialog.history.value == before
    dialog.travel(True)
    assert dialog.framework_fields['difficulty'].toPlainText() == '用粒子间作用解释物质性质。'


def test_generated_docx_pptx_and_notes_have_no_review_reminders_while_json_keeps_them(tmp_path):
    candidate = _candidate()
    candidate['slides'][1]['content'].append('本地整理候选，待教师核对。')
    candidate['slides'][1]['teacher_notes'] += '\n待教师核对。'
    candidate['uncertainties'] = [{'field':'source','description':'待核验来源','teacher_action':'核对原页'}]
    result = NativePreparationRenderer().render(candidate, output_kind='joint', output_dir=tmp_path/'native')
    artifacts = _artifact_map(result)
    stored = json.loads(Path(artifacts['candidate_json']['path']).read_text(encoding='utf8'))
    assert stored == candidate and candidate['uncertainties'] == stored['uncertainties']
    document = Document(artifacts['lesson_plan_docx']['path'])
    text = '\n'.join(p.text for p in document.paragraphs)
    deck = Presentation(artifacts['pptx']['path'])
    text += '\n'.join(shape.text for s in deck.slides for shape in s.shapes if shape.has_text_frame)
    text += '\n'.join(s.notes_slide.notes_text_frame.text for s in deck.slides)
    assert not any(t in text for t in ('待教师', '待核实', '待核验', '本地整理候选'))
    assert all(h in '\n'.join(p.text for p in document.paragraphs) for h in HEADINGS)
    assert stored['candidate_only'] is True and stored['publication_allowed'] is False
