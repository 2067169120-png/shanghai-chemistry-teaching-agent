import gzip
import hashlib
import io
import json

from docx import Document

from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
    PreparationSourcesService,
)
from integrations.deeptutor_shchem_v1.desktop_word_preview_cache import WordPreviewCache


def fixture(tmp_path):
    doc = Document()
    doc.add_paragraph("【例1】写出氯化钠的化学式。")
    doc.add_paragraph("【答案】NaCl")
    stream = io.BytesIO()
    doc.save(stream)
    raw = stream.getvalue()
    name = "演示解析版.docx"
    preview = PreparationSourcesService(tmp_path).word_preview_bytes(raw, name)
    return raw, name, preview


def test_cache_is_lazy_exact_and_source_name_bound(tmp_path):
    raw, name, preview = fixture(tmp_path)
    cache = WordPreviewCache(tmp_path / "cache")
    assert cache.load(raw, name) is None
    assert not cache.root.exists()
    assert cache.save(raw, name, preview)
    assert cache.load(raw, name) == preview
    assert cache.load(raw + b"changed", name) is None
    assert cache.load(raw, "另一文件名.docx") is None


def test_corrupt_or_changed_cache_rebuilds_instead_of_becoming_source(tmp_path):
    raw, name, preview = fixture(tmp_path)
    cache = WordPreviewCache(tmp_path / "cache")
    cache.save(raw, name, preview)
    path = next(cache.root.glob("*.gz"))
    record = json.loads(gzip.decompress(path.read_bytes()))
    record["preview"]["blocks"][0]["text"] = "changed"
    path.write_bytes(gzip.compress(json.dumps(record).encode()))
    assert cache.load(raw, name) is None
    path.write_bytes(b"not gzip")
    assert cache.load(raw, name) is None
    assert cache.save(raw, name, preview)
    assert cache.load(raw, name) == preview
    assert not list(cache.root.glob("*.tmp"))


def test_wrong_source_cannot_be_saved(tmp_path):
    import pytest

    raw, name, preview = fixture(tmp_path)
    cache = WordPreviewCache(tmp_path / "cache")
    preview["source_sha256"] = hashlib.sha256(b"other").hexdigest()
    with pytest.raises(ValueError):
        cache.save(raw, name, preview)
    assert not cache.root.exists()


def test_service_reopen_uses_cache_but_still_checks_archive(desktop_paths, tmp_path, monkeypatch):
    from test_desktop_visual_import_facade import _facade
    from test_word_questions_service import _import

    facade, _, provider = _import(desktop_paths, tmp_path)
    first = facade.word_question_catalog()
    restarted = _facade(desktop_paths, provider)
    monkeypatch.setattr(restarted._word_questions().reader, "word_preview_bytes", lambda *_: (_ for _ in ()).throw(AssertionError("unexpected reread")))
    assert restarted.word_question_catalog() == first
    archived = next((desktop_paths.state_root / "visual-import-v2/sources").glob("*.docx"))
    archived.write_bytes(b"changed")
    assert restarted.word_question_catalog()["items"] == []


from test_desktop_visual_import_facade import (
    desktop_paths as desktop_paths,  # noqa: PLC0414
)
