"""Explicit personal relations retain exact endpoint versions and history."""
from copy import deepcopy
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document

from test_handout_question_handoff import scoped_lecture, desktop_paths
from test_note_page_crops import notes, section_fixture, asset_fixture, shelf_fixture
from test_textbook_concept_cards import concept, supplement, supplement_fixture
from test_personal_textbook_knowledge import native_fixture
from test_personal_textbook_study import chapter, note as personal_note
from test_textbook_reading_hints import write_packet
from test_desktop_visual_import_facade import _png
from integrations.deeptutor_shchem_v1.desktop_composition_sources import import_questions, import_knowledge_note, knowledge_options
from integrations.deeptutor_shchem_v1.desktop_handout_knowledge import handout_knowledge_sources, handout_knowledge_outline, import_handout_knowledge
from integrations.deeptutor_shchem_v1.desktop_knowledge_question_links import (
    KnowledgeQuestionLinks, KNOWLEDGE_QUESTION_LINKS_DRAFT_KEY, validate_collection,
    empty_collection, digest, MAX_REASON)
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import validate_resource
from integrations.deeptutor_shchem_v1.desktop_state import DraftConflictError
from integrations.deeptutor_shchem_v1.reader_cancellation import TaskCancellationEvent, ReadCancelled, read_cancel_scope


def files(root):
    return {str(path.relative_to(root)): (path.read_bytes(), path.stat().st_mtime_ns)
            for path in root.rglob('*') if path.is_file()}


def combine(scoped_lecture, knowledge):
    facade, _scoped, binding, choices, _provider, _path, _batch = scoped_lecture
    facade.paths = replace(facade.paths, workspace_root=knowledge.root)
    service = KnowledgeQuestionLinks(facade)
    row = service.words._catalog({binding['batch_id']: {binding['source_id']}}, read_only=True)[0]['items'][0]
    return SimpleNamespace(facade=facade, service=service, row=row, binding=binding,
                           knowledge=knowledge, card=knowledge.item, choices=choices)


@pytest.fixture
def linked(scoped_lecture, notes):
    return combine(scoped_lecture, notes)


def prepare(f, reason='个人例题关联：核对条件与图示。'):
    return f.service.prepare(f.card, f.row, reason)


def saved(f):
    result = f.service.save(prepare(f), expected=f.service.snapshot())
    return result, result['collection']['records'][0]


def test_prepare_is_read_only_and_save_one_addressed_draft(linked):
    f = linked; before = f.facade.state_store.snapshot(); disk = files(f.facade.paths.state_root)
    value = prepare(f)
    assert f.facade.state_store.snapshot() == before and files(f.facade.paths.state_root) == disk
    assert value['knowledge']['kind'] == 'reading_note'
    assert value['knowledge']['id'] == f.card['original']['evidence']['record_id']
    assert 'source_path' not in str(value['knowledge'])
    snapshot = f.service.save(value, expected=f.service.snapshot())
    records = snapshot['collection']['records']; assert len(records) == 1
    assert f.facade.state_store.snapshot()['drafts'].keys() - before['drafts'].keys() == {KNOWLEDGE_QUESTION_LINKS_DRAFT_KEY}
    row = records[0]; assert row['reason'] == value['reason'] and row['history'][0]['action'] == 'save'
    assert row['question']['batch_id'] == f.binding['batch_id']
    assert row['question']['archive_source_id'] == f.binding['source_id']
    assert row['question']['scope_review_sha256'] and row['question']['ranges']['answer_start']
    assert 'human_reviewed' not in row and 'official' not in row and 'publication_allowed' not in row
    assert f.service.evaluate(row['id'])['status'] == 'current'


def test_duplicate_save_is_idempotent_and_deep_copies(linked):
    f = linked; value = prepare(f); first = f.service.save(value, expected=f.service.snapshot())
    before = files(f.facade.paths.state_root)
    assert f.service.save(value, expected=first) == first and files(f.facade.paths.state_root) == before
    first['collection']['records'][0]['reason'] = 'caller cannot mutate storage'
    assert f.service.snapshot()['collection']['records'][0]['reason'] == value['reason']
    rows = f.service.list(f.card); rows[0]['knowledge']['original']['title'] = 'cannot mutate'
    assert f.service.list()[0]['knowledge']['title'] == f.card['original']['title']


def test_metadata_list_reverse_do_not_read_any_word_inventory(linked, monkeypatch):
    f = linked; snapshot, row = saved(f)
    monkeypatch.setattr(f.service.words, '_inventory', lambda *a, **k: pytest.fail('metadata read enumerated Word'))
    monkeypatch.setattr(f.facade, '_imported_word_source', lambda *a: pytest.fail('metadata read source bytes'))
    assert f.service.list(f.card)[0]['status'] == 'unknown'
    assert f.service.reverse(f.row)[0]['id'] == row['id']
    assert f.service.snapshot() == snapshot


def test_reason_history_and_remove_work_with_missing_endpoints(linked, monkeypatch):
    f = linked; snapshot, row = saved(f)
    monkeypatch.setattr(f.service, 'knowledge_context', lambda *a: (_ for _ in ()).throw(FileNotFoundError()))
    monkeypatch.setattr(f.service, '_question', lambda *a: (_ for _ in ()).throw(FileNotFoundError()))
    assert f.service.evaluate(row['id'])['status'] == 'missing'
    edited = f.service.edit_reason(row['id'], '保存的原理由需要补充。', expected=snapshot)
    removed = f.service.remove(row['id'], expected=edited)
    current = removed['collection']['records'][0]
    assert current['knowledge'] == row['knowledge'] and current['question'] == row['question']
    assert current['removed'] and [entry['action'] for entry in current['history']] == ['save', 'edit_reason', 'remove']
    assert f.service.remove(row['id'], expected=removed) == removed
    assert f.service.edit_reason(row['id'], current['reason'], expected=removed) == removed


def test_removed_or_concurrently_changed_relation_cannot_handoff_old_selection(linked):
    f = linked; snapshot, row = saved(f)
    edited = f.service.edit_reason(row['id'], 'changed reason', expected=snapshot)
    with pytest.raises(ValueError, match='记录已变化'): f.service.resolve_selection(row)
    removed = f.service.remove(row['id'], expected=edited)
    with pytest.raises(ValueError, match='已解除'): f.service.resolve_selection(removed['collection']['records'][0])


@pytest.mark.parametrize('operation', ['save', 'edit', 'remove', 'merge'])
def test_cancelled_commit_keeps_state(linked, operation):
    f = linked; prepared = prepare(f); snapshot, row = saved(f); before = files(f.facade.paths.state_root)
    event = TaskCancellationEvent()
    def cancel(): event.set()
    with pytest.raises(ReadCancelled), read_cancel_scope(event):
        if operation == 'save': f.service.save(prepared, expected=snapshot, guard=cancel)
        elif operation == 'edit': f.service.edit_reason(row['id'], 'new', expected=snapshot, guard=cancel)
        elif operation == 'remove': f.service.remove(row['id'], expected=snapshot, guard=cancel)
        else: f.service.commit_merge(f.service.plan_merge(empty_collection(), expected=snapshot), expected=snapshot, guard=cancel)
    assert files(f.facade.paths.state_root) == before


def test_addressed_cas_rejects_same_draft_change_and_allows_other_drafts(linked):
    f = linked; snapshot, row = saved(f); other = KnowledgeQuestionLinks(f.facade)
    changed = other.edit_reason(row['id'], 'other writer', expected=snapshot)
    before = files(f.facade.paths.state_root)
    with pytest.raises(DraftConflictError): f.service.remove(row['id'], expected=snapshot)
    assert files(f.facade.paths.state_root) == before
    f.facade.state_store.save_draft('unrelated-course', {'kind': 'synthetic-course', 'text': 'other draft'})
    assert f.service.remove(row['id'], expected=changed)['collection']['records'][0]['removed']


@pytest.mark.parametrize('change', ['note', 'pdf', 'scope'])
def test_changed_knowledge_or_selection_is_stale_and_cannot_save(linked, change):
    f = linked; value = prepare(f); snapshot, row = saved(f)
    if change == 'note':
        f.knowledge.note['summary'] += ' source changed'; write_packet(f.knowledge.root, [f.knowledge.note])
    elif change == 'pdf': f.knowledge.book.write_bytes(f.knowledge.book.read_bytes() + b'changed')
    else:
        f.knowledge.registry['volumes'][0]['volume_title'] += ' changed'; f.knowledge.activate()
    before = f.facade.state_store.snapshot()
    assert f.service.evaluate(row['id'])['status'] == 'stale'
    with pytest.raises(ValueError): f.service.save(value, expected=snapshot)
    assert f.facade.state_store.snapshot() == before


def test_late_knowledge_change_during_commit_is_rejected(linked):
    f = linked; value = prepare(f); snapshot = f.service.snapshot(); before = f.facade.state_store.snapshot()
    def change():
        f.knowledge.note['summary'] += ' during callback'; write_packet(f.knowledge.root, [f.knowledge.note])
    with pytest.raises(ValueError): f.service.save(value, expected=snapshot, guard=change)
    assert f.facade.state_store.snapshot() == before


def test_selected_relation_operations_never_enumerate_other_archives(linked, monkeypatch):
    f = linked
    monkeypatch.setattr(f.facade, 'list_imported_word_batches', lambda: pytest.fail('whole Word inventory read'))
    value = prepare(f); result = f.service.save(value)
    assert f.service.evaluate(result['collection']['records'][0]['id'])['status'] == 'current'


def test_current_saved_question_range_change_keeps_old_relation_stale(linked):
    f = linked; value = prepare(f); snapshot, row = saved(f)
    f.service.words.update_range(f.row['key'], f.row['revision'], **{key: f.row[key] for key in
        ('block_start', 'question_end', 'answer_start', 'block_end', 'context_start', 'context_end')})
    before = f.facade.state_store.snapshot()
    assert f.service.evaluate(row['id'])['question_status'] == 'stale'
    with pytest.raises(ValueError): f.service.save(value, expected=snapshot)
    assert f.facade.state_store.snapshot() == before


def test_source_change_in_final_save_callback_cannot_publish(linked):
    f = linked; value = prepare(f); snapshot = f.service.snapshot()
    path = f.facade.imported_word_path(f.row['batch_id'], f.row['archive_source_id'])
    from pathlib import Path
    path = Path(path); before = f.facade.state_store.snapshot()
    with pytest.raises(ValueError):
        f.service.save(value, expected=snapshot, guard=lambda: path.write_bytes(path.read_bytes() + b'changed'))
    assert f.facade.state_store.snapshot() == before


def test_malformed_strict_endpoint_types_raise_value_error(linked):
    f = linked; collection = saved(f)[0]['collection']
    collection['records'][0]['knowledge']['kind'] = []
    with pytest.raises(ValueError): validate_collection(collection)


def test_exact_same_byte_archive_is_pinned_and_absent_never_substituted(linked, tmp_path):
    f = linked; value = prepare(f); snapshot, row = saved(f)
    path = tmp_path / 'copy' / f.row['source_name']; path.parent.mkdir()
    path.write_bytes(f.facade._imported_word_source(f.row['batch_id'], f.row['archive_source_id']).content)
    earlier = tmp_path / 'other.docx'; document = Document(); document.add_paragraph('unrelated'); document.save(earlier)
    other = f.facade.save_visual_import_batch(handout_files=(earlier, path), source_type='教师讲义', group_id='other-archive')
    source = next(source for source in f.facade.imported_word_sources(other.batch_id) if source['source_sha256'] == f.row['source_sha256'])
    f.service.words._catalog({other.batch_id: {source['source_id']}}, read_only=True)
    resolved = f.service.resolve_selection(row)
    assert resolved['endpoint']['batch_id'] == f.row['batch_id']
    assert f.service.words._locations[f.row['key']] == (f.row['batch_id'], f.row['archive_source_id'])
    store = f.facade.state_store; state = store.snapshot()
    original_key = next(key for key, record in state['drafts'].items() if isinstance(record, dict) and record.get('batch_id') == f.row['batch_id'])
    store.save_draft(original_key, {'kind': 'missing-original'})
    assert f.service.evaluate(row['id'])['question_status'] == 'missing'
    with pytest.raises(ValueError): f.service.resolve_selection(row)


def test_course_resource_recovers_actual_revision_and_digest_then_refuses_tamper(linked):
    f = linked; imported = import_questions(f.facade, f.choices, [], expected_reference=f.facade.word_question_reference(f.choices))
    card = imported['resources'][0]
    assert 'question_revision' not in card['original']['evidence']
    before = files(f.facade.paths.state_root)
    result = f.service.resolve_selection(card)
    assert result['selection']['revision'] == f.row['revision']
    assert result['endpoint']['scope_review_sha256'] == card['original']['evidence']['scope_review_sha256']
    assert files(f.facade.paths.state_root) == before
    card['original']['evidence']['scope_review_sha256'] = 'f' * 64; card['original_sha256'] = digest(card['original'])
    with pytest.raises(ValueError): f.service.resolve_selection(card)


def test_course_resource_recomputed_original_cannot_forge_text(linked):
    f = linked; card = import_questions(f.facade, f.choices, [])['resources'][0]
    card['original']['blocks'][0]['value'] += ' forged'; card['blocks'][0]['value'] += ' forged'
    card['original_sha256'] = digest(card['original']); validate_resource(card)
    with pytest.raises(ValueError, match='封存块'): f.service.resolve_selection(card)


def test_course_resource_multiple_identical_archives_requires_explicit_location(linked, tmp_path):
    f = linked; card = import_questions(f.facade, f.choices, [])['resources'][0]
    path = tmp_path / 'copy' / f.row['source_name']; path.parent.mkdir()
    path.write_bytes(f.facade._imported_word_source(f.row['batch_id'], f.row['archive_source_id']).content)
    earlier = tmp_path / 'unrelated.docx'; document = Document(); document.add_paragraph('unrelated'); document.save(earlier)
    f.facade.save_visual_import_batch(handout_files=(earlier, path), source_type='教师讲义', group_id='course-other-copy')
    assert len(f.service.question_locations(card)) == 2
    with pytest.raises(ValueError, match='多个同字节'): f.service.resolve_selection(card)
    result = f.service.resolve_selection(card, location={key: f.row[key] for key in ('batch_id', 'archive_source_id')})
    assert result['endpoint']['archive_source_id'] == f.row['archive_source_id']


@pytest.mark.parametrize('field', ['revision', 'source_sha256', 'block_end', 'archive_source_id'])
def test_question_selection_tamper_fails_without_state_change(linked, field):
    f = linked; selected = deepcopy(f.row); before = f.facade.state_store.snapshot()
    selected[field] = (selected[field] + 1 if field == 'block_end' else 'f' * 64)
    with pytest.raises(ValueError): f.service.prepare(f.card, selected)
    assert f.facade.state_store.snapshot() == before


def test_missing_endpoints_can_be_merged_with_evidence_and_later_assessed(linked, tmp_path):
    f = linked; snapshot, row = saved(f)
    # No live source lookup is needed to restore historical personal relations.
    from integrations.deeptutor_shchem_v1.desktop_state import DesktopStateStore
    fresh = SimpleNamespace(state_store=DesktopStateStore(tmp_path / 'empty-profile'), _word_questions=lambda: f.service.words,
                            _imported_word_source=lambda *a: (_ for _ in ()).throw(FileNotFoundError()))
    restored = KnowledgeQuestionLinks(fresh); expected = restored.snapshot()
    result = restored.commit_merge(restored.plan_merge(snapshot['collection'], expected=expected), expected=expected)
    assert result['collection'] == snapshot['collection'] and restored.list()[0]['status'] == 'unknown'
    assert restored.evaluate(row['id'])['status'] == 'missing'


def test_empty_restore_does_not_create_draft_or_state_file(linked, tmp_path):
    from integrations.deeptutor_shchem_v1.desktop_state import DesktopStateStore
    f = linked; root = tmp_path / 'fresh-profile'
    facade = SimpleNamespace(state_store=DesktopStateStore(root), _word_questions=lambda: f.service.words)
    service = KnowledgeQuestionLinks(facade); snapshot = service.snapshot(); calls = []
    assert service.commit_merge(service.plan_merge(empty_collection(), expected=snapshot), expected=snapshot,
                                guard=lambda: calls.append('guard')) == snapshot
    assert calls == ['guard'] and not root.exists()


@pytest.mark.parametrize('change', ['extra', 'missing', 'reason', 'history', 'id', 'bool_range', 'none_range', 'path'])
def test_strict_collection_rejects_tampered_fields(linked, change):
    f = linked; collection = saved(f)[0]['collection']; row = collection['records'][0]
    if change == 'extra': collection['extra'] = True
    elif change == 'missing': row.pop('question')
    elif change == 'reason': row['reason'] = 'x' * (MAX_REASON + 1)
    elif change == 'history': row['history'][0]['version'] = 2
    elif change == 'id': row['id'] = 'KQL-' + 'f' * 64
    elif change == 'bool_range': row['question']['ranges']['block_start'] = True
    elif change == 'none_range': row['question']['ranges']['question_end'] = None
    else: row['knowledge']['selection']['source_path'] = 'cannot-be-identity'
    with pytest.raises(ValueError): validate_collection(collection)


def test_concept_endpoint_is_fine_identity_not_curriculum_label(scoped_lecture, concept):
    concept.item = import_knowledge_note(concept.facade, concept.option['selection'])
    f = combine(scoped_lecture, concept); value = prepare(f)
    assert value['knowledge']['kind'] == 'registered_concept'
    assert value['knowledge']['id'] == concept.option['selection']['concept_id']
    assert 'source_path' not in str(value['knowledge'])
    snapshot = f.service.save(value); assert f.service.evaluate(snapshot['collection']['records'][0]['id'])['status'] == 'current'


def test_personal_bnote_endpoint_rechecks_saved_note_version(scoped_lecture, native_fixture):
    f = native_fixture; view = personal_note(f, chapter(f), title='个人细知识卡')
    facade = scoped_lecture[0]; facade.paths = replace(facade.paths, workspace_root=f.paths.workspace_root)
    facade.personal_textbook_study = lambda: f.study; facade.textbook_workspace = lambda: f.workspace
    service = KnowledgeQuestionLinks(facade)
    selected = knowledge_options(facade)['reading_notes'][0]['selection']; card = import_knowledge_note(facade, selected)
    row = service.words._catalog({scoped_lecture[2]['batch_id']: {scoped_lecture[2]['source_id']}}, read_only=True)[0]['items'][0]
    prepared = service.prepare(card, row); assert prepared['knowledge']['kind'] == 'personal_note'
    assert prepared['knowledge']['id'].startswith('BNOTE-')
    result = service.save(prepared); identity = result['collection']['records'][0]['id']
    f.study.save_note(f.selection, view['study_revision'], section_id=view['notes'][0]['section_id'],
        note_id=view['notes'][0]['note_id'], title='个人细知识卡', text='已修改范围与内容', pdf_page=2, teacher_confirmed=True)
    assert service.evaluate(identity)['status'] == 'stale'
    with pytest.raises(ValueError): service.save(prepared, expected=result)


@pytest.fixture
def handout_section(scoped_lecture, tmp_path):
    facade = scoped_lecture[0]
    doc = Document(); doc.add_heading('知识点 氢键', 2); doc.add_paragraph('完整原知识条件。')
    p = doc.add_paragraph('两次原图：'); p.add_run().add_picture(BytesIO(_png('blue')))
    p.add_run('中间文字').add_picture(BytesIO(_png('blue')))
    table = doc.add_table(rows=2, cols=2); table.cell(0, 0).merge(table.cell(0, 1)).text = '原表格标题'
    table.cell(1, 0).text = '条件'; table.cell(1, 1).text = '完整结论'
    doc.add_heading('知识点 其他', 2); doc.add_paragraph('其他完整知识。')
    doc.add_paragraph('思维建模  明确的实验方案设计思路')
    doc.add_paragraph('原始方法单元的完整知识正文。')
    path = tmp_path / '完整知识.docx'; doc.save(path)
    batch = facade.save_visual_import_batch(handout_files=(path,), source_type='教师讲义', group_id='section-archive')
    source = next(row for row in handout_knowledge_sources(facade)['sources'] if row['batch_id'] == batch.batch_id)
    outline = handout_knowledge_outline(facade, source)
    request = {'source': outline['source'], 'outline_revision': outline['outline_revision'],
        'section': outline['sections'][0], 'title': '我的氢键知识卡', 'teacher_only': False}
    card = import_handout_knowledge(facade, [request], [], [])['resources'][0]
    service = KnowledgeQuestionLinks(facade)
    row = service.words._catalog({scoped_lecture[2]['batch_id']: {scoped_lecture[2]['source_id']}}, read_only=True)[0]['items'][0]
    return SimpleNamespace(facade=facade, service=service, card=card, row=row, outline=outline)


def test_handout_section_rechecks_table_and_every_picture_read_only(handout_section):
    f = handout_section; before = files(f.facade.paths.state_root)
    value = prepare(f); assert value['knowledge']['kind'] == 'handout_section'
    blocks = value['knowledge']['original']['evidence']['section_snapshot']['blocks']
    assert sum(len(row['images']) for row in blocks) == 2 and any('table_rows' in row for row in blocks)
    assert files(f.facade.paths.state_root) == before
    result = f.service.save(value); row = result['collection']['records'][0]
    assert f.service.evaluate(row['id'])['status'] == 'current'
    assert f.service.knowledge_context(row['knowledge'])['card']['original']['kind'] == 'handout_knowledge'


def test_legacy_styled_card_prepare_save_evaluate_and_context_keep_frozen_directory(handout_section):
    from integrations.deeptutor_shchem_v1 import desktop_handout_knowledge as handout
    f = handout_section
    raw = f.facade._imported_word_source(f.outline['source']['batch_id'],f.outline['source']['source_id']).content
    legacy = handout._headings(Document(BytesIO(raw)),f.outline['preview'])
    old_revision = digest({'index_revision':handout.OUTLINE_REVISION,
        'source_sha256':f.outline['source']['source_sha256'],
        'source_revision':f.outline['source']['source_revision'],'sections':legacy})
    assert old_revision != f.outline['outline_revision']
    f.card = deepcopy(f.card)
    f.card['original']['evidence']['section_snapshot']['outline_revision'] = old_revision
    f.card['original_sha256'] = digest(f.card['original'])
    validate_resource(f.card)
    frozen = deepcopy(f.card['original']); before = files(f.facade.paths.state_root)
    value = prepare(f)
    assert value['knowledge']['selection']['outline_revision'] == old_revision
    assert files(f.facade.paths.state_root) == before
    result = f.service.save(value,expected=f.service.snapshot())
    row = result['collection']['records'][0]; after_save = files(f.facade.paths.state_root)
    assert f.service.evaluate(row['id'])['status'] == 'current'
    context = f.service.knowledge_context(row['knowledge'])
    assert context['card']['original']['evidence']['section_snapshot']['outline_revision'] == old_revision
    assert f.card['original'] == frozen and files(f.facade.paths.state_root) == after_save


@pytest.mark.parametrize('change', ['text', 'table', 'picture', 'role', 'range', 'source_revision', 'outline_revision'])
def test_handout_section_sealed_tamper_is_rejected(handout_section, change):
    f = handout_section; card = deepcopy(f.card); evidence = card['original']['evidence']; snapshot = evidence['section_snapshot']
    if change == 'text':
        snapshot['blocks'][1]['text'] += ' forged'
        next(row for row in card['original']['blocks'] if row['source_index'] == snapshot['blocks'][1]['index'])['value'] += ' forged'
        card['blocks'] = deepcopy(card['original']['blocks'])
    elif change == 'table': next(row for row in snapshot['blocks'] if 'table_rows' in row)['table_rows'][0]['cells'][0]['text'] += ' forged'
    elif change == 'picture': next(row for row in snapshot['blocks'] if row['images'])['images'][0]['container_xml_sha256'] = 'f' * 64
    elif change == 'role': evidence['handout_source']['source_role'] = 'answer'
    elif change == 'source_revision':
        evidence['source_revision'] = 'f'*64; evidence['handout_source']['source_revision'] = 'f'*64
    elif change == 'outline_revision': snapshot['outline_revision'] = 'f'*64
    else: snapshot['section']['end'] -= 1
    card['original_sha256'] = digest(card['original']); before = f.facade.state_store.snapshot()
    with pytest.raises(ValueError): f.service.prepare(card, f.row)
    assert f.facade.state_store.snapshot() == before
