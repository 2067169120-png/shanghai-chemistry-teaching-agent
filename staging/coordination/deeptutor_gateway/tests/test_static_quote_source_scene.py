"""Pure DOCX/XML scene contracts; no facade, StateStore or application window."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
from PIL import Image
import pytest

from integrations.deeptutor_shchem_v1 import word_static_quote as quote
from integrations.deeptutor_shchem_v1.desktop_composition_inline import NativeParagraphDisplay, display_tokens
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import _blocks, block
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import PreparationSourcesService

W, WP, A, PIC, R = quote.W, quote.WP, quote.A, quote.PIC, quote.R


def _picture(document, color):
    raw = BytesIO(); Image.new('RGB', (40, 30), color).save(raw, 'PNG')
    shape = document.add_picture(BytesIO(raw.getvalue()))
    rid = shape._inline.find('.//' + A + 'blip').get(R + 'embed')
    paragraph = shape._inline.getparent().getparent().getparent()
    paragraph.getparent().remove(paragraph)
    return rid


def _drawing(rid, *, anchor=False, cache=False):
    extent = (121920, 693420) if anchor else ((3002280, 2360930) if cache else (4533900, 3566160))
    xextent = (3014062, 2370724) if cache else extent
    tag = 'anchor' if anchor else 'inline'
    attributes = ('distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="251662336" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"'
                  if anchor else 'distT="0" distB="0" distL="0" distR="0"')
    position = ('<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>1400175</wp:posOffset></wp:positionH>'
                '<wp:positionV relativeFrom="paragraph"><wp:posOffset>880110</wp:posOffset></wp:positionV>') if anchor else ''
    effect = '' if anchor else '<a:clrChange><a:clrFrom><a:srgbClr val="FFFFFF"/></a:clrFrom><a:clrTo><a:srgbClr val="FFFFFF"><a:alpha val="0"/></a:srgbClr></a:clrTo></a:clrChange>'
    return etree.fromstring(f'''<w:drawing xmlns:w="{W[1:-1]}" xmlns:wp="{WP[1:-1]}" xmlns:a="{A[1:-1]}" xmlns:pic="{PIC[1:-1]}" xmlns:r="{R[1:-1]}">
      <wp:{tag} {attributes}>{position}<wp:extent cx="{extent[0]}" cy="{extent[1]}"/><wp:effectExtent l="0" t="0" r="0" b="{7620 if anchor else 0}"/>
      {'<wp:wrapNone/>' if anchor else ''}<wp:docPr id="7" name="source"/><wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>
      <a:graphic><a:graphicData uri="{PIC[1:-1]}"><pic:pic><pic:nvPicPr><pic:cNvPr id="7" name="source"/><pic:cNvPicPr/></pic:nvPicPr>
      <pic:blipFill><a:blip r:embed="{rid}" cstate="print">{effect}</a:blip><a:stretch><a:fillRect/></a:stretch></pic:blipFill>
      <pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{xextent[0]}" cy="{xextent[1]}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>
      </pic:pic></a:graphicData></a:graphic></wp:{tag}></w:drawing>''')


def _run(*, boundary=None, instruction=None, drawing=None, position=None):
    run = OxmlElement('w:r')
    if position is not None:
        pr = OxmlElement('w:rPr'); p = OxmlElement('w:position'); p.set(W + 'val', str(position)); pr.append(p); run.append(pr)
    if boundary:
        node = OxmlElement('w:fldChar'); node.set(W + 'fldCharType', boundary); run.append(node)
    if instruction is not None:
        node = OxmlElement('w:instrText'); node.set(quote.XML + 'space', 'preserve'); node.text = instruction; run.append(node)
    if drawing is not None:
        run.append(drawing)
    return run


def _document():
    document = Document(); paragraph = document.add_paragraph()
    rid1 = _picture(document, 'black'); rid2 = _picture(document, 'white')
    ppr = paragraph._p.get_or_add_pPr()
    for name, attrs in [('pStyle', {'val': 'Normal'}), ('snapToGrid', {'val': '0'}),
                        ('ind', {'left': '210', 'right': '0', 'firstLine': '0'}), ('jc', {'val': 'center'})]:
        node = OxmlElement('w:' + name)
        for key, value in attrs.items(): node.set(W + key, value)
        ppr.append(node)
    paragraph._p.extend([_run(drawing=_drawing(rid1, anchor=True)), _run(boundary='begin'),
                         _run(instruction=' QU'), _run(instruction='OTE '),
                         _run(drawing=_drawing(rid2), position=-269), _run(instruction=' '),
                         _run(boundary='separate'), _run(drawing=_drawing(rid2, cache=True), position=-269),
                         _run(boundary='end')])
    return document


def _display(document):
    raw = BytesIO(); document.save(raw); return NativeParagraphDisplay(raw.getvalue()), raw.getvalue()


@pytest.fixture
def source():
    display, raw = _display(_document()); scene = display.source_scene(1)
    assert scene is not None
    return display, raw, scene


def _references(scene):
    return [{key: row[key] for key in ('position', 'container_xml_sha256', 'original_sha256', 'original_mime_type', 'width_mm', 'height_mm')}
            | {'source_branch': 'direct', 'status': 'included'} for row in scene['occurrences']]


def _bound(source):
    display, _raw, scene = source
    text = block('text', 'knowledge', scene['text'], 1)
    images = [block('image', 'knowledge', 'IMG-' + row['original_sha256'], 1) for row in scene['occurrences']]
    objects = [{'container_xml_sha256': row['container_xml_sha256'], 'placeholder': quote.MARKER} for row in scene['occurrences']]
    display.bind(text, images, objects)
    return text, images


def test_source_scene_retains_both_field_occurrences_geometry_effects_and_exact_source(source, tmp_path):
    display, raw, scene = source
    before = etree.tostring(display.elements[0])
    assert scene['projection']['visible_positions'] == [1, 3]
    assert [r['phase'] for r in scene['occurrences']] == ['outside_field', 'instruction', 'cache']
    assert len({r['original_sha256'] for r in scene['occurrences']}) == 2
    assert scene['occurrences'][1]['wp_extent'] != scene['occurrences'][2]['wp_extent']
    assert scene['occurrences'][2]['xfrm']['ext'] != scene['occurrences'][2]['wp_extent']
    assert scene['occurrences'][2]['run_position_half_points'] == -269
    assert scene['occurrences'][2]['transparency'] == {'from': 'FFFFFF', 'to': 'FFFFFF', 'alpha': 0}
    assert scene['frame']['paragraph_left_twips'] == 210
    assert quote.validate_scene(scene, scene['text'], refs=_references(scene)) is scene
    assert etree.tostring(display.elements[0]) == before
    preview = PreparationSourcesService(tmp_path).word_preview_bytes(raw, 'source.docx')
    text, images = _bound(source)
    assert [t['kind'] for t in display_tokens(text, {r['id']: r for r in images})] == ['source_scene']
    assert PreparationSourcesService(tmp_path).word_preview_bytes(raw, 'source.docx') == preview
    assert 'field_code_dependency' in display.reader.read(display.elements[0]).features
    assert 'scene' not in preview['blocks'][0]


@pytest.mark.parametrize('mutation', ['missing_end', 'missing_separate', 'nested', 'cross_boundary', 'simple', 'unknown_command', 'cache_text',
    'outside_text', 'cache_relation', 'external_link', 'unknown_namespace', 'unknown_effect', 'alpha', 'positive_position', 'huge_down_position',
    'mismatched_position', 'anchor_page', 'anchor_wrap', 'anchor_flip', 'group_shape', 'unknown_crop', 'multi_column', 'first_line', 'wrong_alignment',
    'grid', 'inherited_hidden', 'fixed_line_height', 'section_break'])
def test_unknown_or_incomplete_source_is_never_promoted(mutation):
    document = _document(); p = document.paragraphs[0]._p
    runs = p.findall(W + 'r'); drawings = p.findall('.//' + W + 'drawing')
    if mutation == 'missing_end': p.remove(runs[-1])
    elif mutation == 'missing_separate': p.remove(runs[-3])
    elif mutation == 'nested': p.insert(3, _run(boundary='begin'))
    elif mutation == 'cross_boundary': runs[-1][0].set(W + 'fldCharType', 'begin')
    elif mutation == 'simple': p.append(OxmlElement('w:fldSimple'))
    elif mutation == 'unknown_command': runs[2][0].text = ' RE'; runs[3][0].text = 'F '
    elif mutation == 'cache_text': runs[-2].append(OxmlElement('w:t'))
    elif mutation == 'outside_text': runs[0].append(OxmlElement('w:t'))
    elif mutation == 'cache_relation': drawings[2].find('.//' + A + 'blip').set(R + 'embed', drawings[0].find('.//' + A + 'blip').get(R + 'embed'))
    elif mutation == 'external_link': drawings[1].find('.//' + A + 'blip').set(R + 'link', 'external')
    elif mutation == 'unknown_namespace': runs[2][0].tag = '{urn:unknown}instrText'
    elif mutation == 'unknown_effect': drawings[1].find('.//' + A + 'blip').append(etree.Element(A + 'lum'))
    elif mutation == 'alpha': drawings[2].find('.//' + A + 'alpha').set('val', '1')
    elif mutation in {'positive_position', 'huge_down_position'}:
        for run in (runs[4], runs[-2]): run.find(W + 'rPr/' + W + 'position').set(W + 'val', '10' if mutation == 'positive_position' else '-31680')
    elif mutation == 'mismatched_position': runs[-2].find(W + 'rPr/' + W + 'position').set(W + 'val', '-268')
    elif mutation == 'anchor_page': drawings[0].find('.//' + WP + 'positionH').set('relativeFrom', 'page')
    elif mutation == 'anchor_wrap': drawings[0].find('.//' + WP + 'wrapNone').tag = WP + 'wrapSquare'
    elif mutation == 'anchor_flip': drawings[0].find('.//' + A + 'xfrm').set('flipH', '1')
    elif mutation == 'group_shape': drawings[0].find('.//' + A + 'graphicData').set('uri', 'urn:group')
    elif mutation == 'unknown_crop': drawings[1].find('.//' + PIC + 'blipFill').append(etree.Element(A + 'srcRect'))
    elif mutation == 'multi_column': document.element.body.find(W + 'sectPr/' + W + 'cols').set(W + 'num', '2')
    elif mutation == 'first_line': p.find(W + 'pPr/' + W + 'ind').set(W + 'firstLine', '20')
    elif mutation == 'wrong_alignment': p.find(W + 'pPr/' + W + 'jc').set(W + 'val', 'left')
    elif mutation == 'grid': p.find(W + 'pPr/' + W + 'snapToGrid').set(W + 'val', '1')
    elif mutation == 'inherited_hidden': document.styles['Normal'].element.get_or_add_rPr().append(OxmlElement('w:vanish'))
    elif mutation == 'fixed_line_height':
        spacing = OxmlElement('w:spacing'); spacing.set(W + 'lineRule', 'exact'); spacing.set(W + 'line', '240'); p.find(W + 'pPr').append(spacing)
    elif mutation == 'section_break':
        following = document.add_paragraph(); following._p.get_or_add_pPr().append(deepcopy(document.element.body.find(W + 'sectPr')))
    display, _raw = _display(document)
    assert display.source_scene(1) is None


@pytest.mark.parametrize('mutation', ['source_xml', 'paragraph_hash', 'field_hash', 'cache_geometry', 'image_sha', 'phase', 'position',
    'transparency', 'frame', 'policy', 'visible', 'extra', 'bool_index'])
def test_resealed_semantic_tampering_is_rejected(source, mutation):
    scene = deepcopy(source[2])
    if mutation == 'source_xml': scene['paragraph_xml'] += ' '
    elif mutation == 'paragraph_hash': scene['paragraph_xml_sha256'] = '0' * 64
    elif mutation == 'field_hash': scene['quote']['field_xml_sha256'] = '0' * 64
    elif mutation == 'cache_geometry': scene['occurrences'][2]['width_mm'] += 1
    elif mutation == 'image_sha': scene['occurrences'][2]['original_sha256'] = '0' * 64
    elif mutation == 'phase': scene['occurrences'][2]['phase'] = 'instruction'
    elif mutation == 'position': scene['occurrences'][0]['position'] = 3
    elif mutation == 'transparency': scene['occurrences'][2]['transparency']['alpha'] = 2
    elif mutation == 'frame': scene['frame']['column_width_twips'] += 1
    elif mutation == 'policy': scene['projection']['coordinate_policy'] = 'guess'
    elif mutation == 'visible': scene['projection']['visible_positions'] = [3]
    elif mutation == 'extra': scene['unchecked'] = True
    elif mutation == 'bool_index': scene['span']['start'] = False
    with pytest.raises(ValueError): quote.validate_scene(scene, scene['text'], refs=_references(source[2]))


@pytest.mark.parametrize('key,value', [('position', 2), ('container_xml_sha256', 'f' * 64), ('original_sha256', 'f' * 64),
    ('original_mime_type', 'image/jpeg'), ('width_mm', 1), ('height_mm', 1), ('source_branch', 'Choice'), ('status', 'missing')])
def test_each_occurrence_crossref_is_exact(source, key, value):
    refs = _references(source[2]); refs[0][key] = value
    with pytest.raises(ValueError): quote.validate_scene(source[2], source[2]['text'], refs=refs)


def test_binding_keeps_all_original_blocks_and_teacher_added_text(source):
    text, images = _bound(source); original = text['value']; by_id = {r['id']: r for r in images}
    assert len(text['presentation']['images']) == 3
    _blocks([text, *images])
    tokens = display_tokens(text, by_id)
    assert len(tokens) == 1 and len(tokens[0]['images']) == 3
    text['value'] = '先观察分类：' + original + '解释共同特征。'
    _blocks([text, *images])
    tokens = display_tokens(text, by_id, original)
    assert [t['kind'] for t in tokens] == ['text', 'source_scene', 'text']
    assert tokens[0]['text'] == '先观察分类：' and tokens[2]['text'] == '解释共同特征。'
    text['value'] = quote.MARKER + '保留老师在图间的说明' + quote.MARKER * 2
    tokens = display_tokens(text, by_id, original)
    assert not any(t['kind'] == 'source_scene' for t in tokens)
    assert sum(t['kind'] == 'image' for t in tokens) == 3
    assert any(t.get('text') == '保留老师在图间的说明' for t in tokens)


def test_scene_has_no_fresh_composition_ids(source):
    first, images = _bound(source); second, others = _bound(source)
    assert first['id'] != second['id'] and images[0]['id'] != others[0]['id']
    assert first['presentation']['scene'] == second['presentation']['scene']
    assert 'CB-' not in json.dumps(first['presentation']['scene'])


def test_other_bounded_source_coordinates_are_derived_instead_of_sample_hardcoded():
    document = _document(); p = document.paragraphs[0]._p
    p.find('.//' + WP + 'positionH/' + WP + 'posOffset').text = '1700000'
    p.find('.//' + WP + 'positionV/' + WP + 'posOffset').text = '900000'
    p.find(W + 'pPr/' + W + 'ind').set(W + 'left', '350')
    display, _raw = _display(document); scene = display.source_scene(1)
    assert scene is not None and scene['frame']['paragraph_left_twips'] == 350
    assert scene['occurrences'][0]['anchor']['position_h']['offset_emu'] == 1700000
    quote.validate_scene(scene, scene['text'])


@pytest.mark.parametrize('filename,index', [('第08讲 高分子有机物（复习讲义）（原卷版）.docx', 44), ('第08讲 高分子有机物（复习讲义）（解析版）.docx', 45)])
def test_named_readonly_real_source_has_complete_scene_when_available(filename, index):
    path = Path('E:/项目/上海高中化学等级考/sh-chem-db/.intake/2026-07-30-user-teaching-pack/expanded/PKG-072') / filename
    if not path.is_file(): pytest.skip('named local source unavailable')
    original = path.read_bytes(); before = sha256(original).hexdigest()
    display = NativeParagraphDisplay(original); scene = display.source_scene(index)
    assert scene is not None and scene['paragraph_xml_sha256'] == 'fc2b8dd239fb06f2b929cc692a16232d45e6c744a42d2191c6238f4c3ce367cc'
    assert scene['frame']['column_width_twips'] == 9746
    assert scene['frame']['paragraph_left_twips'] == 210
    assert scene['frame']['paragraph_right_twips'] == 0
    quote.validate_scene(scene, scene['text'])
    assert sha256(path.read_bytes()).hexdigest() == before
