"""Explicit archive locations avoid global manifest work, never source checks."""
from collections import Counter
from copy import deepcopy
from dataclasses import replace
from hashlib import sha256
from pathlib import Path
from threading import Event
from types import SimpleNamespace
import json

from docx import Document
import pytest
from test_desktop_visual_import_facade import (
    FakeProviderStore, _facade, desktop_paths as desktop_paths,
)

from integrations.deeptutor_shchem_v1 import desktop_word_questions as words
from integrations.deeptutor_shchem_v1.desktop_facade import _canonical_digest
from integrations.deeptutor_shchem_v1.reader_cancellation import (
    ReadCancelled, read_cancel_scope,
)


def _document(path, text):
    document = Document(); document.add_heading('完整合成讲义', 1)
    document.add_paragraph('【即学即练1】' + text)
    document.add_paragraph('【答案】完整合成答案。')
    document.save(path)
    return path


@pytest.fixture
def archive(desktop_paths, tmp_path):
    provider = FakeProviderStore(configured=False); facade = _facade(desktop_paths, provider)
    chosen = _document(tmp_path / '本课讲义.docx', '请写出水的化学式。')
    neighbor = _document(tmp_path / '同批其他讲义.docx', '请写出氯化钠的化学式。')
    selected = facade.save_visual_import_batch(handout_files=(chosen, neighbor), source_type='教师讲义')
    other = []
    for number in (1, 2):
        path = _document(tmp_path / f'其他批次{number}.docx', f'其他完整题目{number}。')
        other.append(facade.save_visual_import_batch(handout_files=(path,), source_type='教师讲义'))
    rows = facade.imported_word_sources(selected.batch_id)
    row = next(r for r in rows if r['source_name'] == chosen.name)
    root = desktop_paths.state_root / 'visual-import-v2'
    return SimpleNamespace(facade=facade, service=facade._word_questions(), selected=selected,
        row=row, neighbor=next(r for r in rows if r != row), others=other, chosen=chosen,
        root=root, locations={selected.batch_id: {row['source_id']}}, provider=provider)


def _inventory(root):
    return {str(p.relative_to(root)): (p.stat().st_size, p.stat().st_mtime_ns, sha256(p.read_bytes()).hexdigest())
            for p in root.rglob('*') if p.is_file()}


def _track(monkeypatch, archive):
    result = {'list': [], 'details': [], 'saved': [], 'manifest': [], 'restore': [],
              'archive_reads': [], 'preview_load': [], 'preview_parse': [], 'index': []}
    facade, service = archive.facade, archive.service
    for name, key in (('list_imported_word_batches', 'list'), ('import_batch_details', 'details'),
                      ('_saved_visual_import_batch', 'saved'), ('_import_batch_manifest', 'manifest'),
                      ('_restore_visual_import_sources', 'restore')):
        original = getattr(facade, name)
        def call(*args, _key=key, _original=original, **kwargs):
            value = args[0]['batch_id'] if args and isinstance(args[0], dict) else args[0] if args else None
            result[_key].append((value, deepcopy(kwargs.get('source_ids'))))
            return _original(*args, **kwargs)
        monkeypatch.setattr(facade, name, call)
    original_read = Path.read_bytes
    def read(path):
        if path.parent == archive.root / 'sources': result['archive_reads'].append(path.name)
        return original_read(path)
    monkeypatch.setattr(Path, 'read_bytes', read)
    for owner, name, key in ((service.preview_cache, 'load', 'preview_load'),
                             (service.reader, 'word_preview_bytes', 'preview_parse')):
        original = getattr(owner, name)
        def call(data, filename, _key=key, _original=original):
            result[_key].append(filename); return _original(data, filename)
        monkeypatch.setattr(owner, name, call)
    original_index = words.index_word_questions
    def index(preview):
        result['index'].append(preview['source_name']); return original_index(preview)
    monkeypatch.setattr(words, 'index_word_questions', index)
    monkeypatch.setattr(service.preview_cache, 'save', lambda *args: pytest.fail('read-only inventory wrote a preview'))
    return result


def test_cold_and_warm_scoped_calls_read_only_one_batch_and_actual_file(archive, monkeypatch, record_property):
    f = archive; before = _inventory(f.facade.paths.state_root); original = f.chosen.read_bytes()
    with monkeypatch.context() as patch:
        calls = _track(patch, f)
        first, first_inventory = f.service._catalog(f.locations, read_only=True)
        second, second_inventory = f.service._catalog(f.locations, read_only=True)
        assert first == second and len(first['items']) == 1 and not first['warnings']
        assert set(first_inventory) == set(second_inventory) == {f.row['source_sha256']}
        assert calls['list'] == calls['details'] == []
        assert calls['saved'] == calls['manifest'] == [(f.selected.batch_id, None)] * 2
        assert calls['restore'] == [(f.selected.batch_id, {f.row['source_id']})] * 2
        assert calls['archive_reads'] == [f.row['source_sha256'] + '.docx'] * 2
        assert calls['preview_load'] == calls['preview_parse'] == calls['index'] == [f.chosen.name]
        assert first_inventory[f.row['source_sha256']][0].effective_source_file_id == f.row['source_id']
        record_property('actual_call_counts', json.dumps({key: len(rows) for key, rows in calls.items()}, sort_keys=True))
    assert _inventory(f.facade.paths.state_root) == before and f.chosen.read_bytes() == original
    assert f.provider.borrow_calls == 0


def test_global_catalog_keeps_existing_all_batch_manifest_route(archive, monkeypatch, record_property):
    f = archive; before = _inventory(f.facade.paths.state_root)
    with monkeypatch.context() as patch:
        calls = _track(patch, f); catalog, inventory = f.service._catalog(read_only=True)
        ids = {f.selected.batch_id, *(b.batch_id for b in f.others)}
        assert len(catalog['items']) == len(inventory) == 4 and not catalog['warnings']
        assert len(calls['list']) == 1 and {bid for bid, _ in calls['details']} == ids
        assert len(calls['details']) == len(calls['manifest']) == len(calls['restore']) == 3
        assert all(source_ids is None for _, source_ids in calls['restore'])
        assert len(calls['archive_reads']) == len(calls['preview_parse']) == 4
        record_property('actual_call_counts', json.dumps({key: len(rows) for key, rows in calls.items()}, sort_keys=True))
    assert _inventory(f.facade.paths.state_root) == before


def test_multiple_explicit_batches_skip_same_batch_and_other_unselected_archives(archive, monkeypatch, record_property):
    f = archive; extra = f.facade.imported_word_sources(f.others[0].batch_id)[0]
    locations = {**f.locations, f.others[0].batch_id: {extra['source_id']}}
    with monkeypatch.context() as patch:
        calls = _track(patch, f); catalog, inventory = f.service._catalog(locations, read_only=True)
        assert len(catalog['items']) == len(inventory) == 2 and not catalog['warnings']
        assert calls['list'] == calls['details'] == [] and len(calls['saved']) == len(calls['manifest']) == 2
        assert calls['restore'] == [(bid, source_ids) for bid, source_ids in locations.items()]
        assert Counter(calls['archive_reads']) == Counter([f.row['source_sha256'] + '.docx', extra['source_sha256'] + '.docx'])
        record_property('actual_call_counts', json.dumps({key: len(rows) for key, rows in calls.items()}, sort_keys=True))


def test_unrelated_damaged_manifest_and_archive_are_not_required_by_scope(archive):
    f = archive
    (f.root / 'batches' / (f.others[0].batch_id + '.json')).write_bytes(b'broken unrelated fixture')
    (f.root / 'sources' / (f.neighbor['source_sha256'] + '.docx')).write_bytes(b'broken unselected fixture')
    before = _inventory(f.facade.paths.state_root)
    catalog, inventory = f.service._catalog(f.locations, read_only=True)
    assert len(catalog['items']) == len(inventory) == 1 and not catalog['warnings']
    assert _inventory(f.facade.paths.state_root) == before


@pytest.mark.parametrize('locations', [False, 1, [], 'batch', {'': {'source'}}, {'   ': {'source'}},
    {'batch': 'source'}, {'batch': ['source']}, {'batch': None}, {'batch': set()},
    {'batch': {1}}, {'batch': {''}}, {'batch': {'   '}}])
def test_malformed_locations_fail_before_any_facade_read(archive, monkeypatch, locations):
    calls = _track(monkeypatch, archive)
    with pytest.raises(words.WordQuestionError, match='来源位置'): archive.service._inventory(locations, read_only=True)
    assert not any(calls.values())


def test_empty_scope_does_not_fall_back_to_global_catalog(archive, monkeypatch):
    calls = _track(monkeypatch, archive)
    assert archive.service._inventory({}, read_only=True) == ({}, [])
    assert not any(calls.values())


def test_progress_callback_cannot_change_the_selected_location_snapshot(archive):
    f = archive; original = deepcopy(f.locations)
    def progress(_value):
        f.locations[f.selected.batch_id].clear()
        f.locations[f.selected.batch_id].add(f.neighbor['source_id'])
    inventory, warnings = f.service._inventory(f.locations, read_only=True, progress=progress)
    assert set(inventory) == {f.row['source_sha256']} and not warnings
    assert inventory[f.row['source_sha256']][0].effective_source_file_id == next(iter(original[f.selected.batch_id]))


@pytest.mark.parametrize('change', ['missing_archive', 'same_size_bytes', 'larger_bytes', 'closure',
    'size', 'source_id', 'manifest', 'manifest_missing', 'descriptor_field_type'])
def test_selected_source_changes_fail_even_after_preview_cache_warmed(archive, monkeypatch, change):
    f = archive; f.service._catalog(f.locations, read_only=True)
    descriptor = f.facade._saved_visual_import_batch(f.selected.batch_id)
    raw = next(row for row in descriptor['sources'] if row['source_file_id'] == f.row['source_id'])
    path = f.root / raw['archive_relative_path']; manifest_path = f.root / 'batches' / (f.selected.batch_id + '.json')
    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    if change == 'missing_archive': path.unlink()
    elif change == 'same_size_bytes':
        content = bytearray(path.read_bytes()); content[-1] ^= 1; path.write_bytes(content)
    elif change == 'larger_bytes': path.write_bytes(path.read_bytes() + b'changed')
    elif change == 'manifest_missing': manifest_path.unlink()
    elif change == 'manifest':
        manifest['plan']['sources'][0]['filename'] = 'foreign.docx'
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding='utf-8')
    else:
        if change == 'closure': raw['filename'] = 'changed.docx'
        elif change == 'size': raw['size_bytes'] += 1
        elif change == 'source_id': raw['source_file_id'] = 'FOREIGN'
        else: raw['order_index'] = float(raw['order_index'])
        if change != 'closure':
            descriptor['source_closure_sha256'] = _canonical_digest(descriptor['sources'])
            manifest['plan']['sources'] = deepcopy(descriptor['sources'])
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False), encoding='utf-8')
        f.facade.state_store.save_draft(f.facade._visual_import_draft_id(f.selected.batch_id), descriptor)
    before = _inventory(f.facade.paths.state_root)
    monkeypatch.setattr(f.service.reader, 'word_preview_bytes', lambda *args: pytest.fail('changed source was parsed'))
    catalog, inventory = f.service._catalog(f.locations, read_only=True)
    assert not inventory and not catalog['items'] and catalog['warnings']
    assert _inventory(f.facade.paths.state_root) == before


@pytest.mark.parametrize('change', ['unknown_batch', 'unknown_source', 'foreign_restore', 'empty_restore', 'duplicate_restore'])
def test_requested_ids_must_restore_exactly_and_never_fall_back(archive, monkeypatch, change):
    f = archive; locations = deepcopy(f.locations); restore = f.facade._restore_visual_import_sources
    if change == 'unknown_batch': locations = {'DESKTOPBATCH-' + 'f' * 32: {f.row['source_id']}}
    elif change == 'unknown_source': locations[f.selected.batch_id] = {'FOREIGN'}
    else:
        def altered(descriptor, *, source_ids=None):
            sources = restore(descriptor, source_ids=source_ids)
            if change == 'empty_restore': return ()
            if change == 'duplicate_restore': return sources + sources
            return (replace(sources[0], source_file_id='FOREIGN'),)
        monkeypatch.setattr(f.facade, '_restore_visual_import_sources', altered)
    monkeypatch.setattr(f.facade, 'list_imported_word_batches', lambda: pytest.fail('invalid scope fell back to global'))
    catalog, inventory = f.service._catalog(locations, read_only=True)
    assert not inventory and not catalog['items'] and catalog['warnings']


@pytest.mark.parametrize('stage', ['entry', 'progress', 'restored', 'parsed', 'last_progress'])
def test_cancelled_scope_never_returns_partial_inventory_or_writes(archive, monkeypatch, stage):
    f = archive; before = _inventory(f.facade.paths.state_root); event = Event(); reads = []
    restore = f.facade._restore_visual_import_sources; parse = f.service.reader.word_preview_bytes
    def restored(*args, **kwargs):
        reads.append(True); result = restore(*args, **kwargs)
        if stage == 'restored': event.set()
        return result
    def parsed(*args):
        result = parse(*args)
        if stage == 'parsed': event.set()
        return result
    def progress(value):
        if stage == 'progress' or stage == 'last_progress' and value['message'].startswith('已整理'): event.set()
    monkeypatch.setattr(f.facade, '_restore_visual_import_sources', restored)
    monkeypatch.setattr(f.service.reader, 'word_preview_bytes', parsed)
    if stage == 'entry': event.set()
    with read_cancel_scope(event, check_boundaries=False):
        with pytest.raises(ReadCancelled): f.service._inventory(f.locations, read_only=True, progress=progress)
    if stage in {'entry', 'progress'}: assert not reads
    assert not f.service._locations and _inventory(f.facade.paths.state_root) == before


def test_same_bytes_in_another_archive_do_not_redirect_explicit_location(archive, tmp_path):
    f = archive; copied = tmp_path / '复制讲义' / f.chosen.name; copied.parent.mkdir()
    copied.write_bytes(f.chosen.read_bytes())
    earlier = _document(tmp_path / '复制前置文件.docx', '另一条完整题目。')
    batch = f.facade.save_visual_import_batch(handout_files=(earlier, copied), source_type='教师讲义', group_id='another-copy')
    row = next(row for row in f.facade.imported_word_sources(batch.batch_id)
               if row['source_sha256'] == f.row['source_sha256'])
    assert row['source_sha256'] == f.row['source_sha256'] and row['source_id'] != f.row['source_id']
    _, other = f.service._catalog({batch.batch_id: {row['source_id']}}, read_only=True)
    assert other[row['source_sha256']][1] == batch.batch_id
    catalog, original = f.service._catalog(f.locations, read_only=True)
    source, batch_id, preview, _ = original[f.row['source_sha256']]
    assert batch_id == f.selected.batch_id and source.effective_source_file_id == f.row['source_id']
    assert preview['source_sha256'] == f.row['source_sha256']
    assert f.service._locations[catalog['items'][0]['key']] == (f.selected.batch_id, f.row['source_id'])
