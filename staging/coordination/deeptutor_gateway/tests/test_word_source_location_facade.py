"""Source location integration uses synthetic local imports only."""
import io

import pytest
from docx import Document
from test_desktop_visual_import_facade import FakeProviderStore, _facade, _png
from test_desktop_visual_import_facade import desktop_paths as desktop_paths
from test_word_questions_service import _import

from integrations.deeptutor_shchem_v1.desktop_preparation_sources import PreparationSourceError
from integrations.deeptutor_shchem_v1.desktop_word_preview_cache import WordPreviewCache
from integrations.deeptutor_shchem_v1.desktop_word_questions import WordQuestionError


def _state_bytes(root):
    return {p.relative_to(root).as_posix(): p.read_bytes() for p in root.rglob("*") if p.is_file()}


def test_question_location_and_image_are_readonly_source_bound(desktop_paths, tmp_path):
    facade, original, provider = _import(desktop_paths, tmp_path, image=True)
    first, second = facade.word_question_catalog()["items"]
    before = _state_bytes(desktop_paths.state_root)
    source_bytes = original.read_bytes()
    result = facade.word_question_source_locations(first["key"], first["revision"])
    assert {b["block_index"] for b in result["blocks"]} == {b["index"] for b in first["question_blocks"]}
    location = next(loc for b in result["blocks"] for loc in b["locations"] if loc["assets"])
    image = location["assets"][0]
    assert facade.word_question_location_image(
        first["key"], first["revision"], location["location_id"], image["asset_id"]
    )["bytes"] == _png("blue")
    with pytest.raises(WordQuestionError, match="不属于"):
        facade.word_question_location_image(
            second["key"], second["revision"], location["location_id"], image["asset_id"]
        )
    with pytest.raises(WordQuestionError, match="不属于"):
        facade.word_question_location_image(
            first["key"], first["revision"], location["location_id"], image["asset_id"], scope="answer"
        )
    assert original.read_bytes() == source_bytes
    assert _state_bytes(desktop_paths.state_root) == before
    assert provider.borrow_calls == 0


@pytest.mark.parametrize("corrupt", [False, True])
def test_cold_location_calls_do_not_write_or_repair_preview_cache(desktop_paths, tmp_path, monkeypatch, corrupt):
    facade, _, _ = _import(desktop_paths, tmp_path, image=True)
    row = facade.word_question_catalog()["items"][0]
    args = row["batch_id"], row["archive_source_id"]
    old = facade.imported_word_preview(*args)
    cache = WordPreviewCache(desktop_paths.state_root / "word-question-previews")
    path = cache._path(row["source_sha256"], row["source_name"])
    assert path.resolve().is_relative_to(tmp_path.resolve())
    if corrupt:
        path.write_bytes(b"synthetic broken cache")
    else:
        path.unlink()
    facade._word_questions()._cache.clear()
    before = _state_bytes(desktop_paths.state_root)

    def no_save(*args, **kwargs):
        raise AssertionError("Read-only locator must not save a cache")

    monkeypatch.setattr(WordPreviewCache, "save", no_save)
    assert facade.imported_word_preview(*args, read_only=True) == old
    result = facade.word_question_source_locations(row["key"], row["revision"])
    assert result["source_revision"] == old["revision"]
    assert _state_bytes(desktop_paths.state_root) == before


@pytest.mark.parametrize("scope", ["question", "answer"])
def test_same_paragraph_question_answer_cannot_borrow_each_others_objects(desktop_paths, tmp_path, scope):
    doc = Document()
    p = doc.add_paragraph("【例1】根据合成示意作答：")
    p.add_run().add_picture(io.BytesIO(_png("green")))
    p.add_run("【答案】合成答案：")
    p.add_run().add_picture(io.BytesIO(_png("blue")))
    source = tmp_path / "同段题答.docx"
    doc.save(source)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    facade.save_visual_import_batch(handout_files=(source,), source_type="教师讲义")
    row = facade.word_question_catalog()["items"][0]
    assert any(b.get("display_only_split") for b in row[scope + "_blocks"])
    for call in (
        lambda: facade.word_question_source_locations(row["key"], row["revision"], scope=scope),
        lambda: facade.word_question_location_image(row["key"], row["revision"], "forged", "forged", scope=scope),
    ):
        with pytest.raises(WordQuestionError, match="共用原文区块"):
            call()


def test_question_range_change_invalidates_location_and_image(desktop_paths, tmp_path):
    facade, _, _ = _import(desktop_paths, tmp_path, image=True)
    row = facade.word_question_catalog()["items"][0]
    located = facade.word_question_source_locations(row["key"], row["revision"])
    loc = next(loc for b in located["blocks"] for loc in b["locations"] if loc["assets"])
    facade.word_question_update_range(
        row["key"], row["revision"], block_start=3, question_end=5, answer_start=6, block_end=6
    )
    with pytest.raises(WordQuestionError, match="变化"):
        facade.word_question_source_locations(row["key"], row["revision"])
    with pytest.raises(WordQuestionError, match="变化"):
        facade.word_question_location_image(row["key"], row["revision"], loc["location_id"], loc["assets"][0]["asset_id"])


def test_lecture_location_enforces_source_revision_and_bounds(desktop_paths, tmp_path):
    facade, _, _ = _import(desktop_paths, tmp_path, image=True)
    row = facade.word_question_catalog()["items"][0]
    args = row["batch_id"], row["archive_source_id"]
    preview = facade.imported_word_preview(*args)
    for digest, revision, indices in (
        ("f" * 64, preview["revision"], [1]),
        (row["source_sha256"], "f" * 64, [1]),
        (row["source_sha256"], preview["revision"], [9999]),
    ):
        with pytest.raises(PreparationSourceError):
            facade.imported_word_source_locations(*args, digest, revision, indices)
