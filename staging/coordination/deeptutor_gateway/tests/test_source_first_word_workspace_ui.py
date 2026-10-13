"""Source-first Qt choices with production metadata/outline/scoped Word readers.

Only synthetic bytes and Namespace ports in OWN are used. No application,
state store, recovery store, Office or provider is constructed.
"""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from PIL import Image
from PySide6.QtCore import QCoreApplication, QEvent
from PySide6.QtWidgets import QDialog
from test_desktop_library_ui import _ManualBridge, qt_app
from integrations.deeptutor_shchem_v1.desktop_workbench import handout_material_workspace as ui
from integrations.deeptutor_shchem_v1.desktop_word_questions import WordQuestionService
from integrations.deeptutor_shchem_v1.desktop_handout_questions import HandoutQuestionFacade
from integrations.deeptutor_shchem_v1.desktop_handout_material_handoff import validate_material_request

CATALOG = '读取讲义选材资料目录'
OUTLINE = '读取当前讲义选材原文'
PICKER = '读取当前讲义选材完整题目录'
CONFIRM = '核对本次讲义选材来源'


def events(app):
    for _ in range(3):
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        app.processEvents()


@pytest.fixture
def sources(tmp_path, monkeypatch):
    reads, restores, metadata = [], [], []
    stored = {}
    for number, role in enumerate(('handout', 'question', 'answer')):
        doc = Document()
        doc.add_heading('完整课堂讲义', 1)
        doc.add_paragraph('【即学即练1】用完整图形及全部条件说明' + str(number))
        image = BytesIO(); Image.new('RGB', (9, 7), 'blue').save(image, format='PNG')
        doc.add_paragraph().add_run().add_picture(BytesIO(image.getvalue()))
        doc.add_paragraph('【答案】完整参考答案' + str(number))
        image = BytesIO(); Image.new('RGB', (8, 6), 'red').save(image, format='PNG')
        doc.add_paragraph().add_run().add_picture(BytesIO(image.getvalue()))
        doc.add_paragraph('【解析】完整解释与适用条件。')
        stream = BytesIO(); doc.save(stream); content = stream.getvalue()
        source_id = 'OWN-archive-' + str(number)
        # Same display name deliberately cannot choose or substitute an archive.
        stored[source_id] = SimpleNamespace(effective_source_file_id=source_id,
            filename='同名原讲义.docx', role=role, content=content,
            source_sha256=sha256(content).hexdigest())
    root = tmp_path / 'source-first'; root.mkdir()
    state = SimpleNamespace(snapshot=lambda: {})
    facade = SimpleNamespace(paths=SimpleNamespace(workspace_root=root, state_root=root / 'state'),
                             state_store=state)
    def imported(batch, identity):
        assert batch == 'OWN-batch'
        reads.append(identity)
        return stored[identity]
    def restore(descriptor, *, source_ids=None):
        assert descriptor == {'batch_id': 'OWN-batch'}
        restores.append(None if source_ids is None else set(source_ids))
        return tuple(stored[key] for key in stored if source_ids is None or key in source_ids)
    def source_metadata(batch):
        assert batch == 'OWN-batch'; metadata.append(batch)
        return [{'source_id': key, 'source_name': row.filename,
                 'source_sha256': row.source_sha256, 'role': row.role} for key, row in stored.items()]
    facade.list_imported_word_batches = lambda: [SimpleNamespace(batch_id='OWN-batch')]
    facade.imported_word_sources = source_metadata
    facade._imported_word_source = imported
    facade.imported_word_path = lambda batch, identity: root / (identity + '.docx')
    facade._saved_visual_import_batch = lambda batch: {'batch_id': batch}
    facade._restore_visual_import_sources = restore
    monkeypatch.setattr('integrations.deeptutor_shchem_v1.desktop_word_questions.load_attribute_catalog',
                        lambda root: {'knowledge_points': [], 'nodes': []})
    words = WordQuestionService(facade)
    facade._word_questions = lambda: words
    facade.word_question_catalog = words.catalog
    facade.word_question_reference = words.reference
    return SimpleNamespace(facade=facade, words=words, reads=reads, restores=restores,
                           metadata=metadata, stored=stored, root=root)


@pytest.fixture
def workspace(qt_app, sources):
    dialogs = []
    browser = {'original_title': '当前真实教材节', 'curriculum_context': {'book_key': 'BOOK-A',
               'section_key': 'S1'}, 'search_query': '实验现象', 'binding_sha256': 'a' * 64}
    def make(**kwargs):
        tasks = _ManualBridge()
        dialog = ui.HandoutMaterialWorkspace(sources.facade, tasks,
            questions_only=True, question_context=browser, **kwargs)
        dialogs.append(dialog); dialog.show(); events(qt_app)
        return SimpleNamespace(dialog=dialog, tasks=tasks, source=sources, app=qt_app,
                               browser=deepcopy(browser))
    yield make
    for dialog in dialogs:
        dialog.reject(); dialog.deleteLater()
    events(qt_app)


def load(case, number=0):
    if not case.dialog._sources:
        case.tasks.succeed(case.tasks.take(CATALOG), run=True)
    case.dialog.source.setCurrentIndex(number + 1)
    case.tasks.succeed(case.tasks.take(OUTLINE), run=True)


def question_boundary(case, monkeypatch, *, foreign=False):
    opened = []
    class Reader(QDialog):
        def __init__(self, facade, tasks, parent, **options):
            super().__init__(parent)
            assert isinstance(facade, HandoutQuestionFacade)
            catalog = facade.word_question_catalog()
            assert len(catalog['sources']) == 1 and len(catalog['items']) == 1
            row = catalog['items'][0]
            chosen = {'key': row['key'], 'revision': row['revision'], 'points': 3}
            reference = facade.word_question_reference([chosen], include_images=True)
            opened.append((facade, deepcopy(options), deepcopy(reference)))
            self.selection_result = {'selections': [{**chosen,
                'source_sha256': facade.binding['source_sha256'],
                'scoped_source_document': deepcopy(facade.binding)}]}
            if foreign:
                self.selection_result['selections'][0]['scoped_source_document']['source_id'] = 'foreign'
            self._items = {row['key']: row}
    monkeypatch.setattr(ui, 'WordQuestionDialog', Reader)
    case.dialog.questions_button.click()
    case.tasks.succeed(case.tasks.take(PICKER), run=True)
    return opened


def test_metadata_is_native_and_neither_first_row_nor_same_name_reads_a_source(workspace):
    case = workspace(); case.tasks.succeed(case.tasks.take(CATALOG), run=True)
    assert case.source.metadata == ['OWN-batch']
    assert case.dialog.source.count() == 4 and case.dialog.source.currentData() is None
    assert case.source.reads == [] and case.source.restores == [] and not case.source.words._cache
    assert not case.dialog.questions_button.isEnabled() and not case.dialog.confirm.isEnabled()
    assert case.dialog.all_questions_button.isEnabled()
    assert case.dialog.knowledge_button.isHidden() and case.dialog.images_button.isHidden()
    case.dialog.source_search.setText('同名')
    assert case.dialog.source.currentData() is None and case.source.reads == []
    assert case.dialog.request is None


@pytest.mark.parametrize('number', [0, 1, 2], ids=['handout', 'question', 'answer'])
def test_explicit_archive_uses_real_scoped_catalog_full_original_and_answer_roles(workspace, monkeypatch, number):
    case = workspace(); load(case, number)
    document = deepcopy(case.dialog._document)
    assert len(document) == 8 and document['source_id'] == 'OWN-archive-' + str(number)
    assert not case.source.restores  # Reading this outline does not read a global question catalogue.
    opened = question_boundary(case, monkeypatch)
    facade, options, reference = opened[0]
    assert options['required_source_id'] == document['source_sha256']
    assert options['persist_selection'] is False and options['allow_knowledge_links'] is False
    assert options['selection_only'] is True and options['handoff_mode'] == 'composition'
    assert options['initial_curriculum_context'] == case.browser['curriculum_context']
    assert options['initial_search_query'] == case.browser['search_query']
    assert all(row == {document['source_id']} for row in case.source.restores)
    assert set(case.source.reads) == {document['source_id']}
    assert reference['include_images'] is True
    assert '完整参考答案' in reference['materials'] and '完整解释与适用条件' in reference['materials']
    assert len(reference['image_assets']) == 2 and not reference['image_issues']
    assert any('题面原文' in asset['caption'] for asset in reference['image_assets'])
    assert any('原文答案与解析' in asset['caption'] and '教师讲评' in asset['purpose']
               for asset in reference['image_assets'])
    assert not hasattr(facade, 'word_question_save_selection')
    case.dialog._child.accept(); events(case.app)
    case.dialog.confirm.click(); case.tasks.succeed(case.tasks.take(CONFIRM), run=True)
    material = validate_material_request(case.dialog.request)
    assert material['source'] == document
    assert not material['knowledge_requests'] and not material['image_selections']
    assert material['question_selections'][0]['scoped_source_document'] == document
    assert material['selection_order'] == [{'kind': 'question', 'index': 0}]
    assert case.dialog.result() == QDialog.DialogCode.Accepted
    assert not (case.source.root / 'state').exists()  # Read-only native operation never persists tags or preview cache.


def test_scope_options_are_owned_and_foreign_child_selection_is_rejected(workspace, monkeypatch):
    case = workspace(); load(case)
    original = deepcopy(case.browser)
    case.browser['curriculum_context']['section_key'] = 'caller-changed'
    opened = question_boundary(case, monkeypatch, foreign=True)
    assert opened[0][1]['initial_curriculum_context'] == original['curriculum_context']
    case.dialog._child.accept(); events(case.app)
    assert not case.dialog._selections and case.dialog.request is None
    assert '不属于' in case.dialog.status.text()


@pytest.mark.parametrize('action', ['stop', 'close', 'all'])
def test_late_source_callback_cannot_read_or_select_after_cancel_or_explicit_all(workspace, action):
    case = workspace(); case.tasks.succeed(case.tasks.take(CATALOG), run=True)
    case.dialog.source.setCurrentIndex(1); old = case.tasks.take(OUTLINE)
    if action == 'stop': case.dialog.stop_button.click()
    elif action == 'close': case.dialog.reject()
    else: case.dialog.all_questions_button.click()
    assert old.task_id in case.tasks.cancelled
    case.tasks.succeed(old, run=True)
    assert case.dialog._document is None and case.dialog.request is None
    assert not case.source.restores and not case.dialog._selections
    assert case.dialog.all_question_library_requested is (action == 'all')


def test_source_aba_and_late_question_child_keep_new_source_clean(workspace, monkeypatch):
    case = workspace(); load(case)
    question_boundary(case, monkeypatch)
    child = case.dialog._child
    case.dialog.source.setCurrentIndex(2); old = case.tasks.take(OUTLINE)
    case.dialog.source.setCurrentIndex(1); current = case.tasks.take(OUTLINE)
    case.tasks.succeed(current, run=True); case.tasks.succeed(old, run=True)
    child.accept(); events(case.app)
    assert case.dialog._document['source_id'] == 'OWN-archive-0'
    assert not case.dialog._selections and case.dialog.request is None
    assert old.task_id in case.tasks.cancelled


def test_confirm_source_role_change_rejects_without_losing_pending_question(workspace, monkeypatch):
    case = workspace(); load(case)
    question_boundary(case, monkeypatch); case.dialog._child.accept(); events(case.app)
    before = deepcopy(case.dialog._selections)
    case.source.stored['OWN-archive-0'].role = 'answer'
    case.dialog.confirm.click(); task = case.tasks.take(CONFIRM)
    try: task.operation()
    except ValueError as exc: task.on_failure(str(exc))
    else: pytest.fail('Changed original role must reject')
    assert case.dialog.request is None and case.dialog._selections == before
    assert case.dialog.result() != QDialog.DialogCode.Accepted


def test_questions_only_cannot_collect_other_material_and_explicit_all_returns_no_partial(workspace, monkeypatch):
    case = workspace(); load(case)
    with pytest.raises(ValueError, match='只选择完整例题'):
        case.dialog._add_selections('knowledge', [])
    case.dialog.choose_knowledge(); case.dialog.choose_images()
    assert not case.tasks.pending and case.dialog._child is None
    question_boundary(case, monkeypatch); case.dialog._child.accept(); events(case.app)
    case.dialog.all_questions_button.click()
    assert case.dialog.result() == QDialog.DialogCode.Accepted
    assert case.dialog.all_question_library_requested and case.dialog.request is None


def test_late_metadata_after_explicit_all_accept_or_window_close_never_reopens(workspace):
    case = workspace(); late = case.tasks.take(CATALOG)
    case.dialog.all_questions_button.click()
    assert case.dialog.result() == QDialog.DialogCode.Accepted
    case.dialog.close()
    case.tasks.succeed(late, run=True); events(case.app)
    assert case.dialog._closed and case.dialog._sources == []
    assert case.dialog._child is None and case.dialog.request is None
    assert case.source.reads == [] and case.source.restores == []


@pytest.mark.parametrize('width', [430, 1080])
def test_source_first_choice_and_all_fallback_fit_window_without_implicit_read(workspace, width):
    import os
    case = workspace(); case.tasks.succeed(case.tasks.take(CATALOG), run=True)
    case.dialog.resize(width, 760); events(case.app)
    bounds = case.dialog.rect()
    assert case.dialog.splitter.isHidden() and case.dialog.navigation.isHidden()
    for control in (case.dialog.all_questions_button, case.dialog.cancel_button, case.dialog.confirm):
        top = control.mapTo(case.dialog, control.rect().topLeft())
        bottom = control.mapTo(case.dialog, control.rect().bottomRight())
        assert bounds.contains(top) and bounds.contains(bottom)
    assert case.dialog.content_scroll.horizontalScrollBar().maximum() == 0
    assert case.dialog.source.currentData() is None and not case.source.reads and not case.source.restores
    assert case.dialog.all_questions_button.isEnabled() and not case.dialog.confirm.isEnabled()
    screenshot_root = os.environ.get('SECTION_UI_OWN')
    if width == 430 and screenshot_root:
        target = Path(screenshot_root) / 'source-first-word-430.png'
        assert case.dialog.grab().save(str(target), 'PNG')
