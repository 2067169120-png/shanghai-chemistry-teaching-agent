"""Local vector display keeps opaque source bytes, occurrence identity and audiences."""
from copy import deepcopy
import ctypes
from hashlib import sha256
from io import BytesIO
import struct
from threading import Event
from zipfile import ZipFile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
import pytest

from integrations.deeptutor_shchem_v1 import desktop_handout_images as images
from integrations.deeptutor_shchem_v1 import desktop_handout_knowledge as knowledge
from integrations.deeptutor_shchem_v1 import desktop_word_metafile_preview as renderer
from integrations.deeptutor_shchem_v1 import desktop_word_wmf_fonts as fonts
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    add_resource, digest, new_composition, validate_resource)
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import PreparationSourcesService
from integrations.deeptutor_shchem_v1.reader_cancellation import ReadCancelled, read_cancel_scope
from test_handout_knowledge import make_facade, request_for, word_bytes
from test_word_wmf_fonts import _record, _wmf_records, _font, _textout

V = 'urn:schemas-microsoft-com:vml'
O = 'urn:schemas-microsoft-com:office:office'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'


def palette_wmf(*extra):
    palette = struct.pack('<HH', 0x0300, 2) + b'\0\0\0\0\xff\xff\xff\0'
    short_font = _font(b'System')[:25] + b'\0'
    return _wmf_records(_record(0x00F7, palette), _record(0x0234, b'\0\0'), _record(0x0035),
        _record(0x020B, struct.pack('<hh', 0, 0)), _record(0x020C, struct.pack('<hh', 100, 200)),
        _record(0x0214, struct.pack('<hh', 30, 10)), _record(0x0213, struct.pack('<hh', 30, 190)),
        _record(0x0214, struct.pack('<hh', 70, 10)), _record(0x0213, struct.pack('<hh', 70, 190)),
        _record(0x02FB, short_font), *extra)


def test_complete_geometric_stream_accepts_exact_palette_and_unused_compact_font():
    assert fonts.wmf_font_runs(palette_wmf()) == ()


@pytest.mark.parametrize('record', [
    _record(0x00F7),
    _record(0x00F7, struct.pack('<HH', 0x0301, 1) + b'\0'*4),
    _record(0x00F7, struct.pack('<HH', 0x0300, 0)),
    _record(0x00F7, struct.pack('<HH', 0x0300, 2) + b'\0'*4),
    _record(0x00F7, struct.pack('<HH', 0x0300, 1) + b'\0\0\0\x08'),
    _record(0x0234), _record(0x0234, b'\0\0\0\0'), _record(0x0234, b'\x07\0'),
    _record(0x0035, b'\0\0'), _record(0x02FB, b'\0'*26), _record(0x0777),
])
def test_malformed_palette_compact_font_or_unknown_control_never_admitted(record):
    with pytest.raises(fonts.WmfFontError):
        fonts.wmf_font_runs(palette_wmf(record))


@pytest.mark.parametrize('text', [b'', b'Fe3+'])
def test_compact_font_is_refused_if_any_text_draw_record_exists_anywhere(text):
    with pytest.raises(fonts.WmfFontError):
        fonts.wmf_font_runs(palette_wmf(_textout(text)))


@pytest.mark.parametrize('change', ['no_eof', 'truncated', 'after_eof'])
def test_text_free_proof_requires_the_entire_complete_declared_stream(change):
    source = bytearray(palette_wmf())
    if change == 'no_eof': source[-2:] = b'\x35\0'
    elif change == 'truncated': del source[-2:]
    else:
        source += _textout(b'Fe')
        struct.pack_into('<I', source, 28, (len(source)-22)//2)
    with pytest.raises(fonts.WmfFontError): fonts.wmf_font_runs(bytes(source))


def add_object(doc, paragraph, vector, number):
    part = Part(PackURI(f'/word/media/vector{number}.wmf'), 'image/x-wmf', vector, doc.part.package)
    image_rid = doc.part.relate_to(part, RT.IMAGE)
    opaque = b'opaque uninterpreted ChemWindow binary '+str(number).encode()
    ole_part = Part(PackURI(f'/word/embeddings/oleObject{number}.bin'),
        'application/vnd.openxmlformats-officedocument.oleObject', opaque, doc.part.package)
    ole_rid = doc.part.relate_to(ole_part, RT.OLE_OBJECT)
    node = etree.Element('{' + W + '}object')
    shape_id = f'_x0000_i{number}'
    shape = etree.SubElement(node, '{' + V + '}shape', id=shape_id, type='#_x0000_t75',
        style='width:32.25pt;height:8.25pt', filled='f', stroked='f')
    etree.SubElement(shape, '{' + V + '}imagedata', {qn('r:id'): image_rid})
    etree.SubElement(node, '{' + O + '}OLEObject', {
        'Type':'Embed', 'ProgID':'ChemWindow.Document', 'ShapeID':shape_id,
        'DrawAspect':'Content', 'ObjectID':str(number), qn('r:id'):ole_rid})
    paragraph.add_run()._r.append(node)
    return node, opaque


def vector_facade(tmp_path, *, role='handout', repeated=False, mutate=None):
    doc = Document(); doc.add_heading('知识点1 电离平衡', 3)
    paragraph = doc.add_paragraph('先比较CH'); paragraph.add_run('3').font.subscript = True
    paragraph.add_run('COOH的两方向变化。')
    vector = palette_wmf(); node, opaque = add_object(doc, paragraph, vector, 1)
    if repeated:
        paragraph.add_run('再看同图另一处。'); add_object(doc, paragraph, vector, 2)
    paragraph.add_run('完整尾部条件。')
    doc.add_paragraph('下一完整知识段。')
    if mutate: mutate(doc, node)
    facade = make_facade(word_bytes(doc), tmp_path/'source', role=role)
    reader = PreparationSourcesService(facade.paths.workspace_root)
    facade.imported_word_preview = lambda *args, read_only=False: reader.word_preview_bytes(facade.value.data, facade.value.name)
    facade.imported_word_asset = lambda *args: reader.word_asset_bytes(facade.value.data, args[2], render_metafiles=True)
    row = images.handout_image_options(facade)['images'][0]
    return facade, row, vector, opaque


windows = pytest.mark.skipif(not hasattr(ctypes, 'WinDLL'), reason='native Windows GDI+ unavailable')


@windows
def test_source_bound_vector_preview_is_memory_only_and_real_nonblank_png(tmp_path):
    facade, row, vector, _ = vector_facade(tmp_path)
    before = facade.original.stat(); original = facade.original.read_bytes()
    picture = images.handout_image_preview(facade, row)
    assert picture['source'] == row and row['mime_type'] == 'image/x-wmf'
    assert picture['derived_preview'] is True and picture['renderer_revision'] == renderer.RENDERER_REVISION
    assert picture['original_sha256'] == sha256(vector).hexdigest()
    assert picture['preview_sha256'] == sha256(picture['bytes']).hexdigest() != row['sha256']
    with Image.open(BytesIO(picture['bytes'])) as im:
        assert im.format == 'PNG' and any(low != high for low, high in im.getextrema())
    assert facade.original.read_bytes() == original and facade.original.stat().st_mtime_ns == before.st_mtime_ns
    assert not facade.paths.task_root.exists()


@windows
def test_standalone_derivative_exact_evidence_duplicate_new_use_and_teacher_roles(tmp_path):
    facade, row, vector, _ = vector_facade(tmp_path)
    selection = {'source':row, 'title':'两方向如何达到平衡？', 'purpose':'比较原箭头方向。', 'teacher_only':True}
    result = images.import_handout_image(facade, [], [], selection)
    item = result['resources'][0]; asset = result['image_assets'][0]
    evidence = item['original']['evidence']; derivative = evidence['image_derivation']
    assert evidence['handout_source'] == row and validate_resource(item) == item
    assert derivative == {'schema':images.DERIVATION_SCHEMA, 'original_sha256':sha256(vector).hexdigest(),
        'original_mime_type':'image/x-wmf', 'original_bytes_count':len(vector), 'preview_sha256':asset['sha256'],
        'preview_mime_type':'image/png', 'preview_bytes_count':len(images.handout_image_preview(facade,row)['bytes']),
        'renderer_revision':renderer.RENDERER_REVISION}
    assert images.handout_image_roles([item], row) == ['teacher_note']
    again = images.import_handout_image(facade, [asset], [item], {**selection, 'teacher_only':False})
    assert again['duplicate'] and again['resources'] == []
    another = images.import_handout_image(facade, [asset], [item], {**selection, 'purpose':'再次比较。', 'teacher_only':False})
    assert another['image_assets'] == [asset] and all(b['role']=='teacher_note' for b in another['resources'][0]['blocks'])
    assert another['resources'][0]['original']['evidence']['source_roles'] == ['teacher_note']
    assert sha256(PreparationImageStore(facade.paths.task_root/'preparation-v1/images').load(asset)).hexdigest() == asset['sha256']


@windows
@pytest.mark.parametrize('field,value', [('original_sha256','f'*64), ('original_mime_type','image/png'),
    ('original_bytes_count',True), ('preview_sha256','f'*64), ('preview_mime_type','image/jpeg'),
    ('preview_bytes_count',0), ('renderer_revision',''), ('extra',True)])
def test_derivation_evidence_cannot_replace_raw_source_or_display_identity(tmp_path, field, value):
    facade, row, _, _ = vector_facade(tmp_path)
    selection = {'source':row,'title':'问题？','purpose':'','teacher_only':False}
    item = images.import_handout_image(facade,[],[],selection)['resources'][0]
    item['original']['evidence']['image_derivation'][field] = value
    item['original_sha256'] = digest(item['original'])
    with pytest.raises(ValueError): validate_resource(item)


@windows
@pytest.mark.parametrize('field,value', [('original_sha256','f'*64), ('original_mime_type','image/png'),
    ('preview_sha256','f'*64), ('derived_preview',False), ('renderer_revision',None)])
def test_changed_decoder_payload_rejected_without_publishing(tmp_path, field, value):
    facade, row, _, _ = vector_facade(tmp_path)
    original = facade.imported_word_asset
    def altered(*args):
        payload = original(*args); payload[field] = value; return payload
    facade.imported_word_asset = altered
    with pytest.raises(ValueError): images.handout_image_preview(facade,row)
    assert not facade.paths.task_root.exists()


@windows
def test_cancelled_vector_preview_never_creates_image_store(tmp_path):
    facade, row, _, _ = vector_facade(tmp_path); event = Event(); event.set()
    with read_cancel_scope(event, check_boundaries=False):
        with pytest.raises(ReadCancelled): images.handout_image_preview(facade,row)
    assert not facade.paths.task_root.exists()


@windows
@pytest.mark.parametrize('role', ['handout', 'answer'])
def test_whole_knowledge_retains_repeated_inline_objects_and_word_ppt_pixels(tmp_path, app, role):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition, layout_composition, write_pptx
    from test_composition_image_display import word_output, word_image_hashes
    facade, row, vector, _ = vector_facade(tmp_path, role=role, repeated=True)
    request = request_for(facade); raw_before = facade.original.read_bytes()
    result = knowledge.import_handout_knowledge(facade,[request],[],[])
    item = result['resources'][0]; asset = result['image_assets'][0]
    refs = [r for b in item['original']['evidence']['section_snapshot']['blocks'] for r in b['images']]
    assert len(refs) == 2 and len(result['image_assets']) == 1 and validate_resource(item) == item
    assert [r['position'] for r in refs] == [1,2]
    assert all(r['original_sha256']==sha256(vector).hexdigest() and r['preview_sha256']==asset['sha256']
        and r['derived_preview'] and r['renderer_revision']==renderer.RENDERER_REVISION for r in refs)
    assert refs[0]['source_object']['ole']['sha256'] != refs[1]['source_object']['ole']['sha256']
    assert refs[0]['container_xml_sha256'] != refs[1]['container_xml_sha256']
    assert all(r['source_object']['binary_interpreted'] is False for r in refs)
    assert facade.original.read_bytes() == raw_before
    assert images.handout_image_roles([item],row) == (['teacher_note'] if role=='answer' else ['knowledge'])
    data = {asset['asset_id']:PreparationImageStore(facade.paths.task_root/'preparation-v1/images').load(asset)}
    composition = new_composition('16:9')
    composition = add_resource(composition, item,
        {'page_id':composition['pages'][0], 'x':12, 'y':12, 'width':296, 'height':156})
    composition = arrange_composition(composition, [asset], data)
    teacher = layout_composition(composition,[asset],data,student=False)
    assert not teacher['errors']
    path = tmp_path/'native.pptx'; write_pptx(teacher,path,data)
    with ZipFile(path) as package:
        assert {sha256(package.read(n)).hexdigest() for n in package.namelist() if n.startswith('ppt/media/')} == {asset['sha256']}
    teacher_word, student_word = word_output(composition,[asset],data,tmp_path/'word')
    assert word_image_hashes(teacher_word) == [asset['sha256']]*2
    assert word_image_hashes(student_word) == ([] if role=='answer' else [asset['sha256']]*2)
    if role=='handout':
        paragraph = next(p for p in student_word.paragraphs if p.text.startswith('先比较CH'))
        assert any(run.text=='3' and run.font.subscript for run in paragraph.runs)
        assert paragraph.text.endswith('完整尾部条件。')


@windows
@pytest.mark.parametrize('change', ['shape_id', 'unknown_progid', 'link', 'extra_shape', 'extra_relation', 'crop'])
def test_chemwindow_binding_refuses_ambiguous_transformed_or_unrelated_objects(tmp_path, change):
    def mutate(doc,node):
        ole = node.find('.//{' + O + '}OLEObject'); shape = node.find('.//{' + V + '}shape')
        if change=='shape_id': ole.set('ShapeID','different')
        elif change=='unknown_progid': ole.set('ProgID','Unrelated.Document')
        elif change=='link': ole.set('Type','Link')
        elif change=='extra_shape': etree.SubElement(node,'{' + V + '}shape')
        elif change=='extra_relation': shape.set(qn('r:id'),'rId999')
        else: shape.set('rotation','10')
    facade, _, _, _ = vector_facade(tmp_path, mutate=mutate)
    with pytest.raises(ValueError): knowledge.import_handout_knowledge(facade,[request_for(facade)],[],[])
    assert not facade.paths.task_root.exists()


@windows
def test_global_equation_object_review_does_not_admit_chemwindow(tmp_path):
    from integrations.deeptutor_shchem_v1.offline_word_image_label_review import _objects, OfflineWordLabelReviewError
    from integrations.deeptutor_shchem_v1.desktop_preparation_sources import _body_blocks
    facade, row, _, _ = vector_facade(tmp_path)
    doc = Document(BytesIO(facade.value.data)); element = list(_body_blocks(doc._element.body))[1]
    with pytest.raises(OfflineWordLabelReviewError,match='同对象'):
        _objects(element,doc,2,[row])


from test_composition_render import app
