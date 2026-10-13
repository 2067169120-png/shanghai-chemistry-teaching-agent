"""Exact classroom citation projection; native sources and gates stay sealed."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import os
from pathlib import Path
from threading import RLock
from types import SimpleNamespace
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from PIL import Image
import pytest

from integrations.deeptutor_shchem_v1 import desktop_question_classroom_copy as classroom
from integrations.deeptutor_shchem_v1.desktop_composition_inline import NativeParagraphDisplay, display_tokens
from integrations.deeptutor_shchem_v1.desktop_composition_sources import import_questions, merge_question_result
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    FLAGS, block, digest, resource, validate_resource, new_composition, add_resource, visible_blocks)
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import PreparationSourcesService, _word_images
from integrations.deeptutor_shchem_v1.desktop_source_quality import apply_source_quality
from integrations.deeptutor_shchem_v1.desktop_word_question_index import index_word_questions
from integrations.deeptutor_shchem_v1.desktop_word_questions import WordQuestionService
from integrations.deeptutor_shchem_v1.offline_word_image_label_review import _objects


CITATION = '（24-25高三上·上海·阶段练习）'
REAL_SHA = '01f91a44ebbd0376c857cdbf8c5a186764a36f4a011c327a724b037c7c4f9c33'
SAMPLES = [
    (350, 361, '例1（双选）', CITATION,
     '乒乓球底板常用芳纶纤维、玻璃纤维、碳纤维等材料。下列说法正确的是'),
    (339, 348, '【变式训练2·变载体】', CITATION,
     '乒乓球胶水能将胶皮、海绵与底板优化组合成完整的乒乓球拍，对球拍质量起到了决定性作用。国际比赛规定使用更环保的无机胶水，无机胶水以无机物(多为水)作为溶剂，其中黏合剂的成分可能是'),
    (418, 427, '2．', '（2025·四川·高考真题）',
     '材料是科技发展的基础。下列属于金属材料的是'),
]


@pytest.fixture(autouse=True)
def no_application_construction(monkeypatch, tmp_path):
    from integrations.deeptutor_shchem_v1.desktop_facade import DesktopWorkbenchFacade
    from integrations.deeptutor_shchem_v1.desktop_state import DesktopStateStore
    from integrations.deeptutor_shchem_v1.desktop_editor_recovery import PreparationRecoveryStore
    def forbidden(*_args, **_kwargs):
        raise AssertionError('No complete facade, state or recovery constructor in this fixture')
    for cls in (DesktopWorkbenchFacade, DesktopStateStore, PreparationRecoveryStore):
        monkeypatch.setattr(cls, '__init__', forbidden)
    monkeypatch.setattr(tempfile, 'tempdir', str(tmp_path))


@pytest.fixture(scope='module')
def app():
    from PySide6.QtWidgets import QApplication
    from integrations.deeptutor_shchem_v1.desktop_workbench.typography import install_ui_font
    application = QApplication.instance() or QApplication([])
    install_ui_font(application)
    return application


def _evidence():
    return {'source_path': 'private-source.docx', 'source_sha256': 'a'*64,
        'source_revision': 'b'*64, 'question_key': 'synthetic-question',
        'scope_review_sha256': 'c'*64, 'scope_status': 'complete_native_scope',
        'hierarchy': dict.fromkeys(('paper_id', 'theme_big_question_id', 'printed_question_id', 'atomic_part_id'), 'unknown')}


def _plain(text, *, title='原完整题', extra=()):
    return resource('complete_question', title, [block('text', 'question', text, 1),
        *extra, block('text', 'answer', '【答案】A；仅教师答案', 9)], _evidence())


def _assert_original(before, after):
    assert before['original'] == after['original']
    assert before['original_sha256'] == after['original_sha256'] == digest(after['original'])
    assert {k: v for k, v in before.items() if k not in {'title', 'blocks'}} == {
        k: v for k, v in after.items() if k not in {'title', 'blocks'}}
    first = next(b['id'] for b in before['blocks'] if b['kind']=='text' and b['role']=='question')
    assert len(before['blocks']) == len(after['blocks'])
    for left, right in zip(before['blocks'], after['blocks']):
        assert {k: v for k, v in left.items() if k != 'value'} == {k: v for k, v in right.items() if k != 'value'}
        if left['id'] != first:
            assert left == right
    assert all(after[name] is value for name, value in FLAGS.items())
    assert validate_resource(after) == after


@pytest.mark.parametrize('start,end,lead,citation,stem', SAMPLES)
def test_exact_real_sample_wording_only_removes_the_source_parenthesis(start, end, lead, citation, stem):
    item = _plain(lead + citation + stem)
    before = deepcopy(item)
    actual = classroom.classroom_question_copy(item)
    assert item == before
    assert classroom.printed_source_span(item['blocks'][0]['value']) == (len(lead), len(lead+citation))
    assert actual['blocks'][0]['value'] == lead + stem
    assert actual['title'] == {350: '例题1', 339: '例题', 418: '例题2'}[start]
    assert stem not in actual['title'] and citation not in actual['title']
    _assert_original(before, actual)
    assert classroom.classroom_question_copy(actual) == actual


@pytest.mark.parametrize('text', [
    '例1（双选）将1 mol H2通入溶液（25℃、101 kPa）。',
    '例1（2025年高考时消耗1 mol H2）计算物质的量。',
    '例1（高三学生将1 mol H2通入溶液）观察变化。',
    '例1（2025·上海·阶段练习，温度25℃）计算。',
    '例1（2025·上海·高考真题；0.10 mol）计算。',
    '例1（2025·mol·高考真题）计算。',
    '例1（2025·实验消耗氧气·高考真题）计算。',
    '例1（2025·未知学校·期末）计算。',
    '例1（2025·四川·高考真题)计算。',
    '例1（2025·四川·高考真题）',
    '例1（25-24高三上·上海·阶段练习）计算。',
    '例1（2024-2026高三上·上海·阶段练习）计算。',
    '例1（2025·四川·高考真题·数据）计算。',
    '例1（2025四川高考真题）计算。',
    '例1（2025·四川·高考真题） （2024·上海·等级考）计算。',
    '先读取条件（2025·四川·高考真题）再计算。',
    'A．（2025·四川·高考真题）氢气的相对分子质量为2。',
    '例1（2025\n·四川·高考真题）计算。',
    '例1（2025·四川·高考真题（25℃））计算。',
    '例1（双选）（298·四川·高考真题）计算。',
    '(1）（2025·四川·高考真题）计算。',
    '（1)（2025·四川·高考真题）计算。',
])
def test_unknown_or_chemical_parentheses_and_nonopening_citations_remain_exact(text):
    item = _plain(text)
    assert classroom.printed_source_span(text) is None
    assert classroom.classroom_question_copy(item) == item


@pytest.mark.parametrize('lead,citation', [
    ('  1．（多选）', '(2025·四川·高考真题)'),
    ('【例1】', '（2025·全国·高考）'),
    ('（1）', '（24-25高二下·浙江·期末）'),
    ('例题 1．', '（2024-2025学年高三上学期·上海·徐汇区·一模）'),
    ('', '（2025·高考真题）'),
    ('（双选）', '（1999-2000学年高三上·上海·阶段测试）'),
])
def test_bounded_known_forms_keep_number_instruction_spacing_and_chemical_condition(lead, citation):
    stem = '（25℃时）0.10 mol·L⁻¹溶液中有H_{2}O与Na⁺。'
    item = _plain(lead+citation+stem)
    out = classroom.classroom_question_copy(item)
    assert out['blocks'][0]['value'] == lead+stem
    _assert_original(item, out)


def test_only_the_first_question_paragraph_is_eligible_not_shared_material_option_or_answer():
    raw = '例1'+CITATION+'计算。'
    shared = block('text', 'shared_material', raw, 0)
    item = resource('complete_question', '原题', [shared,
        block('text', 'question', '例1：保留这里的题干。', 1),
        block('text', 'question', raw, 2), block('text', 'answer', raw, 3)], _evidence())
    assert classroom.classroom_question_copy(item) == item


@pytest.mark.parametrize('edit', ['body', 'title'])
def test_existing_working_edits_are_never_replaced(edit):
    item = _plain('例1'+CITATION+'计算。')
    if edit == 'body': item['blocks'][0]['value'] = '教师补充条件。'+item['blocks'][0]['value']
    else: item['title'] = '应怎样比较这些材料？'
    assert classroom.classroom_question_copy(item) == item


def _field(p, code):
    for kind, value in [('fldChar', 'begin'), ('instrText', code), ('fldChar', 'separate'), ('fldChar', 'end')]:
        run = OxmlElement('w:r'); node = OxmlElement('w:'+kind)
        if kind == 'fldChar': node.set(qn('w:fldCharType'), value)
        else: node.text = value
        run.append(node); p._p.append(run)


def _native_item(tmp_path, *, rich=True, eq=True, pictures=2, long=False, prefix_script=False):
    buf = BytesIO(); Image.new('RGB', (80, 40), '#2853ab').save(buf, format='PNG'); data = buf.getvalue()
    doc = Document()
    if prefix_script:
        p = doc.add_paragraph('例1（'); p.add_run('2025').font.subscript=True
        p.add_run('·四川·高考真题）')
    else: p = doc.add_paragraph('例1（双选）'+CITATION)
    p.add_run('保留0.10 mol的H')
    p.add_run('2').font.subscript = rich
    p.add_run('O与Na'); p.add_run('+').font.superscript = rich
    if eq: _field(p, r' EQ \o(――→,\s\up7(催化剂)) ')
    for _ in range(pictures):
        p.add_run().add_picture(BytesIO(data), width=Mm(12))
        p.add_run('图后条件不丢失。')
    if long: p.add_run('完整条件。'*4000)
    answer = doc.add_paragraph('【答案】仅教师答案：H'); answer.add_run('2').font.subscript=True
    answer.add_run('O，不能向学生显示。')
    raw = BytesIO(); doc.save(raw); native = NativeParagraphDisplay(raw.getvalue())
    store = PreparationImageStore(tmp_path/'native-images')
    asset = store.import_bytes(data, '原生合成图', '私有测试文档', '验证出现次数')
    all_assets = _word_images(native.elements, native.doc)
    blocks = []
    for index, role in ((1, 'question'), (2, 'answer')):
        element = native.elements[index-1]
        text = block('text', role, native.reader.read(element).text, index)
        objects = _objects(element, native.doc, index, [a for a in all_assets if a['block_index']==index])
        image_blocks = [block('image', role, asset['asset_id'], index) for _ in objects]
        native.bind(text, image_blocks, objects)
        blocks.extend([text, *image_blocks])
    item = resource('complete_question', '原题完整来源', blocks, _evidence())
    return item, [asset] if pictures else [], {asset['asset_id']: data} if pictures else {}, raw.getvalue()


def _tokens(item, *, student):
    by_id = {b['id']: b for b in item['blocks']}
    original = {b['id']: b['value'] for b in item['original']['blocks']}
    return [(b['id'], display_tokens(b, by_id, original[b['id']]))
            for b in visible_blocks(item, student=student) if b['kind']=='text']


def test_native_scripts_eq_and_repeated_same_byte_picture_occurrences_are_proved(tmp_path):
    item, assets, data, raw = _native_item(tmp_path)
    out = classroom.classroom_question_copy(item)
    assert out != item and CITATION not in out['blocks'][0]['value']
    _assert_original(item, out)
    p = item['blocks'][0]['presentation']
    assert len(p['scripts']) == 2 and len(p['eqs']) == 1 and len(p['images']) == 2
    assert p == out['blocks'][0]['presentation']
    before = _tokens(item, student=True)[0][1]; after = _tokens(out, student=True)[0][1]
    assert [t for t in before if t['kind'] != 'text' or t['script'] != 'baseline'] == [
        t for t in after if t['kind'] != 'text' or t['script'] != 'baseline']
    assert len({t['image_block_id'] for t in after if t['kind']=='image'}) == 2
    assert len({t['asset_id'] for t in after if t['kind']=='image'}) == 1
    assert _tokens(item, student=False)[1:] == _tokens(out, student=False)[1:]
    assert all('仅教师答案' not in t.get('text', '') for _, ts in _tokens(out, student=True) for t in ts)


@pytest.mark.parametrize('variant', ['long', 'prefix_script'])
def test_native_rich_edits_without_exact_proof_keep_source_wording(tmp_path, variant):
    item, *_ = _native_item(tmp_path, pictures=0, **{variant: True})
    assert classroom.classroom_question_copy(item) == item


def test_crossing_native_span_is_not_projected(tmp_path):
    item, *_ = _native_item(tmp_path)
    row = item['blocks'][0]; span = row['presentation']['scripts'][0]
    assert not classroom._deletion_preserves_display(row, {b['id']: b for b in item['blocks']}, span['start']+1, span['end'])
    marker = row['presentation']['images'][0]['placeholder']; start = row['value'].index(marker)
    assert not classroom._deletion_preserves_display(row, {b['id']: b for b in item['blocks']}, start, start+len(marker))


def test_renderer_token_disagreement_does_not_silently_drop_rich_content(tmp_path, monkeypatch):
    item, *_ = _native_item(tmp_path)
    native = classroom.display_tokens
    def wrong(row, by_id, original_text=None):
        result = native(row, by_id, original_text)
        return [t for t in result if t['kind'] != 'eq'] if row['value'] != original_text else result
    monkeypatch.setattr(classroom, 'display_tokens', wrong)
    assert classroom.classroom_question_copy(item) == item


def _port(raw, root):
    """Real native compiler ports with an explicit frozen state; no Facade."""
    root.mkdir(parents=True, exist_ok=True)
    archive = root/'lecture.docx'; archive.write_bytes(raw)
    reader = PreparationSourcesService(root/'workspace')
    state = SimpleNamespace(snapshot=lambda: {'drafts': {}})
    cache = {}
    def source(batch, identity):
        assert (batch, identity) == ('BATCH', 'SOURCE')
        data = archive.read_bytes()
        return SimpleNamespace(content=data, filename='第08讲 高分子有机物（复习讲义）（解析版）.docx',
            role='handout', source_sha256=sha256(data).hexdigest(), effective_source_file_id='SOURCE')
    def preview(batch, identity, *, read_only=False):
        assert read_only is True
        value = source(batch, identity)
        if value.source_sha256 not in cache:
            result = reader.word_preview_bytes(value.content, value.filename)
            cache[value.source_sha256] = (result, index_word_questions(result))
        return deepcopy(cache[value.source_sha256][0])
    def catalog(locations=None, *, read_only=False):
        assert read_only is True and locations == {'BATCH': {'SOURCE'}}
        current = source('BATCH', 'SOURCE'); value = preview('BATCH', 'SOURCE', read_only=True)
        indexed = deepcopy(cache[current.source_sha256][1]); rows = []
        for entry in indexed:
            row = apply_source_quality(entry, {})
            row.update(batch_id='BATCH', archive_source_id='SOURCE', source_name=current.filename,
                source_id=current.source_sha256, source_block_count=len(value['blocks']), shared_material_policy='self_contained')
            rows.append(row)
        return {'items': rows, 'sources': [{'source_id': current.source_sha256, 'batch_id': 'BATCH'}]}, {
            current.source_sha256: (current, 'BATCH', value, indexed)}
    words = SimpleNamespace(_lock=RLock(), reader=reader, _catalog=catalog, state=state,
        attribute_store=SimpleNamespace(get=lambda *_a, **_k: None),
        _read_attribute_catalog=lambda: {'knowledge_points': [], 'nodes': []})
    def resolve(selections, *, read_only=True, **_kw):
        rows, inventory = catalog({'BATCH': {'SOURCE'}}, read_only=read_only)
        result = []
        for selection in selections:
            matches = [r for r in rows['items'] if r['key']==selection['key'] and r['revision']==selection['revision']]
            if len(matches) != 1: raise ValueError('exact question version changed')
            result.append({**matches[0], 'points': selection.get('points', 2)})
        return result, inventory
    words._resolve = resolve
    words.reference = lambda selections, include_images=True: WordQuestionService._compile_reference(
        words, selections, include_images=include_images)[0]
    facade = SimpleNamespace(paths=SimpleNamespace(task_root=root/'tasks', workspace_root=root/'workspace'),
        _state=state, _word_questions=lambda: words, _imported_word_source=source,
        imported_word_path=lambda *_a: str(archive), imported_word_preview=preview,
        imported_word_asset=lambda batch, identity, asset: reader.word_asset_bytes(source(batch, identity).content, asset, render_metafiles=True),
        imported_word_sources=lambda _batch: [{'source_id': 'SOURCE', 'source_name': source('BATCH','SOURCE').filename,
            'source_sha256': source('BATCH','SOURCE').source_sha256, 'role': 'handout'}],
        list_imported_word_batches=lambda: [SimpleNamespace(batch_id='BATCH')], word_question_reference=words.reference)
    return facade, words, archive


@pytest.fixture(scope='module')
def real_raw():
    value = os.environ.get('SHCHEM_QUESTION_COPY_SOURCE')
    if not value: pytest.skip('Exact PKG-072 solution source not configured')
    path = Path(value); data = path.read_bytes(); before = (sha256(data).hexdigest(), len(data), path.stat().st_mtime_ns)
    assert before[0] == REAL_SHA
    yield data
    assert (sha256(path.read_bytes()).hexdigest(), path.stat().st_size, path.stat().st_mtime_ns) == before


@pytest.mark.parametrize('start,end,lead,citation,stem', SAMPLES)
def test_real_complete_question_import_preserves_every_source_block_and_picture(real_raw, tmp_path, start, end, lead, citation, stem):
    facade, words, archive = _port(real_raw, tmp_path/'source')
    rows = words._catalog({'BATCH': {'SOURCE'}}, read_only=True)[0]['items']
    chosen = next(r for r in rows if r['block_start']==start)
    assert chosen['block_end'] == end
    selection = [{**{k: chosen[k] for k in ('key', 'revision')}, 'points': 2}]
    ref = words.reference(selection, include_images=True)
    result = import_questions(facade, selection, [], expected_reference=ref)
    item = result['resources'][0]; original = item['original']
    assert original['blocks'][0]['value'] == lead+citation+stem
    assert item['blocks'][0]['value'] == lead+stem and citation not in item['title']
    assert [b['source_index'] for b in original['blocks'] if b['kind']=='text'] == list(range(start,end+1))
    assert [b['value'] for b in original['blocks'] if b['kind']=='text'] == [b['text'] for key in ('context_blocks','question_blocks','answer_blocks') for b in chosen[key]]
    assert original['evidence']['source_sha256'] == REAL_SHA
    assert original['evidence']['source_revision'] == chosen['source_revision']
    assert len([b for b in original['blocks'] if b['kind']=='image']) == sum(len(b['assets']) for key in ('context_blocks','question_blocks','answer_blocks') for b in chosen[key])
    native = {**deepcopy(item), 'title': original['title'], 'blocks': deepcopy(original['blocks'])}
    _assert_original(native, item)
    assert classroom.classroom_question_copy(native) == item
    assert archive.read_bytes() == real_raw
    store = PreparationImageStore(facade.paths.task_root/'preparation-v1/images')
    assert all(sha256(store.load(a)).hexdigest()==a['sha256'] for a in result['image_assets'])
    edited = deepcopy(item); edited['title']='老师的本课问题'; edited['blocks'][0]['value'] += '（教师明确补充）'
    assert merge_question_result(result, [edited], result['image_assets']) == {'resources': [], 'image_assets': result['image_assets']}
    assert edited['title']=='老师的本课问题' and edited['blocks'][0]['value'].endswith('（教师明确补充）')


def test_mixed_publication_recompile_stamp_repeat_and_cancel_keep_the_complete_source(tmp_path):
    from integrations.deeptutor_shchem_v1 import desktop_handout_material_handoff as handoff
    from integrations.deeptutor_shchem_v1.desktop_handout_knowledge import handout_knowledge_sources, handout_knowledge_outline
    doc = Document(); doc.add_heading('知识点1 材料', 1); doc.add_paragraph('材料可根据组成分类。')
    doc.add_heading('例题', 1); doc.add_paragraph('1．'+CITATION+'下列属于金属材料的是？')
    png = BytesIO(); Image.new('RGB', (60,30), '#458ad0').save(png,'PNG')
    doc.add_paragraph().add_run().add_picture(BytesIO(png.getvalue()), width=Mm(20))
    doc.add_paragraph('A．铝合金 B．塑料 C．玻璃 D．橡胶'); doc.add_paragraph('【答案】A。')
    raw=BytesIO(); doc.save(raw); facade, words, archive = _port(raw.getvalue(), tmp_path/'mixed')
    chosen_source = handout_knowledge_sources(facade)['sources'][0]
    outline = handout_knowledge_outline(facade, chosen_source); source = outline['source']
    section = next(s for s in outline['sections'] if s['title']=='知识点1 材料')
    row = next(r for r in words._catalog({'BATCH': {'SOURCE'}}, read_only=True)[0]['items'] if r['question_blocks'][0]['text'].startswith('1．'))
    selected = {k: row[k] for k in ('key','revision','source_sha256')}
    selected.update(points=2, scoped_source_document=deepcopy(source))
    request = {'schema': handoff.SCHEMA, 'source': source,
        'knowledge_requests': [{'source': deepcopy(source), 'outline_revision': outline['outline_revision'],
            'section': deepcopy(section), 'title': '材料怎样分类？', 'teacher_only': False}],
        'image_selections': [], 'question_selections': [selected],
        'selection_order': [{'kind':'question','index':0},{'kind':'knowledge','index':0}]}
    publication = handoff.prepare_handout_materials(facade, request, [], [])
    result = publication.publish()
    assert [r['original']['kind'] for r in result['resources']] == ['complete_question','handout_knowledge']
    question = result['resources'][0]
    assert CITATION in question['original']['blocks'][0]['value'] and CITATION not in question['blocks'][0]['value']
    assert handoff._stamp(result) == handoff._stamp(handoff._compile(facade, request, [], [])[0])
    publication.recheck_published()
    store = PreparationImageStore(facade.paths.task_root/'preparation-v1/images')
    assert all(store.load(a) for a in result['image_assets'])
    publication.discard()
    assert not list(store.root.glob('*.image'))
    assert archive.read_bytes() == raw.getvalue()
    publication = handoff.prepare_handout_materials(facade, request, [], [])
    accepted = publication.publish(); publication.accept()
    duplicate = handoff.prepare_handout_materials(facade, request, accepted['image_assets'], accepted['resources'])
    assert duplicate.result['resources'] == []
    duplicate.discard()
    assert all(store.load(a) for a in accepted['image_assets'])


def test_teacher_and_student_native_word_ppt_outputs_use_work_text_and_preserve_rich_roles(tmp_path, app):
    from pptx import Presentation
    from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition, layout_composition, write_pptx
    from integrations.deeptutor_shchem_v1.desktop_lesson_design import new_design
    from integrations.deeptutor_shchem_v1.desktop_lesson_composition import enable_composition
    from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents
    item, assets, data, _raw = _native_item(tmp_path)
    item = classroom.classroom_question_copy(item)
    value = new_composition(); value = add_resource(value,item,{'page_id':value['pages'][0],'x':12,'y':12,'width':186,'height':273})
    value = arrange_composition(value, assets, data)
    payload = {'topic':'材料的结构与性质', 'audience':'测试课堂', 'lesson_timing':'1课时',
        'objective':'比较材料', 'lesson_route':'复习', 'materials':'私有合成来源',
        'advanced':dict.fromkeys(('learning_and_experiment','template_and_delivery','homework_and_strategy'),'')}
    payload['lesson_design']=enable_composition(new_design(payload)); payload['lesson_design']['composition']=value; payload['image_assets']=assets
    before = deepcopy(payload)
    write_documents(payload, tmp_path, data)
    for student in (False,True):
        native = Document(tmp_path/('student_worksheet.docx' if student else 'lesson_plan.docx'))
        text = '\n'.join(p.text for p in native.paragraphs)
        assert CITATION not in text and '例1（双选）' in text
        assert '0.10 mol的H2O与Na+' in text
        assert ('仅教师答案' in text) is (not student)
        assert len(native.inline_shapes) == 3  # EQ rendering plus two original occurrences.
        paragraph = next(p for p in native.paragraphs if '保留0.10 mol' in p.text and any(r.font.subscript for r in p.runs))
        assert any(r.text=='2' and r.font.subscript for r in paragraph.runs)
        assert any(r.text=='+' and r.font.superscript for r in paragraph.runs)
        layout=layout_composition(value,assets,data,student=student); assert not layout['errors']
        path=tmp_path/('student.pptx' if student else 'teacher.pptx'); write_pptx(layout,path,data)
        deck=Presentation(path); flat=''.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
        assert CITATION not in flat and '0.10 mol的H2O与Na+' in flat
        assert '催化剂' in flat and ('仅教师答案' in flat) is (not student)
        assert sum(s.shape_type==13 for slide in deck.slides for s in slide.shapes)==2
    assert payload == before
