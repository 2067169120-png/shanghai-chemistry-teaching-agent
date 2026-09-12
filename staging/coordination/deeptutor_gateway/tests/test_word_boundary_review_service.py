"""Native source range confirmations persist without modifying originals."""

from hashlib import sha256

import pytest
from docx import Document
from test_desktop_visual_import_facade import FakeProviderStore, _facade
from test_desktop_visual_import_facade import (
    desktop_paths as desktop_paths,  # noqa: PLC0414
)

from integrations.deeptutor_shchem_v1.desktop_word_questions import WordQuestionError


def _source(desktop_paths, tmp_path):
    document = Document()
    for text in (
        "【例1】哪一项符合条件？",
        "A．甲 B．乙 C．丙",
        "(3)C",
        "(3)这是原文独立给出的解析文字，说明丙符合给定条件。",
    ):
        document.add_paragraph(text)
    source = tmp_path / "范围核对样例.docx"
    document.save(source)
    provider = FakeProviderStore(configured=False)
    facade = _facade(desktop_paths, provider)
    facade.save_visual_import_batch(handout_files=(source,), source_type="教师讲义")
    return facade, provider, source


def _choice(item):
    return {"key": item["key"], "revision": item["revision"], "points": 3}


def test_confirmed_range_reopens_and_exports_original_answer_in_teacher_only(
    desktop_paths, tmp_path
):
    facade, provider, source = _source(desktop_paths, tmp_path)
    source_hash = sha256(source.read_bytes()).hexdigest()
    (item,) = facade.word_question_catalog()["items"]
    assert not item["export_ready"]
    with pytest.raises(WordQuestionError, match="边界"):
        facade.word_question_export("原文练习", [_choice(item)])
    updated = facade.word_question_update_range(
        item["key"],
        item["revision"],
        block_start=1,
        question_end=2,
        answer_start=3,
        block_end=4,
        reviewed_issues=["unmarked_answer"],
    )
    assert updated["export_ready"]
    with pytest.raises(WordQuestionError, match="变化"):
        facade.word_question_reference([_choice(item)])
    reopened = _facade(desktop_paths, provider)
    (current,) = reopened.word_question_catalog()["items"]
    assert current["revision"] == updated["revision"]
    assert current["boundary_review"]["issues"] == ["unmarked_answer"]
    exported = reopened.word_question_export("原文练习", [_choice(current)])
    student = Document(exported["student_path"])
    teacher = Document(exported["teacher_path"])
    student_text = "\n".join(p.text for p in student.paragraphs)
    teacher_text = "\n".join(p.text for p in teacher.paragraphs)
    assert "哪一项符合条件" in student_text
    assert "(3)C" not in student_text and "原文独立给出" not in student_text
    assert "(3)C" in teacher_text and "原文独立给出" in teacher_text
    assert "练习 1  3 分" in teacher_text and "3 分" not in student_text
    assert sha256(source.read_bytes()).hexdigest() == source_hash
    assert provider.borrow_calls == 0


def test_unknown_review_rejected_without_changing_saved_ranges(desktop_paths, tmp_path):
    facade, provider, _ = _source(desktop_paths, tmp_path)
    (item,) = facade.word_question_catalog()["items"]
    before = facade.state_store.snapshot()
    with pytest.raises(ValueError, match="确认项"):
        facade.word_question_update_range(
            item["key"],
            item["revision"],
            block_start=1,
            question_end=2,
            answer_start=3,
            block_end=4,
            reviewed_issues=["chemistry_approved"],
        )
    assert facade.state_store.snapshot() == before
    assert provider.borrow_calls == 0
