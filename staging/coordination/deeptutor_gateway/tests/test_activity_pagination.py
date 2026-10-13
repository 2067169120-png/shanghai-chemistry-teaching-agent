"""Short lesson activities retain task context in real composition outputs."""
from copy import deepcopy
import pytest
from docx import Document
from test_composition_render import app, text
from test_composition_flow import personal, composition
from test_lesson_design import sample
from integrations.deeptutor_shchem_v1.desktop_composition_render import (
    _content, arrange_composition, layout_composition, PADDING)
from integrations.deeptutor_shchem_v1.desktop_composition_activities import activity_resource
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    enable_composition, add_resource, block, resource_frames)
from integrations.deeptutor_shchem_v1.desktop_lesson_design import update_node
from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents

TASK='读表1.1，把熔点、沸点与分解炭化分开记录。为什么单凭低熔点不能判定有机物？给出教材例题中的反例。'
TEACHER='教师提示：先收集理由，再请学生用原页反例修正判断。'

def activity(task=TASK):
    p=sample();plan=enable_composition(p['lesson_design'])
    plan=update_node(plan,plan['nodes'][0]['id'],{'title':'独立练习','minutes':None,
        'teacher_action':TEACHER,'student_task':task,'expected_output':'区分三种表述，并指向教材提供的反例。',
        'criteria':'解释使用具体材料与条件，避免全称概括。','material_text':'','image_ids':[],
        'teacher_answer':'','notes':''})
    return p,plan,activity_resource(plan,plan['nodes'][0]['id'],[])

@pytest.mark.parametrize('student',[False,True])
def test_whole_short_activity_moves_to_fresh_page_after_other_material(app,student):
    _,_,item=activity();total=sum(c['height'] for c in _content(item,296,{}, {},student=False,format='16:9'))
    for n in range(1,20):
        filler=personal('已有教材材料',[block('text','personal','已有条件。\n'*n)])
        remaining=156-8-2*PADDING-sum(c['height'] for c in _content(filler,296,{}, {},student=False,format='16:9'))
        if 30<remaining<total:break
    else:pytest.fail('没有找到能放标题但放不下完整短活动的页尾')
    before=composition(filler,item);arranged=arrange_composition(before,[],{})
    layout=layout_composition(arranged,[],{},student=student)
    frames=resource_frames(arranged['placements'][item['id']]);assert len(frames)==1
    assert frames[0]['page_id']!=resource_frames(arranged['placements'][filler['id']])[0]['page_id']
    assert not layout['errors'] and arranged['resources']==before['resources']
    task=next(b for b in item['blocks'] if b['role']=='student_task')
    assert text(layout,task['id'])==task['value']
    assert {e['page_id'] for page in layout['pages'].values() for e in page if e.get('block_id')==task['id']}=={frames[0]['page_id']}
    visible=''.join(e.get('text','') for page in layout['pages'].values() for e in page)
    assert (TEACHER in visible)==(not student)

def test_student_uses_later_frame_for_whole_task_without_altering_authored_layout(app):
    _,_,item=activity();value=composition(item);first=value['pages'][0];later='CP-'+'f'*32;value['pages'].append(later)
    value['placements'][item['id']]={'page_id':first,'x':12,'y':12,'width':296,'height':34,
        'continuations':[{'page_id':later,'x':12,'y':12,'width':296,'height':120}]}
    before=deepcopy(value);layout=layout_composition(value,[],{},student=True)
    task=next(b for b in item['blocks'] if b['role']=='student_task')
    pages={e['page_id'] for page in layout['pages'].values() for e in page if e.get('block_id')==task['id'] and e['kind']=='text'}
    assert pages=={later} and not layout['errors']
    assert text(layout,task['id'])==task['value'] and value==before

def test_long_activity_continues_without_shrinking_or_losing_text(app):
    _,_,item=activity('完整题设和测量条件继续保留。'*140);before=composition(item)
    value=arrange_composition(before,[],{});layout=layout_composition(value,[],{},student=False)
    task=next(b for b in item['blocks'] if b['role']=='student_task')
    assert len(resource_frames(value['placements'][item['id']]))>1 and not layout['errors']
    assert text(layout,task['id'])==task['value'] and value['resources']==before['resources']
    assert all(e['font_px']==32 for page in layout['pages'].values() for e in page if e.get('block_id')==task['id'] and e['kind']=='text')

def test_word_short_activity_keeps_title_context_and_task_as_one_group(app,tmp_path):
    p,plan,item=activity();plan['composition']=add_resource(plan['composition'],item,
        {'page_id':plan['composition']['pages'][0],'x':12,'y':12,'width':186,'height':273})
    p['lesson_design']=plan;p['image_assets']=[];write_documents(p,tmp_path,{})
    teacher=Document(tmp_path/'lesson_plan.docx');student=Document(tmp_path/'student_worksheet.docx')
    for doc,is_student in ((teacher,False),(student,True)):
        start=next(i for i,r in enumerate(doc.paragraphs) if r.text==item['title'])
        # The teacher document continues with consolidation/homework/board
        # sections. A resource keep chain must end before that next heading.
        end=next((i for i,r in enumerate(doc.paragraphs[start+1:],start+1)
            if r.style.name.startswith('Heading ')),len(doc.paragraphs))
        group=doc.paragraphs[start:end]
        assert all(r.paragraph_format.keep_together for r in group)
        assert all(r.paragraph_format.keep_with_next for r in group[:-1])
        assert group[-1].paragraph_format.keep_with_next is False
        expected=[item['title'],*([r['value'] for r in item['blocks']]
            if not is_student else ['学生任务：'+TASK])]
        assert [r.text for r in group]==expected
    teacher_text='\n'.join(r.text for r in teacher.paragraphs)
    assert TEACHER in teacher_text
    assert '3. 巩固小结' in teacher_text and '五、板书设计' in teacher_text
    assert TEACHER not in '\n'.join(r.text for r in student.paragraphs)

def test_word_long_activity_allows_task_to_continue_without_chaining_later_sections(app,tmp_path):
    task='完整题设和测量条件继续保留。'*140
    p,plan,item=activity(task);before=deepcopy(item)
    plan['composition']=add_resource(plan['composition'],item,
        {'page_id':plan['composition']['pages'][0],'x':12,'y':12,'width':186,'height':273})
    p['lesson_design']=plan;p['image_assets']=[];write_documents(p,tmp_path,{})
    for name in ('lesson_plan.docx','student_worksheet.docx'):
        doc=Document(tmp_path/name)
        paragraph=next(r for r in doc.paragraphs if r.text=='学生任务：'+task)
        assert not paragraph.paragraph_format.keep_together
        assert not paragraph.paragraph_format.keep_with_next
    assert item==before and p['lesson_design']['composition']['resources'][0]==before
