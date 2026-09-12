from __future__ import annotations

import hashlib
import io
import json
from copy import deepcopy

import pytest
from docx import Document
from docx.oxml import OxmlElement
from test_desktop_visual_import_facade import FakeProviderStore, _facade, _png
from test_desktop_visual_import_facade import (
    desktop_paths as desktop_paths,  # noqa: PLC0414
)

from integrations.deeptutor_shchem_v1.desktop_preparation_images import (
    MAX_IMAGES,
    PreparationImageStore,
)
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
    PreparationSourceError,
    PreparationSourcesService,
)
from integrations.deeptutor_shchem_v1.desktop_source_quality import QUALITY_PATH
from integrations.deeptutor_shchem_v1.desktop_word_source_reference import (
    WordSourceReferenceError,
    WordSourceReferenceService,
)


def _import(
    desktop_paths, tmp_path, *, images=2, missing=False, answer=False, long=False
):
    doc = Document()
    doc.add_heading("第二章 化学反应", 1)
    p = doc.add_paragraph("原教案完整知识总结：SO")
    p.add_run("4").font.subscript = True
    p.add_run("2−").font.superscript = True
    doc.add_paragraph("教师提问：为什么条件改变时反应结果不同？")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "比较知识点"
    table.cell(1, 0).text = "原始实验条件"
    table.cell(1, 1).text = "原始知识结论"
    for number in range(images):
        color = (number * 17 % 256, number * 11 % 256, number * 7 % 256)
        doc.add_paragraph().add_run().add_picture(io.BytesIO(_png(color)))
    if images:
        doc.add_paragraph("教师解析原文，勿提前展示给学生。")
        doc.add_paragraph().add_run().add_picture(io.BytesIO(_png((0, 0, 0))))
    if missing:
        p = doc.add_paragraph("未支持对象：")
        p.add_run()._r.append(OxmlElement("w:object"))
    if long:
        doc.add_paragraph("原教案必要条件" * 4000)
    doc.add_paragraph("原教案的最后一段不应丢失。")
    path = tmp_path / "已完成教案.docx"
    doc.save(path)
    provider = FakeProviderStore(configured=False)
    facade = _facade(desktop_paths, provider)
    source_files = {"handout_files": (path,)}
    if answer:
        question = Document()
        question.add_paragraph("对应题面：请观察图后分析。")
        question_path = tmp_path / "配套题面.docx"
        question.save(question_path)
        source_files = {"answer_files": (path,), "question_files": (question_path,)}
    receipt = facade.save_visual_import_batch(**source_files, source_type="教师讲义")
    source = next(
        row
        for row in facade.imported_word_sources(receipt.batch_id)
        if row["role"] == ("answer" if answer else "handout")
    )
    preview = facade.imported_word_preview(receipt.batch_id, source["source_id"])
    selection = {
        "batch_id": receipt.batch_id,
        "source_id": source["source_id"],
        "source_sha256": preview["source_sha256"],
        "revision": preview["revision"],
        "block_start": 1,
        "block_end": len(preview["blocks"]),
    }
    return facade, selection, preview, path, provider


def _reference(facade, selection, *, include_images=True):
    return facade.imported_word_image_reference(
        selection["batch_id"],
        selection["source_id"],
        selection["source_sha256"],
        selection["block_start"],
        selection["block_end"],
        selection["revision"],
        include_images=include_images,
    )


def _image_root(facade):
    return facade.paths.task_root / "preparation-v1" / "images"


def _from_document(desktop_paths, tmp_path, document):
    path = tmp_path / "原结构图文测试.docx"
    document.save(path)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    batch = facade.save_visual_import_batch(
        handout_files=(path,), source_type="教师讲义"
    )
    source = facade.imported_word_sources(batch.batch_id)[0]
    preview = facade.imported_word_preview(batch.batch_id, source["source_id"])
    return (
        facade,
        {
            "batch_id": batch.batch_id,
            "source_id": source["source_id"],
            "source_sha256": preview["source_sha256"],
            "revision": preview["revision"],
            "block_start": 1,
            "block_end": len(preview["blocks"]),
        },
        preview,
    )


def test_original_teaching_text_tables_and_every_image_position_are_preserved(
    desktop_paths, tmp_path
):
    facade, selection, preview, path, provider = _import(desktop_paths, tmp_path)
    original = path.read_bytes()
    state = facade.state_store.snapshot()
    ref = _reference(facade, selection)
    assert set(ref) == {
        "materials",
        "warnings",
        "source_selection",
        "include_images",
        "image_assets",
        "image_issues",
        "image_references",
        "reference_issues",
    }
    assert ref["source_selection"] == selection
    assert "SO_{4}^{2−}" in ref["materials"]
    for block in preview["blocks"]:
        assert f"[Word区块{block['index']}]" in ref["materials"]
        assert block["text"] in ref["materials"]
    assert "第1行·第1—2列；横向合并" in ref["materials"]
    assert "原教案的最后一段不应丢失。" in ref["materials"]
    assert len(ref["image_assets"]) == 2
    assert len(ref["image_references"]) == 3
    first, second, repeated = ref["image_references"]
    assert first["asset_id"] == repeated["asset_id"] != second["asset_id"]
    assert first["block_index"] != repeated["block_index"]
    assert all(r["source_role"] == "handout" for r in ref["image_references"])
    assert all(
        "未自动区分题图与答案图" in asset["purpose"] for asset in ref["image_assets"]
    )
    assert all(
        r["source_sha256"] == selection["source_sha256"]
        for r in ref["image_references"]
    )
    assert ref["image_issues"] == []
    for content in (ref["materials"], "\n".join(ref["warnings"])):
        assert "确认导入后仅把原图保存为本地备课素材，不调用模型" in content
        assert (
            "后续是否发送图片像素，以生成时的“本地排版/视觉读取”选择及发送预览为准"
            in content
        )
        assert "不能仅凭图注补写图中条件" in content
        assert "未把原图像素发送给模型" not in content
        assert "模型仅见图注" not in content
    assert str(tmp_path) not in json.dumps(ref)
    assert "selections" not in ref and "points" not in ref
    assert not _image_root(facade).exists()
    assert facade.state_store.snapshot() == state
    assert path.read_bytes() == original
    assert provider.borrow_calls == 0


def test_confirm_revalidates_and_saves_exact_pixels_without_mutating_input(
    desktop_paths, tmp_path
):
    facade, selection, _, _, provider = _import(desktop_paths, tmp_path)
    ref = _reference(facade, selection)
    frozen = deepcopy(ref)
    existing = []
    result = facade.import_word_source_reference(ref, existing)
    assert set(result) == {"materials", "warnings", "image_assets"}
    assert result["materials"] == ref["materials"]
    assert result["warnings"] == ref["warnings"]
    assert result["image_assets"] == ref["image_assets"]
    assert existing == [] and ref == frozen
    store = PreparationImageStore(_image_root(facade))
    assert store.load(result["image_assets"][0]) == _png((0, 0, 0))
    assert len(list(store.root.glob("*.image"))) == 2
    assert not list(store.root.glob(".word-source-*"))
    assert facade.import_word_source_reference(ref, result["image_assets"]) == result
    assert provider.borrow_calls == 0


def test_selection_range_is_inclusive_and_does_not_pull_other_images(
    desktop_paths, tmp_path
):
    facade, selection, preview, _, _ = _import(desktop_paths, tmp_path)
    selection.update(block_start=2, block_end=4)
    ref = _reference(facade, selection)
    assert "SO_{4}" in ref["materials"] and "比较知识点" in ref["materials"]
    assert "第二章 化学反应" not in ref["materials"]
    assert "最后一段" not in ref["materials"]
    assert ref["image_assets"] == [] and ref["image_references"] == []
    selection.update(block_start=5, block_end=5)
    assert preview["assets"][0]["block_index"] == 5
    ref = _reference(facade, selection)
    assert len(ref["image_assets"]) == 1 and len(ref["image_references"]) == 1


@pytest.mark.parametrize(
    "change",
    [
        {"block_start": 0},
        {"block_start": True},
        {"block_end": 99999},
        {"block_start": 6, "block_end": 3},
        {"block_end": 3.5},
        {"source_sha256": "a" * 64},
        {"revision": "b" * 64},
        {"source_id": "does-not-exist"},
        {"batch_id": "missing"},
    ],
)
def test_invalid_or_stale_selection_fails_without_writes(
    desktop_paths, tmp_path, change
):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path)
    selection.update(change)
    with pytest.raises(WordSourceReferenceError):
        _reference(facade, selection)
    assert not _image_root(facade).exists()


@pytest.mark.parametrize(
    "change",
    [
        {"materials": "篡改文字"},
        {"warnings": []},
        {"image_assets": []},
        {"image_references": []},
        {"image_issues": ["篡改"]},
        {"include_images": False},
        {"unexpected": True},
        {"reference_issues": ["篡改"]},
    ],
)
def test_confirm_rejects_any_preview_tampering(desktop_paths, tmp_path, change):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path)
    ref = _reference(facade, selection)
    ref.update(change)
    with pytest.raises(WordSourceReferenceError, match="变化"):
        facade.import_word_source_reference(ref, [])
    assert not _image_root(facade).exists()


def test_confirmation_reads_only_selected_source_and_rejects_changed_archive(
    desktop_paths, tmp_path, monkeypatch
):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path)
    ref = _reference(facade, selection)
    calls = []
    restore = facade._restore_visual_import_sources

    def counted(descriptor, *, source_ids=None):
        calls.append(source_ids)
        return restore(descriptor, source_ids=source_ids)

    monkeypatch.setattr(facade, "_restore_visual_import_sources", counted)
    archive = (
        facade._visual_import_root / "sources" / f"{selection['source_sha256']}.docx"
    )
    archive.write_bytes(b"changed archived test fixture")
    with pytest.raises(WordSourceReferenceError, match="变化"):
        facade.import_word_source_reference(ref, [])
    assert calls == [{selection["source_id"]}]
    assert not _image_root(facade).exists()


def test_unsupported_image_is_explicit_and_text_only_is_opt_in(
    desktop_paths, tmp_path, monkeypatch
):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path)
    original = PreparationSourcesService.word_asset_bytes

    def unsupported(reader, data, asset_id, **kwargs):
        if asset_id == "word-b6-image1":
            raise PreparationSourceError("WMF旧字体尚未可靠验证")
        return original(reader, data, asset_id, **kwargs)

    monkeypatch.setattr(PreparationSourcesService, "word_asset_bytes", unsupported)
    ref = _reference(facade, selection)
    assert len(ref["image_assets"]) == 1
    assert len(ref["image_references"]) == 3
    assert "WMF旧字体" in ref["image_issues"][0]
    with pytest.raises(WordSourceReferenceError, match="未能全部带入"):
        facade.import_word_source_reference(ref, [])
    assert not _image_root(facade).exists()
    monkeypatch.setattr(
        PreparationSourcesService,
        "word_asset_bytes",
        lambda *a, **kw: pytest.fail("text-only must not decode pictures"),
    )
    text = _reference(facade, selection, include_images=False)
    assert text["image_assets"] == [] and text["image_issues"] == []
    assert all(
        r["status"] == "not_included_text_only" for r in text["image_references"]
    )
    result = facade.import_word_source_reference(text, [])
    assert result["image_assets"] == []
    assert "明确仅使用文字" in text["materials"]
    for content in (text["materials"], "\n".join(text["warnings"])):
        assert "本次已明确选择仅文字：未带入任何原图，也未把原图发送给模型" in content
        assert "须由教师补充图中必要条件" in content
        assert "保存为本地备课素材" not in content
        assert "模型仅见图注" not in content
    assert not _image_root(facade).exists()


def test_missing_object_preview_prevents_silent_partial_import(desktop_paths, tmp_path):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path, missing=True)
    ref = _reference(facade, selection)
    assert any("未找到可独立读取" in issue for issue in ref["image_issues"])
    assert "嵌入对象或旧公式" in ref["materials"]
    with pytest.raises(WordSourceReferenceError, match="未能全部带入"):
        facade.import_word_source_reference(ref, [])
    assert not _image_root(facade).exists()


def test_answer_source_role_is_explicit_not_inferred_from_filename(
    desktop_paths, tmp_path
):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path, answer=True)
    ref = _reference(facade, selection)
    assert all(row["source_role"] == "answer" for row in ref["image_references"])
    assert all("只用于教师讲评" in row["purpose"] for row in ref["image_assets"])
    assert "答案来源，仅用于教师讲评" in ref["materials"]


def test_derived_emf_keeps_original_identity_and_renderer_evidence(
    desktop_paths, tmp_path, monkeypatch
):
    facade, selection, preview, _, _ = _import(desktop_paths, tmp_path, images=1)
    image_sha = preview["assets"][0]["sha256"]
    derived = _png("orange")
    digest = hashlib.sha256(derived).hexdigest()
    monkeypatch.setattr(
        PreparationSourcesService,
        "word_asset_bytes",
        lambda *args, **kwargs: {
            "bytes": derived,
            "mime_type": "image/png",
            "original_mime_type": "image/x-emf",
            "original_sha256": image_sha,
            "preview_sha256": digest,
            "derived_preview": True,
            "renderer_revision": "fixture-emf-renderer-v1",
        },
    )
    ref = _reference(facade, selection)
    assert len(ref["image_assets"]) == 1
    for item in ref["image_references"]:
        assert item["original_sha256"] == image_sha
        assert item["preview_sha256"] == digest
        assert item["derived_preview"] is True
        assert item["renderer_revision"] == "fixture-emf-renderer-v1"
    assert image_sha in ref["materials"] and digest in ref["materials"]
    assert "原件未修改" in ref["materials"]
    result = facade.import_word_source_reference(ref, [])
    assert (
        PreparationImageStore(_image_root(facade)).load(result["image_assets"][0])
        == derived
    )


def test_derived_payload_sha_mismatch_is_an_image_issue(
    desktop_paths, tmp_path, monkeypatch
):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path)
    monkeypatch.setattr(
        PreparationSourcesService,
        "word_asset_bytes",
        lambda *args, **kwargs: {
            "bytes": _png("orange"),
            "derived_preview": True,
            "original_sha256": "a" * 64,
            "preview_sha256": "b" * 64,
        },
    )
    ref = _reference(facade, selection)
    assert ref["image_assets"] == []
    assert len(ref["image_issues"]) == 3
    with pytest.raises(WordSourceReferenceError):
        facade.import_word_source_reference(ref, [])


def test_picture_capacity_is_checked_before_any_save(desktop_paths, tmp_path):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path, images=MAX_IMAGES + 1)
    ref = _reference(facade, selection)
    assert len(ref["image_assets"]) == MAX_IMAGES + 1
    assert any("未删减图片" in note for note in ref["warnings"])
    assert ref["reference_issues"]
    with pytest.raises(WordSourceReferenceError, match=f"最多选择{MAX_IMAGES}张"):
        facade.import_word_source_reference(ref, [])
    assert not _image_root(facade).exists()


def test_merged_capacity_keeps_existing_assets_unchanged(desktop_paths, tmp_path):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path)
    ref = _reference(facade, selection)
    store = PreparationImageStore(_image_root(facade))
    existing = [
        store.import_bytes(_png((200, n, 120)), "原有图", "教师", "原有用途")
        for n in range(MAX_IMAGES - 1)
    ]
    before = {p.name: p.read_bytes() for p in store.root.iterdir()}
    with pytest.raises(WordSourceReferenceError, match=f"最多选择{MAX_IMAGES}张"):
        facade.import_word_source_reference(ref, existing)
    assert {p.name: p.read_bytes() for p in store.root.iterdir()} == before
    assert len(existing) == MAX_IMAGES - 1


def test_bad_existing_asset_is_rejected_before_new_files(desktop_paths, tmp_path):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path)
    ref = _reference(facade, selection)
    store = PreparationImageStore(_image_root(facade))
    existing = [store.import_bytes(_png("purple"), "原有图", "教师", "原有用途")]
    (store.root / (existing[0]["sha256"] + ".image")).write_bytes(b"broken")
    before = {p.name: p.read_bytes() for p in store.root.iterdir()}
    with pytest.raises(WordSourceReferenceError, match="变化"):
        facade.import_word_source_reference(ref, existing)
    assert {p.name: p.read_bytes() for p in store.root.iterdir()} == before


def test_write_failure_rolls_back_only_this_batch(desktop_paths, tmp_path, monkeypatch):
    from integrations.deeptutor_shchem_v1 import desktop_word_source_reference as module

    facade, selection, _, _, _ = _import(desktop_paths, tmp_path)
    ref = _reference(facade, selection)
    store = PreparationImageStore(_image_root(facade))
    existing = [store.import_bytes(_png("purple"), "原有图", "教师", "原有用途")]
    before = {p.name: p.read_bytes() for p in store.root.iterdir()}
    link = module.os.link
    calls = []

    def fail_second(src, dst):
        calls.append(dst)
        if len(calls) == 2:
            raise OSError("injected local write failure")
        return link(src, dst)

    monkeypatch.setattr(module.os, "link", fail_second)
    with pytest.raises(WordSourceReferenceError, match="无法完整保存"):
        facade.import_word_source_reference(ref, existing)
    assert len(calls) == 2
    assert {p.name: p.read_bytes() for p in store.root.iterdir()} == before
    assert len(existing) == 1


def test_existing_identical_picture_dedup_keeps_teacher_metadata(
    desktop_paths, tmp_path
):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path)
    ref = _reference(facade, selection)
    store = PreparationImageStore(_image_root(facade))
    existing = [
        store.import_bytes(_png((0, 0, 0)), "教师已修改图题", "已有来源", "已有用途")
    ]
    result = facade.import_word_source_reference(ref, existing)
    assert result["image_assets"][0] == existing[0]
    assert len(result["image_assets"]) == 2
    assert len(ref["image_references"]) == 3


def _quality_note(facade, preview, *, block_index, revision=None):
    path = facade.paths.workspace_root / QUALITY_PATH
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "sources": [
                    {
                        "source_sha256": preview["source_sha256"],
                        "source_revision": revision or preview["revision"],
                        "issues": [
                            {
                                "id": "fixture-erratum",
                                "block_indices": [block_index],
                                "summary": "已知原文知识错误",
                                "suggested_correction": "AI修订建议须教师复核",
                            }
                        ],
                    }
                ],
            }
        ),
        encoding="utf-8",
    )


@pytest.mark.parametrize("text_only", [False, True])
def test_known_source_errata_cannot_bypass_via_full_word_channel(
    desktop_paths, tmp_path, text_only
):
    facade, selection, preview, _, _ = _import(desktop_paths, tmp_path)
    _quality_note(facade, preview, block_index=2)
    ref = _reference(facade, selection, include_images=not text_only)
    assert "已知原文知识错误" in ref["materials"]
    assert any("并非教师确认" in note for note in ref["warnings"])
    assert ref["reference_issues"] and "内容问题" in ref["reference_issues"][0]
    with pytest.raises(WordSourceReferenceError, match="已记录的内容问题"):
        facade.import_word_source_reference(ref, [])
    assert not _image_root(facade).exists()


def test_quality_notes_match_selected_blocks_and_stale_revision_is_not_silent(
    desktop_paths, tmp_path
):
    facade, selection, preview, _, _ = _import(desktop_paths, tmp_path)
    _quality_note(facade, preview, block_index=2)
    selection.update(block_start=3, block_end=4)
    ref = _reference(facade, selection)
    assert "已知原文知识错误" not in ref["materials"]
    facade.import_word_source_reference(ref, [])
    _quality_note(facade, preview, block_index=2, revision="outdated-extractor")
    ref = _reference(facade, selection)
    assert "需重新定位" in ref["materials"]
    with pytest.raises(WordSourceReferenceError, match="内容问题"):
        facade.import_word_source_reference(ref, [])


def test_added_quality_note_after_preview_invalidates_confirmation(
    desktop_paths, tmp_path
):
    facade, selection, preview, _, _ = _import(desktop_paths, tmp_path)
    ref = _reference(facade, selection)
    _quality_note(facade, preview, block_index=2)
    with pytest.raises(WordSourceReferenceError, match="变化"):
        facade.import_word_source_reference(ref, [])
    assert not _image_root(facade).exists()


def test_long_original_is_not_silently_truncated(desktop_paths, tmp_path):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path, images=0, long=True)
    with pytest.raises(WordSourceReferenceError, match="未截断"):
        _reference(facade, selection)
    assert not _image_root(facade).exists()


def test_original_text_only_contract_remains_compatible(desktop_paths, tmp_path):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path)
    old = facade.imported_word_reference(
        selection["batch_id"],
        selection["source_id"],
        selection["source_sha256"],
        selection["block_start"],
        selection["block_end"],
        selection["revision"],
    )
    assert "materials" in old and "warnings" in old
    assert "image_assets" not in old and "source_selection" not in old


@pytest.mark.parametrize(
    "value", [None, [], {"source_id": "missing"}, {"include_images": "yes"}]
)
def test_invalid_reference_schema_is_rejected(desktop_paths, tmp_path, value):
    facade, _, _, _, _ = _import(desktop_paths, tmp_path)
    with pytest.raises(WordSourceReferenceError):
        WordSourceReferenceService(facade).prepare_reference(value, [])


def test_explicit_image_choice_requires_boolean(desktop_paths, tmp_path):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path)
    with pytest.raises(WordSourceReferenceError, match="请选择带入原图"):
        _reference(facade, selection, include_images=1)
    assert not _image_root(facade).exists()


def test_extractor_change_after_preview_requires_new_confirmation(
    desktop_paths, tmp_path, monkeypatch
):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path)
    ref = _reference(facade, selection)
    extract = PreparationSourcesService.word_preview_bytes

    def changed(reader, data, name):
        result = extract(reader, data, name)
        result["revision"] = "f" * 64
        return result

    monkeypatch.setattr(PreparationSourcesService, "word_preview_bytes", changed)
    with pytest.raises(WordSourceReferenceError, match="提取结果已变化"):
        facade.import_word_source_reference(ref, [])
    assert not _image_root(facade).exists()


def test_derivative_pixels_changed_after_preview_require_reconfirmation(
    desktop_paths, tmp_path, monkeypatch
):
    facade, selection, preview, _, _ = _import(desktop_paths, tmp_path, images=1)
    raw = [_png("orange")]

    def convert(*args, **kwargs):
        return {
            "bytes": raw[0],
            "derived_preview": True,
            "original_sha256": preview["assets"][0]["sha256"],
            "preview_sha256": hashlib.sha256(raw[0]).hexdigest(),
            "renderer_revision": "fixture-renderer",
        }

    monkeypatch.setattr(PreparationSourcesService, "word_asset_bytes", convert)
    ref = _reference(facade, selection)
    raw[0] = _png("yellow")
    with pytest.raises(WordSourceReferenceError, match="原图已经变化"):
        facade.import_word_source_reference(ref, [])
    assert not _image_root(facade).exists()


def test_one_successful_picture_does_not_hide_a_separate_missing_object(
    desktop_paths, tmp_path
):
    document = Document()
    paragraph = document.add_paragraph("可读图片：")
    paragraph.add_run().add_picture(io.BytesIO(_png("blue")))
    paragraph.add_run("另有对象：")._r.append(OxmlElement("w:object"))
    facade, selection, _ = _from_document(desktop_paths, tmp_path, document)
    ref = _reference(facade, selection)
    assert len(ref["image_assets"]) == 1
    assert any("图形或对象2：未找到" in issue for issue in ref["image_issues"])
    with pytest.raises(WordSourceReferenceError, match="未能全部带入"):
        facade.import_word_source_reference(ref, [])


def test_readonly_quality_manifest_failure_blocks_reference(desktop_paths, tmp_path):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path)
    path = facade.paths.workspace_root / QUALITY_PATH
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("not-valid-json", encoding="utf-8")
    with pytest.raises(WordSourceReferenceError, match="修订记录无法读取"):
        _reference(facade, selection)
    assert not _image_root(facade).exists()


def test_staging_failure_cannot_publish_the_first_valid_picture(
    desktop_paths, tmp_path, monkeypatch
):
    facade, selection, _, _, _ = _import(desktop_paths, tmp_path)
    ref = _reference(facade, selection)
    save = PreparationImageStore.import_bytes
    calls = []

    def fail_second(store, *args, **kwargs):
        calls.append(store.root)
        if len(calls) == 2:
            raise OSError("injected staging failure")
        return save(store, *args, **kwargs)

    monkeypatch.setattr(PreparationImageStore, "import_bytes", fail_second)
    with pytest.raises(WordSourceReferenceError, match="无法完整保存"):
        facade.import_word_source_reference(ref, [])
    assert len(calls) == 2
    assert list(_image_root(facade).iterdir()) == []


def test_concurrent_valid_destination_is_not_removed_by_failed_batch(
    desktop_paths, tmp_path, monkeypatch
):
    from integrations.deeptutor_shchem_v1 import desktop_word_source_reference as module

    facade, selection, _, _, _ = _import(desktop_paths, tmp_path)
    ref = _reference(facade, selection)
    link = module.os.link
    external = []

    def concurrent_first_then_fail(src, dst):
        if not external:
            # A separate importer wins this destination just before our link.
            link(src, dst)
            external.append(dst)
            raise FileExistsError("concurrent import")
        raise OSError("second destination failed")

    monkeypatch.setattr(module.os, "link", concurrent_first_then_fail)
    with pytest.raises(WordSourceReferenceError, match="无法完整保存"):
        facade.import_word_source_reference(ref, [])
    assert list(_image_root(facade).iterdir()) == external
    assert external[0].read_bytes() == _png((0, 0, 0))


@pytest.mark.parametrize("in_table", [False, True])
def test_repeated_relationship_occurrences_survive_same_block_and_table(
    desktop_paths, tmp_path, in_table
):
    document = Document()
    if in_table:
        table = document.add_table(rows=1, cols=2)
        paragraphs = [table.cell(0, column).paragraphs[0] for column in range(2)]
    else:
        paragraphs = [document.add_paragraph("比较")] * 2
    for paragraph in paragraphs:
        paragraph.add_run().add_picture(io.BytesIO(_png("blue")))
        paragraph.add_run("与")
    facade, selection, preview = _from_document(desktop_paths, tmp_path, document)
    assert len(preview["assets"]) == 1
    frozen = deepcopy(preview)
    ref = _reference(facade, selection)
    assert len(ref["image_assets"]) == 1 and len(ref["image_references"]) == 2
    first, second = ref["image_references"]
    assert first["source_asset_id"] == second["source_asset_id"]
    assert first["block_index"] == second["block_index"]
    assert (first["position"], second["position"]) == (1, 2)
    assert first["asset_id"] == second["asset_id"]
    assert ref["image_issues"] == [] and ref["reference_issues"] == []
    assert (
        facade.imported_word_preview(selection["batch_id"], selection["source_id"])
        == frozen
    )
    facade.import_word_source_reference(ref, [])


@pytest.mark.parametrize("vector_choice", [False, True])
def test_compatibility_choice_and_fallback_keep_explicit_branch_binding(
    desktop_paths, tmp_path, vector_choice
):
    from lxml import etree

    document = Document()
    paragraph = document.add_paragraph("兼容图形")
    choice_run = paragraph.add_run()
    if vector_choice:
        choice_run._r.append(OxmlElement("w:drawing"))
    else:
        choice_run.add_picture(io.BytesIO(_png("blue")))
    fallback_run = paragraph.add_run()
    fallback_run.add_picture(io.BytesIO(_png("blue")))
    namespace = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    alternate = etree.Element("{" + namespace + "}AlternateContent")
    for name, run in (("Choice", choice_run), ("Fallback", fallback_run)):
        branch = etree.SubElement(alternate, "{" + namespace + "}" + name)
        branch.append(run._r)
    paragraph._p.append(alternate)
    facade, selection, preview = _from_document(desktop_paths, tmp_path, document)
    ref = _reference(facade, selection)
    assert len(preview["assets"]) == 1
    assert ref["image_issues"] == []
    assert [row["source_branch"] for row in ref["image_references"]] == (
        ["fallback"] if vector_choice else ["choice", "fallback"]
    )
    assert any("未推断实际排版" in note for note in ref["warnings"])
    assert "兼容回退分支" in ref["materials"]
    facade.import_word_source_reference(ref, [])


def test_literal_image_marker_text_does_not_imply_missing_source_object(
    desktop_paths, tmp_path
):
    document = Document()
    document.add_paragraph("这句说明引用一个文字标签：【待查看原文：图片或图形】")
    facade, selection, _ = _from_document(desktop_paths, tmp_path, document)
    ref = _reference(facade, selection)
    assert ref["image_issues"] == [] and ref["image_assets"] == []
