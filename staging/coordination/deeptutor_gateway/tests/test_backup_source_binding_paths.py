"""Restored provenance paths dedupe only for the same validated source version."""
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
import json
import os
import subprocess

from docx import Document
import pytest

from test_desktop_visual_import_facade import FakeProviderStore, _facade, _png, desktop_paths
from test_desktop_preparation_drafts import _record
from test_lesson_design import sample
from integrations.deeptutor_shchem_v1 import desktop_backup as backup
from integrations.deeptutor_shchem_v1.desktop_facade import DesktopFacadeError, _canonical_digest
from integrations.deeptutor_shchem_v1.desktop_handout_knowledge import (
    handout_knowledge_sources, handout_knowledge_outline, import_handout_knowledge,
)
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    add_resource, digest, enable_composition, new_composition, validate_resource,
)
from integrations.deeptutor_shchem_v1.desktop_state import DesktopStateStore


def _binding(item):
    return item['original']['evidence']['handout_source']


def _with_binding(item, **changes):
    item = deepcopy(item)
    evidence = item['original']['evidence']
    evidence['handout_source'].update(changes)
    source = _binding(item)
    for name in ('source_path', 'source_sha256', 'source_revision'):
        evidence[name] = source[name]
    evidence['source_roles']['document_role'] = source['source_role']
    item['original']['title'] = item['title'] = (
        source['source_name'] + ' · ' + evidence['section_snapshot']['section']['title'])
    role = 'teacher_note' if source['source_role'] == 'answer' else 'knowledge'
    for rows in (item['original']['blocks'], item['blocks']):
        for row in rows:
            row['role'] = role
    item['original_sha256'] = digest(item['original'])
    return item


def _save(fixture, items):
    for position, item in enumerate(items):
        payload = sample()
        payload['lesson_design'] = enable_composition(payload['lesson_design'])
        composition = new_composition()
        composition = add_resource(composition, item, {
            'page_id': composition['pages'][0], 'x': 14, 'y': 14,
            'width': 180, 'height': 120,
        })
        payload['lesson_design']['composition'] = composition
        record = _record(payload)
        record['lesson_design'] = payload['lesson_design']
        fixture.facade._state.save_draft('prep-path-' + str(position), record)


@pytest.fixture
def archive(desktop_paths, tmp_path):
    """Small native full batch, including unselected files and all three roles."""
    provider = FakeProviderStore(configured=False)
    facade = _facade(desktop_paths, provider)
    document = Document()
    document.add_heading('知识点1 合成材料', 3)
    document.add_paragraph('此处是来源路径与备份协议的合成测试材料。')
    chosen = tmp_path / '本课讲义.docx'
    document.save(chosen)
    paths = {}
    for name, color in (('同批资料.png', 'green'), ('题面.png', 'white'), ('答案.png', 'blue')):
        paths[name] = tmp_path / name
        paths[name].write_bytes(_png(color))
    receipt = facade.save_visual_import_batch(
        handout_files=(chosen, paths['同批资料.png']), question_files=(paths['题面.png'],),
        answer_files=(paths['答案.png'],), source_type='合成备份测试资料',
    )
    metadata = next(row for row in handout_knowledge_sources(facade)['sources']
                    if row['source_name'] == chosen.name)
    outline = handout_knowledge_outline(facade, metadata)
    imported = import_handout_knowledge(facade, [{
        'source': outline['source'], 'outline_revision': outline['outline_revision'],
        'section': outline['sections'][0], 'title': '合成来源测试', 'teacher_only': False,
    }], [], [])
    item = imported['resources'][0]
    descriptor = facade._saved_visual_import_batch(receipt.batch_id)
    row = next(row for row in descriptor['sources']
               if row['source_file_id'] == _binding(item)['source_id'])
    return SimpleNamespace(facade=facade, provider=provider, item=item, descriptor=descriptor,
        root=desktop_paths.state_root, chosen=chosen, batch=receipt.batch_id, row=row,
        manifest=desktop_paths.state_root / 'visual-import-v2/batches' / (receipt.batch_id + '.json'),
        archive=desktop_paths.state_root / 'visual-import-v2' / row['archive_relative_path'])


def _plan(fixture):
    return backup.plan_backup(fixture.root, include_images=True, include_tasks=True,
                              include_source_dependencies=True)


def test_two_paths_same_version_roundtrip_preserves_every_original_and_exact_whole_batch(archive, tmp_path, monkeypatch):
    old = _with_binding(archive.item, source_path='X:/unavailable-old-profile/original.docx')
    fresh = deepcopy(archive.item)
    assert _binding(old) != _binding(fresh)
    assert {k: v for k, v in _binding(old).items() if k != 'source_path'} == {
        k: v for k, v in _binding(fresh).items() if k != 'source_path'}
    _save(archive, [old, fresh])
    before = (archive.root / 'desktop-state.v1.json').read_bytes()
    opened = []
    original_open = Path.open
    def tracked(path, *args, **kwargs):
        opened.append(str(path))
        return original_open(path, *args, **kwargs)
    monkeypatch.setattr(Path, 'open', tracked)
    inputs = [old, fresh]
    frozen_inputs = deepcopy(inputs)
    bindings = backup._course_source_bindings(inputs)
    assert len(bindings) == 1 and next(iter(bindings.values())) == _binding(old)
    assert inputs == frozen_inputs
    plan = _plan(archive)
    expected = {'visual-import-v2/batches/' + archive.batch + '.json'} | {
        'visual-import-v2/' + row['archive_relative_path'] for row in archive.descriptor['sources']}
    assert {row.name for row in plan.files if row.name.startswith('visual-import-v2/')} == expected
    assert plan.summary['source_batches'] == 1
    assert plan.summary['source_documents'] == plan.summary['source_archive_files'] == 4
    destination = tmp_path / 'backup.zip'
    backup.create_backup(plan, destination)
    inspected = backup.inspect_backup(destination)
    restored = tmp_path / 'restored'
    backup.restore_backup(destination, restored, expected_manifest=backup.manifest_revision(inspected))
    actual = DesktopStateStore(restored).snapshot()
    for index, item in enumerate((old, fresh)):
        record = actual['drafts']['prep-path-' + str(index)]
        assert record['lesson_design']['composition']['resources'][0] == item
    for relative in expected:
        assert (restored / relative).read_bytes() == (archive.root / relative).read_bytes()
    assert (archive.root / 'desktop-state.v1.json').read_bytes() == before
    assert not any(path.replace('\\', '/').startswith('X:/') for path in opened)
    assert archive.provider.borrow_calls == 0


@pytest.mark.parametrize('field,value', [
    ('source_sha256', 'f' * 64), ('source_revision', 'f' * 64),
    ('source_role', 'answer'), ('source_name', '其他讲义.docx'),
])
def test_same_batch_and_source_id_conflicting_semantics_still_rejected(archive, field, value):
    old = _with_binding(archive.item, source_path='X:/old/original.docx')
    changed = _with_binding(archive.item, **{field: value})
    validate_resource(old)
    validate_resource(changed)
    with pytest.raises(backup.BackupError, match='来源版本不一致'):
        backup._course_source_bindings([old, changed])


@pytest.mark.parametrize('field,value', [
    ('source_id', 'SRC-FOREIGN'), ('batch_id', 'DESKTOPBATCH-' + 'f' * 32),
])
def test_different_native_id_or_batch_cannot_be_substituted_by_matching_path(archive, field, value):
    changed = _with_binding(archive.item, **{field: value})
    _save(archive, [archive.item, changed])
    with pytest.raises((backup.BackupError, DesktopFacadeError)):
        _plan(archive)


@pytest.mark.parametrize('corruption', ['future_schema', 'extra_field', 'missing_field'])
def test_binding_schema_remains_exact(archive, corruption):
    item = deepcopy(archive.item)
    source = _binding(item)
    if corruption == 'future_schema':
        source['schema'] = 'shchem.handout-knowledge-source.future'
    elif corruption == 'extra_field':
        source['arbitrary'] = 'unsupported'
    else:
        del source['source_revision']
    item['original_sha256'] = digest(item['original'])
    with pytest.raises(ValueError):
        backup._course_source_bindings([archive.item, item])


def test_provenance_path_cannot_rescue_changed_current_archive(archive):
    # This real original file still has valid bytes; only the registered archive drifts.
    item = _with_binding(archive.item, source_path=str(archive.chosen))
    _save(archive, [item])
    archive.archive.write_bytes(b'changed current archive')
    with pytest.raises(DesktopFacadeError, match='归档不完整或已变化'):
        _plan(archive)


def test_provenance_path_cannot_rescue_missing_current_archive(archive):
    _save(archive, [_with_binding(archive.item, source_path=str(archive.chosen))])
    archive.archive.unlink()
    with pytest.raises(DesktopFacadeError, match='归档不完整或已变化'):
        _plan(archive)


def test_unselected_same_batch_source_is_still_required(archive):
    _save(archive, [archive.item])
    neighbor = next(row for row in archive.descriptor['sources'] if row['source_file_id'] != archive.row['source_file_id'])
    (archive.root / 'visual-import-v2' / neighbor['archive_relative_path']).write_bytes(b'changed neighbor')
    with pytest.raises(DesktopFacadeError, match='归档不完整或已变化'):
        _plan(archive)


@pytest.mark.parametrize('field,value', [('role', 'answer'), ('filename', '替换讲义.docx'), ('source_file_id', 'SRC-FOREIGN')])
def test_resigned_current_manifest_cannot_override_frozen_native_identity(archive, field, value):
    _save(archive, [archive.item])
    descriptor = deepcopy(archive.descriptor)
    row = next(row for row in descriptor['sources'] if row['source_file_id'] == archive.row['source_file_id'])
    row[field] = value
    descriptor['source_closure_sha256'] = _canonical_digest(descriptor['sources'])
    if field in {'role', 'filename'}:
        receipt_row = next(row for row in descriptor['receipt']['sources']
                           if row['filename'] == archive.row['filename'] and row['role'] == archive.row['role'])
        receipt_row[field] = value
    archive.facade._state.save_draft('visual-import-v2:' + archive.batch, descriptor)
    manifest = json.loads(archive.manifest.read_text(encoding='utf-8'))
    manifest['plan']['sources'] = deepcopy(descriptor['sources'])
    archive.manifest.write_bytes(backup._bytes(manifest))
    with pytest.raises(ValueError):
        _plan(archive)


def test_bytes_changed_after_reviewed_plan_still_rejected(archive, tmp_path):
    _save(archive, [_with_binding(archive.item, source_path='X:/old/original.docx'), archive.item])
    plan = _plan(archive)
    archive.archive.write_bytes(b'changed after review')
    destination = tmp_path / 'must-not-publish.zip'
    with pytest.raises(backup.BackupError, match='变化'):
        backup.create_backup(plan, destination)
    assert not destination.exists()


def test_current_source_directory_link_still_rejected(archive, tmp_path):
    _save(archive, [_with_binding(archive.item, source_path='X:/old/original.docx'), archive.item])
    sources = archive.root / 'visual-import-v2/sources'
    target = tmp_path / 'moved-sources'
    sources.rename(target)
    if os.name == 'nt':
        linked = subprocess.run(['cmd', '/c', 'mklink', '/J', str(sources), str(target)],
                                capture_output=True, check=False)
        assert linked.returncode == 0, linked.stderr.decode(errors='replace')
    else:
        sources.symlink_to(target, target_is_directory=True)
    with pytest.raises(backup.BackupError, match='链接'):
        _plan(archive)


def test_manifest_directory_mismatch_still_rejected(archive):
    _save(archive, [archive.item])
    manifest = json.loads(archive.manifest.read_text(encoding='utf-8'))
    manifest['plan']['sources'][0]['role'] = 'answer'
    archive.manifest.write_bytes(backup._bytes(manifest))
    with pytest.raises(DesktopFacadeError, match='处理结果不完整或已变化'):
        _plan(archive)
