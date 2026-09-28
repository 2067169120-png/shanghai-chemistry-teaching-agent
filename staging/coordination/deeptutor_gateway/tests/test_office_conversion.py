"""Protocol and backend failure tests; Office itself is never activated here."""
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace
from threading import Event
import os

import pytest
from docx import Document

from integrations.deeptutor_shchem_v1 import office_conversion as office
from integrations.deeptutor_shchem_v1.reader_cancellation import ReadCancelled, read_cancel_scope


def document(tmp_path):
    source=tmp_path/"synthetic.docx"
    doc=Document();doc.add_paragraph("Synthetic conversion input");doc.save(source)
    return source


def protocol(env, **overrides):
    return {"session_id":env["SHCHEM_OFFICE_ID"], "input_path":env["SHCHEM_WORD_INPUT"],
            "output_path":env["SHCHEM_WORD_OUTPUT"], "input_sha256":env["SHCHEM_WORD_SHA"],
            "source_path":env["SHCHEM_WORD_ORIGINAL"], "target_path":env["SHCHEM_WORD_TARGET"], **overrides}


def write_record(env,name,**changes):
    Path(env["SHCHEM_OFFICE_SESSION"],name+".json").write_text(json.dumps(protocol(env,**changes)),encoding="utf-8")


@pytest.mark.parametrize("bad_field", [None,"session_id","input_path","output_path","input_sha256",
    "held_handle","cleanup","pid_missing","pid_zero","pid_bool","start_missing","start_zero",
    "private_missing","private_number","cleanup_pid","automation_missing"])
def test_word_accepts_only_bound_new_output_and_verified_process_cleanup(tmp_path,monkeypatch,bad_field):
    source=document(tmp_path);target=tmp_path/"new.pdf"
    def run(command,*,env,**kwargs):
        assert kwargs["contain"] is False and callable(kwargs["cooperative_stop"])
        assert env["SHCHEM_WORD_SHA"]==hashlib.sha256(source.read_bytes()).hexdigest()
        Path(env["SHCHEM_WORD_OUTPUT"]).write_bytes(b"%PDF-synthetic protocol only")
        changes={bad_field:"wrong-task"} if bad_field in {"session_id","input_path","output_path","input_sha256"} else {}
        write_record(env,"result",status="converted",**changes)
        owner={"held_handle":bad_field!="held_handle","pid":2345,"start_filetime":133800000000000000,"private":True}
        if bad_field=="pid_missing":owner.pop("pid")
        if bad_field=="pid_zero":owner["pid"]=0
        if bad_field=="pid_bool":owner["pid"]=True
        if bad_field=="start_missing":owner.pop("start_filetime")
        if bad_field=="start_zero":owner["start_filetime"]=0
        if bad_field=="private_missing":owner.pop("private")
        if bad_field=="private_number":owner["private"]=1
        write_record(env,"owner",**owner)
        if bad_field!="automation_missing":write_record(env,"automation",identity="singleuse_empty_com",pid_verified_before_open=False)
        write_record(env,"cleanup",pid=6789 if bad_field=="cleanup_pid" else 2345,
                     disposition="pending_unverified_cleanup" if bad_field=="cleanup" else "owned_instance_exited")
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(office,"run_owned",run)
    if bad_field:
        with pytest.raises(office.OfficeConversionError) as caught:office.word_pdf(source,target,session_root=tmp_path/"sessions")
        session=Path(caught.value.office_session)
        assert (session/"request.json").is_file() and (session/"caller.json").is_file()
        assert json.loads((session/"caller.json").read_text(encoding="utf-8"))["result_discarded"] is True
        assert not target.exists()
    else:
        assert "已核对该实例退出" in office.word_pdf(source,target,session_root=tmp_path/"sessions")


@pytest.mark.parametrize("backend",["word","libreoffice"])
def test_existing_pdf_is_never_evidence_of_this_conversion(tmp_path,monkeypatch,backend):
    source=document(tmp_path);target=tmp_path/"synthetic.pdf";target.write_bytes(b"%PDF-old frozen output")
    monkeypatch.setattr(office,"run_owned",lambda *a,**k:pytest.fail("must not launch"))
    with pytest.raises(office.OfficeConversionError,match="目标已存在"):
        if backend=="word":office.word_pdf(source,target)
        else:office.libreoffice_pdf(source,tmp_path,"synthetic-soffice")
    assert target.read_bytes()==b"%PDF-old frozen output"


def test_external_relationship_never_changes_word_global_options(tmp_path,monkeypatch):
    from docx.opc.constants import RELATIONSHIP_TYPE
    source=tmp_path/"link.docx";doc=Document();doc.add_paragraph("Synthetic hyperlink")
    doc.part.relate_to("https://example.invalid/synthetic",RELATIONSHIP_TYPE.HYPERLINK,is_external=True);doc.save(source)
    monkeypatch.setattr(office,"run_owned",lambda *a,**k:pytest.fail("external link must be rejected before Office"))
    with pytest.raises(office.OfficeConversionError,match="外部链接"):office.word_pdf(source,tmp_path/"link.pdf")
    assert "Options.UpdateLinksAtOpen" not in office._WORD_SCRIPT


@pytest.mark.parametrize("exited,disposition,complete",[(False,None,False),(True,"taken_over_or_unverified_preserved",False),(True,"owned_instance_exited",True)])
def test_word_cancel_preserves_unproven_cleanup_and_binds_diagnostic(tmp_path,monkeypatch,exited,disposition,complete):
    source=document(tmp_path);target=tmp_path/"output.pdf";observed={}
    def run(command,*,env,cooperative_stop,**kwargs):
        write_record(env,"owner",held_handle=True,private=True,pid=2345,start_filetime=133800000000000000)
        write_record(env,"automation",identity="singleuse_empty_com",pid_verified_before_open=False)
        if disposition:write_record(env,"cleanup",pid=2345,disposition=disposition)
        error=ReadCancelled()
        observed["complete"]=cooperative_stop(SimpleNamespace(wait=lambda timeout:exited),error)
        observed["caller"]=json.loads(Path(env["SHCHEM_OFFICE_SESSION"],"caller.json").read_text(encoding="utf-8"))
        assert Path(env["SHCHEM_OFFICE_SESSION"],"cancel").is_file()
        assert all(observed["caller"][k]==v for k,v in protocol(env).items())
        raise error
    monkeypatch.setattr(office,"run_owned",run)
    with pytest.raises(ReadCancelled):office.word_pdf(source,target,session_root=tmp_path/"sessions")
    assert observed["complete"] is complete
    assert observed["caller"]["result_discarded"] is True
    assert observed["caller"]["status"]=="cancelled"
    assert not target.exists()


def test_libreoffice_uses_private_profile_and_rejects_late_changed_input(tmp_path,monkeypatch):
    source=document(tmp_path);output=tmp_path/"output";output.mkdir()
    def run(command,**kwargs):
        profile=next(arg for arg in command if arg.startswith("-env:UserInstallation="))
        assert profile.startswith("-env:UserInstallation=file://")
        assert "--headless" in command and "--norestore" in command
        (output/"synthetic.pdf").write_bytes(b"%PDF-synthetic")
        source.write_bytes(b"source changed during conversion")
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(office,"run_owned",run)
    with pytest.raises(office.OfficeConversionError):office.libreoffice_pdf(source,output,"synthetic-soffice")


def test_local_backend_cancellation_and_timeout_never_fall_back(tmp_path,monkeypatch):
    from integrations.deeptutor_shchem_v1 import desktop_local_pagination as local
    from integrations.deeptutor_shchem_v1 import paper_export_renderer
    from integrations.deeptutor_shchem_v1.owned_process import ProcessTimeout
    source=document(tmp_path)
    monkeypatch.setattr(local,"find_libreoffice",lambda:Path("synthetic-soffice"))
    monkeypatch.setattr(paper_export_renderer,"_word_com_pdf",lambda *a,**k:pytest.fail("must not fall back after cancellation"))
    for index,error in enumerate((ReadCancelled(),ProcessTimeout(),office.OfficeConversionError())):
        def stop(*a,**k):raise error
        monkeypatch.setattr(office,"libreoffice_pdf",stop)
        with pytest.raises(type(error)):local.render_docx(source,tmp_path/str(index))


def test_cancellation_after_backend_result_does_not_accept_pdf(tmp_path,monkeypatch):
    source=document(tmp_path);event=Event()
    def run(*a,**k):
        (tmp_path/"synthetic.pdf").write_bytes(b"%PDF-late result")
        event.set()
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(office,"run_owned",run)
    with pytest.raises(ReadCancelled),read_cancel_scope(event):
        office.libreoffice_pdf(source,tmp_path,"synthetic-soffice")


@pytest.mark.skipif(os.name!="nt",reason="Production PowerShell ByRef binding")
@pytest.mark.parametrize("failure",["none","close","quit","takeover"])
def test_actual_powershell_cleanup_binding_and_stage_diagnostics_without_com(tmp_path,failure):
    from integrations.deeptutor_shchem_v1.owned_process import run_owned
    # Execute the production finally block against CLR ByRef method doubles.
    # The part that creates Word is never included in this script.
    prefix=office._WORD_SCRIPT.split("try {\n  Stage 'checking_input'",1)[0]
    cleanup=office._WORD_SCRIPT.split("# BEGIN_PRIVATE_CLEANUP:",1)[1].split("\n",1)[1].split("# END_PRIVATE_CLEANUP",1)[0]
    assert "New-Object -ComObject" not in prefix+cleanup
    replacement=r'''
$directory=[Environment]::GetEnvironmentVariable('SHCHEM_SYNTHETIC_CLEANUP','Process')
$testFailure=[Environment]::GetEnvironmentVariable('SHCHEM_SYNTHETIC_FAILURE','Process')
Add-Type -TypeDefinition @'
using System; using System.Runtime.InteropServices;
public class SyntheticClose {
 public bool Fail, Called;
 public void Close(ref object save) {
  if (!(save is int) || (int)save!=0) throw new InvalidOperationException("Wrong argument");
  Called=true; if (Fail) throw new COMException("Synthetic close failure",unchecked((int)0x80020005));
 }
}
public class SyntheticQuit {
 public bool Fail, Called;
 public void Quit(ref object save) {
  if (!(save is int) || (int)save!=0) throw new InvalidOperationException("Wrong argument");
  Called=true; if (Fail) throw new COMException("Synthetic quit failure",unchecked((int)0x80020005));
 }
}
'@
function Private-Word($empty=$false) { return ($testFailure -ne 'takeover') }
$doc=New-Object SyntheticClose; $word=New-Object SyntheticQuit
$doc.Fail=($testFailure -eq 'close'); $word.Fail=($testFailure -eq 'quit')
$owned=$true
'''
    script=tmp_path/"synthetic_cleanup.ps1"
    script.write_text(prefix+replacement+cleanup+"\nRecord 'calls' @{close=$doc.Called;quit=$word.Called}\n",encoding="utf-8-sig")
    executable=Path(os.environ["SystemRoot"])/"System32/WindowsPowerShell/v1.0/powershell.exe"
    result=run_owned([str(executable),"-STA","-NoProfile","-NonInteractive","-ExecutionPolicy","Bypass","-File",str(script)],
        env={**os.environ,"SHCHEM_SYNTHETIC_CLEANUP":str(tmp_path),"SHCHEM_SYNTHETIC_FAILURE":failure},timeout=15)
    assert result.returncode==0,result.stderr
    record=json.loads((tmp_path/"cleanup.json").read_text(encoding="utf-8-sig"))
    calls=json.loads((tmp_path/"calls.json").read_text(encoding="utf-8-sig"))
    if failure in {"close","quit"}:
        error=json.loads((tmp_path/"cleanup-exception.json").read_text(encoding="utf-8-sig"))
        assert error["step"]==("closing_document" if failure=="close" else "quitting_application")
        assert error["hresult"]==-2147352571 and error["error_type"]=="System.Runtime.InteropServices.COMException"
        assert record["exception"]["step"]==error["step"]
        assert calls["close"] and calls["quit"] is (failure=="quit")
    elif failure=="none":
        assert calls["close"] and calls["quit"]
        assert record["disposition"]=="empty_automation_quit_unverified"
        assert record["exception"] is None
    else:
        assert not calls["close"] and not calls["quit"]
        assert record["disposition"]=="taken_over_or_unverified_preserved"
    assert record["steps"][-1]["step"]=="cleanup_finished"


@pytest.mark.parametrize("verified",[True,False])
def test_parent_observation_can_verify_exit_only_after_bound_quit_return(tmp_path,monkeypatch,verified):
    source=document(tmp_path);target=tmp_path/"new.pdf";observed={}
    def run(command,*,env,**kwargs):
        observed.update(env)
        Path(env["SHCHEM_WORD_OUTPUT"]).write_bytes(b"%PDF-synthetic protocol only")
        write_record(env,"result",status="converted")
        write_record(env,"owner",held_handle=True,private=True,pid=2345,start_filetime=133800000000000000)
        write_record(env,"automation",identity="singleuse_empty_com",pid_verified_before_open=False)
        write_record(env,"cleanup",pid=2345,disposition="quit_requested",exception=None)
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(office,"run_owned",run)
    def observer(pid,created,*,timeout):
        assert pid==2345 and created==133800000000000000
        assert timeout==30
        return {"pid":pid,"start_filetime":created,"verified":verified,
                "method":"held_handle_exit" if verified else "still_running"}
    monkeypatch.setattr(office,"_observe_bound_process_exit",observer)
    if verified:office.word_pdf(source,target,session_root=tmp_path/"sessions")
    else:
        with pytest.raises(office.OfficeConversionError) as error:office.word_pdf(source,target,session_root=tmp_path/"sessions")
        assert error.value.cleanup_complete is False and not target.exists()
    session=Path(observed["SHCHEM_OFFICE_SESSION"])
    assert json.loads((session/"cleanup.json").read_text(encoding="utf-8"))["disposition"]=="quit_requested"
    assert json.loads((session/"parent_exit_verified.json").read_text(encoding="utf-8"))["verified"] is verified


@pytest.mark.skipif(os.name!="nt",reason="Windows held process identity observation")
@pytest.mark.parametrize("state",["exited","exits_while_observed","still_running","different_creation","cancelled"])
def test_real_synthetic_process_exit_observation_never_terminates_other_identity(state):
    import ctypes
    import subprocess
    import sys
    from ctypes import wintypes as w
    child=subprocess.Popen([sys.executable,"-B","-c","import time;time.sleep(.4)" if state=="exits_while_observed" else "import time;time.sleep(20)"],
                           creationflags=subprocess.CREATE_NO_WINDOW)
    kernel=ctypes.WinDLL("kernel32",use_last_error=True)
    kernel.GetProcessTimes.argtypes=[w.HANDLE,*[ctypes.POINTER(ctypes.c_uint64)]*4]
    values=[ctypes.c_uint64() for _ in range(4)]
    assert kernel.GetProcessTimes(int(child._handle),*map(ctypes.byref,values))
    created=values[0].value
    try:
        if state=="exited":child.terminate();child.wait(timeout=5)
        if state=="cancelled":
            import time
            from threading import Timer
            event=Event();timer=Timer(.1,event.set);timer.start();started=time.monotonic()
            try:
                with pytest.raises(ReadCancelled),read_cancel_scope(event):
                    office._observe_bound_process_exit(child.pid,created,timeout=30)
                assert time.monotonic()-started<2 and child.poll() is None
            finally:timer.cancel();timer.join()
            return
        observed=office._observe_bound_process_exit(child.pid,created+(1 if state=="different_creation" else 0),timeout=.8 if state=="exits_while_observed" else .1)
        assert observed["verified"] is (state!="still_running")
        if state in {"still_running","different_creation"}:assert child.poll() is None
        if state=="different_creation":assert observed["method"]=="pid_reused"
        if state=="exits_while_observed":assert observed["method"]=="held_handle_exit"
    finally:
        if child.poll() is None:child.terminate()
        child.wait(timeout=5)
