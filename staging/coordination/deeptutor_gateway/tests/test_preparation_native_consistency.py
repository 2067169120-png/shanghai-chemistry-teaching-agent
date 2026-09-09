"""Imported Word and preparation must share native paragraph semantics."""

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
    PreparationSourcesService,
    _paragraph,
    _table_rows,
)
from integrations.deeptutor_shchem_v1.word_native_text import WordNativeTextReader


def test_native_hyphens_default_scripts_and_inherited_hidden_are_consistent():
    document = Document()
    defaults = document.styles.element.find(qn("w:docDefaults"))
    properties = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    vertical = OxmlElement("w:vertAlign")
    vertical.set(qn("w:val"), "subscript")
    properties.append(vertical)
    hidden = document.styles.add_style("HiddenSource", WD_STYLE_TYPE.CHARACTER)
    hidden.font.hidden = True
    paragraph = document.add_paragraph()
    paragraph.add_run("2")
    run = paragraph.add_run("H")
    baseline = OxmlElement("w:vertAlign")
    baseline.set(qn("w:val"), "baseline")
    run._r.get_or_add_rPr().append(baseline)
    run._r.append(OxmlElement("w:noBreakHyphen"))
    text = OxmlElement("w:t")
    text.text = "Cl"
    run._r.append(text)
    run._r.append(OxmlElement("w:softHyphen"))
    paragraph.add_run("HIDDEN", style="HiddenSource")
    warnings = set()
    text = _paragraph(paragraph._p, document, warnings)
    assert text == WordNativeTextReader(document.styles.element).read(paragraph._p).text
    assert text == "_{2}H‑Cl\u00ad"
    assert warnings == {"存在隐藏文字，未加入备课参考。"}


def _wrap(node):
    control, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
    control.append(OxmlElement("w:sdtPr"))
    content.append(node)
    control.append(content)
    return control


def test_wrapped_table_rows_and_cells_remain_in_selected_reference(tmp_path):
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "浓盐酸"
    table.cell(0, 1).text = "1 mol/L"
    table.cell(1, 0).text = "稀硫酸"
    table.cell(1, 1).text = "2 mol/L"
    second = table._tbl.tr_lst[1]
    second.insert(0, OxmlElement("w:tblPrEx"))
    cell = second.tc_lst[1]
    second.remove(cell)
    second.append(_wrap(cell))
    table._tbl.remove(second)
    table._tbl.append(_wrap(second))
    warnings = set()
    rows = _table_rows(table._tbl, document, warnings)
    assert [[cell["text"] for cell in row["cells"]] for row in rows] == [
        ["浓盐酸", "1 mol/L"], ["稀硫酸", "2 mol/L"]
    ]
    assert warnings == set()
    native = WordNativeTextReader(document.styles.element).read(table._tbl)
    assert native.features == ()
    assert "稀硫酸" in native.text
    source = tmp_path / "wrapped-table.docx"
    document.save(source)
    service = PreparationSourcesService(tmp_path)
    preview = service.word_preview(source)
    reference = service.reference(source, preview["source_sha256"], 1, 1, [])
    for text in ("浓盐酸", "1 mol/L", "稀硫酸", "2 mol/L"):
        assert reference["materials"].count(text) == 1
