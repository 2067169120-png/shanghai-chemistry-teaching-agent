"""Source-bound Word pictures enter preparation without model calls or partial results."""

import hashlib
import io
import zipfile
from copy import deepcopy

import pytest
from docx import Document
from PIL import Image
from test_desktop_visual_import_facade import FakeProviderStore, _facade, _png
from test_desktop_visual_import_facade import (
    desktop_paths as desktop_paths,  # noqa: PLC0414
)

from integrations.deeptutor_shchem_v1.desktop_preparation_images import (
    MAX_IMAGES,
    PreparationImageStore,
    normalize_image_assets,
)
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
    PreparationSourceError,
)
from integrations.deeptutor_shchem_v1.desktop_word_questions import WordQuestionError


def _image_bytes(fmt, color="blue"):
    output = io.BytesIO()
    Image.new("RGB", (24, 32), color).save(output, format=fmt)
    return output.getvalue()


def _fixture(desktop_paths, tmp_path, *, count=1, answer_image=False, replace=None):
    doc = Document()
    doc.add_heading("电解质练习", 1)
    for number in range(1, count + 1):
        doc.add_paragraph(f"【即学即练{number}】观察图示后回答问题。")
        doc.add_paragraph().add_run().add_picture(io.BytesIO(_png((number, 10, 20))))
        doc.add_paragraph("【答案】参考答案")
        if answer_image:
            doc.add_paragraph().add_run().add_picture(
                io.BytesIO(_png((number, 10, 20)))
            )
    path = tmp_path / "原图演示讲义.docx"
    doc.save(path)
    if replace:
        source, result = io.BytesIO(path.read_bytes()), io.BytesIO()
        with (
            zipfile.ZipFile(source) as original,
            zipfile.ZipFile(result, "w") as package,
        ):
            for name in original.namelist():
                package.writestr(name, replace.get(name, original.read(name)))
        path.write_bytes(result.getvalue())
    provider = FakeProviderStore(configured=False)
    facade = _facade(desktop_paths, provider)
    facade.save_visual_import_batch(handout_files=(path,), source_type="教师讲义")
    choices = [
        {"key": item["key"], "revision": item["revision"], "points": 3}
        for item in facade.word_question_catalog()["items"]
    ]
    assert len(choices) == count
    return facade, choices, provider


def _store_root(desktop_paths):
    return desktop_paths.task_root / "preparation-v1" / "images"


@pytest.mark.parametrize("fmt", ["PNG", "JPEG", "WEBP"])
def test_preview_is_pure_and_supported_exact_bytes_are_imported(
    desktop_paths, tmp_path, fmt
):
    raw = _image_bytes(fmt)
    facade, choices, provider = _fixture(
        desktop_paths, tmp_path, replace={"word/media/image1.png": raw}
    )
    ref = facade.word_question_reference(choices)
    assert not _store_root(desktop_paths).exists()
    assert set(ref) == {
        "materials",
        "warnings",
        "selections",
        "include_images",
        "image_assets",
        "image_issues",
    }
    assert ref["selections"] == choices
    assert ref["include_images"] is True
    assert not ref["image_issues"]
    assert ref == facade.word_question_reference(choices)
    asset = ref["image_assets"][0]
    assert normalize_image_assets([asset]) == [asset]
    assert asset["sha256"] == hashlib.sha256(raw).hexdigest()
    assert asset["asset_id"] in ref["materials"]
    assert "题面原文 · 区块3" in ref["materials"]
    for content in (ref["materials"], "\n".join(ref["warnings"])):
        assert "确认导入后仅把原图保存为本地备课素材，不调用模型" in content
        assert (
            "后续是否发送图片像素，以生成时的“本地排版/视觉读取”选择及发送预览为准"
            in content
        )
        assert "不能仅凭图注补写图中条件" in content
        assert "未把原图像素发送给模型" not in content
        assert "模型仅见图注" not in content
    assert str(tmp_path) not in ref["materials"]
    result = facade.import_word_question_reference(ref, [])
    assert result == {
        key: ref[key] for key in ("materials", "warnings", "image_assets")
    }
    assert PreparationImageStore(_store_root(desktop_paths)).load(asset) == raw
    assert provider.borrow_calls == 0


def test_same_image_roles_remain_explicit_after_dedup(desktop_paths, tmp_path):
    facade, choices, provider = _fixture(desktop_paths, tmp_path, answer_image=True)
    ref = facade.word_question_reference(choices)
    assert len(ref["image_assets"]) == 1
    asset = ref["image_assets"][0]
    assert ref["materials"].count(asset["asset_id"]) == 2
    assert "题面原文 · 区块3" in ref["materials"]
    assert "原文答案与解析 · 区块5" in ref["materials"]
    assert "兼有题面、答案引用" in asset["purpose"]
    assert "答案引用限后续教师讲评" in asset["purpose"]
    assert asset["caption"] == "Word复用原图1 · 具体题号与角色见参考"
    result = facade.import_word_question_reference(ref, [])
    assert result["image_assets"] == ref["image_assets"]
    assert len(list(_store_root(desktop_paths).glob("*.image"))) == 1
    assert provider.borrow_calls == 0


def test_answer_only_image_is_restricted_to_teacher_explanation(
    desktop_paths, tmp_path, monkeypatch
):
    facade, choices, _ = _fixture(desktop_paths, tmp_path, answer_image=True)
    service = facade._word_questions()
    original_resolve = service._resolve

    def answer_only(values, **kwargs):
        rows, inventory = original_resolve(values, **kwargs)
        rows[0]["question_blocks"] = rows[0]["question_blocks"][:1]
        return rows, inventory

    monkeypatch.setattr(service, "_resolve", answer_only)
    ref = service.reference(choices)
    assert len(ref["image_assets"]) == 1
    assert (
        "仅有答案引用，只用于教师讲评，不放入学生题面"
        in ref["image_assets"][0]["purpose"]
    )
    assert "原文答案与解析" in ref["image_assets"][0]["caption"]


def test_unreadable_shape_is_not_treated_as_image_free_text(
    desktop_paths, tmp_path, monkeypatch
):
    facade, choices, _ = _fixture(desktop_paths, tmp_path)
    service = facade._word_questions()
    original_resolve = service._resolve

    def missing_shape(values, **kwargs):
        rows, inventory = original_resolve(values, **kwargs)
        rows[0]["question_blocks"][1]["assets"] = []
        return rows, inventory

    monkeypatch.setattr(service, "_resolve", missing_shape)
    ref = service.reference(choices)
    assert ref["image_assets"] == []
    assert "未找到可独立读取的原图或对象预览" in ref["image_issues"][0]
    assert "第1题 · 题面原文 · 区块3" in ref["image_issues"][0]
    with pytest.raises(WordQuestionError, match="未能全部带入"):
        service.prepare_reference(ref, [])
    text_ref = service.reference(choices, include_images=False)
    assert not text_ref["image_issues"]
    assert "未附原图或对象" in text_ref["materials"]


def test_repeated_common_material_still_has_each_question_image_reference(
    desktop_paths, tmp_path, monkeypatch
):
    facade, choices, _ = _fixture(desktop_paths, tmp_path, count=2)
    service = facade._word_questions()
    original_resolve = service._resolve

    def with_shared_context(values, **kwargs):
        rows, inventory = original_resolve(values, **kwargs)
        shared = deepcopy(rows[0]["question_blocks"][1:])
        for row in rows:
            row["context_blocks"] = deepcopy(shared)
        return rows, inventory

    monkeypatch.setattr(service, "_resolve", with_shared_context)
    original_image = service.reader.word_asset_bytes
    calls = []

    def image(data, asset_id, **kwargs):
        calls.append(asset_id)
        return original_image(data, asset_id, **kwargs)

    monkeypatch.setattr(service.reader, "word_asset_bytes", image)
    ref = service.reference(choices)
    assert "共同材料：与前面所选题目相同" in ref["materials"]
    assert "第1题 · 共同材料 · 区块3" in ref["materials"]
    assert "第2题 · 共同材料 · 区块3" in ref["materials"]
    assert len(ref["image_assets"]) == 2
    assert len(calls) == 2  # Same source/asset is decoded only once per batch.


@pytest.mark.parametrize("fmt", ["GIF", "BMP", "TIFF"])
def test_unsupported_images_are_not_omitted_and_text_only_is_explicit(
    desktop_paths, tmp_path, fmt
):
    facade, choices, provider = _fixture(
        desktop_paths,
        tmp_path,
        count=2,
        replace={"word/media/image2.png": _image_bytes(fmt)},
    )
    ref = facade.word_question_reference(choices)
    assert len(ref["image_assets"]) == 1
    assert len(ref["image_issues"]) == 1
    assert "第2题 · 题面原文 · 区块6" in ref["image_issues"][0]
    assert "PNG、JPEG和WebP" in ref["image_issues"][0]
    with pytest.raises(WordQuestionError, match="未能全部带入"):
        facade.import_word_question_reference(ref, [])
    assert not _store_root(desktop_paths).exists()
    text_ref = facade.word_question_reference(choices, include_images=False)
    assert text_ref["image_assets"] == [] and text_ref["image_issues"] == []
    for content in (text_ref["materials"], "\n".join(text_ref["warnings"])):
        assert "本次已明确选择仅文字：未带入任何原图，也未把原图发送给模型" in content
        assert "图中条件未识别，须由教师补充文字后再用于讲解或解题" in content
        assert "保存为本地备课素材" not in content
        assert "模型仅见图注" not in content
    assert text_ref["materials"].count("未附原图，本次仅使用文字") == 2
    result = facade.import_word_question_reference(text_ref, [])
    assert result["image_assets"] == []
    assert not _store_root(desktop_paths).exists()
    assert provider.borrow_calls == 0


@pytest.mark.parametrize("broken", [b"not an image", b""])
def test_invalid_image_bytes_are_reported_as_a_batch_issue(
    desktop_paths, tmp_path, broken
):
    facade, choices, _ = _fixture(
        desktop_paths, tmp_path, count=2, replace={"word/media/image2.png": broken}
    )
    ref = facade.word_question_reference(choices)
    assert len(ref["image_assets"]) == 1
    assert ref["image_issues"]
    with pytest.raises(WordQuestionError, match="未能全部带入"):
        facade.import_word_question_reference(ref, [])
    assert not _store_root(desktop_paths).exists()


def test_missing_image_is_reported_with_every_role(
    desktop_paths, tmp_path, monkeypatch
):
    facade, choices, _ = _fixture(desktop_paths, tmp_path, answer_image=True)

    def missing(*args, **kwargs):
        raise PreparationSourceError("未找到对应原图")

    monkeypatch.setattr(facade._word_questions().reader, "word_asset_bytes", missing)
    ref = facade.word_question_reference(choices)
    assert ref["image_assets"] == []
    assert len(ref["image_issues"]) == 2
    assert "题面原文" in ref["image_issues"][0]
    assert "原文答案与解析" in ref["image_issues"][1]
    with pytest.raises(WordQuestionError, match="未能全部带入"):
        facade.import_word_question_reference(ref, [])


def test_preview_exceeding_limit_is_complete_but_import_rejects(
    desktop_paths, tmp_path
):
    facade, choices, _ = _fixture(desktop_paths, tmp_path, count=MAX_IMAGES + 1)
    ref = facade.word_question_reference(choices)
    assert len(ref["image_assets"]) == MAX_IMAGES + 1
    assert ref["image_issues"] == []
    assert "未删减图片" in "".join(ref["warnings"])
    with pytest.raises(WordQuestionError, match="最多选择12张"):
        facade.import_word_question_reference(ref, [])
    assert not _store_root(desktop_paths).exists()


def test_merge_preserves_order_metadata_and_duplicate_image_roles(
    desktop_paths, tmp_path
):
    facade, choices, _ = _fixture(desktop_paths, tmp_path, count=2, answer_image=True)
    ref = facade.word_question_reference(choices)
    store = PreparationImageStore(_store_root(desktop_paths))
    extra = store.import_bytes(_png("yellow"), "教师自己的图注", "教师资料", "教师用途")
    duplicate = store.import_bytes(
        _png((1, 10, 20)), "已修改图注", "已确认来源", "已确认用途"
    )
    existing = [extra, duplicate]
    snapshot = deepcopy(existing)
    result = facade.import_word_question_reference(ref, existing)
    assert existing == snapshot
    assert result["image_assets"] == [extra, duplicate, ref["image_assets"][1]]
    assert duplicate["asset_id"] in result["materials"]
    assert "答案图，只用于教师讲评" in result["materials"]


def test_combined_limit_fails_without_saving_new_images(desktop_paths, tmp_path):
    facade, choices, _ = _fixture(desktop_paths, tmp_path, count=2)
    ref = facade.word_question_reference(choices)
    store = PreparationImageStore(_store_root(desktop_paths))
    existing = [
        store.import_bytes(_png((n, 50, 70)), f"图{n}", "教师", "讲解")
        for n in range(MAX_IMAGES - 1)
    ]
    before = set(store.root.iterdir())
    with pytest.raises(WordQuestionError, match="最多选择12张"):
        facade.import_word_question_reference(ref, existing)
    assert set(store.root.iterdir()) == before


@pytest.mark.parametrize("defect", ["missing", "corrupt", "metadata"])
def test_invalid_existing_image_prevents_all_new_copies(
    desktop_paths, tmp_path, defect
):
    facade, choices, _ = _fixture(desktop_paths, tmp_path, count=2)
    ref = facade.word_question_reference(choices)
    store = PreparationImageStore(_store_root(desktop_paths))
    existing = store.import_bytes(_png("yellow"), "现有图片", "教师资料", "讲解")
    path = store.root / (existing["sha256"] + ".image")
    if defect == "missing":
        path.unlink()
    elif defect == "corrupt":
        path.write_bytes(b"changed test image")
    else:
        existing["width"] += 1
    before = {p.name: p.read_bytes() for p in store.root.iterdir()}
    with pytest.raises(WordQuestionError):
        facade.import_word_question_reference(ref, [existing])
    assert {p.name: p.read_bytes() for p in store.root.iterdir()} == before


@pytest.mark.parametrize(
    "defect", ["duplicate", "extra_key", "blank_caption", "wrong_hash", "whitespace"]
)
def test_existing_metadata_is_validated_without_mutation(
    desktop_paths, tmp_path, defect
):
    facade, choices, _ = _fixture(desktop_paths, tmp_path)
    ref = facade.word_question_reference(choices)
    store = PreparationImageStore(_store_root(desktop_paths))
    existing = store.import_bytes(_png("yellow"), "现有图片", "教师资料", "讲解")
    assets = [existing]
    if defect == "duplicate":
        assets.append(deepcopy(existing))
    elif defect == "extra_key":
        existing["unexpected"] = "extra"
    elif defect == "blank_caption":
        existing["caption"] = ""
    elif defect == "whitespace":
        existing["caption"] = " 现有图片 "
    else:
        existing["sha256"] = "a" * 64
    snapshot = deepcopy(assets)
    before = {p.name: p.read_bytes() for p in store.root.iterdir()}
    with pytest.raises(WordQuestionError):
        facade.import_word_question_reference(ref, assets)
    assert assets == snapshot
    assert {p.name: p.read_bytes() for p in store.root.iterdir()} == before


def test_late_storage_failure_returns_no_partial_form_result(
    desktop_paths, tmp_path, monkeypatch
):
    facade, choices, _ = _fixture(desktop_paths, tmp_path, count=2)
    ref = facade.word_question_reference(choices)
    snapshot = deepcopy(ref)
    original = PreparationImageStore.import_bytes
    calls = []

    def fail_second(store, *args):
        calls.append(True)
        if len(calls) == 2:
            raise OSError("test storage error")
        return original(store, *args)

    monkeypatch.setattr(PreparationImageStore, "import_bytes", fail_second)
    with pytest.raises(WordQuestionError, match="未导入本批参考"):
        facade.import_word_question_reference(ref, [])
    assert ref == snapshot
    # A verified content-addressed cache is allowed to remain, but no successful
    # partial batch is returned to the form and no source is modified.
    assert len(list(_store_root(desktop_paths).glob("*.image"))) == 1


@pytest.mark.parametrize(
    "part",
    ["materials", "warnings", "selections", "image_assets", "include_images", "extra"],
)
def test_entire_confirmed_reference_must_match_recompiled_source(
    desktop_paths, tmp_path, part
):
    facade, choices, _ = _fixture(desktop_paths, tmp_path)
    ref = facade.word_question_reference(choices)
    if part == "materials":
        ref[part] += "更改"
    elif part == "warnings":
        ref[part] = []
    elif part == "selections":
        ref[part][0]["points"] = 4
    elif part == "image_assets":
        ref[part][0]["caption"] = "篡改图注"
    elif part == "include_images":
        ref[part] = False
    else:
        ref[part] = "unexpected"
    with pytest.raises(WordQuestionError, match="重新确认"):
        facade.import_word_question_reference(ref, [])
    assert not _store_root(desktop_paths).exists()


def test_source_change_invalidates_prepared_reference(desktop_paths, tmp_path):
    facade, choices, _ = _fixture(desktop_paths, tmp_path)
    ref = facade.word_question_reference(choices)
    item = facade.word_question_catalog()["items"][0]
    source = (
        desktop_paths.state_root
        / "visual-import-v2"
        / "sources"
        / f"{item['source_sha256']}.docx"
    )
    source.write_bytes(b"changed archive fixture")
    with pytest.raises(WordQuestionError, match="变化"):
        facade.import_word_question_reference(ref, [])
    assert not _store_root(desktop_paths).exists()


@pytest.mark.parametrize("value", [None, 1, "false", [], {}])
def test_include_images_requires_an_actual_boolean(desktop_paths, tmp_path, value):
    facade, choices, _ = _fixture(desktop_paths, tmp_path)
    with pytest.raises(WordQuestionError, match="带入原图"):
        facade.word_question_reference(choices, include_images=value)


def test_prepare_revalidates_inventory_once(desktop_paths, tmp_path, monkeypatch):
    facade, choices, _ = _fixture(desktop_paths, tmp_path, count=3)
    ref = facade.word_question_reference(choices)
    service = facade._word_questions()
    original_inventory = service._inventory
    calls = []

    def inventory(locations=None):
        calls.append(True)
        return original_inventory(locations)

    monkeypatch.setattr(service, "_inventory", inventory)
    facade.import_word_question_reference(ref, [])
    assert len(calls) == 1
