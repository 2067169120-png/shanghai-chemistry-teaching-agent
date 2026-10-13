"""Valid JPEG pixels survive missing or incomplete metadata without re-encoding."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile
import struct
import warnings

from docx import Document
from docx.image.exceptions import UnexpectedEndOfFileError, UnrecognizedImageError
from docx.oxml.ns import qn
from docx.shared import Mm
from PIL import Image
import pytest

from integrations.deeptutor_shchem_v1.desktop_preparation_images import (
    PreparationImageError, image_info, verify_image_bytes)
from integrations.deeptutor_shchem_v1.desktop_word_picture import (
    add_word_picture, add_word_document_picture)


def picture(fmt='JPEG', *, incomplete_exif=False, no_app_markers=False, color='#5c879d', dpi=(72, 72)):
    buffer = BytesIO()
    Image.new('RGB', (80, 40), color).save(buffer, fmt, dpi=dpi)
    data = buffer.getvalue()
    if no_app_markers:
        # This valid JPEG shape occurs in publisher-generated photo thumbnails:
        # SOI is followed directly by quantization tables, with no JFIF/EXIF.
        # Remove only APP0/APP1 metadata; keep every coded pixel byte intact.
        assert fmt == 'JPEG' and data[:2] == b'\xff\xd8'
        original = data; chunks = [data[:2]]; offset = 2
        while data[offset:offset + 2] != b'\xff\xda':
            assert data[offset] == 0xff
            length = int.from_bytes(data[offset + 2:offset + 4], 'big')
            end = offset + 2 + length
            if data[offset + 1] not in (0xe0, 0xe1):
                chunks.append(data[offset:end])
            offset = end
        data = b''.join(chunks) + data[offset:]
        with Image.open(BytesIO(original)) as before, Image.open(BytesIO(data)) as after:
            assert before.size == after.size and before.mode == after.mode == 'RGB'
            assert before.tobytes() == after.tobytes()
    if incomplete_exif:
        # One optional ASCII description promises more bytes than the APP1
        # segment contains. JPEG pixels are untouched and fully decodable.
        tiff = (b'II*\x00' + struct.pack('<I', 8) + struct.pack('<H', 1)
                + struct.pack('<HHII', 0x010E, 2, 100, 26)
                + struct.pack('<I', 0) + b'photo\x00')
        exif = b'Exif\x00\x00' + tiff
        data = data[:2] + b'\xff\xe1' + struct.pack('>H', len(exif)+2) + exif + data[2:]
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        info = image_info(data)
    digest = sha256(data).hexdigest()
    asset = {'asset_id': 'IMG-'+digest, 'sha256': digest, 'caption': '树脂颗粒',
             'source': '合成测试图', 'purpose': '检查原字节图片嵌入', **info}
    return data, asset


@pytest.fixture
def exif_photo():
    data, asset = picture(incomplete_exif=True)
    with pytest.raises(UnexpectedEndOfFileError):
        Document().add_picture(BytesIO(data), width=Mm(40))
    return data, asset


@pytest.fixture
def bare_jpeg_photo():
    data, asset = picture(no_app_markers=True)
    with pytest.raises(UnrecognizedImageError):
        Document().add_picture(BytesIO(data), width=Mm(40))
    return data, asset


def save_reopen(document):
    buffer = BytesIO(); document.save(buffer)
    with ZipFile(BytesIO(buffer.getvalue())) as archive:
        assert archive.testzip() is None
        embedded = {name: archive.read(name) for name in archive.namelist()
                    if name.startswith('word/media/')}
    return Document(BytesIO(buffer.getvalue())), embedded


@pytest.mark.parametrize('size,expected', [
    ({'width': Mm(40)}, (40, 20)),
    ({'height': Mm(20)}, (40, 20)),
    ({'width': Mm(39), 'height': Mm(17)}, (39, 17)),
])
def test_incomplete_exif_embeds_exact_original_bytes_with_requested_geometry(exif_photo, size, expected):
    data, asset = exif_photo
    document = Document()
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        shape = add_word_document_picture(document, data, asset, **size)
    assert (shape.width.mm, shape.height.mm) == expected
    reopened, embedded = save_reopen(document)
    assert list(embedded.values()) == [data]
    assert len(reopened.paragraphs) == 1 and len(reopened.inline_shapes) == 1
    assert (reopened.inline_shapes[0].width.mm, reopened.inline_shapes[0].height.mm) == expected


def test_source_order_styles_image_dedup_and_unique_drawing_ids(exif_photo):
    data, asset = exif_photo
    png, png_asset = picture('PNG', color='#a17732')
    document = Document(); paragraph = document.add_paragraph(style='Heading 2')
    run = paragraph.add_run('前文'); run.bold = True
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        add_word_picture(run, data, asset, width=Mm(40), height=Mm(20))
        run.add_text('中间')
        add_word_picture(run, png, png_asset, width=Mm(20), height=Mm(10))
        run.add_text('后文')
        add_word_picture(run, data, asset, width=Mm(30), height=Mm(15))
    assert paragraph.style.name == 'Heading 2' and run.bold is True
    reopened, embedded = save_reopen(document)
    paragraph = reopened.paragraphs[0]
    assert paragraph.text == '前文中间后文' and paragraph.style.name == 'Heading 2'
    assert paragraph.runs[0].bold is True
    ids = paragraph._p.xpath('.//wp:docPr/@id')
    assert len(ids) == len(set(ids)) == 3
    relationships = paragraph._p.xpath('.//a:blip/@r:embed')
    blobs = [reopened.part.related_parts[identity].blob for identity in relationships]
    assert blobs == [data, png, data]
    assert set(embedded.values()) == {data, png}


@pytest.mark.parametrize('fmt,dpi', [('PNG', (72, 72)), ('JPEG', (72, 144))])
def test_successful_native_path_preserves_existing_dpi_geometry_and_xml(fmt, dpi):
    data, asset = picture(fmt, dpi=dpi)
    before = Document(); before.add_picture(BytesIO(data), width=Mm(40))
    after = Document(); add_word_document_picture(after, data, asset, width=Mm(40))
    assert before._element.xml == after._element.xml
    assert before.inline_shapes[0].height == after.inline_shapes[0].height
    assert list(save_reopen(after)[1].values()) == [data]


@pytest.mark.parametrize('problem', ['changed_bytes', 'wrong_dimensions', 'undecodable_pixels'])
def test_invalid_or_changed_images_remain_rejected_before_embedding(exif_photo, problem):
    data, original = exif_photo; asset = deepcopy(original)
    if problem == 'changed_bytes':
        data = data + b'changed'
    elif problem == 'wrong_dimensions':
        asset['width'] += 1
    else:
        data = data[:-20]
        asset['sha256'] = sha256(data).hexdigest(); asset['asset_id'] = 'IMG-'+asset['sha256']
    document = Document(); run = document.add_paragraph().add_run('保留已有文字')
    with pytest.raises(PreparationImageError), warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        add_word_picture(run, data, asset, width=Mm(40))
    assert run.text == '保留已有文字' and len(document.inline_shapes) == 0
    assert len(document.part.package.image_parts) == 0


def test_other_exception_classes_and_other_formats_are_not_swallowed(monkeypatch):
    data, asset = picture('PNG'); document = Document(); run = document.add_paragraph().add_run()
    def fail(*args, **kwargs):
        raise UnexpectedEndOfFileError()
    monkeypatch.setattr(run, 'add_picture', fail)
    with pytest.raises(UnexpectedEndOfFileError):
        add_word_picture(run, data, asset, width=Mm(40))
    def different_failure(*args, **kwargs):
        raise RuntimeError('unrelated drawing failure')
    monkeypatch.setattr(run, 'add_picture', different_failure)
    with pytest.raises(RuntimeError, match='unrelated drawing failure'):
        add_word_picture(run, data, asset, width=Mm(40))
    assert len(document.part.package.image_parts) == 0


@pytest.mark.parametrize('route', ['node', 'classroom_notes', 'composition'])
def test_lesson_plan_and_student_word_routes_keep_exif_photo(tmp_path, exif_photo, route):
    from test_lesson_design import sample
    from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
        resource, block, add_resource, enable_composition)
    from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents
    data, asset = exif_photo; payload = sample(); payload['image_assets'] = [asset]
    if route == 'composition':
        payload['lesson_design'] = enable_composition(payload['lesson_design'])
        composition = payload['lesson_design']['composition']
        item = resource('personal_text', '原照片', [block('image', 'personal', asset['asset_id'])],
            {'source_path': '个人测试材料', 'source_sha256': 'unknown', 'source_revision': 'unknown'})
        payload['lesson_design']['composition'] = add_resource(composition, item,
            {'page_id': composition['pages'][0], 'x': 12, 'y': 12, 'width': 186, 'height': 220})
    else:
        payload['lesson_design']['nodes'][0]['image_ids'] = [asset['asset_id']]
        if route == 'classroom_notes': payload['lesson_design']['classroom_notes'] = True
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        write_documents(payload, tmp_path, {asset['asset_id']: data})
    for name in ('lesson_plan.docx', 'student_worksheet.docx'):
        document = Document(tmp_path/name)
        assert len(document.inline_shapes) == 1
        with ZipFile(tmp_path/name) as archive:
            embedded = [archive.read(path) for path in archive.namelist() if path.startswith('word/media/')]
        assert embedded == [data]


def test_native_table_cell_keeps_exact_exif_photo_and_text(tmp_path, exif_photo):
    from test_handout_knowledge_tables import row
    from integrations.deeptutor_shchem_v1.desktop_handout_knowledge_tables import write_word_table
    data, asset = exif_photo; document = Document()
    tokens = {(0, 0): [{'kind': 'text', 'text': '原表前文'},
        {'kind': 'image', 'asset_id': asset['asset_id'], 'width_mm': 40, 'height_mm': 20,
         'image_block_id': 'CB-'+'1'*32}, {'kind': 'text', 'text': '原表后文'}]}
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        write_word_table(document, [row('原表前文原表后文')], cell_tokens=tokens,
            assets={asset['asset_id']: asset}, image_data={asset['asset_id']: data})
    reopened, embedded = save_reopen(document)
    assert reopened.tables[0].cell(0, 0).text == '原表前文原表后文'
    assert len(reopened.inline_shapes) == 1
    assert list(embedded.values()) == [data]
    assert (reopened.inline_shapes[0].width.mm, reopened.inline_shapes[0].height.mm) == (40, 20)


@pytest.mark.parametrize('size,expected', [
    ({'width': Mm(40)}, (40, 20)),
    ({'height': Mm(20)}, (40, 20)),
    ({'width': Mm(39), 'height': Mm(17)}, (39, 17)),
])
def test_jpeg_without_app_markers_keeps_raw_bytes_rgb_and_geometry(bare_jpeg_photo, size, expected):
    data, asset = bare_jpeg_photo; frozen = deepcopy(asset)
    document = Document()
    shape = add_word_document_picture(document, data, asset, **size)
    reopened, embedded = save_reopen(document)
    assert list(embedded.values()) == [data]
    assert (shape.width.mm, shape.height.mm) == expected
    assert (reopened.inline_shapes[0].width.mm, reopened.inline_shapes[0].height.mm) == expected
    with Image.open(BytesIO(data)) as before, Image.open(BytesIO(next(iter(embedded.values())))) as after:
        assert before.mode == after.mode == 'RGB' and before.size == after.size
        assert before.tobytes() == after.tobytes()
    assert asset == frozen


def test_jpeg_without_app_markers_deduplicates_and_retains_run_order(bare_jpeg_photo):
    data, asset = bare_jpeg_photo; png, png_asset = picture('PNG')
    document = Document(); paragraph = document.add_paragraph(style='Heading 2')
    run = paragraph.add_run('前文'); run.bold = True
    add_word_picture(run, data, asset, width=Mm(40)); run.add_text('中间')
    add_word_picture(run, png, png_asset, width=Mm(20)); run.add_text('后文')
    add_word_picture(run, data, asset, height=Mm(15))
    reopened, embedded = save_reopen(document)
    p = reopened.paragraphs[0]
    assert p.text == '前文中间后文' and p.style.name == 'Heading 2' and p.runs[0].bold
    ids = p._p.xpath('.//wp:docPr/@id')
    assert len(ids) == len(set(ids)) == 3
    relationships = p._p.xpath('.//a:blip/@r:embed')
    assert [reopened.part.related_parts[rid].blob for rid in relationships] == [data, png, data]
    assert set(embedded.values()) == {data, png}


@pytest.mark.parametrize('fixture', ['exif_photo', 'bare_jpeg_photo'])
def test_metadata_fallback_requires_an_explicit_size(request, fixture):
    data, asset = request.getfixturevalue(fixture)
    document = Document(); run = document.add_paragraph().add_run('保留正文')
    with pytest.raises(ValueError, match='显示尺寸'), warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        add_word_picture(run, data, asset)
    assert run.text == '保留正文' and not document.inline_shapes
    assert not document.part.package.image_parts


@pytest.mark.parametrize('fmt', ['PNG', 'WEBP'])
def test_unrecognized_non_jpeg_is_not_treated_as_native_jpeg(monkeypatch, fmt):
    data, asset = picture(fmt)
    document = Document(); run = document.add_paragraph().add_run('保留正文')
    def fail(*args, **kwargs):
        raise UnrecognizedImageError()
    monkeypatch.setattr(run, 'add_picture', fail)
    with pytest.raises(UnrecognizedImageError):
        add_word_picture(run, data, asset, width=Mm(40))
    assert run.text == '保留正文' and not document.inline_shapes
    assert not document.part.package.image_parts


@pytest.mark.parametrize('problem', ['changed_bytes', 'wrong_dimensions', 'undecodable_pixels'])
def test_bare_jpeg_verification_rejects_before_docx(monkeypatch, bare_jpeg_photo, problem):
    data, asset = bare_jpeg_photo; asset = deepcopy(asset)
    if problem == 'changed_bytes':
        data += b'changed'
    elif problem == 'wrong_dimensions':
        asset['width'] += 1
    else:
        data = data[:-20]
        asset['sha256'] = sha256(data).hexdigest(); asset['asset_id'] = 'IMG-' + asset['sha256']
    document = Document(); run = document.add_paragraph().add_run('保留正文')
    calls = []
    def must_not_parse(*args, **kwargs):
        calls.append(True)
        raise AssertionError('Unverified bytes reached python-docx')
    monkeypatch.setattr(run, 'add_picture', must_not_parse)
    with pytest.raises(PreparationImageError):
        add_word_picture(run, data, asset, width=Mm(40))
    assert not calls and not document.inline_shapes and not document.part.package.image_parts
    assert run.text == '保留正文'
