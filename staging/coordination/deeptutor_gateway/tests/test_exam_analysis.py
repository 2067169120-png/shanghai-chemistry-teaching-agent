from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
import json
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
import pytest
from integrations.deeptutor_shchem_v1.desktop_exam_data import (ExamError,read_xlsx,build_exam,analyse,local_student_advice,ExamStore,digest)
from integrations.deeptutor_shchem_v1.desktop_exam_ai import model_payload,generate_advice,validate_candidate,load_paper
from integrations.deeptutor_shchem_v1.desktop_exam_report import html_report
from integrations.deeptutor_shchem_v1.desktop_exam_fixture import example,candidate,FixtureStore,FixtureTransport

@pytest.fixture
def sample(tmp_path):return example(tmp_path/'example.xlsx')

def test_exact_statistics_and_denominators(sample):
    _,_,e=sample;r=analyse(e)
    assert r['overall']['n']==10 and r['overall']['enrolled']==12
    assert r['overall']['absent']==1 and r['overall']['unavailable']==1
    assert r['overall']['mean']==69 and r['overall']['median']==72
    assert r['overall']['pass_rate']==.7 and r['overall']['high_rate']==.3
    assert sum(b['n'] for b in r['distribution'])==10
    assert r['items'][0]['n']==11 and r['items'][2]['n']==10
    assert r['items'][2]['missing']==1

def test_class_selection_has_independent_population(sample):
    e=sample[2];a=analyse(e,'合成班A');b=analyse(e,'合成班B')
    assert a['overall']['n']==4 and b['overall']['n']==6
    assert len(a['students'])==len(b['students'])==6
    assert analyse(e,'不存在')['overall']['mean'] is None

def test_total_only_has_no_invented_item_mastery(sample):
    b,c,_=sample;c['questions']=[];r=analyse(build_exam(b,c))
    assert r['items']==[] and r['knowledge']==[] and r['overall']['n']==10

def test_item_sum_requires_complete_nonoverlapping_mapping(sample):
    b,c,_=sample;c['total_col']=None;r=analyse(build_exam(b,c));assert r['overall']['mean']==69
    c['questions']=c['questions'][:3]
    with pytest.raises(ExamError,match='全部非重叠'):build_exam(b,c)

@pytest.mark.parametrize('bad',[True,-1,101,float('inf'),'优秀',{'error':'错误公式'}])
def test_invalid_total_never_becomes_zero(sample,bad):
    b,c,_=sample;b['sheets'][0]['rows'][1][3]=bad;e=build_exam(b,c)
    assert e['students'][0]['total'] is None and analyse(e)['overall']['n']==9

def test_explicit_zero_counts(sample):
    b,c,_=sample;b['sheets'][0]['rows'][1][3]=0;c['questions']=[]
    e=build_exam(b,c);assert e['students'][0]['total']==0 and analyse(e)['overall']['n']==10

def test_total_disagreement_keeps_valid_items_but_excludes_total(sample):
    b,c,_=sample;b['sheets'][0]['rows'][1][3]=99;e=build_exam(b,c)
    assert e['students'][0]['scores']['1']==23 and e['students'][0]['total'] is None
    assert any('不一致' in i['detail'] for i in e['issues'])

def test_duplicate_identity_is_blocked_not_silently_overwritten(sample):
    b,c,_=sample;b['sheets'][0]['rows'][2][0]=b['sheets'][0]['rows'][1][0]
    with pytest.raises(ExamError,match='重复'):build_exam(b,c)

def test_duplicate_scores_or_question_ids_blocked(sample):
    b,c,_=sample;c['questions'][1]['column']=4
    with pytest.raises(ExamError):build_exam(b,c)

def test_invalid_thresholds_rejected(sample):
    b,c,_=sample;c['pass_score']=90;c['excellent_score']=60
    with pytest.raises(ExamError,match='达标线'):build_exam(b,c)

def test_missing_id_reports_source_row(sample):
    b,c,_=sample;b['sheets'][0]['rows'][1][0]='';e=build_exam(b,c)
    assert len(e['students'])==11 and any(i['row']==2 and i['field']=='学生' for i in e['issues'])

def test_knowledge_is_weighted_and_unknown_not_invented(sample):
    b,c,_=sample;c['questions'][0]['knowledge']='原理';c['questions'][1]['knowledge']='原理';c['questions'][2]['knowledge']=''
    r=analyse(build_exam(b,c));k=next(x for x in r['knowledge'] if x['label']=='原理')
    assert k['responses']==22 and k['possible_points']==550
    assert all('3' not in x['questions'] for x in r['knowledge'])

def test_xlsx_formula_no_cache_reports_error(sample,tmp_path):
    p=tmp_path/'example.xlsx'
    with ZipFile(p) as z:files={n:z.read(n) for n in z.namelist()}
    from xml.etree import ElementTree as ET
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    data=ET.fromstring(files['xl/worksheets/sheet1.xml'])
    cell=data.find('.//s:c[@r="D2"]',ns);v=cell.find('s:v',ns);assert v is not None;cell.remove(v)
    files['xl/worksheets/sheet1.xml']=ET.tostring(data)
    with ZipFile(p,'w',ZIP_DEFLATED) as z:
        for n,v in files.items():z.writestr(n,v)
    b=read_xlsx(p);assert isinstance(b['sheets'][0]['rows'][1][3],dict)
    e=build_exam(b,sample[1]);assert e['students'][0]['total'] is None

def test_old_xls_rejected(tmp_path):
    with pytest.raises(ExamError,match='.xlsx'):read_xlsx(tmp_path/'x.xls')

def test_model_payload_excludes_names_classes_filename(sample):
    e=sample[2];p=model_payload(analyse(e),{},'',True);s=json.dumps(p,ensure_ascii=False)
    for term in ['示例01','合成班A','example.xlsx','local_label','excel_row']:assert term not in s
    assert p['students'][0]['id']=='S0001'
    assert model_payload(analyse(e),{},'',False)['students']==[]

def test_confirmation_before_secret_and_send(sample):
    store=FixtureStore();t=FixtureTransport();p=model_payload(analyse(sample[2]),{},'')
    with pytest.raises(ExamError):generate_advice(SimpleNamespace(_providers=store),'p','fixture-revision',p,[],confirmed=False,transport=t)
    assert store.borrowed==0 and not t.requests

def test_one_call_structured_result_does_not_modify_statistics(sample):
    report=analyse(sample[2]);before=deepcopy(report);payload=model_payload(report,{},'');t=FixtureTransport();s=FixtureStore()
    result=generate_advice(SimpleNamespace(_providers=s),'p','fixture-revision',payload,[],confirmed=True,transport=t)
    assert result['candidate']==candidate() and len(t.requests)==1 and s.borrowed==1
    assert report==before and result['usage']['total_tokens']==200
    assert b'synthetic-fixture-not-a-real-key' not in t.requests[0].body

def test_text_only_api_warns_before_sending_images(sample):
    p=model_payload(analyse(sample[2]),{},'');s=FixtureStore();t=FixtureTransport()
    with pytest.raises(ExamError,match='读图能力'):generate_advice(SimpleNamespace(_providers=s),'p','fixture-revision',p,[{'width':1,'height':1}],confirmed=True,transport=t)
    assert not t.requests and not s.borrowed

def test_bad_model_references_rejected(sample):
    p=model_payload(analyse(sample[2]),{},'');c=candidate();c['class_actions'][0]['question_ids']=['99']
    with pytest.raises(ExamError,match='题号'):validate_candidate(c,p)
    c=candidate();c['students']=[{'student_id':'S0001','question_ids':[],'action':'a','check':'b'}]
    with pytest.raises(ExamError,match='学生'):validate_candidate(c,p)

def test_model_error_does_not_leak_transport_content(sample):
    p=model_payload(analyse(sample[2]),{},'');s=FixtureStore()
    class Bad:
        def send(self,*a,**k):raise RuntimeError('a-private-token-and-path')
    with pytest.raises(ExamError) as caught:generate_advice(SimpleNamespace(_providers=s),'p','fixture-revision',p,[],confirmed=True,transport=Bad())
    assert 'private' not in str(caught.value)

def test_report_escapes_html_and_separates_names(sample):
    r=analyse(sample[2],'合成班A');r['title']='<script>bad()</script>'
    output=html_report(r)
    assert '<script>' not in output and '&lt;script&gt;' in output
    assert '示例01' not in output and '合成班A' not in output
    assert '示例01' in html_report(r,include_names=True)

def test_local_attention_distinguishes_absence_missing_and_low_score(sample):
    r=analyse(sample[2]);assert '缺考' in local_student_advice(r['students'][4],r)
    assert '缺失' in local_student_advice(r['students'][5],r)
    assert '失' in local_student_advice(r['students'][1],r)

def test_saved_exam_reopens_without_source_mutation(sample,tmp_path):
    e=sample[2];store=ExamStore(tmp_path/'state');b={'exam':e,'paper':{'text':'题目与完整公共材料'},'advice':None}
    original=(tmp_path/'example.xlsx').read_bytes();identity=store.save(b)
    assert store.load(identity)==b and store.entries()[0][1]==e['title']
    assert (tmp_path/'example.xlsx').read_bytes()==original
    with pytest.raises(ExamError):store.load('../escape')

def test_docx_paper_keeps_native_text_table_and_image(tmp_path):
    from docx import Document
    from PIL import Image
    d=Document();d.add_paragraph('共同材料：保持实验条件。');d.add_paragraph('第1题 化学式 H₂SO₄')
    t=d.add_table(rows=1,cols=2);t.cell(0,0).text='条件';t.cell(0,1).text='12.5'
    buf=BytesIO();Image.new('RGB',(180,120),'white').save(buf,format='PNG');d.add_picture(BytesIO(buf.getvalue()))
    path=tmp_path/'试卷.docx';d.save(path);before=path.read_bytes();paper=load_paper(path)
    assert '共同材料' in paper['text'] and '12.5' in paper['text'] and len(paper['pages'])==1
    assert path.read_bytes()==before


def test_declared_visual_api_sends_exact_image_with_fixed_statistics(sample,tmp_path):
    from PIL import Image
    from types import SimpleNamespace
    from integrations.deeptutor_shchem_v1.desktop_exam_fixture import FixtureStore,FixtureTransport
    from integrations.deeptutor_shchem_v1.desktop_exam_ai import load_paper,model_payload,generate_advice
    from integrations.deeptutor_shchem_v1.desktop_exam_data import analyse
    import json,base64
    _,_,exam=sample
    path=tmp_path/'paper.png';Image.new('RGB',(200,300),'white').save(path)
    paper=load_paper(path);payload=model_payload(analyse(exam),paper,'')
    store=FixtureStore(vision=True);transport=FixtureTransport()
    result=generate_advice(SimpleNamespace(_providers=store),'p','fixture-revision',payload,paper['pages'],confirmed=True,transport=transport)
    assert store.borrowed==1 and len(transport.requests)==1
    body=json.loads(transport.requests[0].body)
    images=[b['image_url']['url'] for m in body['messages'] if isinstance(m['content'],list) for b in m['content'] if b.get('type')=='image_url']
    assert len(images)==1 and images[0].split(',')[1]==paper['pages'][0]['data']
    assert result['candidate']['students']==[] and analyse(exam)['overall']['mean']==69
