"""Source-bound tables: complete cells, readable panels and editable Word."""
from copy import deepcopy
from io import BytesIO
import pytest
from docx import Document
from integrations.deeptutor_shchem_v1.desktop_handout_knowledge_tables import (
    source_table, table_chunks, write_word_table, classroom_cell_text)
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import _readable_table, _table_rows
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import block, resource, digest


def row(*values, before=0, after=0):
    return {'grid_before': before, 'grid_after': after, 'cells': [
        {'column_span': value[0], 'vertical_merge': value[2] if len(value) > 2 else 'none', 'text': value[1]}
        if isinstance(value, tuple) else {'column_span': 1, 'vertical_merge': 'none', 'text': value}
        for value in values]}


def comparison():
    # PKG-064 native block 94: exact 6 x 3 source cell shape and text.
    return [row('', '范德华力', '氢键'),
        row('存在', '分子间普遍存在', '已经与电负性很大的原子形成共价键的氢原子与另一个电负性很大的原子之间'),
        row('特征', '无方向性、无饱和性', '具有一定的方向性和饱和性'),
        row('强度', (2, '共价键＞氢键＞范德华力')),
        row('影响其强度的因素', '①组成和结构相似的物质，相对分子质量越大，范德华力越大；\n②分子的极性越大，范德华力越大',
            '对于A—H…B，A、B的电负性越大，B原子的半径越小，氢键键能越大'),
        row('对物质性质的影响', '范德华力主要影响物质的物理性质，如熔、沸点等。范德华力越大，物质的熔、沸点越高',
            '分子间氢键的存在，使物质的熔、沸点升高，分子内氢键使物质的熔、沸点降低，对电离、溶解度等产生影响')]


def experiments():
    # PKG-064 native block 97: full experiment-to-conclusion grid, no transpose.
    return [row('实验', '①', '②', '③', '④', '⑤', '⑥'),
        row('溶质', '蔗糖', 'HCl', 'CH_{4}', 'I_{2}', 'I_{2}', 'P_{4}'),
        row('溶质分子极性', '极性', '极性', '非极性', '非极性', '非极性', '非极性'),
        row('溶剂', '水', '水', '水', '水', 'CCl_{4}', 'CS_{2}'),
        row('溶剂分子极性', '极性', '极性', '极性', '极性', '非极性', '非极性'),
        row('溶解性', '易溶', '易溶', '难溶', '难溶', '易溶', '易溶'),
        row('结论', (6, '非极性溶质一般能溶于非极性溶剂，而极性溶质一般能溶于极性溶剂。'))]


def item_for(rows):
    from integrations.deeptutor_shchem_v1.desktop_handout_knowledge import SOURCE_SCHEMA, SECTION_SCHEMA
    source = dict(schema=SOURCE_SCHEMA, batch_id='synthetic', source_id='source', source_name='合成来源.docx',
        source_sha256='a' * 64, source_role='handout', source_path='合成来源.docx', source_revision='b' * 64)
    snapshot = dict(schema=SECTION_SCHEMA, outline_revision='c' * 64,
        section=dict(id='word-heading-b94', level=3, title='完整比较', start=94, end=94),
        blocks=[dict(index=94, text=_readable_table(rows), images=[], table_rows=deepcopy(rows))])
    return resource('handout_knowledge', '合成来源.docx · 完整比较',
        [block('text', 'knowledge', _readable_table(rows), 94)],
        dict(source_path=source['source_path'], source_sha256=source['source_sha256'], source_revision=source['source_revision'],
            handout_source=source, section_snapshot=snapshot, image_assets=[],
            source_roles=dict(teacher_only=False, document_role='handout', images=[])))


def fake_lines(text, width, *, block_id, font_px, bold=False):
    return [dict(kind='text', text=text, width=width, height=10, block_id=block_id, font_px=font_px, bold=bold)]


def test_projection_requires_exact_unedited_work_and_source_text():
    item = item_for(comparison()); original = deepcopy(item)
    result = source_table(item, item['blocks'][0]); assert result == comparison()
    result[0]['cells'][0]['text'] = '不能影响来源'
    assert item == original
    item['blocks'][0]['value'] += '\n教师补充'
    assert source_table(item, item['blocks'][0]) is None
    assert source_table(original, {**original['blocks'][0], 'source_index': 95}) is None
    assert source_table(original, {**original['blocks'][0], 'presentation': {}}) is None
    item = deepcopy(original); item['original']['evidence']['section_snapshot']['blocks'][0]['text'] += '不一致'
    with pytest.raises(ValueError, match='封存原文'): source_table(item, item['blocks'][0])


def test_other_resource_kinds_and_table_image_placeholders_use_existing_path():
    item = item_for([row('【待查看原文：图片】', '完整文字')])
    assert source_table(item, item['blocks'][0]) is None
    item['original']['kind'] = 'personal_text'
    assert source_table(item, item['blocks'][0]) is None


def test_panels_repeat_first_source_row_without_losing_or_changing_cells():
    rows = comparison(); sealed = deepcopy(rows)
    chunks = table_chunks(rows, 'b', 300, fake_lines, 32, max_height=50)
    assert len(chunks) == 3
    assert [c['_table_source_rows'] for c in chunks] == [[1, 2, 3], [1, 4, 5], [1, 6]]
    assert [c['_table_repeated_first_row'] for c in chunks] == [False, True, True]
    assert rows == sealed
    for chunk in chunks:
        assert chunk['height'] <= 50
        assert all(p['offset_x'] + p['width'] <= 300 + .001 for p in chunk['_parts'])
        assert all(p['offset_y'] + p['height'] <= chunk['height'] + .001 for p in chunk['_parts'])
        for part in chunk['_parts']:
            if part['kind'] == 'text':
                ri, ci, span = part['_table_cell']
                source = next(c for c in rows[ri - 1]['cells']
                    if c['column_span'] == span and c['text'] == part['text'])
                assert source['text'] == part['text'] and part['font_px'] == 32
                assert '表格开始' not in part['text'] and '第1行·' not in part['text']
    strength = next(p for c in chunks for p in c['_parts'] if p['kind'] == 'diagram_box' and p['_table_cell'] == (4, 2, 2))
    assert strength['width'] == 200 and strength['offset_x'] == 100


def test_vertical_merge_chain_is_indivisible_and_all_occurrence_text_remains():
    rows = [row('列甲', '列乙'), row((1, '起点', 'restart'), '第二行'),
        row((1, '续接原文字', 'continue'), '第三行'), row('第四行', '尾部')]
    chunks = table_chunks(rows, 'b', 200, fake_lines, 32, max_height=50)
    assert [c['_table_source_rows'] for c in chunks] == [[1, 2, 3], [1, 4]]
    merged = next(p for p in chunks[0]['_parts'] if p['kind'] == 'diagram_box' and p['_table_cell'] == (2, 1, 1))
    single = next(p for p in chunks[0]['_parts'] if p['kind'] == 'diagram_box' and p['_table_cell'] == (2, 2, 1))
    assert merged['height'] == pytest.approx(2 * single['height'])
    assert not any(p['kind'] == 'diagram_box' and p['_table_cell'] == (3, 1, 1) for p in chunks[0]['_parts'])
    assert any(p.get('text') == '续接原文字' for p in chunks[0]['_parts'])


@pytest.mark.parametrize('rows,error', [
    ([row('甲', '乙'), row('丙')], '网格宽度'),
    ([row((1, '无起点', 'continue'))], '对应起点'),
    ([row('首行'), row('正文')], '完整表格行'),
])
def test_invalid_grid_and_too_small_canvas_fail_without_dropping_content(rows, error):
    with pytest.raises(ValueError, match=error): table_chunks(rows, 'b', 200, fake_lines, 32, max_height=20)


@pytest.mark.parametrize('rows', [comparison(), experiments(),
    [row('标题甲', '标题乙', '标题丙'), row((2, '横纵合并起点', 'restart'), '右侧'),
        row((2, '续接原文字', 'continue'), '尾部'), row('中列', before=1, after=1)]])
def test_word_roundtrip_preserves_exact_source_grid_merge_and_every_cell(rows):
    doc = Document(); table = write_word_table(doc, rows)
    assert _table_rows(table._tbl, doc, set()) == rows
    assert table.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    assert not table._tbl.xpath('.//w:rPr/w:rFonts')
    buffer = BytesIO(); doc.save(buffer); buffer.seek(0); reopened = Document(buffer)
    assert _table_rows(reopened.tables[0]._tbl, reopened, set()) == rows
    xml = table._tbl.xml
    assert '表格开始' not in xml and '第1行·' not in xml


def test_omitted_edge_cells_keep_their_original_grid_positions():
    rows = [row('甲', '乙', '丙'), row('中央', before=1, after=1)]
    chunks = table_chunks(rows, 'b', 300, fake_lines, 32)
    cell = next(p for p in chunks[0]['_parts'] if p['kind'] == 'diagram_box' and p['_table_cell'] == (2, 2, 1))
    assert cell['offset_x'] == 100 and cell['width'] == 100


@pytest.fixture(scope='session')
def app():
    pytest.importorskip('PySide6')
    from PySide6.QtWidgets import QApplication
    from integrations.deeptutor_shchem_v1.desktop_workbench.typography import install_ui_font
    application = QApplication.instance() or QApplication([]); install_ui_font(application)
    return application


@pytest.mark.parametrize('rows', [comparison(), experiments()])
def test_actual_tables_fit_measured_24pt_panels_and_retain_all_cell_text(app, rows):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import _lines
    chunks = table_chunks(rows, 'b', 300, _lines, 32, max_height=136)
    assert 1 <= len(chunks) <= 3
    assert [ri for c in chunks for ri in c['_table_source_rows'] if ri != 1] == list(range(2, len(rows) + 1))
    covered = {}
    for c in chunks:
        assert c['height'] <= 136 + .001
        for p in c['_parts']:
            assert p['offset_x'] >= 0 and p['offset_y'] >= 0
            assert p['offset_x'] + p['width'] <= 300 + .001
            assert p['offset_y'] + p['height'] <= c['height'] + .001
            if p['kind'] == 'text':
                assert p['font_px'] == 32
                if p['_table_cell'][0] != 1: covered.setdefault(p['_table_cell'], []).append(p['text'])
    for ri, r in enumerate(rows[1:], 2):
        ci = r['grid_before'] + 1
        for cell in r['cells']:
            assert ''.join(covered.get((ri, ci, cell['column_span']), [])) == classroom_cell_text(cell['text']).replace('\n', '')
            ci += cell['column_span']


@pytest.mark.parametrize('first_height,expected', [(28, [[1, 2], [1, 3, 4, 5, 6]]),
    (0, [[1, 2, 3, 4, 5, 6]]), (60, [[1, 2, 3, 4], [1, 5, 6]])])
def test_first_panel_uses_remaining_budget_and_never_contains_only_header(first_height, expected):
    chunks = table_chunks(comparison(), 'b', 300, fake_lines, 32, max_height=100, first_height=first_height)
    assert [c['_table_source_rows'] for c in chunks] == expected
    assert len(chunks[0]['_table_source_rows']) >= 2


def test_first_budget_does_not_split_a_vertical_merge_group():
    rows = [row('甲', '乙'), row((1, '起点', 'restart'), '数据'),
        row((1, '续接', 'continue'), '下行'), row('最后', '尾部')]
    chunks = table_chunks(rows, 'b', 200, fake_lines, 32, max_height=100, first_height=45)
    assert chunks[0]['_table_source_rows'] == [1, 2, 3]
    assert chunks[1]['_table_source_rows'] == [1, 4]
    first_merged = [row((1, '起点', 'restart'), '甲'), row((1, '续接', 'continue'), '乙')]
    assert len(table_chunks(first_merged, 'b', 200, fake_lines, 32, max_height=100, first_height=0)) == 1
    restarted = table_chunks(rows, 'b', 200, fake_lines, 32, max_height=100, first_height=28)
    assert restarted[0]['_table_restarted_page'] and restarted[0]['_table_source_rows'] == [1, 2, 3, 4]


def test_numeric_subscripts_have_same_value_in_classroom_and_native_word():
    raw = 'CH_{4}，CCl_{4}，CS_{2}，I_{2}，P_{4}；x_{12}；x_{n}、x^{2}、_{a}'
    assert classroom_cell_text(raw) == 'CH₄，CCl₄，CS₂，I₂，P₄；x₁₂；x_{n}、x²、ₐ'
    rows = [row('化学式', raw)]; doc = Document(); table = write_word_table(doc, rows)
    assert _table_rows(table._tbl, doc, set()) == rows
    assert len(table._tbl.xpath('.//w:rPr/w:vertAlign[@w:val="subscript"]')) == 7
    assert len(table._tbl.xpath('.//w:rPr/w:vertAlign[@w:val="superscript"]')) == 1
    assert any(r.text == '12' and r.font.subscript for r in table.cell(0, 1).paragraphs[0].runs)


@pytest.mark.parametrize('raw,display', [
    ('n(H^{+})、c(H^{＋})、K_{a}', 'n(H⁺)、c(H⁺)、Kₐ'),
    ('SO_{4}^{2−}、Cl^{-}、OH^{－}', 'SO₄²⁻、Cl⁻、OH⁻'),
    ('x^{0123456789}、y_{0123456789}', 'x⁰¹²³⁴⁵⁶⁷⁸⁹、y₀₁₂₃₄₅₆₇₈₉'),
    ('H_{2}O\nK_{a}，H^{+}', 'H₂O\nKₐ，H⁺'),
])
def test_explicit_script_payloads_have_unicode_and_lossless_word_projections(raw, display):
    assert classroom_cell_text(raw) == display
    rows = [row('原单元格', raw)]; sealed = deepcopy(rows)
    doc = Document(); table = write_word_table(doc, rows)
    assert rows == sealed and _table_rows(table._tbl, doc, set()) == sealed
    data = BytesIO(); doc.save(data); data.seek(0); reopened = Document(data)
    assert _table_rows(reopened.tables[0]._tbl, reopened, set()) == sealed
    runs = table.cell(0, 1).paragraphs[0].runs
    assert any(r.font.subscript or r.font.superscript for r in runs)
    assert not table._tbl.xpath('.//w:rPr/w:rFonts')


@pytest.mark.parametrize('raw', [
    'H2O、H+、Ka、x^2、x_2', 'x_{n}、K_{b}、x^{a}',
    'x_{a+n}、x^{2q}、x^{+ }', 'x_{}、x^{}',
    'x_{2、x^{+}尾部', 'x_{2^{+}}、x^{q_{2}}',
    r'x\_{2}、x\^{+}', 'x^{2\n+}、x_{a\n2}',
])
def test_unknown_malformed_or_unmarked_content_is_not_inferred_or_deleted(raw):
    assert classroom_cell_text(raw) == raw
    rows = [row(raw)]; doc = Document(); table = write_word_table(doc, rows)
    assert table.cell(0, 0).text == raw
    assert _table_rows(table._tbl, doc, set()) == rows
    assert not table._tbl.xpath('.//w:rPr/w:vertAlign')


def weak_acid_conditions():
    # PKG-056 native b49: retain the full-width source plus signs, exact 7 x 6
    # source grid and every condition. Tests verify display, not chemistry.
    return [row('改变条件', '平衡移动方向', 'n(H^{＋})', 'c(H^{＋})', '导电能力', 'K_{a}'),
        row('加水稀释', '向右', '增大', '减小', '减弱', '不变'),
        row('加入少量冰醋酸', '向右', '增大', '增大', '增强', '不变'),
        row('通入HCl(g)', '向左', '增大', '增大', '增强', '不变'),
        row('加NaOH(s)', '向右', '减小', '减小', '增强', '不变'),
        row('加CH_{3}COONa(s)', '向左', '减小', '减小', '增强', '不变'),
        row('升高温度', '向右', '增大', '增大', '增强', '增大')]


def test_native_b49_table_exports_correct_scripts_without_changing_work_or_evidence(app, tmp_path):
    from pptx import Presentation
    from test_lesson_design import sample
    from integrations.deeptutor_shchem_v1.desktop_composition_render import (
        arrange_composition, layout_composition, write_pptx, write_pdf)
    from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
        new_composition, add_resource, enable_composition)
    from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents
    rows = weak_acid_conditions(); item = item_for(rows); sealed = deepcopy(item)
    original_hash = item['original_sha256']
    value = new_composition('16:9')
    value = add_resource(value, item, dict(page_id=value['pages'][0], x=12, y=12, width=296, height=156))
    value = arrange_composition(value, [], {})
    for student in (False, True):
        layout = layout_composition(value, [], {}, student=student)
        assert not layout['errors']
        parts = [p for elements in layout['pages'].values() for c in elements
            if '_table_source_rows' in c for p in c['_parts'] if p['kind'] == 'text']
        for header in ('n(H⁺)', 'c(H⁺)', 'Kₐ'):
            assert any(p['text'] == header for p in parts)
        assert ''.join(p['text'] for p in parts if p['_table_cell'] == (6, 1, 1)) == '加CH₃COONa(s)'
        assert all(p['font_px'] == 32 for p in parts)
        path = tmp_path / ('student.pptx' if student else 'teacher.pptx')
        write_pptx(layout, path, {})
        deck = Presentation(path)
        display = ''.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
        assert all(text in display for text in ('n(H⁺)', 'c(H⁺)', 'Kₐ', '加CH₃COONa(s)'))
        assert '^{＋}' not in display and '_{a}' not in display and '_{3}' not in display
        pdf = tmp_path / ('student.pdf' if student else 'teacher.pdf')
        write_pdf(layout, pdf)
        assert pdf.read_bytes().startswith(b'%PDF')
    payload = sample(); payload['lesson_design'] = enable_composition(payload['lesson_design'])
    payload['lesson_design']['composition'] = value; payload['image_assets'] = []
    before = deepcopy(payload)
    write_documents(payload, tmp_path, {})
    for name in ('lesson_plan.docx', 'student_worksheet.docx'):
        document = Document(tmp_path / name)
        table = next(t for t in document.tables if len(t.rows) == 7 and len(t.columns) == 6)
        assert _table_rows(table._tbl, document, set()) == rows
        assert table.cell(0, 2).text == 'n(H＋)' and table.cell(0, 5).text == 'Ka'
        assert len(table._tbl.xpath('.//w:rPr/w:vertAlign[@w:val="superscript"]')) == 2
        assert len(table._tbl.xpath('.//w:rPr/w:vertAlign[@w:val="subscript"]')) == 2
    assert payload == before and item == sealed
    assert value['resources'][0] == sealed and value['resources'][0]['original_sha256'] == original_hash
    assert source_table(item, item['blocks'][0]) == rows


def introduction_item(rows, intro_lines=3):
    item = item_for(rows)
    snapshot = item['original']['evidence']['section_snapshot']
    snapshot['section'].update(id='word-heading-b92', start=92)
    intro = [block('text', 'knowledge', '溶质和溶剂的极性怎样影响溶解性？', 92),
        block('text', 'knowledge', '\n'.join('观察实验数据，比较物质的极性与溶解性。' for _ in range(intro_lines - 1)), 93)]
    snapshot['blocks'] = [dict(index=b['source_index'], text=b['value'], images=[]) for b in intro] + snapshot['blocks']
    item['blocks'] = deepcopy(intro) + item['blocks']; item['original']['blocks'] = deepcopy(item['blocks'])
    item['original_sha256'] = digest(item['original'])
    item['title'] = '哪些物质更容易溶解在水中？'
    return item


def test_measured_classroom_introduction_and_complete_table_rows_share_first_page(app):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition, layout_composition
    from integrations.deeptutor_shchem_v1.desktop_lesson_composition import new_composition, add_resource
    item = introduction_item(experiments()); sealed = deepcopy(item)
    value = new_composition(format='16:9')
    value = add_resource(value, item, dict(page_id=value['pages'][0], x=12, y=12, width=296, height=156))
    value = arrange_composition(value, [], {})
    output = layout_composition(value, [], {}, student=True)
    assert not output['errors']; assert item == sealed
    first = next(iter(output['pages'].values()))
    table = next(c for c in first if '_table_source_rows' in c)
    assert table['_table_source_rows'][:2] == [1, 2]
    assert any(c.get('text', '').startswith('溶质和溶剂') for c in first)
    data_rows = [ri for elements in output['pages'].values() for c in elements
        if '_table_source_rows' in c for ri in c['_table_source_rows'] if ri != 1]
    assert data_rows == list(range(2, 8))
    assert all(p['font_px'] == 32 for elements in output['pages'].values() for c in elements
        if '_table_source_rows' in c for p in c['_parts'] if p['kind'] == 'text')


def test_short_introduction_advances_with_first_table_row_from_partial_page(app):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition, layout_composition
    from integrations.deeptutor_shchem_v1.desktop_lesson_composition import new_composition, add_resource
    previous = resource('personal_text', '课前观察', [block('text', 'personal', '\n'.join('先观察这个现象。' for _ in range(7)))],
        dict(source_path='作者输入', source_sha256='unknown', source_revision='unknown'))
    value = new_composition(format='16:9')
    for item in (previous, introduction_item(experiments())):
        value = add_resource(value, item, dict(page_id=value['pages'][0], x=12, y=12, width=296, height=156))
    value = arrange_composition(value, [], {})
    output = layout_composition(value, [], {}, student=True); assert not output['errors']
    for elements in output['pages'].values():
        if any(c.get('text', '').startswith('溶质和溶剂') for c in elements):
            assert any(c.get('_table_source_rows', [])[:2] == [1, 2] for c in elements)


def test_long_preceding_text_does_not_carry_a_minimum_table_panel_onto_fresh_page(app):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition, layout_composition
    from integrations.deeptutor_shchem_v1.desktop_lesson_composition import new_composition, add_resource
    item = introduction_item(comparison(), intro_lines=10); sealed = deepcopy(item)
    value = new_composition('16:9')
    value = add_resource(value, item, dict(page_id=value['pages'][0], x=12, y=12, width=296, height=156))
    output = layout_composition(arrange_composition(value, [], {}), [], {}, student=True)
    assert not output['errors']; assert item == sealed
    panels = [c for elements in output['pages'].values() for c in elements if '_table_source_rows' in c]
    assert panels[0]['_table_restarted_page']
    assert panels[0]['_table_source_rows'] == [1, 2, 3, 4]
    assert [ri for c in panels for ri in c['_table_source_rows'] if ri != 1] == list(range(2, 7))
    actual = {}
    for c in panels:
        for part in c['_parts']:
            if part['kind'] == 'text' and part['_table_cell'][0] != 1:
                actual.setdefault(part['_table_cell'], []).append(part['text'])
    for ri, source in enumerate(comparison()[1:], 2):
        ci = 1
        for cell in source['cells']:
            assert ''.join(actual[(ri, ci, cell['column_span'])]) == cell['text'].replace('\n', '')
            ci += cell['column_span']


def test_native_rich_numbered_heading_stays_with_following_paragraph(app):
    from integrations.deeptutor_shchem_v1.desktop_composition_inline import NativeParagraphDisplay
    from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition, layout_composition
    from integrations.deeptutor_shchem_v1.desktop_lesson_composition import new_composition, add_resource
    # The exact numbered heading used by the native PKG-064 knowledge paragraph.
    texts = ['\n'.join('先观察分子间作用力的不同表现。' for _ in range(10)),
        '3．影响因素：', '(1)对于组成和结构相似的物质，相对分子质量越大，范德华力越大。']
    document = Document()
    for text in texts: document.add_paragraph(text)
    buffer = BytesIO(); document.save(buffer); display = NativeParagraphDisplay(buffer.getvalue())
    blocks = []
    for index, text in enumerate(texts, 1):
        b = block('text', 'personal', text, index); display.bind(b, [], []); blocks.append(b)
    item = resource('personal_text', '范德华力的大小与什么有关？', blocks,
        dict(source_path='合成原Word.docx', source_sha256='unknown', source_revision='unknown'))
    value = new_composition('16:9')
    value = add_resource(value, item, dict(page_id=value['pages'][0], x=12, y=12, width=296, height=156))
    output = layout_composition(arrange_composition(value, [], {}), [], {}, student=True)
    assert not output['errors']
    heading_pages = [elements for elements in output['pages'].values()
        if any(c.get('block_id') == blocks[1]['id'] and c['kind'] == 'rich_line' for c in elements)]
    assert len(heading_pages) == 1
    assert any(c.get('block_id') == blocks[2]['id'] and c['kind'] == 'rich_line' for c in heading_pages[0])


def test_inline_picture_heading_is_not_inferred_from_only_its_text_parts(app):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import _keep_groups
    chunks = [dict(kind='text', text='开场', height=5, block_id=None),
        dict(kind='rich_line', height=30, block_id='heading', _parts=[
            dict(kind='text', text='3．影响因素：'), dict(kind='image', height=30)]),
        dict(kind='text', text='下一段', height=30, block_id='body')]
    _keep_groups(chunks, 50)
    assert not chunks[1]['_keep_next']


def test_weighted_column_positions_accumulate_original_spans_and_omitted_edges():
    rows = [row('甲', '乙', '丙'), row('标签', (2, '横向合并')),
        row((2, '横纵起点', 'restart'), '右侧'), row((2, '续接', 'continue'), '尾部'),
        row('中列', before=1, after=1)]
    chunks = table_chunks(rows, 'b', 300, fake_lines, 32, max_height=100, column_weights=(1, 2, 2))
    boxes = {p['_table_cell']: p for c in chunks for p in c['_parts'] if p['kind'] == 'diagram_box'}
    assert boxes[(1, 1, 1)]['width'] == 60
    assert boxes[(1, 2, 1)]['offset_x'] == 60 and boxes[(1, 2, 1)]['width'] == 120
    assert boxes[(1, 3, 1)]['offset_x'] == 180 and boxes[(1, 3, 1)]['width'] == 120
    assert boxes[(2, 2, 2)]['offset_x'] == 60 and boxes[(2, 2, 2)]['width'] == 240
    assert boxes[(3, 1, 2)]['offset_x'] == 0 and boxes[(3, 1, 2)]['width'] == 180
    assert (4, 1, 2) not in boxes
    assert boxes[(5, 2, 1)]['offset_x'] == 60 and boxes[(5, 2, 1)]['width'] == 120


@pytest.mark.parametrize('weights', [(1, 2), (0, 2, 2), (1, float('inf'), 2)])
def test_weighted_columns_require_a_complete_positive_finite_grid(weights):
    with pytest.raises(ValueError, match='显示列宽'):
        table_chunks(comparison(), 'b', 300, fake_lines, 32, column_weights=weights)


def test_real_comparison_table_fits_two_weighted_pages_without_losing_cells(app):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import _lines
    rows = comparison(); sealed = deepcopy(rows)
    chunks = table_chunks(rows, 'b', 290, _lines, 32, max_height=134, first_height=0, column_weights=(1, 2, 2))
    assert [c['_table_source_rows'] for c in chunks] == [[1, 2, 3, 4], [1, 5, 6]]
    assert all(c['height'] <= 134 + .001 for c in chunks)
    assert rows == sealed
    actual = {}
    for c in chunks:
        for part in c['_parts']:
            assert part['offset_x'] + part['width'] <= 290 + .001
            if part['kind'] == 'text' and part['_table_cell'][0] != 1:
                assert part['font_px'] == 32
                actual.setdefault(part['_table_cell'], []).append(part['text'])
    for ri, source in enumerate(rows[1:], 2):
        ci = 1
        for cell in source['cells']:
            assert ''.join(actual[(ri, ci, cell['column_span'])]) == cell['text'].replace('\n', '')
            ci += cell['column_span']


def illustrated_table_item(tmp_path, earlier_lines=7, introduction='short'):
    from PIL import Image
    from docx.shared import Mm
    from test_handout_knowledge import make_facade, request_for, word_bytes
    from integrations.deeptutor_shchem_v1.desktop_handout_knowledge import import_handout_knowledge
    from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
    def picture(size, colour):
        buffer = BytesIO(); Image.new('RGB', size, colour).save(buffer, 'PNG'); return buffer.getvalue()
    large = picture((160, 80), '#3276ad'); arrow = picture((80, 10), '#273749')
    document = Document(); document.add_heading('完整知识段', 3)
    document.add_paragraph().add_run().add_picture(BytesIO(large), width=Mm(65))
    document.add_paragraph('\n'.join('早前的完整说明仍按原文顺序讲解。' for _ in range(earlier_lines)))
    p = document.add_paragraph()
    if introduction == 'picture':
        p.add_run().add_picture(BytesIO(large), width=Mm(65))
    else:
        p.add_run('实例\n' if introduction == 'short' else '\n'.join('这一整段长说明应按完整原行继续。' for _ in range(12)) + '\n')
        p.add_run('前项'); p.add_run().add_picture(BytesIO(arrow), width=Mm(8)); p.add_run('后项：比较完整原表。')
    write_word_table(document, weak_acid_conditions())
    facade = make_facade(word_bytes(document), tmp_path)
    result = import_handout_knowledge(facade, [request_for(facade)], [], [])
    store = PreparationImageStore(facade.paths.task_root / 'preparation-v1/images')
    data = {asset['asset_id']: store.load(asset) for asset in result['image_assets']}
    return result['resources'][0], result['image_assets'], data, (large, arrow)


@pytest.mark.parametrize('earlier_lines', [3, 7, 11])
def test_earlier_original_picture_does_not_orphan_complete_short_inline_table_introduction(app, tmp_path, earlier_lines):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition, layout_composition
    from integrations.deeptutor_shchem_v1.desktop_lesson_composition import new_composition, add_resource
    item, assets, data, originals = illustrated_table_item(tmp_path / 'source', earlier_lines)
    sealed = deepcopy(item); intro = next(b for b in item['blocks'] if b['kind'] == 'text' and b['source_index'] == 4)
    table_block = next(b for b in item['blocks'] if b['kind'] == 'text' and b['source_index'] == 5)
    value = new_composition('16:9')
    value = add_resource(value, item, dict(page_id=value['pages'][0], x=12, y=12, width=296, height=156))
    arranged = arrange_composition(value, assets, data)
    for student in (False, True):
        output = layout_composition(arranged, assets, data, student=student); assert not output['errors']
        intro_pages = [(page, elements) for page, elements in output['pages'].items()
            if any(c.get('block_id') == intro['id'] for c in elements)]
        assert len(intro_pages) == 1
        _page, elements = intro_pages[0]
        assert any(c.get('block_id') == table_block['id'] and c.get('_table_source_rows', [])[:2] == [1, 2] for c in elements)
        intro_lines = [c for c in elements if c.get('block_id') == intro['id'] and c['kind'] == 'rich_line']
        assert 1 <= len(intro_lines) <= 3
        assert sum(p['kind'] == 'image' for c in intro_lines for p in c['_parts']) == 1
        assert all(p['font_px'] == 32 for c in intro_lines for p in c['_parts'] if p['kind'] == 'text')
        data_rows = [ri for page in output['pages'].values() for c in page
            if c.get('block_id') == table_block['id'] for ri in c.get('_table_source_rows', []) if ri != 1]
        assert data_rows == list(range(2, 8))
    assert item == sealed and arranged['resources'][0] == sealed
    assert set(data.values()) == set(originals)


@pytest.mark.parametrize('introduction', ['long', 'picture'])
def test_long_or_picture_only_previous_block_is_not_bound_to_sealed_table(app, tmp_path, introduction):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import _content, arrange_composition, layout_composition
    from integrations.deeptutor_shchem_v1.desktop_lesson_composition import new_composition, add_resource
    item, assets, data, _ = illustrated_table_item(tmp_path / 'source', introduction=introduction)
    sealed = deepcopy(item); assets_by_id = {a['asset_id']: a for a in assets}
    chunks = _content(item, 296, assets_by_id, data, student=True, format='16:9')
    table_start = next(i for i, c in enumerate(chunks) if c.get('_table_source_rows', [])[:2] == [1, 2])
    assert chunks[table_start - 1]['block_id'] == next(b['id'] for b in item['blocks'] if b['kind'] == 'text' and b['source_index'] == 4)
    assert not chunks[table_start - 1]['_keep_next']
    value = new_composition('16:9')
    value = add_resource(value, item, dict(page_id=value['pages'][0], x=12, y=12, width=296, height=156))
    arranged = arrange_composition(value, assets, data)
    assert not layout_composition(arranged, assets, data, student=True)['errors']
    assert arranged['resources'][0] == sealed and item == sealed


def test_a_previous_table_is_not_treated_as_short_prose_introduction():
    from integrations.deeptutor_shchem_v1.desktop_composition_render import _table_introduction_start
    chunks = [dict(kind='text', block_id=None, text='作品标题', height=10),
        dict(kind='diagram', block_id='old-table', height=30, _table_source_rows=[1, 2])]
    assert _table_introduction_start(chunks, 32) is None
