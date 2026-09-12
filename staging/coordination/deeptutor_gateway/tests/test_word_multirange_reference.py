"""Disjoint Word source references keep gaps, provenance, and commit checks."""

from __future__ import annotations

import io
from copy import deepcopy

import pytest
from docx import Document
from test_word_source_reference import _from_document, _image_root, _png
from test_word_source_reference import (
    desktop_paths as desktop_paths,  # noqa: PLC0414 - pytest fixture re-export
)

from integrations.deeptutor_shchem_v1.desktop_preparation_images import (
    PreparationImageStore,
)
from integrations.deeptutor_shchem_v1.desktop_word_ranges import (
    MAX_WORD_RANGES,
    normalize_word_ranges,
)
from integrations.deeptutor_shchem_v1.desktop_word_source_reference import (
    WordSourceReferenceError,
)


def _multi_reference(facade, selection, ranges, **kwargs):
    return facade.imported_word_multirange_reference(
        selection["batch_id"],
        selection["source_id"],
        selection["source_sha256"],
        ranges,
        selection["revision"],
        **kwargs,
    )


def _image_document():
    document = Document()
    document.add_paragraph("区块1：不应带入的正文")
    paragraph = document.add_paragraph("区块2：选中的原图")
    paragraph.add_run().add_picture(io.BytesIO(_png("blue")))
    document.add_paragraph("区块3：两个选段之间的原文间隙")
    paragraph = document.add_paragraph("区块4：重复使用同一张原图")
    paragraph.add_run().add_picture(io.BytesIO(_png("blue")))
    document.add_paragraph("区块5：不应带入的末尾")
    return document


def _write_flow_index(workspace, preview):
    folder = workspace / "knowledge" / "lectures"
    folder.mkdir(parents=True, exist_ok=True)
    card = {
        "id": "LECT-SYNTHETIC-MULTIRANGE",
        "source_sha256": preview["source_sha256"],
        "source_name": preview["source_name"],
        "source_preview_revision": preview["revision"],
        "title": "跨段讲练测试索引",
        "human_reviewed": False,
        "review_status": "ai_distilled_pending_teacher_review",
        "knowledge": [],
        "methods": [],
        "pitfalls": [],
        "teaching_flow": [
            {
                "summary": "跨段讲练建议：先讲区块2，再用区块4完成练习闭环。",
                "block_indices": [2, 4],
            }
        ],
        "textbook_links": [],
    }
    for filename in ("part-a.jsonl", "part-b.jsonl", "part-c.jsonl"):
        (folder / filename).write_text("", encoding="utf-8")
    path = folder / "part-a.jsonl"
    import json

    path.write_text(json.dumps(card, ensure_ascii=False) + "\n", encoding="utf-8")
    return card


def test_normalize_word_ranges_sorts_merges_overlap_and_preserves_gaps():
    assert normalize_word_ranges(
        [
            {"start": 8, "end": 9},
            {"start": 2, "end": 3},
            {"start": 3, "end": 5},
            {"start": 12, "end": 12},
            {"start": 9, "end": 10},
        ]
    ) == [
        {"start": 2, "end": 5},
        {"start": 8, "end": 10},
        {"start": 12, "end": 12},
    ]


@pytest.mark.parametrize(
    "value",
    [
        None,
        {},
        [],
        [{"start": 0, "end": 1}],
        [{"start": 1, "end": 0}],
        [{"start": True, "end": 1}],
        [{"start": 1, "end": 1.0}],
        [{"start": 1}],
        [{"start": 1, "end": 1, "extra": True}],
    ],
)
def test_normalize_word_ranges_rejects_bad_values_and_empty_ranges(value):
    with pytest.raises(ValueError):
        normalize_word_ranges(value)


def test_normalize_word_ranges_rejects_more_than_the_disjoint_segment_limit():
    value = [
        {"start": number, "end": number}
        for number in range(1, 2 * MAX_WORD_RANGES + 2, 2)
    ]
    with pytest.raises(ValueError, match=f"最多选择{MAX_WORD_RANGES}段"):
        normalize_word_ranges(value)


@pytest.mark.parametrize(
    "ranges",
    [
        [],
        [{"start": 0, "end": 1}],
        [{"start": 1, "end": 6}],
        [{"start": 6, "end": 6}],
        [{"start": 2, "end": 1}],
        [{"start": 1, "end": 1.5}],
    ],
)
def test_actual_facade_entry_rejects_bad_or_out_of_document_ranges_without_writes(
    desktop_paths, tmp_path, ranges
):
    facade, selection, _ = _from_document(desktop_paths, tmp_path, _image_document())
    with pytest.raises(WordSourceReferenceError):
        _multi_reference(facade, selection, ranges)
    assert not _image_root(facade).exists()


def test_actual_facade_entry_reads_only_selected_segments_and_deduplicates_cross_segment_pixels(
    desktop_paths, tmp_path
):
    facade, selection, preview = _from_document(
        desktop_paths, tmp_path, _image_document()
    )
    before = facade.state_store.snapshot()
    reference = _multi_reference(
        facade,
        selection,
        [
            {"start": 4, "end": 4},
            {"start": 2, "end": 2},
            {"start": 2, "end": 2},
        ],
    )

    assert reference["source_selection"]["block_ranges"] == [
        {"start": 2, "end": 2},
        {"start": 4, "end": 4},
    ]
    assert "区块2：选中的原图" in reference["materials"]
    assert "区块4：重复使用同一张原图" in reference["materials"]
    assert "区块1：不应带入" not in reference["materials"]
    assert "区块3：两个选段之间" not in reference["materials"]
    assert "区块5：不应带入" not in reference["materials"]
    assert len(reference["image_assets"]) == 1
    assert len(reference["image_references"]) == 2
    assert {item["block_index"] for item in reference["image_references"]} == {2, 4}
    assert reference["image_issues"] == []
    assert (
        reference["image_references"][0]["original_sha256"]
        == reference["image_references"][1]["original_sha256"]
    )
    assert (
        reference["image_references"][0]["asset_id"]
        == reference["image_references"][1]["asset_id"]
    )
    assert not _image_root(facade).exists()
    assert facade.state_store.snapshot() == before
    assert all(
        block["text"] not in reference["materials"]
        for block in preview["blocks"]
        if block["index"] in {1, 3, 5}
    )

    result = facade.import_word_source_reference(reference, [])
    store = PreparationImageStore(_image_root(facade))
    assert len(result["image_assets"]) == 1
    assert store.load(result["image_assets"][0]) == _png("blue")
    assert len(list(store.root.glob("*.image"))) == 1


def test_unselected_long_block_is_not_imported_but_selected_long_block_is_rejected(
    desktop_paths, tmp_path
):
    document = Document()
    document.add_paragraph("短选段")
    document.add_paragraph("长正文" * 8000)
    document.add_paragraph("最后一段")
    facade, selection, _ = _from_document(desktop_paths, tmp_path, document)

    short = _multi_reference(
        facade, selection, [{"start": 1, "end": 1}], include_images=False
    )
    assert "长正文" not in short["materials"]
    assert short["image_assets"] == []

    with pytest.raises(WordSourceReferenceError, match="超过备课.*上限"):
        _multi_reference(
            facade, selection, [{"start": 2, "end": 2}], include_images=False
        )
    assert not _image_root(facade).exists()


def test_teaching_flow_requires_complete_cross_segment_support_and_keeps_original_selection(
    desktop_paths, tmp_path
):
    facade, selection, preview = _from_document(
        desktop_paths, tmp_path, _image_document()
    )
    card = _write_flow_index(facade.paths.workspace_root, preview)
    reference = _multi_reference(
        facade,
        selection,
        [{"start": 4, "end": 4}, {"start": 2, "end": 2}],
        include_images=False,
        include_guidance=True,
    )

    study = reference["lecture_study"]
    assert study["note_count"] == 1
    assert study["textbook_concept_count"] == 0
    assert "跨段讲练建议：先讲区块2，再用区块4完成练习闭环。" in study["materials"]
    assert "LECT-SYNTHETIC-MULTIRANGE:teaching_flow:1" in study["materials"]
    assert reference["source_selection"]["block_ranges"] == [
        {"start": 2, "end": 2},
        {"start": 4, "end": 4},
    ]
    assert card["teaching_flow"][0]["block_indices"] == [2, 4]


def test_multirange_confirmation_rejects_tampering_before_writing_images(
    desktop_paths, tmp_path
):
    facade, selection, _ = _from_document(desktop_paths, tmp_path, _image_document())
    reference = _multi_reference(
        facade,
        selection,
        [{"start": 2, "end": 2}, {"start": 4, "end": 4}],
    )
    tampered = deepcopy(reference)
    tampered["source_selection"]["block_ranges"].append({"start": 5, "end": 5})
    with pytest.raises(WordSourceReferenceError, match="参考"):
        facade.import_word_source_reference(tampered, [])
    assert not _image_root(facade).exists()


def test_multirange_confirmation_rejects_changed_archive_before_writing_images(
    desktop_paths, tmp_path
):
    facade, selection, _ = _from_document(desktop_paths, tmp_path, _image_document())
    reference = _multi_reference(
        facade,
        selection,
        [{"start": 2, "end": 2}, {"start": 4, "end": 4}],
    )
    archive = (
        facade._visual_import_root / "sources" / f"{selection['source_sha256']}.docx"
    )
    archive.write_bytes(archive.read_bytes() + b"changed after preview")
    with pytest.raises(WordSourceReferenceError, match="变化"):
        facade.import_word_source_reference(reference, [])
    assert not _image_root(facade).exists()
