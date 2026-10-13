"""Document and picture handoffs retain their exact original archive identity."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from threading import Event
from types import SimpleNamespace

from docx import Document
import pytest

from test_handout_knowledge import make_facade, png, word_bytes
from test_handout_question_handoff import (scoped_lecture as scoped_lecture,
    source_kind_lecture as source_kind_lecture, desktop_paths as desktop_paths)
from test_desktop_visual_import_facade import FakeProviderStore, _facade
from integrations.deeptutor_shchem_v1.desktop_facade import DesktopFacadeError
from integrations.deeptutor_shchem_v1.desktop_handout_images import (
    SCHEMA as PICTURE_SCHEMA, handout_image_context, handout_image_preview,
    validate_source as validate_picture)
from integrations.deeptutor_shchem_v1.desktop_handout_knowledge import (
    handout_knowledge_sources, handout_knowledge_outline)
from integrations.deeptutor_shchem_v1.desktop_handout_sources import (
    handout_source_context, validate_handout_source)
from integrations.deeptutor_shchem_v1 import desktop_handout_sources as sources
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
    PreparationSourcesService, _digest)
from integrations.deeptutor_shchem_v1.desktop_handout_questions import handout_question_facade
from integrations.deeptutor_shchem_v1.desktop_composition_sources import import_questions
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import visible_blocks
from integrations.deeptutor_shchem_v1.reader_cancellation import ReadCancelled, read_cancel_scope


@pytest.fixture(autouse=True)
def isolated_context_preview_cache():
    with sources._PREVIEW_LOCK:
        sources._PREVIEWS.clear()
    yield
    with sources._PREVIEW_LOCK:
        sources._PREVIEWS.clear()


def _document_binding(facade):
    metadata = handout_knowledge_sources(facade)['sources'][0]
    return handout_knowledge_outline(facade, metadata)['source']


def _inventory(root):
    """Include each durable file, including profile, state, caches and originals."""
    return {str(path.relative_to(root)): sha256(path.read_bytes()).hexdigest()
            for path in root.rglob('*') if path.is_file()}


@pytest.fixture
def document_lecture(tmp_path):
    document = Document(); document.add_heading('分子间作用力', 2)
    document.add_heading('知识点1 范德华力', 3)
    document.add_paragraph('完整的讲解文字。')
    paragraph = document.add_paragraph('比较这两幅原图：')
    paragraph.add_run().add_picture(BytesIO(png(2)))
    paragraph.add_run().add_picture(BytesIO(png()))
    document.add_heading('知识点2 氢键', 3)
    document.add_paragraph('保留原条件与结论。')
    facade = make_facade(word_bytes(document), tmp_path / 'document')
    return facade, _document_binding(facade), document


def test_document_context_builds_only_real_picture_rows_and_deep_copies(document_lecture):
    facade, binding, _ = document_lecture; before = _inventory(facade.original.parent)
    result = handout_source_context(facade, binding)
    assert 'source' not in result and result['document_source'] == binding
    assert len(result['images']) == 2
    assert len({row['asset_id'] for row in result['images']}) == 2
    assert len({row['block_index'] for row in result['images']}) == 1
    for row, asset in zip(result['images'], result['preview']['assets']):
        assert validate_picture(row) == row and len(row) == 12
        assert row['schema'] == PICTURE_SCHEMA and 'source_role' not in row
        assert row['batch_id'] == binding['batch_id'] and row['source_id'] == binding['source_id']
        assert all(row[key] == asset[key] for key in (
            'asset_id', 'block_index', 'mime_type', 'sha256', 'bytes_count'))
    result['document_source']['source_name'] = '编辑过的名字'
    result['preview']['blocks'][0]['text'] = '编辑过的文字'
    again = handout_source_context(facade, binding)
    assert again['document_source'] == binding and again['preview']['blocks'][0]['text'] == '分子间作用力'
    assert _inventory(facade.original.parent) == before
    assert not facade.paths.workspace_root.exists() and not facade.paths.task_root.exists()


def _track_preview_parses(monkeypatch):
    calls = []
    read = PreparationSourcesService.word_preview_bytes
    def tracked(reader, content, name):
        calls.append((sha256(content).hexdigest(), name))
        return read(reader, content, name)
    monkeypatch.setattr(PreparationSourcesService, 'word_preview_bytes', tracked)
    return calls


def test_document_context_parses_once_rechecks_bytes_twice_and_never_uses_outline_or_disk_cache(
        document_lecture, monkeypatch):
    facade, binding, _ = document_lecture
    before = _inventory(facade.original.parent); reads = []
    original = facade._imported_word_source
    def exact(*args):
        reads.append(args); return original(*args)
    def forbidden(*args, **kwargs):
        raise AssertionError('Navigation must not read headings or a persistent preview cache.')
    monkeypatch.setattr(facade, '_imported_word_source', exact)
    monkeypatch.setattr(facade, 'imported_word_path', forbidden)
    monkeypatch.setattr(facade, 'imported_word_preview', forbidden, raising=False)
    monkeypatch.setattr('integrations.deeptutor_shchem_v1.desktop_handout_knowledge.handout_knowledge_outline', forbidden)
    monkeypatch.setattr('integrations.deeptutor_shchem_v1.desktop_word_preview_cache.WordPreviewCache.load', forbidden)
    parses = _track_preview_parses(monkeypatch)
    first = handout_source_context(facade, binding)
    second = handout_source_context(facade, binding)
    assert first == second and len(parses) == 1
    assert reads == [(binding['batch_id'], binding['source_id'])] * 4
    assert _inventory(facade.original.parent) == before
    assert not facade.paths.workspace_root.exists() and not facade.paths.task_root.exists()


def test_caller_cannot_mutate_the_private_cached_preview(document_lecture, monkeypatch):
    facade, binding, _ = document_lecture; parses = _track_preview_parses(monkeypatch)
    first = handout_source_context(facade, binding); expected = deepcopy(first)
    first['preview']['blocks'][0]['text'] = '改过的原文'
    first['preview']['assets'][0]['sha256'] = 'f' * 64
    first['images'][0]['source_id'] = 'FOREIGN'
    first['document_source']['source_role'] = 'answer'
    first['contexts'].clear()
    assert handout_source_context(facade, binding) == expected
    assert len(parses) == 1


@pytest.mark.parametrize('change', ['name', 'sha', 'parser', 'revision', 'text', 'resigned_text'])
def test_invalid_memory_preview_is_reparsed_from_current_bytes(document_lecture, monkeypatch, change):
    facade, binding, _ = document_lecture; parses = _track_preview_parses(monkeypatch)
    original = handout_source_context(facade, binding)
    with sources._PREVIEW_LOCK:
        cached, _size = sources._PREVIEWS[sources._preview_key(binding)]
        if change == 'name': cached['source_name'] = '其他原文.docx'
        elif change == 'sha': cached['source_sha256'] = 'f' * 64
        elif change == 'parser': cached['extraction_revision'] = 'old-parser'
        elif change == 'revision': cached['revision'] = 'f' * 64
        else:
            cached['blocks'][0]['text'] = '缓存被改过'
            if change == 'resigned_text':
                cached['revision'] = _digest({key: value for key, value in cached.items() if key != 'revision'})
    assert handout_source_context(facade, binding) == original
    assert len(parses) == 2


def test_current_parser_revision_invalidates_prior_in_memory_preview(document_lecture, monkeypatch):
    facade, binding, _ = document_lecture; handout_source_context(facade, binding)
    new_parser = 'test-current-parser-revision'
    monkeypatch.setattr(sources, 'NATIVE_WORD_TEXT_REVISION', new_parser)
    monkeypatch.setattr('integrations.deeptutor_shchem_v1.desktop_preparation_sources.NATIVE_WORD_TEXT_REVISION', new_parser)
    current = PreparationSourcesService(facade.paths.workspace_root).word_preview_bytes(
        facade.value.data, binding['source_name'])
    parses = _track_preview_parses(monkeypatch)
    with pytest.raises(ValueError, match='提取版本'): handout_source_context(facade, binding)
    updated = {**binding, 'source_revision': current['revision']}
    context = handout_source_context(facade, updated)
    assert context['preview']['extraction_revision'] == new_parser
    assert context['preview'] == current and len(parses) == 2


@pytest.mark.parametrize('mode', ['entries', 'bytes'])
def test_memory_preview_cache_is_bounded(tmp_path, monkeypatch, mode):
    facades, bindings = [], []
    for number in range(3):
        document = Document(); document.add_heading('知识点', 3)
        document.add_paragraph(f'完整知识说明{number}。')
        facade = make_facade(word_bytes(document), tmp_path / str(number))
        facades.append(facade); bindings.append(_document_binding(facade))
    monkeypatch.setattr(sources, '_PREVIEW_LIMIT', 2)
    if mode == 'bytes': monkeypatch.setattr(sources, '_PREVIEW_BYTES_LIMIT', 1)
    parses = _track_preview_parses(monkeypatch)
    for facade, binding in zip(facades, bindings): handout_source_context(facade, binding)
    assert len(parses) == 3
    if mode == 'entries':
        assert len(sources._PREVIEWS) == 2
        assert sources._preview_key(bindings[0]) not in sources._PREVIEWS
    else: assert not sources._PREVIEWS
    assert sum(row[1] for row in sources._PREVIEWS.values()) <= sources._PREVIEW_BYTES_LIMIT


def test_cached_same_byte_archive_cannot_replace_a_missing_selected_archive(document_lecture, monkeypatch):
    facade, binding, _ = document_lecture
    foreign = {**binding, 'batch_id': 'OTHER_BATCH', 'source_id': 'OTHER_SOURCE'}
    available = {(binding['batch_id'], binding['source_id']), ('OTHER_BATCH', 'OTHER_SOURCE')}
    reads = []
    def exact(batch_id, source_id):
        reads.append((batch_id, source_id))
        if (batch_id, source_id) not in available: raise ValueError('指定归档已删除')
        return SimpleNamespace(content=facade.value.data, filename=binding['source_name'],
            role=binding['source_role'], source_sha256=binding['source_sha256'],
            effective_source_file_id=source_id)
    monkeypatch.setattr(facade, '_imported_word_source', exact)
    first = handout_source_context(facade, binding)
    second = handout_source_context(facade, foreign)
    assert len(sources._PREVIEWS) == 2
    assert first['preview'] == second['preview']
    assert second['document_source'] == foreign
    assert all(row['batch_id'] == 'OTHER_BATCH' and row['source_id'] == 'OTHER_SOURCE'
        for row in second['images'])
    available.remove((binding['batch_id'], binding['source_id']))
    with pytest.raises(ValueError, match='指定归档'): handout_source_context(facade, binding)
    assert reads[-1] == (binding['batch_id'], binding['source_id'])
    assert handout_source_context(facade, foreign)['document_source'] == foreign


@pytest.mark.parametrize('field,value', [('effective_source_file_id', 'FOREIGN'),
    ('source_sha256', 'f' * 64)])
def test_declared_archive_identity_must_match_actual_bytes_and_binding(
        document_lecture, monkeypatch, field, value):
    facade, binding, _ = document_lecture; handout_source_context(facade, binding)
    read = facade._imported_word_source
    def altered(*args):
        source = read(*args); setattr(source, field, value); return source
    monkeypatch.setattr(facade, '_imported_word_source', altered)
    with pytest.raises(ValueError): handout_source_context(facade, binding)


@pytest.mark.parametrize('warm', [False, True])
def test_source_change_during_navigation_callback_is_rejected_before_return(
        document_lecture, monkeypatch, warm):
    facade, binding, _ = document_lecture
    if warm: handout_source_context(facade, binding)
    labels = sources._context_labels
    def changed(*args):
        result = labels(*args); facade.value.role = 'answer'; return result
    monkeypatch.setattr(sources, '_context_labels', changed)
    with pytest.raises(ValueError, match='期间变化'): handout_source_context(facade, binding)
    if not warm: assert not sources._PREVIEWS


@pytest.mark.parametrize('stage', ['preview', 'final_read'])
def test_cancelled_context_never_publishes_a_new_memory_preview(document_lecture, monkeypatch, stage):
    facade, binding, _ = document_lecture; event = Event()
    if stage == 'preview':
        read = PreparationSourcesService.word_preview_bytes
        def cancelled(reader, *args):
            preview = read(reader, *args); event.set(); return preview
        monkeypatch.setattr(PreparationSourcesService, 'word_preview_bytes', cancelled)
    else:
        read = facade._imported_word_source; calls = []
        def cancelled(*args):
            value = read(*args); calls.append(args)
            if len(calls) == 2: event.set()
            return value
        monkeypatch.setattr(facade, '_imported_word_source', cancelled)
    with read_cancel_scope(event, check_boundaries=False):
        with pytest.raises(ReadCancelled): handout_source_context(facade, binding)
    assert not sources._PREVIEWS and not facade.paths.task_root.exists()


def test_document_picture_rows_can_use_original_pixel_reader(scoped_lecture):
    facade, _, picture, _, _, _, _ = scoped_lecture
    result = handout_source_context(facade, _document_binding(facade))
    assert handout_image_context(facade, picture)['images'] == result['images']
    assert handout_image_preview(facade, result['images'][0])['bytes']


def test_document_restored_location_preserves_original_provenance(document_lecture, monkeypatch):
    facade, binding, _ = document_lecture; original = deepcopy(binding)
    monkeypatch.setattr(facade, 'imported_word_path', lambda *args: 'D:/恢复/原始讲义.docx')
    context = handout_source_context(facade, binding)
    assert context['document_source'] == original and binding == original
    assert all(row['source_path'] == original['source_path'] for row in context['images'])


@pytest.mark.parametrize('field,value', [('source_name', '其他讲义.docx'),
    ('source_role', 'answer'), ('source_sha256', 'f' * 64),
    ('source_revision', 'f' * 64), ('source_id', 'OTHER'), ('batch_id', 'OTHER')])
def test_document_source_binding_change_is_rejected_before_result(scoped_lecture, field, value):
    facade, *_ = scoped_lecture; binding = _document_binding(facade)
    binding[field] = value
    before = _inventory(facade.paths.state_root)
    with pytest.raises((ValueError, DesktopFacadeError)): handout_source_context(facade, binding)
    assert _inventory(facade.paths.state_root) == before


@pytest.mark.parametrize('change', ['bytes', 'name', 'role', 'deleted'])
def test_actual_document_change_is_rejected(document_lecture, change):
    facade, binding, document = document_lecture
    if change == 'bytes':
        document.add_paragraph('改变的正文'); facade.value.data = word_bytes(document)
    elif change == 'name': facade.value.name = '其他讲义.docx'
    elif change == 'role': facade.value.role = 'answer'
    else: facade.value.data = None
    with pytest.raises(ValueError): handout_source_context(facade, binding)
    assert not facade.paths.task_root.exists()


def test_source_changes_during_document_read_are_rejected(document_lecture, monkeypatch):
    facade, binding, _ = document_lecture; original = facade._imported_word_source
    def changed(*args):
        source = original(*args); facade.value.role = 'answer'; return source
    monkeypatch.setattr(facade, '_imported_word_source', changed)
    with pytest.raises(ValueError, match='期间变化'): handout_source_context(facade, binding)
    assert not facade.paths.task_root.exists()


@pytest.mark.parametrize('field,value', [('asset_id', 'another-picture'),
    ('block_index', 999), ('sha256', 'f' * 64), ('bytes_count', 1),
    ('mime_type', 'image/jpeg')])
def test_legacy_picture_handoff_keeps_occurrence_validation(scoped_lecture, field, value):
    facade, _, picture, *_ = scoped_lecture; picture = deepcopy(picture)
    picture[field] = value
    with pytest.raises(ValueError): handout_source_context(facade, picture)


@pytest.mark.parametrize('change', ['missing', 'extra', 'mixed', 'unknown', 'hash_type', 'unbound'])
def test_strict_union_rejects_incomplete_extra_and_mixed_bindings(source_kind_lecture, change):
    facade, _, binding, *_ = source_kind_lecture; altered = deepcopy(binding)
    if change == 'missing': altered.pop('source_revision')
    elif change == 'extra': altered['foreign'] = True
    elif change == 'mixed':
        if 'source_role' in altered: altered['asset_id'] = 'word-b1-image1'
        else: altered['source_role'] = 'handout'
    elif change == 'unknown': altered['schema'] = 'another-source.v1'
    elif change == 'hash_type': altered['source_sha256'] = int('1' * 64)
    else:
        altered.pop('source_path'); altered.pop('source_revision')
    with pytest.raises(ValueError): validate_handout_source(altered)
    with pytest.raises(ValueError): handout_question_facade(facade, altered)


def test_valid_union_returns_independent_snapshot(source_kind_lecture):
    _, _, binding, *_ = source_kind_lecture
    checked = validate_handout_source(binding); assert checked == binding
    checked['source_name'] = '教师新标题'; assert checked != binding


def test_read_only_context_and_scoped_operations_leave_disk_state_unchanged(source_kind_lecture):
    facade, scoped, binding, choices, provider, *_ = source_kind_lecture
    before = _inventory(facade.paths.state_root); state = facade.state_store.snapshot()
    service = facade._word_questions(); service._cache.clear()
    handout_source_context(facade, binding)
    scoped.word_question_catalog()
    scoped.word_question_reference(choices)
    scoped.word_question_source(key=choices[0]['key'], revision=choices[0]['revision'])
    scoped.word_question_reading_layout(choices[0]['key'], choices[0]['revision'])
    assert _inventory(facade.paths.state_root) == before
    assert facade.state_store.snapshot() == state and provider.borrow_calls == 0
    assert not (facade.paths.state_root / 'word-question-previews').exists()
    assert not (facade.paths.task_root / 'preparation-v1/images').exists()


def test_cancelled_source_and_scoped_question_reads_have_no_writes(source_kind_lecture):
    facade, scoped, binding, *_ = source_kind_lecture
    before = _inventory(facade.paths.state_root); event = Event(); event.set()
    with read_cancel_scope(event, check_boundaries=False):
        with pytest.raises(ReadCancelled): handout_source_context(facade, binding)
        with pytest.raises(ReadCancelled): scoped.word_question_catalog()
    with pytest.raises(ReadCancelled): scoped.word_question_catalog_progress(None, event)
    event.clear()
    with pytest.raises(ReadCancelled):
        scoped.word_question_catalog_progress(lambda _: event.set(), event)
    assert _inventory(facade.paths.state_root) == before


def test_cancellation_during_actual_source_read_has_no_cache_or_images(document_lecture, monkeypatch):
    facade, binding, _ = document_lecture; event = Event(); read = facade._imported_word_source
    def cancelled(*args):
        value = read(*args); event.set(); return value
    monkeypatch.setattr(facade, '_imported_word_source', cancelled)
    with read_cancel_scope(event, check_boundaries=False):
        with pytest.raises(ReadCancelled): handout_source_context(facade, binding)
    assert not facade.paths.task_root.exists() and not facade.paths.workspace_root.exists()


def test_source_change_during_catalog_attribute_callback_is_rejected(source_kind_lecture, monkeypatch):
    facade, scoped, _, *_ = source_kind_lecture
    def changed():
        scoped.binding['source_revision'] = 'f' * 64
        return {'knowledge_points': [], 'nodes': []}
    monkeypatch.setattr(facade._word_questions(), '_read_attribute_catalog', changed)
    with pytest.raises(ValueError): scoped.word_question_catalog()


def test_every_existing_scoped_operation_rechecks_context_under_native_lock(source_kind_lecture, monkeypatch):
    facade, scoped, binding, choices, *_ = source_kind_lecture
    from integrations.deeptutor_shchem_v1 import desktop_handout_questions as questions
    read_context = questions.handout_source_context; calls = []
    service = facade._word_questions(); key, revision = choices[0]['key'], choices[0]['revision']
    def checked_context(*args):
        assert service._lock._is_owned(); calls.append(args[1])
        return read_context(*args)
    monkeypatch.setattr(questions, 'handout_source_context', checked_context)
    operations = {
        'word_question_source': (key, revision),
        'word_question_reading_layout': (key, revision),
        'word_question_source_locations': (key, revision),
        'word_question_location_image': (key, revision, 'LOCATION', 'ASSET'),
        'word_question_select_native': (key, revision, 'LOCATION'),
        'word_question_image': (key, revision, 'ASSET'),
        'word_question_attribute_options': (key, revision),
        'word_question_preview_automatic_scope_labels': (key, revision),
        'word_question_save_automatic_scope_labels': (key, revision),
        'word_question_save_attributes': (key, revision, {}),
        'word_question_update_range': (key, revision),
        'word_question_reference': (choices,),
        'word_question_export': ('课堂练习', choices),
    }
    def callback(*args, **kwargs):
        assert service._lock._is_owned()
        assert service._locations[key] == (binding['batch_id'], binding['source_id'])
        return 'checked result'
    for name, args in operations.items():
        monkeypatch.setattr(facade, name, callback)
        start = len(calls)
        assert getattr(scoped, name)(*args) == 'checked result'
        assert len(calls) == start + 2
    for name in ('word_question_save_selection', 'word_question_saved_selection',
                 'add_word_questions_to_basket', 'word_semantic_tag_preview'):
        assert not hasattr(scoped, name)


def test_document_with_no_images_still_opens_complete_question_and_teacher_answer(
        desktop_paths, tmp_path, monkeypatch):
    monkeypatch.setattr('integrations.deeptutor_shchem_v1.desktop_word_questions.load_attribute_catalog',
        lambda *a, **k: {'knowledge_points': [], 'nodes': []})
    document = Document(); document.add_heading('纯文字讲义', 1)
    document.add_heading('知识点1 分子间作用力', 3)
    document.add_paragraph('完整知识讲解。')
    document.add_heading('考向1 例题', 3)
    document.add_paragraph('【例题1】下列关于分子间作用力的说法正确的是（ ）')
    document.add_paragraph('A．甲 B．乙 C．丙 D．丁')
    document.add_paragraph('【答案】A'); document.add_paragraph('【解析】完整原解析。')
    path = tmp_path / '纯文字讲义.docx'; document.save(path)
    provider = FakeProviderStore(configured=False); facade = _facade(desktop_paths, provider)
    facade.save_visual_import_batch(handout_files=(path,), source_type='教师讲义')
    binding = _document_binding(facade); before = _inventory(facade.paths.state_root)
    context = handout_source_context(facade, binding)
    assert context['images'] == [] and context['contexts'] == {} and 'source' not in context
    scoped = handout_question_facade(facade, binding)
    row = scoped.word_question_catalog()['items'][0]
    choices = [{'key': row['key'], 'revision': row['revision'], 'points': 3}]
    reference = scoped.word_question_reference(choices)
    assert _inventory(facade.paths.state_root) == before
    result = import_questions(scoped, choices, [], expected_reference=reference)
    assert result['image_assets'] == [] and len(result['resources']) == 1
    item = result['resources'][0]
    assert any('完整原解析' in b['value'] for b in item['blocks'] if b['role'] == 'answer')
    assert all(b['role'] != 'answer' for b in visible_blocks(item))
    assert item['original']['evidence']['source_sha256'] == binding['source_sha256']
    assert not item['human_reviewed'] and not item['publication_allowed']
    assert _inventory(facade.paths.state_root) == before and provider.borrow_calls == 0


@pytest.fixture
def closed_eq_document_lecture(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document(); document.add_heading('知识点1 聚合反应条件', 3)
    paragraph = document.add_paragraph('完整原文条件：')
    for kind in ('begin', 'code', 'separate', 'end'):
        run = OxmlElement('w:r')
        node = OxmlElement('w:instrText' if kind == 'code' else 'w:fldChar')
        if kind == 'code': node.text = r' EQ \o(――→,\s\up7(催化剂)) '
        else: node.set(qn('w:fldCharType'), kind)
        run.append(node); paragraph._p.append(run)
    paragraph.add_run('；原条件保持完整。')
    document.add_paragraph().add_run().add_picture(BytesIO(png(2)))
    facade = make_facade(word_bytes(document), tmp_path / 'closed-eq-document')
    return facade, _document_binding(facade), document


def test_closed_eq_document_context_cold_and_warm_keep_exact_binding_and_four_source_reads(
        closed_eq_document_lecture, monkeypatch):
    facade, binding, _ = closed_eq_document_lecture
    before = _inventory(facade.original.parent); facade.value.reads = 0
    parses = _track_preview_parses(monkeypatch)
    first = handout_source_context(facade, binding)
    second = handout_source_context(facade, binding)
    assert first == second and first['document_source'] == binding
    assert first['preview']['extraction_revision'] == sources.NATIVE_WORD_TEXT_REVISION + '+' + sources.STATIC_EQ_REVISION
    assert first['preview']['blocks'][1]['text'] == '完整原文条件：――→〔上移7pt：催化剂〕；原条件保持完整。'
    assert len(first['preview']['blocks'][1]['static_eqs']) == 1
    assert len(first['images']) == 1 and len(parses) == 1 and facade.value.reads == 4
    assert _inventory(facade.original.parent) == before
    assert not facade.paths.workspace_root.exists() and not facade.paths.task_root.exists()


def test_unknown_eq_revision_suffix_is_rejected_even_with_resigned_preview_and_binding(
        closed_eq_document_lecture, monkeypatch):
    facade, binding, _ = closed_eq_document_lecture
    before = _inventory(facade.original.parent); parses = _track_preview_parses(monkeypatch)
    current = handout_source_context(facade, binding)
    forged = deepcopy(current['preview'])
    forged['extraction_revision'] += '+unknown-future-extension'
    forged['revision'] = _digest({key: value for key, value in forged.items() if key != 'revision'})
    altered = {**binding, 'source_revision': forged['revision']}
    assert not sources._valid_preview(forged, altered)
    with sources._PREVIEW_LOCK:
        key = sources._preview_key(altered)
        _preview, size = sources._PREVIEWS[key]
        sources._PREVIEWS[key] = (forged, size)
    with pytest.raises(ValueError, match='提取版本'):
        handout_source_context(facade, altered)
    assert len(parses) == 2 and not sources._PREVIEWS
    assert _inventory(facade.original.parent) == before


def test_old_native_revision_cache_reparses_for_current_eq_binding_but_old_binding_fails(
        closed_eq_document_lecture, monkeypatch):
    facade, binding, _ = closed_eq_document_lecture
    before = _inventory(facade.original.parent); parses = _track_preview_parses(monkeypatch)
    current = handout_source_context(facade, binding)
    old = deepcopy(current['preview'])
    old['extraction_revision'] = sources.NATIVE_WORD_TEXT_REVISION
    old['revision'] = _digest({key: value for key, value in old.items() if key != 'revision'})
    stale = {**binding, 'source_revision': old['revision']}
    assert stale['source_revision'] != binding['source_revision']
    with sources._PREVIEW_LOCK:
        key = sources._preview_key(binding)
        _preview, size = sources._PREVIEWS[key]
        sources._PREVIEWS[key] = (old, size)
    assert handout_source_context(facade, binding) == current
    assert len(parses) == 2
    with pytest.raises(ValueError, match='提取版本'):
        handout_source_context(facade, stale)
    assert len(parses) == 3
    assert _inventory(facade.original.parent) == before


def test_static_eq_implementation_revision_invalidates_cache_key_and_old_binding(
        closed_eq_document_lecture, monkeypatch):
    facade, binding, _ = closed_eq_document_lecture
    before = _inventory(facade.original.parent)
    handout_source_context(facade, binding)
    old_key = sources._preview_key(binding)
    next_revision = 'test-next-static-eq-revision'
    monkeypatch.setattr(sources, 'STATIC_EQ_REVISION', next_revision)
    monkeypatch.setattr('integrations.deeptutor_shchem_v1.desktop_preparation_sources.STATIC_EQ_REVISION', next_revision)
    assert sources._preview_key(binding) != old_key
    assert old_key in sources._PREVIEWS and sources._preview_key(binding) not in sources._PREVIEWS
    preview = PreparationSourcesService(facade.paths.workspace_root).word_preview_bytes(
        facade.value.data, binding['source_name'])
    parses = _track_preview_parses(monkeypatch)
    with pytest.raises(ValueError, match='提取版本'):
        handout_source_context(facade, binding)
    updated = {**binding, 'source_revision': preview['revision']}
    fresh = handout_source_context(facade, updated)
    assert fresh['preview'] == preview
    assert fresh['preview']['extraction_revision'] == sources.NATIVE_WORD_TEXT_REVISION + '+' + next_revision
    assert handout_source_context(facade, updated) == fresh and len(parses) == 2
    assert _inventory(facade.original.parent) == before
