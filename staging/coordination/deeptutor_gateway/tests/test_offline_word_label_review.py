"""Only synthetic Word files, candidate quotes and temporary personal stores."""

from __future__ import annotations

import io
import json
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from test_desktop_visual_import_facade import _png

from integrations.deeptutor_shchem_v1.desktop_word_question_attributes import (
    WordQuestionAttributeStore,
    _seal,
    suggest_attributes,
)
from integrations.deeptutor_shchem_v1.offline_word_label_review import (
    CANDIDATE_SCHEMA_VERSION,
    PROVENANCE,
    RULE_REVISION,
    OfflineWordLabelReviewError,
    OfflineWordLabelReviewService,
)
from runtime.deeptutor_shchem import apply_offline_word_labels as cli
from staging.coordination.deeptutor_gateway.tests.test_desktop_word_auto_attributes import (
    setup as setup,  # noqa: PLC0414 -- existing synthetic fixture
)


def _entry(row, stored):
    evidence = {
        "block_index": row["question_blocks"][0]["index"],
        "quote": row["question_blocks"][0]["text"], "image_sha256": "",
    }
    return {
        "key": row["key"], "question_revision": row["revision"],
        "source_sha256": row["source_sha256"],
        "expected_stored_attribute_revision": stored["revision"],
        "decoded": {
            "primary": {"id": "K11", "evidence": [evidence]},
            "curriculum": [{"section_key": "S1", "evidence": [deepcopy(evidence)]}],
            "note": "合成测试候选，非人工确认。",
        },
    }


@pytest.fixture
def corpus(setup):
    words, facade, source, directory = setup
    rows = words.catalog()["items"]
    catalog = words._read_attribute_catalog()
    old = []
    for row in rows:
        attr = suggest_attributes(row, {"source_name": row["source_name"]}, catalog)
        attr.update(
            primary_knowledge={"id": "unknown", "label": "待确认", "status": "unknown", "evidence": []},
            supporting_knowledge=[], curriculum_candidates=[], curriculum_status="pending_mapping",
            rule_revision="synthetic-old-labels-v1", teacher_note="原有备注必须保留",
        )
        old.append(_seal(attr))
    stored = words.attribute_store.save_many(old)
    batch = {
        "schema_version": CANDIDATE_SCHEMA_VERSION, "provenance": PROVENANCE,
        "candidate_only": True, "human_review": False,
        "entries": [_entry(row, attr) for row, attr in zip(rows, stored, strict=True)],
    }
    return SimpleNamespace(
        words=words, facade=facade, source=source, directory=directory,
        rows=rows, stored=stored, batch=batch, service=OfflineWordLabelReviewService(words),
    )


def _history(corpus):
    return {row["key"]: corpus.words.attribute_store.history(row["key"]) for row in corpus.rows}


def _rewrite(corpus, position, **changes):
    row = corpus.words.attribute_store.save_many([_seal({**corpus.stored[position], **changes})])[0]
    corpus.stored[position] = row
    corpus.batch["entries"][position]["expected_stored_attribute_revision"] = row["revision"]
    return row


def test_preview_is_metadata_only_stable_and_does_not_change_attributes(corpus):
    histories = _history(corpus)
    state = corpus.words.state.snapshot()
    source = corpus.source.read_bytes()
    batch = deepcopy(corpus.batch)
    result = corpus.service.preview(batch)
    assert result == corpus.service.preview(batch)
    assert result["entry_count"] == result["changed_question_count"] == 2
    assert result["primary_filled_count"] == result["curriculum_filled_count"] == 2
    assert result["source_count"] == 1 and result["unchanged_question_count"] == 0
    assert result["before"] == {"question_count": 2, "primary_known": 0, "curriculum_mapped": 0}
    assert result["after_proposed"] == {"question_count": 2, "primary_known": 2, "curriculum_mapped": 2}
    assert result["candidate_only"] is True
    assert result["human_review"] is result["teacher_confirmed"] is result["provider_invoked"] is False
    serial = json.dumps(result, ensure_ascii=False)
    assert "decoded" not in serial and batch["entries"][0]["decoded"]["note"] not in serial
    for row in corpus.rows:
        for block in row["question_blocks"] + row["answer_blocks"]:
            assert block["text"] not in serial
    assert _history(corpus) == histories
    assert corpus.words.state.snapshot() == state
    assert corpus.source.read_bytes() == source
    assert batch == corpus.batch


def test_apply_preserves_all_other_fields_and_prior_history_and_binds_cas(corpus, monkeypatch):
    before = _history(corpus)
    original_save = corpus.words.attribute_store.save_many
    calls = []

    def save(rows, *, expected_revisions=None):
        assert corpus.words._lock._is_owned()
        calls.append(deepcopy(expected_revisions))
        return original_save(rows, expected_revisions=expected_revisions)

    monkeypatch.setattr(corpus.words.attribute_store, "save_many", save)
    preview = corpus.service.preview(corpus.batch)
    result = corpus.service.apply(corpus.batch, expected_plan_sha256=preview["plan_sha256"])
    assert calls == [{row["key"]: row["revision"] for row in corpus.stored}]
    assert result["readback_verified"] is result["applied"] is True
    assert result["after_actual"] == preview["after_proposed"]
    allowed = {"primary_knowledge", "curriculum_candidates", "curriculum_status", "rule_revision", "revision", "edit_version"}
    for old in corpus.stored:
        actual = corpus.words.attribute_store.get(old["key"])
        assert actual["primary_knowledge"]["status"] == "auto_suggested"
        assert actual["curriculum_status"] == "auto_suggested"
        assert actual["rule_revision"] == RULE_REVISION
        assert {key: value for key, value in actual.items() if key not in allowed} == {
            key: value for key, value in old.items() if key not in allowed
        }
        assert actual["edit_version"] == old["edit_version"] + 1
        assert corpus.words.attribute_store.history(old["key"]) == [*before[old["key"]], actual]
        assert result["saved_attribute_revisions"][old["key"]] == actual["revision"]
    with pytest.raises(OfflineWordLabelReviewError, match="已有标签版本已变化"):
        corpus.service.apply(corpus.batch, expected_plan_sha256=preview["plan_sha256"])


def test_noop_keeps_revision_rules_and_history_without_saving(corpus, monkeypatch):
    for entry in corpus.batch["entries"]:
        entry["decoded"] = {"primary": {"id": "unknown", "evidence": []}, "curriculum": [], "note": "无充分依据"}
    before = _history(corpus)
    monkeypatch.setattr(corpus.words.attribute_store, "save_many", lambda *_a, **_k: pytest.fail("noop must not save"))
    preview = corpus.service.preview(corpus.batch)
    assert preview["changed_question_count"] == 0
    receipt = corpus.service.apply(corpus.batch, expected_plan_sha256=preview["plan_sha256"])
    assert receipt["readback_verified"] is True
    assert _history(corpus) == before


def test_existing_nonempty_labels_remain_and_only_the_delta_is_stamped(corpus):
    generated = suggest_attributes(corpus.rows[0], {"source_name": "synthetic.docx"}, corpus.words._read_attribute_catalog())
    _rewrite(corpus, 0, primary_knowledge=generated["primary_knowledge"],
             curriculum_candidates=generated["curriculum_candidates"], curriculum_status="auto_suggested")
    assert corpus.stored[0]["primary_knowledge"]["id"] == "K11"
    assert corpus.stored[0]["curriculum_candidates"]
    before = _history(corpus)
    preview = corpus.service.preview(corpus.batch)
    assert preview["changed_question_count"] == preview["unchanged_question_count"] == 1
    corpus.service.apply(corpus.batch, expected_plan_sha256=preview["plan_sha256"])
    assert corpus.words.attribute_store.get(corpus.stored[0]["key"]) == corpus.stored[0]
    assert corpus.words.attribute_store.history(corpus.stored[0]["key"]) == before[corpus.stored[0]["key"]]


def test_full_question_options_and_attached_context_are_available_as_evidence(corpus):
    row = corpus.rows[0]
    changed = corpus.words.update_range(
        row["key"], row["revision"], block_start=row["block_start"],
        question_end=row["question_end"], answer_start=row["answer_start"],
        block_end=row["block_end"], context_start=1, context_end=1,
    )
    # update_range returns a fresh source-bound question, without saving tags.
    current = corpus.words._resolve([{"key": row["key"], "revision": changed["revision"]}])[0][0]
    attr = suggest_attributes(current, corpus.stored[0]["source"], corpus.words._read_attribute_catalog())
    attr.update(primary_knowledge=deepcopy(corpus.stored[0]["primary_knowledge"]),
                curriculum_candidates=[], curriculum_status="pending_mapping")
    saved = corpus.words.attribute_store.save_many([_seal(attr)])[0]
    batch = deepcopy(corpus.batch)
    batch["entries"] = [_entry(current, saved)]
    decoded = batch["entries"][0]["decoded"]
    last_question = current["question_blocks"][-1]
    context = current["context_blocks"][0]
    decoded["primary"]["evidence"] = [{"block_index": last_question["index"], "quote": last_question["text"], "image_sha256": ""}]
    decoded["curriculum"][0]["evidence"] = [{"block_index": context["index"], "quote": context["text"], "image_sha256": ""}]
    assert corpus.service.preview(batch)["changed_question_count"] == 1


@pytest.mark.parametrize("kind", ["image", "formula_gap"])
def test_answer_assets_or_formula_gaps_do_not_block_a_complete_text_question(corpus, kind):
    document = Document()
    document.add_paragraph("【即学即练1】解释氧化还原反应。")
    answer = document.add_paragraph("【答案】合成参考答案。")
    if kind == "image":
        answer.add_run().add_picture(io.BytesIO(_png("blue")))
    else:
        answer.add_run("【待查看原文公式】")
    stream = io.BytesIO()
    document.save(stream)
    corpus.facade.add("B2", stream.getvalue(), name="synthetic-answer.docx", source_id="SOURCE2")
    row = next(row for row in corpus.words.catalog()["items"] if row["batch_id"] == "B2")
    attr = suggest_attributes(row, {"source_name": row["source_name"]}, corpus.words._read_attribute_catalog())
    attr.update(primary_knowledge=deepcopy(corpus.stored[0]["primary_knowledge"]),
                curriculum_candidates=[], curriculum_status="pending_mapping")
    assert attr["material_status"]["question_image_count"] == 0
    assert attr["material_status"]["answer_image_count"] == (1 if kind == "image" else 0)
    saved = corpus.words.attribute_store.save_many([_seal(attr)])[0]
    batch = {**deepcopy(corpus.batch), "entries": [_entry(row, saved)]}
    preview = corpus.service.preview(batch)
    assert preview["changed_question_count"] == 1
    receipt = corpus.service.apply(batch, expected_plan_sha256=preview["plan_sha256"])
    assert receipt["readback_verified"] is True
    actual = corpus.words.attribute_store.get(row["key"])
    assert actual["answer_status"] == saved["answer_status"]
    assert actual["material_status"] == saved["material_status"]


@pytest.mark.parametrize("change", [
    lambda b: b.update(provenance="provider_api"),
    lambda b: b.update(human_review=True),
    lambda b: b.update(candidate_only=False),
    lambda b: b.update(profile_id="fake"),
    lambda b: b.update(entries=[]),
    lambda b: b.update(entries=b["entries"] * 51),
    lambda b: b.update(entries=[b["entries"][0], b["entries"][0]]),
    lambda b: b["entries"][0].pop("expected_stored_attribute_revision"),
    lambda b: b["entries"][0]["decoded"].update(original_exam_type="grade_exam"),
])
def test_candidate_contract_rejects_missing_fake_or_unbounded_metadata(corpus, change):
    before = _history(corpus)
    change(corpus.batch)
    with pytest.raises(OfflineWordLabelReviewError):
        corpus.service.preview(corpus.batch)
    assert _history(corpus) == before


@pytest.mark.parametrize("case", ["bad_quote", "answer", "image", "bad_primary", "bad_section", "missing_evidence", "unknown_with_evidence"])
def test_invalid_ids_or_non_question_evidence_reject_entire_batch(corpus, case):
    value = corpus.batch["entries"][1]["decoded"]
    if case == "bad_quote":
        value["primary"]["evidence"][0]["quote"] = "原文中不存在的杜撰说明"
    elif case == "answer":
        block = corpus.rows[1]["answer_blocks"][0]
        value["primary"]["evidence"] = [{"block_index": block["index"], "quote": block["text"], "image_sha256": ""}]
    elif case == "image":
        value["primary"]["evidence"][0]["image_sha256"] = "a" * 64
    elif case == "bad_primary":
        value["primary"]["id"] = "K99"
    elif case == "bad_section":
        value["curriculum"][0]["section_key"] = "missing-section"
    elif case == "missing_evidence":
        value["primary"]["evidence"] = []
    else:
        value["primary"]["id"] = "unknown"
    before = _history(corpus)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        corpus.service.preview(corpus.batch)
    assert error.value.code == "invalid_candidate_evidence"
    assert _history(corpus) == before


@pytest.mark.parametrize("case", ["teacher", "pinned", "primary", "supporting", "curriculum_status", "mapping"])
def test_protected_automatic_or_teacher_rows_are_rejected(corpus, case):
    changes = {}
    if case == "teacher":
        changes["annotation_source"] = "teacher_modified"
    elif case == "pinned":
        changes["rule_revision"] = "synthetic-pinned-r1"
    elif case == "curriculum_status":
        changes["curriculum_status"] = "teacher_confirmed"
    elif case == "mapping":
        changes["curriculum_candidates"] = [{
            "section_key": "S1", "chapter_id": "C1", "volume_id": "V1", "label": "测试节",
            "status": "teacher_confirmed", "evidence": [],
        }]
    else:
        tag = {"id": "K11", "label": "测试", "status": "teacher_confirmed", "evidence": []}
        changes["primary_knowledge" if case == "primary" else "supporting_knowledge"] = tag if case == "primary" else [tag]
    _rewrite(corpus, 0, **changes)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        corpus.service.preview(corpus.batch)
    assert error.value.code == "protected_attributes"


@pytest.mark.parametrize("field", ["source_sha256", "question_revision", "source_revision", "index_revision", "extraction_revision"])
def test_stale_stored_identity_is_rejected(corpus, field, tmp_path):
    if field == "source_sha256":
        # Simulate a pre-existing legacy row; the real store correctly refuses
        # changing a current row's source digest through its write interface.
        store = WordQuestionAttributeStore(tmp_path / "legacy-state")
        stale = store.save_many([_seal({**corpus.stored[0], field: "b" * 64}), corpus.stored[1]])
        corpus.words.attribute_store = store
        for entry, attr in zip(corpus.batch["entries"], stale, strict=True):
            entry["expected_stored_attribute_revision"] = attr["revision"]
    else:
        _rewrite(corpus, 0, **{field: "changed-version"})
    with pytest.raises(OfflineWordLabelReviewError):
        corpus.service.preview(corpus.batch)


def test_missing_stored_row_is_never_created(corpus, tmp_path):
    corpus.words.attribute_store = WordQuestionAttributeStore(tmp_path / "empty")
    with pytest.raises(OfflineWordLabelReviewError) as error:
        corpus.service.preview(corpus.batch)
    assert error.value.code == "stored_attributes_missing"
    assert not corpus.words.attribute_store.path.exists()


@pytest.mark.parametrize("case", ["not_ready", "boundary", "gap", "image", "unsupported", "source_issue", "sharing_unknown", "missing_context", "missing_visual"])
def test_incomplete_or_unsupported_sources_are_rejected(corpus, monkeypatch, case):
    if case in {"missing_context", "missing_visual"}:
        _rewrite(corpus, 0, material_status={**corpus.stored[0]["material_status"], case: True})
    else:
        original = corpus.words._resolve

        def resolve(*args, **kwargs):
            rows, inventory = original(*args, **kwargs)
            row = rows[0]
            if case == "not_ready":
                row["selection_ready"] = False
            elif case == "boundary":
                row["boundary_status"] = "needs_review"
            elif case == "gap":
                row["question_blocks"] = row["question_blocks"][:-1]
            elif case == "image":
                row["question_blocks"][0]["assets"] = [{"sha256": "a" * 64}]
            elif case == "unsupported":
                row["question_blocks"][0]["warnings"] = ["合成不可读对象警告"]
            elif case == "source_issue":
                row["content_quality"]["issues"] = [{"id": "known-error"}]
            else:
                row["shared_material_policy"] = "unknown"
            return rows, inventory

        monkeypatch.setattr(corpus.words, "_resolve", resolve)
    before = _history(corpus)
    with pytest.raises(OfflineWordLabelReviewError):
        corpus.service.preview(corpus.batch)
    assert _history(corpus) == before


@pytest.mark.parametrize("changed", ["source", "range", "catalog", "proposal", "attributes", "unselected_preview_block"])
def test_apply_rebuild_rejects_changed_preview_bindings(corpus, monkeypatch, changed):
    preview = corpus.service.preview(corpus.batch)
    before = _history(corpus)
    if changed == "source":
        corpus.source.write_bytes(corpus.source.read_bytes() + b"changed")
    elif changed == "range":
        row = corpus.rows[0]
        corpus.words.update_range(row["key"], row["revision"],
                                 block_start=row["block_start"], question_end=row["question_end"],
                                 answer_start=row["answer_start"], block_end=row["block_end"],
                                 context_start=1, context_end=1)
    elif changed == "catalog":
        data = json.loads(corpus.directory.read_text(encoding="utf-8"))
        data["nodes"][0]["section_title"] = "新的合成目录标题"
        corpus.directory.write_text(json.dumps(data), encoding="utf-8")
    elif changed == "proposal":
        corpus.batch["entries"][0]["decoded"]["note"] = "新的候选说明"
    elif changed == "attributes":
        corpus.words.attribute_store.save_many([_seal({**corpus.stored[0], "teacher_note": "并发修改"})])
        before = _history(corpus)
    else:
        original = corpus.words._resolve

        def changed_preview(*args, **kwargs):
            rows, inventory = original(*args, **kwargs)
            inventory = deepcopy(inventory)
            inventory[rows[0]["source_id"]][2]["blocks"][0]["text"] += "无关段落变更"
            return rows, inventory

        monkeypatch.setattr(corpus.words, "_resolve", changed_preview)
    with pytest.raises(OfflineWordLabelReviewError):
        corpus.service.apply(corpus.batch, expected_plan_sha256=preview["plan_sha256"])
    assert _history(corpus) == before


def test_cas_conflict_rolls_back_earlier_rows_in_same_transaction(corpus, monkeypatch):
    preview = corpus.service.preview(corpus.batch)
    histories = _history(corpus)
    real_save = corpus.words.attribute_store.save_many

    def conflicting_save(rows, *, expected_revisions):
        real_save([_seal({**corpus.stored[1], "teacher_note": "合成并发修改"})])
        return real_save(rows, expected_revisions=expected_revisions)

    monkeypatch.setattr(corpus.words.attribute_store, "save_many", conflicting_save)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        corpus.service.apply(corpus.batch, expected_plan_sha256=preview["plan_sha256"])
    assert error.value.code == "attribute_write_conflict"
    assert corpus.words.attribute_store.get(corpus.stored[0]["key"]) == corpus.stored[0]
    assert corpus.words.attribute_store.history(corpus.stored[0]["key"]) == histories[corpus.stored[0]["key"]]
    other = corpus.words.attribute_store.get(corpus.stored[1]["key"])
    assert other["teacher_note"] == "合成并发修改" and other["primary_knowledge"]["id"] == "unknown"


def _args(corpus, tmp_path, **changes):
    candidates = tmp_path / "candidates.json"
    candidates.write_text(json.dumps(corpus.batch, ensure_ascii=False), encoding="utf-8")
    return SimpleNamespace(
        **{
            "workspace": corpus.facade.paths.workspace_root,
            "state": corpus.facade.paths.state_root,
            "candidates": candidates, "report": tmp_path / "report.json",
            "apply": False, "expected_plan_sha256": None,
            **changes,
        }
    )


def test_cli_default_preview_and_nonoverwrite_report(corpus, tmp_path):
    args = _args(corpus, tmp_path)
    before = _history(corpus)
    result = cli.run(args, service_factory=lambda *_: corpus.words)
    assert result["mode"] == "preview" and json.loads(args.report.read_text(encoding="utf-8")) == result
    with pytest.raises(OfflineWordLabelReviewError) as error:
        cli.run(args, service_factory=lambda *_: pytest.fail("must not initialise service"))
    assert error.value.code == "report_exists"
    assert _history(corpus) == before


def test_cli_apply_uses_application_lock_and_metadata_receipt(corpus, tmp_path):
    pytest.importorskip("PySide6")
    preview = corpus.service.preview(corpus.batch)
    args = _args(corpus, tmp_path, apply=True, expected_plan_sha256=preview["plan_sha256"])

    def factory(_workspace, state):
        assert (state / ".desktop-workbench.lock").exists()
        return corpus.words

    result = cli.run(args, service_factory=factory)
    assert result["mode"] == "apply" and result["readback_verified"]
    assert not (args.state / ".desktop-workbench.lock").exists()


def test_cli_never_removes_even_a_stale_lock(corpus, tmp_path):
    pytest.importorskip("PySide6")
    path = corpus.facade.paths.state_root / ".desktop-workbench.lock"
    path.write_text("synthetic stale lock", encoding="utf-8")
    args = _args(corpus, tmp_path, apply=True, expected_plan_sha256="a" * 64)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        cli.run(args, service_factory=lambda *_: pytest.fail("locked service must not open"))
    assert error.value.code == "desktop_lock_exists"
    assert path.read_text(encoding="utf-8") == "synthetic stale lock"
    report = json.loads(args.report.read_text(encoding="utf-8"))
    assert report["completed"] is False and report["error_code"] == "desktop_lock_exists"


def test_cli_rejects_duplicate_json_and_apply_without_digest(corpus, tmp_path):
    args = _args(corpus, tmp_path, apply=True)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        cli.run(args, service_factory=lambda *_: pytest.fail("must not initialise"))
    assert error.value.code == "expected_plan_required"
    args.apply = False
    args.candidates.write_text('{"provenance":"fake", "provenance":"codex_current_session"}', encoding="utf-8")
    with pytest.raises(OfflineWordLabelReviewError) as error:
        cli.run(args, service_factory=lambda *_: pytest.fail("must not initialise"))
    assert error.value.code == "duplicate_json_key"


def test_read_side_constructor_does_not_initialise_facade_or_write_preview_cache(tmp_path, monkeypatch):
    from integrations.deeptutor_shchem_v1.desktop_facade import DesktopWorkbenchFacade

    monkeypatch.setattr(DesktopWorkbenchFacade, "__init__", lambda *_a, **_k: pytest.fail("full facade forbidden"))
    words = cli.read_side_service(tmp_path, tmp_path / "state")
    assert words.preview_cache.save(b"synthetic", "synthetic.docx", {}) is False
    assert not words.preview_cache.root.exists()
    assert not words.attribute_store.path.exists()
    assert not words.state.path.exists()


def test_cold_lookup_hints_select_only_matching_sources_and_remain_untrusted(corpus, monkeypatch):
    corpus.words._locations.clear()
    descriptor = {
        "kind": "desktop_visual_import_v2", "batch_id": "B1",
        "sources": [{"source_sha256": corpus.rows[0]["source_sha256"], "source_file_id": "SOURCE1"}],
    }
    monkeypatch.setattr(corpus.words.state, "snapshot", lambda: {"drafts": {"synthetic-batch": descriptor}})
    corpus.service.preview(corpus.batch)
    assert corpus.words._locations[corpus.rows[0]["key"]] == ("B1", "SOURCE1")
    descriptor["sources"][0]["source_file_id"] = "missing-source"
    with pytest.raises(OfflineWordLabelReviewError):
        corpus.service.preview(corpus.batch)


@pytest.mark.parametrize("destination", [
    "state_lock", "state_nested", "source_library", "knowledge", "git", "candidate", "reserved_name",
])
@pytest.mark.parametrize("apply", [False, True])
def test_report_reserved_paths_are_rejected_before_mkdir_or_open(corpus, tmp_path, monkeypatch, destination, apply):
    args = _args(corpus, tmp_path, apply=apply, expected_plan_sha256="a" * 64 if apply else None)
    destinations = {
        "state_lock": args.state / ".desktop-workbench.lock",
        "state_nested": args.state / "new-report-folder" / "report.json",
        "source_library": args.workspace / "sh-chem-db" / "new-report-folder" / "report.json",
        "knowledge": args.workspace / "knowledge" / "new-report-folder" / "report.json",
        "git": args.workspace / ".git" / "new-report-folder" / "report.json",
        "candidate": args.candidates,
        "reserved_name": tmp_path / "new-report-folder" / "desktop-state.v1.json",
    }
    args.report = destinations[destination]
    before = _history(corpus)
    candidates = args.candidates.read_bytes()
    monkeypatch.setattr(Path, "mkdir", lambda *_a, **_k: pytest.fail("must reject before mkdir"))
    with pytest.raises(OfflineWordLabelReviewError) as error:
        cli.run(args, service_factory=lambda *_: pytest.fail("must reject before initialisation"))
    assert error.value.code == "reserved_report_path"
    assert error.value.operation["attribute_write_attempted"] is False
    assert error.value.operation["commit_status"] == "not_attempted"
    assert not (args.state / ".desktop-workbench.lock").exists()
    assert not (args.report.parent / "new-report-folder").exists()
    assert args.candidates.read_bytes() == candidates
    assert _history(corpus) == before


@pytest.mark.parametrize("mismatch", [False, True])
def test_service_readback_failure_keeps_confirmed_commit_and_plan(corpus, monkeypatch, mismatch):
    preview = corpus.service.preview(corpus.batch)
    real_get = _fail_readback_after_commit(corpus, monkeypatch, mismatch=mismatch)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        corpus.service.apply(corpus.batch, expected_plan_sha256=preview["plan_sha256"])
    receipt = error.value.operation
    assert receipt["stage"] == "attribute_readback"
    assert receipt["attribute_write_attempted"] is True
    assert receipt["commit_status"] == "committed"
    assert receipt["readback_verified"] is False
    assert receipt["plan_sha256"] == preview["plan_sha256"]
    assert receipt["do_not_retry_automatically"] is True
    assert "不要自动重试" in receipt["recovery_hint"]
    actual = real_get([row["key"] for row in corpus.stored])
    assert all(row["primary_knowledge"]["id"] == "K11" for row in actual.values())


@pytest.mark.parametrize("after_commit", [False, True])
def test_save_exception_reports_unknown_instead_of_claiming_no_write(corpus, monkeypatch, after_commit):
    preview = corpus.service.preview(corpus.batch)
    real_save = corpus.words.attribute_store.save_many

    def broken_save(rows, *, expected_revisions):
        if after_commit:
            real_save(rows, expected_revisions=expected_revisions)
        raise OSError("synthetic save did not return normally")

    monkeypatch.setattr(corpus.words.attribute_store, "save_many", broken_save)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        corpus.service.apply(corpus.batch, expected_plan_sha256=preview["plan_sha256"])
    receipt = error.value.operation
    assert receipt["stage"] == "attribute_write"
    assert receipt["attribute_write_attempted"] is True
    assert receipt["commit_status"] == "unknown"
    assert receipt["plan_sha256"] == preview["plan_sha256"]
    assert receipt["do_not_retry_automatically"] is True
    assert "提交状态未知" in receipt["recovery_hint"]
    actual = corpus.words.attribute_store.get(corpus.stored[0]["key"])
    assert (actual["primary_knowledge"]["id"] == "K11") == after_commit


def _fail_readback_after_commit(corpus, monkeypatch, *, mismatch=False):
    real_get = corpus.words.attribute_store.get_many
    real_save = corpus.words.attribute_store.save_many
    committed = False

    def save(rows, *, expected_revisions):
        nonlocal committed
        result = real_save(rows, expected_revisions=expected_revisions)
        committed = True
        return result

    def get(keys):
        if committed:
            if mismatch:
                return {}
            raise OSError("synthetic readback failure")
        return real_get(keys)

    monkeypatch.setattr(corpus.words.attribute_store, "save_many", save)
    monkeypatch.setattr(corpus.words.attribute_store, "get_many", get)
    return real_get


def _argv(args):
    result = [
        "--workspace", str(args.workspace), "--state", str(args.state),
        "--candidates", str(args.candidates), "--report", str(args.report),
    ]
    if args.apply:
        result.extend(["--apply", "--expected-plan-sha256", args.expected_plan_sha256])
    return result


def test_cli_readback_failure_writes_and_prints_committed_receipt(corpus, tmp_path, monkeypatch, capsys):
    pytest.importorskip("PySide6")
    preview = corpus.service.preview(corpus.batch)
    args = _args(corpus, tmp_path, apply=True, expected_plan_sha256=preview["plan_sha256"])
    _fail_readback_after_commit(corpus, monkeypatch)
    monkeypatch.setattr(cli, "read_side_service", lambda *_: corpus.words)
    assert cli.main(_argv(args)) == 1
    stdout = json.loads(capsys.readouterr().out)
    report = json.loads(args.report.read_text(encoding="utf-8"))
    for receipt in (stdout, report):
        assert receipt["commit_status"] == "committed"
        assert receipt["stage"] == "attribute_readback"
        assert receipt["attribute_write_attempted"] is True
        assert receipt["plan_sha256"] == preview["plan_sha256"]
        assert receipt["do_not_retry_automatically"] is True
    assert not (args.state / ".desktop-workbench.lock").exists()


@pytest.mark.parametrize("fault", ["write", "flush", "close", "write_and_close"])
def test_cli_report_persistence_failure_keeps_commit_facts_in_stdout(corpus, tmp_path, monkeypatch, capsys, fault):
    pytest.importorskip("PySide6")
    preview = corpus.service.preview(corpus.batch)
    args = _args(corpus, tmp_path, apply=True, expected_plan_sha256=preview["plan_sha256"])
    close_calls = []
    real_open = Path.open

    class FaultyReport:
        def __init__(self, stream):
            self.stream = stream

        def write(self, text):
            if fault in {"write", "write_and_close"}:
                raise OSError("synthetic report write failure")
            return self.stream.write(text)

        def flush(self):
            if fault == "flush":
                raise OSError("synthetic report flush failure")
            return self.stream.flush()

        def close(self):
            close_calls.append(1)
            self.stream.close()
            if fault in {"close", "write_and_close"}:
                raise OSError("synthetic report close failure")

    def opening(path, mode="r", *positional, **kwargs):
        stream = real_open(path, mode, *positional, **kwargs)
        return FaultyReport(stream) if path == args.report and mode == "x" else stream

    monkeypatch.setattr(Path, "open", opening)
    monkeypatch.setattr(cli, "read_side_service", lambda *_: corpus.words)
    assert cli.main(_argv(args)) == 1
    receipt = json.loads(capsys.readouterr().out)
    stage = "report_write" if fault == "write_and_close" else "report_" + fault
    assert receipt["stage"] == stage
    assert receipt["error_code"] == stage + "_failed"
    assert receipt["operation_stage"] == "readback_verified"
    assert receipt["commit_status"] == "committed"
    assert receipt["attribute_write_attempted"] is True and receipt["readback_verified"] is True
    assert receipt["plan_sha256"] == preview["plan_sha256"]
    assert receipt["do_not_retry_automatically"] is True
    assert close_calls == [1]
    if fault == "write_and_close":
        assert receipt["report_close_failed"] is True
    for old in corpus.stored:
        actual = corpus.words.attribute_store.get(old["key"])
        assert actual["edit_version"] == old["edit_version"] + 1
        assert len(corpus.words.attribute_store.history(old["key"])) == 2
    assert not (args.state / ".desktop-workbench.lock").exists()


def test_failed_readback_and_failed_report_close_keep_both_failures(corpus, tmp_path, monkeypatch, capsys):
    pytest.importorskip("PySide6")
    preview = corpus.service.preview(corpus.batch)
    args = _args(corpus, tmp_path, apply=True, expected_plan_sha256=preview["plan_sha256"])
    _fail_readback_after_commit(corpus, monkeypatch)
    real_persist = cli._persist_report

    def broken_persist(stream, report):
        real_persist(stream, report)
        raise cli._failure(OSError("synthetic close failure"), report,
                           code="report_close_failed", stage="report_close")

    monkeypatch.setattr(cli, "_persist_report", broken_persist)
    monkeypatch.setattr(cli, "read_side_service", lambda *_: corpus.words)
    assert cli.main(_argv(args)) == 1
    receipt = json.loads(capsys.readouterr().out)
    assert receipt["commit_status"] == "committed" and receipt["readback_verified"] is False
    assert receipt["error_code"] == "report_close_failed"
    assert receipt["operation_error_code"] == "attribute_readback_failed"
    assert receipt["operation_stage"] == "attribute_readback"
