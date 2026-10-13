"""Original Word sections are imported whole, independently of old previews."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from threading import Event, RLock
from types import SimpleNamespace

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
import pytest

from integrations.deeptutor_shchem_v1 import desktop_handout_knowledge as handout
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (block,
    digest, resource, validate_resource)
from integrations.deeptutor_shchem_v1.desktop_preparation_images import (image_info,
    PreparationImageStore, MAX_IMAGES)
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import PreparationSourcesService
from integrations.deeptutor_shchem_v1.desktop_source_quality import QUALITY_PATH
from integrations.deeptutor_shchem_v1.reader_cancellation import ReadCancelled, read_cancel_scope


def png(number=1):
    buffer = BytesIO(); Image.new('RGB', (60, 30), (number, 19, 63)).save(buffer, 'PNG')
    return buffer.getvalue()


def word_bytes(document):
    buffer = BytesIO(); document.save(buffer); return buffer.getvalue()


def make_facade(data, root, role='handout'):
    """Small source-bound fixture for integration/UI tests; no model or state IO.

    ``lecture`` is a test helper alias for an explicitly imported handout role.
    It does not infer source roles from a filename.
    """
    if role == 'lecture': role = 'handout'
    root = Path(root); original = root / 'original.docx'
    root.mkdir(parents=True, exist_ok=True); original.write_bytes(data)
    value = SimpleNamespace(data=data, name='分子间作用力讲义.docx', role=role,
        reads=0, source_hash=sha256(data).hexdigest())
    def source(*args):
        value.reads += 1
        if value.data is None: raise ValueError('原Word已删除')
        return SimpleNamespace(content=value.data, filename=value.name, role=value.role,
            source_sha256=sha256(value.data).hexdigest(), effective_source_file_id='SOURCE')
    facade = SimpleNamespace(paths=SimpleNamespace(workspace_root=root/'workspace',
        task_root=root/'task'), value=value, original=original,
        list_imported_word_batches=lambda: [SimpleNamespace(batch_id='BATCH')],
        imported_word_sources=lambda _: [{'source_id':'SOURCE', 'source_name':value.name,
            'source_sha256':value.source_hash, 'role':value.role}],
        _imported_word_source=source,
        imported_word_path=lambda *args: str(original))
    return facade


def request_for(facade, section=None, title='分子间的作用力怎样影响性质？', teacher_only=False):
    source = handout.handout_knowledge_sources(facade)['sources'][0]
    outline = handout.handout_knowledge_outline(facade, source)
    if isinstance(section, str): chosen = next(s for s in outline['sections'] if s['title'] == section)
    elif isinstance(section, dict): chosen = section
    elif type(section) is int: chosen = outline['sections'][section]
    else: chosen = next((s for s in outline['sections'] if s['level'] == 3), outline['sections'][0])
    return {'source': outline['source'], 'outline_revision': outline['outline_revision'],
        'section': deepcopy(chosen), 'title': title, 'teacher_only': teacher_only}


def test_whole_knowledge_unit_retains_memory_blanks_pictures_and_comparison_table(tmp_path):
    doc = Document(); doc.add_heading('知识点2 通用高分子材料', 3)
    doc.add_paragraph('得分速记'); doc.add_paragraph('2．合成纤维')
    doc.add_paragraph('（1）纤维的分类')
    doc.add_paragraph().add_run().add_picture(BytesIO(png()))
    doc.add_paragraph('（2）合成纤维的性能和应用')
    doc.add_paragraph('①六大纶包括：       、       、       。')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '纤维'; table.cell(0, 1).text = '性能'
    table.cell(1, 0).text = '聚酯纤维'; table.cell(1, 1).text = '耐磨'
    doc.add_paragraph('3．合成橡胶'); doc.add_paragraph('（1）橡胶的分类')
    doc.add_paragraph().add_run().add_picture(BytesIO(png(2)))
    doc.add_paragraph('硫化后保留弹性与强度。')
    doc.add_heading('考向1 高分子材料', 3)
    doc.add_paragraph('【例1】下列性质正确的是____。'); doc.add_paragraph('【答案】A')
    facade = make_facade(word_bytes(doc), tmp_path/'whole-knowledge')
    request = request_for(facade, '知识点2 通用高分子材料')
    outline = handout.handout_knowledge_outline(facade, request['source'])
    result = handout.import_handout_knowledge(facade, [request], [], [])
    snapshot = result['resources'][0]['original']['evidence']['section_snapshot']
    assert snapshot['section'] == request['section']
    assert [b['index'] for b in snapshot['blocks']] == list(range(1, 13))
    assert [b['text'] for b in snapshot['blocks']] == [b['text'] for b in outline['preview']['blocks'][:12]]
    assert '       ' in snapshot['blocks'][6]['text']
    assert snapshot['blocks'][7]['table_rows'][1]['cells'][0]['text'] == '聚酯纤维'
    assert len(result['image_assets']) == 2


def test_real_task_inside_a_numbered_knowledge_unit_still_blocks_student_knowledge(tmp_path):
    doc = Document(); doc.add_heading('知识点2 通用高分子材料', 3)
    doc.add_paragraph('得分速记'); doc.add_paragraph('2．合成纤维')
    doc.add_paragraph('（1）纤维的分类')
    doc.add_paragraph('（2）请写出这种纤维的单体。')
    doc.add_paragraph('【答案】原来的完整参考。')
    facade = make_facade(word_bytes(doc), tmp_path/'real-task')
    with pytest.raises(ValueError, match='已识别的题目或答案'):
        handout.import_handout_knowledge(facade, [request_for(facade)], [], [])
    assert not facade.paths.task_root.exists()


@pytest.mark.parametrize('filename,bounds', [
    ('第08讲 高分子有机物（复习讲义）（原卷版）.docx', (211, 255)),
    ('第08讲 高分子有机物（复习讲义）（解析版）.docx', (254, 297)),
])
def test_actual_pkg072_full_materials_section_passes_unchanged_role_guard(filename, bounds):
    root = Path('E:/项目/上海高中化学等级考')
    source = root/'sh-chem-db/.intake/2026-07-30-user-teaching-pack/expanded/PKG-072'/filename
    if not source.is_file():
        pytest.skip('Exact local source is unavailable')
    raw = source.read_bytes(); before = (sha256(raw).hexdigest(), source.stat().st_mtime_ns)
    preview = PreparationSourcesService(root).word_preview_bytes(raw, filename)
    original = deepcopy(preview)
    sections = handout._headings(Document(BytesIO(raw)), preview)
    section = next(s for s in sections if s['title'] == '知识点2 通用高分子材料')
    assert (section['start'], section['end']) == bounds
    handout._check_roles(SimpleNamespace(), None, preview, section)
    assert preview == original
    assert (sha256(source.read_bytes()).hexdigest(), source.stat().st_mtime_ns) == before


@pytest.fixture
def lecture(tmp_path):
    doc = Document(); doc.add_heading('考点二 分子间作用力', 2)
    doc.add_heading('知识点1 范德华力', 3)
    doc.add_paragraph('1. 特征：分子之间存在相互作用力。')
    p = doc.add_paragraph('SO'); p.add_run('4').font.subscript = True
    p.add_run('2−').font.superscript = True
    math = OxmlElement('m:oMath'); fraction = OxmlElement('m:f')
    for tag, text in [('m:num', '1'), ('m:den', '2')]:
        part = OxmlElement(tag); run = OxmlElement('m:r'); token = OxmlElement('m:t')
        token.text = text; run.append(token); part.append(run); fraction.append(part)
    math.append(fraction); p._p.append(math)
    p = doc.add_paragraph('完整原图：'); p.add_run().add_picture(BytesIO(png()))
    p.add_run('并比较：').add_picture(BytesIO(png()))
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = '分子间作用力比较'
    table.cell(1, 0).text = '范德华力'; table.cell(1, 1).text = '化学键'
    doc.add_heading('知识点2 氢键', 3)
    doc.add_paragraph('完整第二知识段，保留必要的条件和结论。')
    doc.add_paragraph().add_run().add_picture(BytesIO(png(2)))
    doc.add_heading('考向1 例题', 3)
    doc.add_paragraph('【例题1】下列关于氢键的说法正确的是（ ）')
    doc.add_paragraph('A．甲 B．乙 C．丙 D．丁')
    doc.add_paragraph('【答案】A'); doc.add_paragraph('【解析】完整解析。')
    facade = make_facade(word_bytes(doc), tmp_path/'case')
    return SimpleNamespace(facade=facade, doc=doc, first=request_for(facade),
        second=request_for(facade, '知识点2 氢键'), original=facade.original.read_bytes())


def images_on_disk(facade):
    root = facade.paths.task_root / 'preparation-v1' / 'images'
    return list(root.glob('*.image')) if root.exists() else []


def test_catalog_is_metadata_only_and_target_outline_does_not_change_old_preview(lecture):
    f = lecture; f.facade.value.reads = 0
    catalog = handout.handout_knowledge_sources(f.facade)
    assert f.facade.value.reads == 0 and catalog['issues'] == []
    raw = f.facade.value.data
    old = PreparationSourcesService(f.facade.paths.workspace_root).word_preview_bytes(raw, f.facade.value.name)
    result = handout.handout_knowledge_outline(f.facade, catalog['sources'][0])
    assert result['preview'] == old and old['sections'] == [{'start':1, 'level':2,
        'title':'考点二 分子间作用力', 'end':14}]
    assert [s['level'] for s in result['sections']] == [2, 3, 3, 3]
    assert handout.handout_knowledge_outline(f.facade, result['source']) == result
    assert not f.facade.paths.task_root.exists() and f.facade.original.read_bytes() == f.original


def test_actual_heading_levels_inheritance_direct_outline_and_image_only_headings(tmp_path):
    from docx.enum.style import WD_STYLE_TYPE
    doc = Document(); p = doc.add_paragraph(style='Heading 1'); p.add_run().add_picture(BytesIO(png()))
    doc.add_paragraph().add_run().add_picture(BytesIO(png(2)))
    doc.add_paragraph(style='Heading 1')  # Empty heading still delimits the previous section.
    style = doc.styles.add_style('我的三级栏目', WD_STYLE_TYPE.PARAGRAPH); style.base_style = doc.styles['Heading 3']
    doc.add_paragraph('真实三级栏目', style=style); doc.add_paragraph('完整正文。')
    p = doc.add_paragraph('真实第九级栏目'); node = OxmlElement('w:outlineLvl'); node.set(qn('w:val'), '8')
    p._p.get_or_add_pPr().append(node); doc.add_paragraph('第九级正文。')
    toc = doc.styles.add_style('TOC 1', WD_STYLE_TYPE.PARAGRAPH); toc.base_style = doc.styles['Heading 1']
    doc.add_paragraph('目录条目', style=toc)
    facade = make_facade(word_bytes(doc), tmp_path/'heading')
    outline = handout.handout_knowledge_outline(facade, handout.handout_knowledge_sources(facade)['sources'][0])
    assert [(s['start'],s['end'],s['level']) for s in outline['sections']] == [(1,2,1),(4,8,3),(6,8,9)]
    assert outline['sections'][0]['title'] == '【待查看原文：图片或图形】'
    result = handout.import_handout_knowledge(facade, [request_for(facade, 0)], [], [])
    assert len(result['image_assets']) == 2 and len(result['resources'][0]['blocks']) == 4


def test_multiple_whole_sections_preserve_original_text_scripts_math_tables_and_image_occurrences(lecture):
    f = lecture; result = handout.import_handout_knowledge(f.facade, [f.first,f.second], [], [])
    assert set(result) == {'resources','image_assets'} and len(result['resources']) == 2
    first = result['resources'][0]; original = first['original']; snapshot = original['evidence']['section_snapshot']
    assert first['title'] == f.first['title'] and original['title'].endswith('知识点1 范德华力')
    assert [(b['source_index'],b['value']) for b in original['blocks'] if b['kind']=='text'] == [
        (b['index'], b['text']) for b in snapshot['blocks']]
    chemical = next(b for b in original['blocks'] if 'SO' in b['value'])
    assert 'SO_{4}^{2−}' in chemical['value'] and '\\frac{1}{2}' in chemical['value']
    assert [s['script'] for s in chemical['presentation']['scripts']] == ['subscript','superscript']
    table = next(b for b in snapshot['blocks'] if 'table_rows' in b)
    assert table['table_rows'][0]['cells'][0]['column_span'] == 2
    picture = next(b for b in snapshot['blocks'] if b['images'])
    assert len(picture['images']) == 2 and picture['images'][0]['asset_id'] == picture['images'][1]['asset_id']
    refs = next(b for b in original['blocks'] if b['source_index']==picture['index'] and b['kind']=='text')['presentation']['images']
    assert len(refs) == 2 and len(result['image_assets']) == 2
    assert all(b['role']=='knowledge' for item in result['resources'] for b in item['blocks'])
    assert all('原Word教案图文参考' not in b['value'] and '原文缺口与待核对提醒' not in b['value']
        for b in original['blocks'] if b['kind']=='text')
    for item in result['resources']:
        assert validate_resource(item) == item
        assert item['candidate_only'] and not item['human_reviewed'] and not item['publication_allowed']
    store = PreparationImageStore(f.facade.paths.task_root/'preparation-v1/images')
    assert {store.load(a) for a in result['image_assets']} == {png(),png(2)}
    assert f.facade.original.read_bytes() == f.original


def test_duplicate_batch_and_reimport_keep_existing_edited_resource_and_picture_metadata(lecture):
    f = lecture; result = handout.import_handout_knowledge(f.facade, [f.first,f.first], [], [])
    assert len(result['resources']) == 1
    item = result['resources'][0]; item['locked'] = False; item['title'] = '教师已编排的问题'
    next(b for b in item['blocks'] if b['kind']=='text')['value'] = '教师编辑的讲解。'
    assets = deepcopy(result['image_assets']); assets[0]['caption'] = '已编辑的图题'
    again = handout.import_handout_knowledge(f.facade, [{**f.first,'title':'另一个问题标题'}], assets, [item])
    assert again == {'resources':[],'image_assets':assets}
    assert item['title'] == '教师已编排的问题' and item['blocks'][0]['value'] == '教师编辑的讲解。'


@pytest.mark.parametrize('teacher', ['explicit','document'])
def test_original_teacher_scope_is_retained(lecture, teacher):
    f = lecture; request = deepcopy(f.first); existing, resources = [], []
    if teacher == 'explicit': request['teacher_only'] = True
    if teacher == 'document':
        f.facade.value.role = 'answer'; request = request_for(f.facade)
    result = handout.import_handout_knowledge(f.facade,[request],existing,resources)
    assert all(b['role']=='teacher_note' for b in result['resources'][0]['blocks'])


@pytest.mark.parametrize('section', ['考点二 分子间作用力','考向1 例题'])
def test_known_question_or_answer_range_requires_complete_question_entry(lecture, section):
    f = lecture
    with pytest.raises(ValueError,match='完整题'):
        handout.import_handout_knowledge(f.facade,[request_for(f.facade,section)],[],[])
    assert not images_on_disk(f.facade) and not f.facade.paths.task_root.exists()


def test_unindexed_answer_marker_cannot_become_student_knowledge(tmp_path):
    doc = Document(); doc.add_heading('知识点1',3); doc.add_paragraph('【解析】独立答案内容。')
    facade = make_facade(word_bytes(doc),tmp_path/'answer')
    with pytest.raises(ValueError,match='完整题'):
        handout.import_handout_knowledge(facade,[request_for(facade)],[],[])


@pytest.mark.parametrize('change', ['bytes','deleted','name','role','outline_revision','section_end','section_title','revision'])
def test_source_or_heading_change_fails_before_any_image_file(lecture,change):
    f = lecture; request = deepcopy(f.first)
    if change == 'bytes':
        f.doc.add_paragraph('改变后的正文'); f.facade.value.data = word_bytes(f.doc)
    elif change == 'deleted': f.facade.value.data = None
    elif change == 'name': f.facade.value.name = '另一个Word.docx'
    elif change == 'role': f.facade.value.role = 'answer'
    elif change == 'outline_revision': request['outline_revision'] = 'a'*64
    elif change == 'section_end': request['section']['end'] -= 1
    elif change == 'section_title': request['section']['title'] = '改过的原标题'
    else: request['source']['source_revision'] = 'a'*64
    with pytest.raises(ValueError): handout.import_handout_knowledge(f.facade,[request],[],[])
    assert not f.facade.paths.task_root.exists() and f.facade.original.read_bytes()==f.original


@pytest.mark.parametrize('change', ['text','index','image','table','teacher_role'])
def test_resealing_resource_cannot_replace_original_blocks_or_scope(lecture,change):
    f = lecture; item = handout.import_handout_knowledge(f.facade,[{**f.first,'teacher_only':True}],[],[])['resources'][0]
    altered = deepcopy(item); original = altered['original']
    if change == 'text': original['blocks'][0]['value'] = '篡改后的原文'
    elif change == 'index': original['blocks'][0]['source_index'] += 1
    elif change == 'image': next(b for b in original['blocks'] if b['kind']=='image')['value'] = 'IMG-'+'a'*64
    elif change == 'table': next(b for b in original['evidence']['section_snapshot']['blocks'] if 'table_rows' in b)['table_rows'][0]['cells'][0]['text'] = '改后表格'
    else:
        for b in original['blocks']: b['role']='knowledge'
    altered['original_sha256'] = digest(original); altered['blocks'] = deepcopy(original['blocks'])
    with pytest.raises(ValueError): validate_resource(altered)


def test_whole_batch_rejects_known_source_quality_before_writing_first_valid_section(lecture):
    f = lecture; target = f.facade.paths.workspace_root / QUALITY_PATH
    target.parent.mkdir(parents=True)
    target.write_text(json.dumps({'schema_version':1,'sources':[{
        'source_sha256':f.first['source']['source_sha256'],'source_revision':f.first['source']['source_revision'],
        'issues':[{'id':'known','block_indices':[f.second['section']['start']+1],
            'summary':'原来源有已知问题','suggested_correction':'保持原文件并检查内容'}]}]},ensure_ascii=False),encoding='utf-8')
    with pytest.raises(ValueError,match='来源问题'):
        handout.import_handout_knowledge(f.facade,[f.first,f.second],[],[])
    assert not images_on_disk(f.facade) and not f.facade.paths.task_root.exists()


def test_missing_visual_is_not_a_silent_text_only_import(tmp_path):
    doc=Document();doc.add_heading('知识点1',3)
    doc.add_paragraph('缺失原图：').add_run()._r.append(OxmlElement('w:object'))
    facade=make_facade(word_bytes(doc),tmp_path/'missing')
    with pytest.raises(ValueError,match='原图缺口'):
        handout.import_handout_knowledge(facade,[request_for(facade)],[],[])
    assert not facade.paths.task_root.exists()


def test_cancelled_outline_or_import_never_creates_cache_or_picture_files(lecture):
    f=lecture;event=Event();event.set()
    with read_cancel_scope(event,check_boundaries=False):
        for operation in (lambda: handout.handout_knowledge_outline(f.facade,f.first['source']),
                lambda: handout.import_handout_knowledge(f.facade,[f.first],[],[])):
            with pytest.raises(ReadCancelled): operation()
    assert not f.facade.paths.task_root.exists()


def test_merged_capacity_is_checked_before_any_file(lecture):
    f=lecture; existing=[]
    for number in range(40,40+MAX_IMAGES):
        data=png(number); token=sha256(data).hexdigest()
        existing.append({'asset_id':'IMG-'+token,'sha256':token,'caption':'当前图',
            'source':'当前来源','purpose':'讲解',**image_info(data)})
    with pytest.raises(ValueError,match='最多'):
        handout.import_handout_knowledge(f.facade,[f.first],existing,[])
    assert not f.facade.paths.task_root.exists()


def test_image_commit_failure_rolls_back_only_new_batch_files(lecture,monkeypatch):
    from integrations.deeptutor_shchem_v1 import desktop_word_source_reference as reference
    f=lecture; original_link=reference.os.link; calls=[]
    def fail_second(a,b):
        calls.append(b)
        if len(calls)==2: raise OSError('fixture write failure')
        return original_link(a,b)
    monkeypatch.setattr(reference.os,'link',fail_second)
    with pytest.raises(OSError):handout.import_handout_knowledge(f.facade,[f.first,f.second],[],[])
    assert len(calls)==2 and not images_on_disk(f.facade) and f.facade.original.read_bytes()==f.original


def test_final_check_detects_source_change_during_batch_without_partial_handoff(lecture,monkeypatch):
    f=lecture; original=handout.WordSourceReferenceService._compile; calls=[]
    def changed(service,*args,**kwargs):
        result=original(service,*args,**kwargs);calls.append(1)
        if len(calls)==1:
            f.doc.add_paragraph('确认期间来源变化');f.facade.value.data=word_bytes(f.doc)
        return result
    monkeypatch.setattr(handout.WordSourceReferenceService,'_compile',changed)
    with pytest.raises(ValueError):handout.import_handout_knowledge(f.facade,[f.first],[],[])
    assert not f.facade.paths.task_root.exists()


def test_scoped_saved_question_range_cannot_be_imported_as_knowledge(lecture):
    f=lecture; preview=handout.handout_knowledge_outline(f.facade,f.first['source'])['preview']
    original=f.facade._imported_word_source('BATCH','SOURCE'); calls=[]
    def catalog(locations,read_only=False):
        calls.append((locations,read_only))
        row={'batch_id':'BATCH','archive_source_id':'SOURCE','source_sha256':preview['source_sha256'],
            'source_revision':preview['revision'],'context_blocks':[],
            'question_blocks':[preview['blocks'][2]],'answer_blocks':[]}
        return {'items':[row]}, {preview['source_sha256']:(original,'BATCH',preview,[])}
    words=SimpleNamespace(_lock=RLock(),_catalog=catalog); f.facade._word_questions=lambda:words
    with pytest.raises(ValueError,match='完整题'):
        handout.import_handout_knowledge(f.facade,[f.first],[],[])
    assert calls==[({'BATCH':{'SOURCE'}},True)] and not f.facade.paths.task_root.exists()


def test_crop_and_effects_keep_complete_embedded_pixels_and_source_display_evidence(lecture,tmp_path):
    f=lecture; node=next(n for n in f.doc._element.body.iter() if n.tag==qn('w:drawing'))
    blip=next(n for n in node.iter() if n.tag==qn('a:blip'))
    bright=OxmlElement('a:lum');bright.set('bright','-20000');blip.append(bright)
    fill=next(n for n in node.iter() if n.tag==qn('pic:blipFill'))
    crop=OxmlElement('a:srcRect');crop.set('b','4995');fill.insert(1,crop)
    facade=make_facade(word_bytes(f.doc),tmp_path/'effects')
    result=handout.import_handout_knowledge(facade,[request_for(facade)],[],[])
    item=result['resources'][0];refs=[r for b in item['original']['evidence']['section_snapshot']['blocks'] for r in b['images']]
    assert len(refs)==2 and refs[0]['container_xml_sha256']!=refs[1]['container_xml_sha256']
    assert all(r['width_mm']>0 and r['height_mm']>0 for r in refs)
    assert len(result['image_assets'])==1
    assert PreparationImageStore(facade.paths.task_root/'preparation-v1/images').load(result['image_assets'][0])==png()


def test_actual_archive_facade_uses_scoped_question_catalog_readonly(tmp_path,desktop_paths):
    from test_desktop_visual_import_facade import FakeProviderStore, _facade
    doc=Document();doc.add_heading('知识点1 范德华力',3)
    doc.add_paragraph('1. 特征：广泛存在。');doc.add_paragraph('2. 影响因素：相对分子质量。')
    original=tmp_path/'明确知识讲义.docx';doc.save(original);raw=original.read_bytes()
    facade=_facade(desktop_paths,FakeProviderStore(configured=False))
    facade.save_visual_import_batch(handout_files=(original,),source_type='教师讲义')
    request=request_for(facade); before=deepcopy(facade._state.snapshot())
    files_before={str(p) for p in desktop_paths.state_root.rglob('*') if p.is_file()}
    result=handout.import_handout_knowledge(facade,[request],[],[])
    assert len(result['resources'])==1 and not result['image_assets']
    assert facade._state.snapshot()==before and original.read_bytes()==raw
    assert {str(p) for p in desktop_paths.state_root.rglob('*') if p.is_file()}==files_before


from test_desktop_visual_import_facade import desktop_paths as desktop_paths
