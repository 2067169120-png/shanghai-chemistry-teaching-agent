"""Whole lecture materials compile in isolation and publish as one batch."""
from copy import deepcopy
from dataclasses import replace
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from test_desktop_visual_import_facade import (FakeProviderStore, _facade, _png,
    desktop_paths as desktop_paths)
from test_knowledge_question_links import linked, scoped_lecture, notes, section_fixture, asset_fixture, shelf_fixture
from integrations.deeptutor_shchem_v1 import desktop_handout_material_handoff as handoff
from integrations.deeptutor_shchem_v1.desktop_handout_knowledge import handout_knowledge_sources, handout_knowledge_outline
from integrations.deeptutor_shchem_v1.desktop_handout_sources import handout_source_context
from integrations.deeptutor_shchem_v1.desktop_handout_questions import HandoutQuestionFacade
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
from integrations.deeptutor_shchem_v1.reader_cancellation import TaskCancellationEvent, ReadCancelled, read_cancel_scope


@pytest.fixture(params=[True, False], ids=['with-images', 'text-only'])
def mixed(desktop_paths, tmp_path, monkeypatch, request):
    monkeypatch.setattr('integrations.deeptutor_shchem_v1.desktop_word_questions.load_attribute_catalog',
                        lambda *a, **k: {'knowledge_points': [], 'nodes': []})
    doc = Document(); doc.add_heading('课堂讲义', 1)
    doc.add_heading('知识点1 分子间的作用', 2)
    doc.add_paragraph('完整知识说明及适用条件。')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '作用'; table.cell(0, 1).text = '性质'
    table.cell(1, 0).text = '氢键'; table.cell(1, 1).text = '物性'
    if request.param: doc.add_paragraph().add_run().add_picture(BytesIO(_png('blue')))
    doc.add_heading('例题', 2)
    doc.add_paragraph('【即学即练1】根据完整条件回答问题。')
    if request.param: doc.add_paragraph().add_run().add_picture(BytesIO(_png('red')))
    doc.add_paragraph('【答案】完整参考答案及解析。')
    if request.param: doc.add_paragraph().add_run().add_picture(BytesIO(_png('red')))
    path = tmp_path / '混合课堂讲义.docx'; doc.save(path)
    provider = FakeProviderStore(configured=False); facade = _facade(desktop_paths, provider)
    batch = facade.save_visual_import_batch(handout_files=(path,), source_type='教师讲义')
    metadata = next(s for s in handout_knowledge_sources(facade)['sources'] if s['batch_id'] == batch.batch_id)
    outline = handout_knowledge_outline(facade, metadata); source = outline['source']
    section = next(s for s in outline['sections'] if s['title'].startswith('知识点'))
    knowledge = {'source': deepcopy(source), 'outline_revision': outline['outline_revision'],
        'section': section, 'title': '分子间的作用怎样影响物性？', 'teacher_only': False}
    context = handout_source_context(facade, source)
    images = [{'source': context['images'][0], 'title': '这张知识网怎样串起本课？',
        'purpose': '比较作用与物性', 'teacher_only': False}] if request.param else []
    row = HandoutQuestionFacade(facade, source).word_question_catalog()['items'][0]
    question = {'key': row['key'], 'revision': row['revision'], 'points': 2,
        'source_sha256': row['source_sha256'], 'scoped_source_document': deepcopy(source)}
    value = {'schema': handoff.SCHEMA, 'source': source, 'knowledge_requests': [knowledge],
        'image_selections': images, 'question_selections': [question]}
    return facade, value, provider


def images(facade):
    return {p.name: (p.read_bytes(), p.stat().st_mtime_ns)
        for p in (facade.paths.task_root / 'preparation-v1' / 'images').glob('*.image')}


@pytest.fixture
def ordered_mixed(desktop_paths, tmp_path, monkeypatch):
    monkeypatch.setattr('integrations.deeptutor_shchem_v1.desktop_word_questions.load_attribute_catalog',
                        lambda *a, **k: {'knowledge_points': [], 'nodes': []})
    doc = Document(); doc.add_heading('有两节知识的完整课堂讲义', 1)
    for index, color in enumerate(('blue', 'green'), 1):
        doc.add_heading(f'知识点{index} 完整知识栏目', 2)
        doc.add_paragraph(f'第{index}节知识的完整条件和说明。')
        doc.add_paragraph().add_run().add_picture(BytesIO(_png(color)))
    doc.add_heading('例题', 2)
    doc.add_paragraph('【即学即练1】根据完整条件回答问题。')
    doc.add_paragraph().add_run().add_picture(BytesIO(_png('red')))
    doc.add_paragraph('【答案】完整参考答案及解析。')
    doc.add_paragraph().add_run().add_picture(BytesIO(_png('purple')))
    path = tmp_path/'交错选材讲义.docx'; doc.save(path)
    provider = FakeProviderStore(configured=False); facade = _facade(desktop_paths, provider)
    batch = facade.save_visual_import_batch(handout_files=(path,), source_type='教师讲义')
    metadata = next(s for s in handout_knowledge_sources(facade)['sources'] if s['batch_id'] == batch.batch_id)
    outline = handout_knowledge_outline(facade, metadata); source = outline['source']
    knowledge = [{'source': deepcopy(source), 'outline_revision': outline['outline_revision'],
        'section': section, 'title': '相同的教师问题标题', 'teacher_only': False}
        for section in outline['sections'] if section['title'].startswith('知识点')]
    context = handout_source_context(facade, source)
    selected_images = [{'source': image, 'title': '相同的教师问题标题',
        'purpose': f'课堂读图用途{index}', 'teacher_only': False}
        for index, image in enumerate(context['images'][:2])]
    row = HandoutQuestionFacade(facade, source).word_question_catalog()['items'][0]
    question = {'key': row['key'], 'revision': row['revision'], 'points': 2,
        'source_sha256': row['source_sha256'], 'scoped_source_document': deepcopy(source)}
    request = {'schema': handoff.SCHEMA, 'source': source, 'knowledge_requests': knowledge,
        'image_selections': selected_images, 'question_selections': [question],
        'selection_order': [{'kind': 'image', 'index': 0}, {'kind': 'knowledge', 'index': 0},
            {'kind': 'question', 'index': 0}, {'kind': 'image', 'index': 1}, {'kind': 'knowledge', 'index': 1}]}
    return facade, request, provider


def test_interleaved_source_order_is_preserved_after_single_native_batch_compilation(ordered_mixed, monkeypatch):
    facade, request, provider = ordered_mixed; before = deepcopy(request)
    state = facade.state_store.snapshot(); counts = {'knowledge': 0, 'question': 0}
    knowledge, questions = handoff.import_handout_knowledge, handoff.import_questions
    def counted_knowledge(*args, **kwargs):
        counts['knowledge'] += 1
        assert len(args[1]) == 2
        return knowledge(*args, **kwargs)
    def counted_questions(*args, **kwargs):
        counts['question'] += 1
        assert len(args[1]) == 1
        return questions(*args, **kwargs)
    monkeypatch.setattr(handoff, 'import_handout_knowledge', counted_knowledge)
    monkeypatch.setattr(handoff, 'import_questions', counted_questions)
    publication = handoff.prepare_handout_materials(facade, request, [], [])
    assert counts == {'knowledge': 1, 'question': 1}
    assert not images(facade)
    result = publication.publish(); publication.accept()
    assert counts == {'knowledge': 2, 'question': 2}
    resources = result['resources']
    assert [r['original']['kind'] for r in resources] == [
        'handout_image', 'handout_knowledge', 'complete_question', 'handout_image', 'handout_knowledge']
    assert resources[0]['original']['evidence']['handout_source'] == request['image_selections'][0]['source']
    assert resources[3]['original']['evidence']['handout_source'] == request['image_selections'][1]['source']
    assert [resources[i]['original']['evidence']['section_snapshot']['section'] for i in (1, 4)] == [
        r['section'] for r in request['knowledge_requests']]
    assert any(b['role'] == 'answer' for b in resources[2]['blocks'])
    assert request == before and facade.state_store.snapshot() == state and provider.borrow_calls == 0


def test_interleaved_duplicate_inputs_skip_existing_work_copies_in_place(ordered_mixed):
    facade, request, _ = ordered_mixed
    first = handoff.prepare_handout_materials(facade, request, [], [])
    complete = first.publish(); first.accept()
    # Existing picture and complete question can be edited without creating a
    # second occurrence or moving them among the newly selected knowledge.
    existing = deepcopy([complete['resources'][0], complete['resources'][2]])
    for row in existing:
        row['title'] = '教师已经修改的旧课堂问题'
        for b in row['blocks']:
            if b['kind'] == 'text': b['value'] += '\n已有讲解不变。'
    before = deepcopy(existing); files = images(facade)
    publication = handoff.prepare_handout_materials(facade, request, complete['image_assets'], existing)
    result = publication.publish(); publication.accept()
    assert [r['original']['kind'] for r in result['resources']] == [
        'handout_knowledge', 'handout_image', 'handout_knowledge']
    assert result['resources'][1]['original']['evidence']['handout_source'] == request['image_selections'][1]['source']
    assert existing == before and images(facade) == files


def test_interleaved_same_picture_different_purposes_remain_two_ordered_uses(ordered_mixed):
    facade, request, _ = ordered_mixed
    request['image_selections'][1]['source'] = deepcopy(request['image_selections'][0]['source'])
    publication = handoff.prepare_handout_materials(facade, request, [], [])
    result = publication.publish(); publication.accept()
    first, later = result['resources'][0], result['resources'][3]
    assert first['blocks'][0]['value'] == later['blocks'][0]['value']
    assert first['original']['evidence']['handout_source'] == later['original']['evidence']['handout_source']
    assert [b['value'] for b in first['blocks'] if b['kind'] == 'text'] == ['课堂读图用途0']
    assert [b['value'] for b in later['blocks'] if b['kind'] == 'text'] == ['课堂读图用途1']


def test_duplicate_inputs_inside_one_batch_own_empty_groups_without_reordering_other_materials(ordered_mixed):
    facade, request, _ = ordered_mixed
    request['image_selections'][1] = deepcopy(request['image_selections'][0])
    duplicate = deepcopy(request['knowledge_requests'][0])
    duplicate['title'] = '重复选择的标题不替换第一份作品'
    request['knowledge_requests'].append(duplicate)
    request['selection_order'].insert(0, {'kind': 'knowledge', 'index': 2})
    publication = handoff.prepare_handout_materials(facade, request, [], [])
    result = publication.publish(); publication.accept()
    assert [r['original']['kind'] for r in result['resources']] == [
        'handout_image', 'handout_knowledge', 'complete_question', 'handout_knowledge']
    assert result['resources'][1]['title'] == request['knowledge_requests'][0]['title']
    assert result['resources'][3]['original']['evidence']['section_snapshot']['section'] == request['knowledge_requests'][1]['section']


def test_image_order_before_complete_question_still_keeps_answer_picture_teacher_only(ordered_mixed):
    facade, request, _ = ordered_mixed
    context = handout_source_context(facade, request['source'])
    request['image_selections'][0]['source'] = context['images'][-1]
    publication = handoff.prepare_handout_materials(facade, request, [], [])
    result = publication.publish(); publication.accept()
    answer_picture = result['resources'][0]
    assert answer_picture['original']['kind'] == 'handout_image'
    assert all(b['role'] == 'teacher_note' for b in answer_picture['blocks'])
    assert all(not r['publication_allowed'] and not r['human_reviewed'] for r in result['resources'])


def test_exact_compiler_group_with_several_resources_keeps_internal_order(ordered_mixed):
    facade, request, _ = ordered_mixed
    publication = handoff.prepare_handout_materials(facade, request, [], [])
    rows = publication.result['resources']
    extra = deepcopy(rows[1]); extra['id'] = 'CR-'+'f'*32
    ordered = handoff._ordered_additions(request, [rows[1], extra, rows[4]],
        [[rows[0]], [rows[3]]], [rows[2]])
    assert [r['id'] for r in ordered] == [rows[0]['id'], rows[1]['id'], extra['id'],
        rows[2]['id'], rows[3]['id'], rows[4]['id']]
    publication.discard()


@pytest.mark.parametrize('mutation', ['not_list', 'missing', 'duplicate', 'out_of_range', 'negative',
    'boolean_index', 'wrong_kind', 'extra_field', 'missing_field'])
def test_selection_order_is_a_strict_complete_unique_permutation(mixed, mutation):
    facade, request, _ = mixed
    request['selection_order'] = [{'kind': kind, 'index': index}
        for kind, name in handoff._SELECTION_ARRAYS.items() for index in range(len(request[name]))]
    order = request['selection_order']
    if mutation == 'not_list': request['selection_order'] = {'kind': 'knowledge', 'index': 0}
    elif mutation == 'missing': order.pop()
    elif mutation == 'duplicate': order[-1] = deepcopy(order[0])
    elif mutation == 'out_of_range': order[0]['index'] = 99
    elif mutation == 'negative': order[0]['index'] = -1
    elif mutation == 'boolean_index': order[0]['index'] = False
    elif mutation == 'wrong_kind': order[0]['kind'] = 'other'
    elif mutation == 'extra_field': order[0]['title'] = 'cannot infer from title'
    else: order[0].pop('index')
    before = deepcopy(request)
    with pytest.raises(ValueError): handoff.prepare_handout_materials(facade, request, [], [])
    assert request == before and not images(facade)


def test_explicit_question_first_order_also_supports_text_only_materials(mixed):
    facade, request, _ = mixed
    request['selection_order'] = [{'kind': 'question', 'index': 0}, {'kind': 'knowledge', 'index': 0}]
    if request['image_selections']: request['selection_order'].append({'kind': 'image', 'index': 0})
    publication = handoff.prepare_handout_materials(facade, request, [], [])
    result = publication.publish(); publication.accept()
    assert [r['original']['kind'] for r in result['resources']] == [
        'complete_question', 'handout_knowledge'] + (['handout_image'] if request['image_selections'] else [])
    if not request['image_selections']: assert not result['image_assets'] and not images(facade)


@pytest.mark.parametrize('change', ['request_order', 'output_order', 'source', 'cancel'])
def test_ordered_publication_rechecks_request_output_source_and_cancellation(mixed, change, monkeypatch):
    facade, request, _ = mixed
    request['selection_order'] = [{'kind': kind, 'index': index}
        for kind, name in handoff._SELECTION_ARRAYS.items() for index in range(len(request[name]))]
    publication = handoff.prepare_handout_materials(facade, request, [], [])
    event = TaskCancellationEvent()
    if change == 'request_order': publication.request['selection_order'].reverse()
    elif change == 'output_order': publication.result['resources'].reverse()
    elif change == 'source':
        read = facade._imported_word_source
        monkeypatch.setattr(facade, '_imported_word_source', lambda *a, **kw: replace(read(*a, **kw), role='answer'))
    else: event.set()
    with pytest.raises((ValueError, ReadCancelled)), read_cancel_scope(event): publication.publish()
    assert not images(facade)


def test_mixed_prepare_is_read_only_then_one_owned_publication(mixed):
    facade, request, provider = mixed; state = facade.state_store.snapshot()
    published = handoff.prepare_handout_materials(facade, request, [], [])
    assert not images(facade) and facade.state_store.snapshot() == state
    kinds = [r['original']['kind'] for r in published.result['resources']]
    assert kinds == ['handout_knowledge'] + (['handout_image'] if request['image_selections'] else []) + ['complete_question']
    result = published.publish()
    assert all(b['role'] != 'answer' for b in result['resources'][0]['blocks'])
    assert any(b['role'] == 'answer' for b in result['resources'][-1]['blocks'])
    assert result['resources'][0]['original']['evidence']['section_snapshot']['blocks'][-1].get('images') is not None
    assert facade.state_store.snapshot() == state and provider.borrow_calls == 0
    assert len(images(facade)) == (2 if request['image_selections'] else 0)
    published.discard(); assert not images(facade)


def test_duplicate_keeps_work_titles_explanation_and_teacher_scope(mixed):
    facade, request, _ = mixed
    first = handoff.prepare_handout_materials(facade, request, [], []); result = first.publish(); first.accept()
    for row in result['resources']:
        row['title'] = '教师修改后的课堂问题'
        for block in row['blocks']:
            if block['kind'] == 'text': block['value'] += '\n教师补充讲解。'
    old = deepcopy(result); files = images(facade)
    duplicate_request = deepcopy(request)
    duplicate_request['knowledge_requests'][0]['teacher_only'] = True
    repeated = handoff.prepare_handout_materials(facade, duplicate_request,
        result['image_assets'], result['resources'])
    assert repeated.result['resources'] == []
    repeated.publish(); repeated.accept()
    assert old == result and images(facade) == files


@pytest.mark.parametrize('mutation', ['schema', 'source_role', 'source_version', 'knowledge_archive',
    'image_archive', 'question_archive', 'question_version', 'question_hash', 'unknown_field', 'empty'])
def test_invalid_or_foreign_mixed_batch_never_publishes(mixed, mutation):
    facade, value, _ = mixed; request = deepcopy(value)
    if mutation == 'schema': request['schema'] = 'unknown'
    elif mutation == 'source_role': request['source']['source_role'] = 'answer'
    elif mutation == 'source_version': request['source']['source_revision'] = 'f' * 64
    elif mutation == 'knowledge_archive': request['knowledge_requests'][0]['source']['source_id'] = 'FOREIGN'
    elif mutation == 'image_archive':
        if not request['image_selections']: pytest.skip('source has no image')
        request['image_selections'][0]['source']['batch_id'] = 'FOREIGN'
    elif mutation == 'question_archive': request['question_selections'][0]['scoped_source_document']['batch_id'] = 'FOREIGN'
    elif mutation == 'question_version': request['question_selections'][0]['revision'] = 'f' * 64
    elif mutation == 'question_hash': request['question_selections'][0]['source_sha256'] = 'f' * 64
    elif mutation == 'unknown_field': request['question_selections'][0]['borrow_source'] = True
    else:
        for key in ('knowledge_requests', 'image_selections', 'question_selections'): request[key] = []
    with pytest.raises(ValueError): handoff.prepare_handout_materials(facade, request, [], [])
    assert not images(facade)


def test_source_change_after_prepare_blocks_final_publication(mixed, monkeypatch):
    facade, request, _ = mixed
    publication = handoff.prepare_handout_materials(facade, request, [], [])
    read = facade._imported_word_source
    monkeypatch.setattr(facade, '_imported_word_source', lambda *a, **kw: replace(read(*a, **kw), role='answer'))
    with pytest.raises(ValueError): publication.publish()
    assert not images(facade)


@pytest.mark.parametrize('when', ['before', 'final_guard', 'after_commit'])
def test_cancel_or_late_guard_rolls_back_whole_owned_batch(mixed, when):
    facade, request, _ = mixed
    publication = handoff.prepare_handout_materials(facade, request, [], [])
    event = TaskCancellationEvent(); calls = []
    def guard():
        calls.append(True)
        if when == 'final_guard' and len(calls) == 2: event.set()
        if when == 'after_commit' and len(calls) == 3: raise ValueError('late lesson change')
    if when == 'before': event.set()
    with pytest.raises((ReadCancelled, ValueError)), read_cancel_scope(event): publication.publish(guard)
    assert not images(facade)


def test_existing_owned_image_survives_late_discard_and_accept_keeps_new(mixed):
    facade, request, _ = mixed
    store = PreparationImageStore(facade.paths.task_root / 'preparation-v1' / 'images')
    old = store.import_bytes(_png('green'), '旧图题', '旧来源', '旧用途'); before = images(facade)
    publication = handoff.prepare_handout_materials(facade, request, [old], [])
    result = publication.publish(); publication.discard()
    assert images(facade) == before
    publication = handoff.prepare_handout_materials(facade, request, [old], [])
    publication.publish(); publication.accept(); publication.discard()
    assert set(before) <= set(images(facade)) and result['image_assets'][0] == old


def test_another_identical_archive_cannot_redirect_this_batch(mixed, tmp_path):
    facade, request, _ = mixed
    path = tmp_path / 'same-copy.docx'
    path.write_bytes(Path(request['source']['source_path']).read_bytes())
    other = facade.save_visual_import_batch(handout_files=(path,), source_type='教师讲义', group_id='other-copy')
    other_source = next(s for s in handout_knowledge_sources(facade)['sources'] if s['batch_id'] == other.batch_id)
    other_bound = handout_knowledge_outline(facade, other_source)['source']
    facade._word_questions()._catalog({other.batch_id: {other_bound['source_id']}}, read_only=True)
    publication = handoff.prepare_handout_materials(facade, request, [], [])
    assert all(r['original']['evidence']['source_path'] == request['source']['source_path']
        for r in publication.result['resources'])
    publication.publish(); publication.accept()


def test_source_change_during_owned_copy_rolls_back(mixed, monkeypatch):
    facade, request, _ = mixed
    if not request['image_selections']: pytest.skip('copy race requires image')
    publication = handoff.prepare_handout_materials(facade, request, [], [])
    read = facade._imported_word_source; link = handoff.os.link
    root = facade.paths.task_root / 'preparation-v1' / 'images'
    def changed(source, target):
        result = link(source, target)
        if Path(target).parent == root:
            monkeypatch.setattr(facade, '_imported_word_source', lambda *a, **kw: replace(read(*a, **kw), role='answer'))
        return result
    monkeypatch.setattr(handoff.os, 'link', changed)
    with pytest.raises(ValueError): publication.publish()
    assert not images(facade)


def test_concurrent_successful_image_owner_is_never_removed(mixed, monkeypatch):
    facade, request, _ = mixed
    if not request['image_selections']: pytest.skip('copy race requires image')
    publication = handoff.prepare_handout_materials(facade, request, [], [])
    link = handoff.os.link; root = facade.paths.task_root / 'preparation-v1' / 'images'; concurrent = []
    def other_wins(source, target):
        if Path(target).parent == root and not concurrent:
            Path(target).write_bytes(Path(source).read_bytes()); concurrent.append(Path(target))
            raise FileExistsError('another importer won the exclusive link')
        return link(source, target)
    monkeypatch.setattr(handoff.os, 'link', other_wins)
    publication.publish(); publication.discard()
    assert list(images(facade)) == [concurrent[0].name]


def linked_request(f):
    prepared = f.service.prepare(f.card, f.row, '具体的知识与完整题关系。')
    snapshot = f.service.save(prepared, expected=f.service.snapshot())
    relation = snapshot['collection']['records'][0]
    resolved = f.service.resolve_selection(relation)
    selection = {**resolved['selection'], 'points': 2}
    return {'link_id': relation['id'], 'snapshot': snapshot, 'selection': selection,
        'reference': resolved['scoped_facade'].word_question_reference([selection], include_images=True),
        'scoped_facade': resolved['scoped_facade'], 'knowledge_endpoint': prepared['knowledge']}


def test_linked_question_exact_endpoints_before_and_after_native_compile(linked):
    request = linked_request(linked); state = linked.facade.state_store.snapshot()
    publication = handoff.prepare_linked_question(linked.facade, request, [], [])
    assert not images(linked.facade)
    result = publication.publish(); publication.accept()
    assert len(result['resources']) == 1 and result['resources'][0]['original']['kind'] == 'complete_question'
    assert linked.facade.state_store.snapshot() == state


@pytest.mark.parametrize('change', ['endpoint', 'snapshot', 'removed', 'after_prepare'])
def test_linked_question_rejects_wrong_knowledge_or_changed_relation(linked, change):
    request = linked_request(linked)
    if change == 'endpoint': request['knowledge_endpoint']['id'] = 'WRONG'
    elif change == 'snapshot': request['snapshot']['revision'] = 'f' * 64
    elif change == 'removed': linked.service.remove(request['link_id'], expected=request['snapshot'])
    if change == 'after_prepare':
        publication = handoff.prepare_linked_question(linked.facade, request, [], [])
        linked.service.edit_reason(request['link_id'], '读取后改动', expected=request['snapshot'])
        with pytest.raises(ValueError): publication.publish()
    else:
        with pytest.raises(ValueError): handoff.prepare_linked_question(linked.facade, request, [], [])
    assert not images(linked.facade)
