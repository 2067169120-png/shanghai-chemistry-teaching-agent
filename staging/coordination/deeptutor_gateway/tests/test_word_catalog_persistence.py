"""Real isolated Word archives exercise the actual listing/detail service path."""
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from threading import Event
import json

from docx import Document
import pytest
from test_desktop_visual_import_facade import (
    _facade, FakeProviderStore, desktop_paths as desktop_paths,
)
from test_word_questions_service import _import, _choice
from test_word_question_dialog import qt_app as qt_app, _Tasks, _cleanup

from integrations.deeptutor_shchem_v1 import desktop_word_questions as words
from integrations.deeptutor_shchem_v1.desktop_word_catalog_index import (
    WordCatalogIndex, compact_questions, expand_questions, projection_matches,
    search_projection,
)
from integrations.deeptutor_shchem_v1.desktop_word_question_filters import (
    compile_question_matcher,
)
from integrations.deeptutor_shchem_v1.desktop_explorer_index import PersonalSearchIndex
from integrations.deeptutor_shchem_v1.reader_cancellation import ReadCancelled, read_cancel_scope


def _files(root):
    return {str(path.relative_to(root)): (path.stat().st_mtime_ns, sha256(path.read_bytes()).hexdigest())
            for path in root.rglob('*') if path.is_file()}


def _rows(catalog):
    return [{key: value for key, value in row.items() if key != 'search_provider'} for row in catalog['items']]


def test_second_service_reuses_directory_without_reading_preview_or_parsing(desktop_paths, tmp_path, monkeypatch):
    facade, source, provider = _import(desktop_paths, tmp_path, image=True)
    first = facade.word_question_listing()
    restarted = _facade(desktop_paths, provider)
    service = restarted._word_questions()
    monkeypatch.setattr(service.preview_cache, 'load', lambda *a: pytest.fail('listing loaded full preview'))
    monkeypatch.setattr(service.reader, 'word_preview_bytes', lambda *a: pytest.fail('listing parsed Word'))
    monkeypatch.setattr(words, 'index_word_questions', lambda *a: pytest.fail('listing recognized boundaries again'))
    second = restarted.word_question_listing()
    assert first['revision'] == second['revision'] and _rows(first) == _rows(second)
    assert service.catalog_index_statistics()['persistent_hits'] == 1
    assert service._cache == {}
    assert not second['items'][0].get('question_blocks')
    assert not any('assets' in row or 'search_projection' in row for row in second['items'])
    assert second['items'][0]['source_role'] == 'handout'
    assert second['items'][0]['source_asset_count'] == 1
    assert second['sources'][0]['source_revision'] == second['items'][0]['source_revision']
    assert provider.borrow_calls == 0


def test_selected_detail_rechecks_source_and_only_expands_one_source(desktop_paths, tmp_path, monkeypatch):
    facade, _, provider = _import(desktop_paths, tmp_path, image=True)
    other = tmp_path / '其他来源.docx'
    document = Document(); document.add_paragraph('【例1】另一题面。'); document.add_paragraph('【答案】另一答案。'); document.save(other)
    facade.save_visual_import_batch(handout_files=(other,), source_type='教师讲义')
    rows = facade.word_question_listing()['items']; chosen = next(row for row in rows if row['source_name'] != other.name)
    restarted = _facade(desktop_paths, provider)  # persisted locations also avoid a cold global catalogue
    calls, original = [], restarted._restore_visual_import_sources
    def restore(descriptor, *, source_ids=None):
        calls.append(source_ids)
        return original(descriptor, source_ids=source_ids)
    monkeypatch.setattr(restarted, '_restore_visual_import_sources', restore)
    detail = restarted.word_question_detail(chosen['key'], chosen['revision'])
    assert detail['question_blocks'] and detail['key'] == chosen['key']
    assert calls == [{chosen['archive_source_id']}]
    assert len(restarted._word_questions()._cache) == 1
    archive = desktop_paths.state_root / 'visual-import-v2' / 'sources' / (chosen['source_sha256'] + '.docx')
    archive.write_bytes(b'changed isolated source')
    with pytest.raises(words.WordQuestionError, match='变化'):
        restarted.word_question_detail(chosen['key'], chosen['revision'])
    assert calls == [{chosen['archive_source_id']}] * 2


def test_read_only_cold_listing_never_writes_and_its_body_search_still_works(desktop_paths, tmp_path):
    facade, _, _ = _import(desktop_paths, tmp_path)
    before = _files(desktop_paths.state_root)
    listing = facade._word_questions().catalog_listing(read_only=True)
    index = PersonalSearchIndex(listing, 'word_native')
    assert index.search({}, '氯化钠')['total'] == 1
    assert index.search({}, '氯化钠属于电解质')['total'] == 0  # answer/analysis only
    assert _files(desktop_paths.state_root) == before
    assert not facade._word_questions().catalog_index.path.exists()
    assert facade._word_questions()._cache == {}


def test_light_index_roundtrip_keeps_split_body_and_block_sharing(desktop_paths, tmp_path):
    facade, _, _ = _import(desktop_paths, tmp_path)
    full, inventory = facade._word_questions()._catalog(read_only=True)
    source, _, preview, original = next(iter(inventory.values()))
    compact = compact_questions(preview, original)
    restored = expand_questions(preview, compact)
    assert restored == original
    for row in compact:
        assert all(type(ref) is int or set(ref) == {'index', 'text_range'} for group in ('question_blocks', 'answer_blocks', 'context_blocks') for ref in row[group])
    assert restored[0]['question_blocks'][0]['text'] is preview['blocks'][2]['text']
    assert not any(isinstance(ref, dict) and 'text' in ref for row in compact for group in ('question_blocks', 'answer_blocks', 'context_blocks') for ref in row[group])
    preview = deepcopy(preview)
    preview['blocks'][2]['text'] += '同段【答案】B'
    from integrations.deeptutor_shchem_v1.desktop_preparation_sources import _digest
    preview['revision'] = _digest({key: value for key, value in preview.items() if key != 'revision'})
    actual = words.index_word_questions(preview)
    assert any(block.get('display_only_split') for row in actual for block in row['question_blocks'])
    assert expand_questions(preview, compact_questions(preview, actual)) == actual


@pytest.mark.parametrize('query', ['水', '化钠', '氯化钠', '氯化钠氯', 'A．铜', '铜  B', '化钠\nC', '甲甲', '甲甲甲', '甲乙甲', '乙甲乙', 'ZZ', 'σ', 'SS'])
def test_postings_preserve_short_ordered_overlapping_and_casefold_substrings(query):
    body = '氯化钠氯化钠\nA．铜  B．氯化钠\nC．水\n甲甲甲乙甲乙\nß σ zz'
    assert projection_matches(search_projection(body), query) == (query.casefold() in body.casefold())


@pytest.mark.parametrize('query', ['', '水', '氯化钠', '写出水', 'A．铜  B', 'H2O', '氯化钠属于电解质', '知识点01', '电解质', '另一来源未命中'])
def test_listing_filters_and_explorer_results_match_complete_catalog(desktop_paths, tmp_path, query):
    facade, _, _ = _import(desktop_paths, tmp_path)
    full = facade.word_question_catalog(); listing = facade.word_question_listing()
    matcher = compile_question_matcher({'query': query}, full['attribute_catalog'])
    expected = [row['key'] for row in full['items'] if matcher(row)]
    light_matcher = compile_question_matcher({'query': query}, listing['attribute_catalog'])
    assert [row['key'] for row in listing['items'] if light_matcher(row)] == expected
    assert [entry['key'] for entry in PersonalSearchIndex(listing, 'word_native').search({}, query)['entries']] == expected
    assert not any(row.get('question_blocks') or row.get('answer_blocks') or row.get('context_blocks') for row in listing['items'])


def test_snapshot_current_attributes_become_stale_after_current_range_change(desktop_paths, tmp_path, monkeypatch):
    facade, _, _ = _import(desktop_paths, tmp_path)
    item = facade.word_question_catalog()['items'][0]
    store = facade._word_questions().attribute_store
    old = store.save_many([facade.word_question_attribute_options(item['key'], item['revision'])['attributes']])[0]
    first = facade.word_question_listing()['items'][0]
    assert first['attributes'] == old
    updated = facade.word_question_update_range(item['key'], item['revision'], block_start=3, question_end=4, answer_start=5, block_end=5)
    listing = facade.word_question_listing(); current = listing['items'][0]
    assert current['revision'] == updated['revision'] and current['attribute_stale'] and 'attributes' not in current
    snapshots = facade._word_questions().observed_catalog_snapshots()
    observed = snapshots[0]['questions'][0]
    assert observed['key'] == current['key'] and observed['revision'] == current['revision']
    assert observed['attribute_state'] == 'stale' and observed['stored_question_revision'] == old['question_revision']
    assert observed['observed_attributes'] == old
    assert snapshots[0]['freshness']['live_revalidated'] is False
    assert snapshots[0]['freshness']['timestamp_proves_current_validity'] is False
    restarted = _facade(desktop_paths, FakeProviderStore(configured=False)); service = restarted._word_questions()
    monkeypatch.setattr(service.preview_cache, 'load', lambda *a: pytest.fail('saved range listing loaded preview'))
    assert restarted.word_question_listing()['items'][0]['revision'] == updated['revision']
    new = store.save_many([facade.word_question_attribute_options(current['key'], current['revision'])['attributes']])[0]
    again = restarted.word_question_listing()['items'][0]
    assert again['attributes'] == new and not again.get('attribute_stale')


def test_quality_changes_are_live_and_never_frozen_in_directory(desktop_paths, tmp_path, monkeypatch):
    facade, _, _ = _import(desktop_paths, tmp_path)
    row = facade.word_question_listing()['items'][0]
    issue = {'id': 'fixture-issue', 'summary': '隔离质量证据', 'suggested_correction': '隔离建议', 'block_indices': [row['block_start']]}
    notes = {row['source_sha256']: {'source_revision': row['source_revision'], 'notes_revision': 'fixture-revision', 'issues': [issue]}}
    monkeypatch.setattr(words, 'source_quality_notes', lambda root: notes)
    blocked = facade.word_question_listing()['items'][0]
    assert not blocked['selection_ready'] and not blocked['export_ready']
    assert blocked['content_quality']['issues'] == [issue]
    notes.clear()
    assert facade.word_question_listing()['items'][0]['selection_ready']


def test_new_batch_and_same_filename_never_reuse_old_source_identity(desktop_paths, tmp_path):
    facade, path, _ = _import(desktop_paths, tmp_path)
    old = facade.word_question_listing()['items'][0]
    document = Document(); document.add_paragraph('【例1】新的独立题面。'); document.add_paragraph('【答案】新答案。'); document.save(path)
    facade.save_visual_import_batch(handout_files=(path,), source_type='教师讲义')
    listing = facade.word_question_listing()
    assert len(listing['items']) == 3
    new = next(row for row in listing['items'] if row['source_sha256'] != old['source_sha256'])
    assert new['key'] != old['key'] and new['revision'] != old['revision']
    archive = desktop_paths.state_root / 'visual-import-v2' / 'sources' / (old['source_sha256'] + '.docx')
    archive.unlink()
    remaining = facade.word_question_listing()
    assert len(remaining['items']) == 1 and remaining['warnings']
    with pytest.raises(words.WordQuestionError, match='变化'):
        facade.word_question_detail(old['key'], old['revision'])


def test_cas_and_corrupt_index_rebuild_only_derived_store(desktop_paths, tmp_path):
    facade, path, provider = _import(desktop_paths, tmp_path)
    _catalog, inventory = facade._word_questions()._catalog()
    source, _, preview, items = next(iter(inventory.values()))
    store = facade._word_questions().catalog_index; compact = compact_questions(preview, items)
    revision = store.revision(preview)
    assert not store.save(source.content, source.filename, preview, compact, expected_revision='wrong-cas')
    assert not store.save(source.content, source.filename, preview, compact, expected_revision=None)
    assert store.save(source.content, source.filename, preview, compact, expected_revision=revision)
    original = path.read_bytes()
    store.path.write_bytes(b'isolated unreadable sqlite fixture')
    restarted = _facade(desktop_paths, provider)
    listing = restarted.word_question_listing()
    assert len(listing['items']) == 2 and restarted._word_questions().catalog_index.recovered_databases
    assert path.read_bytes() == original
    assert len(restarted._word_questions().observed_catalog_snapshots()[0]['questions']) == 2


def test_cancellation_before_projection_never_publishes_snapshot(desktop_paths, tmp_path):
    facade, _, _ = _import(desktop_paths, tmp_path)
    event = Event()
    def progress(value):
        if '轻量目录' in value['message']:
            event.set()
    with read_cancel_scope(event, check_boundaries=False):
        with pytest.raises(ReadCancelled):
            facade._word_questions().catalog_listing(progress=progress)
    assert facade._word_questions().observed_catalog_snapshots() == []
    assert not facade._word_questions()._locations


def test_actual_word_dialog_loads_only_current_body_and_source_failure_disables_selection(desktop_paths, tmp_path, qt_app, monkeypatch):
    from integrations.deeptutor_shchem_v1.desktop_workbench.word_question_dialog import WordQuestionDialog
    facade, _, _ = _import(desktop_paths, tmp_path)
    original, calls = facade.word_question_detail, []
    def detail(key, revision):
        calls.append(key)
        return original(key, revision)
    monkeypatch.setattr(facade, 'word_question_detail', detail)
    tasks = _Tasks(); dialog = WordQuestionDialog(facade, tasks)
    try:
        tasks.finish()  # catalogue; current detail is still pending
        assert len(dialog._items) == 2 and all(row.get('lightweight') for row in dialog._items.values())
        assert not dialog.current_question_readiness()[0] and calls == []
        tasks.flush()
        current = dialog._current_key
        assert calls == [current] and dialog.current_question_readiness()[0]
        assert sum(bool(row.get('question_blocks')) for row in dialog._items.values()) == 1
        dialog.question_list.setCurrentRow(1); tasks.flush()
        assert calls[-1] == dialog._current_key != current
        assert sum(bool(row.get('question_blocks')) for row in dialog._items.values()) == 1
        row = dialog._items[dialog._current_key]
        archive = desktop_paths.state_root / 'visual-import-v2' / 'sources' / (row['source_sha256'] + '.docx')
        archive.write_bytes(b'isolated changed source before reopening')
        dialog._show_question(dialog._current_key); tasks.flush()
        assert not dialog.current_question_readiness()[0]
        assert dialog._items[dialog._current_key].get('lightweight')
        assert not dialog._items[dialog._current_key].get('question_blocks')
    finally:
        _cleanup(dialog, tasks); facade.shutdown()


@pytest.mark.parametrize('progress_port', [True, False])
def test_source_restore_progress_uses_optional_port_and_keeps_legacy_port(desktop_paths, tmp_path, monkeypatch, progress_port):
    facade, _, _ = _import(desktop_paths, tmp_path)
    batch = facade.list_imported_word_batches()[0]
    source = facade.imported_word_sources(batch.batch_id)[0]
    service = facade._word_questions(); calls, messages = [], []
    before = _files(desktop_paths.state_root)
    restore = facade._restore_visual_import_sources
    def old_port(descriptor, *, source_ids=None):
        calls.append(('legacy', source_ids))
        return restore(descriptor, source_ids=source_ids)
    def new_port(descriptor, progress, *, source_ids=None):
        calls.append(('progress', source_ids))
        progress({'message': 'fixture source validation progress'})
        return restore(descriptor, source_ids=source_ids)
    monkeypatch.setattr(facade, '_restore_visual_import_sources', old_port)
    monkeypatch.setattr(facade, '_restore_visual_import_sources_progress', new_port if progress_port else None, raising=False)
    inventory, warnings = service._inventory({batch.batch_id: {source['source_id']}}, read_only=True, progress=messages.append)
    assert len(inventory) == 1 and not warnings
    assert calls == [('progress' if progress_port else 'legacy', {source['source_id']})]
    assert any(row['message'] == 'fixture source validation progress' for row in messages) is progress_port
    assert _files(desktop_paths.state_root) == before
