"""Student handouts omit answer-only pages without rewriting teacher placement."""
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

import pytest
from docx import Document
from pptx import Presentation
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from test_composition_render import app, text
from test_composition_ui import desk, window
from test_desktop_studio_ui import settle
from test_lesson_design import sample
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    FORMATS, add_resource, block, enable_composition, new_composition, resource)
from integrations.deeptutor_shchem_v1.desktop_lesson_design import digest
from integrations.deeptutor_shchem_v1.desktop_composition_render import (
    arrange_composition, layout_composition, write_pdf, write_pptx)
from integrations.deeptutor_shchem_v1.desktop_lesson_output import (
    checked_file, export_design, write_documents)


def question(format='16:9'):
    item = resource('complete_question', '题面标题', [
        block('text', 'question', '完整学生条件。后续条件仍需保留。', 1),
        block('text', 'answer', '\n'.join('教师答案第%d行，须完整保留。' % i for i in range(80)), 2)],
        {'source_path': '合成题.docx', 'source_sha256': 'a'*64, 'source_revision': 'b'*64,
         'question_key': 'synthetic-student-page-scope', 'scope_review_sha256': 'c'*64,
         'scope_status': 'complete_native_scope',
         'hierarchy': dict.fromkeys(('paper_id','theme_big_question_id','printed_question_id','atomic_part_id'), 'unknown')})
    value = new_composition(format)
    w, h = FORMATS[format]
    value = add_resource(value, item, {'page_id': value['pages'][0], 'x': 12, 'y': 12,
                                     'width': w-24, 'height': h-24})
    return arrange_composition(value, [], {})


def teacher_note(value):
    item = resource('personal_text', '仅教师可见标题', [block('text', 'teacher_note', '仅教师内容')],
                    {'source_path': '作者输入', 'source_sha256': 'unknown', 'source_revision': 'unknown'})
    page = 'CP-'+uuid4().hex
    value = deepcopy(value)
    value['pages'].append(page)
    return add_resource(value, item, {'page_id': page, 'x': 12, 'y': 12, 'width': 186, 'height': 80})


def payload(value):
    p = sample()
    p['lesson_design'] = enable_composition(p['lesson_design'])
    p['lesson_design']['composition'] = value
    return p


@pytest.mark.parametrize('format', FORMATS)
def test_answer_only_continuation_pages_are_absent_from_student_and_teacher_is_intact(app, format):
    value = question(format)
    before = deepcopy(value)
    teacher = layout_composition(value, [], {}, student=False)
    student = layout_composition(value, [], {}, student=True)
    assert not teacher['errors'] and not student['errors']
    assert list(teacher['pages']) == value['pages']
    assert len(student['pages']) == 1 < len(teacher['pages'])
    assert text(student) == '题面标题完整学生条件。后续条件仍需保留。'
    assert '教师答案第79行' in text(teacher) and '教师答案' not in text(student)
    assert value == before and teacher['composition'] == before


def test_author_inserted_blank_page_remains_in_student_printout(app, tmp_path):
    value = question('A4')
    blank = 'CP-'+uuid4().hex
    value['pages'].append(blank)
    before = deepcopy(value)
    student = layout_composition(value, [], {}, student=True)
    assert not student['errors']
    assert list(student['pages']) == [value['pages'][0], blank]
    assert not student['pages'][blank]
    import pypdfium2 as pdfium
    path = tmp_path/'intentional-blank.pdf'
    write_pdf(student, path)
    with pdfium.PdfDocument(str(path)) as pdf:
        assert len(pdf) == 2
    assert value == before


def test_teacher_only_resource_title_does_not_leak_into_any_student_artifact(app, tmp_path):
    value = teacher_note(question('A4'))
    p = payload(value)
    student = layout_composition(value, [], {}, student=True)
    teacher = layout_composition(value, [], {}, student=False)
    assert '仅教师可见标题' not in text(student)
    assert '仅教师可见标题' in text(teacher)
    path = tmp_path/'student.pptx'
    write_pptx(student, path, {})
    words = ''.join(shape.text for slide in Presentation(path).slides for shape in slide.shapes if shape.has_text_frame)
    assert '仅教师可见标题' not in words and '教师答案' not in words
    write_documents(p, tmp_path, {})
    assert '仅教师可见标题' not in ''.join(s.text for s in Document(tmp_path/'student_worksheet.docx').paragraphs)
    assert '仅教师可见标题' in ''.join(s.text for s in Document(tmp_path/'lesson_plan.docx').paragraphs)


def test_only_teacher_material_requires_student_content_before_output(app):
    value = teacher_note(new_composition())
    value['pages'] = value['pages'][1:]
    result = layout_composition(value, [], {}, student=True)
    assert any('学生版没有可见' in e for e in result['errors'])
    assert not result['pages']


def test_output_pdf_ppt_and_audit_agree_on_each_audience_page_list(app, tmp_path):
    import json
    import pypdfium2 as pdfium
    value = question()
    p = payload(value)
    facade = SimpleNamespace(paths=SimpleNamespace(task_root=tmp_path/'tasks'))
    before = deepcopy(p)
    record = export_design(facade, p)
    p['lesson_design']['exports'].append(record)
    directory = checked_file(facade, p['lesson_design'], record['id'], 'student_canvas.pdf').parent
    audit = json.loads((directory/'canvas-output.json').read_text(encoding='utf-8'))
    assert audit['page_ids'] == value['pages']
    assert audit['page_ids_by_audience']['teacher'] == value['pages']
    assert audit['page_ids_by_audience']['student'] == [value['pages'][0]]
    for scope, pdfname, pptname in [('teacher','lesson_canvas.pdf','lesson_presentation.pptx'),
                                   ('student','student_canvas.pdf','student_presentation.pptx')]:
        with pdfium.PdfDocument(str(directory/pdfname)) as pdf:
            assert len(pdf) == len(audit['page_ids_by_audience'][scope])
        assert len(Presentation(directory/pptname).slides) == len(audit['page_ids_by_audience'][scope])
    p['lesson_design']['exports'].pop()
    assert p == before


def test_native_student_teacher_switch_matches_export_without_history_edit(desk):
    dialog, panel, app, page = desk
    original = question()
    panel.commit(original)
    before = deepcopy(dialog.history.value)
    panel.student.setChecked(True)
    settle(app)
    assert len(panel.layout_result['pages']) == 1
    assert '学生版 1页' in panel.status.text()
    hidden = original['placements'][original['resources'][0]['id']]['continuations'][-1]['page_id']
    panel.selected = original['resources'][0]['id']
    panel.frame_index = len(original['placements'][panel.selected]['continuations'])
    panel.refresh()
    panel.fit_view()  # Selecting a hidden answer continuation cannot lose the view.
    assert panel.page_picker.currentData() == hidden
    assert panel.page_picker.currentText().startswith('教师编排第')
    QTest.mouseClick(panel.student, Qt.MouseButton.LeftButton)
    settle(app)
    assert len(panel.layout_result['pages']) == len(original['pages'])
    assert dialog.history.value == before
