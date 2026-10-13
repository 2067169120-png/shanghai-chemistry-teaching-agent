"""Private reference disclosure defaults and actual native Qt reading geometry."""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ.setdefault("QT_SCALE_FACTOR", "1")

import pytest

pytest.importorskip("PySide6")
from docx import Document
from docx.shared import Mm
from PIL import Image
from PySide6.QtCore import QPoint, QRect, Qt

from test_desktop_visual_import_facade import (
    FakeProviderStore, _facade, desktop_paths as desktop_paths,
)
from test_word_question_dialog import (
    WORKBENCH_STYLE, _AttributeFacade, _Facade, _Tasks, _isolated_qt_app,
)
from runtime.deeptutor_shchem.word_question_reader_layout_qa import (
    footer_evidence, fully_visible, reveal, settle,
)
from integrations.deeptutor_shchem_v1.desktop_workbench.app import create_application
from integrations.deeptutor_shchem_v1.desktop_workbench.word_inline_reader import WordInlineReader
from integrations.deeptutor_shchem_v1.desktop_workbench.word_question_dialog import WordQuestionDialog


@pytest.fixture
def qt_app():
    with _isolated_qt_app(stylesheet=WORKBENCH_STYLE) as app:
        create_application(["private-reference-reader-contract"])
        yield app


def _open(facade, qt_app, **kwargs):
    tasks = _Tasks()
    dialog = WordQuestionDialog(facade, tasks, **kwargs)
    dialog.resize(900, 760)
    dialog.show()
    tasks.flush()
    settle(qt_app)
    return dialog, tasks


@pytest.mark.parametrize("initial_key", [None, "Q1"])
@pytest.mark.parametrize("width", [320, 420, 900])
def test_private_reference_starts_with_filters_and_settings_collapsed(qt_app, initial_key, width):
    saved = [{"key": "Q2", "revision": "revision-Q2", "points": 7}]
    facade = _Facade(saved=saved)
    dialog, tasks = _open(facade, qt_app, persist_selection=False,
        handoff_mode="reference", required_source_id="A", initial_source_id="A",
        initial_question_key=initial_key, initial_search_query="电解质")
    dialog.resize(width, 780)
    settle(qt_app)
    assert dialog.filter_toggle.isVisible() and not dialog.filter_toggle.isChecked()
    assert dialog.filter_host.isHidden()
    assert not dialog.question_tools_button.isChecked() and dialog.question_tools.isHidden()
    assert dialog.question_navigation_button.isVisible() and dialog.question_navigation.isHidden()
    assert dialog.body_scroll.horizontalScrollBar().maximum() == 0
    assert fully_visible(dialog.filter_toggle, dialog.body_scroll.viewport())
    assert fully_visible(dialog.question_navigation_button, dialog.body_scroll.viewport())
    assert dialog.search.text() == "电解质" and dialog.source_combo.currentData() == "A"
    assert dialog._required_source_id == "A" and dialog._initial_question_key == initial_key
    assert dialog._current_key == "Q1" and dialog.selections == []
    assert not any(call[0] in {"saved", "save", "reference", "export"} for call in facade.calls)
    dialog.close_button.click()
    tasks.flush()
    assert dialog._closed and facade.saved == saved


@pytest.mark.parametrize("mode", ["reference", "composition", "both"])
def test_persistent_initial_question_keeps_automatic_settings_and_filter_access(qt_app, mode):
    dialog, tasks = _open(_Facade(), qt_app, handoff_mode=mode, initial_question_key="Q1")
    assert dialog.filter_toggle.isHidden() and dialog.filter_host.isVisible()
    assert dialog.question_navigation_button.isHidden() and dialog.question_navigation.isVisible()
    assert dialog.question_tools_button.isChecked() and dialog.question_tools.isVisible()
    assert dialog.points.isEnabled() and dialog.range_button.isEnabled()
    dialog.close()
    tasks.flush()


@pytest.mark.parametrize("mode,filters_collapsed", [("composition", True), ("both", False)])
def test_existing_private_composition_and_both_disclosures_are_preserved(qt_app, mode, filters_collapsed):
    dialog, tasks = _open(_Facade(), qt_app, handoff_mode=mode,
        persist_selection=False, initial_question_key="Q1")
    assert dialog.filter_toggle.isVisible() == filters_collapsed
    assert dialog.filter_host.isHidden() == filters_collapsed
    assert dialog.question_navigation_button.isHidden() and dialog.question_navigation.isVisible()
    assert dialog.question_tools_button.isChecked() and dialog.question_tools.isVisible()
    dialog.close()
    tasks.flush()


def test_private_reference_disclosures_preserve_filters_source_points_and_actions(qt_app, tmp_path):
    facade = _AttributeFacade(tmp_path / "isolated-attributes")
    source_before = deepcopy(facade.catalog)
    dialog, tasks = _open(facade, qt_app, handoff_mode="reference",
        persist_selection=False, required_source_id="A", initial_question_key="Q1",
        initial_search_query="电解质")
    reveal(dialog, dialog.filter_toggle, qt_app)
    dialog.filter_toggle.click()
    settle(qt_app)
    assert dialog.filter_host.isVisible()
    reveal(dialog, dialog.search, qt_app)
    assert fully_visible(dialog.search, dialog.body_scroll.viewport())
    dialog.search.setText("Q1")
    dialog._filter_items()
    assert dialog.question_list.count() == 1
    reveal(dialog, dialog.filter_toggle, qt_app)
    dialog.filter_toggle.click()
    settle(qt_app)
    assert dialog.filter_host.isHidden() and dialog.search.text() == "Q1"
    assert dialog.source_combo.currentData() == "A" and dialog._required_source_id == "A"
    assert dialog._current_key == "Q1" and dialog.question_list.count() == 1
    reveal(dialog, dialog.question_navigation_button, qt_app)
    dialog.question_navigation_button.click()
    settle(qt_app)
    assert dialog.question_navigation.isVisible()
    assert not dialog.question_list.visibleRegion().isEmpty()
    for control in (dialog.previous_button, dialog.next_button, dialog.reload_button):
        reveal(dialog, control, qt_app)
        assert control.isVisible()
    reveal(dialog, dialog.question_navigation_button, qt_app)
    dialog.question_navigation_button.click()
    settle(qt_app)
    assert dialog.question_navigation.isHidden() and dialog._current_key == "Q1"
    reveal(dialog, dialog.question_tools_button, qt_app)
    dialog.question_tools_button.click()
    settle(qt_app)
    for control in (dialog.points, dialog.range_button, dialog.attributes_button):
        reveal(dialog, control, qt_app)
        assert control.isVisible() and control.isEnabled()
    dialog.points.setValue(9)
    reveal(dialog, dialog.question_tools_button, qt_app)
    dialog.question_tools_button.click()
    settle(qt_app)
    assert dialog.question_tools.isHidden() and dialog.points.value() == 9
    reveal(dialog, dialog.select_current_button, qt_app)
    dialog.select_current_button.click()
    assert dialog.selections == [{"key": "Q1", "revision": "revision-Q1", "points": 9}]
    assert dialog.export_button.isEnabled()
    reveal(dialog, dialog.preparation_toggle, qt_app)
    dialog.preparation_toggle.click()
    reveal(dialog, dialog.preview_button, qt_app)
    dialog.preview_button.click()
    tasks.flush()
    settle(qt_app)
    assert dialog.import_button.isEnabled() and dialog.preparation_reference is None
    assert facade.catalog == source_before and not any(call[0] == "save" for call in facade.calls)
    dialog.close_button.click()
    tasks.flush()
    assert dialog._closed and facade.saved == []


def _png(color):
    output = BytesIO()
    Image.new("RGB", (240, 100), color).save(output, format="PNG")
    return output.getvalue()


@pytest.fixture
def native_long_question(desktop_paths, tmp_path):
    document = Document()
    document.add_paragraph("1．请回答下列问题：核对完整记录、条件和原图。（合成阅读验收材料）")
    for index in range(12):
        document.add_paragraph(f"题面记录第{index + 1}段：保留材料、实验步骤、条件与单位；逐项阅读完整记录。")
        if index in (2, 5, 8):
            document.add_paragraph().add_run().add_picture(BytesIO(_png((30 + index, 80, 160))), width=Mm(40))
    document.add_paragraph("题面末段定位：完整题最后一段仍可滚动阅读。")
    document.add_paragraph().add_run().add_picture(BytesIO(_png("#cc5522")), width=Mm(40))
    document.add_paragraph("【答案】合成阅读验收的独立参考文字。")
    for index in range(12):
        document.add_paragraph(f"解析记录第{index + 1}段：核对原文和必要条件，保留完整答案说明。")
        if index in (2, 5, 8):
            document.add_paragraph().add_run().add_picture(BytesIO(_png((120, 40 + index, 90))), width=Mm(40))
    document.add_paragraph("答案末段定位：独立答案最后一段仍可滚动阅读。")
    document.add_paragraph().add_run().add_picture(BytesIO(_png("#337744")), width=Mm(40))
    path = tmp_path / "合成长题四图四答案图.docx"
    document.save(path)
    provider = FakeProviderStore(configured=False)
    facade = _facade(desktop_paths, provider)
    facade.save_visual_import_batch(handout_files=(path,), source_type="教师讲义")
    catalog = facade.word_question_catalog()
    assert len(catalog["items"]) == 1
    row = catalog["items"][0]
    return facade, row, provider


def _visible_document_rect(dialog, reader, scroll):
    rect = QRect(reader.viewport().mapTo(dialog, QPoint(0, 0)), reader.viewport().size())
    rect = rect.intersected(QRect(scroll.viewport().mapTo(dialog, QPoint(0, 0)), scroll.viewport().size()))
    return rect.intersected(QRect(dialog.body_scroll.viewport().mapTo(dialog, QPoint(0, 0)),
                                dialog.body_scroll.viewport().size()))


def _block_rect(dialog, reader, block):
    rect = reader.document().documentLayout().blockBoundingRect(block).toAlignedRect()
    rect.translate(reader.viewport().mapTo(dialog, QPoint(0, 0)))
    return rect


def _image_blocks(reader):
    blocks = []
    block = reader.document().begin()
    while block.isValid():
        iterator = block.begin()
        while not iterator.atEnd():
            fragment = iterator.fragment()
            if fragment.charFormat().isImageFormat():
                blocks.append(block)
            iterator += 1
        block = block.next()
    return blocks


@pytest.mark.parametrize("width,height", [(900, 780), (420, 780)])
def test_native_private_reference_first_screen_and_scroll_reach_full_text_and_last_images(
    qt_app, native_long_question, width, height, record_property,
):
    facade, row, provider = native_long_question
    before = facade.state_store.snapshot()
    dialog, tasks = _open(facade, qt_app, handoff_mode="reference", persist_selection=False,
        required_source_id=row["source_id"], initial_source_id=row["source_id"],
        initial_question_key=row["key"])
    dialog.resize(width, height)
    settle(qt_app)
    assert (dialog.width(), dialog.height()) == (width, height)
    assert dialog.body_scroll.verticalScrollBar().value() == 0
    assert dialog.body_scroll.horizontalScrollBar().maximum() == 0
    reader = dialog._panels[0].findChild(WordInlineReader)
    assert reader is not None and dialog.current_question_readiness()[0]
    assert "答案末段定位" not in reader.toPlainText()
    area = _visible_document_rect(dialog, reader, dialog.tabs.widget(0))
    record_property("window", f"{width}x{height}")
    record_property("initial_visible_body_height", area.height())
    record_property("footer_height", dialog.reader_footer.height())
    assert area.height() >= 180, (width, area, footer_evidence(dialog))
    first = reader.document().begin()
    for _ in range(3):
        assert area.contains(_block_rect(dialog, reader, first)), (width, first.text(), area)
        first = first.next()
    for tab, sentinel, last_color in ((0, "题面末段定位", "#cc5522"), (1, "答案末段定位", "#337744")):
        dialog.tabs.setCurrentIndex(tab)
        tasks.flush()
        settle(qt_app)
        scroll = dialog.tabs.widget(tab)
        reader = dialog._panels[tab].findChild(WordInlineReader)
        assert reader is not None and sentinel in reader.toPlainText()
        pictures = _image_blocks(reader)
        assert len(pictures) == 4 and len(reader.reading["images"]) == 4
        last = pictures[-1]
        reveal(dialog, dialog.tabs, qt_app)
        marker = reader.document().find(sentinel).block()
        for block in (marker, last):
            block_rect = reader.document().documentLayout().blockBoundingRect(block).toAlignedRect()
            origin = reader.mapTo(scroll.widget(), QPoint(0, 0))
            scroll.ensureVisible(origin.x() + block_rect.center().x(),
                                 origin.y() + block_rect.center().y(), 0, block_rect.height() // 2 + 2)
            settle(qt_app)
            assert _visible_document_rect(dialog, reader, scroll).contains(_block_rect(dialog, reader, block))
        assert scroll.verticalScrollBar().maximum() > 0
        record_property(f"tab_{tab}_scroll_maximum", scroll.verticalScrollBar().maximum())
        last_image_id = next(t["asset_id"] for p in reversed(reader.reading["paragraphs"])
                             for t in p["tokens"] if t["kind"] == "image")
        assert reader._images[last_image_id].pixelColor(0, 0).name() == last_color
        assert footer_evidence(dialog)["close_complete"] and footer_evidence(dialog)["status_unclipped"]
    dialog.close_button.click()
    tasks.flush()
    assert dialog._closed and dialog.selections == []
    assert facade.state_store.snapshot() == before and provider.borrow_calls == 0
