"""Actual Qt table geometry, using in-memory synthetic Word only."""

from __future__ import annotations

import os
from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
pytest.importorskip("PySide6")
from PySide6.QtCore import QUrl
from PySide6.QtGui import QFont, QFontDatabase, QPixmap, QTextCursor, QTextTable
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QDialog

from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
    _readable_table,
    _table_rows,
)
from integrations.deeptutor_shchem_v1.desktop_workbench.import_word_dialog import (
    ImportWordDialog,
)
from integrations.deeptutor_shchem_v1.desktop_workbench.word_lesson_reader import (
    WordLessonReader,
)
from integrations.deeptutor_shchem_v1.desktop_workbench.word_table_layout import (
    word_table_html,
)


def _rows(kind="merged"):
    document = Document()
    columns = 7 if kind == "wide" else 3
    table = document.add_table(rows=4, cols=columns)
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.text = f"第{r + 1}行第{c + 1}列合成内容"
    if kind == "merged":
        table.cell(0, 0).merge(table.cell(0, 1)).text = "氧化还原反应 · 知识点对照"
        table.cell(0, 2).text = "课堂练习"
        table.cell(1, 0).merge(table.cell(2, 0)).text = "电子转移"
        table.cell(1, 1).text = "失去电子\n化合价升高"
        table.cell(2, 1).text = "得到电子\n化合价降低"
        table.cell(1, 2).text = "辨认变化过程"
        table.cell(2, 2).text = "核对反应条件"
        # Explicit omitted edge cells, not empty or inferred source cells.
        row = table.rows[3]._tr
        cells = list(row.tc_lst)
        row.remove(cells[0])
        row.remove(cells[2])
        for name in ("gridBefore", "gridAfter"):
            value = OxmlElement("w:" + name)
            value.set(qn("w:val"), "1")
            row.get_or_add_trPr().append(value)
    elif kind == "nested":
        table.cell(1, 1).add_table(rows=1, cols=2).cell(
            0, 0
        ).text = "内嵌表格原文不得丢弃"
    elif kind == "invalid_merge":
        value = OxmlElement("w:vMerge")
        value.set(qn("w:val"), "continue")
        table.cell(0, 0)._tc.get_or_add_tcPr().append(value)
    elif kind == "continuation_text":
        table.cell(0, 0).merge(table.cell(1, 0)).text = "合并起点"
        continuation = table.rows[1]._tr.tc_lst[0]
        paragraph = continuation.p_lst[0]
        run, text = OxmlElement("w:r"), OxmlElement("w:t")
        text.text = "续格仍有原文，不能吞掉"
        run.append(text)
        paragraph.append(run)
    elif kind == "html":
        table.cell(
            0, 0
        ).text = '<img src="https://example.invalid/private.png"><a href="file:///C:/private">原文链接</a><script>不执行</script>'
    warnings = set()
    rows = _table_rows(table._tbl, document, warnings)
    return rows


def _source(rows=None, *, sha="a" * 64, revision="1" * 64):
    rows = rows if rows is not None else _rows()
    return {
        "source_name": "合成教案 · 氧化还原反应.docx",
        "source_sha256": sha,
        "revision": revision,
        "blocks": [
            {
                "index": 1,
                "label": "区块 1 · 导入",
                "text": "课堂导入：按原教案顺序阅读，不把表格变成知识摘要。",
                "warnings": [],
            },
            {
                "index": 2,
                "label": "区块 2 · 原教案表格",
                "text": _readable_table(rows),
                "warnings": [],
            },
            {
                "index": 3,
                "label": "区块 3 · 总结",
                "text": "课后总结：电子转移。末段完整保留。",
                "warnings": [],
            },
        ],
        "sections": [{"start": 2, "end": 3, "title": "原表与课堂小结"}],
        "assets": [
            {
                "asset_id": "synthetic-image",
                "label": "合成图，不来自真实教材",
                "block_index": 2,
                "preview_supported": True,
                "mime_type": "image/png",
            }
        ],
        "warnings": [],
    }


def _overlay(rows, source):
    return {
        "source_sha256": source["source_sha256"],
        "source_revision": source["revision"],
        "tables": {2: rows},
    }


def _tables(reader):
    def visit(frame):
        for child in frame.childFrames():
            if isinstance(child, QTextTable):
                yield child
            yield from visit(child)

    return list(visit(reader.browser.document().rootFrame()))


def _cell_text(table, row, col):
    cell = table.cellAt(row, col)
    cursor = cell.firstCursorPosition()
    cursor.setPosition(
        cell.lastCursorPosition().position(), QTextCursor.MoveMode.KeepAnchor
    )
    return cursor.selectedText()


def _image_formats(reader):
    values = []
    block = reader.browser.document().begin()
    while block.isValid():
        iterator = block.begin()
        while not iterator.atEnd():
            fragment = iterator.fragment()
            if fragment.isValid() and fragment.charFormat().isImageFormat():
                values.append(fragment.charFormat().toImageFormat())
            iterator += 1
        block = block.next()
    return values


@pytest.fixture
def qt_app():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def reader(qt_app):
    widget = WordLessonReader()
    widget.resize(420, 800)
    widget.show()
    yield widget
    widget.close()
    widget.deleteLater()
    qt_app.processEvents()


def test_actual_qtexttable_preserves_horizontal_vertical_merges_and_omitted_edges(
    reader,
):
    rows = _rows()
    source = _source(rows)
    reader.set_source(source)
    assert reader.set_table_previews(_overlay(rows, source))
    tables = _tables(reader)
    assert len(tables) == 1
    table = tables[0]
    assert (table.rows(), table.columns()) == (4, 3)
    assert table.cellAt(0, 0).columnSpan() == 2
    assert table.cellAt(0, 0).rowSpan() == 1
    assert table.cellAt(0, 1).firstPosition() == table.cellAt(0, 0).firstPosition()
    assert table.cellAt(1, 0).rowSpan() == 2
    assert table.cellAt(2, 0).firstPosition() == table.cellAt(1, 0).firstPosition()
    assert "电子转移" == _cell_text(table, 1, 0)
    assert "（原行省略）" == _cell_text(table, 3, 0) == _cell_text(table, 3, 2)
    assert "第4行第2列合成内容" == _cell_text(table, 3, 1)
    assert reader._blocks[2]["text"] == _readable_table(rows)


@pytest.mark.parametrize(
    "kind", ["wide", "nested", "invalid_merge", "continuation_text"]
)
def test_wide_nested_or_ambiguous_merge_retains_exact_source_text(reader, kind):
    rows = _rows(kind)
    expected = _readable_table(rows)
    assert word_table_html(rows, expected) is None
    source = _source(rows)
    reader.set_source(source)
    assert reader.set_table_previews(_overlay(rows, source))
    assert not _tables(reader)
    assert expected in reader.browser.toPlainText()
    assert "保留逐格原文" in reader.browser.toPlainText()
    assert not reader.table_mode_button.isVisible()


def test_geometry_cannot_replace_mismatched_native_cell_text(reader):
    rows = _rows()
    source = _source(rows)
    source["blocks"][1]["text"] += "\n保留的额外来源原文"
    reader.set_source(source)
    assert reader.set_table_previews(_overlay(rows, source))
    assert not _tables(reader)
    assert source["blocks"][1]["text"] in reader.browser.toPlainText()


def test_cell_html_is_literal_and_never_creates_images_or_external_links(reader):
    rows = _rows("html")
    source = _source(rows)
    reader.set_source(source)
    assert reader.set_table_previews(_overlay(rows, source))
    table = _tables(reader)[0]
    assert rows[0]["cells"][0]["text"] == _cell_text(table, 0, 0)
    assert not _image_formats(reader)
    assert not reader.browser.openLinks() and not reader.browser.openExternalLinks()
    assert (
        reader.browser.document().resource(
            2, QUrl("https://example.invalid/private.png")
        )
        is None
    )


@pytest.mark.parametrize("changed", ["source_sha256", "source_revision"])
def test_overlay_requires_exact_current_source_binding(reader, changed):
    rows = _rows()
    source = _source(rows)
    reader.set_source(source)
    stale = _overlay(rows, source)
    stale[changed] = "f" * 64
    assert not reader.set_table_previews(stale)
    assert not _tables(reader)
    assert _readable_table(rows) in reader.browser.toPlainText()
    assert reader.set_table_previews(_overlay(rows, source))
    assert not reader.set_table_previews(stale)
    assert len(_tables(reader)) == 1


def test_toggle_preserves_blocks_search_position_image_and_block_actions(reader):
    rows = _rows()
    source = _source(rows)
    reader.set_source(source)
    reader.set_table_previews(_overlay(rows, source))
    pixmap = QPixmap(360, 120)
    pixmap.fill(0xFF8899AA)
    reader.show_image("synthetic-image", pixmap, "合成图像预览")
    reader.search_edit.setText("电子转移")
    reader.find_next()
    previous_index = reader._match_index
    requests = []
    reader.block_requested.connect(requests.append)
    reader.table_mode_button.click()
    assert not _tables(reader)
    assert _readable_table(rows) in reader.browser.toPlainText()
    assert reader.search_edit.text() == "电子转移"
    assert reader._match_index == previous_index
    assert len(reader._matches) == 2
    assert set(reader._blocks) == {1, 2, 3}
    assert reader._image_asset_id == "synthetic-image"
    assert reader._image_pixmap.size() == pixmap.size()
    assert len(_image_formats(reader)) == 1
    reader.table_mode_button.click()
    assert len(_tables(reader)) == 1
    assert len(_image_formats(reader)) == 1
    assert not requests
    action = next(
        url for url, value in reader._actions.items() if value == ("block", 2)
    )
    reader.browser.anchorClicked.emit(QUrl(action))
    assert requests == [2]
    assert "末段完整保留" in reader.browser.toPlainText()


def test_source_change_discards_old_grid_search_image_and_stale_overlay(reader):
    rows = _rows()
    first = _source(rows)
    reader.set_source(first)
    reader.set_table_previews(_overlay(rows, first))
    pixmap = QPixmap(360, 120)
    pixmap.fill(0xFF8899AA)
    assert reader.show_image("synthetic-image", pixmap, "合成图像预览")
    assert reader._image_pixmap is not None
    reader.search_edit.setText("电子转移")
    reader.table_mode_button.click()
    replacement = _source(_rows("wide"), sha="b" * 64, revision="2" * 64)
    reader.set_source(replacement)
    assert reader.search_edit.text() == ""
    assert reader._image_pixmap is None
    assert not reader._table_grids and not reader._table_fallbacks
    assert reader._show_table_grids
    assert "网格预览 · 切换逐格原文" in reader.table_mode_button.text()
    assert not reader.set_table_previews(_overlay(rows, first))
    assert not _tables(reader)
    assert replacement["blocks"][1]["text"] in reader.browser.toPlainText()


class SyntheticFacade:
    def __init__(self, mode="ok"):
        self.rows = _rows()
        self.sources = {
            "synthetic-a": _source(self.rows),
            "synthetic-b": _source(_rows("wide"), sha="b" * 64, revision="2" * 64),
        }
        self.table_calls = []
        self.reference_calls = []
        self.mode = mode

    def imported_word_sources(self, *_args):
        return [
            {"source_id": key, "source_name": value["source_name"], "role": "handout"}
            for key, value in self.sources.items()
        ]

    def imported_word_preview(self, _batch, source_id):
        return deepcopy(self.sources[source_id])

    def imported_word_table_previews(self, batch, source_id, digest, revision):
        self.table_calls.append((batch, source_id, digest, revision))
        if self.mode == "raise":
            raise ValueError("internal-file-path-must-not-be-visible")
        result = _overlay(
            self.rows if source_id == "synthetic-a" else _rows("wide"),
            self.sources[source_id],
        )
        if self.mode == "stale":
            result["source_revision"] = "e" * 64
        return result

    def imported_word_reference(
        self, batch, source_id, digest, start, end, *, expected_revision
    ):
        self.reference_calls.append(
            (batch, source_id, digest, start, end, expected_revision)
        )
        return {
            "materials": "\n".join(
                f"[Word区块{block['index']}]\n{block['text']}"
                for block in self.sources[source_id]["blocks"]
                if start <= block["index"] <= end
            ),
            "warnings": [],
        }


@pytest.mark.parametrize("mode", ["ok", "raise", "stale"])
def test_import_dialog_attaches_optional_grid_but_text_selection_confirm_remains_authoritative(
    qt_app, mode
):
    facade = SyntheticFacade(mode)
    dialog = ImportWordDialog(facade, "synthetic-batch")
    try:
        dialog.show()
        QTest.qWait(20)
        assert facade.table_calls == [
            ("synthetic-batch", "synthetic-a", "a" * 64, "1" * 64)
        ]
        assert bool(_tables(dialog.reader)) is (mode == "ok")
        assert facade.reference_calls == []
        dialog._set_range(2, 3)
        selection = dialog._selection_key()
        if mode == "ok":
            dialog.reader.table_mode_button.click()
        else:
            assert _readable_table(facade.rows) in dialog.reader.browser.toPlainText()
        assert dialog._selection_key() == selection
        assert "末段完整保留" in dialog.reader.browser.toPlainText()
        assert (
            "internal-file-path"
            not in dialog.reader.browser.toPlainText()
            + dialog.reader.search_status.text()
        )
        dialog._compile_preview()
        assert dialog.import_button.isEnabled()
        assert _readable_table(facade.rows) in dialog._preview_reference["materials"]
        dialog._confirm()
        assert dialog.result() == QDialog.DialogCode.Accepted
        assert dialog.reference is not None
        assert len(facade.reference_calls) == 2
        assert all(call[3:5] == (2, 3) for call in facade.reference_calls)
    finally:
        dialog.close()
        dialog.deleteLater()
        qt_app.processEvents()


@pytest.mark.parametrize("width", [420, 900])
def test_grid_and_fallback_fit_actual_dialog_without_hidden_horizontal_content(
    qt_app, width
):
    facade = SyntheticFacade()
    dialog = ImportWordDialog(facade, "synthetic-batch")
    try:
        dialog.resize(width, 900)
        dialog.show()
        for source in (0, 1):
            dialog.source_combo.setCurrentIndex(source)
            dialog.reader.scroll_to_block(2)
            QTest.qWait(40)
            browser = dialog.reader.browser
            assert dialog.width() == width
            assert browser.document().size().width() <= browser.viewport().width()
            assert browser.horizontalScrollBar().maximum() == 0
            if source == 0:
                assert _tables(dialog.reader)[0].cellAt(1, 0).rowSpan() == 2
            else:
                assert not _tables(dialog.reader)
                assert (
                    facade.sources["synthetic-b"]["blocks"][1]["text"]
                    in browser.toPlainText()
                )
    finally:
        dialog.close()
        dialog.deleteLater()
        qt_app.processEvents()


def write_synthetic_screenshots(output_directory):
    """Optional owned-widget QA capture; never opens the real workbench."""
    app = QApplication.instance() or QApplication([])
    for font in ("C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/msyhbd.ttc"):
        if Path(font).is_file():
            QFontDatabase.addApplicationFont(font)
    app.setFont(QFont("Microsoft YaHei UI", 10))
    output = Path(output_directory)
    output.mkdir(parents=True, exist_ok=True)
    dialog = ImportWordDialog(SyntheticFacade(), "synthetic-batch")
    files = []
    try:
        dialog.show()
        for width in (900, 420):
            dialog.resize(width, 900)
            dialog.reader.scroll_to_block(2)
            QTest.qWait(50)
            browser = dialog.reader.browser
            assert browser.document().size().width() <= browser.viewport().width()
            assert browser.horizontalScrollBar().maximum() == 0
            path = output / f"synthetic-word-grid-{width}.png"
            assert dialog.grab().save(str(path))
            files.append(str(path.resolve()))
    finally:
        dialog.close()
        dialog.deleteLater()
        app.processEvents()
    return files
