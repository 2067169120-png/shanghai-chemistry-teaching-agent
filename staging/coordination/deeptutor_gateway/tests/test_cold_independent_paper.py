"""Real default facade over owned synthetic Word archives; no Office/provider."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import hashlib
import json

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw

from integrations.deeptutor_shchem_v1.desktop_facade import build_default_facade
from integrations.deeptutor_shchem_v1.desktop_paths import DesktopPaths
from integrations.deeptutor_shchem_v1.desktop_mixed_paper_drafts import DRAFT_ID
from integrations.deeptutor_shchem_v1.desktop_mixed_paper_service import MixedPaperError
from integrations.deeptutor_shchem_v1.desktop_state import DesktopStateError, DraftConflictError
from integrations.deeptutor_shchem_v1.desktop_word_question_attributes import WordQuestionAttributeError


def sha(raw):
    return hashlib.sha256(raw).hexdigest()


def json_file(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


@pytest.fixture
def cold(tmp_path):
    workspace = tmp_path / "cold-install"
    (workspace / "integrations/deeptutor_shchem_v1").mkdir(parents=True)
    state = tmp_path / "external-userdata"
    paths = DesktopPaths.from_workspace(workspace, state_root=state)
    facade = build_default_facade(paths)
    assert type(facade._providers).__name__ == "ModelProviderSettingsStore"
    assert type(facade._providers._backend).__name__ == "WindowsCredentialManagerBackend"
    assert paths.uses_personal_library and not (workspace / "sh-chem-db").exists()
    source = tmp_path / "authored-source.docx"
    picture = Image.new("RGB", (160, 90), "white")
    ImageDraw.Draw(picture).line((20, 65, 80, 20, 140, 65), fill="black", width=3)
    stream = BytesIO()
    picture.save(stream, format="PNG")
    image_bytes = stream.getvalue()
    doc = Document()
    doc.add_heading("合成软件验收材料", 0)
    doc.add_paragraph("【例1】软件测试题面：观察下列合成线图并记录图中折线段数。")
    doc.add_picture(BytesIO(image_bytes), width=Inches(1.6))
    doc.add_paragraph("【答案】软件测试参考：2段。")
    doc.add_paragraph("【例2】软件测试题面：在空格中写出测试编号。")
    doc.add_paragraph("【答案】软件测试参考：2。")
    doc.save(source)
    raw = source.read_bytes()
    receipt = facade.save_visual_import_batch(handout_files=(source,), source_type="教师讲义")
    descriptor = deepcopy(facade._saved_visual_import_batch(receipt.batch_id))
    source_rows = facade.imported_word_sources(receipt.batch_id)
    assert len(source_rows) == 1 and source_rows[0]["role"] == "handout"
    assert source_rows[0]["source_sha256"] == sha(raw)
    restored = facade._restore_visual_import_sources(descriptor)
    assert len(restored) == 1 and restored[0].content == raw
    rows = facade.word_question_catalog()["items"]
    assert len(rows) == 2 and all(r["selection_ready"] and r["export_ready"] for r in rows)
    selections = [{"key": row["key"], "revision": row["revision"], "points": 2} for row in rows]
    assert facade.add_word_questions_to_basket(selections) == 2
    archive = facade._visual_import_root / descriptor["sources"][0]["archive_relative_path"]
    return {"facade": facade, "paths": paths, "source": source, "raw": raw,
            "descriptor": descriptor, "rows": rows, "batch_id": receipt.batch_id,
            "archive": archive, "image_bytes": image_bytes,
            "tax": workspace / "sh-chem-db/kb/knowledge_taxonomy.json",
            "directory": workspace / "sh-chem-db/kb/classification/supplemental_wechat_textbook_tagging_v1_2026-08-27/textbook_directory_nodes.json"}


def preview_request(session, editor):
    payload = deepcopy(editor.current.record["payload"])
    return {"schema_version": "shchem.desktop-mixed-paper-request.v1", **payload["settings_ui"],
            "basket_sha256": session.paper_basket_projection()["basket_sha256"],
            "section_order": payload["order"],
            "settings_by_key": {key: payload["settings"][key] for key in payload["order"]},
            "draft_binding": editor.check()}


def create_preview(cold):
    facade = cold["facade"]
    session = facade.independent_paper_library().create()
    editor = session.open_mixed_paper_draft_session()
    editor.read(session.paper_basket_projection())
    return session, editor, session.create_paper_preview(preview_request(session, editor))


def test_baseline_cold_failure_is_strict_catalog_before_commit(cold):
    facade = cold["facade"]
    frozen = facade.state_store.path.read_bytes()
    with pytest.raises(WordQuestionAttributeError, match="教材目录尚未就绪"):
        facade.independent_paper_library().create()
    assert facade.state_store.path.read_bytes() == frozen
    assert not facade.independent_paper_library().list()


def test_real_cold_create_save_reopen_content_preview_preserves_source_and_basket(cold):
    facade = cold["facade"]
    public = deepcopy(facade.state_store.basket_snapshot())
    session, editor, preview = create_preview(cold)
    payload = deepcopy(editor.current.record["payload"])
    payload["order"].reverse()
    payload["settings_ui"]["title"] = "合成冷启动保存卷"
    payload["settings"][payload["order"][0]]["points"] = 7.5
    editor.save(payload)
    # Fresh default service objects read the persisted archive and saved scope.
    fresh = build_default_facade(cold["paths"])
    reopened = fresh.independent_paper_library().open(session.paper_id)
    second = reopened.open_mixed_paper_draft_session()
    second.read(reopened.paper_basket_projection())
    assert second.current.record["payload"] == payload
    current = reopened.create_paper_preview(preview_request(reopened, second))
    assert [s["key"] for s in current.preview_model["sections"]] == payload["order"]
    student = "\n".join(b.get("text", "") for s in current.preview_model["sections"] for b in s["student_blocks"])
    teacher = "\n".join(b.get("text", "") for s in current.preview_model["sections"] for b in s["teacher_blocks"])
    assert "软件测试题面" in student and "软件测试参考" not in student
    assert "软件测试参考" in teacher
    images = [b for s in current.preview_model["sections"] for b in s["student_blocks"] if b["kind"] == "image"]
    assert len(images) == 1
    assert reopened.paper_preview_image(current.preview_id, images[0]["image_id"])["data"] == cold["image_bytes"]
    assert not current.export_ready and not preview.export_ready
    assert facade.state_store.basket_snapshot() == public
    assert fresh._saved_visual_import_batch(cold["batch_id"]) == cold["descriptor"]
    assert cold["source"].read_bytes() == cold["raw"] == cold["archive"].read_bytes()
    catalog = fresh._word_questions()._read_dependency_attribute_catalog()
    assert len(catalog["knowledge_points"]) == 19 and catalog["nodes"] == []
    assert len(catalog["warnings"]) == 2
    assert not cold["tax"].exists() and not cold["directory"].exists()
    # Clearing the shared basket cannot remove the independently saved items.
    fresh.state_store.clear_basket()
    again = fresh.independent_paper_library().open(session.paper_id)
    assert len(again.paper_basket_projection()["items"]) == 2
    assert fresh.state_store.basket() == []


def test_catalog_and_bulk_annotation_remain_strict_without_optional_files(cold):
    words = cold["facade"]._word_questions()
    frozen = cold["facade"].state_store.path.read_bytes()
    with pytest.raises(WordQuestionAttributeError, match="教材目录尚未就绪"):
        words._read_attribute_catalog()
    with pytest.raises(Exception, match="未提交标签") as blocked:
        words.annotate_imported_batch(cold["batch_id"])
    assert isinstance(blocked.value.__cause__, WordQuestionAttributeError)
    assert "教材目录尚未就绪" in str(blocked.value.__cause__)
    catalog = words.catalog()
    assert catalog["attribute_catalog"] is None and len(catalog["items"]) == 2
    assert cold["facade"].state_store.path.read_bytes() == frozen


@pytest.mark.parametrize("which", ["taxonomy_json", "directory_json", "taxonomy_shape", "directory_shape"])
def test_installed_corrupt_catalog_never_falls_back_or_partially_creates_paper(cold, which):
    if which == "taxonomy_json":
        cold["tax"].parent.mkdir(parents=True)
        cold["tax"].write_text("{", encoding="utf-8")
    elif which == "directory_json":
        cold["directory"].parent.mkdir(parents=True)
        cold["directory"].write_text("{", encoding="utf-8")
    elif which == "taxonomy_shape":
        json_file(cold["tax"], {"dimensions": {"knowledge_points": ["broken-row"]}})
    else:
        json_file(cold["directory"], {"nodes": ["broken-row"]})
    facade = cold["facade"]
    frozen = facade.state_store.path.read_bytes()
    with pytest.raises((ValueError, TypeError, KeyError)):
        facade.independent_paper_library().create()
    assert facade.state_store.path.read_bytes() == frozen
    assert facade.independent_paper_library().list() == []


@pytest.mark.parametrize("change", ["archive", "labels", "catalog_install", "catalog_change", "catalog_corrupt"])
def test_real_changes_block_saved_open_projection_and_new_preview_without_state_write(cold, change):
    if change == "catalog_change":
        json_file(cold["tax"], {"dimensions": {"knowledge_points": [{"id": "K01", "name": "软件测试分类"}]}})
        json_file(cold["directory"], {"nodes": []})
    session, editor, _ = create_preview(cold)
    request = preview_request(session, editor)
    facade = cold["facade"]
    if change == "archive":
        cold["archive"].write_bytes(cold["raw"] + b"owned synthetic corruption")
    elif change == "labels":
        row = cold["rows"][0]
        options = facade.word_question_attribute_options(row["key"], row["revision"])
        facade.word_question_save_attributes(row["key"], row["revision"],
            {"teacher_note": "合成软件测试标签变更"},
            expected_attribute_revision=options["attributes"]["revision"],
            expected_stored_revision=options["stored_revision"])
    elif change == "catalog_install":
        # Even matching K IDs change dependency provenance and warnings.
        points = facade._word_questions()._read_dependency_attribute_catalog()["knowledge_points"]
        json_file(cold["tax"], {"dimensions": {"knowledge_points": points}})
        json_file(cold["directory"], {"nodes": []})
    elif change == "catalog_change":
        json_file(cold["directory"], {"nodes": [{"id": "synthetic-node", "name": "软件测试节点"}]})
    else:
        cold["tax"].parent.mkdir(parents=True)
        cold["tax"].write_text("{", encoding="utf-8")
    frozen = facade.state_store.path.read_bytes()
    for operation in (lambda: facade.independent_paper_library().open(session.paper_id),
                      session.paper_basket_projection,
                      lambda: session.create_paper_preview(request)):
        with pytest.raises(Exception):
            operation()
        assert facade.state_store.path.read_bytes() == frozen


@pytest.mark.parametrize("when", ["before", "after"])
def test_real_save_uncertainty_is_not_retried_and_reopen_observes_disk(cold, monkeypatch, when):
    session, editor, _ = create_preview(cold)
    store = cold["facade"].state_store
    original = store._write_unlocked
    calls = []
    def failed(value, **kwargs):
        calls.append(1)
        if when == "after":
            original(value, **kwargs)
        raise OSError("owned synthetic uncertain write")
    monkeypatch.setattr(store, "_write_unlocked", failed)
    payload = deepcopy(editor.current.record["payload"])
    payload["settings_ui"]["title"] = "合成不确定保存"
    with pytest.raises(DesktopStateError):
        editor.save(payload)
    with pytest.raises(DesktopStateError):
        editor.save(payload)
    assert calls == [1]
    monkeypatch.setattr(store, "_write_unlocked", original)
    reopened = cold["facade"].independent_paper_library().open(session.paper_id)
    assert (reopened.state_store.draft_snapshot(DRAFT_ID).record["payload"] == payload) == (when == "after")
