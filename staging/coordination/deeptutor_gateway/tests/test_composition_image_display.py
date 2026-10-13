"""Per-occurrence image limits change display, never complete source content."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document
from docx.shared import Mm
from PIL import Image
from test_composition_render import app, text

from integrations.deeptutor_shchem_v1.desktop_composition_image_display import (
    SCHEMA, arrange_image_resource, cap_image_size, image_display_rows, set_image_display,
)
from integrations.deeptutor_shchem_v1.desktop_composition_inline import NativeParagraphDisplay
from integrations.deeptutor_shchem_v1.desktop_composition_render import (
    PADDING, _content, _fit_end, _keep_groups, arrange_composition, layout_composition,
)
from integrations.deeptutor_shchem_v1.desktop_composition_rich_render import flatten_elements
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    FORMATS, RENDER_REVISION, add_resource, block, digest, new_composition, resource,
    resource_frames, set_click_reveal, set_student_arrangement, validate_resource,
)
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import _word_images
from integrations.deeptutor_shchem_v1.offline_word_image_label_review import _objects


def personal(title, blocks):
    return resource('personal_text', title, blocks,
        {'source_path': '教师输入', 'source_sha256': 'unknown', 'source_revision': 'unknown'})


def composition(*items, format='16:9'):
    value = new_composition(format); width, height = FORMATS[format]
    for item in items:
        value = add_resource(value, item, {'page_id': value['pages'][0], 'x': 12, 'y': 12,
            'width': width-24, 'height': height-24})
    return value


@pytest.fixture
def picture(tmp_path, app):
    buffer = BytesIO(); Image.new('RGB', (240, 120), '#4b7c9a').save(buffer, 'PNG')
    raw = buffer.getvalue(); store = PreparationImageStore(tmp_path/'images')
    asset = store.import_bytes(raw, '合成原图', '教师', '显示尺寸核对')
    item = personal('从完整图中可以观察到什么？', [block('image', 'personal', asset['asset_id'], 1),
        block('text', 'personal', '完整的条件和解释保持不变。', 2)])
    return composition(item), [asset], {asset['asset_id']: raw}, store


@pytest.fixture
def repeated_inline(tmp_path, app):
    buffer = BytesIO(); Image.new('RGB', (240, 120), '#597b92').save(buffer, 'PNG')
    raw = buffer.getvalue(); store = PreparationImageStore(tmp_path/'images')
    asset = store.import_bytes(raw, '合成重复原图', '教师', '每次出现分别设置')
    doc = Document(); p = doc.add_paragraph('先看H'); p.add_run('2').font.subscript = True; p.add_run('O。')
    p.add_run().add_picture(BytesIO(raw), width=Mm(35)); p.add_run('再看同一原图。')
    p.add_run().add_picture(BytesIO(raw), width=Mm(35)); p.add_run('完整尾部条件。')
    answer = doc.add_paragraph('仅教师原图及解析。')
    answer.add_run().add_picture(BytesIO(raw), width=Mm(35))
    document = BytesIO(); doc.save(document); display = NativeParagraphDisplay(document.getvalue())
    blocks = []
    for index, role in ((1, 'question'), (2, 'answer')):
        element = display.elements[index-1]
        row = block('text', role, display.reader.read(element).text, index)
        originals = [a for a in _word_images(display.elements, display.doc) if a['block_index'] == index]
        objects = _objects(element, display.doc, index, originals)
        pictures = [block('image', role, asset['asset_id'], index) for _ in objects]
        display.bind(row, pictures, objects); blocks.extend([row, *pictures])
    item = resource('complete_question', '原图的两次出现分别说明什么？', blocks,
        {'source_path': '合成原题.docx', 'source_sha256': sha256(document.getvalue()).hexdigest(),
         'source_revision': 'b'*64, 'question_key': 'synthetic-repeat', 'scope_review_sha256': 'c'*64,
         'scope_status': 'complete_native_scope',
         'hierarchy': dict.fromkeys(('paper_id', 'theme_big_question_id', 'printed_question_id', 'atomic_part_id'), 'unknown')})
    return composition(item), [asset], {asset['asset_id']: raw}, store


def pictures(layout):
    return [p for elements in layout['pages'].values() for p in flatten_elements(elements)
            if p['kind'] == 'image']


def test_rows_bind_each_original_occurrence_and_return_detached_limits(repeated_inline):
    value, _, _, _ = repeated_inline; item = value['resources'][0]
    rows = image_display_rows(item)
    assert len(rows) == 3 and len({r['image_block_id'] for r in rows}) == 3
    assert len({r['asset_id'] for r in rows}) == 1
    assert [r['source_index'] for r in rows] == [1, 1, 2]
    assert [r['role'] for r in rows] == ['question', 'question', 'answer']
    assert all(r['limits'] == {} for r in rows)
    changed = set_image_display(value, item['id'], {rows[1]['image_block_id']: {'max_width_mm': 20}})
    returned = image_display_rows(changed['resources'][0]); returned[1]['limits']['max_width_mm'] = 999
    assert changed['resources'][0]['image_display']['limits'][rows[1]['image_block_id']] == {'max_width_mm': 20}
    assert 'image_display' not in item


@pytest.mark.parametrize('invalid', [None, [], {'max_width_mm': 0}, {'max_width_mm': -1},
    {'max_width_mm': True}, {'max_width_mm': float('nan')}, {'max_height_mm': float('inf')},
    {'max_width_mm': 400.1}, {'max_width_mm': 10**1000}, {}, {'width': 20},
    {'max_width_mm': 20, 'crop': True}])
def test_invalid_axis_limits_are_value_errors_without_source_or_draft_mutation(picture, invalid):
    value, _, _, _ = picture; item = value['resources'][0]; before = deepcopy(value)
    with pytest.raises(ValueError):
        set_image_display(value, item['id'], {item['blocks'][0]['id']: invalid})
    assert value == before


@pytest.mark.parametrize('invalid', ['text_block', 'foreign_block', 'schema', 'extra', 'empty_field', 'resource'])
def test_display_schema_and_original_image_binding_are_strict(picture, invalid):
    value, _, _, _ = picture; item = value['resources'][0]; before = deepcopy(value)
    if invalid in {'text_block', 'foreign_block', 'resource'}:
        key = item['blocks'][1]['id'] if invalid == 'text_block' else 'CB-'+'e'*32
        identity = 'CR-'+'e'*32 if invalid == 'resource' else item['id']
        with pytest.raises(ValueError): set_image_display(value, identity, {key: {'max_height_mm': 20}})
    else:
        copied = deepcopy(item)
        copied['image_display'] = {'schema': SCHEMA, 'limits': {item['blocks'][0]['id']: {'max_height_mm': 20}}}
        if invalid == 'schema': copied['image_display']['schema'] = 'unknown'
        elif invalid == 'extra': copied['image_display']['source'] = 'not allowed'
        else: copied['image_display']['limits'] = {}
        with pytest.raises(ValueError): validate_resource(copied)
    assert value == before


@pytest.mark.parametrize('limits', [{'max_width_mm': 30}, {'max_height_mm': 12},
    {'max_width_mm': 30, 'max_height_mm': 12}, {'max_width_mm': 400}])
def test_standalone_limits_apply_after_default_fit_and_keep_whole_picture_ratio(picture, limits):
    value, assets, data, store = picture; item = value['resources'][0]; before = deepcopy(value)
    default = arrange_composition(value, assets, data)
    original = pictures(layout_composition(default, assets, data, student=False))[0]
    changed = set_image_display(value, item['id'], {item['blocks'][0]['id']: limits})
    changed = arrange_image_resource(changed, item['id'], assets, data)
    for student in (False, True):
        layout = layout_composition(changed, assets, data, student=student)
        assert not layout['errors']
        image = pictures(layout)[0]
        expected = cap_image_size(original['width'], original['height'], limits)
        assert (image['width'], image['height']) == pytest.approx(expected)
        assert image['width']/image['height'] == pytest.approx(2)
        assert text(layout, item['blocks'][1]['id']) == item['blocks'][1]['value']
    assert changed['resources'][0]['original'] == item['original']
    assert changed['resources'][0]['original_sha256'] == item['original_sha256']
    assert changed['resources'][0]['locked'] == item['locked'] and value == before
    assert store.load(assets[0]) == data[assets[0]['asset_id']]
    assert sha256(data[assets[0]['asset_id']]).hexdigest() == assets[0]['sha256']


def test_same_inline_bytes_get_independent_caps_and_teacher_only_picture_stays_hidden(repeated_inline):
    value, assets, data, store = repeated_inline; item = value['resources'][0]; before = deepcopy(value)
    rows = image_display_rows(item)
    changed = set_image_display(value, item['id'], {
        rows[0]['image_block_id']: {'max_width_mm': 18},
        rows[1]['image_block_id']: {'max_height_mm': 5},
        rows[2]['image_block_id']: {'max_width_mm': 12}})
    changed = arrange_image_resource(changed, item['id'], assets, data)
    teacher = layout_composition(changed, assets, data, student=False)
    student = layout_composition(changed, assets, data, student=True)
    assert not teacher['errors'] and not student['errors']
    assert [(p['width'], p['height']) for p in pictures(teacher)] == pytest.approx([(18, 9), (10, 5), (12, 6)])
    assert [p['image_block_id'] for p in pictures(teacher)] == [r['image_block_id'] for r in rows]
    assert len(pictures(student)) == 2 and '仅教师' not in text(student)
    flat = list(flatten_elements([e for p in student['pages'].values() for e in p]))
    assert ''.join(e.get('text', '') for e in flat).endswith('完整尾部条件。')
    assert any(e.get('text') == '2' and e['font_px'] < 32 for e in flat)
    assert changed['resources'][0]['blocks'] == item['blocks']
    assert changed['resources'][0]['original'] == item['original'] and value == before
    assert store.load(assets[0]) == data[assets[0]['asset_id']]


def test_reset_removes_optional_field_and_returns_exact_default_display(repeated_inline):
    value, assets, data, _ = repeated_inline; item = value['resources'][0]
    base = arrange_composition(value, assets, data)
    changed = set_image_display(value, item['id'], {image_display_rows(item)[0]['image_block_id']: {'max_width_mm': 10}})
    reset = set_image_display(changed, item['id'], {})
    assert reset == value and 'image_display' not in reset['resources'][0]
    default = pictures(layout_composition(base, assets, data, student=False))
    after = pictures(layout_composition(arrange_image_resource(reset, item['id'], assets, data), assets, data, student=False))
    assert [(p['width'], p['height']) for p in default] == [(p['width'], p['height']) for p in after]
    assert default[0]['width'] == pytest.approx(35*32/15)


@pytest.mark.parametrize('format', FORMATS)
def test_local_reflow_keeps_unrelated_resources_shared_pages_and_intentional_blanks(picture, format):
    value, assets, data, _ = picture; item = value['resources'][0]
    other = personal('教师已有环节', [block('text', 'personal', '未完成的其他素材。'*50)])
    original = composition(item, other, format=format)
    pages = ['CP-'+uuid4().hex for _ in range(5)]; original['pages'] = pages
    width, height = FORMATS[format]
    original['placements'] = {
        item['id']: {'page_id': pages[1], 'x': 12, 'y': 12, 'width': width-24, 'height': 20,
                     'continuations': [{'page_id': pages[3], 'x': 12, 'y': 12, 'width': width-24, 'height': height-24}]},
        other['id']: {'page_id': pages[1], 'x': 12, 'y': 40, 'width': width-24, 'height': 12}}
    original = set_student_arrangement(original, 1)
    if format == '16:9': original = set_click_reveal(original, item['id'], True)
    before = deepcopy(original)
    changed = set_image_display(original, item['id'], {item['blocks'][0]['id']: {'max_height_mm': 20}})
    changed = arrange_image_resource(changed, item['id'], assets, data)
    assert original == before
    assert changed['resources'][1] == before['resources'][1]
    assert changed['placements'][other['id']] == before['placements'][other['id']]
    assert changed['pages'][0] == pages[0] and changed['pages'][1] not in pages
    assert [p for p in changed['pages'] if p in pages] == [pages[0], pages[1], pages[2], pages[4]]
    assert changed['student_arrangement'] == before['student_arrangement']
    assert changed.get('click_reveals') == before.get('click_reveals')
    assert layout_composition(changed, assets, data, student=False)['errors']  # unrelated unfinished draft stays so


@pytest.mark.parametrize('hidden', ['teacher_only', 'click'])
def test_selected_resource_can_have_a_legitimate_empty_student_scope(picture, hidden):
    value, assets, data, _ = picture; source = value['resources'][0]
    if hidden == 'teacher_only':
        item = personal('仅教师的完整图', [block('image', 'teacher_note', assets[0]['asset_id'])])
        value = composition(item)
    else:
        item = source; value = set_click_reveal(value, item['id'], True)
    changed = set_image_display(value, item['id'], {image_display_rows(item)[0]['image_block_id']: {'max_height_mm': 20}})
    changed = arrange_image_resource(changed, item['id'], assets, data)
    teacher = layout_composition(changed, assets, data, student=False)
    student = layout_composition(changed, assets, data, student=True)
    assert not teacher['errors'] and len(pictures(teacher)) == 1
    assert not pictures(student) and not student['pages']
    assert changed.get('click_reveals') == value.get('click_reveals')
    assert changed['resources'][0]['blocks'] == item['blocks']


def knowledge_item(texts):
    from integrations.deeptutor_shchem_v1.desktop_handout_knowledge import SOURCE_SCHEMA, SECTION_SCHEMA
    source = dict(schema=SOURCE_SCHEMA, batch_id='synthetic', source_id='source', source_name='合成讲义.docx',
        source_sha256='a'*64, source_role='handout', source_path='合成讲义.docx', source_revision='b'*64)
    document = Document()
    for value in texts: document.add_paragraph(value)
    raw = BytesIO(); document.save(raw); display = NativeParagraphDisplay(raw.getvalue())
    blocks = []
    for index, value in enumerate(texts, 1):
        b = block('text', 'knowledge', value, index); display.bind(b, [], []); blocks.append(b)
    snapshot = dict(schema=SECTION_SCHEMA, outline_revision='c'*64,
        section=dict(id='word-heading-b1', level=3, title='完整知识', start=1, end=len(texts)),
        blocks=[dict(index=i, text=t, images=[]) for i, t in enumerate(texts, 1)])
    return resource('handout_knowledge', '合成讲义.docx · 完整知识', blocks,
        dict(source_path=source['source_path'], source_sha256=source['source_sha256'], source_revision=source['source_revision'],
             handout_source=source, section_snapshot=snapshot, image_assets=[],
             source_roles=dict(teacher_only=False, document_role='handout', images=[])))


def test_short_numbered_original_heading_without_colon_moves_with_next_complete_line_group(app):
    for count in range(5, 17):
        texts = ['\n'.join('前序完整条件保持原文。' for _ in range(count)),
            '2．与化合物类型的关系', '后一段完整的适用条件和解释。'*18]
        item = knowledge_item(texts)
        chunks = _content(item, 296, {}, {}, student=False, format='16:9')
        old = _keep_groups([dict(c) for c in chunks], 130, numbered_headings=False)
        end = _fit_end(old, 0, 156-2*PADDING, strict=True)
        heading = [i for i, c in enumerate(old) if c.get('block_id') == item['blocks'][1]['id'] and c['kind'] == 'rich_line']
        body = [i for i, c in enumerate(old) if c.get('block_id') == item['blocks'][2]['id'] and c['kind'] == 'rich_line']
        if heading and heading[-1] < end <= body[0]: break
    else: pytest.fail('未构造到短编号标题留在旧页尾的条件')
    before = deepcopy(item)
    arranged = arrange_composition(composition(item), [], {})
    for student in (False, True):
        layout = layout_composition(arranged, [], {}, student=student)
        assert not layout['errors']
        heading_page = next(p for p, elements in layout['pages'].items()
            if any(e.get('block_id') == item['blocks'][1]['id'] and e['kind'] == 'rich_line' for e in elements))
        first_body_page = next(p for p, elements in layout['pages'].items()
            if any(e.get('block_id') == item['blocks'][2]['id'] and e['kind'] == 'rich_line' for e in elements))
        assert heading_page == first_body_page
        for b in item['blocks']:
            parts = [p for elements in layout['pages'].values() for p in flatten_elements(elements)
                     if p['kind'] == 'text' and p.get('block_id') == b['id']]
            assert ''.join(p['text'] for p in parts) == b['value'].replace('\n', '')
            assert all(p['font_px'] == 32 for p in parts)
    assert item == before and arranged['resources'][0] == before


def test_numbered_heading_group_does_not_exceed_capacity_or_infer_picture_heading():
    def chunks(next_height):
        return [dict(kind='text', text='2．完整标题', height=10, block_id='a', _paragraph=('a', 0)),
                dict(kind='text', text='完整下一行', height=next_height, block_id='b', _paragraph=('b', 0))]
    ordinary = chunks(10); _keep_groups(ordinary, 25)
    explicit = chunks(10); _keep_groups(explicit, 25, numbered_headings=True)
    too_tall = chunks(30); _keep_groups(too_tall, 25, numbered_headings=True)
    assert not ordinary[0]['_keep_next'] and explicit[0]['_keep_next'] and not too_tall[0]['_keep_next']
    picture_heading = [dict(kind='rich_line', height=10, block_id='a', _paragraph=('a', 0),
        _parts=[dict(kind='text', text='2．标题'), dict(kind='image', height=10)]),
        *[dict(kind='text', text='下一段完整行', height=10, block_id='b', _paragraph=('b', 0)) for _ in range(4)]]
    _keep_groups(picture_heading, 25, numbered_headings=True)
    assert not picture_heading[0]['_keep_next']


@pytest.mark.parametrize('kind', ['personal_text', 'complete_question'])
def test_numbered_no_colon_rule_does_not_expand_to_personal_text_or_question_options(app, kind):
    texts = ['\n'.join('前序完整条件。' for _ in range(12)), '2．短编号条件', 'A．选项或其他原文。'*20]
    role = 'question' if kind == 'complete_question' else 'personal'
    blocks = [block('text', role, t, i) for i, t in enumerate(texts, 1)]
    evidence = {'source_path': '合成来源', 'source_sha256': 'unknown', 'source_revision': 'unknown'}
    if kind == 'complete_question': evidence.update(source_sha256='a'*64, source_revision='b'*64,
        question_key='synthetic-option', scope_review_sha256='c'*64, scope_status='complete_native_scope',
        hierarchy=dict.fromkeys(('paper_id', 'theme_big_question_id', 'printed_question_id', 'atomic_part_id'), 'unknown'))
    item = resource(kind, '完整条件', blocks, evidence)
    chunks = _content(item, 296, {}, {}, student=False, format='16:9')
    heading = next(c for c in chunks if c.get('block_id') == blocks[1]['id'] and c['kind'] == 'text')
    assert not heading['_keep_next']


def test_renderer_revision_changes_new_fingerprint_but_old_files_remain_readable(picture, tmp_path, monkeypatch):
    from types import SimpleNamespace
    from test_lesson_design import sample
    from integrations.deeptutor_shchem_v1.desktop_lesson_design import content_fingerprint
    from integrations.deeptutor_shchem_v1.desktop_lesson_composition import enable_composition
    from integrations.deeptutor_shchem_v1 import desktop_lesson_composition as composition_module
    from integrations.deeptutor_shchem_v1.desktop_lesson_output import checked_file
    value, assets, _, _ = picture
    payload = sample(); payload['lesson_design'] = enable_composition(payload['lesson_design'])
    payload['lesson_design']['composition'] = value; payload['image_assets'] = assets
    before = content_fingerprint(payload)
    original_value = deepcopy(value)
    item = value['resources'][0]
    payload['lesson_design']['composition'] = set_image_display(value, item['id'], {item['blocks'][0]['id']: {'max_height_mm': 20}})
    display_changed = content_fingerprint(payload)
    assert display_changed != before
    assert value == original_value
    assert payload['lesson_design']['composition']['resources'][0]['original'] == item['original']
    assert payload['lesson_design']['composition']['resources'][0]['blocks'] == item['blocks']
    assert composition_module.RENDER_REVISION == RENDER_REVISION and RENDER_REVISION
    with monkeypatch.context() as revision_probe:
        revision_probe.setattr(composition_module, 'RENDER_REVISION', RENDER_REVISION + '.synthetic-fixture-probe')
        assert content_fingerprint(payload) != display_changed
    assert content_fingerprint(payload) == display_changed
    facade = SimpleNamespace(paths=SimpleNamespace(task_root=tmp_path/'tasks'))
    root = facade.paths.task_root/'preparation-v1/node-exports'/('OUT-'+'a'*32)
    root.mkdir(parents=True); raw = b'original historical output'; (root/'lesson_plan.docx').write_bytes(raw)
    payload['lesson_design']['exports'].append({'id': 'OUT-'+'a'*32, 'fingerprint': 'b'*64,
        'created_at': '2026-10-07T00:00:00Z', 'files': [{'name': 'lesson_plan.docx', 'sha256': sha256(raw).hexdigest()}]})
    assert checked_file(facade, payload['lesson_design'], 'OUT-'+'a'*32, 'lesson_plan.docx').read_bytes() == raw


def word_output(value, assets, data, directory, *, notes=False, node_image=False):
    from test_lesson_design import sample
    from integrations.deeptutor_shchem_v1.desktop_lesson_composition import enable_composition
    from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents
    payload = sample(); payload['lesson_design'] = enable_composition(payload['lesson_design'])
    payload['lesson_design']['composition'] = deepcopy(value)
    payload['lesson_design']['classroom_notes'] = notes
    payload['image_assets'] = assets
    if node_image: payload['lesson_design']['nodes'][0]['image_ids'] = [assets[0]['asset_id']]
    directory.mkdir(); before = deepcopy(payload)
    write_documents(payload, directory, data)
    assert payload == before
    return Document(directory/'lesson_plan.docx'), Document(directory/'student_worksheet.docx')


def word_image_sizes(document):
    # These are actual wp:extent values serialized in the DOCX package.
    return [(int(e.get('cx'))/36000, int(e.get('cy'))/36000)
            for e in document.element.xpath('.//wp:inline/wp:extent')]


def word_image_hashes(document):
    hashes = []
    for node in document.element.xpath('.//a:blip'):
        from docx.oxml.ns import qn
        hashes.append(sha256(document.part.rels[node.get(qn('r:embed'))].target_part.blob).hexdigest())
    return hashes


@pytest.mark.parametrize('notes', [False, True], ids=['student-worksheet', 'classroom-notes'])
def test_word_inline_caps_use_exact_occurrences_and_keep_source_bytes_scripts_and_teacher_roles(repeated_inline, tmp_path, notes):
    value, assets, data, _ = repeated_inline; item = value['resources'][0]
    before = deepcopy(value); rows = image_display_rows(item)
    changed = set_image_display(value, item['id'], {
        rows[0]['image_block_id']: {'max_width_mm': 18},
        rows[1]['image_block_id']: {'max_height_mm': 5},
        rows[2]['image_block_id']: {'max_width_mm': 12}})
    teacher, student = word_output(changed, assets, data, tmp_path/'inline', notes=notes)
    assert word_image_sizes(teacher) == pytest.approx([(18, 9), (10, 5), (12, 6)])
    assert word_image_sizes(student) == pytest.approx([(18, 9), (10, 5)])
    assert word_image_hashes(teacher) == [assets[0]['sha256']]*3
    assert word_image_hashes(student) == [assets[0]['sha256']]*2
    assert any('仅教师原图及解析' in p.text for p in teacher.paragraphs)
    assert not any('仅教师原图及解析' in p.text for p in student.paragraphs)
    paragraph = next(p for p in student.paragraphs if p.text.startswith('先看H'))
    assert paragraph.text == '先看H2O。再看同一原图。完整尾部条件。'
    assert next(r for r in paragraph.runs if r.text == '2').font.subscript
    assert changed['resources'][0]['original'] == item['original'] and value == before


@pytest.mark.parametrize('notes', [False, True], ids=['student-worksheet', 'classroom-notes'])
def test_word_standalone_caps_fit_after_existing_geometry_and_reset_to_old_sizes(picture, tmp_path, notes):
    value, assets, data, _ = picture; item = value['resources'][0]
    defaults = word_output(value, assets, data, tmp_path/'default', notes=notes)
    changed = set_image_display(value, item['id'], {item['blocks'][0]['id']: {'max_height_mm': 20}})
    limited = word_output(changed, assets, data, tmp_path/'limited', notes=notes)
    reset = set_image_display(changed, item['id'], {})
    restored = word_output(reset, assets, data, tmp_path/'reset', notes=notes)
    assert word_image_sizes(defaults[0]) == pytest.approx([(165, 82.5)])
    assert word_image_sizes(defaults[1]) == pytest.approx([(130, 65)] if notes else [(165, 82.5)])
    for original, result, restored_doc in zip(defaults, limited, restored):
        assert word_image_sizes(result) == pytest.approx([(40, 20)])
        assert word_image_sizes(restored_doc) == word_image_sizes(original)
        assert word_image_hashes(result) == word_image_hashes(original) == [assets[0]['sha256']]
    assert reset == value and limited[0].paragraphs[-1].text == defaults[0].paragraphs[-1].text


def test_word_inline_default_and_reset_retain_old_authored_dimensions(repeated_inline, tmp_path):
    value, assets, data, _ = repeated_inline; item = value['resources'][0]
    baseline = word_output(value, assets, data, tmp_path/'base')
    changed = set_image_display(value, item['id'], {image_display_rows(item)[0]['image_block_id']: {'max_height_mm': 4}})
    restored = word_output(set_image_display(changed, item['id'], {}), assets, data, tmp_path/'restored')
    assert word_image_sizes(baseline[0]) == pytest.approx([(35, 17.5)]*3)
    assert word_image_sizes(baseline[1]) == pytest.approx([(35, 17.5)]*2)
    assert [word_image_sizes(d) for d in restored] == [word_image_sizes(d) for d in baseline]


def test_word_standalone_same_bytes_keep_per_block_caps_and_existing_student_repeat_policy(picture, tmp_path):
    value, assets, data, _ = picture
    first = block('image', 'personal', assets[0]['asset_id'], 1)
    second = block('image', 'personal', assets[0]['asset_id'], 2)
    hidden = block('image', 'teacher_note', assets[0]['asset_id'], 3)
    item = personal('三个原图块的不同用途', [first, second, hidden])
    changed = set_image_display(composition(item), item['id'], {
        first['id']: {'max_width_mm': 24}, second['id']: {'max_height_mm': 8},
        hidden['id']: {'max_height_mm': 5}})
    teacher, worksheet = word_output(changed, assets, data, tmp_path/'worksheet')
    assert word_image_sizes(teacher) == pytest.approx([(24, 12), (16, 8), (10, 5)])
    assert word_image_sizes(worksheet) == pytest.approx([(24, 12), (16, 8)])
    teacher_notes, notes = word_output(changed, assets, data, tmp_path/'notes', notes=True)
    assert word_image_sizes(teacher_notes) == pytest.approx([(24, 12), (16, 8), (10, 5)])
    assert word_image_sizes(notes) == pytest.approx([(24, 12)])
    assert word_image_hashes(teacher) == [assets[0]['sha256']]*3
    assert changed['resources'][0]['original'] == item['original']


@pytest.mark.parametrize('notes', [False, True])
def test_word_resource_cap_does_not_change_old_node_material_image_path(picture, tmp_path, notes):
    value, assets, data, _ = picture; item = value['resources'][0]
    changed = set_image_display(value, item['id'], {item['blocks'][0]['id']: {'max_height_mm': 10}})
    teacher, student = word_output(changed, assets, data, tmp_path/'node', notes=notes, node_image=True)
    assert word_image_sizes(teacher) == pytest.approx([(165, 82.5), (20, 10)])
    # Classroom notes already omit repeated standalone image bytes printed in a
    # teaching node. The resource cap must not silently resize that node image.
    assert word_image_sizes(student) == pytest.approx([(130, 65)] if notes else [(165, 82.5), (20, 10)])
