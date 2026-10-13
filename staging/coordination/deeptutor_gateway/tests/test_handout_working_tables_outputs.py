"""Native cell edits through the common canvas/Word projection and saved drafts.

All writes use pytest's explicitly assigned backend-worker basetemp. Only a
minimal QApplication is constructed; no workbench window or recovery service.
"""
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation
import pytest

from integrations.deeptutor_shchem_v1.desktop_handout_table_projection import source_table_projection
from integrations.deeptutor_shchem_v1.desktop_handout_knowledge_tables import write_word_table
from integrations.deeptutor_shchem_v1.desktop_handout_working_tables import editable_table_cells
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import edit_resource, enable_composition
from integrations.deeptutor_shchem_v1.desktop_lesson_design import new_design, content_fingerprint
from integrations.deeptutor_shchem_v1.desktop_preparation_images import verify_image_bytes
from test_handout_working_tables import (
    REAL_ROOT, REAL, row, synthetic_item, composition_for, changed, edited, text_block,
    read_real, assert_manifests_unchanged)


@pytest.fixture(scope='session')
def app():
    from PySide6.QtWidgets import QApplication
    from integrations.deeptutor_shchem_v1.desktop_workbench.typography import install_ui_font
    application = QApplication.instance() or QApplication([])
    install_ui_font(application)
    return application


def payload_for(value, assets, *, classroom_notes=False):
    p = {'output_kind': 'joint', 'topic': '原生表格编辑验证', 'audience': '合成验收',
        'lesson_route': '复习', 'lesson_timing': '1课时×40分钟', 'objective': '观察表格原文与原图',
        'materials': '合成验证材料', 'advanced': dict.fromkeys(
            ('learning_and_experiment', 'template_and_delivery', 'homework_and_strategy'), '')}
    p['lesson_design'] = enable_composition(new_design(p))
    p['lesson_design']['composition'] = value
    p['image_assets'] = deepcopy(assets)
    if classroom_notes:
        p['lesson_design']['classroom_notes'] = True
    return p


def word_cells(document, table):
    cells, occurrences = [], []
    for ri, tr in enumerate(table._tbl.tr_lst, 1):
        starts = tr.xpath('./w:trPr/w:gridBefore')
        column = int(starts[0].get(qn('w:val'))) + 1 if starts else 1
        for tc in tr.tc_lst:
            merge = tc.xpath('./w:tcPr/w:vMerge')
            cells.append((ri, column, tc.grid_span, merge[0].get(qn('w:val'), 'continue') if merge else 'none'))
            for image in tc.xpath('.//a:blip'):
                raw = document.part.related_parts[image.get(qn('r:embed'))].blob
                occurrences.append((ri, column, sha256(raw).hexdigest()))
            column += tc.grid_span
    return cells, occurrences


def source_cells(rows):
    cells = []
    for ri, r in enumerate(rows, 1):
        column = r['grid_before'] + 1
        for c in r['cells']:
            cells.append((ri, column, c['column_span'], c['vertical_merge']))
            column += c['column_span']
    return cells


def pdf_text(path):
    import pypdfium2 as pdfium
    result = []
    document = pdfium.PdfDocument(str(path))
    try:
        for index in range(len(document)):
            page = document[index]; textpage = page.get_textpage()
            try: result.append(textpage.get_text_range())
            finally: textpage.close(); page.close()
    finally: document.close()
    return re.sub(r'\s', '', ''.join(result))


@pytest.mark.parametrize('classroom_notes', [False, True])
def test_edited_common_projection_canvas_pdf_ppt_and_native_word(app, tmp_path, classroom_notes):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import (
        arrange_composition, layout_composition, page_image, write_pdf, write_pptx)
    from integrations.deeptutor_shchem_v1.desktop_composition_rich_render import flatten_elements
    from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents
    item, ids, assets, data = synthetic_item(); before = deepcopy(item)
    cells = changed(item, ids[0], cell_index=4, text='图前CH_{3} ')
    cells[4]['segments'][1] = 'Fe^{2+} '
    cells[4]['segments'][2] = '\n图后检索原文'
    out = edited(item, ids[0], cells)
    projection = source_table_projection(out, text_block(out, ids[0]))
    value = arrange_composition(composition_for(out), assets, data)
    for student in (False, True):
        layout = layout_composition(value, assets, data, student=student)
        assert not layout['errors']
        leaves = [p for page in layout['pages'].values() for p in flatten_elements(page)]
        images = [p for p in leaves if p['kind'] == 'image']
        assert [p['image_block_id'] for p in images] == projection['image_block_ids']
        assert [p['_table_cell'][:2] for p in images] == [(o['row'], o['column']) for o in projection['occurrences']]
        words = ''.join(p.get('text', '') for p in leaves)
        assert '图前CH₃' in words and 'Fe²⁺' in words and '图后检索原文' in words
        assert all(p['font_px'] == 32 for p in leaves if p['kind'] == 'text' and '_table_cell' in p)
        assert '表格开始' not in words and '待查看原文' not in words
        for page in layout['pages']:
            assert not page_image(layout, page).isNull()
        pdf = tmp_path / f'canvas-{student}.pdf'; write_pdf(layout, pdf)
        assert '图后检索原文' in pdf_text(pdf)
        ppt = tmp_path / f'presentation-{student}.pptx'; write_pptx(layout, ppt, data)
        deck = Presentation(ppt)
        pictures = [s for slide in deck.slides for s in slide.shapes if s.shape_type == 13]
        assert [sha256(s.image.blob).hexdigest() for s in pictures] == [o['preview_sha256'] for o in projection['occurrences']]
        ppt_words = ''.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
        assert '图前CH₃' in ppt_words and 'Fe²⁺' in ppt_words and '图后检索原文' in ppt_words
        for shape, image in zip(pictures, images):
            assert (shape.left.mm, shape.top.mm, shape.width.mm, shape.height.mm) == pytest.approx(
                (image['x'], image['y'], image['width'], image['height']), abs=.001)
    write_documents(payload_for(value, assets, classroom_notes=classroom_notes), tmp_path, data)
    for filename in ('lesson_plan.docx', 'student_worksheet.docx'):
        doc = Document(tmp_path / filename)
        assert len(doc.tables) == 1 and len(doc.inline_shapes) == len(projection['occurrences'])
        table = doc.tables[0]
        grid, occurrences = word_cells(doc, table)
        assert grid == source_cells(projection['rows'])
        assert occurrences == [(o['row'], o['column'], o['preview_sha256']) for o in projection['occurrences']]
        assert '图后检索原文' in ''.join(table._tbl.xpath('.//w:t/text()'))
        assert table._tbl.xpath('.//w:vertAlign[@w:val="subscript"]')
        assert table._tbl.xpath('.//w:vertAlign[@w:val="superscript"]')
        assert not table._tbl.xpath('.//w:t[contains(text(), "待查看原文")]')
        assert any(p.text == '完整表后原文。' for p in doc.paragraphs)
    assert item == before


def test_teacher_scope_stays_out_of_student_output_after_cell_edit(app, tmp_path):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition, layout_composition
    from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents
    item, ids, assets, data = synthetic_item(teacher=True, legacy=True)
    out = edited(item, ids[0]); value = arrange_composition(composition_for(out), assets, data)
    student = layout_composition(value, assets, data, student=True)
    assert not any(student['pages'].values())
    write_documents(payload_for(value, assets), tmp_path, data)
    teacher = Document(tmp_path / 'lesson_plan.docx'); student_doc = Document(tmp_path / 'student_worksheet.docx')
    assert len(teacher.tables) == 1 and len(teacher.inline_shapes) == 3
    assert not student_doc.tables and not student_doc.inline_shapes


@pytest.mark.parametrize('code,index', [('062', 45), ('072', 108), ('072', 112), ('072', 270), ('072', 305), ('072', 307)])
def test_actual_cell_edit_reopens_native_word_with_original_grid_and_picture_bytes(tmp_path, code, index):
    payload, manifests = read_real(code)
    item, b = next((r, b) for r in payload['lesson_design']['composition']['resources']
        if r['original']['kind'] == 'handout_knowledge' and r['original']['evidence']['source_sha256'] == REAL[code][2]
        for b in r['blocks'] if b['kind'] == 'text' and b['source_index'] == index)
    original = source_table_projection(item, b)
    cells = editable_table_cells(item, b['id'])
    picture_cell = next(c for c in cells if len(c['segments']) > 1)
    picture_cell['segments'][0] += '观察 '
    picture_cell['segments'][-1] += '\n图后文字'
    out = edited(item, b['id'], cells)
    projection = source_table_projection(out, text_block(out, b['id']))
    store = REAL_ROOT / ('2026-10-08-native-image-table-workbench/extra-protein-visuals/candidate-course-v8/tasks/preparation-v1/images'
        if code == '062' else '2026-10-08-polymer-paint-preview-workbench/native-complete-source-insertion-v5/isolated-task/profile/tasks/preparation-v1/images')
    needed = {o['asset_id'] for o in projection['occurrences']}
    assets = {a['asset_id']: a for a in payload['image_assets'] if a['asset_id'] in needed}
    assert set(assets) == needed
    data = {}
    for identity, asset in assets.items():
        path = store / (asset['sha256'] + '.image')
        assert path.is_file()
        manifests[path] = (sha256(path.read_bytes()).hexdigest(), path.stat().st_size, path.stat().st_mtime_ns)
        data[identity] = verify_image_bytes(asset, path.read_bytes())
    doc = Document()
    write_word_table(doc, projection['rows'], cell_tokens=projection['cell_tokens'], assets=assets, image_data=data)
    filename = tmp_path / f'actual-{code}-B{index}-edited.docx'; doc.save(filename)
    reopened = Document(filename)
    grid, occurrences = word_cells(reopened, reopened.tables[0])
    assert grid == source_cells(original['rows'])
    assert occurrences == [(o['row'], o['column'], o['preview_sha256']) for o in original['occurrences']]
    assert len(reopened.inline_shapes) == len(original['occurrences'])
    assert '图后文字' in ''.join(reopened.tables[0]._tbl.xpath('.//w:t/text()'))
    assert not reopened.tables[0]._tbl.xpath('.//w:t[contains(text(), "待查看原文")]')
    (tmp_path / 'read-only-source-manifest.json').write_text(json.dumps([
        {'path': str(path), 'sha256': record[0], 'bytes': record[1], 'mtime_ns': record[2]}
        for path, record in manifests.items()], ensure_ascii=False, indent=2), encoding='utf-8')
    assert_manifests_unchanged(manifests)


def test_long_work_text_continues_without_truncation_and_unsafe_merge_fails(app):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import _lines
    from integrations.deeptutor_shchem_v1.desktop_handout_knowledge_tables import table_chunks
    item, ids, _, _ = synthetic_item([row('标题甲', '标题乙'), row('起始', '右侧')])
    lines = '\n'.join(f'完整行{i:03d}' for i in range(80))
    out = edited(item, ids[0], changed(item, ids[0], cell_index=2, text=lines))
    projection = source_table_projection(out, text_block(out, ids[0]))
    panels = table_chunks(projection['rows'], ids[0], 296, _lines, 32, max_height=136,
        cell_tokens=projection['cell_tokens'])
    assert len(panels) > 1
    words = ''.join(p.get('text', '') for panel in panels for p in panel['_parts'])
    assert all(f'完整行{i:03d}' in words for i in range(80))
    assert all(p['font_px'] == 32 for panel in panels for p in panel['_parts'] if p['kind'] == 'text')
    merged, mids, _, _ = synthetic_item([row((1, '首行合并', 'restart'), '右侧'),
        row((1, '', 'continue'), '第二行')])
    out = edited(merged, mids[0], changed(merged, mids[0], text=lines))
    projection = source_table_projection(out, text_block(out, mids[0]))
    with pytest.raises(ValueError, match='首行纵向合并'):
        table_chunks(projection['rows'], mids[0], 296, _lines, 32, max_height=136)


def test_draft_json_and_regular_zip_keep_sidecar_original_images_and_fingerprint(tmp_path, monkeypatch):
    """Verify independent unpacking; product restore instantiates Recovery and is excluded here."""
    from integrations.deeptutor_shchem_v1 import desktop_backup as backup
    from integrations.deeptutor_shchem_v1.desktop_preparation_drafts import PreparationDraftService
    from integrations.deeptutor_shchem_v1.desktop_state import DesktopStateStore
    def no_recovery(*_args, **_kwargs):
        raise AssertionError('This test must never instantiate Recovery')
    monkeypatch.setattr(backup, 'PreparationRecoveryStore', no_recovery)
    item, ids, assets, data = synthetic_item(second=True)
    value = composition_for(item)
    original_payload = payload_for(value, assets)
    value = edit_resource(value, item['id'], table_cells={ids[0]: changed(item, ids[0])})
    payload = deepcopy(original_payload); payload['lesson_design']['composition'] = value
    fingerprint = content_fingerprint(payload)
    assert fingerprint != content_fingerprint(original_payload)
    filename = tmp_path / 'edited-payload.json'
    filename.write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
    assert json.loads(filename.read_text('utf-8')) == payload
    profile = tmp_path / 'profile'; state = DesktopStateStore(profile)
    record = {'kind': 'preparation', 'contract_version': 'lesson-blueprint/2.0.0', 'status': 'draft',
        'output_kind': payload['output_kind'], 'created_at': '2026-10-09T00:00:00Z',
        'core_fields': {k: payload[k] for k in ('topic', 'audience', 'lesson_route', 'lesson_timing', 'objective', 'materials')},
        'advanced': deepcopy(payload['advanced']), 'lesson_design': deepcopy(payload['lesson_design']),
        'image_assets': deepcopy(assets)}
    state.save_draft('table-cells', record)
    assert PreparationDraftService._payload(state.snapshot()['drafts']['table-cells']) == payload
    image_root = profile / 'tasks/preparation-v1/images'; image_root.mkdir(parents=True)
    for asset in assets:
        (image_root / (asset['sha256'] + '.image')).write_bytes(data[asset['asset_id']])
    before = {str(p.relative_to(profile)): p.read_bytes() for p in profile.rglob('*') if p.is_file()}
    plan = backup.plan_backup(profile, include_images=True)
    archive = tmp_path / 'native-table.zip'; backup.create_backup(plan, archive)
    report = backup.inspect_backup(archive)
    assert report['summary']['images_included'] == len(assets)
    with ZipFile(archive) as z:
        restored_state = json.loads(z.read('files/desktop-state.v1.json'))
        reopened = PreparationDraftService._payload(restored_state['drafts']['table-cells'])
        assert reopened == payload and content_fingerprint(reopened) == fingerprint
        for asset in assets:
            raw = z.read('files/tasks/preparation-v1/images/' + asset['sha256'] + '.image')
            assert verify_image_bytes(asset, raw) == data[asset['asset_id']]
    assert {str(p.relative_to(profile)): p.read_bytes() for p in profile.rglob('*') if p.is_file()} == before
    assert not (profile / 'recovery').exists()
