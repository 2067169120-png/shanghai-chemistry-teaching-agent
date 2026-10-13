"""Knowledge cards retain scientific conditions and immutable source evidence."""
from copy import deepcopy
import pytest
from docx import Document
from test_composition_render import app, text
from test_composition_flow import personal, composition
from test_lesson_design import sample
from integrations.deeptutor_shchem_v1.desktop_composition_render import (
    _content, arrange_composition, layout_composition, PADDING)
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    resource, block, resource_frames, enable_composition)
from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents

SUMMARY = ('甲烷中的碳形成四个碳氢键，任意两键夹角109°28′；同类以四个原子通过共价单键连接的碳称饱和碳。'
           '乙烯、乙炔中的相应碳连接原子数少于4，称不饱和碳。双键碳相邻两键角约120°，'
           '三键碳相邻两键角180°；图1.6同时标乙烯具体角度121.7°与116.6°。')
LIMIT = '适用范围与待核对处：不把约120°写成所有乙烯键角精确相等；不由一个局部碳的构型判断整分子所有原子共平面。'
ADVICE = '课堂核对建议：把原子连接数、键型和局部空间形状分开核对，区分近似角度与图中具体示例值。'
# The classroom heading omits the review wording; both complete scientific
# conditions after the colon must survive exactly, for either audience.
DISPLAY_LIMIT = '适用范围：'+LIMIT.split('：',1)[1]

def note(summary=SUMMARY):
    return resource('knowledge_note','饱和与不饱和碳的连接数和局部空间角度',[
        block('text','knowledge',summary), block('text','knowledge',LIMIT),
        block('text','teacher_note',ADVICE)],
        {'source_path':'合成教材.pdf','source_sha256':'a'*64,'source_revision':'b'*64,
         'record_id':'READ-test-card','pdf_pages':[12,13],'printed_pages':[8,9]})

@pytest.mark.parametrize('student',[False,True])
def test_short_card_moves_with_all_conditions_to_next_page(app,student):
    item=note();source=deepcopy(item)
    total=sum(c['height'] for c in _content(item,296,{}, {},student=False,format='16:9'))
    for n in range(1,20):
        filler=personal('已有材料',[block('text','personal','保留已有材料。\n'*n)])
        remaining=156-8-2*PADDING-sum(c['height'] for c in _content(filler,296,{}, {},student=False,format='16:9'))
        if 35<remaining<total:break
    else:pytest.fail('没有找到能放标题和部分正文的页尾')
    before=composition(filler,item);arranged=arrange_composition(before,[],{})
    layout=layout_composition(arranged,[],{},student=student)
    frames=resource_frames(arranged['placements'][item['id']])
    assert len(frames)==1 and not layout['errors']
    assert frames[0]['page_id']!=resource_frames(arranged['placements'][filler['id']])[0]['page_id']
    assert arranged['resources']==before['resources']
    for row,expected in zip(item['blocks'][:2],(SUMMARY,DISPLAY_LIMIT)):
        assert text(layout,row['id'])==expected
        assert {e['page_id'] for page in layout['pages'].values() for e in page if e.get('block_id')==row['id']}=={frames[0]['page_id']}
    assert text(layout,item['blocks'][2]['id'])=='' and ADVICE not in text(layout)
    assert item==source

def test_long_card_continues_with_complete_text_at_projection_font(app):
    summary='完整条件与适用范围。'*200
    item=note(summary);source=deepcopy(item);before=composition(item)
    arranged=arrange_composition(before,[],{});layout=layout_composition(arranged,[],{},student=False)
    assert len(resource_frames(arranged['placements'][item['id']]))>1 and not layout['errors']
    assert arranged['resources']==before['resources']
    for row,expected in zip(item['blocks'][:2],(summary,DISPLAY_LIMIT)):
        assert text(layout,row['id'])==expected
    assert text(layout,item['blocks'][2]['id'])=='' and ADVICE not in text(layout)
    assert item==source
    assert all(e['font_px']==32 for page in layout['pages'].values() for e in page if e.get('block_id') and e['kind']=='text')

def test_explicit_small_frames_can_split_card_without_rewriting_layout(app):
    item=note();source=deepcopy(item);value=composition(item);later='CP-'+'e'*32;value['pages'].append(later)
    value['placements'][item['id']]={'page_id':value['pages'][0],'x':12,'y':12,'width':296,'height':48,
        'continuations':[{'page_id':later,'x':12,'y':12,'width':296,'height':130}]}
    before=deepcopy(value);layout=layout_composition(value,[],{},student=False)
    assert not layout['errors'] and value==before
    assert len({f['frame']['page_id'] for f in layout['fragments']})==2
    for row,expected in zip(item['blocks'][:2],(SUMMARY,DISPLAY_LIMIT)):
        assert text(layout,row['id'])==expected
    assert text(layout,item['blocks'][2]['id'])=='' and ADVICE not in text(layout)
    assert item==source

def documents(tmp_path,item):
    payload=sample();payload['lesson_design']=enable_composition(payload['lesson_design'])
    payload['lesson_design']['composition']=composition(item,format='A4')
    payload['image_assets']=[];write_documents(payload,tmp_path,{})
    return Document(tmp_path/'lesson_plan.docx'),Document(tmp_path/'student_worksheet.docx')

def test_word_short_card_keeps_heading_body_and_scientific_conditions_together(app,tmp_path):
    item=note();before=deepcopy(item);teacher,student=documents(tmp_path,item)
    for doc in (teacher,student):
        start=next(i for i,p in enumerate(doc.paragraphs) if p.text==item['title'])
        # Framework sections after the resource are outside its keep chain.
        end=next((i for i,p in enumerate(doc.paragraphs[start+1:],start+1)
            if p.style.name.startswith('Heading ')),len(doc.paragraphs))
        group=doc.paragraphs[start:end]
        assert all(p.paragraph_format.keep_together for p in group)
        assert all(p.paragraph_format.keep_with_next for p in group[:-1])
        assert group[-1].paragraph_format.keep_with_next is False
        assert [p.text for p in group]==[item['title'],SUMMARY,
            '适用范围：'+LIMIT.split('：',1)[1]]
        assert ADVICE not in '\n'.join(p.text for p in doc.paragraphs)
    # Source evidence and review advice remain in the immutable saved card;
    # classroom Word outputs retain the chemistry and its scientific limits.
    assert item==before
    assert item['original']['blocks'][-1]['value']==ADVICE
    assert item['original']['evidence']['source_sha256']=='a'*64
    assert item['original']['evidence']['pdf_pages']==[12,13]
    assert '3. 巩固小结' in '\n'.join(p.text for p in teacher.paragraphs)

def test_word_long_card_keeps_normal_continuation(app,tmp_path):
    item=note('完整条件与适用范围。'*200);teacher,student=documents(tmp_path,item)
    for doc in (teacher,student):
        paragraph=next(p for p in doc.paragraphs if p.text==item['blocks'][0]['value'])
        assert not paragraph.paragraph_format.keep_together
        assert not paragraph.paragraph_format.keep_with_next
