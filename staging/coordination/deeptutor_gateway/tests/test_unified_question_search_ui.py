"""Actual Qt five-source acceptance on isolated authored source files.

The production facade, Qt task bridge, unified adapter, source readers and basket
compilers run against synthetic Word/visual imports and formal reader fixtures.
DeferredTasks controls only deterministic late-delivery and small fault cases.
No teacher profile, real source collection, Office process or API is used.
"""
from __future__ import annotations

import hashlib
import json
import os
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from time import monotonic
from types import SimpleNamespace

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
pytest.importorskip("PySide6")
from PySide6.QtCore import QObject, QThread, Qt, Signal
from PySide6.QtGui import QColor, QFont, QFontDatabase, QPalette
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QLabel
from docx import Document
from PIL import Image

from test_desktop_library import Reader
from test_personal_visual_questions import imported_visual_batch as imported_visual_batch
from test_unified_question_search import FakeFacade, LANES

from integrations.deeptutor_shchem_v1.candidate_review import CandidateCropPayload
from integrations.deeptutor_shchem_v1.desktop_unified_question_search import UnifiedQuestionSearch
from integrations.deeptutor_shchem_v1 import desktop_unified_question_search as unified
from integrations.deeptutor_shchem_v1.desktop_workbench.explorer_reader import PersonalQuestionReader
from integrations.deeptutor_shchem_v1.desktop_workbench.library_detail import LibraryDetailDialog
from integrations.deeptutor_shchem_v1.desktop_workbench.main_window import WORKBENCH_STYLE
from integrations.deeptutor_shchem_v1.desktop_workbench.question_explorer_page import QuestionExplorerPage
from integrations.deeptutor_shchem_v1.desktop_workbench.tasks import DesktopTaskBridge
from integrations.deeptutor_shchem_v1.question_search_workbench import QuestionSearchWorkbench


@pytest.fixture(scope="module")
def app():
    value = QApplication.instance() or QApplication([])
    previous_font = QFont(value.font())
    # The Windows offscreen plugin may omit system font discovery. Register
    # the installed Chinese font for readable synthetic screenshot evidence.
    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    if os.name == "nt" and font_path.is_file():
        identifier = QFontDatabase.addApplicationFont(str(font_path))
        families = QFontDatabase.applicationFontFamilies(identifier)
        if families:
            value.setFont(QFont(families[0], 10))
    yield value
    value.setFont(previous_font)


@pytest.fixture(autouse=True)
def temporary_storage_policy(tmp_path, monkeypatch):
    # The task's Windows storage policy requires an explicit E: --basetemp.
    if os.name == "nt" and tmp_path.drive.upper() != "E:":
        pytest.skip("Run synthetic Qt acceptance with --basetemp on E:")
    if os.name == "nt":
        scratch = Path("E:/临时文件/UserTemp/shchem-full-closure-20261013")
        scratch.mkdir(parents=True, exist_ok=True)
        monkeypatch.setenv("TEMP", str(scratch))
        monkeypatch.setenv("TMP", str(scratch))


class DeferredTasks(QObject):
    """Run real operations on demand and deliberately deliver obsolete results."""
    task_finished = Signal(str)
    task_cancelled = Signal(str)

    def __init__(self):
        super().__init__()
        self.jobs = {}
        self.cancelled = set()
        self.errors = []

    def _submit(self, title, operation, progress, callbacks):
        key = f"synthetic-task-{len(self.jobs) + 1}"
        self.jobs[key] = {"title": title, "operation": operation,
                          "progress": progress, "callbacks": callbacks,
                          "finished": False}
        return key

    def submit(self, title, operation, **callbacks):
        return self._submit(title, operation, False, callbacks)

    def submit_progress(self, title, operation, **callbacks):
        return self._submit(title, operation, True, callbacks)

    def cancel(self, key):
        self.cancelled.add(key)
        self.task_cancelled.emit(key)

    def execute(self, key, *, honour_cancel=True):
        job = self.jobs[key]
        report = job["callbacks"].get("on_progress", lambda _: None)
        cancelled = lambda: honour_cancel and key in self.cancelled
        return (job["operation"](report, cancelled) if job["progress"]
                else job["operation"]())

    def deliver(self, key, value):
        job = self.jobs[key]
        job["finished"] = True
        job["callbacks"].get("on_success", lambda _: None)(value)
        self.task_finished.emit(key)

    def finish(self, key):
        job = self.jobs[key]
        try:
            value = self.execute(key)
        except Exception as exc:
            job["finished"] = True
            self.errors.append((job["title"], str(exc)))
            job["callbacks"].get("on_failure", lambda _: None)(str(exc))
            self.task_finished.emit(key)
            return None
        self.deliver(key, value)
        return value

    def flush(self):
        while True:
            pending = next((key for key, job in self.jobs.items()
                            if not job["finished"] and key not in self.cancelled), None)
            if pending is None:
                return
            self.finish(pending)


def _production_search_ports(facade):
    """Use real facade wrappers when available; bind identical ports otherwise."""
    if callable(getattr(facade, "search_unified_questions", None)):
        return
    adapter = UnifiedQuestionSearch(facade)
    facade.search_unified_questions = adapter.search
    facade.unified_question_search_is_current = adapter.is_current
    facade.invalidate_unified_question_search = adapter.invalidate


class AuthoredFormalReader(Reader):
    def __init__(self):
        super().__init__()
        raw = BytesIO()
        Image.new("RGB", (240, 80), "#58768b").save(raw, format="PNG")
        self.raw = raw.getvalue()
        self.sha = hashlib.sha256(self.raw).hexdigest()
        theme = self.value["papers"][0]["theme_groups"][0]
        theme["theme"]["title"] = "合成电化学主题"
        theme["shared_context"]["context_summary_zh"] = "合成电化学共同材料及条件"
        for atom in theme["atomic_chain"]:
            atom["label_summary"]["primary_K"] = "K10"
        for scan in self.details.values():
            for descriptor in scan["evidence_descriptors"]:
                descriptor["sha256"] = self.sha

    def question_crop(self, node, crop):
        self.crop_calls.append((node, crop))
        return CandidateCropPayload(b"damaged" if self.bad_bytes else self.raw, self.sha)


@pytest.fixture
def real_sources(imported_visual_batch, tmp_path):
    context = imported_visual_batch
    facade = context["facade"]
    formal = AuthoredFormalReader()
    facade._themes = facade._supplemental = formal
    facade._wave_visual = facade._wave_crops = formal
    facade._master_workbench = facade._master_direct = formal
    facade._search = QuestionSearchWorkbench()
    facade._curriculum = SimpleNamespace(catalog=lambda: {"volumes": []})
    taxonomy = facade.paths.workspace_root / "sh-chem-db/kb/knowledge_taxonomy.json"
    catalog = json.loads(taxonomy.read_text(encoding="utf-8"))
    catalog["dimensions"]["knowledge_points"].append({"id": "K10", "name": "电化学（合成标签）"})
    taxonomy.write_text(json.dumps(catalog, ensure_ascii=False), encoding="utf-8")
    for visual_row in facade.personal_visual_questions()["items"]:
        visual_options = facade.personal_visual_question_attribute_options(
            visual_row["batch_id"], visual_row["key"], visual_row["revision"])
        facade.personal_visual_question_save_attributes(
            visual_row["batch_id"], visual_row["key"], visual_row["revision"],
            {"primary_knowledge": {"id": "K10", "label": "电化学（合成标签）",
                                   "status": "teacher_confirmed", "evidence": []}},
            expected_attribute_revision=visual_options["attributes"]["revision"])
    doc = Document()
    doc.add_heading("电化学（合成题面，非真实试卷）", 1)
    doc.add_paragraph("共同材料：合成电化学公共条件，铜和银电极标签保持。")
    doc.add_paragraph("【例1】合成电化学原题：根据共同材料选择电子流向（　）。")
    doc.add_paragraph("A．铜流向银　B．银流向铜")
    doc.add_paragraph("【答案】A")
    doc.add_paragraph("【解析】合成教师原答案完整保留。")
    source = tmp_path / "合成电化学讲义.docx"
    doc.save(source)
    facade.save_visual_import_batch(handout_files=(source,), source_type="合成教师讲义")
    word = facade.word_question_catalog()["items"]
    assert len(word) == 1
    row = word[0]
    options = facade.word_question_attribute_options(row["key"], row["revision"])
    facade.word_question_save_attributes(row["key"], row["revision"], {
        "primary_knowledge": {"id": "K10", "label": "电化学（合成标签）",
                              "status": "teacher_confirmed", "evidence": []},
    }, expected_attribute_revision=options["attributes"]["revision"],
        expected_stored_revision=options["stored_revision"])
    _production_search_ports(facade)
    assert callable(facade.search_unified_questions)
    assert "search_unified_questions" not in facade.__dict__
    return facade, formal, context, row


def _settle(app, predicate, timeout=15):
    deadline = monotonic() + timeout
    while not predicate() and monotonic() < deadline:
        app.processEvents()
        QTest.qWait(5)
    assert predicate()
    app.processEvents()


def _flush(tasks, app, predicate=lambda: True):
    if isinstance(tasks, DeferredTasks):
        tasks.flush()
        app.processEvents()
    else:
        _settle(app, lambda: predicate() and not tasks._workers)


def _page(facade, app, *, sources=LANES, knowledge="K01", query="合成题面", tasks=None):
    tasks = tasks or DeferredTasks()
    page = QuestionExplorerPage(facade, tasks)
    page.setObjectName("UnifiedAcceptancePage")
    page.setStyleSheet(WORKBENCH_STYLE + "\nQWidget#UnifiedAcceptancePage {background: #F4F6F1;}")
    # A standalone page has no production main-window backing surface. Fill
    # that surface explicitly so its synthetic PNG is opaque and readable.
    palette = page.palette()
    palette.setColor(QPalette.ColorRole.Window, QColor("#F4F6F1"))
    page.setPalette(palette)
    page.setAutoFillBackground(True)
    page._started = True
    payload = {"sources": list(sources), "query": query,
               "filters": {"knowledge": [knowledge], "knowledge_mode": "any"},
               "practice_task_id": "synthetic-practice-task"}
    assert page.apply_unified_handoff(payload)
    page.resize(1280, 860)
    page.show()
    app.processEvents()
    page.search_button.click()
    try:
        _flush(tasks, app, lambda: not page._loading)
    except AssertionError:
        detail = (page._loading, page.result_summary.text(), getattr(tasks, "errors", []),
                  list(getattr(tasks, "_workers", {})), getattr(tasks, "cancel_details", []))
        _dispose(page, app)
        pytest.fail(f"Initial real Qt search did not settle: {detail!r}")
    return page, tasks, payload


def _dispose(page, app):
    page.close()
    if isinstance(page.tasks, DesktopTaskBridge):
        assert page.tasks.shutdown(5000)
    page.deleteLater()
    app.processEvents()


def _labels(widget):
    return "\n".join(label.text() for label in widget.findChildren(QLabel))


def _memory_facade(*, count=1):
    facade = FakeFacade()
    state = {}
    basket = [{"key": "already-selected-before-search"}]
    facade.state_store = SimpleNamespace(snapshot=lambda: deepcopy(state),
        _update=lambda operation: operation(state))
    facade.basket = lambda: deepcopy(basket)
    facade.saved_basket = basket
    for index in range(count):
        for lane in LANES[:3]:
            facade.add_formal(lane, f"{lane}-{index}")
        facade.add_word(f"word-{index}")
        facade.add_visual(f"visual-{index}")
    _production_search_ports(facade)
    return facade


def test_real_five_sources_same_query_knowledge_open_original_readers_and_basket(app, real_sources, monkeypatch):
    facade, formal, context, word = real_sources
    thread_reads = {"metadata": [], "sha": []}
    for name, key in (("_source_metadata", "metadata"), ("_file_sha", "sha")):
        original = getattr(unified, name)
        def observed(*args, _original=original, _key=key, **kwargs):
            thread_reads[_key].append(QThread.currentThread() == app.thread())
            return _original(*args, **kwargs)
        monkeypatch.setattr(unified, name, observed)
    tasks = DesktopTaskBridge()
    tasks.errors = []
    tasks.cancel_details = []
    tasks.task_failed.connect(lambda _key, title, message: tasks.errors.append((title, message)))
    tasks.task_cancelled_detail.connect(lambda _key, title, message: tasks.cancel_details.append((title, message)))
    page, tasks, _ = _page(facade, app, knowledge="K10", query="电化学", tasks=tasks)
    try:
        assert not tasks.errors
        assert {card.entry["lane"] for card in page.cards} == set(LANES)
        assert len(page.cards) == 6  # Three themes, one Word segment, two visual hits.
        assert page._reader is None and not any(card.ready for card in page.cards)
        assert not formal.crop_calls  # Unified listing never reads source pictures.
        assert "完整主题 3" in page.result_summary.text()
        assert "Word完整题段 1" in page.result_summary.text()
        assert "图片印刷题命中（入篮为整主题） 2" in page.result_summary.text()
        screenshot = os.environ.get("SHCHEM_QT_ACCEPTANCE_SCREENSHOT")
        if screenshot:
            target = Path(screenshot)
            assert os.name != "nt" or target.drive.upper() == "E:"
            target.parent.mkdir(parents=True, exist_ok=True)
            assert page.grab().save(str(target))
        for lane in LANES:
            card = next(card for card in page.cards if card.entry["lane"] == lane)
            QTest.mouseClick(card.open, Qt.MouseButton.LeftButton)
            _flush(tasks, app, lambda: card.ready or bool(tasks.errors))
            assert card.ready, (lane, tasks.errors, card.note.text())
            reader = page._reader
            assert isinstance(reader, PersonalQuestionReader if lane in LANES[3:] else LibraryDetailDialog)
            reader.tabs.setCurrentIndex(1)
            _flush(tasks, app)
            if lane in LANES[:3]:
                assert "合成电化学共同材料" in _labels(reader)
                assert "来源答案原文\n第二行" in _labels(reader)
            elif lane == "word_native":
                assert "合成电化学公共条件" in _labels(reader)
                assert "合成教师原答案完整保留" in _labels(reader)
                assert reader.entry["payload"]["key"] == word["key"]
            else:
                assert reader.entry["payload"]["shared_text"] in _labels(reader)
                assert reader.entry["payload"]["answer_text"] in _labels(reader)
                assert reader.images and all(image.loaded for image in reader.images)
            before = len(facade.basket())
            card.add.click()
            _flush(tasks, app, lambda: not page._adding)
            assert len(facade.basket()) == before + 1, (lane, tasks.errors)
        basket = facade.basket()
        assert len(basket) == 5 and not tasks.errors
        projection = []
        tasks.submit("验收完整题篮投影", facade.paper_basket_projection,
                     on_success=projection.append)
        _flush(tasks, app, lambda: bool(projection))
        visual = next(item for item in projection[0]["items"] if item.get("kind") == "personal_visual_theme")
        assert len(visual["content"]["theme"]["printed_questions"]) == 2
        assert visual["content"]["theme"]["shared_materials"]
        assert context["provider"].requests  # Synthetic strict provider only, no configured API.
        assert thread_reads["metadata"] and thread_reads["sha"]
        assert not any(thread_reads["metadata"] + thread_reads["sha"])
    finally:
        _dispose(page, app)


def test_actual_qt_pages_keep_cross_source_order_and_source_subset_is_lazy(app):
    facade = _memory_facade(count=4)
    page, tasks, _ = _page(facade, app)
    try:
        assert not tasks.errors and len(page.cards) == page.PAGE_SIZE
        first = [card.entry["entry_id"] for card in page.cards]
        page.next.click()
        tasks.flush()
        second = [card.entry["entry_id"] for card in page.cards]
        assert page._page == 1 and not set(first).intersection(second)
        page.previous.click()
        tasks.flush()
        assert [card.entry["entry_id"] for card in page.cards] == first
        page.next.click()
        tasks.flush()
        assert page._page == 1
        conditions = page._reading_context()[:4]
        page.reload_button.click()
        tasks.flush()
        assert page._page == 0 and page._cursors == [None]
        assert [card.entry["entry_id"] for card in page.cards] == first
        assert page._reading_context()[:4] == conditions
        assert not facade.body_calls
        subset = _memory_facade()
        other, other_tasks, _ = _page(subset, app, sources=("word_native",))
        try:
            assert [card.entry["lane"] for card in other.cards] == ["word_native"]
            assert not subset.formal_calls and subset.visual_calls == 0
            assert set(subset.snapshot_calls) == {"word_native"}
            for box in other.source_checks.values():
                box.setChecked(False)
            other.search_button.click()
            other_tasks.flush()
            assert not other.cards and "尚未选择资料来源" in other.result_summary.text()
            assert subset.visual_calls == 0 and not subset.formal_calls
            assert other.facade.basket() == subset.saved_basket
        finally:
            _dispose(other, app)
    finally:
        _dispose(page, app)


def test_failed_source_keeps_conditions_task_and_existing_basket_with_unknown_total(app):
    facade = _memory_facade()
    facade.failures["word_native"] = RuntimeError("合成来源不可用")
    page, tasks, payload = _page(facade, app)
    try:
        assert {card.entry["lane"] for card in page.cards} == set(LANES) - {"word_native"}
        assert "Word完整题段 未知" in page.result_summary.text()
        assert "部分来源未完成核对" in page.result_summary.text()
        assert page.filters == {"knowledge": {"K01"}} and page.query.text() == "合成题面"
        assert page.practice_context == payload and facade.basket() == facade.saved_basket
        assert "来源" in _labels(page) and not facade.body_calls
    finally:
        _dispose(page, app)


def test_late_search_progress_and_success_after_query_change_cannot_replace_new_cards(app):
    facade = _memory_facade()
    page, tasks, _ = _page(facade, app)
    try:
        page.search_button.click()
        old = page._search_task
        result = tasks.execute(old)
        page.query.setText("无合成匹配")
        assert old in tasks.cancelled
        tasks.jobs[old]["callbacks"]["on_progress"]({"message": "过期进度"})
        tasks.deliver(old, result)
        assert not page.cards and "过期进度" not in page.result_summary.text()
        page.search_button.click()
        tasks.flush()
        assert not page.cards and not page._loading
        assert facade.basket() == facade.saved_basket
    finally:
        _dispose(page, app)


@pytest.mark.parametrize("change", ["source_generation", "cache_epoch"])
def test_obsolete_version_cannot_open_read_or_add_and_preserves_teacher_conditions(app, change):
    facade = _memory_facade()
    page, tasks, payload = _page(facade, app)
    try:
        card = page.cards[0]
        if change == "source_generation":
            facade.invalidate_unified_question_search(lanes=(card.entry["lane"],))
        else:
            facade.cache_epochs["registry"] += 1
        before = len(tasks.jobs)
        card.open.click()
        assert len(tasks.jobs) == before and not page.cards and page._reader is None
        assert page._needs_reload and "来源或缓存版本已改变" in page.result_summary.text()
        assert page.filters == {"knowledge": {"K01"}} and page.query.text() == "合成题面"
        assert page.practice_context == payload and facade.basket() == facade.saved_basket
        page.add(card)
        assert facade.basket() == facade.saved_basket and not facade.body_calls
    finally:
        _dispose(page, app)


def test_source_checkbox_change_cancels_old_worker_and_retains_only_selected_source(app):
    facade = _memory_facade()
    page, tasks, payload = _page(facade, app)
    try:
        page.search_button.click()
        old = page._search_task
        old_result = tasks.execute(old)
        for lane, box in page.source_checks.items():
            box.setChecked(lane == "visual_native")
        assert old in tasks.cancelled and not page.cards
        tasks.deliver(old, old_result)
        assert not page.cards
        page.search_button.click()
        tasks.flush()
        assert [card.entry["lane"] for card in page.cards] == ["visual_native"]
        assert page.filters == {"knowledge": {"K01"}} and page.practice_context == payload
        assert facade.basket() == facade.saved_basket and not facade.body_calls
    finally:
        _dispose(page, app)


def test_real_word_archive_change_after_read_blocks_original_basket_guard(app, real_sources):
    facade, _, _, word = real_sources
    page, tasks, payload = _page(facade, app, sources=("word_native",),
        knowledge="K10", query="电化学")
    try:
        card = page.cards[0]
        card.open.click()
        tasks.flush()
        assert card.ready
        archive = facade.paths.state_root / "visual-import-v2/sources" / (word["source_sha256"] + ".docx")
        archive.write_bytes(b"synthetic source changed after reading")
        card.add.click()
        tasks.flush()
        assert not facade.basket() and not page.cards
        assert page._needs_reload and page.practice_context == payload
        # Independently call the unchanged production compiler as well as the
        # Qt version gate: the old source identity must fail on its own.
        with pytest.raises(Exception, match="变化|读取|来源|核对"):
            facade.add_word_questions_to_basket([{"key": word["key"], "revision": word["revision"]}])
        assert not facade.basket()
    finally:
        _dispose(page, app)


def test_required_formal_picture_failure_never_enables_qt_basket_action(app, real_sources):
    facade, formal, _, _ = real_sources
    formal.bad_bytes = True
    page, tasks, _ = _page(facade, app, sources=("master",),
        knowledge="K10", query="电化学")
    try:
        card = page.cards[0]
        card.open.click()
        tasks.flush()
        assert isinstance(page._reader, LibraryDetailDialog)
        assert not card.ready and not card.add.isEnabled()
        page.add(card)
        assert not facade.basket()
        assert any(widget.load_state == "failed" for widget in page._reader._question_image_widgets)
    finally:
        _dispose(page, app)


@pytest.mark.parametrize("lane", ["master", "word_native", "visual_native"])
def test_late_original_reader_delivery_after_query_change_cannot_attach_or_enable_add(app, real_sources, lane):
    facade, _, _, _ = real_sources
    page, tasks, payload = _page(facade, app, sources=(lane,), knowledge="K10", query="电化学")
    try:
        card = page.cards[0]
        card.open.click()
        old = page._detail_task
        result = tasks.execute(old)
        page.query.setText("无合成匹配")
        assert old in tasks.cancelled and page._reader is None
        tasks.deliver(old, result)
        assert page._reader is None and not page.cards
        page.search_button.click()
        tasks.flush()
        assert not page.cards and not facade.basket()
        assert page.practice_context == payload
    finally:
        _dispose(page, app)
