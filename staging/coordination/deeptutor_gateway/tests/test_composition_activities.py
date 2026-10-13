"""Source revision, audience scope, explicit refresh and whole-file coverage."""
from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace

import pytest

from test_lesson_design import sample
from test_composition_render import app
from integrations.deeptutor_shchem_v1.desktop_composition_activities import *
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (FLAGS, add_resource, edit_resource,
    enable_composition, validate_resource, visible_blocks, digest)
from integrations.deeptutor_shchem_v1.desktop_lesson_design import update_node, update_goal, DesignHistory


def linked():
    p = sample(); plan = enable_composition(p['lesson_design'])
    item = activity_resource(plan, plan['nodes'][0]['id'], [])
    value = plan['composition']
    plan['composition'] = add_resource(value, item,
        {'page_id': value['pages'][0], 'x': 12, 'y': 12, 'width': 186, 'height': 273})
    p['lesson_design'] = plan; p['image_assets'] = []
    return p, item


def test_activity_snapshot_keeps_full_source_and_uses_typed_audience():
    p, item = linked(); evidence = item['original']['evidence']
    assert evidence['source_sha256'] == 'unknown'
    assert evidence['activity_snapshot']['node'] == p['lesson_design']['nodes'][0]
    assert evidence['source_revision'] == digest(evidence['activity_snapshot'])
    assert item['locked'] and all(item[k] is v for k, v in FLAGS.items())
    student = '\n'.join(b['value'] for b in visible_blocks(item))
    teacher = '\n'.join(b['value'] for b in visible_blocks(item, student=False))
    assert '根据正逆速率' in student and '合成反应A⇌B' in student
    for key in ('teacher_action', 'teacher_answer', 'expected_output', 'criteria'):
        assert p['lesson_design']['nodes'][0][key] in teacher
        assert p['lesson_design']['nodes'][0][key] not in student


@pytest.mark.parametrize('changes', [{'student_task': '新学生任务'}, {'teacher_answer': '新教师答案'},
    {'material_text': '新材料'}, {'minutes': 15}, {'student_material': False}, {'notes': '新的教师备注'}])
def test_source_edits_are_stale_without_mutating_canvas_copy(changes):
    p, item = linked(); plan = p['lesson_design']; before = deepcopy(plan['composition'])
    updated = update_node(plan, plan['nodes'][0]['id'], changes)
    assert activity_status(updated, item, []) == 'changed'
    assert updated['composition'] == before
    with pytest.raises(ValueError, match='关联已变化'): checked_activity_links(updated, [])


def test_goal_wording_is_part_of_source_revision_and_unrelated_goals_are_not():
    p, item = linked(); plan = p['lesson_design']
    other = update_goal(plan, plan['objectives'][1]['id'], '不相关目标的改动')
    assert activity_status(other, item, []) == 'current'
    changed = update_goal(plan, plan['objectives'][0]['id'], '相关目标的新表述')
    assert activity_status(changed, item, []) == 'changed'


def test_refresh_keeps_every_frame_order_and_refuses_edited_work_copy_without_review():
    p, item = linked(); plan = p['lesson_design']; value = plan['composition']; identity = item['id']
    value = edit_resource(value, identity, locked=False)
    row = next(b for b in item['blocks'] if b['role'] == 'student_task')
    value = edit_resource(value, identity, title='教师编辑的标题', values={row['id']: '画布上的修改任务'})
    plan['composition'] = value; history = DesignHistory(plan)
    changed = update_node(plan, plan['nodes'][0]['id'], {'student_task': '来自新环节的任务'})
    with pytest.raises(ValueError, match='已有编辑'): refresh_activity(value, identity, changed, [])
    with pytest.raises(ValueError, match='副本已修改'):
        refresh_activity(value, identity, changed, [], expected='a'*64, replace_edited=True)
    updated = refresh_activity(value, identity, changed, [], expected=digest(value), replace_edited=True)
    assert updated['placements'] == value['placements'] and updated['pages'] == value['pages']
    fresh = updated['resources'][0]
    assert fresh['id'] == identity and fresh['locked'] is False
    assert '来自新环节的任务' in '\n'.join(b['value'] for b in fresh['blocks'])
    assert activity_status(changed, fresh, []) == 'current'
    changed['composition'] = updated; history.put(changed); history.travel()
    assert history.value['composition'] == value


def test_explicit_independent_copy_preserves_sealed_original_after_source_deleted():
    p, item = linked(); plan = p['lesson_design']; plan['nodes'] = []
    assert activity_status(plan, item, []) == 'missing'
    with pytest.raises(ValueError, match='关联已变化'): checked_activity_links(plan, [])
    kept = keep_activity_snapshot(plan['composition'], item['id'])
    assert kept['resources'][0]['original'] == item['original']
    assert kept['resources'][0]['blocks'] == item['blocks'] and kept['placements'] == plan['composition']['placements']
    plan['composition'] = kept
    assert checked_activity_links(plan, [])[0]['status'] == 'snapshot'


def test_same_node_cannot_be_silently_added_twice_and_order_follows_design():
    p, item = linked(); plan = p['lesson_design']; identity = plan['nodes'][0]['id']
    assert not activity_options(plan, [])
    with pytest.raises(ValueError, match='已经加入'): import_activities(plan, [identity], [])
    duplicate = activity_resource(plan, identity, [])
    with pytest.raises(ValueError, match='已经关联'): add_resource(plan['composition'], duplicate, plan['composition']['placements'][item['id']])


def test_resealed_snapshot_cannot_demote_teacher_fields_to_student_scope():
    p, item = linked(); invalid = deepcopy(item)
    row = next(b for b in invalid['original']['blocks'] if b['role'] == 'answer')
    row['role'] = 'student_task'
    invalid['blocks'] = deepcopy(invalid['original']['blocks'])
    invalid['original_sha256'] = digest(invalid['original'])
    with pytest.raises(ValueError, match='教师／学生范围'): validate_resource(invalid)


def test_real_six_files_have_activity_once_and_hide_teacher_fields(tmp_path, app):
    pytest.importorskip('PySide6')
    from PySide6.QtWidgets import QApplication
    from PIL import Image
    from docx import Document
    from pptx import Presentation
    import pypdfium2 as pdfium
    from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition
    from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
    from integrations.deeptutor_shchem_v1.desktop_lesson_output import export_design, checked_file, COMPOSITION_FILES
    p, item = linked(); plan = p['lesson_design']
    store = PreparationImageStore(tmp_path/'tasks/preparation-v1/images')
    image = BytesIO(); Image.new('RGB', (150, 80), '#82a49b').save(image, format='PNG')
    asset = store.import_bytes(image.getvalue(), '合成关联图', '测试作者输入', '学生范围核对')
    plan = update_node(plan, plan['nodes'][0]['id'], {'image_ids': [asset['asset_id']]})
    p['image_assets'] = [asset]; data = {asset['asset_id']: image.getvalue()}
    plan['composition'] = refresh_activity(plan['composition'], item['id'], plan, p['image_assets'])
    plan['composition'] = arrange_composition(plan['composition'], p['image_assets'], data)
    p['lesson_design'] = plan
    facade = SimpleNamespace(paths=SimpleNamespace(task_root=tmp_path/'tasks'), preparation_image_bytes=store.load)
    record = export_design(facade, p); plan['exports'].append(record)
    assert set(f['name'] for f in record['files']) == set(COMPOSITION_FILES)
    source = plan['nodes'][0]
    for name, student in [('lesson_plan.docx', False), ('student_worksheet.docx', True),
                          ('lesson_presentation.pptx', False), ('student_presentation.pptx', True),
                          ('lesson_canvas.pdf', False), ('student_canvas.pdf', True)]:
        path = checked_file(facade, plan, record['id'], name)
        if name.endswith('.docx'):
            doc = Document(path); text = ''.join(p.text for p in doc.paragraphs); images = len(doc.inline_shapes)
            assert images == 1
        elif name.endswith('.pptx'):
            deck = Presentation(path); text = ''.join(s.text for sl in deck.slides for s in sl.shapes if s.has_text_frame)
            assert sum(s.shape_type == 13 for sl in deck.slides for s in sl.shapes) == 1
        else:
            pdf = pdfium.PdfDocument(str(path)); parts = []
            for i in range(len(pdf)):
                page = pdf[i]; tp = page.get_textpage(); parts.append(tp.get_text_range()); tp.close(); page.close()
            pdf.close(); text = ''.join(parts)
        canonical = ''.join(text.split())
        assert canonical.count(''.join(source['student_task'].split())) == 1
        assert (source['teacher_answer'] not in text) == student
        for field in ('teacher_action', 'expected_output', 'criteria'):
            assert (''.join(source[field].split()) not in canonical) == student
    changed = update_node(plan, source['id'], {'student_material': False})
    p['lesson_design'] = changed
    with pytest.raises(ValueError, match='关联已变化'): export_design(facade, p)
    assert checked_file(facade, changed, record['id'], 'student_canvas.pdf').is_file()
    fresh = refresh_activity(changed['composition'], item['id'], changed, p['image_assets'])
    assert not any(b['kind'] == 'image' for b in visible_blocks(fresh['resources'][0]))
