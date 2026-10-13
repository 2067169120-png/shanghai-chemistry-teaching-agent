"""Automatic own-stem references do not replace persisted material decisions."""

from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from PIL import Image

from test_desktop_word_auto_attributes import Facade

from integrations.deeptutor_shchem_v1 import desktop_word_question_attributes as attributes
from integrations.deeptutor_shchem_v1.desktop_handout_knowledge import _check_roles
from integrations.deeptutor_shchem_v1.desktop_word_question_reading import question_reading_layout
from integrations.deeptutor_shchem_v1.desktop_word_questions import WordQuestionService
from integrations.deeptutor_shchem_v1.offline_word_label_review import OfflineWordLabelReviewError


STEM = "例1（双选）某制品使用合成材料甲、合成材料乙、合成材料丙等材料。下列说法正确的是"
OPTIONS = [
    "A．合成材料甲适用于所述用途",
    "B．合成材料乙具有所述性质",
    "C．合成材料丙具有所述结构",
    "D．上述材料具有相同的性质",
]


def question(stem=STEM, *, image=True, options=OPTIONS):
    texts = [stem, *(["【待查看原文：图片或图形】"] if image else []), *options]
    return {
        "key": "word-question-context-fixture",
        "source_sha256": "a" * 64,
        "revision": "question-version",
        "source_revision": "source-version",
        "index_revision": "index-version",
        "extraction_revision": "extraction-version",
        "chapter": "",
        "question_blocks": [
            {"index": index, "text": text, "assets": [{"asset_id": "synthetic"}] if image and index == 21 else []}
            for index, text in enumerate(texts, 20)
        ],
        "answer_blocks": [{"index": 26, "text": "【答案】合成参考答案。", "assets": []}],
        "context_blocks": [],
        "warnings": [],
    }


def material(item):
    return attributes.suggest_attributes(item, {})["material_status"]


@pytest.mark.parametrize("layout", ["separate_blocks", "multiline_block", "inline_options"])
@pytest.mark.parametrize("image", [False, True])
def test_named_own_stem_resolves_later_option_without_relying_on_image_count(layout, image):
    item = question(image=image)
    if layout != "separate_blocks":
        joiner = "\n" if layout == "multiline_block" else "  "
        blocks = item["question_blocks"]
        item["question_blocks"] = [{
            "index": 20,
            "text": joiner.join(block["text"] for block in blocks),
            "assets": [asset for block in blocks for asset in block["assets"]],
        }]
    before = deepcopy(item)
    result = material(item)
    assert result["missing_context"] is False
    assert result["has_shared_context"] is False
    assert result["context_block_indices"] == []
    assert result["question_image_count"] == int(image)
    assert item == before


@pytest.mark.parametrize("reference", ["上述材料", "根据上述材料"])
def test_reference_in_later_stem_sentence_has_same_own_antecedent(reference):
    item = question("某制品使用材料甲甲、材料乙乙等材料。" + reference + "，回答问题。", image=False, options=[])
    assert material(item)["missing_context"] is False


@pytest.mark.parametrize("dependency", ["前题", "上一题", "前一题", "上题", "上一小题", "前一小题", "上一问", "前一问"])
def test_previous_question_dependency_is_not_resolved_by_own_materials(dependency):
    item = question(options=[*OPTIONS[:3], f"D．上述材料是否满足{dependency}的结论"])
    assert material(item)["missing_context"] is True


@pytest.mark.parametrize("stem", [
    "上述材料用于制品。下列说法正确的是",
    "根据上述材料，判断合成材料甲、合成材料乙等材料。",
    "根据上述实验结果，判断合成材料甲、合成材料乙等材料。",
    "",
    "下列关于材料的说法正确的是",
    "材料的应用十分广泛。请判断下列说法。",
    "【待查看原文：图片或图形】",
    "【未提取到文字】",
    "某制品使用____、合成材料乙等材料。下列说法正确的是",
    "某制品使用合成材料甲、……等材料。下列说法正确的是",
    "某制品使用合成材料甲、合成材料乙等材料，所需条件略去。下列说法正确的是",
    "上述对象包括合成材料甲、合成材料乙等材料。下列说法正确的是",
    "某制品使用合成材料甲、合成材料乙等材料",
])
def test_sparse_leading_dependent_or_broken_stem_stays_unresolved(stem):
    assert material(question(stem))["missing_context"] is True


@pytest.mark.parametrize("reference", ["根据上述实验", "根据上述数据", "根据上述"])
def test_material_enumeration_does_not_resolve_other_unspecified_referents(reference):
    assert material(question(options=[*OPTIONS[:3], "D．" + reference + "判断结果"]))["missing_context"] is True


@pytest.mark.parametrize("leading", ["上述实验结果已知。", "上述数据如表所示。", "根据上述条件。"])
def test_unresolved_leading_reference_is_not_hidden_by_later_own_materials(leading):
    assert material(question(leading + STEM))["missing_context"] is True


def test_earlier_options_answer_and_chapter_cannot_supply_missing_stem():
    item = question("下列说法正确的是", options=["A．材料甲甲、材料乙乙等材料。", *OPTIONS[1:]])
    item["answer_blocks"][0]["text"] = "【答案】" + STEM
    item["chapter"] = STEM
    assert material(item)["missing_context"] is True


def test_shared_material_warning_is_not_cleared_by_own_enumeration():
    item = question()
    item["warnings"] = ["共同材料范围待核对。"]
    assert material(item)["missing_context"] is True


def test_explicit_shared_context_keeps_existing_bound_context_behavior():
    item = question("根据上述材料，回答问题。")
    item["warnings"] = ["共同材料范围已绑定。"]
    item["context_blocks"] = [{"index": 19, "text": "【共同材料】材料甲甲、材料乙乙的已知条件。", "assets": []}]
    result = material(item)
    assert result["missing_context"] is False
    assert result["has_shared_context"] is True
    assert result["context_block_indices"] == [19]


def test_identity_filename_and_source_metadata_do_not_resolve_deictics():
    complete, broken = question(), question("下列说法正确的是")
    for item in (complete, broken):
        original = material(item)
        item.update(key="word-question-unrelated-identity", source_sha256="b" * 64)
        item["source_name"] = "完整材料-另一个地区-2020.docx"
        proposed = attributes.suggest_attributes(item, {"source_name": STEM, "lecture_topic": STEM})
        assert proposed["material_status"] == original
    assert material(complete)["missing_context"] is False
    assert material(broken)["missing_context"] is True


@pytest.fixture
def native(tmp_path, monkeypatch):
    """Real paragraph/image extraction with disposable archive and state only."""
    monkeypatch.setattr("integrations.deeptutor_shchem_v1.desktop_word_questions.source_quality_notes", lambda _root: {})
    document = Document()
    document.add_paragraph(STEM)
    png = BytesIO()
    Image.new("RGB", (48, 24), "blue").save(png, format="PNG")
    document.add_paragraph().add_run().add_picture(BytesIO(png.getvalue()))
    for option in OPTIONS:
        document.add_paragraph(option)
    document.add_paragraph("【答案】合成参考答案。")
    document.add_paragraph("【解析】本段是教师答案，只能在答案入口显示。")
    output = BytesIO()
    document.save(output)
    facade = Facade(tmp_path)
    source = facade.add("CONTEXT-BATCH", output.getvalue(), "合成材料课堂练习.docx")
    words = WordQuestionService(facade)
    catalog, inventory = words._catalog(read_only=True)
    assert len(catalog["items"]) == 1
    row = catalog["items"][0]
    return words, row, inventory[row["source_id"]][2], source


def test_real_extraction_reads_own_stem_image_and_options_with_answer_isolation(native):
    words, row, preview, source = native
    before = source.read_bytes()
    state_before = words.state.snapshot()
    proposed = attributes.suggest_attributes(row, {"source_name": row["source_name"]})
    assert proposed["material_status"]["missing_context"] is False
    result = question_reading_layout(words, row["key"], row["revision"])
    assert len(result["images"]) == 1
    assert [p["source_index"] for p in result["paragraphs"]] == [b["index"] for b in row["question_blocks"]]
    assert all(p["role"] == "question" for p in result["paragraphs"])
    assert "教师答案" not in str(result["paragraphs"])
    answer = question_reading_layout(words, row["key"], row["revision"], scope="answer")
    assert "教师答案" in str(answer["paragraphs"])
    assert not answer["images"]
    with pytest.raises(ValueError, match="已识别的题目或答案"):
        _check_roles(SimpleNamespace(), None, preview, {"start": row["block_start"], "end": row["block_end"]})
    assert source.read_bytes() == before
    assert words.state.snapshot() == state_before
    assert not words.attribute_store.path.exists()
    assert not words.preview_cache.root.exists()


@pytest.mark.parametrize("teacher_modified", [False, True])
def test_persisted_material_decision_survives_new_suggestion_and_reading(native, teacher_modified):
    words, row, _preview, source = native
    new = attributes.suggest_attributes(row, {"source_name": row["source_name"]})
    assert new["material_status"]["missing_context"] is False
    legacy = deepcopy(new)
    legacy["material_status"]["missing_context"] = True
    legacy["rule_revision"] = "legacy-automatic-material-decision"
    old = words.attribute_store.save_many([attributes._seal(legacy)])[0]
    if teacher_modified:
        old = words.attribute_store.save_teacher_edit(row["key"], {"teacher_note": "保留原材料复核判断。"}, expected_revision=old["revision"])
    unrelated = question()
    unrelated["key"] = "word-question-another-saved-decision"
    words.attribute_store.save_many([attributes.suggest_attributes(unrelated, {})])
    history = words.attribute_store.history(row["key"])
    database = words.attribute_store.path.read_bytes()
    original = source.read_bytes()
    for merge in (attributes.complete_missing_attributes, attributes.recheck_automatic_attributes):
        assert merge(old, new)["material_status"] == old["material_status"]
    with pytest.raises(OfflineWordLabelReviewError, match="缺失材料"):
        question_reading_layout(words, row["key"], row["revision"])
    assert words.attribute_store.get(row["key"]) == old
    assert words.attribute_store.history(row["key"]) == history
    assert words.attribute_store.get(unrelated["key"])["material_status"]["missing_context"] is False
    assert words.attribute_store.path.read_bytes() == database
    assert source.read_bytes() == original
