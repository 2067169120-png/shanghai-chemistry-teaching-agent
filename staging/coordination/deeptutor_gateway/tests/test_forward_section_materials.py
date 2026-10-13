"""Forward section selection uses real source readers and owned fixture paths."""
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
from io import BytesIO
from types import MethodType, SimpleNamespace

import pytest

from test_textbook_section_reader import section_fixture
from test_textbook_asset_catalog import asset_fixture
from test_textbook_bookshelf import shelf_fixture, colored_pdf
from test_textbook_supplements import supplement_fixture
from test_textbook_reading_hints import note as reading_note, write_packet
from test_textbook_crop_reading import files

from integrations.deeptutor_shchem_v1 import desktop_forward_section_materials as batch
from integrations.deeptutor_shchem_v1.desktop_composition_sources import (
    import_knowledge_note, knowledge_selection, knowledge_textbook_options)
from integrations.deeptutor_shchem_v1.desktop_facade import DesktopWorkbenchFacade
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import FLAGS
from integrations.deeptutor_shchem_v1.desktop_note_page_crops import note_page_selection
from integrations.deeptutor_shchem_v1.desktop_section_material_handoff import prepare_section_material_requests
from integrations.deeptutor_shchem_v1.desktop_section_materials import (
    recheck_textbook_section_material_context, textbook_section_material_context)
from integrations.deeptutor_shchem_v1.reader_cancellation import (
    ReadCancelled, TaskCancellationEvent, read_cancel_scope)


@pytest.fixture
def forward(shelf_fixture, tmp_path):
    f = shelf_fixture
    f.note = reading_note(volume_id=f.row['volume_id'], section_key=f.row['section_key'],
        pdf_pages=[8, 9], printed_pages=[4, 5],
        textbook_evidence={'source_sha256': f.row['source_sha256'],
                          'pdf_pages': [8, 9], 'printed_pages': [4, 5]})
    write_packet(f.root, [f.note])
    f.task_root = tmp_path / 'owned-forward-state'
    f.facade = SimpleNamespace(paths=SimpleNamespace(workspace_root=f.root, task_root=f.task_root))
    for name in ('textbook_bookshelf', 'textbook_reading_reference', '_textbook_crop_call',
            'preparation_textbook_crop_page', 'preparation_textbook_crop_preview',
            'import_textbook_preparation_crops', 'preparation_textbook_crop_original',
            'preparation_image_bytes', 'preparation_textbook_asset_options'):
        setattr(f.facade, name, MethodType(getattr(DesktopWorkbenchFacade, name), f.facade))
    f.facade.personal_textbook_study = lambda: SimpleNamespace(books=lambda: [])
    f.current = f.shelf.snapshot()['reading_notes'][0]
    f.book_key = 'registered:' + f.current['volume_id']
    f.section_key = f.current['section_key']
    f.selection = {'mode': 'registered', 'section_key': f.section_key,
                   'revision': f.current['directory_revision']}
    f.context = textbook_section_material_context(f.facade, f.book_key, f.section_key, f.selection)
    selected = next(row for row in f.context['knowledge_options'] if row.get('note_id') == f.current['note_id'])
    card = import_knowledge_note(f.facade, knowledge_selection(selected))
    f.request = {'kind': 'knowledge', 'context': deepcopy(f.context),
        'knowledge_selection': knowledge_selection(selected), 'knowledge_card': card,
        'knowledge_original': deepcopy(card['original'])}
    return f


def crop_request(f, caption='原书第一个用途', purpose='核对全部条件'):
    request = deepcopy(f.request)
    selection = note_page_selection(f.current, 8)
    page = f.facade.preparation_textbook_crop_page(selection)
    bounds = {'left': 100, 'top': 100, 'right': 610, 'bottom': 630}
    preview = f.facade.preparation_textbook_crop_preview(selection, page['page_revision'], bounds, caption, purpose)
    request.update(kind='textbook_images',
        crop_binding=knowledge_textbook_options(f.facade, request['knowledge_card'])['binding'],
        crop_requests=[{'selection': selection, 'page_revision': page['page_revision'], 'bounds': bounds,
            'caption': caption, 'purpose': purpose, 'expected_preview_revision': preview['preview_revision'],
            'expected_image_sha256': preview['image_asset']['sha256'], 'selection_confirmed': True}])
    return request


def prepare(f, requests, resources=None, assets=None):
    return prepare_section_material_requests(f.facade, None, resources or [], assets or [], requests)


def test_forward_context_is_explicit_real_section_without_classroom_or_auto_knowledge(forward):
    f = forward
    before = files(f.root, f.task_root)
    value = textbook_section_material_context(f.facade, f.book_key, f.section_key, f.selection)
    assert value['entry_kind'] == 'textbook_section' and value['resource_id'] is None
    assert value['knowledge_selection'] is None and value['knowledge_options']
    assert value['section_selection'] == f.selection and value['section_kind'] == 'registered'
    assert value['source_binding']['source_sha256'] == sha256(f.book.read_bytes()).hexdigest()
    assert recheck_textbook_section_material_context(f.facade, value)
    assert all(value[key] is expected for key, expected in FLAGS.items())
    json.dumps(value, ensure_ascii=False)
    assert files(f.root, f.task_root) == before and not f.task_root.exists()


@pytest.mark.parametrize('change', ['book', 'section', 'revision', 'extra_field'])
def test_forward_rejects_wrong_or_incomplete_explicit_scope(forward, change):
    f = forward
    book, section, selection = f.book_key, f.section_key, deepcopy(f.selection)
    if change == 'book': book = 'registered:another-real-looking-book'
    elif change == 'section': section = 'other-section'
    elif change == 'revision': selection['revision'] = 'f' * 64
    else: selection['title'] = '不能由标题猜归属'
    with pytest.raises(ValueError):
        textbook_section_material_context(f.facade, book, section, selection)
    assert not f.task_root.exists()


@pytest.mark.parametrize('change', ['source', 'directory', 'note', 'asset'])
def test_context_recheck_rejects_late_source_directory_note_or_asset_change(forward, change):
    f = forward
    if change == 'source': f.book.write_bytes(f.book.read_bytes() + b'% changed')
    elif change == 'directory':
        f.registry['volumes'][0]['volume_title'] += '新版本'; f.activate()
    elif change == 'note':
        f.note['summary'] += '原文笔记版本变化'; write_packet(f.root, [f.note])
    else:
        asset = deepcopy(f.asset); asset['label'] += '新图表版本'; f.write_assets([asset])
    before = files(f.root, f.task_root)
    with pytest.raises(ValueError): recheck_textbook_section_material_context(f.facade, f.context)
    assert files(f.root, f.task_root) == before and not f.task_root.exists()


def test_no_knowledge_candidate_remains_browsable_and_cannot_invent_anchor(forward):
    f = forward
    f.row['section_key'] = None
    f.catalog.write_text(json.dumps(f.row), encoding='utf8'); write_packet(f.root, [])
    context = textbook_section_material_context(f.facade, f.book_key, f.section_key, f.selection)
    assert context['status'] == 'ready' and context['knowledge_options'] == []
    forged = deepcopy(f.request); forged['context'] = context
    with pytest.raises(ValueError): prepare(f, [forged])
    assert not f.task_root.exists()


def test_ordered_mixed_alias_crops_publish_once_and_preserve_all_receipts(forward):
    f = forward
    first, second = crop_request(f), crop_request(f, '相同像素第二用途', '另一项完整教学说明')
    ordered = [second, f.request, first]
    before = files(f.root, f.task_root)
    inputs = deepcopy(ordered)
    publication = prepare(f, ordered)
    assert files(f.root, f.task_root) == before and not f.task_root.exists()
    result = publication.publish()
    assert [r['original']['kind'] for r in result['resources']] == ['knowledge_note', 'textbook_image', 'textbook_image']
    assert [r['title'] for r in result['resources'][1:]] == ['相同像素第二用途', '原书第一个用途']
    assert len(result['image_assets']) == 1
    assert len(list((f.task_root / 'preparation-v1' / 'textbook-crops').glob('*.json'))) == 2
    for row in result['resources'][1:]:
        assert f.facade.preparation_textbook_crop_original(row, result['image_assets'])['crop_pixels_verified']
        assert row['original']['selected_from_knowledge']['resource_id'] == result['resources'][0]['id']
    publication.recheck_published()
    publication.accept(); publication.discard()
    assert ordered == inputs
    assert {k: v for k, v in files(f.root).items() if not Path(k).is_relative_to(f.task_root)} == before
    assert all(all(row[key] is flag for key, flag in FLAGS.items()) for row in result['resources'])
    with pytest.raises(ValueError): publication.publish()


def test_retained_teacher_knowledge_is_anchor_without_rewriting_work_copy(forward):
    f = forward
    retained = import_knowledge_note(f.facade, f.request['knowledge_selection'])
    retained['title'] = '教师已编辑标题'; retained['blocks'][0]['value'] = '教师已有完整工作内容'
    original = deepcopy(retained)
    publication = prepare(f, [crop_request(f), f.request], [retained])
    result = publication.publish(); publication.accept()
    assert len(result['resources']) == 1 and result['resources'][0]['original']['kind'] == 'textbook_image'
    linked = result['resources'][0]['original']['selected_from_knowledge']
    assert linked['resource_id'] == retained['id'] and linked['original_sha256'] == retained['original_sha256']
    assert retained == original


def test_cancellation_during_crop_staging_never_marks_final_publication(forward, monkeypatch):
    f = forward
    request = crop_request(f)
    event = TaskCancellationEvent()
    original = batch.TextbookPreparationCropService._save_preview
    def stop(service, preview):
        result = original(service, preview)
        event.set()
        return result
    monkeypatch.setattr(batch.TextbookPreparationCropService, '_save_preview', stop)
    before = files(f.root, f.task_root)
    with pytest.raises(ReadCancelled), read_cancel_scope(event): prepare(f, [request, f.request])
    assert event.is_set() and event._published is False
    assert files(f.root, f.task_root) == before and not f.task_root.exists()


def test_guard_late_change_rolls_back_complete_file_set(forward):
    f = forward
    publication = prepare(f, [crop_request(f)])
    before = files(f.root, f.task_root)
    def guard():
        if list((f.task_root / 'preparation-v1' / 'textbook-crops').glob('*.json')):
            raise ValueError('课堂目标发生 ABA 变化')
    with pytest.raises(ValueError, match='ABA'): publication.publish(guard)
    assert files(f.root, f.task_root) == before and not f.task_root.exists()


def test_link_failure_after_first_file_removes_only_owned_output(forward, monkeypatch):
    f = forward
    publication = prepare(f, [crop_request(f)])
    before = files(f.root, f.task_root)
    original, calls = batch.os.link, []
    def fail(source, destination):
        if Path(destination).is_relative_to(f.task_root):
            calls.append(str(destination))
            if len(calls) == 2: raise OSError('synthetic late link failure')
        return original(source, destination)
    monkeypatch.setattr(batch.os, 'link', fail)
    with pytest.raises(OSError, match='synthetic'): publication.publish()
    assert len(calls) == 2 and files(f.root, f.task_root) == before and not f.task_root.exists()


def test_accepted_shared_claim_survives_first_publishers_discard(forward):
    f = forward
    request = crop_request(f)
    first, second = prepare(f, [request]), prepare(f, [request])
    result = first.publish()
    second.publish(); second.accept()
    published = files(f.task_root)
    first.discard()
    assert files(f.task_root) == published
    assert f.facade.preparation_textbook_crop_original(result['resources'][-1], result['image_assets'])['crop_pixels_verified']
    assert not batch._FILE_CLAIMS


def test_pending_edit_late_source_change_rejects_and_owned_discard_cleans(forward):
    f = forward
    publication = prepare(f, [crop_request(f)])
    publication.publish()
    f.note['summary'] += '编辑期间新版本'; write_packet(f.root, [f.note])
    with pytest.raises(ValueError): publication.recheck_published()
    publication.discard()
    assert not f.task_root.exists()


def test_mismatched_contexts_and_missing_explicit_knowledge_fail_before_any_publish(forward):
    f = forward
    other = deepcopy(f.request); other['context']['section_filter_key'] = 'different'
    with pytest.raises(ValueError): prepare(f, [f.request, other])
    other = deepcopy(f.request); other.pop('knowledge_selection')
    with pytest.raises(ValueError): prepare(f, [other])
    assert not f.task_root.exists()


def test_explicit_knowledge_identity_collision_rejects_before_publishing(forward):
    f = forward
    option = next(row for row in f.context['knowledge_options'] if row.get('concept_id'))
    card = import_knowledge_note(f.facade, knowledge_selection(option))
    card['id'] = f.request['knowledge_card']['id']
    other = {'kind': 'knowledge', 'context': deepcopy(f.context),
        'knowledge_selection': knowledge_selection(option), 'knowledge_card': card,
        'knowledge_original': deepcopy(card['original'])}
    with pytest.raises(ValueError): prepare(f, [f.request, other])
    assert not f.task_root.exists()


def test_two_unaccepted_shared_batches_clean_after_the_last_discard(forward):
    f = forward
    request = crop_request(f)
    first, second = prepare(f, [request]), prepare(f, [request])
    first.publish(); second.publish()
    first.discard()
    assert files(f.task_root)
    second.discard()
    assert not f.task_root.exists() and not batch._FILE_CLAIMS and not batch._DIRECTORY_CLAIMS


@pytest.fixture
def lecture_forward(forward, monkeypatch):
    """Real native DOCX services over a read-only snapshot; no desktop stores."""
    from docx import Document
    from test_handout_knowledge import make_facade, png, word_bytes
    from integrations.deeptutor_shchem_v1 import desktop_handout_knowledge as knowledge
    from integrations.deeptutor_shchem_v1.desktop_handout_sources import handout_source_context
    from integrations.deeptutor_shchem_v1.desktop_handout_questions import HandoutQuestionFacade
    from integrations.deeptutor_shchem_v1.desktop_word_questions import WordQuestionService
    f = forward
    doc = Document(); doc.add_heading('知识点1 全部知识网与完整栏目', 3)
    doc.add_paragraph('完整条件 A：温度、单位和物质名称均须保留。')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '条件'; table.cell(0, 1).text = '结论'
    table.cell(1, 0).text = '状态'; table.cell(1, 1).text = '保留比较关系'
    doc.add_paragraph().add_run().add_picture(BytesIO(png(12)))
    doc.add_heading('考向1 完整题', 3)
    doc.add_paragraph('【例1】依据全部条件回答：________。')
    doc.add_paragraph().add_run().add_picture(BytesIO(png(40)))
    doc.add_paragraph('【答案】完整教师答案和全部解析。')
    doc.add_paragraph().add_run().add_picture(BytesIO(png(85)))
    archive = make_facade(word_bytes(doc), f.root / 'own-lecture-source')
    for name in ('list_imported_word_batches', 'imported_word_sources', '_imported_word_source', 'imported_word_path'):
        setattr(f.facade, name, getattr(archive, name))
    f.facade._saved_visual_import_batch = lambda _: {'batch_id': 'BATCH'}
    f.facade._restore_visual_import_sources = lambda *_args, **_kwargs: [f.facade._imported_word_source('BATCH', 'SOURCE')]
    f.facade.state_store = SimpleNamespace(snapshot=lambda: {'drafts': {}})
    f.facade.paths.state_root = f.root / 'owned-word-metadata'
    f.words = WordQuestionService(f.facade)
    f.facade._word_questions = lambda: f.words
    f.facade.word_question_reference = lambda choices, **kwargs: batch._SectionReferenceWords(f.words).reference(choices, **kwargs)
    for name in ('textbook_question_context', 'validate_textbook_question_context',
                 'imported_word_preview', 'imported_word_asset'):
        setattr(f.facade, name, MethodType(getattr(DesktopWorkbenchFacade, name), f.facade))
    monkeypatch.setattr('integrations.deeptutor_shchem_v1.desktop_word_questions.load_attribute_catalog',
        lambda *_args, **_kwargs: {'knowledge_points': [], 'nodes': []})
    meta = knowledge.handout_knowledge_sources(f.facade)['sources'][0]
    outline = knowledge.handout_knowledge_outline(f.facade, meta)
    f.lecture_source = outline['source']
    f.lecture_knowledge = {'source': deepcopy(f.lecture_source), 'outline_revision': outline['outline_revision'],
        'section': next(row for row in outline['sections'] if row['title'].startswith('知识点1')),
        'title': '教师选择的完整知识网', 'teacher_only': False}
    image = handout_source_context(f.facade, f.lecture_source)['images'][0]
    f.lecture_image = {'source': image, 'title': '整张知识网原图', 'purpose': '核对联系', 'teacher_only': False}
    row = HandoutQuestionFacade(f.facade, f.lecture_source).word_question_catalog()['items'][0]
    f.question_choice = {key: row[key] for key in ('key', 'revision')}
    f.question_choice['points'] = 2
    f.lecture_question = {**f.question_choice, 'source_sha256': row['source_sha256'],
                         'scoped_source_document': deepcopy(f.lecture_source)}
    f.archive = archive
    return f


def lecture_request(f, **choices):
    from integrations.deeptutor_shchem_v1.desktop_handout_material_handoff import SCHEMA
    request = deepcopy(f.request)
    request.update(kind='handout_materials', course_handout_source=None,
        material_request={'schema': SCHEMA, 'source': deepcopy(f.lecture_source),
            'knowledge_requests': choices.get('knowledge', []),
            'image_selections': choices.get('images', []),
            'question_selections': choices.get('questions', [])})
    return request


def test_real_knowledge_crop_whole_handout_and_complete_question_keep_requested_order(lecture_forward):
    f = lecture_forward
    requests = [f.request, lecture_request(f, images=[f.lecture_image]), crop_request(f),
        lecture_request(f, knowledge=[f.lecture_knowledge]), lecture_request(f, questions=[f.lecture_question])]
    before = files(f.root)
    publication = prepare(f, requests)
    assert files(f.root) == before and not f.task_root.exists()
    result = publication.publish(); publication.accept()
    assert [row['original']['kind'] for row in result['resources']] == [
        'knowledge_note', 'handout_image', 'textbook_image', 'handout_knowledge', 'complete_question']
    handout = result['resources'][3]['original']
    assert handout['evidence']['section_snapshot']['section'] == f.lecture_knowledge['section']
    assert any('table_rows' in row for row in handout['evidence']['section_snapshot']['blocks'])
    assert len([row for row in result['resources'][-1]['original']['blocks'] if row['kind'] == 'image']) == 2
    assert any(row['role'] == 'answer' for row in result['resources'][-1]['original']['blocks'])
    assert not f.facade.paths.state_root.exists()
    assert all(all(row[key] is flag for key, flag in FLAGS.items()) for row in result['resources'])


def test_real_browse_question_rechecks_complete_reference_without_preview_writes(lecture_forward):
    from integrations.deeptutor_shchem_v1.desktop_knowledge_questions import knowledge_question_context
    f = lecture_forward
    request = deepcopy(f.request)
    request.update(kind='questions', question_mode='browse', question_selections=[f.question_choice],
        question_context=knowledge_question_context(f.facade, request['knowledge_card']),
        reference=f.facade.word_question_reference([f.question_choice], include_images=True))
    before = files(f.root)
    publication = prepare(f, [request])
    assert files(f.root) == before
    result = publication.publish(); publication.accept()
    assert [row['original']['kind'] for row in result['resources']] == ['knowledge_note', 'complete_question']
    assert len(result['image_assets']) == 2 and not f.facade.paths.state_root.exists()


def test_second_handout_failure_and_late_source_reject_entire_mixed_batch(lecture_forward):
    f = lecture_forward
    good = lecture_request(f, knowledge=[f.lecture_knowledge])
    bad = lecture_request(f, questions=[f.lecture_question])
    bad['material_request']['question_selections'][0]['revision'] = 'f' * 64
    before = files(f.root)
    with pytest.raises(ValueError): prepare(f, [good, bad])
    assert files(f.root) == before and not f.task_root.exists()
    publication = prepare(f, [good])
    f.archive.value.data += b'changed source bytes'
    with pytest.raises(ValueError): publication.publish()
    assert files(f.root) == before and not f.task_root.exists()


def test_forward_supplement_preserves_declared_noncontiguous_concept_and_full_section(supplement_fixture):
    f = supplement_fixture
    facade = SimpleNamespace(paths=SimpleNamespace(workspace_root=f.root),
        textbook_bookshelf=f.shelf.snapshot,
        personal_textbook_study=lambda: SimpleNamespace(books=lambda: []))
    selected = next(row for row in f.shelf.snapshot()['supplements'] if row['supplement_node_key'] == 'declared-project')
    selection = {'mode': 'supplement', 'supplement_node_key': selected['supplement_node_key'],
                 'revision': selected['revision']}
    before = files(f.root)
    context = textbook_section_material_context(facade, 'registered:' + f.row['volume_id'],
        'declared-project', selection)
    assert context['section_kind'] == 'supplement' and context['source_binding']['pdf_pages'] == [43, 44, 45]
    assert context['knowledge_options'][0]['pdf_pages'] == [43, 45]
    assert recheck_textbook_section_material_context(facade, context) and files(f.root) == before
    f.product['supplement_nodes'][0]['title'] += '新版本'; f.write_product()
    with pytest.raises(ValueError): recheck_textbook_section_material_context(facade, context)


def test_forward_personal_exact_book_and_section_without_registered_directory(tmp_path):
    from integrations.deeptutor_shchem_v1.desktop_textbook_personal_study import (
        PersonalTextbookStudyService, FLAGS as PERSONAL_FLAGS, STUDY_SCHEMA)
    from integrations.deeptutor_shchem_v1.desktop_textbook_workspace import TextbookWorkspaceService, SCHEMA
    root, state_root = tmp_path / 'workspace', tmp_path / 'owned-personal-state'
    root.mkdir(); (state_root / 'textbook-sources').mkdir(parents=True)
    raw = colored_pdf(4); source_hash = sha256(raw).hexdigest(); source_id = 'BOOK-' + source_hash
    (state_root / 'textbook-sources' / (source_hash + '.pdf')).write_bytes(raw)
    book = {'source_id': source_id, 'source_sha256': source_hash, 'source_name': '无任何映射提示.pdf',
        'page_count': 4, 'human_reviewed': False, 'review_status': 'local_archived_pending_review'}
    section = {'section_id': 'BSEC-' + '1' * 32, 'source_id': source_id, 'source_sha256': source_hash,
        'title': '教师明确登记章节', 'pdf_start': 2, 'pdf_end': 3, 'printed_pages': None,
        'range_confirmed': True, 'created_at': 'fixture', 'updated_at': 'fixture', **PERSONAL_FLAGS}
    record = {'textbook_workspace': {'schema': SCHEMA, 'books': [book], 'candidate_snapshot': None,
        'personal_study': {'schema': STUDY_SCHEMA, 'sections': [section], 'notes': []}}}
    state = SimpleNamespace(snapshot=lambda: deepcopy(record))
    paths = SimpleNamespace(workspace_root=root, state_root=state_root, task_root=state_root / 'tasks')
    study = PersonalTextbookStudyService(TextbookWorkspaceService(paths, state))
    facade = SimpleNamespace(paths=paths, personal_textbook_study=lambda: study,
        textbook_bookshelf=lambda: {'sections': [], 'concepts': [], 'reading_notes': []})
    book_selection = {'source_id': source_id, 'revision': source_hash}
    selection = {'mode': 'personal', 'selection': book_selection, 'section_id': section['section_id'],
        'study_revision': study.book_study(book_selection)['study_revision']}
    before = files(tmp_path)
    context = textbook_section_material_context(facade, 'personal:' + source_id,
        'personal:' + source_id + ':' + section['section_id'], selection)
    assert context['source_binding']['pdf_pages'] == [2, 3]
    assert context['source_binding']['printed_pages'] == 'unknown'
    assert context['section_kind'] == 'personal' and context['knowledge_options'] == []
    assert recheck_textbook_section_material_context(facade, context) and files(tmp_path) == before
    record['textbook_workspace']['personal_study']['sections'][0]['pdf_end'] = 4
    with pytest.raises(ValueError): recheck_textbook_section_material_context(facade, context)


@pytest.mark.parametrize('creator', ['forward', 'legacy'])
def test_shared_claim_between_real_legacy_material_and_forward_accept(lecture_forward, creator):
    from integrations.deeptutor_shchem_v1.desktop_handout_material_handoff import prepare_handout_materials
    f = lecture_forward
    request = lecture_request(f, images=[f.lecture_image])
    forward_publication = prepare(f, [request])
    legacy_publication = prepare_handout_materials(f.facade, request['material_request'], [], [])
    first, accepted = ((forward_publication, legacy_publication) if creator == 'forward'
                       else (legacy_publication, forward_publication))
    first.publish(); result = accepted.publish(); accepted.accept()
    before = files(f.task_root)
    first.discard()
    assert files(f.task_root) == before
    assert all(f.facade.preparation_image_bytes(asset) for asset in result['image_assets'])
    assert not batch._FILE_CLAIMS


@pytest.mark.parametrize('creator', ['forward', 'image'])
def test_shared_claim_between_real_image_batch_and_forward_keeps_receipt(forward, creator):
    from integrations.deeptutor_shchem_v1.desktop_composition_image_batch import prepare_image_batch
    f = forward
    request = crop_request(f)
    forward_publication = prepare(f, [request])
    image_publication = prepare_image_batch(f.facade, [], [], [
        {'kind': 'textbook', 'request': deepcopy(request['crop_requests'][0])}])
    first, accepted = ((forward_publication, image_publication) if creator == 'forward'
                       else (image_publication, forward_publication))
    first.publish(); result = accepted.publish(); accepted.accept()
    before = files(f.task_root)
    first.discard()
    assert files(f.task_root) == before
    crop = next(row for row in result['resources'] if row['original']['kind'] == 'textbook_image')
    assert f.facade.preparation_textbook_crop_original(crop, result['image_assets'])['crop_pixels_verified']
    assert not batch._FILE_CLAIMS


@pytest.mark.parametrize('creator', ['legacy', 'image'])
def test_shared_claim_between_real_legacy_material_and_image_batch(lecture_forward, creator):
    from integrations.deeptutor_shchem_v1.desktop_handout_material_handoff import prepare_handout_materials
    from integrations.deeptutor_shchem_v1.desktop_composition_image_batch import prepare_image_batch
    f = lecture_forward
    request = lecture_request(f, images=[f.lecture_image])
    raw = f.facade.imported_word_asset('BATCH', 'SOURCE', f.lecture_image['source']['asset_id'])['bytes']
    picture = f.root / 'owned-photo.png'; picture.write_bytes(raw)
    choice = {'kind': 'image', 'selection': {'local': {'file_path': str(picture),
        'sha256': sha256(raw).hexdigest(), 'caption': '明确选择的本机图', 'source': str(picture), 'purpose': '比较'},
        'title': '本机用途', 'purpose': '比较', 'teacher_only': False}}
    legacy_publication = prepare_handout_materials(f.facade, request['material_request'], [], [])
    image_publication = prepare_image_batch(f.facade, [], [], [choice])
    first, accepted = ((legacy_publication, image_publication) if creator == 'legacy'
                       else (image_publication, legacy_publication))
    first.publish(); result = accepted.publish(); accepted.accept()
    before = files(f.task_root)
    first.discard()
    assert files(f.task_root) == before
    assert all(f.facade.preparation_image_bytes(asset) == raw for asset in result['image_assets'])
    assert not batch._FILE_CLAIMS


def test_legacy_material_partial_link_failure_still_cleans_only_owned_images(lecture_forward, monkeypatch):
    from integrations.deeptutor_shchem_v1 import desktop_handout_material_handoff as legacy
    f = lecture_forward
    request = lecture_request(f, knowledge=[f.lecture_knowledge], questions=[f.lecture_question])
    publication = legacy.prepare_handout_materials(f.facade, request['material_request'], [], [])
    original, calls = legacy.os.link, []
    def fail(source, destination):
        if Path(destination).parent == f.task_root / 'preparation-v1' / 'images':
            calls.append(str(destination))
            if len(calls) == 2: raise OSError('old publisher partial failure')
        return original(source, destination)
    monkeypatch.setattr(legacy.os, 'link', fail)
    with pytest.raises(OSError, match='old publisher'): publication.publish()
    assert len(calls) == 2 and not list((f.task_root / 'preparation-v1' / 'images').glob('*.image'))
    assert not batch._FILE_CLAIMS


@pytest.mark.parametrize('change', ['request', 'payload', 'receipt'])
def test_forward_mutated_staging_never_publishes(forward, change):
    f = forward
    publication = prepare(f, [crop_request(f)])
    if change == 'request': publication.requests[0]['knowledge_card']['title'] += '改变'
    elif change == 'payload':
        key = next(iter(publication._payloads)); publication._payloads[key] = b'changed pixels'
    else:
        key = next(iter(publication._receipts)); publication._receipts[key] += b'\n'
    with pytest.raises(ValueError): publication.publish()
    assert not f.task_root.exists() and not batch._FILE_CLAIMS


def test_forward_published_receipt_bytes_are_pinned_during_detached_edit(forward):
    f = forward
    publication = prepare(f, [crop_request(f)])
    publication.publish()
    receipt = next((f.task_root / 'preparation-v1' / 'textbook-crops').glob('*.json'))
    value = json.loads(receipt.read_text(encoding='utf8'))
    value['created_at_utc'] = 'external receipt change'
    receipt.write_text(json.dumps(value), encoding='utf8')
    with pytest.raises(ValueError): publication.recheck_published()
    publication.discard()
    assert receipt.exists()  # Preserve an externally changed file, even on discard.
    assert not list((f.task_root / 'preparation-v1' / 'images').glob('*.image'))


def test_forward_post_link_replacement_rejects_and_preserves_foreign_bytes(forward):
    f = forward
    publication = prepare(f, [crop_request(f)])
    replaced = []
    def guard():
        images = list((f.task_root / 'preparation-v1' / 'images').glob('*.image'))
        if images and not replaced:
            images[0].write_bytes(b'external replacement'); replaced.append(images[0])
    with pytest.raises(ValueError): publication.publish(guard)
    assert len(replaced) == 1 and replaced[0].read_bytes() == b'external replacement'
    assert not list((f.task_root / 'preparation-v1' / 'textbook-crops').glob('*.json'))
    assert not batch._FILE_CLAIMS


@pytest.mark.parametrize('kind', ['forward', 'legacy'])
def test_post_link_same_bytes_foreign_inode_is_never_owned(lecture_forward, monkeypatch, kind):
    from integrations.deeptutor_shchem_v1.desktop_handout_material_handoff import prepare_handout_materials
    f = lecture_forward
    request = lecture_request(f, images=[f.lecture_image])
    publication = (prepare(f, [request]) if kind == 'forward' else
                   prepare_handout_materials(f.facade, request['material_request'], [], []))
    real, replaced = batch.os.link, []
    def replace_after_link(source, destination):
        result = real(source, destination)
        target = Path(destination)
        if target.parent == f.task_root / 'preparation-v1' / 'images' and not replaced:
            raw = target.read_bytes(); old_inode = target.stat().st_ino
            target.unlink(); target.write_bytes(raw)
            assert target.stat().st_ino != old_inode
            replaced.append(target)
        return result
    monkeypatch.setattr(batch.os, 'link', replace_after_link)
    publication.publish(); publication.discard()
    assert len(replaced) == 1 and replaced[0].is_file()
    assert replaced[0].read_bytes() == f.facade.imported_word_asset('BATCH', 'SOURCE', f.lecture_image['source']['asset_id'])['bytes']
    assert not batch._FILE_CLAIMS
