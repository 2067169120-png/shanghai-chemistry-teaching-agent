"""Bound Office output, isolation and actual-preview backend selection; no Office activation."""
from pathlib import Path
from types import SimpleNamespace
from threading import Event
import base64,json,os
import pytest
from pptx import Presentation
from integrations.deeptutor_shchem_v1 import office_conversion as office
from integrations.deeptutor_shchem_v1.reader_cancellation import ReadCancelled,read_cancel_scope
from test_office_conversion import write_record


def slide(tmp_path):
    p=tmp_path/'synthetic.pptx';deck=Presentation();deck.slides.add_slide(deck.slide_layouts[6]);deck.save(p);return p


@pytest.mark.parametrize('bad',[None,'word_identity','unbound_before_open','owner','cleanup_pid','pdf','source','foreign_record'])
def test_powerpoint_requires_its_own_preopen_pid_protocol(tmp_path,monkeypatch,bad):
    p=slide(tmp_path);target=tmp_path/'new.pdf'
    def run(command,*,env,**kwargs):
        script=base64.b64decode(command[-1]).decode('utf-16le')
        assert 'PowerPoint.Application' in script and kwargs['contain'] is False
        assert Path(env['SHCHEM_WORD_INPUT']).suffix=='.pptx'
        Path(env['SHCHEM_WORD_OUTPUT']).write_bytes(b'bad' if bad=='pdf' else b'%PDF-synthetic')
        write_record(env,'result',status='converted',**({'source_path':'foreign'} if bad=='foreign_record' else {}))
        write_record(env,'automation',identity='singleuse_empty_com' if bad=='word_identity' else 'new_pid_empty_com',pid_verified_before_open=bad!='unbound_before_open')
        write_record(env,'owner',held_handle=bad!='owner',private=True,pid=2345,start_filetime=133800000000000000)
        write_record(env,'cleanup',pid=6789 if bad=='cleanup_pid' else 2345,disposition='owned_instance_exited')
        if bad=='source':p.write_bytes(b'changed')
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(office,'run_owned',run)
    if bad:
        with pytest.raises(office.OfficeConversionError):office.powerpoint_pdf(p,target,session_root=tmp_path/'sessions')
        assert not target.exists()
    else:assert '已核对该实例退出' in office.powerpoint_pdf(p,target,session_root=tmp_path/'sessions')


@pytest.mark.parametrize('backend',['word','powerpoint'])
def test_source_changed_during_exit_wait_is_not_published(tmp_path,monkeypatch,backend):
    from test_office_conversion import document
    p=document(tmp_path) if backend=='word' else slide(tmp_path);target=tmp_path/'new.pdf'
    def run(command,*,env,**kwargs):
        Path(env['SHCHEM_WORD_OUTPUT']).write_bytes(b'%PDF-synthetic')
        write_record(env,'result',status='converted')
        write_record(env,'automation',identity='singleuse_empty_com' if backend=='word' else 'new_pid_empty_com',pid_verified_before_open=backend=='powerpoint')
        write_record(env,'owner',held_handle=True,private=True,pid=2345,start_filetime=133800000000000000)
        write_record(env,'cleanup',pid=2345,disposition='quit_requested',exception=None)
        return SimpleNamespace(returncode=0)
    def observe(pid,created,**kwargs):
        p.write_bytes(b'changed while waiting for exit')
        return {'pid':pid,'start_filetime':created,'verified':True,'method':'held_handle_exit'}
    monkeypatch.setattr(office,'run_owned',run);monkeypatch.setattr(office,'_observe_bound_process_exit',observe)
    with pytest.raises(office.OfficeConversionError,match='原文件已变化'):
        getattr(office,backend+'_pdf')(p,target,session_root=tmp_path/'sessions')
    assert not target.exists()


def test_powerpoint_in_use_is_specific_and_cannot_publish(tmp_path,monkeypatch):
    p=slide(tmp_path)
    def run(command,*,env,**kwargs):write_record(env,'error',code='powerpoint_in_use');return SimpleNamespace(returncode=0)
    monkeypatch.setattr(office,'run_owned',run)
    with pytest.raises(office.OfficeConversionError,match='正在使用') as e:office.powerpoint_pdf(p,tmp_path/'new.pdf',session_root=tmp_path/'sessions')
    assert e.value.code=='powerpoint_in_use' and not (tmp_path/'new.pdf').exists()


def test_powerpoint_cancel_never_publishes_late_output(tmp_path,monkeypatch):
    p=slide(tmp_path);event=Event();target=tmp_path/'new.pdf'
    def run(command,*,env,cooperative_stop,**kwargs):
        error=ReadCancelled();assert cooperative_stop(SimpleNamespace(wait=lambda _:False),error) is False
        Path(env['SHCHEM_WORD_OUTPUT']).write_bytes(b'%PDF-late retained session only')
        raise error
    monkeypatch.setattr(office,'run_owned',run)
    with pytest.raises(ReadCancelled),read_cancel_scope(event):office.powerpoint_pdf(p,target,session_root=tmp_path/'sessions')
    assert not target.exists()


def test_powerpoint_external_or_macro_content_never_launches(tmp_path,monkeypatch):
    from pptx.opc.constants import RELATIONSHIP_TYPE
    from zipfile import ZipFile,ZIP_DEFLATED
    p=slide(tmp_path);deck=Presentation(p);deck.slides[0].part.relate_to('https://example.invalid/local-only',RELATIONSHIP_TYPE.HYPERLINK,is_external=True);deck.save(p)
    monkeypatch.setattr(office,'run_owned',lambda *a,**k:pytest.fail('must reject before COM'))
    with pytest.raises(office.OfficeConversionError,match='外部链接'):office.powerpoint_pdf(p,tmp_path/'new.pdf')
    p=slide(tmp_path);copy=tmp_path/'macro.pptx'
    with ZipFile(p) as old,ZipFile(copy,'w',ZIP_DEFLATED) as new:
        for n in old.namelist():new.writestr(n,old.read(n))
        new.writestr('ppt/vbaProject.bin',b'synthetic')
    with pytest.raises(office.OfficeConversionError,match='宏内容'):office.powerpoint_pdf(copy,tmp_path/'new.pdf')


@pytest.mark.skipif(os.name!='nt',reason='Production PowerShell cleanup method binding')
@pytest.mark.parametrize('failure',['none','close','quit','takeover'])
def test_powerpoint_production_cleanup_without_com(tmp_path,failure):
    from integrations.deeptutor_shchem_v1.powerpoint_conversion_script import POWERPOINT_SCRIPT
    from integrations.deeptutor_shchem_v1.owned_process import run_owned
    prefix=POWERPOINT_SCRIPT.split("try {\n  Stage 'checking_input'",1)[0]
    cleanup=POWERPOINT_SCRIPT.split('# BEGIN_PRIVATE_CLEANUP:',1)[1].split('\n',1)[1].split('# END_PRIVATE_CLEANUP',1)[0]
    assert 'New-Object -ComObject' not in prefix+cleanup
    setup=r'''
$directory=[Environment]::GetEnvironmentVariable('SHCHEM_SYNTHETIC_CLEANUP','Process')
$failure=[Environment]::GetEnvironmentVariable('SHCHEM_SYNTHETIC_FAILURE','Process')
Add-Type -TypeDefinition @'
using System; using System.Runtime.InteropServices;
public class SyntheticClose { public bool Fail,Called; public void Close() { Called=true; if(Fail) throw new COMException("close",unchecked((int)0x80020005)); } }
public class SyntheticQuit { public bool Fail,Called; public void Quit() { Called=true; if(Fail) throw new COMException("quit",unchecked((int)0x80020005)); } }
public class ShChemConversionOwner { public static uint WaitForSingleObject(IntPtr h,uint w) { return 0; } }
'@
function Private-PowerPoint($empty=$false) { return ($failure -ne 'takeover') }
$doc=New-Object SyntheticClose;$app=New-Object SyntheticQuit
$doc.Fail=($failure -eq 'close');$app.Fail=($failure -eq 'quit');$owned=$true;$pidVerified=$true
'''
    script=tmp_path/'cleanup.ps1';script.write_text(prefix+setup+cleanup+"\nRecord 'calls' @{close=$doc.Called;quit=$app.Called}\n",encoding='utf-8-sig')
    exe=Path(os.environ['SystemRoot'])/'System32/WindowsPowerShell/v1.0/powershell.exe'
    result=run_owned([str(exe),'-STA','-NoProfile','-NonInteractive','-ExecutionPolicy','Bypass','-File',str(script)],env={**os.environ,'SHCHEM_SYNTHETIC_CLEANUP':str(tmp_path),'SHCHEM_SYNTHETIC_FAILURE':failure},timeout=15)
    assert result.returncode==0,result.stderr
    r=json.loads((tmp_path/'cleanup.json').read_text(encoding='utf-8-sig'));calls=json.loads((tmp_path/'calls.json').read_text(encoding='utf-8-sig'))
    if failure=='takeover':assert not calls['close'] and not calls['quit'] and r['disposition']=='taken_over_or_unverified_preserved'
    elif failure=='none':assert calls['close'] and calls['quit'] and r['disposition']=='owned_instance_exited'
    else:
        assert calls['close'] and calls['quit'] is (failure=='quit')
        assert r['exception']['step']==('closing_document' if failure=='close' else 'quitting_application')
        assert r['disposition']=='cleanup_unverified_preserved'


@pytest.mark.parametrize('engine',['powerpoint','libreoffice','missing','cancelled'])
def test_actual_preview_selects_local_engine_without_error_fallback(tmp_path,monkeypatch,engine):
    from integrations.deeptutor_shchem_v1 import desktop_lesson_output as output,desktop_local_pagination as local
    p=slide(tmp_path);before=p.read_bytes();called=[]
    monkeypatch.setattr(output,'checked_file',lambda *a:p)
    monkeypatch.setattr(local,'find_libreoffice',lambda:Path('synthetic-office') if engine in {'libreoffice','cancelled'} else None)
    monkeypatch.setattr(office,'powerpoint_available',lambda:engine!='missing')
    def lo(source,target,*a,**k):
        called.append('libreoffice')
        if engine=='cancelled':raise ReadCancelled()
        result=target/'synthetic.pdf';result.write_bytes(b'%PDF-synthetic');return result
    def ppt(source,target,*a,**k):called.append('powerpoint');target.write_bytes(b'%PDF-synthetic')
    monkeypatch.setattr(office,'libreoffice_pdf',lo);monkeypatch.setattr(office,'powerpoint_pdf',ppt)
    if engine=='missing':
        with pytest.raises(ValueError,match='PowerPoint或LibreOffice'):output.actual_ppt_preview(None,None,'synthetic')
        assert called==[]
    elif engine=='cancelled':
        with pytest.raises(ReadCancelled):output.actual_ppt_preview(None,None,'synthetic')
        assert called==['libreoffice']
    else:
        r=output.actual_ppt_preview(None,None,'synthetic');assert called==[engine]
        assert r['engine']==('Microsoft PowerPoint' if engine=='powerpoint' else 'LibreOffice Impress')
    assert p.read_bytes()==before
