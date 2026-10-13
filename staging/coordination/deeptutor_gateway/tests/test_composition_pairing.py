"""Selected teaching pairs preserve complete material and the surrounding lesson."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from uuid import uuid4

import pytest
from PIL import Image
from test_composition_render import app, material, text

from integrations.deeptutor_shchem_v1.desktop_composition_pairing import compose_teaching_pair
from integrations.deeptutor_shchem_v1.desktop_composition_render import (
    arrange_composition, layout_composition, write_pptx,
)
from integrations.deeptutor_shchem_v1.desktop_composition_rich_render import flatten_elements
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    FORMATS, add_resource, block, digest, new_composition, resource, resource_frames,
    set_click_reveal, set_student_arrangement,
)
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore


def personal(title, body):
    return resource('personal_text', title, [block('text', 'personal', body)],
        {'source_path': '教师输入', 'source_sha256': 'unknown', 'source_revision': 'unknown'})


def composition(*items, format='16:9'):
    value = new_composition(format)
    width, height = FORMATS[format]
    for item in items:
        value = add_resource(value, item, {'page_id': value['pages'][0],
            'x': 12, 'y': 12, 'width': width-24, 'height': height-24})
    return arrange_composition(value, [], {})


@pytest.mark.parametrize('mode', ['side_by_side', 'stacked'])
@pytest.mark.parametrize('format', FORMATS)
def test_short_pair_shares_one_page_without_changing_format_or_sources(app, mode, format):
    first = personal('观察什么？', '先观察实验现象。')
    second = personal('怎样解释？', '用微粒变化解释现象。')
    original = composition(first, second, format=format)
    before = deepcopy(original)
    paired = compose_teaching_pair(original, [], {}, first['id'], second['id'], mode=mode)
    left, right = [resource_frames(paired['placements'][r['id']]) for r in (first, second)]
    assert len(left) == len(right) == len(paired['pages']) == 1
    assert left[0]['page_id'] == right[0]['page_id']
    assert (left[0]['x'] < right[0]['x']) if mode == 'side_by_side' else (left[0]['y'] < right[0]['y'])
    assert paired['format'] == format and paired['resources'] == original['resources']
    assert original == before
    for student in (False, True):
        layout = layout_composition(paired, [], {}, student=student)
        assert not layout['errors']
        assert text(layout, first['blocks'][0]['id']) == first['blocks'][0]['value']
        assert text(layout, second['blocks'][0]['id']) == second['blocks'][0]['value']


@pytest.mark.parametrize('mode', ['side_by_side', 'stacked'])
def test_long_resource_continues_every_character_and_whole_original_image(app, tmp_path, mode):
    buffer = BytesIO()
    Image.new('RGB', (150, 900), '#547987').save(buffer, format='PNG')
    raw = buffer.getvalue()
    store = PreparationImageStore(tmp_path/'images')
    asset = store.import_bytes(raw, '合成整图', '教师', '组合完整性验证')
    picture = resource('preparation_image', '怎样读懂整幅图？',
        [block('image', 'material', asset['asset_id'])],
        {'source_path': asset['source'], 'source_sha256': asset['sha256'],
         'source_revision': digest(asset), 'image_asset': deepcopy(asset), 'source_roles': []})
    body = '完整讲解条件甲乙丙丁。'*280+'最后条件必须保留'
    note = personal('原理如何贯通？', body)
    original = new_composition('16:9')
    for item in (picture, note):
        original = add_resource(original, item, {'page_id': original['pages'][0],
            'x': 12, 'y': 12, 'width': 296, 'height': 156})
    data = {asset['asset_id']: raw}
    original = arrange_composition(original, [asset], data)
    before = deepcopy(original)
    paired = compose_teaching_pair(original, [asset], data, picture['id'], note['id'], mode=mode)
    assert len(resource_frames(paired['placements'][note['id']])) > 1
    for placement in paired['placements'].values():
        frames = resource_frames(placement)
        assert len({f['width'] for f in frames}) == 1
        assert len({f['page_id'] for f in frames}) == len(frames)
    for student in (False, True):
        layout = layout_composition(paired, [asset], data, student=student)
        assert not layout['errors'] and text(layout, note['blocks'][0]['id']) == body
        assert all(e['font_px'] == 32 for page in layout['pages'].values() for e in page
                   if e['kind'] == 'text' and e.get('block_id') == note['blocks'][0]['id'])
        images = [e for page in layout['pages'].values() for e in flatten_elements(page)
                  if e['kind'] == 'image']
        assert len(images) == 1
        assert abs(images[0]['width']/images[0]['height']-150/900) < .0001
    assert paired['resources'] == before['resources'] and original == before
    assert sha256(data[asset['asset_id']]).hexdigest() == asset['sha256']
    assert store.load(asset) == raw


def test_only_vacated_target_pages_are_removed_and_shared_and_blank_pages_stay(app):
    left, right, other = (personal('看什么？', '观察。'), personal('怎么解释？', '解释。'),
                           personal('后续环节', '其他条件不动。'))
    original = composition(left, right, other)
    pages = ['CP-'+uuid4().hex for _ in range(6)]
    original['pages'] = pages
    original['placements'] = {
        left['id']: {'page_id': pages[1], 'x': 12, 'y': 12, 'width': 140, 'height': 60},
        right['id']: {'page_id': pages[4], 'x': 12, 'y': 12, 'width': 140, 'height': 60},
        other['id']: {'page_id': pages[1], 'x': 164, 'y': 12, 'width': 140, 'height': 60}}
    before = deepcopy(original)
    paired = compose_teaching_pair(original, [], {}, left['id'], right['id'])
    assert paired['pages'][0] == pages[0]
    assert paired['pages'][1] not in pages
    assert paired['pages'][2:] == [pages[1], pages[2], pages[3], pages[5]]
    assert paired['placements'][other['id']] == before['placements'][other['id']]
    assert original == before and paired['resources'] == before['resources']


@pytest.mark.parametrize('mode', ['side_by_side', 'stacked'])
def test_complete_long_question_and_teacher_answer_keep_both_audience_scopes(material, mode):
    payload, _, _, data, body = material
    original = deepcopy(payload['lesson_design']['composition'])
    first, second = original['resources']
    before = deepcopy(original)
    paired = compose_teaching_pair(original, payload['image_assets'], data,
                                   first['id'], second['id'], mode=mode)
    assert paired['resources'] == before['resources'] and original == before
    for student in (False, True):
        layout = layout_composition(paired, payload['image_assets'], data, student=student)
        assert not layout['errors']
        assert text(layout, first['blocks'][1]['id']) == body
        assert ('仅教师答案唯一标记' not in text(layout)) == student
        elements = [e for page in layout['pages'].values() for e in flatten_elements(page)]
        assert sum(e['kind'] == 'image' for e in elements) == (1 if student else 2)
        assert all(e['font_px'] == (32 if paired['format'] == '16:9' else 15) for e in elements
                   if e['kind'] == 'text' and e.get('block_id') == first['blocks'][1]['id'])


def test_click_reveal_and_teacher_only_answer_keep_their_roles_and_native_animation(app, tmp_path):
    first = personal('为什么会有变化？', '先观察，再思考原因。')
    first['blocks'].append(block('text', 'teacher_note', '教师原答案保留。'))
    first['original']['blocks'] = deepcopy(first['blocks'])
    first['original_sha256'] = digest(first['original'])
    answer = personal('怎样解释现象？', '微粒变化给出课堂解释。')
    original = set_click_reveal(composition(first, answer), answer['id'], True)
    paired = compose_teaching_pair(original, [], {}, first['id'], answer['id'])
    assert paired['click_reveals'] == original['click_reveals']
    assert paired['resources'] == original['resources']
    static = layout_composition(paired, [], {}, student=True)
    classroom = layout_composition(paired, [], {}, student=True, include_reveals=True)
    teacher = layout_composition(paired, [], {}, student=False)
    assert '课堂解释' not in text(static) and '课堂解释' in text(classroom)
    assert '教师原答案' in text(teacher) and '教师原答案' not in text(classroom)
    assert not static['errors'] and not classroom['errors'] and not teacher['errors']
    target = tmp_path/'pair.pptx'
    write_pptx(classroom, target, {})
    from pptx import Presentation
    presentation = Presentation(target)
    assert len(presentation.slides) == 1
    assert 'nodeType="clickEffect"' in presentation.slides[0]._element.xml
    assert 'shchem-click-'+answer['id'] in presentation.slides[0]._element.xml


def test_student_independent_flow_requires_explicit_keep_choice(app):
    first, second = personal('原图观察', '完整观察条件。'), personal('讲解', '完整解释。')
    original = set_student_arrangement(composition(first, second), 1)
    kept = compose_teaching_pair(original, [], {}, first['id'], second['id'], classroom_uses_pair=False)
    shared = compose_teaching_pair(original, [], {}, first['id'], second['id'], classroom_uses_pair=True)
    assert kept['student_arrangement'] == original['student_arrangement']
    assert 'student_arrangement' not in shared
    assert layout_composition(kept, [], {}, student=True)['independent_student_arrangement']
    assert not layout_composition(shared, [], {}, student=True)['independent_student_arrangement']
    assert original['student_arrangement'] == {'columns': 1}


@pytest.mark.parametrize('option', ['duplicate', 'missing', 'unhashable', 'mode', 'boolean'])
def test_invalid_selection_or_options_are_value_errors_without_mutation(app, option):
    first, second = personal('观察', '条件。'), personal('解释', '结论。')
    original = composition(first, second)
    before = deepcopy(original)
    identities = [first['id'], second['id']]
    options = {}
    if option == 'duplicate': identities[1] = first['id']
    elif option == 'missing': identities[1] = 'CR-'+uuid4().hex
    elif option == 'unhashable': identities[0] = []
    elif option == 'mode': options['mode'] = 'automatic'
    else: options['classroom_uses_pair'] = 1
    with pytest.raises(ValueError):
        compose_teaching_pair(original, [], {}, *identities, **options)
    assert original == before


def test_unrelated_resource_layout_error_is_reported_without_rearranging_it(app):
    first, second, other = personal('观察', '条件。'), personal('解释', '结论。'), personal('保留环节', '完整条件。'*100)
    original = composition(first, second, other)
    original['placements'][other['id']] = {'page_id': original['pages'][-1],
        'x': 12, 'y': 12, 'width': 296, 'height': 12}
    before = deepcopy(original)
    with pytest.raises(ValueError, match='保留环节'):
        compose_teaching_pair(original, [], {}, first['id'], second['id'])
    assert original == before


@pytest.mark.parametrize('role,body', [('personal', ''), ('personal', ' \n '), ('teacher_note', '')])
def test_empty_selected_resource_requires_content_without_mutating_a_valid_draft(app, role, body):
    empty = resource('personal_text', '还没有内容', [block('text', role, body)],
        {'source_path': '教师输入', 'source_sha256': 'unknown', 'source_revision': 'unknown'})
    note = personal('解释', '已填写的完整讲解。')
    original = composition(empty, note)
    before = deepcopy(original)
    with pytest.raises(ValueError, match='请先补充内容'):
        compose_teaching_pair(original, [], {}, empty['id'], note['id'])
    assert original == before


def test_new_pair_pages_respect_existing_160_page_limit(app):
    first, second = personal('观察', '条件。'), personal('解释', '结论。')
    original = composition(first, second)
    original['pages'].extend('CP-'+uuid4().hex for _ in range(159))
    # One vacated page is replaced, so exactly 160 including deliberate blanks
    # remains valid. A longer pair cannot displace those authored blank pages.
    paired = compose_teaching_pair(original, [], {}, first['id'], second['id'])
    assert len(paired['pages']) == 160
    long = deepcopy(original)
    text_block = long['resources'][1]['blocks'][0]
    text_block['value'] = '长讲解完整条件。'*150
    with pytest.raises(ValueError, match='160页'):
        compose_teaching_pair(long, [], {}, first['id'], second['id'])
