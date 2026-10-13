"""Native click entrances stay separate from static student handouts and raw answers."""
from copy import deepcopy
from types import SimpleNamespace
import pytest
from docx import Document
from pptx import Presentation
from test_composition_render import app as render_app
from test_composition_ui import desk, personal
from test_question_explorer_ui import window
from test_desktop_studio_ui import settle
from test_lesson_design import sample
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    new_composition, add_resource, resource, block, set_click_reveal, validate_composition,
    enable_composition, digest)
from integrations.deeptutor_shchem_v1.desktop_composition_render import layout_composition
from integrations.deeptutor_shchem_v1.desktop_lesson_output import export_design, checked_file


def example():
    c=new_composition('16:9')
    question=resource('complete_question','先独立选择',[
        block('text','question','完整条件：判断氢键是否属于化学键。',3),
        block('text','answer','SOURCE_ANSWER_PRIVATE',4)],
        dict(source_path='原讲义.docx',source_sha256='a'*64,source_revision='b'*64,
             question_key='example',scope_review_sha256='c'*64,scope_status='complete_native_scope',
             hierarchy=dict.fromkeys(('paper_id','theme_big_question_id','printed_question_id','atomic_part_id'),'unknown')))
    explanation=personal('CLICK_ANSWER\n分子间氢键与化学键分别判断。')
    c=add_resource(c,question,dict(page_id=c['pages'][0],x=12,y=12,width=296,height=80))
    c=add_resource(c,explanation,dict(page_id=c['pages'][0],x=12,y=100,width=296,height=60))
    return c,question,explanation


def test_click_choice_changes_presentation_only_without_changing_source_or_roles(render_app):
    c,q,a=example();original=deepcopy(c)
    animated=set_click_reveal(c,a['id'],True)
    initial=layout_composition(animated,[],{},student=True)
    completed=layout_composition(animated,[],{},student=True,include_reveals=True)
    assert not initial['errors'] and not completed['errors']
    body=lambda result:'\n'.join(e.get('text','') for es in result['pages'].values() for e in es)
    assert '完整条件' in body(initial) and 'CLICK_ANSWER' not in body(initial)
    assert 'CLICK_ANSWER' in body(completed) and 'SOURCE_ANSWER_PRIVATE' not in body(completed)
    assert c==original and animated['resources']==original['resources']
    assert set_click_reveal(animated,a['id'],False)==original


@pytest.mark.parametrize('change',[
    {'click_reveals':['missing']}, {'click_reveals':True}, {'click_reveals':[{}]},
])
def test_invalid_click_references_are_not_silently_dropped(change):
    c,_,_=example()
    with pytest.raises(ValueError):validate_composition({**c,**change})


def test_registered_classroom_ppt_has_one_click_group_and_static_student_files_hide_it(tmp_path,render_app):
    c,q,a=example();c=set_click_reveal(c,a['id'],True)
    p=sample();p['lesson_design']=enable_composition(p['lesson_design']);p['lesson_design']['composition']=c
    facade=SimpleNamespace(paths=SimpleNamespace(task_root=tmp_path/'tasks'))
    before=deepcopy(p);record=export_design(facade,p);p['lesson_design']['exports'].append(record)
    deck=Presentation(checked_file(facade,p['lesson_design'],record['id'],'student_presentation.pptx'))
    assert len(deck.slides)==1
    slide=deck.slides[0]
    groups=[s for s in slide.shapes if s.name=='shchem-click-'+a['id']]
    assert len(groups)==1 and len(groups[0].shapes)>1
    assert 'CLICK_ANSWER' in ''.join(s.text for s in groups[0].shapes if s.has_text_frame)
    targets=slide._element.xpath('.//p:timing//p:spTgt')
    assert [t.get('spid') for t in targets]==[str(groups[0].shape_id)]
    effects=slide._element.xpath('.//p:cTn[@presetClass="entr"]')
    assert len(effects)==1 and effects[0].get('nodeType')=='clickEffect'
    assert effects[0].get('presetID')=='1'
    student=Document(checked_file(facade,p['lesson_design'],record['id'],'student_worksheet.docx'))
    assert not any('CLICK_ANSWER' in x.text or 'SOURCE_ANSWER_PRIVATE' in x.text for x in student.paragraphs)
    teacher=Document(checked_file(facade,p['lesson_design'],record['id'],'lesson_plan.docx'))
    assert 'SOURCE_ANSWER_PRIVATE' in '\n'.join(x.text for x in teacher.paragraphs)
    assert p['lesson_design']['composition']==before['lesson_design']['composition']


def test_native_click_control_undo_remove_and_recovery_keep_raw_resources(desk):
    dialog,panel,app,page=desk
    panel.change_format('16:9');settle(app)
    panel.add_result({'resources':[personal('课堂分析步骤')]},digest(dialog.history.value));settle(app)
    identity=panel.composition()['resources'][0]['id'];panel.selected=identity;panel.refresh();settle(app)
    original=deepcopy(panel.composition())
    assert panel.click_reveal.isEnabled()
    panel.click_reveal.click();settle(app)
    marked=deepcopy(panel.composition())
    assert marked['click_reveals']==[identity] and marked['resources']==original['resources']
    assert page._payload()['lesson_design']['composition']==marked
    dialog.travel(False);settle(app);assert panel.composition()==original
    dialog.travel(True);settle(app);assert panel.composition()==marked
    panel.selected=identity;panel.remove_selected();settle(app)
    assert 'click_reveals' not in panel.composition()
    dialog.travel(False);settle(app);assert panel.composition()==marked
