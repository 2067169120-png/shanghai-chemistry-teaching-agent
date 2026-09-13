"""Synthetic-only DOCX layout hints; these tests do not certify real pagination."""

from copy import deepcopy
import hashlib
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.shared import Pt
from PIL import Image
import pytest

from integrations.deeptutor_shchem_v1.desktop_personal_visual_questions import _digest
from integrations.deeptutor_shchem_v1.desktop_personal_visual_theme_export import (
    ITEM_KIND,
    SCHEMA_VERSION,
    PersonalVisualThemeExportError,
)
from integrations.deeptutor_shchem_v1.desktop_personal_visual_theme_writer import (
    append_personal_visual_theme,
    build_theme_blocks,
)


class SyntheticTheme:
    """In-memory typed fixture, with real PNG bytes and explicit page geometry."""

    def __init__(self):
        self.assets = {}
        self.item = {
            "kind": ITEM_KIND, "title_zh": "合成布局测试", "source_zh": "合成来源",
            "settings": {"use_source_scores": True}, "source_ref": {},
            "content": {
                "schema_version": SCHEMA_VERSION, "candidate_only": True,
                "teacher_reviewed": False, "publication_allowed": False,
                "images": [], "warnings": ["合成测试警示", "合成测试警示"],
                "theme": {"context": "识别的主题摘要", "shared_materials": [], "printed_questions": []},
            },
        }

    @property
    def theme(self):
        return self.item["content"]["theme"]

    def image(self, box, *, role="question", page_size=(1000, 1400), page_number=1):
        index = len(self.assets) + 1
        output = BytesIO()
        Image.new("RGB", (box[2] - box[0], box[3] - box[1]), (index, 120, 180)).save(output, "PNG")
        raw = output.getvalue()
        digest = hashlib.sha256(raw).hexdigest()
        binding = {
            "source_file_id": "S-ANSWER" if role == "answer" else "S-QUESTION",
            "source_role": "answer" if role == "answer" else "question",
            "page_number": page_number, "page_sha256": str(page_number % 10) * 64,
            "page_size": list(page_size), "pixel_xyxy": list(box), "evidence_id": f"EV-{index}",
        }
        key = "PVTHEMEIMG-" + _digest({"binding": binding, "role": role, "sha256": digest})
        self.assets[key] = raw
        self.item["content"]["images"].append({
            "asset_id": key, "sha256": digest, "role": role, "teacher_only": role == "answer",
            "caption": f"合成裁片{index}", "binding": binding,
            "width": box[2] - box[0], "height": box[3] - box[1], "content_type": "image/png",
        })
        return key

    def node(self, refs=(), **values):
        images = {image["asset_id"]: image for image in self.item["content"]["images"]}
        return {"image_refs": list(refs),
                "evidence_refs": [images[key]["binding"]["evidence_id"] for key in refs], **values}

    def atomic(self, refs=(), answer_refs=(), **values):
        return self.node(refs, part_label=None, stem="合成作答要求", answer=self.node(
            answer_refs, status="present", answer_body="合成答案文字", analysis="",
            chemical_expressions=[], scoring_points=[], max_score=2), **values)

    def question(self, refs=(), atomics=None):
        result = self.node(refs, question_number=len(self.theme["printed_questions"]) + 1,
                           stem="合成题干", atomic_parts=atomics or [self.atomic()])
        self.theme["printed_questions"].append(result)
        return result

    def freeze(self):
        self.item["source_ref"]["content_sha256"] = _digest(self.item["content"])
        return self.item, self.assets


def _render(fixture, audience="student"):
    document = Document()
    item, assets = fixture.freeze()
    blocks = append_personal_visual_theme(document, item, assets, audience)
    assert len(document.paragraphs) == len(blocks) + 1
    return document, blocks, list(zip(blocks, document.paragraphs[1:], strict=True))


def _image_sizes(document, blocks):
    return {block["asset_id"]: shape.width for block, shape in zip(
        (block for block in blocks if block["kind"] == "image"), document.inline_shapes, strict=True)}


def test_narrow_question_and_short_answer_use_source_density_not_full_width():
    fixture = SyntheticTheme()
    stem = fixture.image((100, 100, 350, 140))
    options = fixture.image((100, 150, 900, 210))
    answer = fixture.image((100, 100, 200, 130), role="answer")
    fixture.question([stem], [fixture.atomic([options], [answer])])
    student, student_blocks, _ = _render(fixture)
    teacher, teacher_blocks, _ = _render(fixture, "teacher")
    sizes = _image_sizes(teacher, teacher_blocks)
    section = teacher.sections[-1]
    available = section.page_width - section.left_margin - section.right_margin
    assert sizes[stem] == pytest.approx(available * .25, abs=1)
    assert sizes[options] == pytest.approx(available * .8, abs=1)
    assert sizes[answer] == pytest.approx(available * .1, abs=1)
    assert _image_sizes(student, student_blocks) == {key: sizes[key] for key in (stem, options)}


def test_different_scan_pixel_sizes_use_the_same_normalized_page_density():
    fixture = SyntheticTheme()
    first = fixture.image((100, 100, 350, 140))
    second = fixture.image((200, 200, 700, 280), page_size=(2000, 2800), page_number=2)
    fixture.question([first, second])
    document, blocks, _ = _render(fixture)
    sizes = _image_sizes(document, blocks)
    assert sizes[first] == sizes[second]


def test_tall_image_reduces_the_whole_source_page_scale_without_cropping():
    fixture = SyntheticTheme()
    tall = fixture.image((100, 50, 900, 1950), page_size=(1000, 2000))
    short = fixture.image((100, 50, 300, 90), page_size=(1000, 2000))
    fixture.question([tall, short])
    document, blocks, _ = _render(fixture)
    sizes = _image_sizes(document, blocks)
    assert sizes[tall] / 800 == pytest.approx(sizes[short] / 200, abs=.01)
    section = document.sections[-1]
    assert document.inline_shapes[0].height < section.page_height - section.top_margin - section.bottom_margin


def test_whole_printed_question_keeps_stem_options_and_atomic_fragments_together():
    fixture = SyntheticTheme()
    stem = fixture.image((100, 100, 350, 140))
    options = fixture.image((100, 150, 900, 210))
    fixture.question([stem], [fixture.atomic([options])])
    fixture.question([fixture.image((100, 300, 900, 350))])
    _, _, rows = _render(fixture)
    question = [(block, paragraph) for block, paragraph in rows if block["group_id"] == "question:0"]
    assert [block["kind"] for block, _ in question] == ["text", "image", "image"]
    assert [paragraph.paragraph_format.keep_with_next for _, paragraph in question] == [True, True, False]
    assert all(paragraph.paragraph_format.keep_together for _, paragraph in question)
    assert {block["segment_id"] for block, _ in question} == {"question:0:stem", "question:0:atomic:0"}


def test_oversized_question_breaks_between_complete_atomic_groups_not_infinite_keep():
    fixture = SyntheticTheme()
    atomics = []
    for index in range(3):
        refs = [fixture.image((100, 100 + 700 * index + 300 * part, 900, 400 + 700 * index + 300 * part),
                              page_size=(1000, 4000)) for part in range(2)]
        atomics.append(fixture.atomic(refs))
    fixture.question([fixture.image((100, 3500, 900, 3540), page_size=(1000, 4000))], atomics)
    _, _, rows = _render(fixture)
    question = [(block, paragraph) for block, paragraph in rows if block["group_id"] == "question:0"]
    assert sum(paragraph.paragraph_format.keep_with_next is False for _, paragraph in question) >= 2
    for index in range(3):
        atomic = [paragraph for block, paragraph in question if block["segment_id"] == f"question:0:atomic:{index}"]
        assert len(atomic) == 2
        assert atomic[0].paragraph_format.keep_with_next is True


def test_single_oversized_atomic_can_break_between_fragments():
    fixture = SyntheticTheme()
    refs = [fixture.image((100, 100 + 900 * i, 900, 900 + 900 * i), page_size=(1000, 4000)) for i in range(3)]
    fixture.question([], [fixture.atomic(refs)])
    _, _, rows = _render(fixture)
    atomic = [paragraph for block, paragraph in rows if block["segment_id"] == "question:0:atomic:0"]
    assert len(atomic) == 3
    assert all(paragraph.paragraph_format.keep_with_next is False for paragraph in atomic)


def test_teacher_answer_heading_stays_with_image_and_none_labels_are_omitted():
    fixture = SyntheticTheme()
    for _ in range(2):
        question = fixture.image((100, 100, 900, 160))
        answer = fixture.image((100, 100, 200, 130), role="answer")
        fixture.question([question], [fixture.atomic([], [answer])])
    _, blocks, rows = _render(fixture, "teacher")
    for index in range(2):
        answer = [(block, paragraph) for block, paragraph in rows if block["group_id"] == f"answer:{index}"]
        assert [block["kind"] for block, _ in answer] == ["text", "image"]
        assert answer[0][0]["text"] == "来源参考答案"
        assert answer[0][1].paragraph_format.keep_with_next is True
        assert answer[1][1].paragraph_format.keep_with_next is False
    text = "\n".join(block.get("text", "") for block in blocks)
    assert "None" not in text
    assert text.count("非官方评分认定") == 1
    assert text.count("合成测试警示") == 1
    assert text.count("来源参考分值：2分") == 2


def test_missing_answer_and_text_only_atomic_do_not_print_none_label():
    fixture = SyntheticTheme()
    question = fixture.question()
    question["atomic_parts"][0]["answer"].update(status="missing", max_score=0)
    blocks = build_theme_blocks(*fixture.freeze(), "teacher")
    text = "\n".join(block.get("text", "") for block in blocks)
    assert "合成作答要求" in text and "来源答案待补充" in text and "来源分值待核对" in text
    assert "None" not in text


def test_shared_body_images_replace_synopsis_but_text_only_material_is_retained():
    fixture = SyntheticTheme()
    image = fixture.image((100, 100, 900, 300), role="shared_material")
    fixture.theme["shared_materials"] = [
        fixture.node([image], content="原图中的正文识别", chemical_expressions=[]),
        fixture.node(content="只有文字的重要条件", chemical_expressions=[{"raw": "Fe³⁺ + e⁻ → Fe²⁺"}]),
    ]
    fixture.question()
    before = deepcopy(fixture.freeze())
    document, blocks, _ = _render(fixture)
    text = "\n".join(block.get("text", "") for block in blocks)
    assert "识别的主题摘要" not in text and "原图中的正文识别" not in text
    assert "只有文字的重要条件" in text and "Fe³⁺ + e⁻ → Fe²⁺" in text
    assert any(block.get("asset_id") == image for block in blocks)
    assert fixture.freeze() == before
    output = BytesIO()
    document.save(output)
    with ZipFile(BytesIO(output.getvalue())) as archive:
        media = {archive.read(name) for name in archive.namelist() if name.startswith("word/media/")}
    assert media == {fixture.assets[image]}


@pytest.mark.parametrize("with_figure", [False, True])
def test_context_and_text_remain_without_shared_body_image(with_figure):
    fixture = SyntheticTheme()
    refs = [fixture.image((100, 100, 400, 300), role="shared_material")] if with_figure else []
    material = fixture.node(refs, content="重要条件", chemical_expressions=[])
    material["evidence_refs"] = []  # A figure-only edge, not source-body coverage.
    fixture.theme["shared_materials"] = [material]
    fixture.question()
    blocks = build_theme_blocks(*fixture.freeze(), "student")
    text = "\n".join(block.get("text", "") for block in blocks)
    assert "识别的主题摘要" in text and "重要条件" in text
    assert sum(block["kind"] == "image" for block in blocks) == int(with_figure)


@pytest.mark.parametrize("field,value", [
    ("page_size", None), ("page_size", [True, 1400]), ("pixel_xyxy", [100, 100, 1200, 140]),
    ("pixel_xyxy", [100, 100, 351, 140]),
])
def test_invalid_frozen_geometry_fails_without_guessing_dpi(field, value):
    fixture = SyntheticTheme()
    key = fixture.image((100, 100, 350, 140))
    question = fixture.question([key])
    descriptor = fixture.item["content"]["images"][0]
    descriptor["binding"][field] = value
    new_key = "PVTHEMEIMG-" + _digest({"binding": descriptor["binding"], "role": descriptor["role"],
                                      "sha256": descriptor["sha256"]})
    descriptor["asset_id"] = new_key
    fixture.assets[new_key] = fixture.assets.pop(key)
    question["image_refs"] = [new_key]
    with pytest.raises(PersonalVisualThemeExportError, match="像素|几何"):
        build_theme_blocks(*fixture.freeze(), "student")


def test_same_frozen_page_cannot_have_conflicting_dimensions():
    fixture = SyntheticTheme()
    first = fixture.image((100, 100, 350, 140))
    second = fixture.image((100, 150, 350, 190), page_size=(2000, 2800))
    fixture.question([first, second])
    with pytest.raises(PersonalVisualThemeExportError, match="同一冻结原页"):
        build_theme_blocks(*fixture.freeze(), "student")


def test_tiny_output_page_is_rejected_instead_of_silent_picture_fallback():
    fixture = SyntheticTheme()
    fixture.question([fixture.image((100, 100, 350, 140))])
    document = Document()
    section = document.sections[-1]
    section.page_height = section.top_margin + section.bottom_margin + Pt(50)
    with pytest.raises(PersonalVisualThemeExportError, match="页面尺寸"):
        append_personal_visual_theme(document, *fixture.freeze(), "student")
