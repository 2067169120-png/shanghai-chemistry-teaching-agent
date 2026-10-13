"""Source-bound classroom answers never bypass static student role filtering."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4
import json

import pytest
from PIL import Image
from PySide6.QtWidgets import QApplication
from docx import Document
from pptx import Presentation

from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    FORMATS, add_resource, block, digest, edit_resource, enable_composition,
    new_composition, place_resource, resource, resource_frames, set_click_reveal,
    set_student_arrangement, validate_composition)
from integrations.deeptutor_shchem_v1.desktop_question_answer_reveal import (
    answer_reveal_options, clear_question_answer_reveal, configure_question_answer_reveal,
    question_answer_reveal_errors)
from integrations.deeptutor_shchem_v1.desktop_composition_render import (
    arrange_composition, layout_composition, write_pdf, write_pptx)
from integrations.deeptutor_shchem_v1.desktop_composition_resource_layout import arrange_resource_layout
from integrations.deeptutor_shchem_v1.desktop_composition_rich_render import flatten_elements
from integrations.deeptutor_shchem_v1.desktop_preparation_images import image_info
from integrations.deeptutor_shchem_v1.desktop_lesson_design import new_design
from integrations.deeptutor_shchem_v1.desktop_lesson_output import export_composition, write_documents


@pytest.fixture(scope='module')
def qt_app():
    from integrations.deeptutor_shchem_v1.desktop_workbench.typography import install_ui_font
    app = QApplication.instance() or QApplication([])
    app.setStyle('Fusion'); install_ui_font(app)
    return app


def personal(text, role='personal'):
    return resource('personal_text', text, [block('text', role, text)],
        {'source_path': '作者输入', 'source_sha256': 'unknown', 'source_revision': 'unknown'})


def question(*, picture=False, answer='【答案】AB', title='依据完整题面判断'):
    rows = [block('text', 'shared_material', '共同条件保持不变。', 0),
        block('text', 'question', '题干完整文字：根据给定材料判断下列说法。', 1)]
    assets, data = [], {}
    if picture:
        buffer = BytesIO(); Image.new('RGB', (1269, 298), '#5789af').save(buffer, format='PNG')
        raw = buffer.getvalue(); checksum = sha256(raw).hexdigest()
        asset = dict(asset_id='IMG-'+checksum, sha256=checksum, caption='合成题图',
            source='合成测试', purpose='完整图片', **image_info(raw))
        rows.append(block('image', 'question', asset['asset_id'], 2))
        assets.append(asset); data[asset['asset_id']] = raw
    rows.extend(block('text', 'question', letter+'．完整选项条件。', index+3)
                for index, letter in enumerate('ABCD'))
    rows.extend([block('text', 'answer', answer, 10),
        block('text', 'answer', 'TEACHER_REASON_PRIVATE：这一块解释全部判断理由。', 11),
        block('text', 'teacher_note', 'TEACHER_NOTE_PRIVATE', 12)])
    item = resource('complete_question', title, rows, dict(source_path='合成题.docx',
        source_sha256='a'*64, source_revision='b'*64, question_key='synthetic-complete',
        scope_review_sha256='c'*64, scope_status='complete_native_scope',
        hierarchy=dict.fromkeys(('paper_id', 'theme_big_question_id', 'printed_question_id', 'atomic_part_id'), 'unknown')))
    value = new_composition('16:9')
    value = add_resource(value, item, dict(page_id=value['pages'][0], x=12, y=12, width=296, height=156))
    return value, item, rows[-3]['id'], assets, data


def text(layout):
    return ''.join(e.get('text', '') for page in layout['pages'].values() for e in flatten_elements(page))


def ppt_text(shapes):
    return ''.join(ppt_text(s.shapes) if hasattr(s, 'shapes') else s.text if s.has_text_frame else '' for s in shapes)


def payload(value, assets=(), *, notes=False):
    result = dict(topic='课堂完整题', audience='合成验证', lesson_timing='一课时', objective='保留完整条件', image_assets=list(assets))
    result['lesson_design'] = enable_composition(new_design(result))
    result['lesson_design'].update(composition=value, classroom_notes=notes)
    return result


def configured(qt_app, **kwargs):
    value, item, answer_id, assets, data = question(**kwargs)
    result = configure_question_answer_reveal(value, item['id'], [answer_id], assets, data)
    return value, result, item, answer_id, assets, data


def test_explicit_answer_identity_preserves_original_and_only_projects_the_selected_answer(qt_app):
    original, value, item, answer_id, assets, data = configured(qt_app)
    assert [r['selected'] for r in answer_reveal_options(item)] == [False, False]
    selected = value['resources'][0]
    for key in ('original', 'original_sha256', 'blocks', 'title', 'locked'):
        assert selected[key] == item[key]
    assert answer_reveal_options(selected)[0]['selected']
    static = layout_composition(value, assets, data, student=True)
    classroom = layout_composition(value, assets, data, student=True, include_reveals=True)
    teacher = layout_composition(value, assets, data, student=False)
    assert not (static['errors'] or classroom['errors'] or teacher['errors'])
    assert len(static['pages']) == len(classroom['pages']) == 1
    assert 'D．完整选项条件。' in text(static)
    assert '【答案】AB' not in text(static) and '【答案】AB' in text(classroom)
    assert 'TEACHER_REASON_PRIVATE' not in text(classroom) and 'TEACHER_NOTE_PRIVATE' not in text(classroom)
    assert 'TEACHER_REASON_PRIVATE' in text(teacher) and 'TEACHER_NOTE_PRIVATE' in text(teacher)
    answers = [e for p in classroom['pages'].values() for e in flatten_elements(p) if e.get('reveal_id')]
    assert len(answers) == 1 and answers[0]['block_id'] == answer_id and answers[0]['role'] == 'answer'
    assert answers[0]['original_sha256'] == item['original_sha256']
    assert answers[0]['source_index'] == 10 and answers[0]['font_px'] == 32
    cleared = clear_question_answer_reveal(value, item['id'])
    assert cleared['resources'] == original['resources'] and cleared['placements'] == value['placements']
    assert '【答案】' not in text(layout_composition(cleared, assets, data, include_reveals=True))
    assert validate_composition(json.loads(json.dumps(value))) == value


def test_teacher_work_answer_edit_is_projected_by_original_identity_without_changing_original(qt_app):
    before, item, answer_id, assets, data = question()
    value = edit_resource(before, item['id'], locked=False)
    value = edit_resource(value, item['id'], values={answer_id: '【答案】AC', item['blocks'][1]['id']: '教师明确修改的题干。'})
    assert answer_reveal_options(value['resources'][0])[0]['working_value_changed']
    value = configure_question_answer_reveal(value, item['id'], [answer_id], assets, data)
    assert value['resources'][0]['original'] == item['original']
    classroom = layout_composition(value, assets, data, include_reveals=True)
    assert '【答案】AC' in text(classroom) and '【答案】AB' not in text(classroom)
    assert '教师明确修改的题干。' in text(classroom)
    static = layout_composition(value, assets, data)
    assert '【答案】' not in text(static)


@pytest.mark.parametrize('columns', [1, 2])
def test_configured_question_is_full_width_exclusive_in_flow_and_local_reflow(qt_app, columns):
    original, item, answer_id, assets, data = question()
    left, right = personal('之前素材'), personal('之后素材')
    original['resources'].insert(0, left); original['resources'].append(right)
    for row in (left, right):
        original['placements'][row['id']] = dict(page_id=original['pages'][0], x=12, y=12, width=100, height=60)
    unrelated = {identity: deepcopy(original['placements'][identity]) for identity in (left['id'], right['id'])}
    value = configure_question_answer_reveal(original, item['id'], [answer_id], assets, data)
    assert all(value['placements'][key] == frame for key, frame in unrelated.items())
    assert original['pages'][0] in value['pages']  # Shared old page is retained.
    value = set_student_arrangement(value, columns)
    laid = layout_composition(value, assets, data)
    assert not laid['errors']
    containing = [es for es in laid['pages'].values() if any(e['resource_id'] == item['id'] for e in es)]
    assert len(containing) == 1 and {e['resource_id'] for e in containing[0]} == {item['id']}
    frame = next(f['frame'] for f in laid['fragments'] if f['resource_id'] == item['id'])
    assert (frame['x'], frame['y'], frame['width'], frame['height']) == (12, 7, 296, 166)
    assert '之前素材' in text(laid) and '之后素材' in text(laid)
    again = arrange_resource_layout(value, item['id'], assets, data)
    assert again['placements'][item['id']]['height'] == 166
    assert all(again['placements'][key] == frame for key, frame in unrelated.items())
    teacher = arrange_composition(value, assets, data, columns=columns)
    assert teacher['placements'][item['id']]['width'] == 296
    assert teacher['placements'][item['id']]['height'] == 166


@pytest.mark.parametrize('format', ['A4', 'A3'])
def test_print_formats_preserve_inactive_classroom_intent_and_restore_it(qt_app, format):
    _, value, item, answer_id, assets, data = configured(qt_app)
    config = deepcopy(value['resources'][0]['question_answer_reveal'])
    value['format'] = format
    w, h = FORMATS[format]
    value['placements'][item['id']] = dict(page_id=value['pages'][0], x=12, y=12, width=w-24, height=h-24)
    printed = arrange_composition(value, assets, data)
    assert printed['resources'][0]['question_answer_reveal'] == config
    assert '【答案】AB' not in text(layout_composition(printed, assets, data, include_reveals=True))
    printed['format'] = '16:9'
    printed['placements'][item['id']] = dict(page_id=printed['pages'][0], x=12, y=12, width=296, height=156)
    restored = arrange_composition(printed, assets, data)
    assert restored['resources'][0]['question_answer_reveal'] == config
    assert restored['placements'][item['id']]['height'] == 166
    assert '【答案】AB' in text(layout_composition(restored, assets, data, include_reveals=True))


@pytest.mark.parametrize('malice', ['digest', 'question', 'shared', 'teacher', 'missing', 'duplicates', 'role', 'source', 'extra'])
def test_malicious_or_stale_bindings_do_not_reveal_any_answer(qt_app, malice):
    _, value, item, answer_id, assets, data = configured(qt_app)
    target = value['resources'][0]; config = target['question_answer_reveal']
    if malice == 'digest': config['original_sha256'] = '0'*64
    elif malice in {'question', 'shared', 'teacher'}:
        role = {'question': 'question', 'shared': 'shared_material', 'teacher': 'teacher_note'}[malice]
        config['answer_block_ids'] = [next(b['id'] for b in item['blocks'] if b['role'] == role)]
    elif malice == 'missing': config['answer_block_ids'] = ['CB-'+'f'*32]
    elif malice == 'duplicates': config['answer_block_ids'] = [answer_id, answer_id]
    elif malice == 'role': target['blocks'][-3]['role'] = 'question'
    elif malice == 'source': target['original']['blocks'][-3]['value'] = '替换来源'
    else: config['teacher_only'] = False
    with pytest.raises(ValueError): layout_composition(value, assets, data, include_reveals=True)


@pytest.mark.parametrize('mode', ['long_title', 'long_answer', 'multiline_answer', 'too_tall'])
def test_unfittable_title_answer_and_full_question_are_not_shortened(qt_app, mode):
    kwargs = {'title': '过长标题'*30} if mode == 'long_title' else {'answer': '长答案'*30} if mode == 'long_answer' else {'answer': '第一行\n第二行'} if mode == 'multiline_answer' else {}
    value, item, answer_id, assets, data = question(**kwargs)
    if mode == 'too_tall':
        value = edit_resource(value, item['id'], locked=False)
        value = edit_resource(value, item['id'], values={item['blocks'][1]['id']: '题干完整文字。'*400})
    before = deepcopy(value)
    with pytest.raises(ValueError): configure_question_answer_reveal(value, item['id'], [answer_id], assets, data)
    assert value == before


@pytest.mark.parametrize('mode', ['missing', 'corrupt', 'wrong_answer_image'])
def test_missing_corrupt_images_and_answer_picture_identity_are_rejected(qt_app, mode):
    value, item, answer_id, assets, data = question(picture=True)
    if mode == 'missing': data = {}
    elif mode == 'corrupt': data = {assets[0]['asset_id']: b'not the original bytes'}
    else: answer_id = next(b['id'] for b in item['blocks'] if b['kind'] == 'image')
    before = deepcopy(value)
    with pytest.raises(ValueError): configure_question_answer_reveal(value, item['id'], [answer_id], assets, data)
    assert value == before


def test_manual_frames_are_measured_and_overlap_is_not_ignored_by_student_flow(qt_app):
    _, value, item, answer_id, assets, data = configured(qt_app)
    short = place_resource(value, item['id'], {**value['placements'][item['id']], 'height': 40})
    assert question_answer_reveal_errors(short, item['id'], assets, data)
    short = set_student_arrangement(short, 2)
    assert question_answer_reveal_errors(short, item['id'], assets, data)
    overlap = personal('共享页素材')
    value = add_resource(value, overlap, dict(page_id=value['placements'][item['id']]['page_id'], x=20, y=20, width=70, height=40))
    assert any('重叠' in e for e in question_answer_reveal_errors(value, item['id'], assets, data))
    with pytest.raises(ValueError): place_resource(value, item['id'], {**value['placements'][item['id']], 'y': 30})


def test_configuring_clears_only_same_resource_whole_click_and_rejects_teacher_only(qt_app):
    value, item, answer_id, assets, data = question()
    other = personal('另一个点击步骤')
    page = 'CP-'+uuid4().hex; value['pages'].append(page)
    value = add_resource(value, other, dict(page_id=page, x=12, y=12, width=296, height=80))
    value = set_click_reveal(set_click_reveal(value, item['id'], True), other['id'], True)
    result = configure_question_answer_reveal(value, item['id'], [answer_id], assets, data)
    assert result['click_reveals'] == [other['id']]
    with pytest.raises(ValueError): set_click_reveal(result, item['id'], True)
    with pytest.raises(ValueError): answer_reveal_options(personal('仅教师内容', 'teacher_note'))
    no_answer = configure_question_answer_reveal(value, item['id'], [], assets, data)
    assert '【答案】' not in text(layout_composition(no_answer, assets, data, include_reveals=True))


def test_work_answer_retains_native_subscript_and_real_title_textbox_clearance(qt_app, tmp_path):
    value, item, old_id, assets, data = question(answer='H_{2}O')
    rows = deepcopy(item['blocks'])
    rows[-3]['presentation'] = {'schema': 'shchem.native-word-paragraph.v1', 'images': [],
        'scripts': [{'start': 1, 'end': 5, 'script': 'subscript'}]}
    rich = resource('complete_question', item['title'], rows, item['original']['evidence'])
    value['resources'] = [rich]; value['placements'] = {rich['id']: value['placements'][item['id']]}
    value = configure_question_answer_reveal(value, rich['id'], [old_id], assets, data)
    classroom = layout_composition(value, assets, data, include_reveals=True)
    answer = [e for es in classroom['pages'].values() for e in flatten_elements(es) if e.get('reveal_id')]
    assert ''.join(e['text'] for e in answer) == 'H2O'
    assert any(e['text'] == '2' and e['script'] == 'AlignSubScript' for e in answer)
    title = next(e for es in classroom['pages'].values() for e in es if e.get('text') == rich['title'])
    assert title['x']+title['width'] <= min(e['x'] for e in answer)-10+.001
    write_pptx(classroom, tmp_path/'rich.pptx', data)
    group = next(s for s in Presentation(tmp_path/'rich.pptx').slides[0].shapes if hasattr(s, 'shapes'))
    assert ppt_text(group.shapes) == 'H2O' and len(group.shapes) == 3


def test_source_answer_entrance_coexists_with_process_steps_and_other_clicks(qt_app, tmp_path):
    value, item, answer_id, assets, data = question()
    process = resource('personal_text', '讲解过程', [block('text', 'personal', '第一步'), block('text', 'personal', '第二步')],
        dict(source_path='教师输入', source_sha256='unknown', source_revision='unknown'))
    process['exposition'] = {'schema': 'shchem.classroom-exposition.v1', 'kind': 'process'}
    other = personal('独立点击内容')
    for row in (process, other):
        page = 'CP-'+uuid4().hex; value['pages'].append(page)
        value = add_resource(value, row, dict(page_id=page, x=12, y=12, width=296, height=156))
        value = set_click_reveal(value, row['id'], True)
    value = configure_question_answer_reveal(value, item['id'], [answer_id], assets, data)
    assert value['click_reveals'] == [process['id'], other['id']]
    classroom = layout_composition(value, assets, data, include_reveals=True)
    write_pptx(classroom, tmp_path/'mixed.pptx', data)
    deck = Presentation(tmp_path/'mixed.pptx')
    groups = [s for slide in deck.slides for s in slide.shapes if s.name.startswith('shchem-click-')]
    assert len(groups) == 4
    assert [ppt_text(s.shapes) for s in groups if '-answer-' in s.name] == ['【答案】AB']
    assert sum(len(slide._element.xpath('.//p:cTn[@nodeType="clickEffect"]')) for slide in deck.slides) == 4
    assert '完整学生' not in ''.join(ppt_text(s.shapes) for s in groups)


def test_configured_frame_preserves_manual_geometry_and_fingerprint_revision(qt_app, monkeypatch):
    from integrations.deeptutor_shchem_v1 import desktop_lesson_composition as module
    from integrations.deeptutor_shchem_v1.desktop_lesson_design import content_fingerprint
    _, value, item, answer_id, assets, data = configured(qt_app)
    value = place_resource(value, item['id'], {**value['placements'][item['id']], 'y': 5, 'height': 170})
    assert not question_answer_reveal_errors(value, item['id'], assets, data)
    reflow = arrange_composition(value, assets, data)
    assert (reflow['placements'][item['id']]['y'], reflow['placements'][item['id']]['height']) == (5, 170)
    with pytest.raises(ValueError): place_resource(value, item['id'], {'page_id': value['pages'][0]})
    p = payload(value, assets); current = content_fingerprint(p)
    assert module.RENDER_REVISION.endswith('.v41')
    monkeypatch.setattr(module, 'RENDER_REVISION', 'shchem.composition-classroom-exposition.v38')
    assert content_fingerprint(p) != current


def test_blank_work_question_and_answer_only_source_cannot_unlock_classroom_answer(qt_app):
    value, item, answer_id, assets, data = question()
    value = edit_resource(value, item['id'], locked=False)
    value = edit_resource(value, item['id'], values={b['id']: '  ' for b in item['blocks'] if b['role'] == 'question'})
    with pytest.raises(ValueError): configure_question_answer_reveal(value, item['id'], [answer_id], assets, data)
    raw = deepcopy(item)
    for rows in (raw['blocks'], raw['original']['blocks']):
        for row in rows:
            if row['role'] == 'question': row['role'] = 'answer'
    raw['original_sha256'] = digest(raw['original'])
    with pytest.raises(ValueError): answer_reveal_options(raw)


@pytest.mark.parametrize('notes', [False, True])
def test_real_office_and_static_pdf_outputs_keep_answers_in_only_the_classroom_entrance(qt_app, tmp_path, notes):
    _, value, item, answer_id, assets, data = configured(qt_app)
    p = payload(value, assets, notes=notes)
    facade = SimpleNamespace(paths=SimpleNamespace(task_root=tmp_path/'tasks'), preparation_image_bytes=lambda asset: data[asset['asset_id']])
    record = export_composition(facade, p)
    output = tmp_path/'tasks/preparation-v1/node-exports'/record['id']
    student = Document(output/'student_worksheet.docx')
    teacher = Document(output/'lesson_plan.docx')
    assert '【答案】' not in ''.join(p.text for p in student.paragraphs)
    assert 'TEACHER_REASON_PRIVATE' not in ''.join(p.text for p in student.paragraphs)
    assert 'TEACHER_REASON_PRIVATE' in ''.join(p.text for p in teacher.paragraphs)
    deck = Presentation(output/'student_presentation.pptx'); assert len(deck.slides) == 1
    slide = deck.slides[0]
    groups = [s for s in slide.shapes if s.name.startswith('shchem-click-')]
    assert len(groups) == 1 and ppt_text(groups[0].shapes) == '【答案】AB'
    assert 'D．完整选项条件。' in ppt_text([s for s in slide.shapes if s not in groups])
    assert 'TEACHER_REASON_PRIVATE' not in ppt_text(slide.shapes)
    assert [t.get('spid') for t in slide._element.xpath('.//p:timing//p:spTgt')] == [str(groups[0].shape_id)]
    assert len(slide._element.xpath('.//p:cTn[@nodeType="clickEffect"]')) == 1
    audit = json.loads((output/'canvas-output.json').read_text(encoding='utf-8'))
    assert audit['classroom_answer_reveals'][0]['answer_block_ids'] == [answer_id]
    static = layout_composition(value, assets, data)
    write_pptx(static, tmp_path/'static.pptx', data)
    assert '【答案】' not in ppt_text(Presentation(tmp_path/'static.pptx').slides[0].shapes)
    completed = layout_composition(value, assets, data, include_reveals=True)
    write_pdf(completed, tmp_path/'completed-static.pdf')
    import pypdfium2 as pdfium
    for path in (output/'student_canvas.pdf', tmp_path/'completed-static.pdf'):
        with pdfium.PdfDocument(str(path)) as pdf:
            body = ''.join(page.get_textpage().get_text_range() for page in pdf)
        assert '【答案】' not in body and 'TEACHER_REASON_PRIVATE' not in body
