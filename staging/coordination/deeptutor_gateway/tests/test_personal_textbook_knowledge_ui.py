"""Native candidate review, resumable edits and current-lesson handoff."""
from copy import deepcopy
from types import SimpleNamespace

import pytest
from PySide6.QtCore import Qt, QRect
from PySide6.QtWidgets import QDialog

from test_personal_textbook_knowledge import native_fixture as native_fixture, native_pdf, draft_for
from test_personal_textbook_study import chapter
from test_personal_textbook_ui import app as app, facade_for, Bridge, complete
from test_desktop_studio_ui import window as window, settle
from test_textbook_bookshelf_ui import painted
from integrations.deeptutor_shchem_v1.desktop_workbench.textbook_knowledge_dialog import TextbookKnowledgeDialog


@pytest.fixture
def dialog(app, native_fixture):
    f = native_fixture; view = chapter(f); bridge = Bridge()
    value = TextbookKnowledgeDialog(facade_for(f), bridge, f.selection, view, view['sections'][0])
    value.resize(1060, 810); value.show(); complete(bridge, '读取教材知识候选批次'); app.processEvents()
    yield value, bridge, f
    if not value._closed:
        value._edits.clear(); value._draft = None; value.reject()
    value.deleteLater(); app.processEvents()


def generate(value, bridge):
    value.generate_button.click(); complete(bridge, '整理教材知识候选')
    complete(bridge, '保存教材知识候选批次')


def check_all(value):
    for index in range(value.items.count()):
        value.items.setCurrentRow(index); value.confirmed.setChecked(True)
    value.select_button.click()


def test_native_preview_durable_candidate_and_each_row_confirmation(dialog):
    value, bridge, f = dialog; generate(value, bridge)
    assert value.items.count() == 2 and not value.save_button.isEnabled()
    assert f.study.book_study(f.selection)['notes'] == [] and len(value._draft_rows) == 1
    value.items.item(0).setCheckState(Qt.CheckState.Checked)
    assert not value.save_button.isEnabled()
    value.confirmed.setChecked(True); assert value.save_button.isEnabled()
    value.text.setPlainText('教师修订：A → B（条件X）；2.0 g。')
    assert not value.confirmed.isChecked() and value.items.item(0).checkState() == Qt.CheckState.Unchecked
    assert not value.save_button.isEnabled()
    value.confirmed.setChecked(True); value.select_button.click()
    assert value.items.item(1).checkState() == Qt.CheckState.Unchecked
    check_all(value); emitted = []; value.notes_saved.connect(emitted.append)
    value.save_button.click(); complete(bridge, '保存勾选教材知识笔记')
    assert value._closed and emitted[0]['added'] == 2
    assert emitted[0]['view']['notes'][0]['text'] == '教师修订：A → B（条件X）；2.0 g。'
    assert f.source.read_bytes() == f.original


def test_page_location_reads_and_paints_real_source_inside_review(dialog, app):
    value, bridge, f = dialog; generate(value, bridge)
    value.read_button.click(); app.processEvents(); complete(bridge, '读取教材原页'); app.processEvents()
    assert value.tabs.currentIndex() == 1
    assert value._reader.view.pageNavigator().currentPage() == 1
    assert [value._reader.pages.itemData(i) for i in range(value._reader.pages.count())] == [2, 3]
    assert painted(app, value._reader) > 10 and 'Synthetic' in value._draft['items'][0]['text']
    value._reader.next.click(); app.processEvents()
    assert value._reader.view.pageNavigator().currentPage() == 2 and f.source.read_bytes() == f.original


def test_closing_saves_edits_and_reopening_requires_fresh_checks(dialog, app):
    value, bridge, f = dialog; generate(value, bridge)
    value.title.setText('保留修订标题'); value.text.setPlainText('保留修订内容与条件X。')
    value.confirmed.setChecked(True); value.select_button.click()
    value.close_button.click(); assert not value._closed
    complete(bridge, '保存教材知识核对进度'); assert value._closed
    view = f.study.book_study(f.selection); other_bridge = Bridge()
    reopened = TextbookKnowledgeDialog(facade_for(f), other_bridge, f.selection, view, view['sections'][0])
    reopened.show(); complete(other_bridge, '读取教材知识候选批次')
    reopened.saved_batches.setCurrentIndex(1); app.processEvents()
    assert reopened.title.text() == '保留修订标题' and reopened.text.toPlainText() == '保留修订内容与条件X。'
    assert not reopened.confirmed.isChecked() and not reopened.save_button.isEnabled()
    assert reopened._draft['items'][0]['title'] != '保留修订标题'
    assert all(reopened.items.item(i).checkState() == Qt.CheckState.Unchecked for i in range(reopened.items.count()))
    before = f.state.path.read_bytes(); reopened.reject()
    assert reopened._closed and f.state.path.read_bytes() == before
    reopened.deleteLater(); app.processEvents()


def test_invalid_progress_keeps_window_and_previous_draft(dialog):
    value, bridge, f = dialog; generate(value, bridge)
    value.text.setPlainText('字' * 1201); before = f.state.path.read_bytes()
    value.close_button.click(); task = bridge.take('保存教材知识核对进度')
    with pytest.raises(Exception) as exc: task.operation()
    bridge.fail(task, str(exc.value)); bridge.task_finished.emit(task.task_id)
    assert not value._closed and not value._busy and '未覆盖' in value.status.text()
    assert f.state.path.read_bytes() == before


def test_model_requires_fresh_checkbox_after_each_selection_change(dialog):
    value, _bridge, _f = dialog
    value.profiles.addItem('合成模型', {'profile_id': 'fixture', 'profile_revision': 'revision'})
    value.profiles.addItem('合成模型二', {'profile_id': 'fixture2', 'profile_revision': 'revision2'})
    value.profiles.setCurrentIndex(0)
    value.mode.setCurrentIndex(1)
    assert not value.send_confirmed.isChecked() and not value.generate_button.isEnabled()
    value.send_confirmed.setChecked(True); assert value.generate_button.isEnabled()
    value.profiles.setCurrentIndex(1); assert not value.send_confirmed.isChecked()
    value.send_confirmed.setChecked(True); value.end.setValue(2)
    assert not value.send_confirmed.isChecked() and not value.generate_button.isEnabled()
    value.mode.setCurrentIndex(0); assert value.generate_button.isEnabled()


def test_cancel_and_late_result_do_not_leave_busy_or_save_candidate(dialog):
    value, bridge, f = dialog
    value.generate_button.click(); task = bridge.take('整理教材知识候选')
    draft = task.operation(); before = f.state.path.read_bytes()
    bridge.task_cancelled.emit(task.task_id, task.label); bridge.task_finished.emit(task.task_id)
    bridge.succeed(task, draft, run=False)
    assert not value._busy and value._draft is None and value.generate_button.isEnabled()
    assert '已取消' in value.status.text() and f.state.path.read_bytes() == before
    value.generate_button.click(); task = bridge.take('整理教材知识候选'); value.reject()
    bridge.succeed(task, draft, run=False); bridge.task_finished.emit(task.task_id)
    assert value._closed and value._draft is None and f.state.path.read_bytes() == before


@pytest.mark.parametrize('width', [420, 1060])
def test_review_actions_fit_desktop_narrow_and_wide(dialog, app, width):
    from integrations.deeptutor_shchem_v1.desktop_workbench.studio_style import WORKBENCH_STYLE
    value, bridge, _f = dialog; value.setStyleSheet(WORKBENCH_STYLE); generate(value, bridge)
    value.resize(width, 980); app.processEvents()
    assert value.width() == width
    assert value.splitter.orientation() == (Qt.Orientation.Vertical if width < 720 else Qt.Orientation.Horizontal)
    for control in (value.settings_toggle, value.saved_batches, value.save_button, value.close_button, value.status):
        assert control.isVisible() and value.rect().contains(QRect(control.mapTo(value, control.rect().topLeft()), control.size()))


def test_stale_chapter_batch_is_readable_but_cannot_be_adopted(dialog, app):
    value, bridge, f = dialog; generate(value, bridge); value.reject()
    old = f.study.book_study(f.selection)
    view = f.study.save_section(f.selection, old['study_revision'], section_id=old['sections'][0]['section_id'],
        title='新章名与范围', pdf_start=3, pdf_end=3, range_confirmed=True)
    tasks = Bridge(); reopened = TextbookKnowledgeDialog(facade_for(f), tasks, f.selection, view, view['sections'][0])
    reopened.show(); complete(tasks, '读取教材知识候选批次'); reopened.saved_batches.setCurrentIndex(1)
    assert reopened._draft_stale and not reopened.confirmed.isEnabled() and not reopened.save_button.isEnabled()
    assert '章节已变化' in reopened.saved_batches.currentText()
    reopened.read_button.click(); app.processEvents(); complete(tasks, '读取教材原页'); app.processEvents()
    assert reopened._reader.view.pageNavigator().currentPage() == 1
    assert reopened._reader.pages.count() == 3
    before = f.state.path.read_bytes(); reopened.reject(); assert f.state.path.read_bytes() == before
    reopened.deleteLater(); app.processEvents()


def test_new_batch_saves_current_edits_before_replacing_editor(dialog):
    value, bridge, f = dialog; generate(value, bridge)
    first_id = value._draft['draft_id']
    value.text.setPlainText('切换批次前保留的修订。')
    value.settings_toggle.setChecked(True); value.start.setValue(3); value.generate_button.click()
    complete(bridge, '保存教材知识核对进度')
    complete(bridge, '整理教材知识候选'); complete(bridge, '保存教材知识候选批次')
    assert value.items.count() == 1 and value._draft['draft_id'] != first_id
    rows = f.state.snapshot()['textbook_workspace']['knowledge_drafts']
    first = next(r for r in rows if r['draft_id'] == first_id)
    assert first['review_edits'][0]['text'] == '切换批次前保留的修订。'
    value.text.setPlainText('第二批切换前也保存。')
    index = next(i for i in range(1, value.saved_batches.count()) if value.saved_batches.itemData(i)['draft_id'] == first_id)
    value.saved_batches.setCurrentIndex(index); complete(bridge, '保存教材知识核对进度')
    assert value._draft['draft_id'] == first_id and value.text.toPlainText() == '切换批次前保留的修订。'
    assert not value.confirmed.isChecked() and not value.save_button.isEnabled()


def test_actual_review_to_current_lesson_preserves_manual_note_and_exports(window, tmp_path, monkeypatch):
    from docx import Document
    from pptx import Presentation
    from integrations.deeptutor_shchem_v1.desktop_workbench.lesson_design_dialog import LessonDesignDialog
    from integrations.deeptutor_shchem_v1.desktop_lesson_output import checked_file, FILES
    win, app = window; source = native_pdf(tmp_path / '合成知识流程教材.pdf')
    workspace = win.facade.textbook_workspace(); workspace.commit_books(workspace.preview_books([source]))
    study = win.facade.personal_textbook_study(); book = study.books()[0]
    selection = {k: book[k] for k in ('source_id', 'revision')}
    view = study.save_section(selection, study.book_study(selection)['study_revision'],
        title='合成备课章节', pdf_start=2, pdf_end=3, range_confirmed=True)
    view = study.save_note(selection, view['study_revision'], section_id=view['sections'][0]['section_id'],
        title='已有人工作品', text='保留原有人工笔记与条件Z。', pdf_page=2, teacher_confirmed=True)
    page = win.preparation_page
    page.topic.setText('教师原课题'); page.audience.setText('教师自定对象')
    page.objective.setPlainText('教师原目标'); page.materials.setPlainText('原有备课参考')
    frozen_images = deepcopy(page._payload().get('image_assets', []))
    win.navigate('textbooks'); settle(app, lambda: not win.textbook_page._loading)
    win.textbook_page.book.setCurrentIndex(win.textbook_page.book.findData(book['source_id']))
    personal = win.textbook_page.personal; settle(app, lambda: not personal._busy)
    personal.note_list.item(0).setCheckState(Qt.CheckState.Checked)
    def execute(dialog):
        dialog.show(); settle(app, lambda: not dialog._busy)
        dialog.generate_button.click(); settle(app, lambda: not dialog._busy and dialog.items.count() == 2)
        dialog.items.setCurrentRow(0)
        dialog.title.setText('用于本课的合成知识')
        dialog.text.setPlainText('A → B（条件X）；用量 2.0 g。')
        check_all(dialog); dialog.save_button.click(); settle(app, lambda: dialog._closed)
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(TextbookKnowledgeDialog, 'exec', execute)
    personal.knowledge_button.click(); settle(app)
    assert len(personal._checked) == 3 and study.book_study(selection)['notes'][0] == view['notes'][0]
    personal.preview_button.click(); settle(app, lambda: not personal._busy)
    assert '知识笔记' in personal.reference_preview.toPlainText()
    personal.confirm_button.click(); settle(app, lambda: not personal._busy)
    assert page.topic.text() == '教师原课题' and page.objective.toPlainText() == '教师原目标'
    assert page.materials.toPlainText().startswith('原有备课参考')
    assert '保留原有人工笔记' in page.materials.toPlainText() and 'A → B（条件X）；用量 2.0 g。' in page.materials.toPlainText()
    assert page._payload().get('image_assets', []) == frozen_images
    page._save(); settle(app, lambda: not page._save_task_id)
    saved = [win.facade.load_preparation_draft(r['draft_id'], r['revision']) for r in win.facade.preparation_draft_options()]
    assert any('教师修订的知识候选' in r['payload']['materials'] for r in saved)
    lesson = LessonDesignDialog(page); lesson.show(); settle(app)
    lesson.add_node(); lesson.bind_source()
    assert 'A → B' in lesson.history.value['nodes'][0]['material_text']
    lesson.locked.setChecked(False); lesson.student_material.setChecked(True)
    lesson.fields['material_text'].setPlainText('用于本课的合成知识\nA → B（条件X）；用量 2.0 g。\n合成知识流程教材.pdf · PDF第2页')
    lesson.fields['teacher_answer'].setPlainText('教师答案哨兵-知识候选')
    lesson.generate.click(); settle(app, lambda: not lesson.busy)
    assert lesson.history.value['exports'], lesson.status.text()
    identity = lesson.history.value['exports'][-1]['id']
    files = {name: checked_file(win.facade, lesson.history.value, identity, name) for name in FILES}
    presentation = Presentation(files['lesson_presentation.pptx'])
    slides = '\n'.join(s.text for p in presentation.slides for s in p.shapes if s.has_text_frame)
    notes = '\n'.join(p.notes_slide.notes_text_frame.text for p in presentation.slides)
    def body(path):
        doc = Document(path)
        return '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    teacher, student = body(files['lesson_plan.docx']), body(files['student_worksheet.docx'])
    assert all('A → B（条件X）；用量 2.0 g。' in text for text in (slides, teacher, student))
    assert '教师答案哨兵-知识候选' in teacher and '教师答案哨兵-知识候选' in notes
    assert '教师答案哨兵-知识候选' not in slides and '教师答案哨兵-知识候选' not in student
    assert source.read_bytes() == workspace.read_book(book['source_id'], book['source_sha256'])['pdf_bytes']
    lesson.reject(); lesson.deleteLater(); settle(app)
