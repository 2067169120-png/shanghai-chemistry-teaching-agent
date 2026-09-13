from __future__ import annotations

import hashlib
import io
import json
from copy import deepcopy
from zipfile import ZipFile

import pytest
from docx import Document
from PIL import Image

from integrations.deeptutor_shchem_v1 import paper_export_renderer as renderer
from integrations.deeptutor_shchem_v1 import paper_export_workbench as workbench
from integrations.deeptutor_shchem_v1 import paper_format_presets as formats
from integrations.deeptutor_shchem_v1.candidate_review import CandidateCropPayload
from staging.coordination.deeptutor_gateway.tests.test_paper_export_workbench_api import (
    _atomic,
    _capture_submitted_futures,
    _catalog,
    _details,
    _fake_render_factory,
    _request_payload,
)


def _png(color: str, size: tuple[int, int] = (200, 80)) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", size, color).save(output, format="PNG")
    return output.getvalue()


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _fixture():
    catalog = _catalog()
    theme = catalog["papers"][0]["theme_groups"][0]
    theme["atomic_chain"] = [_atomic(f"Q{n}", f"P{n}", n) for n in (2, 3, 5)]
    theme["counts"] = {"atomic": 3}
    details = {}
    answers = {}
    questions = {}
    for n, color in ((2, "red"), (3, "green"), (5, "blue")):
        node = f"Q{n}"
        raw = _png(color)
        answers[node] = raw
        questions[node] = _png(color, (210, 85))
        detail = deepcopy(_details()["A1"])
        detail["master_node_id"] = node
        detail["evidence_descriptors"] = [
            {"crop_id": f"QUESTION-{node}", "evidence_role": "question"}
        ]
        detail["reference_answer_images"] = [
            {
                "crop_id": f"ANSWER-{node}",
                "evidence_role": "answer",
                "sha256": _sha(raw),
                "source_sha256": "a" * 64,
                "source_page": 6,
                "bytes": len(raw),
                "width": 200,
                "height": 80,
                "access": "teacher_reference_answer_only",
                "display_mode": "preview_only" if n == 2 else "inline_required",
                "presentation_revision_id": "synthetic-answer-r1",
                "caption_zh": f"本题来源结构图 {n}",
            }
        ]
        details[node] = detail
    return catalog, details, answers, questions


def _resolver(tmp_path, details, answers, *, answer_loader=True):
    calls = []

    def load(node, crop):
        assert crop == f"ANSWER-{node}"
        calls.append((node, crop))
        raw = answers[node]
        return CandidateCropPayload(data=raw, sha256=_sha(raw))

    resolver = workbench._WorkbenchContentResolver(
        scope="master",
        rows=[{"atomic_part_id": node} for node in details],
        details=details,
        crop_loader=lambda node, crop: CandidateCropPayload(
            data=b"question", sha256=_sha(b"question")
        ),
        answer_crop_loader=load if answer_loader else None,
        asset_root=tmp_path,
        score_per_atomic=2,
        default_answer_lines=0,
    )
    return resolver, calls


def test_old_text_only_answers_keep_exact_contract(tmp_path):
    details = _details()
    resolver, calls = _resolver(tmp_path, details, {})
    answer = resolver.resolve_atomic_part({"atomic_part_id": "A1"})["reference_answer"]
    assert "content_blocks" not in answer
    assert formats._normalize_reference_answer(answer) == answer
    assert calls == []


@pytest.mark.parametrize(
    "field,value",
    [
        ("evidence_role", "question"),
        ("access", "teacher_loopback_read_only"),
        ("display_mode", "auto"),
        ("sha256", "b" * 64),
        ("source_sha256", "unknown"),
        ("bytes", 1),
        ("width", 1),
        ("source_page", 0),
        ("presentation_revision_id", ""),
    ],
)
def test_invalid_answer_image_descriptors_fail_closed(tmp_path, field, value):
    _, details, answers, _ = _fixture()
    details["Q3"]["reference_answer_images"][0][field] = value
    resolver, _ = _resolver(tmp_path, details, answers)
    with pytest.raises(workbench.PaperExportWorkbenchError):
        resolver.resolve_atomic_part({"atomic_part_id": "Q3"})


def test_required_image_never_falls_back_to_question_loader(tmp_path):
    _, details, answers, _ = _fixture()
    resolver, calls = _resolver(tmp_path, details, answers, answer_loader=False)
    # A preview-only image is optional for this export and must not be fetched.
    assert (
        "content_blocks"
        not in resolver.resolve_atomic_part({"atomic_part_id": "Q2"})[
            "reference_answer"
        ]
    )
    with pytest.raises(
        workbench.PaperExportWorkbenchError, match="专用读取接口"
    ) as caught:
        resolver.resolve_atomic_part({"atomic_part_id": "Q3"})
    assert caught.value.code == "paper_export_answer_image_loader_missing"
    assert calls == []


@pytest.mark.parametrize(
    "mutation",
    ["detail_identity", "crop_binding", "payload_sha", "payload_bytes", "duplicate"],
)
def test_cross_question_or_drifted_answer_images_fail(tmp_path, mutation):
    _, details, answers, _ = _fixture()
    resolver, _ = _resolver(tmp_path, details, answers)
    if mutation == "detail_identity":
        details["Q3"]["master_node_id"] = "Q5"
    elif mutation == "crop_binding":
        details["Q3"]["reference_answer_images"][0]["crop_id"] = "ANSWER-Q5"

        def strict_route(node, crop):
            raise workbench.PaperExportWorkbenchError(
                "answer_crop_not_bound", "跨题答案图", 404
            )

        resolver.answer_crop_loader = strict_route
    elif mutation == "payload_sha":
        resolver.answer_crop_loader = lambda node, crop: CandidateCropPayload(
            data=answers[node], sha256="b" * 64
        )
    elif mutation == "payload_bytes":
        resolver.answer_crop_loader = lambda node, crop: CandidateCropPayload(
            data=answers["Q5"], sha256=_sha(answers[node])
        )
    else:
        details["Q3"]["reference_answer_images"] *= 2
    with pytest.raises(workbench.PaperExportWorkbenchError):
        resolver.resolve_atomic_part({"atomic_part_id": "Q3"})


@pytest.mark.parametrize("status", ["absent", "unaligned"])
def test_unaligned_images_are_not_exported(tmp_path, status):
    _, details, answers, _ = _fixture()
    details["Q3"]["reference_answer"].update(
        availability=status, reference_answer_text=None
    )
    resolver, calls = _resolver(tmp_path, details, answers)
    with pytest.raises(workbench.PaperExportWorkbenchError, match="不能嵌入"):
        resolver.resolve_atomic_part({"atomic_part_id": "Q3"})
    assert calls == []


@pytest.mark.parametrize(
    "status,blocks",
    [
        (
            "absent",
            [
                {
                    "block_type": "image",
                    "asset_ref": "images/a.png",
                    "alt_text_zh": "答案图",
                }
            ],
        ),
        (
            "unaligned",
            [
                {
                    "block_type": "image",
                    "asset_ref": "images/a.png",
                    "alt_text_zh": "答案图",
                }
            ],
        ),
        ("aligned", []),
        ("aligned", [{"block_type": "text", "text_zh": "not an image"}]),
        ("aligned", [{"block_type": "image", "text_zh": "missing asset"}]),
    ],
)
def test_reference_answer_normalizer_rejects_invalid_image_blocks(status, blocks):
    value = {
        "status": status,
        "text_zh": "来源文字" if status == "aligned" else None,
        "authority_label": "nonofficial" if status == "aligned" else "none",
        "source_label_zh": "来源" if status == "aligned" else None,
        "independently_verified": False,
        "content_blocks": blocks,
    }
    with pytest.raises(formats.PaperFormatContractError):
        formats._normalize_reference_answer(value)


def test_start_routes_only_two_required_images_and_student_docx_has_none(
    tmp_path, monkeypatch
):
    catalog, details, answers, questions = _fixture()
    captured = {}
    calls = []
    manager = workbench.PaperExportJobManager(tmp_path)
    futures = _capture_submitted_futures(manager, monkeypatch)
    monkeypatch.setattr(
        workbench, "render_export_bundle", _fake_render_factory(captured)
    )
    monkeypatch.setattr(workbench, "_locate_toolchain", lambda: None)

    def question_loader(scope, node, crop):
        assert scope == "master" and crop == f"QUESTION-{node}"
        return CandidateCropPayload(data=questions[node], sha256=_sha(questions[node]))

    def answer_loader(scope, node, crop):
        assert scope == "master" and crop == f"ANSWER-{node}"
        calls.append((scope, node, crop))
        return CandidateCropPayload(data=answers[node], sha256=_sha(answers[node]))

    try:
        request = _request_payload()
        request["answer_space_lines"] = 0
        started = manager.start(
            request,
            theme_catalog_loader=lambda scope: catalog,
            detail_loader=lambda scope, node: details[node],
            crop_loader=question_loader,
            answer_crop_loader=answer_loader,
        )
        futures[0].result(timeout=10)
        completed = manager.get(started["job_id"])
        assert completed["status"] == "completed", completed
        assert calls == [("master", "Q3", "ANSWER-Q3"), ("master", "Q5", "ANSWER-Q5")]
        bundle = captured["bundle"]
        assert bundle["publication_allowed"] is False
        assert "teacher_notes" not in json.dumps(bundle["student_plan"])
        assets = tmp_path / "paper-export-workbench" / started["job_id"] / "assets"
        teacher_parts = [
            p
            for q in bundle["teacher_plan"]["visible"]["theme_sections"][0][
                "printed_questions"
            ]
            for p in q["atomic_parts"]
        ]
        assert (
            "content_blocks"
            not in teacher_parts[0]["teacher_notes"]["source_reference_answer"]
        )
        for part in teacher_parts[1:]:
            answer = part["teacher_notes"]["source_reference_answer"]
            assert answer["authority_label"] == "nonofficial"
            assert answer["independently_verified"] is False
            assert len(answer["content_blocks"]) == 1
            assert "supplemental_answer" not in part["teacher_notes"]
        for audience in ("student", "teacher"):
            path = tmp_path / f"{audience}.docx"
            renderer.build_docx_from_plan(
                bundle[f"{audience}_plan"],
                preset=bundle["preset"],
                output_path=path,
                asset_root=assets,
                blueprint=bundle["blueprint"],
            )
            with ZipFile(path) as archive:
                media = {
                    _sha(archive.read(name))
                    for name in archive.namelist()
                    if name.startswith("word/media/")
                }
                document = archive.read("word/document.xml").decode()
            assert {_sha(raw) for raw in questions.values()} <= media
            assert _sha(answers["Q2"]) not in media
            if audience == "student":
                assert not ({_sha(raw) for raw in answers.values()} & media)
                assert "非官方参考答案图" not in document
            else:
                assert {_sha(answers[node]) for node in ("Q3", "Q5")} <= media
                assert "非官方参考答案图" in document
                assert all(f"第{n}题评分" in document for n in (2, 3, 5))
                assert "2 分" in document
                answer_rows = [
                    row
                    for table in Document(path).tables
                    for row in table.rows
                    if row._tr.xpath(".//w:drawing")
                ]
                assert len(answer_rows) == 2
                assert all(
                    not row._tr.xpath("./w:trPr/w:cantSplit") for row in answer_rows
                )
                plain_rows = [
                    row
                    for table in Document(path).tables
                    for row in table.rows
                    if not row._tr.xpath(".//w:drawing")
                ]
                assert plain_rows
                assert all(
                    row.cells[0].paragraphs[0].paragraph_format.keep_with_next is False
                    for row in plain_rows
                )
    finally:
        manager.shutdown()


@pytest.mark.parametrize("failure", ["missing_loader", "wrong_sha", "wrong_role"])
def test_required_image_failure_never_reaches_renderer(tmp_path, monkeypatch, failure):
    catalog, details, answers, questions = _fixture()
    manager = workbench.PaperExportJobManager(tmp_path)
    futures = _capture_submitted_futures(manager, monkeypatch)
    captured = {}
    monkeypatch.setattr(
        workbench, "render_export_bundle", _fake_render_factory(captured)
    )
    monkeypatch.setattr(workbench, "_locate_toolchain", lambda: None)
    if failure == "wrong_role":
        details["Q3"]["reference_answer_images"][0]["evidence_role"] = "question"
    loader = (
        None
        if failure == "missing_loader"
        else lambda scope, node, crop: CandidateCropPayload(
            data=answers[node],
            sha256="b" * 64 if failure == "wrong_sha" else _sha(answers[node]),
        )
    )
    try:
        started = manager.start(
            _request_payload(),
            theme_catalog_loader=lambda scope: catalog,
            detail_loader=lambda scope, node: details[node],
            crop_loader=lambda scope, node, crop: CandidateCropPayload(
                data=questions[node], sha256=_sha(questions[node])
            ),
            answer_crop_loader=loader,
        )
        futures[0].result(timeout=10)
        job = manager.get(started["job_id"])
        assert job["status"] == "failed"
        assert job["error"]["code"].startswith("paper_export_answer_image_")
        assert job["artifacts"] == []
        assert captured == {}
    finally:
        manager.shutdown()
