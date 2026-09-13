from __future__ import annotations

import json
import re
import sys
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from jsonschema import Draft202012Validator
from PIL import Image

from integrations.deeptutor_shchem_v1 import paper_export_renderer as renderer


def test_external_renderer_commands_request_no_visible_windows(monkeypatch) -> None:
    observed: dict[str, object] = {}

    def fake_run(command, **kwargs):
        observed["command"] = list(command)
        observed.update(kwargs)
        return renderer.subprocess.CompletedProcess(command, 0, stdout="", stderr="")

    monkeypatch.setattr(renderer.subprocess, "run", fake_run)
    completed = renderer._run_command(["synthetic-render-tool"], timeout=1)
    assert completed.returncode == 0
    assert observed["command"] == ["synthetic-render-tool"]
    assert observed["creationflags"] == getattr(
        renderer.subprocess, "CREATE_NO_WINDOW", 0
    )


def _all_docx_text(path: Path) -> str:
    doc = Document(path)
    parts = [paragraph.text for paragraph in doc.paragraphs]

    def add_table(table) -> None:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
                for nested in cell.tables:
                    add_table(nested)

    for table in doc.tables:
        add_table(table)
    return "\n".join(parts)


def _toolchain() -> renderer.RendererToolchain:
    existing = Path(renderer.__file__).resolve()
    return renderer.RendererToolchain(
        python_exe=Path(sys.executable).resolve(),
        render_docx_script=existing,
        pdftoppm_exe=existing,
        dpi=300,
        conversion_backend="word_com",
    )


def _review_records(report: dict) -> list[dict]:
    page_paths = sorted(
        {
            page["path"]
            for artifact in report["artifacts"]
            if artifact["format"] == "docx"
            for key in ("docx_render_pages", "pdf_render_pages")
            for page in artifact[key]
        }
    )
    return [
        {
            "path": path,
            "status": "pass",
            "notes_zh": "页眉页脚、题号、分值与正文均无截断、重叠或溢出。",
        }
        for path in page_paths
    ]


def _patch_page_tools(monkeypatch: pytest.MonkeyPatch) -> None:
    def fake_docx(docx_path, output_dir, *, toolchain):
        del toolchain
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page-1.png"
        Image.new("RGB", (1240, 1754), "white").save(page)
        pdf = output_dir / f"{docx_path.stem}.pdf"
        pdf.write_bytes(b"%PDF-FAKE-DOCX-" + docx_path.name.encode("utf-8"))
        return [page], pdf, {
            "tool": "render_docx.py",
            "script_sha256": "0" * 64,
            "invocation_mode": "test-double",
            "conversion_backend": "word_com",
            "dpi": 300,
        }

    def fake_pdf(pdf_path, output_dir, *, toolchain):
        del pdf_path, toolchain
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page-1.png"
        Image.new("RGB", (1240, 1754), "white").save(page)
        return [page], {
            "tool": "pdftoppm",
            "executable_sha256": "1" * 64,
            "stderr_sha256": "2" * 64,
            "dpi": 300,
        }

    def fake_pdf_audit(path, *, plan, preset=None, blueprint=None):
        del path, plan, preset, blueprint
        return {
            "status": "pass",
            "checks": {
                "title_present": True,
                "shared_material_once": True,
                "secret_scan_pass": True,
            },
            "text_sha256": "3" * 64,
        }

    monkeypatch.setattr(renderer, "_render_with_canonical_docx_tool", fake_docx)
    monkeypatch.setattr(renderer, "_render_pdf_with_poppler", fake_pdf)
    monkeypatch.setattr(renderer, "_pdf_page_count", lambda path: 1)
    monkeypatch.setattr(renderer, "_pdf_text_audit", fake_pdf_audit)


def test_synthetic_bundle_is_theme_first_and_has_prior_dependency():
    bundle = renderer.build_synthetic_demo_bundle()

    assert bundle["preflight_report"]["status"] == "ready_for_renderer"
    assert bundle["publication_allowed"] is False
    assert bundle["blueprint"]["top_level_unit"] == "theme_big_question"
    assert bundle["blueprint"]["standalone_choice_section_allowed"] is False
    assert bundle["blueprint"]["counts"] == {
        "theme_count": 1,
        "printed_question_count": 3,
        "atomic_part_count": 3,
        "shared_material_count": 1,
    }
    atomic_parts = [
        atomic
        for printed in bundle["blueprint"]["theme_bundles"][0]["printed_questions"]
        for atomic in printed["atomic_parts"]
    ]
    assert [atomic["item_type"] for atomic in atomic_parts] == [
        "choice_single",
        "fill_blank",
        "short_answer",
    ]
    assert atomic_parts[2]["dependency"] == {
        "kind": "one_prior_part",
        "prior_atomic_part_ids": ["SYNTH-CO2-A2"],
        "status": "validated_explicit",
    }


def test_render_bundle_rejects_request_drift():
    bundle = renderer.build_synthetic_demo_bundle()
    tampered = deepcopy(bundle)
    tampered["render_request"]["render_request_digest"] = "0" * 64

    with pytest.raises(renderer.PaperExportRendererError) as exc_info:
        renderer.validate_render_bundle(tampered)

    assert exc_info.value.code == "render_request_drift"


def test_student_and_teacher_docx_share_layout_without_answer_leak(tmp_path: Path):
    bundle = renderer.build_synthetic_demo_bundle()
    student = tmp_path / "student.docx"
    teacher = tmp_path / "teacher.docx"
    renderer.build_docx_from_plan(
        bundle["student_plan"], preset=bundle["preset"], output_path=student
    )
    renderer.build_docx_from_plan(
        bundle["teacher_plan"], preset=bundle["preset"], output_path=teacher
    )

    student_audit = renderer.audit_docx(
        student, plan=bundle["student_plan"], preset=bundle["preset"]
    )
    teacher_audit = renderer.audit_docx(
        teacher, plan=bundle["teacher_plan"], preset=bundle["preset"]
    )
    assert student_audit["status"] == "pass"
    assert teacher_audit["status"] == "pass"

    student_text = _all_docx_text(student)
    teacher_text = _all_docx_text(teacher)
    shared_material = "某学习小组将二氧化碳通入澄清石灰水。通入少量二氧化碳时产生白色浑浊；继续通入过量二氧化碳后，浑浊逐渐消失。"
    assert student_text.count(shared_material) == 1
    assert teacher_text.count(shared_material) == 1
    assert "参考答案（非官方，未独立核验）" not in student_text
    assert "教师解析" not in student_text
    assert "易错点：" not in student_text
    assert "参考答案（非官方，未独立核验）" in teacher_text
    assert "易错点：" in teacher_text
    assert "来源标签：" in teacher_text
    assert sum(
        line.startswith("来源标签：") for line in teacher_text.splitlines()
    ) == 1
    assert "选择题" not in student_text
    assert "姓名：" in student_text and "班级：" in student_text
    assert "姓名：" not in teacher_text and "班级：" not in teacher_text

    student_document = Document(student)
    assert student_document.sections[0].different_first_page_header_footer
    assert not any(
        paragraph.text.strip()
        for paragraph in student_document.sections[0].first_page_header.paragraphs
    )

    with ZipFile(student) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        footer_xml = "".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/footer") and name.endswith(".xml")
        )
    with ZipFile(teacher) as archive:
        teacher_document_xml = archive.read("word/document.xml").decode("utf-8")
    assert document_xml.count("<w:tcBorders>") >= 8
    assert not re.search(
        r'<w:bottom\b[^>]*w:val="single"', teacher_document_xml
    )
    assert re.search(r"<w:instrText[^>]*>\s*PAGE\s*</w:instrText>", footer_xml)
    assert re.search(r"<w:instrText[^>]*>\s*NUMPAGES\s*</w:instrText>", footer_xml)


def test_docx_builder_honors_editable_layout_tokens(tmp_path: Path):
    bundle = renderer.build_synthetic_demo_bundle()
    preset = deepcopy(bundle["preset"])
    preset["page_layout"]["margins_mm"] = {
        "top": 21,
        "bottom": 19,
        "left": 25,
        "right": 22,
    }
    preset["page_layout"]["title"].update(
        {"size_pt": 18, "bold": False, "alignment": "left"}
    )
    preset["page_layout"]["theme_heading"].update(
        {"size_pt": 13, "bold": False}
    )
    preset["page_layout"]["body"].update(
        {"size_pt": 11, "line_spacing": 1.3}
    )
    output = tmp_path / "custom-layout.docx"

    renderer.build_docx_from_plan(
        bundle["student_plan"], preset=preset, output_path=output
    )
    audit = renderer.audit_docx(
        output, plan=bundle["student_plan"], preset=preset
    )

    assert audit["status"] == "pass"
    expected_width = round((210 - 25 - 22) / 25.4 * 1440)
    # Shared context is now inline prose/figures, not a forced full-width box.
    # Remaining answer-line tables may be narrower but cannot exceed the body.
    table_widths = [
        int(table._tbl.tblPr.find(renderer.qn("w:tblW")).get(renderer.qn("w:w")))
        for table in Document(output).tables
    ]
    assert table_widths and all(width <= expected_width for width in table_widths)


def test_source_numbered_images_keep_one_visible_number_and_compact_line_caps(
    tmp_path: Path,
):
    bundle = renderer.build_synthetic_demo_bundle()
    plan = deepcopy(bundle["student_plan"])
    asset_root = tmp_path / "assets"
    image_dir = asset_root / "images"
    image_dir.mkdir(parents=True)
    source_image = image_dir / "question.png"
    Image.new("RGB", (820, 120), "white").save(source_image)
    source_sha256 = renderer._sha256_file(source_image)

    atomics = [
        atomic
        for section in plan["visible"]["theme_sections"]
        for printed in section["printed_questions"]
        for atomic in printed["atomic_parts"]
    ]
    for index, atomic in enumerate(atomics, start=1):
        atomic["question_blocks"] = [
            {
                "block_type": "image",
                "text_zh": None,
                "asset_ref": "images/question.png",
                "alt_text_zh": f"源题第{index}题，裁片自带题号。",
            }
        ]
        atomic["answer_space"]["lines"] = 4

    output = tmp_path / "image-numbering.docx"
    renderer.build_docx_from_plan(
        plan,
        preset=bundle["preset"],
        output_path=output,
        asset_root=asset_root,
        blueprint=bundle["blueprint"],
    )

    assert renderer._sha256_file(source_image) == source_sha256
    text = re.sub(r"\s+", "", _all_docx_text(output))
    assert all(f"{number}.（" not in text for number in (1, 2, 3))
    assert all(f"第{number}题题图：" in "\n".join(renderer._docx_image_alt_texts(output)) for number in (1, 2, 3))

    doc = Document(output)
    score_paragraphs = [
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text.strip() in {"（2 分）", "（3 分）", "（5 分）"}
    ]
    assert len(score_paragraphs) == 3
    assert all(paragraph.alignment == 2 for paragraph in score_paragraphs)

    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert len(re.findall(r'<w:bottom\b[^>]*w:val="single"', document_xml)) == 3
    content_width = renderer._content_width_dxa(bundle["preset"])
    assert f'w:w="{round(content_width * 0.62)}"' in document_xml
    assert f'w:w="{round(content_width * 0.84)}"' in document_xml
    assert document_xml.count('w:left="113"') >= 3

    audit = renderer.audit_docx(output, plan=plan, preset=bundle["preset"])
    assert audit["status"] == "pass"
    assert audit["checks"]["source_image_numbers_not_repeated_as_text_headings"]


@pytest.mark.parametrize("summary,shown", [
    ("以卷面主题‘氯气’为共同语境；具体化学语义标签仍待人审。", False),
    ("以卷面主题‘氯气’为共同语境；具体化学语义标签仍待入库。", False),
    ("观察氯水的颜色变化，并解释消毒作用。", True),
])
def test_classroom_copy_keeps_real_context_but_not_internal_placeholders(tmp_path, summary, shown):
    bundle = renderer.build_synthetic_demo_bundle()
    plan = deepcopy(bundle["student_plan"])
    plan["visible"]["theme_sections"][0]["context_summary_zh"] = summary
    output = tmp_path / "student.docx"
    renderer.build_docx_from_plan(plan, preset=bundle["preset"], output_path=output)
    assert (summary in _all_docx_text(output)) is shown


def test_shared_question_answer_slots_keep_question_number_at_page_boundaries(tmp_path):
    bundle = renderer.build_synthetic_demo_bundle()
    plan = deepcopy(bundle["student_plan"])
    printed = plan["visible"]["theme_sections"][0]["printed_questions"][0]
    image_dir = tmp_path / "images"
    image_dir.mkdir()
    Image.new("RGB", (830, 287), "white").save(image_dir / "whole-question.png")
    atomic = printed["atomic_parts"][0]
    atomic["question_blocks"] = [{
        "block_type": "image", "text_zh": None,
        "asset_ref": "images/whole-question.png", "alt_text_zh": "整题与共同条件",
    }]
    printed["atomic_parts"].append(deepcopy(atomic))
    output = tmp_path / "student.docx"
    renderer.build_docx_from_plan(plan, preset=bundle["preset"], output_path=output, asset_root=tmp_path)
    text = _all_docx_text(output)
    assert "第1题（1）" in text and "第1题（2）" in text
    assert "作答单元" not in text
    assert len(Document(output).inline_shapes) == 1


def test_shared_image_order_padding_and_edge_warning_are_nonblocking(tmp_path: Path):
    bundle = renderer.build_synthetic_demo_bundle()
    plan = deepcopy(bundle["student_plan"])
    blueprint = deepcopy(bundle["blueprint"])
    asset_root = tmp_path / "assets"
    image_dir = asset_root / "images"
    image_dir.mkdir(parents=True)

    early = Image.new("RGB", (820, 80), "white")
    for x in range(120):
        early.putpixel((x, 0), (0, 0, 0))
    early.save(image_dir / "early.png")
    late = Image.new("RGB", (820, 80), "white")
    for x in range(90):
        late.putpixel((x, 79), (0, 0, 0))
    late.save(image_dir / "late.png")

    section = plan["visible"]["theme_sections"][0]
    section["shared_materials"] = [
        {
            "render_once_key": "late",
            "content_blocks": [
                {
                    "block_type": "image",
                    "text_zh": None,
                    "asset_ref": "images/late.png",
                    "alt_text_zh": "LATE-MATERIAL",
                }
            ],
        },
        {
            "render_once_key": "early",
            "content_blocks": [
                {
                    "block_type": "image",
                    "text_zh": None,
                    "asset_ref": "images/early.png",
                    "alt_text_zh": "EARLY-MATERIAL",
                }
            ],
        },
    ]
    blueprint["theme_bundles"][0]["shared_materials"] = [
        {"render_once_key": "late", "page": 2},
        {"render_once_key": "early", "page": 1},
    ]

    output = tmp_path / "shared-materials.docx"
    renderer.build_docx_from_plan(
        plan,
        preset=bundle["preset"],
        output_path=output,
        asset_root=asset_root,
        blueprint=blueprint,
    )
    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert document_xml.index("EARLY-MATERIAL") < document_xml.index("LATE-MATERIAL")
    assert document_xml.count('w:left="113"') >= 2

    notes = renderer._source_asset_edge_review_notes(plan, asset_root)
    assert len(notes) == 1
    assert "2 张源裁片中 2 张" in notes[0]["note_zh"]
    assert "上缘合计 120" in notes[0]["note_zh"]
    assert "下缘合计 90" in notes[0]["note_zh"]
    assert notes[0]["human_reviewed"] is False
    assert notes[0]["chemistry_reviewed"] is False


def _bound_material_fixture(tmp_path, *, audience="student"):
    bundle = renderer.build_synthetic_demo_bundle()
    plan = deepcopy(bundle[f"{audience}_plan"])
    blueprint = deepcopy(bundle["blueprint"])
    section = plan["visible"]["theme_sections"][0]
    source = blueprint["theme_bundles"][0]
    section["shared_materials"] = []
    source["shared_materials"] = []
    for index, printed in enumerate(section["printed_questions"]):
        for kind, height in (("M", 130), ("Q", 60)):
            asset = f"{kind}{index}.png"
            Image.new("RGB", (820, height), (245 - index, 250, 250)).save(tmp_path / asset)
        printed["atomic_parts"][0]["question_blocks"] = [{
            "block_type": "image", "text_zh": None, "asset_ref": f"Q{index}.png",
            "alt_text_zh": f"QUESTION-{index}",
        }]
        printed["atomic_parts"][0]["answer_space"] = {"mode": "ruled_lines_exact", "lines": 0}
        section["shared_materials"].append({
            "render_once_key": f"material-{index}",
            "content_blocks": [{
                "block_type": "image", "text_zh": None, "asset_ref": f"M{index}.png",
                "alt_text_zh": f"MATERIAL-{index}",
            }],
        })
        source["shared_materials"].append({
            "render_once_key": f"material-{index}", "page": 1,
            "used_by_atomic_count": 99,  # Counts are deliberately not placement evidence.
            "used_by_atomic_ids": [source["printed_questions"][index]["atomic_parts"][0]["atomic_part_id"]],
        })
    return bundle, plan, blueprint


def test_bound_materials_are_inline_at_first_use_and_do_not_keep_entire_theme(tmp_path):
    bundle, plan, blueprint = _bound_material_fixture(tmp_path)
    before = deepcopy((plan, blueprint))
    target = tmp_path / "flow.docx"
    renderer.build_docx_from_plan(
        plan, preset=bundle["preset"], output_path=target, asset_root=tmp_path,
        blueprint=blueprint, suppressed_shared_keys=set(),
    )
    document = Document(target)
    assert not document.tables  # No giant adjacent single-cell material boxes.
    assert len(document.inline_shapes) == 6
    xml = document._element.xml
    markers = [f"{kind}-{index}" for index in range(3) for kind in ("MATERIAL", "QUESTION")]
    assert [xml.index(marker) for marker in markers] == sorted(xml.index(marker) for marker in markers)
    assert xml.count("【共同材料】") == 3
    assert 'w:type="page"' not in xml and "w:pageBreakBefore" not in xml
    labels = [p for p in document.paragraphs if "【共同材料】" in p.text]
    assert all(p.paragraph_format.keep_with_next is True for p in labels)
    for paragraph in document.paragraphs:
        descriptions = paragraph._p.xpath(".//wp:docPr/@descr")
        if any("MATERIAL-" in text for text in descriptions):
            assert paragraph.paragraph_format.keep_with_next is True
        if any("QUESTION-" in text for text in descriptions):
            assert paragraph.paragraph_format.keep_with_next is False
    assert (plan, blueprint) == before  # Layout projection never rewrites frozen bindings.


def test_legacy_materials_without_member_ids_stay_at_theme_opening_without_long_keep_chain(tmp_path):
    bundle, plan, blueprint = _bound_material_fixture(tmp_path)
    for material in blueprint["theme_bundles"][0]["shared_materials"]:
        material.pop("used_by_atomic_ids")
    target = tmp_path / "legacy-flow.docx"
    renderer.build_docx_from_plan(
        plan, preset=bundle["preset"], output_path=target, asset_root=tmp_path,
        blueprint=blueprint, suppressed_shared_keys=set(),
    )
    document = Document(target)
    xml = document._element.xml
    assert max(xml.index(f"MATERIAL-{index}") for index in range(3)) < xml.index("QUESTION-0")
    images = [p for p in document.paragraphs if p._p.xpath(".//wp:docPr")]
    assert [p.paragraph_format.keep_with_next for p in images[:3]] == [False, False, True]


@pytest.mark.parametrize("members", [[], ["unknown"], ["SYNTH-CO2-A1", "SYNTH-CO2-A1"], "SYNTH-CO2-A1"])
def test_invalid_material_membership_is_not_guessed_or_dropped(tmp_path, members):
    _, plan, blueprint = _bound_material_fixture(tmp_path)
    blueprint["theme_bundles"][0]["shared_materials"][0]["used_by_atomic_ids"] = members
    with pytest.raises(renderer.PaperExportRendererError) as caught:
        renderer._shared_material_schedule(plan["visible"]["theme_sections"][0], blueprint["theme_bundles"][0])
    assert caught.value.code == "shared_material_membership_invalid"


def test_duplicate_legacy_material_key_is_rejected_without_silent_dedup(tmp_path):
    _, plan, blueprint = _bound_material_fixture(tmp_path)
    section = plan["visible"]["theme_sections"][0]
    section["shared_materials"].append(deepcopy(section["shared_materials"][0]))
    with pytest.raises(renderer.PaperExportRendererError) as caught:
        renderer._shared_material_schedule(section, blueprint["theme_bundles"][0])
    assert caught.value.code == "shared_material_duplicate"


@pytest.mark.parametrize("keys", [["other"], ["visible", "visible"]])
def test_single_material_must_match_blueprint_key_instead_of_falling_back(keys):
    section = {"shared_materials": [{"render_once_key": "visible"}], "printed_questions": [{}, {}]}
    source = {
        "shared_materials": [{"render_once_key": key, "used_by_atomic_ids": ["A2"]} for key in keys],
        "printed_questions": [
            {"atomic_parts": [{"atomic_part_id": "A1"}]},
            {"atomic_parts": [{"atomic_part_id": "A2"}]},
        ],
    }
    with pytest.raises(renderer.PaperExportRendererError) as caught:
        renderer._shared_material_schedule(section, source)
    assert caught.value.code == "render_hint_alignment_invalid"


@pytest.mark.parametrize(
    ("title", "sequence", "expected"),
    [
        ("五 有机合成药物", 5, "二、有机合成药物"),
        ("五、 有机合成药物", 5, "二、有机合成药物"),
        ("五 有机合成药物", None, "二、五 有机合成药物"),
        ("五 有机合成药物", 4, "二、五 有机合成药物"),
        ("1,2-二氯乙烷", 5, "二、1,2-二氯乙烷"),
        ("五氯化磷", 5, "二、五氯化磷"),
    ],
)
def test_theme_display_strips_only_bound_source_ordinal(title, sequence, expected):
    section = {"theme_number": 2, "heading_zh": f"二、{title}"}
    source = {"source": {"theme_title": title, "source_theme_sequence": sequence}}
    assert renderer._display_theme_heading(section, source) == expected
    assert source["source"]["theme_title"] == title


def test_source_numbered_teacher_scores_follow_original_labels_and_keep_global_anchors(tmp_path):
    bundle, plan, blueprint = _bound_material_fixture(tmp_path, audience="teacher")
    section = plan["visible"]["theme_sections"][0]
    source = blueprint["theme_bundles"][0]
    for index, (printed, original) in enumerate(zip(section["printed_questions"], source["printed_questions"], strict=True)):
        printed["question_number"] = 10 + index
        original["source_number"] = ("1", "2a", None)[index]
    target = tmp_path / "source-scores.docx"
    renderer.build_docx_from_plan(
        plan, preset=bundle["preset"], output_path=target, asset_root=tmp_path,
        blueprint=blueprint, suppressed_shared_keys=set(),
    )
    text = _all_docx_text(target)
    assert "第1题评分" in text and "第2a题评分" in text and "第12题评分" in text
    assert "第10题评分" not in text and "第11题评分" not in text
    assert renderer._score_labels_present(text, plan, True, blueprint)
    assert not renderer._score_labels_present(text.replace("第2a题评分", "第11题评分"), plan, True, blueprint)
    alts = renderer._docx_image_alt_texts(target)
    assert any("第10题题图" in alt for alt in alts)
    assert section["printed_questions"][0]["question_number"] == 10


def test_tall_inline_question_is_scaled_whole_not_cropped_or_split(tmp_path):
    Image.new("RGB", (300, 2000), "white").save(tmp_path / "tall.png")
    document = Document()
    renderer._add_asset_block(document, {"asset_ref": "tall.png"}, tmp_path)
    assert len(document.inline_shapes) == 1
    shape = document.inline_shapes[0]
    assert shape.height.mm <= 210.01
    assert shape.width / shape.height == pytest.approx(300 / 2000, abs=0.001)
    assert document.paragraphs[-1].paragraph_format.keep_together is True


def test_shared_question_visual_dedup_is_exact_and_keeps_merely_similar_images(
    tmp_path: Path,
):
    bundle = renderer.build_synthetic_demo_bundle()
    plan = deepcopy(bundle["student_plan"])
    blueprint = deepcopy(bundle["blueprint"])
    asset_root = tmp_path / "assets"
    image_dir = asset_root / "images"
    image_dir.mkdir(parents=True)

    def patterned(offset: int) -> Image.Image:
        image = Image.new("RGB", (40, 20), "white")
        for y in range(20):
            color = (
                (offset + y * 3) % 256,
                (offset * 2 + y * 5) % 256,
                (offset * 3 + y * 7) % 256,
            )
            image.paste(color, (0, y, 40, y + 1))
        return image

    same = patterned(11)
    same.save(image_dir / "same.png")
    sub = patterned(73)
    sub.save(image_dir / "sub.png")
    container = Image.new("RGB", (40, 24), "white")
    container.paste(sub, (0, 2))
    container.save(image_dir / "container.png")
    keep = patterned(149)
    keep.save(image_dir / "keep.png")
    similar = keep.copy()
    similar.paste((0, 0, 0), (0, 17, 40, 20))
    similar.save(image_dir / "similar.png")

    section = plan["visible"]["theme_sections"][0]
    section["shared_materials"] = [
        {
            "render_once_key": "same-key",
            "content_blocks": [
                {
                    "block_type": "image",
                    "text_zh": None,
                    "asset_ref": "images/same.png",
                    "alt_text_zh": "SAME-SHARED",
                }
            ],
        },
        {
            "render_once_key": "sub-key",
            "content_blocks": [
                {
                    "block_type": "image",
                    "text_zh": None,
                    "asset_ref": "images/sub.png",
                    "alt_text_zh": "SUB-SHARED",
                }
            ],
        },
        {
            "render_once_key": "keep-key",
            "content_blocks": [
                {
                    "block_type": "image",
                    "text_zh": None,
                    "asset_ref": "images/keep.png",
                    "alt_text_zh": "KEEP-SHARED",
                }
            ],
        },
    ]
    blueprint["theme_bundles"][0]["shared_materials"] = [
        {"render_once_key": "same-key", "page": 1},
        {"render_once_key": "sub-key", "page": 1},
        {"render_once_key": "keep-key", "page": 1},
    ]
    question_refs = ["same.png", "container.png", "similar.png"]
    for index, printed in enumerate(section["printed_questions"]):
        printed["atomic_parts"][0]["question_blocks"] = [
            {
                "block_type": "image",
                "text_zh": None,
                "asset_ref": f"images/{question_refs[index]}",
                "alt_text_zh": f"QUESTION-{index + 1}",
            }
        ]

    dedup = renderer._shared_question_visual_dedup(plan, asset_root)
    assert dedup["suppressed_shared_keys"] == frozenset(
        {"same-key", "sub-key"}
    )
    methods = {
        match["method"]
        for record in dedup["records"]
        for match in record["matches"]
    }
    assert "same_sha256" in methods
    assert "exact_subimage" in methods
    notes = renderer._shared_question_dedup_review_notes(dedup)
    assert len(notes) == 1
    assert "共抑制 2 个" in notes[0]["note_zh"]

    output = tmp_path / "dedup.docx"
    renderer.build_docx_from_plan(
        plan,
        preset=bundle["preset"],
        output_path=output,
        asset_root=asset_root,
        blueprint=blueprint,
    )
    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "SAME-SHARED" not in document_xml
    assert "SUB-SHARED" not in document_xml
    assert "KEEP-SHARED" in document_xml
    assert all(f"QUESTION-{number}" in document_xml for number in (1, 2, 3))


def test_missing_toolchain_fails_closed(tmp_path: Path):
    toolchain = renderer.RendererToolchain(
        python_exe=tmp_path / "missing-python.exe",
        render_docx_script=tmp_path / "missing-render.py",
        pdftoppm_exe=tmp_path / "missing-pdftoppm.exe",
    )

    with pytest.raises(renderer.PaperExportRendererError) as exc_info:
        toolchain.validated()

    assert exc_info.value.code == "renderer_tool_missing"


def test_renderer_rejects_dpi_below_frozen_preset(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    _patch_page_tools(monkeypatch)
    existing = Path(renderer.__file__).resolve()
    low_dpi = renderer.RendererToolchain(
        python_exe=Path(sys.executable).resolve(),
        render_docx_script=existing,
        pdftoppm_exe=existing,
        dpi=150,
        conversion_backend="word_com",
    )

    with pytest.raises(renderer.PaperExportRendererError) as exc_info:
        renderer.render_export_bundle(
            renderer.build_synthetic_demo_bundle(),
            output_dir=tmp_path,
            toolchain=low_dpi,
        )

    assert exc_info.value.code == "renderer_dpi_below_preset"


def test_render_stays_pending_until_every_page_is_visually_reviewed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    _patch_page_tools(monkeypatch)
    report = renderer.render_export_bundle(
        renderer.build_synthetic_demo_bundle(),
        output_dir=tmp_path,
        toolchain=_toolchain(),
    )

    assert report["status"] == "awaiting_visual_review"
    assert report["visual_review"]["status"] == "pending"
    assert report["receipts"] == []
    assert {item["artifact_id"] for item in report["artifacts"]} == {
        "student_docx",
        "student_pdf",
        "teacher_docx",
        "teacher_pdf",
    }
    assert all((tmp_path / name).is_file() for name in renderer.ARTIFACT_FILENAMES.values())

    one_page = report["artifacts"][0]["docx_render_pages"][0]["path"]
    with pytest.raises(renderer.PaperExportRendererError) as exc_info:
        renderer.finalize_visual_review(
            output_dir=tmp_path,
            reviewed_pages=[
                {"path": one_page, "status": "pass", "notes_zh": "本页无截断。"}
            ],
            reviewed_at="2026-08-27T12:00:00+00:00",
        )
    assert exc_info.value.code == "visual_review_incomplete"


def test_complete_visual_review_builds_contract_receipts_and_hash_manifest(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    _patch_page_tools(monkeypatch)
    pending = renderer.render_export_bundle(
        renderer.build_synthetic_demo_bundle(),
        output_dir=tmp_path,
        toolchain=_toolchain(),
    )
    reviews = _review_records(pending)
    additional_note = {
        "reviewer_kind": "delegating_root_task_independent_visual_review",
        "note_zh": "上游任务独立查看学生版与教师版页面，未见答案泄漏、断题或溢出。",
        "scope_zh": "额外版面复核，不替代化学人工审校。",
        "human_reviewed": False,
        "chemistry_reviewed": False,
    }

    final = renderer.finalize_visual_review(
        output_dir=tmp_path,
        reviewed_pages=reviews,
        reviewed_at="2026-08-27T12:00:00+00:00",
        additional_review_notes=[additional_note],
    )

    assert final["status"] == "layout_smoke_passed"
    assert final["visual_review"]["layout_defect_count"] == 0
    assert final["visual_review"]["human_reviewed"] is False
    assert final["visual_review"]["chemistry_reviewed"] is False
    assert final["visual_review"]["additional_review_notes"] == [additional_note]
    assert len(final["receipts"]) == 4
    assert final["contract_validation_report"]["status"] == "local_delivery_candidate"
    assert final["contract_validation_report"]["publication_allowed"] is False

    manifest = json.loads((tmp_path / "hash_manifest.json").read_text(encoding="utf-8"))
    assert manifest["file_count"] == len(manifest["files"])
    for item in manifest["files"]:
        path = tmp_path / item["path"]
        assert path.is_file()
        assert renderer._sha256_file(path) == item["sha256"]


def test_finalize_rejects_artifact_hash_drift(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    _patch_page_tools(monkeypatch)
    pending = renderer.render_export_bundle(
        renderer.build_synthetic_demo_bundle(),
        output_dir=tmp_path,
        toolchain=_toolchain(),
    )
    student_docx = tmp_path / renderer.ARTIFACT_FILENAMES["student_docx"]
    student_docx.write_bytes(student_docx.read_bytes() + b"tampered")

    with pytest.raises(renderer.PaperExportRendererError) as exc_info:
        renderer.finalize_visual_review(
            output_dir=tmp_path,
            reviewed_pages=_review_records(pending),
            reviewed_at="2026-08-27T12:00:00+00:00",
        )

    assert exc_info.value.code == "pending_artifact_hash_drift"


def test_finalize_rejects_page_hash_drift(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    _patch_page_tools(monkeypatch)
    pending = renderer.render_export_bundle(
        renderer.build_synthetic_demo_bundle(),
        output_dir=tmp_path,
        toolchain=_toolchain(),
    )
    page_path = tmp_path / pending["artifacts"][0]["docx_render_pages"][0]["path"]
    Image.new("RGB", (1240, 1754), "black").save(page_path)

    with pytest.raises(renderer.PaperExportRendererError) as exc_info:
        renderer.finalize_visual_review(
            output_dir=tmp_path,
            reviewed_pages=_review_records(pending),
            reviewed_at="2026-08-27T12:00:00+00:00",
        )

    assert exc_info.value.code == "pending_page_hash_drift"


def test_renderer_schema_accepts_bundle_pending_report_and_manifest(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
):
    schema_path = (
        Path(renderer.__file__).resolve().parent
        / "schemas"
        / "paper_export_renderer_v1.schema.json"
    )
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    Draft202012Validator.check_schema(schema)
    validator = Draft202012Validator(schema)
    bundle = renderer.build_synthetic_demo_bundle()
    validator.validate(bundle)

    _patch_page_tools(monkeypatch)
    report = renderer.render_export_bundle(
        bundle,
        output_dir=tmp_path,
        toolchain=_toolchain(),
    )
    validator.validate(report)
    final = renderer.finalize_visual_review(
        output_dir=tmp_path,
        reviewed_pages=_review_records(report),
        reviewed_at="2026-08-27T12:00:00+00:00",
    )
    manifest = json.loads((tmp_path / "hash_manifest.json").read_text(encoding="utf-8"))
    validator.validate(final)
    validator.validate(manifest)
