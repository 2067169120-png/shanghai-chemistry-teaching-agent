import importlib.util
import json
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[4]
spec = importlib.util.spec_from_file_location("local_material_import", ROOT / "runtime/deeptutor_shchem/import_local_materials.py")
cli = importlib.util.module_from_spec(spec)
spec.loader.exec_module(cli)


def test_question_bank_cli_previews_then_imports_and_reuses_without_changing_source(tmp_path, capsys):
    folder = tmp_path / "materials"
    folder.mkdir()
    document = Document()
    document.add_paragraph("【例1】合成回归材料：选择标记A。")
    document.add_paragraph("A. A B. B")
    document.add_paragraph("【答案】A")
    source = folder / "练习.docx"
    document.save(source)
    original = source.read_bytes()
    state = tmp_path / "state"
    arguments = ["--kind", "question-bank", "--folder", str(folder), "--state-root", str(state), "--workspace", str(ROOT)]
    assert cli.main(arguments) == 0
    preview = json.loads(capsys.readouterr().out)
    assert preview["applied"] is False and preview["model_calls"] == 0
    assert not list((state / "tasks").rglob("manifest.json"))
    assert cli.main([*arguments, "--apply"]) == 0
    imported = json.loads(capsys.readouterr().out)
    assert imported["batches"][0]["native_quick_count"] >= 1
    assert cli.main([*arguments, "--apply"]) == 0
    reused = json.loads(capsys.readouterr().out)
    assert reused["batches"][0]["batch_id"] == imported["batches"][0]["batch_id"]
    assert source.read_bytes() == original
