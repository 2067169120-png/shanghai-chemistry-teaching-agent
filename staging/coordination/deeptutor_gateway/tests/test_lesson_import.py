from copy import deepcopy
from types import SimpleNamespace
from pathlib import Path
import json
import pytest
from test_desktop_preparation import _payload, _raw_candidate
from test_lesson_design import sample
from integrations.deeptutor_shchem_v1.desktop_preparation import normalize_preparation_candidate
from integrations.deeptutor_shchem_v1.desktop_lesson_design import new_design, coverage, update_node, DesignHistory
from integrations.deeptutor_shchem_v1.desktop_lesson_import import prepare_import, apply_import
from integrations.deeptutor_shchem_v1.desktop_lesson_output import export_design, checked_file, FILES


def source():
    return {'task_id':'PREP-'+'a'*32, 'source_revision':'b'*64,
            'candidate':normalize_preparation_candidate(_raw_candidate(), _payload())}


def imported(plan=None):
    p=plan or new_design(_payload());s=source();proposal=prepare_import(p,s)
    selected=[o['id'] for o in proposal['options']]
    return apply_import(p,proposal,selected)[0]


def test_source_and_existing_teacher_nodes_do_not_change():
    p=sample()['lesson_design'];s=source();before=deepcopy((p,s))
    proposal=prepare_import(p,s);out,_=apply_import(p,proposal,[proposal['options'][0]['id']])
    assert (p,s)==before and out['nodes'][0]==p['nodes'][0]
    assert len(out['nodes'])==2 and out['exports']==p['exports']


def test_stage_activity_evidence_and_exact_material_are_connected():
    s=source();p=new_design(_payload());proposal=prepare_import(p,s)
    n=proposal['options'][0]['node'];c=s['candidate']
    assert n['teacher_action']==c['lesson_stages'][0]['teacher_action']
    assert n['student_task']==c['lesson_stages'][0]['student_action']
    assert c['assessments'][0]['evidence_of_learning'] in n['expected_output']
    assert c['assessments'][0]['success_criteria'][0] in n['criteria']
    assert '教师提供的反应实例' in n['material_text']
    assert c['slides'][1]['content'][0] in n['material_text']
    assert s['task_id'] in n['notes'] and s['source_revision'] in n['notes']
    assert '课后安排' in n['notes'] and '待核对' in n['notes']


def test_no_automatic_confirmation_or_student_material_exposure():
    out=imported()
    assert all(not n['confirmed'] and n['locked'] and not n['student_material'] for n in out['nodes'])
    assert all(g['status']!='已明确关联' for g in coverage(out)['objectives'])
    for n in out['nodes']:
        with pytest.raises(ValueError):update_node(out,n['id'],{'material_text':'change','locked':False})


def test_only_selected_nodes_and_required_goals_added():
    p=new_design({'objective':''});proposal=prepare_import(p,source())
    out,_=apply_import(p,proposal,[proposal['options'][0]['id']])
    assert len(out['nodes'])==1 and len(out['objectives'])==1


def test_exact_goal_text_reuses_id_without_fuzzy_matching():
    p=new_design({'objective':source()['candidate']['objectives'][0]['statement']})
    first=p['objectives'][0]['id'];out=imported(p)
    assert out['nodes'][0]['objective_ids']==[first] and len(out['objectives'])==2


def test_reimport_does_not_overwrite_or_duplicate():
    out=imported();before=deepcopy(out);proposal=prepare_import(out,source())
    assert all(o['already_present'] for o in proposal['options'])
    with pytest.raises(ValueError,match='已经加入'):apply_import(out,proposal,[proposal['options'][0]['id']])
    assert out==before


def test_preview_is_invalid_after_another_edit():
    p=new_design(_payload());proposal=prepare_import(p,source());p['objectives'][0]['text']='Changed'
    with pytest.raises(ValueError,match='已变化'):apply_import(p,proposal,[proposal['options'][0]['id']])


def test_unbound_slides_are_not_silently_dropped_or_turned_into_student_tasks():
    proposal=prepare_import(new_design(_payload()),source())
    orphan=proposal['options'][-1]
    assert orphan['node']['title']=='课题与目标'
    assert orphan['node']['student_task']=='' and orphan['warnings']
    assert '识别角色' in orphan['node']['material_text']


def test_unassigned_activity_is_retained():
    s=source();s['candidate']['lesson_stages']=[]
    proposal=prepare_import(new_design(_payload()),s)
    assert proposal['options'][0]['node']['title']=='证据辨析'
    assert proposal['options'][0]['node']['student_task']=='标注化合价并说明判断依据。'


def test_complex_visual_reference_stays_teacher_only_with_warning():
    s=source();s['candidate']['slides'][1]['visual']={'kind':'comparison','table':{'values':['12.5','Fe³⁺']}}
    proposal=prepare_import(new_design(_payload()),s)
    assert '12.5' in proposal['options'][0]['node']['notes']
    assert proposal['options'][0]['warnings']


def test_dangling_objective_and_assets_are_not_guessed():
    s=source();s['candidate']['lesson_stages'][0]['objective_ids']=['missing']
    with pytest.raises(ValueError,match='不存在'):prepare_import(new_design(_payload()),s)
    s=source();s['candidate']['slides'][1]['image']={'asset_id':'IMG-'+'c'*64}
    with pytest.raises(ValueError,match='图片'):prepare_import(new_design(_payload()),s)


def test_import_can_be_undone_and_reapplied_without_old_file_changes():
    p=new_design(_payload());history=DesignHistory(p);out=imported(p);history.put(out)
    assert history.travel()==p and history.travel(True)==out


def test_imported_design_exports_three_files_and_keeps_teacher_notes_off_student_pages(tmp_path):
    from docx import Document
    from pptx import Presentation
    p=_payload();p['lesson_design']=imported()
    f=SimpleNamespace(paths=SimpleNamespace(task_root=tmp_path))
    out=export_design(f,p);p['lesson_design']['exports'].append(out)
    student=Document(checked_file(f,p['lesson_design'],out['id'],FILES[2]))
    text='\n'.join(r.text for r in student.paragraphs)
    assert '完成标注、交流与修正。' in text and '来源任务' not in text
    assert '教师提供的反应实例' not in text
    ppt=Presentation(checked_file(f,p['lesson_design'],out['id'],FILES[0]))
    screen='\n'.join(sh.text for sl in ppt.slides for sh in sl.shapes if sh.has_text_frame)
    assert '来源任务' not in screen
    assert all(checked_file(f,p['lesson_design'],out['id'],n).is_file() for n in FILES)


def test_no_source_or_unknown_source_rejected():
    with pytest.raises(ValueError):prepare_import(new_design(_payload()),{})
    s=source();s['candidate']['schema_version']='future'
    with pytest.raises(ValueError):prepare_import(new_design(_payload()),s)
