"""EQ occurrence IDs normalize only as typed references during atomic handoff."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Mm, Pt
import pytest

from test_handout_material_handoff import FakeProviderStore, _facade, _png, images
from test_handout_material_handoff import desktop_paths as desktop_paths
from test_handout_source_context import _inventory
from test_static_eq_boundaries import _field, _run
from integrations.deeptutor_shchem_v1 import desktop_handout_material_handoff as handoff
from integrations.deeptutor_shchem_v1.desktop_handout_knowledge import handout_knowledge_sources, handout_knowledge_outline
from integrations.deeptutor_shchem_v1.desktop_handout_sources import handout_source_context
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import validate_resource
from integrations.deeptutor_shchem_v1.reader_cancellation import TaskCancellationEvent, ReadCancelled, read_cancel_scope


@pytest.fixture
def eq_handoff(desktop_paths, tmp_path, monkeypatch):
    monkeypatch.setattr('integrations.deeptutor_shchem_v1.desktop_word_questions.load_attribute_catalog',
                        lambda *a, **k: {'knowledge_points': [], 'nodes': []})
    document = Document(); document.add_heading('完整静态条件讲义', 1)
    document.add_heading('知识点1 聚合反应条件', 2)
    paragraph = document.add_paragraph('反应条件：')
    paragraph._p.extend(_field(
        _run(instruction=r' EQ \o(――→,\s\up7(H', font_pt=12),
        _run(instruction='＋', script='superscript', font_pt=12),
        _run(instruction=r'),\s\do5(△)) ', font_pt=12),
    ))
    paragraph.add_run('；保留完整原条件。')
    native_png = _png('blue')
    paragraph = document.add_paragraph('图片条件：')
    paragraph._p.extend([_run(boundary='begin'), _run(instruction=r' EQ \o(', font_pt=12)])
    run = paragraph.add_run(); run.font.size = Pt(12)
    picture = run.add_picture(BytesIO(native_png), width=Mm(18), height=Mm(4))
    fill = picture._inline.find('.//{http://schemas.openxmlformats.org/drawingml/2006/picture}blipFill')
    crop = OxmlElement('a:srcRect')
    for key, value in {'l': 10000, 't': 5000, 'r': 20000, 'b': 5000}.items():
        crop.set(key, str(value))
    fill.insert(1, crop)
    paragraph._p.extend([
        _run(instruction=r',\s\up7(催化剂)) ', font_pt=12),
        _run(boundary='separate'), _run(boundary='end'),
    ])
    document.add_paragraph('同字节原图的独立出现：').add_run().add_picture(BytesIO(native_png), width=Mm(18))
    document.add_paragraph('字面 CB 标识保持原文。')
    path = tmp_path / '静态条件交接讲义.docx'; document.save(path)
    provider = FakeProviderStore(configured=False); facade = _facade(desktop_paths, provider)
    batch = facade.save_visual_import_batch(handout_files=(path,), source_type='教师讲义')
    metadata = next(row for row in handout_knowledge_sources(facade)['sources'] if row['batch_id'] == batch.batch_id)
    outline = handout_knowledge_outline(facade, metadata)
    source = outline['source']; context = handout_source_context(facade, source)
    section = next(row for row in outline['sections'] if row['title'].startswith('知识点'))
    knowledge = {'source': deepcopy(source), 'outline_revision': outline['outline_revision'],
                 'section': deepcopy(section), 'title': '这些条件怎样保留？', 'teacher_only': False}
    request = {'schema': handoff.SCHEMA, 'source': source, 'knowledge_requests': [knowledge],
               'image_selections': [], 'question_selections': []}
    archive = Path(source['source_path'])
    return SimpleNamespace(facade=facade, request=request, provider=provider, context=context,
                           archive=archive, raw=archive.read_bytes(), png=native_png)


def _image_field(result, location='work'):
    resource = result['resources'][0]
    if location == 'snapshot':
        spans = [span for row in resource['original']['evidence']['section_snapshot']['blocks']
                 for span in row.get('static_eqs', [])]
    else:
        rows = resource['original']['blocks'] if location == 'original' else resource['blocks']
        spans = [span for row in rows for span in row.get('presentation', {}).get('eqs', [])]
    return next(span['field'] for span in spans if span['field']['base']['kind'] == 'image')


def _independent_picture(result):
    field = _image_field(result)
    pictures = [row for row in result['resources'][0]['original']['blocks'] if row['kind'] == 'image']
    bound = next(row for row in pictures if row['id'] == field['base']['image_block_id'])
    return next(row for row in pictures if row['id'] != bound['id'] and row['value'] == bound['value'])


def _plain_row(result, location):
    resource = result['resources'][0]
    if location == 'snapshot':
        return next(row for row in resource['original']['evidence']['section_snapshot']['blocks'] if row['text'].startswith('字面 CB'))
    rows = resource['original']['blocks'] if location == 'original' else resource['blocks']
    return next(row for row in rows if row['kind'] == 'text' and row['value'].startswith('字面 CB'))


def test_fresh_native_compiles_have_equal_semantic_stamps_without_mutating_full_snapshots(eq_handoff):
    case = eq_handoff
    state = case.facade.state_store.snapshot(); before = _inventory(case.facade.paths.state_root)
    first, payloads = handoff._compile(case.facade, case.request, [], [])
    second, repeated_payloads = handoff._compile(case.facade, case.request, [], [])
    original_first, original_second = deepcopy(first), deepcopy(second)
    assert first['resources'][0]['id'] != second['resources'][0]['id']
    assert _image_field(first)['base']['image_block_id'] != _image_field(second)['base']['image_block_id']
    assert handoff._stamp(first) == handoff._stamp(second)
    assert first == original_first and second == original_second
    assert payloads == repeated_payloads == {sha256(case.png).hexdigest(): case.png}
    item = first['resources'][0]
    assert validate_resource(item) == item
    snapshot = item['original']['evidence']['section_snapshot']
    source_rows = case.context['preview']['blocks'][snapshot['section']['start'] - 1:snapshot['section']['end']]
    assert [(row['index'], row['text']) for row in snapshot['blocks']] == [(row['index'], row['text']) for row in source_rows]
    assert sum(len(row.get('static_eqs', [])) for row in snapshot['blocks']) == 2
    assert all(_image_field(first, location) == _image_field(first) for location in ('original', 'snapshot'))
    assert _independent_picture(first)['id'] != _image_field(first)['base']['image_block_id']
    assert _inventory(case.facade.paths.state_root) == before and not images(case.facade)
    assert case.facade.state_store.snapshot() == state and case.provider.borrow_calls == 0
    assert case.archive.read_bytes() == case.raw


def test_prepare_publish_accept_keeps_source_and_state_and_commits_only_expected_image(eq_handoff):
    case = eq_handoff
    state = case.facade.state_store.snapshot(); before = _inventory(case.facade.paths.state_root)
    request = deepcopy(case.request)
    publication = handoff.prepare_handout_materials(case.facade, case.request, [], [])
    prepared = deepcopy(publication.result)
    assert not images(case.facade) and _inventory(case.facade.paths.state_root) == before
    result = publication.publish(); publication.accept(); publication.discard()
    assert result == prepared == publication.result
    assert len(result['image_assets']) == 1
    expected_path = case.facade.paths.task_root / 'preparation-v1/images' / (sha256(case.png).hexdigest() + '.image')
    expected = {**before, str(expected_path.relative_to(case.facade.paths.state_root)): sha256(case.png).hexdigest()}
    assert _inventory(case.facade.paths.state_root) == expected
    assert {raw for raw, _mtime in images(case.facade).values()} == {case.png}
    assert case.request == request and case.facade.state_store.snapshot() == state
    assert case.archive.read_bytes() == case.raw and case.provider.borrow_calls == 0


@pytest.mark.parametrize('location', ['original', 'work', 'snapshot'])
def test_only_occurrence_ids_are_normalized_and_all_eq_evidence_changes_affect_stamp(eq_handoff, location):
    result, _ = handoff._compile(eq_handoff.facade, eq_handoff.request, [], [])
    before = deepcopy(result); original_stamp = handoff._stamp(result)
    for change in ('condition', 'crop', 'container_sha', 'other_occurrence', 'script', 'instruction', 'field_sha'):
        altered = deepcopy(result); field = _image_field(altered, location)
        if change == 'condition': field['overlays'][0]['runs'][0]['text'] = '新催化条件'
        elif change == 'crop': field['base']['crop']['l'] += 1
        elif change == 'container_sha': field['base']['container_xml_sha256'] = 'f' * 64
        elif change == 'other_occurrence':
            replacement = _independent_picture(altered)
            assert replacement['id'] != field['base']['image_block_id']
            field['base']['image_block_id'] = replacement['id']
        elif change == 'script': field['overlays'][0]['runs'][0]['script'] = 'superscript'
        elif change == 'instruction': field['instruction'] += ' '
        else: field['field_sha256'] = 'f' * 64
        untouched = deepcopy(altered)
        assert handoff._stamp(altered) != original_stamp, (location, change)
        assert altered == untouched, (location, change)
    assert result == before


def test_literal_block_ids_in_plain_text_are_not_recursively_replaced(eq_handoff):
    first, _ = handoff._compile(eq_handoff.facade, eq_handoff.request, [], [])
    second, _ = handoff._compile(eq_handoff.facade, eq_handoff.request, [], [])
    assert handoff._stamp(first) == handoff._stamp(second)
    for location in ('original', 'work', 'snapshot'):
        left, right = deepcopy(first), deepcopy(second)
        key = 'text' if location == 'snapshot' else 'value'
        left_literal = _image_field(left)['base']['image_block_id']
        right_literal = _image_field(right)['base']['image_block_id']
        assert left_literal.startswith('CB-') and right_literal.startswith('CB-')
        assert left_literal != right_literal
        _plain_row(left, location)[key] = left_literal
        _plain_row(right, location)[key] = right_literal
        left_before, right_before = deepcopy(left), deepcopy(right)
        assert handoff._stamp(left) != handoff._stamp(right), location
        assert left == left_before and right == right_before


@pytest.mark.parametrize('change', ['prepared_result', 'fresh_recompile', 'source_bytes', 'cancel_guard'])
def test_eq_handoff_mutation_or_cancellation_rejects_atomically_without_images(eq_handoff, monkeypatch, change):
    case = eq_handoff
    before = _inventory(case.facade.paths.state_root); state = case.facade.state_store.snapshot()
    publication = handoff.prepare_handout_materials(case.facade, case.request, [], [])
    event = TaskCancellationEvent(); guard_calls = []
    if change == 'prepared_result':
        _image_field(publication.result)['overlays'][0]['runs'][0]['text'] = '暂存条件被改'
    elif change == 'fresh_recompile':
        compile_native = handoff._compile
        def different(*args, **kwargs):
            result, payloads = compile_native(*args, **kwargs)
            _image_field(result, 'snapshot')['base']['crop']['r'] += 1
            return result, payloads
        monkeypatch.setattr(handoff, '_compile', different)
    elif change == 'source_bytes':
        document = Document(BytesIO(case.raw)); document.add_paragraph('原文件在准备后发生变化。')
        document.save(case.archive)
    def guard():
        guard_calls.append(True)
        if change == 'cancel_guard' and len(guard_calls) == 2:
            event.set()
    try:
        with pytest.raises((ValueError, ReadCancelled)), read_cancel_scope(event):
            publication.publish(guard)
    finally:
        publication.discard()
        if change == 'source_bytes': case.archive.write_bytes(case.raw)
    assert not images(case.facade)
    assert _inventory(case.facade.paths.state_root) == before
    assert case.facade.state_store.snapshot() == state and case.provider.borrow_calls == 0
