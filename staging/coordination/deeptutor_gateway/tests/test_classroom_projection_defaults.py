"""Measured classroom defaults without an application, state or source workspace.

These are geometry/source regressions, not a classroom visual-quality approval.
All pictures are synthetic, and saved output goes only to pytest's own tmp tree.
"""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO

import pytest
from PIL import Image
from PySide6.QtWidgets import QApplication
from pptx import Presentation

from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    FLAGS, FORMATS, add_resource, block, digest, edit_resource, new_composition,
    resource, suggested_resource_title, validate_resource)
from integrations.deeptutor_shchem_v1.desktop_question_classroom_copy import classroom_question_copy
from integrations.deeptutor_shchem_v1.desktop_composition_render import (
    PADDING, PX_MM, arrange_composition, layout_composition, write_pptx)
from integrations.deeptutor_shchem_v1.desktop_composition_rich_render import flatten_elements
from integrations.deeptutor_shchem_v1.desktop_preparation_images import image_info, normalize_image_assets
from integrations.deeptutor_shchem_v1.desktop_question_answer_reveal import configure_question_answer_reveal


@pytest.fixture(autouse=True)
def no_complete_application_constructors(monkeypatch):
    from integrations.deeptutor_shchem_v1.desktop_facade import DesktopWorkbenchFacade
    from integrations.deeptutor_shchem_v1.desktop_state import DesktopStateStore
    from integrations.deeptutor_shchem_v1.desktop_editor_recovery import PreparationRecoveryStore
    from integrations.deeptutor_shchem_v1.desktop_workbench.main_window import TeacherWorkbenchWindow
    def forbidden(*_args, **_kwargs):
        raise AssertionError('This regression has no facade, state or recovery workspace')
    for cls in (DesktopWorkbenchFacade, DesktopStateStore, PreparationRecoveryStore, TeacherWorkbenchWindow):
        monkeypatch.setattr(cls, '__init__', forbidden)


@pytest.fixture(scope='module')
def qt_app():
    from integrations.deeptutor_shchem_v1.desktop_workbench.typography import install_ui_font
    app = QApplication.instance() or QApplication([])
    install_ui_font(app)
    return app


def _question(stem, *, extra=()):
    return resource('complete_question', '来源讲义2025.docx · 原文区块350—361', [
        block('text', 'shared_material', '共同材料与题目编号相互独立。', 349),
        block('text', 'question', stem, 350), *extra,
        block('text', 'answer', '【答案】AB', 356),
        block('text', 'analysis', '仅教师讲评：保持完整解释。', 361)],
        dict(source_path='synthetic-source.docx', source_sha256='a'*64,
            source_revision='b'*64, question_key='synthetic-key-999',
            scope_review_sha256='c'*64, scope_status='complete_native_scope',
            hierarchy=dict.fromkeys(('paper_id', 'theme_big_question_id',
                'printed_question_id', 'atomic_part_id'), 'unknown')))


def _composition(item, format='16:9', *, inset=12):
    value = new_composition(format)
    w, h = FORMATS[format]
    return add_resource(value, item, dict(page_id=value['pages'][0], x=inset,
        y=12, width=w-2*inset, height=h-24))


def _picture(size=(1620, 1250)):
    buffer = BytesIO()
    Image.new('RGB', size, '#47758a').save(buffer, format='PNG')
    data = buffer.getvalue(); checksum = sha256(data).hexdigest()
    asset = normalize_image_assets([dict(asset_id='IMG-'+checksum, sha256=checksum,
        caption='合成完整原图', source='synthetic fixture', purpose='观察原图',
        **image_info(data))])[0]
    return asset, {asset['asset_id']: data}


def _figure(kind, asset, data):
    rows = [block('image', 'material', asset['asset_id'], 5),
        block('text', 'material', '观察完整表头和图中数据。', 5)]
    evidence = dict(source_path='synthetic-source.pdf', source_sha256='a'*64,
        source_revision='b'*64)
    if kind == 'textbook_image':
        evidence.update(pdf_pages=[5], printed_pages=[5])
    elif kind == 'preparation_image':
        evidence.update(source_path=asset['source'], source_sha256=asset['sha256'],
            source_revision=digest(asset), image_asset=deepcopy(asset), source_roles=[])
    else:
        evidence.update(handout_source=dict(schema='shchem.handout-picture.v1',
            batch_id='synthetic-batch', source_id='synthetic-source',
            source_name='合成讲义.docx', source_path=evidence['source_path'],
            source_sha256=evidence['source_sha256'], source_revision=evidence['source_revision'],
            asset_id='native-picture-5', block_index=5, mime_type=asset['content_type'],
            sha256=asset['sha256'], bytes_count=len(data[asset['asset_id']])),
            image_asset=deepcopy(asset), source_roles=[])
    return resource(kind, '观察结构与表格', rows, evidence)


def _elements(layout):
    return [e for page in layout['pages'].values() for e in flatten_elements(page)]


def _body_text(layout, identity):
    return ''.join(e.get('text', '') for e in _elements(layout) if e.get('block_id') == identity)


@pytest.mark.parametrize('stem,expected', [
    ('例1（双选）乒乓球底板常用不同材料。下列说法正确的是', '例题1'),
    ('14．比较各气体装置的目的与条件。', '例题14'),
    ('【例题 12】完整条件和题面必须保留。', '例题12'),
    ('【变式训练2·变载体】完整材料及条件。', '例题'),
    ('将1 mol气体通入溶液，根据现象判断。', '例题'),
    ('3．【待查看原文：图片或图形】', '例题3'),
])
def test_new_title_is_neutral_navigation_and_seals_complete_question(stem, expected):
    item = _question(stem); before = deepcopy(item)
    assert suggested_resource_title(item) == expected
    assert item == before  # A suggestion does not edit either text scope.
    item['title'] = suggested_resource_title(item)
    assert item['blocks'] == before['blocks'] and item['original'] == before['original']
    assert item['original_sha256'] == before['original_sha256'] == digest(item['original'])
    assert set(item['original']['evidence']['hierarchy'].values()) == {'unknown'}
    assert all(item[name] is value for name, value in FLAGS.items())
    assert validate_resource(item) == item


def test_exact_printed_source_removal_does_not_promote_the_stem_to_a_title():
    citation = '（24-25高三上·上海·阶段练习）'
    stem = '例1（双选）'+citation+'完整材料、温度、浓度和末尾限制。'
    item = _question(stem); before = deepcopy(item)
    out = classroom_question_copy(item)
    assert out['title'] == '例题1'
    assert out['blocks'][1]['value'] == stem.replace(citation, '')
    assert out['original'] == before['original'] and out['original']['blocks'][1]['value'] == stem
    assert out['original_sha256'] == before['original_sha256']
    assert [b['id'] for b in out['blocks']] == [b['id'] for b in before['blocks']]
    assert [b['role'] for b in out['blocks']] == [b['role'] for b in before['blocks']]
    assert item == before and classroom_question_copy(out) == out


def test_teacher_question_title_and_working_edits_are_preserved(qt_app):
    item = _question('例1（2025·上海·等级考）完整条件。')
    value = _composition(item)
    teacher_title = '这些结构如何影响材料的性质？'
    changed = edit_resource(value, item['id'], title=teacher_title)
    current = changed['resources'][0]; before = deepcopy(current)
    assert suggested_resource_title(current) == teacher_title
    assert classroom_question_copy(current) == before
    assert current['blocks'] == item['blocks'] and current['original'] == item['original']
    laid = layout_composition(changed, [], {}, student=True)
    assert not laid['errors']
    title = [e for e in _elements(laid) if e.get('block_id') is None and e['kind'] == 'text']
    assert ''.join(e['text'] for e in title) == teacher_title
    assert current == before and '【答案】' not in ''.join(e.get('text', '') for e in _elements(laid))


@pytest.mark.parametrize('format', ['A4', 'A3', '16:9'])
@pytest.mark.parametrize('student', [False, True])
@pytest.mark.parametrize('columns', [1, 2])
def test_continuation_uses_measured_single_line_navigation_with_full_body(qt_app, format, student, columns):
    body = '完整条件甲乙丙丁。'*600 + '末尾限制必须保留'
    item = _question(body)
    item['title'] = ('比较各材料在相同条件下的性质与结构，再结合实验现象说明判断依据。'*3
        + '\n教师完整教学问题保持可编辑')
    value = _composition(item, format); before = deepcopy(value)
    flowed = arrange_composition(value, [], {}, columns=columns)
    laid = layout_composition(flowed, [], {}, student=student)
    assert not laid['errors'] and value == before and flowed['resources'] == before['resources']
    fragments = [f for f in laid['fragments'] if f['resource_id'] == item['id']]
    assert len(fragments) > 1 and _body_text(laid, item['blocks'][1]['id']) == body
    navigation = [e for e in _elements(laid) if e.get('navigation_heading') and e['kind'] != 'space']
    if format == '16:9':
        assert len(navigation) == len(fragments)-1
        for fragment in fragments[1:]:
            headings = [e for e in fragment['elements'] if e.get('navigation_heading') and e['kind'] == 'text']
            assert len(headings) == 1
            heading = headings[0]
            assert heading['text'].endswith('（续）') and '…' in heading['text']
            assert heading['font_px'] == 32 and heading['block_id'] is None
            assert heading['_line'].naturalTextWidth()/PX_MM <= fragment['frame']['width']-2*PADDING+.001
            assert heading['height'] < 13  # One actual 24pt line, including leading.
            body_elements = [e for e in fragment['elements'] if e.get('block_id') is not None and e['kind'] != 'space']
            assert min(e['y'] for e in body_elements) >= heading['y']+heading['height']
            assert max(e['y']+e['height'] for e in fragment['elements']) <= fragment['frame']['y']+fragment['frame']['height']-PADDING+.001
        assert all(e['font_px'] == 32 for e in _elements(laid) if e.get('block_id') == item['blocks'][1]['id'] and e['kind'] == 'text')
    else:
        assert not navigation
        assert all(e['font_px'] == 15 for e in _elements(laid) if e.get('block_id') == item['blocks'][1]['id'] and e['kind'] == 'text')
    assert ('仅教师讲评' in _body_text(laid, item['blocks'][-1]['id'])) is (not student)
    assert flowed['resources'][0]['title'] == item['title']


@pytest.mark.parametrize('kind', ['textbook_image', 'preparation_image', 'handout_image'])
@pytest.mark.parametrize('format', ['A4', 'A3', '16:9'])
def test_one_independent_whole_figure_centers_only_in_projection(qt_app, kind, format):
    asset, data = _picture(); item = _figure(kind, asset, data)
    value = _composition(item, format, inset=17); before = deepcopy(value)
    original_bytes = data[asset['asset_id']]
    laid = layout_composition(value, [asset], data, student=True)
    assert not laid['errors'] and len(laid['fragments']) == 1
    image = next(e for e in _elements(laid) if e['kind'] == 'image')
    frame = laid['fragments'][0]['frame']; inner_left = frame['x']+PADDING
    inner_right = frame['x']+frame['width']-PADDING
    assert image['width']/image['height'] == pytest.approx(1620/1250)
    assert image['x'] >= inner_left and image['x']+image['width'] <= inner_right+.001
    if format == '16:9':
        assert image['x']+image['width']/2 == pytest.approx((inner_left+inner_right)/2)
        assert image['x'] > inner_left+30  # A tall whole figure remains height-limited.
    else:
        assert image['x'] == inner_left
        assert not image.get('center_in_frame')
    assert _body_text(laid, item['blocks'][1]['id']) == item['blocks'][1]['value']
    assert value == before and data[asset['asset_id']] == original_bytes
    assert sha256(original_bytes).hexdigest() == asset['sha256']


def test_inline_formula_picture_remains_in_its_text_line_at_native_scale(qt_app):
    asset, data = _picture((80, 40))
    marker = '【待查看原文：图片或图形】'
    row = block('text', 'question', '例2 图前条件'+marker+'图后条件。', 350)
    picture = block('image', 'question', asset['asset_id'], 350)
    row['presentation'] = dict(schema='shchem.native-word-paragraph.v1', images=[
        dict(image_block_id=picture['id'], placeholder=marker, width_mm=12, height_mm=6)])
    item = _question('占位题面')
    item['original']['blocks'][1:2] = [deepcopy(row), deepcopy(picture)]
    item['blocks'][1:2] = [deepcopy(row), deepcopy(picture)]
    item['original_sha256'] = digest(item['original'])
    item = validate_resource(item); item['title'] = suggested_resource_title(item)
    before = deepcopy(item)
    laid = layout_composition(_composition(item), [asset], data)
    assert not laid['errors']
    line = next(e for p in laid['pages'].values() for e in p if e['kind'] == 'rich_line' and any(p['kind']=='image' for p in e['_parts']))
    parts = line['_parts']; image = next(p for p in parts if p['kind'] == 'image')
    left = next(p for p in parts if '图前条件' in p.get('text', ''))
    right = next(p for p in parts if '图后条件' in p.get('text', ''))
    assert left['offset_x'] < image['offset_x'] < right['offset_x']
    assert image['width'] == pytest.approx(12*32/15) and image['height'] == pytest.approx(6*32/15)
    assert not image.get('center_in_frame') and item == before
    assert item['blocks'][1]['presentation']['images'][0]['width_mm'] == 12


def test_neutral_title_still_allows_selected_short_answer_same_page_and_click(qt_app, tmp_path):
    asset, data = _picture((1269, 298))
    item = _question('例1 完整题干条件。', extra=[block('image', 'question', asset['asset_id'], 351)])
    item['title'] = suggested_resource_title(item)
    value = _composition(item); before = deepcopy(value)
    answer_id = next(b['id'] for b in item['blocks'] if b['role'] == 'answer')
    configured = configure_question_answer_reveal(value, item['id'], [answer_id], [asset], data)
    static = layout_composition(configured, [asset], data, student=True)
    classroom = layout_composition(configured, [asset], data, student=True, include_reveals=True)
    assert not static['errors'] and not classroom['errors']
    assert len(static['pages']) == len(classroom['pages']) == 1
    assert _body_text(static, answer_id) == '' and _body_text(classroom, answer_id) == '【答案】AB'
    assert _body_text(classroom, item['blocks'][-1]['id']) == ''
    assert sum(e['kind']=='image' for e in _elements(static)) == 1
    reveals = [e for e in _elements(classroom) if e.get('reveal_id')]
    assert len(reveals) == 1 and reveals[0]['role'] == 'answer' and reveals[0]['font_px'] == 32
    assert configured['resources'][0]['blocks'] == item['blocks']
    assert configured['resources'][0]['original'] == item['original'] and value == before
    ppt_path = tmp_path/'synthetic-selected-answer.pptx'
    write_pptx(classroom, ppt_path, data)
    deck = Presentation(ppt_path)
    assert len(deck.slides) == 1
    assert sum(s.shape_type == 13 for s in deck.slides[0].shapes) == 1
    assert len(deck.slides[0]._element.xpath('.//p:cTn[@nodeType="clickEffect"]')) == 1
    reveal_group = next(s for s in deck.slides[0].shapes if s.name.startswith('shchem-click-'))
    assert [e.get('spid') for e in deck.slides[0]._element.xpath('.//p:spTgt')] == [str(reveal_group.shape_id)]
    assert '仅教师讲评' not in ''.join(s.text for s in deck.slides[0].shapes if s.has_text_frame)
