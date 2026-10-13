"""Whole text/image coverage in measured pages, real Office files and backups."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from types import SimpleNamespace
import re
import pytest
pytest.importorskip('PySide6')
from PySide6.QtWidgets import QApplication
from PIL import Image
from test_lesson_design import sample
from test_desktop_preparation_drafts import _record
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import *
from integrations.deeptutor_shchem_v1.desktop_composition_render import *
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
from integrations.deeptutor_shchem_v1.desktop_lesson_output import export_design, COMPOSITION_FILES, checked_file, copy_output_bundle
from integrations.deeptutor_shchem_v1.desktop_state import DesktopStateStore
from integrations.deeptutor_shchem_v1.desktop_backup import plan_backup,create_backup,inspect_backup,restore_backup,manifest_revision
from integrations.deeptutor_shchem_v1.desktop_workbench.lesson_composition_panel import ResourceEditor


@pytest.fixture(scope='session')
def app():
    from integrations.deeptutor_shchem_v1.desktop_workbench.typography import install_ui_font
    app=QApplication.instance() or QApplication([]);install_ui_font(app)
    return app


@pytest.fixture
def material(tmp_path,app):
    state=DesktopStateStore(tmp_path/'state')
    store=PreparationImageStore(state.root/'tasks/preparation-v1/images')
    buffer=BytesIO();Image.new('RGB',(120,90),'#769f89').save(buffer,format='PNG')
    asset=store.import_bytes(buffer.getvalue(),'合成图','测试作者输入','完整图片输出核对')
    body='完整长题条件甲乙丙丁。'*600+'尾部不可丢失条件'
    item=resource('complete_question','合成完整题',[
        block('text','shared_material','公共条件：全程温度保持固定。',1),
        block('text','question',body,2),block('image','question',asset['asset_id'],2),
        block('text','answer','仅教师答案唯一标记',3),block('image','answer',asset['asset_id'],3)],
        {'source_path':'合成原题.docx','source_sha256':'a'*64,'source_revision':'c'*64,
        'question_key':'synthetic-full-question','scope_review_sha256':'d'*64,'scope_status':'complete_native_scope',
        'hierarchy':dict.fromkeys(('paper_id','theme_big_question_id','printed_question_id','atomic_part_id'),'unknown')})
    personal=resource('personal_text','个人段落',[block('text','personal','个人文字需同时进入教师和学生文件')],
        {'source_path':'作者输入','source_sha256':'unknown','source_revision':'unknown'})
    value=new_composition()
    for row in (item,personal):value=add_resource(value,row,{'page_id':value['pages'][0],'x':12,'y':12,'width':186,'height':273})
    data={asset['asset_id']:buffer.getvalue()}
    value=arrange_composition(value,[asset],data)
    p=sample();p['lesson_design']=enable_composition(p['lesson_design']);p['lesson_design']['composition']=value;p['image_assets']=[asset]
    facade=SimpleNamespace(paths=SimpleNamespace(task_root=state.root/'tasks'),preparation_image_bytes=store.load)
    return p,facade,state,data,body


def text(layout,block_id=None):
    return ''.join(e.get('text','') for page in layout['pages'].values() for e in page if block_id is None or e.get('block_id')==block_id)


@pytest.mark.parametrize('format',FORMATS)
@pytest.mark.parametrize('columns',[1,2])
def test_whole_lines_pictures_and_tail_survive_continuations(material,format,columns):
    p,facade,state,data,body=material;value=deepcopy(p['lesson_design']['composition']);value['format']=format
    w,h=FORMATS[format]
    value['placements']={r['id']:{'page_id':value['pages'][0],'x':12,'y':12,'width':w-24,'height':h-24} for r in value['resources']}
    value=arrange_composition(value,p['image_assets'],data,columns=columns)
    teacher=layout_composition(value,p['image_assets'],data,student=False)
    student=layout_composition(value,p['image_assets'],data,student=True)
    assert not teacher['errors'] and not student['errors']
    key=value['resources'][0]['blocks'][1]['id']
    assert text(teacher,key)==body and text(student,key)==body
    assert '仅教师答案唯一标记' in text(teacher) and '仅教师答案唯一标记' not in text(student)
    assert sum(e['kind']=='image' for page in teacher['pages'].values() for e in page)==2
    assert sum(e['kind']=='image' for page in student['pages'].values() for e in page)==1
    assert len(value['pages'])>1
    assert value['resources']==p['lesson_design']['composition']['resources']


def test_small_or_overlapping_frames_block_export_without_losing_original(material,tmp_path):
    p,facade,state,data,body=material;value=deepcopy(p['lesson_design']['composition'])
    first,second=value['resources'];frame=deepcopy(value['placements'][first['id']]);frame.pop('continuations',None);frame['height']=15
    value=place_resource(value,first['id'],frame)
    assert layout_composition(value,p['image_assets'],data)['errors']
    p['lesson_design']['composition']=value
    with pytest.raises(ValueError,match='未放完'):export_design(facade,p)
    assert not (facade.paths.task_root/'preparation-v1/node-exports').exists()
    assert value['resources'][0]['blocks'][1]['value']==body
    value=place_resource(value,second['id'],{**frame,'height':30})
    assert any('重叠' in e for e in layout_composition(value,p['image_assets'],data)['errors'])


def test_actual_pdf_pptx_documents_and_six_file_backup(material,tmp_path):
    import pypdfium2 as pdfium
    from pptx import Presentation
    from docx import Document
    p,facade,state,data,body=material
    before=deepcopy(p);record=export_design(facade,p);assert p==before
    p['lesson_design']['exports'].append(record)
    assert {f['name'] for f in record['files']}==set(COMPOSITION_FILES)
    def pdf_text(name):
        result=[];pdf=pdfium.PdfDocument(str(checked_file(facade,p['lesson_design'],record['id'],name)))
        try:
            for i in range(len(pdf)):
                page=pdf[i];textpage=page.get_textpage()
                try:result.append(textpage.get_text_range())
                finally:textpage.close();page.close()
        finally:pdf.close()
        return re.sub(r'\s','', ''.join(result))
    assert '尾部不可丢失条件' in pdf_text('lesson_canvas.pdf')
    assert '尾部不可丢失条件' in pdf_text('student_canvas.pdf')
    assert '仅教师答案唯一标记' in pdf_text('lesson_canvas.pdf')
    assert '仅教师答案唯一标记' not in pdf_text('student_canvas.pdf')
    for name,student in (('lesson_presentation.pptx',False),('student_presentation.pptx',True)):
        deck=Presentation(checked_file(facade,p['lesson_design'],record['id'],name))
        words=''.join(shape.text for slide in deck.slides for shape in slide.shapes
                      if shape.has_text_frame and shape.name!='shchem-continuation-heading')
        assert body in words
        assert ('仅教师答案唯一标记' not in words)==student
        assert len(deck.slides)==len(p['lesson_design']['composition']['pages'])
        assert sum(shape.shape_type==13 for slide in deck.slides for shape in slide.shapes)==(1 if student else 2)
    for name,student in (('lesson_plan.docx',False),('student_worksheet.docx',True)):
        doc=Document(checked_file(facade,p['lesson_design'],record['id'],name));words=''.join(x.text for x in doc.paragraphs)
        assert body in words and ('仅教师答案唯一标记' not in words)==student
    stored=_record(p);stored['lesson_design']=p['lesson_design'];state.save_draft('preparation-v1',stored)
    plan=plan_backup(state.root,include_images=True,include_tasks=True)
    assert plan.summary['node_output_files']==6
    backup=tmp_path/'composition.zip';create_backup(plan,backup);manifest=inspect_backup(backup)
    restored=tmp_path/'restored';restore_backup(backup,restored,expected_manifest=manifest_revision(manifest))
    for f in record['files']:
        assert sha256((restored/'tasks/preparation-v1/node-exports'/record['id']/f['name']).read_bytes()).hexdigest()==f['sha256']
    folder=copy_output_bundle(facade,p['lesson_design'],record['id'],tmp_path)
    from pathlib import Path
    assert all((Path(folder)/name).is_file() for name in COMPOSITION_FILES)


def test_editor_unlock_edit_relock_retains_original_and_answer_scope(material):
    p,facade,state,data,body=material;value=p['lesson_design']['composition'];item=value['resources'][0]
    editor=ResourceEditor(value,item['id'],None)
    assert editor.body.isReadOnly()
    editor.locked.setChecked(False);editor.blocks.setCurrentIndex(1);editor.body.setPlainText(body+'教师修改的作品条件')
    editor.title.setText('新作品标题');editor.locked.setChecked(True);editor.apply()
    changed=editor.result_value['resources'][0]
    assert changed['locked'] and changed['original']==item['original']
    assert changed['blocks'][1]['value'].endswith('教师修改的作品条件')
    assert changed['title']=='新作品标题' and value['resources'][0]==item
    assert changed['blocks'][-2]['role']=='answer'
    editor.deleteLater()
