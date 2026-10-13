"""Projection keeps full source content, readable type and reversible purpose changes."""
from copy import deepcopy
from docx import Document
from pptx import Presentation
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from test_composition_render import app, material, text
from test_composition_inline_personal import inline
from test_composition_ui import desk, window, personal
from test_desktop_studio_ui import settle
from test_lesson_design import sample
from integrations.deeptutor_shchem_v1.desktop_lesson_design import digest
from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition, layout_composition, write_pptx
from integrations.deeptutor_shchem_v1.desktop_composition_rich_render import flatten_elements
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import FORMATS, enable_composition
from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents
from integrations.deeptutor_shchem_v1.desktop_composition_source_summary import source_summary


def projection(value, assets, data):
    value=deepcopy(value);value['format']='16:9';w,h=FORMATS['16:9']
    value['placements']={r['id']:{'page_id':value['pages'][0],'x':12,'y':12,'width':w-24,'height':h-24} for r in value['resources']}
    return arrange_composition(value,assets,data)


def test_projected_long_question_remains_complete_and_student_answers_stay_hidden(material,tmp_path):
    payload,facade,state,data,body=material;before=deepcopy(payload['lesson_design']['composition'])
    value=projection(before,payload['image_assets'],data)
    teacher=layout_composition(value,payload['image_assets'],data,student=False)
    student=layout_composition(value,payload['image_assets'],data,student=True)
    assert not teacher['errors'] and not student['errors']
    key=value['resources'][0]['blocks'][1]['id']
    assert text(teacher,key)==text(student,key)==body
    assert '仅教师答案唯一标记' not in text(student)
    assert value['resources']==before['resources']
    path=tmp_path/'projection.pptx';write_pptx(student,path,data);deck=Presentation(path)
    assert abs(deck.slide_width/deck.slide_height-16/9)<.001
    runs=[r for slide in deck.slides for shape in slide.shapes if shape.has_text_frame for p in shape.text_frame.paragraphs for r in p.runs]
    assert all(r.font.size.pt>=24 for r in runs)
    assert ''.join(r.text for r in runs).count('尾部不可丢失条件')==1


def test_projection_enlarges_inline_scripts_and_pictures_without_rewriting_originals(inline):
    value,assets,data,_,_=inline;before=deepcopy(value)
    projected=projection(value,assets,data);layout=layout_composition(projected,assets,data,student=True)
    assert not layout['errors']
    elements=[e for page in layout['pages'].values() for e in flatten_elements(page)]
    base=next(e for e in elements if e.get('text')=='图前H')
    sub=next(e for e in elements if e.get('text')=='2');sup=next(e for e in elements if e.get('text')=='+')
    assert base['font_px']*72/96>=24
    assert sub['font_px']<base['font_px'] and sub['y']>base['y']
    assert sup['font_px']<base['font_px'] and sup['y']<base['y']
    assert next(e for e in elements if e['kind']=='image')['width']>18
    assert projected['resources']==before['resources'] and value==before


def test_purpose_buttons_are_one_undoable_transaction(desk):
    d,panel,app,page=desk
    panel.add_result({'resources':[personal('完整条件。'*200)]},digest(d.history.value));before=deepcopy(d.history.value)
    QTest.mouseClick(panel.projection_button,Qt.MouseButton.LeftButton);settle(app)
    projected=deepcopy(d.history.value)
    assert panel.format.currentText()=='16:9' and not panel.layout_result['errors']
    assert projected['composition']['resources']==before['composition']['resources']
    d.travel(False);assert d.history.value==before
    d.travel(True);assert d.history.value==projected
    QTest.mouseClick(panel.print_button,Qt.MouseButton.LeftButton);settle(app)
    assert panel.format.currentText()=='A4' and not panel.layout_result['errors']
    d.travel(False);assert d.history.value==projected


def test_word_body_omits_source_notes_while_original_evidence_stays_complete(inline,tmp_path):
    import json
    from integrations.deeptutor_shchem_v1.desktop_composition_source_summary import raw_source_text
    value,assets,data,_,_=inline;value=arrange_composition(value,assets,data)
    payload=sample();payload['lesson_design']=enable_composition(payload['lesson_design']);payload['lesson_design']['composition']=value;payload['image_assets']=assets
    before=deepcopy(payload);original=deepcopy(value['resources'][0]['original'])
    summary=source_summary(value['resources'][0])
    write_documents(payload,tmp_path,data)
    teacher=Document(tmp_path/'lesson_plan.docx');student=Document(tmp_path/'student_worksheet.docx')
    for doc in (teacher,student):
        body='\n'.join(p.text for p in doc.paragraphs)
        assert all(label not in body for label in ('来源与内容：','素材类型：','原资料标题：','来源图片：'))
        assert original['evidence']['source_sha256'] not in body
    item=payload['lesson_design']['composition']['resources'][0]
    assert payload==before and item['original']==original
    assert source_summary(item)==summary
    assert '素材类型：完整题' in summary and '题面：原文区块1' in summary
    assert '教师答案与解析：原文区块2' in summary and '来源图片：1处' in summary
    assert json.loads(raw_source_text(item))==original['evidence']
    assert json.loads(raw_source_text(item,full=True))==original
    assert '仅教师答案' not in ''.join(p.text for p in student.paragraphs)
