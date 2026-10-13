"""Native table images keep their original cells, order and editable text."""
from collections import Counter
from copy import deepcopy
from hashlib import sha256
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm
from PIL import Image
from pptx import Presentation
import pytest

from integrations.deeptutor_shchem_v1.desktop_handout_table_projection import source_table_projection
from integrations.deeptutor_shchem_v1.desktop_handout_knowledge_tables import table_chunks, write_word_table
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import digest, validate_resource
from integrations.deeptutor_shchem_v1.desktop_composition_rich_render import flatten_elements
from test_handout_knowledge_tables import app, row


POSITIONS = [(3, 2), (4, 2), (5, 2), (5, 2), (6, 2), (6, 2), (7, 2), (7, 3), (7, 3)]


def illustrated(tmp_path, *, teacher_only=False, merged=False, preface=False):
    from test_handout_knowledge import make_facade, request_for, word_bytes
    from integrations.deeptutor_shchem_v1.desktop_handout_knowledge import import_handout_knowledge
    from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
    document = Document(); document.add_heading('知识点1 完整图文比较', 3)
    if preface:
        for text in ('得分速记', '1．观察结构', '（1）完整定义保留。'):
            document.add_paragraph(text)
        buffer = BytesIO(); Image.new('RGB', (60, 30), '#4a7290').save(buffer, 'PNG')
        document.add_paragraph().add_run().add_picture(BytesIO(buffer.getvalue()), width=Mm(45))
        document.add_paragraph('（2）完整比较表')
    table = document.add_table(rows=7, cols=3)
    for ri in range(7):
        for ci in range(3): table.cell(ri, ci).text = f'原文{ri + 1}/{ci + 1}'
    paragraph = table.cell(1, 1).paragraphs[0]
    paragraph.add_run('H'); paragraph.add_run('2').font.subscript = True; paragraph.add_run('O')
    if merged:
        table.cell(2, 0).merge(table.cell(3, 0))
        table.cell(4, 1).merge(table.cell(4, 2))
    pictures = []
    for index in range(5):
        buffer = BytesIO(); Image.new('RGB', (60 + index * 10, 25), (50 * index, 80, 100)).save(buffer, 'PNG')
        pictures.append(buffer.getvalue())
    for (ri, ci), identity in zip(POSITIONS, (0, 0, 1, 0, 2, 0, 0, 3, 4)):
        p = table.cell(ri - 1, ci - 1).paragraphs[0]
        p.add_run().add_picture(BytesIO(pictures[identity]), width=Mm(12))
        p.add_run(f'第{len(p.runs)}段')
    document.add_paragraph('完整表格之后的原文也必须保留。')
    facade = make_facade(word_bytes(document), tmp_path)
    request = request_for(facade, teacher_only=teacher_only)
    result = import_handout_knowledge(facade, [request], [], [])
    store = PreparationImageStore(facade.paths.task_root / 'preparation-v1/images')
    data = {a['asset_id']: store.load(a) for a in result['image_assets']}
    item = result['resources'][0]
    table_index = next(b['index'] for b in item['original']['evidence']['section_snapshot']['blocks'] if 'table_rows' in b)
    text = next(b for b in item['blocks'] if b['kind'] == 'text' and b['source_index'] == table_index)
    return item, text, result['image_assets'], data


def composition(item, assets, data):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition
    from integrations.deeptutor_shchem_v1.desktop_lesson_composition import new_composition, add_resource
    value = new_composition('16:9')
    value = add_resource(value, item, dict(page_id=value['pages'][0], x=12, y=12, width=296, height=156))
    return arrange_composition(value, assets, data)


def payload_for(value, assets, *, classroom_notes=False):
    from test_lesson_design import sample
    from integrations.deeptutor_shchem_v1.desktop_lesson_composition import enable_composition
    payload = sample(); payload['lesson_design'] = enable_composition(payload['lesson_design'])
    payload['lesson_design']['composition'] = value; payload['image_assets'] = assets
    if classroom_notes: payload['lesson_design']['classroom_notes'] = True
    return payload


def test_projection_binds_nine_occurrences_in_cells_even_with_five_unique_assets(tmp_path):
    item, text, assets, data = illustrated(tmp_path); before = deepcopy(item)
    projection = source_table_projection(item, text)
    occurrences = projection['occurrences']
    assert len(occurrences) == 9 and len({r['asset_id'] for r in occurrences}) == 5
    assert [(r['row'], r['column']) for r in occurrences] == POSITIONS
    assert [r['position'] for r in occurrences] == list(range(1, 10))
    assert len(set(projection['image_block_ids'])) == 9
    assert Counter(r['asset_id'] for r in occurrences).most_common(1)[0][1] == 5
    for occurrence in occurrences:
        assert sha256(data[occurrence['asset_id']]).hexdigest() == occurrence['preview_sha256']
        assert occurrence['width_mm'] == pytest.approx(12)
        cell = projection['rows'][occurrence['row'] - 1]['cells'][occurrence['column'] - 1]
        assert cell['text'][occurrence['text_start']:occurrence['text_end']] == occurrence['placeholder']
    assert item == before


@pytest.mark.parametrize('field,value', [('position', 7), ('preview_sha256', 'd' * 64),
    ('container_xml_sha256', 'unknown'), ('width_mm', -1), ('source_role', 'answer'), ('source_sha256', 'c' * 64)])
def test_corrupt_occurrence_evidence_never_unlocks_table_projection(tmp_path, field, value):
    item, text, _, _ = illustrated(tmp_path)
    snapshot = item['original']['evidence']['section_snapshot']['blocks'][1]
    snapshot['images'][0][field] = value; item['original_sha256'] = digest(item['original'])
    with pytest.raises(ValueError): source_table_projection(item, text)


def test_unknown_placeholder_and_changed_source_digest_fail_closed(tmp_path):
    from integrations.deeptutor_shchem_v1.desktop_preparation_sources import _readable_table
    item, text, _, _ = illustrated(tmp_path)
    bad = deepcopy(item); bad['original_sha256'] = 'd' * 64
    with pytest.raises(ValueError, match='摘要'): source_table_projection(bad, bad['blocks'][1])
    snapshot = item['original']['evidence']['section_snapshot']['blocks'][1]
    cell = snapshot['table_rows'][2]['cells'][1]
    cell['text'] = cell['text'].replace('【待查看原文：图片或图形】', '【待查看原文：未知对象】')
    snapshot['text'] = _readable_table(snapshot['table_rows'])
    for rows in (item['blocks'], item['original']['blocks']):
        next(b for b in rows if b['id'] == text['id'])['value'] = snapshot['text']
    item['original_sha256'] = digest(item['original'])
    assert source_table_projection(item, text) is None


@pytest.mark.parametrize('change', ['asset', 'delete', 'order', 'role', 'source_index'])
def test_work_image_identity_cannot_silently_restore_or_consume_old_binding(tmp_path, change):
    item, text, _, _ = illustrated(tmp_path)
    indices = [i for i, b in enumerate(item['blocks']) if b['kind'] == 'image']
    first, second = indices[:2]
    if change == 'delete': del item['blocks'][first]
    elif change == 'order': item['blocks'][first], item['blocks'][second] = item['blocks'][second], item['blocks'][first]
    elif change == 'asset': item['blocks'][first]['value'] = item['blocks'][indices[-1]]['value']
    elif change == 'role': item['blocks'][first]['role'] = 'teacher_note'
    else: item['blocks'][first]['source_index'] += 1
    with pytest.raises(ValueError): source_table_projection(item, text)


@pytest.mark.parametrize('classroom_notes', [False, True])
def test_teacher_edit_keeps_work_text_and_all_nine_unbound_occurrences(app, tmp_path, classroom_notes):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import layout_composition
    from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents
    item, text, assets, data = illustrated(tmp_path / 'source'); source = deepcopy(item['original'])
    text['value'] = '教师重新组织后的完整工作文字。'
    validate_resource(item); assert source_table_projection(item, text) is None
    value = composition(item, assets, data)
    layout = layout_composition(value, assets, data, student=True)
    parts = [p for page in layout['pages'].values() for p in flatten_elements(page)]
    assert sum(p['kind'] == 'image' for p in parts) == 9
    assert not any('_table_cell' in p for p in parts)
    assert any(p.get('text') == text['value'] for p in parts)
    output = tmp_path / 'out'; output.mkdir()
    write_documents(payload_for(value, assets, classroom_notes=classroom_notes), output, data)
    for name in ('lesson_plan.docx', 'student_worksheet.docx'):
        doc = Document(output / name)
        assert len(doc.inline_shapes) == 9 and not doc.tables
        assert any(p.text == text['value'] for p in doc.paragraphs)
    assert item['original'] == source


@pytest.mark.parametrize('merged', [False, True])
def test_full_rich_cells_canvas_pdf_ppt_and_native_word_preserve_positions(app, tmp_path, merged):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import layout_composition, page_image, write_pdf, write_pptx
    from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents
    item, text, assets, data = illustrated(tmp_path / 'source', merged=merged); before = deepcopy(item)
    projection = source_table_projection(item, text); value = composition(item, assets, data)
    for student in (False, True):
        layout = layout_composition(value, assets, data, student=student); assert not layout['errors']
        flattened = [p for page in layout['pages'].values() for p in flatten_elements(page)]
        for page in layout['pages'].values():
            for panel in page:
                if panel.get('block_id') != text['id'] or '_table_source_rows' not in panel: continue
                leaves = list(flatten_elements([panel]))
                for cell in (p for p in leaves if p['kind'] == 'diagram_box'):
                    for part in (p for p in leaves if p['kind'] in {'image','text'} and p.get('_table_cell')==cell['_table_cell']):
                        assert part['x'] >= cell['x'] and part['y'] >= cell['y']
                        assert part['x']+part['width'] <= cell['x']+cell['width']+.001
                        assert part['y']+part['height'] <= cell['y']+cell['height']+.001
        images = [p for p in flattened if p['kind'] == 'image']
        assert [p['image_block_id'] for p in images] == projection['image_block_ids']
        assert [p['_table_cell'][:2] for p in images] == POSITIONS
        assert all(p['font_px'] == 32 for p in flattened if p['kind'] == 'text' and '_table_cell' in p)
        for page in layout['pages']:
            assert not page_image(layout, page).isNull()
        pdf = tmp_path / f'canvas-{student}.pdf'; write_pdf(layout, pdf)
        assert pdf.read_bytes().startswith(b'%PDF')
        ppt = tmp_path / f'presentation-{student}.pptx'; write_pptx(layout, ppt, data)
        deck = Presentation(ppt)
        pictures = [s for slide in deck.slides for s in slide.shapes if s.shape_type == 13]
        assert [sha256(s.image.blob).hexdigest() for s in pictures] == [r['preview_sha256'] for r in projection['occurrences']]
        all_text = ''.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
        assert '完整表格之后的原文也必须保留。' in all_text and 'H₂O' in all_text
        assert '待查看原文' not in all_text and '表格开始' not in all_text
        for shape, part in zip(pictures, images):
            assert shape.left.mm == pytest.approx(part['x'], abs=.001)
            assert shape.top.mm == pytest.approx(part['y'], abs=.001)
            assert shape.width.mm == pytest.approx(part['width'], abs=.001)
    write_documents(payload_for(value, assets), tmp_path, data)
    for name in ('lesson_plan.docx', 'student_worksheet.docx'):
        doc = Document(tmp_path / name); table = doc.tables[0]
        assert len(table._tbl.tr_lst) == 7 and len(doc.inline_shapes) == 9
        word_occurrences = []
        for ri, tr in enumerate(table._tbl.tr_lst, 1):
            column = 1
            for tc in tr.tc_lst:
                for image in tc.xpath('.//a:blip'):
                    word_occurrences.append((ri, column, sha256(doc.part.related_parts[image.get(qn('r:embed'))].blob).hexdigest()))
                column += tc.grid_span
        assert word_occurrences == [(r['row'], r['column'], r['preview_sha256']) for r in projection['occurrences']]
        assert not table._tbl.xpath('.//w:t[contains(text(), "待查看原文")]')
        assert table._tbl.xpath('.//w:vertAlign[@w:val="subscript"]')
        if merged:
            assert table._tbl.xpath('.//w:gridSpan[@w:val="2"]')
            assert table._tbl.xpath('.//w:vMerge[@w:val="restart"]')
            assert table._tbl.xpath('.//w:vMerge[@w:val="continue"]')
    assert item == before and value['resources'][0] == before


def test_teacher_scope_table_remains_out_of_student_canvas_and_word(app, tmp_path):
    from integrations.deeptutor_shchem_v1.desktop_composition_render import layout_composition
    from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents
    item, text, assets, data = illustrated(tmp_path / 'source', teacher_only=True)
    assert len(source_table_projection(item, text)['occurrences']) == 9
    value = composition(item, assets, data)
    student = layout_composition(value, assets, data, student=True)
    assert not any(p['kind'] == 'image' for page in student['pages'].values() for p in flatten_elements(page))
    write_documents(payload_for(value, assets), tmp_path, data)
    assert len(Document(tmp_path / 'lesson_plan.docx').inline_shapes) == 9
    assert not Document(tmp_path / 'student_worksheet.docx').inline_shapes


def test_recursive_offsets_reach_run_and_image_leaves():
    leaves = list(flatten_elements([dict(kind='diagram', x=10, y=20, width=80, height=50, resource_id='r', _parts=[
        dict(kind='rich_line', offset_x=5, offset_y=7, width=60, height=15, block_id='b', _table_cell=(2, 3, 1), _parts=[
            dict(kind='image', offset_x=3, offset_y=4, width=12, height=9, image_block_id='i'),
            dict(kind='text', offset_x=18, offset_y=2, width=30, height=8, text='原文')])])]))
    assert [(e['x'], e['y']) for e in leaves] == [(18, 31), (33, 29)]
    assert all(e['_table_cell'] == (2, 3, 1) and e['resource_id'] == 'r' and e['block_id'] == 'b' for e in leaves)


@pytest.mark.parametrize('earlier_lines', [0, 7, 11])
def test_arranged_rich_table_keeps_first_page_budget_after_save_and_reopen(app, tmp_path, earlier_lines):
    import json
    from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition, layout_composition
    from integrations.deeptutor_shchem_v1.desktop_lesson_composition import new_composition, add_resource, resource, block
    item, text, assets, data = illustrated(tmp_path / 'source', preface=True)
    projection = source_table_projection(item, text)
    resources = []
    if earlier_lines:
        resources.append(resource('personal_text', '前一资源', [block('text', 'personal',
            '\n'.join('前一资源的完整内容。' for _ in range(earlier_lines)))],
            dict(source_path='作者输入', source_sha256='unknown', source_revision='unknown')))
    resources.append(item); value = new_composition('16:9')
    for resource_item in resources:
        value = add_resource(value, resource_item, dict(page_id=value['pages'][0], x=12, y=12, width=296, height=156))
    saved = json.loads(json.dumps(arrange_composition(value, assets, data)))
    for student in (False, True):
        layout = layout_composition(saved, assets, data, student=student)
        assert not layout['errors']
        images = [p for page in layout['pages'].values() for p in flatten_elements(page)
            if p['kind'] == 'image' and p.get('block_id') == text['id']]
        assert [p['image_block_id'] for p in images] == projection['image_block_ids']
    first = saved['placements'][item['id']]
    assert first['height'] == pytest.approx(180 - 12 - first['y'])


def test_rich_cell_text_stays_visible_on_white_output_with_dark_system_palette(app):
    from PySide6.QtGui import QColor, QPalette, QImage, QPainter
    from integrations.deeptutor_shchem_v1.desktop_composition_rich_render import rich_lines
    from integrations.deeptutor_shchem_v1.desktop_composition_render import paint_elements
    original = app.palette(); dark = QPalette(original)
    dark.setColor(QPalette.ColorRole.Text, QColor('white')); app.setPalette(dark)
    try:
        chunks = rich_lines({'id':'cell'}, {}, 80, {}, {}, 60, font_px=32,
            tokens=[{'kind':'text','text':'CH₃－COOH','script':'baseline'}])
        canvas = QImage(400,100,QImage.Format.Format_RGB32); canvas.fill(QColor('white'))
        painter = QPainter(canvas); painter.scale(4,4)
        paint_elements(painter,[{**c,'x':0,'y':0} for c in chunks]); painter.end()
        assert sum(canvas.pixelColor(x,y).lightness()<160 for x in range(400) for y in range(100)) > 100
    finally:
        app.setPalette(original)


def test_word_picture_fits_merged_cell_width_with_full_original_bytes(app, tmp_path):
    item, text, assets, data = illustrated(tmp_path / 'source')
    projection = source_table_projection(item, text)
    token = deepcopy(projection['cell_tokens'][(2,1)][1])
    assert token['kind']=='image'
    token.update(width_mm=400, height_mm=200)
    document = Document()
    rows = [row('甲','乙','丙'), row('左侧',(2,'【待查看原文：图片或图形】'))]
    table = write_word_table(document, rows, cell_tokens={(1,1):[token]},
        assets={a['asset_id']:a for a in assets}, image_data=data)
    shape = document.inline_shapes[0]
    assert shape.width.mm == pytest.approx(table.cell(1,1).width.mm - 2*108*25.4/1440, abs=.001)
    assert shape.width / shape.height == pytest.approx(2, abs=.001)
    image = table._tbl.xpath('.//a:blip')[0]
    assert document.part.related_parts[image.get(qn('r:embed'))].blob == data[token['asset_id']]


def test_new_render_revision_changes_fingerprint_without_rewriting_saved_design(monkeypatch):
    from test_lesson_design import sample
    from integrations.deeptutor_shchem_v1 import desktop_lesson_composition as composition_module
    from integrations.deeptutor_shchem_v1.desktop_lesson_design import content_fingerprint, validate_design
    payload = sample(); payload['lesson_design'] = composition_module.enable_composition(payload['lesson_design'])
    saved = deepcopy(payload); assert composition_module.RENDER_REVISION == 'shchem.composition-classroom-exposition.v37'
    new = content_fingerprint(payload)
    monkeypatch.setattr(composition_module, 'RENDER_REVISION', 'shchem.composition-classroom-exposition.v36')
    old = content_fingerprint(payload)
    assert old != new and payload == saved
    assert validate_design(saved['lesson_design']) == saved['lesson_design']
