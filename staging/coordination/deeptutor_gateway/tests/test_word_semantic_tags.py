from __future__ import annotations

import io
import json
from contextlib import contextmanager
from copy import deepcopy
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from test_desktop_visual_import_facade import _facade, _png
from test_desktop_visual_import_facade import (
    desktop_paths as desktop_paths,  # noqa: PLC0414
)

from integrations.deeptutor_shchem_v1.desktop_facade import DesktopFacadeError
from integrations.deeptutor_shchem_v1.desktop_word_question_attributes import (
    WordQuestionAttributeError,
    WordQuestionAttributeStore,
    _seal,
    suggest_attributes,
)
from integrations.deeptutor_shchem_v1.desktop_word_questions import WordQuestionError
from integrations.deeptutor_shchem_v1.desktop_word_semantic_tags import (
    OUTPUT_SCHEMA,
    REVISION,
    WordSemanticTagError,
    compact_catalog,
    merge_response,
    tag_prompt,
)
from integrations.deeptutor_shchem_v1.model_provider_probe import (
    ProbeTransportResponse,
)
from integrations.deeptutor_shchem_v1.model_provider_settings import (
    ModelProviderProbeContext,
)

CATALOG = {
    "knowledge_points": [
        {
            "id": "K05",
            "name": "金属及其化合物",
            "subtopics": ["金属及其化合物"],
        },
        {
            "id": "K10",
            "name": "水溶液中的离子平衡",
            "subtopics": ["电解质"],
        },
    ],
    "nodes": [
        {
            "node_key": "n-metal",
            "volume_id": "v1",
            "chapter_id": "c1",
            "section_title": "重要的金属化合物",
        },
        {
            "node_key": "n-solution",
            "volume_id": "v1",
            "chapter_id": "c2",
            "section_title": "酸碱中和与盐类水解",
        },
    ],
}


def _catalog() -> dict[str, Any]:
    return deepcopy(CATALOG)


def _question(key: str = "word-question-synthetic") -> dict[str, Any]:
    return {
        "key": key,
        "source_sha256": "a" * 64,
        "source_revision": "source-revision",
        "index_revision": "index-revision",
        "extraction_revision": "extraction-revision",
        "revision": "question-revision",
        "source_name": "synthetic.docx",
        "chapter": "",
        "question_blocks": [
            {
                "index": 1,
                "text": "金属钠与水反应，写出化学方程式。",
                "assets": [],
            }
        ],
        "answer_blocks": [],
        "context_blocks": [],
        "warnings": [],
    }


def _attributes(key: str = "word-question-synthetic") -> dict[str, Any]:
    return suggest_attributes(
        _question(key), {"source_name": "synthetic.docx"}, _catalog()
    )


def _unit(attributes: dict[str, Any] | None = None) -> dict[str, Any]:
    return {
        "attributes": deepcopy(attributes or _attributes()),
        "input": {
            "blocks": [
                {
                    "index": 1,
                    "kind": "question_text",
                    "text": "金属钠与水反应，写出化学方程式。",
                    "image_sha256s": [],
                }
            ],
            "current_primary": (attributes or _attributes())["primary_knowledge"],
            "current_curriculum": (attributes or _attributes())[
                "curriculum_candidates"
            ],
            "source_sha256": "a" * 64,
        },
    }


def _text_evidence(index: int = 1, quote: str = "金属钠与水") -> dict[str, Any]:
    return {"block_index": index, "quote": quote, "image_sha256": ""}


def _response(
    *,
    primary: str = "K05",
    section: str = "n-metal",
    evidence: dict[str, Any] | None = None,
    note: str = "synthetic candidate",
) -> dict[str, Any]:
    evidence = deepcopy(evidence or _text_evidence())
    return {
        "primary": {"id": primary, "evidence": [] if primary == "unknown" else [evidence]},
        "curriculum": []
        if section is None
        else [{"section_key": section, "evidence": [deepcopy(evidence)]}],
        "note": note,
    }


class SemanticProvider:
    def __init__(self, *, vision: bool = True, revision: str = "REV-SEMANTIC-1"):
        self.vision = vision
        self.revision = revision
        self.borrow_calls = 0
        self.policy_calls = 0

    def list_metadata(self) -> list[dict[str, Any]]:
        capabilities = ["text", "structured_output"]
        data_classes = ["synthetic_only", "question_text_redacted"]
        if self.vision:
            capabilities.append("vision")
            data_classes.append("source_page_image")
        return [
            {
                "profile_id": "semantic",
                "revision": self.revision,
                "credential_state": "configured",
                "display_name": "Synthetic semantic model",
                "model_id": "semantic-fixture",
                "api_style": "responses",
                "effective_capabilities": capabilities,
                "allowed_data_classes": data_classes,
                "image_egress": (
                    "teacher_confirmed_source_pages" if self.vision else "deny"
                ),
            }
        ]

    @contextmanager
    def borrow_invocation_context(
        self, profile_id: str, *, expected_revision: str
    ):
        assert profile_id == "semantic"
        assert expected_revision == self.revision
        self.borrow_calls += 1
        yield ModelProviderProbeContext(
            profile_id=profile_id,
            provider_id="openai_compatible",
            model_id="semantic-fixture",
            base_url_policy="openai_compatible_public_https_v1",
            base_url="https://example.com/v1",
            revision=self.revision,
            api_key="test-only",
            provider_kind="openai_compatible",
            api_style="responses",
            local_endpoint_policy="deny",
        )

    def invocation_policy(self, profile_id: str, *, expected_revision: str):
        assert profile_id == "semantic"
        assert expected_revision == self.revision
        self.policy_calls += 1
        capabilities = ["text", "structured_output"]
        declared = ["text", "structured_output"]
        data_classes = ["synthetic_only", "question_text_redacted"]
        if self.vision:
            capabilities.append("vision")
            declared.append("vision")
            data_classes.append("source_page_image")
        return {
            "effective_capabilities": capabilities,
            "allowed_data_classes": data_classes,
            "image_egress": (
                "teacher_confirmed_source_pages" if self.vision else "deny"
            ),
            "capability_evidence": {"declared": declared, "catalog": declared},
        }


class SemanticTransport:
    def __init__(
        self,
        *,
        primary_ids: list[str] | None = None,
        section_keys: list[str | None] | None = None,
        fail_calls: set[int] | None = None,
        image_evidence: bool = False,
    ) -> None:
        self.primary_ids = primary_ids or ["K05"]
        self.section_keys = section_keys or ["n-metal"]
        self.fail_calls = fail_calls or set()
        self.image_evidence = image_evidence
        self.calls: list[dict[str, Any]] = []

    def send(
        self,
        request: Any,
        *,
        cancel_event: Any,
        deadline_monotonic: float,
    ) -> ProbeTransportResponse:
        del cancel_event, deadline_monotonic
        body = json.loads(request.body)
        self.calls.append({"request": request, "body": body})
        call_number = len(self.calls)
        if call_number in self.fail_calls:
            return ProbeTransportResponse(
                http_status=503,
                content_type="application/json",
                content_encoding=None,
                body=b"{}",
                latency_ms=1,
                model_invoked=True,
            )
        content = body["input"][0]["content"]
        prompt = content[0]["text"]
        input_payload = json.loads(prompt.split("本题资料：", 1)[1])
        block = input_payload["blocks"][0]
        primary = self.primary_ids[min(call_number - 1, len(self.primary_ids) - 1)]
        section = self.section_keys[min(call_number - 1, len(self.section_keys) - 1)]
        if self.image_evidence and block["image_sha256s"]:
            evidence = {
                "block_index": block["index"],
                "quote": "图中原题示意",
                "image_sha256": block["image_sha256s"][0],
            }
        else:
            evidence = {
                "block_index": block["index"],
                "quote": block["text"][: min(12, len(block["text"]))],
                "image_sha256": "",
            }
        decoded = _response(
            primary=primary, section=section, evidence=evidence, note="synthetic"
        )
        payload = {
            "status": "completed",
            "error": None,
            "incomplete_details": None,
            "output_text": json.dumps(decoded, ensure_ascii=False),
            "usage": {"input_tokens": 11, "output_tokens": 7, "total_tokens": 18},
        }
        return ProbeTransportResponse(
            http_status=200,
            content_type="application/json",
            content_encoding=None,
            body=json.dumps(payload, ensure_ascii=False).encode(),
            latency_ms=1,
            model_invoked=True,
        )


def _make_facade(
    desktop_paths,
    tmp_path: Path,
    provider: SemanticProvider,
    transport: SemanticTransport,
    *,
    image: bool = False,
    missing_image: bool = False,
    question_count: int = 3,
):
    doc = Document()
    doc.add_heading("知识点01 资料", 1)
    questions = [
        ("【即学即练1】下列物质中属于电解质的是（ ）", "A．铜  B．氯化钠  C．乙醇  D．蔗糖"),
        ("【典例2】请写出金属钠与水反应的化学方程式。", ""),
        ("【练习3】计算溶液的物质的量浓度。", ""),
    ]
    for number, (stem, options) in enumerate(questions[:question_count]):
        doc.add_paragraph(stem)
        if options:
            doc.add_paragraph(options)
        if missing_image and number == 1:
            doc.add_paragraph("【待查看原文：图片或图形】")
        elif image and number == 1:
            doc.add_paragraph().add_run().add_picture(io.BytesIO(_png("blue")))
        doc.add_paragraph("【答案】B")
    path = tmp_path / "semantic-tags.docx"
    doc.save(path)
    facade = _facade(desktop_paths, provider, transport=transport)
    facade._preparation_transport = transport
    facade.save_visual_import_batch(handout_files=(path,), source_type="教师讲义")
    words = facade._word_questions()
    words._read_attribute_catalog = lambda: _catalog()
    return facade


def _choices(facade, count: int | None = None) -> list[dict[str, Any]]:
    rows = facade.word_question_catalog()["items"]
    if count is not None:
        rows = rows[:count]
    return [
        {"key": row["key"], "revision": row["revision"], "points": 3}
        for row in rows
    ]


def test_prompt_binds_visual_input_order_to_actual_hashes() -> None:
    unit = _unit()
    unit["images"] = [{"sha256": "d" * 64}, {"sha256": "c" * 64}]
    prompt = tag_prompt(unit, CATALOG)
    image_manifest = prompt.split("本次图片顺序：", 1)[1].split("\n", 1)[0]
    assert json.loads(image_manifest) == unit["images"]
    assert "相同SHA的重复区块引用" in prompt


def test_tag_prompt_contains_only_exact_input_evidence_and_closed_catalog() -> None:
    unit = _unit()
    prompt = tag_prompt(unit, _catalog())
    assert "金属钠与水反应，写出化学方程式。" in prompt
    assert "quote逐字摘录" in prompt
    assert "只从给定目录选择ID" in prompt
    assert "api_key" not in prompt and "C:\\Users" not in prompt
    payload = json.loads(prompt.split("本题资料：", 1)[1])
    assert payload["blocks"][0]["text"] == unit["input"]["blocks"][0]["text"]
    compact = compact_catalog(_catalog())
    assert {item["id"] for item in compact["knowledge_points"]} == {"K05", "K10"}
    assert {item["section_key"] for item in compact["textbook_sections"]} == {
        "n-metal",
        "n-solution",
    }


@pytest.mark.parametrize(
    ("mutator", "message"),
    [
        (lambda value: value["primary"].update(id="K99"), "目录之外的知识点"),
        (
            lambda value: value["curriculum"][0].update(section_key="n-unknown"),
            "目录之外的教材章节",
        ),
        (
            lambda value: value["primary"]["evidence"][0].update(quote="不在题面"),
            "文字依据不在本题原文",
        ),
        (
            lambda value: value["primary"]["evidence"][0].update(image_sha256="b" * 64),
            "未发送的题图",
        ),
    ],
)
def test_merge_response_rejects_untrusted_taxonomy_or_evidence(mutator, message) -> None:
    value = _response()
    mutator(value)
    with pytest.raises(WordSemanticTagError, match=message):
        merge_response(value, _unit(), _catalog())


def test_merge_response_unknown_is_explicit_and_does_not_invent_evidence() -> None:
    unit = _unit()
    unknown = _response(primary="unknown", section=None)
    assert merge_response(unknown, unit, _catalog()) == unit["attributes"]
    with pytest.raises(WordSemanticTagError, match="未知主考点"):
        merge_response(
            {
                **unknown,
                "primary": {"id": "unknown", "evidence": [_text_evidence()]},
            },
            unit,
            _catalog(),
        )


def test_missing_only_merge_preserves_nonempty_teacher_exam_fields_and_revision() -> None:
    existing = _attributes()
    existing["curriculum_candidates"] = [
        {
            "section_key": "n-solution",
            "chapter_id": "c2",
            "volume_id": "v1",
            "label": "酸碱中和与盐类水解",
            "status": "teacher_confirmed",
            "evidence": [
                {
                    "kind": "teacher_note",
                    "quote": "教师已保留的教材映射",
                    "block_index": 1,
                }
            ],
        }
    ]
    existing["curriculum_status"] = "teacher_confirmed"
    existing["original_source"]["exam_type"] = {
        "value": "first_mock",
        "status": "source_observed",
        "evidence": [],
    }
    existing["original_source"]["display_label"] = "讲义收录题·一模"
    existing["teacher_note"] = "教师保留的说明"
    existing = _seal(existing)
    unit = _unit(existing)
    result = merge_response(_response(), unit, _catalog())
    assert result["primary_knowledge"]["id"] == "K05"
    assert result["curriculum_candidates"][0]["section_key"] == "n-solution"
    assert result["original_source"]["exam_type"]["value"] == "first_mock"
    assert result["original_source"]["display_label"] == "讲义收录题·一模"
    assert result["teacher_note"] == "教师保留的说明"
    assert result["rule_revision"] == REVISION

    protected = deepcopy(existing)
    protected["annotation_source"] = "teacher_modified"
    protected = _seal(protected)
    assert merge_response(_response(), _unit(protected), _catalog()) == protected


def test_attribute_store_compare_and_swap_is_atomic_and_keeps_history(tmp_path: Path) -> None:
    store = WordQuestionAttributeStore(tmp_path / "attributes")
    first = _attributes("q1")
    second = _attributes("q2")
    assert store.save_many([first, second])
    first_current = store.get("q1")
    second_current = store.get("q2")
    assert first_current is not None and second_current is not None

    first_update = deepcopy(first_current)
    first_update["teacher_note"] = "updated q1"
    first_update = _seal(first_update)
    saved = store.save_many(
        [first_update], expected_revisions={"q1": first_current["revision"]}
    )
    assert saved[0]["edit_version"] == 2
    assert len(store.history("q1")) == 2

    stale_first = deepcopy(first_update)
    stale_first["teacher_note"] = "should not commit"
    stale_first = _seal(stale_first)
    stale_second = deepcopy(second_current)
    stale_second["teacher_note"] = "should not commit"
    stale_second = _seal(stale_second)
    with pytest.raises(WordQuestionAttributeError, match="标签已被修改"):
        store.save_many(
            [stale_first, stale_second],
            expected_revisions={"q1": first_current["revision"], "q2": "stale"},
        )
    assert store.get("q1")["teacher_note"] == "updated q1"
    assert store.get("q2")["teacher_note"] == second_current["teacher_note"]
    assert len(store.history("q1")) == 2
    assert len(store.history("q2")) == 1


def test_preview_is_read_only_and_does_not_transport_or_create_attribute_store(
    desktop_paths, tmp_path: Path
) -> None:
    provider = SemanticProvider()
    transport = SemanticTransport()
    facade = _make_facade(desktop_paths, tmp_path, provider, transport, question_count=1)
    choices = _choices(facade, 1)
    store_path = facade._word_questions().attribute_store.path
    plan = facade.word_semantic_tag_preview(choices, "semantic", provider.revision)
    assert plan["request_count"] == 1
    assert transport.calls == []
    assert provider.borrow_calls == 0
    assert not store_path.exists()


def test_confirmed_text_run_uses_existing_structured_text_builder_and_apply_is_explicit(
    desktop_paths, tmp_path: Path
) -> None:
    provider = SemanticProvider()
    transport = SemanticTransport()
    facade = _make_facade(desktop_paths, tmp_path, provider, transport, question_count=2)
    choices = _choices(facade, 2)
    plan = facade.word_semantic_tag_preview(choices, "semantic", provider.revision)
    with pytest.raises(WordSemanticTagError, match="查看发送预览"):
        facade.word_semantic_tag_run(
            plan["plan_id"], plan["revision"], confirmed=False
        )
    assert transport.calls == []
    result = facade.word_semantic_tag_run(
        plan["plan_id"], plan["revision"], confirmed=True
    )
    assert [item["status"] for item in result["items"]] == ["ready", "ready"]
    body = transport.calls[0]["body"]
    content = body["input"][0]["content"]
    assert [item["type"] for item in content] == ["input_text"]
    assert "金属钠" in content[0]["text"] or "电解质" in content[0]["text"]
    assert body["text"]["format"]["schema"] == OUTPUT_SCHEMA
    assert facade._word_questions().attribute_store.get(choices[0]["key"]) is None

    saved = facade.word_semantic_tag_apply(plan["plan_id"], [choices[0]["key"]])
    assert [item["key"] for item in saved] == [choices[0]["key"]]
    assert facade._word_questions().attribute_store.get(choices[0]["key"]) is not None
    assert facade._word_questions().attribute_store.get(choices[1]["key"]) is None
    with pytest.raises(WordSemanticTagError, match="所选结果没有可保存"):
        facade.word_semantic_tag_apply(plan["plan_id"], ["not-current-result"])


def test_confirmed_image_run_uses_visual_builder_and_actual_image_sha(
    desktop_paths, tmp_path: Path
) -> None:
    provider = SemanticProvider()
    transport = SemanticTransport(image_evidence=True)
    facade = _make_facade(
        desktop_paths, tmp_path, provider, transport, image=True, question_count=2
    )
    choices = _choices(facade, 2)
    image_choice = choices[1]
    plan = facade.word_semantic_tag_preview(choices, "semantic", provider.revision)
    image_units = [unit for unit in plan["units"] if unit["key"] == image_choice["key"]]
    assert image_units and image_units[0]["images"]
    sha = image_units[0]["images"][0]["sha256"]
    assert facade.word_semantic_tag_image(plan["plan_id"], sha) == _png("blue")
    facade.word_semantic_tag_run(plan["plan_id"], plan["revision"], confirmed=True)
    image_call = transport.calls[1]
    content = image_call["body"]["input"][0]["content"]
    assert [item["type"] for item in content] == ["input_text", "input_image"]
    assert sha in content[0]["text"]
    assert content[1]["image_url"].startswith("data:image/png;base64,")


def test_missing_vision_policy_blocks_image_without_text_fallback(
    desktop_paths, tmp_path: Path
) -> None:
    provider = SemanticProvider(vision=False)
    transport = SemanticTransport()
    facade = _make_facade(
        desktop_paths, tmp_path, provider, transport, image=True, question_count=2
    )
    plan = facade.word_semantic_tag_preview(
        _choices(facade, 2), "semantic", provider.revision
    )
    blocked = [unit for unit in plan["units"] if unit["status"] == "blocked"]
    assert blocked
    assert any("读图" in unit["reason"] for unit in blocked)
    assert plan["request_count"] == 1
    result = facade.word_semantic_tag_run(
        plan["plan_id"], plan["revision"], confirmed=True
    )
    assert result["items"] and all(item["status"] == "ready" for item in result["items"])
    assert len(transport.calls) == 1
    assert all(
        not any(part.get("type") == "input_image" for part in call["body"]["input"][0]["content"])
        for call in transport.calls
    )


def test_missing_image_placeholder_blocks_without_text_fallback(
    desktop_paths, tmp_path: Path
) -> None:
    provider = SemanticProvider()
    transport = SemanticTransport()
    facade = _make_facade(
        desktop_paths,
        tmp_path,
        provider,
        transport,
        missing_image=True,
        question_count=2,
    )
    choices = _choices(facade, 2)
    plan = facade.word_semantic_tag_preview(choices, "semantic", provider.revision)
    blocked = [unit for unit in plan["units"] if unit["status"] == "blocked"]
    assert blocked and any("核对原Word" in unit["reason"] for unit in blocked)
    assert plan["request_count"] == 1
    result = facade.word_semantic_tag_run(
        plan["plan_id"], plan["revision"], confirmed=True
    )
    assert result["finished"] is True
    assert len(transport.calls) == 1
    assert all(
        not any(
            part.get("type") == "input_image"
            for part in call["body"]["input"][0]["content"]
        )
        for call in transport.calls
    )


def test_stale_profile_catalog_and_attributes_reject_before_transport(
    desktop_paths, tmp_path: Path
) -> None:
    provider = SemanticProvider()
    transport = SemanticTransport()
    facade = _make_facade(desktop_paths, tmp_path, provider, transport, question_count=1)
    choice = _choices(facade, 1)[0]
    plan = facade.word_semantic_tag_preview([choice], "semantic", provider.revision)

    provider.revision = "REV-SEMANTIC-2"
    with pytest.raises(DesktopFacadeError, match="模型配置已变化"):
        facade.word_semantic_tag_run(plan["plan_id"], plan["revision"], confirmed=True)
    assert transport.calls == []

    provider.revision = "REV-SEMANTIC-1"
    plan = facade.word_semantic_tag_preview([choice], "semantic", provider.revision)
    facade._word_questions()._read_attribute_catalog = lambda: {
        **_catalog(),
        "nodes": [{**_catalog()["nodes"][0], "section_title": "changed"}],
    }
    with pytest.raises(WordSemanticTagError, match="重新预览"):
        facade.word_semantic_tag_run(plan["plan_id"], plan["revision"], confirmed=True)
    assert transport.calls == []

    facade._word_questions()._read_attribute_catalog = lambda: _catalog()
    plan = facade.word_semantic_tag_preview([choice], "semantic", provider.revision)
    row, _ = facade._word_questions()._resolve([choice])
    attr = suggest_attributes(row[0], {"source_name": row[0]["source_name"]}, _catalog())
    facade._word_questions().attribute_store.save_many([attr])
    with pytest.raises(WordSemanticTagError, match="重新预览"):
        facade.word_semantic_tag_run(plan["plan_id"], plan["revision"], confirmed=True)
    assert transport.calls == []


def test_stale_source_is_rejected_without_transport(desktop_paths, tmp_path: Path) -> None:
    provider = SemanticProvider()
    transport = SemanticTransport()
    facade = _make_facade(desktop_paths, tmp_path, provider, transport, question_count=1)
    choice = _choices(facade, 1)[0]
    plan = facade.word_semantic_tag_preview([choice], "semantic", provider.revision)
    archive = (
        desktop_paths.state_root
        / "visual-import-v2"
        / "sources"
        / f"{facade.word_question_catalog()['items'][0]['source_sha256']}.docx"
    )
    archive.write_bytes(b"stale source fixture")
    with pytest.raises((WordQuestionError, WordSemanticTagError)):
        facade.word_semantic_tag_run(plan["plan_id"], plan["revision"], confirmed=True)
    assert transport.calls == []


def test_run_stops_on_failure_and_retains_earlier_candidates(
    desktop_paths, tmp_path: Path
) -> None:
    provider = SemanticProvider()
    transport = SemanticTransport(
        primary_ids=["K05", "K05", "K05"],
        section_keys=["n-metal", "n-metal", "n-metal"],
        fail_calls={2},
    )
    facade = _make_facade(desktop_paths, tmp_path, provider, transport, question_count=3)
    choices = _choices(facade, 3)
    plan = facade.word_semantic_tag_preview(choices, "semantic", provider.revision)
    result = facade.word_semantic_tag_run(
        plan["plan_id"], plan["revision"], confirmed=True
    )
    assert len(transport.calls) == 2
    assert [item["status"] for item in result["items"]] == ["ready", "failed"]
    assert result["items"][0]["changed"] is True
    assert result["items"][1]["changed"] is False
    assert result["finished"] is True
    assert facade._word_questions().attribute_store.get(choices[0]["key"]) is None
