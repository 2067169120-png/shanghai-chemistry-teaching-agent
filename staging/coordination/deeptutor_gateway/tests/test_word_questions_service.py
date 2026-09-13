import io
from pathlib import Path

import pytest
from docx import Document
from test_desktop_visual_import_facade import FakeProviderStore, _facade, _png
from test_desktop_visual_import_facade import (
    desktop_paths as desktop_paths,  # noqa: PLC0414
)

from integrations.deeptutor_shchem_v1.desktop_word_questions import WordQuestionError


def _import(desktop_paths, tmp_path, *, image=False, long=False):
    doc = Document()
    doc.add_heading("知识点01 电解质", 1)
    doc.add_paragraph("1．电解质定义知识总结，不是题目。")
    doc.add_paragraph("【即学即练1】下列物质中属于电解质的是（ ）")
    doc.add_paragraph("A．铜  B．氯化钠  C．乙醇  D．蔗糖")
    if image:
        doc.add_paragraph().add_run().add_picture(io.BytesIO(_png("blue")))
    if long:
        doc.add_paragraph("需保留的必要条件" * 3000)
    doc.add_paragraph("【答案】B")
    doc.add_paragraph("【解析】氯化钠属于电解质。")
    doc.add_heading("知识点02 非电解质", 1)
    doc.add_paragraph("【典例2】请写出水的化学式。")
    doc.add_paragraph("【答案】H2O")
    path = tmp_path / "演示讲义.docx"
    doc.save(path)
    provider = FakeProviderStore(configured=False)
    facade = _facade(desktop_paths, provider)
    facade.save_visual_import_batch(handout_files=(path,), source_type="教师讲义")
    return facade, path, provider


def _choice(item):
    return {"key": item["key"], "revision": item["revision"], "points": 3}


def test_specific_questions_persist_and_compile_without_api(desktop_paths, tmp_path):
    facade, path, provider = _import(desktop_paths, tmp_path)
    items = facade.word_question_catalog()["items"]
    assert len(items) == 2
    selected = [_choice(items[1])]
    facade.word_question_save_selection(selected)
    restarted = _facade(desktop_paths, provider)
    assert restarted.word_question_saved_selection() == selected
    ref = restarted.word_question_reference(selected)
    assert "写出水" in ref["materials"] and "H2O" in ref["materials"]
    assert "下列物质中" not in ref["materials"]
    assert "本次练习分值：3分" in ref["materials"]
    assert str(tmp_path) not in ref["materials"]
    assert provider.borrow_calls == 0
    both = restarted.word_question_reference([_choice(item) for item in items])
    assert "共同材料：与前面" not in both["materials"]
    # Reimporting exactly the same source does not double the question count.
    facade.save_visual_import_batch(handout_files=(path,), source_type="教师讲义")
    assert len(facade.word_question_catalog()["items"]) == 2


def test_revised_ranges_invalidate_previous_selections(desktop_paths, tmp_path):
    facade, _, provider = _import(desktop_paths, tmp_path)
    item = facade.word_question_catalog()["items"][0]
    updated = facade.word_question_update_range(
        item["key"],
        item["revision"],
        block_start=3,
        question_end=4,
        answer_start=5,
        block_end=5,
    )
    assert updated["key"] == item["key"] and updated["revision"] != item["revision"]
    with pytest.raises(WordQuestionError, match="变化"):
        facade.word_question_reference([_choice(item)])
    restarted = _facade(desktop_paths, provider)
    assert (
        restarted.word_question_catalog()["items"][0]["revision"] == updated["revision"]
    )
    assert (
        "【解析】"
        not in restarted.word_question_reference([_choice(updated)])["materials"]
    )


def test_archive_changes_are_not_masked_by_cache(desktop_paths, tmp_path):
    facade, _, _ = _import(desktop_paths, tmp_path)
    item = facade.word_question_catalog()["items"][0]
    archive = (
        desktop_paths.state_root
        / "visual-import-v2"
        / "sources"
        / f"{item['source_sha256']}.docx"
    )
    archive.write_bytes(b"changed test fixture")
    assert facade.word_question_catalog()["warnings"]
    with pytest.raises(WordQuestionError, match="变化"):
        facade.word_question_reference([_choice(item)])


def test_images_stay_bound_to_their_question_and_reference_warns(
    desktop_paths, tmp_path
):
    facade, _, _ = _import(desktop_paths, tmp_path, image=True)
    first, second = facade.word_question_catalog()["items"]
    asset = next(a for block in first["question_blocks"] for a in block["assets"])
    assert facade.word_question_image(
        first["key"], first["revision"], asset["asset_id"]
    )["bytes"] == _png("blue")
    with pytest.raises(WordQuestionError, match="不属于"):
        facade.word_question_image(second["key"], second["revision"], asset["asset_id"])
    reference = facade.word_question_reference([_choice(first)])
    assert any(
        "不调用模型" in warning and "后续是否发送图片像素" in warning
        and "发送预览" in warning
        for warning in reference["warnings"]
    )


def test_known_question_reads_only_its_source_but_rechecks_bytes(
    desktop_paths, tmp_path, monkeypatch
):
    facade, _, _ = _import(desktop_paths, tmp_path, image=True)
    other = Document()
    other.add_paragraph("【例1】请写出二氧化碳的化学式。")
    other.add_paragraph("【答案】CO2")
    other_path = tmp_path / "另一份教案.docx"
    other.save(other_path)
    facade.save_visual_import_batch(handout_files=(other_path,), source_type="教师讲义")
    items = facade.word_question_catalog()["items"]
    chosen = next(row for row in items if row["source_name"] == "演示讲义.docx")
    unrelated = next(row for row in items if row["source_name"] == other_path.name)
    calls = []
    restore = facade._restore_visual_import_sources

    def counted(descriptor, *, source_ids=None):
        calls.append(source_ids)
        return restore(descriptor, source_ids=source_ids)

    monkeypatch.setattr(facade, "_restore_visual_import_sources", counted)
    unrelated_archive = (
        desktop_paths.state_root
        / "visual-import-v2"
        / "sources"
        / f"{unrelated['source_sha256']}.docx"
    )
    unrelated_archive.write_bytes(b"unrelated damaged fixture")
    preview = facade.word_question_source(chosen["key"], chosen["revision"])
    assert preview["source_sha256"] == chosen["source_sha256"]
    assert calls == [{chosen["archive_source_id"]}]
    calls.clear()
    source_archive = (
        desktop_paths.state_root
        / "visual-import-v2"
        / "sources"
        / f"{chosen['source_sha256']}.docx"
    )
    source_archive.write_bytes(b"selected source changed")
    with pytest.raises(WordQuestionError, match="变化"):
        facade.word_question_source(chosen["key"], chosen["revision"])
    assert calls == [{chosen["archive_source_id"]}]


def test_clear_selection_does_not_read_archives(desktop_paths, tmp_path, monkeypatch):
    facade, _, _ = _import(desktop_paths, tmp_path)
    monkeypatch.setattr(
        facade,
        "_restore_visual_import_sources",
        lambda *args, **kwargs: pytest.fail("clearing choices must not read a DOCX"),
    )
    assert facade.word_question_save_selection([]) == []


def test_selection_limits_empty_clear_and_no_truncation(desktop_paths, tmp_path):
    facade, _, _ = _import(desktop_paths, tmp_path, long=True)
    item = facade.word_question_catalog()["items"][0]
    with pytest.raises(WordQuestionError, match="未截断"):
        facade.word_question_reference([_choice(item)])
    with pytest.raises(WordQuestionError, match="重复"):
        facade.word_question_save_selection([_choice(item), _choice(item)])
    with pytest.raises(WordQuestionError, match="分值"):
        facade.word_question_save_selection([dict(_choice(item), points=-1)])
    assert facade.word_question_save_selection([]) == []


def test_empty_source_does_not_hide_readable_questions(desktop_paths, tmp_path):
    facade, _, _ = _import(desktop_paths, tmp_path)
    empty = tmp_path / "空白讲义.docx"
    Document().save(empty)
    facade.save_visual_import_batch(handout_files=(empty,), source_type="教师讲义")
    catalog = facade.word_question_catalog()
    assert len(catalog["items"]) == 2
    assert any("空白讲义.docx" in text for text in catalog["warnings"])


def test_actual_selected_export_returns_openable_student_and_teacher_files(
    desktop_paths, tmp_path
):
    facade, _, provider = _import(desktop_paths, tmp_path)
    selected = facade.word_question_catalog()["items"][1]
    result = facade.word_question_export("水的化学式练习", [_choice(selected)])
    student = Document(result["student_path"])
    teacher = Document(result["teacher_path"])
    student_text = "\n".join(p.text for p in student.paragraphs)
    teacher_text = "\n".join(p.text for p in teacher.paragraphs)
    assert "请写出水" in student_text and "H2O" not in student_text
    assert "H2O" in teacher_text and "3 分" in teacher_text
    assert "属于电解质" not in student_text
    assert Path(result["student_path"]).is_relative_to(desktop_paths.state_root)
    assert provider.borrow_calls == 0
