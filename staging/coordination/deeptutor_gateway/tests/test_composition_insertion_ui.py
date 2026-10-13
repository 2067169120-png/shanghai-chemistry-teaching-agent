"""Native selectors and modal source completion exercise one batch transaction."""
from copy import deepcopy
from hashlib import sha256
import json
import os
from pathlib import Path
from types import SimpleNamespace

import pytest
from PySide6.QtCore import QTimer, Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QDialog, QLabel, QLineEdit, QVBoxLayout

from test_composition_render import app
from test_lesson_design import sample
from test_composition_insertion import fixture_compositions, text_resource
from test_textbook_crop_composition_handoff import (
    crop_fixture, section_fixture, asset_fixture, shelf_fixture, DeferredTasks)
from integrations.deeptutor_shchem_v1 import desktop_handout_material_handoff as handoff
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import enable_composition, resource_frames
from integrations.deeptutor_shchem_v1.desktop_lesson_design import DesignHistory, digest, new_node
from integrations.deeptutor_shchem_v1.desktop_lesson_output import COMPOSITION_FILES, checked_file, export_design
from integrations.deeptutor_shchem_v1.desktop_workbench.lesson_activity_dialogs import ActivityPicker
from integrations.deeptutor_shchem_v1.desktop_workbench.lesson_composition_panel import LessonCompositionPanel
from integrations.deeptutor_shchem_v1.desktop_workbench.studio_preparation import StudioPreparationPage


def evidence(widget, name, value=None):
    folder = os.getenv('SHCHEM_INSERTION_EVIDENCE')
    if folder:
        root = Path(folder); root.mkdir(parents=True, exist_ok=True)
        assert widget.grab().save(str(root/(name+'.png')))
        if value is not None: (root/(name+'.json')).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


class Assets:
    def __init__(self): self.value = []
    def assets(self): return deepcopy(self.value)
    def set_assets(self, rows): self.value = deepcopy(rows)


@pytest.fixture
def panel_desk(app, tmp_path):
    old, _, anchor = fixture_compositions()
    tails = old['placements'][anchor]['continuations']
    old['placements'][anchor] = {**tails[0], 'continuations': [
        {key: value for key, value in old['placements'][anchor].items() if key != 'continuations'}]}
    old['placements'][anchor]['width'] = old['placements'][anchor]['continuations'][0]['width']
    payload = sample(); payload['lesson_design'] = enable_composition(payload['lesson_design'])
    payload['lesson_design']['composition'] = old; payload['image_assets'] = []
    # Two genuine activity-picker rows, with full content requiring continuations.
    payload['lesson_design']['nodes'][0]['student_task'] = '讲义完整任务：观察各项条件并给出证据。\n'*35
    second = new_node('观察与提问'); second.update(title='第二项完整材料',
        student_task='第二材料全文，逐项讨论。\n'*32, teacher_action='完整教师讲解。\n'*8,
        minutes=10)
    payload['lesson_design']['nodes'].append(second)
    owner = SimpleNamespace(history=DesignHistory(payload['lesson_design']), status=QLabel(),
        tasks=SimpleNamespace(is_shutting_down=False), _stopping=False, _close_when_finished=False,
        facade=SimpleNamespace(paths=SimpleNamespace(task_root=tmp_path/'tasks'),
                               preparation_image_bytes=lambda _: pytest.fail('unexpected image read')),
        page=SimpleNamespace(image_assets_widget=Assets(), topic=QLineEdit('原课题')))
    window = QDialog(); window.resize(1480, 950); layout = QVBoxLayout(window)
    panel = LessonCompositionPanel(owner); layout.addWidget(panel)
    def sync(): panel.set_payload({**payload, 'lesson_design': owner.history.value})
    owner.sync_page = sync
    owner.travel = lambda redo=False: (owner.history.travel(redo), sync())
    sync(); panel.selected = anchor; panel.refresh(); window.show(); app.processEvents()
    yield owner, panel, window, app, payload, anchor
    window.close(); window.deleteLater(); app.processEvents()


def choose_mode(panel, mode):
    panel.insertion_position.setCurrentIndex(panel.insertion_position.findData(mode))
    assert panel.insertion_position.currentData() == mode


def assert_preserved(before, after):
    old = before['composition']; value = after['composition']
    assert [row for row in value['resources'] if row['id'] in old['placements']] == old['resources']
    assert {key: value['placements'][key] for key in old['placements']} == old['placements']
    assert [page for page in value['pages'] if page in old['pages']] == old['pages']
    for key in before:
        if key != 'composition': assert after[key] == before[key]


@pytest.mark.parametrize('mode', ['before', 'after', 'end'])
def test_visible_position_and_actual_activity_picker_add_full_batch_once(panel_desk, mode):
    owner, panel, window, app, _, anchor = panel_desk
    assert panel.insertion_position.isVisible()
    assert [panel.insertion_position.itemText(i) for i in range(3)] == ['作品末尾', '所选素材之前', '所选素材之后']
    choose_mode(panel, mode); before = deepcopy(owner.history.value); errors = []
    def complete_picker():
        try:
            picker = app.activeModalWidget(); assert isinstance(picker, ActivityPicker)
            for index in range(picker.list.count()): picker.list.item(index).setCheckState(Qt.CheckState.Checked)
            evidence(picker, 'actual-activity-picker-'+mode)
            QTest.mouseClick(picker.apply_button, Qt.MouseButton.LeftButton)
        except BaseException as exc:
            errors.append(exc)
            if app.activeModalWidget() is not None: app.activeModalWidget().reject()
    QTimer.singleShot(20, complete_picker)
    button = next(b for b in panel.findChildren(type(panel.step_flow_button)) if b.text() == '教学环节…')
    QTest.mouseClick(button, Qt.MouseButton.LeftButton); app.processEvents()
    assert not errors, errors
    after = deepcopy(owner.history.value); assert_preserved(before, after)
    old = before['composition']; value = after['composition']
    incoming = [row for row in value['resources'] if row['id'] not in old['placements']]
    assert len(incoming) == 2
    assert all(len(resource_frames(value['placements'][row['id']])) > 1 for row in incoming)
    new_pages = [page for page in value['pages'] if page not in old['pages']]
    resource_index = len(old['resources']) if mode == 'end' else 1 + (mode == 'after')
    page_index = len(old['pages']) if mode == 'end' else 1 if mode == 'before' else 3
    assert value['resources'][resource_index:resource_index+2] == incoming
    assert value['pages'][page_index:page_index+len(new_pages)] == new_pages
    assert len(owner.history.back) == 1
    assert not panel.layout_result['errors']
    evidence(window, 'native-batch-insertion-'+mode, {'mode': mode, 'old': old, 'new': value})
    owner.travel(False); assert owner.history.value == before
    owner.travel(True); assert owner.history.value == after


def test_selected_resource_is_frozen_before_picker_and_not_read_on_return(panel_desk):
    owner, panel, _, app, _, anchor = panel_desk
    choose_mode(panel, 'after'); before = deepcopy(owner.history.value)
    def complete():
        picker = app.activeModalWidget(); assert isinstance(picker, ActivityPicker)
        # A navigation change cannot silently move the captured destination.
        panel.selected = before['composition']['resources'][-1]['id']
        picker.list.item(0).setCheckState(Qt.CheckState.Checked); picker.apply_button.click()
    QTimer.singleShot(20, complete); panel.add_activities(); app.processEvents()
    value = owner.history.value['composition']
    assert value['resources'][2]['original']['kind'] == 'lesson_activity'
    assert value['pages'][3] not in before['composition']['pages']
    assert value['resources'][1]['id'] == anchor


@pytest.mark.parametrize('change', ['position', 'anchor-removed', 'history', 'closing', 'stopping', 'assets', 'new-source-action'])
def test_late_batch_is_blocked_without_overwriting_newer_state(panel_desk, change):
    owner, panel, _, _, _, anchor = panel_desk
    choose_mode(panel, 'after'); insertion = panel.freeze_insertion(); expected = digest(owner.history.value)
    if change == 'position': choose_mode(panel, 'before')
    elif change == 'anchor-removed': panel.remove_selected()
    elif change == 'history':
        later = deepcopy(owner.history.value); later['nodes'][0]['notes'] = '读取后的教师编辑'; owner.history.put(later)
    elif change == 'closing': owner._close_when_finished = True
    elif change == 'stopping': owner._stopping = True
    elif change == 'new-source-action': panel.freeze_insertion()
    else: owner.page.image_assets_widget.value = [{'changed': True}]
    before = deepcopy(owner.history.value); assets = owner.page.image_assets_widget.assets(); back = deepcopy(owner.history.back)
    with pytest.raises(ValueError): panel.add_result({'resources': [text_resource('迟到内容')]}, expected,
        expected_assets=[], insertion=insertion)
    assert owner.history.value == before and owner.history.back == back and owner.page.image_assets_widget.assets() == assets


def test_native_picker_cancel_and_duplicate_zero_result_keep_edits(panel_desk):
    owner, panel, _, app, _, _ = panel_desk
    choose_mode(panel, 'after'); before = deepcopy(owner.history.value)
    QTimer.singleShot(20, lambda: app.activeModalWidget().reject()); panel.add_activities()
    assert owner.history.value == before and not owner.history.back
    panel.add_result({'resources': []}, digest(before), insertion=panel.freeze_insertion())
    assert owner.history.value == before and not owner.history.back


def test_teacher_only_anchor_uses_physical_pages_while_student_flow_is_active(panel_desk):
    owner, panel, _, _, _, _ = panel_desk
    before = deepcopy(owner.history.value)
    teacher = before['composition']['resources'][3]['id']
    panel.selected = teacher; panel.refresh(); choose_mode(panel, 'after')
    assert panel.student_flow_active() and not panel.resource_display_frames(teacher)
    incoming = text_resource('新增课堂可见材料', '完整教师页后加入。\n'*36)
    panel.add_result({'resources': [incoming]}, digest(before), insertion=panel.freeze_insertion())
    value = owner.history.value['composition']; assert_preserved(before, owner.history.value)
    assert value['resources'][4] == incoming
    assert value['pages'][4] not in before['composition']['pages']
    assert len(owner.history.back) == 1
    owner.travel(False); assert owner.history.value == before


def test_batch_keeps_prior_six_outputs_and_history_record(panel_desk):
    from pptx import Presentation
    owner, panel, _, _, payload, _ = panel_desk
    payload = {**payload, 'lesson_design': deepcopy(owner.history.value)}
    record = export_design(owner.facade, payload)
    assert {row['name'] for row in record['files']} == set(COMPOSITION_FILES)
    owner.history.value['exports'].append(record)
    paths = {name: checked_file(owner.facade, owner.history.value, record['id'], name) for name in COMPOSITION_FILES}
    hashes = {name: sha256(path.read_bytes()).hexdigest() for name, path in paths.items()}
    before = deepcopy(owner.history.value); choose_mode(panel, 'after')
    panel.add_result({'resources': [text_resource('新批材料', '完整新材料。\n'*45), text_resource('新批第二项')]},
        digest(before), insertion=panel.freeze_insertion())
    assert owner.history.value['exports'] == before['exports']
    current = {**payload, 'lesson_design': deepcopy(owner.history.value)}
    latest = export_design(owner.facade, current)
    current['lesson_design']['exports'].append(latest)
    deck = Presentation(checked_file(owner.facade, current['lesson_design'], latest['id'], 'lesson_presentation.pptx'))
    value = owner.history.value['composition']
    texts = [''.join(shape.text for shape in slide.shapes if shape.has_text_frame) for slide in deck.slides]
    assert len(texts) == len(value['pages'])
    assert '新批材料' in texts[3]
    assert '新批第二项' not in texts[-2:]  # The later old teacher/last pages remain after the new batch.
    assert ''.join(texts).count('完整新材料。') == 45
    assert {name: sha256(path.read_bytes()).hexdigest() for name, path in paths.items()} == hashes
    owner.travel(False); assert owner.history.value == before


class Publication:
    def __init__(self, rows):
        self.result = {'resources': deepcopy(rows), 'image_assets': []}
        self.accepted = False; self.discarded = False
    def publish(self, guard): guard(); return deepcopy(self.result)
    def accept(self): self.accepted = True
    def discard(self): self.discarded = True


def material_request():
    source = {'schema': 'shchem.handout-knowledge-source.v1', 'batch_id': 'BATCH',
        'source_id': 'SOURCE', 'source_name': '合成讲义.docx', 'source_path': '合成讲义.docx',
        'source_sha256': 'a'*64, 'source_revision': 'b'*64, 'source_role': 'handout'}
    return {'schema': handoff.SCHEMA, 'source': source, 'knowledge_requests': [], 'image_selections': [],
        'question_selections': [{'key': 'Q1', 'revision': 'c'*64, 'points': 2,
            'source_sha256': 'a'*64, 'scoped_source_document': source}]}


@pytest.fixture
def studio_desk(app, crop_fixture, monkeypatch):
    tasks = DeferredTasks(); page = StudioPreparationPage(crop_fixture.facade, tasks)
    page._availability_timer.stop(); assert page.recovery is None
    old, _, anchor = fixture_compositions()
    tails = old['placements'][anchor]['continuations']
    old['placements'][anchor] = {**tails[0], 'continuations': [
        {key: value for key, value in old['placements'][anchor].items() if key != 'continuations'}]}
    old['placements'][anchor]['width'] = old['placements'][anchor]['continuations'][0]['width']
    page._lesson_design = sample()['lesson_design']; page._lesson_design = enable_composition(page._lesson_design)
    page._lesson_design['composition'] = old; page._editor_draft_id = 'existing-draft'
    page._lesson_editor_session = {'design_id': page._lesson_design['id'], 'resource_id': anchor}
    page.refresh_material_insertion_targets(); page.resize(1320, 920); page.show(); app.processEvents()
    publication = Publication([text_resource('完整讲义材料甲', '全文保留，不按小题截断。\n'*42), text_resource('完整讲义材料乙', '另一材料全文。\n'*38)])
    monkeypatch.setattr(handoff, 'prepare_handout_materials', lambda *_: publication)
    yield page, tasks, app, publication, anchor
    if page._material_handoff_dialog is not None: page._material_handoff_dialog.reject()
    page.close(); page.deleteLater(); app.processEvents()


def studio_choose_after(page, anchor):
    page.material_insertion_position.setCurrentIndex(page.material_insertion_position.findData('after'))
    page.material_insertion_target.setCurrentIndex(page.material_insertion_target.findData(anchor))


def test_visible_source_entry_freezes_named_target_and_adds_batch_after_all_teacher_pages(studio_desk):
    page, tasks, app, publication, anchor = studio_desk
    studio_choose_after(page, anchor)
    assert page.material_insertion_position.isVisible() and page.material_insertion_target.isVisible()
    assert page.material_insertion_target.currentText().startswith('2 · ')
    assert page.material_insertion_target.count() == 5
    evidence(page, 'source-entry-position')
    page.handout_button.click()
    frozen = deepcopy(page._source_material_insertion); assert frozen['anchor_resource_id'] == anchor
    assert page.import_handout_materials(material_request())
    dialog = page._material_handoff_dialog; before = deepcopy(dialog.history.value)
    key = dialog._task; tasks.deliver(key, tasks.calculate(key)); app.processEvents()
    after = deepcopy(dialog.history.value); assert_preserved(before, after)
    assert after['composition']['resources'][2:4] == publication.result['resources']
    assert after['composition']['pages'][3] not in before['composition']['pages']
    assert len(dialog.history.back) == 1 and publication.accepted
    evidence(dialog, 'source-entry-batch-after', {'frozen': frozen, 'after': after['composition']})
    dialog.travel(False); assert dialog.history.value == before
    dialog.travel(True); assert dialog.history.value == after


@pytest.mark.parametrize('change', ['position', 'target-material', 'target-draft', 'anchor-removed', 'target-signal', 'cancel', 'close'])
def test_source_late_changed_target_cancel_or_close_adds_nothing(studio_desk, change):
    page, tasks, app, publication, anchor = studio_desk
    studio_choose_after(page, anchor); page.handout_button.click()
    assert page.import_handout_materials(material_request())
    dialog = page._material_handoff_dialog; key = dialog._task; value = tasks.calculate(key)
    if change == 'position': page.material_insertion_position.setCurrentIndex(1)
    elif change == 'target-material': page.material_insertion_target.setCurrentIndex(4)
    elif change == 'target-draft': page._editor_draft_id = 'other-draft'
    elif change == 'anchor-removed':
        dialog.composition_panel.selected = anchor; dialog.composition_panel.remove_selected()
    elif change == 'target-signal': page.preparation_target_changed.emit()
    elif change == 'cancel': dialog.cancel()
    else: dialog.reject()
    before = deepcopy(page._lesson_design); tasks.deliver(key, value); app.processEvents()
    assert not publication.accepted
    assert page._lesson_design == before
    assert not any(row['id'] in {item['id'] for item in publication.result['resources']}
        for row in page._lesson_design['composition']['resources'])
    assert page.last_material_handoff['status'] == 'not_added'


def test_source_destination_changed_while_picker_open_rejects_before_any_task(studio_desk):
    page, tasks, _, publication, anchor = studio_desk
    studio_choose_after(page, anchor); page.handout_button.click()
    before = deepcopy(page._lesson_design)
    page.material_insertion_target.setCurrentIndex(4)
    assert not page.import_handout_materials(material_request())
    assert not tasks.pending and page._lesson_design == before and not publication.accepted


def test_first_source_batch_enables_position_choices_without_reopening_preparation(studio_desk):
    page, tasks, app, publication, _ = studio_desk
    page._lesson_design = None; page.refresh_material_insertion_targets()
    assert not page.material_insertion_position.model().item(1).isEnabled()
    assert page.import_handout_materials(material_request())
    dialog = page._material_handoff_dialog; key = dialog._task
    tasks.deliver(key, tasks.calculate(key)); app.processEvents()
    assert publication.accepted and page.material_insertion_target.count() == 2
    assert all(page.material_insertion_position.model().item(index).isEnabled() for index in (1, 2))


@pytest.fixture
def source_lifecycle_desk(studio_desk, tmp_path):
    """Real bookshelf and studio pages with the exact destination signal wiring."""
    from PySide6.QtWidgets import QStackedWidget
    from test_desktop_library_ui import _ManualBridge
    from integrations.deeptutor_shchem_v1.desktop_workbench.textbook_bookshelf_page import TextbookBookshelfPage
    from integrations.deeptutor_shchem_v1.desktop_state import DesktopStateStore
    page, _, app, _, anchor = studio_desk
    page.facade.paths.state_root = tmp_path/'source-lifecycle-isolated-state'
    page.facade._state = DesktopStateStore(page.facade.paths.state_root)
    bridge = _ManualBridge(); shelf = TextbookBookshelfPage(page.facade, bridge)
    stack = QStackedWidget(); stack.resize(1400, 950)
    stack.addWidget(page); stack.addWidget(shelf); stack.setCurrentWidget(page); stack.show()
    emitted = []; imported = []
    # Mirrors MainWindow._textbook_preparation_changed and navigate('textbooks').
    page.preparation_target_changed.connect(lambda: shelf.set_preparation_context(
        page._payload(), busy=page.studio_busy(), target_changed=True))
    def navigate_to_bookshelf():
        shelf.set_preparation_context(page._payload(), busy=page.studio_busy(), target_changed=False)
        stack.setCurrentWidget(shelf)
    page.textbook_requested.connect(navigate_to_bookshelf)
    def cards_to_preparation(result):
        emitted.append(deepcopy(result)); stack.setCurrentWidget(page)
        # Use the actual receiving method and actual native composition dialog.
        # Close the editor only after its modal loop starts; no panel is mocked.
        QTimer.singleShot(20, lambda: app.activeModalWidget().reject() if app.activeModalWidget() else None)
        accepted = page.import_textbook_cards(result); imported.append(accepted)
        (shelf.cards_appended if accepted else shelf.cards_rejected)()
    shelf.preparation_cards_requested.connect(cards_to_preparation)
    shelf.refresh(); bridge.succeed(bridge.take('读取教材书架'), run=True)
    available = next(row for row in shelf._snapshot['volumes'] if row['local_file_status'] == 'found')
    shelf.book.setCurrentIndex(shelf.book.findData(available['volume_id'])); app.processEvents()
    shelf.knowledge_list.item(0).setCheckState(Qt.CheckState.Checked)
    shelf.preview_button.click(); bridge.succeed(bridge.take('核对教材备课参考'), run=True)
    yield page, shelf, bridge, stack, emitted, imported, app, anchor
    shelf.close(); shelf.deleteLater(); stack.removeWidget(page); page.setParent(None)
    stack.close(); stack.deleteLater(); app.processEvents()


@pytest.mark.parametrize('change', ['position', 'target-material', 'same-position-reentry'])
def test_real_bookshelf_old_card_callback_cannot_borrow_reentered_destination(source_lifecycle_desk, change):
    page, shelf, bridge, stack, emitted, imported, app, anchor = source_lifecycle_desk
    studio_choose_after(page, anchor); page.textbook_button.click()
    assert stack.currentWidget() is shelf
    shelf.cards_button.click(); old = bridge.take('核对教材备课参考')
    result = old.operation(); assert result['cards']['resources']
    request_before, selection_before = shelf._reference_request, shelf._selection_key()
    epoch_before = shelf._preparation_target_epoch
    stack.setCurrentWidget(page)
    if change == 'position': page.material_insertion_position.setCurrentIndex(1)
    elif change == 'target-material': page.material_insertion_target.setCurrentIndex(4)
    page.textbook_button.click(); assert stack.currentWidget() is shelf
    assert shelf._reference_request == request_before and shelf._selection_key() == selection_before
    frozen = deepcopy(page._source_material_insertion); before = deepcopy(page._lesson_design)
    bridge.succeed(old, result); app.processEvents()
    assert emitted == [] and imported == []
    assert page._lesson_design == before and page._source_material_insertion == frozen
    assert shelf._preparation_target_epoch > epoch_before
    assert '备课已切换' in shelf.status.text()
    evidence(stack, 'source-old-callback-revoked-'+change, {'epoch_before': epoch_before,
        'epoch_after': shelf._preparation_target_epoch, 'frozen': frozen})


def test_real_bookshelf_single_entry_delivers_actual_knowledge_card_once(source_lifecycle_desk):
    page, shelf, bridge, stack, emitted, imported, app, anchor = source_lifecycle_desk
    studio_choose_after(page, anchor); before = deepcopy(page._lesson_design)
    page.textbook_button.click(); shelf.cards_button.click()
    task = bridge.take('核对教材备课参考'); result = task.operation()
    assert result['cards']['resources'][0]['original']['kind'] == 'knowledge_note'
    bridge.succeed(task, result); app.processEvents()
    assert len(emitted) == 1 and imported == [True]
    value = page._lesson_design['composition']
    assert len(value['resources']) == len(before['composition']['resources']) + 1
    assert value['resources'][2]['original'] == result['cards']['resources'][0]['original']
    assert_preserved(before, page._lesson_design)
    assert page._source_material_insertion is None
    assert shelf._cards_added_key == shelf._selection_key() and not shelf.cards_button.isEnabled()
