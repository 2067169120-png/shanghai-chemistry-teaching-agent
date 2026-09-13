from __future__ import annotations

import copy
import json
from pathlib import Path
from typing import Any

import pytest
from jsonschema import Draft202012Validator
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from integrations.deeptutor_shchem_v1.desktop_preparation import (
    DesktopPreparationError,
    _canonical_candidate_digest,
    _validate_canonical_candidate,
    normalize_preparation_candidate,
    preparation_candidate_schema,
)
from integrations.deeptutor_shchem_v1.desktop_preparation_renderer import (
    SLIDE_HEIGHT_PX,
    SLIDE_WIDTH_PX,
    NativePreparationRenderer,
    _FontBook,
    _json_bytes,
    _make_layouts,
    _wrap_text,
    _write_pptx,
)
from integrations.deeptutor_shchem_v1.desktop_preparation_visual import (
    ComparisonRowCountError,
    normalize_slide_visual,
)
from staging.coordination.deeptutor_gateway.tests.test_desktop_preparation import (
    _payload,
    _raw_candidate,
)


def _comparison_visual() -> dict[str, Any]:
    return {
        "kind": "comparison",
        "comparison": {
            "dimension_label": "维度",
            "columns": ["方案甲", "方案乙"],
            "rows": [
                {"label": "依据", "values": ["观察证据", "定量证据"]},
                {"label": "结论", "values": ["先描述", "再解释"]},
            ],
        },
        "steps": [],
    }


def _process_visual() -> dict[str, Any]:
    return {
        "kind": "process",
        "comparison": None,
        "steps": [
            {"label": "观察", "detail": "记录可见现象"},
            {"label": "建模", "detail": "用粒子模型解释"},
            {"label": "表达", "detail": "写出规范结论"},
        ],
    }


def _visual_slide(
    *,
    title: str,
    visual: dict[str, Any],
    content: list[str] | None = None,
) -> dict[str, Any]:
    return {
        "id": title,
        "title": title,
        "purpose": "检查显式可编辑版式",
        "content": content or ["保留普通文字内容"],
        "visual": visual,
        "teacher_notes": "教师复核后使用。",
    }


def _visual_pairs(layout: Any) -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    for element in layout.elements:
        if element.element_type != "text":
            continue
        assert len(element.source_texts) == len(element.source_paths)
        pairs.extend(zip(element.source_texts, element.source_paths, strict=True))
    return pairs


def _assert_layout_bounds(layout: Any) -> None:
    for element in layout.elements:
        assert element.width > 0
        assert element.height > 0
        assert 0 <= element.x <= SLIDE_WIDTH_PX - element.width
        assert 0 <= element.y <= SLIDE_HEIGHT_PX - element.height


def test_normalize_slide_visual_trims_without_mutating_input() -> None:
    source = {
        "kind": "comparison",
        "comparison": {
            "dimension_label": "  维度  ",
            "columns": ["  方案甲 ", "方案乙"],
            "rows": [
                {"label": "  依据 ", "values": ["  观察证据 ", "定量证据"]},
                {"label": "结论", "values": ["先描述", " 再解释 "]},
            ],
        },
        "steps": [],
    }
    original = copy.deepcopy(source)

    normalized = normalize_slide_visual(source)

    assert normalized == _comparison_visual()
    assert normalized is not source
    assert normalized["comparison"] is not source["comparison"]
    assert source == original

    normalized["comparison"]["columns"][0] = "已修改"
    assert source["comparison"]["columns"][0] == "  方案甲 "


def test_normalize_slide_visual_accepts_none_and_process() -> None:
    assert normalize_slide_visual(None) is None

    source = _process_visual()
    original = copy.deepcopy(source)
    normalized = normalize_slide_visual(source)

    assert normalized == source
    assert normalized is not source
    assert normalized["steps"] is not source["steps"]
    assert source == original


@pytest.mark.parametrize("cell_count", [0, 1, 3])
def test_comparison_row_count_error_is_precise_and_does_not_repair(cell_count) -> None:
    source = _comparison_visual()
    source["comparison"]["rows"][1]["values"] = ["不应回显的模型文本"] * cell_count
    original = copy.deepcopy(source)
    with pytest.raises(ComparisonRowCountError) as caught:
        normalize_slide_visual(source)
    assert f"第2行有{cell_count}格内容，但表头有2个数据列" in caught.value.message_zh
    assert "不应回显" not in caught.value.message_zh
    assert source == original


def test_candidate_reports_slide_and_row_for_three_headers_two_cells() -> None:
    candidate = _raw_candidate()
    visual = _comparison_visual()
    visual["comparison"]["columns"] = ["概念/对象", "判断依据或条件", "典型例式"]
    candidate["slides"][1]["visual"] = visual
    original = copy.deepcopy(candidate)
    with pytest.raises(DesktopPreparationError) as caught:
        normalize_preparation_candidate(candidate, _payload())
    assert caught.value.code == "preparation_candidate_visual_invalid"
    assert (
        "PPT页面2：比较表第1行有2格内容，但表头有3个数据列" in caught.value.message_zh
    )
    assert "拆分后重试" not in caught.value.message_zh
    assert candidate == original


def test_frozen_candidate_reports_comparison_row_count_without_weakening_hash_gate() -> (
    None
):
    candidate = normalize_preparation_candidate(_raw_candidate(), _payload())
    visual = _comparison_visual()
    visual["comparison"]["columns"] = ["概念/对象", "判断依据或条件", "典型例式"]
    candidate["slides"][1]["visual"] = visual
    with pytest.raises(DesktopPreparationError) as caught:
        _validate_canonical_candidate(candidate, _payload())
    assert caught.value.code == "preparation_seed_drift"
    candidate["candidate_id"] = (
        "PREPCAND-" + _canonical_candidate_digest(candidate)[:32]
    )
    original = copy.deepcopy(candidate)
    with pytest.raises(DesktopPreparationError) as caught:
        _validate_canonical_candidate(candidate, _payload())
    assert caught.value.code == "preparation_seed_invalid"
    assert (
        "冻结候选PPT页面2：比较表第1行有2格内容，但表头有3个数据列"
        in caught.value.message_zh
    )
    assert candidate == original


@pytest.mark.parametrize(
    "value",
    [
        [],
        "comparison",
        1,
        True,
        {"kind": "comparison", "comparison": None, "steps": []},
        {
            "kind": "comparison",
            "comparison": _comparison_visual()["comparison"],
            "steps": [],
            "extra": "拒绝未知字段",
        },
        {
            "kind": "comparison",
            "comparison": _comparison_visual()["comparison"],
            "steps": [{"label": "不应出现", "detail": "比较页不能有流程节点"}],
        },
        {
            "kind": "process",
            "comparison": _comparison_visual()["comparison"],
            "steps": _process_visual()["steps"],
        },
        {"kind": "unknown", "comparison": None, "steps": []},
        {"kind": "comparison", "comparison": {}, "steps": []},
        {
            "kind": "comparison",
            "comparison": {
                "dimension_label": "维度",
                "columns": ["仅一列"],
                "rows": [
                    {"label": "依据", "values": ["内容"]},
                    {"label": "结论", "values": ["内容"]},
                ],
            },
            "steps": [],
        },
        {
            "kind": "comparison",
            "comparison": {
                "dimension_label": "维度",
                "columns": ["一", "二", "三", "四"],
                "rows": [
                    {"label": "依据", "values": ["1", "2", "3", "4"]},
                    {"label": "结论", "values": ["1", "2", "3", "4"]},
                ],
            },
            "steps": [],
        },
        {
            "kind": "comparison",
            "comparison": {
                "dimension_label": "维度",
                "columns": ["一", "二"],
                "rows": [{"label": "只有一行", "values": ["1", "2"]}],
            },
            "steps": [],
        },
        {
            "kind": "comparison",
            "comparison": {
                "dimension_label": "维度",
                "columns": ["一", "二"],
                "rows": [
                    {"label": "1", "values": ["1", "2"]},
                    {"label": "2", "values": ["1", "2"]},
                    {"label": "3", "values": ["1", "2"]},
                    {"label": "4", "values": ["1", "2"]},
                    {"label": "5", "values": ["1", "2"]},
                ],
            },
            "steps": [],
        },
        {
            "kind": "comparison",
            "comparison": {
                "dimension_label": "维度",
                "columns": ["一", "二"],
                "rows": [
                    {"label": "依据", "values": ["只给一格"]},
                    {"label": "结论", "values": ["1", "2"]},
                ],
            },
            "steps": [],
        },
        {
            "kind": "comparison",
            "comparison": {
                "dimension_label": "维度",
                "columns": ["一", "二"],
                "rows": [
                    {"label": "依据", "values": ["1", "2", "3"]},
                    {"label": "结论", "values": ["1", "2"]},
                ],
            },
            "steps": [],
        },
        {
            "kind": "process",
            "comparison": None,
            "steps": [],
        },
        {
            "kind": "process",
            "comparison": None,
            "steps": [{"label": "只有一步", "detail": "不构成流程"}],
        },
        {
            "kind": "process",
            "comparison": None,
            "steps": [{"label": str(index), "detail": "节点"} for index in range(5)],
        },
        {
            "kind": "process",
            "comparison": None,
            "steps": [
                {"label": "观察", "detail": "节点"},
                {"label": "建模", "detail": "节点", "extra": "拒绝未知字段"},
            ],
        },
    ],
)
def test_normalize_slide_visual_rejects_closed_shape_and_cardinality_violations(
    value: Any,
) -> None:
    with pytest.raises(ValueError):
        normalize_slide_visual(value)


@pytest.mark.parametrize(
    "value",
    [
        {
            "kind": "comparison",
            "comparison": {
                "dimension_label": "",
                "columns": ["一", "二"],
                "rows": [
                    {"label": "依据", "values": ["1", "2"]},
                    {"label": "结论", "values": ["1", "2"]},
                ],
            },
            "steps": [],
        },
        {
            "kind": "comparison",
            "comparison": {
                "dimension_label": "x" * 33,
                "columns": ["一", "二"],
                "rows": [
                    {"label": "依据", "values": ["1", "2"]},
                    {"label": "结论", "values": ["1", "2"]},
                ],
            },
            "steps": [],
        },
        {
            "kind": "comparison",
            "comparison": {
                "dimension_label": "维度",
                "columns": ["一" * 33, "二"],
                "rows": [
                    {"label": "依据", "values": ["1", "2"]},
                    {"label": "结论", "values": ["1", "2"]},
                ],
            },
            "steps": [],
        },
        {
            "kind": "comparison",
            "comparison": {
                "dimension_label": "维度",
                "columns": ["一", "二"],
                "rows": [
                    {"label": "依" * 33, "values": ["1", "2"]},
                    {"label": "结论", "values": ["1", "2"]},
                ],
            },
            "steps": [],
        },
        {
            "kind": "comparison",
            "comparison": {
                "dimension_label": "维度",
                "columns": ["一", "二"],
                "rows": [
                    {"label": "依据", "values": ["x" * 121, "2"]},
                    {"label": "结论", "values": ["1", "2"]},
                ],
            },
            "steps": [],
        },
        {
            "kind": "process",
            "comparison": None,
            "steps": [
                {"label": "观察", "detail": "x" * 121},
                {"label": "建模", "detail": "节点"},
            ],
        },
        {
            "kind": "process",
            "comparison": None,
            "steps": [
                {"label": "x" * 33, "detail": "节点"},
                {"label": "建模", "detail": "节点"},
            ],
        },
    ],
)
def test_normalize_slide_visual_rejects_label_and_detail_limits(value: Any) -> None:
    with pytest.raises(ValueError):
        normalize_slide_visual(value)


def test_candidate_schema_requires_nullable_visual_and_accepts_explicit_null() -> None:
    schema = preparation_candidate_schema()
    Draft202012Validator.check_schema(schema)
    slide_schema = schema["$defs"]["slide"]
    assert "visual" in slide_schema["required"]
    visual_schema = slide_schema["properties"]["visual"]
    assert {branch.get("type") for branch in visual_schema["anyOf"]} >= {
        "null",
        "object",
    }
    visual_object = next(
        branch for branch in visual_schema["anyOf"] if branch.get("type") == "object"
    )
    assert visual_object["additionalProperties"] is False
    assert set(visual_object["required"]) == {"kind", "comparison", "steps"}

    candidate = copy.deepcopy(_raw_candidate())
    for activity in candidate["activities"]:
        activity["worksheet"] = None
    for slide in candidate["slides"]:
        slide["visual"] = None
        slide["image"] = None
    Draft202012Validator(schema).validate(candidate)


def test_normalize_candidate_defaults_legacy_visual_without_changing_canonical_shape() -> (
    None
):
    candidate = _raw_candidate()
    original = copy.deepcopy(candidate)

    normalized = normalize_preparation_candidate(candidate, _payload())

    assert candidate == original
    assert all("visual" not in slide for slide in normalized["slides"])


def test_normalize_candidate_keeps_only_nonempty_normalized_visual() -> None:
    candidate = _raw_candidate()
    candidate["slides"][0]["visual"] = {
        "kind": "process",
        "comparison": None,
        "steps": [
            {"label": "  观察 ", "detail": "记录现象"},
            {"label": "解释", "detail": "形成模型"},
        ],
    }
    original = copy.deepcopy(candidate)

    normalized = normalize_preparation_candidate(candidate, _payload())

    assert candidate == original
    assert normalized["slides"][0]["visual"] == {
        "kind": "process",
        "comparison": None,
        "steps": [
            {"label": "观察", "detail": "记录现象"},
            {"label": "解释", "detail": "形成模型"},
        ],
    }
    assert all("visual" not in slide for slide in normalized["slides"][1:])


def test_make_layouts_honors_visual_kind_on_first_and_last_pages_and_binds_sources() -> (
    None
):
    candidate = {
        "title": "显式版式测试",
        "topic": "化学推理",
        "slides": [
            _visual_slide(title="总结：比较两种证据", visual=_comparison_visual()),
            _visual_slide(title="练习：按步骤解释", visual=_process_visual()),
        ],
    }

    layouts = _make_layouts(candidate)

    assert [layout.slide_type for layout in layouts] == ["comparison", "process"]
    assert [layout.public()["layout_id"] for layout in layouts] == [
        "native.comparison.v1",
        "native.process.v1",
    ]
    for layout in layouts:
        _assert_layout_bounds(layout)

    expected_pairs = {
        # Page 1: comparison.
        ("维度", "/slides/0/visual/comparison/dimension_label"),
        ("方案甲", "/slides/0/visual/comparison/columns/0"),
        ("方案乙", "/slides/0/visual/comparison/columns/1"),
        ("依据", "/slides/0/visual/comparison/rows/0/label"),
        ("观察证据", "/slides/0/visual/comparison/rows/0/values/0"),
        ("定量证据", "/slides/0/visual/comparison/rows/0/values/1"),
        ("结论", "/slides/0/visual/comparison/rows/1/label"),
        ("先描述", "/slides/0/visual/comparison/rows/1/values/0"),
        ("再解释", "/slides/0/visual/comparison/rows/1/values/1"),
        # Page 2: process.
        ("观察", "/slides/1/visual/steps/0/label"),
        ("记录可见现象", "/slides/1/visual/steps/0/detail"),
        ("建模", "/slides/1/visual/steps/1/label"),
        ("用粒子模型解释", "/slides/1/visual/steps/1/detail"),
        ("表达", "/slides/1/visual/steps/2/label"),
        ("写出规范结论", "/slides/1/visual/steps/2/detail"),
    }
    actual_pairs = set(_visual_pairs(layouts[0]) + _visual_pairs(layouts[1]))
    assert expected_pairs <= actual_pairs
    assert all(
        path.startswith("/slides/")
        for _text, path in actual_pairs
        if "/visual/" in path
    )


def test_write_pptx_uses_editable_text_and_native_right_arrows(tmp_path: Path) -> None:
    candidate = {
        "title": "流程图可编辑性测试",
        "topic": "化学解释流程",
        "slides": [
            _visual_slide(title="流程页置于首尾也不降级", visual=_process_visual()),
        ],
    }
    layouts = _make_layouts(candidate)
    path = tmp_path / "visual.pptx"

    _write_pptx(path, layouts, candidate, root=tmp_path)

    presentation = Presentation(path)
    assert len(presentation.slides) == 1
    slide = presentation.slides[0]
    editable_text = [
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    ]
    all_text = "\n".join(editable_text)
    for text in (
        "流程页置于首尾也不降级",
        "观察",
        "记录可见现象",
        "建模",
        "用粒子模型解释",
        "表达",
        "写出规范结论",
    ):
        assert text in all_text
    assert editable_text
    assert all(
        shape.shape_type in {MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.AUTO_SHAPE}
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )

    right_arrows = [
        shape
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
        and shape.auto_shape_type == MSO_SHAPE.RIGHT_ARROW
    ]
    assert len(right_arrows) >= len(_process_visual()["steps"]) - 1


@pytest.mark.parametrize(
    "source",
    ["学习证据，需要核对。", "学习（观察记录）与解释。", "先观察，再讨论，最后核对。"],
)
def test_chinese_wrap_preserves_text_and_punctuation(source: str) -> None:
    draw = ImageDraw.Draw(Image.new("RGB", (500, 500)))
    font = _FontBook().get(32)
    width = int(draw.textlength("学习证据", font=font)) + 1
    wrapped = _wrap_text(draw, source, font, width)
    lines = wrapped.splitlines()
    assert len(lines) > 1
    assert "".join(lines) == source
    assert all(line[0] not in "，。）" for line in lines)
    assert all(line[-1] != "（" for line in lines)
    assert all(draw.textlength(line, font=font) <= width for line in lines)


def test_native_shapes_override_office_theme_effects(tmp_path: Path) -> None:
    candidate = {
        "title": "无主题阴影",
        "slides": [
            _visual_slide(title="比较", visual=_comparison_visual()),
            _visual_slide(title="顺序", visual=_process_visual()),
        ],
    }
    path = tmp_path / "effects.pptx"
    _write_pptx(path, _make_layouts(candidate), candidate, root=tmp_path)
    checked = 0
    for slide in Presentation(path).slides:
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                effects = shape._element.spPr.find(qn("a:effectLst"))
                assert effects is not None and len(effects) == 0
                checked += 1
    assert checked > 4


def test_normalized_visual_candidate_renders_without_rewriting_it(
    tmp_path: Path,
) -> None:
    raw = _raw_candidate()
    raw["slides"][0]["visual"] = _comparison_visual()
    raw["slides"][1]["visual"] = _process_visual()
    candidate = normalize_preparation_candidate(raw, _payload(output_kind="ppt"))
    before = copy.deepcopy(candidate)
    result = NativePreparationRenderer().render(
        candidate,
        output_kind="ppt",
        output_dir=tmp_path / "artifacts",
    )
    saved_bytes = (tmp_path / "artifacts" / "candidate.json").read_bytes()
    assert candidate == before == json.loads(saved_bytes)
    assert saved_bytes == _json_bytes(before)
    assert result["quality"]["machine_checks_passed"] is True
    assert not list((tmp_path / "artifacts").rglob("*.docx"))
    qa = json.loads(
        (tmp_path / "artifacts" / "qa_report.json").read_text(encoding="utf-8")
    )
    binding = next(
        row for row in qa["checks"] if row["check"] == "visible_candidate_text_binding"
    )
    assert binding["passed"] is True and binding["unbound"] == []
    deck = json.loads(
        (tmp_path / "artifacts" / "deck.json").read_text(encoding="utf-8")
    )
    assert [slide["slide_type"] for slide in deck["slides"][:2]] == [
        "comparison",
        "process",
    ]
