from __future__ import annotations

import hashlib
import io
from copy import deepcopy
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from integrations.deeptutor_shchem_v1 import desktop_preparation_sources as sources
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
    PreparationSourceError,
    PreparationSourcesService,
    _digest,
)
from integrations.deeptutor_shchem_v1.desktop_word_table_preview import (
    WordTablePreviewService,
)


def _bytes(document):
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _document(*, nested=False, tables=True):
    document = Document()
    document.add_paragraph("合成正文段落，不重提取。")
    if tables:
        table = document.add_table(rows=3, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "合成比较标题"
        table.cell(0, 2).text = "原条件"
        table.cell(1, 0).merge(table.cell(2, 0)).text = "同一条件"
        table.cell(1, 1).text = "甲"
        table.cell(1, 2).text = "乙"
        table.cell(2, 1).text = "丙"
        table.cell(2, 2).text = "丁"
        if nested:
            inner = table.cell(1, 1).add_table(rows=1, cols=2)
            inner.cell(0, 0).text = "内层左"
            inner.cell(0, 1).text = "内层右"
    document.add_paragraph("末段仍在原位置。")
    return document


class _Facade:
    def __init__(self, data, tmp_path):
        self.data = data
        self.filename = "合成原教案.docx"
        self.value = PreparationSourcesService(tmp_path).word_preview_bytes(
            data, self.filename
        )
        self.calls = []

    def _imported_word_source(self, batch_id, source_id):
        self.calls.append(("source", batch_id, source_id))
        return SimpleNamespace(content=self.data, filename=self.filename)

    def imported_word_preview(self, batch_id, source_id):
        self.calls.append(("preview", batch_id, source_id))
        return deepcopy(self.value)


def _preview(facade, *, sha=None, revision=None):
    return WordTablePreviewService(facade).preview(
        "batch-selected",
        "source-selected",
        sha or hashlib.sha256(facade.data).hexdigest(),
        revision or facade.value["revision"],
    )


def _rehash(facade):
    facade.value["revision"] = _digest(
        {key: value for key, value in facade.value.items() if key != "revision"}
    )


def test_selected_table_preserves_source_merges_and_original_cache(tmp_path):
    facade = _Facade(_bytes(_document()), tmp_path)
    original, cached = facade.data, deepcopy(facade.value)
    result = _preview(facade)

    assert result["source_sha256"] == hashlib.sha256(original).hexdigest()
    assert result["source_revision"] == cached["revision"]
    assert list(result["tables"]) == [2]
    rows = result["tables"][2]
    assert rows[0]["cells"][0] == {
        "column_span": 2,
        "vertical_merge": "none",
        "text": "合成比较标题",
    }
    assert rows[1]["cells"][0]["vertical_merge"] == "restart"
    assert rows[2]["cells"][0]["vertical_merge"] == "continue"
    assert rows[2]["cells"][0]["text"] == ""
    assert len(rows[2]["cells"]) == 3
    assert facade.calls == [
        ("source", "batch-selected", "source-selected"),
        ("preview", "batch-selected", "source-selected"),
    ]
    rows[0]["cells"][0]["text"] = "只修改独立显示对象"
    assert facade.data == original and facade.value == cached


def test_omitted_grid_edges_remain_source_metadata_not_invented_cells(tmp_path):
    document = _document()
    row = document.tables[0].rows[1]._tr
    row.remove(row.tc_lst[-1])
    row.remove(row.tc_lst[0])
    properties = row.get_or_add_trPr()
    for name in ("gridBefore", "gridAfter"):
        item = OxmlElement("w:" + name)
        item.set(qn("w:val"), "1")
        properties.append(item)
    facade = _Facade(_bytes(document), tmp_path)
    row = _preview(facade)["tables"][2][1]
    assert row["grid_before"] == row["grid_after"] == 1
    assert [cell["text"] for cell in row["cells"]] == ["甲"]


def test_nested_table_stays_explicit_in_cell_text_without_fake_flat_grid(tmp_path):
    facade = _Facade(_bytes(_document(nested=True)), tmp_path)
    result = _preview(facade)
    assert list(result["tables"]) == [2]
    nested_text = result["tables"][2][1]["cells"][1]["text"]
    assert nested_text.startswith("甲\n【表格开始")
    assert "内层左" in nested_text and "内层右" in nested_text
    assert "【表格结束】" in nested_text
    assert facade.value["blocks"][1]["text"].count("【表格开始") == 2


def test_paragraphs_are_not_reextracted_when_only_table_rows_are_requested(
    tmp_path, monkeypatch
):
    facade = _Facade(_bytes(_document()), tmp_path)
    reader = sources._paragraph
    seen = []

    def only_cell_paragraphs(element, *args, **kwargs):
        assert any(parent.tag == qn("w:tc") for parent in element.iterancestors())
        seen.append(element)
        return reader(element, *args, **kwargs)

    monkeypatch.setattr(sources, "_paragraph", only_cell_paragraphs)
    assert _preview(facade)["tables"]
    assert seen


def test_no_tables_returns_empty_projection_without_changing_preview(tmp_path):
    facade = _Facade(_bytes(_document(tables=False)), tmp_path)
    before = deepcopy(facade.value)
    assert _preview(facade)["tables"] == {}
    assert facade.value == before


def test_changed_source_hash_fails_before_reading_preview(tmp_path):
    facade = _Facade(_bytes(_document()), tmp_path)
    with pytest.raises(PreparationSourceError, match="原教案已变化"):
        _preview(facade, sha="0" * 64)
    assert [call[0] for call in facade.calls] == ["source"]


def test_changed_preview_revision_fails_before_table_read(tmp_path, monkeypatch):
    facade = _Facade(_bytes(_document()), tmp_path)
    with pytest.raises(PreparationSourceError, match="预览已变化"):
        _preview(facade, revision="0" * 64)


@pytest.mark.parametrize("field", ["source_sha256", "source_name", "revision"])
def test_preview_identity_or_digest_mismatch_is_rejected(tmp_path, field):
    facade = _Facade(_bytes(_document()), tmp_path)
    facade.value[field] = "changed"
    expected = facade.value["revision"] if field != "revision" else "a" * 64
    with pytest.raises(PreparationSourceError, match="预览已变化"):
        _preview(facade, revision=expected)


def test_source_change_between_archive_and_cached_preview_is_rejected(tmp_path):
    facade = _Facade(_bytes(_document()), tmp_path)
    expected = facade.value["revision"]

    def changed_preview(*_):
        value = deepcopy(facade.value)
        value["source_sha256"] = "0" * 64
        return value

    facade.imported_word_preview = changed_preview
    with pytest.raises(PreparationSourceError, match="预览已变化"):
        _preview(facade, revision=expected)


def test_table_text_must_match_even_a_self_consistent_preview(tmp_path):
    facade = _Facade(_bytes(_document()), tmp_path)
    facade.value["blocks"][1]["text"] += "不在原件的内容"
    _rehash(facade)
    before = deepcopy(facade.value)
    with pytest.raises(PreparationSourceError, match="内容与当前原文预览不一致"):
        _preview(facade)
    assert facade.value == before


@pytest.mark.parametrize("change", ["duplicate", "missing", "bool_index"])
def test_ambiguous_or_changed_block_positions_are_rejected(tmp_path, change):
    facade = _Facade(_bytes(_document()), tmp_path)
    if change == "missing":
        facade.value["blocks"].pop()
    elif change == "duplicate":
        facade.value["blocks"][1]["index"] = 1
    else:
        facade.value["blocks"][0]["index"] = True
    _rehash(facade)
    with pytest.raises(PreparationSourceError, match="区块位置已变化"):
        _preview(facade)


def test_invalid_word_container_has_safe_message_without_raw_details(tmp_path):
    facade = _Facade(_bytes(_document()), tmp_path)
    facade.data = b"not a DOCX container; private fixture detail"
    facade.value["source_sha256"] = hashlib.sha256(facade.data).hexdigest()
    _rehash(facade)
    with pytest.raises(PreparationSourceError, match="原表格暂时无法读取") as caught:
        _preview(facade)
    assert "private fixture" not in str(caught.value)


@pytest.mark.parametrize(
    "field,value",
    [
        ("batch_id", ""),
        ("source_id", None),
        ("source_sha256", "A" * 64),
        ("source_sha256", "g" * 64),
        ("expected_revision", True),
    ],
)
def test_invalid_selection_is_rejected_without_reading_source(tmp_path, field, value):
    facade = _Facade(_bytes(_document()), tmp_path)
    selection = {
        "batch_id": "batch-selected",
        "source_id": "source-selected",
        "source_sha256": hashlib.sha256(facade.data).hexdigest(),
        "expected_revision": facade.value["revision"],
    }
    selection[field] = value
    with pytest.raises(PreparationSourceError, match="选择信息不完整"):
        WordTablePreviewService(facade).preview(**selection)
    assert facade.calls == []
