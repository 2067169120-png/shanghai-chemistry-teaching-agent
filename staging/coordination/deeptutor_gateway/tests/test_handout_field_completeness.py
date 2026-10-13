"""A picture-complete section must not silently lose Word field instructions."""
from copy import deepcopy
from hashlib import sha256

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest

from integrations.deeptutor_shchem_v1 import desktop_handout_knowledge as handout
from test_handout_knowledge import make_facade, png, request_for, word_bytes
from io import BytesIO


def field(paragraph, *, simple=False, result=''):
    if simple:
        node=OxmlElement('w:fldSimple');node.set(qn('w:instr'), ' EQ \\o(→,\\s\\up7(催化剂)) ')
        run=OxmlElement('w:r');text=OxmlElement('w:t');text.text=result;run.append(text);node.append(run)
        paragraph._p.append(node)
        return
    for kind in ('begin','code','separate','result','end'):
        run=OxmlElement('w:r')
        if kind in {'begin','separate','end'}:
            node=OxmlElement('w:fldChar');node.set(qn('w:fldCharType'),kind)
        else:
            node=OxmlElement('w:instrText' if kind=='code' else 'w:t')
            node.text=' EQ \\o(→,\\s\\up7(催化剂)) ' if kind=='code' else result
        run.append(node);paragraph._p.append(run)


@pytest.mark.parametrize('simple,result', [(False,''),(False,'→'),(True,''),(True,'→')])
@pytest.mark.parametrize('teacher_only',[False,True])
def test_field_is_rejected_even_when_all_pictures_and_cached_text_are_present(tmp_path,simple,result,teacher_only):
    doc=Document();doc.add_heading('知识点1 聚合反应',3)
    paragraph=doc.add_paragraph('聚合反应条件：');field(paragraph,simple=simple,result=result)
    doc.add_paragraph().add_run().add_picture(BytesIO(png()))
    facade=make_facade(word_bytes(doc),tmp_path/'source')
    request=request_for(facade,teacher_only=teacher_only)
    source_bytes=facade.original.read_bytes();before=sha256(source_bytes).hexdigest()
    with pytest.raises(ValueError,match='Word公式域或字段'):
        handout.import_handout_knowledge(facade,[request],[],[])
    assert not facade.paths.task_root.exists()
    assert sha256(facade.original.read_bytes()).hexdigest()==before


def test_field_outside_selected_whole_section_does_not_block_that_section(tmp_path):
    doc=Document();doc.add_heading('知识点1 聚合反应',3)
    doc.add_paragraph('双键打开并连接成链。');doc.add_heading('知识点2 旧公式',3)
    field(doc.add_paragraph('条件：'))
    facade=make_facade(word_bytes(doc),tmp_path/'source')
    request=request_for(facade,'知识点1 聚合反应')
    result=handout.import_handout_knowledge(facade,[request],[],[])
    assert len(result['resources'])==1
    snapshot=result['resources'][0]['original']['evidence']['section_snapshot']
    assert [b['index'] for b in snapshot['blocks']]==[1,2]
    assert snapshot['blocks'][1]['text']=='双键打开并连接成链。'


def test_mixed_batch_does_not_publish_earlier_safe_images_before_later_field_failure(tmp_path):
    doc=Document();doc.add_heading('知识点1 安全栏目',3)
    doc.add_paragraph('完整文字。');doc.add_paragraph().add_run().add_picture(BytesIO(png()))
    doc.add_heading('知识点2 旧公式',3);field(doc.add_paragraph('条件：'))
    facade=make_facade(word_bytes(doc),tmp_path/'source')
    requests=[request_for(facade,'知识点1 安全栏目'),request_for(facade,'知识点2 旧公式')]
    original=deepcopy(requests)
    with pytest.raises(ValueError,match='Word公式域或字段'):
        handout.import_handout_knowledge(facade,requests,[],[])
    assert requests==original and not facade.paths.task_root.exists()
