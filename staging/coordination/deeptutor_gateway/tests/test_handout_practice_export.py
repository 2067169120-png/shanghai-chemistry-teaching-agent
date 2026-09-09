"""Run with bundled Documents Python; all sources and personal state are fixtures."""

from __future__ import annotations

import hashlib
from copy import deepcopy

import pytest
from docx import Document

from integrations.deeptutor_shchem_v1.desktop_handout_candidates import (
    HandoutCandidateError,
    HandoutCandidateService,
)
from integrations.deeptutor_shchem_v1.desktop_handout_practice import (
    DRAFT_ID,
    SOURCE_ROOT,
    HandoutPracticeService,
    NativeParagraphs,
)
from integrations.deeptutor_shchem_v1.desktop_state import DesktopStateStore


@pytest.fixture
def practice(tmp_path):
    folder = tmp_path / SOURCE_ROOT / "expanded/PKG-TEST"
    folder.mkdir(parents=True)
    for role in ("question", "answer"):
        doc = Document()
        paragraph = doc.add_paragraph("2．下列哪项符合所给条件？H")
        paragraph.add_run("2").font.subscript = True
        paragraph.add_run("O")
        table = doc.add_table(rows=2, cols=2)
        for cell, text in zip(
            (c for r in table.rows for c in r.cells),
            ("选项", "条件", "A", "条件一"),
            strict=True,
        ):
            cell.text = text
        if role == "answer":
            doc.add_paragraph("【答案】A")
        doc.save(folder / (role + ".docx"))
    locators = ["word/document.xml#/w:document/w:body/w:p[1]"] + [
        f"word/document.xml#/w:document/w:body/w:tbl[1]/w:tr[{r}]/w:tc[{c}]/w:p[1]"
        for r in (1, 2)
        for c in (1, 2)
    ]
    question_text = "2．下列哪项符合所给条件？H2O\n选项\n条件\nA\n条件一"
    item = {
        "key": "first",
        "revision": "a" * 64,
        "batch_id": "BATCH-TEST",
        "package_id": "PKG-TEST",
        "title": "第一题组 第2题",
        "classification": "native_text_complete",
        "question_text": question_text,
        "answer_text": "【非官方参考答案，正确性未核验】\n"
        + question_text
        + "\n【答案】A",
        "printed_number": "2",
        "parent_title": "第一题组",
        "source_name": "question.docx",
        "atomic_count": 1,
        "blockers": [],
        "candidate_only": True,
        "question_pages": [],
        "answer_pages": [],
        "source_document": {
            "relative_path": "expanded/PKG-TEST/question.docx",
            "sha256": hashlib.sha256(
                (folder / "question.docx").read_bytes()
            ).hexdigest(),
        },
        "editable_source": {
            "candidate_id": "Q2",
            "parent_id": "G1",
            "question_locators": locators,
            "answer_locators": locators
            + ["word/document.xml#/w:document/w:body/w:p[2]"],
            "answer_source_sha256": hashlib.sha256(
                (folder / "answer.docx").read_bytes()
            ).hexdigest(),
        },
    }

    class Reader:
        def catalog(self):
            return {"items": [deepcopy(item)], "warnings": [], "batches": []}

    state = DesktopStateStore(tmp_path / "state")
    service = HandoutPracticeService(
        HandoutCandidateService(tmp_path, state, reader_factory=lambda _: Reader())
    )
    return service, item, folder


def selection(item):
    return [{"key": item["key"], "revision": item["revision"]}]


def test_export_preserves_native_subscripts_table_source_and_roles(practice):
    service, item, folder = practice
    before = {p.name: p.read_bytes() for p in folder.glob("*.docx")}
    result = service.export("化学讲义练习", selection(item), 1)
    student = Document(service.artifact_path(result["export_id"], "student"))
    teacher = Document(service.artifact_path(result["export_id"], "teacher"))
    assert "【答案】" not in "\n".join(p.text for p in student.paragraphs)
    assert "【答案】A" in "\n".join(p.text for p in teacher.paragraphs)
    assert len(student.tables) == len(teacher.tables) == 1
    assert student.tables[0].cell(1, 1).text == "条件一"
    assert any(
        r.text == "2" and r.font.subscript is True
        for p in student.paragraphs
        for r in p.runs
    )
    assert any("原题号 2" in p.text for p in student.paragraphs)
    assert student.paragraphs[0].style.name == "Title"
    assert all(
        result[key] is False
        for key in ("answer_correctness_verified", "paper_structure_claimed")
    )
    assert result["items"][0]["review_state"] == "pending"
    assert all(p.read_bytes() == before[p.name] for p in folder.glob("*.docx"))
    reopened = HandoutPracticeService(service.candidates)
    assert reopened.history()[0] == result
    assert reopened.draft()["selections"] == selection(item)
    assert service.state.basket() == []


@pytest.mark.parametrize("mutation", ["stale", "incomplete", "missing", "duplicate"])
def test_invalid_selection_never_creates_output(practice, mutation):
    service, item, _ = practice
    selected = selection(item)
    if mutation == "stale":
        item["revision"] = "b" * 64
    elif mutation == "incomplete":
        item["classification"] = "hybrid_visual_required"
    elif mutation == "missing":
        selected[0]["key"] = "missing"
    else:
        selected += selected
    with pytest.raises(HandoutCandidateError):
        service.export("测试练习", selected)
    assert not service.root.exists()
    assert not service.history()


def test_source_hash_change_and_wrong_native_locator_fail(practice):
    service, item, folder = practice
    native = NativeParagraphs(service.candidates.workspace)
    wrong = deepcopy(item)
    wrong["editable_source"]["question_locators"] = [
        "word/document.xml#/w:document/w:body/w:p[2]"
    ]
    with pytest.raises(HandoutCandidateError):
        native.read(wrong, "question")
    path = folder / "question.docx"
    path.write_bytes(path.read_bytes() + b"changed")
    with pytest.raises(HandoutCandidateError):
        NativeParagraphs(service.candidates.workspace).read(item, "question")


def test_partial_table_or_numeric_cell_is_not_an_independent_question(practice):
    service, item, _ = practice
    item["editable_source"]["question_locators"] = item["editable_source"][
        "question_locators"
    ][1:]
    with pytest.raises(HandoutCandidateError):
        NativeParagraphs(service.candidates.workspace).read(item, "question")
    item["editable_source"]["question_locators"].insert(
        0, "word/document.xml#/w:document/w:body/w:p[1]"
    )
    item["editable_source"]["question_locators"].pop()
    with pytest.raises(HandoutCandidateError):
        NativeParagraphs(service.candidates.workspace).read(item, "question")


def test_missing_editable_binding_hidden_text_and_mismatch_fail(practice):
    service, item, folder = practice
    no_binding = {**item, "editable_source": None}
    with pytest.raises(HandoutCandidateError):
        NativeParagraphs(service.candidates.workspace).read(no_binding, "question")
    doc = Document(folder / "question.docx")
    doc.paragraphs[0].runs[0].font.hidden = True
    doc.save(folder / "question.docx")
    item["source_document"]["sha256"] = hashlib.sha256(
        (folder / "question.docx").read_bytes()
    ).hexdigest()
    with pytest.raises(HandoutCandidateError):
        NativeParagraphs(service.candidates.workspace).read(item, "question")


@pytest.mark.parametrize(
    "title,lines",
    [("", 2), ("a" * 101, 2), ("标题\n换行", 2), ("标题", True), ("标题", 9)],
)
def test_invalid_options_are_not_saved(practice, title, lines):
    service, item, _ = practice
    with pytest.raises(HandoutCandidateError):
        service.save_draft(title, selection(item), lines)
    assert DRAFT_ID not in service.state.snapshot()["drafts"]


def test_only_registered_generated_files_can_be_opened(practice):
    service, item, _ = practice
    with pytest.raises(HandoutCandidateError):
        service.artifact_path("../../unrelated", "student")
    with pytest.raises(HandoutCandidateError):
        service.artifact_path("HANDOUT-PRACTICE-" + "a" * 32, "student")
    result = service.export("化学练习", selection(item), 0)
    path = service.artifact_path(result["export_id"], "student")
    # Word outputs are intentionally editable; an edited file can still reopen.
    doc = Document(path)
    doc.add_paragraph("教师自行补充")
    doc.save(path)
    assert service.artifact_path(result["export_id"], "student") == path
    assert path.read_bytes()

def test_export_retains_ordered_multipage_answer_provenance(practice):
    service, item, _ = practice
    item["answer_pages"] = [
        {"page": 8, "path": "batches/test/solution/page-0008.png", "sha256": "8" * 64},
        {"page": 9, "path": "batches/test/solution/page-0009.png", "sha256": "9" * 64},
    ]
    result = service.export("跨页答案来源保留", selection(item), 0)
    assert [page["page"] for page in result["items"][0]["answer_pages"]] == [8, 9]
    assert result["candidate_only"] is True
    assert result["answer_correctness_verified"] is False
