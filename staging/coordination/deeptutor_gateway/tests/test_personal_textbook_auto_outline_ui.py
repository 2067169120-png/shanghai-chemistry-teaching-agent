"""Native outline selection, real PDF positioning and stale callback isolation."""
from copy import deepcopy

import pytest
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QTreeWidgetItemIterator
from test_desktop_studio_ui import settle
from test_desktop_studio_ui import window as window
from test_personal_textbook_auto_outline import outline_fixture as outline_fixture
from test_personal_textbook_auto_outline import source_pdf
from test_personal_textbook_ui import Bridge, complete, facade_for
from test_personal_textbook_ui import app as app
from test_textbook_bookshelf_ui import painted

from integrations.deeptutor_shchem_v1.desktop_workbench.personal_textbook_panel import (
    PersonalTextbookPanel,
)


@pytest.fixture
def automatic_panel(app, outline_fixture):
    f = outline_fixture; bridge = Bridge(); value = PersonalTextbookPanel(facade_for(f), bridge)
    value.resize(1080, 860); value.show(); value.set_book(f.study.books()[0])
    complete(bridge, '读取个人教材章节与笔记'); app.processEvents()
    yield value, bridge, f
    value.close(); value.deleteLater(); app.processEvents()


def detected(value, bridge):
    value.detect_outline_button.click(); complete(bridge, '识别个人教材目录')


def all_items(tree):
    iterator = QTreeWidgetItemIterator(tree); items = []
    while iterator.value():
        item = iterator.value()
        if item.data(0, Qt.ItemDataRole.UserRole): items.append(item)
        iterator += 1
    return items


def test_detection_preview_and_original_page_are_read_only_and_not_confirmed(automatic_panel, app):
    value, bridge, f = automatic_panel; before = deepcopy(f.state.snapshot())
    detected(value, bridge)
    assert len(all_items(value.outline_tree)) == 3 and len(value._outline_selected()) == 3
    assert value.outline_tree.topLevelItemCount() == 2
    assert f.state.snapshot() == before and value.tabs.currentIndex() == value._outline_tab
    value.outline_read_button.click(); app.processEvents(); complete(bridge, '读取教材原页'); app.processEvents()
    assert value._reader._loaded_pages and value._reader.pages.currentData() == 2
    assert value._reader.view.pageNavigator().currentPage() == 1 and painted(app, value._reader) > 10
    assert f.state.snapshot() == before


def test_selected_batch_saves_groups_and_can_make_lesson_reference(automatic_panel):
    value, bridge, f = automatic_panel; detected(value, bridge)
    all_items(value.outline_tree)[1].setCheckState(0, Qt.CheckState.Unchecked)
    value.outline_save_button.click(); complete(bridge, '保存个人教材自动目录')
    view = f.study.book_study(f.selection)
    assert len(view['sections']) == 2 and all(row['range_confirmed'] is False for row in view['sections'])
    assert len(all_items(value.tree)) == 2 and value.tree.topLevelItemCount() == 2
    assert value._outline_preview is None
    value.note_title.setText('自动目录下的教师笔记'); value.note_text.setPlainText('A → B（条件X）；用量 1.0 g。')
    value.note_confirmed.setChecked(True); value.save_note.click(); complete(bridge, '保存个人教材笔记')
    value.note_list.item(0).setCheckState(Qt.CheckState.Checked)
    emitted = []; value.reference_requested.connect(emitted.append)
    value.preview_button.click(); complete(bridge, '核对个人教材备课参考')
    value.confirm_button.click(); complete(bridge, '核对个人教材备课参考')
    assert len(emitted) == 1 and 'A → B' in emitted[0]['materials']
    assert all(row['human_reviewed'] is False for row in f.study.book_study(f.selection)['sections'])


def test_cancel_and_switch_cannot_restore_previous_directory(automatic_panel):
    value, bridge, f = automatic_panel
    value.detect_outline_button.click(); task = bridge.take('识别个人教材目录')
    payload = task.operation(); bridge.task_cancelled.emit(task.task_id, task.label); bridge.task_finished.emit(task.task_id)
    bridge.succeed(task, payload, run=False)
    assert value._outline_preview is None and not value.outline_save_button.isEnabled()
    value.detect_outline_button.click(); task = bridge.take('识别个人教材目录'); payload = task.operation()
    value.set_book(f.study.books()[0]); complete(bridge, '读取个人教材章节与笔记')
    bridge.succeed(task, payload, run=False)
    assert value._outline_preview is None and not all_items(value.outline_tree)


def test_changes_after_preview_keep_existing_chapters_and_pending_batch(automatic_panel):
    value, bridge, f = automatic_panel; detected(value, bridge)
    before = f.study.book_study(f.selection)
    f.study.save_section(f.selection, before['study_revision'], title='后台新增教师章节',
                         pdf_start=1, pdf_end=2, range_confirmed=True)
    snapshot = deepcopy(f.state.snapshot())
    value.outline_save_button.click(); task = bridge.take('保存个人教材自动目录')
    try: task.operation()
    except Exception as exc: bridge.fail(task, str(exc))
    bridge.task_finished.emit(task.task_id)
    assert f.state.snapshot() == snapshot and '预览后变化' in value.status.text()
    assert not value._busy


def test_outline_controls_fit_narrow_window(automatic_panel, app):
    from integrations.deeptutor_shchem_v1.desktop_workbench.studio_style import (
        WORKBENCH_STYLE,
    )
    from integrations.deeptutor_shchem_v1.desktop_workbench.typography import (
        install_ui_font,
    )
    value, bridge, _f = automatic_panel
    install_ui_font(app); value.setStyleSheet(WORKBENCH_STYLE)
    value.resize(420, 920); detected(value, bridge); app.processEvents()
    assert value.width() == 420
    assert value.splitter.orientation() == Qt.Orientation.Vertical
    for widget in (value.outline_tree, value.outline_read_button, value.outline_save_button):
        assert widget.isVisible() and widget.width() <= value.width() and widget.height() > 20
    assert value.outline_tree.verticalScrollBar().isVisible() or value.outline_tree.height() > 100


def test_repeat_detection_does_not_check_already_saved_rows(automatic_panel):
    value, bridge, f = automatic_panel; detected(value, bridge)
    value.outline_save_button.click(); complete(bridge, '保存个人教材自动目录')
    before = deepcopy(f.state.snapshot()); detected(value, bridge)
    assert len(all_items(value.outline_tree)) == 3 and not value._outline_selected()
    assert not value.outline_save_button.isEnabled() and f.state.snapshot() == before


def test_batch_append_preserves_unsaved_teacher_note_and_chapter_form(automatic_panel):
    value, bridge, f = automatic_panel
    value.section_title.setText('原教师章节');value.start.setValue(1);value.end.setValue(3)
    value.range_confirmed.setChecked(True);value.save_section.click();complete(bridge, '保存个人教材章节')
    value.section_title.setText('尚未保存的范围修订');value.end.setValue(4)
    value.note_title.setText('尚未保存的笔记');value.note_text.setPlainText('保留输入条件 X；用量 2.0 g。')
    value.note_page.setValue(2)
    old = f.study.book_study(f.selection)['sections'][0]
    detected(value, bridge);value.outline_save_button.click();complete(bridge, '保存个人教材自动目录')
    assert value.section_title.text() == '尚未保存的范围修订' and value.end.value() == 4
    assert value.note_title.text() == '尚未保存的笔记' and '2.0 g' in value.note_text.toPlainText()
    assert value.note_page.value() == 2 and not value.note_confirmed.isChecked()
    after = f.study.book_study(f.selection)
    assert after['sections'][0] == old and not after['notes']


@pytest.mark.parametrize('page', [True, 0, 7, '2'])
def test_original_page_position_rejects_non_integer_or_out_of_book(automatic_panel, app, page):
    value, bridge, _f = automatic_panel
    value._read(False, initial_page=page); app.processEvents()
    complete(bridge, '读取教材原页'); app.processEvents()
    assert not value._reader._loaded_pages and value._reader._source is None
    assert '个人教材来源或页范围未能核对' in value._reader.status.text()


def exercise_automatic_handoff(win, app, source, *, capture=None):
    """Exercise real workers and output files; optional captures use synthetic data."""
    from docx import Document
    from pptx import Presentation

    from integrations.deeptutor_shchem_v1.desktop_lesson_output import (
        FILES,
        checked_file,
    )
    from integrations.deeptutor_shchem_v1.desktop_workbench.lesson_design_dialog import (
        LessonDesignDialog,
    )

    def screenshot(name, widget):
        if capture:
            settle(app); capture(name, widget)

    original = source.read_bytes()
    workspace = win.facade.textbook_workspace()
    workspace.commit_books(workspace.preview_books([source]))
    study = win.facade.personal_textbook_study(); book = study.books()[0]
    selection = {key: book[key] for key in ('source_id', 'revision')}
    page = win.preparation_page
    page.topic.setText('自动目录联动合成课题'); page.objective.setPlainText('保留既有目标')
    page.audience.setText('合成授课对象')
    page.materials.setPlainText('保留既有参考')
    previous = deepcopy(page._payload())
    win.navigate('textbooks'); settle(app, lambda: not win.textbook_page._loading)
    win.textbook_page.book.setCurrentIndex(win.textbook_page.book.findData(book['source_id']))
    personal = win.textbook_page.personal; settle(app, lambda: not personal._busy)
    personal.detect_outline_button.click(); settle(app, lambda: not personal._busy)
    assert len(personal._outline_selected()) == 3
    screenshot('synthetic-directory-main-1280.png', win)
    personal.outline_save_button.click(); settle(app, lambda: not personal._busy)
    settle(app, lambda: personal._reader is not None and personal._reader._loaded_pages)
    view = study.book_study(selection)
    assert len(view['sections']) == 3 and all(row['range_confirmed'] is False for row in view['sections'])
    assert [personal._reader.pages.itemData(i) for i in range(personal._reader.pages.count())] == [2, 3]
    assert painted(app, personal._reader) > 10
    screenshot('synthetic-directory-reading-main-1280.png', win)
    personal.new_note.click()
    personal.note_title.setText('合成课堂笔记'); personal.note_text.setPlainText('A → B（条件X）；用量 2.0 g。')
    personal.note_page.setValue(2); personal.note_confirmed.setChecked(True)
    personal.save_note.click(); settle(app, lambda: not personal._busy)
    personal.note_list.item(0).setCheckState(Qt.CheckState.Checked)
    personal.preview_button.click(); settle(app, lambda: not personal._busy)
    screenshot('synthetic-directory-note-preview-main-1280.png', win)
    personal.confirm_button.click(); settle(app, lambda: not personal._busy)
    payload = page._payload()
    for key in ('topic', 'audience', 'objective', 'image_input_mode', 'image_assets'):
        assert payload.get(key) == previous.get(key)
    assert payload['materials'].startswith('保留既有参考') and 'A → B' in payload['materials']
    page._save(); settle(app, lambda: not page._save_task_id)
    assert any('A → B' in win.facade.load_preparation_draft(row['draft_id'], row['revision'])['payload']['materials']
               for row in win.facade.preparation_draft_options()), page.status.text()
    dialog = LessonDesignDialog(page); dialog.show(); settle(app)
    dialog.add_node(); dialog.bind_source()
    assert 'A → B' in dialog.history.value['nodes'][0]['material_text']
    dialog.locked.setChecked(False); dialog.student_material.setChecked(True)
    dialog.fields['material_text'].setPlainText('合成课堂笔记\nA → B（条件X）；用量 2.0 g。\n合成自动目录教材.pdf · PDF第2页')
    dialog.fields['teacher_answer'].setPlainText('仅教师答案哨兵-自动目录')
    screenshot('synthetic-directory-lesson-design-1100.png', dialog)
    dialog.generate.click(); settle(app, lambda: not dialog.busy)
    assert dialog.history.value['exports'], dialog.status.text()
    identity = dialog.history.value['exports'][-1]['id']
    files = {name: checked_file(win.facade, dialog.history.value, identity, name) for name in FILES}
    presentation = Presentation(files['lesson_presentation.pptx'])
    body = '\n'.join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)
    notes = '\n'.join(slide.notes_slide.notes_text_frame.text for slide in presentation.slides)
    def doc_text(path):
        doc = Document(path)
        return '\n'.join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    teacher, student = doc_text(files['lesson_plan.docx']), doc_text(files['student_worksheet.docx'])
    assert all('A → B' in text and 'PDF第2页' in text for text in (body, teacher, student))
    assert '仅教师答案哨兵-自动目录' in teacher and '仅教师答案哨兵-自动目录' in notes
    assert '仅教师答案哨兵-自动目录' not in body and '仅教师答案哨兵-自动目录' not in student
    assert source.read_bytes() == original
    assert all(row['human_reviewed'] is False for row in study.book_study(selection)['sections'])
    screenshot('synthetic-directory-lesson-generated-1100.png', dialog)
    dialog.reject(); dialog.deleteLater(); settle(app)
    return files


def test_auto_directory_actual_window_note_draft_and_three_outputs(window, tmp_path):
    win, app = window
    exercise_automatic_handoff(win, app, source_pdf(tmp_path / '合成自动目录教材.pdf'))
