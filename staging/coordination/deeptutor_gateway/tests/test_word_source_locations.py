"""Exact original-XML occurrences remain independent of cached native text."""

from __future__ import annotations

import hashlib
import io
from copy import deepcopy
from types import SimpleNamespace

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image

from integrations.deeptutor_shchem_v1 import desktop_word_source_locations as source_locations
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
    PreparationSourceError,
    PreparationSourcesService,
    _digest,
)
from integrations.deeptutor_shchem_v1.desktop_word_source_locations import (
    A, M, MC, O, R, V, W, WordSourceLocationService,
)
from integrations.deeptutor_shchem_v1.word_native_text import WordNativeTextReader


def _bytes(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _png(color="red"):
    output = io.BytesIO()
    Image.new("RGB", (5, 4), color).save(output, "PNG")
    return output.getvalue()


def _node(tag, text=None, **attributes):
    value = etree.Element(tag)
    value.text = text
    for key, item in attributes.items():
        value.set(key, item)
    return value


def _picture(paragraph, color="red"):
    run = paragraph.add_run()
    run.add_picture(io.BytesIO(_png(color)))
    return run, next(node for node in run._r.iter() if node.tag == A + "blip")


class _Facade:
    def __init__(self, document, tmp_path):
        self.paths = SimpleNamespace(workspace_root=tmp_path)
        self.data = _bytes(document)
        self.filename = "结构定位合成样本.docx"
        self.value = PreparationSourcesService(tmp_path).word_preview_bytes(self.data, self.filename)
        self.calls = []

    def _imported_word_source(self, batch_id, source_id):
        self.calls.append(("source", batch_id, source_id))
        return SimpleNamespace(content=self.data, filename=self.filename)

    def imported_word_preview(self, batch_id, source_id, *, read_only=False):
        assert read_only is True, "The location service must never populate the preview cache"
        self.calls.append(("preview", batch_id, source_id, read_only))
        return deepcopy(self.value)


def _args(facade):
    return {"batch_id": "batch-selected", "source_id": "source-selected",
            "source_sha256": hashlib.sha256(facade.data).hexdigest(),
            "expected_revision": facade.value["revision"]}


def _preview(facade, **kwargs):
    return WordSourceLocationService(facade).preview(**{**_args(facade), **kwargs})


def _locations(preview):
    return [location for block in preview["blocks"] for location in block["locations"]]


def _rehash(facade):
    facade.value["revision"] = _digest({key: value for key, value in facade.value.items() if key != "revision"})


def _ole(document, paragraph, *, with_preview=False, shape_id="ole-fixture", external=False, wrong_rid=False):
    image_rid = None
    if with_preview:
        run, image = _picture(paragraph, "blue")
        image_rid = image.get(R + "embed")
        paragraph._p.remove(run._r)
    if external:
        rid = document.part.relate_to("https://example.invalid/object.bin", RT.OLE_OBJECT, is_external=True)
    else:
        part = Part(PackURI("/word/embeddings/ole-fixture.bin"),
                    "application/vnd.openxmlformats-officedocument.oleObject", b"original synthetic OLE bytes", document.part.package)
        rid = document.part.relate_to(part, RT.OLE_OBJECT)
    host = _node(W + "object")
    if image_rid:
        shape = _node(V + "shape", id=shape_id)
        image = _node(V + "imagedata")
        image.set(R + "id", image_rid)
        shape.append(image)
        host.append(shape)
    ole = _node(O + "OLEObject", ShapeID=shape_id, ProgID="Equation.3")
    ole.set(R + "id", "rId-does-not-exist" if wrong_rid else rid)
    host.append(ole)
    paragraph.add_run()._r.append(host)
    return host, ole


def test_multiple_objects_and_repeated_relationship_are_distinct_exact_occurrences(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("同段两个相同原图和其他对象：")
    _picture(paragraph)
    _picture(paragraph)
    math = _node(M + "oMath")
    math_run = _node(M + "r")
    math_run.append(_node(M + "t", "x"))
    math.append(math_run)
    paragraph._p.append(math)
    symbol = _node(W + "sym")
    symbol.set(W + "font", "Symbol")
    symbol.set(W + "char", "F061")
    paragraph.add_run()._r.append(symbol)
    facade = _Facade(doc, tmp_path)
    saved = deepcopy(facade.value)
    source_bytes = facade.data
    result = _preview(facade)
    locations = _locations(result)
    assert [item["kind"] for item in locations] == ["image", "image", "omml", "symbol"]
    pictures = locations[:2]
    assert pictures[0]["assets"][0]["asset_id"] == pictures[1]["assets"][0]["asset_id"]
    assert pictures[0]["location_id"] != pictures[1]["location_id"]
    assert len({item["xml_locator"] for item in locations}) == 4
    parsed = Document(io.BytesIO(facade.data))
    for item in locations:
        path = item["xml_locator"].split("#", 1)[1]
        nodes = parsed._element.getroottree().xpath(path, namespaces={"w": W[1:-1]})
        assert len(nodes) == 1
        assert "正文第1段" in item["position_text"]
    assert "源对象4" in locations[3]["position_text"]
    locations[0]["assets"][0]["sha256"] = "mutated display copy"
    assert facade.value == saved and facade.data == source_bytes
    assert facade.calls[-1] == ("preview", "batch-selected", "source-selected", True)


def test_literal_placeholder_does_not_invent_source_objects(tmp_path):
    doc = Document()
    doc.add_paragraph("字面文字【待查看原文：图片或图形】和【未提取到文字】，不是图片。")
    facade = _Facade(doc, tmp_path)
    result = _preview(facade)
    assert _locations(result) == []
    assert result["blocks"][0]["context_text"] == facade.value["blocks"][0]["text"]


def test_nested_merged_tables_sdt_and_real_child_indices_are_retained(tmp_path):
    doc = Document()
    doc.add_paragraph("正文引入")
    table = doc.add_table(rows=3, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "合并标题"
    _picture(merged.paragraphs[0])
    table.cell(1, 0).merge(table.cell(2, 0)).text = "纵合并"
    inner = table.cell(1, 1).add_table(rows=1, cols=2)
    target = inner.cell(0, 1).paragraphs[0]
    target.add_run("格内原文")
    _picture(target, "blue")
    cell = target._p.getparent()
    cell.remove(target._p)
    sdt, content = _node(W + "sdt"), _node(W + "sdtContent")
    content.append(target._p)
    sdt.append(_node(W + "sdtPr"))
    sdt.append(content)
    cell.append(sdt)
    continuation = table._tbl.tr_lst[2].tc_lst[0]
    p = continuation.find(W + "p")
    sym = _node(W + "sym")
    sym.set(W + "char", "F061")
    run = _node(W + "r")
    run.append(sym)
    p.append(run)
    # Wrap an actual source row; coordinates must still include the wrapper's
    # raw child positions rather than pretend the row was a direct child.
    row = table._tbl.tr_lst[1]
    original_position = table._tbl.index(row)
    table._tbl.remove(row)
    wrapper, row_content = _node(W + "sdt"), _node(W + "sdtContent")
    row_content.append(row)
    wrapper.append(row_content)
    table._tbl.insert(original_position, wrapper)
    facade = _Facade(doc, tmp_path)
    result = _preview(facade, block_indices=[2])
    locations = _locations(result)
    assert len(locations) == 3
    assert "源第1行·源第1—2列（横向合并2列）" in locations[0]["position_text"]
    nested = locations[1]
    assert "源第2行·源第2列" in nested["position_text"]
    assert "格内第1张嵌套表" in nested["position_text"]
    assert "源第1行·源第2列" in nested["position_text"]
    assert "格内第1段" in nested["position_text"]
    assert "格内原文" in nested["context_text"]
    assert "纵向合并续接" in locations[2]["position_text"]
    parsed = Document(io.BytesIO(facade.data))
    for location in locations:
        selected = parsed._element.getroottree().xpath(location["xml_locator"].split("#", 1)[1], namespaces={"w": W[1:-1]})
        assert len(selected) == 1 and selected[0].tag in {A + "blip", W + "sym"}


def test_omitted_table_columns_are_source_grid_columns(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    _picture(table.cell(0, 1).paragraphs[0])
    row = table.rows[0]._tr
    row.remove(row.tc_lst[0])
    before = OxmlElement("w:gridBefore")
    before.set(qn("w:val"), "1")
    row.get_or_add_trPr().append(before)
    result = _preview(_Facade(doc, tmp_path))
    assert "源第1行·源第2列" in _locations(result)[0]["position_text"]


def test_hidden_deleted_and_compatibility_branches_are_never_claimed_visible(tmp_path):
    doc = Document()
    hidden_style = doc.styles.add_style("HiddenFixture", WD_STYLE_TYPE.CHARACTER)
    hidden_style.font.hidden = True
    paragraph = doc.add_paragraph("状态样本")
    hidden, _ = _picture(paragraph)
    hidden.style = hidden_style
    deleted_run, _ = _picture(paragraph)
    paragraph._p.remove(deleted_run._r)
    deletion = _node(W + "del")
    deletion.append(deleted_run._r)
    paragraph._p.append(deletion)
    branch_run, _ = _picture(paragraph)
    paragraph._p.remove(branch_run._r)
    compatible = _node(MC + "AlternateContent")
    choice, fallback = _node(MC + "Choice", Requires="wps"), _node(MC + "Fallback")
    choice.append(deepcopy(branch_run._r))
    fallback.append(deepcopy(branch_run._r))
    compatible.extend([choice, fallback])
    paragraph._p.append(compatible)
    facade = _Facade(doc, tmp_path)
    locations = _locations(_preview(facade))
    assert len(locations) == 4
    assert [item["source_states"] for item in locations] == [["hidden"], ["deleted_or_moved_from"], ["compatibility_choice"], ["compatibility_fallback"]]
    assert all(item["current_visibility_confirmed"] is False for item in locations)
    assert all(item["context_text"].startswith("【原结构内容；显示状态待核对】") for item in locations)
    assert len({item["location_id"] for item in locations}) == 4
    assert all(len(item["assets"]) == 1 for item in locations)


def test_complex_fields_are_individual_nodes_not_a_fabricated_complete_span(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("复杂域前文")
    for kind in ("begin", "separate", "end"):
        node = _node(W + "fldChar")
        node.set(W + "fldCharType", kind)
        paragraph.add_run()._r.append(node)
        if kind == "begin":
            paragraph.add_run()._r.append(_node(W + "instrText", " REF fixture "))
    # An unmatched code in another paragraph cannot be silently paired.
    doc.add_paragraph().add_run()._r.append(_node(W + "instrText", " PAGE "))
    facade = _Facade(doc, tmp_path)
    locations = _locations(_preview(facade))
    assert len(locations) == 5
    assert all(item["kind"] == "field_code" for item in locations)
    assert all(item["complete_field_span_verified"] is False for item in locations)
    assert all(any("完整起止" in notice for notice in item["notices"]) for item in locations)
    assert {item.get("field_instruction") for item in locations if item.get("field_instruction")} == {" REF fixture ", " PAGE "}


def test_ole_links_exact_host_binary_and_shape_bound_replacement(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("旧公式")
    _ole(doc, paragraph, with_preview=True)
    facade = _Facade(doc, tmp_path)
    locations = _locations(_preview(facade))
    ole = next(item for item in locations if item["kind"] == "ole_object")
    image = next(item for item in locations if item["kind"] == "image")
    assert ole["assets"] == image["assets"]
    assert ole["host_xml_locator"] != ole["xml_locator"]
    reference = ole["relationship_references"][0]
    assert reference["package_part"] == "/word/embeddings/ole-fixture.bin"
    assert reference["sha256"] == hashlib.sha256(b"original synthetic OLE bytes").hexdigest()
    assert reference["xml_locator"] == ole["xml_locator"]


@pytest.mark.parametrize("problem", ["no_preview", "external", "wrong_rid", "wrong_shape", "other_branch"])
def test_ole_without_its_own_verified_replacement_never_uses_neighbor_image(tmp_path, problem):
    doc = Document()
    paragraph = doc.add_paragraph("相邻图不是旧公式的替代图")
    _picture(paragraph, "green")
    host, ole = _ole(doc, paragraph, with_preview=problem != "no_preview",
                     external=problem == "external", wrong_rid=problem == "wrong_rid")
    if problem == "wrong_shape":
        ole.set("ShapeID", "not-the-preview-shape")
    elif problem == "other_branch":
        shape = next(child for child in host if child.tag == V + "shape")
        host.remove(shape)
        compatible, fallback = _node(MC + "AlternateContent"), _node(MC + "Fallback")
        fallback.append(shape)
        compatible.append(fallback)
        host.insert(0, compatible)
    facade = _Facade(doc, tmp_path)
    location = next(item for item in _locations(_preview(facade)) if item["kind"] == "ole_object")
    assert location["assets"] == []
    assert any("不能用邻图" in notice for notice in location["notices"])
    with pytest.raises(PreparationSourceError, match="不属于所选对象"):
        WordSourceLocationService(facade).image(**_args(facade), location_id=location["location_id"], asset_id="word-b1-image1")


def test_external_and_broken_image_references_do_not_borrow_valid_image(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("关系验证")
    run, _ = _picture(paragraph)
    for rid in ("missing-picture-id", doc.part.relate_to("https://example.invalid/picture.png", RT.IMAGE, is_external=True)):
        copied = deepcopy(run._r)
        next(node for node in copied.iter() if node.tag == A + "blip").set(R + "embed", rid)
        paragraph._p.append(copied)
    locations = _locations(_preview(_Facade(doc, tmp_path)))
    assert len(locations) == 3
    assert len(locations[0]["assets"]) == 1
    assert locations[1]["assets"] == locations[2]["assets"] == []


def test_altchunk_and_special_xml_are_located_without_following_references(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("原文")
    paragraph.add_run()._r.append(_node("{urn:fixture:special}payload", "opaque"))
    chunk = _node(W + "altChunk")
    chunk.set(R + "id", doc.part.relate_to("https://example.invalid/fragment.html", RT.A_F_CHUNK, is_external=True))
    doc._element.body.insert(len(doc._element.body) - 1, chunk)
    locations = _locations(_preview(_Facade(doc, tmp_path)))
    assert [item["kind"] for item in locations] == ["special_object", "alt_chunk"]
    assert locations[1]["relationship_references"][0]["status"] == "external_not_loaded"
    assert all(item["assets"] == [] for item in locations)


def test_image_rebinds_exact_occurrence_asset_source_revision_and_expected_image_sha(tmp_path, monkeypatch):
    doc = Document()
    paragraph = doc.add_paragraph("第一个原图")
    _picture(paragraph)
    _picture(paragraph)
    _picture(doc.add_paragraph("另一位置"), "blue")
    facade = _Facade(doc, tmp_path)
    locations = _locations(_preview(facade))
    service = WordSourceLocationService(facade)
    original = PreparationSourcesService.word_asset_bytes
    calls = []
    def checked(self, content, asset_id, *, render_metafiles=False, expected_sha256=None):
        calls.append((hashlib.sha256(content).hexdigest(), asset_id, render_metafiles, expected_sha256))
        return original(self, content, asset_id, render_metafiles=render_metafiles, expected_sha256=expected_sha256)
    monkeypatch.setattr(PreparationSourcesService, "word_asset_bytes", checked)
    for location in locations[:2]:
        result = service.image(**_args(facade), location_id=location["location_id"], asset_id=location["assets"][0]["asset_id"])
        assert result["bytes"] == _png() and result["location_id"] == location["location_id"]
    assert all(call[2] is True and call[3] == hashlib.sha256(_png()).hexdigest() for call in calls)
    with pytest.raises(PreparationSourceError, match="不属于所选对象"):
        service.image(**_args(facade), location_id=locations[0]["location_id"], asset_id=locations[-1]["assets"][0]["asset_id"])
    assert len(calls) == 2
    with pytest.raises(PreparationSourceError, match="预览已变化"):
        service.image(**{**_args(facade), "expected_revision": "0" * 64}, location_id=locations[0]["location_id"], asset_id="word-b1-image1")
    assert len(calls) == 2


@pytest.mark.parametrize("field,value", [
    ("batch_id", ""), ("source_id", "bad\nsource"), ("source_sha256", "A" * 64),
    ("expected_revision", True), ("block_indices", [True]), ("block_indices", [1, 1]),
    ("block_indices", []), ("block_indices", "1"),
])
def test_invalid_input_is_rejected_before_any_source_read(tmp_path, field, value):
    doc = Document()
    doc.add_paragraph("source")
    facade = _Facade(doc, tmp_path)
    with pytest.raises(PreparationSourceError):
        _preview(facade, **{field: value})
    assert facade.calls == []


def test_source_and_self_consistent_but_wrong_projection_are_rejected(tmp_path):
    doc = Document()
    doc.add_paragraph("真实文本")
    facade = _Facade(doc, tmp_path)
    with pytest.raises(PreparationSourceError, match="原Word已变化"):
        _preview(facade, source_sha256="0" * 64)
    assert [entry[0] for entry in facade.calls] == ["source"]
    facade.value["blocks"][0]["text"] = "伪造的自洽预览"
    _rehash(facade)
    with pytest.raises(PreparationSourceError, match="内容与当前原文预览不一致"):
        _preview(facade)


def test_forged_asset_catalogue_and_changed_block_identity_are_rejected(tmp_path):
    doc = Document()
    _picture(doc.add_paragraph("原图"))
    facade = _Facade(doc, tmp_path)
    facade.value["assets"][0]["sha256"] = "0" * 64
    _rehash(facade)
    with pytest.raises(PreparationSourceError, match="图片目录"):
        _preview(facade)
    facade = _Facade(doc, tmp_path)
    facade.value["blocks"][0]["index"] = True
    _rehash(facade)
    with pytest.raises(PreparationSourceError, match="区块位置已变化"):
        _preview(facade)


def test_nonselected_or_path_like_image_identifiers_are_rejected(tmp_path):
    doc = Document()
    _picture(doc.add_paragraph("原图"))
    facade = _Facade(doc, tmp_path)
    location = _locations(_preview(facade))[0]
    service = WordSourceLocationService(facade)
    with pytest.raises(PreparationSourceError, match="定位标识不正确"):
        service.image(**_args(facade), location_id=location["location_id"], asset_id="C:/private/image.png")
    with pytest.raises(PreparationSourceError, match="不属于所选对象"):
        service.image(**_args(facade), location_id="word-location-" + "0" * 64, asset_id="word-b1-image1")


def test_source_changes_between_initial_bytes_and_preview_are_rejected(tmp_path):
    doc = Document()
    doc.add_paragraph("old")
    facade = _Facade(doc, tmp_path)
    args = _args(facade)
    def changed(*_, read_only=False):
        value = deepcopy(facade.value)
        value["source_sha256"] = "0" * 64
        return value
    facade.imported_word_preview = changed
    with pytest.raises(PreparationSourceError, match="预览已变化"):
        WordSourceLocationService(facade).preview(**args)


def test_duplicate_ole_shape_ownership_is_ambiguous_not_a_replacement_binding(tmp_path):
    doc = Document()
    host, ole = _ole(doc, doc.add_paragraph("重复宿主声明"), with_preview=True)
    host.append(deepcopy(ole))
    locations = _locations(_preview(_Facade(doc, tmp_path)))
    objects = [item for item in locations if item["kind"] == "ole_object"]
    assert len(objects) == 2
    assert objects[0]["location_id"] != objects[1]["location_id"]
    assert all(item["assets"] == [] for item in objects)


def test_deleted_table_row_keeps_source_row_not_current_visible_row_number(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=1)
    _picture(table.cell(0, 0).paragraphs[0])
    _picture(table.cell(1, 0).paragraphs[0], "blue")
    deleted_row = table.rows[0]._tr
    position = table._tbl.index(deleted_row)
    table._tbl.remove(deleted_row)
    deletion = _node(W + "del")
    deletion.append(deleted_row)
    table._tbl.insert(position, deletion)
    locations = _locations(_preview(_Facade(doc, tmp_path)))
    assert len(locations) == 2
    assert "deleted_or_moved_from" in locations[0]["source_states"]
    assert "源第1行" in locations[0]["position_text"]
    assert "源第2行" in locations[1]["position_text"]


@pytest.mark.parametrize("field", ["source_sha256", "source_name", "extraction_revision", "revision"])
def test_each_preview_binding_must_match_the_original_and_expected_revision(tmp_path, field):
    doc = Document()
    doc.add_paragraph("source")
    facade = _Facade(doc, tmp_path)
    args = _args(facade)
    facade.value[field] = "changed"
    with pytest.raises(PreparationSourceError, match="预览已变化"):
        WordSourceLocationService(facade).preview(**args)


@pytest.mark.parametrize("image_type", ["drawing", "ole_replacement"])
def test_png_mime_does_not_make_a_nonimage_relationship_an_image_binding(tmp_path, image_type):
    doc = Document()
    paragraph = doc.add_paragraph("关系类型也必须正确")
    if image_type == "drawing":
        _, image = _picture(paragraph)
        attribute = R + "embed"
    else:
        host, _ = _ole(doc, paragraph, with_preview=True)
        image = next(node for node in host.iterdescendants() if node.tag == V + "imagedata")
        attribute = R + "id"
    image_part = doc.part.rels[image.get(attribute)].target_part
    wrong_rid = doc.part.relate_to(image_part, RT.OLE_OBJECT)
    image.set(attribute, wrong_rid)
    facade = _Facade(doc, tmp_path)
    original_preview = deepcopy(facade.value)
    assert len(original_preview["assets"]) == 1  # Existing cache format is preserved.
    locations = _locations(_preview(facade))
    objects = [item for item in locations if item["kind"] in {"image", "ole_object"}]
    assert objects and all(item["assets"] == [] for item in objects)
    image_location = next(item for item in locations if item["kind"] == "image")
    assert image_location["relationship_references"][0]["relationship_type"] == RT.OLE_OBJECT
    for location in objects:
        with pytest.raises(PreparationSourceError, match="不属于所选对象"):
            WordSourceLocationService(facade).image(
                **_args(facade), location_id=location["location_id"], asset_id="word-b1-image1",
            )
    assert facade.value == original_preview


@pytest.mark.parametrize("nesting", ["shape", "object", "shape_and_object"])
def test_outer_ole_cannot_borrow_an_image_from_a_nested_shape_or_object(tmp_path, nesting):
    doc = Document()
    paragraph = doc.add_paragraph("外层旧公式没有自身替代图")
    picture_run, image = _picture(paragraph, "blue")
    image_rid = image.get(R + "embed")
    paragraph._p.remove(picture_run._r)
    host, outer_ole = _ole(doc, paragraph, with_preview=False, shape_id="outer")
    outer_shape = _node(V + "shape", id="outer")
    textbox, text_content = _node(V + "textbox"), _node(W + "txbxContent")
    inner_paragraph, inner_run = _node(W + "p"), _node(W + "r")
    target = inner_run
    if "object" in nesting:
        inner_host = _node(W + "object")
        target.append(inner_host)
        target = inner_host
    if "shape" in nesting:
        inner_shape = _node(V + "shape", id="inner")
        target.append(inner_shape)
        target = inner_shape
    inner_image = _node(V + "imagedata")
    inner_image.set(R + "id", image_rid)
    target.append(inner_image)
    if nesting == "shape_and_object":
        inner_ole = _node(O + "OLEObject", ShapeID="inner")
        inner_ole.set(R + "id", outer_ole.get(R + "id"))
        inner_host.append(inner_ole)
    inner_paragraph.append(inner_run)
    text_content.append(inner_paragraph)
    textbox.append(text_content)
    outer_shape.append(textbox)
    host.insert(0, outer_shape)
    facade = _Facade(doc, tmp_path)
    locations = _locations(_preview(facade))
    ole_locations = [item for item in locations if item["kind"] == "ole_object"]
    # The outer OLE node follows its shape in the actual source tree.
    outer = ole_locations[-1]
    assert outer["assets"] == []
    picture = next(item for item in locations if item["kind"] == "image")
    assert len(picture["assets"]) == 1
    if nesting == "shape_and_object":
        assert ole_locations[0]["assets"] == picture["assets"]
    with pytest.raises(PreparationSourceError, match="不属于所选对象"):
        WordSourceLocationService(facade).image(
            **_args(facade), location_id=outer["location_id"], asset_id=picture["assets"][0]["asset_id"],
        )


def test_native_context_for_repeated_objects_in_one_paragraph_is_read_once(tmp_path, monkeypatch):
    doc = Document()
    paragraph = doc.add_paragraph("同段上下文只投影一次")
    for _ in range(3):
        _picture(paragraph)
    facade = _Facade(doc, tmp_path)
    original = WordNativeTextReader.read
    paragraphs_read = []
    def counted(self, element):
        if element.tag == W + "p":
            paragraphs_read.append(element)
        return original(self, element)
    monkeypatch.setattr(WordNativeTextReader, "read", counted)
    locations = _locations(_preview(facade))
    assert len(locations) == 3
    assert len(paragraphs_read) == 1
    assert len({item["context_text"] for item in locations}) == 1
    assert source_locations.MAX_LOCATIONS >= 601


def test_object_count_limit_fails_closed_and_allows_a_smaller_selected_range(tmp_path, monkeypatch):
    doc = Document()
    paragraph = doc.add_paragraph("两个对象超过测试阈值")
    _picture(paragraph)
    _picture(paragraph)
    _picture(doc.add_paragraph("一个对象在阈值内"), "blue")
    facade = _Facade(doc, tmp_path)
    original = deepcopy(facade.value)
    first_location = _locations(_preview(facade, block_indices=[1]))[0]
    monkeypatch.setattr(source_locations, "MAX_LOCATIONS", 1)
    with pytest.raises(PreparationSourceError, match="对象数量超过限制.*限制区块范围或拆分"):
        _preview(facade)
    result = _preview(facade, block_indices=[2])
    assert len(_locations(result)) == 1
    location = _locations(result)[0]
    asset_id = location["assets"][0]["asset_id"]
    service = WordSourceLocationService(facade)
    image = service.image(**_args(facade), location_id=location["location_id"], asset_id=asset_id)
    assert image["bytes"] == _png("blue")
    with pytest.raises(PreparationSourceError, match="不属于所选对象"):
        service.image(**_args(facade), location_id=first_location["location_id"], asset_id=asset_id)
    with pytest.raises(PreparationSourceError, match="区块范围不正确"):
        service.image(**_args(facade), location_id=location["location_id"], asset_id="word-b99-image1")
    assert facade.value == original


def test_total_context_output_limit_fails_closed_with_small_fixture(tmp_path, monkeypatch):
    doc = Document()
    _picture(doc.add_paragraph("甲段原图"))
    _picture(doc.add_paragraph("乙段原图"), "blue")
    facade = _Facade(doc, tmp_path)
    original = deepcopy(facade.value)
    first_location = _locations(_preview(facade, block_indices=[1]))[0]
    context_size = len(facade.value["blocks"][1]["text"])
    monkeypatch.setattr(source_locations, "MAX_CONTEXT_CHARS", 2 * context_size)
    with pytest.raises(PreparationSourceError, match="上下文超过输出限制.*限制区块范围或拆分"):
        _preview(facade)
    result = _preview(facade, block_indices=[2])
    location = _locations(result)[0]
    asset_id = location["assets"][0]["asset_id"]
    service = WordSourceLocationService(facade)
    image = service.image(**_args(facade), location_id=location["location_id"], asset_id=asset_id)
    assert image["bytes"] == _png("blue")
    with pytest.raises(PreparationSourceError, match="不属于所选对象"):
        service.image(**_args(facade), location_id=first_location["location_id"], asset_id=asset_id)
    with pytest.raises(PreparationSourceError, match="不属于所选对象"):
        service.image(**_args(facade), location_id=location["location_id"], asset_id=first_location["assets"][0]["asset_id"])
    # Both the block text and its occurrence context count against the budget.
    monkeypatch.setattr(source_locations, "MAX_CONTEXT_CHARS", context_size)
    with pytest.raises(PreparationSourceError, match="上下文超过输出限制"):
        _preview(facade, block_indices=[2])
    assert facade.value == original


def test_hidden_outer_paragraph_style_is_not_overridden_by_normal_textbox_style(tmp_path):
    doc = Document()
    style = doc.styles.add_style("OuterHiddenFixture", WD_STYLE_TYPE.PARAGRAPH)
    style.font.hidden = True
    outer = doc.add_paragraph("外层隐藏", style=style)
    inner = doc.add_paragraph("文本框内使用Normal", style="Normal")
    _picture(inner)
    doc._element.body.remove(inner._p)
    pict, shape = _node(W + "pict"), _node(V + "shape", id="textbox")
    textbox, content = _node(V + "textbox"), _node(W + "txbxContent")
    content.append(inner._p)
    textbox.append(content)
    shape.append(textbox)
    pict.append(shape)
    outer.add_run()._r.append(pict)
    locations = _locations(_preview(_Facade(doc, tmp_path)))
    picture = next(item for item in locations if item["kind"] == "image")
    assert "hidden" in picture["source_states"]
    assert picture["current_visibility_confirmed"] is False
    assert "文本框内第1段" in picture["position_text"]
    assert picture["context_text"].startswith("【原结构内容；显示状态待核对】")
    assert any("隐藏" in notice for notice in picture["notices"])
