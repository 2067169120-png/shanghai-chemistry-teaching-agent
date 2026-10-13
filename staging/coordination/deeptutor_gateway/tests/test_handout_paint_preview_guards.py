"""Synthetic Paint previews keep local object identity and closed display guards."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from types import SimpleNamespace

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml.ns import qn
from docx.shared import Mm
from lxml import etree
from PIL import Image, ImageDraw
import pytest

from integrations.deeptutor_shchem_v1 import desktop_handout_knowledge as knowledge
from integrations.deeptutor_shchem_v1 import desktop_handout_paint_preview as paint
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    block, digest, resource, validate_resource)
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
    PreparationSourcesService, _body_blocks, _word_images)
from integrations.deeptutor_shchem_v1.offline_word_image_label_review import (
    OfflineWordLabelReviewError, _objects)
from test_handout_knowledge import make_facade, request_for, word_bytes

W, V, O = paint.W, paint.V, paint.O


def paint_png(*, animated=False, color='#275c91'):
    """Asymmetric, nonblank pixels make a vertical flip observable."""
    image = Image.new('RGB', (80, 40), 'white')
    drawing = ImageDraw.Draw(image)
    drawing.rectangle((2, 2, 30, 10), fill=color)
    drawing.line((12, 32, 65, 25), fill='black', width=3)
    output = BytesIO()
    if animated:
        second = Image.new('RGB', image.size, '#91a73c')
        image.save(output, 'PNG', save_all=True, append_images=[second], duration=100, loop=0)
    else:
        image.save(output, 'PNG')
    return output.getvalue()


def add_paint_object(doc, paragraph, png_bytes, number):
    """Same-container PBrush/PNG object; OLE stays opaque and occurrence-specific."""
    image_part = Part(PackURI(f'/word/media/paint{number}.png'), 'image/png',
        png_bytes, doc.part.package)
    image_rid = doc.part.relate_to(image_part, RT.IMAGE)
    opaque = b'opaque uninterpreted PBrush binary ' + str(number).encode('ascii')
    ole_part = Part(PackURI(f'/word/embeddings/paint{number}.bin'),
        'application/vnd.openxmlformats-officedocument.oleObject', opaque, doc.part.package)
    ole_rid = doc.part.relate_to(ole_part, RT.OLE_OBJECT)
    node = etree.Element(f'{{{W}}}object')
    shape_id = f'_x0000_paint{number}'
    shape = etree.SubElement(node, f'{{{V}}}shape', id=shape_id, type='#_x0000_t75',
        style='width:32.25pt;height:8.25pt', filled='f', stroked='f')
    etree.SubElement(shape, f'{{{V}}}imagedata', {qn('r:id'): image_rid})
    etree.SubElement(node, f'{{{O}}}OLEObject', {
        'Type': 'Embed', 'ProgID': 'PBrush', 'ShapeID': shape_id,
        'DrawAspect': 'Content', 'ObjectID': str(number), qn('r:id'): ole_rid})
    paragraph.add_run()._r.append(node)
    return node, opaque


def paint_document(*, repeated=False, with_flip=False, png_bytes=None):
    """Return document, shared PNG, independent object nodes and opaque binaries."""
    raw = paint_png() if png_bytes is None else png_bytes
    doc = Document()
    doc.add_heading('知识点1 完整图形栏目', 3)
    paragraph = doc.add_paragraph('先比较原图。')
    node, opaque = add_paint_object(doc, paragraph, raw, 1)
    nodes, binaries = [node], [opaque]
    if repeated:
        paragraph.add_run('再比较同图的另一处。')
        node, opaque = add_paint_object(doc, paragraph, raw, 2)
        nodes.append(node); binaries.append(opaque)
    paragraph.add_run('完整尾部条件。')
    if with_flip:
        paragraph = doc.add_paragraph('这张原生图片保持原字节并垂直翻转。')
        paragraph.add_run().add_picture(BytesIO(raw), width=Mm(24))
        paragraph._p.xpath('.//pic:spPr/a:xfrm')[0].set('flipV', '1')
    doc.add_paragraph('下一完整知识段。')
    return doc, raw, nodes, binaries


def paint_facade(tmp_path, *, role='handout', repeated=False, with_flip=False,
                 png_bytes=None, mutate=None):
    """Owned source-bound facade for guard and downstream export tests."""
    doc, raw, nodes, binaries = paint_document(
        repeated=repeated, with_flip=with_flip, png_bytes=png_bytes)
    if mutate:
        mutate(doc, nodes[0])
    facade = make_facade(word_bytes(doc), tmp_path/'source', role=role)
    service = PreparationSourcesService(facade.paths.workspace_root)
    facade.imported_word_preview = lambda *args, read_only=False: service.word_preview_bytes(
        facade.value.data, facade.value.name)
    facade.imported_word_asset = lambda *args: service.word_asset_bytes(
        facade.value.data, args[2], render_metafiles=True)
    return SimpleNamespace(facade=facade, document=doc, png=raw,
        nodes=nodes, ole_bytes=binaries)


def frozen_refs(item):
    return [ref for native in item['original']['evidence']['section_snapshot']['blocks']
            for ref in native['images']]


def assert_no_publication(facade, original, mtime):
    assert not facade.paths.task_root.exists()
    assert facade.original.read_bytes() == original
    assert facade.original.stat().st_mtime_ns == mtime


def test_two_same_png_paint_objects_keep_independent_opaque_occurrences(tmp_path):
    fixture = paint_facade(tmp_path, repeated=True)
    facade = fixture.facade; before = facade.original.read_bytes()
    result = knowledge.import_handout_knowledge(facade, [request_for(facade)], [], [])
    item = result['resources'][0]; refs = frozen_refs(item)
    assert validate_resource(item) == item
    assert len(refs) == 2 and len(result['image_assets']) == 1
    assert [r['position'] for r in refs] == [1, 2]
    assert len({r['container_xml_sha256'] for r in refs}) == 2
    assert len({r['source_object']['shape_id'] for r in refs}) == 2
    assert len({r['source_object']['ole']['sha256'] for r in refs}) == 2
    assert {r['source_object']['ole']['sha256'] for r in refs} == {
        sha256(raw).hexdigest() for raw in fixture.ole_bytes}
    assert all(r['original_sha256'] == r['preview_sha256'] == sha256(fixture.png).hexdigest()
               and r['derived_preview'] is False and r['renderer_revision'] is None for r in refs)
    assert all(r['source_object']['prog_id'] == 'PBrush'
               and r['source_object']['binary_interpreted'] is False
               and r['source_object']['derived_preview'] is False for r in refs)
    asset = result['image_assets'][0]
    store = PreparationImageStore(facade.paths.task_root/'preparation-v1/images')
    assert store.load(asset) == fixture.png and facade.original.read_bytes() == before


@pytest.mark.parametrize('change', [
    'shape_id', 'unknown_progid', 'link', 'icon', 'image_external', 'ole_external',
    'extra_shape', 'vml_rotation_attribute', 'vml_rotation_style'])
def test_paint_binding_rejects_ambiguous_external_or_transformed_objects(tmp_path, change):
    def mutate(doc, node):
        shape = node.find(f'{{{V}}}shape')
        image = shape.find(f'{{{V}}}imagedata')
        ole = node.find(f'{{{O}}}OLEObject')
        if change == 'shape_id': ole.set('ShapeID', 'another-object')
        elif change == 'unknown_progid': ole.set('ProgID', 'Unrelated.Document')
        elif change == 'link': ole.set('Type', 'Link')
        elif change == 'icon': ole.set('DrawAspect', 'Icon')
        elif change in {'image_external', 'ole_external'}:
            element = image if change == 'image_external' else ole
            reltype = RT.IMAGE if change == 'image_external' else RT.OLE_OBJECT
            element.set(qn('r:id'), doc.part.relate_to(
                'https://example.invalid/never-requested', reltype, is_external=True))
        elif change == 'extra_shape': etree.SubElement(node, f'{{{V}}}shape')
        elif change == 'vml_rotation_attribute': shape.set('rotation', '10')
        else: shape.set('style', shape.get('style') + ';rotation:10')
    fixture = paint_facade(tmp_path, mutate=mutate); facade = fixture.facade
    before, mtime = facade.original.read_bytes(), facade.original.stat().st_mtime_ns
    with pytest.raises(ValueError):
        knowledge.import_handout_knowledge(facade, [request_for(facade)], [], [])
    assert_no_publication(facade, before, mtime)


@pytest.mark.parametrize('raw', [b'broken PNG fixture', paint_png(animated=True)], ids=['corrupt', 'animated'])
def test_invalid_or_animated_png_never_becomes_a_paint_preview(tmp_path, raw):
    fixture = paint_facade(tmp_path, png_bytes=raw); facade = fixture.facade
    before, mtime = facade.original.read_bytes(), facade.original.stat().st_mtime_ns
    with pytest.raises(ValueError):
        knowledge.import_handout_knowledge(facade, [request_for(facade)], [], [])
    assert_no_publication(facade, before, mtime)


@pytest.mark.parametrize('attribute,value', [('rot', '900000'), ('flipH', '1'), ('flipV', 'unknown')])
def test_drawing_unproved_transforms_block_the_whole_knowledge_import(tmp_path, attribute, value):
    def mutate(doc, _node):
        transform = doc._element.body.xpath('.//pic:spPr/a:xfrm')[0]
        transform.attrib.clear(); transform.set(attribute, value)
    fixture = paint_facade(tmp_path, with_flip=True, mutate=mutate); facade = fixture.facade
    before, mtime = facade.original.read_bytes(), facade.original.stat().st_mtime_ns
    with pytest.raises(ValueError):
        knowledge.import_handout_knowledge(facade, [request_for(facade)], [], [])
    assert_no_publication(facade, before, mtime)


@pytest.mark.parametrize('part', ['png', 'ole'])
def test_changed_frozen_png_or_opaque_ole_bytes_reject_before_publication(tmp_path, part):
    fixture = paint_facade(tmp_path); facade = fixture.facade
    request = request_for(facade)
    before, mtime = facade.original.read_bytes(), facade.original.stat().st_mtime_ns
    doc = Document(BytesIO(facade.value.data))
    object_node = doc._element.body.xpath('.//w:object')[0]
    local = object_node.find(f'.//{{{V}}}imagedata') if part == 'png' else object_node.find(f'{{{O}}}OLEObject')
    target = doc.part.rels[local.get(qn('r:id'))].target_part
    target._blob = paint_png(color='#bf3319') if part == 'png' else b'changed opaque PBrush binary'
    facade.value.data = word_bytes(doc)
    with pytest.raises(ValueError):
        knowledge.import_handout_knowledge(facade, [request], [], [])
    assert_no_publication(facade, before, mtime)


def test_invalid_second_section_cannot_publish_images_from_valid_first_section(tmp_path):
    doc, raw, _nodes, _binaries = paint_document()
    doc.add_heading('知识点2 第二完整图形栏目', 3)
    paragraph = doc.add_paragraph('第二栏目的完整条件。')
    node, _ = add_paint_object(doc, paragraph, raw, 2)
    node.find(f'{{{O}}}OLEObject').set('ShapeID', 'unrelated-second-shape')
    facade = make_facade(word_bytes(doc), tmp_path/'mixed-sections')
    before, mtime = facade.original.read_bytes(), facade.original.stat().st_mtime_ns
    requests = [request_for(facade, title) for title in (
        '知识点1 完整图形栏目', '知识点2 第二完整图形栏目')]
    with pytest.raises(ValueError): knowledge.import_handout_knowledge(facade, requests, [], [])
    assert_no_publication(facade, before, mtime)


@pytest.mark.parametrize('change', ['binary_interpreted', 'derived_preview', 'coverage',
    'image_sha', 'ole_external', 'object_type', 'draw_aspect', 'container', 'unknown_field'])
def test_resealed_paint_evidence_cannot_replace_bound_display_proof(tmp_path, change):
    fixture = paint_facade(tmp_path); facade = fixture.facade
    item = knowledge.import_handout_knowledge(facade, [request_for(facade)], [], [])['resources'][0]
    changed = deepcopy(item); obj = frozen_refs(changed)[0]['source_object']
    if change == 'binary_interpreted': obj[change] = True
    elif change == 'derived_preview': obj[change] = True
    elif change == 'coverage': obj[change] = 'whole_binary_interpreted'
    elif change == 'image_sha': obj['image']['sha256'] = 'f'*64
    elif change == 'ole_external': obj['ole']['is_external'] = True
    elif change == 'object_type': obj[change] = 'Link'
    elif change == 'draw_aspect': obj[change] = 'Icon'
    elif change == 'container': obj['container_xml_sha256'] = 'f'*64
    else: obj['extra'] = True
    changed['original_sha256'] = digest(changed['original'])
    with pytest.raises(ValueError): validate_resource(changed)


@pytest.mark.parametrize('change', ['flip_false', 'original_sha', 'container', 'extra',
    'missing_presentation', 'presentation_disagrees'])
def test_resealed_vertical_flip_must_match_its_original_occurrence_and_presentation(tmp_path, change):
    fixture = paint_facade(tmp_path, with_flip=True); facade = fixture.facade
    item = knowledge.import_handout_knowledge(facade, [request_for(facade)], [], [])['resources'][0]
    changed = deepcopy(item)
    ref = next(r for r in frozen_refs(changed) if 'transform' in r)
    paragraph = next(b for b in changed['original']['blocks'] if b['kind'] == 'text'
        and any('transform' in r for r in b.get('presentation', {}).get('images', [])))
    if change == 'flip_false': ref['transform']['flip_vertical'] = False
    elif change == 'original_sha': ref['transform']['original_sha256'] = 'f'*64
    elif change == 'container': ref['transform']['container_xml_sha256'] = 'f'*64
    elif change == 'extra': ref['transform']['crop'] = True
    elif change == 'missing_presentation': paragraph['presentation']['images'] = []
    else: paragraph['presentation']['images'][0].pop('transform')
    changed['blocks'] = deepcopy(changed['original']['blocks'])
    changed['original_sha256'] = digest(changed['original'])
    with pytest.raises(ValueError): validate_resource(changed)


def test_global_equation_object_review_still_rejects_paint_same_object_preview(tmp_path):
    fixture = paint_facade(tmp_path); doc = Document(BytesIO(fixture.facade.value.data))
    elements = list(_body_blocks(doc._element.body))
    original_assets = _word_images(elements, doc)
    with pytest.raises(OfflineWordLabelReviewError): _objects(elements[1], doc, 2, original_assets)


def test_old_ordinary_png_snapshot_still_validates_without_new_transform_fields(tmp_path):
    doc = Document(); doc.add_heading('知识点1 普通原图', 3)
    doc.add_paragraph('完整原文。').add_run().add_picture(BytesIO(paint_png()))
    facade = make_facade(word_bytes(doc), tmp_path/'ordinary')
    item = knowledge.import_handout_knowledge(facade, [request_for(facade)], [], [])['resources'][0]
    assert all('transform' not in r and 'source_object' not in r for r in frozen_refs(item))
    assert validate_resource(deepcopy(item)) == item


def test_paint_cannot_masquerade_as_old_chemwindow_evidence_after_resealing(tmp_path):
    fixture = paint_facade(tmp_path); facade = fixture.facade
    item = knowledge.import_handout_knowledge(facade, [request_for(facade)], [], [])['resources'][0]
    changed = deepcopy(item); obj = frozen_refs(changed)[0]['source_object']
    obj['prog_id'] = 'ChemWindow.Document'
    for key in ('object_type', 'draw_aspect', 'derived_preview'): obj.pop(key)
    changed['original_sha256'] = digest(changed['original'])
    with pytest.raises(ValueError): validate_resource(changed)


def test_non_handout_resource_cannot_add_a_vertical_flip_without_source_evidence():
    marker = '【待查看原文：图片或图形】'
    row = block('text', 'personal', '完整原文。' + marker, 1)
    transform = {'schema': paint.TRANSFORM_SCHEMA, 'flip_vertical': True,
        'container_xml_sha256': 'a'*64, 'original_sha256': 'b'*64}
    row['presentation'] = {'schema': 'shchem.native-word-paragraph.v1',
        'images': [{'image_block_id': 'CB-'+'c'*32,
        'placeholder': marker, 'width_mm': 20, 'height_mm': 10}], 'scripts': []}
    picture = block('image', 'personal', 'IMG-'+'b'*64, 1); picture['id'] = 'CB-'+'c'*32
    item = resource('personal_text', '原生行内图', [row, picture],
        {'source_path': '教师输入', 'source_sha256': 'unknown', 'source_revision': 'unknown'})
    assert validate_resource(item) == item
    for blocks in (item['original']['blocks'], item['blocks']):
        blocks[0]['presentation']['images'][0]['transform'] = deepcopy(transform)
    item['original_sha256'] = digest(item['original'])
    with pytest.raises(ValueError, match='完整讲义栏目的来源证据'):
        validate_resource(item)
