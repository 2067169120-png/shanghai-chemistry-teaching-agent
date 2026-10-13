"""Page breaks keep readable starts and nearby picture context without data loss."""
from copy import deepcopy
from io import BytesIO
import pytest
from PIL import Image
from docx import Document
from docx.shared import Mm
from test_composition_render import app, text
from integrations.deeptutor_shchem_v1.desktop_composition_render import (
    _content, arrange_composition, layout_composition, PX_MM, PADDING, write_pptx)
from integrations.deeptutor_shchem_v1.desktop_composition_inline import NativeParagraphDisplay
from integrations.deeptutor_shchem_v1.desktop_composition_rich_render import flatten_elements
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    FORMATS, new_composition, add_resource, resource, block, resource_frames)
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
from integrations.deeptutor_shchem_v1.offline_word_image_label_review import _objects
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import _word_images


def personal(title, blocks):
    return resource('personal_text', title, blocks,
        {'source_path':'作者输入','source_sha256':'unknown','source_revision':'unknown'})


def composition(*items, format='16:9'):
    value=new_composition();value['format']=format;w,h=FORMATS[format]
    for item in items:
        value=add_resource(value,item,{'page_id':value['pages'][0],
            'x':12,'y':12,'width':w-24,'height':h-24})
    return value


def visible_fragments(layout, identity):
    return [[e for e in f['elements'] if e['kind']!='space']
        for f in layout['fragments'] if f['resource_id']==identity]


def test_title_moves_with_two_opening_lines_instead_of_filling_a_short_page_tail(app):
    target=personal('必须随正文移动的标题',[block('text','personal','完整正文条件。'*35)])
    chunks=_content(target,296,{}, {},student=False,format='16:9')
    opening=next(i for i,c in enumerate(chunks) if c.get('block_id'))
    title_height=sum(c['height'] for c in chunks[:opening])+2*PADDING
    required=title_height+sum(c['height'] for c in chunks[opening:opening+2])
    for n in range(1,16):
        filler=personal('前序任务',[block('text','personal','已完成。\n'*n)])
        remaining=156-8-2*PADDING-sum(c['height'] for c in _content(filler,296,{}, {},student=False,format='16:9'))
        if title_height<=remaining<required:break
    else:pytest.fail('未构造到恰好只能容纳标题的页尾')
    original=composition(filler,target);arranged=arrange_composition(original,[],{})
    layout=layout_composition(arranged,[],{},student=False);assert not layout['errors']
    first=visible_fragments(layout,target['id'])[0]
    assert len([e for e in first if e.get('block_id')==target['blocks'][0]['id']])>=2
    assert first[0]['page_id']!=resource_frames(arranged['placements'][filler['id']])[0]['page_id']
    assert text(layout,target['blocks'][0]['id'])==target['blocks'][0]['value']
    assert arranged['resources']==original['resources']


@pytest.mark.parametrize('inline_picture',[False,True])
def test_title_and_tall_first_picture_share_a_page_with_whole_original_image(app,tmp_path,inline_picture):
    buf=BytesIO();Image.new('RGB',(120,900),'#53866a').save(buf,format='PNG');raw=buf.getvalue()
    store=PreparationImageStore(tmp_path/'images');asset=store.import_bytes(raw,'合成图','作者','分页验证')
    picture=block('image','personal',asset['asset_id'],1);blocks=[picture]
    if inline_picture:
        doc=Document();doc.add_paragraph().add_run().add_picture(BytesIO(raw),width=Mm(45))
        package=BytesIO();doc.save(package);display=NativeParagraphDisplay(package.getvalue())
        element=display.elements[0];paragraph=block('text','personal',display.reader.read(element).text,1)
        display.bind(paragraph,[picture],_objects(element,display.doc,1,_word_images(display.elements,display.doc)))
        blocks=[paragraph,picture]
    item=personal('标题与首图同页',blocks);original=composition(item)
    arranged=arrange_composition(original,[asset],{asset['asset_id']:raw})
    layout=layout_composition(arranged,[asset],{asset['asset_id']:raw},student=False)
    assert not layout['errors']
    parts=list(flatten_elements(visible_fragments(layout,item['id'])[0]))
    images=[p for p in parts if p['kind']=='image'];assert len(images)==1
    assert abs(images[0]['width']/images[0]['height']-120/900)<.0001
    assert arranged['resources']==original['resources'] and store.load(asset)==raw


@pytest.mark.parametrize('inline_picture',[False,True])
def test_short_picture_intro_moves_with_its_picture(app,tmp_path,inline_picture):
    buf=BytesIO();Image.new('RGB',(120,70),'#476d9b').save(buf,format='PNG');raw=buf.getvalue()
    store=PreparationImageStore(tmp_path/'images');asset=store.import_bytes(raw,'合成图','作者','图注续页验证')
    caption=block('text','personal','读图任务：保留完整坐标与条件。',2)
    image=block('image','personal',asset['asset_id'],2)
    for n in range(1,16):
        prelude=block('text','personal','前序条件。\n'*n,1)
        blocks=[prelude,caption,image]
        if inline_picture:
            doc=Document();doc.add_paragraph(caption['value']);doc.add_paragraph().add_run().add_picture(BytesIO(raw),width=Mm(32))
            package=BytesIO();doc.save(package);display=NativeParagraphDisplay(package.getvalue())
            element=display.elements[1];paragraph=block('text','personal',display.reader.read(element).text,2)
            display.bind(paragraph,[image],_objects(element,display.doc,2,_word_images(display.elements,display.doc)))
            blocks=[prelude,caption,paragraph,image]
        item=personal('连续任务',blocks)
        chunks=_content(item,296,{asset['asset_id']:asset},{asset['asset_id']:raw},student=False,format='16:9')
        caption_start=next(i for i,c in enumerate(chunks) if c.get('block_id')==caption['id'])
        used=sum(c['height'] for c in chunks[:caption_start])+2*PADDING
        caption_height=sum(c['height'] for c in chunks if c.get('block_id')==caption['id'])
        picture_height=next(c['height'] for c in chunks if c['kind']=='image' or
                            any(p['kind']=='image' for p in c.get('_parts',[])))
        if used+caption_height<156<used+caption_height+picture_height:break
    else:pytest.fail('未构造到图注可放入而图片放不入的页尾')
    value=arrange_composition(composition(item),[asset],{asset['asset_id']:raw})
    layout=layout_composition(value,[asset],{asset['asset_id']:raw},student=False)
    captions=[e for page in layout['pages'].values() for e in page if e.get('block_id')==caption['id'] and e['kind']=='text']
    pictures=[e for page in layout['pages'].values() for e in flatten_elements(page) if e['kind']=='image']
    assert not layout['errors'] and len(pictures)==1
    assert {e['page_id'] for e in captions}=={pictures[0]['page_id']}
    assert text(layout,prelude['id'])==prelude['value'].replace('\n','')


@pytest.mark.parametrize('inline_picture',[False,True])
def test_short_opening_stem_and_tall_picture_stay_with_title_without_shrinking_text(app,tmp_path,inline_picture):
    buf=BytesIO();Image.new('RGB',(120,900),'#64748b').save(buf,format='PNG');raw=buf.getvalue()
    store=PreparationImageStore(tmp_path/'images');asset=store.import_bytes(raw,'合成图','作者','短题干及首图验证')
    stem=block('text','personal','读图任务：保留完整的题设条件。',1)
    image=block('image','personal',asset['asset_id'],2);blocks=[stem,image]
    if inline_picture:
        doc=Document();doc.add_paragraph(stem['value']);doc.add_paragraph().add_run().add_picture(BytesIO(raw),width=Mm(45))
        package=BytesIO();doc.save(package);display=NativeParagraphDisplay(package.getvalue())
        element=display.elements[1];paragraph=block('text','personal',display.reader.read(element).text,2)
        display.bind(paragraph,[image],_objects(element,display.doc,2,_word_images(display.elements,display.doc)))
        blocks=[stem,paragraph,image]
    item=personal('题干与首图同页',blocks);original=composition(item)
    arranged=arrange_composition(original,[asset],{asset['asset_id']:raw})
    layout=layout_composition(arranged,[asset],{asset['asset_id']:raw},student=False)
    assert not layout['errors']
    fragments=visible_fragments(layout,item['id']);parts=list(flatten_elements(fragments[0]))
    images=[p for p in parts if p['kind']=='image'];assert len(images)==1
    assert text(layout,stem['id'])==stem['value']
    assert all(p['font_px']==32 for p in parts if p.get('block_id')==stem['id'] and p['kind']=='text')
    assert abs(images[0]['width']/images[0]['height']-120/900)<.0001
    assert arranged['resources']==original['resources'] and store.load(asset)==raw


def test_student_picture_uses_later_full_frame_after_teacher_only_short_continuation(app,tmp_path):
    buf=BytesIO();Image.new('RGB',(120,900),'#6b7890').save(buf,format='PNG');raw=buf.getvalue()
    store=PreparationImageStore(tmp_path/'images');asset=store.import_bytes(raw,'合成图','作者','受众续页验证')
    intro=block('text','personal','给定图表的完整条件。'*7)
    hidden=block('text','teacher_note','仅教师的备课说明。'*6)
    picture=block('image','personal',asset['asset_id'])
    target=personal('图表任务',[intro,hidden,picture])
    for n in range(1,16):
        filler=personal('前序任务',[block('text','personal','前序条件。\n'*n)])
        value=arrange_composition(composition(filler,target),[asset],{asset['asset_id']:raw})
        teacher=layout_composition(value,[asset],{asset['asset_id']:raw},student=False)
        small=[f for f in teacher['fragments'] if f['resource_id']==target['id'] and
               0<f['frame']['height']<60 and
               {e['block_id'] for e in f['elements'] if e['kind']!='space'
                and not e.get('navigation_heading')}=={hidden['id']}]
        if small:break
    else:pytest.fail('未构造到仅含教师说明的短续页外框')
    before=deepcopy(value);student=layout_composition(value,[asset],{asset['asset_id']:raw},student=True)
    assert not teacher['errors'] and not student['errors']
    assert text(student,intro['id'])==intro['value'] and not text(student,hidden['id'])
    images=[e for page in student['pages'].values() for e in flatten_elements(page) if e['kind']=='image']
    assert len(images)==1 and abs(images[0]['width']/images[0]['height']-120/900)<.0001
    assert small[0]['frame']['page_id'] in student['omitted_teacher_page_ids']
    assert value==before and store.load(asset)==raw


@pytest.mark.parametrize('opening_lines',[1,3])
def test_title_short_stem_and_picture_in_one_native_paragraph_share_first_page(app,tmp_path,opening_lines):
    buf=BytesIO();Image.new('RGB',(120,900),'#527b87').save(buf,format='PNG');raw=buf.getvalue()
    store=PreparationImageStore(tmp_path/'images');asset=store.import_bytes(raw,'合成图','作者','段内首图验证')
    doc=Document();paragraph=doc.add_paragraph()
    stem='\n'.join(['读图任务：完整题设条件。']*opening_lines)
    run=paragraph.add_run(stem);run.add_break();run.add_picture(BytesIO(raw),width=Mm(90))
    package=BytesIO();doc.save(package);display=NativeParagraphDisplay(package.getvalue());element=display.elements[0]
    row=block('text','personal',display.reader.read(element).text,1)
    image=block('image','personal',asset['asset_id'],1)
    display.bind(row,[image],_objects(element,display.doc,1,_word_images(display.elements,display.doc)))
    target=personal('同一段落中的题干与首图',[row,image]);original=composition(target)
    arranged=arrange_composition(original,[asset],{asset['asset_id']:raw})
    for student in (False,True):
        layout=layout_composition(arranged,[asset],{asset['asset_id']:raw},student=student)
        assert not layout['errors']
        first=list(flatten_elements(visible_fragments(layout,target['id'])[0]))
        pictures=[p for p in first if p['kind']=='image'];assert len(pictures)==1
        assert ''.join(p['text'] for p in first if p['kind']=='text' and p.get('block_id')==row['id'])==stem.replace('\n','')
        assert all(p['font_px']==32 for p in first if p['kind']=='text' and p.get('block_id')==row['id'])
        assert abs(pictures[0]['width']/pictures[0]['height']-120/900)<.0001
    assert arranged['resources']==original['resources'] and store.load(asset)==raw


def test_two_font_height_badge_does_not_separate_short_stem_from_next_native_picture(app,tmp_path):
    store=PreparationImageStore(tmp_path/'images');raws=[];assets=[]
    for size,label in [((240,90),'徽标'),((120,900),'整图')]:
        buf=BytesIO();Image.new('RGB',size,'#627b93').save(buf,format='PNG');raw=buf.getvalue()
        raws.append(raw);assets.append(store.import_bytes(raw,'合成'+label,'作者','徽标浮点边界验证'))
    doc=Document();p=doc.add_paragraph();p.add_run().add_picture(BytesIO(raws[0]),height=Mm(7.9375))
    stem='完整原题条件与来源说明。\n仍须保留第二行条件。';p.add_run(stem)
    doc.add_paragraph().add_run().add_picture(BytesIO(raws[1]),width=Mm(90))
    package=BytesIO();doc.save(package);display=NativeParagraphDisplay(package.getvalue());blocks=[];rows=[]
    for index,(element,asset) in enumerate(zip(display.elements,assets),1):
        row=block('text','personal',display.reader.read(element).text,index);picture=block('image','personal',asset['asset_id'],index)
        source_images=[a for a in _word_images(display.elements,display.doc) if a['block_index']==index]
        display.bind(row,[picture],_objects(element,display.doc,index,source_images))
        rows.append(row);blocks.extend([row,picture])
    target=personal('来源标题保留完整\n并保留第二行',[*blocks]);original=composition(target)
    data={a['asset_id']:raw for a,raw in zip(assets,raws)};arranged=arrange_composition(original,assets,data)
    for student in (False,True):
        layout=layout_composition(arranged,assets,data,student=student);assert not layout['errors']
        first=list(flatten_elements(visible_fragments(layout,target['id'])[0]));pictures=[p for p in first if p['kind']=='image']
        assert {p['asset_id'] for p in pictures}==set(data)
        assert ''.join(p['text'] for p in first if p['kind']=='text' and p.get('block_id')==rows[0]['id'])==stem.replace('\n','')
        assert all(p['font_px']==32 for p in first if p['kind']=='text' and p.get('block_id')==rows[0]['id'])
    assert arranged['resources']==original['resources'] and all(store.load(a)==raw for a,raw in zip(assets,raws))


def test_long_same_paragraph_stem_keeps_full_picture_on_continuation(app,tmp_path):
    buf=BytesIO();Image.new('RGB',(120,900),'#627193').save(buf,format='PNG');raw=buf.getvalue()
    store=PreparationImageStore(tmp_path/'images');asset=store.import_bytes(raw,'合成图','作者','长段内题干验证')
    doc=Document();paragraph=doc.add_paragraph();stem='\n'.join(['完整条件继续保留。']*7)
    run=paragraph.add_run(stem);run.add_break();run.add_picture(BytesIO(raw),width=Mm(90))
    package=BytesIO();doc.save(package);display=NativeParagraphDisplay(package.getvalue());element=display.elements[0]
    row=block('text','personal',display.reader.read(element).text,1);image=block('image','personal',asset['asset_id'],1)
    display.bind(row,[image],_objects(element,display.doc,1,_word_images(display.elements,display.doc)))
    target=personal('长段落不压缩题干字号',[row,image]);original=composition(target)
    arranged=arrange_composition(original,[asset],{asset['asset_id']:raw})
    layout=layout_composition(arranged,[asset],{asset['asset_id']:raw},student=False)
    assert not layout['errors']
    fragments=visible_fragments(layout,target['id']);assert len(fragments)>=2
    assert not any(p['kind']=='image' for p in flatten_elements(fragments[0]))
    parts=[p for page in layout['pages'].values() for p in flatten_elements(page)]
    pictures=[p for p in parts if p['kind']=='image'];assert len(pictures)==1 and pictures[0]['height']>100
    assert ''.join(p['text'] for p in parts if p['kind']=='text' and p.get('block_id')==row['id'])==stem.replace('\n','')
    assert all(p['font_px']==32 for p in parts if p['kind']=='text' and p.get('block_id')==row['id'])
    assert arranged['resources']==original['resources'] and store.load(asset)==raw


def test_long_paragraph_final_line_does_not_continue_alone_and_ppt_keeps_all_text(app,tmp_path):
    item=None
    # Find a paragraph whose old greedy break would leave exactly one last line.
    for n in range(65,160):
        candidate=personal('长段落',[block('text','personal','完整条件。'*n+'结尾证据')])
        chunks=_content(candidate,296,{}, {},student=False,format='16:9')
        used=0;cursor=0
        while cursor<len(chunks) and used+chunks[cursor]['height']<=150:
            used+=chunks[cursor]['height'];cursor+=1
        if sum(c['kind']=='text' for c in chunks[cursor:])==1:item=candidate;break
    assert item is not None
    original=composition(item);value=arrange_composition(original,[],{})
    layout=layout_composition(value,[],{},student=False)
    fragments=visible_fragments(layout,item['id']);key=item['blocks'][0]['id']
    assert not layout['errors'] and len(fragments)>=2
    assert len([e for e in fragments[-1] if e.get('block_id')==key])>=2
    assert text(layout,key)==item['blocks'][0]['value']
    path=tmp_path/'complete.pptx';write_pptx(layout,path,{})
    from pptx import Presentation
    words=''.join(s.text for slide in Presentation(path).slides for s in slide.shapes
                  if s.has_text_frame and s.name!='shchem-continuation-heading')
    assert item['blocks'][0]['value'] in words and value['resources']==original['resources']
