"""Native paragraph order, inline picture positions, script visibility and personal books."""
from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace
import json
import pytest
from PIL import Image
from docx import Document
from docx.shared import Mm
from test_composition_render import app as app
from test_lesson_design import sample
from test_personal_textbook_knowledge import native_fixture as native_fixture
from test_personal_textbook_study import chapter,note,selected
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import *
from integrations.deeptutor_shchem_v1.desktop_composition_inline import NativeParagraphDisplay,bound_image_ids,display_tokens
from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition,layout_composition
from integrations.deeptutor_shchem_v1.desktop_composition_rich_render import flatten_elements
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import _word_images
from integrations.deeptutor_shchem_v1.offline_word_image_label_review import _objects
from integrations.deeptutor_shchem_v1.desktop_lesson_output import export_design,COMPOSITION_FILES,checked_file
from integrations.deeptutor_shchem_v1.desktop_composition_sources import import_knowledge_note,import_textbook_crop,knowledge_options
from integrations.deeptutor_shchem_v1.desktop_textbook_preparation_crops import TextbookPreparationCropService
from integrations.deeptutor_shchem_v1.desktop_personal_textbook_crops import book_page_selection

@pytest.fixture
def inline(tmp_path,app):
    buf=BytesIO();Image.new('RGB',(80,40),'#1071cc').save(buf,format='PNG');data=buf.getvalue()
    document=Document();paragraph=document.add_paragraph('图前H')
    paragraph.add_run('2').font.subscript=True;paragraph.add_run('O')
    paragraph.add_run().add_picture(BytesIO(data),width=Mm(18));paragraph.add_run('图后Na')
    paragraph.add_run('+').font.superscript=True;paragraph.add_run('尾部条件。')
    answer_paragraph=document.add_paragraph('仅教师答案H');answer_paragraph.add_run('2').font.subscript=True;answer_paragraph.add_run('O')
    raw=BytesIO();document.save(raw);display=NativeParagraphDisplay(raw.getvalue())
    native=display.reader.read(display.elements[0]).text
    assert native=='图前H_{2}O【待查看原文：图片或图形】图后Na^{+}尾部条件。'
    store=PreparationImageStore(tmp_path/'state/tasks/preparation-v1/images')
    asset=store.import_bytes(data,'合成行内图','测试作者','行内位置核对')
    text=block('text','question',native,1);picture=block('image','question',asset['asset_id'],1)
    objects=_objects(display.elements[0],display.doc,1,_word_images(display.elements,display.doc))
    display.bind(text,[picture],objects)
    assert abs(text['presentation']['images'][0]['width_mm']-18)<.01
    answer=block('text','answer','仅教师答案H_{2}O',2)
    display.bind(answer,[],[])
    item=resource('complete_question','合成原生题',[text,picture,answer],
        {'source_path':'合成原题.docx','source_sha256':'a'*64,'source_revision':'b'*64,'question_key':'synthetic-inline',
         'scope_review_sha256':'c'*64,'scope_status':'complete_native_scope',
         'hierarchy':dict.fromkeys(('paper_id','theme_big_question_id','printed_question_id','atomic_part_id'),'unknown')})
    value=new_composition();value=add_resource(value,item,{'page_id':value['pages'][0],'x':12,'y':12,'width':186,'height':273})
    facade=SimpleNamespace(paths=SimpleNamespace(task_root=tmp_path/'state/tasks'),preparation_image_bytes=store.load)
    return value,[asset],{asset['asset_id']:data},facade,native

@pytest.mark.parametrize('format',FORMATS)
@pytest.mark.parametrize('columns',[1,2])
def test_source_picture_stays_between_text_and_scripts_keep_their_positions(inline,format,columns):
    value,assets,data,facade,native=inline;value['format']=format;w,h=FORMATS[format]
    for rid in value['placements']:value['placements'][rid]={'page_id':value['pages'][0],'x':12,'y':12,'width':w-24,'height':h-24}
    value=arrange_composition(value,assets,data,columns=columns)
    layout=layout_composition(value,assets,data,student=True);assert not layout['errors']
    rows=[e for elements in layout['pages'].values() for e in elements if e['kind']=='rich_line'];assert rows
    parts=rows[0]['_parts'];image=next(p for p in parts if p['kind']=='image')
    left=next(p for p in parts if p.get('text')=='图前H');right=next(p for p in parts if p.get('text')=='图后Na')
    assert left['offset_x']<image['offset_x']<right['offset_x']
    assert abs(image['width']-18*(32/15 if format=='16:9' else 1))<.01
    assert value['resources'][0]['blocks'][0]['presentation']['images'][0]['width_mm']==18
    sub=next(p for p in parts if p.get('text')=='2');sup=next(p for p in parts if p.get('text')=='+')
    assert sub['font_px']<left['font_px'] and sup['font_px']<right['font_px']
    assert sub['offset_y']>left['offset_y'] and sup['offset_y']<right['offset_y']
    flat=[e for elements in layout['pages'].values() for e in flatten_elements(elements)]
    assert sum(e['kind']=='image' for e in flat)==1
    text=''.join(e.get('text','') for e in flat)
    assert '图前H2O' in text and '图后Na+尾部条件。' in text and '待查看原文' not in text and '仅教师答案' not in text
    assert value['resources'][0]['blocks'][0]['value']==native

def test_unlocked_edits_preserve_inline_bindings_and_cannot_drop_markers(inline):
    value,_,_,_,native=inline;item=value['resources'][0];text=item['blocks'][0]
    value=edit_resource(value,item['id'],locked=False)
    changed=edit_resource(value,item['id'],values={text['id']:'新增条件。'+native})
    assert changed['resources'][0]['original']==item['original']
    with pytest.raises(ValueError,match='标记'):edit_resource(value,item['id'],values={text['id']:'删去了原图位置'})
    bad=deepcopy(item);bad['blocks'][0]['presentation']['images'][0]['width_mm']=90
    with pytest.raises(ValueError,match='不能被静默替换'):validate_resource(bad)
    bad=deepcopy(item);bad['blocks'][1]['role']='answer'
    with pytest.raises(ValueError,match='跨答案范围'):validate_resource(bad)


def test_literal_braces_and_work_copy_additions_never_acquire_native_script_format():
    doc=Document();p=doc.add_paragraph('字面_{2}、^{+}；H');p.add_run('2').font.subscript=True;p.add_run('O')
    buf=BytesIO();doc.save(buf);native=NativeParagraphDisplay(buf.getvalue())
    value=native.reader.read(native.elements[0]).text;row=block('text','question',value,1);native.bind(row,[],[])
    tokens=display_tokens(row,{})
    assert ''.join(t['text'] for t in tokens)=='字面_{2}、^{+}；H2O'
    assert [t['text'] for t in tokens if t['script']=='subscript']==['2']
    changed=deepcopy(row);changed['value']='新增_{2}。'+value
    tokens=display_tokens(changed,{},value)
    assert ''.join(t['text'] for t in tokens)=='新增_{2}。字面_{2}、^{+}；H2O'
    assert [t['text'] for t in tokens if t['script']=='subscript']==['2']
    changed['value']=value.replace('；H_{2}O','；H_{3}O')
    assert not any(t['script']!='baseline' for t in display_tokens(changed,{},value))


def test_omml_scripts_use_authored_operands_and_keep_literal_math_text():
    from docx.oxml import parse_xml
    doc=Document();p=doc.add_paragraph('公式：')
    p._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        '<m:r><m:t>字面_{8}；</m:t></m:r><m:sSubSup><m:e><m:r><m:t>SO</m:t></m:r></m:e>'
        '<m:sub><m:r><m:t>4</m:t></m:r></m:sub><m:sup><m:r><m:t>2−</m:t></m:r></m:sup></m:sSubSup></m:oMath>'))
    buf=BytesIO();doc.save(buf);native=NativeParagraphDisplay(buf.getvalue())
    value=native.reader.read(native.elements[0]).text;row=block('text','question',value,1);native.bind(row,[],[])
    tokens=display_tokens(row,{})
    assert ''.join(t['text'] for t in tokens)=='公式：字面_{8}；SO42−'
    assert [(t['text'],t['script']) for t in tokens if t['script']!='baseline']==[('4','subscript'),('2−','superscript')]


def test_spaces_between_source_runs_do_not_discard_script_positions():
    doc=Document();p=doc.add_paragraph('  (4)加入无水CaCl');p.add_run('2').font.subscript=True
    p.add_run('的目的是');p.add_run('           ');p.add_run('。  ')
    buf=BytesIO();doc.save(buf);native=NativeParagraphDisplay(buf.getvalue())
    value=native.reader.read(native.elements[0]).text;row=block('text','question',value,1);native.bind(row,[],[])
    tokens=display_tokens(row,{})
    assert ''.join(t['text'] for t in tokens)=='(4)加入无水CaCl2的目的是           。'
    assert [(t['text'],t['script']) for t in tokens if t['script']!='baseline']==[('2','subscript')]

def test_actual_pdf_pptx_and_docx_have_inline_condition_order_without_duplicate_images(inline,tmp_path):
    from pptx import Presentation
    import pypdfium2 as pdfium
    value,assets,data,facade,native=inline;value=arrange_composition(value,assets,data)
    p=sample();p['lesson_design']=enable_composition(p['lesson_design']);p['lesson_design']['composition']=value;p['image_assets']=assets
    record=export_design(facade,p);p['lesson_design']['exports'].append(record)
    for name in ('lesson_presentation.pptx','student_presentation.pptx'):
        deck=Presentation(checked_file(facade,p['lesson_design'],record['id'],name));text=''.join(s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame)
        assert '图前H2O' in text and '图后Na+尾部条件。' in text and '待查看原文' not in text
        assert sum(s.shape_type==13 for slide in deck.slides for s in slide.shapes)==1
        assert ('仅教师答案' in text)==name.startswith('lesson_')
    for name in ('lesson_canvas.pdf','student_canvas.pdf'):
        pdf=pdfium.PdfDocument(str(checked_file(facade,p['lesson_design'],record['id'],name)))
        page=pdf[0];tp=page.get_textpage();text=tp.get_text_range();tp.close();page.close();pdf.close()
        assert '图前H' in text and '尾部条件。' in text and '待查看原文' not in text
        assert ('仅教师答案' in text)==name.startswith('lesson_')
    doc=Document(checked_file(facade,p['lesson_design'],record['id'],'student_worksheet.docx'))
    para=next(p for p in doc.paragraphs if '图前H' in p.text)
    assert para.text=='图前H2O图后Na+尾部条件。'
    assert len(doc.inline_shapes)==1 and len(para._p.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing'))==1
    assert next(r for r in para.runs if r.text=='2').font.subscript
    assert next(r for r in para.runs if r.text=='+').font.superscript


def test_wrapped_pdf_paragraph_is_searchable_once_with_native_scripts(inline):
    import pypdfium2 as pdfium
    value,assets,data,facade,_=inline;item=value['resources'][0];source=item['blocks'][0]['value']
    extra='完整条件只出现一次，不能因为跨行排版而在PDF里重复。'*16
    value=edit_resource(value,item['id'],locked=False)
    value=edit_resource(value,item['id'],values={item['blocks'][0]['id']:source+extra})
    value=arrange_composition(value,assets,data);p=sample();p['lesson_design']=enable_composition(p['lesson_design'])
    p['lesson_design']['composition']=value;p['image_assets']=assets;record=export_design(facade,p)
    p['lesson_design']['exports'].append(record)
    pdf=pdfium.PdfDocument(str(checked_file(facade,p['lesson_design'],record['id'],'student_canvas.pdf')));texts=[]
    for page in pdf:
        tp=page.get_textpage();texts.append(tp.get_text_range());tp.close();page.close()
    pdf.close();text=''.join(''.join(texts).split())
    assert text.count('完整条件只出现一次，不能因为跨行排版而在PDF里重复。')==16
    assert '图前H2O图后Na+尾部条件。'+extra in text

def personal_facade(f):
    crops=TextbookPreparationCropService(f.paths.workspace_root,f.paths.task_root,personal_books=f.workspace)
    return SimpleNamespace(paths=f.paths,textbook_bookshelf=lambda:{'reading_notes':[]},
        personal_textbook_study=lambda:f.study,textbook_workspace=lambda:f.workspace,
        personal_textbook_reference=f.study.reference,preparation_textbook_crop_preview=crops.preview,
        import_textbook_preparation_crops=crops.confirm_batch)

def test_personal_note_and_crop_keep_book_sha_unknown_printed_pages_and_closed_flags(native_fixture):
    f=native_fixture;view=note(f,chapter(f));facade=personal_facade(f)
    options=knowledge_options(facade)['reading_notes'];assert len(options)==1 and options[0]['origin']=='personal'
    before=f.state.snapshot();card=import_knowledge_note(facade,options[0]['selection'])
    assert f.state.snapshot()==before
    assert card['original']['evidence']['printed_pages']=='unknown' and card['original']['evidence']['source_sha256']==f.selection['revision']
    assert card['blocks'][0]['value']==view['notes'][0]['text'] and not card['generation_allowed']
    selection=book_page_selection(view['book'],2)
    page=TextbookPreparationCropService(f.paths.workspace_root,f.paths.task_root,personal_books=f.workspace).page(selection)
    bounds={'left':20,'top':20,'right':300,'bottom':300}
    preview=facade.preparation_textbook_crop_preview(selection,page['page_revision'],bounds,'合成裁图','课堂观察')
    request={'selection':selection,'page_revision':page['page_revision'],'bounds':bounds,'caption':'合成裁图','purpose':'课堂观察',
        'expected_preview_revision':preview['preview_revision'],'expected_image_sha256':preview['image_asset']['sha256'],'selection_confirmed':True}
    crop=import_textbook_crop(facade,request,[]);assert len(crop['image_assets'])==1
    assert crop['resources'][0]['original']['evidence']['printed_pages']=='unknown'
    assert f.source.read_bytes()==f.original and f.state.snapshot()==before
    assert not crop['resources'][0]['human_reviewed'] and not crop['resources'][0]['publication_allowed']

def test_stale_personal_note_or_changed_book_cannot_enter_composition(native_fixture):
    f=native_fixture;view=note(f,chapter(f));facade=personal_facade(f);choice=knowledge_options(facade)['reading_notes'][0]['selection']
    note(f,view,text='已修改的条件。')
    with pytest.raises(ValueError,match='已变化'):import_knowledge_note(facade,choice)
    fresh=knowledge_options(facade)['reading_notes'][0]['selection']
    (f.workspace.root/(f.selection['revision']+'.pdf')).write_bytes(b'corrupted')
    with pytest.raises(ValueError):import_knowledge_note(facade,fresh)
