"""Generated classroom content stays direct; real chemistry and source evidence stay exact."""
from copy import deepcopy
import json
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from test_desktop_preparation_renderer import _candidate, _artifact_map
from integrations.deeptutor_shchem_v1.desktop_teaching_framework import classroom_text
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import visible_blocks, resource, block
from integrations.deeptutor_shchem_v1.desktop_preparation_renderer import NativePreparationRenderer


@pytest.mark.parametrize('disclaimer', [
    '仅供参考。', '本教案仅供教学参考。', '请教师自行检查。', '课前请教师核对。',
    'AI可能有误。', 'AI生成内容可能存在错误。', '人工智能生成的内容可能不准确。',
    '不保证准确性和完整性。', '本内容不能替代教师判断。',
])
def test_disclaimer_is_removed_while_next_supported_sentence_stays(disclaimer):
    chemical = '组成和结构相似时，比较分子间作用力与沸点的关系。'
    assert classroom_text(disclaimer + '\n' + chemical) == chemical


@pytest.mark.parametrize('chemical', [
    '不能用范德华力判断分子的稳定性。',
    '不要把范德华力误认为化学键。',
    '比较方程式两边的原子数，核对方程式系数。',
    '请学生核对计算单位，再解释结果。',
    '组成和结构相似时才能作此比较；结果与温度、压强有关。',
])
def test_pedagogical_checks_and_chemical_limitations_are_preserved(chemical):
    assert classroom_text(chemical) == chemical


def test_appended_reference_clause_does_not_erase_supported_chemistry():
    assert classroom_text('比较时保持温度相同，仅供参考。') == '比较时保持温度相同'
    assert classroom_text('HF沸点为50℃（待核实）。') == ''


def test_complete_source_question_is_not_rewritten_by_disclaimer_cleanup():
    item = resource('complete_question', '原题', [block('text', 'question', '题设原文：仅供参考。请教师核查实验记录。', 2)],
                    {'source_path': '原卷.docx', 'source_sha256': 'a'*64, 'source_revision': 'b'*64,
                     'question_key': 'source-question', 'scope_review_sha256': 'c'*64,
                     'scope_status': 'complete_native_scope',
                     'hierarchy': dict.fromkeys(('paper_id', 'theme_big_question_id', 'printed_question_id', 'atomic_part_id'), 'unknown')})
    before = deepcopy(item)
    assert visible_blocks(item)[0]['value'] == item['blocks'][0]['value']
    assert item['original'] == before['original'] and item['original_sha256'] == before['original_sha256']


def test_real_docx_ppt_text_tables_and_speaker_notes_are_direct_without_mutating_candidate(tmp_path):
    candidate = _candidate(); body = '保持组成和结构相似这一比较条件。'
    disclaimer = '本教案仅供参考。请教师自行检查。AI可能有误。'
    candidate['slides'][1]['content'].append(body + '\n' + disclaimer)
    candidate['slides'][1]['teacher_notes'] += '\n' + body + '\n' + disclaimer
    candidate['activities'][0]['teacher_action'] += '\n' + disclaimer
    candidate['uncertainties'] = [{'field': 'source', 'description': '待核验原始数据', 'teacher_action': '检查原页'}]
    before = deepcopy(candidate)
    artifacts = _artifact_map(NativePreparationRenderer().render(candidate, output_kind='joint', output_dir=tmp_path/'native'))
    stored = json.loads(Path(artifacts['candidate_json']['path']).read_text(encoding='utf-8'))
    assert candidate == before == stored
    doc = Document(artifacts['lesson_plan_docx']['path'])
    word_text = '\n'.join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    deck = Presentation(artifacts['pptx']['path'])
    ppt_text = '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    notes = '\n'.join(s.notes_slide.notes_text_frame.text for s in deck.slides)
    assert body in ppt_text and body in notes
    for text in (word_text, ppt_text, notes):
        assert not any(v in text for v in ('仅供参考', '教师自行检查', 'AI可能有误', '待核验原始数据'))
    assert stored['candidate_only'] is True and stored['publication_allowed'] is False
