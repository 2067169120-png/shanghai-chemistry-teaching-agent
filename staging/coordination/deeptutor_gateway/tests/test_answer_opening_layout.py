from io import BytesIO
from copy import deepcopy
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from test_word_question_export import item, picture, output_doc
from integrations.deeptutor_shchem_v1.desktop_word_question_export import export_word_questions
from integrations.deeptutor_shchem_v1.desktop_mixed_paper_export import build_mixed_paper_docx
from integrations.deeptutor_shchem_v1.desktop_answer_layout import keep_answer_opening


def chain_value(p):
    node=p.find('./'+qn('w:pPr')+'/'+qn('w:keepNext'))
    return node.get(qn('w:val'),'1') if node is not None else None


@pytest.mark.parametrize('mixed',[False,True])
def test_answer_label_is_linked_to_picture_but_not_to_next_question(mixed):
    doc=Document();doc.add_paragraph('27. 合成题');doc.add_paragraph('【答案】27. 合成答案')
    picture(doc,'white');doc.add_paragraph('补充说明')
    row=item(doc,[1],[2,3,4]);before=doc.element.xml
    result=(build_mixed_paper_docx('合成',[{'kind':'word_question','question':row}]) if mixed
            else export_word_questions('合成',[row]))
    teacher=output_doc(result,'teacher')
    label=next(p for p in teacher.paragraphs if p.text.startswith('【答案】'))
    assert label.paragraph_format.keep_with_next is True
    next_p=label._p.getnext()
    assert next_p.xpath('.//w:drawing') and chain_value(next_p) is None
    assert not teacher.paragraphs[-1].paragraph_format.keep_with_next
    assert doc.element.xml==before
    if mixed:
        heading=next(p for p in teacher.paragraphs if p.text.startswith('参考答案 ·'))
        assert heading.paragraph_format.keep_with_next is True


def test_short_label_and_blank_lines_bind_to_first_image_only():
    doc=Document();doc.add_paragraph('答案1');doc.add_paragraph();doc.add_paragraph();picture(doc,'white');doc.add_paragraph('尾段')
    blocks=list(doc.element.body)[:-1]
    keep_answer_opening(blocks)
    assert [chain_value(p) for p in blocks]==['1','1','1',None,None]


def test_long_answer_can_flow_over_pages():
    doc=Document();doc.add_paragraph('答案'*100);doc.add_paragraph('下一段')
    blocks=list(doc.element.body)[:-1];before=doc.element.xml;keep_answer_opening(blocks)
    assert doc.element.xml==before


def test_last_short_answer_does_not_chain_into_another_question():
    doc=Document();p=doc.add_paragraph('答案');keep_answer_opening([p._p]);assert chain_value(p._p) is None


def test_explicit_page_break_is_not_removed_or_linked_across():
    doc=Document();doc.add_paragraph('答案');p=doc.add_paragraph('另页说明');p.paragraph_format.page_break_before=True
    before=doc.element.xml;keep_answer_opening(list(doc.element.body)[:-1]);assert doc.element.xml==before


def test_first_formula_is_real_content_not_blank_spacer():
    doc=Document();p=doc.add_paragraph();p._p.append(OxmlElement('m:oMath'));doc.add_paragraph('说明')
    before=doc.element.xml;keep_answer_opening(list(doc.element.body)[:-1]);assert doc.element.xml==before


def test_table_rows_are_not_forced_into_one_unbroken_answer():
    doc=Document();doc.add_paragraph('答案表');table=doc.add_table(rows=30,cols=1)
    for cell in table.column_cells(0):cell.text='数据'
    before=table._tbl.xml;keep_answer_opening(list(doc.element.body)[:-1])
    assert chain_value(doc.paragraphs[0]._p)=='1' and table._tbl.xml==before
