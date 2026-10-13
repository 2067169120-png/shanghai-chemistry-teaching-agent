"""Independent adversarial cases for source-bound answer and actual geometry."""
from copy import deepcopy
from uuid import uuid4

import pytest
from pptx import Presentation

from test_question_answer_reveal import qt_app, question, personal, ppt_text, text
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    add_resource, block, digest, edit_resource, place_resource, resource,
    set_click_reveal, set_student_arrangement, validate_composition)
from integrations.deeptutor_shchem_v1.desktop_question_answer_reveal import (
    answer_reveal_options, clear_question_answer_reveal,
    configure_question_answer_reveal, question_answer_reveal_errors)
from integrations.deeptutor_shchem_v1.desktop_composition_render import (
    layout_composition, write_pptx)
from integrations.deeptutor_shchem_v1.desktop_composition_rich_render import flatten_elements


def rich_answer(answer, scripts=()):
    value, old, aid, assets, data = question(answer=answer)
    rows = deepcopy(old['blocks'])
    rows[-3]['presentation'] = {'schema': 'shchem.native-word-paragraph.v1',
        'images': [], 'scripts': list(scripts)}
    item = resource('complete_question', old['title'], rows, old['original']['evidence'])
    value['resources'] = [item]
    value['placements'] = {item['id']: value['placements'][old['id']]}
    return value, item, aid, assets, data


def test_native_one_line_answer_rejects_editable_boxes_outside_reserved_slot(qt_app):
    value, item, aid, assets, data = rich_answer('AB\t\t\t\t')
    before = deepcopy(value)
    with pytest.raises(ValueError, match='原答案超出标题旁空间'):
        configure_question_answer_reveal(value, item['id'], [aid], assets, data)
    assert value == before
    # The same intact source remains usable when the teacher selects no reveal.
    full_page = configure_question_answer_reveal(value, item['id'], [], assets, data)
    assert full_page['resources'][0]['blocks'] == item['blocks']
    assert not question_answer_reveal_errors(full_page, item['id'], assets, data)


@pytest.mark.parametrize('separator', ['\u2028', '\u2029'])
def test_native_unicode_line_breaks_cannot_bypass_one_line_answer_measurement(qt_app, separator):
    value, item, aid, assets, data = rich_answer('答案A'+separator+'答案B')
    before = deepcopy(value)
    with pytest.raises(ValueError, match='一行'):
        configure_question_answer_reveal(value, item['id'], [aid], assets, data)
    assert value == before


def test_current_teacher_rich_edit_preserves_native_subscript_and_only_exact_original_id(qt_app, tmp_path):
    value, item, aid, assets, data = rich_answer('H_{2}O',
        [{'start': 1, 'end': 5, 'script': 'subscript'}])
    value = edit_resource(value, item['id'], locked=False)
    value = edit_resource(value, item['id'], values={aid: 'NaH_{2}O'})
    value = configure_question_answer_reveal(value, item['id'], [aid], assets, data)
    current = value['resources'][0]
    assert current['original'] == item['original']
    option = answer_reveal_options(current)[0]
    assert option['text'] == 'NaH_{2}O' and option['original_text'] == 'H_{2}O'
    assert option['working_value_changed'] and option['selected']
    classroom = layout_composition(value, assets, data, include_reveals=True)
    runs = [e for page in classroom['pages'].values() for e in flatten_elements(page) if e.get('reveal_id')]
    assert ''.join(e['text'] for e in runs) == 'NaH2O'
    assert {e['block_id'] for e in runs} == {aid} and {e['role'] for e in runs} == {'answer'}
    assert any(e['text'] == '2' and e['script'] == 'AlignSubScript' for e in runs)
    title = next(e for page in classroom['pages'].values() for e in page if e.get('text') == item['title'])
    assert title['x']+title['width']+10 <= min(e['x'] for e in runs)+.001
    assert max(e['x']+e['width'] for e in runs) <= 305+.001
    write_pptx(classroom, tmp_path/'current-rich.pptx', data)
    slide = Presentation(tmp_path/'current-rich.pptx').slides[0]
    groups = [s for s in slide.shapes if s.name.startswith('shchem-click-')]
    assert len(groups) == 1 and groups[0].name.endswith('-answer-'+aid)
    assert ppt_text(groups[0].shapes) == 'NaH2O'
    assert {node.get('spid') for node in slide._element.xpath('.//p:timing//p:spTgt')} == {str(groups[0].shape_id)}
    assert 'TEACHER_REASON_PRIVATE' not in ppt_text(slide.shapes)


@pytest.mark.parametrize('scope', ['teacher_only', 'answer_only'])
def test_complete_label_cannot_unlock_source_scope_without_complete_native_evidence(qt_app, scope):
    value, item, aid, assets, data = question()
    original = value['resources'][0]['original']
    original['evidence']['scope_status'] = scope
    value['resources'][0]['original_sha256'] = digest(original)
    before = deepcopy(value)
    with pytest.raises(ValueError, match='完整范围复核'):
        configure_question_answer_reveal(value, item['id'], [aid], assets, data)
    assert value == before


def test_teacher_continuation_overflow_is_checked_even_when_independent_student_fits(qt_app):
    value, item, aid, assets, data = question()
    value = edit_resource(value, item['id'], locked=False)
    reason_id = item['blocks'][-2]['id']
    value = edit_resource(value, item['id'], values={reason_id: '教师完整推理与条件。'*80})
    value = configure_question_answer_reveal(value, item['id'], [aid], assets, data)
    placement = deepcopy(value['placements'][item['id']])
    assert placement['continuations']
    placement['continuations'][0]['height'] = 20
    value = place_resource(value, item['id'], placement)
    value = set_student_arrangement(value, 2)
    for reveal in (False, True):
        assert not layout_composition(value, assets, data, include_reveals=reveal)['errors']
    assert layout_composition(value, assets, data, student=False)['errors']
    assert question_answer_reveal_errors(value, item['id'], assets, data)


def test_shared_teacher_only_frame_overlap_is_checked_when_student_flow_hides_it(qt_app):
    value, item, aid, assets, data = question()
    value = configure_question_answer_reveal(value, item['id'], [aid], assets, data)
    teacher_only = personal('隐含教师备课说明', 'teacher_note')
    value = add_resource(value, teacher_only, dict(page_id=value['placements'][item['id']]['page_id'],
        x=20, y=20, width=80, height=25))
    value = set_student_arrangement(value, 1)
    before = deepcopy(value)
    assert not layout_composition(value, assets, data, include_reveals=True)['errors']
    assert any('其他素材重叠' in e for e in question_answer_reveal_errors(value, item['id'], assets, data))
    assert value == before


def test_three_existing_click_resources_keep_nine_native_groups_and_frames(qt_app, tmp_path):
    value, item, aid, assets, data = question()
    existing = []
    for ordinal, size in enumerate((4, 4, 1)):
        rows = [block('text', 'personal', f'原点击{ordinal}的步骤{n}') for n in range(size)]
        old = resource('personal_text', f'旧点击素材{ordinal}', rows,
            dict(source_path='作者输入', source_sha256='unknown', source_revision='unknown'))
        if size > 1: old['exposition'] = {'schema': 'shchem.classroom-exposition.v1', 'kind': 'process'}
        page = 'CP-'+uuid4().hex; value['pages'].append(page)
        value = add_resource(value, old, dict(page_id=page, x=12, y=12, width=296, height=156))
        value = set_click_reveal(value, old['id'], True); existing.append(old)
    before = deepcopy(value)
    baseline = layout_composition(value, assets, data, include_reveals=True)
    assert not baseline['errors']
    write_pptx(baseline, tmp_path/'baseline.pptx', data)
    value = configure_question_answer_reveal(value, item['id'], [aid], assets, data)
    assert value['click_reveals'] == before['click_reveals']
    assert all(value['placements'][r['id']] == before['placements'][r['id']] for r in existing)
    assert all(r == next(a for a in value['resources'] if a['id'] == r['id']) for r in existing)
    write_pptx(layout_composition(value, assets, data, include_reveals=True), tmp_path/'configured.pptx', data)
    def groups(path):
        return {s.name: ppt_text(s.shapes) for slide in Presentation(path).slides for s in slide.shapes
            if s.name.startswith('shchem-click-')}
    prior, now = groups(tmp_path/'baseline.pptx'), groups(tmp_path/'configured.pptx')
    assert len(prior) == 9 and len(now) == 10
    assert {name: now[name] for name in prior} == prior
    assert list(set(now)-set(prior)) == ['shchem-click-'+item['id']+'-answer-'+aid]
    cleared = clear_question_answer_reveal(set_student_arrangement(value, 2), item['id'])
    assert cleared['placements'] == value['placements'] and cleared['click_reveals'] == before['click_reveals']
    assert cleared['resources'][0] == item
    assert validate_composition(cleared) == cleared
