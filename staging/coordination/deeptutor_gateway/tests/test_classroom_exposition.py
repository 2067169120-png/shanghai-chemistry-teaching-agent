"""Native diagrams and complete Word notes share the editable work copy."""
from copy import deepcopy
from pathlib import Path
import pytest
from docx import Document
from pptx import Presentation
from test_composition_render import app as app
from test_teaching_output_purposes import bundle
from test_lesson_design_ui import desk
from test_question_explorer_ui import window
from test_desktop_library_ui import qt_app
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import *
from integrations.deeptutor_shchem_v1.desktop_composition_render import layout_composition,write_pptx,page_image
from integrations.deeptutor_shchem_v1.desktop_composition_rich_render import flatten_elements
from integrations.deeptutor_shchem_v1.desktop_composition_exposition import SCHEMA as EXPOSITION
from integrations.deeptutor_shchem_v1.desktop_lesson_output import export_design,copy_primary_bundle,checked_file,write_documents

def diagram(kind='process',long=False):
    rows=['比较条件\n组成和结构相似','分子间作用\n强弱不同','物理性质\n沸点不同']
    if long:rows[1]+='完整条件不可截断。'*120
    item=resource('personal_text','条件、作用与性质',[block('text','personal',t) for t in rows],
        {'source_path':'合成作者输入','source_sha256':'unknown','source_revision':'unknown'})
    item['exposition']={'schema':EXPOSITION,'kind':kind}
    c=new_composition('16:9');return add_resource(c,item,{'page_id':c['pages'][0],'x':12,'y':12,'width':296,'height':156})

@pytest.mark.parametrize('kind',['process','process_vertical','comparison'])
def test_native_text_arrows_and_preview_follow_complete_edited_blocks(app,tmp_path,kind):
    c=diagram(kind);original=deepcopy(c['resources'][0]['original'])
    r=c['resources'][0];c=edit_resource(c,r['id'],locked=False)
    c=edit_resource(c,r['id'],values={r['blocks'][0]['id']:'实际比较条件\n保留完整条件'})
    result=layout_composition(c,[],{},student=True,include_reveals=True);assert not result['errors']
    assert not page_image(result,next(iter(result['pages']))).isNull()
    elements=list(flatten_elements(next(iter(result['pages'].values()))))
    assert sum(e['kind']=='diagram_box' for e in elements)==3
    assert sum(e['kind']=='diagram_arrow' for e in elements)==(0 if kind=='comparison' else 2)
    assert '实际比较条件' in ''.join(e.get('text','') for e in elements)
    target=tmp_path/'diagram.pptx';write_pptx(result,target,{})
    deck=Presentation(target);assert '实际比较条件' in ''.join(s.text for s in deck.slides[0].shapes if s.has_text_frame)
    assert sum(s.shape_type==1 for s in deck.slides[0].shapes)==(3 if kind=='comparison' else 5)
    assert c['resources'][0]['original']==original

def test_long_diagram_blocks_export_instead_of_clipping_source(app):
    c=diagram(long=True);snapshot=deepcopy(c)
    assert layout_composition(c,[],{},student=True)['errors'] and c==snapshot

def test_process_clicks_reveal_each_node_and_its_incoming_arrow(app,tmp_path):
    c=diagram('process_vertical');item=c['resources'][0];c=set_click_reveal(c,item['id'],True)
    result=layout_composition(c,[],{},student=True,include_reveals=True);target=tmp_path/'steps.pptx'
    write_pptx(result,target,{})
    slide=Presentation(target).slides[0]
    groups=[s for s in slide.shapes if s.shape_type==6]
    assert len(groups)==len(slide._element.xpath('.//p:cTn[@nodeType="clickEffect"]'))==3
    assert '条件、作用与性质' in ''.join(s.text for s in slide.shapes if s.has_text_frame)
    assert '条件、作用与性质' not in ''.join(s.text for group in groups for s in group.shapes if s.has_text_frame)
    for index in (1,2):
        assert any(s.auto_shape_type==36 for s in groups[index].shapes if s.shape_type==1)

def test_teacher_roles_cannot_be_smuggled_into_a_visible_diagram():
    c=diagram();c['resources'][0]['blocks'][0]['role']='teacher_note'
    with pytest.raises(ValueError):validate_composition(c)

def test_classroom_word_notes_include_revealed_explanation_and_keep_old_version(tmp_path,app):
    facade,p,old=bundle(tmp_path/'source',True);old_state=deepcopy(p['lesson_design'])
    c=p['lesson_design']['composition'];r=c['resources'][0]
    c['format']='16:9';c['placements'][r['id']].update(width=296,height=70)
    regular=resource('personal_text','常规课堂说明',[block('text','personal','先观察，再解释。')],
        {'source_path':'合成作者输入','source_sha256':'unknown','source_revision':'unknown'})
    c=add_resource(c,regular,{'page_id':c['pages'][0],'x':12,'y':90,'width':296,'height':78})
    c=set_click_reveal(c,r['id'],True);p['lesson_design']['composition']=c;p['lesson_design']['classroom_notes']=True
    record=export_design(facade,p);p['lesson_design']['exports'].append(record)
    word=Document(checked_file(facade,p['lesson_design'],record['id'],'student_worksheet.docx'))
    text='\n'.join(t.text for t in word.paragraphs)
    assert '课堂讲义' in text and '课堂正文唯一标记' in text and '教师活动：' not in text
    folder=Path(copy_primary_bundle(facade,p['lesson_design'],record['id'],tmp_path))
    assert {f.name for f in folder.iterdir()}=={'教案.docx','课堂演示.pptx','课堂讲义.docx'}
    # The selected historical file retains its actual learning-sheet identity.
    historical=Path(copy_primary_bundle(facade,p['lesson_design'],old['id'],tmp_path))
    assert {f.name for f in historical.iterdir()}=={'教案.docx','课堂演示.pptx'}
    assert old_state['exports'][0]==old

def test_native_notes_switch_changes_one_design_step_and_undo_restores(desk):
    d,app,page=desk;before=deepcopy(d.history.value)
    d.classroom_notes.setChecked(True);assert d.history.value['classroom_notes'] is True
    d.travel(False);assert d.history.value==before and not d.classroom_notes.isChecked()

@pytest.mark.parametrize('kind',['process','process_vertical','comparison'])
def test_classroom_handout_keeps_narrative_and_complete_diagrams_in_word_tables(tmp_path,kind):
    from integrations.deeptutor_shchem_v1.desktop_lesson_design import new_design,new_node
    p={'topic':'结构与性质','audience':'高二','objective':'解释完整条件'}
    plan=new_design(p);plan['schema_version']=DESIGN_SCHEMA;plan['composition']=diagram(kind)
    plan['classroom_notes']=True
    visible=new_node();visible.update(title='知识讲解',student_material=True,
        material_text='先区分分子内与分子间的作用。\n\n再把结构、作用和性质联系起来。')
    hidden=new_node();hidden.update(material_text='教师专用材料不进入学生讲义。')
    plan['nodes']=[visible,hidden];p['lesson_design']=plan;snapshot=deepcopy(p)
    write_documents(p,tmp_path,{})
    student=Document(tmp_path/'student_worksheet.docx')
    body='\n'.join(t.text for t in student.paragraphs)
    assert visible['material_text'].split('\n\n')[0] in body
    assert visible['material_text'].split('\n\n')[1] in body
    assert hidden['material_text'] not in body and '教师活动：' not in body
    assert len(student.tables)==1
    table=student.tables[0]
    values=[c.text for r in table.rows for c in r.cells]
    assert all(b['value'] in values for b in plan['composition']['resources'][0]['blocks'])
    assert p==snapshot

def test_notes_preserve_selected_node_network_image_once_and_exclude_teacher_image(tmp_path):
    from io import BytesIO
    from hashlib import sha256
    from zipfile import ZipFile
    from PIL import Image
    from integrations.deeptutor_shchem_v1.desktop_lesson_design import new_design,new_node
    from integrations.deeptutor_shchem_v1.desktop_preparation_images import image_info
    def image(color):
        stream=BytesIO();Image.new('RGB',(400,200),color).save(stream,format='PNG');data=stream.getvalue()
        sha=sha256(data).hexdigest()
        return {'asset_id':'IMG-'+sha,'sha256':sha,'caption':'合成知识网','source':'合成测试图片',
                'purpose':'测试材料选择与字节保存',**image_info(data)},data
    visible,data=image('blue');hidden,private=image('red')
    p={'topic':'结构与性质','audience':'高二','objective':'整理知识','image_assets':[visible,hidden]}
    plan=new_design(p);plan['classroom_notes']=True
    first=new_node();first.update(title='知识联系',student_material=True,image_ids=[visible['asset_id']])
    duplicate=new_node();duplicate.update(title='归纳',student_material=True,image_ids=[visible['asset_id']],material_text='回扣联系。')
    teacher=new_node();teacher.update(title='教师材料',image_ids=[hidden['asset_id']])
    plan['nodes']=[first,duplicate,teacher];p['lesson_design']=plan;snapshot=deepcopy(p)
    write_documents(p,tmp_path,{visible['asset_id']:data,hidden['asset_id']:private})
    doc=Document(tmp_path/'student_worksheet.docx')
    assert len(doc.inline_shapes)==1 and '回扣联系。' in '\n'.join(x.text for x in doc.paragraphs)
    with ZipFile(tmp_path/'student_worksheet.docx') as package:
        media=[package.read(n) for n in package.namelist() if '/media/' in n]
    assert data in media and private not in media and p==snapshot

def test_full_notes_preserve_source_copy_hidden_roles_and_native_text(app,tmp_path):
    c=diagram();r=c['resources'][0]
    original=deepcopy(r['original'])
    r['exposition']={'schema':EXPOSITION,'kind':'notes'}
    c=edit_resource(c,r['id'],values={r['blocks'][0]['id']:'定义\n分子间作用力与分子内化学键的作用对象不同。\n\n条件\n比较组成和结构相似的分子物质。'})
    result=layout_composition(c,[],{},student=True,include_reveals=True)
    assert not result['errors']
    path=tmp_path/'notes.pptx';write_pptx(result,path,{})
    text=''.join(s.text for s in Presentation(path).slides[0].shapes if s.has_text_frame)
    assert '作用对象不同' in text and '组成和结构相似' in text
    assert c['resources'][0]['original']==original

def test_full_notes_continue_or_report_overflow_instead_of_reducing_text(app):
    c=diagram(long=True);c['resources'][0]['exposition']={'schema':EXPOSITION,'kind':'notes'}
    snapshot=deepcopy(c)
    assert layout_composition(c,[],{},student=True)['errors'] and c==snapshot

def test_full_notes_word_keeps_all_paragraphs_and_headings(tmp_path):
    from integrations.deeptutor_shchem_v1.desktop_lesson_design import new_design
    p={'topic':'知识讲解','audience':'高二','objective':'解释条件'}
    plan=new_design(p);plan['schema_version']=DESIGN_SCHEMA;plan['composition']=diagram()
    plan['classroom_notes']=True;plan['composition']['resources'][0]['exposition']={'schema':EXPOSITION,'kind':'notes'}
    p['lesson_design']=plan;write_documents(p,tmp_path,{})
    doc=Document(tmp_path/'student_worksheet.docx')
    assert not doc.tables
    body='\n'.join(p.text for p in doc.paragraphs)
    assert all(b['value'] in body for b in plan['composition']['resources'][0]['blocks'])
