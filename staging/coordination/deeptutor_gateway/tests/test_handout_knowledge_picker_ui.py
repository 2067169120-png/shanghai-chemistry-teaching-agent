"""Source-bound Word section selection with controllable, isolated Qt callbacks."""
from copy import deepcopy
from hashlib import sha256
from types import SimpleNamespace

import pytest

pytest.importorskip("PySide6")
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QDialog

from test_desktop_library_ui import _ManualBridge, _png_bytes, qt_app
from test_desktop_studio_ui import settle


CATALOG = "读取讲义知识资料目录"
OUTLINE = "读取讲义知识栏目"
IMAGE = "读取讲义知识原图"
CONFIRM = "核对所选讲义知识栏目"
PLACEHOLDER = "【待查看原文：图片或图形】"


def _digest(value):
    return sha256(value.encode("utf-8")).hexdigest()


def _source(number, *, role="handout"):
    return {
        "schema": "shchem.handout-knowledge-source.v1",
        "batch_id": "synthetic-word-batch",
        "source_id": f"synthetic-source-{number:03}",
        "source_name": f"合成讲义第{number:03}份.docx",
        "source_sha256": _digest(f"source-bytes-{number}"),
        "source_role": role,
    }


def _outline(source, *, count=3, image_bytes=None):
    """The H3 section really begins at block 64, not at its list position."""
    sections = [
        {
            "id": f"word-heading-b{64 + index * 14}",
            "level": 3,
            "title": ("知识点1范德华力" if index == 0 else f"知识点{index + 1}合成栏目"),
            "start": 64 + index * 14,
            "end": 77 + index * 14,
        }
        for index in range(count)
    ]
    headings = {section["start"]: section["title"] for section in sections}
    bound = {
        **deepcopy(source),
        "source_path": f"C:/synthetic-test-data/{source['source_name']}",
        "source_revision": _digest(source["source_id"] + "-revision"),
    }
    blocks = [
        {
            "index": index,
            "text": headings.get(index, f"{source['source_name']}：合成原文区块{index}，保留完整段落。"),
            "label": f"原文区块{index}",
            "warnings": [],
        }
        for index in range(1, 64 + count * 14)
    ]
    assets = []
    if image_bytes is not None:
        assets.append({
            "asset_id": "word-b65-image1",
            "block_index": 65,
            "mime_type": "image/png",
            "sha256": sha256(image_bytes).hexdigest(),
            "bytes_count": len(image_bytes),
            "preview_supported": True,
            "label": "合成知识结构原图",
        })
    return {
        "source": bound,
        "outline_revision": _digest(source["source_id"] + f"-outline-{count}"),
        "sections": sections,
        "preview": {
            "source_name": source["source_name"],
            "source_sha256": source["source_sha256"],
            "revision": bound["source_revision"],
            "blocks": blocks,
            "assets": assets,
            "sections": deepcopy(sections),
            "warnings": [],
        },
    }


class _SyntheticLibrary:
    def __init__(self, *, source_count=2, section_count=3, role="handout"):
        self.sources = [_source(index, role=role) for index in range(source_count)]
        self.image_bytes = _png_bytes()
        self.outlines = {
            source["source_id"]: _outline(source, count=section_count, image_bytes=self.image_bytes)
            for source in self.sources
        }
        self.catalog_calls = 0
        self.outline_calls = []
        self.image_calls = []
        self.context_calls = []

    def catalog(self, _facade):
        self.catalog_calls += 1
        return {"sources": deepcopy(self.sources), "issues": []}

    def outline(self, _facade, source):
        self.outline_calls.append(deepcopy(source))
        return deepcopy(self.outlines[source["source_id"]])

    def image(self, _facade, source):
        self.image_calls.append(deepcopy(source))
        return {
            "source": deepcopy(source),
            "bytes": self.image_bytes,
            "info": {"width": 1200, "height": 800, "content_type": "image/png"},
        }

    def context(self, _facade, source):
        from integrations.deeptutor_shchem_v1.desktop_handout_images import handout_image_context

        self.context_calls.append(deepcopy(source))

        def preview(_batch_id, source_id, *, read_only=False):
            assert read_only is True
            return deepcopy(self.outlines[source_id]["preview"])

        return handout_image_context(SimpleNamespace(imported_word_preview=preview), source)


def _legacy_image_source(library, number=0):
    outline = library.outlines[library.sources[number]["source_id"]]
    bound = outline["source"]
    asset = outline["preview"]["assets"][0]
    return {
        "schema": "shchem.handout-picture.v1",
        **{key: bound[key] for key in (
            "batch_id", "source_id", "source_name", "source_path", "source_sha256", "source_revision"
        )},
        **{key: asset[key] for key in ("asset_id", "block_index", "mime_type", "sha256", "bytes_count")},
    }


@pytest.fixture
def make_picker(qt_app, monkeypatch):
    from integrations.deeptutor_shchem_v1.desktop_workbench import handout_knowledge_picker as module

    created = []

    def make(*, library=None, initial_source=None):
        library = library or _SyntheticLibrary()
        bridge = _ManualBridge()
        bridge.shutdown = lambda _timeout: None
        monkeypatch.setattr(module, "DesktopTaskBridge", lambda _parent: bridge)
        monkeypatch.setattr(module, "handout_knowledge_sources", library.catalog)
        monkeypatch.setattr(module, "handout_knowledge_outline", library.outline)
        monkeypatch.setattr(module, "handout_image_preview", library.image)
        monkeypatch.setattr(module, "handout_image_context", library.context, raising=False)
        dialog = module.HandoutKnowledgePicker(SimpleNamespace(), initial_source=initial_source)
        dialog.show()
        settle(qt_app)
        created.append(dialog)
        return SimpleNamespace(dialog=dialog, bridge=bridge, library=library, app=qt_app)

    yield make
    for dialog in created:
        if dialog.result() != QDialog.DialogCode.Accepted:
            dialog.reject()
        dialog.deleteLater()
    settle(qt_app)


def _catalog(case):
    case.bridge.succeed(case.bridge.take(CATALOG), run=True)


def _select_source(case, number=0):
    wanted = case.library.sources[number]
    index = next(
        index for index in range(case.dialog.source.count())
        if case.dialog.source.itemData(index) == wanted
    )
    case.dialog.source.setCurrentIndex(index)


def _ready(case, number=0):
    _catalog(case)
    _select_source(case, number)
    case.bridge.succeed(case.bridge.take(OUTLINE), run=True)


def _check(case, index, title=None):
    item = case.dialog.sections.item(index)
    case.dialog.sections.setCurrentItem(item)
    item.setCheckState(Qt.CheckState.Checked)
    if title is not None:
        case.dialog.title.setText(title)


def _confirm(case):
    case.dialog.choose()
    case.bridge.succeed(case.bridge.take(CONFIRM), run=True)


def test_catalog_of_98_sources_reads_no_outlines_until_a_source_is_chosen(make_picker):
    case = make_picker(library=_SyntheticLibrary(source_count=98))
    assert [task.label for task in case.bridge.pending] == [CATALOG]
    _catalog(case)
    assert case.library.catalog_calls == 1
    assert case.library.outline_calls == [] and not case.bridge.pending
    assert case.dialog.source.count() == 99 and case.dialog.sections.count() == 0
    case.dialog.source_search.setText("第097份")
    assert case.dialog.source.count() == 2
    assert case.library.outline_calls == [] and not case.bridge.pending
    _select_source(case, 97)
    assert [task.label for task in case.bridge.pending] == [OUTLINE]
    case.bridge.succeed(case.bridge.take(OUTLINE), run=True)
    assert case.library.outline_calls == [case.library.sources[97]]
    assert "合成原文区块105" in case.dialog.reader.browser.toPlainText()


def test_h3_ranges_require_questions_and_return_original_sections_in_document_order(make_picker):
    case = make_picker()
    _ready(case)
    value = case.dialog
    original = deepcopy(case.library.outlines[case.library.sources[0]["source_id"]])
    assert "H3" in value.sections.item(0).text()
    assert "64" in value.sections.item(0).text() and "77" in value.sections.item(0).text()
    assert value.sections.item(0).data(Qt.ItemDataRole.UserRole) == original["sections"][0]
    _check(case, 1)
    value.title.setText("  ")
    assert not value.confirm.isEnabled()
    value.choose()
    assert value.requests is None and not case.bridge.pending
    value.title.setText("第二段的知识怎样联系？")
    _check(case, 0, "如何用第一段解释分子间作用？")
    assert value.confirm.isEnabled()
    _confirm(case)
    assert value.result() == QDialog.DialogCode.Accepted
    assert [request["section"] for request in value.requests] == original["sections"][:2]
    assert [request["title"] for request in value.requests] == [
        "如何用第一段解释分子间作用？", "第二段的知识怎样联系？"
    ]
    for request in value.requests:
        assert set(request) == {"source", "outline_revision", "section", "title", "teacher_only"}
        assert request["source"] == original["source"]
        assert request["outline_revision"] == original["outline_revision"]
        assert request["teacher_only"] is False
    assert case.library.outlines[case.library.sources[0]["source_id"]] == original


def test_section_search_preserves_selection_and_full_original_text(make_picker):
    case = make_picker()
    _ready(case)
    value = case.dialog
    _check(case, 0, "作用力怎样影响性质？")
    full_text = value.reader.browser.toPlainText()
    value.section_search.setText("无匹配的合成栏目")
    assert value.sections.count() == 0 and value.confirm.isEnabled()
    assert value.reader.browser.toPlainText() == full_text
    value.section_search.clear()
    assert value.sections.item(0).checkState() == Qt.CheckState.Checked
    assert value.title.text() == "作用力怎样影响性质？"
    _confirm(case)
    assert len(value.requests) == 1


def test_source_switch_clears_choices_and_ignores_old_outline_success_and_failure(make_picker):
    case = make_picker()
    _ready(case)
    _check(case, 0, "保留在第一份讲义的问题")
    case.dialog.read_source()
    old = case.bridge.take(OUTLINE)
    _select_source(case, 1)
    current = case.bridge.take(OUTLINE)
    assert not case.dialog.confirm.isEnabled() and case.dialog.requests is None
    case.bridge.succeed(current, run=True)
    before = case.dialog.reader.browser.toPlainText()
    case.bridge.succeed(old, case.library.outlines[case.library.sources[0]["source_id"]])
    case.bridge.fail(old, "过时读取失败")
    assert old.task_id in case.bridge.cancelled
    assert case.dialog.reader.browser.toPlainText() == before
    assert case.library.sources[1]["source_name"] in before
    assert all(case.dialog.sections.item(i).checkState() == Qt.CheckState.Unchecked
               for i in range(case.dialog.sections.count()))
    assert case.dialog.title.text() == "" and not case.dialog.confirm.isEnabled()
    assert "过时读取失败" not in case.dialog.status.text()


@pytest.mark.parametrize("phase", ["catalog", "outline", "image", "confirm"])
def test_closed_dialog_never_accepts_late_callback(make_picker, phase):
    case = make_picker()
    if phase == "catalog":
        pending = case.bridge.take(CATALOG)
    else:
        _catalog(case)
        _select_source(case)
        if phase == "outline":
            pending = case.bridge.take(OUTLINE)
        else:
            case.bridge.succeed(case.bridge.take(OUTLINE), run=True)
            if phase == "image":
                case.dialog.show_image("word-b65-image1")
                pending = case.bridge.take(IMAGE)
            else:
                _check(case, 0, "本课问题")
                case.dialog.choose()
                pending = case.bridge.take(CONFIRM)
    result = pending.operation()
    case.dialog.reject()
    text_after_close = case.dialog.reader.browser.toPlainText()
    case.bridge.succeed(pending, result)
    case.bridge.fail(pending, "关闭后的过时失败")
    assert pending.task_id in case.bridge.cancelled
    assert case.dialog.requests is None and case.dialog.result() == QDialog.DialogCode.Rejected
    assert not case.dialog.confirm.isEnabled() and not case.dialog.isVisible()
    assert case.dialog.reader.browser.toPlainText() == text_after_close
    assert "关闭后的过时失败" not in case.dialog.status.text()
    assert not case.bridge.pending


@pytest.mark.parametrize("phase", ["image", "confirm"])
def test_source_change_invalidates_pending_image_and_confirmation(make_picker, phase):
    case = make_picker()
    _ready(case)
    if phase == "image":
        case.dialog.show_image("word-b65-image1")
        pending = case.bridge.take(IMAGE)
    else:
        _check(case, 0, "第一份讲义的问题")
        case.dialog.choose()
        pending = case.bridge.take(CONFIRM)
    result = pending.operation()
    _select_source(case, 1)
    case.bridge.succeed(case.bridge.take(OUTLINE), run=True)
    before = case.dialog.reader.browser.toPlainText()
    case.bridge.succeed(pending, result)
    assert pending.task_id in case.bridge.cancelled
    assert case.dialog.requests is None and case.dialog.result() != QDialog.DialogCode.Accepted
    assert case.dialog.reader.browser.toPlainText() == before
    assert case.dialog.reader._image_pixmap is None and not case.dialog.confirm.isEnabled()


@pytest.mark.parametrize("change", ["outline_revision", "source_sha256", "source_revision", "source_role", "section_range"])
def test_fresh_outline_or_source_change_refuses_confirmation(make_picker, change):
    case = make_picker()
    _ready(case)
    _check(case, 0, "本课引导问题")
    case.dialog.choose()
    pending = case.bridge.take(CONFIRM)
    fresh = pending.operation()
    if change == "outline_revision":
        fresh[change] = _digest("different-outline")
    elif change == "section_range":
        fresh["sections"][0]["end"] -= 1
    elif change == "source_role":
        fresh["source"][change] = "answer"
    else:
        fresh["source"][change] = _digest("different-source")
        fresh["preview"]["revision" if change == "source_revision" else change] = fresh["source"][change]
    case.bridge.succeed(pending, fresh)
    assert case.dialog.requests is None and case.dialog.result() != QDialog.DialogCode.Accepted
    assert not case.dialog.confirm.isEnabled()


def test_editing_question_during_confirmation_requires_new_confirmation(make_picker):
    case = make_picker()
    _ready(case)
    _check(case, 0, "原引导问题")
    case.dialog.choose()
    pending = case.bridge.take(CONFIRM)
    case.dialog.title.setText("核对期间修改的问题")
    case.bridge.succeed(pending, run=True)
    assert case.dialog.requests is None and case.dialog.result() != QDialog.DialogCode.Accepted
    assert case.dialog.confirm.isEnabled()
    _confirm(case)
    assert case.dialog.requests[0]["title"] == "核对期间修改的问题"


def test_twenty_section_limit_never_emits_twenty_one_requests(make_picker):
    case = make_picker(library=_SyntheticLibrary(section_count=21))
    _ready(case)
    for index in range(21):
        _check(case, index, f"第{index + 1}段的本课问题")
    selected = [case.dialog.sections.item(i) for i in range(case.dialog.sections.count())
                if case.dialog.sections.item(i).checkState() == Qt.CheckState.Checked]
    assert len(selected) in {20, 21}
    if len(selected) == 21:
        assert not case.dialog.confirm.isEnabled()
        case.dialog.choose()
        assert not case.bridge.pending and case.dialog.requests is None
        selected[-1].setCheckState(Qt.CheckState.Unchecked)
    assert case.dialog.confirm.isEnabled()
    _confirm(case)
    assert len(case.dialog.requests) == 20


@pytest.mark.parametrize("invalid", ["empty_body", "missing_start", "missing_end", "reverse_range", "heading_only"])
def test_empty_or_invalid_sections_are_not_selectable(make_picker, invalid):
    library = _SyntheticLibrary()
    snapshot = library.outlines[library.sources[0]["source_id"]]
    snapshot["sections"] = [snapshot["sections"][0]]
    section = snapshot["sections"][0]
    if invalid == "empty_body":
        snapshot["preview"]["blocks"] = []
        snapshot["preview"]["assets"] = []
    elif invalid == "missing_start":
        section["start"] = 0
    elif invalid == "missing_end":
        section["end"] = 10000
    elif invalid == "reverse_range":
        section["start"], section["end"] = 77, 64
    else:
        snapshot["preview"]["assets"] = []
        for block in snapshot["preview"]["blocks"]:
            if section["start"] < block["index"] <= section["end"]:
                block["text"] = " "
    case = make_picker(library=library)
    _ready(case)
    for index in range(case.dialog.sections.count()):
        flags = case.dialog.sections.item(index).flags()
        assert not (flags & Qt.ItemFlag.ItemIsEnabled and flags & Qt.ItemFlag.ItemIsUserCheckable)
    case.dialog.choose()
    assert case.dialog.requests is None and not case.dialog.confirm.isEnabled()
    assert not case.bridge.pending


def test_answer_source_stays_teacher_only_even_if_checkbox_is_programmatically_cleared(make_picker):
    case = make_picker(library=_SyntheticLibrary(role="answer"))
    _ready(case)
    _check(case, 0, "教师怎样使用这段答案？")
    assert case.dialog.teacher_only.isChecked() and not case.dialog.teacher_only.isEnabled()
    case.dialog.teacher_only.setChecked(False)
    _confirm(case)
    assert case.dialog.requests[0]["teacher_only"] is True
    assert case.dialog.requests[0]["source"]["source_role"] == "answer"


def test_image_heading_is_readable_but_request_keeps_original_placeholder(make_picker):
    library = _SyntheticLibrary()
    snapshot = library.outlines[library.sources[0]["source_id"]]
    snapshot["sections"][0]["title"] = PLACEHOLDER
    for block in snapshot["preview"]["blocks"]:
        if 64 <= block["index"] <= 77:
            block["text"] = PLACEHOLDER
    heading_image = deepcopy(snapshot['preview']['assets'][0])
    heading_image.update(asset_id='word-b64-image1', block_index=64)
    snapshot['preview']['assets'].append(heading_image)
    original = deepcopy(snapshot)
    case = make_picker(library=library)
    _ready(case)
    assert "原图栏目" in case.dialog.sections.item(0).text()
    assert PLACEHOLDER not in case.dialog.sections.item(0).text()
    assert '原区块 64—77 · 可查看2张原图' in case.dialog.sections.item(0).text()
    assert case.dialog.reader._source == original['preview']
    case.dialog.show_image("word-b65-image1")
    case.bridge.succeed(case.bridge.take(IMAGE), run=True)
    assert case.dialog.reader._image_pixmap is not None
    row = library.image_calls[0]
    assert row["schema"] == "shchem.handout-picture.v1" and row["block_index"] == 65
    assert row["source_revision"] == snapshot["source"]["source_revision"]
    assert row["source_sha256"] == snapshot["source"]["source_sha256"]
    _check(case, 0, "这张知识结构图怎样串联概念？")
    _confirm(case)
    assert case.dialog.requests[0]["section"] == snapshot["sections"][0]
    assert case.dialog.requests[0]["section"]["title"] == PLACEHOLDER
    assert snapshot == original


def test_duplicate_readable_headings_keep_independent_edits_ranges_and_full_navigation(make_picker, monkeypatch):
    library = _SyntheticLibrary()
    snapshot = library.outlines[library.sources[0]['source_id']]
    snapshot['sections'][0]['title'] = snapshot['sections'][1]['title'] = '完整知识网'
    snapshot['sections'][1]['level'] = 9
    snapshot['preview']['sections'] = []
    original = deepcopy(snapshot)
    case = make_picker(library=library)
    _ready(case)
    assert case.dialog.reader._source == original['preview']
    assert case.dialog.reader.section_picker.count() == 4
    assert case.dialog.sections.item(0).text() != case.dialog.sections.item(1).text()
    visited = []
    monkeypatch.setattr(case.dialog.reader, 'scroll_to_block', lambda index: visited.append(index))
    case.dialog.show_section(case.dialog.sections.item(0))
    _check(case, 0, '如何解释第一幅图？')
    _check(case, 1, '第二处知识怎样联系？')
    assert 64 in visited and 78 in visited
    _confirm(case)
    assert [r['section'] for r in case.dialog.requests] == original['sections'][:2]
    assert [r['title'] for r in case.dialog.requests] == ['如何解释第一幅图？', '第二处知识怎样联系？']
    assert all(r['source'] == original['source'] for r in case.dialog.requests)
    assert snapshot == original


@pytest.mark.parametrize("source_shape", ["metadata", "bound"])
def test_initial_source_limits_selection_to_matching_batch_source_and_sha(make_picker, source_shape):
    library = _SyntheticLibrary(source_count=3)
    initial = deepcopy(library.sources[1] if source_shape == "metadata"
                       else library.outlines[library.sources[1]["source_id"]]["source"])
    case = make_picker(library=library, initial_source=initial)
    _catalog(case)
    assert case.dialog.source.count() == 2
    assert case.dialog.source.currentData() == library.sources[1]
    case.bridge.succeed(case.bridge.take(OUTLINE), run=True)
    assert case.library.outline_calls == [library.sources[1]]
    assert case.dialog.sections.count() == 3


@pytest.mark.parametrize("mismatch", ["batch_id", "source_id", "source_sha256", "source_revision"])
def test_initial_source_mismatch_never_falls_back_to_another_version(make_picker, mismatch):
    library = _SyntheticLibrary()
    initial = deepcopy(library.outlines[library.sources[0]["source_id"]]["source"])
    initial[mismatch] = _digest("different-initial") if mismatch.endswith(("sha256", "revision")) else "different-initial"
    case = make_picker(library=library, initial_source=initial)
    _catalog(case)
    if mismatch == "source_revision":
        case.bridge.succeed(case.bridge.take(OUTLINE), run=True)
    else:
        assert not case.bridge.pending and not library.outline_calls
    assert case.dialog.sections.count() == 0 and not case.dialog.confirm.isEnabled()
    assert case.dialog.requests is None


def test_legacy_image_entry_uses_catalog_role_and_rechecks_exact_picture_before_each_outline(make_picker):
    library = _SyntheticLibrary(role="answer")
    initial = _legacy_image_source(library)
    untouched = deepcopy(initial)
    case = make_picker(library=library, initial_source=initial)
    _catalog(case)
    assert case.dialog.source.count() == 2 and case.dialog.source.currentData() == library.sources[0]
    assert not library.context_calls and not library.outline_calls
    case.bridge.succeed(case.bridge.take(OUTLINE), run=True)
    assert library.context_calls == [untouched]
    assert library.outline_calls == [library.sources[0]]
    _check(case, 0, "教师怎样借助这段原答案讲解？")
    assert case.dialog.teacher_only.isChecked() and not case.dialog.teacher_only.isEnabled()
    case.dialog.teacher_only.setChecked(False)
    _confirm(case)
    assert library.context_calls == [untouched, untouched]
    assert initial == untouched and "source_role" not in initial
    request = case.dialog.requests[0]
    assert request["source"]["schema"] == "shchem.handout-knowledge-source.v1"
    assert request["source"]["source_role"] == "answer" and request["teacher_only"] is True
    assert "asset_id" not in request["source"] and len(request["source"]) == 8


@pytest.mark.parametrize("invalid", [
    "missing_asset", "extra_role", "bad_schema", "bad_block_index", "bad_image_hash", "not_mapping"
])
def test_invalid_legacy_initial_record_cannot_start_outline(make_picker, invalid):
    library = _SyntheticLibrary()
    initial = _legacy_image_source(library)
    if invalid == "missing_asset":
        initial.pop("asset_id")
    elif invalid == "extra_role":
        initial["source_role"] = "handout"
    elif invalid == "bad_schema":
        initial["schema"] = "shchem.handout-picture.v2"
    elif invalid == "bad_block_index":
        initial["block_index"] = True
    elif invalid == "bad_image_hash":
        initial["sha256"] = "not-an-image-hash"
    else:
        initial = "untrusted-legacy-record"
    case = make_picker(library=library, initial_source=initial)
    _catalog(case)
    assert not library.context_calls and not library.outline_calls and not case.bridge.pending
    assert case.dialog.source.currentData() is None and case.dialog.sections.count() == 0
    case.dialog.choose()
    assert case.dialog.requests is None and not case.dialog.confirm.isEnabled()


@pytest.mark.parametrize("change", ["same_bytes_other_archive", "different_name"])
def test_legacy_initial_record_does_not_match_by_file_hash_alone(make_picker, change):
    library = _SyntheticLibrary(source_count=1)
    initial = _legacy_image_source(library)
    if change == "same_bytes_other_archive":
        library.sources[0]["batch_id"] = "different-archive"
        library.sources[0]["source_id"] = "different-file-id"
    else:
        library.sources[0]["source_name"] = "同字节但另一个文件名.docx"
    assert library.sources[0]["source_sha256"] == initial["source_sha256"]
    case = make_picker(library=library, initial_source=initial)
    _catalog(case)
    assert not library.context_calls and not library.outline_calls and not case.bridge.pending
    assert case.dialog.source.currentData() is None and case.dialog.sections.count() == 0
    assert case.dialog.requests is None and not case.dialog.confirm.isEnabled()


@pytest.mark.parametrize("phase", ["outline", "confirmation"])
@pytest.mark.parametrize("change", ["source_sha256", "source_revision", "asset_id", "block_index"])
def test_stale_legacy_picture_is_rejected_before_outline_or_confirmation(make_picker, phase, change):
    library = _SyntheticLibrary()
    initial = _legacy_image_source(library)
    case = make_picker(library=library, initial_source=initial)
    _catalog(case)
    if phase == "outline":
        pending = case.bridge.take(OUTLINE)
        outline_count = 0
    else:
        case.bridge.succeed(case.bridge.take(OUTLINE), run=True)
        _check(case, 0, "核对旧图来源后的本课问题")
        case.dialog.choose()
        pending = case.bridge.take(CONFIRM)
        outline_count = 1
    preview = library.outlines[library.sources[0]["source_id"]]["preview"]
    if change == "source_sha256":
        preview["source_sha256"] = _digest("source-changed-after-catalog")
    elif change == "source_revision":
        preview["revision"] = _digest("preview-changed-after-catalog")
    elif change == "asset_id":
        preview["assets"][0]["asset_id"] = "different-picture-occurrence"
    else:
        preview["assets"][0]["block_index"] += 1
    with pytest.raises(ValueError, match="变化") as error:
        pending.operation()
    case.bridge.fail(pending, str(error.value))
    assert len(library.context_calls) == outline_count + 1
    assert all(source == initial for source in library.context_calls)
    assert len(library.outline_calls) == outline_count
    assert case.dialog.requests is None and case.dialog.result() != QDialog.DialogCode.Accepted
    assert not case.dialog.confirm.isEnabled()
    case.dialog.title.setText("改标题也不能绕过旧图来源失效")
    assert not case.dialog.confirm.isEnabled()
    if phase == "outline":
        assert case.dialog.sections.count() == 0


@pytest.mark.parametrize("width", [420, 1080])
def test_responsive_layout_keeps_confirmation_reachable_without_horizontal_scroll(make_picker, width):
    case = make_picker()
    _ready(case)
    _check(case, 0, "可达的本课问题")
    value = case.dialog
    value.resize(width, 740)
    settle(case.app)
    assert value.width() == width
    assert value.split.orientation() == (Qt.Orientation.Vertical if width < 750 else Qt.Orientation.Horizontal)
    assert value.content_scroll.horizontalScrollBar().maximum() == 0
    assert value.reader.browser.horizontalScrollBar().maximum() == 0
    assert value.sections.horizontalScrollBar().maximum() == 0
    for widget in (value.source_search, value.source, value.section_search, value.title, value.details_button):
        value.content_scroll.ensureWidgetVisible(widget)
        settle(case.app)
        pos = widget.mapTo(value.content_scroll.viewport(), widget.rect().topLeft())
        assert 0 <= pos.x() and pos.x() + widget.width() <= value.content_scroll.viewport().width()
        assert 0 <= pos.y() and pos.y() + widget.height() <= value.content_scroll.viewport().height()
    value.details_button.click()
    value.content_scroll.ensureWidgetVisible(value.details)
    settle(case.app)
    assert value.details.isVisible() and "source_sha256" in value.details.toPlainText()
    point = value.confirm.mapTo(value, value.confirm.rect().topLeft())
    assert 0 <= point.x() and point.x() + value.confirm.width() <= value.width()
    assert 0 <= point.y() and point.y() + value.confirm.height() <= value.height()
    assert value.confirm.isEnabled()


@pytest.mark.parametrize("entry", ["manual_source", "legacy_image"])
def test_actual_docx_outline_lazy_original_picture_and_fresh_confirmation(
    qt_app, monkeypatch, tmp_path, entry
):
    """Exercise real knowledge/image APIs against a temporary, generated Word."""
    from io import BytesIO

    from docx import Document

    from test_handout_knowledge import make_facade, png, word_bytes
    from integrations.deeptutor_shchem_v1.desktop_handout_images import handout_image_options
    from integrations.deeptutor_shchem_v1.desktop_preparation_sources import PreparationSourcesService
    from integrations.deeptutor_shchem_v1.desktop_workbench import handout_knowledge_picker as module

    doc = Document()
    for index in range(1, 64):
        doc.add_paragraph(f"合成前置原文区块{index}。")
    doc.add_heading("知识点1 范德华力", 3)
    doc.add_paragraph("分子间存在相互作用；本段用于检验完整原文读取。")
    doc.add_paragraph("知识结构原图：").add_run().add_picture(BytesIO(png()))
    doc.add_heading("知识点2 氢键", 3)
    doc.add_paragraph("第二段的完整合成正文。")
    original = word_bytes(doc)
    facade = make_facade(original, tmp_path / "original-word")
    service = PreparationSourcesService(facade.paths.workspace_root)
    asset_reads = []
    preview_reads = []

    def original_preview(batch_id, source_id, *, read_only=False):
        assert read_only is True
        preview_reads.append((batch_id, source_id))
        source = facade._imported_word_source(batch_id, source_id)
        return service.word_preview_bytes(source.content, source.filename)

    def original_asset(batch_id, source_id, asset_id):
        asset_reads.append((batch_id, source_id, asset_id))
        source = facade._imported_word_source(batch_id, source_id)
        return service.word_asset_bytes(source.content, asset_id, render_metafiles=True)

    facade.imported_word_preview = original_preview
    facade.imported_word_asset = original_asset
    bridge = _ManualBridge()
    bridge.shutdown = lambda _timeout: None
    monkeypatch.setattr(module, "DesktopTaskBridge", lambda _parent: bridge)
    initial_source = None
    if entry == "legacy_image":
        initial_source = handout_image_options(facade)["images"][0]
        assert len(initial_source) == 12 and "source_role" not in initial_source
        facade.value.reads = 0
        preview_reads.clear()
    value = module.HandoutKnowledgePicker(facade, initial_source=initial_source)
    try:
        value.show()
        settle(qt_app)
        bridge.succeed(bridge.take(CATALOG), run=True)
        assert facade.value.reads == 0 and not preview_reads and not asset_reads
        if entry == "manual_source":
            assert not bridge.pending
            value.source.setCurrentIndex(1)
        else:
            assert value.source.count() == 2 and value.source.currentData()["source_role"] == "handout"
        pending = bridge.take(OUTLINE)
        outline = pending.operation()
        bridge.succeed(pending, outline)
        first = outline["sections"][0]
        assert (first["start"], first["end"], first["level"]) == (64, 66, 3)
        assert value.sections.item(0).data(Qt.ItemDataRole.UserRole) == first
        assert "第二段的完整合成正文。" in value.reader.browser.toPlainText()
        assert not asset_reads
        assert preview_reads == ([] if entry == "manual_source" else [("BATCH", "SOURCE")])

        asset_id = outline["preview"]["assets"][0]["asset_id"]
        value.show_image(asset_id)
        assert not asset_reads
        bridge.succeed(bridge.take(IMAGE), run=True)
        assert asset_reads == [("BATCH", "SOURCE", asset_id)]
        assert preview_reads == [("BATCH", "SOURCE")] * (2 if entry == "manual_source" else 3)
        assert value.reader._image_pixmap is not None

        value.sections.item(0).setCheckState(Qt.CheckState.Checked)
        value.title.setText("这段知识与原图怎样互相解释？")
        assert value.confirm.isEnabled()
        reads_before_confirmation = facade.value.reads
        value.choose()
        bridge.succeed(bridge.take(CONFIRM), run=True)
        assert facade.value.reads > reads_before_confirmation
        assert preview_reads == [("BATCH", "SOURCE")] * (2 if entry == "manual_source" else 4)
        assert value.result() == QDialog.DialogCode.Accepted
        assert value.requests == [{
            "source": outline["source"],
            "outline_revision": outline["outline_revision"],
            "section": first,
            "title": "这段知识与原图怎样互相解释？",
            "teacher_only": False,
        }]
        assert not facade.paths.task_root.exists()
        assert facade.original.read_bytes() == original
    finally:
        if value.result() != QDialog.DialogCode.Accepted:
            value.reject()
        value.deleteLater()
        settle(qt_app)
