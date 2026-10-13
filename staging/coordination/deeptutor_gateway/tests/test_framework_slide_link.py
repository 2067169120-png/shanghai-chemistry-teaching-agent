"""Framework editing reaches the same slide work, without replacing source work."""
from copy import deepcopy
import json
import pytest
from PySide6.QtCore import QTimer
from PySide6.QtWidgets import QApplication, QDialog, QDialogButtonBox
from docx import Document
from pptx import Presentation
from test_lesson_design import sample
from test_lesson_design_ui import desk
from test_question_explorer_ui import window
from integrations.deeptutor_shchem_v1.desktop_teaching_framework import FIELDS, SCHEMA
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    DESIGN_SCHEMA, FORMATS, block, digest, new_composition, resource, validate_composition,
    validate_resource, visible_blocks)
from integrations.deeptutor_shchem_v1.desktop_framework_composition import *
from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents
from integrations.deeptutor_shchem_v1.desktop_composition_render import layout_composition, write_pptx
from integrations.deeptutor_shchem_v1.desktop_workbench.framework_slide_dialog import FrameworkSlideDialog

_APP = None


@pytest.fixture(autouse=True)
def qt_app():
    global _APP
    _APP = QApplication.instance() or QApplication([])
    return _APP


def source_plan():
    payload = sample()
    plan = payload['lesson_design']
    plan['teaching_framework'] = {'schema_version': SCHEMA,
        **{key: '本课' + label + '完整内容' for key, label in FIELDS}}
    comp = new_composition('16:9')
    item = resource('personal_text', '原有素材', [block('text', 'personal', '保留原有内容、顺序及位置。')],
                    {'source_path': '个人文字', 'source_sha256': 'unknown', 'source_revision': 'unknown'})
    w, h = FORMATS[comp['format']]
    comp['resources'] = [item]
    comp['placements'] = {item['id']: {'page_id': comp['pages'][0], 'x': 12, 'y': 12,
                                     'width': w-24, 'height': h-24}}
    plan['composition'] = validate_composition(comp)
    plan['schema_version'] = DESIGN_SCHEMA
    return payload


def linked_plan():
    payload = source_plan()
    payload['lesson_design'] = synchronize_framework(payload['lesson_design'], list(BY_SECTION))['plan']
    return payload


def card(plan, section):
    return next(r for r in plan['composition']['resources'] if r.get('framework_link', {}).get('section') == section)


def test_all_sections_link_without_changing_original_resource_frames_nodes_or_goals():
    payload = source_plan(); original = deepcopy(payload['lesson_design'])
    result = synchronize_framework(original, list(BY_SECTION)); plan = result['plan']
    original_item = original['composition']['resources'][0]
    assert next(r for r in plan['composition']['resources'] if r['id'] == original_item['id']) == original_item
    assert plan['composition']['placements'][original_item['id']] == original['composition']['placements'][original_item['id']]
    assert plan['nodes'] == original['nodes'] and plan['objectives'] == original['objectives']
    assert [r['title'] for r in plan['composition']['resources']] == [
        '本课目标','导入新课','原有素材','巩固小结','板书设计','课后任务']
    assert all(row['action'] == 'added' for row in result['changes'])
    assert original == payload['lesson_design']
    assert all(row['status'] == 'current' for row in checked_framework_links(plan))


def test_repeat_sync_is_exactly_idempotent_without_new_pages_or_cards():
    plan = linked_plan()['lesson_design']; before = deepcopy(plan)
    result = synchronize_framework(plan, list(BY_SECTION))
    assert result['plan'] == before and all(r['action'] == 'unchanged' for r in result['changes'])


def test_board_edit_updates_only_board_and_keeps_resource_id_and_fitting_frame():
    plan = linked_plan()['lesson_design']; before = deepcopy(plan)
    old = card(plan, 'board_design'); plan['teaching_framework']['board_design'] = '结构 → 作用 → 性质'
    updated = refresh_linked_framework(plan)
    current = card(updated, 'board_design')
    assert current['id'] == old['id'] and current['original_sha256'] != old['original_sha256']
    assert updated['composition']['placements'] == before['composition']['placements']
    assert current['blocks'][0]['value'] == '结构 → 作用 → 性质'
    assert [r for r in updated['composition']['resources'] if r['id'] != old['id']] == [
        r for r in before['composition']['resources'] if r['id'] != old['id']]


@pytest.mark.parametrize('field', ['blocks', 'title'])
def test_edited_slide_copy_is_retained_and_explicit_replacement_relinks(field):
    plan = linked_plan()['lesson_design']; old = card(plan, 'board_design')
    if field == 'blocks': old['blocks'][0]['value'] = '课件中的个人板书文字'
    else: old['title'] = '课件中的个人标题'
    work = deepcopy(old)
    plan['teaching_framework']['board_design'] = '教案里的新板书'
    kept = refresh_linked_framework(plan); current = card(kept, 'board_design')
    assert current['blocks'] == work['blocks'] and current['title'] == work['title']
    assert current['original'] == work['original'] and current['framework_link']['mode'] == 'snapshot'
    assert checked_framework_links(kept)[-2]['status'] == 'snapshot'
    updated = synchronize_framework(kept, ['board_design'], replace_edited=True)['plan']
    current = card(updated, 'board_design')
    assert current['id'] == work['id'] and current['blocks'][0]['value'] == '教案里的新板书'
    assert current['framework_link']['mode'] == 'linked'


def test_empty_column_removes_only_its_unedited_card_and_owned_empty_page():
    plan = linked_plan()['lesson_design']; old = card(plan, 'board_design')
    original_item = next(r for r in plan['composition']['resources'] if r['title'] == '原有素材')
    frame = deepcopy(plan['composition']['placements'][original_item['id']])
    plan['teaching_framework']['board_design'] = ''
    result = synchronize_framework(plan, ['board_design'])
    assert result['changes'][0]['action'] == 'removed_empty'
    assert not any(r['id'] == old['id'] for r in result['plan']['composition']['resources'])
    assert result['plan']['composition']['placements'][original_item['id']] == frame


def test_long_board_repaginates_its_card_without_moving_original_material(tmp_path):
    plan = linked_plan()['lesson_design']; original = next(r for r in plan['composition']['resources'] if r['title'] == '原有素材')
    frame = deepcopy(plan['composition']['placements'][original['id']])
    lines = [f'第{i}步：保留这行完整的板书内容。' for i in range(1, 81)]
    plan['teaching_framework']['board_design'] = '\n'.join(lines)
    updated = refresh_linked_framework(plan); current = card(updated, 'board_design')
    assert updated['composition']['placements'][original['id']] == frame
    assert len(resource_frames(updated['composition']['placements'][current['id']])) > 1
    layout = layout_composition(updated['composition'], [], {}, student=True)
    assert layout['errors'] == []
    write_pptx(layout, tmp_path/'board.pptx', {})
    text = '\n'.join(shape.text for slide in Presentation(tmp_path/'board.pptx').slides for shape in slide.shapes if shape.has_text_frame)
    assert all(line in text for line in lines)


@pytest.mark.parametrize('change', ['source_revision', 'snapshot_value', 'student_role', 'official_flag', 'link_design'])
def test_bad_provenance_or_scope_cannot_be_accepted(change):
    item = framework_resource(source_plan()['lesson_design'], 'goals')
    if change == 'source_revision': item['original']['evidence']['source_revision'] = 'a'*64
    if change == 'snapshot_value': item['original']['evidence']['framework_snapshot']['fields']['values'] = '篡改'
    if change == 'student_role': item['blocks'][1]['role'] = 'material'
    if change == 'official_flag': item['official_claim_allowed'] = True
    if change == 'link_design': item['framework_link']['design_id'] = 'LD-'+'a'*32
    with pytest.raises(ValueError): validate_resource(item)


def test_student_scope_omits_teacher_focus_difficulty_methods_and_review_reminders():
    plan = source_plan()['lesson_design']
    plan['teaching_framework']['knowledge_skills'] = '本地整理候选，待教师核对。\n条件相同，数值为12.5 mol·L⁻¹。'
    item = framework_resource(plan, 'goals'); before = deepcopy(item)
    student = '\n'.join(b['value'] for b in visible_blocks(item, student=True))
    teacher = '\n'.join(b['value'] for b in visible_blocks(item, student=False))
    assert '待教师' not in student and '12.5 mol·L⁻¹' in student
    assert '本课教学重点完整内容' not in student and '本课教学重点完整内容' in teacher
    assert item == before and not item['human_reviewed'] and not item['publication_allowed']


def test_stale_link_blocks_export_until_updated_and_preview_change_guard_is_atomic():
    plan = linked_plan()['lesson_design']; expected = digest(plan)
    plan['teaching_framework']['board_design'] = '新板书'; before = deepcopy(plan)
    with pytest.raises(ValueError, match='旧教案'): checked_framework_links(plan)
    with pytest.raises(ValueError, match='重新预览'): synchronize_framework(plan, ['board_design'], expected=expected)
    assert plan == before


def test_empty_new_framework_does_not_create_blank_cover_or_empty_composition():
    plan = sample()['lesson_design']
    plan['teaching_framework'] = {'schema_version': SCHEMA, **dict.fromkeys(LABELS, '')}
    assert 'composition' not in synchronize_framework(plan, list(BY_SECTION))['plan']
    plan['teaching_framework']['board_design'] = '唯一板书'
    result = synchronize_framework(plan, ['board_design'])['plan']
    assert len(result['composition']['pages']) == 1


def test_generated_word_and_both_ppts_share_framework_without_duplicating_word_headings(tmp_path):
    payload = linked_plan(); write_documents(payload, tmp_path, {})
    teacher = '\n'.join(p.text for p in Document(tmp_path/'lesson_plan.docx').paragraphs)
    student = '\n'.join(p.text for p in Document(tmp_path/'student_worksheet.docx').paragraphs)
    assert teacher.count('本课板书设计完整内容') == 1 and '本课板书设计完整内容' in student
    assert '本课教学重点完整内容' in teacher and '本课教学重点完整内容' not in student
    for audience in (False, True):
        layout = layout_composition(payload['lesson_design']['composition'], [], {}, student=audience)
        assert not layout['errors']
        name = 'student.pptx' if audience else 'teacher.pptx'; write_pptx(layout, tmp_path/name, {})
        text = '\n'.join(s.text for page in Presentation(tmp_path/name).slides for s in page.shapes if s.has_text_frame)
        assert '本课板书设计完整内容' in text and '本课布置作业完整内容' in text
        assert ('本课教学重点完整内容' in text) is (not audience)


@pytest.mark.parametrize('accepted', [False, True])
def test_actual_preview_dialog_cancel_or_apply_and_undo(desk, accepted):
    dialog, app, page = desk; before = deepcopy(dialog.history.value)
    dialog.framework_fields['board_design'].setPlainText('结构 → 作用 → 性质')
    observed = []
    def choose():
        preview = app.activeModalWidget()
        assert isinstance(preview, FrameworkSlideDialog)
        assert 'board_design' in preview.chosen_sections()
        preview.sections.setCurrentRow(3)
        assert '结构 → 作用 → 性质' in preview.preview.toPlainText()
        observed.append(True)
        (preview.accept if accepted else preview.reject)()
    QTimer.singleShot(0, choose); dialog.framework_slides_button.click()
    assert observed
    if not accepted:
        assert dialog.history.value == before and dialog.framework_fields['board_design'].toPlainText() == '结构 → 作用 → 性质'
    else:
        assert card(page._payload()['lesson_design'], 'board_design')['blocks'][0]['value'] == '结构 → 作用 → 性质'
        after = deepcopy(dialog.history.value); dialog.travel(False)
        assert dialog.history.value == before
        dialog.travel(True); assert dialog.history.value == after


@pytest.mark.parametrize('action', ['save_draft', 'generate_outputs'])
def test_save_and_generate_apply_pending_field_edits_and_update_linked_card(desk, monkeypatch, action):
    dialog, app, page = desk
    dialog.framework_fields['board_design'].setPlainText('原板书'); dialog.apply_framework()
    dialog.history.put(synchronize_framework(dialog.history.value, ['board_design'])['plan']); dialog.sync_page()
    dialog.framework_fields['board_design'].setPlainText('未另点应用的新板书')
    captured = []
    monkeypatch.setattr(dialog, 'run', lambda *args, **kwargs: captured.append(deepcopy(page._payload())))
    getattr(dialog, action)()
    assert captured[0]['lesson_design']['teaching_framework']['board_design'] == '未另点应用的新板书'
    assert card(captured[0]['lesson_design'], 'board_design')['blocks'][0]['value'] == '未另点应用的新板书'


def test_return_keeps_pending_framework_edit_in_real_recovery_without_extra_apply(desk):
    dialog, app, page = desk
    dialog.framework_fields['board_design'].setPlainText('返回后仍能恢复的板书')
    dialog.reject(); app.processEvents()
    restored = page.recovery.store.load()['payload']['lesson_design']
    assert restored['teaching_framework']['board_design'] == '返回后仍能恢复的板书'
