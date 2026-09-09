"""Readable preparation tables retain source content and grid relationships."""

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from integrations.deeptutor_shchem_v1.desktop_preparation_sources import (
    PreparationSourcesService,
    _block_text,
    _readable_table,
)


def test_merged_cells_are_not_repeated_or_filled():
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(1, 1)).text = "共同条件"
    table.cell(0, 2).text = "水溶液"
    table.cell(1, 2).text = "熔融"
    table.cell(2, 0).text = "下一项"
    text = _block_text(table._tbl, document, set())
    assert text.count("共同条件") == 1
    assert "〔第1行·第1—2列；横向合并；纵向合并起点〕\n共同条件" in text
    assert "〔第2行·第1—2列；横向合并；纵向合并续接上方〕\n【未提取到文字】" in text
    assert "〔第2行·第3列〕\n熔融" in text
    assert "〔第3行·第1列〕\n下一项" in text
    assert '"column_span"' not in text


def test_omitted_edge_cells_keep_column_coordinates():
    document = Document()
    table = document.add_table(rows=2, cols=4)
    row = table._tbl.tr_lst[1]
    for cell in [row.tc_lst[0], *row.tc_lst[2:]]:
        row.remove(cell)
    properties = row.get_or_add_trPr()
    for name, count in (("gridBefore", 1), ("gridAfter", 2)):
        node = OxmlElement("w:" + name)
        node.set(qn("w:val"), str(count))
        properties.append(node)
    text = _block_text(table._tbl, document, set())
    assert "〔第2行：行首省略1列；行尾省略2列〕" in text
    assert "〔第2行·第2列〕\n【未提取到文字】" in text
    assert "〔第2行·第1列〕" not in text
    assert "〔第2行·第3列〕" not in text


def test_multiline_equations_conditions_and_gaps_reach_reference_exactly(tmp_path):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "条件\n水溶液\t与熔融分开\n不可省略"
    paragraph = table.cell(0, 1).paragraphs[0]
    paragraph.add_run("SO")
    paragraph.add_run("4").font.subscript = True
    paragraph.add_run("2−").font.superscript = True
    paragraph.add_run("\n= ≠ ⇌\n")
    paragraph.add_run()._r.append(OxmlElement("w:object"))
    paragraph.add_run("\n末尾条件")
    path = tmp_path / "knowledge-table.docx"
    document.save(path)
    before = path.read_bytes()
    service = PreparationSourcesService(tmp_path)
    preview = service.word_preview(path)
    text = preview["blocks"][0]["text"]
    reference = service.reference(path, preview["source_sha256"], 1, 1, [])
    assert "条件\n水溶液\t与熔融分开\n不可省略" in text
    assert "SO_{4}^{2−}\n= ≠ ⇌\n【待查看原文：嵌入对象或旧公式】\n末尾条件" in text
    assert text in reference["materials"]
    assert reference["materials"].count("末尾条件") == 1
    assert reference["warnings"] == [
        "Word区块1：嵌入对象或旧公式未提取；需查看原文件，不得根据残句补写。"
    ]
    assert path.read_bytes() == before


def test_source_text_is_not_summarized_or_treated_as_formatting():
    source_text = '第一行\n\n  保留缩进\t原文 "text": 1\n【表格结束】\n最后一行'
    rows = [
        {
            "grid_before": 0,
            "grid_after": 0,
            "cells": [
                {"column_span": 1, "vertical_merge": "none", "text": source_text}
            ],
        }
    ]
    assert source_text in _readable_table(rows)
    # The format is human-readable source data, not an unambiguous parser input.
    assert _readable_table(rows).count("最后一行") == 1


def test_empty_rows_and_unknown_merge_markers_are_not_inferred():
    rows = [
        {"grid_before": 0, "grid_after": 0, "cells": []},
        {
            "grid_before": 0,
            "grid_after": 0,
            "cells": [
                {"column_span": 1, "vertical_merge": "unrecognized", "text": "原内容"}
            ],
        },
    ]
    text = _readable_table(rows)
    assert "〔第1行：无单元格〕" in text
    assert "〔第2行·第1列；原纵向合并标记：unrecognized〕\n原内容" in text
