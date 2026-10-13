"""Projection navigation repeats titles without rewriting or hiding source content."""
from copy import deepcopy
import re

import pytest
from pptx import Presentation
from test_composition_render import app, material, text
from test_composition_student_pages import question, teacher_note
from integrations.deeptutor_shchem_v1.desktop_composition_render import (
    PADDING, arrange_composition, layout_composition, write_pdf, write_pptx)
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import FORMATS
from integrations.deeptutor_shchem_v1.desktop_workbench.composition_page_navigator import (
    CompositionPageNavigator, visible_fragments)


def projection(material):
    payload, facade, state, data, body = material
    before = deepcopy(payload)
    value = deepcopy(payload['lesson_design']['composition'])
    value['format'] = '16:9'
    width, height = FORMATS[value['format']]
    value['placements'] = {r['id']: {'page_id': value['pages'][0], 'x': 12, 'y': 12,
                                    'width': width-24, 'height': height-24}
                           for r in value['resources']}
    value = arrange_composition(value, payload['image_assets'], data)
    assert payload == before
    return value, payload['image_assets'], data, body


@pytest.mark.parametrize('student', [False, True])
def test_every_projection_continuation_has_a_title_and_complete_body(material, student):
    value, assets, data, body = projection(material)
    before = deepcopy(value)
    layout = layout_composition(value, assets, data, student=student)
    assert not layout['errors'] and value == before
    item = value['resources'][0]
    fragments = [f for f in visible_fragments(layout) if f['resource_id'] == item['id']]
    assert len(fragments) > 2
    assert not any(e.get('navigation_heading') for e in fragments[0]['elements'])
    for fragment in fragments[1:]:
        headings = [e for e in fragment['elements'] if e.get('navigation_heading') and e['kind'] == 'text']
        assert ''.join(e['text'] for e in headings) == item['title']+'（续）'
        assert len(headings) == 1
        assert all(e['block_id'] is None and e['font_px'] == 32 for e in headings)
        source = [e for e in fragment['elements'] if e.get('block_id') is not None and e['kind'] != 'space']
        assert source and min(e['y'] for e in source) >= max(e['y']+e['height'] for e in headings)
        assert max(e['y']+e['height'] for e in fragment['elements']) <= (
            fragment['frame']['y']+fragment['frame']['height']-PADDING+.001)
    assert text(layout, item['blocks'][1]['id']) == body
    assert not any(e.get('navigation_heading') and e.get('block_id') is not None
                   for page in layout['pages'].values() for e in page)


def test_A4_keeps_normal_flow_without_projection_headings(material):
    payload, facade, state, data, body = material
    value = deepcopy(payload['lesson_design']['composition'])
    before = deepcopy(value)
    layout = layout_composition(value, payload['image_assets'], data, student=True)
    assert not layout['errors'] and value == before
    assert not any(e.get('navigation_heading') for page in layout['pages'].values() for e in page)
    assert text(layout, value['resources'][0]['blocks'][1]['id']) == body


def test_hidden_answers_do_not_create_headings_or_student_pages(app):
    value = teacher_note(question())
    before = deepcopy(value)
    teacher = layout_composition(value, [], {}, student=False)
    student = layout_composition(value, [], {}, student=True)
    assert not teacher['errors'] and not student['errors'] and value == before
    assert len(student['pages']) == 1 < len(teacher['pages'])
    assert not any(e.get('navigation_heading') for page in student['pages'].values() for e in page)
    assert '教师答案' not in text(student) and '仅教师可见标题' not in text(student)
    assert '教师答案第79行' in text(teacher)


def test_editable_PPT_and_PDF_keep_navigation_separate_from_full_source(material, tmp_path):
    import pypdfium2 as pdfium
    value, assets, data, body = projection(material)
    layout = layout_composition(value, assets, data, student=True)
    ppt = tmp_path/'projection.pptx'
    pdf = tmp_path/'projection.pdf'
    write_pptx(layout, ppt, data)
    write_pdf(layout, pdf)
    deck = Presentation(ppt)
    source = ''.join(shape.text for slide in deck.slides for shape in slide.shapes
                     if shape.has_text_frame and shape.name != 'shchem-continuation-heading')
    headings = [shape for slide in deck.slides for shape in slide.shapes
                if shape.has_text_frame and shape.name == 'shchem-continuation-heading']
    assert headings and body in source and '仅教师答案唯一标记' not in source
    assert all(shape.text == '合成完整题（续）' for shape in headings)
    assert all(run.font.size.pt == 24 for shape in headings
               for paragraph in shape.text_frame.paragraphs for run in paragraph.runs)
    assert all(run.font.size.pt >= 24 for slide in deck.slides for shape in slide.shapes
               if shape.has_text_frame for paragraph in shape.text_frame.paragraphs for run in paragraph.runs)
    with pdfium.PdfDocument(str(pdf)) as doc:
        words = []
        for page in doc:
            tp = page.get_textpage()
            words.append(tp.get_text_range())
            tp.close()
            page.close()
        assert len(doc) == len(deck.slides) == len(layout['pages'])
    assert '合成完整题（续）' in re.sub(r'\s', '', ''.join(words))


@pytest.mark.parametrize('student', [False, True])
def test_native_directory_labels_actual_audience_fragments_and_keeps_raw_frame_identity(material, student):
    value, assets, data, body = projection(material)
    layout = layout_composition(value, assets, data, student=student)
    before = deepcopy(value)
    navigator = CompositionPageNavigator()
    kinds = {r['original']['kind']: r['original']['kind'] for r in value['resources']}
    navigator.set_layout(layout, value['resources'], kinds, student=student)
    item = value['resources'][0]
    fragments = [f for f in visible_fragments(layout) if f['resource_id'] == item['id']]
    for position, fragment in enumerate(fragments, 1):
        key = (fragment['frame']['page_id'], item['id'], fragment['frame_index'])
        child = navigator._fragment_items[key]
        assert f'本素材第{position}/{len(fragments)}页' in child.text(0)
        if student:
            assert '题面' in child.text(0) and '讲评' not in child.text(0)
        navigator.set_current(*key)
        assert navigator.tree.currentItem() is child
    assert value == before
    navigator.deleteLater()
