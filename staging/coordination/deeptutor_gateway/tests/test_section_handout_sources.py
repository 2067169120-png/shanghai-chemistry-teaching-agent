"""Exact course-used lecture identities, using real DOCX and no state store."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from threading import Event
from types import SimpleNamespace
from uuid import uuid4

from docx import Document
import pytest

from integrations.deeptutor_shchem_v1 import desktop_handout_knowledge as handout
from integrations.deeptutor_shchem_v1 import desktop_section_handout_sources as used
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    FLAGS, block, digest, resource, validate_resource)
from integrations.deeptutor_shchem_v1.reader_cancellation import ReadCancelled, read_cancel_scope


SOURCE_FIELDS = {'schema', 'batch_id', 'source_id', 'source_name', 'source_path',
    'source_sha256', 'source_revision', 'source_role'}


def _word_bytes(note='保留原文字、条件与空格。'):
    doc = Document()
    doc.add_heading('知识点1 完整知识栏目', 3)
    doc.add_paragraph(note)
    doc.add_heading('知识点2 第二知识栏目', 3)
    doc.add_paragraph('另一个完整段落。')
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _facade(root):
    """An exact archive resolver, with no profile, database or mutable state API."""
    records, reads = {}, []

    def source(batch_id, source_id):
        reads.append((batch_id, source_id))
        current = records[(batch_id, source_id)]
        data = current.path.read_bytes()
        return SimpleNamespace(content=data, filename=current.name,
            role=current.role, effective_source_file_id=source_id,
            source_sha256=sha256(data).hexdigest())

    def forbidden_catalog(*args):
        pytest.fail('Course-used sources must not enumerate the Word catalogue')

    return SimpleNamespace(paths=SimpleNamespace(workspace_root=root / 'workspace',
        task_root=root / 'task'), records=records, reads=reads,
        _imported_word_source=source,
        imported_word_path=lambda batch_id, source_id: str(records[(batch_id, source_id)].path),
        list_imported_word_batches=forbidden_catalog, imported_word_sources=forbidden_catalog)


def _add_document(facade, root, batch='BATCH', source_id='SOURCE', data=None,
                  name='同名讲义.docx', role='handout'):
    path = root / batch / source_id / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data if data is not None else _word_bytes())
    facade.records[(batch, source_id)] = SimpleNamespace(path=path, name=name, role=role)
    return {'schema': handout.SOURCE_SCHEMA, 'batch_id': batch, 'source_id': source_id,
        'source_name': name, 'source_sha256': sha256(path.read_bytes()).hexdigest(),
        'source_role': role}


def _import_sections(facade, metadata):
    outline = handout.handout_knowledge_outline(facade, metadata)
    requests = [{'source': outline['source'], 'outline_revision': outline['outline_revision'],
        'section': section, 'title': '教师自己的问题标题', 'teacher_only': False}
        for section in outline['sections']]
    result = handout.import_handout_knowledge(facade, requests, [], [])
    assert result['image_assets'] == []
    return result['resources']


@pytest.fixture
def course(tmp_path):
    facade = _facade(tmp_path)
    metadata = _add_document(facade, tmp_path)
    resources = _import_sections(facade, metadata)
    facade.reads.clear()
    return SimpleNamespace(facade=facade, resources=resources, root=tmp_path,
        path=facade.records[('BATCH', 'SOURCE')].path)


def _files(root):
    return {str(path.relative_to(root)): (sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns)
        for path in root.rglob('*') if path.is_file()}


def _rebind(row, **changes):
    result = deepcopy(row)
    evidence = result['original']['evidence']
    evidence['handout_source'].update(changes)
    for field in ('source_path', 'source_sha256', 'source_revision'):
        evidence[field] = evidence['handout_source'][field]
    if 'source_name' in changes:
        result['original']['title'] = changes['source_name'] + ' · ' + evidence['section_snapshot']['section']['title']
    if 'source_role' in changes:
        evidence['source_roles']['document_role'] = changes['source_role']
    result['original_sha256'] = digest(result['original'])
    return validate_resource(result)


def test_real_docx_groups_sealed_sections_keeps_eight_fields_and_is_read_only(course):
    before, saved = _files(course.root), deepcopy(course.resources)
    result = used.course_handout_sources(course.facade, course.resources)
    assert len(result['sources']) == 1 and result['unavailable'] == []
    candidate = result['sources'][0]
    assert set(candidate) == {'source', 'resource_ids', 'label', 'binding_sha256'}
    assert set(candidate['source']) == SOURCE_FIELDS
    assert candidate['source'] == saved[0]['original']['evidence']['handout_source']
    assert candidate['resource_ids'] == sorted(row['id'] for row in saved)
    assert candidate['label'] == '同名讲义.docx'
    assert all(result[key] is value for key, value in FLAGS.items())
    assert used.recheck_course_handout_source(course.facade, course.resources, candidate) is True
    assert course.facade.reads and set(course.facade.reads) == {('BATCH', 'SOURCE')}
    assert course.resources == saved and _files(course.root) == before


def test_working_edits_and_course_order_do_not_replace_original_or_binding(course):
    candidate = used.course_handout_sources(course.facade, course.resources)['sources'][0]
    originals = [deepcopy(row['original']) for row in course.resources]
    for row in course.resources:
        row['title'] = '教师已改的课堂问题'
        row['locked'] = False
        row['blocks'][1]['value'] = '教师自己的课堂讲解。'
        validate_resource(row)
    assert used.recheck_course_handout_source(course.facade, list(reversed(course.resources)), candidate)
    assert [row['original'] for row in course.resources] == originals
    assert all(row['blocks'][1]['value'] == '教师自己的课堂讲解。' for row in course.resources)


@pytest.mark.parametrize('batch,source_id,same_bytes', [
    ('BATCH', 'OTHER', False), ('OTHER-BATCH', 'SOURCE', True)])
def test_same_name_or_same_bytes_in_another_archive_never_merge(course, batch, source_id, same_bytes):
    data = course.path.read_bytes() if same_bytes else _word_bytes('同名但不同原文。')
    metadata = _add_document(course.facade, course.root, batch, source_id, data)
    other = _import_sections(course.facade, metadata)
    resources = [*course.resources, *other]
    result = used.course_handout_sources(course.facade, resources)
    assert len(result['sources']) == 2 and result['unavailable'] == []
    assert {tuple(row['source'][key] for key in ('batch_id', 'source_id'))
        for row in result['sources']} == {('BATCH', 'SOURCE'), (batch, source_id)}
    assert {frozenset(row['resource_ids']) for row in result['sources']} == {
        frozenset(row['id'] for row in course.resources), frozenset(row['id'] for row in other)}


def test_old_extraction_revision_is_unavailable_and_not_merged_with_current(course):
    stale = _rebind(course.resources[0], source_revision='0' * 64)
    stale['id'] = 'CR-' + uuid4().hex
    result = used.course_handout_sources(course.facade, [*course.resources, stale])
    assert len(result['sources']) == 1
    assert result['sources'][0]['resource_ids'] == sorted(row['id'] for row in course.resources)
    assert result['unavailable'][0]['resource_ids'] == [stale['id']]
    assert '版本' in result['unavailable'][0]['reason']


def test_provenance_path_is_preserved_and_part_of_exact_candidate_identity(course):
    historical = _rebind(course.resources[0], source_path='D:/封存/原始资料/同名讲义.docx')
    historical['id'] = 'CR-' + uuid4().hex
    rows = [course.resources[0], historical]
    before = deepcopy(rows)
    result = used.course_handout_sources(course.facade, rows)
    assert len(result['sources']) == 2
    paths = {row['source']['source_path'] for row in result['sources']}
    assert paths == {str(course.path), 'D:/封存/原始资料/同名讲义.docx'}
    candidate = result['sources'][0]
    changed = deepcopy(candidate)
    changed['source']['source_path'] = 'D:/另一个同名文件.docx'
    with pytest.raises(ValueError, match='已失效'):
        used.recheck_course_handout_source(course.facade, rows, changed)
    assert rows == before


@pytest.mark.parametrize('change', ['bytes', 'missing', 'role', 'name', 'wrong_id'])
def test_changed_or_wrong_exact_archive_is_unavailable_and_recheck_rejects(course, change):
    candidate = used.course_handout_sources(course.facade, course.resources)['sources'][0]
    original_reader = course.facade._imported_word_source
    if change == 'bytes':
        course.path.write_bytes(_word_bytes('源文件已经变化。'))
    elif change == 'missing':
        course.facade.records.pop(('BATCH', 'SOURCE'))
    elif change in {'role', 'name'}:
        setattr(course.facade.records[('BATCH', 'SOURCE')], change,
            'question' if change == 'role' else '另一份同名候选.docx')
    else:
        def wrong_source(batch_id, source_id):
            value = original_reader(batch_id, source_id)
            value.effective_source_file_id = 'OTHER'
            return value
        course.facade._imported_word_source = wrong_source
    before = _files(course.root)
    result = used.course_handout_sources(course.facade, course.resources)
    assert result['sources'] == [] and len(result['unavailable']) == 1
    with pytest.raises(ValueError, match='已失效'):
        used.recheck_course_handout_source(course.facade, course.resources, candidate)
    assert _files(course.root) == before


@pytest.mark.parametrize('change', ['missing_field', 'extra_field', 'invalid_hash_type', 'original_digest', 'promoted_gate'])
def test_incomplete_or_corrupt_sealed_resource_never_produces_a_source(course, change):
    row = deepcopy(course.resources[0])
    source = row['original']['evidence']['handout_source']
    if change == 'missing_field':
        source.pop('source_path')
    elif change == 'extra_field':
        source['inferred_section'] = '不要猜节'
    elif change == 'invalid_hash_type':
        source['source_revision'] = int('1' * 64)
        row['original']['evidence']['source_revision'] = source['source_revision']
    elif change == 'original_digest':
        row['original']['blocks'][1]['value'] = '改写封存原文'
    else:
        row['publication_allowed'] = True
    if change != 'original_digest':
        row['original_sha256'] = digest(row['original'])
    result = used.course_handout_sources(course.facade, [row])
    assert result['sources'] == [] and len(result['unavailable']) == 1
    assert course.facade.reads == []


@pytest.mark.parametrize('change', ['removed', 'original_changed', 'duplicate_id'])
def test_changed_classroom_bindings_reject_saved_candidate(course, change):
    candidate = used.course_handout_sources(course.facade, course.resources)['sources'][0]
    resources = deepcopy(course.resources)
    if change == 'removed':
        resources.pop()
    elif change == 'original_changed':
        resources[0] = _rebind(resources[0], source_path='D:/真实封存路径变化.docx')
    else:
        resources.append(deepcopy(resources[0]))
    with pytest.raises(ValueError, match='已失效'):
        used.recheck_course_handout_source(course.facade, resources, candidate)
    if change == 'duplicate_id':
        result = used.course_handout_sources(course.facade, resources)
        assert any('不唯一' in row['reason'] for row in result['unavailable'])


@pytest.mark.parametrize('field', ['document_source', 'preview'])
def test_independent_bad_reader_response_cannot_unlock_candidate(course, monkeypatch, field):
    actual = used.handout_source_context
    def mismatched(facade, source):
        context = actual(facade, source)
        if field == 'document_source':
            context[field]['source_path'] = 'D:/无关同名文档.docx'
        else:
            context[field]['revision'] = 'f' * 64
        return context
    monkeypatch.setattr(used, 'handout_source_context', mismatched)
    result = used.course_handout_sources(course.facade, course.resources)
    assert result['sources'] == [] and '读取结果' in result['unavailable'][0]['reason']


def test_only_actual_handout_knowledge_originals_are_candidates(course):
    personal = resource('personal_text', '同名讲义.docx', [block('text', 'personal', '教师文字')],
        {'source_path': '教师手工文字', 'source_sha256': 'unknown', 'source_revision': 'unknown'})
    course.facade.reads.clear()
    result = used.course_handout_sources(course.facade, [personal])
    assert result['sources'] == [] and result['unavailable'] == []
    assert course.facade.reads == []


def test_course_changes_during_read_are_rejected(course, monkeypatch):
    actual = used.handout_source_context
    def course_changed(facade, source):
        context = actual(facade, source)
        course.resources.pop()
        return context
    monkeypatch.setattr(used, 'handout_source_context', course_changed)
    with pytest.raises(ValueError, match='读取期间变化'):
        used.course_handout_sources(course.facade, course.resources)


def test_actual_word_changes_during_final_read_are_unavailable(course):
    original_reader = course.facade._imported_word_source
    def source_changed(batch_id, source_id):
        if course.facade.reads:
            course.path.write_bytes(_word_bytes('原Word在读取过程中改变。'))
        return original_reader(batch_id, source_id)
    course.facade._imported_word_source = source_changed
    result = used.course_handout_sources(course.facade, course.resources)
    assert result['sources'] == [] and '读取期间变化' in result['unavailable'][0]['reason']


def test_read_cancellation_is_propagated_without_writes(course):
    before, saved = _files(course.root), deepcopy(course.resources)
    stop = Event()
    original_reader = course.facade._imported_word_source
    def cancelled(batch_id, source_id):
        value = original_reader(batch_id, source_id)
        stop.set()
        return value
    course.facade._imported_word_source = cancelled
    with pytest.raises(ReadCancelled), read_cancel_scope(stop):
        used.course_handout_sources(course.facade, course.resources)
    assert _files(course.root) == before and course.resources == saved
