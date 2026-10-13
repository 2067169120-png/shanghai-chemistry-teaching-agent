"""Literal method units stay source-bound and outside all exercise scopes."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from threading import RLock

from docx import Document
import pytest

from integrations.deeptutor_shchem_v1 import desktop_handout_knowledge as handout
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import digest
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import PreparationSourcesService
from integrations.deeptutor_shchem_v1.desktop_word_question_index import index_word_questions
from test_handout_knowledge import make_facade, png, request_for, word_bytes


def outline_for(facade):
    source = handout.handout_knowledge_sources(facade)['sources'][0]
    return handout.handout_knowledge_outline(facade, source)


def method_sections(outline):
    return [s for s in outline['sections'] if s['id'].startswith('word-method-b')]


def method_lecture():
    doc = Document()
    doc.add_heading('考向1 实验方案的设计', 2)
    doc.add_paragraph('思维建模  完整实验方案设计思路')
    doc.add_paragraph('围绕目的，依次核对原理、装置、操作和结论。')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = '原始比较表'
    table.cell(1, 0).text = '实验目的'; table.cell(1, 1).text = '证据与结论'
    doc.add_paragraph().add_run().add_picture(BytesIO(png()))
    doc.add_paragraph()
    doc.add_paragraph('【变式训练1】下列实验方案正确的是________。')
    doc.add_paragraph('【答案】A')
    doc.add_paragraph('【解析】独立题的完整解析。')
    return doc


def test_whole_unstyled_unit_keeps_text_table_picture_and_blank_in_original_order(tmp_path):
    doc = method_lecture(); facade = make_facade(word_bytes(doc), tmp_path/'method')
    outline = outline_for(facade); method = method_sections(outline)[0]
    assert method == {'id':'word-method-b2', 'level':3,
        'title':'思维建模 完整实验方案设计思路', 'start':2, 'end':6}
    assert outline['sections'][0] == {'id':'word-heading-b1', 'level':2,
        'title':'考向1 实验方案的设计', 'start':1, 'end':9}
    assert not facade.paths.task_root.exists()
    result = handout.import_handout_knowledge(facade, [request_for(facade, method)], [], [])
    original = result['resources'][0]['original']
    snapshot = original['evidence']['section_snapshot']['blocks']
    assert [b['index'] for b in snapshot] == [2,3,4,5,6]
    assert snapshot[2]['table_rows'][0]['cells'][0]['column_span'] == 2
    assert [b['source_index'] for b in original['blocks'] if b['kind']=='image'] == [5]
    assert len(result['image_assets']) == 1
    assert original['evidence']['source_sha256'] == sha256(facade.value.data).hexdigest()
    assert all(b['role']=='knowledge' for b in original['blocks'])
    assert result['resources'][0]['candidate_only'] and not result['resources'][0]['human_reviewed']


@pytest.mark.parametrize('text', [
    '本节使用思维建模 实验设计思路', '思维建模', '思维建模实验设计思路',
    '思维建模 实验设计思路。这里仍是正文。', '思维建模 下列方案正确吗？',
    '【思维建模】实验设计思路', '思维建模 实验设计思路\n接着叙述正文',
])
def test_sentence_mentions_and_nonliteral_titles_are_not_promoted(tmp_path, text):
    doc = Document(); doc.add_heading('知识点1',3)
    doc.add_paragraph(text); doc.add_paragraph('完整正文。')
    facade = make_facade(word_bytes(doc), tmp_path/'mention')
    assert method_sections(outline_for(facade)) == []


@pytest.mark.parametrize('where', ['question','answer'])
def test_same_literal_title_inside_known_question_or_answer_never_becomes_knowledge(tmp_path, where):
    doc = Document(); doc.add_heading('考向1',3)
    doc.add_paragraph('【例题1】请回答下列问题：________。')
    if where == 'answer': doc.add_paragraph('【答案】A')
    doc.add_paragraph('思维建模  完整实验方案设计思路')
    doc.add_paragraph('这是该题原文的一部分，不能转成知识栏目。')
    if where == 'question': doc.add_paragraph('【答案】A')
    facade = make_facade(word_bytes(doc), tmp_path/'inside')
    if where == 'answer':
        # The automatic recognizer treats the literal label as a heading. An
        # explicit saved answer scope proves these paragraphs still belong to
        # the answer and must outrank that automatic heading interpretation.
        source = facade._imported_word_source('BATCH','SOURCE')
        preview = PreparationSourcesService(facade.paths.workspace_root).word_preview_bytes(
            source.content, source.filename)
        row = index_word_questions(preview)[0]
        row.update(answer_blocks=deepcopy(preview['blocks'][2:]), batch_id='BATCH',
            archive_source_id='SOURCE', source_revision=preview['revision'])
        class Words:
            _lock = RLock()
            def _catalog(self, locations, *, read_only):
                assert read_only
                return {'items':[row]}, {source.source_sha256:(source,'BATCH',preview,[])}
        facade._word_questions = lambda: Words()
    assert method_sections(outline_for(facade)) == []
    with pytest.raises(ValueError,match='完整题'):
        handout.import_handout_knowledge(facade,[request_for(facade,0)],[],[])
    assert not facade.paths.task_root.exists()


def test_saved_shared_context_excludes_whole_method_and_catalog_is_read_only(tmp_path):
    facade = make_facade(word_bytes(method_lecture()), tmp_path/'context')
    source = facade._imported_word_source('BATCH','SOURCE')
    preview = PreparationSourcesService(facade.paths.workspace_root).word_preview_bytes(
        source.content, source.filename)
    row = index_word_questions(preview)[0]
    row.update(context_blocks=deepcopy(preview['blocks'][1:6]), batch_id='BATCH',
        archive_source_id='SOURCE', source_sha256=preview['source_sha256'],
        source_revision=preview['revision'])
    calls = []
    class Words:
        _lock = RLock()
        def _catalog(self, locations, *, read_only):
            calls.append((locations, read_only))
            return {'items':[row]}, {preview['source_sha256']:(source,'BATCH',preview,[])}
    facade._word_questions = lambda: Words()
    assert method_sections(outline_for(facade)) == []
    assert calls == [({'BATCH':{'SOURCE'}},True)]
    assert not facade.paths.task_root.exists()


def test_changed_saved_catalog_version_is_rejected_before_any_publication(tmp_path):
    facade = make_facade(word_bytes(method_lecture()), tmp_path/'catalog-change')
    source = facade._imported_word_source('BATCH','SOURCE')
    preview = PreparationSourcesService(facade.paths.workspace_root).word_preview_bytes(
        source.content, source.filename)
    preview['revision'] = '0'*64
    class Words:
        _lock = RLock()
        def _catalog(self, locations, *, read_only):
            return {'items':[]}, {source.source_sha256:(source,'BATCH',preview,[])}
    facade._word_questions = lambda: Words()
    with pytest.raises(ValueError,match='同一原文版本'):
        outline_for(facade)
    assert not facade.paths.task_root.exists()


def test_adjacent_titles_empty_real_heading_and_source_order_delimit_units(tmp_path):
    doc = Document(); doc.add_heading('考向1',2)
    doc.add_paragraph('思维建模  第一个思路')
    doc.add_paragraph('思维建模：第二个完整思路')
    doc.add_paragraph('第二个完整单元正文。')
    doc.add_paragraph(style='Heading 7')
    doc.add_paragraph('这一段已在空标题之后。')
    doc.add_paragraph('思维建模  第三个完整思路')
    doc.add_paragraph('第三个完整单元正文。')
    facade = make_facade(word_bytes(doc), tmp_path/'adjacent')
    assert [(s['start'],s['end']) for s in method_sections(outline_for(facade))] == [(3,4),(7,8)]


def test_long_complete_body_has_no_fixed_scan_limit(tmp_path):
    doc = Document(); doc.add_paragraph('思维建模  完整设计流程')
    for n in range(140): doc.add_paragraph(f'第{n+1}条原始正文，不省略条件。')
    doc.add_paragraph('【例题1】请填写________。'); doc.add_paragraph('【答案】原答案。')
    facade = make_facade(word_bytes(doc), tmp_path/'long')
    assert [(s['start'],s['end']) for s in method_sections(outline_for(facade))] == [(1,141)]


def test_legacy_heading_ids_ranges_and_unaffected_outline_digest_are_exact(tmp_path):
    doc = Document(); doc.add_heading('一级栏目',1); doc.add_heading('知识点1',3)
    doc.add_paragraph('完整正文。'); doc.add_heading('二级栏目',2); doc.add_paragraph('后续正文。')
    facade = make_facade(word_bytes(doc), tmp_path/'legacy'); outline = outline_for(facade)
    styled = handout._headings(doc,outline['preview'])
    assert outline['sections'] == styled
    assert outline['outline_revision'] == digest({'index_revision':handout.OUTLINE_REVISION,
        'source_sha256':outline['source']['source_sha256'],
        'source_revision':outline['source']['source_revision'],'sections':styled})
    mixed = method_lecture(); mixed_facade = make_facade(word_bytes(mixed), tmp_path/'mixed')
    mixed_outline = outline_for(mixed_facade)
    assert [s for s in mixed_outline['sections'] if s['id'].startswith('word-heading-b')] == \
        handout._headings(mixed,mixed_outline['preview'])
    # A new method unit does not make its old, exercise-containing parent valid.
    with pytest.raises(ValueError,match='完整题'):
        handout.import_handout_knowledge(mixed_facade,[request_for(mixed_facade,0)],[],[])


def test_legacy_styled_request_and_saved_resource_still_validate_and_deduplicate(tmp_path):
    doc = Document(); doc.add_heading('知识点1 既有完整栏目',3)
    doc.add_paragraph('既有原始知识正文。')
    doc.add_paragraph().add_run().add_picture(BytesIO(png()))
    doc.add_heading('考向1',3); doc.add_paragraph('思维建模  完整方法流程')
    doc.add_paragraph('新方法单元原文。')
    facade = make_facade(word_bytes(doc),tmp_path/'compat'); outline = outline_for(facade)
    styled = handout._headings(doc,outline['preview'])
    legacy_revision = digest({'index_revision':handout.OUTLINE_REVISION,
        'source_sha256':outline['source']['source_sha256'],
        'source_revision':outline['source']['source_revision'],'sections':styled})
    assert legacy_revision != outline['outline_revision']
    legacy_request = request_for(facade,styled[0]); legacy_request['outline_revision'] = legacy_revision
    old = handout.import_handout_knowledge(facade,[legacy_request],[],[])
    resource = old['resources'][0]
    assert resource['original']['evidence']['section_snapshot']['outline_revision'] == legacy_revision
    resource['title'] = '教师已编辑的课件'; resource['locked'] = False
    resource['blocks'][0]['value'] = '教师工作文字。'
    old['image_assets'][0]['caption'] = '教师工作图题'
    current_request = request_for(facade,styled[0])
    again = handout.import_handout_knowledge(facade,[current_request],old['image_assets'],[resource])
    assert again == {'resources':[],'image_assets':old['image_assets']}
    assert resource['title'] == '教师已编辑的课件' and resource['blocks'][0]['value'] == '教师工作文字。'
    # The same legacy digest never authorizes a new method ID.
    new_request = request_for(facade,method_sections(outline)[0])
    new_request['outline_revision'] = legacy_revision
    with pytest.raises(ValueError,match='范围已变化'):
        handout.import_handout_knowledge(facade,[new_request],old['image_assets'],[resource])


@pytest.mark.parametrize('change', ['source_bytes','source_revision','outline_revision','section_range'])
def test_method_request_stale_source_version_or_tampered_range_fails_before_images(tmp_path, change):
    doc = method_lecture(); facade = make_facade(word_bytes(doc), tmp_path/'stale')
    request = request_for(facade, method_sections(outline_for(facade))[0])
    if change == 'source_bytes':
        doc.add_paragraph('新版本正文。'); facade.value.data = word_bytes(doc)
    elif change == 'source_revision': request['source']['source_revision'] = '0'*64
    elif change == 'outline_revision': request['outline_revision'] = '0'*64
    else: request['section']['end'] -= 1
    with pytest.raises(ValueError):
        handout.import_handout_knowledge(facade,[request],[],[])
    assert not facade.paths.task_root.exists()


PKG066 = Path(r'E:/项目/上海高中化学等级考/sh-chem-db/.intake/2026-07-30-user-teaching-pack/expanded/PKG-066/第07讲 化学实验方案的设计和评价（复习讲义）（上海专用）（解析版）.docx')


@pytest.mark.skipif(not PKG066.is_file(), reason='Local source is not available on this host')
def test_actual_pkg066_five_complete_method_units_keep_every_original_image(tmp_path):
    raw = PKG066.read_bytes(); before = (sha256(raw).hexdigest(), PKG066.stat().st_mtime_ns)
    assert before[0] == 'edd4696d06e0d3ad762319eb3f17f609729fec086781b393f212fc90b9eaef6e'
    facade = make_facade(raw, tmp_path/'pkg066'); facade.value.name = PKG066.name
    outline = outline_for(facade); methods = method_sections(outline)
    assert [(s['start'],s['end']) for s in methods] == [(56,61),(107,109),(257,258),(290,292),(311,313)]
    assert outline['source']['source_revision'] == '78d50c8cb3ef8af5f867e63c651c419aaccf0be290d275ee6d125f119d46184d'
    styled = handout._headings(Document(BytesIO(raw)),outline['preview'])
    assert [s for s in outline['sections'] if s['id'].startswith('word-heading-b')] == styled
    requests = [request_for(facade, s, title=f'怎样运用第{n+1}个原始方法单元？') for n,s in enumerate(methods)]
    result = handout.import_handout_knowledge(facade,requests,[],[])
    assert len(result['resources']) == 5
    for section, res in zip(methods,result['resources'],strict=True):
        original = res['original']; snapshot = original['evidence']['section_snapshot']
        assert snapshot['section'] == section
        assert [b['index'] for b in snapshot['blocks']] == list(range(section['start'],section['end']+1))
        expected = [a for a in outline['preview']['assets'] if section['start'] <= a['block_index'] <= section['end']]
        refs = [r for b in snapshot['blocks'] for r in b['images']]
        assert [(r['block_index'],r['source_asset_id'],r['original_sha256']) for r in refs] == \
            [(a['block_index'],a['asset_id'],a['sha256']) for a in expected]
        assert original['evidence']['source_sha256'] == before[0]
        assert original['evidence']['source_revision'] == outline['source']['source_revision']
        assert res['candidate_only'] and not res['human_reviewed'] and not res['publication_allowed']
    assert (sha256(PKG066.read_bytes()).hexdigest(),PKG066.stat().st_mtime_ns) == before
    assert not (tmp_path/'pkg066'/'desktop-state.v1.json').exists()
