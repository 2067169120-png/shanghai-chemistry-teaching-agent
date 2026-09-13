"""Structural regression only; chemistry and page appearance need direct review."""

import importlib.util
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

SCRIPT = (
    Path(__file__).resolve().parents[1]
    / "scripts/build_source_led_electrolyte_lesson.py"
)
SPEC = importlib.util.spec_from_file_location("electrolyte_exemplar", SCRIPT)
BUILDER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(BUILDER)


class SourceLedLessonTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="shchem-exemplar-test-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        # Layout fixture only; no source/chemistry fidelity claim.
        self.figure = self.root / "fixture.png"
        Image.new("RGB", (1400, 864), "white").save(self.figure)
        self.deck = BUILDER.build_deck(self.figure)

    def test_lesson_time_and_order(self):
        self.assertEqual(21, len(self.deck.prs.slides))
        self.assertEqual(40, sum(s["minutes"] for s in self.deck.records))
        self.assertEqual(list(range(1, 22)), [s["page"] for s in self.deck.records])

    def test_response_pairs_have_separate_question_and_answer_slides(self):
        for key in ("boundary", "strength", "equation", "recall", "exit"):
            question = [
                r["page"]
                for r in self.deck.records
                if r["response_pair"] == key + "-question"
            ]
            answer = [
                r["page"]
                for r in self.deck.records
                if r["response_pair"] == key + "-answer"
            ]
            self.assertEqual(1, len(question))
            self.assertEqual([question[0] + 1], answer)

    def test_actual_notes_and_source_on_each_slide(self):
        for slide, record in zip(self.deck.prs.slides, self.deck.records):
            self.assertTrue(record["source"])
            self.assertIn(
                record["teacher_notes"], slide.notes_slide.notes_text_frame.text
            )
            self.assertGreater(len(record["teacher_notes"]), 45)
        notes = [r["page"] for r in self.deck.records if r["notebook_section"]]
        self.assertEqual([5, 8, 10, 17, 19], notes)

    def test_source_figure_is_real_picture_shape_only_on_observation_slide(self):
        locations = [
            i + 1
            for i, slide in enumerate(self.deck.prs.slides)
            for shape in slide.shapes
            if shape.shape_type == 13
        ]
        self.assertEqual([6], locations)

    def test_equations_are_editable_with_super_and_subscripts(self):
        path = self.root / "deck.pptx"
        self.deck.prs.save(path)
        deck = Presentation(path)
        baselines = []
        for slide in deck.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    self.assertNotIn("^{", shape.text)
                    self.assertNotIn("_{", shape.text)
                    baselines += [
                        run._r.get_or_add_rPr().get("baseline")
                        for p in shape.text_frame.paragraphs
                        for run in p.runs
                    ]
        self.assertIn("30000", baselines)
        self.assertIn("-20000", baselines)
        question = "\n".join(s.text for s in deck.slides[14].shapes if s.has_text_frame)
        self.assertNotIn("Ba2+", question)
        answer = "\n".join(s.text for s in deck.slides[15].shapes if s.has_text_frame)
        self.assertIn("Ba2+", answer)

    def test_text_boxes_stay_on_page(self):
        for box in self.deck.boxes:
            x, y, w, h = box["bbox_inches"]
            self.assertGreaterEqual(min(x, y, w, h), 0)
            self.assertLessEqual(x + w, BUILDER.W + 0.01)
            self.assertLessEqual(y + h, BUILDER.H + 0.01)

    def test_notebook_is_two_writable_pages_without_answers(self):
        path = self.root / "notebook.pdf"
        BUILDER.build_notebook(path)
        pdf = PdfReader(path)
        self.assertEqual(2, len(pdf.pages))
        content = "\n".join(page.extract_text() for page in pdf.pages)
        self.assertIn("比较维度", content)
        self.assertIn("独立重建", content)
        self.assertNotIn("全部电离", content)
        self.assertNotIn("Ba2+", content)


if __name__ == "__main__":
    unittest.main()
