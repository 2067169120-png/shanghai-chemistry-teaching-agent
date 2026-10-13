"""A classroom map resumes only its exact archived lecture's complete questions."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from test_desktop_visual_import_facade import FakeProviderStore, _facade, _png, desktop_paths as desktop_paths

from integrations.deeptutor_shchem_v1.desktop_handout_images import handout_image_options
from integrations.deeptutor_shchem_v1.desktop_handout_questions import handout_question_facade
from integrations.deeptutor_shchem_v1.desktop_handout_knowledge import (
    handout_knowledge_sources, handout_knowledge_outline)
from integrations.deeptutor_shchem_v1.desktop_composition_sources import import_questions, merge_question_result
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import visible_blocks
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore


@pytest.fixture
def scoped_lecture(desktop_paths, tmp_path, monkeypatch):
    monkeypatch.setattr('integrations.deeptutor_shchem_v1.desktop_word_questions.load_attribute_catalog',
                        lambda *a, **k: {'knowledge_points': [], 'nodes': []})
    doc = Document(); doc.add_heading('合成课堂讲义', 1)
    doc.add_paragraph('【共同材料】根据下列实验图回答各题。')
    doc.add_paragraph().add_run().add_picture(BytesIO(_png('blue')))
    doc.add_paragraph('【即学即练1】观察完整图示后回答问题。')
    doc.add_paragraph().add_run().add_picture(BytesIO(_png('red')))
    doc.add_paragraph('【答案】参考答案及完整解析。')
    doc.add_paragraph().add_run().add_picture(BytesIO(_png('red')))
    path = tmp_path / '合成课堂讲义.docx'; doc.save(path)
    provider = FakeProviderStore(configured=False); facade = _facade(desktop_paths, provider)
    batch = facade.save_visual_import_batch(handout_files=(path,), source_type='教师讲义')
    binding = handout_image_options(facade)['images'][0]
    scoped = handout_question_facade(facade, binding)
    row = scoped.word_question_catalog()['items'][0]
    choices = [{'key': row['key'], 'revision': row['revision'], 'points': 3}]
    return facade, scoped, binding, choices, provider, path, batch


@pytest.fixture(params=['picture', 'document'])
def source_kind_lecture(scoped_lecture, request):
    facade, scoped, binding, choices, provider, path, batch = scoped_lecture
    if request.param == 'document':
        metadata = next(row for row in handout_knowledge_sources(facade)['sources']
                        if row['batch_id'] == binding['batch_id']
                        and row['source_id'] == binding['source_id'])
        binding = handout_knowledge_outline(facade, metadata)['source']
        scoped = handout_question_facade(facade, binding)
    return facade, scoped, binding, choices, provider, path, batch


def test_complete_shared_material_images_and_teacher_answer_survive_scope(source_kind_lecture):
    facade, scoped, binding, choices, provider, *_ = source_kind_lecture
    before = facade.state_store.snapshot()
    reference = scoped.word_question_reference(choices)
    assert facade.state_store.snapshot() == before
    assert not (facade.paths.task_root / 'preparation-v1/images').exists()
    result = import_questions(scoped, choices, [], expected_reference=reference)
    item = result['resources'][0]
    assert [b['role'] for b in item['blocks'] if b['kind'] == 'image'] == ['shared_material', 'question', 'answer']
    assert any('共同材料' in b['value'] for b in item['blocks'] if b['role'] == 'shared_material' and b['kind'] == 'text')
    assert any('完整解析' in b['value'] for b in item['blocks'] if b['role'] == 'answer' and b['kind'] == 'text')
    assert all(b['role'] != 'answer' for b in visible_blocks(item))
    assert item['original']['evidence']['source_sha256'] == binding['source_sha256']
    assert item['original']['evidence']['source_revision'] == binding['source_revision']
    assert item['original']['evidence']['question_key'] == choices[0]['key']
    assert item['locked'] and not item['human_reviewed'] and not item['publication_allowed']
    store = PreparationImageStore(facade.paths.task_root / 'preparation-v1/images')
    assert all(store.load(a) for a in result['image_assets'])
    assert facade.state_store.snapshot() == before and provider.borrow_calls == 0


def test_other_same_byte_archive_cannot_redirect_reads_or_import(source_kind_lecture, tmp_path):
    facade, scoped, binding, choices, _, path, batch = source_kind_lecture
    copied = tmp_path / '另一个目录' / path.name; copied.parent.mkdir(); copied.write_bytes(path.read_bytes())
    earlier = tmp_path / '前一份资料.docx'; document = Document(); document.add_paragraph('合成其他资料。'); document.save(earlier)
    other = facade.save_visual_import_batch(handout_files=(earlier, copied), source_type='教师讲义', group_id='copied-archive')
    assert other.batch_id != batch.batch_id
    other_binding = next(row for row in handout_image_options(facade)['images']
                         if row['batch_id'] == other.batch_id and row['source_sha256'] == binding['source_sha256'])
    assert other_binding['source_sha256'] == binding['source_sha256']
    assert other_binding['source_id'] != binding['source_id']
    service = facade._word_questions()
    service._catalog({other.batch_id: {other_binding['source_id']}}, read_only=True)
    assert service._locations[choices[0]['key']] == (other.batch_id, other_binding['source_id'])
    original = scoped.word_question_reference(choices)
    assert service._locations[choices[0]['key']] == (batch.batch_id, binding['source_id'])
    service._catalog({other.batch_id: {other_binding['source_id']}}, read_only=True)
    scoped.word_question_source(choices[0]['key'], choices[0]['revision'])
    assert service._locations[choices[0]['key']] == (batch.batch_id, binding['source_id'])
    service._catalog({other.batch_id: {other_binding['source_id']}}, read_only=True)
    layout = scoped.word_question_reading_layout(choices[0]['key'], choices[0]['revision'])
    assert all(p['role'] != 'answer' for p in layout['paragraphs'])
    assert service._locations[choices[0]['key']] == (batch.batch_id, binding['source_id'])
    service._catalog({other.batch_id: {other_binding['source_id']}}, read_only=True)
    result = import_questions(scoped, choices, [], expected_reference=original)
    assert len(result['resources']) == 1
    assert service._locations[choices[0]['key']] == (batch.batch_id, binding['source_id'])


@pytest.mark.parametrize('change', ['source_revision', 'picture', 'question_revision', 'foreign_key', 'catalog_archive'])
def test_stale_binding_or_foreign_question_cannot_copy_images(source_kind_lecture, monkeypatch, change):
    facade, scoped, binding, choices, *_ = source_kind_lecture
    reference = scoped.word_question_reference(choices)
    if change == 'source_revision': scoped.binding['source_revision'] = 'f' * 64
    elif change == 'picture': scoped.binding['sha256'] = 'f' * 64
    elif change == 'question_revision': choices[0]['revision'] = 'f' * 64
    elif change == 'foreign_key': choices[0]['key'] = 'another-source-question'
    else:
        service = facade._word_questions(); original = service._catalog
        def altered(*a, **k):
            catalog, inventory = original(*a, **k)
            catalog['items'][0]['archive_source_id'] = 'FOREIGN'
            return catalog, inventory
        monkeypatch.setattr(service, '_catalog', altered)
    with pytest.raises(ValueError): import_questions(scoped, choices, [], expected_reference=reference)
    assert not (facade.paths.task_root / 'preparation-v1/images').exists()


def test_scope_reads_original_pixels_and_rejects_change_during_read(source_kind_lecture, monkeypatch):
    facade, scoped, binding, choices, *_ = source_kind_lecture
    row = scoped.word_question_catalog()['items'][0]
    asset = next(b['assets'][0] for b in row['question_blocks'] if b.get('assets'))
    result = scoped.word_question_image(choices[0]['key'], choices[0]['revision'], asset['asset_id'])
    assert result['bytes']
    read = facade.word_question_image
    def changed(*a, **k):
        result = read(*a, **k); scoped.binding['source_revision'] = 'f' * 64
        return result
    monkeypatch.setattr(facade, 'word_question_image', changed)
    with pytest.raises(ValueError): scoped.word_question_image(choices[0]['key'], choices[0]['revision'], asset['asset_id'])
    assert not (facade.paths.task_root / 'preparation-v1/images').exists()


def test_same_question_use_deduplicates_and_preserves_teacher_title(source_kind_lecture):
    _, scoped, _, choices, *_ = source_kind_lecture
    reference = scoped.word_question_reference(choices)
    first = import_questions(scoped, choices, [], expected_reference=reference)
    first['resources'][0]['title'] = '教师整理的课堂问题标题'
    before = deepcopy(first)
    second = import_questions(scoped, choices, [], expected_reference=reference)
    assert merge_question_result(second, first['resources'], first['image_assets']) == {'resources': [], 'image_assets': first['image_assets']}
    assert first == before
