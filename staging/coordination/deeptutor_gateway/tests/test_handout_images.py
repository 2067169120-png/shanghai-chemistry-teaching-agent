"""Original pixels, document identity and roles survive explicit lecture handoff."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from random import Random
from threading import Event
from types import SimpleNamespace
from PIL import Image
import pytest

from integrations.deeptutor_shchem_v1 import desktop_handout_images as handout
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import validate_resource, digest
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
from integrations.deeptutor_shchem_v1.reader_cancellation import ReadCancelled, read_cancel_scope
from test_composition_teaching_images import source_question
from test_composition_render import app


@pytest.fixture
def lecture(tmp_path):
    buffer = BytesIO(); Image.frombytes('RGB',(160,80),Random(11).randbytes(160*80*3)).save(buffer,format='PNG')
    data = buffer.getvalue(); image_hash = sha256(data).hexdigest(); calls = []
    asset = {'asset_id':'word-b20-image1','label':'原图','block_index':20,
        'mime_type':'image/png','sha256':image_hash,'bytes_count':len(data),'preview_supported':True}
    value = {'source_name':'合成讲义.docx','source_sha256':'a'*64,'revision':'b'*64,'assets':[asset]}
    def preview(*args,read_only=False): calls.append(read_only); return deepcopy(value)
    facade = SimpleNamespace(paths=SimpleNamespace(task_root=tmp_path/'tasks'),
        list_imported_word_batches=lambda:[SimpleNamespace(batch_id='BATCH')],
        imported_word_sources=lambda _: [{'source_id':'SOURCE','source_name':value['source_name'],'source_sha256':value['source_sha256']}],
        imported_word_preview=preview, imported_word_path=lambda *args:str(tmp_path/'原始讲义.docx'),
        imported_word_asset=lambda *args:{'bytes':data})
    row = handout.handout_image_options(facade)['images'][0]
    return SimpleNamespace(facade=facade,row=row,value=value,data=data,calls=calls,
        selection={'source':row,'title':'这些知识点怎样联系起来？','purpose':'沿分支说明概念关系。','teacher_only':False})


def test_catalog_and_full_preview_are_read_only_without_image_copy(lecture):
    f=lecture; report=handout.handout_image_options(f.facade); picture=handout.handout_image_preview(f.facade,f.row)
    assert report['sources_read']==1 and not report['issues'] and report['images']==[f.row]
    assert picture['bytes']==f.data and all(f.calls) and not f.facade.paths.task_root.exists()


@pytest.mark.parametrize('change',['document','revision','block','image','bytes','disappear','during_read'])
def test_stale_sources_fail_before_any_handoff_copy(lecture,change):
    f=lecture
    if change=='document':f.value['source_sha256']='c'*64
    elif change=='revision':f.value['revision']='c'*64
    elif change=='block':f.value['assets'][0]['block_index']=21
    elif change=='image':f.value['assets'][0]['sha256']='c'*64
    elif change=='bytes':f.value['assets'][0]['bytes_count']+=1
    elif change=='disappear':f.value['assets']=[]
    else:
        def changed(*args):f.value['revision']='c'*64;return {'bytes':f.data}
        f.facade.imported_word_asset=changed
    with pytest.raises(ValueError):handout.import_handout_image(f.facade,[],[],f.selection)
    assert not f.facade.paths.task_root.exists()


def test_exact_original_picture_and_word_provenance_are_kept_with_new_use_and_duplicate(lecture):
    f=lecture; result=handout.import_handout_image(f.facade,[],[],f.selection); item=result['resources'][0]
    asset=result['image_assets'][0]; store=PreparationImageStore(f.facade.paths.task_root/'preparation-v1/images')
    assert store.load(asset)==f.data and item['original']['evidence']['source_sha256']=='a'*64
    assert item['original']['evidence']['handout_source']==f.row
    assert item['original']['blocks'][0]['source_index']==20 and item['title']==f.selection['title']
    assert validate_resource(item)==item and not item['publication_allowed'] and not item['human_reviewed']
    assert handout.import_handout_image(f.facade,result['image_assets'],[item],f.selection)['duplicate']
    other=handout.import_handout_image(f.facade,result['image_assets'],[item],{**f.selection,'purpose':'课末重新连接分支。'})
    assert len(other['resources'])==1 and other['image_assets']==result['image_assets']


def test_teacher_answer_picture_keeps_role_even_if_requested_for_classroom(lecture):
    f=lecture; imported=handout.import_handout_image(f.facade,[],[],f.selection); asset=imported['image_assets'][0]
    answer=source_question(asset,'answer'); result=handout.import_handout_image(f.facade,[asset],[answer],f.selection)
    item=result['resources'][0]; assert all(b['role']=='teacher_note' for b in item['blocks'])
    assert item['original']['evidence']['source_roles']==['answer']
    altered=deepcopy(item); altered['original']['blocks'][0]['role']='material'
    altered['original_sha256']=digest(altered['original']);altered['blocks']=deepcopy(altered['original']['blocks'])
    with pytest.raises(ValueError):validate_resource(altered)


def test_duplicate_original_use_preserves_edited_work_title_and_explanation(lecture):
    f = lecture; imported = handout.import_handout_image(f.facade, [], [], f.selection)
    item = deepcopy(imported['resources'][0]); item['locked'] = False
    item['title'] = '教师修改后的引导问题？'
    item['blocks'][1]['value'] = '教师重新安排的讲解与笔记。'
    assert validate_resource(item) == item
    before = deepcopy(item)
    result = handout.import_handout_image(f.facade, imported['image_assets'], [item], f.selection)
    assert result['duplicate'] and result['resources'] == []
    assert result['image_assets'] == imported['image_assets'] and item == before
    # Changing only the new title does not create a second copy of the same use.
    again = handout.import_handout_image(f.facade, imported['image_assets'], [item],
        {**f.selection, 'title': '重新输入的课堂问题？'})
    assert again['duplicate'] and item == before
    another_use = handout.import_handout_image(f.facade, imported['image_assets'], [item],
        {**f.selection, 'purpose': '复习时对照另一条知识分支。'})
    assert not another_use['duplicate'] and len(another_use['resources']) == 1
    another_audience = handout.import_handout_image(f.facade, imported['image_assets'], [item],
        {**f.selection, 'teacher_only': True})
    assert not another_audience['duplicate']
    assert all(b['role'] == 'teacher_note' for b in another_audience['resources'][0]['blocks'])


def test_bound_original_image_and_block_cannot_be_replaced_by_resealing_resource(lecture):
    f=lecture; item=handout.import_handout_image(f.facade,[],[],f.selection)['resources'][0]
    for field,value in [('value','IMG-'+'c'*64),('source_index',21)]:
        altered=deepcopy(item);altered['original']['blocks'][0][field]=value
        altered['original_sha256']=digest(altered['original']);altered['blocks']=deepcopy(altered['original']['blocks'])
        with pytest.raises(ValueError):validate_resource(altered)


def test_cancelled_read_or_import_never_creates_cached_preview_or_image_file(lecture):
    f=lecture; event=Event();event.set()
    with read_cancel_scope(event,check_boundaries=False):
        for operation in (lambda:handout.handout_image_options(f.facade),
                lambda:handout.import_handout_image(f.facade,[],[],f.selection)):
            with pytest.raises(ReadCancelled):operation()
    assert not f.facade.paths.task_root.exists()


def test_unsupported_original_image_is_not_silently_transformed(lecture):
    f=lecture;f.value['assets'][0]['mime_type']='image/tiff';row=handout.handout_image_options(f.facade)['images'][0]
    with pytest.raises(ValueError,match='原Word'):handout.handout_image_preview(f.facade,row)
    assert not f.facade.paths.task_root.exists()


def test_single_original_figure_fits_authored_frame_whole_without_clipping_or_shrinking_text(lecture,app,tmp_path):
    from integrations.deeptutor_shchem_v1.desktop_lesson_composition import new_composition,add_resource
    from integrations.deeptutor_shchem_v1.desktop_composition_render import layout_composition,write_pptx
    from pptx import Presentation
    f=lecture; result=handout.import_handout_image(f.facade,[],[],f.selection);item=result['resources'][0]
    c=new_composition('16:9');c=add_resource(c,item,{'page_id':c['pages'][0],'x':12,'y':12,'width':150,'height':72})
    value=layout_composition(c,result['image_assets'],{result['image_assets'][0]['asset_id']:f.data},student=True)
    assert not value['errors'];elements=list(value['pages'].values())[0]
    image=next(e for e in elements if e['kind']=='image')
    assert image['width']/image['height']==pytest.approx(2) and image['y']+image['height']<=84
    assert all(e['font_px']==32 for e in elements if e['kind']=='text' and not e['bold'])
    target=tmp_path/'lecture.pptx';write_pptx(value,target,{result['image_assets'][0]['asset_id']:f.data})
    picture=next(s for s in Presentation(target).slides[0].shapes if s.shape_type==13)
    assert picture.image.blob==f.data
