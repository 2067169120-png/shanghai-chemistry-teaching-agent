import json
from copy import deepcopy
from types import SimpleNamespace

import pytest
from docx import Document
from test_desktop_prompt_blueprint import PAYLOAD, SECTIONS, compiler

from integrations.deeptutor_shchem_v1 import (
    desktop_handout_prompt_reference as reference,
)
from integrations.deeptutor_shchem_v1.desktop_blueprint_generation import (
    blueprint_prompt,
)
from integrations.deeptutor_shchem_v1.desktop_handout_practice import (
    DRAFT_ID,
    EXPORT_KIND,
)
from integrations.deeptutor_shchem_v1.desktop_prompt_blueprint import (
    PromptPreviewError,
    build_request,
    compile_preview,
)


@pytest.fixture
def source(tmp_path, monkeypatch):
    item = {
        "key": "q1",
        "revision": "r1",
        "practice_eligible": True,
        "source_name": "离子反应讲义.docx",
        "parent_title": "专项训练",
        "printed_number": "6",
        "source_document": {
            "relative_path": "expanded/private/source.docx",
            "sha256": "source-sha",
        },
        "editable_source": {"question_locators": ["local-only-locator"]},
    }
    draft = {
        "kind": "native_handout_practice_selection",
        "title": "离子反应参考",
        "selections": [{"key": "q1", "revision": "r1"}],
        "created_at": "2026-09-08",
    }
    records = {DRAFT_ID: draft}
    state = SimpleNamespace(snapshot=lambda: {"drafts": deepcopy(records)})
    # In-memory paragraphs only; no DOCX artifact is created or changed.
    document = Document()
    question = document.add_paragraph("选择正确的离子式：SO")
    question.add_run("4").font.subscript = True
    question.add_run("2-").font.superscript = True
    answer = document.add_paragraph("非官方解答示例")
    calls = []

    class Native:
        def __init__(self, root):
            assert root == tmp_path

        def read(self, chosen, role):
            calls.append(role)
            assert chosen == item
            return [question if role == "question" else answer]

    monkeypatch.setattr(reference, "NativeParagraphs", Native)
    candidates = SimpleNamespace(
        workspace=tmp_path, state=state, catalog=lambda: {"items": [item]}
    )
    selection = {
        key: reference.reference_options(state)[0][key]
        for key in ("reference_id", "revision")
    }
    selection["include_answers"] = False
    return candidates, records, item, selection, calls


def test_saved_selection_revision_tracks_title_order_and_content_not_print_spacing(
    source,
):
    service, records, item, selected, _ = source
    original = reference.reference_options(service.state)[0]
    records[DRAFT_ID]["answer_lines"] = 4
    assert (
        reference.reference_options(service.state)[0]["revision"]
        == original["revision"]
    )
    records[DRAFT_ID]["title"] += "修改"
    assert (
        reference.reference_options(service.state)[0]["revision"]
        != original["revision"]
    )
    records["export1"] = {"kind": EXPORT_KIND, "items": [item], "title": "导出过的练习"}
    assert len(reference.reference_options(service.state)) == 2


def test_actual_question_not_only_structure_and_no_answer_read_without_opt_in(source):
    service, records, _, selection, calls = source
    before = deepcopy(records)
    result = reference.compile_handout_reference(service, selection)
    assert result["question_count"] == 1
    assert "SO_{4}^{2-}" in json.dumps(result["evidence"], ensure_ascii=False)
    assert "非官方解答示例" not in json.dumps(result["evidence"], ensure_ascii=False)
    assert calls == ["question"]
    assert records == before
    selection["include_answers"] = True
    result = reference.compile_handout_reference(service, selection)
    assert calls == ["question", "question", "answer"]
    assert "非官方解答示例" in json.dumps(result["evidence"], ensure_ascii=False)


@pytest.mark.parametrize(
    "change",
    ["missing", "stale_selection", "stale_item", "incomplete", "duplicate", "size"],
)
def test_changed_or_incomplete_sources_do_not_silently_drop_or_truncate(
    source, monkeypatch, change
):
    service, records, item, selection, _ = source
    if change == "missing":
        records.clear()
    elif change == "stale_selection":
        records[DRAFT_ID]["title"] = "changed"
    elif change == "stale_item":
        item["revision"] = "changed"
    elif change == "incomplete":
        item["practice_eligible"] = False
    elif change == "duplicate":
        records[DRAFT_ID]["selections"] *= 2
        selection["revision"] = reference.reference_options(service.state)[0][
            "revision"
        ]
    else:
        monkeypatch.setattr(reference, "MAX_REFERENCE_CHARS", 5)
    with pytest.raises(reference.HandoutReferenceError):
        reference.compile_handout_reference(service, selection)


@pytest.mark.parametrize(
    "handout",
    [
        False,
        [],
        {},
        {"reference_id": "a", "revision": "r"},
        {"reference_id": "a", "revision": "r", "include_answers": "yes"},
    ],
)
def test_bad_handout_input_rejected(handout):
    with pytest.raises(PromptPreviewError):
        build_request({**PAYLOAD, "handout_reference": handout}, SECTIONS)


def test_compilation_sends_only_selected_evidence_and_keeps_binding_local(source):
    service, _, _, selection, _ = source
    payload = {**PAYLOAD, "handout_reference": selection}
    preview = compile_preview(
        service.workspace,
        payload,
        SECTIONS,
        compiler=compiler,
        handout_loader=lambda value: reference.compile_handout_reference(
            service, value
        ),
    )
    outbound = blueprint_prompt(preview)
    assert "SO_{4}^{2-}" in outbound
    assert outbound.count("SO_{4}^{2-}") == 1
    assert "借鉴哪一道讲义选题" in outbound
    assert "非官方解答示例" not in outbound
    for private in ("local-only-locator", "expanded/private", "source-sha", DRAFT_ID):
        assert private not in outbound
    assert preview["handout_reference"]["local_provenance"]
    assert preview["external_call_performed"] is False
    assert preview["eligibility"] == "blueprint_only"
    with pytest.raises(PromptPreviewError):
        compile_preview(service.workspace, payload, SECTIONS, compiler=compiler)


def test_script_style_inheritance_and_merged_table_preserved():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    paragraph = doc.add_paragraph("Cl")
    paragraph.style.font.superscript = True
    run = paragraph.add_run("-")
    assert reference.paragraph_text(paragraph) == "^{Cl-}"
    baseline = OxmlElement("w:vertAlign")
    baseline.set(qn("w:val"), "baseline")
    run._r.get_or_add_rPr().append(baseline)
    assert reference.paragraph_text(paragraph) == "^{Cl}-"
    paragraph.style.font.superscript = False
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "共同条件"
    table.cell(1, 0).text = "单位"
    p = table.cell(1, 1).paragraphs[0]
    p.add_run("mol·L")
    p.add_run("-1").font.superscript = True
    text = reference.native_text([table])
    assert '"column_span": 2' in text
    assert text.count("共同条件") == 1
    assert "mol·L^{-1}" in text


def test_dialog_selects_handout_confirmed_content_and_fits_small_window(
    source, monkeypatch
):
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QApplication, QMessageBox

    from integrations.deeptutor_shchem_v1.desktop_workbench.prompt_blueprint_dialog import (
        PromptBlueprintDialog,
    )

    app = QApplication.instance() or QApplication([])
    service, _, _, _, _ = source
    payloads, confirmations = [], []

    def compile_payload(payload):
        payloads.append(payload)
        result = compile_preview(
            service.workspace,
            payload,
            SECTIONS,
            compiler=compiler,
            handout_loader=lambda value: reference.compile_handout_reference(
                service, value
            ),
        )
        return {**result, "preview_id": "TEST-PREVIEW"}

    class Tasks:
        def submit(self, label, operation, *, on_success, on_failure):
            on_success(operation())

    facade = SimpleNamespace(
        prompt_curriculum_sections=lambda: [
            SimpleNamespace(section_key=k, display_label_zh=v)
            for k, v in SECTIONS.items()
        ],
        prompt_handout_references=lambda: reference.reference_options(service.state),
        preparation_profiles=lambda: [
            SimpleNamespace(provider_name="测试模型", model_id="test")
        ],
        compile_prompt_blueprint=compile_payload,
    )
    dialog = PromptBlueprintDialog(facade, Tasks())
    dialog.resize(420, 760)
    dialog.show()
    app.processEvents()
    assert dialog.width() == 420 and dialog.height() == 760
    assert dialog.inputs_scroll.horizontalScrollBar().maximum() == 0
    assert dialog.inputs_scroll.verticalScrollBar().maximum() > 0
    dialog.title_input.setText(PAYLOAD["title"])
    dialog.goal.setPlainText(PAYLOAD["learning_goal"])
    dialog.sections.item(0).setCheckState(Qt.CheckState.Checked)
    dialog.handout_reference.setCurrentIndex(1)
    assert dialog.include_handout_answers.isEnabled()
    dialog.include_handout_answers.setChecked(True)
    dialog.compile_button.click()
    assert payloads[-1]["handout_reference"]["include_answers"] is True
    assert "非官方解答示例" in dialog.evidence_output.toPlainText()

    def cancel(*args):
        confirmations.append(args[2])
        return QMessageBox.StandardButton.No

    monkeypatch.setattr(QMessageBox, "question", cancel)
    dialog.generate_button.click()
    assert "1 道讲义" in confirmations[0] and "非官方解答" in confirmations[0]
    saved_preview = dialog._preview
    saved_input = deepcopy(payloads[-1])
    dialog.include_handout_answers.setChecked(False)
    assert not dialog.generate_button.isEnabled()
    assert not dialog.task_output.toPlainText()
    from test_desktop_blueprint_generation import candidate

    saved_input["handout_reference"]["revision"] = "older-selection-revision"
    dialog._history = [{
        "input": saved_input, "preview": saved_preview,
        "result": {"candidate": candidate(), "model_id": "test", "latency_ms": 1},
    }]
    dialog._load_history(1)
    assert dialog.handout_reference.currentData()["revision"] == "older-selection-revision"
    assert dialog.include_handout_answers.isChecked()
    assert not dialog.generate_button.isEnabled()
    assert "共同语境" in dialog.generated_output.toPlainText()
    dialog.close()
