"""Explicit lecture selection, private full-question handoff and Qt lifetimes."""
from copy import deepcopy
from types import SimpleNamespace

import pytest

pytest.importorskip('PySide6')
from PySide6.QtCore import QPoint, Qt
from PySide6.QtWidgets import QDialog

from test_desktop_library_ui import _ManualBridge, qt_app
from test_desktop_studio_ui import settle
from test_handout_knowledge_picker_ui import _SyntheticLibrary
from test_word_question_dialog import _Facade, _Tasks, _check
from test_question_explorer_ui import window

from integrations.deeptutor_shchem_v1.desktop_workbench import handout_material_workspace as module
from integrations.deeptutor_shchem_v1.desktop_workbench.word_question_dialog import WordQuestionDialog


CATALOG = '读取讲义选材资料目录'
OUTLINE = '读取当前讲义选材原文'
CONFIRM = '核对本次讲义选材来源'


def _knowledge(outline, number=0):
    return {'source': deepcopy(outline['source']), 'outline_revision': outline['outline_revision'],
        'section': deepcopy(outline['sections'][number]), 'title': '分子之间有哪些作用？', 'teacher_only': False}


def _question(document):
    return {'key': 'Q1', 'revision': 'revision-Q1', 'points': 2,
        'source_sha256': document['source_sha256'], 'scoped_source_document': deepcopy(document)}


@pytest.fixture
def workspace(qt_app, monkeypatch):
    library = _SyntheticLibrary()
    bridge = _ManualBridge()
    calls = []
    def context(_facade, binding):
        calls.append(deepcopy(binding))
        outline = library.outlines[binding['source_id']]
        return {'document_source': deepcopy(outline['source']), 'preview': deepcopy(outline['preview']),
                'images': [], 'contexts': {}}
    monkeypatch.setattr(module, 'handout_knowledge_sources', library.catalog)
    monkeypatch.setattr(module, 'handout_knowledge_outline', library.outline)
    monkeypatch.setattr(module, 'handout_source_context', context)
    dialog = module.HandoutMaterialWorkspace(SimpleNamespace(), bridge)
    dialog.show()
    settle(qt_app)
    case = SimpleNamespace(dialog=dialog, library=library, bridge=bridge, context_calls=calls, app=qt_app)
    yield case
    dialog.reject()
    dialog.deleteLater()
    settle(qt_app)


def _ready(case, number=0):
    if not case.dialog._sources:
        case.bridge.succeed(case.bridge.take(CATALOG), run=True)
    case.dialog.source.setCurrentIndex(number + 1)
    case.bridge.succeed(case.bridge.take(OUTLINE), run=True)


def test_directory_reads_metadata_only_and_reading_does_not_select(workspace):
    case = workspace
    case.bridge.succeed(case.bridge.take(CATALOG), run=True)
    assert case.library.catalog_calls == 1
    assert not case.library.outline_calls and not case.context_calls
    _ready(case)
    case.dialog.navigation.setCurrentRow(1)
    assert len(case.library.outline_calls) == 1
    assert case.dialog._document == case.library.outlines[case.library.sources[0]['source_id']]['source']
    assert len(case.dialog._document) == 8
    assert case.dialog.navigation.count() == 3
    assert not case.dialog._selections and not case.dialog.confirm.isEnabled()
    assert case.dialog.request is None


def test_readable_picture_and_duplicate_titles_use_original_full_heading_ranges(workspace, monkeypatch):
    case = workspace
    outline = case.library.outlines[case.library.sources[0]['source_id']]
    marker = '【待查看原文：图片或图形】'
    outline['sections'][0]['title'] = marker
    for section in outline['sections'][1:]:
        section['title'] = '知识点2 完整原知识'
    outline['sections'][2]['level'] = 9
    heading = next(b for b in outline['preview']['blocks'] if b['index'] == 64)
    heading.update(text=marker, label='64 · ' + marker)
    asset = deepcopy(outline['preview']['assets'][0])
    asset.update(asset_id='word-b64-image1', block_index=64)
    outline['preview']['assets'].append(asset)
    outline['preview']['sections'] = []  # The independent full outline includes H3 and H9.
    original = deepcopy(outline)
    _ready(case)
    dialog = case.dialog
    assert dialog._outline == outline == original
    assert dialog.reader._source == original['preview']
    assert dialog.navigation.count() == 3 and dialog.reader.section_picker.count() == 4
    first = dialog.navigation.item(0)
    assert first.data(Qt.ItemDataRole.UserRole) == original['sections'][0]
    assert '原图栏目 · 区块64' in first.text() and marker not in first.text()
    assert '原区块 64—77 · 可查看2张原图' in first.text()
    visited = []
    monkeypatch.setattr(dialog.reader, 'scroll_to_block', lambda index: visited.append(index))
    for index in (0, 1, 2):
        dialog.navigation.setCurrentRow(index)
        assert dialog.navigation.item(index).data(Qt.ItemDataRole.UserRole) == original['sections'][index]
    assert visited == [64, 78, 92]
    assert dialog.navigation.item(1).text() != dialog.navigation.item(2).text()
    assert dialog.reader.section_picker.itemData(2) == 78
    assert dialog.reader.section_picker.itemData(3) == 92
    assert not dialog._selections and dialog.request is None


def test_mixed_selection_is_one_explicit_source_request_after_async_recheck(workspace):
    _ready(workspace)
    dialog = workspace.dialog
    outline = dialog._outline
    dialog._add_selections('knowledge', [_knowledge(outline)])
    image = {'source': deepcopy(next(iter(dialog._picture_rows.values()))),
             'title': '这张知识网怎样串起概念？', 'purpose': '联系主干概念', 'teacher_only': False}
    dialog._add_selections('image', [image])
    dialog._add_selections('question', [_question(dialog._document)], {'Q1': '为什么水结冰体积增大？'})
    assert dialog.selected.count() == 3 and dialog.confirm.isEnabled()
    dialog.confirm.click()
    assert dialog.request is None and not dialog.confirm.isEnabled()
    workspace.bridge.succeed(workspace.bridge.take(CONFIRM), run=True)
    assert dialog.result() == QDialog.DialogCode.Accepted
    assert set(dialog.request) == {'schema', 'source', 'knowledge_requests', 'image_selections', 'question_selections', 'selection_order'}
    assert dialog.request['selection_order'] == [{'kind': 'knowledge', 'index': 0},
        {'kind': 'image', 'index': 0}, {'kind': 'question', 'index': 0}]
    assert dialog.request['schema'] == module.SCHEMA
    assert dialog.request['source'] == outline['source']
    assert dialog.request['knowledge_requests'] == [_knowledge(outline)]
    assert dialog.request['image_selections'] == [image]
    assert dialog.request['question_selections'] == [_question(outline['source'])]
    assert len(workspace.context_calls) == 1


def test_switching_source_clears_mixed_pending_and_old_outline_cannot_return(workspace):
    case = workspace
    _ready(case)
    case.dialog._add_selections('knowledge', [_knowledge(case.dialog._outline)])
    case.dialog.source.setCurrentIndex(2)
    pending = case.bridge.take(OUTLINE)
    case.dialog.source.setCurrentIndex(1)
    current = case.bridge.take(OUTLINE)
    case.bridge.succeed(current, run=True)
    case.bridge.succeed(pending, run=True)
    assert not case.dialog._selections and case.dialog.request is None
    assert case.dialog._document['source_id'] == case.library.sources[0]['source_id']
    assert pending.task_id in case.bridge.cancelled


def _interleaved_selection(case):
    dialog = case.dialog
    source = deepcopy(next(iter(dialog._picture_rows.values())))
    first = {'source': source, 'title': '先观察分类图', 'purpose': '引出强弱区别', 'teacher_only': False}
    second = {'source': deepcopy(source), 'title': '再读分布图', 'purpose': '联系电离常数', 'teacher_only': False}
    dialog._add_selections('image', [first])
    dialog._add_selections('knowledge', [_knowledge(dialog._outline)])
    dialog._add_selections('question', [_question(dialog._document)], {'Q1': '用实验说明电离程度'})
    dialog._add_selections('image', [second])
    dialog._add_selections('knowledge', [_knowledge(dialog._outline, 1)])
    return first, second


def test_interleaved_order_moves_keep_selection_and_confirm_exact_teaching_sequence(workspace):
    _ready(workspace)
    dialog = workspace.dialog
    first, second = _interleaved_selection(workspace)
    before = deepcopy(dialog._selections)
    dialog.selected.setCurrentRow(3)
    dialog.move_up_button.click()
    assert dialog.selected.currentRow() == 2
    assert dialog._selections[2] == before[3]
    dialog.move_down_button.click()
    assert dialog.selected.currentRow() == 3 and dialog._selections == before
    dialog.selected.setCurrentRow(0)
    assert not dialog.move_up_button.isEnabled() and dialog.move_down_button.isEnabled()
    dialog.selected.setCurrentRow(4)
    assert dialog.move_up_button.isEnabled() and not dialog.move_down_button.isEnabled()
    dialog.confirm.click()
    assert not dialog.move_up_button.isEnabled() and not dialog.move_down_button.isEnabled()
    dialog.move_selection(-1)
    assert dialog._selections == before
    workspace.bridge.succeed(workspace.bridge.take(CONFIRM), run=True)
    assert dialog.request['selection_order'] == [
        {'kind': 'image', 'index': 0}, {'kind': 'knowledge', 'index': 0},
        {'kind': 'question', 'index': 0}, {'kind': 'image', 'index': 1},
        {'kind': 'knowledge', 'index': 1}]
    assert dialog.request['image_selections'] == [first, second]
    assert dialog.request['knowledge_requests'] == [_knowledge(dialog._outline), _knowledge(dialog._outline, 1)]
    assert dialog._selections == before


def test_reorder_after_stopping_recheck_does_not_accept_old_order(workspace):
    _ready(workspace)
    dialog = workspace.dialog
    _interleaved_selection(workspace)
    dialog.confirm.click()
    old = workspace.bridge.take(CONFIRM)
    dialog.stop_button.click()
    dialog.selected.setCurrentRow(0)
    dialog.move_down_button.click()
    assert dialog._selections[0]['kind'] == 'knowledge'
    workspace.bridge.succeed(old, run=True)
    assert dialog.request is None
    dialog.confirm.click()
    workspace.bridge.succeed(workspace.bridge.take(CONFIRM), run=True)
    assert dialog.request['selection_order'] == [
        {'kind': 'knowledge', 'index': 0}, {'kind': 'image', 'index': 0},
        {'kind': 'question', 'index': 0}, {'kind': 'image', 'index': 1},
        {'kind': 'knowledge', 'index': 1}]


def test_remove_and_replace_pending_material_keep_reordered_neighbours(workspace):
    _ready(workspace)
    dialog = workspace.dialog
    _interleaved_selection(workspace)
    dialog.selected.setCurrentRow(4)
    dialog.move_up_button.click()
    original = deepcopy(dialog._selections)
    replacement = {**original[3]['value'], 'title': '教师修订后的常数问题'}
    dialog._add_selections('knowledge', [replacement])
    assert [r['key'] for r in dialog._selections] == [r['key'] for r in original]
    assert dialog._selections[3]['value'] == replacement
    dialog.selected.setCurrentRow(2)
    dialog.remove_button.click()
    assert dialog.selected.currentRow() == 2
    assert [r['key'] for r in dialog._selections] == [original[i]['key'] for i in (0, 1, 3, 4)]
    dialog.confirm.click()
    workspace.bridge.succeed(workspace.bridge.take(CONFIRM), run=True)
    assert dialog.request['selection_order'] == [
        {'kind': 'image', 'index': 0}, {'kind': 'knowledge', 'index': 0},
        {'kind': 'knowledge', 'index': 1}, {'kind': 'image', 'index': 1}]


@pytest.mark.parametrize('width', [420, 1120])
def test_reorder_actions_fit_both_widths_without_horizontal_scroll(workspace, width):
    _ready(workspace)
    dialog = workspace.dialog
    _interleaved_selection(workspace)
    dialog.resize(width, 780)
    settle(workspace.app)
    dialog.content_scroll.ensureWidgetVisible(dialog.move_up_button)
    settle(workspace.app)
    assert dialog.width() == width and dialog.content_scroll.horizontalScrollBar().maximum() == 0
    for button in (dialog.move_up_button, dialog.move_down_button, dialog.remove_button, dialog.clear_button):
        assert button.width() > 0 and button.isVisible()
    dialog.selected.setCurrentRow(1)
    dialog.move_up_button.click()
    assert dialog.selected.currentRow() == 0 and dialog._selections[0]['kind'] == 'knowledge'


@pytest.mark.parametrize('action', ['stop', 'close'])
def test_stop_and_close_ignore_late_confirmation(workspace, action):
    _ready(workspace)
    dialog = workspace.dialog
    dialog._add_selections('knowledge', [_knowledge(dialog._outline)])
    dialog.confirm.click()
    pending = workspace.bridge.take(CONFIRM)
    if action == 'stop':
        dialog.stop_button.click()
    else:
        dialog.reject()
    workspace.bridge.succeed(pending, run=True)
    assert dialog.request is None and dialog.result() != QDialog.DialogCode.Accepted
    assert pending.task_id in workspace.bridge.cancelled
    if action == 'stop':
        assert dialog.confirm.isEnabled() and len(dialog._selections) == 1


def test_changed_source_fails_confirm_without_losing_pending_selection(workspace):
    _ready(workspace)
    dialog = workspace.dialog
    dialog._add_selections('knowledge', [_knowledge(dialog._outline)])
    dialog.confirm.click()
    changed = deepcopy(dialog._document)
    changed['source_revision'] = 'f' * 64
    workspace.bridge.succeed(workspace.bridge.take(CONFIRM), {'document_source': changed})
    assert dialog.request is None and len(dialog._selections) == 1
    assert '已变化' in dialog.status.text()


def test_same_knowledge_updates_pending_question_title_and_limit_is_atomic(workspace):
    _ready(workspace)
    dialog = workspace.dialog
    first = _knowledge(dialog._outline)
    dialog._add_selections('knowledge', [first])
    edited = {**first, 'title': '这种作用怎样影响沸点？'}
    dialog._add_selections('knowledge', [edited])
    assert dialog.selected.count() == 1 and dialog._selections[0]['value'] == edited
    before = deepcopy(dialog._selections)
    many = [{**edited, 'section': {**edited['section'], 'id': f'extra-{i}'}} for i in range(21)]
    with pytest.raises(ValueError, match='最多20'):
        dialog._add_selections('knowledge', many)
    assert dialog._selections == before


@pytest.mark.parametrize('width', [420, 1120])
def test_fixed_footer_is_reachable_at_narrow_and_wide_widths(workspace, width):
    _ready(workspace)
    dialog = workspace.dialog
    dialog.resize(width, 710)
    settle(workspace.app)
    assert dialog.width() == width
    for button in (dialog.confirm, dialog.stop_button, dialog.cancel_button):
        point = button.mapTo(dialog, QPoint(0, 0))
        assert point.y() >= 0 and point.y() + button.height() <= dialog.height()
        assert point.x() >= 0 and point.x() + button.width() <= dialog.width()
    assert dialog.content_scroll.horizontalScrollBar().maximum() == 0


class _KnowledgeChild(QDialog):
    def __init__(self, facade, parent=None, initial_source=None):
        super().__init__(parent)
        self.bound = deepcopy(initial_source)
        self.requests = None


def test_source_bound_child_collects_only_after_manual_confirmation(workspace, monkeypatch):
    _ready(workspace)
    monkeypatch.setattr(module, 'HandoutKnowledgePicker', _KnowledgeChild)
    dialog = workspace.dialog
    dialog.choose_knowledge()
    child = dialog._child
    assert child.bound == dialog._document and not dialog._selections
    child.requests = [_knowledge(dialog._outline)]
    child.accept()
    assert dialog._child is None and len(dialog._selections) == 1
    dialog.choose_knowledge()
    child = dialog._child
    child.requests = [_knowledge(dialog._outline, 1)]
    child.reject()
    assert len(dialog._selections) == 1


def test_child_cannot_add_a_different_file_or_survive_source_switch(workspace, monkeypatch):
    _ready(workspace)
    monkeypatch.setattr(module, 'HandoutKnowledgePicker', _KnowledgeChild)
    dialog = workspace.dialog
    dialog.choose_knowledge()
    child = dialog._child
    child.requests = [_knowledge(workspace.library.outlines[workspace.library.sources[1]['source_id']])]
    child.accept()
    assert not dialog._selections and '不属于当前' in dialog.status.text()
    dialog.choose_knowledge()
    child = dialog._child
    child.requests = [_knowledge(dialog._outline)]
    dialog.source.setCurrentIndex(2)
    child.finished.emit(QDialog.DialogCode.Accepted)
    assert not dialog._selections and dialog.request is None


def _word_facade():
    facade = _Facade()
    facade.binding = {'schema': 'shchem.handout-knowledge-source.v1', 'batch_id': 'batch-synthetic',
        'source_id': 'archive-synthetic', 'source_name': '合成完整讲义.docx', 'source_sha256': 'a' * 64,
        'source_revision': 'b' * 64, 'source_path': 'C:/synthetic/lecture.docx', 'source_role': 'handout'}
    item = facade.catalog['items'][0]
    item.update(source_id='a' * 64, batch_id=facade.binding['batch_id'],
        archive_source_id=facade.binding['source_id'], source_name=facade.binding['source_name'],
        source_sha256=facade.binding['source_sha256'], source_revision=facade.binding['source_revision'])
    facade.catalog['items'] = [item]
    facade.catalog['sources'] = [{'source_id': 'a' * 64, 'source_name': facade.binding['source_name']}]
    return facade


def _selection_dialog(qt_app):
    facade, tasks = _word_facade(), _Tasks()
    dialog = WordQuestionDialog(facade, tasks, required_source_id='a' * 64,
        handoff_mode='composition', selection_only=True)
    tasks.flush()
    _check(dialog, 'Q1')
    dialog._preview_selected()
    tasks.finish('预览 Word 选题')
    assert dialog.composition_button.isEnabled()
    return dialog, facade, tasks


def test_selection_only_word_rechecks_original_and_never_imports_or_persists(qt_app, monkeypatch):
    from integrations.deeptutor_shchem_v1 import desktop_composition_sources
    def forbidden(*_args, **_kwargs):
        raise AssertionError('selection-only must not copy assets or compile imported cards')
    monkeypatch.setattr(desktop_composition_sources, 'import_questions', forbidden)
    dialog, facade, tasks = _selection_dialog(qt_app)
    assert dialog.range_button.isHidden() and dialog.attributes_button.isHidden()
    assert not dialog._persist_choices and not dialog._allow_knowledge_links
    dialog.composition_button.click()
    tasks.finish('核对完整题并加入课件')
    assert dialog.result() == QDialog.DialogCode.Accepted
    assert dialog.composition_result is None and dialog.preparation_reference is None
    assert dialog.selection_result['selections'] == [_question(facade.binding)]
    assert len([call for call in facade.calls if call[0] == 'reference']) == 2
    assert not any(call[0] in {'saved', 'save', 'export', 'range'} for call in facade.calls)
    assert facade.saved == []
    dialog.deleteLater()


@pytest.mark.parametrize('change', ['reference', 'binding', 'close'])
def test_selection_only_rejects_changed_reference_foreign_row_and_late_callback(qt_app, change):
    dialog, facade, tasks = _selection_dialog(qt_app)
    dialog.composition_button.click()
    if change == 'reference':
        facade.reference_changed = True
    elif change == 'binding':
        dialog._items['Q1']['archive_source_id'] = 'foreign-archive'
    else:
        dialog.reject()
    tasks.finish('核对完整题并加入课件', allow_cancelled=True)
    assert dialog.selection_result is None and dialog.composition_result is None
    assert dialog.result() != QDialog.DialogCode.Accepted
    assert not any(call[0] in {'saved', 'save', 'export', 'range'} for call in facade.calls)
    dialog.reject()
    tasks.flush()
    dialog.deleteLater()


def test_answer_document_figures_are_teacher_only_even_before_collection(qt_app):
    library = _SyntheticLibrary(role='answer')
    outline = library.outlines[library.sources[0]['source_id']]
    bound = outline['source']
    asset = outline['preview']['assets'][0]
    row = {'schema': 'shchem.handout-picture.v1', **{key: bound[key] for key in
           ('batch_id', 'source_id', 'source_name', 'source_path', 'source_sha256', 'source_revision')},
           **{key: asset[key] for key in ('asset_id', 'block_index', 'mime_type', 'sha256', 'bytes_count')}}
    picker = module._BoundImagePicker(SimpleNamespace(), {'document_source': bound, 'images': [row]}, [],
                                     tasks=_ManualBridge())
    assert picker.teacher_only.isChecked() and not picker.teacher_only.isEnabled()
    picker.reject()
    picker.deleteLater()


class _WorkspaceChild(QDialog):
    def __init__(self, facade, tasks=None, parent=None):
        super().__init__(parent)
        self.request = None


def test_true_question_page_has_source_entry_and_emits_only_confirmed_request(window, monkeypatch):
    win, app = window
    page = win.library_page
    monkeypatch.setattr(module, 'HandoutMaterialWorkspace', _WorkspaceChild)
    calls = []
    page.handout_materials_requested.connect(calls.append)
    win.resize(800, 700)
    settle(app)
    assert page.handout_materials_button.isVisible()
    page.handout_materials_button.click()
    child = page._handout_materials_dialog
    page.open_handout_materials()
    assert page._handout_materials_dialog is child
    child.request = {'schema': module.SCHEMA, 'source': 'synthetic-request-only'}
    child.reject()
    assert calls == [] and page._handout_materials_dialog is None
    page.open_handout_materials()
    child = page._handout_materials_dialog
    # The parent signal is the subject here; avoid invoking the root's real
    # preparation importer with a deliberately synthetic incomplete request.
    page.handout_materials_requested.disconnect(win._handout_materials_to_preparation)
    child.request = {'schema': module.SCHEMA, 'source': 'synthetic-request-only'}
    child.accept()
    assert calls == [child.request] and page._handout_materials_dialog is None


def test_legacy_library_uses_the_same_source_material_request(qt_app, monkeypatch):
    from test_desktop_library_ui import _Facade as LibraryFacade
    from integrations.deeptutor_shchem_v1.desktop_workbench.library_page import LibraryPage
    monkeypatch.setattr(module, 'HandoutMaterialWorkspace', _WorkspaceChild)
    page = LibraryPage(LibraryFacade(), _ManualBridge())
    calls = []
    page.handout_materials_requested.connect(calls.append)
    page.open_handout_materials()
    child = page._handout_materials_dialog
    child.request = {'schema': module.SCHEMA, 'source': 'synthetic-request-only'}
    child.accept()
    assert calls == [child.request]
    page.open_handout_materials()
    child = page._handout_materials_dialog
    child.request = {'schema': module.SCHEMA, 'source': 'late'}
    page.close()
    child.finished.emit(QDialog.DialogCode.Accepted)
    assert len(calls) == 1
    page.deleteLater()


def test_real_archived_word_selection_reads_images_and_answers_without_profile_write(window, monkeypatch):
    from hashlib import sha256
    from io import BytesIO
    from docx import Document
    from PIL import Image
    from test_desktop_library_ui import _png_bytes
    from integrations.deeptutor_shchem_v1.intake_imports import LocalPageRenderer, RenderedPage
    from integrations.deeptutor_shchem_v1.desktop_handout_knowledge import handout_knowledge_sources, handout_knowledge_outline
    from integrations.deeptutor_shchem_v1.desktop_handout_questions import handout_question_facade
    win, app = window
    facade = win.facade
    # The fixture is wholly under tmp_path and uses authored bitmap previews;
    # this is a native Word/Qt workflow check, not Office layout validation.
    def authored_preview(_renderer, source_path, *, mime_type, work_root):
        work_root.mkdir(parents=True, exist_ok=False)
        output = work_root / 'authored-preview.png'
        Image.new('RGB', (204, 284), 'white').save(output)
        return [RenderedPage(output, 'image/png')]
    document = Document()
    document.add_heading('知识点1 完整原图阅读（合成测试）', 3)
    document.add_paragraph('完整原文知识段。')
    document.add_paragraph('【例1】观察下图，比较两种情形。')
    document.add_paragraph().add_run().add_picture(BytesIO(_png_bytes()))
    document.add_paragraph('【答案】A。')
    document.add_paragraph('【解析】以下图示仅供教师讲解。')
    document.add_paragraph().add_run().add_picture(BytesIO(_png_bytes()))
    source_path = facade.paths.workspace_root / '原图完整选材合成.docx'
    document.save(source_path)
    with monkeypatch.context() as seed_patch:
        seed_patch.setattr(LocalPageRenderer, 'render', authored_preview)
        facade.save_visual_import_batch(handout_files=(source_path,), source_type='教师讲义')
    catalog = handout_knowledge_sources(facade)
    source = next(row for row in catalog['sources'] if row['source_name'] == source_path.name)
    binding = handout_knowledge_outline(facade, source)['source']
    scoped = handout_question_facade(facade, binding)
    def manifest():
        return {str(path.relative_to(facade.paths.state_root)): (path.stat().st_size,
                path.stat().st_mtime_ns, sha256(path.read_bytes()).hexdigest())
                for path in facade.paths.state_root.rglob('*') if path.is_file()}
    before = manifest()
    tasks = _Tasks()
    dialog = WordQuestionDialog(scoped, tasks, handoff_mode='composition',
        required_source_id=binding['source_sha256'], selection_only=True)
    tasks.flush()
    ready = next(row for row in dialog._items.values() if row['export_ready'] and row['selection_ready'])
    _check(dialog, ready['key'])
    dialog.tabs.setCurrentIndex(1)
    tasks.flush()
    assert ready['answer_blocks'] and any(block.get('assets') for block in ready['question_blocks'])
    dialog._preview_selected()
    tasks.finish('预览 Word 选题')
    assert dialog.composition_button.isEnabled()
    dialog.composition_button.click()
    tasks.finish('核对完整题并加入课件')
    assert dialog.result() == QDialog.DialogCode.Accepted and dialog.composition_result is None
    assert dialog.selection_result['selections'][0]['scoped_source_document'] == binding
    assert dialog.selection_result['reference']['image_assets']
    assert manifest() == before
    dialog.deleteLater()
    settle(app)
