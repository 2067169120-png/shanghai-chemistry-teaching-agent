from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import pytest

from test_word_question_export import item, output_doc
from test_personal_visual_theme_writer import SyntheticTheme
from integrations.deeptutor_shchem_v1.desktop_word_question_export import export_word_questions
from integrations.deeptutor_shchem_v1.desktop_mixed_paper_export import build_mixed_paper_docx
from integrations.deeptutor_shchem_v1.desktop_paper_numbering import leading_label, RASTER_NOTICE
from integrations.deeptutor_shchem_v1.paper_export_renderer import build_synthetic_demo_bundle


def paragraphs(data, audience='student'):
    doc = Document(BytesIO(data[audience + '_bytes']))
    return [p.text for p in doc.paragraphs]


@pytest.mark.parametrize('value', ['27. 题干', '27．题干', '27、题干', '第27题：题干', '【例27】题干', '【例题27】题干', '27 （3分）题干'])
def test_supported_source_labels_are_replaced(value):
    doc = Document(); doc.add_paragraph(value)
    result = export_word_questions('合成编号', [item(doc, [1])])
    assert any(p.startswith('1. ') and '27' not in p for p in paragraphs(result))
    assert doc.paragraphs[0].text == value


@pytest.mark.parametrize('value', ['1.25 mol/L', '2025. 试题背景', '2H2 + O2', '1 mol', '（1）氧化剂', 'A. 硫酸'])
def test_chemistry_numbers_and_subquestions_are_not_source_labels(value):
    assert leading_label(value) is None


def test_cross_source_order_answer_identity_and_reorder():
    a, b = Document(), Document()
    a.add_paragraph('27. 氧化剂是_____。');a.add_paragraph('27. 参考结论甲')
    b.add_paragraph('38. 还原剂是_____。');b.add_paragraph('38. 参考结论乙')
    qa,qb=item(a,[1],[2],key='a'),item(b,[1],[2],key='b')
    before=deepcopy([qa,qb])
    for rows, expected in (([qa,qb],['1. 氧化剂','2. 还原剂']),([qb,qa],['1. 还原剂','2. 氧化剂'])):
        result=build_mixed_paper_docx('合成顺序', [{'kind':'word_question','question':q} for q in rows])
        for audience in ('student','teacher'):
            text='\n'.join(paragraphs(result,audience))
            assert all(value in text for value in expected)
            assert '27. ' not in text and '38. ' not in text
        answers='\n'.join(paragraphs(result,'teacher'))
        assert ('1. 参考结论甲' in answers) == (rows[0] is qa)
        assert '参考结论' not in '\n'.join(paragraphs(result))
    assert [qa,qb] == before


def test_multiple_questions_and_references_share_one_sequence_without_cascading():
    doc=Document()
    for value in ('12. 计算第一步。','13. 利用第12题的结果，注意浓度为12.5 mol/L。',
                  '12. 结论为12。','13. 根据第12题得到结论。'):
        doc.add_paragraph(value)
    one=item(doc,[1],[3],key='one')
    # Both selections must refer to the same saved DOCX bytes. Serializing the
    # Document twice can change ZIP timestamps at the two-second boundary.
    two={**deepcopy(one), 'key':'two', 'question_blocks':[{'index':2}], 'answer_blocks':[{'index':4}]}
    result=export_word_questions('合成引用',[one,two])
    text='\n'.join(paragraphs(result,'teacher'))
    assert '2. 利用第1题的结果，注意浓度为12.5 mol/L。' in text
    assert '2. 根据第1题得到结论。' in text
    result=export_word_questions('合成多问',[item(doc,[1,2],[3,4])])
    assert '2. 利用第1题' in '\n'.join(paragraphs(result))


def test_split_runs_and_omml_survive():
    doc=Document();p=doc.add_paragraph()
    p.add_run('2');p.add_run('7').bold=True;p.add_run('. SO')
    p.add_run('4').font.subscript=True
    p._p.append(parse_xml(f'<m:oMath {nsdecls("m")}><m:r><m:t>n/V</m:t></m:r></m:oMath>'))
    original=doc.element.xml
    result=output_doc(export_word_questions('公式',[item(doc,[1])]))
    assert any(p.text.startswith('1. SO') for p in result.paragraphs)
    assert result.element.xpath('.//m:t')[0].text=='n/V'
    assert result.element.xpath('.//w:vertAlign/@w:val')==['subscript']
    assert doc.element.xml==original


def test_word_numbers_continue_after_core_printed_questions_not_atomic_parts():
    bundle=build_synthetic_demo_bundle();doc=Document();doc.add_paragraph('99. 最后一题')
    count=len(bundle['student_plan']['visible']['theme_sections'][0]['printed_questions'])
    before=deepcopy(bundle)
    result=build_mixed_paper_docx('混合编号',[
        {'kind':'core_plan','bundle':bundle}, {'kind':'word_question','question':item(doc,[1])}])
    assert f'{count+1}. 最后一题' in '\n'.join(paragraphs(result))
    assert bundle==before


def test_visual_projected_labels_continue_and_raster_limitation_is_visible():
    fixture=SyntheticTheme();ref=fixture.image((10,10,220,90))
    fixture.question([ref]); fixture.theme['printed_questions'][0]['question_number']=37
    fixture.item['content']['source_scores']=[{'max_score':2}]
    visual, assets=fixture.freeze();before=deepcopy(visual)
    doc=Document();doc.add_paragraph('23. Word题')
    result=build_mixed_paper_docx('图文编号',[
        {'kind':'word_question','question':item(doc,[1])},
        {'kind':'personal_visual_theme','item':visual,'assets':assets}])
    for role in ('student','teacher'):
        body='\n'.join(paragraphs(result,role))
        assert '第2题' in body and '第37题' not in body
        assert RASTER_NOTICE in body
        with ZipFile(BytesIO(result[role+'_bytes'])) as z:
            images=[z.read(n) for n in z.namelist() if n.startswith('word/media/')]
        assert assets[ref] in images
    assert RASTER_NOTICE in result['warnings']
    assert visual==before


def test_example_steps_are_not_reclassified_as_new_questions():
    doc=Document()
    for value in ('【例27】实验任务。','1. 加入试剂。','2. 加热并观察。','（1）描述现象。','【答案】27. 对应答案。'):
        doc.add_paragraph(value)
    second=Document();second.add_paragraph('38. 下一题。')
    result=export_word_questions('步骤保留',[item(doc,[1,2,3,4],[5]),item(second,[1],key='next')])
    text='\n'.join(paragraphs(result,'teacher'))
    assert '1. 实验任务。' in text and '2. 下一题。' in text
    assert '1. 加入试剂。' in text and '2. 加热并观察。' in text
    assert '【答案】1. 对应答案。' in text


def test_known_nested_printed_questions_keep_parent_material_and_own_answers():
    doc=Document()
    for value in ('27. 阅读下列共同材料，按要求回答所有问题。','10. 问题甲_____。','11. 问题乙_____。',
                  '【答案】10. 答案甲。','11. 答案乙。'):
        doc.add_paragraph(value)
    grouped=item(doc,[1,2,3],[4,5]);grouped['nested_section_starts']=[2,3]
    result=export_word_questions('完整主题',[grouped])
    text='\n'.join(paragraphs(result,'teacher'))
    assert '27. ' not in text and '阅读下列共同材料' in text
    assert '1. 问题甲' in text and '2. 问题乙' in text
    assert '【答案】1. 答案甲' in text and '2. 答案乙' in text
    assert doc.paragraphs[0].text.startswith('27.')


@pytest.mark.parametrize('source', ['【变式3-1】条件与任务。','例27：条件与任务。','【典例一】条件与任务。','【随堂练习２７】条件与任务。'])
def test_editorial_exercise_labels_use_current_number(source):
    doc=Document();doc.add_paragraph(source)
    output=export_word_questions('教辅题号',[item(doc,[1])])
    assert '1. 条件与任务。' in paragraphs(output)
    assert doc.paragraphs[0].text==source


def test_full_width_number_delimiter_does_not_consume_numeric_conditions():
    doc=Document();doc.add_paragraph('27．25 ℃时，计算浓度。')
    result=export_word_questions('温度条件',[item(doc,[1])])
    assert '1. 25 ℃时，计算浓度。' in paragraphs(result)
    assert leading_label('27.25 mol/L') is None


def test_automatic_numbering_does_not_duplicate_current_prefix():
    doc=Document();doc.add_paragraph('按条件填写答案。',style='List Number')
    result=output_doc(export_word_questions('自动编号',[item(doc,[1])]))
    p=next(p for p in result.paragraphs if '按条件' in p.text)
    assert p.text=='1. 按条件填写答案。'
    assert p._p.xpath('./w:pPr/w:numPr/w:numId/@w:val')==['0']


def test_core_only_composer_uses_current_sequence_and_matching_audit(tmp_path):
    from integrations.deeptutor_shchem_v1 import paper_export_renderer as r
    bundle=build_synthetic_demo_bundle();before=deepcopy(bundle)
    for audience in ('student','teacher'):
        plan=deepcopy(bundle[audience+'_plan']);blueprint=deepcopy(bundle['blueprint'])
        plan['visible']['version_label_zh']='本地题篮导出 · 题目来源见教师版 · 不可发布'
        for ti,theme in enumerate(plan['visible']['theme_sections']):
            for qi,q in enumerate(theme['printed_questions']):
                old=27+qi
                q['question_number']=old
                blueprint['theme_bundles'][ti]['printed_questions'][qi]['source_number']=str(old)
                q['atomic_parts'][0]['question_blocks'][0]['text_zh']=f'{old}. 浓度12.5 mol/L。参考第27题。'
        saved=deepcopy(plan);out=tmp_path/(audience+'.docx')
        r.build_docx_from_plan(plan,preset=bundle['preset'],blueprint=blueprint,output_path=out)
        text='\n'.join(p.text for p in Document(out).paragraphs)
        assert '参考第1题' in text and '参考第27题' not in text
        assert '12.5 mol/L' in text and '27. 浓度' not in text
        audit=r.audit_docx(out,plan=plan,preset=bundle['preset'],blueprint=blueprint)
        assert audit['status']=='pass',audit['checks']
        assert plan==saved
    assert bundle==before


def test_core_references_never_jump_between_original_papers():
    from integrations.deeptutor_shchem_v1.desktop_paper_numbering import rewrite_core_text
    b=build_synthetic_demo_bundle();v=b['student_plan']['visible'];bp=b['blueprint']
    first=v['theme_sections'][0];first['printed_questions']=first['printed_questions'][:1]
    bp['theme_bundles'][0]['printed_questions']=bp['theme_bundles'][0]['printed_questions'][:1]
    bp['theme_bundles'][0]['printed_questions'][0]['source_number']='27'
    first['printed_questions'][0]['question_number']=27
    first['printed_questions'][0]['atomic_parts'][0]['question_blocks'][0]['text_zh']='27. 参考第38题，不应指到另一试卷。'
    second=deepcopy(first);second['printed_questions'][0]['question_number']=38
    second['printed_questions'][0]['atomic_parts'][0]['question_blocks'][0]['text_zh']='38. 本卷另一来源。'
    secondbp=deepcopy(bp['theme_bundles'][0]);secondbp['source']['paper_id']='ANOTHER-PAPER'
    secondbp['printed_questions'][0]['source_number']='38'
    v['theme_sections']=[first,second];bp['theme_bundles']=[bp['theme_bundles'][0],secondbp]
    rewrite_core_text(v,bp,1)
    assert '第38题' in first['printed_questions'][0]['atomic_parts'][0]['question_blocks'][0]['text_zh']


def test_explicit_reference_after_math_keeps_formula_and_run_properties():
    doc=Document();doc.add_paragraph('27. 先求浓度。')
    p=doc.add_paragraph('38. 使用公式：')
    p._p.append(parse_xml(f'<m:oMath {nsdecls("m")}><m:r><m:t>n/V</m:t></m:r></m:oMath>'))
    p.add_run('，利用第27题结果，12.5不变。').bold=True
    before=doc.element.xml
    exported=output_doc(export_word_questions('公式引用',[item(doc,[1,2])]))
    assert any('利用第1题结果，12.5不变。' in p.text for p in exported.paragraphs)
    assert exported.element.xpath('.//m:t')[0].text=='n/V'
    assert doc.element.xml==before
