"""One source crop, readable purpose and teacher-only provenance stay usable."""
from copy import deepcopy
from io import BytesIO
from hashlib import sha256
import pytest
from PIL import Image
from docx import Document
from pptx import Presentation
from test_composition_render import app, text
from test_composition_flow import composition, personal
from test_lesson_design import sample
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import block, resource, resource_frames, enable_composition
from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition, layout_composition, write_pptx
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents

PURPOSE = '读图用途：对照装置与产物照片，保留标签与完整图形。\n仅教师用途标记，不补写原图未给出的结论。'

def figure(tmp_path, purpose=PURPOSE):
    buffer = BytesIO(); Image.new('RGB', (815,668), '#548093').save(buffer, format='PNG')
    raw = buffer.getvalue(); store = PreparationImageStore(tmp_path/'images')
    asset = store.import_bytes(raw,'合成教材图','合成原页','仅用于排版回归')
    item = resource('textbook_image','教材图：完整标签和用途',[
        block('image','material',asset['asset_id']), block('text','teacher_note',purpose)],
        {'source_path':'合成教材.pdf','source_sha256':'a'*64,'source_revision':'b'*64,
         'pdf_pages':[39],'printed_pages':'unknown'})
    return item, asset, raw, store

@pytest.mark.parametrize('format',['A4','A3','16:9'])
def test_short_figure_moves_as_whole_group_with_readable_purpose_and_exact_bytes(app,tmp_path,format):
    item, asset, raw, store = figure(tmp_path)
    filler = personal('前序完整材料',[block('text','personal','前序条件。\n'*12)])
    original = composition(filler,item,format=format)
    arranged = arrange_composition(original,[asset],{asset['asset_id']:raw})
    layout = layout_composition(arranged,[asset],{asset['asset_id']:raw},student=False)
    fragments = [f for f in layout['fragments'] if f['resource_id']==item['id']]
    assert not layout['errors'] and len(fragments)==1
    parts = fragments[0]['elements']; pictures = [e for e in parts if e['kind']=='image']
    assert len(pictures)==1 and abs(pictures[0]['width']/pictures[0]['height']-815/668)<1e-6
    assert text(layout,item['blocks'][1]['id'])==PURPOSE.replace('\n','')
    assert all(e['font_px']==(32 if format=='16:9' else 15) for e in parts if e.get('block_id')==item['blocks'][1]['id'] and e['kind']=='text')
    assert arranged['resources']==original['resources'] and store.load(asset)==raw

def test_student_picture_excludes_teacher_purpose_in_canvas_and_editable_ppt(app,tmp_path):
    item,asset,raw,store=figure(tmp_path)
    value=arrange_composition(composition(item),[asset],{asset['asset_id']:raw})
    teacher=layout_composition(value,[asset],{asset['asset_id']:raw},student=False)
    student=layout_composition(value,[asset],{asset['asset_id']:raw},student=True)
    assert not teacher['errors'] and not student['errors']
    assert not text(student,item['blocks'][1]['id'])
    for name,layout in [('teacher',teacher),('student',student)]:
        path=tmp_path/(name+'.pptx');write_pptx(layout,path,{asset['asset_id']:raw})
        ppt=Presentation(path);slides=list(ppt.slides)
        assert len(slides)==1
        pictures=[s for s in slides[0].shapes if s.shape_type==13]
        assert len(pictures)==1 and pictures[0].image.blob==raw
        words=''.join(s.text for s in slides[0].shapes if s.has_text_frame)
        assert ('仅教师用途标记' in words)==(name=='teacher')
    assert store.load(asset)==raw

def test_long_purpose_continues_completely_without_shrinking_projection_text(app,tmp_path):
    purpose='完整用途、条件和适用范围。'*120
    item,asset,raw,store=figure(tmp_path,purpose)
    original=composition(item);value=arrange_composition(original,[asset],{asset['asset_id']:raw})
    layout=layout_composition(value,[asset],{asset['asset_id']:raw},student=False)
    assert not layout['errors'] and len(resource_frames(value['placements'][item['id']]))>1
    assert text(layout,item['blocks'][1]['id'])==purpose
    assert all(e['font_px']==32 for p in layout['pages'].values() for e in p if e.get('block_id')==item['blocks'][1]['id'] and e['kind']=='text')
    assert value['resources']==original['resources'] and store.load(asset)==raw

def test_manual_smaller_frames_split_soft_figure_group_without_rewriting_layout(app,tmp_path):
    item,asset,raw,store=figure(tmp_path)
    value=arrange_composition(composition(item),[asset],{asset['asset_id']:raw})
    second='CP-'+'c'*32;value['pages'].append(second)
    frame=value['placements'][item['id']];frame['height']-=16
    frame['continuations']=[{'page_id':second,'x':12,'y':12,'width':frame['width'],'height':156}]
    before=deepcopy(value);layout=layout_composition(value,[asset],{asset['asset_id']:raw},student=False)
    assert not layout['errors'] and len(layout['fragments'])==2 and value==before
    assert text(layout,item['blocks'][1]['id'])==PURPOSE.replace('\n','') and store.load(asset)==raw

@pytest.mark.parametrize('long',[False,True])
def test_word_bounded_group_includes_source_but_long_purpose_retains_normal_flow(app,tmp_path,long):
    item,asset,raw,_=figure(tmp_path,PURPOSE if not long else '完整条件与用途。'*200)
    payload=sample();payload['lesson_design']=enable_composition(payload['lesson_design'])
    payload['lesson_design']['composition']=composition(item,format='A4');payload['image_assets']=[asset]
    original = deepcopy(payload)
    write_documents(payload,tmp_path,{asset['asset_id']:raw})
    teacher=Document(tmp_path/'lesson_plan.docx');student=Document(tmp_path/'student_worksheet.docx')
    start=next(i for i,p in enumerate(teacher.paragraphs) if p.text==item['title'])
    end=next(i for i,p in enumerate(teacher.paragraphs[start:], start) if p.text==item['blocks'][1]['value'])
    group=teacher.paragraphs[start:end+1]
    purpose=group[-1]
    assert not purpose.paragraph_format.keep_with_next
    if not long:
        assert all(p.paragraph_format.keep_with_next for p in group[:-1])
        assert all(p.paragraph_format.keep_together for p in group)
    for document in (teacher,student):
        words='\n'.join(p.text for p in document.paragraphs)
        assert 'PDF文件页：39' not in words and '印刷页：待核对' not in words
        assert '合成教材.pdf' not in words
    assert '仅教师用途标记' not in ''.join(p.text for p in student.paragraphs)
    assert teacher.inline_shapes[0]._inline.xpath('.//a:blip')[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
    for doc in (teacher,student):
        blobs=[r.target_part.blob for r in doc.part.rels.values() if 'image' in r.reltype]
        assert raw in blobs
    assert payload == original
