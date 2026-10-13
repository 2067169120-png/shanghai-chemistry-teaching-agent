"""Actual book-to-crop-to-lesson/PPT handoff, with source and stale UI checks."""
from copy import deepcopy

import pytest
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QDialog
from pptx import Presentation

from test_personal_textbook_ui import (app as app, panel as panel, complete, enter_chapter)
from test_personal_textbook_crops import request
from test_personal_textbook_study import book_fixture as book_fixture
from test_desktop_studio_ui import window as window
from test_desktop_library_ui import _ManualBridge
from integrations.deeptutor_shchem_v1.desktop_personal_textbook_crops import book_page_selection
from integrations.deeptutor_shchem_v1.desktop_workbench.textbook_preparation_crop_dialog import TextbookPreparationCropDialog
from integrations.deeptutor_shchem_v1.desktop_textbook_preparation_crops import TextbookPreparationCropService


def test_graphic_page_entry_works_before_directory_and_narrow_controls_are_reachable(panel,app):
    value,_,f=panel;value.tabs.setCurrentIndex(value._graphics_tab);value.resize(420,800);app.processEvents()
    assert value.crop_page.maximum()==4 and value.crop_button.isEnabled()
    assert value.splitter.orientation()==Qt.Orientation.Vertical and value.width()==420
    value.crop_page.setValue(4)
    assert book_page_selection(value._book,value.crop_page.value())['visual_asset_id'].endswith('-4')
    scroll=value.tabs.widget(value._graphics_tab)
    scroll.ensureWidgetVisible(value.crop_button);app.processEvents()
    p=value.crop_button.mapTo(value,value.crop_button.rect().bottomRight())
    assert 0 <= p.x() < value.width() and 0 <= p.y() < value.height()
    assert not f.study.book_study(f.selection)['sections']


def test_graphic_entry_starts_at_chapter_page_and_locates_exact_whole_book_page(panel,app):
    value,tasks,_=panel;enter_chapter(value,tasks)
    assert value.crop_page.value()==2
    value.crop_page.setValue(4);value.crop_read_button.click();app.processEvents()
    complete(tasks,'读取教材原页');app.processEvents()
    assert value._reader.pages.currentData()==4 and value._reader.pages.count()==4


@pytest.mark.parametrize('result',['accepted','cancelled','changed-book','changed-page','closed'])
def test_personal_crop_dialog_handoff_rejects_stale_context(panel,monkeypatch,result):
    value,_,_=panel;seen=[]
    choice={'selection':book_page_selection(value._book,1),'expected_image_sha256':'a'*64}
    from integrations.deeptutor_shchem_v1.desktop_workbench import textbook_preparation_crop_dialog as module
    class TestDialog:
        selection_request=choice
        def __init__(self,facade,tasks,selection,parent,*,allow_composition_handoff=False):
            assert allow_composition_handoff; seen.append(selection)
        def exec(self):
            if result=='changed-book':value.deactivate()
            elif result=='changed-page':value.crop_page.setValue(2)
            elif result=='closed':value.close()
            return QDialog.DialogCode.Rejected if result=='cancelled' else QDialog.DialogCode.Accepted
        def deleteLater(self):pass
    monkeypatch.setattr(module,'TextbookPreparationCropDialog',TestDialog)
    value.image_requested.connect(lambda x:seen.append(x));value.crop_button.click()
    assert len(seen)==(2 if result=='accepted' else 1)
    assert seen[0]==choice['selection']


def test_native_personal_crop_canvas_previews_book_pixels(panel,app):
    value,_,_=panel
    bridge=_ManualBridge()
    dialog=TextbookPreparationCropDialog(value.facade,bridge,book_page_selection(value._book,2),value)
    dialog.show();bridge.succeed(bridge.take('读取教材框选原页'),run=True);app.processEvents()
    assert dialog._page['source_kind']=='personal_textbook_page' and dialog.canvas.scene().items()
    dialog.whole_page.click();dialog.purpose.setText('展示合成页色块')
    dialog.preview_button.click();bridge.succeed(bridge.take('核对教材裁片'),run=True);app.processEvents()
    assert dialog._preview and not dialog.confirm.isEnabled()
    dialog.checked.setChecked(True);dialog.confirm.click()
    assert dialog.result()==QDialog.DialogCode.Accepted and dialog.selection_request
    dialog.deleteLater();app.processEvents()


def test_main_window_personal_crop_signal_preserves_lesson_and_exports_same_image(window,book_fixture,monkeypatch):
    from integrations.deeptutor_shchem_v1.desktop_workbench.lesson_design_dialog import LessonDesignDialog
    from integrations.deeptutor_shchem_v1.desktop_lesson_output import export_design,checked_file
    win,app=window;f=book_fixture;prep=win.preparation_page;tasks=_ManualBridge();prep.tasks=tasks
    service=TextbookPreparationCropService(f.root,win.facade.paths.task_root,personal_books=f.workspace)
    monkeypatch.setattr(win.facade,'import_textbook_preparation_crop',lambda choice:service.confirm(**choice))
    choice,_,value=request(f)
    prep.topic.setText('新书图表课');prep.objective.setPlainText('比较图示边界')
    prep.materials.setPlainText('原有材料，条件2.0 g');before=deepcopy(prep._payload())
    win.textbook_page.personal.image_requested.emit(choice)
    assert win.stack.currentWidget() is prep and not prep.isEnabled()
    tasks.succeed(tasks.take('添加教材备课图表'),run=True);app.processEvents()
    after=prep._payload();asset=after.pop('image_assets')[0];assert after==before
    assert win.facade.preparation_image_bytes(asset)==value['image_bytes']
    design=LessonDesignDialog(prep);design.show();design.add_node();design.bind.click()
    design.locked.setChecked(False);design.student_material.setChecked(True)
    design.fields['student_task'].setPlainText('观察合成图表的边界。')
    design.fields['teacher_answer'].setPlainText('教师答案-ONLY');app.processEvents()
    payload=prep._payload();export=export_design(win.facade,payload);payload['lesson_design']['exports'].append(export)
    ppt=Presentation(checked_file(win.facade,payload['lesson_design'],export['id'],'lesson_presentation.pptx'))
    pictures=[shape for slide in ppt.slides for shape in slide.shapes if shape.shape_type==13]
    assert len(pictures)==1 and pictures[0].image.blob==value['image_bytes']
    text='\n'.join(shape.text for slide in ppt.slides for shape in slide.shapes if shape.has_text_frame)
    assert '教师答案-ONLY' not in text and 'PDF第2页' not in text
    assert 'PDF第2页' in asset['source']
    design.close();design.deleteLater();app.processEvents()
