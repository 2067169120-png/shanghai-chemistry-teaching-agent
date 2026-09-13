"""Whole-file previews are local, lazy, cancellable before commit, and source-bound."""

import csv
import hashlib
import io
from copy import deepcopy
from dataclasses import replace

import pytest
from docx import Document
from test_desktop_visual_import_facade import FakeProviderStore, _facade, _png
from test_desktop_visual_import_facade import (
    desktop_paths as desktop_paths,  # noqa: PLC0414
)

from integrations.deeptutor_shchem_v1 import desktop_import_preview as preview_module
from integrations.deeptutor_shchem_v1.desktop_facade import DesktopFacadeError
from integrations.deeptutor_shchem_v1.desktop_preparation_images import (
    PreparationImageStore,
)


def _word_bytes(*, picture=False, text="物质分类练习"):
    doc = Document()
    doc.add_paragraph("【即学即练1】" + text)
    if picture:
        doc.add_paragraph().add_run().add_picture(io.BytesIO(_png("blue")))
    doc.add_paragraph("【答案】教师参考答案")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def _files(tmp_path, *, picture=False):
    word = tmp_path / "教师讲义.docx"
    image = tmp_path / "题面.png"
    pdf = tmp_path / "补充资料.pdf"
    word.write_bytes(_word_bytes(picture=picture))
    image.write_bytes(_png("green"))
    pdf.write_bytes(b"%PDF-1.4\n% synthetic byte fixture\n%%EOF")
    return word, image, pdf


def _stored(root):
    return {
        str(path.relative_to(root)): path.read_bytes()
        for path in root.rglob("*")
        if path.is_file()
    }


def _source_args(preview, index=0):
    return (
        preview["preview_id"],
        preview["revision"],
        preview["sources"][index]["source_id"],
    )


def test_preview_sources_and_cancel_do_not_write_or_invoke_provider(
    desktop_paths, tmp_path
):
    word, image, pdf = _files(tmp_path, picture=True)
    provider = FakeProviderStore(configured=False)
    facade = _facade(desktop_paths, provider)
    before = _stored(desktop_paths.state_root)
    preview = facade.preview_import_files(
        question_files=(image,),
        answer_files=(pdf,),
        handout_files=(word,),
        source_type="教师自有资料",
    )
    assert [row["role"] for row in preview["sources"]] == [
        "question",
        "answer",
        "handout",
    ]
    assert [row["kind"] for row in preview["sources"]] == ["image", "pdf", "docx"]
    assert all(
        set(row)
        == {"source_id", "role", "kind", "source_name", "source_sha256", "order_index"}
        for row in preview["sources"]
    )
    assert str(tmp_path) not in str(preview)
    assert "完整内容及其全部候选题" in preview["warnings"][0]
    image_result = facade.preview_import_source(*_source_args(preview))
    pdf_result = facade.preview_import_source(*_source_args(preview, 1))
    assert image_result["bytes"] == image.read_bytes()
    assert pdf_result["bytes"] == pdf.read_bytes()
    assert image_result["mime_type"] == "image/png"
    assert pdf_result["mime_type"] == "application/pdf"
    word_result = facade.preview_import_source(*_source_args(preview, 2))
    assert word_result["kind"] == "docx"
    assert len(word_result["questions"]) == 1
    assert word_result["source_name"] == word.name
    assert word_result["source_sha256"] == hashlib.sha256(word.read_bytes()).hexdigest()
    asset = word_result["preview"]["assets"][0]
    assert facade.preview_import_asset(*_source_args(preview, 2), asset["asset_id"])[
        "bytes"
    ] == _png("blue")
    facade.discard_import_preview(preview["preview_id"])
    assert _stored(desktop_paths.state_root) == before
    assert provider.borrow_calls == 0
    with pytest.raises(DesktopFacadeError, match="失效"):
        facade.preview_import_source(*_source_args(preview))


def test_98_file_inventory_does_not_parse_all_word_files(
    desktop_paths, tmp_path, monkeypatch
):
    raw = _word_bytes()
    paths = []
    for number in range(98):
        path = tmp_path / f"解析讲义{number}.docx"
        path.write_bytes(raw)
        paths.append(path)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    calls = []
    original = preview_module.PreparationSourcesService.word_preview_bytes

    def parse(service, *args):
        calls.append(True)
        return original(service, *args)

    monkeypatch.setattr(
        preview_module.PreparationSourcesService, "word_preview_bytes", parse
    )
    preview = facade.preview_import_files(
        handout_files=tuple(paths), source_type="一轮解析版"
    )
    assert len(preview["sources"]) == 98 and not calls
    first = facade.preview_import_source(*_source_args(preview))
    first["questions"].clear()
    again = facade.preview_import_source(*_source_args(preview))
    assert len(calls) == 1 and len(again["questions"]) == 1
    assert not (desktop_paths.state_root / "desktop-state.v1.json").exists()


def test_sessions_are_bounded_and_old_session_expires(desktop_paths, tmp_path):
    word, _, _ = _files(tmp_path)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    previews = [
        facade.preview_import_files(handout_files=(word,), source_type="讲义")
        for _ in range(preview_module.MAX_PREVIEW_SESSIONS + 1)
    ]
    assert (
        len(facade._import_preview_service._sessions)
        == preview_module.MAX_PREVIEW_SESSIONS
    )
    with pytest.raises(DesktopFacadeError, match="失效"):
        facade.preview_import_source(*_source_args(previews[0]))
    assert facade.preview_import_source(*_source_args(previews[-1]))["questions"]


def test_busy_sessions_are_not_evicted_or_unbounded(desktop_paths, tmp_path):
    word, _, _ = _files(tmp_path)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    for _ in range(preview_module.MAX_PREVIEW_SESSIONS):
        preview = facade.preview_import_files(handout_files=(word,), source_type="讲义")
        facade._import_preview_service._sessions[preview["preview_id"]]["busy"] = True
    before = tuple(facade._import_preview_service._sessions)
    with pytest.raises(DesktopFacadeError, match="确认导入"):
        facade.preview_import_files(handout_files=(word,), source_type="讲义")
    assert tuple(facade._import_preview_service._sessions) == before


def test_word_cache_is_bounded(desktop_paths, tmp_path):
    raw = _word_bytes()
    paths = []
    for number in range(preview_module.MAX_CACHED_WORD_SOURCES + 1):
        path = tmp_path / f"讲义{number}.docx"
        path.write_bytes(raw)
        paths.append(path)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    preview = facade.preview_import_files(
        handout_files=tuple(paths), source_type="讲义"
    )
    for index in range(len(paths)):
        facade.preview_import_source(*_source_args(preview, index))
    cache = facade._import_preview_service._sessions[preview["preview_id"]]["cache"]
    assert len(cache) == preview_module.MAX_CACHED_WORD_SOURCES
    assert preview["sources"][0]["source_id"] not in cache


@pytest.mark.parametrize(
    "defect", ["revision", "source", "asset", "other_source_asset", "mutated_return"]
)
def test_preview_is_bound_to_its_session_and_source(desktop_paths, tmp_path, defect):
    word, image, _ = _files(tmp_path, picture=True)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    preview = facade.preview_import_files(
        question_files=(image,), handout_files=(word,), source_type="讲义"
    )
    preview_id, revision, source_id = _source_args(preview, 1)
    if defect == "revision":
        with pytest.raises(DesktopFacadeError, match="失效"):
            facade.preview_import_source(preview_id, "changed", source_id)
    elif defect == "source":
        with pytest.raises(DesktopFacadeError, match="不属于"):
            facade.preview_import_source(preview_id, revision, "other")
    elif defect == "asset":
        with pytest.raises(DesktopFacadeError, match="不属于"):
            facade.preview_import_asset(
                preview_id, revision, source_id, "word-b999-image1"
            )
    elif defect == "other_source_asset":
        with pytest.raises(DesktopFacadeError, match="Word预览"):
            facade.preview_import_asset(*_source_args(preview), "word-b2-image1")
    else:
        preview["sources"][1]["source_name"] = "替代名.docx"
        assert (
            facade.preview_import_source(preview_id, revision, source_id)["source_name"]
            == word.name
        )


def test_changed_file_is_detected_even_if_word_result_was_cached(
    desktop_paths, tmp_path
):
    word, _, _ = _files(tmp_path)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    preview = facade.preview_import_files(handout_files=(word,), source_type="讲义")
    facade.preview_import_source(*_source_args(preview))
    word.write_bytes(_word_bytes(text="已改变条件"))
    with pytest.raises(DesktopFacadeError, match="变化"):
        facade.preview_import_source(*_source_args(preview))
    with pytest.raises(DesktopFacadeError, match="变化"):
        facade.commit_import_preview(
            *_source_args(preview)[:2], [preview["sources"][0]["source_id"]]
        )
    assert not (desktop_paths.state_root / "desktop-state.v1.json").exists()


def test_commit_imports_only_checked_whole_files_in_original_role_order(
    desktop_paths, tmp_path
):
    word, image, pdf = _files(tmp_path)
    provider = FakeProviderStore(configured=False)
    facade = _facade(desktop_paths, provider)
    preview = facade.preview_import_files(
        question_files=(image, pdf), handout_files=(word,), source_type="勾选导入"
    )
    selected = [preview["sources"][2]["source_id"], preview["sources"][1]["source_id"]]
    result = facade.commit_import_preview(
        preview["preview_id"], preview["revision"], selected
    )
    assert [(row.role, row.order_index, row.filename) for row in result.sources] == [
        ("question", 1, pdf.name),
        ("handout", 1, word.name),
    ]
    assert result.source_count == 2
    archived = desktop_paths.state_root / "visual-import-v2" / "sources"
    assert sorted(path.read_bytes() for path in archived.iterdir()) == sorted(
        [pdf.read_bytes(), word.read_bytes()]
    )
    assert len(facade.word_question_catalog()["items"]) == 1
    assert provider.borrow_calls == 0
    with pytest.raises(DesktopFacadeError, match="失效"):
        facade.preview_import_source(*_source_args(preview))


def test_repeat_confirm_keeps_existing_batch_semantics(desktop_paths, tmp_path):
    word, _, _ = _files(tmp_path)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    results = []
    for _ in range(2):
        preview = facade.preview_import_files(handout_files=(word,), source_type="讲义")
        results.append(
            facade.commit_import_preview(
                preview["preview_id"],
                preview["revision"],
                [preview["sources"][0]["source_id"]],
            )
        )
    assert results[0].batch_id == results[1].batch_id
    assert len(facade.list_imported_word_batches()) == 1


@pytest.mark.parametrize("selection", [[], ["other"], ["same", "same"], "same", None])
def test_invalid_selection_has_no_import_record(desktop_paths, tmp_path, selection):
    word, _, _ = _files(tmp_path)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    preview = facade.preview_import_files(handout_files=(word,), source_type="讲义")
    with pytest.raises(DesktopFacadeError):
        facade.commit_import_preview(
            preview["preview_id"], preview["revision"], selection
        )
    assert not (desktop_paths.state_root / "desktop-state.v1.json").exists()


def test_answer_only_selection_is_rejected_before_saving(desktop_paths, tmp_path):
    word, image, _ = _files(tmp_path)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    preview = facade.preview_import_files(
        question_files=(image,), answer_files=(word,), source_type="答案"
    )
    with pytest.raises(DesktopFacadeError, match="不能只导入答案"):
        facade.commit_import_preview(
            preview["preview_id"],
            preview["revision"],
            [preview["sources"][1]["source_id"]],
        )
    assert not (desktop_paths.state_root / "desktop-state.v1.json").exists()


@pytest.mark.parametrize("change", ["bytes", "role", "order", "name"])
def test_saver_verifies_final_loaded_snapshot_before_first_write(
    desktop_paths, tmp_path, monkeypatch, change
):
    word, _, _ = _files(tmp_path)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    preview = facade.preview_import_files(handout_files=(word,), source_type="讲义")
    original = facade._build_visual_import_sources

    def changed(**kwargs):
        loaded = original(**kwargs)
        updates = {
            "bytes": {"content": _word_bytes(text="替换内容")},
            "role": {"role": "answer"},
            "order": {"order_index": 2},
            "name": {"filename": "替换.docx"},
        }
        return (replace(loaded[0], **updates[change]),)

    def no_archive(**kwargs):
        pytest.fail("archive coordinator must not be created for changed inputs")

    monkeypatch.setattr(facade, "_build_visual_import_sources", changed)
    monkeypatch.setattr(facade, "_visual_import_coordinator", no_archive)
    with pytest.raises(DesktopFacadeError, match="发生变化"):
        facade.commit_import_preview(
            preview["preview_id"],
            preview["revision"],
            [preview["sources"][0]["source_id"]],
        )
    assert not (desktop_paths.state_root / "desktop-state.v1.json").exists()


def test_files_changed_after_loading_cannot_replace_snapshot_bytes(
    desktop_paths, tmp_path, monkeypatch
):
    word, _, _ = _files(tmp_path)
    raw = word.read_bytes()
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    preview = facade.preview_import_files(handout_files=(word,), source_type="讲义")
    original = facade._build_visual_import_sources

    def changed_after_loading(**kwargs):
        loaded = original(**kwargs)
        word.write_bytes(_word_bytes(text="加载之后替换"))
        return loaded

    monkeypatch.setattr(facade, "_build_visual_import_sources", changed_after_loading)
    facade.commit_import_preview(
        preview["preview_id"], preview["revision"], [preview["sources"][0]["source_id"]]
    )
    archive = desktop_paths.state_root / "visual-import-v2" / "sources"
    assert next(archive.iterdir()).read_bytes() == raw


@pytest.mark.parametrize("when", ["before_loading", "after_planning"])
def test_cancellation_before_persistence_has_no_records(desktop_paths, tmp_path, when):
    word, _, _ = _files(tmp_path)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    preview = facade.preview_import_files(handout_files=(word,), source_type="讲义")
    before = _stored(desktop_paths.state_root)
    calls = []

    def cancel():
        calls.append(True)
        return when == "before_loading" or len(calls) == 2

    with pytest.raises(DesktopFacadeError, match="取消"):
        facade.commit_import_preview(
            preview["preview_id"],
            preview["revision"],
            [preview["sources"][0]["source_id"]],
            should_cancel=cancel,
        )
    assert _stored(desktop_paths.state_root) == before
    assert len(calls) == (1 if when == "before_loading" else 2)


def test_persistence_finishes_if_cancel_is_requested_after_final_gate(
    desktop_paths, tmp_path
):
    word, _, _ = _files(tmp_path)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    preview = facade.preview_import_files(handout_files=(word,), source_type="讲义")
    calls, state = [], {"cancel": False}

    def cancel():
        calls.append(True)
        return state["cancel"]

    def progress(_event):
        state["cancel"] = True

    receipt = facade.commit_import_preview(
        preview["preview_id"],
        preview["revision"],
        [preview["sources"][0]["source_id"]],
        should_cancel=cancel,
        progress_callback=progress,
    )
    assert state["cancel"] and len(calls) == 2
    assert receipt.source_count == 1
    assert facade.list_imported_word_batches()[0].batch_id == receipt.batch_id


@pytest.mark.parametrize(
    "filename,raw",
    [
        ("wrong.docx", b"not zip"),
        ("wrong.png", b"not png"),
        ("wrong.pdf", b"not pdf"),
        ("empty.docx", b""),
        ("unsupported.txt", b"text"),
    ],
)
def test_invalid_files_fail_before_any_session_or_record(
    desktop_paths, tmp_path, filename, raw
):
    path = tmp_path / filename
    path.write_bytes(raw)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    with pytest.raises(DesktopFacadeError):
        facade.preview_import_files(handout_files=(path,), source_type="讲义")
    assert not facade._import_preview_service._sessions
    assert not (desktop_paths.state_root / "desktop-state.v1.json").exists()


@pytest.mark.parametrize(
    "defect",
    ["empty", "too_many", "file_size", "batch_size", "source_type", "string_files"],
)
def test_initial_metadata_limits_are_checked(
    desktop_paths, tmp_path, monkeypatch, defect
):
    word, _, _ = _files(tmp_path)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    files, source_type = (word,), "讲义"
    if defect == "empty":
        files = ()
    elif defect == "too_many":
        files = (word,) * (preview_module.MAX_FILES + 1)
    elif defect == "file_size":
        monkeypatch.setattr(preview_module, "MAX_WORD_BYTES", 10)
    elif defect == "batch_size":
        monkeypatch.setattr(preview_module, "MAX_BATCH_BYTES", 10)
    elif defect == "source_type":
        source_type = "讲义\n非正常控制符"
    else:
        files = str(word)
    with pytest.raises(DesktopFacadeError):
        facade.preview_import_files(handout_files=files, source_type=source_type)
    assert not facade._import_preview_service._sessions
    assert not (desktop_paths.state_root / "desktop-state.v1.json").exists()


def _pack_inventory(desktop_paths, *, count=2):
    root = desktop_paths.shchem_root / preview_module.TEACHING_PACK
    expanded = root / "expanded" / "PKG-001"
    expanded.mkdir(parents=True)
    rows, expected = [], []
    for number in range(count):
        for role in ("原卷版", "解析版"):
            path = expanded / f"教师讲义{number}（{role}）.docx"
            raw = _word_bytes()
            path.write_bytes(raw)
            rows.append(
                {
                    "document_role": role,
                    "output_relative_path": path.relative_to(root).as_posix(),
                    "bytes": str(len(raw)),
                    "sha256": hashlib.sha256(raw).hexdigest(),
                }
            )
            if role == "解析版":
                expected.append(str(path))
    inventory = root / "archive_inventory.csv"
    _write_inventory(inventory, rows)
    return inventory, rows, expected


def _write_inventory(path, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=("document_role", "output_relative_path", "bytes", "sha256"),
        )
        writer.writeheader()
        writer.writerows(rows)


def test_teaching_pack_uses_only_recorded_analysis_files_without_state_writes(
    desktop_paths,
):
    _, _, expected = _pack_inventory(desktop_paths, count=49)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    before = _stored(desktop_paths.state_root)
    # There are 98 files in this fixture, but only its 49 recorded analysis files
    # are selected. The real inventory has 98 analysis + 98 original files.
    result = facade.teaching_pack_analysis_files()
    assert result == tuple(expected)
    assert _stored(desktop_paths.state_root) == before


@pytest.mark.parametrize(
    "defect",
    ["changed_bytes", "wrong_size", "outside", "absolute", "duplicate", "missing"],
)
def test_teaching_pack_rejects_changed_or_outside_inventory(desktop_paths, defect):
    inventory, rows, expected = _pack_inventory(desktop_paths)
    if defect == "changed_bytes":
        from pathlib import Path

        Path(expected[0]).write_bytes(_word_bytes(text="已变化"))
    elif defect == "wrong_size":
        rows[1]["bytes"] = "1"
    elif defect == "outside":
        rows[1]["output_relative_path"] = "../../outside.docx"
    elif defect == "absolute":
        rows[1]["output_relative_path"] = expected[0]
    elif defect == "duplicate":
        rows.append(deepcopy(rows[1]))
    else:
        inventory.unlink()
    if defect not in {"changed_bytes", "missing"}:
        _write_inventory(inventory, rows)
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    with pytest.raises(DesktopFacadeError, match="已有讲义清单"):
        facade.teaching_pack_analysis_files()
    assert not (desktop_paths.state_root / "desktop-state.v1.json").exists()


def test_preparation_image_bytes_reads_local_pixels_and_rejects_changed_content(
    desktop_paths,
):
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    store = PreparationImageStore(desktop_paths.task_root / "preparation-v1" / "images")
    raw = _png("blue")
    asset = store.import_bytes(raw, "本地原图", "教师资料", "课堂使用")
    assert facade.preparation_image_bytes(asset) == raw
    changed = dict(asset, width=asset["width"] + 1)
    with pytest.raises(DesktopFacadeError, match="尺寸或格式"):
        facade.preparation_image_bytes(changed)
    (store.root / (asset["sha256"] + ".image")).write_bytes(_png("red"))
    with pytest.raises(DesktopFacadeError, match="内容已经变化"):
        facade.preparation_image_bytes(asset)


def test_preparation_image_bytes_validates_metadata_before_store_creation(
    desktop_paths,
):
    facade = _facade(desktop_paths, FakeProviderStore(configured=False))
    with pytest.raises(DesktopFacadeError, match="不完整"):
        facade.preparation_image_bytes({"path": "unbound"})
    assert not (desktop_paths.task_root / "preparation-v1" / "images").exists()
