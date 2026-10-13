"""Real import/source/teacher-note UI and handoff, including narrow layouts."""
from copy import deepcopy

import pytest
from PySide6.QtCore import Qt, QPointF, QRect
from PySide6.QtWidgets import QApplication, QMessageBox

from test_personal_textbook_study import book_fixture as book_fixture, pdf
from test_desktop_library_ui import _ManualBridge
from test_textbook_bookshelf_ui import painted
from test_desktop_studio_ui import window as window, settle
from integrations.deeptutor_shchem_v1.desktop_facade import DesktopWorkbenchFacade
from integrations.deeptutor_shchem_v1.desktop_workbench.personal_textbook_panel import PersonalTextbookPanel
from integrations.deeptutor_shchem_v1.desktop_workbench.textbook_bookshelf_page import TextbookBookshelfPage
from integrations.deeptutor_shchem_v1.desktop_workbench.textbook_import_dialog import TextbookImportDialog
from integrations.deeptutor_shchem_v1.desktop_workbench.tasks import DesktopTaskBridge


class Event:
    def __init__(self):self.callbacks = []
    def connect(self, callback):self.callbacks.append(callback)
    def emit(self, *args):
        for callback in self.callbacks:callback(*args)


class Bridge(_ManualBridge):
    def __init__(self):
        super().__init__()
        self.task_finished, self.task_cancelled = Event(), Event()
    def submit_progress(self, label, operation, *, on_progress, on_success, on_failure):
        return self.submit(label, lambda: operation(on_progress, lambda: False), on_success=on_success, on_failure=on_failure)
    def cancel(self, task_id):
        super().cancel(task_id);return True


def complete(bridge, label, *, value=None, run=True):
    task = bridge.take(label)
    bridge.succeed(task, value, run=run)
    bridge.task_finished.emit(task.task_id)
    return task


@pytest.fixture
def app():
    return QApplication.instance() or QApplication([])


def facade_for(f):
    facade = DesktopWorkbenchFacade.__new__(DesktopWorkbenchFacade)
    facade.paths, facade._state = f.paths, f.state
    return facade


@pytest.fixture
def panel(app, book_fixture):
    f = book_fixture;bridge = Bridge()
    value = PersonalTextbookPanel(facade_for(f), bridge)
    value.resize(1000, 760);value.show()
    value.set_book(f.study.books()[0]);complete(bridge, '读取个人教材章节与笔记')
    app.processEvents()
    yield value, bridge, f
    value.close();value.deleteLater();app.processEvents()


def enter_chapter(value, bridge):
    value.section_title.setText('合成章一');value.start.setValue(2);value.end.setValue(3)
    value.range_confirmed.setChecked(True);value.save_section.click()
    complete(bridge, '保存个人教材章节')


def enter_note(value, bridge):
    value.new_note.click()
    value.note_title.setText('合成笔记');value.note_text.setPlainText('A → B（条件X）；用量 2.0 g。')
    value.note_page.setValue(2);value.note_confirmed.setChecked(True);value.save_note.click()
    complete(bridge, '保存个人教材笔记')


def test_native_chapter_and_note_save_keep_unknowns_and_require_each_confirmation(panel):
    value, bridge, f = panel
    assert value._reader is None and not value.tree.topLevelItemCount() and not value.save_section.isEnabled()
    value.section_title.setText('章');value.range_confirmed.setChecked(True)
    value.end.setValue(3)
    assert not value.range_confirmed.isChecked() and not value.save_section.isEnabled()
    enter_chapter(value, bridge)
    assert value.tree.topLevelItemCount() == 1 and not value._checked
    enter_note(value, bridge)
    assert value.note_list.count() == 1 and not value._checked and not value.preview_button.isEnabled()
    view = f.study.book_study(f.selection)
    assert view['sections'][0]['printed_pages'] is None and view['notes'][0]['human_reviewed'] is False
    assert f.source.read_bytes() == f.original and f.workspace.books()[0]['human_reviewed'] is False


def test_read_book_and_teacher_section_paint_original_pdf_and_clamp_scope(panel, app):
    value, bridge, f = panel
    value.read_book_button.click();app.processEvents()
    complete(bridge, '读取教材原页');app.processEvents()
    assert value._reader._loaded_pages and value._reader.pages.count() == 4
    assert painted(app, value._reader) > 10
    enter_chapter(value, bridge)
    value.read_section_button.click();app.processEvents();complete(bridge, '读取教材原页');app.processEvents()
    reader = value._reader
    assert [reader.pages.itemData(i) for i in range(reader.pages.count())] == [2, 3]
    assert '印刷页码：待核对' in reader.summary.toPlainText() and painted(app, reader) > 10
    reader.view.pageNavigator().jump(0, QPointF());app.processEvents()
    assert reader.view.pageNavigator().currentPage() == 1
    reader._jump(3);assert reader.view.pageNavigator().currentPage() == 1
    reader.next.click();app.processEvents();assert reader.view.pageNavigator().currentPage() == 2
    assert not reader.excerpt_button.isVisible() and f.source.read_bytes() == f.original


def test_reference_must_preview_then_fresh_check_and_can_retry_rejected_handoff(panel):
    value, bridge, f = panel;enter_chapter(value, bridge);enter_note(value, bridge)
    emitted = [];value.reference_requested.connect(emitted.append)
    value.note_list.item(0).setCheckState(Qt.CheckState.Checked)
    assert not value.confirm_button.isEnabled()
    value.preview_button.click();complete(bridge, '核对个人教材备课参考')
    assert value.confirm_button.isEnabled() and 'A → B' in value.reference_preview.toPlainText()
    value.confirm_button.click();complete(bridge, '核对个人教材备课参考')
    assert len(emitted) == 1 and not f.state.snapshot()['textbook_workspace']['books'][0]['human_reviewed']
    value.reference_rejected();assert value.confirm_button.isEnabled()
    value.reference_appended();value._controls();assert not value.confirm_button.isEnabled()


def test_edits_and_selection_invalidate_reference_and_existing_note_revision(panel):
    value, bridge, f = panel;enter_chapter(value, bridge);enter_note(value, bridge)
    value.note_list.item(0).setCheckState(Qt.CheckState.Checked)
    value.preview_button.click();complete(bridge, '核对个人教材备课参考')
    value.note_text.setPlainText('尚未保存的修改')
    assert not value.confirm_button.isEnabled() and value._reference is None
    value.note_list.setCurrentRow(0);value.note_text.setPlainText('改变原笔记条件 Y。')
    value.note_confirmed.setChecked(True);value.save_note.click();complete(bridge, '保存个人教材笔记')
    assert value._checked == {} and '改变原笔记' in f.study.book_study(f.selection)['notes'][0]['text']


def test_cancelled_and_closed_callbacks_cannot_restore_a_reference(panel):
    value, bridge, f = panel;enter_chapter(value, bridge);enter_note(value, bridge)
    value.note_list.item(0).setCheckState(Qt.CheckState.Checked)
    value.preview_button.click();task = bridge.take('核对个人教材备课参考')
    bridge.task_cancelled.emit(task.task_id, task.label);bridge.task_finished.emit(task.task_id)
    bridge.succeed(task, run=True)
    assert not value._busy and value._reference is None
    value.preview_button.click();task = bridge.take('核对个人教材备课参考');value.close()
    bridge.succeed(task, run=True)
    assert task.task_id in bridge.cancelled and value._reference is None


def test_source_change_after_preview_cannot_enter_preparation(panel):
    value, bridge, f = panel;enter_chapter(value, bridge);enter_note(value, bridge)
    value.note_list.item(0).setCheckState(Qt.CheckState.Checked)
    value.preview_button.click();complete(bridge, '核对个人教材备课参考')
    archive = f.workspace.root / (f.selection['revision'] + '.pdf');archive.write_bytes(b'changed')
    value.confirm_button.click();task = bridge.take('核对个人教材备课参考')
    emitted = [];value.reference_requested.connect(emitted.append)
    try:task.operation()
    except Exception as exc:bridge.fail(task, str(exc))
    assert emitted == [] and value._reference is None and not value.confirm_button.isEnabled()


def test_whole_book_import_dialog_uses_original_contract_and_recovers_stop(app, book_fixture):
    f = book_fixture;bridge = Bridge();dialog = TextbookImportDialog(facade_for(f), bridge)
    dialog.files.file_list._append_paths([str(f.source)]);dialog._preview();complete(bridge, '预览整本教材')
    assert dialog.preview['sources'][0]['duplicate'] and dialog.save_button.isEnabled()
    dialog.save_button.click();complete(bridge, '导入整本教材')
    assert dialog.imported_source_ids == [f.selection['source_id']] and dialog.preview is None
    assert not dialog._busy and not dialog.save_button.isEnabled()
    dialog._preview();task = bridge.take('预览整本教材')
    dialog._stop();bridge.task_cancelled.emit(task.task_id, task.label);bridge.task_finished.emit(task.task_id)
    assert not dialog._busy and dialog.preview_button.isEnabled()
    dialog.done(0);bridge.succeed(task, run=True)
    assert dialog.preview is None


def test_shelf_shows_personal_books_without_any_official_directory(app, book_fixture):
    f = book_fixture;bridge = Bridge();facade = facade_for(f)
    value = TextbookBookshelfPage(facade, bridge);value.resize(1000, 820);value.show()
    value.refresh();complete(bridge, '读取教材书架');complete(bridge, '读取个人教材章节与笔记')
    app.processEvents()
    assert value.book.count() == 1 and value.book.currentData() == f.selection['source_id']
    assert value.personal.isVisible() and not value.splitter.isVisible() and value.import_button.isEnabled()
    assert value._snapshot['volumes'] == [] and len(value._snapshot['local_books']) == 1
    assert f.state.snapshot()['textbook_workspace']['books'][0]['human_reviewed'] is False
    value.close();value.deleteLater();app.processEvents()


@pytest.mark.parametrize('width', [420, 1100])
def test_personal_shelf_keeps_fixed_actions_in_bounds_at_small_width(app, book_fixture, width):
    from integrations.deeptutor_shchem_v1.desktop_workbench.studio_style import WORKBENCH_STYLE
    from integrations.deeptutor_shchem_v1.desktop_workbench.typography import install_ui_font
    f = book_fixture;bridge = Bridge();value = TextbookBookshelfPage(facade_for(f), bridge)
    install_ui_font(app);value.setStyleSheet(WORKBENCH_STYLE)
    value.resize(width, 920);value.show();value.refresh()
    complete(bridge, '读取教材书架');complete(bridge, '读取个人教材章节与笔记');app.processEvents()
    assert value.width() == width
    assert value.personal.splitter.orientation() == (Qt.Orientation.Vertical if width < 720 else Qt.Orientation.Horizontal)
    for control in (value.book, value.import_button, value.personal.read_book_button,
                    value.personal.preview_button, value.personal.confirm_button, value.personal.status):
        assert control.isVisible() and value.rect().contains(QRect(control.mapTo(value, control.rect().topLeft()), control.size()))
    value.close();value.deleteLater();app.processEvents()


def test_real_worker_reads_and_saves_teacher_notes(app, book_fixture):
    f = book_fixture;bridge = DesktopTaskBridge()
    value = PersonalTextbookPanel(facade_for(f), bridge);value.show();value.set_book(f.study.books()[0])
    settle(app, lambda: not value._busy)
    value.section_title.setText('真实后台任务的合成章节');value.start.setValue(2);value.end.setValue(3)
    value.range_confirmed.setChecked(True);value.save_section.click();settle(app, lambda: not value._busy)
    assert value.tree.topLevelItemCount() == 1
    value.note_title.setText('合成笔记');value.note_text.setPlainText('合成流程笔记 1.0 g。')
    value.note_confirmed.setChecked(True);value.save_note.click();settle(app, lambda: not value._busy)
    assert len(f.study.book_study(f.selection)['notes']) == 1
    value.close();assert bridge.shutdown(5000);value.deleteLater();app.processEvents()


def test_actual_main_window_handoff_preserves_form_and_saves_note_snapshot(window, tmp_path):
    win, app = window;source = pdf(tmp_path / '合成导入教材.pdf')
    workspace = win.facade.textbook_workspace();workspace.commit_books(workspace.preview_books([source]))
    study = win.facade.personal_textbook_study();book = study.books()[0]
    selection = {key: book[key] for key in ('source_id', 'revision')}
    view = study.save_section(selection, study.book_study(selection)['study_revision'],
        title='合成章', pdf_start=2, pdf_end=3, range_confirmed=True)
    view = study.save_note(selection, view['study_revision'], section_id=view['sections'][0]['section_id'],
        title='合成备课笔记', text='A → B（条件X）；用量 2.0 g。', pdf_page=2, teacher_confirmed=True)
    page = win.preparation_page
    page.topic.setText('保留当前课题');page.audience.setText('教师自定对象')
    page.objective.setPlainText('保留教学目标');page.materials.setPlainText('原有参考')
    mode, images = page._payload().get('image_input_mode', 'local_only'), deepcopy(page._payload().get('image_assets', []))
    win.navigate('textbooks');settle(app, lambda: not win.textbook_page._loading)
    win.textbook_page.book.setCurrentIndex(win.textbook_page.book.findData(book['source_id']))
    personal = win.textbook_page.personal;settle(app, lambda: not personal._busy)
    personal.note_list.item(0).setCheckState(Qt.CheckState.Checked)
    personal.preview_button.click();settle(app, lambda: not personal._busy)
    personal.confirm_button.click();settle(app, lambda: not personal._busy)
    assert page.topic.text() == '保留当前课题' and page.objective.toPlainText() == '保留教学目标'
    assert page.materials.toPlainText().startswith('原有参考') and 'A → B' in page.materials.toPlainText()
    assert page._payload().get('image_input_mode', 'local_only') == mode and page._payload().get('image_assets', []) == images
    page._save();settle(app, lambda: not page._save_task_id)
    assert any('A → B' in win.facade.load_preparation_draft(row['draft_id'], row['revision'])['payload']['materials'] for row in win.facade.preparation_draft_options())
    from integrations.deeptutor_shchem_v1.desktop_workbench.lesson_design_dialog import LessonDesignDialog
    from integrations.deeptutor_shchem_v1.desktop_lesson_output import checked_file, FILES
    from pptx import Presentation
    from docx import Document
    dialog = LessonDesignDialog(page);dialog.show();settle(app)
    dialog.add_node();dialog.bind_source()
    assert 'A → B' in dialog.history.value['nodes'][0]['material_text']
    dialog.locked.setChecked(False);dialog.student_material.setChecked(True)
    # The teacher explicitly chooses the classroom text; the full source stays
    # in the preparation draft rather than being silently cut into a slide.
    dialog.fields['material_text'].setPlainText('合成备课笔记\nA → B（条件X）；用量 2.0 g。\n合成导入教材.pdf · PDF第2页')
    dialog.fields['teacher_answer'].setPlainText('仅教师答案哨兵-合成材料')
    dialog.generate.click();settle(app, lambda: not dialog.busy)
    assert dialog.history.value['exports'], dialog.status.text()
    identity = dialog.history.value['exports'][-1]['id']
    files = {name: checked_file(win.facade, dialog.history.value, identity, name) for name in FILES}
    presentation = Presentation(files['lesson_presentation.pptx'])
    body = '\n'.join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)
    notes = '\n'.join(slide.notes_slide.notes_text_frame.text for slide in presentation.slides)
    def doc_text(path):
        doc = Document(path)
        return '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    teacher, student = doc_text(files['lesson_plan.docx']), doc_text(files['student_worksheet.docx'])
    assert all('A → B' in text for text in (body, teacher, student))
    assert '仅教师答案哨兵-合成材料' in teacher and '仅教师答案哨兵-合成材料' in notes
    assert '仅教师答案哨兵-合成材料' not in body and '仅教师答案哨兵-合成材料' not in student
    dialog.reject();dialog.deleteLater();settle(app)


def test_busy_lesson_rejects_textbook_reference_without_form_change(window, monkeypatch):
    win, _app = window;page = win.preparation_page;before = deepcopy(page._payload())
    page._lesson_busy = True
    monkeypatch.setattr(QMessageBox, 'information', lambda *args: None)
    assert not page.import_word_reference({'materials': '合成教材参考', 'warnings': [], 'source_kind': 'textbook'})
    assert page._payload() == before
    page._lesson_busy = False


def test_native_backup_includes_original_and_teacher_notes(panel, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog
    from zipfile import ZipFile
    value, bridge, f = panel;enter_chapter(value, bridge);enter_note(value, bridge)
    destination = tmp_path / 'backup.zip'
    monkeypatch.setattr(QFileDialog, 'getSaveFileName', lambda *args: (str(destination), ''))
    value.backup_button.click();complete(bridge, '备份个人教材与笔记')
    with ZipFile(destination) as archive:
        manifest = archive.read('manifest.json').decode()
        assert 'personal_study' in manifest
        assert archive.read('textbook-sources/' + f.selection['revision'] + '.pdf') == f.original
    before = destination.read_bytes()
    value.backup_button.click();task = bridge.take('备份个人教材与笔记')
    try:task.operation()
    except Exception as exc:bridge.fail(task, str(exc))
    assert destination.read_bytes() == before and '未覆盖' in value.status.text()


def test_generic_intake_has_whole_book_path_and_prefills_only_question_pdfs(window, tmp_path, monkeypatch):
    from integrations.deeptutor_shchem_v1.desktop_workbench.dialogs import ImportDialog
    win, app = window;source = pdf(tmp_path / '本书.pdf')
    from PySide6.QtWidgets import QDialog
    parent = ImportDialog(win.facade, win.tasks, win)
    parent.source_type.setCurrentText('教材与课程文件')
    parent.question_files.file_list._append_paths([str(source)])
    observed = []
    def execute(dialog):
        observed.extend(dialog.files.paths())
        workspace = win.facade.textbook_workspace()
        dialog._saved(workspace.commit_books(workspace.preview_books(dialog.files.paths())))
        return QDialog.DialogCode.Rejected
    monkeypatch.setattr(TextbookImportDialog, 'exec', execute)
    parent.textbook_import_button.click();app.processEvents()
    assert observed == [str(source)] and parent.open_textbooks_requested
    assert win.facade.textbook_workspace().books()[0]['source_name'] == '本书.pdf'
    assert parent.question_files.paths() == [str(source)]
    parent.deleteLater();settle(app)
