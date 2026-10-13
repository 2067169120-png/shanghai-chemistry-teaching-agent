"""Real PDF pixels, local source receipts, and real lesson output files."""
from copy import deepcopy
import hashlib
import io
import json
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from PIL import Image

from test_textbook_section_reader import section_fixture as section_fixture
from test_textbook_asset_catalog import asset_fixture as asset_fixture, selection
from test_textbook_bookshelf import shelf_fixture as shelf_fixture, disk_bytes, replace_book
from test_lesson_design import sample
from integrations.deeptutor_shchem_v1 import desktop_textbook_preparation_crops as crops
from integrations.deeptutor_shchem_v1.desktop_facade import DesktopWorkbenchFacade, DesktopFacadeError
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
from integrations.deeptutor_shchem_v1.desktop_lesson_output import export_design, checked_file
from integrations.deeptutor_shchem_v1.reader_cancellation import TaskCancellationEvent, ReadCancelled, read_cancel_scope


@pytest.fixture
def crop_fixture(shelf_fixture, tmp_path):
    f = shelf_fixture
    f.task_root = tmp_path.parent / (tmp_path.name+'-crop-state')
    f.crops = crops.TextbookPreparationCropService(f.root, f.task_root)
    f.selection = selection(f)
    f.page = f.crops.page(f.selection)
    f.bounds = {'left':100,'top':100,'right':610,'bottom':630}
    f.facade = DesktopWorkbenchFacade.__new__(DesktopWorkbenchFacade)
    f.facade.paths = SimpleNamespace(workspace_root=f.root, task_root=f.task_root)
    return f


def preview(f):
    return f.crops.preview(f.selection,f.page['page_revision'],f.bounds,'合成图表','观察合成色块与边缘')


def request(f, value=None):
    value = value or preview(f)
    return {'selection':deepcopy(f.selection),'page_revision':f.page['page_revision'],
            'bounds':deepcopy(f.bounds),'caption':value['image_asset']['caption'],
            'purpose':value['image_asset']['purpose'],'expected_preview_revision':value['preview_revision'],
            'expected_image_sha256':value['image_asset']['sha256'],'selection_confirmed':True}


def test_preview_is_read_only_and_exact_pixels_then_confirmation_has_separate_receipt(crop_fixture):
    f = crop_fixture
    before = disk_bytes(f.root)
    value = preview(f)
    assert not f.task_root.exists() and disk_bytes(f.root) == before
    assert f.page['width']*f.page['height'] <= crops.MAX_PAGE_PIXELS
    assert f.page['effective_dpi'] <= crops.PAGE_DPI
    with Image.open(io.BytesIO(f.page['page_data'])) as original, Image.open(io.BytesIO(value['image_bytes'])) as cut:
        assert cut.size == (510,530)
        assert cut.tobytes() == original.crop((100,100,610,630)).tobytes()
        assert cut.getpixel((0,0)) != cut.getpixel((400,400))  # Page margin and vector drawing.
    asset = f.facade.import_textbook_preparation_crop(request(f,value))
    assert asset == value['image_asset']
    assert set(asset) == {'asset_id','sha256','caption','source','purpose','width','height','content_type'}
    store = PreparationImageStore(f.task_root/'preparation-v1'/'images')
    assert store.load(asset) == value['image_bytes']
    receipts = list((f.task_root/'preparation-v1'/'textbook-crops').glob('*.json'))
    assert len(receipts) == 1
    record = json.loads(receipts[0].read_text(encoding='utf-8'))
    assert record['source_sha256'] == hashlib.sha256(f.book.read_bytes()).hexdigest()
    assert record['pdf_page'] == 9 and record['printed_page'] == 5
    assert record['bounds'] == f.bounds and record['recipe'] == f.page['recipe']
    assert record['candidate_only'] is True
    assert all(record[key] is False for key in ('human_reviewed','retrieval_ready','generation_allowed','publication_allowed'))
    assert record['crop_id'] in asset['source'] and 'PDF第9页' in asset['source']
    assert str(f.root) not in json.dumps(record) and str(f.task_root) not in json.dumps(asset)
    saved = disk_bytes(f.task_root)
    assert f.facade.import_textbook_preparation_crop(request(f,value)) == asset
    assert disk_bytes(f.task_root) == saved and disk_bytes(f.root) == before


@pytest.mark.parametrize('bounds',[
    None, {}, {'left':True,'top':0,'right':20,'bottom':20},
    {'left':0.0,'top':0,'right':20,'bottom':20},
    {'left':-1,'top':0,'right':20,'bottom':20},
    {'left':20,'top':0,'right':20,'bottom':20},
    {'left':0,'top':0,'right':7000,'bottom':20},
    {'left':0,'top':0,'right':20,'bottom':20,'x':0},
])
def test_invalid_bounds_do_not_create_state(crop_fixture,bounds):
    f=crop_fixture
    with pytest.raises(crops.TextbookPreparationCropError):
        f.crops.preview(f.selection,f.page['page_revision'],bounds,'合成图','合成用途')
    assert not f.task_root.exists()


@pytest.mark.parametrize('field,value',[
    ('selection_confirmed',False), ('selection_confirmed',1),
    ('expected_preview_revision','a'*64), ('expected_image_sha256','a'*64),
    ('caption','修改后的图题'), ('purpose','修改后的用途'),
    ('page_revision','a'*64),
])
def test_changed_preview_or_missing_confirmation_cannot_write(crop_fixture,field,value):
    f=crop_fixture; choice=request(f); choice[field]=value
    with pytest.raises(DesktopFacadeError):f.facade.import_textbook_preparation_crop(choice)
    assert not f.task_root.exists()


@pytest.mark.parametrize('change',['pdf','catalog'])
def test_original_pdf_or_catalog_change_invalidates_confirmation(crop_fixture,change):
    f=crop_fixture; choice=request(f)
    if change=='pdf':f.book.write_bytes(f.book.read_bytes()+b'\n% changed original\n')
    else:
        f.asset['description']='changed source metadata'; f.write_assets([f.asset])
    with pytest.raises(DesktopFacadeError):f.facade.import_textbook_preparation_crop(choice)
    assert not f.task_root.exists()


@pytest.mark.parametrize('field,value',[('caption',''),('purpose',''),('caption','字'*161),('purpose','字'*501)])
def test_invalid_image_metadata_does_not_write(crop_fixture,field,value):
    f=crop_fixture; choice=request(f); choice[field]=value
    with pytest.raises(DesktopFacadeError):f.facade.import_textbook_preparation_crop(choice)
    assert not f.task_root.exists()


@pytest.mark.parametrize('corrupt',['receipt','image'])
def test_corrupt_existing_record_or_image_is_preserved(crop_fixture,corrupt):
    f=crop_fixture; choice=request(f); f.facade.import_textbook_preparation_crop(choice)
    if corrupt=='receipt':target=next((f.task_root/'preparation-v1'/'textbook-crops').glob('*.json'))
    else:target=next((f.task_root/'preparation-v1'/'images').glob('*.image'))
    target.write_bytes(b'changed-existing-content')
    before=disk_bytes(f.task_root)
    with pytest.raises(DesktopFacadeError):f.facade.import_textbook_preparation_crop(choice)
    assert disk_bytes(f.task_root)==before


def test_cancellation_wins_before_final_write(crop_fixture,monkeypatch):
    f=crop_fixture; choice=request(f); stop=TaskCancellationEvent()
    actual=f.crops.preview
    def stopped(*args,**kwargs):
        value=actual(*args,**kwargs); stop.set(); return value
    monkeypatch.setattr(f.crops,'preview',stopped)
    with pytest.raises(ReadCancelled), read_cancel_scope(stop):f.crops.confirm(**choice)
    assert not f.task_root.exists()


def test_receipt_publication_failure_creates_no_untraceable_image(crop_fixture,monkeypatch):
    f=crop_fixture; choice=request(f)
    def fail_link(*_):raise OSError('synthetic receipt link failure')
    monkeypatch.setattr(crops.os,'link',fail_link)
    with pytest.raises(DesktopFacadeError):f.facade.import_textbook_preparation_crop(choice)
    assert not list(f.task_root.rglob('*.image')) and not list(f.task_root.rglob('*.json'))
    assert not list(f.task_root.rglob('*.tmp'))


def test_image_write_failure_keeps_only_a_closed_traceable_selection(crop_fixture,monkeypatch):
    f=crop_fixture; choice=request(f)
    def fail_image(*_,**__):raise OSError('synthetic image store failure')
    monkeypatch.setattr(crops.PreparationImageStore,'import_bytes',fail_image)
    with pytest.raises(DesktopFacadeError):f.facade.import_textbook_preparation_crop(choice)
    assert not list(f.task_root.rglob('*.image'))
    receipt=json.loads(next(f.task_root.rglob('*.json')).read_text(encoding='utf-8'))
    assert receipt['candidate_only'] and not receipt['human_reviewed']
    assert not receipt['publication_allowed'] and receipt['source_sha256']==f.asset['source_sha256']


def test_facade_rejects_unknown_request_fields(crop_fixture):
    f=crop_fixture; choice=request(f); choice['source_path']=str(f.book)
    with pytest.raises(DesktopFacadeError):f.facade.import_textbook_preparation_crop(choice)
    assert not f.task_root.exists()


def test_long_valid_source_identity_is_compact_without_losing_receipt(crop_fixture,monkeypatch):
    f=crop_fixture;source=f.assets.source(**f.selection)
    source['source_title']='册'*200;source['visual_asset_id']='SYNTHETIC-'+('x'*170)
    monkeypatch.setattr(f.crops.catalog,'source',lambda **_:source)
    page=f.crops.page(f.selection)
    value=f.crops.preview(f.selection,page['page_revision'],f.bounds,'合成图表','观察合成色块')
    assert len(value['image_asset']['source'])<160
    assert value['provenance']['source_title']==source['source_title']
    assert value['provenance']['visual_asset_id']==source['visual_asset_id']
    assert value['provenance']['source_sha256']==source['source_sha256']
    assert not f.task_root.exists()


def test_rotated_page_and_whole_page_use_displayed_pixel_coordinates(shelf_fixture,tmp_path):
    from pypdf import PdfReader, PdfWriter
    f=shelf_fixture; reader=PdfReader(io.BytesIO(f.book.read_bytes())); reader.pages[8].rotate(90)
    writer=PdfWriter(); writer.append(reader); output=io.BytesIO(); writer.write(output)
    replace_book(f,output.getvalue())
    service=crops.TextbookPreparationCropService(f.root,tmp_path/'rotated-state')
    page=service.page(selection(f)); assert page['width']>page['height']
    value=service.preview(selection(f),page['page_revision'],
        {'left':0,'top':0,'right':page['width'],'bottom':page['height']},'整页合成图','确认使用整页')
    with Image.open(io.BytesIO(page['page_data'])) as source, Image.open(io.BytesIO(value['image_bytes'])) as cut:
        assert source.size==cut.size and source.tobytes()==cut.tobytes()
    assert not service.task_root.exists()


@pytest.mark.parametrize('student_visible',[True,False])
def test_crop_enters_real_lesson_files_only_for_selected_visibility(crop_fixture,student_visible):
    from docx import Document
    from pptx import Presentation
    f=crop_fixture; value=preview(f); asset=f.facade.import_textbook_preparation_crop(request(f,value))
    payload=sample(); payload['image_assets']=[asset]
    node=payload['lesson_design']['nodes'][0]
    node['image_ids']=[asset['asset_id']]; node['student_material']=student_visible
    original_payload = deepcopy(payload)
    receipt_files = {path: path.read_bytes() for path in (f.task_root/'preparation-v1/textbook-crops').glob('*.json')}
    output=export_design(f.facade,payload)
    payload['lesson_design']['exports'].append(output)
    teacher=checked_file(f.facade,payload['lesson_design'],output['id'],'lesson_plan.docx')
    worksheet=checked_file(f.facade,payload['lesson_design'],output['id'],'student_worksheet.docx')
    ppt=checked_file(f.facade,payload['lesson_design'],output['id'],'lesson_presentation.pptx')
    assert len(Document(teacher).inline_shapes)==1
    assert len(Document(worksheet).inline_shapes)==int(student_visible)
    with ZipFile(teacher) as z:
        assert value['image_bytes'] in [z.read(n) for n in z.namelist() if n.startswith('word/media/')]
    with ZipFile(ppt) as z:
        media=[z.read(n) for n in z.namelist() if n.startswith('ppt/media/')]
        assert (value['image_bytes'] in media)==student_visible
    presentation=Presentation(ppt)
    screen='\n'.join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)
    assert '仅教师答案-SECRET' not in screen
    if student_visible:
        assert asset['caption'] in screen
    for words in (screen, '\n'.join(p.text for p in Document(teacher).paragraphs),
                  '\n'.join(p.text for p in Document(worksheet).paragraphs)):
        assert ''.join(asset['source'].split()) not in ''.join(words.split())
        assert '裁选记录' not in words and 'PDF第9页' not in words
    assert payload['image_assets'] == original_payload['image_assets']
    assert payload['lesson_design']['nodes'] == original_payload['lesson_design']['nodes']
    assert {path: path.read_bytes() for path in receipt_files} == receipt_files
    for saved in output['files']:
        path=checked_file(f.facade,payload['lesson_design'],output['id'],saved['name'])
        assert hashlib.sha256(path.read_bytes()).hexdigest()==saved['sha256']
