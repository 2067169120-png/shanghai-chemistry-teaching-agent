"""Student flow removes teacher-created gaps and remains one saved work choice."""
from copy import deepcopy
from types import SimpleNamespace
import json

import pytest
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QGraphicsItem
from pptx import Presentation
from test_composition_render import app, text
from test_composition_ui import desk, personal, window
from test_composition_student_pages import question, payload, teacher_note
from test_desktop_studio_ui import settle
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    FORMATS, add_resource, block, resource, new_composition, edit_resource, reorder_resources, set_student_arrangement, validate_composition)
from integrations.deeptutor_shchem_v1.desktop_composition_render import arrange_composition, layout_composition
from integrations.deeptutor_shchem_v1.desktop_lesson_output import export_design, checked_file
from integrations.deeptutor_shchem_v1.desktop_workbench.lesson_composition_panel import ResourceFrame


def gap_example(format='16:9'):
    value=question(format)
    card=personal('完整的后续学生材料。')
    # The teacher answer continues over many frames; next material is authored
    # far down the final teacher page. Student flow must measure it afresh.
    w,h=FORMATS[format]
    page='CP-'+'c'*32;value['pages'].append(page)
    value=add_resource(value,card,{'page_id':page,'x':12,'y':h-65,'width':w-24,'height':53})
    return value,card


@pytest.mark.parametrize('format',FORMATS)
@pytest.mark.parametrize('columns',[1,2])
def test_student_material_reflows_after_hidden_answers_and_teacher_positions_survive(app,format,columns):
    value,card=gap_example(format);before=deepcopy(value)
    teacher=layout_composition(value,[],{},student=False)
    chosen=set_student_arrangement(value,columns)
    student=layout_composition(chosen,[],{},student=True)
    teacher_after=layout_composition(chosen,[],{},student=False)
    assert not student['errors']
    assert '教师答案' not in text(student) and '完整的后续学生材料。' in text(student)
    assert text(teacher)==text(teacher_after)
    assert chosen['placements']==before['placements'] and chosen['resources']==before['resources']
    assert list(teacher_after['pages'])==before['pages']
    assert len(student['pages'])<len(teacher_after['pages'])
    assert student['independent_student_arrangement'] and value==before
    assert list(student['pages'])==list(layout_composition(chosen,[],{},student=True)['pages'])
    assert set_student_arrangement(chosen)==before


def test_visible_edit_and_reorder_update_student_flow_without_cached_old_content(app):
    value,card=gap_example();value=set_student_arrangement(value,1)
    first=layout_composition(value,[],{},student=True)
    changed=edit_resource(value,card['id'],title='新作品标题')
    changed=reorder_resources(changed,[r['id'] for r in changed['resources']][::-1])
    result=layout_composition(changed,[],{},student=True)
    assert not result['errors'] and '新作品标题' in text(result)
    assert result['fragments'][0]['resource_id']==card['id']
    assert first['fragments'][0]['resource_id']!=card['id']
    assert changed['placements']==value['placements']
    assert changed['resources'][0]['original']==card['original']


@pytest.mark.parametrize('format',FORMATS)
def test_short_student_question_keeps_stem_and_every_option_on_one_page(app,format):
    example=question(format)['resources'][0]
    item=resource('complete_question','完整选项题',[
        block('text','question','请依据全部条件作答，原文须完整保留。',1),
        *[block('text','question',letter+'．这是完整的选项条件，不能因分页省略。',index)
          for index,letter in enumerate('ABCD',2)],block('text','answer','仅教师答案',6)],example['original']['evidence'])
    value=new_composition(format);w,h=FORMATS[format]
    for row in (personal('\n'.join('先读前置材料和任务。'*3 for _ in range(8))),item):
        value=add_resource(value,row,{'page_id':value['pages'][0],'x':12,'y':12,'width':w-24,'height':h-24})
    value=arrange_composition(value,[],{});value=set_student_arrangement(value,1)
    result=layout_composition(value,[],{},student=True)
    fragments=[f for f in result['fragments'] if f['resource_id']==item['id']]
    assert not result['errors'] and len(fragments)==1
    assert {e['block_id'] for e in fragments[0]['elements'] if e.get('block_id')}=={b['id'] for b in item['blocks'] if b['role']=='question'}
    assert '仅教师答案' not in text(result)


def test_long_student_question_still_continues_without_shortening(app):
    value=question();item=value['resources'][0]
    value=edit_resource(value,item['id'],locked=False)
    content='\n'.join('第%d条完整条件，不能删减。'%i for i in range(60))
    value=edit_resource(value,item['id'],values={item['blocks'][0]['id']:content})
    result=layout_composition(set_student_arrangement(value,1),[],{},student=True)
    assert not result['errors'] and len(result['pages'])>1
    assert '第59条完整条件，不能删减。' in text(result) and '教师答案' not in text(result)


@pytest.mark.parametrize('policy',[{}, {'columns':True},{'columns':0},{'columns':3},{'columns':1,'extra':0}])
def test_invalid_student_layout_policy_is_rejected(app,policy):
    value,_=gap_example();value['student_arrangement']=policy
    with pytest.raises(ValueError,match='学生版'):validate_composition(value)


def test_teacher_only_title_and_intentional_shared_blanks_follow_selected_mode(app):
    value=teacher_note(question('A4'));blank='CP-'+'f'*32;value['pages'].append(blank)
    shared=layout_composition(value,[],{},student=True)
    independent=layout_composition(set_student_arrangement(value,1),[],{},student=True)
    assert blank in shared['pages'] and blank not in independent['pages']
    assert '仅教师可见标题' not in text(independent) and not independent['errors']
    value['resources']=[value['resources'][-1]]
    value['placements']={value['resources'][0]['id']:value['placements'][value['resources'][0]['id']]}
    result=layout_composition(set_student_arrangement(value,1),[],{},student=True)
    assert not result['pages'] and any('学生版没有可见内容' in e for e in result['errors'])


def test_export_pdf_ppt_and_audit_share_derived_student_pages_and_full_text(app,tmp_path):
    import pypdfium2 as pdfium
    value,_=gap_example();chosen=set_student_arrangement(value,1);p=payload(chosen)
    facade=SimpleNamespace(paths=SimpleNamespace(task_root=tmp_path/'tasks'))
    original=deepcopy(p);record=export_design(facade,p);p['lesson_design']['exports'].append(record)
    directory=checked_file(facade,p['lesson_design'],record['id'],'student_canvas.pdf').parent
    audit=json.loads((directory/'canvas-output.json').read_text(encoding='utf-8'))
    student=layout_composition(chosen,[],{},student=True)
    assert audit['student_arrangement']=={'columns':1} and audit['student_independent_arrangement']
    assert audit['page_ids_by_audience']['student']==list(student['pages'])
    assert audit['page_ids_by_audience']['teacher']==value['pages']
    with pdfium.PdfDocument(str(directory/'student_canvas.pdf')) as doc:assert len(doc)==len(student['pages'])
    slides=Presentation(directory/'student_presentation.pptx').slides
    assert len(slides)==len(student['pages'])
    actual=''.join(s.text for slide in slides for s in slide.shapes if s.has_text_frame)
    assert '完整的后续学生材料。' in actual and '教师答案' not in actual
    p['lesson_design']['exports'].pop();assert p==original


def test_native_choice_navigation_undo_and_teacher_geometry_remain_separate(desk):
    dialog,panel,qtapp,_page=desk
    value,card=gap_example();panel.commit(value);settle(qtapp)
    before=deepcopy(dialog.history.value)
    panel.student_arrangement.setCurrentIndex(1);settle(qtapp)
    chosen=deepcopy(dialog.history.value)
    assert chosen['composition']['student_arrangement']=={'columns':1}
    assert chosen['composition']['placements']==before['composition']['placements']
    assert '学生版独立1栏' in panel.status.text()
    assert panel.page_navigation.pages.count()==len(panel.layout_result['pages'])
    panel.selected=card['id'];panel.refresh();settle(qtapp)
    assert panel.page_picker.currentText().startswith('学生编排') and not panel.geometry_apply.isEnabled()
    assert not any(i.flags() & QGraphicsItem.GraphicsItemFlag.ItemIsMovable for i in panel.scene.items() if isinstance(i,ResourceFrame))
    panel.move_frame(card['id'],0,30,30);panel.apply_geometry();assert dialog.history.value==chosen
    dialog.resize(1050,830);settle(qtapp)
    from PySide6.QtCore import QRectF
    w,h=FORMATS[value['format']];pages=list(panel.layout_result['pages']);index=pages.index(panel.visible_page())
    visible=panel.view.mapFromScene(QRectF(0,index*(h+12),w,h)).boundingRect()
    assert visible.left()>=-1 and visible.right()<=panel.view.viewport().width()+1
    QTest.mouseClick(panel.student,Qt.MouseButton.LeftButton);settle(qtapp)
    assert panel.geometry_apply.isEnabled() and panel.page_picker.currentText().startswith('教师编排')
    assert panel.composition()['placements']==value['placements']
    dialog.travel(False);assert dialog.history.value==before
    dialog.travel(True);assert dialog.history.value==chosen
    panel.student_arrangement.setCurrentIndex(0);settle(qtapp)
    assert panel.composition()==value
