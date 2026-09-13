from __future__ import annotations

import copy
import hashlib
import json
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

from integrations.deeptutor_shchem_v1.desktop_preparation_renderer import (
    NativePreparationRenderCancelled,
    NativePreparationRenderer,
    NativePreparationRenderError,
    _draw_multiline_text,
    _fallback_multiline_plan,
    _fit_text,
    _FontBook,
    _glyph_runs,
    _GlyphFallback,
    _multiline_textbbox,
    _numbered_title,
    _text_length,
    _wrap_text,
)


def test_glyph_fallback_keeps_primary_runs_and_reports_all_fonts_missing():
    primary, symbols = object(), object()
    resolver = _GlyphFallback(primary, None, ())
    resolver.faces = [primary, symbols]
    resolver.loaded = True
    supported = {primary: set("中文Na"), symbols: set("₂⁺⇌")}
    resolver.supports = lambda face, char: char in supported[face]
    text = "中文Na₂⁺⇌中文"
    assert resolver.runs(text) == [
        ("中文Na", primary),
        ("₂⁺⇌", symbols),
        ("中文", primary),
    ]
    assert "".join(run for run, _ in resolver.runs(text)) == text
    with pytest.raises(NativePreparationRenderError) as raised:
        resolver.runs("\u0378")
    assert raised.value.code == "preparation_preview_glyph_missing"
    assert "U+0378" in raised.value.message_zh
    assert "未替换原文" in raised.value.message_zh


@pytest.mark.parametrize("align", ["left", "center", "right"])
def test_fallback_measurement_and_drawing_share_exact_runs_and_baselines(align):
    # Bundled Pillow fonts make this geometry test independent of OS font paths.
    primary = ImageFont.load_default(size=32)
    symbols = ImageFont.load_default(size=26)
    resolver = _GlyphFallback(primary, ImageFont, ())
    resolver.faces = [primary, symbols]
    resolver.loaded = True
    resolver.supports = lambda face, char: (face is symbols) == (char == "B")
    primary._shchem_glyph_fallback = resolver
    text = "AABAA\nAB"
    canvas = Image.new("L", (300, 160), 0)
    draw = ImageDraw.Draw(canvas)
    runs = _glyph_runs("AABAA", primary)
    assert _text_length(draw, "AABAA", primary) == sum(
        draw.textlength(run, font=face) for run, face in runs
    )
    plan = _fallback_multiline_plan(draw, text, primary, 8, align)
    assert plan is not None
    assert _multiline_textbbox(draw, text, primary, 8, align) == plan[0]
    _draw_multiline_text(
        draw, (20, 20), text, font=primary, fill=255, spacing=8, align=align
    )
    expected = Image.new("L", canvas.size, 0)
    expected_draw = ImageDraw.Draw(expected)
    for run, face, x, baseline in plan[1]:
        expected_draw.text(
            (20 + x, 20 + baseline), run, font=face, fill=255, anchor="ls"
        )
    assert canvas.tobytes() == expected.tobytes()
    ink = canvas.getbbox()
    assert ink is not None
    box = plan[0]
    assert ink[0] >= int(20 + box[0]) - 1
    assert ink[1] >= int(20 + box[1]) - 1
    assert ink[2] <= int(20 + box[2]) + 1
    assert ink[3] <= int(20 + box[3]) + 1
    wrapped = _wrap_text(draw, text, primary, 80)
    assert wrapped.replace("\n", "") == text.replace("\n", "")


@pytest.mark.parametrize("bold", [False, True])
def test_available_cjk_and_chemistry_fonts_render_real_glyphs_without_tofu(bold):
    font = _FontBook().get(36, bold)
    resolver = font._shchem_glyph_fallback
    if not all(resolver.supports(font, char) for char in "电离"):
        pytest.skip("No CJK primary font installed")
    text = "电离：H₂CO₃ ⇌ H⁺ + HCO₃⁻；SO₄²⁻"
    resolver._load()
    if not all(
        any(resolver.supports(face, c) for face in resolver.faces) for c in text
    ):
        pytest.skip("No installed font covers all chemistry test glyphs")
    runs = _glyph_runs(text, font)
    assert "".join(run for run, _ in runs) == text
    for run, face in runs:
        assert all(resolver.supports(face, char) for char in run)
        if "电离" in run:
            assert face is font
    canvas = Image.new("L", (1100, 150), 0)
    draw = ImageDraw.Draw(canvas)
    _draw_multiline_text(draw, (20, 20), text, font=font, fill=255, spacing=8)
    assert canvas.getbbox() is not None
    if any(not resolver.supports(font, char) for char in text):
        old = Image.new("L", canvas.size, 0)
        ImageDraw.Draw(old).text((20, 20), text, font=font, fill=255)
        assert canvas.tobytes() != old.tobytes()


@pytest.mark.parametrize(
    "token",
    [
        "H2SO3",
        "Ca(OH)2",
        "NaCl(aq)",
        "CO3^2-",
        "H₂SO₃",
        "Fe³⁺",
        "CuSO4·5H2O",
        "PowerPoint",
    ],
)
def test_wrap_preserves_formula_and_latin_tokens(token):
    class UnitWidthDraw:
        def textlength(self, text, *, font):
            return len(text)

    text = "说明依据是" + token + "。"
    width = max(7, len(token) + 2)
    wrapped = _wrap_text(UnitWidthDraw(), text, None, width)
    assert token in wrapped
    assert wrapped.replace("\n", "") == text
    assert all(len(line) <= width for line in wrapped.splitlines())
    assert not any(line.startswith("。") for line in wrapped.splitlines())


def test_unbreakable_token_reports_horizontal_overflow():
    draw = ImageDraw.Draw(Image.new("RGB", (800, 600)))
    text, _, overflow = _fit_text(
        draw,
        _FontBook(),
        "H2SO3",
        width=1,
        height=500,
        preferred_px=20,
        minimum_px=20,
    )
    assert text == "H2SO3"
    assert overflow is True


def test_wrap_preserves_cjk_punctuation_and_explicit_lines():
    class UnitWidthDraw:
        def textlength(self, text, *, font):
            return len(text)

    text = "观察（实验），解释。\n\n再讨论【依据】。"
    wrapped = _wrap_text(UnitWidthDraw(), text, None, 5)
    assert wrapped.replace("\n", "") == text.replace("\n", "")
    assert "\n\n" in wrapped
    assert all(
        not line.startswith(("，", "。", "）", "】")) for line in wrapped.splitlines()
    )
    assert all(not line.endswith(("（", "【")) for line in wrapped.splitlines())


def test_classroom_cover_title_remains_complete_and_readable(tmp_path):
    candidate = _candidate("ppt")
    title = "水溶液导电，溶质一定是电解质吗？"
    candidate["slides"][0]["title"] = title
    candidate["topic"] = title
    original = copy.deepcopy(candidate)
    result = NativePreparationRenderer().render(
        candidate, output_kind="ppt", output_dir=tmp_path / "balanced-cover"
    )
    deck = json.loads(
        Path(_artifact_map(result)["deck_json"]["path"]).read_text("utf-8")
    )
    element = next(
        item
        for item in deck["slides"][0]["elements"]
        if item.get("source_paths") == ["/topic"]
    )
    lines = element["text"].splitlines()
    assert "".join(lines) == title
    assert 1 <= len(lines) <= 2
    assert min(map(len, lines)) >= 6
    assert element["font_size_pt"] >= 39.6
    assert element["align"] == "center"
    assert element["overflow"] is False
    presentation = Presentation(_artifact_map(result)["pptx"]["path"])
    assert element["text"] in [
        shape.text for shape in presentation.slides[0].shapes if shape.has_text_frame
    ]
    assert candidate == original


@pytest.mark.parametrize(
    "text,width",
    [
        ("短标题", 1260),
        ("观察（实验），解释。\n\n再讨论【依据】。", 450),
        ("H2SO3", 1),
        ("NaCl(aq)与CO2水溶液的导电性比较", 600),
    ],
)
def test_title_balance_preserves_explicit_breaks_tokens_and_overflow(text, width):
    draw = ImageDraw.Draw(Image.new("RGB", (1600, 900)))
    fonts = _FontBook()
    options = {"width": width, "height": 700, "preferred_px": 60, "minimum_px": 60}
    ordinary = _fit_text(draw, fonts, text, **options)
    balanced = _fit_text(draw, fonts, text, balance_lines=True, **options)
    assert balanced[0].replace("\n", "") == text.replace("\n", "")
    assert balanced[1:] == ordinary[1:]
    assert len(balanced[0].splitlines()) == len(ordinary[0].splitlines())
    if "\n" in text or width == 1 or len(ordinary[0].splitlines()) == 1:
        assert balanced == ordinary
    if "NaCl(aq)" in text:
        assert "NaCl(aq)" in balanced[0] and "CO2" in balanced[0]
    assert all(
        not line.startswith(("，", "。", "）", "】", "？"))
        for line in balanced[0].splitlines()
    )


@pytest.mark.parametrize(
    ("title", "expected"),
    [
        ("活动1：概念辨析", "活动1 概念辨析"),
        ("活动 1 概念辨析", "活动1 概念辨析"),
        ("活动1", "活动1"),
        ("概念辨析", "活动1 概念辨析"),
        ("活动10：概念辨析", "活动1 活动10：概念辨析"),
        ("活动1号材料", "活动1 活动1号材料"),
        ("比较活动1的结果", "活动1 比较活动1的结果"),
    ],
)
def test_numbered_title_only_removes_matching_display_prefix(title, expected):
    assert _numbered_title("活动", 1, title) == expected


def test_docx_matrix_and_sections_do_not_repeat_numbered_titles(tmp_path):
    candidate = _candidate("lesson_plan")
    candidate["activities"][0]["title"] = "活动1：证据标注"
    candidate["activities"][0]["teacher_action"] = "额外准备证据材料"
    candidate["assessments"][0]["title"] = "评价1：证据链检查"
    original = copy.deepcopy(candidate)
    result = NativePreparationRenderer().render(
        candidate, output_kind="lesson_plan", output_dir=tmp_path / "numbered"
    )
    document = Document(_artifact_map(result)["lesson_plan_docx"]["path"])
    matrix = "\n".join(
        cell.text for row in document.tables[1].rows for cell in row.cells
    )
    body = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for text in (matrix, body):
        assert "活动1 证据标注" in text
        assert "评价1 证据链检查" in text
        assert "活动1 活动1" not in text and "评价1 评价1" not in text
    assert candidate == original


def _candidate(artifact_mode: str = "linked_bundle") -> dict:
    return {
        "schema_version": "shchem.desktop-preparation-candidate.v1",
        "candidate_id": "PREPCAND-fixture-001",
        "title": "氧化还原反应证据推理",
        "topic": "氧化还原反应",
        "audience": "高一学生",
        "artifact_mode": artifact_mode,
        "lesson_route": "new_lesson",
        "timing": {
            "periods": 1,
            "minutes_per_period": 40,
            "total_minutes": 40,
        },
        "source_basis": {
            "mode": "teacher_input_only",
            "evidence_ids": [],
            "statement_zh": "教师输入驱动的个人备课候选",
        },
        "objectives": [
            {"id": "O1", "statement": "依据化合价变化说明电子转移方向"},
            {"id": "O2", "statement": "用证据判断氧化剂与还原剂"},
        ],
        "activities": [
            {
                "id": "A1",
                "title": "证据标注",
                "objective_ids": ["O1"],
                "minutes": 15,
                "teacher_action": "展示反应信息并追问化合价变化",
                "student_action": "标注化合价并说明电子得失",
                "materials": ["反应信息"],
            },
            {
                "id": "A2",
                "title": "迁移判断",
                "objective_ids": ["O2"],
                "minutes": 20,
                "teacher_action": "组织独立作答和同伴核对",
                "student_action": "写出判断依据并修正错误",
                "materials": ["迁移练习"],
            },
        ],
        "assessments": [
            {
                "id": "E1",
                "title": "证据链检查",
                "objective_ids": ["O1"],
                "activity_ids": ["A1"],
                "evidence_of_learning": "化合价与电子转移标注",
                "success_criteria": "方向和理由一致",
            },
            {
                "id": "E2",
                "title": "迁移检测",
                "objective_ids": ["O2"],
                "activity_ids": ["A2"],
                "evidence_of_learning": "判断结论与理由",
                "success_criteria": "结论引用可见证据",
            },
        ],
        "slides": [
            {
                "id": "S1",
                "order": 1,
                "title": "氧化还原反应证据推理",
                "purpose": "建立本课核心问题",
                "objective_ids": ["O1", "O2"],
                "activity_ids": [],
                "assessment_ids": [],
                "minutes": 2,
                "content": ["如何从可见信息走到电子转移解释"],
                "teacher_notes": "先收集学生的初始判断",
            },
            {
                "id": "S2",
                "order": 2,
                "title": "化合价变化与电子转移",
                "purpose": "连接证据与模型",
                "objective_ids": ["O1"],
                "activity_ids": ["A1"],
                "assessment_ids": ["E1"],
                "minutes": 10,
                "content": ["标出反应前后的化合价", "根据升降方向判断电子得失"],
                "teacher_notes": "保留学生标注痕迹",
            },
            {
                "id": "S3",
                "order": 3,
                "title": "迁移练习",
                "purpose": "独立完成判断并写出证据",
                "objective_ids": ["O2"],
                "activity_ids": ["A2"],
                "assessment_ids": ["E2"],
                "minutes": 8,
                "content": ["判断氧化剂和还原剂", "写出支持结论的化合价变化"],
                "teacher_notes": "计时作答",
            },
            {
                "id": "S4",
                "order": 4,
                "title": "答案与讲评",
                "purpose": "核对证据链中的推理节点",
                "objective_ids": ["O2"],
                "activity_ids": ["A2"],
                "assessment_ids": ["E2"],
                "minutes": 8,
                "content": ["先判断化合价升降", "再据电子得失确定物质角色"],
                "teacher_notes": "作答后再展示",
            },
            {
                "id": "S5",
                "order": 5,
                "title": "方法总结",
                "purpose": "回顾证据到结论的路径",
                "objective_ids": ["O1", "O2"],
                "activity_ids": [],
                "assessment_ids": [],
                "minutes": 2,
                "content": ["找变化", "判得失", "说明物质角色"],
                "teacher_notes": "用学生语言完成总结",
            },
        ],
        "lesson_stages": [
            {
                "id": "L1",
                "title": "证据与模型",
                "objective_ids": ["O1"],
                "activity_ids": ["A1"],
                "assessment_ids": ["E1"],
                "minutes": 15,
                "teacher_actions": ["展示反应信息并追问化合价变化"],
                "student_actions": ["标注化合价并说明电子得失"],
                "materials": ["反应信息"],
                "assessment": "检查标注与解释是否一致",
            },
            {
                "id": "L2",
                "title": "迁移与评价",
                "objective_ids": ["O2"],
                "activity_ids": ["A2"],
                "assessment_ids": ["E2"],
                "minutes": 20,
                "teacher_actions": ["组织独立作答和同伴核对"],
                "student_actions": ["写出判断依据并修正错误"],
                "materials": ["迁移练习"],
                "assessment": "检查结论是否引用可见证据",
            },
        ],
        "homework": {
            "title": "课后练习",
            "tasks": [
                {
                    "id": "HW1",
                    "instruction": "完成一题氧化剂与还原剂判断",
                    "objective_ids": ["O2"],
                }
            ],
            "estimated_minutes": 8,
        },
        "uncertainties": ["练习难度需结合班级情况确认"],
        "candidate_only": True,
        "teacher_review_required": True,
        "publication_allowed": False,
        "official_claim_allowed": False,
    }


def _artifact_map(result: dict) -> dict[str, dict]:
    return {row["artifact_id"]: row for row in result["artifacts"]}


def _all_candidate_strings(value) -> set[str]:
    result: set[str] = set()

    def walk(current) -> None:
        if isinstance(current, str):
            result.add(current.strip())
        elif isinstance(current, dict):
            for child in current.values():
                walk(child)
        elif isinstance(current, list):
            for child in current:
                walk(child)

    walk(value)
    return result


def test_joint_renderer_creates_editable_pptx_matching_pngs_docx_and_machine_qa(
    tmp_path: Path,
) -> None:
    progress: list[tuple[int, str]] = []

    def report(value: int, stage: str) -> None:
        progress.append((value, stage))

    output = tmp_path / "attempt-0001"
    result = NativePreparationRenderer().render(
        _candidate(),
        output_kind="joint",
        output_dir=output,
        report_progress=report,
        is_cancelled=lambda: False,
    )

    assert result["output_kind"] == "linked_bundle"
    assert result["slide_count"] == 5
    assert progress[0] == (2, "validating")
    assert progress[-1] == (100, "completed")
    assert [value for value, _stage in progress] == sorted(
        value for value, _stage in progress
    )

    artifacts = _artifact_map(result)
    assert set(artifacts) == {
        "candidate_json",
        "deck_json",
        "pptx",
        "preview_montage",
        "lesson_plan_docx",
        "qa_report",
    }
    for row in artifacts.values():
        path = Path(row["path"])
        assert path.is_file()
        assert path.resolve().is_relative_to(output.resolve())
        payload = path.read_bytes()
        assert row["size"] == len(payload) > 0
        assert row["sha256"] == hashlib.sha256(payload).hexdigest()
        assert row["filename"] == path.relative_to(output).as_posix()

    presentation = Presentation(artifacts["pptx"]["path"])
    assert presentation.slide_width * 9 == presentation.slide_height * 16
    assert len(presentation.slides) == 5
    editable_text = [
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    ]
    assert _candidate()["topic"] in "\n".join(editable_text)
    assert "再据电子得失确定物质角色" in "\n".join(editable_text)

    slide_root = Path(result["rendered_slides"])
    slide_paths = sorted(
        slide_root.glob("slide-*.png"), key=lambda path: int(path.stem.split("-")[1])
    )
    assert [path.name for path in slide_paths] == [
        f"slide-{number}.png" for number in range(1, 6)
    ]
    for path in slide_paths:
        with Image.open(path) as image:
            assert image.size == (1600, 900)
            assert image.mode == "RGB"
    with Image.open(artifacts["preview_montage"]["path"]) as image:
        assert image.size == (1600, 900)

    deck = json.loads(Path(artifacts["deck_json"]["path"]).read_text(encoding="utf-8"))
    assert deck["slide_size"]["aspect_ratio"] == "16:9"
    assert [slide["slide_type"] for slide in deck["slides"]] == [
        "cover",
        "content",
        "practice",
        "answer",
        "summary",
    ]
    candidate_strings = _all_candidate_strings(_candidate())
    for slide in deck["slides"]:
        for element in slide["elements"]:
            for source_text in element.get("source_texts", []):
                assert source_text in candidate_strings

    document = Document(artifacts["lesson_plan_docx"]["path"])
    document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    document_text += "\n" + "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "氧化还原反应证据推理" in document_text
    assert "依据化合价变化说明电子转移方向" in document_text
    assert "完成一题氧化剂与还原剂判断" in document_text

    qa = json.loads(Path(artifacts["qa_report"]["path"]).read_text(encoding="utf-8"))
    assert qa["scope"] == "machine_layout_check_only"
    assert qa["machine_checks_passed"] is True
    assert qa["candidate_only"] is True
    assert qa["teacher_review_required"] is True
    assert qa["publication_allowed"] is False
    assert qa["human_visual_review_performed"] is False
    assert "chemistry_correctness" in qa["not_checked"]
    assert result["quality"]["rendered_slides"] == qa["rendered_slides"]


def test_glyph_fallback_does_not_rewrite_pptx_docx_or_candidate_text(tmp_path):
    candidate = _candidate()
    expression = "H₂CO₃ ⇌ H⁺ + HCO₃⁻"
    candidate["slides"][1]["content"] = [expression]
    candidate["lesson_stages"][0]["teacher_actions"] = [expression]
    before = copy.deepcopy(candidate)
    result = NativePreparationRenderer().render(
        candidate, output_kind="joint", output_dir=tmp_path / "unicode-originals"
    )
    artifacts = _artifact_map(result)
    presentation = Presentation(artifacts["pptx"]["path"])
    assert expression in [
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    document = Document(artifacts["lesson_plan_docx"]["path"])
    texts = [p.text for p in document.paragraphs]
    texts.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    assert any(expression in text for text in texts)
    saved = json.loads(Path(artifacts["candidate_json"]["path"]).read_text("utf-8"))
    assert saved["slides"][1]["content"] == [expression]
    assert candidate == before


def test_cover_four_supporting_items_fit_below_chapter_title_at_safe_size(
    tmp_path: Path,
) -> None:
    candidate = copy.deepcopy(_candidate("ppt"))
    candidate["topic"] = "氧化还原反应复习"
    candidate["audience"] = "高三等级考复习班"
    candidate["slides"][0].update(
        {
            "title": "课题与目标",
            "purpose": "建立学习方向",
            "content": ["识别角色", "应用电子守恒", "规范表达", "说明判断依据"],
        }
    )

    result = NativePreparationRenderer().render(
        candidate,
        output_kind="ppt",
        output_dir=tmp_path / "five-item-cover",
    )
    artifacts = _artifact_map(result)
    deck = json.loads(Path(artifacts["deck_json"]["path"]).read_text(encoding="utf-8"))
    cover = deck["slides"][0]
    source_paths = [
        "/slides/0/content/0",
        "/slides/0/content/1",
        "/slides/0/content/2",
        "/slides/0/content/3",
    ]
    supporting = next(
        element
        for element in cover["elements"]
        if element.get("source_paths") == source_paths
    )
    assert supporting["bbox_px"] == [150, 490, 1300, 235]
    assert supporting["text"].replace("\n\n", "\n") == (
        "识别角色\n应用电子守恒\n规范表达\n说明判断依据"
    )
    assert supporting["font_size_pt"] >= 20
    assert supporting["overflow"] is False
    assert (
        deck["slide_size"]["height_px"]
        - sum((supporting["bbox_px"][1], supporting["bbox_px"][3]))
        >= 100
    )

    qa = json.loads(Path(artifacts["qa_report"]["path"]).read_text(encoding="utf-8"))
    overflow_check = next(
        check
        for check in qa["checks"]
        if check["check"] == "text_box_overflow_estimate"
    )
    shared_layout_check = next(
        check
        for check in qa["checks"]
        if check["check"] == "pptx_and_png_share_layout_plan"
    )
    assert overflow_check == {
        "check": "text_box_overflow_estimate",
        "items": [],
        "method": "Pillow metrics with 0.4-em line-end reserve and 1.18-em line boxes",
        "passed": True,
    }
    assert shared_layout_check["passed"] is True
    assert qa["machine_checks_passed"] is True

    presentation = Presentation(artifacts["pptx"]["path"])
    editable_cover_text = "\n".join(
        shape.text
        for shape in presentation.slides[0].shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )
    for text in supporting["source_texts"]:
        assert text in editable_cover_text
    with Image.open(Path(result["rendered_slides"]) / "slide-1.png") as image:
        assert image.size == (1600, 900)


def test_docx_title_has_no_border_and_uncertainties_hide_internal_field_codes(
    tmp_path: Path,
) -> None:
    candidate = copy.deepcopy(_candidate("lesson_plan"))
    candidate["uncertainties"] = [
        {
            "field": "example_data",
            "description": "例题数据需要教师核对。",
            "teacher_action": "授课前核对题面、答案和单位。",
        },
        {
            "field": "source_basis",
            "description": "材料出处需要教师确认。",
            "teacher_action": "使用前检查原始材料。",
        },
    ]

    result = NativePreparationRenderer().render(
        candidate,
        output_kind="lesson_plan",
        output_dir=tmp_path / "readable-docx",
    )
    docx_path = Path(_artifact_map(result)["lesson_plan_docx"]["path"])
    with ZipFile(docx_path) as archive:
        styles_xml = archive.read("word/styles.xml")
        document_xml = archive.read("word/document.xml")

    word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    styles_root = ElementTree.fromstring(styles_xml)
    title_style = styles_root.find(f".//{word}style[@{word}styleId='Title']")
    assert title_style is not None
    assert title_style.find(f"./{word}pPr/{word}pBdr") is None
    style_color = title_style.find(f"./{word}rPr/{word}color")
    assert style_color is not None
    assert style_color.get(f"{word}val") == "000000"

    document_root = ElementTree.fromstring(document_xml)
    title_paragraph = document_root.find(f"./{word}body/{word}p")
    assert title_paragraph is not None
    assert title_paragraph.find(f"./{word}pPr/{word}pBdr") is None
    title_color = title_paragraph.find(f"./{word}r/{word}rPr/{word}color")
    assert title_color is not None
    assert title_color.get(f"{word}val") == "000000"

    assert b"example_data" not in document_xml
    assert b"source_basis" not in document_xml
    document = Document(docx_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "例题数据需要教师核对。 教师处理：授课前核对题面、答案和单位。" in text
    assert "材料出处需要教师确认。 教师处理：使用前检查原始材料。" in text


@pytest.mark.parametrize(
    ("output_kind", "expected", "forbidden"),
    [
        (
            "ppt",
            {"candidate_json", "deck_json", "pptx", "preview_montage", "qa_report"},
            {"lesson_plan_docx"},
        ),
        (
            "lesson_plan",
            {"candidate_json", "lesson_plan_docx", "qa_report"},
            {"deck_json", "pptx", "preview_montage"},
        ),
    ],
)
def test_output_kind_only_creates_corresponding_artifacts(
    tmp_path: Path, output_kind: str, expected: set[str], forbidden: set[str]
) -> None:
    candidate = _candidate(output_kind)
    output = tmp_path / output_kind
    events: list[dict] = []
    result = NativePreparationRenderer()(
        candidate,
        output_kind=output_kind,
        output_dir=output,
        report_progress=events.append,
    )
    assert set(_artifact_map(result)) == expected
    assert not (set(_artifact_map(result)) & forbidden)
    assert events[-1]["progress"] == 100
    assert events[-1]["stage"] == "completed"
    if output_kind == "lesson_plan":
        assert result["rendered_slides"] is None
        assert not (output / "rendered_slides").exists()
        assert not (output / "deck.json").exists()
    else:
        assert Path(result["rendered_slides"]).is_dir()
        assert not (output / "lesson_plan.docx").exists()


def test_renderer_cancellation_and_existing_attempt_are_fail_closed(
    tmp_path: Path,
) -> None:
    renderer = NativePreparationRenderer()
    cancelled = tmp_path / "cancelled"
    with pytest.raises(NativePreparationRenderCancelled) as error:
        renderer.render(
            _candidate("ppt"),
            output_kind="ppt",
            output_dir=cancelled,
            is_cancelled=lambda: True,
        )
    assert error.value.code == "preparation_render_cancelled"
    assert not cancelled.exists()

    occupied = tmp_path / "occupied"
    occupied.mkdir()
    sentinel = occupied / "candidate.json"
    sentinel.write_text("keep", encoding="utf-8")
    with pytest.raises(NativePreparationRenderError) as error:
        renderer.render(_candidate("ppt"), output_kind="ppt", output_dir=occupied)
    assert error.value.code == "preparation_render_artifact_exists"
    assert sentinel.read_text(encoding="utf-8") == "keep"


def test_renderer_rejects_authority_drift_and_output_kind_mismatch(
    tmp_path: Path,
) -> None:
    renderer = NativePreparationRenderer()
    unsafe = _candidate("ppt")
    unsafe["publication_allowed"] = True
    with pytest.raises(NativePreparationRenderError) as error:
        renderer.render(unsafe, output_kind="ppt", output_dir=tmp_path / "unsafe")
    assert error.value.code == "preparation_candidate_authority_invalid"

    with pytest.raises(NativePreparationRenderError) as error:
        renderer.render(
            _candidate("lesson_plan"),
            output_kind="ppt",
            output_dir=tmp_path / "mismatch",
        )
    assert error.value.code == "preparation_output_kind_mismatch"


def test_teacher_notes_are_preserved_off_projection_and_cover_uses_subject(tmp_path):
    candidate = _candidate()
    candidate["slides"][0]["title"] = "封面"
    candidate["slides"][0]["purpose"] = "仅给教师的教学意图"
    candidate["slides"][0]["teacher_notes"] = (
        "参考 E24 的范围尚待核验，请勿当作已核验答案。"
    )
    output = tmp_path / "notes"
    result = NativePreparationRenderer().render(
        candidate, output_kind="joint", output_dir=output
    )
    presentation = Presentation(_artifact_map(result)["pptx"]["path"])
    visible = "\n".join(
        s.text for s in presentation.slides[0].shapes if s.has_text_frame
    )
    assert candidate["topic"] in visible.replace("\n", "")
    assert "封面" not in visible
    assert "仅给教师的教学意图" not in visible
    notes = presentation.slides[0].notes_slide.notes_text_frame.text
    assert candidate["slides"][0]["purpose"] in notes
    assert candidate["slides"][0]["teacher_notes"] in notes
    document = Document(_artifact_map(result)["lesson_plan_docx"]["path"])
    matrix = document.tables[1]
    assert "tblHeader" in matrix.rows[0]._tr.xml
    assert all(row._tr.xpath("./w:trPr/w:cantSplit") for row in matrix.rows[1:])
    assert not matrix._tbl.xpath("./w:tblPr/w:tblpPr")
    matrix_text = "\n".join(cell.text for row in matrix.rows for cell in row.cells)
    assert "展示反应信息并追问化合价变化" not in matrix_text
    text = "\n".join(p.text for p in document.paragraphs)
    assert "展示反应信息并追问化合价变化" in text
    assert "化合价与电子转移标注" in text
    assert "教学过程" in text and "评价标准" in text
    assert "活动说明" not in text and "补充活动" not in text
    assert text.count("展示反应信息并追问化合价变化") == 1


def test_candidate_title_never_controls_artifact_paths(tmp_path: Path) -> None:
    candidate = copy.deepcopy(_candidate("lesson_plan"))
    candidate["title"] = "..\\..\\outside"
    output = tmp_path / "safe"
    result = NativePreparationRenderer().render(
        candidate, output_kind="lesson_plan", output_dir=output
    )
    assert all(
        Path(row["path"]).resolve().is_relative_to(output.resolve())
        for row in result["artifacts"]
    )
    assert not (tmp_path / "outside.docx").exists()


def test_lesson_timeline_retains_linked_activity_supplements_and_input(tmp_path):
    candidate = _candidate("lesson_plan")
    candidate["activities"][0]["teacher_action"] = (
        "先请学生独立写出判断依据，再比较两种解释。"
    )
    original = copy.deepcopy(candidate)
    result = NativePreparationRenderer().render(
        candidate, output_kind="lesson_plan", output_dir=tmp_path / "supplement"
    )
    document = Document(_artifact_map(result)["lesson_plan_docx"]["path"])
    text = "\n".join(p.text for p in document.paragraphs)
    assert text.count(candidate["activities"][0]["teacher_action"]) == 1
    assert text.count(candidate["lesson_stages"][0]["teacher_actions"][0]) == 1
    assert text.index("教学过程") < text.index(
        candidate["activities"][0]["teacher_action"]
    )
    assert text.index("教学过程") < text.index("评价标准")
    assert "补充活动" not in text
    assert "活动补充资料" in text
    assert text.index(candidate["lesson_stages"][-1]["title"]) < text.index(
        "活动补充资料"
    )
    assert candidate == original


def test_repeated_stage_action_is_not_removed_as_duplicate(tmp_path):
    candidate = _candidate("lesson_plan")
    repeated = copy.deepcopy(candidate["lesson_stages"][0])
    repeated["title"] = "课末重新检查初始判断"
    candidate["lesson_stages"].append(repeated)
    result = NativePreparationRenderer().render(
        candidate, output_kind="lesson_plan", output_dir=tmp_path / "repeated"
    )
    document = Document(_artifact_map(result)["lesson_plan_docx"]["path"])
    text = "\n".join(p.text for p in document.paragraphs)
    # The same action can intentionally recur at a different point in time.
    assert text.count(repeated["teacher_actions"][0]) == 2
    assert "课末重新检查初始判断" in text


def test_unlinked_legacy_activity_is_preserved_without_guessing_placement(tmp_path):
    candidate = _candidate("lesson_plan")
    candidate["lesson_stages"][0]["activity_ids"] = []
    candidate["activities"][0]["teacher_action"] = "未定位活动仍需保留的教师准备说明。"
    result = NativePreparationRenderer().render(
        candidate, output_kind="lesson_plan", output_dir=tmp_path / "legacy"
    )
    document = Document(_artifact_map(result)["lesson_plan_docx"]["path"])
    text = "\n".join(p.text for p in document.paragraphs)
    assert "补充活动" in text
    assert text.count(candidate["activities"][0]["teacher_action"]) == 1


def test_cover_does_not_inject_teacher_audience_notes(tmp_path):
    candidate = _candidate("ppt")
    candidate["slides"][0]["title"] = "本节课的核心问题"
    candidate["audience"] = "高二化学复习；实际班级学情尚未提供"
    result = NativePreparationRenderer().render(
        candidate, output_kind="ppt", output_dir=tmp_path / "cover-audience"
    )
    presentation = Presentation(_artifact_map(result)["pptx"]["path"])
    visible = "\n".join(
        s.text for s in presentation.slides[0].shapes if s.has_text_frame
    )
    assert candidate["topic"] in visible.replace("\n", "")
    assert candidate["audience"] not in visible
    assert "本节课的核心问题" not in visible


def test_ppt_suggested_minutes_are_in_presenter_notes_only(tmp_path):
    candidate = _candidate("ppt")
    original = copy.deepcopy(candidate)
    result = NativePreparationRenderer().render(
        candidate, output_kind="ppt", output_dir=tmp_path / "timing-notes"
    )
    presentation = Presentation(_artifact_map(result)["pptx"]["path"])
    for slide, row in zip(presentation.slides, candidate["slides"], strict=True):
        notes = slide.notes_slide.notes_text_frame.text
        assert f"建议用时：{row['minutes']} 分钟" in notes
        assert row["teacher_notes"] in notes
        visible = "\n".join(
            shape.text for shape in slide.shapes if shape.has_text_frame
        )
        assert "建议用时" not in visible
    assert candidate == original
