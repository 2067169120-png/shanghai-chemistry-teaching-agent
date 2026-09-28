"""Synthetic Word/PNG/WMF and temporary state only; no production writes."""

from __future__ import annotations

import hashlib
import io
import json
import os
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from PIL import Image, ImageDraw

from integrations.deeptutor_shchem_v1.desktop_word_question_attributes import _seal, suggest_attributes
from integrations.deeptutor_shchem_v1.offline_word_label_review import (
    OfflineWordLabelReviewError, OfflineWordLabelReviewService, _digest,
)
from integrations.deeptutor_shchem_v1.offline_word_image_label_review import (
    CANDIDATE_SCHEMA_VERSION, ROOT_REVIEW_SCHEMA_VERSION, RULE_REVISION,
    OfflineWordImageLabelReviewService,
)
from runtime.deeptutor_shchem import apply_offline_word_image_labels as cli
from staging.coordination.deeptutor_gateway.tests.test_desktop_word_auto_attributes import setup as setup
from staging.coordination.deeptutor_gateway.tests.test_word_metafile_preview import _wmf


def _sha(raw):
    return hashlib.sha256(raw).hexdigest()


def _ref(path):
    raw = path.read_bytes()
    return {"path": str(path.resolve()), "sha256": _sha(raw), "bytes": len(raw)}


def _json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    return _ref(path)


def _image(*, color="black"):
    image = Image.new("RGBA", (90, 45), (255, 255, 255, 0))
    ImageDraw.Draw(image).line((10, 10, 70, 35), fill=color, width=2)
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def _document(*, wmf=False, wrong_shape=False, extra_field=False, answer_image=False,
              invisible_outline=False, visible_vml_stroke=False, repeated_image=False, crop=False):
    document = Document()
    document.add_heading("合成课堂图像练习", 1)
    document.add_paragraph("【即学即练1】解释氧化还原反应及图中现象。")
    if wmf:
        package = document.part.package
        image_part = Part(PackURI("/word/media/fixture.wmf"), "image/x-wmf", _wmf(), package)
        image_id = document.part.relate_to(image_part, RT.IMAGE)
        # The route verifies the OOXML preview relation, never binary content.
        ole_part = Part(PackURI("/word/embeddings/fixture.bin"),
                        "application/vnd.openxmlformats-officedocument.oleObject",
                        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(504), package)
        ole_id = document.part.relate_to(ole_part, RT.OLE_OBJECT)
        shape = "different" if wrong_shape else "_x0000_i1001"
        invisible_vml = ('<v:stroke joinstyle="miter"/>' if invisible_outline or visible_vml_stroke else '')
        invisible_vml += ('<o:lock v:ext="edit" aspectratio="t"/><w10:anchorlock/>' if invisible_outline else '')
        stroked = 't' if visible_vml_stroke else 'f'
        xml = f'''<w:object xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
          xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"
          xmlns:w10="urn:schemas-microsoft-com:office:word"
          xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
          <v:shape id="_x0000_i1001" type="#_x0000_t75" style="width:40pt;height:20pt" stroked="{stroked}">
            <v:imagedata r:id="{image_id}"/>{invisible_vml}</v:shape>
          <o:OLEObject Type="Embed" ProgID="Equation.DSMT4" ShapeID="{shape}"
            DrawAspect="Content" r:id="{ole_id}"/></w:object>'''
        document.add_paragraph().add_run()._r.append(parse_xml(xml))
    else:
        document.add_picture(io.BytesIO(_image()))
        paragraph = document.paragraphs[-1]
        if invisible_outline:
            sppr = paragraph._p.xpath('.//pic:spPr')[0]
            sppr.append(parse_xml('<a:noFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'))
            sppr.append(parse_xml('<a:ln xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:noFill/></a:ln>'))
        if crop:
            fill = paragraph._p.xpath('.//pic:blipFill')[0]
            fill.append(parse_xml('<a:srcRect xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" l="50000"/>'))
        if repeated_image:
            paragraph._p.append(deepcopy(paragraph._p.xpath('./w:r')[0]))
    if extra_field:
        document.add_paragraph().add_run()._r.append(parse_xml(
            '<w:instrText xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"> EQ hidden </w:instrText>'
        ))
    document.add_paragraph("A．氧化 B．还原 C．两者均无 D．信息不足")
    document.add_paragraph("【答案】B")
    if answer_image:
        document.add_picture(io.BytesIO(_image(color="blue")))
    document.add_paragraph("【解析】合成答案文字，测试只验证证据约束。")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _corpus(setup, tmp_path, **options):
    words, facade, _source, _directory = setup
    source = facade.add("B2", _document(**options), source_id="IMAGE1")
    row = next(row for row in words.catalog()["items"] if row["batch_id"] == "B2")
    attr = suggest_attributes(row, {"source_name": "合成图像.docx"}, words._read_attribute_catalog())
    attr.update(primary_knowledge={"id": "unknown", "label": "待确认", "status": "unknown", "evidence": []},
                supporting_knowledge=[], curriculum_candidates=[], curriculum_status="pending_mapping",
                teacher_note="保留备注", rule_revision="synthetic-before-image-v1")
    old = words.attribute_store.save_many([_seal(attr)])[0]
    service = OfflineWordImageLabelReviewService(words)
    identity = {"key": row["key"], "question_revision": row["revision"], "source_sha256": row["source_sha256"],
                "expected_stored_attribute_revision": old["revision"]}
    folder = tmp_path / "image-evidence"
    folder.mkdir()
    return SimpleNamespace(words=words, facade=facade, row=row, old=old, source=source,
                           service=service, identity=identity, folder=folder)


def _review(corpus, *, white=False):
    inspected = corpus.service.describe_entry(corpus.identity)
    scope = inspected["reading_scope"]
    rows, inventory = corpus.words._resolve([{"key": corpus.row["key"], "revision": corpus.row["revision"]}], read_only=True)
    row = rows[0]
    source = inventory[row["source_id"]][0]
    document = Document(io.BytesIO(source.content))
    media, viewed, composites = [], [], []
    for obj in scope["objects"]:
        for relation in (obj["image"], obj["ole"]):
            if relation is None:
                continue
            original = corpus.folder / Path(relation["package_part"]).name
            original.write_bytes(document.part.rels[relation["relationship_id"]].target_part.blob)
            media.append({**relation, "block_index": obj["block_index"], "local_original_path": str(original)})
    for image in scope["images"]:
        data = corpus.words.reader.word_asset_bytes(source.content, image["asset_id"], render_metafiles=True,
                                                   expected_sha256=image["sha256"])["bytes"]
        display = corpus.folder / (image["asset_id"] + ".png")
        display.write_bytes(data)
        if white:
            with Image.open(io.BytesIO(data)) as raw_image:
                rgba = raw_image.convert("RGBA")
            white_image = Image.alpha_composite(Image.new("RGBA", rgba.size, "white"), rgba).convert("RGB")
            composite = corpus.folder / (image["asset_id"] + "-white.png")
            white_image.save(composite)
            composites.append({"original": _ref(display), "display": _ref(composite), "pixel_equality": True})
            display = composite
        viewed.append({**_ref(display), "viewed_by_root_model": True, "human_reviewed": False,
                       "observation": "合成测试的黑色线段；此声明只作为测试输入。"})
    reading = corpus.folder / "reading.md"
    reading.write_text("\n".join(block["text"] for group in scope["blocks"].values() for block in group), encoding="utf-8")
    evidence = {"question": row, "media": media, "document_xml_sha256": scope["document_xml_sha256"]}
    evidence_ref = _json(corpus.folder / "evidence.json", evidence)
    first = scope["images"][0]
    proof = {"block_index": first["block_index"], "quote": "合成图中黑色线段的模型观察",
             "image_sha256": viewed[0]["sha256"]}
    decoded = {"primary": {"id": "K11", "evidence": [proof]},
               "curriculum": [{"section_key": "S1", "evidence": [deepcopy(proof)]}],
               "note": "临时合成AI候选，非人工确认。"}
    item = {**corpus.identity, "source_evidence": evidence_ref, "reading_file": _ref(reading),
            "full_extracted_question_answer_context_text_read": True, "neighbor_text_read": True,
            "block_indices_read": {group: [block["index"] for block in blocks] for group, blocks in scope["blocks"].items()},
            "images_actually_viewed": viewed, "decision": "candidate", "candidate_only": True,
            "human_reviewed": False, "teacher_confirmed": False, "automatic_source_gate_cleared": False,
            "decoded_label_proposal_reviewed": deepcopy(decoded)}
    root = {"schema_version": ROOT_REVIEW_SCHEMA_VERSION, "candidate_only": True, "human_reviewed": False,
            "reviewer": "synthetic root model declaration", "white_composites_verified": composites, "items": [item]}
    root_path = corpus.folder / "root.json"
    corpus.batch = {"schema_version": CANDIDATE_SCHEMA_VERSION, "provenance": "codex_current_session",
                    "candidate_only": True, "human_review": False, "entries": [
                        {**corpus.identity, **inspected, "decoded": decoded, "root_review": _json(root_path, root)}]}
    corpus.root, corpus.root_path, corpus.evidence = root, root_path, evidence
    return corpus


@pytest.fixture
def image_corpus(setup, tmp_path):
    return _review(_corpus(setup, tmp_path))


def _update_root(corpus):
    corpus.batch["entries"][0]["root_review"] = _json(corpus.root_path, corpus.root)


@pytest.mark.parametrize("mode", ["missing_only", "recheck_automatic"])
def test_real_png_full_text_preview_and_atomic_apply(image_corpus, mode):
    c = image_corpus
    before = c.words.attribute_store.history(c.row["key"])
    state = c.words.state.snapshot()
    preview = c.service.preview(c.batch, mode=mode)
    assert preview == c.service.preview(c.batch, mode=mode)
    assert c.words.state.snapshot() == state and c.words.attribute_store.history(c.row["key"]) == before
    assert preview["human_review"] is preview["teacher_confirmed"] is preview["provider_invoked"] is False
    assert preview["changed_question_count"] == 1
    assert "合成图中黑色线段" not in json.dumps(preview, ensure_ascii=False)
    receipt = c.service.apply(c.batch, expected_plan_sha256=preview["plan_sha256"], mode=mode)
    assert receipt["commit_status"] == "committed" and receipt["readback_verified"] is True
    actual = c.words.attribute_store.get(c.row["key"])
    assert actual["primary_knowledge"]["status"] == "auto_suggested"
    assert actual["primary_knowledge"]["evidence"][0]["kind"] == "model_image_observation"
    assert actual["teacher_note"] == c.old["teacher_note"]
    assert actual["material_status"] == c.old["material_status"]
    assert actual["edit_version"] == c.old["edit_version"] + 1
    assert c.words.attribute_store.history(c.row["key"]) == [*before, actual]
    with pytest.raises(OfflineWordLabelReviewError) as error:
        c.service.apply(c.batch, expected_plan_sha256=preview["plan_sha256"], mode=mode)
    assert error.value.code == "stored_revision_changed"


def test_legacy_text_only_gate_still_refuses_png(image_corpus):
    c = image_corpus
    old_batch = deepcopy(c.batch)
    old_batch["schema_version"] = "shchem.offline-word-label-candidates.v1"
    for name in ("root_review", "reading_scope", "reading_scope_sha256"):
        old_batch["entries"][0].pop(name)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        OfflineWordLabelReviewService(c.words).preview(old_batch)
    assert error.value.code == "image_not_supported"


def test_white_png_composition_is_reproduced(setup, tmp_path):
    c = _review(_corpus(setup, tmp_path), white=True)
    assert c.service.preview(c.batch)["entries"][0]["white_display_bindings"]
    display = Path(c.root["white_composites_verified"][0]["display"]["path"])
    Image.new("RGB", (90, 45), "white").save(display)
    changed = _ref(display)
    c.root["white_composites_verified"][0]["display"] = changed
    c.root["items"][0]["images_actually_viewed"][0].update(changed)
    _update_root(c)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        c.service.preview(c.batch)
    assert error.value.code == "white_display_pixels_changed"


@pytest.mark.skipif(os.name != "nt", reason="The real WMF renderer requires Windows GDI+")
def test_real_wmf_same_object_preview_without_interpreting_ole(setup, tmp_path):
    c = _review(_corpus(setup, tmp_path, wmf=True, invisible_outline=True))
    preview = c.service.preview(c.batch)
    scope = c.batch["entries"][0]["reading_scope"]
    assert scope["objects"][0]["ole"]["binary_interpreted"] is False
    assert scope["images"][0]["display"]["derived_preview"] is True
    assert "wmf-font-preflight" in scope["images"][0]["display"]["renderer_revision"]
    assert preview["entries"][0]["uninterpreted_ole_with_same_object_preview_count"] == 1
    c.service.apply(c.batch, expected_plan_sha256=preview["plan_sha256"])


def test_ole_shape_mismatch_fails_before_render(setup, tmp_path):
    c = _corpus(setup, tmp_path, wmf=True, wrong_shape=True)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        c.service.describe_entry(c.identity)
    assert error.value.code == "ole_same_object_unproven"


@pytest.mark.parametrize("options", [{"crop": True}, {"wmf": True, "visible_vml_stroke": True}])
def test_crop_or_visible_vml_effect_is_not_ignored(setup, tmp_path, options):
    c = _corpus(setup, tmp_path, **options)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        c.service.describe_entry(c.identity)
    assert error.value.code == "transformed_image_not_supported"


def test_explicit_no_fill_and_repeated_same_image_keep_every_native_occurrence(setup, tmp_path):
    c = _review(_corpus(setup, tmp_path, invisible_outline=True, repeated_image=True))
    scope = c.batch["entries"][0]["reading_scope"]
    assert len(scope["objects"]) == 2 and len(scope["images"]) == 1
    assert c.service.preview(c.batch)["entries"][0]["native_object_count"] == 2
    c.evidence["media"].pop()
    c.root["items"][0]["source_evidence"] = _json(c.folder / "evidence.json", c.evidence)
    _update_root(c)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        c.service.preview(c.batch)
    assert error.value.code == "reviewed_object_inventory_changed"


def test_unread_field_cannot_be_covered_by_a_picture(setup, tmp_path):
    c = _corpus(setup, tmp_path, extra_field=True)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        c.service.describe_entry(c.identity)
    assert error.value.code == "unsupported_source_content"


@pytest.mark.parametrize("mutation,code", [
    ("no_view", "actual_image_read_required"), ("false_view", "actual_image_read_unbound"),
    ("held", "root_review_binding_mismatch"), ("human", "root_model_review_required"),
    ("answer_not_read", "root_review_binding_mismatch"), ("changed_decoded", "root_review_binding_mismatch"),
    ("missing_native_object", "image_reading_scope_changed"), ("changed_answer", "image_reading_scope_changed"),
])
def test_incomplete_or_forged_reading_declarations_refused(image_corpus, mutation, code):
    c = image_corpus
    item = c.root["items"][0]
    if mutation == "no_view": item["images_actually_viewed"] = []
    if mutation == "false_view": item["images_actually_viewed"][0]["viewed_by_root_model"] = False
    if mutation == "held": item["decision"] = "held"
    if mutation == "human": c.root["human_reviewed"] = True
    if mutation == "answer_not_read": item["block_indices_read"]["answer_blocks"] = []
    if mutation == "changed_decoded": c.batch["entries"][0]["decoded"]["note"] = "different proposal"
    if mutation == "missing_native_object": c.batch["entries"][0]["reading_scope"]["objects"] = []
    if mutation == "changed_answer": c.batch["entries"][0]["reading_scope"]["blocks"]["answer_blocks"][0]["text"] = "other answer"
    _update_root(c)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        c.service.preview(c.batch)
    assert error.value.code == code
    assert c.words.attribute_store.get(c.row["key"]) == c.old


@pytest.mark.parametrize("file", ["display", "original", "reading", "root", "source"])
def test_missing_or_tampered_bound_files_fail_without_write(image_corpus, file):
    c = image_corpus
    preview = c.service.preview(c.batch)
    paths = {"display": Path(c.root["items"][0]["images_actually_viewed"][0]["path"]),
             "original": Path(c.evidence["media"][0]["local_original_path"]),
             "reading": Path(c.root["items"][0]["reading_file"]["path"]),
             "root": c.root_path, "source": c.source}
    paths[file].write_bytes(b"changed")
    with pytest.raises(OfflineWordLabelReviewError) as error:
        c.service.apply(c.batch, expected_plan_sha256=preview["plan_sha256"])
    assert error.value.operation["attribute_write_attempted"] is False
    assert c.words.attribute_store.get(c.row["key"]) == c.old


def test_missing_display_file_is_not_a_reading_declaration(image_corpus):
    c = image_corpus
    Path(c.root["items"][0]["images_actually_viewed"][0]["path"]).unlink()
    with pytest.raises(OfflineWordLabelReviewError) as error:
        c.service.preview(c.batch)
    assert error.value.code == "invalid_evidence_file"


def test_complete_attached_context_is_bound_and_must_have_been_read(image_corpus):
    c = image_corpus
    row = c.row
    changed = c.words.update_range(row["key"], row["revision"], block_start=row["block_start"],
                                   question_end=row["question_end"], answer_start=row["answer_start"],
                                   block_end=row["block_end"], context_start=1, context_end=1)
    c.row = c.words._resolve([{"key": row["key"], "revision": changed["revision"]}], read_only=True)[0][0]
    attr = suggest_attributes(c.row, c.old["source"], c.words._read_attribute_catalog())
    attr.update(primary_knowledge=deepcopy(c.old["primary_knowledge"]), curriculum_candidates=[], curriculum_status="pending_mapping")
    c.old = c.words.attribute_store.save_many([_seal(attr)])[0]
    c.identity.update(question_revision=c.row["revision"], expected_stored_attribute_revision=c.old["revision"])
    _review(c)
    assert c.service.preview(c.batch)["changed_question_count"] == 1
    assert c.batch["entries"][0]["reading_scope"]["blocks"]["context_blocks"][0]["text"] == "合成课堂图像练习"
    c.root["items"][0]["block_indices_read"]["context_blocks"] = []
    _update_root(c)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        c.service.preview(c.batch)
    assert error.value.code == "root_review_binding_mismatch"


def test_changed_old_label_after_preview_fails_before_write(image_corpus):
    c = image_corpus
    preview = c.service.preview(c.batch)
    updated = c.words.attribute_store.save_many([_seal({**c.old, "teacher_note": "changed after preview"})])[0]
    with pytest.raises(OfflineWordLabelReviewError) as error:
        c.service.apply(c.batch, expected_plan_sha256=preview["plan_sha256"])
    assert error.value.code == "stored_revision_changed"
    assert error.value.operation["attribute_write_attempted"] is False
    assert c.words.attribute_store.get(c.row["key"]) == updated


def test_answer_image_is_required_even_when_unused_as_label_evidence(setup, tmp_path):
    c = _review(_corpus(setup, tmp_path, answer_image=True))
    assert c.service.preview(c.batch)["changed_question_count"] == 1
    c.root["items"][0]["images_actually_viewed"] = c.root["items"][0]["images_actually_viewed"][:1]
    _update_root(c)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        c.service.preview(c.batch)
    assert error.value.code == "image_not_actually_read"


def test_old_label_change_teacher_protection_and_fresh_noop(image_corpus, monkeypatch):
    c = image_corpus
    preview = c.service.preview(c.batch)
    c.service.apply(c.batch, expected_plan_sha256=preview["plan_sha256"])
    actual = c.words.attribute_store.get(c.row["key"])
    c.identity["expected_stored_attribute_revision"] = actual["revision"]
    _review(c)
    before = c.words.attribute_store.history(c.row["key"])
    preview = c.service.preview(c.batch)
    assert preview["changed_question_count"] == 0
    monkeypatch.setattr(c.words.attribute_store, "save_many", lambda *_a, **_k: pytest.fail("no-op cannot save"))
    assert c.service.apply(c.batch, expected_plan_sha256=preview["plan_sha256"])["readback_verified"] is True
    assert c.words.attribute_store.history(c.row["key"]) == before


@pytest.mark.parametrize("protection", ["teacher_modified", "teacher_confirmed", "pinned"])
def test_teacher_attributes_never_overwritten(image_corpus, protection):
    c = image_corpus
    changed = deepcopy(c.old)
    if protection == "teacher_modified": changed["annotation_source"] = "teacher_modified"
    if protection == "pinned": changed["rule_revision"] = "teacher-pinned-v1"
    if protection == "teacher_confirmed":
        changed["primary_knowledge"] = {"id": "K11", "label": "氧化还原", "status": "teacher_confirmed", "evidence": []}
    actual = c.words.attribute_store.save_many([_seal(changed)])[0]
    c.batch["entries"][0]["expected_stored_attribute_revision"] = actual["revision"]
    with pytest.raises(OfflineWordLabelReviewError) as error:
        c.service.preview(c.batch)
    assert error.value.code == "protected_attributes"
    assert c.words.attribute_store.get(c.row["key"]) == actual


@pytest.mark.parametrize("after_commit", [False, True])
def test_write_exception_preserves_unknown_commit_status(image_corpus, monkeypatch, after_commit):
    c = image_corpus
    plan = c.service.preview(c.batch)
    original = c.words.attribute_store.save_many
    def fail(rows, *, expected_revisions):
        assert expected_revisions == {c.old["key"]: c.old["revision"]}
        if after_commit: original(rows, expected_revisions=expected_revisions)
        raise OSError("synthetic write failure")
    monkeypatch.setattr(c.words.attribute_store, "save_many", fail)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        c.service.apply(c.batch, expected_plan_sha256=plan["plan_sha256"])
    assert error.value.operation["commit_status"] == "unknown"
    assert error.value.operation["do_not_retry_automatically"] is True


def test_readback_failure_reports_committed(image_corpus, monkeypatch):
    c = image_corpus
    plan = c.service.preview(c.batch)
    original = c.words.attribute_store.save_many
    def save(rows, **kwargs):
        result = original(rows, **kwargs)
        monkeypatch.setattr(c.words.attribute_store, "get_many", lambda _keys: {})
        return result
    monkeypatch.setattr(c.words.attribute_store, "save_many", save)
    with pytest.raises(OfflineWordLabelReviewError) as error:
        c.service.apply(c.batch, expected_plan_sha256=plan["plan_sha256"])
    assert error.value.code == "readback_mismatch"
    assert error.value.operation["commit_status"] == "committed"
    assert error.value.operation["do_not_retry_automatically"] is True


def test_cli_preview_default_and_nonoverwrite_report(image_corpus, tmp_path):
    c = image_corpus
    candidates = tmp_path / "image-candidates.json"
    _json(candidates, c.batch)
    args = SimpleNamespace(workspace=c.facade.paths.workspace_root, state=c.facade.paths.state_root,
                           candidates=candidates, report=tmp_path / "image-preview.json", apply=False,
                           expected_plan_sha256=None, label_mode="missing_only")
    result = cli.run(args, service_factory=lambda *_args: c.words)
    assert result["schema_version"] == "shchem.offline-word-image-label-plan.v1"
    assert result["commit_status"] == "not_attempted"
    with pytest.raises(OfflineWordLabelReviewError) as error:
        cli.run(args, service_factory=lambda *_args: pytest.fail("must not initialize"))
    assert error.value.code == "report_exists"
