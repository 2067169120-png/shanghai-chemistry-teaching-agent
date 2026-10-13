"""Exact compile read projections remain detached and independently rechecked."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from threading import RLock
from types import SimpleNamespace

import pytest
from docx import Document

from integrations.deeptutor_shchem_v1 import desktop_composition_sources as sources
from integrations.deeptutor_shchem_v1.desktop_forward_section_materials import _SectionStagingFacade
from integrations.deeptutor_shchem_v1.desktop_handout_material_handoff import _StagingFacade, _ReadOnlyReferenceWords
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import block, digest, resource
from integrations.deeptutor_shchem_v1.desktop_section_source_stage import SectionSourceStage
from integrations.deeptutor_shchem_v1.desktop_word_questions import WordQuestionService
from integrations.deeptutor_shchem_v1.desktop_visual_import_v2 import DesktopSourceFile
from integrations.deeptutor_shchem_v1.reader_cancellation import ReadCancelled, TaskCancellationEvent, read_cancel_scope


@pytest.mark.parametrize('domain', ['saved_batch', 'batch_manifest', 'word_catalog', 'knowledge_textbook_options'])
def test_new_domain_detached_hit_native_tail_and_memory_change(domain):
    value, calls = {'nested': ['first']}, []
    stage = SectionSourceStage(SimpleNamespace())
    def native():
        calls.append(stage._reads_enabled)
        return deepcopy(value)
    before = stage.read_memo(domain, {'scope': 'a', 'role': 'handout', 'revision': 'r'}, native)
    before['nested'][0] = 'caller edit'
    assert stage.read_memo(domain, {'scope': 'a', 'role': 'handout', 'revision': 'r'}, native) == value
    assert calls == [True]
    value['nested'][0] = 'late port change'
    with pytest.raises(ValueError, match='精确读取结果'):
        stage.finish()
    assert calls == [True, False]
    assert stage.closed and not stage.reads and not stage._read_keys


def test_new_memo_domains_have_typed_scope_and_independent_stage_lifetime():
    calls = []
    def native():
        calls.append(True)
        return {'value': 'source'}
    first = SectionSourceStage(SimpleNamespace())
    first.read_memo('descriptor', {'record_id': True}, native)
    first.read_memo('descriptor', {'record_id': 1}, native)
    first.finish()
    assert len(calls) == 4 and first.statistics['native_read_tail']['descriptor'] == 2
    with pytest.raises(ValueError):
        first.read_memo('descriptor', {'record_id': 1}, native)
    second = SectionSourceStage(SimpleNamespace())
    second.read_memo('descriptor', {'record_id': 1}, native)
    second.finish()
    assert len(calls) == 6


def test_cache_hit_checks_cancellation_and_close_clears_readers():
    stage = SectionSourceStage(SimpleNamespace())
    stage.read_memo('descriptor', {'batch': 'a'}, lambda: {'value': 'first'})
    event = TaskCancellationEvent(); event.set()
    with read_cancel_scope(event, check_boundaries=False), pytest.raises(ReadCancelled):
        stage.read_memo('descriptor', {'batch': 'a'}, lambda: {'value': 'second'})
    stage.close()
    assert not event._published and not stage.reads and not stage._read_keys


def test_saved_descriptor_and_manifest_share_nested_stage_but_rebind_rejects(tmp_path):
    calls, value = {'saved': 0, 'manifest': 0}, {'batch_id': 'A', 'sources': [{'role': 'handout'}]}
    class Port:
        paths = SimpleNamespace(task_root=tmp_path, workspace_root=tmp_path, state_root=tmp_path)
        _visual_import_root = tmp_path / 'original-profile'
        def _saved_visual_import_batch(self, batch):
            calls['saved'] += 1
            assert batch == 'A'
            return deepcopy(value)
        def _import_batch_manifest(self, descriptor):
            calls['manifest'] += 1
            assert descriptor == value
            return {'plan': deepcopy(descriptor)}, 'a' * 64
    port = Port(); staged = _StagingFacade(port, tmp_path / 'staged')
    stage = SectionSourceStage(staged, excluded_roots=(tmp_path,)); staged._section_source_stage = stage
    nested = _StagingFacade(staged, tmp_path / 'nested')
    for facade in (staged, nested, staged):
        descriptor = facade._saved_visual_import_batch('A')
        facade._import_batch_manifest(descriptor)
        descriptor['sources'][0]['role'] = 'caller changed'
    assert calls == {'saved': 1, 'manifest': 1}
    port._saved_visual_import_batch = lambda _batch: {'batch_id': 'A', 'sources': [{'role': 'answer'}]}
    with pytest.raises(ValueError, match='精确读取结果'):
        stage.finish()
    assert stage.closed and calls == {'saved': 1, 'manifest': 1}


@pytest.mark.parametrize('change', ['content', 'path'])
def test_descriptor_domain_tracks_memory_value_and_original_root_identity(tmp_path, change):
    value = {'batch_id': 'A', 'sources': [{'role': 'handout'}]}
    port = SimpleNamespace(paths=SimpleNamespace(task_root=tmp_path, workspace_root=tmp_path),
        _visual_import_root=tmp_path / 'original',
        _saved_visual_import_batch=lambda _batch: deepcopy(value))
    staged = _StagingFacade(port, tmp_path / 'staged')
    stage = SectionSourceStage(staged, excluded_roots=(tmp_path,)); staged._section_source_stage = stage
    staged._saved_visual_import_batch('A')
    if change == 'content': value['sources'][0]['role'] = 'answer'
    else: port._visual_import_root = tmp_path / 'different-profile'
    with pytest.raises(ValueError, match='精确读取结果'):
        stage.finish()


@pytest.mark.parametrize('change', ['semantic', 'method'])
def test_private_inventory_manifest_verification_seals_full_semantics_and_rebinds(tmp_path, change):
    calls, value = [], {'metadata': {'status': 'first'}, 'unused': ['large processing payload']}
    class Port:
        paths = SimpleNamespace(workspace_root=tmp_path, task_root=tmp_path)
        def _import_batch_manifest(self, _descriptor):
            calls.append('native')
            return deepcopy(value), 'a' * 64
    port = Port(); staged = _StagingFacade(port, tmp_path / 'staged')
    stage = SectionSourceStage(staged, excluded_roots=(tmp_path,)); staged._section_source_stage = stage
    first = staged._verify_batch_manifest({'batch_id': 'A', 'sources': []})
    assert first['verified'] and first['raw_revision'] == 'a' * 64 and 'unused' not in first
    first['verified'] = False
    assert staged._verify_batch_manifest({'batch_id': 'A', 'sources': []})['verified']
    assert calls == ['native']
    # The real manifest API continues returning every field independently.
    complete, revision = staged._import_batch_manifest({'batch_id': 'A', 'sources': []})
    assert complete == value and revision == 'a' * 64 and calls == ['native', 'native']
    if change == 'semantic': value['metadata']['status'] = 'late semantic change with same claimed raw revision'
    else: port._import_batch_manifest = lambda _descriptor: ({'different': 'port'}, 'a' * 64)
    with pytest.raises(ValueError, match='精确读取结果'): stage.finish()
    assert stage.closed and not stage.reads


def test_file_proof_keeps_native_read_inventory_for_new_memo(tmp_path):
    original = tmp_path / 'source.custom'
    original.write_bytes(b'full exact original')
    stage = SectionSourceStage(SimpleNamespace())
    stage.read_memo('descriptor', {'original': str(original)}, lambda: original.read_text())
    original.write_bytes(b'changed exact original')
    with pytest.raises(ValueError): stage.finish()
    assert str(original).lower() in {str(path).lower() for path, _proof in stage.files.files.values()}
    assert stage.closed


@pytest.fixture
def word_port(tmp_path):
    document = Document()
    document.add_paragraph('【即学即练1】合成题目：请选择完整的化学条件。')
    document.add_paragraph('【答案】合成解析保留在教师范围。')
    stream = BytesIO(); document.save(stream); content = stream.getvalue()
    counts, state = {'snapshot': 0, 'restore': 0}, {'drafts': {}}
    class State:
        def snapshot(self):
            counts['snapshot'] += 1
            return deepcopy(state)
    class Port:
        paths = SimpleNamespace(workspace_root=tmp_path, state_root=tmp_path / 'state', task_root=tmp_path / 'tasks')
        state_store = State()
        _visual_import_root = tmp_path / 'archive'
        def _saved_visual_import_batch(self, batch):
            assert batch in {'A', 'B'}
            return {'batch_id': batch, 'sources': [{'source_file_id': batch + '-source',
                'source_sha256': sha256(content).hexdigest(), 'filename': 'synthetic.docx',
                'role': 'answer' if batch == 'B' else 'handout'}]}
        def _restore_visual_import_sources(self, descriptor, *, source_ids):
            counts['restore'] += 1
            batch = descriptor['batch_id']
            assert source_ids == {batch + '-source'}
            return (DesktopSourceFile(role='answer' if batch == 'B' else 'handout', order_index=1,
                filename='synthetic.docx', mime_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                content=content, group_id=batch, source_file_id=batch + '-source'),)
        def list_imported_word_batches(self):
            raise AssertionError('Exact catalog cannot fall back to all imported sources')
        def _word_questions(self):
            return self.words
    port = Port(); port.words = WordQuestionService(port)
    port.words._lock = RLock()
    staged = _SectionStagingFacade(port, tmp_path / 'staged')
    stage = SectionSourceStage(staged, excluded_roots=(tmp_path,)); staged._section_source_stage = stage
    return SimpleNamespace(port=port, staged=staged, stage=stage, counts=counts, state=state)


def test_catalog_hit_restores_same_byte_archive_locations_and_selection_validation(word_port):
    f = word_port; words = f.staged._word_questions()
    left, inventory = words._catalog({'A': {'A-source'}}, read_only=True)
    assert left['items']
    row = left['items'][0]; chosen = {'key': row['key'], 'revision': row['revision'], 'points': 2}
    words._catalog({'B': {'B-source'}}, read_only=True)
    assert words._locations[row['key']] == ('B', 'B-source')
    again, _ = words._catalog({'A': {'A-source'}}, read_only=True)
    assert words._locations[row['key']] == ('A', 'A-source') and again == left
    assert f.counts['snapshot'] == 2
    selected, _ = words._resolve([chosen], read_only=True)
    assert selected[0]['points'] == 2 and selected[0]['archive_source_id'] == 'A-source'
    with pytest.raises(ValueError): words._resolve([{**chosen, 'revision': 'f' * 64}], read_only=True)
    with pytest.raises(ValueError): words._resolve([{**chosen, 'points': 0}], read_only=True)
    assert f.counts['snapshot'] == 2
    f.stage.finish()
    assert f.counts['snapshot'] == 4 and f.stage.statistics['native_read_tail']['word_catalog'] == 2
    assert f.port.words._locations == {} and f.port.words._cache == {}


@pytest.mark.parametrize('change', ['ranges', 'attributes', 'reader', 'known_issue'])
def test_catalog_tail_freshly_rechecks_state_attributes_and_native_extractor(word_port, change, monkeypatch):
    f = word_port; words = f.staged._word_questions()
    catalog, _ = words._catalog({'A': {'A-source'}}, read_only=True)
    row = catalog['items'][0]
    if change == 'ranges':
        from integrations.deeptutor_shchem_v1.desktop_word_questions import RANGES_DRAFT
        f.state['drafts'][RANGES_DRAFT] = {'ranges': {row['key']: {'source_revision': 'f' * 64, 'range': {}}}}
    elif change == 'attributes':
        f.port.words.attribute_store.get_many = lambda _keys: {row['key']: {
            'source_sha256': row['source_sha256'], 'question_revision': row['revision'], 'new': 'changed port'}}
    elif change == 'reader':
        native = f.port.words.reader.word_preview_bytes
        def changed(*args, **kwargs):
            preview = native(*args, **kwargs)
            preview['warnings'].append('late extraction port changed')
            return preview
        f.port.words.reader.word_preview_bytes = changed
    else:
        monkeypatch.setattr('integrations.deeptutor_shchem_v1.desktop_word_questions.source_quality_notes',
            lambda _root: {row['source_sha256']: {'source_revision': row['source_revision'], 'notes_revision': 'f' * 64,
                'issues': [{'id': 'late-issue', 'block_indices': [row['question_blocks'][0]['index']],
                           'summary': 'Late known source issue', 'suggested_correction': 'Hold source'}]}})
    with pytest.raises(ValueError, match='精确读取结果'):
        f.stage.finish()
    assert f.stage.closed and not f.stage.reads


def test_catalog_scope_includes_native_service_and_attribute_port(word_port):
    f = word_port; first = f.staged._word_questions()
    first_catalog, _ = first._catalog({'A': {'A-source'}}, read_only=True)
    row = first_catalog['items'][0]
    second_native = WordQuestionService(f.port)
    second_native.attribute_store.get_many = lambda _keys: {row['key']: {
        'source_sha256': row['source_sha256'], 'question_revision': row['revision'], 'different': 'profile attributes'}}
    second = _ReadOnlyReferenceWords(second_native, f.staged)
    second_catalog, _ = second._catalog({'A': {'A-source'}}, read_only=True)
    assert 'attributes' not in first_catalog['items'][0]
    assert second_catalog['items'][0]['attributes']['different'] == 'profile attributes'
    assert f.stage.statistics['read_checks']['word_catalog'] == 2
    f.stage.finish()
    assert f.stage.statistics['native_read_tail']['word_catalog'] == 2


def test_current_word_service_replacement_cannot_reuse_old_catalog_or_native_tail(word_port):
    f = word_port; words = f.staged._word_questions()
    words._catalog({'A': {'A-source'}}, read_only=True)
    replacement = WordQuestionService(f.port)
    f.port.words = replacement
    with pytest.raises(ValueError, match='来源端口'):
        f.staged._word_questions()
    with pytest.raises(ValueError, match='来源端口'):
        words._catalog({'A': {'A-source'}}, read_only=True)
    with pytest.raises(ValueError, match='来源端口'):
        f.stage.finish()
    assert f.stage.closed and not f.stage.reads and f.port.words is replacement


def test_native_catalog_manifest_discard_adapter_keeps_public_complete_contract(word_port):
    f = word_port; calls = []
    complete = {'native_processing_result': {'unused': ['complete payload preserved']}}
    def manifest(_descriptor):
        calls.append('native validated')
        return deepcopy(complete), 'a' * 64
    f.port._import_batch_manifest = manifest
    f.staged._word_questions()._catalog({'A': {'A-source'}}, read_only=True)
    cache = [value for (domain, _scope, _seal), (value, _reader) in f.stage.reads.items()
             if domain == 'batch_manifest']
    assert len(cache) == 1 and cache[0]['verified'] and 'native_processing_result' not in cache[0]
    assert calls == ['native validated']
    # Reading through the public contract still returns all original fields.
    public = f.staged._import_batch_manifest({'batch_id': 'A', 'sources': []})
    assert public == (complete, 'a' * 64)
    f.stage.finish()
    assert f.stage.statistics['native_read_tail']['batch_manifest'] == 2


def test_textbook_options_keep_original_binding_and_new_domain_native_tail(tmp_path, monkeypatch):
    card = resource('knowledge_note', 'Explicit source', [block('text', 'knowledge', 'Complete authored source')], {
        'source_path': 'memory:source', 'source_sha256': 'a' * 64, 'source_revision': 'b' * 64,
        'record_id': 'NOTE', 'pdf_pages': [1], 'printed_pages': 'unknown'})
    facade = SimpleNamespace(); stage = SectionSourceStage(facade); facade._section_source_stage = stage
    value, calls = {'rows': [{'selection': {'visual_asset_id': 'one', 'revision': 'r'}}], 'binding': 'b'}, []
    def native(_facade, chosen):
        calls.append(stage._reads_enabled)
        assert chosen['original'] == card['original']
        return deepcopy(value)
    monkeypatch.setattr(sources, '_knowledge_textbook_options_native', native)
    current = sources.knowledge_textbook_options(facade, card); current['rows'].clear()
    assert sources.knowledge_textbook_options(facade, card) == value and calls == [True]
    value['binding'] = 'changed source'
    with pytest.raises(ValueError, match='精确读取结果'): stage.finish()
    assert calls == [True, False]
