"""Complete authored steps, real measured Office outputs and one guarded UI commit."""
from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document
from PIL import Image
from pptx import Presentation
from PySide6.QtCore import QPoint,Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QDialog,QLabel,QVBoxLayout
from test_composition_render import app
from test_lesson_design import sample
from integrations.deeptutor_shchem_v1.desktop_composition_exposition import validate_exposition
from integrations.deeptutor_shchem_v1.desktop_composition_render import layout_composition,write_pptx
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    add_resource,block,enable_composition,new_composition,resource,validate_composition)
from integrations.deeptutor_shchem_v1.desktop_lesson_design import DesignHistory
from integrations.deeptutor_shchem_v1.desktop_lesson_design import digest
from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
from integrations.deeptutor_shchem_v1.desktop_workbench.lesson_composition_panel import LessonCompositionPanel,ResourceEditor
from integrations.deeptutor_shchem_v1.desktop_workbench.step_flow_dialog import (
    StepFlowDialog,step_flow_resource,step_flow_composition,step_flow_result)


def fill(dialog,rows):
    dialog.question_title.setText('怎样把假设转化为可验证的实验？')
    for index,(heading,description) in enumerate(rows):
        if index>=len(dialog.steps):
            dialog.step_list.setCurrentRow(len(dialog.steps)-1);dialog.add_button.click()
        dialog.step_list.setCurrentRow(index);dialog.step_title.setText(heading);dialog.description.setPlainText(description)
    dialog.update_preview()


@pytest.fixture
def desk(app,tmp_path):
    stream=BytesIO();Image.new('RGB',(160,90),'#467c63').save(stream,format='PNG');raw=stream.getvalue()
    store=PreparationImageStore(tmp_path/'images');asset=store.import_bytes(raw,'合成原图','合成测试作者','来源角色保留测试')
    image=resource('textbook_image','保留教材原图',[block('image','material',asset['asset_id'])],
        {'source_path':'合成教材.pdf','source_sha256':'a'*64,'source_revision':'b'*64,'pdf_pages':[3],'printed_pages':'unknown'})
    question=resource('complete_question','保留完整原题与答案',[
        block('text','question','合成原题完整条件',4),block('text','answer','合成教师答案',5)],
        {'source_path':'合成原题.docx','source_sha256':'c'*64,'source_revision':'d'*64,
         'question_key':'synthetic-flow-original','scope_review_sha256':'e'*64,'scope_status':'complete_native_scope',
         'hierarchy':dict.fromkeys(('paper_id','theme_big_question_id','printed_question_id','atomic_part_id'),'unknown')})
    value=new_composition('16:9')
    for row in (image,question):
        value=add_resource(value,row,{'page_id':value['pages'][0],'x':12,'y':12,'width':296,'height':156})
    payload=sample();payload['lesson_design']=enable_composition(payload['lesson_design'])
    payload['lesson_design']['composition']=value;payload['image_assets']=[asset]
    current=[deepcopy(payload['image_assets'])]
    images=SimpleNamespace(assets=lambda:deepcopy(current[0]),set_assets=lambda assets:current.__setitem__(0,deepcopy(assets)))
    owner=SimpleNamespace(history=DesignHistory(payload['lesson_design']),page=SimpleNamespace(image_assets_widget=images),
        facade=SimpleNamespace(preparation_image_bytes=lambda value:store.load(value)),status=QLabel())
    window=QDialog();layout=QVBoxLayout(window);panel=LessonCompositionPanel(owner);layout.addWidget(panel)
    def sync():panel.set_payload({**payload,'image_assets':current[0],'lesson_design':owner.history.value})
    owner.sync_page=sync;owner.travel=lambda redo=False:(owner.history.travel(redo),sync());sync()
    yield owner,panel,window,current
    window.close();window.deleteLater();app.processEvents()


@pytest.mark.parametrize('count',[2,5,6])
def test_all_supported_steps_preserve_authored_strings_and_unknown_provenance(count):
    rows=[{'title':f' {index+1} 完整标题 ','description':'说明前后空格与换行\n末尾完整标记  '} for index in range(count)]
    item=step_flow_resource('  导学问题？  ',rows,'process_vertical');validate_exposition(item)
    assert len(item['blocks'])==count
    assert [b['value'] for b in item['blocks']]==[r['title']+'\n'+r['description'] for r in rows]
    assert item['title']=='  导学问题？  ' and item['original']['blocks']==item['blocks']
    assert all(b['role']=='personal' and b['source_index'] is None for b in item['blocks'])
    assert item['original']['evidence']['source_sha256']==item['original']['evidence']['source_revision']=='unknown'


@pytest.mark.parametrize('count',[1,7])
def test_invalid_counts_rejected_without_recasting_source(count):
    with pytest.raises(ValueError):step_flow_resource('导学问题？',[{'title':'完整标题','description':''}]*count,'process')


def test_defaults_cancel_and_edit_controls_preserve_complete_order(app):
    dialog=StepFlowDialog();fill(dialog,[('观察完整条件','说明一完整保留'),('比较实验条件','说明二完整保留')])
    assert dialog.direction.currentData()=='process_vertical' and dialog.click_steps.isChecked()
    first=deepcopy(dialog.steps);dialog.step_list.setCurrentRow(1);dialog.add_button.click()
    dialog.step_title.setText('推导结论');dialog.description.setPlainText('末尾不能丢失\n保留第二行')
    dialog.up_button.click();assert [s['title'] for s in dialog.steps]==['观察完整条件','推导结论','比较实验条件']
    dialog.down_button.click();assert dialog.steps[:2]==first
    dialog.remove_button.click();assert dialog.steps==first
    dialog.reject();assert dialog.result_value is None;dialog.deleteLater()


def test_print_format_keeps_full_steps_without_requiring_classroom_switch(app):
    dialog=StepFlowDialog('A4');fill(dialog,[('观察','完整观察条件'),('解释','完整解释步骤')])
    assert not dialog.click_steps.isChecked() and not dialog.click_steps.isEnabled()
    assert dialog.apply_button.isEnabled();dialog.apply()
    assert 'click_reveals' not in dialog.result_value
    assert len(dialog.result_value['resources'][0]['blocks'])==2;dialog.deleteLater()


def test_six_step_bound_buttons_and_existing_resource_editor(app):
    dialog=StepFlowDialog();fill(dialog,[(f'步骤{i}','') for i in range(1,7)])
    assert len(dialog.steps)==6 and not dialog.add_button.isEnabled();dialog.add_step();assert len(dialog.steps)==6
    dialog.step_list.setCurrentRow(0);assert not dialog.up_button.isEnabled()
    dialog.step_list.setCurrentRow(5);assert not dialog.down_button.isEnabled()
    item=step_flow_resource('怎样完成六个步骤？',dialog.steps,'process')
    value=new_composition('16:9');value=add_resource(value,item,{'page_id':value['pages'][0],'x':12,'y':12,'width':296,'height':156})
    editor=ResourceEditor(value,item['id'],None)
    assert editor.exposition.findData('process')>=0 and editor.exposition.findData('process_vertical')>=0
    assert editor.exposition.findData('comparison')<0
    editor.apply();assert len(editor.result_value['resources'][0]['blocks'])==6
    editor.deleteLater();dialog.reject();dialog.deleteLater()


def test_overflow_keeps_all_steps_and_disables_apply(app):
    dialog=StepFlowDialog();rows=[(f'第{i}步完整标题','长说明与尾部条件。'*160) for i in range(1,6)]
    fill(dialog,rows);before=deepcopy(dialog.steps)
    assert not dialog.apply_button.isEnabled() and dialog._candidate is None
    dialog.apply();assert dialog.result_value is None and dialog.steps==before
    assert len(dialog.steps)==5 and '完整内容' in dialog.status.text()
    dialog.reject();dialog.deleteLater()


def test_narrow_dialog_buttons_and_full_step_editor_remain_reachable(app):
    dialog=StepFlowDialog();fill(dialog,[('观察','完整条件'),('解释','完整理由')]);dialog.resize(420,820);dialog.show()
    app.processEvents();QTest.qWait(20);app.processEvents()
    assert dialog.width()==420 and dialog.split.orientation()==Qt.Orientation.Vertical
    for button in (dialog.apply_button,dialog.cancel_button):
        assert button.isVisible() and dialog.rect().contains(button.mapTo(dialog,QPoint(0,0)))
        assert dialog.rect().contains(button.mapTo(dialog,button.rect().bottomRight()))
    dialog.controls_scroll.ensureWidgetVisible(dialog.description);app.processEvents()
    assert dialog.description.toPlainText()=='完整理由'
    dialog.reject();dialog.deleteLater()


def test_guard_rechecked_after_preview_measurement_and_before_apply(app,monkeypatch):
    from integrations.deeptutor_shchem_v1.desktop_workbench import step_flow_dialog as module
    alive=[True];checks=[]
    def guard():
        checks.append(alive[0])
        if not alive[0]:raise ValueError('当前作品已切换')
    dialog=StepFlowDialog(guard=guard);fill(dialog,[('观察',''),('解释','')]);assert dialog.apply_button.isEnabled()
    original=module.layout_composition
    def late(*args,**kwargs):
        result=original(*args,**kwargs);alive[0]=False;return result
    monkeypatch.setattr(module,'layout_composition',late);dialog.update_preview()
    assert dialog._candidate is None and not dialog.apply_button.isEnabled()
    dialog.apply();assert dialog.result_value is None and '已切换' in dialog.status.text()
    assert False in checks;dialog.reject();dialog.deleteLater()


def test_cancel_zero_changes_and_single_commit_undo_redo_preserve_original_roles(desk,monkeypatch):
    owner,panel,window,_current=desk;before=deepcopy(owner.history.value)
    monkeypatch.setattr(StepFlowDialog,'exec',lambda self:QDialog.DialogCode.Rejected)
    panel.step_flow_button.click();assert owner.history.value==before and not owner.history.back
    def accept(self):
        fill(self,[('观察完整条件',''),('比较作用',''),('得出结论','')]);self.apply();return self.result()
    monkeypatch.setattr(StepFlowDialog,'exec',accept);panel.step_flow_button.click()
    after=deepcopy(owner.history.value)
    assert len(owner.history.back)==1 and len(after['composition']['resources'])==3
    assert after['composition']['resources'][:2]==before['composition']['resources']
    item=after['composition']['resources'][-1]
    assert after['composition']['click_reveals']==[item['id']] and len(item['blocks'])==3
    owner.travel(False);assert owner.history.value==before
    owner.travel(True);assert owner.history.value==after


@pytest.mark.parametrize('target',['design','assets'])
def test_late_dialog_accept_cannot_replace_changed_design_or_assets(desk,monkeypatch,target):
    owner,panel,window,current=desk;expected=[]
    def accept_then_change(self):
        fill(self,[('观察',''),('解释','')]);self.apply()
        if target=='design':
            value=deepcopy(owner.history.value);value['nodes'][0]['teacher_action']='教师刚修改当前目标';owner.history.put(value)
        else:current[0][0]['caption']='刚修改的图片用途'
        owner.sync_page();expected.append((deepcopy(owner.history.value),deepcopy(current[0])))
        return self.result()
    monkeypatch.setattr(StepFlowDialog,'exec',accept_then_change);panel.step_flow_button.click()
    assert (owner.history.value,current[0])==expected[0]
    assert len(owner.history.value['composition']['resources'])==2 and '变化' in panel.status.text()


@pytest.mark.parametrize('target',['design','assets'])
def test_late_measurement_rechecks_target_before_any_asset_or_history_update(desk,monkeypatch,target):
    from integrations.deeptutor_shchem_v1.desktop_workbench import lesson_composition_panel as module
    owner,panel,window,current=desk
    request=step_flow_result('怎样完成？',[{'title':'观察','description':''},{'title':'解释','description':''}])
    before=digest(owner.history.value);assets=deepcopy(current[0]);expected=[]
    original=module.layout_composition
    def late(*args,**kwargs):
        rendered=original(*args,**kwargs)
        if target=='design':
            value=deepcopy(owner.history.value);value['nodes'][0]['teacher_action']='测量期间切换的目标';owner.history.put(value)
        else:current[0][0]['caption']='测量期间改写的图片说明'
        expected.append((deepcopy(owner.history.value),deepcopy(current[0])))
        return rendered
    monkeypatch.setattr(module,'layout_composition',late)
    with pytest.raises(ValueError,match='已变化'):panel.add_result(request,before,expected_assets=assets)
    assert (owner.history.value,current[0])==expected[0]
    assert len(owner.history.value['composition']['resources'])==2


@pytest.mark.parametrize('rows',[
    [('明确实验目的',''),('提出合理假设',''),('选择验证方法',''),('安排操作步骤',''),('分析结果与结论','')],
    # Complete b108 transcription was supplied and directly verified by root.
    # This remains teacher personal text; neither source SHA field is fabricated.
    [('明确目的原理','认真审题，明确实验目的和要求，明确实验原理'),
     ('提出合理假设','根据物质的性质提出合理的假设'),
     ('寻求验证方法','寻求能够证明或检验假设的方法'),
     ('设计操作步骤','根据设想的方法设计合理的操作方法与步骤'),
     ('分析得出结论','依据实验现象和数据，综合分析得出实验结论')]],ids=['synthetic-short','root-reviewed-b108-complete'])
def test_complete_five_step_native_click_order_and_editable_word_table(app,tmp_path,rows):
    dialog=StepFlowDialog();fill(dialog,rows);dialog.direction.setCurrentIndex(dialog.direction.findData('process'))
    dialog.update_preview();assert dialog.apply_button.isEnabled(),dialog.status.text();dialog.apply()
    item=dialog.result_value['resources'][0];value=step_flow_composition(dialog.result_value)
    layout=layout_composition(value,[],{},student=True,include_reveals=True);assert not layout['errors']
    target=tmp_path/'five-steps.pptx';write_pptx(layout,target,{})
    slide=Presentation(target).slides[0];groups=[shape for shape in slide.shapes if shape.shape_type==6]
    assert len(groups)==len(slide._element.xpath('.//p:cTn[@nodeType="clickEffect"]'))==5
    assert [''.join(shape.text for shape in group.shapes if shape.has_text_frame) for group in groups]==[title+body for title,body in rows]
    assert dialog.question_title.text() in ''.join(shape.text for shape in slide.shapes if shape.has_text_frame)
    targets=[int(node.get('spid')) for node in slide._element.xpath('.//p:cTn[@nodeType="clickEffect"]//p:spTgt')]
    assert targets==[group.shape_id for group in groups]
    body_runs=[run for group in groups for shape in group.shapes if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if not run.font.bold]
    if any(body for title,body in rows):assert body_runs and all(run.font.size.pt==27 for run in body_runs)
    p=sample();p['lesson_design']=enable_composition(p['lesson_design']);p['lesson_design']['composition']=value
    p['lesson_design']['classroom_notes']=True;before=deepcopy(p);write_documents(p,tmp_path,{})
    word=Document(tmp_path/'student_worksheet.docx');assert len(word.tables)==1
    assert [word.tables[0].cell(0,index*2).text for index in range(5)]==[title+('\n'+body if body else '') for title,body in rows]
    assert p==before;dialog.deleteLater()


def test_body_remains_27pt_in_real_powerpoint(app,tmp_path):
    dialog=StepFlowDialog();fill(dialog,[('观察','条件完整'),('解释','依据完整')]);dialog.apply()
    item=dialog.result_value['resources'][0];value=new_composition('16:9')
    value=add_resource(value,item,{'page_id':value['pages'][0],'x':12,'y':12,'width':296,'height':156})
    layout=layout_composition(value,[],{},student=False);assert not layout['errors']
    target=tmp_path/'body-27pt.pptx';write_pptx(layout,target,{})
    runs=[run for shape in Presentation(target).slides[0].shapes if shape.has_text_frame for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
    body=[run for run in runs if run.text in ('条件完整','依据完整')]
    assert len(body)==2 and all(run.font.size.pt==27 for run in body);dialog.deleteLater()


def test_production_handoff_frame_keeps_complete_flow_and_rejects_forged_placement(app):
    request=step_flow_result('怎样完成？',[{'title':'观察','description':'完整条件'},{'title':'解释','description':'完整依据'}])
    value=step_flow_composition(request);frame=value['placements'][request['resources'][0]['id']]
    assert {key:frame[key] for key in ('x','y','width','height')}=={'x':10,'y':10,'width':300,'height':160}
    assert not layout_composition(value,[],{},student=False)['errors']
    changed=deepcopy(request);changed['step_flow_placement']['frame']['height']=180
    with pytest.raises(ValueError):step_flow_composition(changed)
