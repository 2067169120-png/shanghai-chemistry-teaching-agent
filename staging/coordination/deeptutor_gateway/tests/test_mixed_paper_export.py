from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches
from PIL import Image

from integrations.deeptutor_shchem_v1.desktop_mixed_paper_export import (
    build_mixed_paper_docx,
)
from integrations.deeptutor_shchem_v1.desktop_word_question_export import (
    R,
    V,
    WordQuestionExportError,
    _prepare,
    _Writer,
)
from integrations.deeptutor_shchem_v1.paper_export_renderer import (
    PaperExportRendererError,
    _SyntheticDemoResolver,
    build_document_plans,
    build_render_request,
    build_synthetic_demo_bundle,
    run_export_preflight,
)


def _png(color="red"):
    stream = BytesIO()
    Image.new("RGB", (40, 20), color).save(stream, "PNG")
    return stream.getvalue()


def _word(
    key="甲", *, content=None, question=None, answer=None, context=None, points=3.5
):
    if content is None:
        doc = Document()
        doc.add_paragraph("共同材料：温度与反应。")  # 1
        paragraph = doc.add_paragraph(key + "题：填写化学式 __________")  # 2
        paragraph.add_run("SO")
        paragraph.add_run("4").font.subscript = True
        paragraph._p.append(
            parse_xml(f"<m:oMath {nsdecls('m')}><m:r><m:t>n/V</m:t></m:r></m:oMath>")
        )
        table = doc.add_table(rows=1, cols=2)  # 3
        table.cell(0, 0).text, table.cell(0, 1).text = key + "实验数据", "mol/L"
        doc.add_paragraph().add_run().add_picture(BytesIO(_png()), width=Inches(1))  # 4
        doc.add_paragraph("【答案】" + key + "答案仅教师可见")  # 5
        stream = BytesIO()
        doc.save(stream)
        content = stream.getvalue()
    return {
        "key": key,
        "revision": "test-v1",
        "source_name": key + "教案.docx",
        "source_bytes": content,
        "source_sha256": sha256(content).hexdigest(),
        "export_ready": True,
        "points": points,
        "question_blocks": [
            {"index": i} for i in ([2, 3, 4] if question is None else question)
        ],
        "answer_blocks": [{"index": i} for i in ([5] if answer is None else answer)],
        "context_blocks": [{"index": i} for i in ([1] if context is None else context)],
    }


def _doc(result, audience="student"):
    return Document(BytesIO(result[audience + "_bytes"]))


def _text(doc):
    return "\n".join(n.text or "" for n in doc.element.xpath(".//w:t"))


def _sections(*words):
    return [{"kind": "word_question", "question": word} for word in words]


def test_interleaving_preserves_full_word_content_and_core_dependency_order():
    bundle = build_synthetic_demo_bundle()
    first, second = _word(), _word("乙", points=4)
    sections = [
        _sections(first)[0],
        {"kind": "core_plan", "bundle": bundle},
        _sections(second)[0],
    ]
    frozen = deepcopy(sections)
    result = build_mixed_paper_docx("合成混合练习", sections, duration_minutes=80)
    student, teacher = _doc(result), _doc(result, "teacher")
    body, answer = _text(student), _text(teacher)
    assert body.index("甲题：") < body.index("某学习小组") < body.index("乙题：")
    assert body.count("共同材料：温度与反应。") == 2
    assert "甲答案仅教师可见" not in body and "乙答案仅教师可见" not in body
    assert (
        "甲题：" in answer
        and "甲答案仅教师可见" in answer
        and "乙答案仅教师可见" in answer
    )
    assert (
        "本题 3.5 分" in answer
        and "本题 4 分" in answer
        and "本次分值 17.5 分" in answer
    )
    assert "本次分值" not in body
    assert "80 分钟 · 3 个题组" in body
    assert body.count("填写化学式 __________") == 2
    assert [
        table.cell(0, 0).text
        for table in student.tables
        if table.cell(0, 0).text.endswith("实验数据")
    ] == ["甲实验数据", "乙实验数据"]
    assert len(student.element.xpath(".//m:oMath")) == 2
    assert len(student.element.xpath(".//w:vertAlign[@w:val='subscript']")) == 2
    with ZipFile(BytesIO(result["student_bytes"])) as archive:
        images = [
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/media/")
        ]
    assert _png() in images
    assert sections == frozen
    assert set(result) == {"student_bytes", "teacher_bytes", "warnings"}


def test_student_added_score_switch_does_not_change_teacher_scoring():
    sections = _sections(_word(points=6))
    for show in (False, True):
        result = build_mixed_paper_docx("练习", sections, show_student_scores=show)
        assert ("题组 1（6 分）" in _text(_doc(result))) is show
        assert "参考答案 · 本题 6 分" in _text(_doc(result, "teacher"))


def test_shapes_remain_unique_when_other_renderer_inserts_images_between_word_blocks():
    row = _word()
    item = _prepare([row])[0]
    document = Document()
    writer = _Writer(document)
    for _ in range(3):
        document.add_picture(BytesIO(_png("blue")), width=Inches(1))
        writer.append(item.source, (4,))
    ids = document.element.xpath(".//wp:docPr/@id")
    assert len(ids) == 6 and len(set(ids)) == 6


def test_native_vml_image_rectangles_get_distinct_ids_without_changing_geometry():
    document = Document()
    picture = (
        document.add_paragraph().add_run().add_picture(BytesIO(_png()), width=Inches(1))
    )
    rid = picture._inline.xpath(".//a:blip/@r:embed")[0]
    paragraph = document.paragraphs[0]
    paragraph._p.remove(paragraph.runs[0]._r)
    paragraph._p.append(
        parse_xml(
            f'<w:r {nsdecls("w", "r")} xmlns:v="{V}"><w:pict>'
            f'<v:rect id="original_rect" style="width:72pt;height:36pt">'
            f'<v:imagedata r:id="{rid}"/></v:rect></w:pict></w:r>'
        )
    )
    stream = BytesIO()
    document.save(stream)
    rows = [
        _word(key, content=stream.getvalue(), question=[1], answer=[], context=[])
        for key in ("甲", "乙")
    ]
    result = _doc(build_mixed_paper_docx("矩形图混排", _sections(*rows)))
    rectangles = list(result.element.iter(f"{{{V}}}rect"))
    assert len(rectangles) == 2 and len({rect.get("id") for rect in rectangles}) == 2
    assert all(rect.get("style") == "width:72pt;height:36pt" for rect in rectangles)
    assert all(
        result.part.rels[node.get(f"{{{R}}}id")].target_part.blob == _png()
        for node in result.element.iter(f"{{{V}}}imagedata")
    )


def _bundle_with_image():
    bundle = build_synthetic_demo_bundle()

    class Resolver(_SyntheticDemoResolver):
        def resolve_atomic_part(self, reference):
            value = super().resolve_atomic_part(reference)
            if reference["atomic_part_id"] == "SYNTH-CO2-A1":
                value["question_blocks"].append(
                    {
                        "block_type": "image",
                        "asset_ref": "diagram.png",
                        "alt_text_zh": "合成示意图",
                        "text_zh": None,
                    }
                )
            return value

    plans = build_document_plans(
        bundle["blueprint"],
        preset=bundle["preset"],
        content_resolver=Resolver(),
        paper_metadata={
            "title_zh": "合成图文测试",
            "subtitle_zh": "混合布局",
            "version_label_zh": "合成测试",
        },
    )
    bundle["student_plan"], bundle["teacher_plan"] = plans["student"], plans["teacher"]
    bundle["preflight_report"] = run_export_preflight(
        preset=bundle["preset"],
        blueprint=bundle["blueprint"],
        student_plan=bundle["student_plan"],
        teacher_plan=bundle["teacher_plan"],
    )
    bundle["render_request"] = build_render_request(
        preset=bundle["preset"],
        student_plan=bundle["student_plan"],
        teacher_plan=bundle["teacher_plan"],
        preflight_report=bundle["preflight_report"],
    )
    return bundle


def test_real_core_picture_and_word_images_keep_distinct_relationships(tmp_path):
    (tmp_path / "diagram.png").write_bytes(_png("blue"))
    sections = [
        _sections(_word())[0],
        {"kind": "core_plan", "bundle": _bundle_with_image(), "asset_root": tmp_path},
        _sections(_word("乙"))[0],
    ]
    result = build_mixed_paper_docx("图文混合测试", sections)
    for audience in ("student", "teacher"):
        document = _doc(result, audience)
        ids = document.element.xpath(".//wp:docPr/@id")
        assert len(ids) == len(set(ids)) == 3
        refs = document.element.xpath(".//a:blip/@r:embed")
        assert len(refs) == 3
        assert all(
            document.part.rels[ref].target_part.blob in {_png(), _png("blue")}
            for ref in refs
        )


def test_missing_core_image_fails_without_partial_outputs(tmp_path):
    with pytest.raises(PaperExportRendererError):
        build_mixed_paper_docx(
            "缺图测试",
            [
                {
                    "kind": "core_plan",
                    "bundle": _bundle_with_image(),
                    "asset_root": tmp_path,
                }
            ],
        )


def test_cross_selected_answer_overlap_is_rejected_even_with_interleaved_core():
    row = _word()
    row["answer_blocks"] = [{"index": 3}]
    row["question_blocks"] = [{"index": 2}]
    other = _word(
        "乙", content=row["source_bytes"], question=[3], answer=[], context=[]
    )
    with pytest.raises(WordQuestionExportError, match="另一道题的答案"):
        build_mixed_paper_docx(
            "练习",
            [
                _sections(row)[0],
                {"kind": "core_plan", "bundle": build_synthetic_demo_bundle()},
                _sections(other)[0],
            ],
        )


@pytest.mark.parametrize(
    "kwargs",
    [
        {"title": ""},
        {"sections": []},
        {"sections": [{"kind": "unknown"}]},
        {"duration_minutes": True},
        {"duration_minutes": 0},
        {"show_student_scores": 1},
        {"subtitle": "x" * 301},
    ],
)
def test_invalid_input_fails_closed(kwargs):
    values = {"title": "练习", "sections": _sections(_word())}
    values.update(kwargs)
    with pytest.raises(WordQuestionExportError):
        build_mixed_paper_docx(**values)


def test_stale_word_source_and_tampered_core_preflight_are_rejected():
    row = _word()
    row["source_bytes"] += b"changed"
    with pytest.raises(WordQuestionExportError, match="预览版本不一致"):
        build_mixed_paper_docx("练习", _sections(row))
    bundle = build_synthetic_demo_bundle()
    bundle["student_plan"]["visible"]["theme_sections"][0]["theme_score"] = 100
    with pytest.raises((PaperExportRendererError, ValueError)):
        build_mixed_paper_docx("练习", [{"kind": "core_plan", "bundle": bundle}])
