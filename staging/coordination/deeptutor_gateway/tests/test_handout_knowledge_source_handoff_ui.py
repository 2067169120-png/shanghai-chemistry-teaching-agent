"""Whole knowledge sections resume their real Word without inventing a picture."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from random import Random
from types import SimpleNamespace

from docx import Document
from PIL import Image
import pytest
from PySide6.QtCore import QPoint, QRect
from PySide6.QtWidgets import QDialog

from test_composition_ui import desk as desk, personal
from test_question_explorer_ui import window as window
from test_desktop_library_ui import _ManualBridge, qt_app as qt_app
from test_desktop_studio_ui import settle
from test_handout_knowledge import make_facade, request_for, word_bytes
from test_word_question_dialog import _Facade, _preparation_asset
from test_composition_resource_titles import question

from integrations.deeptutor_shchem_v1 import desktop_handout_images as pictures
from integrations.deeptutor_shchem_v1 import desktop_handout_sources as sources
from integrations.deeptutor_shchem_v1 import desktop_handout_questions as questions
from integrations.deeptutor_shchem_v1.desktop_handout_knowledge import import_handout_knowledge
from integrations.deeptutor_shchem_v1.desktop_lesson_design import digest
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
from integrations.deeptutor_shchem_v1.desktop_preparation_sources import PreparationSourcesService
from integrations.deeptutor_shchem_v1.desktop_workbench.handout_image_picker import HandoutImagePicker
from integrations.deeptutor_shchem_v1.desktop_workbench.word_question_dialog import WordQuestionDialog


@pytest.fixture
def knowledge_source(tmp_path):
    data = []
    for seed in (11, 22):
        buffer = BytesIO()
        Image.frombytes('RGB', (160, 80), Random(seed).randbytes(160 * 80 * 3)).save(buffer, 'PNG')
        data.append(buffer.getvalue())
    document = Document(); document.add_heading('知识点1 分子间作用力', 3)
    document.add_paragraph('完整原讲义知识段，保留条件与结论。')
    for raw in data: document.add_paragraph().add_run().add_picture(BytesIO(raw))
    facade = make_facade(word_bytes(document), tmp_path / 'original-source')
    service = PreparationSourcesService(facade.paths.workspace_root)
    def preview(*_args, read_only=False):
        assert read_only
        return service.word_preview_bytes(facade.value.data, facade.value.name)
    facade.imported_word_preview = preview
    pixels = {sha256(raw).hexdigest(): raw for raw in data}
    def asset(_batch, _source, identity):
        row = next(a for a in preview(read_only=True)['assets'] if a['asset_id'] == identity)
        return {'bytes': pixels[row['sha256']]}
    facade.imported_word_asset = asset
    request = request_for(facade)
    imported = import_handout_knowledge(facade, [request], [], [])
    return SimpleNamespace(facade=facade, binding=request['source'], request=request,
        imported=imported, pixels=pixels, original=facade.original.read_bytes())


def latest_preview(picker, bridge):
    pending = bridge.pending.pop()
    assert pending.label == '读取讲义原图'
    bridge.succeed(pending, run=True)
    assert picker.confirm.isEnabled()


def test_document_picker_focuses_actual_file_and_returns_real_picture(qt_app, knowledge_source):
    f = knowledge_source; snapshot = sources.handout_source_context(f.facade, f.binding)
    assert 'source' not in snapshot
    bridge = _ManualBridge(); picker = HandoutImagePicker(f.facade, snapshot, [], tasks=bridge)
    picker.show(); settle(qt_app)
    try:
        assert picker.document_source == f.binding and len(picker.document_source) == 8
        assert picker.source.currentData() == (f.binding['batch_id'], f.binding['source_id'])
        assert picker.list.count() == 2 and all(len(row) == 12 for row in picker.rows)
        latest_preview(picker, bridge)
        picker.list.setCurrentRow(1); latest_preview(picker, bridge)
        selected = deepcopy(picker.current); picker.choose()
        assert picker.selection['source'] == selected and selected == snapshot['images'][1]
        assert 'asset_id' not in f.binding and f.facade.original.read_bytes() == f.original
    finally:
        picker.reject(); picker.deleteLater(); settle(qt_app)


def test_document_without_images_keeps_exact_file_and_cannot_accept(qt_app, tmp_path):
    document = Document(); document.add_heading('知识点1 范德华力', 3)
    document.add_paragraph('只有完整文字，没有原图。')
    facade = make_facade(word_bytes(document), tmp_path / 'text-only')
    binding = request_for(facade)['source']; bridge = _ManualBridge()
    picker = HandoutImagePicker(facade, sources.handout_source_context(facade, binding), [], tasks=bridge)
    try:
        assert picker.document_source == binding and picker.source.currentIndex() == 1
        assert picker.source.currentText().endswith('（0图）') and not picker.rows
        assert picker.current is None and not picker.confirm.isEnabled() and not picker.read_button.isEnabled()
        assert '没有可选原图' in picker.status.text() and not bridge.pending
        picker.choose(); assert picker.selection is None and not facade.paths.task_root.exists()
    finally: picker.reject(); picker.deleteLater(); settle(qt_app)


@pytest.mark.parametrize('change', ['picture_as_document', 'missing_field', 'foreign_image'])
def test_document_snapshot_rejects_wrong_binding_or_cross_file_rows(qt_app, knowledge_source, change):
    f = knowledge_source; snapshot = sources.handout_source_context(f.facade, f.binding)
    if change == 'picture_as_document': snapshot['document_source'] = snapshot['images'][0]
    elif change == 'missing_field': snapshot['document_source'].pop('source_role')
    else: snapshot['images'][0]['source_id'] = 'OTHER'
    with pytest.raises(ValueError): HandoutImagePicker(f.facade, snapshot, [], tasks=_ManualBridge())
    assert f.facade.original.read_bytes() == f.original


def test_document_picker_keeps_teacher_picture_scope(qt_app, knowledge_source):
    f = knowledge_source; teacher_request = {**f.request, 'teacher_only': True}
    teacher = import_handout_knowledge(f.facade, [teacher_request], [], [])
    bridge = _ManualBridge()
    picker = HandoutImagePicker(f.facade, sources.handout_source_context(f.facade, f.binding),
        teacher['resources'], tasks=bridge)
    try:
        latest_preview(picker, bridge)
        assert picker.teacher_only.isChecked() and not picker.teacher_only.isEnabled()
        picker.choose(); assert picker.selection['teacher_only'] is True
    finally: picker.reject(); picker.deleteLater(); settle(qt_app)


def test_document_navigation_preserves_edits_and_ignores_cancelled_preview(qt_app, knowledge_source):
    f = knowledge_source; bridge = _ManualBridge()
    picker = HandoutImagePicker(f.facade, sources.handout_source_context(f.facade, f.binding), [], tasks=bridge)
    try:
        latest_preview(picker, bridge)
        picker.title.setText('这些分子之间怎样相互作用？'); picker.purpose.setPlainText('联系完整原文。')
        picker.list.setCurrentRow(1); obsolete = bridge.pending.pop()
        picker.list.setCurrentRow(0)
        bridge.succeed(obsolete, run=True)
        assert picker._previewed is None and not picker.confirm.isEnabled()
        latest_preview(picker, bridge)
        assert picker.title.text() == '这些分子之间怎样相互作用？'
        assert picker.purpose.toPlainText() == '联系完整原文。'
        picker.list.setCurrentRow(1); late = bridge.pending.pop(); picker.reject()
        bridge.succeed(late, run=True)
        assert picker.selection is None and picker._previewed is None
    finally: picker.reject(); picker.deleteLater(); settle(qt_app)


def setup_panel(desk, knowledge_source, monkeypatch):
    d, panel, app, page = desk; f = knowledge_source
    store = PreparationImageStore(f.facade.paths.task_root / 'preparation-v1/images')
    monkeypatch.setattr(d.facade, 'preparation_image_bytes', store.load)
    panel.add_result(f.imported, digest(d.history.value)); panel.refresh(); settle(app)
    assert panel.handout_source_button.isVisible() and panel.handout_question_button.isVisible()
    seen = []; context = sources.handout_source_context
    def same_document(_facade, binding):
        seen.append(deepcopy(binding)); return context(f.facade, binding)
    monkeypatch.setattr(sources, 'handout_source_context', same_document)
    preview = pictures.handout_image_preview
    monkeypatch.setattr(pictures, 'handout_image_preview', lambda _facade, row: preview(f.facade, row))
    monkeypatch.setattr(questions, 'handout_question_facade',
        lambda _facade, binding: seen.append(deepcopy(binding)) or _Facade())
    monkeypatch.setattr(d, 'run', lambda _title, operation, done, **kw: done(operation()))
    return d, panel, app, page, f, seen


@pytest.mark.parametrize('size', [(1400, 900), (1000, 850)])
def test_reading_and_same_document_actions_are_reachable_without_scrolling(desk, knowledge_source, monkeypatch, size):
    d, panel, app, page, _f, _seen = setup_panel(desk, knowledge_source, monkeypatch)
    before = deepcopy(page._payload())
    d.resize(*size); settle(app)
    panel.inspector_panel.verticalScrollBar().setValue(0); settle(app)
    viewport = panel.inspector_panel.viewport()
    for button in (panel.edit_button, panel.handout_source_button,
                   panel.handout_knowledge_button, panel.handout_question_button):
        visible_rect = QRect(button.mapTo(viewport, QPoint()), button.size())
        assert button.isVisible() and button.isEnabled()
        assert viewport.rect().contains(visible_rect), (button.text(), visible_rect, viewport.rect())
    assert page._payload() == before


@pytest.mark.parametrize('route', ['images', 'questions'])
def test_knowledge_entry_passes_all_eight_source_fields_and_cancel_keeps_work(desk, knowledge_source, monkeypatch, route):
    d, panel, app, page, f, seen = setup_panel(desk, knowledge_source, monkeypatch)
    edited = panel.composition(); edited['resources'][0]['title'] = '本课教师自拟问题？'; panel.commit(edited)
    before = deepcopy(d.history.value); assets = page.image_assets_widget.assets(); dialogs = []
    def cancel_images(picker):
        dialogs.append(deepcopy(picker.document_source)); return QDialog.DialogCode.Rejected
    def cancel_questions(dialog):
        dialogs.append((dialog._required_source_id, dialog._persist_choices, dialog._handoff_mode))
        dialog.reject(); return QDialog.DialogCode.Rejected
    monkeypatch.setattr(HandoutImagePicker, 'exec', cancel_images)
    monkeypatch.setattr(WordQuestionDialog, 'exec', cancel_questions)
    (panel.handout_source_button if route == 'images' else panel.handout_question_button).click(); settle(app)
    assert seen == [f.binding] and len(seen[0]) == 8 and 'asset_id' not in seen[0]
    assert dialogs == ([f.binding] if route == 'images' else [(f.binding['source_sha256'], False, 'composition')])
    assert d.history.value == before and page.image_assets_widget.assets() == assets


@pytest.mark.parametrize('route', ['images', 'questions'])
@pytest.mark.parametrize('change', ['draft', 'images', 'closing', 'stopping'])
def test_late_document_scope_does_not_open_after_work_or_lifecycle_changes(desk, knowledge_source, monkeypatch, route, change):
    d, panel, app, page, f, seen = setup_panel(desk, knowledge_source, monkeypatch)
    pending = []; opened = []
    monkeypatch.setattr(d, 'run', lambda _title, operation, done, **kw: pending.append((operation, done)))
    monkeypatch.setattr(HandoutImagePicker, 'exec', lambda _: opened.append(True))
    monkeypatch.setattr(WordQuestionDialog, 'exec', lambda _: opened.append(True))
    (panel.add_handout_from_selected if route == 'images' else panel.add_handout_questions)()
    snapshot = pending[0][0]()
    if change == 'draft': panel.add_result({'resources': [personal('选择后又编辑的内容')]}, digest(d.history.value))
    elif change == 'images': page.image_assets_widget.set_assets(page.image_assets_widget.assets() + [_preparation_asset(777)])
    elif change == 'closing': d._close_when_finished = True
    else: d._stopping = True
    before = deepcopy(d.history.value); assets = page.image_assets_widget.assets()
    pending[0][1](snapshot)
    assert not opened and d.history.value == before and page.image_assets_widget.assets() == assets
    d._close_when_finished = d._stopping = False


def changed_work(d, panel, page, change):
    if change == 'draft': panel.add_result({'resources': [personal('选择期间又编辑的内容')]}, digest(d.history.value))
    elif change == 'images': page.image_assets_widget.set_assets(page.image_assets_widget.assets() + [_preparation_asset(777)])
    elif change == 'closing': d._close_when_finished = True
    else: d._stopping = True


@pytest.mark.parametrize('change', ['draft', 'images', 'closing', 'stopping'])
def test_document_picture_confirmation_cannot_import_after_work_changes(desk, knowledge_source, monkeypatch, change):
    d, panel, app, page, f, _ = setup_panel(desk, knowledge_source, monkeypatch)
    imports = []; expected = []
    def accept(picker):
        picker.selection = {'source': deepcopy(picker.rows[1]), 'title': '为什么这样排列？',
            'purpose': '联系完整讲义知识段。', 'teacher_only': False}
        changed_work(d, panel, page, change)
        expected.append((deepcopy(d.history.value), page.image_assets_widget.assets()))
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(HandoutImagePicker, 'exec', accept)
    monkeypatch.setattr(pictures, 'import_handout_image', lambda *args: imports.append(True))
    panel.add_handout_from_selected()
    assert not imports and d.history.value == expected[0][0]
    assert page.image_assets_widget.assets() == expected[0][1]
    d._close_when_finished = d._stopping = False


@pytest.mark.parametrize('change', ['draft', 'images', 'closing', 'stopping'])
def test_late_document_picture_import_keeps_current_work(desk, knowledge_source, monkeypatch, change):
    d, panel, app, page, f, _ = setup_panel(desk, knowledge_source, monkeypatch)
    pending = []; imported = pictures.import_handout_image
    monkeypatch.setattr(d, 'run', lambda _title, operation, done, **kw: pending.append((operation, done)))
    monkeypatch.setattr(pictures, 'import_handout_image',
        lambda _facade, assets, resources, selection: imported(f.facade, assets, resources, selection))
    def accept(picker):
        picker.selection = {'source': deepcopy(picker.rows[1]), 'title': '为什么这样排列？',
            'purpose': '联系完整讲义知识段。', 'teacher_only': False}
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(HandoutImagePicker, 'exec', accept)
    panel.add_handout_from_selected(); pending[0][1](pending[0][0]())
    assert len(pending) == 2
    result = pending[1][0]()
    changed_work(d, panel, page, change)
    before = deepcopy(d.history.value); assets = page.image_assets_widget.assets()
    # The owner task wrapper reports this known error; a late result never commits.
    message = ('选材窗口已关闭或停止，本批未加入。' if change in {'closing', 'stopping'}
               else '选择讲义图期间变化')
    with pytest.raises(ValueError, match=message): pending[1][1](result)
    assert d.history.value == before and page.image_assets_widget.assets() == assets
    assert f.facade.original.read_bytes() == f.original
    d._close_when_finished = d._stopping = False


def test_document_picture_confirmation_is_one_undo_navigates_and_keeps_duplicate_edit(desk, knowledge_source, monkeypatch):
    d, panel, app, page, f, _ = setup_panel(desk, knowledge_source, monkeypatch)
    before = deepcopy(d.history.value); calls = []; imported = pictures.import_handout_image
    def accept(picker):
        picker.selection = {'source': deepcopy(picker.rows[1]), 'title': '怎样解释这幅结构图？',
            'purpose': '完整观察，再联系原文。', 'teacher_only': False}
        return QDialog.DialogCode.Accepted
    def import_real(_facade, assets, resources, selection):
        calls.append(deepcopy(selection)); return imported(f.facade, assets, resources, selection)
    monkeypatch.setattr(HandoutImagePicker, 'exec', accept)
    monkeypatch.setattr(pictures, 'import_handout_image', import_real)
    panel.add_handout_from_selected(); settle(app)
    current = deepcopy(d.history.value); added = panel.composition()['resources'][-1]
    assert added['original']['kind'] == 'handout_image' and len(calls) == 1
    assert added['original']['evidence']['handout_source'] == calls[0]['source']
    assert panel.composition()['placements'][added['id']]['page_id'] == panel.visible_page()
    d.travel(False); assert d.history.value == before
    d.travel(True); assert d.history.value == current
    edited = panel.composition(); edited['resources'][-1]['title'] = '教师保留的课堂标题？'; panel.commit(edited)
    saved = deepcopy(d.history.value); panel.selected = f.imported['resources'][0]['id']; panel.refresh()
    panel.add_handout_from_selected(); assert d.history.value == saved and '已有编辑保持' in d.status.text()


def test_document_complete_question_batch_is_one_undo_and_preserves_answer_and_duplicate(desk, knowledge_source, monkeypatch):
    d, panel, app, page, f, _ = setup_panel(desk, knowledge_source, monkeypatch)
    first, second = question(), question('15．比较分子之间的相互作用。')
    second['original']['evidence']['question_key'] = 'synthetic-second-question'
    second['original_sha256'] = digest(second['original'])
    incoming = {'resources': [first, second], 'image_assets': []}
    def accept(dialog):
        dialog.composition_result = deepcopy(incoming); dialog.reject(); return QDialog.DialogCode.Accepted
    monkeypatch.setattr(WordQuestionDialog, 'exec', accept)
    before = deepcopy(d.history.value); panel.add_handout_questions(); settle(app)
    after = deepcopy(d.history.value)
    assert len(panel.composition()['resources']) == 3
    selected_question = panel.composition()['resources'][1]
    assert panel.selected == selected_question['id']
    assert all(any(b['role'] == 'answer' for b in item['blocks']) for item in panel.composition()['resources'][1:])
    assert panel.composition()['placements'][selected_question['id']]['page_id'] == panel.visible_page()
    d.travel(False); assert d.history.value == before
    d.travel(True); assert d.history.value == after
    edited = panel.composition(); edited['resources'][1]['title'] = '已编辑的例题问题？'; panel.commit(edited)
    saved = deepcopy(d.history.value); panel.selected = f.imported['resources'][0]['id']; panel.refresh()
    panel.add_handout_questions(); assert d.history.value == saved and '已有编辑保持' in d.status.text()
