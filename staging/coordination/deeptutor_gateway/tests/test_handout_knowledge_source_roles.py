"""Visibility belongs to a verified source occurrence, not shared image bytes."""
from copy import deepcopy
from io import BytesIO
import json

from docx import Document
import pytest

from integrations.deeptutor_shchem_v1 import desktop_handout_knowledge as handout
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    block, digest, resource, validate_resource, visible_blocks,
)
from test_handout_knowledge import make_facade, png, request_for, word_bytes


def shared_picture_source(tmp_path, *, role='handout', different_source=False):
    document = Document()
    document.add_heading('知识点1 纤维的性能', 3)
    document.add_paragraph('纤维具有一定的强度和弹性。')
    document.add_paragraph().add_run().add_picture(BytesIO(png()))
    document.add_heading('知识点2 橡胶的性能', 3)
    document.add_paragraph('橡胶的弹性与结构有关。' if not different_source else '另一份来源保留完整原文。')
    paragraph = document.add_paragraph()
    paragraph.add_run().add_picture(BytesIO(png()))
    paragraph.add_run('同一图片的再次出现：').add_picture(BytesIO(png()))
    return make_facade(word_bytes(document), tmp_path, role=role)


@pytest.mark.parametrize('previous_use', ['same_source_other_section', 'other_source', 'standalone_image'])
def test_reused_teacher_bytes_do_not_hide_current_verified_knowledge(tmp_path, previous_use):
    first_facade = shared_picture_source(tmp_path / 'first')
    prior = handout.import_handout_knowledge(first_facade,
        [request_for(first_facade, 0, teacher_only=True)], [], [])
    old_resources, old_assets = prior['resources'], prior['image_assets']
    if previous_use == 'other_source':
        current = shared_picture_source(tmp_path / 'current', different_source=True)
        current.paths.task_root = first_facade.paths.task_root
    else:
        current = first_facade
    if previous_use == 'standalone_image':
        asset = old_assets[0]
        old_resources = [resource('preparation_image', '已有教师图片',
            [block('image', 'teacher_note', asset['asset_id'])],
            dict(source_path=asset['source'], source_sha256=asset['sha256'],
                source_revision=digest(asset), image_asset=asset, source_roles=[]))]
    before = deepcopy((old_resources, old_assets))
    result = handout.import_handout_knowledge(current,
        [request_for(current, 1, teacher_only=False)], old_assets, old_resources)
    item = result['resources'][0]
    assert all(row['role'] == 'knowledge' for row in item['blocks'])
    assert visible_blocks(item, student=True) == item['blocks']
    assert len([row for row in visible_blocks(item) if row['kind'] == 'image']) == 2
    assert result['image_assets'] == old_assets and len(old_assets) == 1
    assert (old_resources, old_assets) == before
    assert all(visible_blocks(old, student=True) == [] for old in old_resources)
    evidence = item['original']['evidence']
    assert evidence['source_roles'] == dict(schema=handout.SOURCE_ROLES_SCHEMA,
        teacher_only=False, document_role='handout',
        images=[dict(asset_id=old_assets[0]['asset_id'], roles=['knowledge'])])
    assert all(ref['source_sha256'] == evidence['source_sha256']
        and ref['source_role'] == 'handout' and ref['block_index'] >= 4
        for row in evidence['section_snapshot']['blocks'] for ref in row['images'])
    assert validate_resource(json.loads(json.dumps(item))) == item


@pytest.mark.parametrize('boundary', ['answer_source', 'teacher_only'])
def test_real_source_teacher_boundary_survives_reused_student_bytes(tmp_path, boundary):
    first = shared_picture_source(tmp_path / 'student')
    prior = handout.import_handout_knowledge(first, [request_for(first, 0)], [], [])
    current = shared_picture_source(tmp_path / 'teacher',
        role='answer' if boundary == 'answer_source' else 'handout')
    current.paths.task_root = first.paths.task_root
    result = handout.import_handout_knowledge(current,
        [request_for(current, 1, teacher_only=boundary == 'teacher_only')],
        prior['image_assets'], prior['resources'])
    item = result['resources'][0]
    assert all(row['role'] == 'teacher_note' for row in item['blocks'])
    assert visible_blocks(item, student=True) == []
    assert all(row['roles'] == ['teacher_note'] for row in item['original']['evidence']['source_roles']['images'])
    assert validate_resource(item) == item
    assert visible_blocks(prior['resources'][0], student=True)


def test_actual_answer_occurrence_with_reused_bytes_still_requires_complete_question(tmp_path):
    document = Document()
    document.add_heading('知识点1 纤维的性能', 3)
    document.add_paragraph('纤维具有一定的强度。')
    document.add_paragraph().add_run().add_picture(BytesIO(png()))
    document.add_heading('考向1 完整例题', 3)
    document.add_paragraph('【例1】下列纤维的性质正确的是（ ）')
    document.add_paragraph('A．甲 B．乙 C．丙 D．丁')
    document.add_paragraph('【答案】A')
    document.add_paragraph().add_run().add_picture(BytesIO(png()))
    facade = make_facade(word_bytes(document), tmp_path)
    prior = handout.import_handout_knowledge(facade, [request_for(facade, 0)], [], [])
    before = deepcopy(prior)
    with pytest.raises(ValueError, match='题目或答案'):
        handout.import_handout_knowledge(facade, [request_for(facade, 1)],
            prior['image_assets'], prior['resources'])
    assert prior == before


def test_same_batch_reuses_pixels_but_preserves_each_source_scope(tmp_path):
    facade = shared_picture_source(tmp_path)
    result = handout.import_handout_knowledge(facade, [
        request_for(facade, 0, teacher_only=True), request_for(facade, 1)], [], [])
    assert len(result['image_assets']) == 1
    assert not visible_blocks(result['resources'][0], student=True)
    assert visible_blocks(result['resources'][1], student=True) == result['resources'][1]['blocks']


@pytest.mark.parametrize('change', ['schema', 'record_role', 'document_role', 'occurrence_role', 'occurrence_source', 'duplicate_asset', 'missing_asset', 'teacher_flag'])
def test_source_role_evidence_rejects_partial_or_resealed_tampering(tmp_path, change):
    facade = shared_picture_source(tmp_path)
    item = handout.import_handout_knowledge(facade, [request_for(facade, 1)], [], [])['resources'][0]
    evidence = item['original']['evidence']; scope = evidence['source_roles']
    ref = next(ref for row in evidence['section_snapshot']['blocks'] for ref in row['images'])
    if change == 'schema': scope['schema'] = 'invented-scope'
    elif change == 'record_role': scope['images'][0]['roles'] = ['teacher_note']
    elif change == 'document_role': scope['document_role'] = 'answer'
    elif change == 'occurrence_role': ref['source_role'] = 'answer'
    elif change == 'occurrence_source': ref['source_sha256'] = 'a' * 64
    elif change == 'duplicate_asset': scope['images'].append(deepcopy(scope['images'][0]))
    elif change == 'missing_asset': scope['images'].clear()
    else: scope['teacher_only'] = True
    item['original_sha256'] = digest(item['original'])
    with pytest.raises(ValueError):
        validate_resource(item)


@pytest.mark.parametrize('boundary', ['teacher_only', 'answer_source'])
def test_current_teacher_source_cannot_be_resealed_as_visible_knowledge(tmp_path, boundary):
    facade = shared_picture_source(tmp_path, role='answer' if boundary == 'answer_source' else 'handout')
    item = handout.import_handout_knowledge(facade,
        [request_for(facade, 1, teacher_only=boundary == 'teacher_only')], [], [])['resources'][0]
    for row in item['original']['blocks']:
        row['role'] = 'knowledge'
    item['blocks'] = deepcopy(item['original']['blocks'])
    item['original_sha256'] = digest(item['original'])
    with pytest.raises(ValueError):
        validate_resource(item)


@pytest.mark.parametrize('legacy_teacher_image', [False, True])
def test_legacy_saved_resources_keep_their_frozen_role_semantics(tmp_path, legacy_teacher_image):
    facade = shared_picture_source(tmp_path)
    item = handout.import_handout_knowledge(facade, [request_for(facade, 1)], [], [])['resources'][0]
    scope = item['original']['evidence']['source_roles']
    scope.pop('schema')
    scope['images'][0]['roles'] = ['teacher_note'] if legacy_teacher_image else []
    if legacy_teacher_image:
        for row in item['original']['blocks']:
            row['role'] = 'teacher_note'
        item['blocks'] = deepcopy(item['original']['blocks'])
    item['original_sha256'] = digest(item['original'])
    frozen = deepcopy(item)
    assert validate_resource(item) == frozen and item == frozen
    assert bool(visible_blocks(item, student=True)) is not legacy_teacher_image
    if legacy_teacher_image:
        # A version marker alone must not reinterpret a frozen teacher snapshot.
        scope['schema'] = handout.SOURCE_ROLES_SCHEMA
        item['original_sha256'] = digest(item['original'])
        with pytest.raises(ValueError):
            validate_resource(item)
