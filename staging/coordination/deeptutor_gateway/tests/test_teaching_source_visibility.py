"""Teacher deliverables use real materials while evidence stays in source records."""
from copy import deepcopy
from pathlib import Path
import json

from docx import Document
from pptx import Presentation
from test_desktop_preparation import _payload
from test_desktop_preparation_images import _import, _image_candidate
from test_desktop_preparation_renderer import _artifact_map
from test_lesson_design import sample
from test_composition_render import app as render_app
from integrations.deeptutor_shchem_v1.desktop_preparation import normalize_preparation_candidate
from integrations.deeptutor_shchem_v1.desktop_preparation_images import PreparationImageStore
from integrations.deeptutor_shchem_v1.desktop_preparation_renderer import NativePreparationRenderer
from integrations.deeptutor_shchem_v1.desktop_lesson_composition import (
    enable_composition, add_resource, resource, block)
from integrations.deeptutor_shchem_v1.desktop_lesson_output import write_documents


def test_generated_image_slide_has_material_and_task_without_source_or_defensive_notes(tmp_path):
    store = PreparationImageStore(tmp_path/'images')
    original, asset = _import(tmp_path, store)
    asset['source'] = 'PRIVATE_SOURCE_SENTINEL：原页与内部SHA记录'
    candidate = normalize_preparation_candidate(_image_candidate(asset), {**_payload(), 'image_assets':[asset]})
    before = deepcopy(candidate)
    artifacts = _artifact_map(NativePreparationRenderer().render(candidate, output_kind='joint',
        output_dir=tmp_path/'native', image_data={asset['asset_id']:store.load(asset)}))
    deck = Presentation(artifacts['pptx']['path'])
    slide = deck.slides[1]
    body = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
    notes = slide.notes_slide.notes_text_frame.text
    assert any(s.shape_type == 13 and s.image.blob == original.read_bytes() for s in slide.shapes)
    assert asset['caption'] in body and candidate['slides'][1]['image']['observation_prompt'] in body
    assert asset['purpose'] in notes
    word = Document(artifacts['lesson_plan_docx']['path'])
    word_text = '\n'.join([p.text for p in word.paragraphs]+[c.text for t in word.tables for r in t.rows for c in r.cells])
    assert not any(marker in body+notes+word_text for marker in (
        'PRIVATE_SOURCE_SENTINEL','图片来源：','图片内容需教师复核','未调用模型/未查看像素'))
    stored = json.loads(Path(artifacts['candidate_json']['path']).read_text(encoding='utf-8'))
    assert candidate == before == stored
    assert stored['image_assets'][0]['source'] == asset['source']
    assert stored['candidate_only'] is True and stored['publication_allowed'] is False


def test_composition_word_shows_original_image_caption_and_question_without_audit_citations(tmp_path, render_app):
    store = PreparationImageStore(tmp_path/'images')
    _, asset = _import(tmp_path, store)
    asset['source'] = 'PRIVATE_IMAGE_SOURCE_SENTINEL'
    payload = sample(); payload['image_assets'] = [asset]
    payload['lesson_design'] = enable_composition(payload['lesson_design'])
    node = payload['lesson_design']['nodes'][0]
    node['image_ids'] = [asset['asset_id']]
    composition = payload['lesson_design']['composition']
    image = resource('textbook_image','读图任务',[
        block('image','material',asset['asset_id']), block('text','material','找出曲线上的异常点。')],
        dict(source_path='PRIVATE_BOOK_SOURCE_SENTINEL.pdf',source_sha256='a'*64,source_revision='b'*64,
             pdf_pages=[55],printed_pages='unknown'))
    composition = add_resource(composition,image,dict(page_id=composition['pages'][0],x=12,y=12,width=186,height=100))
    question = resource('complete_question','独立作答',[
        block('text','question','完整题面：依据结构判断分子间作用力。',3),
        block('text','answer','教师原答案。',4)],
        dict(source_path='PRIVATE_HANDOUT_SOURCE_SENTINEL.docx',source_sha256='c'*64,source_revision='d'*64,
             question_key='question',scope_review_sha256='e'*64,scope_status='complete_native_scope',
             hierarchy=dict.fromkeys(('paper_id','theme_big_question_id','printed_question_id','atomic_part_id'),'unknown')))
    composition = add_resource(composition,question,dict(page_id=composition['pages'][0],x=12,y=118,width=186,height=120))
    payload['lesson_design']['composition'] = composition
    before = deepcopy(payload)
    output = tmp_path/'delivery'; output.mkdir()
    write_documents(payload,output,{asset['asset_id']:store.load(asset)})
    for name in ('lesson_plan.docx','student_worksheet.docx'):
        doc = Document(output/name)
        text = '\n'.join(p.text for p in doc.paragraphs)
        assert '完整题面：' in text and len(doc.inline_shapes) > 0
        assert not any(marker in text for marker in ('PRIVATE_IMAGE_SOURCE_SENTINEL',
            'PRIVATE_BOOK_SOURCE_SENTINEL','PRIVATE_HANDOUT_SOURCE_SENTINEL','来源：'))
    assert payload == before
    assert payload['lesson_design']['composition']['resources'][1]['original']['evidence']['source_sha256'] == 'c'*64
