"""Isolated authored sample for UI tests; never install it in a teacher's library."""
from __future__ import annotations

import json
from pathlib import Path
from docx import Document


def seed_demo(workspace: Path, state: Path):
    from integrations.deeptutor_shchem_v1.desktop_facade import build_default_facade
    from integrations.deeptutor_shchem_v1.desktop_paths import DesktopPaths

    (workspace / "integrations/deeptutor_shchem_v1").mkdir(parents=True)
    kb = workspace / "sh-chem-db/kb"
    kb.mkdir(parents=True)
    knowledge = [{"id": "K09", "name": "化学平衡"}, {"id": "K10", "name": "电离与离子反应"},
                 {"id": "K01", "name": "物质分类"}, {"id": "K05", "name": "原子结构"}]
    (kb / "knowledge_taxonomy.json").write_text(json.dumps({"dimensions": {"knowledge_points": knowledge}}, ensure_ascii=False), encoding="utf-8")
    folder = kb / "classification/supplemental_wechat_textbook_tagging_v1_2026-08-27"
    folder.mkdir(parents=True)
    nodes = [{"node_key": "DEMO-EQ", "volume_id": "DEMO-BOOK", "volume_title": "化学原理（合成目录）",
              "chapter_id": "DEMO-C1", "chapter_title": "化学平衡", "section_title": "平衡状态的判断", "section_number": "1.1"},
             {"node_key": "DEMO-ION", "volume_id": "DEMO-BOOK", "volume_title": "化学原理（合成目录）",
              "chapter_id": "DEMO-C2", "chapter_title": "水溶液中的反应", "section_title": "电离与离子反应", "section_number": "2.1"}]
    (folder / "textbook_directory_nodes.json").write_text(json.dumps({"nodes": nodes}, ensure_ascii=False), encoding="utf-8")
    facade = build_default_facade(DesktopPaths.from_workspace(workspace, state_root=state))
    sources = []
    for exam, title in (("second_mock", "二模筛选演示"), ("school_exam", "校考筛选演示")):
        doc = Document()
        doc.add_heading(f"选题中心 · {title}（合成内容，非真实试卷）", 0)
        doc.add_heading("化学平衡", 1)
        doc.add_paragraph("【例1】恒温密闭容器中存在可逆反应 A(g) ⇌ B(g)。下列说法正确的是（　）")
        doc.add_paragraph("A．平衡时反应停止　B．平衡时正、逆反应速率相等且不为零")
        doc.add_paragraph("C．平衡时A和B的浓度必定相等　D．加入催化剂一定改变平衡组成")
        doc.add_paragraph("【答案】B")
        doc.add_paragraph("【解析】动态平衡不是反应停止；同一反应的正逆速率相等。")
        doc.add_paragraph("【例2】对于 A(g) ⇌ B(g)，某温度下平衡时 c(A)=0.20 mol/L，c(B)=0.60 mol/L。计算该反应的平衡常数。")
        doc.add_paragraph("答：____________________________________________________")
        doc.add_paragraph("【答案】K=c(B)/c(A)=3.0。")
        doc.add_heading("电离与离子反应", 1)
        doc.add_paragraph("【例3】下列物质中属于电解质的是（　）")
        doc.add_paragraph("A．铜　B．氯化钠　C．乙醇　D．蔗糖")
        doc.add_paragraph("【答案】B")
        doc.add_paragraph("【解析】氯化钠在熔融状态下能够导电，属于电解质。")
        path = workspace / (title + ".docx")
        doc.save(path)
        facade.save_visual_import_batch(handout_files=(path,), source_type="教师讲义")
        sources.append((path.name, exam))
    exams = dict(sources)
    for row in facade.word_question_catalog()["items"]:
        options = facade.word_question_attribute_options(row["key"], row["revision"])
        attrs = options["attributes"]
        ion = "电解质" in str(row.get("question_blocks"))
        node = nodes[1 if ion else 0]
        point = knowledge[1 if ion else 0]
        updates = {"primary_knowledge": {"id": point["id"], "label": point["name"], "status": "teacher_confirmed", "evidence": []},
                   "applicable_grades": {**attrs["applicable_grades"], "values": ["grade_12"], "status": "teacher_confirmed", "evidence": []},
                   "original_source": {**attrs["original_source"], "exam_type": {"value": exams[row["source_name"]], "status": "teacher_confirmed", "evidence": []}},
                   "curriculum_status": "teacher_confirmed", "curriculum_candidates": [{"section_key": node["node_key"], "volume_id": node["volume_id"], "chapter_id": node["chapter_id"], "label": node["section_title"], "status": "teacher_confirmed", "evidence": []}],
                   "teacher_note": "自动化界面测试合成标签，不代表真实考试、教材映射或教师审核。"}
        facade.word_question_save_attributes(row["key"], row["revision"], updates,
            expected_attribute_revision=attrs["revision"], expected_stored_revision=options["stored_revision"])
    return facade
