"""Generated Word fixtures in caller-owned temporary directories only."""
from pathlib import Path
import json
from docx import Document

from integrations.deeptutor_shchem_v1.desktop_facade import DesktopWorkbenchFacade
from integrations.deeptutor_shchem_v1.desktop_paths import DesktopPaths
from integrations.deeptutor_shchem_v1.desktop_state import DesktopStateStore


class NoProviders:
    calls = 0

    def list_metadata(self):
        return []

    def borrow_invocation_context(self, *args, **kwargs):
        self.calls += 1
        raise AssertionError("This synthetic acceptance must not call a provider")


def seed(root):
    root = Path(root)
    workspace = root / "workspace"
    (workspace / "integrations" / "deeptutor_shchem_v1").mkdir(parents=True)
    taxonomy = workspace / "sh-chem-db" / "kb"
    taxonomy.mkdir(parents=True)
    (taxonomy / "knowledge_taxonomy.json").write_text(json.dumps({"dimensions": {
        "knowledge_points": [{"id": "K03", "name": "合成标签：晶体类型"}]}}), encoding="utf-8")
    nodes = taxonomy / "classification" / "supplemental_wechat_textbook_tagging_v1_2026-08-27"
    nodes.mkdir(parents=True)
    (nodes / "textbook_directory_nodes.json").write_text('{"nodes": []}', encoding="utf-8")
    paths = DesktopPaths.from_workspace(workspace, state_root=root / "state")
    providers, unused = NoProviders(), object()
    facade = DesktopWorkbenchFacade(paths, theme_reader=unused, supplemental_reader=unused,
        curriculum_reader=unused, search_reader=unused, provider_store=providers,
        state_store=DesktopStateStore(paths.state_root), paper_export_jobs=unused)
    source = root / "合成晶体类型来源.docx"
    doc = Document()
    doc.add_heading("晶体类型 · 合成软件验收", 0)
    doc.add_paragraph("【例1】合成材料甲：请比较记录甲与记录乙，在下方填写用于判断晶体类型的证据。")
    doc.add_paragraph("记录甲：待核对。记录乙：待核对。本页仅检查题面、范围与答题区的保存。")
    doc.add_paragraph("【答案】合成答案甲：教师版专用文字；不代表化学判断或官方评分。")
    doc.add_paragraph("【例2】合成材料乙：将两组观测记录分别填入表格，说明还缺少哪些判断依据。")
    doc.add_paragraph("观测记录：尚未提供。本题仅供软件版式验收。")
    doc.add_paragraph("【答案】合成答案乙：待核验数据来源与完整性，不补造结论。")
    doc.save(source)
    facade.save_visual_import_batch(handout_files=(source,), source_type="教师讲义")
    rows = facade.word_question_catalog()["items"]
    assert len(rows) == 2
    for row in rows:
        facade.add_word_questions_to_basket([{"key": row["key"], "revision": row["revision"], "points": 3}])
    assert providers.calls == 0
    return facade, source, rows, providers
