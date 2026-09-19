"""Windows synthetic task-set acceptance with real LibreOffice pagination."""
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))


def main(destination):
    from copy import deepcopy
    from datetime import date
    import hashlib
    import json
    import tempfile
    import time
    import shutil
    from integrations.deeptutor_shchem_v1.desktop_workbench.app import create_application
    from integrations.deeptutor_shchem_v1.desktop_workbench.exam_dashboard import ExamDashboard
    from integrations.deeptutor_shchem_v1.desktop_workbench.tasks import DesktopTaskBridge
    from integrations.deeptutor_shchem_v1.desktop_review_fixture import seed_review
    from integrations.deeptutor_shchem_v1.desktop_closure_fixture import seed_word
    from integrations.deeptutor_shchem_v1.desktop_exam_fixture import example
    from integrations.deeptutor_shchem_v1.desktop_exam_followup import create_followup
    from integrations.deeptutor_shchem_v1.desktop_exam_practice import freeze_selection
    from integrations.deeptutor_shchem_v1.desktop_mixed_paper_service import _ACTIVE
    from docx import Document

    output = Path(destination); output.mkdir(parents=True, exist_ok=True)
    app = create_application(['independent-practice-probe'])
    errors = []; previous_hook = sys.excepthook
    sys.excepthook = lambda typ, value, tb: errors.append(str(value))
    def settle(test=lambda: True):
        for _ in range(8): app.processEvents(); time.sleep(.025)
        limit = time.monotonic() + 150
        while not test() and time.monotonic() < limit:
            app.processEvents(); time.sleep(.03)
        assert test(), 'Task operation failed or did not finish'
    with tempfile.TemporaryDirectory(prefix='exam-practice-') as directory:
        temp = Path(directory)
        f, _, _, transport, _ = seed_review(ROOT, temp/'state', temp/'sources', shared_paper=True)
        bridge = DesktopTaskBridge(); d = None
        try:
            path, rows = seed_word(f, temp/'word'); source = path.read_bytes(); calls = transport.calls
            for row in rows:
                f.add_word_questions_to_basket([{'key': row['key'], 'revision': row['revision'], 'points': 2}])
            basket = f.basket(); chosen = basket[0]['key']
            exam = example(temp/'synthetic.xlsx')[2]
            task = create_followup(exam, '1', ['S0001'], '核对动态平衡的判断依据', date.today().isoformat())
            task = freeze_selection(task, basket, [chosen])
            d = ExamDashboard(f, bridge); d.accept_exam(exam)
            assert d.followup_panel.store(task)
            f.state_store.remove_basket_item(chosen)
            assert len(f.basket()) == 1 and f.basket()[0]['key'] != chosen
            remaining = deepcopy(f.basket())
            sentinel = {'preview_id': 'unrelated-global', 'preview_hash': 'unchanged'}
            f.state_store.save_draft(_ACTIVE, sentinel)
            f.state_store.save_draft('paper-current', {'title': '保留教师原组卷草稿'})
            d.open_saved(exam['id'], task['id']); d.show(); d.resize(1280, 860)
            panel = d.followup_panel; d.tabs.setCurrentWidget(panel); settle()
            assert panel.current()['practice_set'] == task['practice_set']
            assert d.grab().save(str(output/'task-set-reopened.png'))
            panel.preview_button.click(); settle(lambda: d._task is None)
            assert panel.preview is not None, d.status.text()
            viewer = panel._preview_dialog
            for index, audience in enumerate(('student', 'teacher')):
                viewer.tabs.setCurrentIndex(index)
                for number in range(1, len(viewer.review.pages[audience])+1):
                    viewer.page_selector.setValue(number)
                    settle(lambda: viewer.review_page_button.isEnabled())
                    viewer.review_page_button.click()
            settle(lambda: viewer.confirm_button.isEnabled())
            assert viewer.grab().save(str(output/'task-set-real-pagination.png'))
            viewer.confirm_button.click(); settle(lambda: d._task is None and panel.approved)
            viewer.reject(); panel.export.click(); settle(lambda: d._task is None)
            assert panel.current()['exports'], d.status.text()
            artifacts = panel.current()['exports'][-1]['artifacts']
            copied = []
            for artifact in artifacts:
                ext = Path(artifact['path']).suffix
                role = 'student' if artifact['artifact_id'].startswith('student') else 'teacher'
                target = output/(role+ext); shutil.copyfile(artifact['path'], target)
                copied.append({'file': target.name, 'sha256': hashlib.sha256(target.read_bytes()).hexdigest()})
            def text(name):
                doc = Document(output/name)
                return '\n'.join([p.text for p in doc.paragraphs] +
                    [c.text for t in doc.tables for r in t.rows for c in r.cells])
            student = text('student.docx'); teacher = text('teacher.docx')
            assert '共同材料' in student and '动态平衡' in student
            assert '不是浓度必定相等' not in student and '不是浓度必定相等' in teacher
            assert 'c(B)=0.60' not in student
            assert f.basket() == remaining and path.read_bytes() == source and transport.calls == calls
            assert f.state_store.snapshot()['drafts'][_ACTIVE] == sentinel
            assert f.state_store.snapshot()['drafts']['paper-current']['title'] == '保留教师原组卷草稿'
            saved = d.store.load(exam['id'])
            assert saved['followups'] == d.followups and not errors
            d.resize(800, 700); settle(); assert panel.text.height() >= 150
            assert d.grab().save(str(output/'task-set-compact.png'))
            (output/'acceptance.json').write_text(json.dumps({
                'scope': 'Windows source; synthetic data; real Office pagination; no new EXE',
                'checks': {'saved_and_reopened': True, 'selected_question_removed_from_global_basket': True,
                    'unrelated_basket_and_global_preview_unchanged': True, 'original_source_unchanged': True,
                    'real_docx_pdf_export': True, 'teacher_student_answer_separation': True,
                    'unrelated_question_excluded': True, 'task_export_persisted': True},
                'network_model_calls': 0, 'artifacts': copied, 'uncaught_errors': errors,
            }, ensure_ascii=False, indent=2), encoding='utf-8')
        finally:
            if d is not None: d.dirty = False; d.close(); d.deleteLater()
            bridge.wait_for_done(5000); bridge.shutdown(); f.shutdown(); app.processEvents()
            sys.excepthook = previous_hook
    return 0


if __name__ == '__main__':
    raise SystemExit(main(sys.argv[1] if len(sys.argv) > 1 else 'candidate-qa/practice-office'))
