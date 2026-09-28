"""Real Qt and Office acceptance; generated sources and temporary state only."""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import traceback

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QPoint, QRect, Qt, qVersion
from PySide6.QtWidgets import QDialog, QVBoxLayout
from docx import Document
from integrations.deeptutor_shchem_v1.desktop_workbench.app import create_application
from integrations.deeptutor_shchem_v1.desktop_workbench.scan_paper_page import ScanPaperPage
from integrations.deeptutor_shchem_v1.desktop_mixed_paper_drafts import DRAFT_ID
from integrations.deeptutor_shchem_v1.desktop_state import DesktopStateStore
from runtime.deeptutor_shchem.independent_paper_fixture import seed
from runtime.deeptutor_shchem.word_import_identity_qa import QueuedTasks
from runtime.deeptutor_shchem.mixed_paper_draft_undo_qa import settle, rect_in, reachable_controls


def geometry(workspace, window):
    names = ("scope_label", "saved_papers", "refresh_button", "new_button", "open_button")
    result = {}
    for name in names:
        control = getattr(workspace, name)
        rect = rect_in(control, window)
        assert window.rect().contains(rect), (name, rect, window.size())
        result[name] = [rect.x(), rect.y(), rect.width(), rect.height()]
    panel = workspace._mixed_panel
    if panel is not None:
        for name in ("status", "save_button", "undo_button", "back_button"):
            control = getattr(panel, name)
            rect = rect_in(control, window)
            assert window.rect().contains(rect), (name, rect, window.size())
            result[name] = [rect.x(), rect.y(), rect.width(), rect.height()]
        label = panel.status
        needed = label.fontMetrics().boundingRect(QRect(0, 0, label.contentsRect().width(), 10000),
            Qt.TextFlag.TextWordWrap, label.text()).height()
        assert label.contentsRect().height() >= needed
        assert panel.scroll.horizontalScrollBar().maximum() == 0
        assert panel.scroll.viewport().height() >= 80
        result["body_viewport_height"] = panel.scroll.viewport().height()
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--ui-only", action="store_true", help="复核界面时保留原Office成品，不重复转换。")
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    app = create_application(["independent-paper-qa"])
    manifest = {"schema": "shchem.independent-paper-qa.v1", "synthetic_only": True,
                "real_personal_state_accessed": False, "provider_calls": 0, "qt_version": qVersion(),
                "real_office_pagination": False, "images": [], "artifacts": [], "checks": {}}
    errors = []
    previous_hook = sys.excepthook
    def report_error(typ, value, tb):
        errors.append(str(value))
        traceback.print_exception(typ, value, tb)
    sys.excepthook = report_error

    def screenshot(name, window, *, workspace=None):
        settle(app)
        checks = geometry(workspace, window) if workspace is not None else {}
        path = output / (name + ".png")
        shot = window.grab()
        assert shot.save(str(path))
        manifest["images"].append({"file": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "logical_size": [window.width(), window.height()], "pixel_size": [shot.width(), shot.height()],
            "device_pixel_ratio": shot.devicePixelRatio(), "geometry": checks})

    with tempfile.TemporaryDirectory(prefix="independent-paper-qa-") as directory:
        root = Path(directory)
        facade, source, rows, providers = seed(root)
        original = source.read_bytes()
        tasks = QueuedTasks()
        window = QDialog()
        window.setWindowTitle("独立当前卷 · 合成软件验收")
        layout = QVBoxLayout(window)
        layout.setContentsMargins(0, 0, 0, 0)
        page = ScanPaperPage(facade, tasks)
        layout.addWidget(page)
        workspace = page._composer
        window.resize(360, 520)
        window.show()
        settle(app)
        assert workspace._mixed_panel is None
        screenshot("01-choose-or-create-360x520", window, workspace=workspace)
        workspace.new_button.click()
        tasks.flush()
        panel = workspace._mixed_panel
        assert panel is not None and panel._loaded_once and not panel._restore_failed
        assert len(panel.model.order) == 2
        panel.title.setText("晶体类型 · 独立当前卷（合成验收）")
        panel.sections.setCurrentRow(1)
        panel.points.setValue(5)
        panel.up_button.click()
        panel.save_button.click()
        assert "晶体类型 · 独立当前卷（合成验收） · 2 段" in workspace.saved_papers.currentText()
        identity = panel.facade.paper_id
        original_order = list(panel.model.order)
        saved_payload = deepcopy(panel._draft_session.current.record["payload"])
        screenshot("02-current-top-360x520", window, workspace=workspace)
        for width in (360, 420):
            window.resize(width, 520)
            settle(app)
            reachable_controls(panel, app)
            panel.scroll.ensureWidgetVisible(panel.remove_button, 0, 0)
            settle(app)
            screenshot(f"03-current-edit-{width}x520", window, workspace=workspace)
        window.resize(1000, 780)
        panel.scroll.verticalScrollBar().setValue(0)
        screenshot("04-current-wide-1000x780", window, workspace=workspace)

        # Change the public basket to an actually imported, unrelated source.
        extra = root / "合成新增公共题篮来源.docx"
        doc = Document()
        doc.add_paragraph("【例1】新公共题篮材料丙：此题不得混入已保存的卷。")
        doc.add_paragraph("【答案】新题答案不得进入旧卷教师版。")
        doc.save(extra)
        facade.save_visual_import_batch(handout_files=(extra,), source_type="教师讲义")
        new_row = next(row for row in facade.word_question_catalog()["items"] if row["source_name"] == extra.name)
        facade.state_store.clear_basket()
        facade.add_word_questions_to_basket([{"key": new_row["key"], "revision": new_row["revision"], "points": 8}])
        basket = deepcopy(facade.basket())
        page.update_basket_count()
        assert panel.model.order == original_order
        workspace.refresh(selected=identity)
        workspace.open_button.click()
        tasks.flush()
        panel = workspace._mixed_panel
        assert panel._draft_session.current.record["payload"] == saved_payload
        assert panel.model.order == original_order and len(panel.model.items) == 2
        window.resize(420, 520)
        screenshot("05-reopened-after-basket-change-420x520", window, workspace=workspace)
        facade.state_store.clear_basket()
        page.update_basket_count()
        workspace.open_button.click()
        tasks.flush()
        panel = workspace._mixed_panel
        assert panel.model.order == original_order
        screenshot("06-reopened-empty-basket-420x520", window, workspace=workspace)
        manifest["checks"].update(reopen_after_public_basket_change=True, reopen_after_empty_basket=True,
                                 public_basket_never_rebuilds_current_paper=True)

        # Keep only the first source question for actual student/teacher output.
        panel.sections.setCurrentRow(0)
        panel.remove_button.click()
        assert len(panel.model.order) == 1
        panel.sections.setCurrentRow(0)
        panel.points.setValue(7.5)
        panel.details_button.click()
        tasks.flush()
        details = panel._details_dialog
        assert details.details["section_count"] == 1 and details.details["current_score"]["total"] == 7.5
        details.resize(420, 520)
        screenshot("07-current-details-420x520", details)
        details.close()
        panel.remove_button.click()
        window.resize(360, 520)
        screenshot("08-empty-current-360x520", window, workspace=workspace)
        assert panel.undo_button.isEnabled() and not panel.preview_button.isEnabled()
        panel.undo_button.click()
        assert len(panel.model.order) == 1
        if not args.ui_only:
            panel.preview_button.click()
            tasks.flush()
            preview = panel._preview_dialog
            assert preview is not None and preview.review is not None, panel.status.text()
            preview.resize(900, 740)
            for index, audience in enumerate(("student", "teacher")):
                preview.tabs.setCurrentIndex(index)
                for number in range(1, len(preview.review.pages[audience]) + 1):
                    preview.page_selector.setValue(number)
                    tasks.flush()
                    settle(app)
                    assert preview.review_page_button.isEnabled()
                    preview.review_page_button.click()
                screenshot(f"09-actual-{audience}-pagination-900x740", preview)
            assert preview.confirm_button.isEnabled()
            preview.confirm_button.click()
            tasks.flush()
            assert panel._approved and panel.export_button.isEnabled()
            preview.close()
            panel.export_button.click()
            tasks.flush()
            assert len(panel._artifact_paths) == 4, panel.status.text()
            for kind, original_path in panel._artifact_paths.items():
                target = output / (kind + Path(original_path).suffix)
                shutil.copyfile(original_path, target)
                manifest["artifacts"].append({"file": target.name, "sha256": hashlib.sha256(target.read_bytes()).hexdigest()})
            def text(name):
                doc = Document(output / name)
                return "\n".join(doc.element.xpath("//w:t/text()"))
            student, teacher = text("student_docx.docx"), text("teacher_docx.docx")
            assert "合成材料甲" in student and "合成材料乙" not in student and "材料丙" not in student
            assert "合成答案甲" not in student and "合成答案甲" in teacher
            assert "合成答案乙" not in teacher and "新题答案" not in teacher
            manifest["real_office_pagination"] = True
            manifest["checks"].update(real_four_files_exported=True, student_teacher_separation=True,
                excluded_and_new_basket_questions_absent=True)
            window.resize(420, 520)
            panel.scroll.verticalScrollBar().setValue(panel.scroll.verticalScrollBar().maximum())
            screenshot("10-exported-420x520", window, workspace=workspace)
        assert source.read_bytes() == original, "synthetic original source changed"
        assert providers.calls == 0, "provider unexpectedly called"
        assert not facade.basket(), "public basket unexpectedly changed"
        manifest["checks"].update(saved_order_points_and_details=True, original_source_unchanged=True,
                                 saved_selector_matches_committed_title_and_count=True)

        # Same-paper CAS failure must clear approval and preserve external edit.
        physical = identity + ":" + DRAFT_ID
        store = DesktopStateStore(facade.state_store.root)
        changed = store.draft_snapshot(physical).record
        changed["payload"]["settings_ui"]["title"] = "另一窗口的合成标题"
        store.save_draft(physical, changed)
        before = store.path.read_bytes()
        panel.save_button.click()
        assert panel._restore_failed and store.path.read_bytes() == before
        assert not panel.export_button.isEnabled()
        window.resize(360, 520)
        screenshot("11-concurrent-error-360x520", window, workspace=workspace)
        panel.reload_button.click()
        tasks.flush()
        assert panel.title.text() == "另一窗口的合成标题"
        window.resize(420, 520)
        screenshot("12-recovered-420x520", window, workspace=workspace)
        manifest["checks"]["concurrent_write_preserved_and_recoverable"] = True
        changed = store.draft_snapshot(physical).record
        changed["payload"]["schema_version"] = "future-version"
        store.save_draft(physical, changed)
        before = store.path.read_bytes()
        panel.reload_button.click()
        tasks.flush()
        assert panel._restore_failed and not panel.restart_button.isVisible() and not panel.restart_button.isEnabled()
        assert store.path.read_bytes() == before and "顶部题篮新建" in panel.status.text()
        window.resize(360, 520)
        screenshot("13-incompatible-preserved-360x520", window, workspace=workspace)
        manifest["checks"]["incompatible_draft_preserved_without_unusable_restart"] = True
        assert not errors
        manifest["uncaught_errors"] = errors
        window.close()
        tasks.flush()
        facade.shutdown()
        settle(app)
    sys.excepthook = previous_hook
    (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"images": len(manifest["images"]), "artifacts": len(manifest["artifacts"]),
                      "real_office_pagination": manifest["real_office_pagination"], "provider_calls": 0}))


if __name__ == "__main__":
    main()
