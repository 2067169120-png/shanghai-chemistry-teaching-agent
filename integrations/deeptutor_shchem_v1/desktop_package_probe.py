"""Opt-in frozen-app smoke using only generated, isolated test material.

Run the same shipping executable with --verify-package NEW_OUTPUT_DIRECTORY.
No teacher state, credentials, network request or external Python is used.
"""
from __future__ import annotations
import hashlib
import json
import os
from pathlib import Path
import platform
import sys
import tempfile
import time


def synthetic_candidate():
    """A renderer fixture, not a teacher-ready lesson or a claimed original exam."""
    slides = []
    for i, (title, content) in enumerate([
        ('合成打包自检', ['仅验证软件输出，不作为授课材料。']),
        ('化合价与电子转移', ['先标注化合价，再判断电子得失。']),
        ('独立练习', ['说明判断所用的证据。']),
        ('答案与讲评', ['核对化合价变化与电子得失方向。']),
        ('方法总结', ['找变化，判得失，说明依据。']),
    ], 1):
        slides.append(dict(id=f'S{i}', order=i, title=title, purpose='合成软件输出检查',
            objective_ids=['O1'], activity_ids=['A1'] if 1<i<5 else [],
            assessment_ids=['E1'] if i in (3,4) else [], minutes=8,
            content=content, teacher_notes='合成数据，不是完整课堂内容。'))
    return dict(schema_version='shchem.desktop-preparation-candidate.v1',
        candidate_id='PREPCAND-package-synthetic', title='合成打包自检', topic='合成打包自检',
        audience='软件验收', artifact_mode='linked_bundle', lesson_route='new_lesson',
        timing=dict(periods=1, minutes_per_period=40, total_minutes=40),
        source_basis=dict(mode='teacher_input_only', evidence_ids=[], statement_zh='离线软件合成验收材料'),
        objectives=[dict(id='O1', statement='依据化合价变化说明电子转移')],
        activities=[dict(id='A1', title='合成检查', objective_ids=['O1'], minutes=30,
            teacher_action='提出问题', student_action='说明依据', materials=['合成文字'])],
        assessments=[dict(id='E1', title='合成检测', objective_ids=['O1'], activity_ids=['A1'],
            evidence_of_learning='解释文字', success_criteria='依据和结论一致')],
        slides=slides,
        lesson_stages=[dict(id='L1', title='合成输出阶段', objective_ids=['O1'], activity_ids=['A1'],
            assessment_ids=['E1'], minutes=40, teacher_actions=['提出问题'], student_actions=['说明依据'],
            materials=['合成文字'], assessment='核对说明')],
        homework=dict(title='合成课后任务', tasks=[dict(id='HW1', instruction='复核解释', objective_ids=['O1'])], estimated_minutes=5),
        uncertainties=['只验证软件输出链，不验证课堂质量。'], candidate_only=True,
        teacher_review_required=True, publication_allowed=False, official_claim_allowed=False)


def run_probe(output: Path) -> int:
    output = output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtCore import qVersion
    from PySide6.QtGui import QIcon
    from docx import Document
    from pptx import Presentation
    from .desktop_paths import DesktopPaths
    from .desktop_facade import build_default_facade
    from .desktop_environment import collect_environment_report
    from .desktop_version import DESKTOP_VERSION
    from .desktop_preparation_renderer import NativePreparationRenderer
    from .desktop_local_pagination import render_docx, find_libreoffice
    from .desktop_workbench.app import create_application
    from .desktop_workbench.main_window import TeacherWorkbenchWindow, ALL_ROUTES
    from .desktop_workbench.dialogs import SettingsDialog

    app = create_application(['package-verification'])
    errors = []
    previous_hook = sys.excepthook
    sys.excepthook = lambda kind, value, tb: errors.append(kind.__name__)
    core = Path(__file__).resolve().parent
    workspace = DesktopPaths.discover(core).workspace_root
    frozen = bool(getattr(sys, 'frozen', False))
    if frozen:
        assert core.is_relative_to(Path(sys._MEIPASS).resolve()), 'Core did not load from bundle'
        assert workspace == Path(sys._MEIPASS).resolve(), 'Unexpected external source workspace'
    screenshots = []
    def settle(predicate=lambda: True):
        deadline = time.monotonic() + 30
        for _ in range(6):
            app.processEvents(); time.sleep(.03)
        while not predicate() and time.monotonic() < deadline:
            app.processEvents(); time.sleep(.03)
        assert predicate(), 'Native operation did not finish'
        assert not errors, errors
    def capture(widget, filename):
        path = output / filename
        assert widget.grab().save(str(path))
        screenshots.append(dict(file=filename, sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
    try:
        with tempfile.TemporaryDirectory(prefix='shchem-package-state-') as state:
            paths = DesktopPaths.from_workspace(workspace, state_root=state)
            facade = build_default_facade(paths)
            window = TeacherWorkbenchWindow(facade)
            window.resize(1360, 900)
            window.show()
            try:
                settle(lambda: not window.home_page._loading)
                for route in ALL_ROUTES:
                    window.navigate(route)
                    settle()
                    assert window.stack.currentWidget() is window.pages[route]
                    capture(window, 'packaged-' + route + '.png')
                expected_mode = '打包版' if frozen else '源码版'
                assert window.mode_label.text() == expected_mode
                settings = SettingsDialog(facade, window.tasks, window)
                settings.show()
                settle(lambda: settings._active_task_id is None)
                assert settings.advanced.content.isHidden()
                capture(settings, 'packaged-settings.png')
                settings.close(); settings.deleteLater(); settle()
                svg = core / 'desktop_workbench/studio_assets/down.svg'
                assert not QIcon(str(svg)).pixmap(24, 24).isNull()
                environment = collect_environment_report(paths)
                check = {row['key']:row for row in environment['checks']}
                assert all(check[key]['status']=='ready' for key in ('qt_widgets','qt_svg','qt_pdf','qt_pdf_widgets','docx','pptx','pillow','pdfium','pypdf','jsonschema','icons','resources','state'))
                prep = window.preparation_page
                assert prep.apply_studio_template('concept', '合成打包草稿', '软件验收')
                prep.materials.setPlainText('这是打包验收用合成资料。')
                facade.create_preparation_draft(prep._payload())
                assert facade.preparation_draft_options()[0]['title'] == '合成打包草稿'
                from .desktop_work_organization_probe import exercise
                organization = exercise(window, settle, capture)
                result = NativePreparationRenderer().render(synthetic_candidate(), output_kind='joint', output_dir=output/'synthetic-render')
                artifacts = {row['artifact_id']:row for row in result['artifacts']}
                assert len(Presentation(artifacts['pptx']['path']).slides)==5
                assert Document(artifacts['lesson_plan_docx']['path']).paragraphs
                file_checks = [{k:row[k] for k in ('artifact_id','filename','size','sha256')} for row in result['artifacts']]
                assert all(Path(row['path']).is_file() for row in result['artifacts'])
                pdf_pages = None
                if find_libreoffice():
                    pages, pdf, metadata = render_docx(artifacts['lesson_plan_docx']['path'], output/'sample-pages')
                    assert pages and pdf.is_file()
                    pdf_pages = len(pages)
            finally:
                window.close(); app.processEvents()
        report = dict(version=DESKTOP_VERSION, frozen=frozen, module_origin='bundle' if frozen else 'source',
            platform=platform.platform(), python=platform.python_version(), qt=qVersion(),
            qt_platform=os.environ.get('QT_QPA_PLATFORM'), routes_opened=list(ALL_ROUTES),
            local_draft_saved=True, work_organization=organization, editable_pptx_slides=5, lesson_docx_created=True, pdf_pages=pdf_pages,
            environment=environment, screenshots=screenshots, output_files=file_checks,
            uncaught_errors=errors, scope='Isolated native bundle navigation, settings, saved draft and production renderer DOCX/PPTX creation. Not live API, Office-PPTX visual parity, original library or classroom acceptance.')
        (output/'package-probe.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        return 0
    finally:
        sys.excepthook = previous_hook
