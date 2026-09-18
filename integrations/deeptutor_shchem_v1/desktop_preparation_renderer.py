from __future__ import annotations

"""Native, dependency-light renderer for desktop preparation candidates.

The renderer deliberately consumes only a normalized preparation candidate.  It
does not read the question bank, call a provider, perform OCR, or infer chemistry
content.  PowerPoint and PNG output are both produced from the same immutable
layout plan so the desktop preview is a faithful representation of the editable
deck structure.
"""

import hashlib
import inspect
import io
import json
import math
import os
import re
import zipfile
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from uuid import uuid4

from .desktop_preparation_images import (
    PreparationImageError,
    normalize_image_assets,
    normalize_slide_image,
    verify_image_bytes,
)
from .desktop_preparation_review import classroom_review
from .desktop_preparation_visual import normalize_slide_visual
from .desktop_preparation_worksheet import build_worksheet_document, worksheet_records

RENDERER_SCHEMA_VERSION = "shchem.native-preparation-renderer.v1"
DECK_SCHEMA_VERSION = "shchem.native-preparation-deck.v1"
QA_SCHEMA_VERSION = "shchem.native-preparation-machine-layout-qa.v1"

SLIDE_WIDTH_PX = 1600
SLIDE_HEIGHT_PX = 900
SLIDE_WIDTH_IN = 13.333333333333334
SLIDE_HEIGHT_IN = 7.5

PPTX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)
DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

_OUTPUT_KIND_ALIASES = {
    "ppt": "ppt",
    "lesson_plan": "lesson_plan",
    "joint": "linked_bundle",
    "linked_bundle": "linked_bundle",
}

_FILENAMES = {
    "candidate_json": "candidate.json",
    "deck_json": "deck.json",
    "pptx": "lesson_presentation.pptx",
    "preview_montage": "rendered_montage.png",
    "lesson_plan_docx": "lesson_plan.docx",
    "student_worksheet_docx": "student_worksheet.docx",
    "qa_report": "qa_report.json",
}

_CONTENT_TYPES = {
    "candidate_json": "application/json",
    "deck_json": "application/json",
    "pptx": PPTX_CONTENT_TYPE,
    "preview_montage": "image/png",
    "lesson_plan_docx": DOCX_CONTENT_TYPE,
    "student_worksheet_docx": DOCX_CONTENT_TYPE,
    "qa_report": "application/json",
}

_PALETTE = {
    "navy": "17324D",
    "navy_deep": "10283D",
    "teal": "138A86",
    "teal_light": "DDF2F0",
    "blue_light": "E8F0F6",
    "paper": "F7F9FB",
    "white": "FFFFFF",
    "ink": "172B3A",
    "muted": "617583",
    "line": "CCD7DE",
    "answer": "EAF3F6",
    "practice": "EDF7F5",
}

_FONT_FALLBACK = (
    "Microsoft YaHei",
    "Microsoft YaHei UI",
    "Noto Sans CJK SC",
    "Source Han Sans SC",
    "SimHei",
    "Arial",
)

_LESSON_ROUTE_LABELS = {
    "new_lesson": "新授课",
    "review": "复习课",
    "experiment": "实验课",
    "exercise_review": "习题讲评",
    "paper_review": "试卷讲评",
    "special_topic": "专题课",
    "hotspot": "热点专题",
    "other": "其他",
}


class NativePreparationRenderError(RuntimeError):
    """A stable renderer failure that the desktop manager can persist."""

    def __init__(
        self,
        code: str,
        message_zh: str,
        *,
        retryable: bool = False,
        details: Mapping[str, Any] | None = None,
    ) -> None:
        super().__init__(message_zh)
        self.code = code
        self.message_zh = message_zh
        self.retryable = retryable
        self.details = dict(details or {})


class NativePreparationRenderCancelled(NativePreparationRenderError):
    def __init__(self) -> None:
        super().__init__(
            "preparation_render_cancelled",
            "备课产物生成已取消。",
            retryable=True,
        )


@dataclass(frozen=True)
class _LayoutElement:
    element_type: str
    x: int
    y: int
    width: int
    height: int
    fill: str | None = None
    line: str | None = None
    text: str = ""
    color: str = _PALETTE["ink"]
    font_px: int = 32
    bold: bool = False
    align: str = "left"
    valign: str = "top"
    source_paths: tuple[str, ...] = ()
    source_texts: tuple[str, ...] = ()
    structural: bool = False
    overflow: bool = False
    image_data: bytes | None = None
    image_asset_id: str = ""
    image_source_sha256: str = ""
    image_alt: str = ""

    def public(self) -> dict[str, Any]:
        value: dict[str, Any] = {
            "element_type": self.element_type,
            "bbox_px": [self.x, self.y, self.width, self.height],
        }
        if self.fill is not None:
            value["fill"] = self.fill
        if self.line is not None:
            value["line"] = self.line
        if self.element_type == "image":
            value.update(
                {
                    "asset_id": self.image_asset_id,
                    "source_sha256": self.image_source_sha256,
                    "embedded_sha256": _sha256(self.image_data or b""),
                    "alt_text": self.image_alt,
                    "fit": "contain",
                }
            )
        if self.element_type == "text":
            value.update(
                {
                    "text": self.text,
                    "font_family": _FONT_FALLBACK[0],
                    "font_size_pt": round(self.font_px * 72 / 120, 1),
                    "bold": self.bold,
                    "color": self.color,
                    "align": self.align,
                    "valign": self.valign,
                    "source_paths": list(self.source_paths),
                    "source_texts": list(self.source_texts),
                    "structural": self.structural,
                    "overflow": self.overflow,
                }
            )
        return value


@dataclass(frozen=True)
class _SlideLayout:
    page_no: int
    slide_id: str
    slide_type: str
    title: str
    elements: tuple[_LayoutElement, ...]
    layout_variant: str = "v1"

    def public(self) -> dict[str, Any]:
        return {
            "page_no": self.page_no,
            "slide_id": self.slide_id,
            "slide_type": self.slide_type,
            "title": self.title,
            "layout_id": f"native.{self.slide_type}.{self.layout_variant}",
            "elements": [element.public() for element in self.elements],
        }


def _json_bytes(value: Any) -> bytes:
    try:
        return (
            json.dumps(
                value,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
                allow_nan=False,
            ).encode("utf-8")
            + b"\n"
        )
    except (TypeError, ValueError) as exc:
        raise NativePreparationRenderError(
            "preparation_candidate_json_invalid",
            "备课候选无法写入规范 JSON。",
            details={"exception_type": type(exc).__name__},
        ) from exc


def _plain_json(value: Mapping[str, Any]) -> dict[str, Any]:
    return json.loads(_json_bytes(value).decode("utf-8"))


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _contained(root: Path, candidate: Path) -> Path:
    root_resolved = root.resolve(strict=True)
    candidate_resolved = candidate.resolve(strict=False)
    try:
        candidate_resolved.relative_to(root_resolved)
    except ValueError as exc:
        raise NativePreparationRenderError(
            "preparation_render_path_escape",
            "备课产物路径超出当前生成目录。",
        ) from exc
    return candidate_resolved


def _safe_output_root(output_dir: str | os.PathLike[str]) -> Path:
    if not isinstance(output_dir, (str, os.PathLike)):
        raise NativePreparationRenderError(
            "preparation_render_output_invalid", "备课生成目录不正确。"
        )
    raw = Path(output_dir)
    if raw.exists() and raw.is_symlink():
        raise NativePreparationRenderError(
            "preparation_render_output_symlink_forbidden",
            "备课生成目录不能是符号链接。",
        )
    try:
        raw.mkdir(parents=True, exist_ok=True)
        root = raw.resolve(strict=True)
    except OSError as exc:
        raise NativePreparationRenderError(
            "preparation_render_output_unavailable",
            "无法创建备课生成目录。",
            retryable=True,
            details={"exception_type": type(exc).__name__},
        ) from exc
    if not root.is_dir():
        raise NativePreparationRenderError(
            "preparation_render_output_invalid", "备课生成目录不正确。"
        )
    return root


def _atomic_write(path: Path, data: bytes, *, root: Path) -> None:
    target = _contained(root, path)
    if target.exists():
        raise NativePreparationRenderError(
            "preparation_render_artifact_exists",
            "本次生成目录中已有同名产物，已停止覆盖。",
            details={"filename": target.name},
        )
    temporary = _contained(root, target.with_name(f".{target.name}.{uuid4().hex}.tmp"))
    try:
        temporary.write_bytes(data)
        os.replace(temporary, target)
    except OSError as exc:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass
        raise NativePreparationRenderError(
            "preparation_render_write_failed",
            "备课产物写入失败。",
            retryable=True,
            details={"filename": target.name, "exception_type": type(exc).__name__},
        ) from exc


def _as_text(value: Any) -> str:
    if not isinstance(value, str):
        return ""
    return value.strip()


def _string_values(value: Any, *, skip_ids: bool = True) -> list[str]:
    result: list[str] = []

    def visit(current: Any, key: str = "") -> None:
        if isinstance(current, str):
            text = current.strip()
            if text and not (skip_ids and (key == "id" or key.endswith("_id"))):
                result.append(text)
            return
        if isinstance(current, Mapping):
            preferred = (
                "text",
                "statement",
                "title",
                "instruction",
                "content",
                "prompt",
                "answer",
                "description",
            )
            visited: set[str] = set()
            for name in preferred:
                if name in current:
                    visit(current[name], name)
                    visited.add(name)
            for name, child in current.items():
                key_text = str(name)
                if name in visited or key_text.endswith(("_ids", "_refs")):
                    continue
                visit(child, key_text)
            return
        if isinstance(current, Sequence) and not isinstance(
            current, (bytes, bytearray)
        ):
            for child in current:
                visit(child, key)

    visit(value)
    deduplicated: list[str] = []
    seen: set[str] = set()
    for item in result:
        if item not in seen:
            seen.add(item)
            deduplicated.append(item)
    return deduplicated


def _candidate_strings(candidate: Mapping[str, Any]) -> set[str]:
    found: set[str] = set()

    def visit(value: Any) -> None:
        if isinstance(value, str):
            found.add(value.strip())
        elif isinstance(value, Mapping):
            for child in value.values():
                visit(child)
        elif isinstance(value, Sequence) and not isinstance(value, (bytes, bytearray)):
            for child in value:
                visit(child)

    visit(candidate)
    return found


def _callback_arity(callback: Callable[..., Any]) -> tuple[int, bool]:
    try:
        signature = inspect.signature(callback)
    except (TypeError, ValueError):
        return 1, False
    positional = 0
    varargs = False
    for parameter in signature.parameters.values():
        if parameter.kind == inspect.Parameter.VAR_POSITIONAL:
            varargs = True
        elif parameter.kind in (
            inspect.Parameter.POSITIONAL_ONLY,
            inspect.Parameter.POSITIONAL_OR_KEYWORD,
        ):
            positional += 1
    return positional, varargs


def _report_progress(
    callback: Callable[..., Any] | None,
    progress: int,
    stage: str,
    message_zh: str,
) -> None:
    if callback is None:
        return
    payload = {
        "progress": int(progress),
        "stage": stage,
        "message_zh": message_zh,
    }
    positional, varargs = _callback_arity(callback)
    if varargs or positional >= 3:
        callback(payload["progress"], payload["stage"], payload["message_zh"])
    elif positional == 2:
        callback(payload["progress"], payload["stage"])
    elif positional == 1:
        callback(payload)
    else:
        callback()


def _check_cancelled(callback: Callable[[], bool] | None) -> None:
    if callback is not None and callback():
        raise NativePreparationRenderCancelled()


def _hex_rgb(value: str) -> tuple[int, int, int]:
    return tuple(int(value[index : index + 2], 16) for index in (0, 2, 4))  # type: ignore[return-value]


def _font_paths() -> tuple[tuple[Path, ...], tuple[Path, ...]]:
    windows = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    regular = (
        windows / "msyh.ttc",
        windows / "msyh.ttf",
        windows / "msyhui.ttc",
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    )
    bold = (
        windows / "msyhbd.ttc",
        windows / "msyhbd.ttf",
        windows / "msyh.ttc",
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    )
    return regular, bold


class _FontBook:
    def __init__(self) -> None:
        try:
            from PIL import ImageFont
        except ImportError as exc:
            raise NativePreparationRenderError(
                "preparation_render_dependency_missing",
                "缺少 Pillow，无法生成备课逐页预览。",
            ) from exc
        self._image_font = ImageFont
        self._regular, self._bold = _font_paths()
        self._cache: dict[tuple[int, bool], Any] = {}

    def get(self, size: int, bold: bool = False) -> Any:
        key = (int(size), bool(bold))
        if key in self._cache:
            return self._cache[key]
        candidates = self._bold if bold else self._regular
        for path in candidates:
            if not path.is_file():
                continue
            try:
                font = self._image_font.truetype(str(path), size=int(size))
            except OSError:
                continue
            font._shchem_glyph_fallback = _GlyphFallback(
                font, self._image_font, (*candidates, *_symbol_font_paths(bold))
            )
            self._cache[key] = font
            return font
        font = self._image_font.load_default()
        font._shchem_glyph_fallback = _GlyphFallback(font, self._image_font, ())
        self._cache[key] = font
        return font


def _symbol_font_paths(bold: bool = False) -> tuple[Path, ...]:
    windows = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    return (
        windows / "seguisym.ttf",
        windows / ("cambriab.ttf" if bold else "cambria.ttc"),
        windows / ("arialbd.ttf" if bold else "arial.ttf"),
        Path(
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
            if bold
            else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
        ),
        Path("/usr/share/fonts/truetype/noto/NotoSansSymbols2-Regular.ttf"),
    )


class _GlyphFallback:
    """Resolve actual missing glyphs, without rewriting any source character.

    Pillow does not perform desktop font linking. Compare its missing-glyph mask
    using two unassigned codepoints; cache coverage per loaded face and character.
    Measurement and painting below consume precisely the same contiguous runs.
    """

    def __init__(self, primary: Any, loader: Any, paths: Sequence[Path]) -> None:
        self.primary = primary
        self.loader = loader
        self.paths = tuple(dict.fromkeys(paths))
        self.faces = [primary]
        self.loaded = False
        self.coverage: dict[tuple[int, str], bool] = {}
        self.missing_masks: dict[int, set[tuple[Any, bytes]]] = {}
        self.selected: dict[str, Any] = {}

    def supports(self, face: Any, char: str) -> bool:
        if char.isspace():
            return True
        key = (id(face), char)
        if key not in self.coverage:

            def mask_signature(value: str) -> tuple[Any, bytes]:
                mask = face.getmask(value)
                return mask.size, bytes(mask)

            if id(face) not in self.missing_masks:
                self.missing_masks[id(face)] = {
                    mask_signature("\U0010ffff"),
                    mask_signature("\u0378"),
                }
            self.coverage[key] = (
                mask_signature(char) not in self.missing_masks[id(face)]
            )
        return self.coverage[key]

    def _load(self) -> None:
        if self.loaded:
            return
        self.loaded = True
        primary_path = str(getattr(self.primary, "path", ""))
        for path in self.paths:
            if not path.is_file() or str(path) == primary_path:
                continue
            try:
                self.faces.append(
                    self.loader.truetype(str(path), size=self.primary.size)
                )
            except OSError:
                continue

    def runs(self, text: str) -> list[tuple[str, Any]]:
        runs: list[tuple[str, Any]] = []
        for char in text:
            if char not in self.selected:
                if self.supports(self.primary, char):
                    self.selected[char] = self.primary
                else:
                    self._load()
                    face = next((f for f in self.faces if self.supports(f, char)), None)
                    if face is None:
                        raise NativePreparationRenderError(
                            "preparation_preview_glyph_missing",
                            f"预览字体均缺少字符 U+{ord(char):04X}（{char}），请安装支持该字符的字体后重试；未替换原文。",
                        )
                    self.selected[char] = face
            face = self.selected[char]
            if runs and runs[-1][1] is face:
                runs[-1] = (runs[-1][0] + char, face)
            else:
                runs.append((char, face))
        return runs


def _glyph_runs(text: str, font: Any) -> list[tuple[str, Any]]:
    resolver = getattr(font, "_shchem_glyph_fallback", None)
    return resolver.runs(text) if resolver is not None else [(text, font)]


def _text_length(draw: Any, text: str, font: Any) -> float:
    return sum(draw.textlength(run, font=face) for run, face in _glyph_runs(text, font))


def _fallback_multiline_plan(
    draw: Any, text: str, font: Any, spacing: int, align: str
) -> (
    tuple[tuple[float, float, float, float], list[tuple[str, Any, float, float]]] | None
):
    lines = [_glyph_runs(line, font) for line in text.split("\n")]
    if all(face is font for runs in lines for _, face in runs):
        return None
    widths = [
        sum(draw.textlength(run, font=face) for run, face in runs) for runs in lines
    ]
    width = max(widths, default=0)
    ascent = font.getmetrics()[0]
    # Use one baseline for the mixed fonts. Increase line advance only when an
    # actual glyph descends farther than the primary font's native line measure.
    boxes = [[face.getbbox(run, anchor="ls") for run, face in runs] for runs in lines]
    advance = (
        max(
            font.getbbox("A")[3],
            max((ascent + box[3] for row in boxes for box in row), default=0),
        )
        + spacing
    )
    placements: list[tuple[str, Any, float, float]] = []
    bounds: list[tuple[float, float, float, float]] = []
    for index, runs in enumerate(lines):
        x = (width - widths[index]) * ({"center": 0.5, "right": 1}.get(align, 0))
        baseline = ascent + index * advance
        for (run, face), box in zip(runs, boxes[index]):
            placements.append((run, face, x, baseline))
            bounds.append(
                (x + box[0], baseline + box[1], x + box[2], baseline + box[3])
            )
            x += draw.textlength(run, font=face)
    if not bounds:
        return (0, 0, 0, 0), placements
    return (
        min(b[0] for b in bounds),
        min(b[1] for b in bounds),
        max(b[2] for b in bounds),
        max(b[3] for b in bounds),
    ), placements


def _multiline_textbbox(
    draw: Any, text: str, font: Any, spacing: int, align: str = "left"
) -> tuple[float, float, float, float]:
    plan = _fallback_multiline_plan(draw, text, font, spacing, align)
    return (
        plan[0]
        if plan is not None
        else draw.multiline_textbbox(
            (0, 0), text, font=font, spacing=spacing, align=align
        )
    )


def _draw_multiline_text(
    draw: Any,
    xy: tuple[float, float],
    text: str,
    *,
    font: Any,
    fill: Any,
    spacing: int,
    align: str = "left",
) -> None:
    plan = _fallback_multiline_plan(draw, text, font, spacing, align)
    if plan is None:
        draw.multiline_text(
            xy, text, font=font, fill=fill, spacing=spacing, align=align
        )
        return
    for run, face, x, baseline in plan[1]:
        draw.text((xy[0] + x, xy[1] + baseline), run, font=face, fill=fill, anchor="ls")


# These are typographic units, not a chemistry parser. Preserve Latin words,
# formulas with simple parenthesized groups, hydrate dots and charge notation
# verbatim. Never infer subscripts, chemical validity or equation semantics.
_LATIN_UNIT = r"[A-Za-z0-9₀-₉⁰¹²³⁴⁵⁶⁷⁸⁹]+"
_WRAP_UNIT = re.compile(
    rf"{_LATIN_UNIT}(?:\({_LATIN_UNIT}\)(?:{_LATIN_UNIT})?|·{_LATIN_UNIT})*"
    r"(?:\^[0-9]*[+-]|[⁺⁻]+|[+-])?|."
)


def _wrap_text(draw: Any, text: str, font: Any, max_width: int) -> str:
    closing = set("，。！？；：、）》】〕〉”’…,.!?;:)]}")
    opening = set("（《【〔〈“‘([{")
    wrapped: list[str] = []
    for raw_line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw_line.strip()
        if not line:
            wrapped.append("")
            continue
        current: list[str] = []
        for unit in _WRAP_UNIT.findall(line):
            proposed = "".join(current) + unit
            if current and _text_length(draw, proposed, font) > max_width:
                # Carry a whole preceding unit with closing punctuation;
                # never leave an opening bracket at the previous line end.
                split = len(current)
                while split > 0:
                    following = current[split] if split < len(current) else unit
                    if (
                        current[split - 1][-1] not in opening
                        and following[0] not in closing
                    ):
                        break
                    split -= 1
                if split:
                    wrapped.append("".join(current[:split]).rstrip())
                    current = current[split:]
                    while current and current[0].isspace():
                        current.pop(0)
                # An unbreakable unit can exceed the box. Keep it intact and
                # let _fit_text report overflow instead of silently clipping.
                if current or not unit.isspace():
                    current.append(unit)
            else:
                current.append(unit)
        if current or not wrapped:
            wrapped.append("".join(current).rstrip())
    return "\n".join(wrapped)


def _wrap_title_text(draw: Any, text: str, font: Any, max_width: int) -> str:
    """Rebalance an orphaned title line without shrinking or rewriting text."""
    wrapped = _wrap_text(draw, text, font, max_width)
    lines = wrapped.split("\n")
    # Respect intentional author breaks and keep ordinary/body wrapping intact.
    if "\n" in text or "\r" in text or len(lines) < 2:
        return wrapped
    widths = [_text_length(draw, line, font) for line in lines]
    if max(widths) > max_width or widths[-1] >= max(widths) * 0.4:
        return wrapped
    best, best_spread = wrapped, max(widths) - min(widths)
    low, high = 1, max_width
    # Find a narrower measure with the same line count, reusing the existing
    # formula/Latin-token and CJK-punctuation break rules. Never add a line.
    while low <= high:
        measure = (low + high) // 2
        trial = _wrap_text(draw, text, font, measure)
        trial_lines = trial.split("\n")
        trial_widths = [_text_length(draw, line, font) for line in trial_lines]
        if len(trial_lines) > len(lines) or max(trial_widths) > measure:
            low = measure + 1
            continue
        spread = max(trial_widths) - min(trial_widths)
        if len(trial_lines) == len(lines) and spread < best_spread:
            best, best_spread = trial, spread
        high = measure - 1
    return best


def _text_height(draw: Any, text: str, font: Any, spacing: int) -> int:
    if not text:
        return 0
    box = _multiline_textbbox(draw, text, font, spacing)
    return max(0, int(box[3] - box[1]))


def _text_fits(
    draw: Any, text: str, font: Any, spacing: int, width: int, height: int
) -> bool:
    # PowerPoint uses explicit 1.18-em paragraph line advance below. Pillow's
    # ink bounding box can be shorter than this, especially for CJK/mixed glyphs.
    # Reserve the same complete line boxes in both outputs, not just ink height.
    frame_height = math.ceil(len(text.split("\n")) * getattr(font, "size", 0) * 1.18)
    return max(_text_height(draw, text, font, spacing), frame_height) <= height and all(
        _text_length(draw, line, font) <= width for line in text.split("\n")
    )


def _fit_text(
    draw: Any,
    fonts: _FontBook,
    text: str,
    *,
    width: int,
    height: int,
    preferred_px: int,
    minimum_px: int,
    bold: bool = False,
    balance_lines: bool = False,
) -> tuple[str, int, bool]:
    chosen = max(minimum_px, preferred_px)
    wrap = _wrap_title_text if balance_lines else _wrap_text
    for size in range(chosen, minimum_px - 1, -2):
        font = fonts.get(size, bold)
        # Read-only Office measurements of the retained teaching deck showed
        # up to ~0.4 em of additional horizontal extent versus Pillow. Reserve
        # that at the line end before wrapping; do not rewrite source characters
        # or enable Office-only wrapping that would diverge from PNG previews.
        measure_width = max(1, width - math.ceil(size * 0.4))
        wrapped = wrap(draw, text, font, measure_width)
        spacing = max(5, int(size * 0.22))
        if _text_fits(draw, wrapped, font, spacing, measure_width, height):
            return wrapped, size, False
    font = fonts.get(minimum_px, bold)
    measure_width = max(1, width - math.ceil(minimum_px * 0.4))
    wrapped = wrap(draw, text, font, measure_width)
    spacing = max(5, int(minimum_px * 0.22))
    return (
        wrapped,
        minimum_px,
        not _text_fits(draw, wrapped, font, spacing, measure_width, height),
    )


def _slide_type(slide: Mapping[str, Any], page_no: int, slide_count: int) -> str:
    try:
        visual = normalize_slide_visual(slide.get("visual"))
    except ValueError:
        raise NativePreparationRenderError(
            "preparation_candidate_visual_invalid", "页面比较结构或流程节点不正确。"
        ) from None
    if visual is not None:
        return visual["kind"]
    explicit = _as_text(slide.get("slide_type") or slide.get("kind")).lower()
    aliases = {
        "cover": "cover",
        "title": "cover",
        "content": "content",
        "practice": "practice",
        "exercise": "practice",
        "answer": "answer",
        "solution": "answer",
        "summary": "summary",
        "homework": "summary",
    }
    if explicit in aliases:
        return aliases[explicit]
    haystack = " ".join(
        _string_values([slide.get("title"), slide.get("purpose")], skip_ids=False)
    ).lower()
    if page_no == 1:
        return "cover"
    if any(
        token in haystack for token in ("答案", "讲评", "解析", "answer", "solution")
    ):
        return "answer"
    if any(
        token in haystack
        for token in ("练习", "任务", "检测", "作答", "practice", "exercise")
    ):
        return "practice"
    if page_no == slide_count or any(
        token in haystack for token in ("总结", "回顾", "作业", "summary", "homework")
    ):
        return "summary"
    return "content"


def _slide_source_lines(
    slide: Mapping[str, Any], index: int
) -> tuple[list[str], list[str]]:
    lines: list[str] = []
    paths: list[str] = []
    # Teaching intent belongs in presenter notes, not student-facing bullets.
    content = slide.get("content")
    content_values = _string_values(content)
    for item_index, text in enumerate(content_values):
        if text not in lines:
            lines.append(text)
            paths.append(f"/slides/{index}/content/{item_index}")
    return lines, paths


def _source_element(
    *,
    draw: Any,
    fonts: _FontBook,
    x: int,
    y: int,
    width: int,
    height: int,
    source_texts: Sequence[str],
    source_paths: Sequence[str],
    preferred_px: int,
    minimum_px: int,
    color: str,
    bold: bool = False,
    align: str = "left",
    valign: str = "top",
    numbered: bool = False,
    paragraph_separator: str = "\n\n",
    balance_lines: bool = False,
) -> _LayoutElement:
    raw_lines = [text.strip() for text in source_texts if text.strip()]
    if numbered and len(raw_lines) > 1:
        visible = paragraph_separator.join(
            f"{number}. {text}" for number, text in enumerate(raw_lines, 1)
        )
    else:
        visible = paragraph_separator.join(raw_lines)
    wrapped, size, overflow = _fit_text(
        draw,
        fonts,
        visible,
        width=width,
        height=height,
        preferred_px=preferred_px,
        minimum_px=minimum_px,
        bold=bold,
        balance_lines=balance_lines,
    )
    return _LayoutElement(
        "text",
        x,
        y,
        width,
        height,
        text=wrapped,
        color=color,
        font_px=size,
        bold=bold,
        align=align,
        valign=valign,
        source_paths=tuple(source_paths),
        source_texts=tuple(raw_lines),
        overflow=overflow,
    )


def _visual_elements(
    slide: Mapping[str, Any], *, index: int, page_no: int, draw: Any, fonts: _FontBook
) -> list[_LayoutElement]:
    """Place semantic structures as editable shapes with exact source pointers."""
    visual = normalize_slide_visual(slide["visual"])
    assert visual is not None
    base = f"/slides/{index}"
    elements = [
        _LayoutElement("rect", 0, 0, 1600, 900, fill=_PALETTE["paper"]),
        _LayoutElement("rect", 0, 0, 1600, 14, fill=_PALETTE["teal"]),
    ]

    def text(
        value: str,
        path: str,
        box: tuple[int, int, int, int],
        *,
        size: int = 34,
        bold: bool = False,
        color: str = _PALETTE["ink"],
        center: bool = False,
    ) -> None:
        x, y, width, height = box
        elements.append(
            _source_element(
                draw=draw,
                fonts=fonts,
                x=x,
                y=y,
                width=width,
                height=height,
                source_texts=[value],
                source_paths=[path],
                preferred_px=size,
                minimum_px=26,
                color=color,
                bold=bold,
                align="center" if center else "left",
                valign="middle",
            )
        )

    text(
        _as_text(slide.get("title")),
        base + "/title",
        (100, 48, 1380, 122),
        size=56,
        bold=True,
        color=_PALETTE["navy"],
    )
    lines, paths = _slide_source_lines(slide, index)
    content_height = 98
    comparison_table_y = 320
    if visual["kind"] == "comparison":
        # Comparison pages need enough vertical room for a source-bound
        # introduction (often a textbook sentence plus a short distinction).
        # Measure with the same wrapping/font metrics used by _source_element;
        # this keeps the editable PPTX and the PNG preview on one layout plan.
        visible = "\n".join(text.strip() for text in lines if text.strip())
        preferred_font = fonts.get(34, False)
        wrapped = _wrap_text(draw, visible, preferred_font, 1380)
        spacing = max(5, int(34 * 0.22))
        required_height = _text_height(draw, wrapped, preferred_font, spacing) + 12
        content_height = min(148, max(98, required_height))
        comparison_table_y += content_height - 98
    elements.append(
        _source_element(
            draw=draw,
            fonts=fonts,
            x=104,
            y=190,
            width=1380,
            height=content_height,
            source_texts=lines,
            source_paths=paths,
            preferred_px=34,
            minimum_px=26,
            color=_PALETTE["muted"],
            paragraph_separator="\n",
        )
    )

    if visual["kind"] == "comparison":
        table = visual["comparison"]
        assert isinstance(table, Mapping)
        widths = [224] + [int(1176 / len(table["columns"]))] * len(table["columns"])
        row_height = 352 // len(table["rows"])
        table_base = base + "/visual/comparison"
        headers = [(table["dimension_label"], table_base + "/dimension_label")] + [
            (label, table_base + f"/columns/{column}")
            for column, label in enumerate(table["columns"])
        ]
        x = 100
        for column, ((label, path), width) in enumerate(
            zip(headers, widths, strict=True)
        ):
            elements.append(
                _LayoutElement(
                    "rect",
                    x,
                    comparison_table_y,
                    width,
                    66,
                    fill=_PALETTE["navy"] if column == 0 else _PALETTE["teal"],
                    line=_PALETTE["white"],
                )
            )
            text(
                label,
                path,
                (x + 16, comparison_table_y + 5, width - 32, 56),
                size=32,
                bold=True,
                color=_PALETTE["white"],
                center=True,
            )
            x += width
        for row_index, row in enumerate(table["rows"]):
            y = comparison_table_y + 66 + row_index * row_height
            x = 100
            values = [(row["label"], table_base + f"/rows/{row_index}/label")] + [
                (value, table_base + f"/rows/{row_index}/values/{column}")
                for column, value in enumerate(row["values"])
            ]
            for column, ((value, path), width) in enumerate(
                zip(values, widths, strict=True)
            ):
                fill = (
                    _PALETTE["blue_light"]
                    if column == 0
                    else (_PALETTE["white"] if row_index % 2 == 0 else "EEF4F5")
                )
                elements.append(
                    _LayoutElement(
                        "rect",
                        x,
                        y,
                        width,
                        row_height,
                        fill=fill,
                        line=_PALETTE["line"],
                    )
                )
                text(
                    value,
                    path,
                    (x + 18, y + 10, width - 36, row_height - 20),
                    size=32,
                    bold=column == 0,
                    center=column == 0,
                )
                x += width
    else:
        steps = visual["steps"]
        gap = 54
        width = (1400 - gap * (len(steps) - 1)) // len(steps)
        for step_index, step in enumerate(steps):
            x = 100 + step_index * (width + gap)
            step_base = base + f"/visual/steps/{step_index}"
            elements.extend(
                [
                    _LayoutElement(
                        "rect",
                        x,
                        344,
                        width,
                        362,
                        fill=_PALETTE["white"],
                        line=_PALETTE["line"],
                    ),
                    _LayoutElement("rect", x, 344, width, 100, fill=_PALETTE["navy"]),
                ]
            )
            text(
                step["label"],
                step_base + "/label",
                (x + 20, 356, width - 40, 76),
                size=34,
                bold=True,
                color=_PALETTE["white"],
                center=True,
            )
            text(
                step["detail"],
                step_base + "/detail",
                (x + 24, 468, width - 48, 212),
                size=32,
            )
            if step_index < len(steps) - 1:
                elements.append(
                    _LayoutElement(
                        "right_arrow",
                        x + width + 10,
                        505,
                        gap - 20,
                        34,
                        fill=_PALETTE["teal"],
                        structural=True,
                    )
                )
    elements.append(
        _LayoutElement(
            "text",
            1370,
            802,
            120,
            44,
            text=str(page_no),
            font_px=24,
            color=_PALETTE["muted"],
            align="right",
            structural=True,
        )
    )
    return elements


def _layout_slide(
    slide: Mapping[str, Any],
    *,
    index: int,
    page_no: int,
    slide_count: int,
    candidate: Mapping[str, Any],
    draw: Any,
    fonts: _FontBook,
    image_data: Mapping[str, bytes] | None = None,
) -> _SlideLayout:
    slide_id = _as_text(slide.get("id") or slide.get("slide_id")) or f"S{page_no:02d}"
    title = _as_text(slide.get("title")) or _as_text(candidate.get("title"))
    kind = _slide_type(slide, page_no, slide_count)
    lines, paths = _slide_source_lines(slide, index)
    elements: list[_LayoutElement] = []

    if slide.get("image") is not None:
        kind = "image"
        elements.extend(
            _image_elements(
                slide,
                index=index,
                candidate=candidate,
                image_data=image_data or {},
                draw=draw,
                fonts=fonts,
            )
        )
    elif kind in {"comparison", "process"}:
        elements.extend(
            _visual_elements(
                slide, index=index, page_no=page_no, draw=draw, fonts=fonts
            )
        )
    elif kind == "cover":
        elements.extend(
            [
                _LayoutElement("rect", 0, 0, 1600, 900, fill=_PALETTE["navy_deep"]),
                _LayoutElement("rect", 0, 0, 36, 900, fill=_PALETTE["teal"]),
                _LayoutElement("rect", 112, 158, 180, 9, fill=_PALETTE["teal"]),
            ]
        )
        # The teacher's chapter/topic is the cover authority. A model-generated
        # hook question belongs after the cover, not in place of its title.
        topic_title = _as_text(candidate.get("topic"))
        title_source = topic_title or title
        title_path = (
            "/topic"
            if topic_title
            else (
                f"/slides/{index}/title" if _as_text(slide.get("title")) else "/title"
            )
        )
        if title_source.strip().lower() in {
            "封面",
            "封面页",
            "cover",
            "title",
            "本节课的核心问题",
            "本课核心问题",
        }:
            title_source = _as_text(candidate.get("title")) or _as_text(
                candidate.get("topic")
            )
            title_path = "/title" if _as_text(candidate.get("title")) else "/topic"
        if title_source:
            elements.append(
                _source_element(
                    draw=draw,
                    fonts=fonts,
                    x=112,
                    y=204,
                    width=1260,
                    height=270,
                    source_texts=[title_source],
                    source_paths=[title_path],
                    preferred_px=82,
                    minimum_px=54,
                    balance_lines=True,
                    color=_PALETTE["white"],
                    bold=True,
                    valign="middle",
                )
            )
        supporting: list[str] = []
        supporting_paths: list[str] = []
        # Audience may contain private class assumptions and preparation gaps.
        # It remains in the teacher document, never auto-injected on projection.
        for key in ("topic",):
            text = _as_text(candidate.get(key))
            if text and text != title_source and text not in supporting:
                supporting.append(text)
                supporting_paths.append(f"/{key}")
        if lines:
            for text, source_path in zip(lines, paths):
                if text != title_source and text not in supporting:
                    supporting.append(text)
                    supporting_paths.append(source_path)
        if supporting:
            elements.append(
                _source_element(
                    draw=draw,
                    fonts=fonts,
                    x=116,
                    y=500,
                    width=1120,
                    height=244,
                    source_texts=supporting,
                    source_paths=supporting_paths,
                    preferred_px=34,
                    minimum_px=25,
                    color="C7D8E5",
                    paragraph_separator="\n",
                )
            )
    elif kind == "practice":
        elements.extend(
            [
                _LayoutElement("rect", 0, 0, 1600, 900, fill=_PALETTE["practice"]),
                _LayoutElement("rect", 0, 0, 218, 900, fill=_PALETTE["navy"]),
                _LayoutElement("rect", 218, 0, 16, 900, fill=_PALETTE["teal"]),
            ]
        )
        if title:
            elements.append(
                _source_element(
                    draw=draw,
                    fonts=fonts,
                    x=292,
                    y=80,
                    width=1160,
                    height=126,
                    source_texts=[title],
                    source_paths=[f"/slides/{index}/title"],
                    preferred_px=58,
                    minimum_px=42,
                    color=_PALETTE["navy"],
                    bold=True,
                    valign="middle",
                )
            )
        if lines:
            elements.append(
                _source_element(
                    draw=draw,
                    fonts=fonts,
                    x=302,
                    y=258,
                    width=1070,
                    height=510,
                    source_texts=lines,
                    source_paths=paths,
                    preferred_px=39,
                    minimum_px=26,
                    color=_PALETTE["ink"],
                    numbered=len(lines) > 1,
                )
            )
        elements.append(
            _LayoutElement(
                "text",
                72,
                672,
                94,
                90,
                text=str(page_no),
                color=_PALETTE["white"],
                font_px=58,
                bold=True,
                align="center",
                valign="middle",
                structural=True,
            )
        )
    elif kind == "answer":
        elements.extend(
            [
                _LayoutElement("rect", 0, 0, 1600, 900, fill=_PALETTE["answer"]),
                _LayoutElement("rect", 0, 0, 1600, 22, fill=_PALETTE["teal"]),
                _LayoutElement("rect", 94, 230, 10, 476, fill=_PALETTE["teal"]),
            ]
        )
        if title:
            elements.append(
                _source_element(
                    draw=draw,
                    fonts=fonts,
                    x=96,
                    y=72,
                    width=1320,
                    height=124,
                    source_texts=[title],
                    source_paths=[f"/slides/{index}/title"],
                    preferred_px=58,
                    minimum_px=42,
                    color=_PALETTE["navy"],
                    bold=True,
                    valign="middle",
                )
            )
        if lines:
            elements.append(
                _source_element(
                    draw=draw,
                    fonts=fonts,
                    x=148,
                    y=246,
                    width=1260,
                    height=444,
                    source_texts=lines,
                    source_paths=paths,
                    preferred_px=40,
                    minimum_px=26,
                    color=_PALETTE["ink"],
                    numbered=len(lines) > 1,
                    valign="middle",
                )
            )
        elements.append(
            _LayoutElement(
                "text",
                1400,
                780,
                96,
                48,
                text=str(page_no),
                color=_PALETTE["muted"],
                font_px=25,
                align="right",
                structural=True,
            )
        )
    elif kind == "summary":
        elements.extend(
            [
                _LayoutElement("rect", 0, 0, 1600, 900, fill=_PALETTE["navy"]),
                _LayoutElement("rect", 0, 810, 1600, 90, fill=_PALETTE["teal"]),
            ]
        )
        if title:
            elements.append(
                _source_element(
                    draw=draw,
                    fonts=fonts,
                    x=106,
                    y=74,
                    width=1320,
                    height=132,
                    source_texts=[title],
                    source_paths=[f"/slides/{index}/title"],
                    preferred_px=60,
                    minimum_px=44,
                    color=_PALETTE["white"],
                    bold=True,
                    valign="middle",
                )
            )
        if lines:
            elements.append(
                _source_element(
                    draw=draw,
                    fonts=fonts,
                    x=112,
                    y=262,
                    width=1250,
                    height=420,
                    source_texts=lines,
                    source_paths=paths,
                    preferred_px=39,
                    minimum_px=26,
                    color="E9F1F6",
                    numbered=len(lines) > 1,
                )
            )
        elements.append(
            _LayoutElement(
                "text",
                1412,
                826,
                82,
                44,
                text=str(page_no),
                color=_PALETTE["white"],
                font_px=24,
                align="right",
                valign="middle",
                structural=True,
            )
        )
    else:
        elements.extend(
            [
                _LayoutElement("rect", 0, 0, 1600, 900, fill=_PALETTE["paper"]),
                _LayoutElement("rect", 0, 0, 26, 900, fill=_PALETTE["navy"]),
                _LayoutElement("rect", 94, 196, 126, 8, fill=_PALETTE["teal"]),
            ]
        )
        if title:
            elements.append(
                _source_element(
                    draw=draw,
                    fonts=fonts,
                    x=94,
                    y=62,
                    width=1320,
                    height=118,
                    source_texts=[title],
                    source_paths=[f"/slides/{index}/title"],
                    preferred_px=56,
                    minimum_px=42,
                    color=_PALETTE["navy"],
                    bold=True,
                    valign="middle",
                )
            )
        if lines:
            elements.append(
                _source_element(
                    draw=draw,
                    fonts=fonts,
                    x=98,
                    y=248,
                    width=1270,
                    height=470,
                    source_texts=lines,
                    source_paths=paths,
                    preferred_px=38,
                    minimum_px=25,
                    color=_PALETTE["ink"],
                    numbered=len(lines) > 1,
                )
            )
        elements.append(
            _LayoutElement(
                "text",
                1398,
                790,
                96,
                44,
                text=str(page_no),
                color=_PALETTE["muted"],
                font_px=24,
                align="right",
                structural=True,
            )
        )

    return _SlideLayout(page_no, slide_id, kind, title, tuple(elements))


def _image_elements(slide, *, index, candidate, image_data, draw, fonts):
    try:
        assets = normalize_image_assets(candidate.get("image_assets", []))
        binding = normalize_slide_image(slide["image"], assets)
        if binding is None or slide.get("visual") is not None:
            raise PreparationImageError("图片页不能同时包含其他图形版式。")
        asset_index, asset = next(
            (i, a) for i, a in enumerate(assets) if a["asset_id"] == binding["asset_id"]
        )
        data = image_data.get(asset["asset_id"])
        if not isinstance(data, bytes):
            raise PreparationImageError("图片页缺少本地图片副本。")
        verify_image_bytes(asset, data)
    except PreparationImageError as exc:
        raise NativePreparationRenderError(exc.code, exc.message_zh) from exc
    if asset["content_type"] == "image/webp":
        from PIL import Image

        stream = io.BytesIO()
        with Image.open(io.BytesIO(data)) as picture:
            picture.save(stream, format="PNG")
        data = stream.getvalue()
    ratio = min(950 / asset["width"], 500 / asset["height"])
    width, height = round(asset["width"] * ratio), round(asset["height"] * ratio)
    base = f"/slides/{index}"
    asset_base = f"/image_assets/{asset_index}"
    elements = [
        _LayoutElement("rect", 0, 0, 1600, 900, fill=_PALETTE["paper"]),
        _LayoutElement(
            "image",
            84 + (950 - width) // 2,
            240 + (500 - height) // 2,
            width,
            height,
            image_data=data,
            image_asset_id=asset["asset_id"],
            image_source_sha256=asset["sha256"],
            image_alt=asset["caption"],
        ),
    ]
    elements.append(
        _source_element(
            draw=draw,
            fonts=fonts,
            x=84,
            y=76,
            width=1430,
            height=128,
            source_texts=[_as_text(slide.get("title"))],
            source_paths=[base + "/title"],
            preferred_px=58,
            minimum_px=40,
            color=_PALETTE["navy"],
            bold=True,
        )
    )
    content, paths = _slide_source_lines(slide, index)
    elements.append(
        _source_element(
            draw=draw,
            fonts=fonts,
            x=1090,
            y=256,
            width=426,
            height=468,
            source_texts=[binding["observation_prompt"], *content],
            source_paths=[base + "/image/observation_prompt", *paths],
            preferred_px=36,
            minimum_px=28,
            color=_PALETTE["ink"],
            paragraph_separator="\n\n",
        )
    )
    elements.append(
        _source_element(
            draw=draw,
            fonts=fonts,
            x=84,
            y=760,
            width=1430,
            height=96,
            source_texts=[asset["caption"], asset["source"]],
            source_paths=[asset_base + "/caption", asset_base + "/source"],
            preferred_px=26,
            minimum_px=22,
            color=_PALETTE["muted"],
            paragraph_separator="\n",
        )
    )
    return elements


def _make_layouts(
    candidate: Mapping[str, Any],
    image_data: Mapping[str, bytes] | None = None,
    *,
    classroom_projection: bool = False,
) -> list[_SlideLayout]:
    slides = candidate.get("slides")
    if not isinstance(slides, list) or not slides:
        raise NativePreparationRenderError(
            "preparation_candidate_slides_invalid", "备课候选缺少可生成的幻灯片页纲。"
        )
    if any(not isinstance(slide, Mapping) for slide in slides):
        raise NativePreparationRenderError(
            "preparation_candidate_slides_invalid", "备课候选幻灯片页纲不正确。"
        )
    try:
        from PIL import Image, ImageDraw
    except ImportError as exc:
        raise NativePreparationRenderError(
            "preparation_render_dependency_missing",
            "缺少 Pillow，无法生成备课逐页预览。",
        ) from exc
    image = Image.new("RGB", (SLIDE_WIDTH_PX, SLIDE_HEIGHT_PX), "white")
    draw = ImageDraw.Draw(image)
    fonts = _FontBook()
    layouts = [
        _layout_slide(
            slide,
            index=index,
            page_no=index + 1,
            slide_count=len(slides),
            candidate=candidate,
            draw=draw,
            fonts=fonts,
            image_data=image_data,
        )
        for index, slide in enumerate(slides)
    ]
    if classroom_projection:
        from .desktop_preparation_classroom_layout import classroom_layout

        layouts = [classroom_layout(item, draw=draw, fonts=fonts) for item in layouts]
    return layouts


def _pptx_coord(px: int) -> int:
    from pptx.util import Inches

    return Inches(px / 120)


def _set_pptx_run_font(run: Any, *, size_px: int, bold: bool, color: str) -> None:
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Pt

    run.font.name = _FONT_FALLBACK[0]
    run.font.size = Pt(size_px * 72 / 120)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    properties = run._r.get_or_add_rPr()
    east_asia = properties.find(qn("a:ea"))
    if east_asia is None:
        east_asia = OxmlElement("a:ea")
        properties.append(east_asia)
    east_asia.set("typeface", _FONT_FALLBACK[0])


def _add_pptx_element(slide: Any, element: _LayoutElement) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Pt

    x = _pptx_coord(element.x)
    y = _pptx_coord(element.y)
    width = _pptx_coord(element.width)
    height = _pptx_coord(element.height)
    if element.element_type == "image":
        shape = slide.shapes.add_picture(
            io.BytesIO(element.image_data), x, y, width=width, height=height
        )
        shape._element.nvPicPr.cNvPr.set("descr", element.image_alt)
        return
    if element.element_type in {"rect", "right_arrow"}:
        shape_type = (
            MSO_SHAPE.RIGHT_ARROW
            if element.element_type == "right_arrow"
            else MSO_SHAPE.RECTANGLE
        )
        shape = slide.shapes.add_shape(shape_type, x, y, width, height)
        # Override theme effects so Office does not introduce shadows absent
        # from the shared preview layout.
        shape._element.spPr.append(OxmlElement("a:effectLst"))
        if element.element_type == "right_arrow":
            shape.adjustments[0] = 0.5
            shape.adjustments[1] = 0.5
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(
            element.fill or _PALETTE["paper"]
        )
        if element.line:
            shape.line.color.rgb = RGBColor.from_string(element.line)
            shape.line.width = Pt(0.6)
        else:
            shape.line.fill.background()
        return
    if element.element_type != "text":
        return
    shape = slide.shapes.add_textbox(x, y, width, height)
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = {
        "top": MSO_ANCHOR.TOP,
        "middle": MSO_ANCHOR.MIDDLE,
        "bottom": MSO_ANCHOR.BOTTOM,
    }.get(element.valign, MSO_ANCHOR.TOP)
    text_lines = element.text.split("\n")
    for line_index, line in enumerate(text_lines):
        paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
        paragraph.alignment = {
            "left": PP_ALIGN.LEFT,
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
        }.get(element.align, PP_ALIGN.LEFT)
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = Pt(element.font_px * 72 / 120 * 1.18)
        run = paragraph.add_run()
        run.text = line
        _set_pptx_run_font(
            run,
            size_px=element.font_px,
            bold=element.bold,
            color=element.color,
        )


def _write_pptx(
    path: Path,
    layouts: Sequence[_SlideLayout],
    candidate: Mapping[str, Any],
    *,
    root: Path,
) -> None:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise NativePreparationRenderError(
            "preparation_render_dependency_missing",
            "缺少 python-pptx，无法生成可编辑 PPTX。",
        ) from exc
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(SLIDE_HEIGHT_IN)
    while presentation.slides:
        slide_id = presentation.slides._sldIdLst[0]
        presentation.part.drop_rel(slide_id.rId)
        del presentation.slides._sldIdLst[0]
    presentation.core_properties.title = _as_text(candidate.get("title"))
    presentation.core_properties.subject = _as_text(candidate.get("topic"))
    for layout_index, layout in enumerate(layouts):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for element in layout.elements:
            _add_pptx_element(slide, element)
        original = candidate["slides"][layout_index]
        notes = []
        minutes = original.get("minutes")
        if type(minutes) is int and minutes >= 0:
            notes.append(f"建议用时：{minutes} 分钟")
        purpose = _as_text(original.get("purpose"))
        if purpose:
            notes.append("本页教学意图：" + purpose)
        teacher_notes = _as_text(original.get("teacher_notes"))
        if teacher_notes:
            notes.append(teacher_notes)
        if original.get("image"):
            binding = original["image"]
            asset = next(
                a
                for a in candidate["image_assets"]
                if a["asset_id"] == binding["asset_id"]
            )
            notes.extend(
                [
                    "看图任务：" + binding["observation_prompt"],
                    "图片教学用途：" + asset["purpose"],
                    "图片来源：" + asset["source"],
                    ("图片内容需教师复核；本页由本地教学环节生成，未调用模型。"
                     if candidate.get("source_basis", {}).get("mode") == "teacher_nodes_local"
                     else "图片内容需教师复核；模型只收到图片说明，未查看像素。"),
                ]
            )
        if notes:
            slide.notes_slide.notes_text_frame.text = "\n\n".join(notes)
    temporary = _contained(root, path.with_name(f".{path.name}.{uuid4().hex}.tmp"))
    try:
        presentation.save(temporary)
        if path.exists():
            raise NativePreparationRenderError(
                "preparation_render_artifact_exists",
                "本次生成目录中已有同名产物，已停止覆盖。",
                details={"filename": path.name},
            )
        os.replace(temporary, path)
    except NativePreparationRenderError:
        temporary.unlink(missing_ok=True)
        raise
    except OSError as exc:
        temporary.unlink(missing_ok=True)
        raise NativePreparationRenderError(
            "preparation_pptx_write_failed",
            "可编辑 PPTX 写入失败。",
            retryable=True,
            details={"exception_type": type(exc).__name__},
        ) from exc


def _draw_layout(layout: _SlideLayout, fonts: _FontBook) -> Any:
    try:
        from PIL import Image, ImageDraw
    except ImportError as exc:
        raise NativePreparationRenderError(
            "preparation_render_dependency_missing",
            "缺少 Pillow，无法生成备课逐页预览。",
        ) from exc
    image = Image.new(
        "RGB", (SLIDE_WIDTH_PX, SLIDE_HEIGHT_PX), _hex_rgb(_PALETTE["paper"])
    )
    draw = ImageDraw.Draw(image)
    for element in layout.elements:
        if element.element_type == "image":
            with Image.open(io.BytesIO(element.image_data)) as picture:
                rgba = picture.convert("RGBA").resize(
                    (element.width, element.height), Image.Resampling.LANCZOS
                )
                image.paste(rgba, (element.x, element.y), rgba)
            continue
        bounds = (
            element.x,
            element.y,
            element.x + element.width,
            element.y + element.height,
        )
        if element.element_type == "rect":
            draw.rectangle(
                bounds,
                fill=_hex_rgb(element.fill or _PALETTE["paper"]),
                outline=_hex_rgb(element.line) if element.line else None,
                width=1,
            )
            continue
        if element.element_type == "right_arrow":
            x, y, right, bottom = bounds
            neck = right - element.width / 2
            draw.polygon(
                [
                    (x, y + element.height / 4),
                    (neck, y + element.height / 4),
                    (neck, y),
                    (right, y + element.height / 2),
                    (neck, bottom),
                    (neck, bottom - element.height / 4),
                    (x, bottom - element.height / 4),
                ],
                fill=_hex_rgb(element.fill or _PALETTE["teal"]),
            )
            continue
        if element.element_type != "text" or not element.text:
            continue
        font = fonts.get(element.font_px, element.bold)
        spacing = max(5, int(element.font_px * 0.22))
        text_box = _multiline_textbbox(draw, element.text, font, spacing, element.align)
        text_width = max(0, text_box[2] - text_box[0])
        text_height = max(0, text_box[3] - text_box[1])
        text_x = element.x
        if element.align == "center":
            text_x = element.x + max(0, (element.width - text_width) // 2)
        elif element.align == "right":
            text_x = element.x + max(0, element.width - text_width)
        text_y = element.y
        if element.valign == "middle":
            text_y = element.y + max(0, (element.height - text_height) // 2)
        elif element.valign == "bottom":
            text_y = element.y + max(0, element.height - text_height)
        _draw_multiline_text(
            draw,
            (text_x, text_y),
            element.text,
            font=font,
            fill=_hex_rgb(element.color),
            spacing=spacing,
            align=element.align,
        )
    return image


def _save_png(image: Any, path: Path, *, root: Path) -> None:
    temporary = _contained(root, path.with_name(f".{path.name}.{uuid4().hex}.tmp"))
    try:
        image.save(temporary, format="PNG", optimize=True)
        if path.exists():
            raise NativePreparationRenderError(
                "preparation_render_artifact_exists",
                "本次生成目录中已有同名产物，已停止覆盖。",
                details={"filename": path.name},
            )
        os.replace(temporary, path)
    except NativePreparationRenderError:
        temporary.unlink(missing_ok=True)
        raise
    except OSError as exc:
        temporary.unlink(missing_ok=True)
        raise NativePreparationRenderError(
            "preparation_png_write_failed",
            "备课逐页预览写入失败。",
            retryable=True,
            details={"filename": path.name, "exception_type": type(exc).__name__},
        ) from exc


def _montage(images: Sequence[Any]) -> Any:
    from PIL import Image, ImageDraw

    canvas = Image.new(
        "RGB", (SLIDE_WIDTH_PX, SLIDE_HEIGHT_PX), _hex_rgb(_PALETTE["paper"])
    )
    if not images:
        return canvas
    count = len(images)
    columns = max(1, min(4, math.ceil(math.sqrt(count * 16 / 9))))
    rows = math.ceil(count / columns)
    margin_x, margin_y, gap = 42, 42, 20
    cell_width = (SLIDE_WIDTH_PX - 2 * margin_x - (columns - 1) * gap) // columns
    cell_height = (SLIDE_HEIGHT_PX - 2 * margin_y - (rows - 1) * gap) // rows
    thumb_width = min(cell_width, int(cell_height * 16 / 9))
    thumb_height = int(thumb_width * 9 / 16)
    if thumb_height > cell_height:
        thumb_height = cell_height
        thumb_width = int(thumb_height * 16 / 9)
    draw = ImageDraw.Draw(canvas)
    for index, source in enumerate(images):
        row, column = divmod(index, columns)
        x = margin_x + column * (cell_width + gap) + (cell_width - thumb_width) // 2
        y = margin_y + row * (cell_height + gap) + (cell_height - thumb_height) // 2
        resized = source.resize((thumb_width, thumb_height), Image.Resampling.LANCZOS)
        draw.rectangle(
            (x - 3, y - 3, x + thumb_width + 3, y + thumb_height + 3),
            fill=_hex_rgb(_PALETTE["line"]),
        )
        canvas.paste(resized, (x, y))
    return canvas


def _set_docx_font(
    run: Any, *, size: float = 11, bold: bool = False, color: str = "000000"
) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = _FONT_FALLBACK[0]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), _FONT_FALLBACK[0])


def _set_cell_shading(cell: Any, color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_table_borders(table: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = table._tbl.tblPr
    existing = properties.find(qn("w:tblBorders"))
    if existing is not None:
        properties.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")
        borders.append(element)
    properties.append(borders)


def _add_docx_lines(document: Any, values: Any) -> None:
    for text in _string_values(values):
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = 0
        _set_docx_font(paragraph.add_run(text), size=11)


def _remove_docx_paragraph_borders(properties: Any) -> None:
    from docx.oxml.ns import qn

    for border in properties.findall(qn("w:pBdr")):
        properties.remove(border)


def _uncertainty_lines(value: Any) -> list[str]:
    items = (
        value
        if isinstance(value, Sequence)
        and not isinstance(value, (str, bytes, bytearray))
        else [value]
    )
    result: list[str] = []
    for item in items:
        if isinstance(item, Mapping):
            description = _as_text(item.get("description"))
            teacher_action = _as_text(item.get("teacher_action"))
            if description and teacher_action:
                text = f"{description} 教师处理：{teacher_action}"
            elif description:
                text = description
            elif teacher_action:
                text = f"教师处理：{teacher_action}"
            else:
                continue
        else:
            text = _as_text(item)
        if text and text not in result:
            result.append(text)
    return result


def _numbered_title(label: str, number: int, title: Any) -> str:
    text = _as_text(title).strip()
    # Remove only this exact redundant display prefix. A different number or
    # a word such as 活动10/活动1号 remains visible for review.
    prefix = rf"^{re.escape(label)}\s*{number}(?=$|[\s:：、.．])\s*[:：、.．]?\s*"
    text = re.sub(prefix, "", text, count=1)
    return f"{label}{number}" + (f" {text}" if text else "")


def _semantic_text(value: Any) -> str:
    values = _string_values(value)
    return "；".join(values)


def _metadata_text(candidate: Mapping[str, Any], key: str) -> str:
    value = candidate.get(key)
    if key == "source_basis" and isinstance(value, Mapping):
        statement = _as_text(value.get("statement_zh"))
        return statement or _semantic_text(value)
    if key == "lesson_route":
        route = _as_text(value)
        return _LESSON_ROUTE_LABELS.get(route, route)
    if key == "timing" and isinstance(value, Mapping):
        periods = value.get("periods")
        minutes_per_period = value.get("minutes_per_period")
        total_minutes = value.get("total_minutes")
        if (
            type(periods) is int
            and type(minutes_per_period) is int
            and type(total_minutes) is int
        ):
            return f"{periods} 课时  每课时 {minutes_per_period} 分钟  共 {total_minutes} 分钟"
    return _semantic_text(value)


def _write_docx(path: Path, candidate: Mapping[str, Any], *, root: Path) -> None:
    try:
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise NativePreparationRenderError(
            "preparation_render_dependency_missing",
            "缺少 python-docx，无法生成可编辑教案。",
        ) from exc

    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = document.styles["Normal"]
    normal.font.name = _FONT_FALLBACK[0]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12.5)):
        style = document.styles[style_name]
        style.font.name = _FONT_FALLBACK[0]
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), _FONT_FALLBACK[0])
        if style_name == "Title":
            _remove_docx_paragraph_borders(style._element.get_or_add_pPr())

    title = _as_text(candidate.get("title")) or _as_text(candidate.get("topic"))
    title_paragraph = document.add_paragraph(style="Title")
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_paragraph.paragraph_format.space_after = Pt(16)
    _remove_docx_paragraph_borders(title_paragraph._p.get_or_add_pPr())
    _set_docx_font(title_paragraph.add_run(title), size=22, bold=True)

    metadata: list[tuple[str, str]] = []
    for label, key in (
        ("课题", "topic"),
        ("对象", "audience"),
        ("课型", "lesson_route"),
        ("课时", "timing"),
        ("内容依据", "source_basis"),
    ):
        text = _metadata_text(candidate, key)
        if text and text != title:
            metadata.append((label, text))
    if metadata:
        table = document.add_table(rows=0, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        for column, width in zip(table.columns, (1.05, 5.85), strict=True):
            column.width = Inches(width)
        _set_table_borders(table)
        for row_index, (label, value) in enumerate(metadata):
            cells = table.add_row().cells
            cells[0].width = Inches(1.05)
            cells[1].width = Inches(5.85)
            cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_shading(cells[0], "E8F0F6")
            if row_index % 2:
                _set_cell_shading(cells[1], "F7F9FB")
            _set_docx_font(cells[0].paragraphs[0].add_run(label), size=10.5, bold=True)
            _set_docx_font(cells[1].paragraphs[0].add_run(value), size=10.5)
        document.add_paragraph().paragraph_format.space_after = Pt(2)

    objectives = candidate.get("objectives")
    if isinstance(objectives, list) and objectives:
        document.add_heading("教学目标", level=1)
        for objective in objectives:
            if not isinstance(objective, Mapping):
                continue
            statement = _as_text(objective.get("statement"))
            if statement:
                paragraph = document.add_paragraph(style="List Number")
                _set_docx_font(paragraph.add_run(statement), size=11)

    activities = candidate.get("activities")
    assessments = candidate.get("assessments")
    if isinstance(activities, list) and activities:
        document.add_heading("目标 活动与评价", level=1)
        table = document.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        for column, width in zip(table.columns, (1.0, 2.95, 2.95), strict=True):
            column.width = Inches(width)
        _set_table_borders(table)
        headers = ("教学目标", "学习活动", "评价证据")
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell, header, column in zip(
            table.rows[0].cells, headers, table.columns, strict=True
        ):
            # Changing tblGrid does not update cells created by add_table.
            # Keep the header's preferred widths consistent with the body.
            cell.width = column.width
            _set_cell_shading(cell, _PALETTE["navy"])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_docx_font(
                cell.paragraphs[0].add_run(header), size=10.5, bold=True, color="FFFFFF"
            )
        objective_by_id = {
            _as_text(item.get("id")): f"目标{index}"
            for index, item in enumerate(objectives or [], 1)
            if isinstance(item, Mapping)
        }
        assessment_rows = [
            item for item in assessments or [] if isinstance(item, Mapping)
        ]
        for row_index, activity in enumerate(activities):
            if not isinstance(activity, Mapping):
                continue
            objective_ids = activity.get("objective_ids")
            objective_text = "\n".join(
                objective_by_id.get(_as_text(item), _as_text(item))
                for item in objective_ids or []
                if _as_text(item)
            )
            activity_text = _numbered_title(
                "活动", row_index + 1, activity.get("title")
            )
            activity_id = _as_text(activity.get("id"))
            assessment_texts: list[str] = []
            for assessment_index, assessment in enumerate(assessment_rows, 1):
                ids = assessment.get("activity_ids")
                if isinstance(ids, list) and activity_id in ids:
                    assessment_texts.append(
                        _numbered_title(
                            "评价", assessment_index, assessment.get("title")
                        )
                    )
            # One mapping row is one record: do not leave its remaining
            # objective IDs alone under the repeated header on the next page.
            # Keep the table itself splittable between records.
            row = table.add_row()
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            cells = row.cells
            if row_index % 2:
                for cell in cells:
                    _set_cell_shading(cell, "F4F8FA")
            for cell, text, width in zip(
                cells,
                (objective_text, activity_text, "\n".join(assessment_texts)),
                (1.0, 2.95, 2.95),
            ):
                cell.width = Inches(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                _set_docx_font(cell.paragraphs[0].add_run(text), size=10.2)

    activity_rows = [item for item in activities or [] if isinstance(item, Mapping)]
    assessment_rows = [item for item in assessments or [] if isinstance(item, Mapping)]
    written_activities: set[int] = set()
    activity_coverage: dict[int, dict[str, set[str]]] = {}

    def write_activity_details(
        index: int,
        activity: Mapping[str, Any],
        covered: Mapping[str, set[str]],
        section_title: str | None = None,
    ) -> bool:
        # Deduplicate only identical, same-role text for an explicitly linked
        # activity. Never infer equivalence from titles or paraphrase content.
        fields = []
        for label, key in (
            ("教师活动", "teacher_action"),
            ("学生活动", "student_action"),
            ("材料", "materials"),
        ):
            values = [
                value
                for value in _string_values(activity.get(key))
                if value not in covered.get(label, set())
            ]
            if values:
                fields.append((label, values))
        worksheet = activity.get("worksheet")
        if not fields and not isinstance(worksheet, Mapping):
            return False
        if section_title:
            document.add_heading(section_title, level=1)
        document.add_heading(
            _numbered_title("活动", index + 1, activity.get("title")), level=2
        )
        if isinstance(worksheet, Mapping):
            paragraph = document.add_paragraph()
            _set_docx_font(
                paragraph.add_run(
                    "配套学习单："
                    + _as_text(worksheet.get("title"))
                    + "。见同批产物 student_worksheet.docx。"
                )
            )
        for label, values in fields:
            paragraph = document.add_paragraph()
            _set_docx_font(paragraph.add_run(label + "："), bold=True)
            _set_docx_font(paragraph.add_run("\n".join(values)))
        return True

    lesson_stages = candidate.get("lesson_stages")
    if isinstance(lesson_stages, list) and lesson_stages:
        document.add_heading("教学过程", level=1)
        for stage in lesson_stages:
            if not isinstance(stage, Mapping):
                continue
            stage_title = _as_text(stage.get("title") or stage.get("name"))
            minutes = stage.get("minutes")
            if isinstance(minutes, int) and not isinstance(minutes, bool):
                stage_title = (
                    f"{stage_title}  {minutes} 分钟"
                    if stage_title
                    else f"{minutes} 分钟"
                )
            if stage_title:
                document.add_heading(stage_title, level=2)
            activity_ids = stage.get("activity_ids") or []
            linked = [
                (index, activity)
                for index, activity in enumerate(activity_rows)
                if activity.get("id") and activity.get("id") in activity_ids
            ]
            if linked:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.keep_with_next = True
                _set_docx_font(
                    paragraph.add_run(
                        "关联活动："
                        + "、".join(f"活动{index + 1}" for index, _ in linked)
                    ),
                    size=10,
                )
            stage_fields = (
                (
                    "教师活动",
                    stage.get("teacher_actions") or stage.get("teacher_action"),
                ),
                (
                    "学生活动",
                    stage.get("student_actions") or stage.get("student_action"),
                ),
                ("材料", stage.get("materials")),
                ("评价", stage.get("assessment") or stage.get("assessments")),
            )
            covered: dict[str, set[str]] = {}
            for label, value in stage_fields:
                values = _string_values(value)
                covered[label] = set(values)
                if not values:
                    continue
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.space_after = Pt(5)
                _set_docx_font(paragraph.add_run(label + "："), size=11, bold=True)
                _set_docx_font(paragraph.add_run("\n".join(values)), size=11)

            for index, activity in linked:
                written_activities.add(index)
                all_covered = activity_coverage.setdefault(index, {})
                for label, values in covered.items():
                    all_covered.setdefault(label, set()).update(values)

    # A single activity may span introduction, practice and feedback stages.
    # Do not insert its entire summary under the first (often introductory)
    # stage. Keep additional resources separate from the executable timeline.
    supplements_started = False
    for index in sorted(written_activities):
        written = write_activity_details(
            index,
            activity_rows[index],
            activity_coverage[index],
            None if supplements_started else "活动补充资料",
        )
        supplements_started = supplements_started or written

    remaining = [
        (index, activity)
        for index, activity in enumerate(activity_rows)
        if index not in written_activities
    ]
    if remaining:
        # Preserve older or incomplete candidates without guessing their place
        # in the timeline; normalized linked candidates need no second process.
        document.add_heading("补充活动" if lesson_stages else "教学活动", level=1)
        for index, activity in remaining:
            write_activity_details(index, activity, {})

    if assessment_rows:
        document.add_heading("评价标准", level=1)
        for index, assessment in enumerate(assessment_rows, 1):
            document.add_heading(
                _numbered_title("评价", index, assessment.get("title")), level=2
            )
            _add_docx_lines(
                document, _string_values(assessment.get("evidence_of_learning"))
            )
            _add_docx_lines(
                document, _string_values(assessment.get("success_criteria"))
            )

    homework = candidate.get("homework")
    if isinstance(homework, Mapping):
        homework_values = _string_values(homework.get("tasks"))
        homework_title = _as_text(homework.get("title"))
        if homework_title or homework_values:
            document.add_heading(homework_title or "课后任务", level=1)
            _add_docx_lines(document, homework_values)

    uncertainties = candidate.get("uncertainties")
    uncertainty_values = _uncertainty_lines(uncertainties)
    if uncertainty_values:
        document.add_heading("待教师确认", level=1)
        _add_docx_lines(document, uncertainty_values)

    document.core_properties.title = title
    _save_docx_document(document, path, root=root)


def _save_docx_document(document: Any, path: Path, *, root: Path) -> None:
    temporary = _contained(root, path.with_name(f".{path.name}.{uuid4().hex}.tmp"))
    try:
        document.save(temporary)
        if path.exists():
            raise NativePreparationRenderError(
                "preparation_render_artifact_exists",
                "本次生成目录中已有同名产物，已停止覆盖。",
                details={"filename": path.name},
            )
        os.replace(temporary, path)
    except NativePreparationRenderError:
        temporary.unlink(missing_ok=True)
        raise
    except OSError as exc:
        temporary.unlink(missing_ok=True)
        raise NativePreparationRenderError(
            "preparation_docx_write_failed",
            "可编辑教案写入失败。",
            retryable=True,
            details={"exception_type": type(exc).__name__},
        ) from exc


def _artifact_row(root: Path, artifact_id: str, path: Path) -> dict[str, Any]:
    safe_path = _contained(root, path)
    if safe_path.is_symlink() or not safe_path.is_file():
        raise NativePreparationRenderError(
            "preparation_render_artifact_missing",
            "备课产物缺失或路径不安全。",
            details={"artifact_id": artifact_id},
        )
    data = safe_path.read_bytes()
    return {
        "artifact_id": artifact_id,
        "filename": safe_path.relative_to(root).as_posix(),
        "path": str(safe_path),
        "content_type": _CONTENT_TYPES[artifact_id],
        "size": len(data),
        "sha256": _sha256(data),
    }


def _zip_check(path: Path, required: Sequence[str]) -> tuple[bool, list[str]]:
    try:
        with zipfile.ZipFile(path) as archive:
            bad_member = archive.testzip()
            names = set(archive.namelist())
    except (OSError, zipfile.BadZipFile):
        return False, list(required)
    missing = [name for name in required if name not in names]
    return bad_member is None and not missing, missing


def _deck_json(
    candidate: Mapping[str, Any], layouts: Sequence[_SlideLayout]
) -> dict[str, Any]:
    return {
        "schema_version": DECK_SCHEMA_VERSION,
        "candidate_id": _as_text(candidate.get("candidate_id")),
        "title": _as_text(candidate.get("title")),
        "topic": _as_text(candidate.get("topic")),
        "slide_size": {
            "aspect_ratio": "16:9",
            "width_px": SLIDE_WIDTH_PX,
            "height_px": SLIDE_HEIGHT_PX,
            "width_in": SLIDE_WIDTH_IN,
            "height_in": SLIDE_HEIGHT_IN,
        },
        "font_fallback": list(_FONT_FALLBACK),
        "palette": dict(_PALETTE),
        "layout_source": "shared_native_layout_plan",
        "slides": [layout.public() for layout in layouts],
        "candidate_only": True,
        "teacher_review_required": True,
        "publication_allowed": False,
        "official_claim_allowed": False,
    }


def _machine_qa(
    *,
    candidate: Mapping[str, Any],
    output_kind: str,
    layouts: Sequence[_SlideLayout],
    pptx_path: Path | None,
    slide_paths: Sequence[Path],
    montage_path: Path | None,
    docx_path: Path | None,
    worksheet_path: Path | None = None,
) -> dict[str, Any]:
    candidate_texts = _candidate_strings(candidate)
    unbound: list[dict[str, Any]] = []
    overflow: list[dict[str, Any]] = []
    for layout in layouts:
        for element_index, element in enumerate(layout.elements):
            if element.element_type != "text":
                continue
            if element.overflow:
                overflow.append(
                    {"page_no": layout.page_no, "element_index": element_index}
                )
            if element.structural:
                continue
            for source_text in element.source_texts:
                if source_text not in candidate_texts:
                    unbound.append(
                        {
                            "page_no": layout.page_no,
                            "element_index": element_index,
                            "source_text": source_text,
                        }
                    )

    checks: list[dict[str, Any]] = []
    if pptx_path is not None:
        pptx_ok, pptx_missing = _zip_check(
            pptx_path, ["[Content_Types].xml", "ppt/presentation.xml"]
        )
        from pptx import Presentation

        presentation = Presentation(pptx_path)
        slide_count_ok = len(presentation.slides) == len(layouts)
        dimensions_ok = presentation.slide_width * 9 == presentation.slide_height * 16
        editable_text_shapes = sum(
            1
            for slide in presentation.slides
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        )
        picture_checks = []
        for layout, rendered in zip(layouts, presentation.slides):
            planned = [e for e in layout.elements if e.element_type == "image"]
            pictures = [s for s in rendered.shapes if s.shape_type == 13]
            matched = len(planned) == len(pictures)
            for element, picture in zip(planned, pictures):
                matched = matched and (
                    _sha256(picture.image.blob) == _sha256(element.image_data)
                    and abs(
                        picture.left
                        - element.x * presentation.slide_width / SLIDE_WIDTH_PX
                    )
                    <= 2
                    and abs(
                        picture.top
                        - element.y * presentation.slide_height / SLIDE_HEIGHT_PX
                    )
                    <= 2
                    and abs(
                        picture.width
                        - element.width * presentation.slide_width / SLIDE_WIDTH_PX
                    )
                    <= 2
                    and abs(
                        picture.height
                        - element.height * presentation.slide_height / SLIDE_HEIGHT_PX
                    )
                    <= 2
                )
            picture_checks.append(
                {
                    "page_no": layout.page_no,
                    "image_count": len(planned),
                    "passed": matched,
                }
            )
        checks.extend(
            [
                {
                    "check": "pptx_ooxml_package",
                    "passed": pptx_ok,
                    "missing_members": pptx_missing,
                },
                {"check": "pptx_slide_count", "passed": slide_count_ok},
                {"check": "pptx_widescreen_16_9", "passed": dimensions_ok},
                {
                    "check": "pptx_local_images_match_layout_and_bytes",
                    "passed": all(row["passed"] for row in picture_checks),
                    "pages": picture_checks,
                },
                {
                    "check": "pptx_editable_native_text",
                    "passed": editable_text_shapes > 0,
                    "editable_text_shape_count": editable_text_shapes,
                },
                {
                    "check": "pptx_presenter_notes_preserved",
                    "passed": all(
                        all(
                            _as_text(original.get(key))
                            in rendered.notes_slide.notes_text_frame.text
                            for key in ("purpose", "teacher_notes")
                            if _as_text(original.get(key))
                        )
                        for original, rendered in zip(
                            candidate["slides"], presentation.slides
                        )
                    ),
                },
            ]
        )
        png_dimensions_ok = True
        png_rows: list[dict[str, Any]] = []
        try:
            from PIL import Image

            for page_no, path in enumerate(slide_paths, 1):
                with Image.open(path) as image:
                    dimensions = list(image.size)
                    image.verify()
                data = path.read_bytes()
                if dimensions != [SLIDE_WIDTH_PX, SLIDE_HEIGHT_PX]:
                    png_dimensions_ok = False
                png_rows.append(
                    {
                        "page_no": page_no,
                        "filename": path.name,
                        "dimensions": dimensions,
                        "size": len(data),
                        "size_bytes": len(data),
                        "sha256": _sha256(data),
                    }
                )
            montage_ok = False
            montage_dimensions: list[int] = []
            if montage_path is not None:
                with Image.open(montage_path) as image:
                    montage_dimensions = list(image.size)
                    image.verify()
                montage_ok = montage_dimensions == [SLIDE_WIDTH_PX, SLIDE_HEIGHT_PX]
        except (OSError, ValueError):
            png_dimensions_ok = False
            montage_ok = False
            montage_dimensions = []
        checks.extend(
            [
                {
                    "check": "rendered_slide_dimensions_1600x900",
                    "passed": png_dimensions_ok and len(slide_paths) == len(layouts),
                },
                {
                    "check": "montage_dimensions_1600x900",
                    "passed": montage_ok,
                    "dimensions": montage_dimensions,
                },
                {
                    "check": "pptx_and_png_share_layout_plan",
                    "passed": True,
                    "layout_source": "shared_native_layout_plan",
                },
            ]
        )
    else:
        png_rows = []

    if docx_path is not None:
        docx_ok, docx_missing = _zip_check(
            docx_path, ["[Content_Types].xml", "word/document.xml"]
        )
        checks.append(
            {
                "check": "docx_ooxml_package",
                "passed": docx_ok,
                "missing_members": docx_missing,
            }
        )

    if worksheet_path is not None:
        from docx import Document

        worksheet_ok, worksheet_missing = _zip_check(
            worksheet_path, ["[Content_Types].xml", "word/document.xml"]
        )
        document = Document(worksheet_path)
        visible = {paragraph.text for paragraph in document.paragraphs}
        visible.update(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        expected: set[str] = set()
        for _index, _activity, worksheet in worksheet_records(candidate):
            expected.add(worksheet["title"])
            expected.update(worksheet["instructions"])
            for section in worksheet["sections"]:
                expected.update([section["heading"], section["prompt"]])
                expected.update(section["columns"])
                expected.update(section["row_labels"])
        checks.extend(
            [
                {
                    "check": "worksheet_ooxml_package",
                    "passed": worksheet_ok,
                    "missing_members": worksheet_missing,
                },
                {
                    "check": "worksheet_student_text_preserved",
                    "passed": expected <= visible,
                    "worksheet_count": len(worksheet_records(candidate)),
                    "missing_text_count": len(expected - visible),
                },
            ]
        )

    checks.extend(
        [
            {
                "check": "visible_candidate_text_binding",
                "passed": not unbound,
                "unbound": unbound,
            },
            {
                "check": "text_box_overflow_estimate",
                "passed": not overflow,
                "items": overflow,
                "method": "Pillow metrics with 0.4-em line-end reserve and 1.18-em line boxes",
            },
        ]
    )
    passed = all(item.get("passed") is True for item in checks)
    return {
        "schema_version": QA_SCHEMA_VERSION,
        "candidate_id": _as_text(candidate.get("candidate_id")),
        "output_kind": output_kind,
        "qa_status": (
            "pass_layout_candidate_teacher_review_pending"
            if passed
            else "layout_warning_candidate_teacher_review_pending"
        ),
        "scope": "machine_layout_check_only",
        "machine_checks_passed": passed,
        "human_visual_review_performed": False,
        "checks": checks,
        "rendered_slides": png_rows,
        # Advisory only; keep layout PASS separate from semantic/teaching QA.
        "classroom_review": classroom_review(candidate),
        "not_checked": [
            "chemistry_correctness",
            "source_authority",
            "answer_correctness",
            "classroom_fit",
            "publication_readiness",
        ],
        "candidate_only": True,
        "teacher_review_required": True,
        "publication_allowed": False,
        "official_claim_allowed": False,
    }


class NativePreparationRenderer:
    """Generate native desktop preparation artifacts inside one attempt root.

    ``joint`` is accepted as a UI alias for the canonical ``linked_bundle``
    output kind.  The instance is callable so it can be injected directly into
    ``DesktopPreparationManager``; ``render`` remains available for explicit use.
    """

    def __call__(
        self,
        candidate: Mapping[str, Any],
        *,
        output_kind: str,
        output_dir: str | os.PathLike[str],
        report_progress: Callable[..., Any] | None = None,
        is_cancelled: Callable[[], bool] | None = None,
        image_data: Mapping[str, bytes] | None = None,
    ) -> dict[str, Any]:
        return self.render(
            candidate,
            output_kind=output_kind,
            output_dir=output_dir,
            report_progress=report_progress,
            is_cancelled=is_cancelled,
            image_data=image_data,
        )

    def render(
        self,
        candidate: Mapping[str, Any],
        *,
        output_kind: str,
        output_dir: str | os.PathLike[str],
        report_progress: Callable[..., Any] | None = None,
        is_cancelled: Callable[[], bool] | None = None,
        image_data: Mapping[str, bytes] | None = None,
    ) -> dict[str, Any]:
        if not isinstance(candidate, Mapping):
            raise NativePreparationRenderError(
                "preparation_candidate_invalid", "备课候选结构不正确。"
            )
        normalized_kind = _OUTPUT_KIND_ALIASES.get(output_kind)
        if normalized_kind is None:
            raise NativePreparationRenderError(
                "preparation_output_kind_invalid", "备课输出类型不正确。"
            )
        canonical = _plain_json(candidate)
        candidate_kind = _OUTPUT_KIND_ALIASES.get(
            _as_text(canonical.get("artifact_mode"))
        )
        if candidate_kind is not None and candidate_kind != normalized_kind:
            raise NativePreparationRenderError(
                "preparation_output_kind_mismatch",
                "备课候选与本次输出类型不一致。",
            )
        if (
            canonical.get("candidate_only") is not True
            or canonical.get("teacher_review_required") is not True
            or canonical.get("publication_allowed") is not False
            or canonical.get("official_claim_allowed", False) is not False
        ):
            raise NativePreparationRenderError(
                "preparation_candidate_authority_invalid",
                "备课候选未保持教师审核和禁止发布边界。",
            )
        if not _as_text(canonical.get("title")):
            raise NativePreparationRenderError(
                "preparation_candidate_title_invalid", "备课候选缺少标题。"
            )

        _check_cancelled(is_cancelled)
        _report_progress(report_progress, 2, "validating", "正在检查备课候选。")
        root = _safe_output_root(output_dir)
        layouts = _make_layouts(
            canonical, image_data=image_data, classroom_projection=True
        )
        wants_ppt = normalized_kind in {"ppt", "linked_bundle"}
        wants_docx = normalized_kind in {"lesson_plan", "linked_bundle"}
        try:
            wants_worksheet = bool(worksheet_records(canonical))
        except (ValueError, TypeError):
            raise NativePreparationRenderError(
                "preparation_candidate_worksheet_invalid",
                "学习单任务或填写区域不正确。",
            ) from None

        reserved = [root / _FILENAMES["candidate_json"], root / _FILENAMES["qa_report"]]
        if wants_ppt:
            reserved.extend(
                [
                    root / _FILENAMES["deck_json"],
                    root / _FILENAMES["pptx"],
                    root / _FILENAMES["preview_montage"],
                    root / "rendered_slides",
                ]
            )
        if wants_docx:
            reserved.append(root / _FILENAMES["lesson_plan_docx"])
        if wants_worksheet:
            reserved.append(root / _FILENAMES["student_worksheet_docx"])
        for path in reserved:
            safe = _contained(root, path)
            if safe.exists() or safe.is_symlink():
                raise NativePreparationRenderError(
                    "preparation_render_artifact_exists",
                    "本次生成目录中已有保留产物，已停止覆盖。",
                    details={"filename": path.name},
                )

        candidate_path = root / _FILENAMES["candidate_json"]
        _atomic_write(candidate_path, _json_bytes(canonical), root=root)
        _report_progress(report_progress, 9, "candidate_frozen", "备课候选副本已冻结。")
        _check_cancelled(is_cancelled)

        deck_path: Path | None = None
        pptx_path: Path | None = None
        montage_path: Path | None = None
        slide_paths: list[Path] = []
        docx_path: Path | None = None
        worksheet_path: Path | None = None

        if wants_ppt:
            deck = _deck_json(canonical, layouts)
            deck_path = root / _FILENAMES["deck_json"]
            _atomic_write(deck_path, _json_bytes(deck), root=root)
            _report_progress(
                report_progress, 17, "deck_planned", "PPT 共享版式已生成。"
            )
            _check_cancelled(is_cancelled)

            pptx_path = root / _FILENAMES["pptx"]
            _write_pptx(pptx_path, layouts, canonical, root=root)
            _report_progress(
                report_progress, 42, "pptx_written", "可编辑 PPTX 已生成。"
            )
            _check_cancelled(is_cancelled)

            slide_root = _contained(root, root / "rendered_slides")
            slide_root.mkdir(parents=False, exist_ok=False)
            fonts = _FontBook()
            images: list[Any] = []
            for index, layout in enumerate(layouts):
                _check_cancelled(is_cancelled)
                image = _draw_layout(layout, fonts)
                slide_path = _contained(root, slide_root / f"slide-{index + 1}.png")
                _save_png(image, slide_path, root=root)
                images.append(image)
                slide_paths.append(slide_path)
                progress = 44 + round(28 * (index + 1) / len(layouts))
                _report_progress(
                    report_progress,
                    progress,
                    "rendering_slides",
                    f"正在生成第 {index + 1} 页预览。",
                )
            montage_path = root / _FILENAMES["preview_montage"]
            _save_png(_montage(images), montage_path, root=root)
            _report_progress(
                report_progress, 76, "montage_written", "PPT 总览图已生成。"
            )
            _check_cancelled(is_cancelled)

        if wants_docx:
            docx_path = root / _FILENAMES["lesson_plan_docx"]
            _write_docx(docx_path, canonical, root=root)
            _report_progress(
                report_progress, 88, "lesson_plan_written", "可编辑教案已生成。"
            )
            _check_cancelled(is_cancelled)

        if wants_worksheet:
            worksheet_path = root / _FILENAMES["student_worksheet_docx"]
            _save_docx_document(
                build_worksheet_document(canonical), worksheet_path, root=root
            )
            _report_progress(
                report_progress, 92, "worksheet_written", "活动配套学习单已生成。"
            )
            _check_cancelled(is_cancelled)

        qa = _machine_qa(
            candidate=canonical,
            output_kind=normalized_kind,
            layouts=layouts,
            pptx_path=pptx_path,
            slide_paths=slide_paths,
            montage_path=montage_path,
            docx_path=docx_path,
            worksheet_path=worksheet_path,
        )
        qa_path = root / _FILENAMES["qa_report"]
        _atomic_write(qa_path, _json_bytes(qa), root=root)

        artifact_paths: list[tuple[str, Path]] = [("candidate_json", candidate_path)]
        if wants_ppt:
            assert (
                deck_path is not None
                and pptx_path is not None
                and montage_path is not None
            )
            artifact_paths.extend(
                [
                    ("deck_json", deck_path),
                    ("pptx", pptx_path),
                    ("preview_montage", montage_path),
                ]
            )
        if wants_docx:
            assert docx_path is not None
            artifact_paths.append(("lesson_plan_docx", docx_path))
        if worksheet_path is not None:
            artifact_paths.append(("student_worksheet_docx", worksheet_path))
        artifact_paths.append(("qa_report", qa_path))
        artifacts = [
            _artifact_row(root, artifact_id, path)
            for artifact_id, path in artifact_paths
        ]
        _report_progress(report_progress, 100, "completed", "备课候选产物已生成。")
        return {
            "schema_version": RENDERER_SCHEMA_VERSION,
            "output_kind": normalized_kind,
            "output_dir": str(root),
            "artifacts": artifacts,
            "slide_count": len(layouts),
            "rendered_slides": str(root / "rendered_slides") if wants_ppt else None,
            "quality": {
                "qa_status": qa["qa_status"],
                "status": qa["qa_status"],
                "machine_checks_passed": qa["machine_checks_passed"],
                "machine_layout_check_only": True,
                "human_visual_review_performed": False,
                "rendered_slides": qa["rendered_slides"],
                "candidate_only": True,
                "teacher_review_required": True,
                "publication_allowed": False,
                "official_claim_allowed": False,
            },
        }


__all__ = [
    "DECK_SCHEMA_VERSION",
    "QA_SCHEMA_VERSION",
    "RENDERER_SCHEMA_VERSION",
    "SLIDE_HEIGHT_PX",
    "SLIDE_WIDTH_PX",
    "NativePreparationRenderCancelled",
    "NativePreparationRenderError",
    "NativePreparationRenderer",
]
