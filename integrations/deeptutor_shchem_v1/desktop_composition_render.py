"""Measured local composition. Canvas, PDF and PPTX share positioned elements.

Text lines and pictures are indivisible at page boundaries. Nothing is shortened
to fit a frame. A too-small or overlapping frame is an editable draft, not an
exportable artifact. All pictures are supplied as verified local preparation
bytes; text is plain text, never HTML or a resource URL.
"""
from __future__ import annotations
from copy import deepcopy
from io import BytesIO
from uuid import uuid4
import re

from PySide6.QtCore import QPointF, QRectF, Qt, QMarginsF, QSizeF
from PySide6.QtGui import QColor, QFont, QFontMetricsF, QImage, QPainter, QPageSize, QPageLayout, QPdfWriter, QTextLayout, QTextOption, QGuiApplication
from .desktop_lesson_composition import FORMATS, digest, resource_frames, validate_composition, visible_blocks
from .desktop_preparation_images import verify_image_bytes
from .reader_cancellation import check_read_cancelled
from .desktop_question_answer_reveal import active_question_page

PX_MM = 96 / 25.4
PADDING = 3.0
MAX_ELEMENTS = 30000


def _lines(text, width, *, bold=False, block_id=None, font_px=None):
    if QGuiApplication.instance() is None:raise ValueError('页面排版需要桌面Qt运行环境。')
    font = QFont(QGuiApplication.font()); font.setPixelSize(font_px or (17 if bold else 15)); font.setBold(bold)
    result = []
    for paragraph_index, paragraph in enumerate(text.split('\n')):
        layout = QTextLayout(paragraph, font)
        option = QTextOption(); option.setWrapMode(QTextOption.WrapMode.WrapAtWordBoundaryOrAnywhere)
        layout.setTextOption(option); layout.beginLayout()
        while True:
            line = layout.createLine()
            if not line.isValid(): break
            line.setLineWidth(width * PX_MM)
            line.setPosition(QPointF(0, 0))
            # Keep layout alive along with the native line used by QPainter.
            result.append({'kind': 'text', 'text': paragraph[line.textStart():line.textStart()+line.textLength()],
                'height': (line.height() + 2) / PX_MM, 'width': width,
                'bold': bold, 'font_px': font.pixelSize(), 'font_family':font.family(), 'block_id': block_id,
                '_layout': layout, '_line': line, '_paragraph': (block_id, paragraph_index)})
        layout.endLayout()
        if any(index == 0 for run in layout.glyphRuns() for index in run.glyphIndexes()):
            raise ValueError('存在无法显示的文字字形，请载入支持的本地字体；原文字保留。')
        if not paragraph:
            result.append({'kind': 'space', 'height': 4 / PX_MM, 'width': width, 'block_id': block_id})
        result.append({'kind': 'space', 'height': 1.0, 'width': width, 'block_id': block_id})
    return result


def _picture_height(chunk):
    if chunk['kind'] == 'image': return chunk['height']
    return max((p['height'] for p in chunk.get('_parts', []) if p['kind'] == 'image'), default=0.0)


def _keep_groups(chunks, capacity, *, numbered_headings=False, example_introductions=False, body_px=15):
    """Soft typographic groups bounded by one complete usable page.

    These derived links do not infer question dependencies or edit source blocks.
    Larger paragraphs continue in whole lines; a group is never made too tall.
    """
    totals = [0.0]
    for chunk in chunks:
        totals.append(totals[-1] + chunk['height'])
        chunk['_keep_next'] = False

    def join(start, end):
        left, right = start, end
        while left and chunks[left-1]['_keep_next']: left -= 1
        while right < len(chunks) and chunks[right-1]['_keep_next']: right += 1
        if totals[right]-totals[left] > capacity + .001: return False
        for index in range(left, right-1): chunks[index]['_keep_next'] = True
        return True

    visible = [i for i, c in enumerate(chunks) if c['kind'] != 'space']
    body = [i for i in visible if chunks[i].get('block_id') is not None]
    if body:
        first = body[0]; opening_end = first+1
        if len(body) > 1 and chunks[first].get('_paragraph') == chunks[body[1]].get('_paragraph'):
            opening_end = body[1]+1
        if not join(0, opening_end): join(0, first+1)

    paragraphs = {}
    for index in body:
        key = chunks[index].get('_paragraph')
        if key is not None: paragraphs.setdefault(key, []).append(index)
    for indices in paragraphs.values():
        if len(indices) >= 2:
            join(indices[0], indices[1]+1)
            join(indices[-2], indices[-1]+1)

    if numbered_headings:
        # Only a complete short paragraph in imported knowledge can be a
        # numbered heading without a colon. Keep its next complete opening
        # line group with it; full questions and options keep their old rules.
        for indices in paragraphs.values():
            parts = [chunks[i] for i in indices]
            if any(c['kind'] not in {'text', 'rich_line'} or
                   (c['kind'] == 'rich_line' and any(p['kind'] != 'text' for p in c['_parts']))
                   for c in parts):
                continue
            text = ''.join(c['text'] if c['kind'] == 'text' else ''.join(p['text'] for p in c['_parts'])
                           for c in parts).strip()
            if not text or len(text) > 40 or not re.match(r'^(?:\d+[．.、]|[一二三四五六七八九十]+[、．.])', text):
                continue
            following = next((i for i in body if i > indices[-1]), None)
            if following is not None:
                join(indices[0], following+1)

    for position,index in enumerate(visible[:-1]):
        candidate=chunks[index]
        text=candidate.get('text','').strip()
        if candidate['kind']=='rich_line' and all(p['kind']=='text' for p in candidate['_parts']):
            text=''.join(p['text'] for p in candidate['_parts']).strip()
        if (len(text)<=40 and re.match(r'^(?:\d+[．.、]|[一二三四五六七八九十]+[、．.]|知识点\d)',text)
                and (text.endswith(('：',':')) or text.startswith('知识点'))):
            join(index,visible[position+1]+1)

    # Explicit example introductions in projection knowledge belong to the
    # immediately following standalone picture. Resolve this narrow direction
    # before the ordinary two-sided caption pass, so the preceding picture
    # cannot claim it first. Oversized pairs keep the existing whole-line flow.
    following_examples = {}
    if example_introductions:
        def empty_text(chunk):
            return ((chunk['kind'] == 'text' and not chunk['text'].strip())
                or (chunk['kind'] == 'rich_line' and all(p['kind'] == 'text' and not p['text'].strip()
                    for p in chunk['_parts'])))
        for key, indices in paragraphs.items():
            parts = [chunks[i] for i in indices]
            if (not 1 <= len(parts) <= 3 or any(c['kind'] not in {'text', 'rich_line'}
                    or (c['kind'] == 'rich_line' and any(p['kind'] != 'text' for p in c['_parts']))
                    for c in parts)):
                continue
            text = ''.join(c['text'] if c['kind'] == 'text' else
                ''.join(p['text'] for p in c['_parts']) for c in parts).strip()
            if text not in {'例如：', '例如:', '如：', '如:', '举例：', '举例:'}:
                continue
            following = next((i for i in visible if i > indices[-1] and not empty_text(chunks[i])), None)
            if following is None:
                continue
            picture = chunks[following]; pictures = picture.get('_parts', [])
            standalone = (picture['kind'] == 'image' or (picture['kind'] == 'rich_line'
                and len(pictures) == 1 and pictures[0]['kind'] == 'image'
                and pictures[0]['height'] > 2*body_px/PX_MM + .001
                and paragraphs.get(picture.get('_paragraph')) == [following]))
            if (standalone and totals[following+1]-totals[indices[0]] <= capacity + .001
                    and join(indices[0], following+1)):
                following_examples[key] = following

    # A short adjacent paragraph is a useful picture intro/caption. Keep it
    # together only when the complete combined group fits; never shrink text.
    for position, index in enumerate(visible):
        if not _picture_height(chunks[index]): continue
        for neighbor_position in (position-1, position+1):
            if not 0 <= neighbor_position < len(visible): continue
            neighbor = visible[neighbor_position]
            key = chunks[neighbor].get('_paragraph')
            indices = paragraphs.get(key, [])
            if key in following_examples and following_examples[key] != index:
                continue
            if 1 <= len(indices) <= 3:
                join(min(index, indices[0]), max(index, indices[-1])+1)
    return chunks


def _fit_end(chunks, start, capacity, *, strict=False):
    end = start; used = 0.0
    while end < len(chunks) and used + chunks[end]['height'] <= capacity + .001:
        used += chunks[end]['height']; end += 1
    if start < end < len(chunks) and chunks[end-1].get('_keep_next'):
        boundary = end-1
        while boundary > start and chunks[boundary-1].get('_keep_next'): boundary -= 1
        # Automatic flow advances to a fresh page. A deliberately authored
        # smaller frame may split a soft group, while still keeping whole lines.
        if boundary > start or strict: end = boundary
    return end


def _continuation_heading(item, width, format):
    """Derived projection navigation; never part of the saved source blocks."""
    if format != '16:9':
        return []
    if QGuiApplication.instance() is None:raise ValueError('页面排版需要桌面Qt运行环境。')
    inner = width-2*PADDING
    font = QFont(QGuiApplication.font()); font.setPixelSize(32); font.setBold(True)
    metrics = QFontMetricsF(font)
    suffix = '（续）'
    title_width = inner*PX_MM-metrics.horizontalAdvance(suffix)
    if title_width < metrics.horizontalAdvance('…'):
        raise ValueError('投影续页导航没有足够的一行空间，请加宽外框；完整标题保留。')
    # Only this derived navigation is elided. The exact work/source titles and
    # every source block still use their complete saved text.
    title = metrics.elidedText(' '.join(item['title'].split()), Qt.TextElideMode.ElideRight, title_width)
    chunks = _lines(title + suffix, inner, bold=True, font_px=32)
    chunks.append({'kind': 'space', 'height': 1.5, 'width': width-2*PADDING, 'block_id': None})
    return [{**chunk, 'navigation_heading': True} for chunk in chunks]


def _table_remaining_height(chunks, first_capacity, flow_capacity):
    """Measure preceding complete groups using the same page-fit rules."""
    prefix = _keep_groups([dict(c) for c in chunks], flow_capacity)
    cursor = 0; available = max(0.0, first_capacity)
    while cursor < len(prefix):
        end = _fit_end(prefix, cursor, available, strict=True)
        if end == cursor:
            if available == flow_capacity:
                return 0.0  # The ordinary paginator reports an oversized group.
            available = flow_capacity; continue
        available -= sum(c['height'] for c in prefix[cursor:end]); cursor = end
        if cursor < len(prefix): available = flow_capacity
    return max(0.0, available)


def _table_introduction_start(chunks, body_px):
    """Find the complete short body block immediately before a sealed table.

    Earlier source text/pictures do not affect this local boundary. Only small
    inline pictures may accompany the text; standalone pictures, other tables
    and longer blocks retain ordinary complete-line pagination.
    """
    if not chunks or chunks[-1].get('block_id') is None:
        return None
    block_id = chunks[-1]['block_id']; start = len(chunks) - 1
    while start and chunks[start - 1].get('block_id') == block_id:
        start -= 1
    body = [c for c in chunks[start:] if c['kind'] != 'space']
    if (not 1 <= len(body) <= 3
            or any(c['kind'] not in {'text', 'rich_line'}
                or _picture_height(c) > 2 * body_px / PX_MM + .001 for c in body)):
        return None
    text = ''.join(c.get('text', '') if c['kind'] == 'text' else
        ''.join(p.get('text', '') for p in c['_parts'] if p['kind'] == 'text') for c in body)
    if not text.strip():
        return None
    # Preserve the existing complete opening group when the entire preceding
    # text is short, including its resource title. This never pulls an earlier
    # picture/table or a long preceding explanation into the local group.
    prefix_body = [c for c in chunks if c.get('block_id') is not None and c['kind'] != 'space']
    if len(prefix_body) <= 3 and all(c['kind'] in {'text', 'rich_line'} and not _picture_height(c) for c in prefix_body):
        return 0
    return start


def _question_header(item, width, assets, image_data):
    """Measure one explicit work-copy answer against an actual short title box."""
    inner = width - 2*PADDING
    ids = item['question_answer_reveal']['answer_block_ids']
    if not ids:
        return inner, []
    answer_width, gap = 67.0, 10.0
    title_width = inner - answer_width - gap
    if title_width <= 5:
        raise ValueError('题目标题与原答案没有足够的并排空间，请加宽外框或仅保整题同页。')
    title = _lines(item['title'], title_width, bold=True, font_px=38)
    title_lines = [c for c in title if c['kind'] != 'space']
    if (len(title_lines) != 1
            or title_lines[0]['_line'].naturalTextWidth()/PX_MM > title_width+.001):
        raise ValueError('题目标题过长，不能与原答案并排；请编辑作品标题或仅保整题同页。')
    by_id = {b['id']: b for b in item['blocks']}
    originals = {b['id']: b for b in item['original']['blocks']}
    row = by_id[ids[0]]
    if 'presentation' in row:
        from .desktop_composition_rich_render import rich_lines
        answer = rich_lines(row, by_id, answer_width, assets, image_data, 14,
            originals[row['id']]['value'], font_px=32)
    else:
        answer = _lines(row['value'], answer_width, block_id=row['id'], font_px=32)
    answer = [c for c in answer if c['kind'] != 'space']
    if (len(answer) != 1 or answer[0]['kind'] not in {'text', 'rich_line'}
            or answer[0]['height'] > title_lines[0]['height']+.001):
        raise ValueError('所选原答案在24pt字号下不能完整放入标题旁的一行，请仅保整题同页或另选短答案。')
    # Native rich lines can keep trailing tabs on one Qt line while their
    # editable run boxes extend past the slot. Measure the actual child boxes
    # as well as the line; do not let Office silently place them outside it.
    if (answer[0]['_line'].naturalTextWidth()/PX_MM > answer_width+.001
            or any(part['offset_x'] < -.001
                or part['offset_x']+part['width'] > answer_width+.001
                for part in answer[0].get('_parts', []))):
        raise ValueError('所选原答案超出标题旁空间，不能截断或缩小显示。')
    for chunk in answer:
        chunk.update(role='answer', source_index=row['source_index'],
            original_sha256=item['original_sha256'], answer_block_ids=list(ids),
            reveal_id=item['id']+'-answer-'+row['id'], answer_text_scope='working_copy',
            offset_x=PADDING+inner-answer_width, offset_y=4.0)
    return title_width, answer


def _measure_question_page(item, frame, assets, image_data):
    """Fail closed before any audience can split an explicitly complete page."""
    chunks = _content(item, frame['width'], assets, image_data, student=True,
        image_height=150, format='16:9', keep_short_question=True,
        figure_frame_height=frame['height'], first_flow_height=frame['height']-2*PADDING)
    if not chunks or _fit_end(chunks, 0, frame['height']-2*PADDING) != len(chunks):
        raise ValueError('完整学生题面放不进这一页；请扩大外框或清除整题同页配置，正文与原图不会缩小或删去。')
    _, answer = _question_header(item, frame['width'], assets, image_data)
    return chunks, answer


def _content(item, width, assets, image_data, *, student, image_height=155, format='A4', keep_short_question=False, figure_frame_height=None, first_flow_height=None):
    if width <= 2 * PADDING + 5: raise ValueError('资源太窄，请加宽外框。')
    blocks = visible_blocks(item, student=student)
    if student and not blocks:
        return []  # A teacher-only resource must not disclose its title either.
    inner = width - 2 * PADDING
    # Projection uses 24pt body text. Measure with the same font used by the
    # canvas/PDF/editable PPT; large content continues instead of shrinking.
    body_px, title_px = (32, 38) if format == '16:9' else (15, 17)
    title_width = _question_header(item, width, assets, image_data)[0] if active_question_page(item, format) else inner
    chunks = _lines(item['title'], title_width, bold=True, font_px=title_px)
    chunks.append({'kind': 'space', 'height': 1.5, 'width': inner, 'block_id': None})
    if 'exposition' in item and format == '16:9':
        if item['exposition']['kind']=='notes':
            from .desktop_composition_exposition import notes_chunks
            return chunks+notes_chunks(blocks,inner,_lines,body_px=body_px)
        from .desktop_composition_exposition import diagram_chunks
        return chunks + diagram_chunks(item, inner, _lines)
    capacity = FORMATS[format][1] - 24 - 2*PADDING
    continuation_reserve = sum(c['height'] for c in _continuation_heading(item, width, format))
    flow_capacity = capacity-continuation_reserve
    if flow_capacity <= body_px/PX_MM:
        raise ValueError('投影续页标题超出可用画幅，请调整作品标题或画幅；完整内容保留。')
    # A source crop has one whole picture followed by the teacher's purpose.
    # Reserve a short purpose at its readable font before fitting that picture,
    # so the picture does not consume a full page and orphan the purpose.
    # This affects physical placement only, never the crop or original blocks.
    short_figure = False
    figure_height = image_height
    if (item['original']['kind'] in {'textbook_image', 'preparation_image', 'handout_image'} and blocks
            and blocks[0]['kind'] == 'image'
            and all(b['kind'] == 'text' and 'presentation' not in b for b in blocks[1:])):
        tail = [c for b in blocks[1:] for c in _lines(b['value'], inner, font_px=body_px)]
        reserved = sum(c['height'] for c in chunks + tail) + 2.0
        if sum(c['kind'] == 'text' for c in tail) <= 6 and reserved <= capacity / 2:
            short_figure = True
            figure_capacity = capacity if figure_frame_height is None else min(capacity, figure_frame_height - 2*PADDING)
            figure_height = min(image_height, max(0.0, figure_capacity - reserved))
    from .desktop_composition_inline import bound_image_ids
    from .desktop_composition_rich_render import rich_lines
    inline_images=bound_image_ids(item);by_id={b['id']:b for b in item['blocks']}
    figures = [b for b in blocks if b['kind'] == 'image']
    center_figure = (format == '16:9'
        and item['original']['kind'] in {'textbook_image', 'preparation_image', 'handout_image'}
        and len(figures) == 1 and figures[0]['id'] not in inline_images)
    image_limits = item.get('image_display', {}).get('limits', {})
    originals={b['id']:b['value'] for b in item['original']['blocks'] if b['kind']=='text'}
    table_introductions = []
    for block in blocks:
        check_read_cancelled()
        if block['id'] in inline_images:continue
        if item['original']['kind']=='handout_knowledge' and block['kind']=='text':
            from .desktop_handout_knowledge_tables import table_chunks
            from .desktop_handout_table_projection import source_table_projection
            projection=source_table_projection(item,block)
            if projection is not None:
                table=projection['rows']
                rich_cells=dict(cell_tokens=projection['cell_tokens'], assets=assets, image_data=image_data, image_limits=image_limits)
                grid_columns=table[0]['grid_before']+table[0]['grid_after']+sum(c['column_span'] for c in table[0]['cells'])
                column_weights=(1,2,2) if format=='16:9' and grid_columns==3 else None
                introduction_budget = None
                first_height = (_table_remaining_height(chunks,
                    min(capacity, capacity if first_flow_height is None else first_flow_height), flow_capacity)
                    if format == '16:9' else None)
                introduction_start = _table_introduction_start(chunks, body_px) if first_height is not None else None
                if introduction_start is not None:
                    introduction_height = sum(c['height'] for c in chunks[introduction_start:])
                    if introduction_height < flow_capacity:
                        # Reserve only the immediately adjacent complete short
                        # source block, even after earlier pictures or tables.
                        introduction_budget = flow_capacity-introduction_height
                        first_height = min(first_height, introduction_budget)
                panels = table_chunks(table,block['id'],inner,_lines,body_px,
                    max_height=flow_capacity,first_height=first_height,column_weights=column_weights,**rich_cells)
                if introduction_budget is not None and panels[0]['_table_restarted_page']:
                    # A complete short introduction can advance with its first
                    # panel. Use its fresh-page remainder if the first whole
                    # table group fits there; longer preceding content uses the
                    # full new-page table budget returned above.
                    candidate = table_chunks(table,block['id'],inner,_lines,body_px,
                        max_height=flow_capacity,first_height=introduction_budget,column_weights=column_weights,**rich_cells)
                    if not candidate[0]['_table_restarted_page']: panels = candidate
                if (introduction_budget is not None
                        and panels[0]['height'] <= introduction_budget + .001):
                    table_introductions.append((introduction_start, len(chunks)))
                chunks.extend(panels)
                inline_images.update(projection['image_block_ids'])
                continue
        body = [c for c in chunks if c.get('block_id') is not None and c['kind'] != 'space']
        first_body = not body
        # A short opening stem also belongs beside the first whole picture.
        # Small inline formula images may be part of that stem. Larger pictures
        # and longer introductions continue normally; text is never reduced.
        # Native Word millimetres can round a two-font-height badge just above
        # this boundary; use the same small geometry tolerance as page fitting.
        short_intro = len(body) <= 3 and all(_picture_height(c) <= 2*body_px/PX_MM + .001 for c in body)
        reserved = sum(c['height'] for c in chunks) if short_intro else 0.0
        # Inline Qt lines need baseline/descender space beside the whole image.
        picture_height = min(image_height, capacity-reserved
                             -(body_px/PX_MM if block['kind']=='text' else 2.0),
                             flow_capacity-(body_px/PX_MM if block['kind']=='text' else 2.0))
        if short_figure and block['kind'] == 'image':
            picture_height = min(picture_height, figure_height)
        if picture_height <= 0:
            if first_body:
                raise ValueError('标题超出可用画幅，请调整标题或画幅；完整内容保留。')
            # An unusually tall opening cannot fit with the picture. Keep the
            # picture whole on a continuation instead of rejecting its text.
            picture_height = min(image_height, flow_capacity-(body_px/PX_MM if block['kind']=='text' else 2.0))
        if block['kind'] == 'text':
            if 'presentation' in block:
                chunks.extend(rich_lines(block,by_id,inner,assets,image_data,picture_height,originals[block['id']],font_px=body_px,
                                         opening_line_budget=3-len(body) if short_intro else 0,image_limits=image_limits))
            else:chunks.extend(_lines(block['value'], inner, block_id=block['id'], font_px=body_px))
        else:
            identity = block['value']
            if identity not in assets or identity not in image_data:
                raise ValueError('资源引用的本地图片缺失，请重新带入完整素材。')
            data = verify_image_bytes(assets[identity], image_data[identity])
            image = QImage.fromData(data)
            if image.isNull(): raise ValueError('本地图片无法显示，未省略原图。')
            # Fit whole pictures without altering their aspect ratio or bytes.
            w = min(inner, image.width()/PX_MM*(body_px/15), picture_height * image.width() / image.height())
            h = w * image.height() / image.width()
            from .desktop_composition_image_display import cap_image_size
            w, h = cap_image_size(w, h, image_limits.get(block['id']))
            chunks.append({'kind': 'image', 'asset_id': identity, 'width': w, 'height': h,
                           'block_id': block['id'], 'image_block_id': block['id'], '_image': image,
                           **({'center_in_frame': True} if center_figure else {})})
            chunks.append({'kind': 'space', 'height': 2.0, 'width': inner, 'block_id': block['id']})
        if len(chunks) > MAX_ELEMENTS: raise ValueError('资源排版超出容量，请按课时拆分；未截断文字。')
    chunks = _keep_groups(chunks, flow_capacity, numbered_headings=item['original']['kind']=='handout_knowledge',
        example_introductions=format=='16:9' and item['original']['kind']=='handout_knowledge', body_px=body_px)
    if item['original']['kind']=='handout_knowledge' and format=='16:9':
        # A source-bound table takes precedence over a preceding soft picture
        # caption link when that older link would make the local group too tall.
        # The complete short block and first header/data-row group advance
        # together; earlier source blocks retain their own complete content.
        for start, index in table_introductions:
            left = start
            while left and chunks[left-1]['_keep_next']: left -= 1
            if start and sum(c['height'] for c in chunks[left:index+1]) > flow_capacity + .001:
                chunks[start-1]['_keep_next'] = False
            for c in chunks[start:index]: c['_keep_next'] = True
    # Short cards and activities retain their title and all conditions together.
    # Longer resources still flow in complete lines at the same readable font.
    if (short_figure or item['original']['kind'] in {'lesson_activity', 'knowledge_note', 'handout_knowledge', 'lesson_framework'}
            or (keep_short_question and item['original']['kind']=='complete_question')) and sum(c['height'] for c in chunks) <= capacity + .001:
        for chunk in chunks[:-1]: chunk['_keep_next'] = True
    return chunks


def _overlaps(a, b):
    return (min(a['x']+a['width'], b['x']+b['width']) - max(a['x'], b['x']) > .05
        and min(a['y']+a['height'], b['y']+b['height']) - max(a['y'], b['y']) > .05)


def _layout_composition(value, assets, image_data, *, student, include_reveals=False):
    pages = {key: [] for key in value['pages']}; occupied = {key: [] for key in value['pages']}
    errors = []; fragments = []
    for item in value['resources']:
        if student and not include_reveals and item['id'] in value.get('click_reveals', []):
            continue
        frames = resource_frames(value['placements'][item['id']])
        # A continuation uses the same width: line breaking remains stable.
        if any(abs(f['width']-frames[0]['width']) > .01 for f in frames):
            errors.append(item['title'] + '：续页宽度须与首页相同。'); continue
        answer_chunks = []
        try:
            if active_question_page(item, value['format']):
                _, answer_chunks = _measure_question_page(item, frames[0], assets, image_data)
            chunks = _content(item, frames[0]['width'], assets, image_data, student=student,
                image_height=min(155,FORMATS[value['format']][1]-24-2*PADDING),format=value['format'],
                keep_short_question=student and 'student_arrangement' in value,
                figure_frame_height=frames[0]['height'],first_flow_height=frames[0]['height']-2*PADDING)
        except ValueError as exc:
            errors.append(item['title'] + '：' + str(exc)); continue
        cursor = 0
        continuation_heading = _continuation_heading(item, frames[0]['width'], value['format'])
        for frame_index, frame in enumerate(frames):
            if cursor == len(chunks): break
            selected = []; top = PADDING
            heading = continuation_heading if cursor and _has_visible_content(chunks[cursor:]) else []
            reserve = sum(c['height'] for c in heading)
            available = frame['height']-2*PADDING-reserve
            end = _fit_end(chunks, cursor, available)
            if student and end > cursor and _fit_end(chunks, cursor, available, strict=True) == cursor:
                # Audience filtering can leave the first teacher frame too short
                # for the complete student task. Use a later existing frame when
                # it holds this group, without changing the authored placements.
                if any(_fit_end(chunks, cursor, later['height']-2*PADDING-reserve, strict=True) > cursor
                       for later in frames[frame_index+1:]):
                    continue
            # Do not turn a hidden or empty fragment into a visible heading.
            if _has_visible_content(chunks[cursor:end]):
                for chunk in heading:
                    selected.append({**chunk, 'x': frame['x']+PADDING, 'y': frame['y']+top,
                                     'resource_id': item['id'], 'page_id': frame['page_id']})
                    top += chunk['height']
            while cursor < end:
                chunk = chunks[cursor]
                selected.append({**chunk, 'x': frame['x']+PADDING +
                                 ((frame['width']-2*PADDING-chunk['width'])/2 if chunk.get('center_in_frame') else 0),
                                 'y': frame['y']+top,
                                 'resource_id': item['id'], 'page_id': frame['page_id']})
                top += chunk['height']; cursor += 1
            if frame_index == 0 and student and include_reveals and selected:
                selected.extend({**chunk, 'x': frame['x']+chunk['offset_x'], 'y': frame['y']+chunk['offset_y'],
                    'resource_id': item['id'], 'page_id': frame['page_id']} for chunk in answer_chunks)
            if not selected:
                # Hidden teacher blocks can leave a short continuation frame
                # before a whole student picture. Use the later authored frame
                # when it can hold this chunk; the cursor never skips content.
                if student and any(chunks[cursor]['height'] <= later['height']-2*PADDING-reserve+.001
                                   for later in frames[frame_index+1:]):
                    continue
                errors.append(item['title'] + '：外框放不下完整文字行或图片，请扩大尺寸。')
            occupied[frame['page_id']].append((item['id'], frame))
            pages[frame['page_id']].extend(selected)
            fragments.append({'resource_id': item['id'], 'frame_index': frame_index,
                              'frame': deepcopy(frame), 'elements': selected, 'used_height': top+PADDING})
        if cursor < len(chunks):
            errors.append(item['title'] + '：完整内容尚未放完，请“按顺序排满并续页”或增加外框；导出不会截断。')
    for page_id, frames in occupied.items():
        for i, (identity, frame) in enumerate(frames):
            for other_id, other in frames[i+1:]:
                if identity != other_id and _overlaps(frame, other):
                    errors.append(f'第{value["pages"].index(page_id)+1}页资源外框重叠，请调整位置。')
    if sum(len(p) for p in pages.values()) > MAX_ELEMENTS:
        errors.append('本作品排版超出容量，请按课时拆分；未截断文字。')
    return {'composition': value, 'student': student, 'include_reveals': include_reveals, 'pages': pages,
            'fragments': fragments, 'errors': list(dict.fromkeys(errors))}


def _has_visible_content(elements):
    return any(e['kind'] in {'image', 'rich_line', 'diagram'}
               or (e['kind'] == 'text' and e['text'].strip()) for e in elements)


def layout_composition(composition, assets, image_data, *, student=True, include_reveals=False):
    """Render the saved audience choice without mutating source or teacher frames.

    Shared placement keeps intentional blank pages and existing page IDs. An
    explicitly selected student flow measures visible content independently.
    """
    value = validate_composition(composition)
    assets = {a['asset_id']: a for a in assets}
    if student and 'student_arrangement' in value:
        projected = _arrange_ordered(value, list(assets.values()), image_data,
                                     columns=value['student_arrangement']['columns'], student=True)
        result = _layout_composition(projected, assets, image_data, student=True, include_reveals=include_reveals)
        result['composition'] = value
        result['independent_student_arrangement'] = True
        result['omitted_teacher_page_ids'] = []
        if not any(_has_visible_content(p) for p in result['pages'].values()):
            result['pages'] = {}; result['fragments'] = []
            result['errors'].append('学生版没有可见内容，请加入学生任务或材料后再生成。')
        return result
    result = _layout_composition(value, assets, image_data, student=student, include_reveals=include_reveals)
    omitted = []
    empty = [p for p, elements in result['pages'].items() if not _has_visible_content(elements)]
    if student and empty:
        # Only remeasure the teacher scope when there are empty student pages.
        # This distinguishes hidden answers from deliberately authored blanks.
        teacher = _layout_composition(value, assets, image_data, student=False)
        omitted = [p for p in empty if _has_visible_content(teacher['pages'][p])]
        result['pages'] = {p: elements for p, elements in result['pages'].items() if p not in omitted}
        result['fragments'] = [f for f in result['fragments'] if f['frame']['page_id'] not in omitted]
        if not result['pages']:
            result['errors'].append('学生版没有可见内容，请加入学生任务或材料后再生成。')
    result['omitted_teacher_page_ids'] = omitted
    result['independent_student_arrangement'] = False
    return result


def arrange_composition(composition, assets, image_data, *, columns=1):
    """Explicit teacher action: replace placements with measured ordered frames."""
    return _arrange_ordered(composition, assets, image_data, columns=columns, student=False)


def _arrange_ordered(composition, assets, image_data, *, columns, student):
    value = validate_composition(composition)
    if type(columns) is not int or columns not in (1, 2): raise ValueError('请选择一栏或两栏。')
    if student:
        value['resources'] = [r for r in value['resources'] if visible_blocks(r, student=True)]
    # Read-only student previews and outputs need repeatable page identities.
    # Teacher-authored frames and page IDs are retained in the saved composition.
    seed = digest({'format': value['format'], 'columns': columns,
                   'resources': [{'id': r['id'], 'title': r['title'], 'blocks': visible_blocks(r, student=True),
                       **({'question_answer_reveal': r['question_answer_reveal']} if active_question_page(r, value['format']) else {})}
                                 for r in value['resources']]}) if student else None
    def page_id(index):
        return 'CP-' + (digest({'student_flow': seed, 'page': index})[:32] if student else uuid4().hex)
    page_w, page_h = FORMATS[value['format']]; margin, gap = 12.0, 8.0
    width = (page_w - 2*margin - gap*(columns-1)) / columns
    height = page_h - 2*margin
    assets_by_id = {a['asset_id']: a for a in assets}
    pages = [page_id(0)]; placements = {}; column = 0; y = margin
    def advance():
        nonlocal column, y
        column += 1; y = margin
        if column == columns:
            column = 0; pages.append(page_id(len(pages)))
            if len(pages) > 160: raise ValueError('完整作品超过160页，请按课时拆分。')
    for item in value['resources']:
        if active_question_page(item, value['format']):
            # Only an explicitly configured full question owns a fresh full
            # page, including when surrounding material uses two columns.
            frame = deepcopy(item['question_answer_reveal']['classroom_frame'])
            _measure_question_page(item, frame, assets_by_id, image_data)
            if column or y != margin:
                column = columns-1; advance()
            full_width = frame['width']
            chunks = _content(item, full_width, assets_by_id, image_data, student=student,
                image_height=150, format=value['format'], keep_short_question=student,
                figure_frame_height=frame['height'], first_flow_height=frame['height']-2*PADDING)
            heading_height = sum(c['height'] for c in _continuation_heading(item, full_width, value['format']))
            end = _fit_end(chunks, 0, frame['height']-2*PADDING)
            if not end:
                raise ValueError(item['title'] + '：整题首框放不下完整文字或图片。')
            frames = [{**frame, 'page_id': pages[-1]}]
            while end < len(chunks):
                column = columns-1; advance()
                start = end
                reserve = heading_height if _has_visible_content(chunks[start:]) else 0.0
                end = _fit_end(chunks, start, height-2*PADDING-reserve)
                if end == start:
                    raise ValueError(item['title'] + '：教师解析续页放不下完整文字或图片。')
                frames.append({'page_id': pages[-1], 'x': frame['x'], 'y': margin,
                    'width': full_width, 'height': height})
            placements[item['id']] = {**frames[0], **({'continuations': frames[1:]} if len(frames)>1 else {})}
            # The next ordinary resource starts on the next page; do not add
            # an unused trailing page when this was the final resource.
            column = columns-1; y = page_h-margin
            continue
        chunks = _content(item, width, assets_by_id, image_data, student=student,image_height=min(155,height-2*PADDING),format=value['format'],
                          keep_short_question=student,first_flow_height=page_h-margin-y-2*PADDING)
        has_table = any('_table_source_rows' in c for c in chunks)
        heading_height = sum(c['height'] for c in _continuation_heading(item, width, value['format']))
        if any(chunk['height'] + 2*PADDING > height for chunk in chunks):
            raise ValueError(item['title'] + '含放不下一整页的图片或文字行，请调整画幅。')
        frames = []; cursor = 0
        while cursor < len(chunks):
            available = page_h - margin - y; start = cursor
            reserve = heading_height if start and _has_visible_content(chunks[start:]) else 0.0
            cursor = _fit_end(chunks, start, available-2*PADDING-reserve, strict=True)
            used = 2*PADDING + reserve + sum(c['height'] for c in chunks[start:cursor])
            if cursor == start:
                if y == margin and available == height:
                    raise ValueError(item['title'] + '含放不下投影标题与完整文字行或图片的内容，请调整画幅。')
                advance()
                if start == 0 and has_table:
                    # No content was placed in the preceding partial page.
                    # Recompute complete table panels for the actual first
                    # page budget before recording any resource frames.
                    chunks = _content(item, width, assets_by_id, image_data, student=student,
                        image_height=min(155,height-2*PADDING),format=value['format'],
                        keep_short_question=student,first_flow_height=height-2*PADDING)
                continue
            # The model keeps one resource per page. Its continuation goes to
            # the next page even in two-column mode; another resource can use
            # remaining column space independently after this resource ends.
            # Native table panel grouping depends on the first page budget.
            # Preserve that budget when this resource continues; shrinking the
            # saved first frame to used content would repartition its cells on
            # reopening and make the later authored frames inconsistent.
            if start == 0 and cursor < len(chunks) and has_table: used = available
            frames.append({'page_id': pages[-1], 'x': margin+column*(width+gap), 'y': y,
                           'width': width, 'height': used})
            y += used + gap
            if cursor < len(chunks):
                column = columns-1; advance()
        placements[item['id']] = {**frames[0], **({'continuations': frames[1:]} if len(frames)>1 else {})}
    value.update(pages=pages, placements=placements)
    return validate_composition(value)


def paint_elements(painter, elements):
    for element in elements:
        if element['kind'] == 'space': continue
        painter.save(); painter.translate(element['x'], element['y'])
        if element['kind'] == 'image':
            painter.drawImage(QRectF(0, 0, element['width'], element['height']), element['_image'])
        elif element['kind'] == 'diagram':
            from .desktop_composition_rich_render import positioned_parts
            paint_elements(painter, positioned_parts({**element,'x':0,'y':0}))
        elif element['kind'] == 'diagram_box':
            from PySide6.QtGui import QPen
            painter.setPen(QPen(QColor(element['stroke']), .5));painter.setBrush(QColor(element['fill']))
            painter.drawRect(QRectF(0, 0, element['width'], element['height']))
        elif element['kind'] == 'diagram_arrow':
            from PySide6.QtGui import QPolygonF
            w,h=element['width'],element['height']
            if element.get('vertical'):painter.translate(w,0);painter.rotate(90);w,h=h,w
            painter.setPen(Qt.PenStyle.NoPen);painter.setBrush(QColor(element['fill']))
            painter.drawPolygon(QPolygonF([QPointF(0,h*.35),QPointF(w*.6,h*.35),QPointF(w*.6,0),
                QPointF(w,h*.5),QPointF(w*.6,h),QPointF(w*.6,h*.65),QPointF(0,h*.65)]))
        elif element['kind']=='rich_line':
            # drawContents with a clip still emits the whole paragraph into a
            # vector PDF for every line. Draw just this measured native line,
            # then its verified inline objects at the same measured positions.
            painter.scale(1/PX_MM,1/PX_MM);painter.setPen(QColor('#17242f'))
            line=element['_line'];line.draw(painter,QPointF(0,-line.y()))
            for part in element['_parts']:
                if part['kind']=='image':painter.drawImage(QRectF(part['offset_x']*PX_MM,part['offset_y']*PX_MM,
                    part['width']*PX_MM,part['height']*PX_MM),part['_image'])
                elif part['kind']=='diagram':painter.drawImage(QRectF(part['offset_x']*PX_MM,part['offset_y']*PX_MM,
                    part['width']*PX_MM,part['height']*PX_MM),part['_image'])
        else:
            painter.scale(1/PX_MM, 1/PX_MM); painter.setPen(QColor('#17242f'))
            element['_line'].draw(painter, QPointF(0, 0))
        painter.restore()


def page_image(layout, page_id, *, scale=4):
    w, h = FORMATS[layout['composition']['format']]
    image = QImage(round(w*scale), round(h*scale), QImage.Format.Format_RGB32); image.fill(Qt.GlobalColor.white)
    painter = QPainter(image); painter.scale(scale, scale)
    paint_elements(painter, layout['pages'][page_id]); painter.end()
    return image


def write_pdf(layout, path):
    if layout['errors']: raise ValueError('\n'.join(layout['errors']))
    w, h = FORMATS[layout['composition']['format']]
    writer = QPdfWriter(str(path)); writer.setResolution(96)
    writer.setPageSize(QPageSize(QSizeF(w, h), QPageSize.Unit.Millimeter))
    writer.setPageMargins(QMarginsF(0,0,0,0), QPageLayout.Unit.Millimeter)
    painter = QPainter(writer)
    if not painter.isActive(): raise ValueError('PDF输出无法打开，已有成品保留。')
    try:
        for index, elements in enumerate(layout['pages'].values()):
            check_read_cancelled()
            if index and not writer.newPage(): raise ValueError('PDF续页无法写入。')
            # Even a completed classroom preview cannot inject its dedicated
            # answer entrances into a static student PDF.
            if layout.get('student'):
                elements = [e for e in elements if not e.get('reveal_id')]
            painter.save(); painter.scale(PX_MM, PX_MM); paint_elements(painter, elements); painter.restore()
    finally: painter.end()


def write_pptx(layout, path, image_data):
    if layout['errors']: raise ValueError('\n'.join(layout['errors']))
    from pptx import Presentation
    from pptx.util import Mm, Pt
    from pptx.dml.color import RGBColor
    presentation = Presentation(); w,h = FORMATS[layout['composition']['format']]
    presentation.slide_width, presentation.slide_height = Mm(w), Mm(h)
    for elements in layout['pages'].values():
        check_read_cancelled()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        reveal_ids = layout['composition'].get('click_reveals', []) if layout['student'] else []
        reveal_shapes = {identity: [] for identity in reveal_ids}
        answer_shapes = {}
        classroom = layout['student'] and layout.get('include_reveals', False) and layout['composition']['format'] == '16:9'
        stepped={r['id']:r for r in layout['composition']['resources'] if r['id'] in reveal_ids
                 and r.get('exposition',{}).get('kind') in {'process','process_vertical'}}
        node_shapes={identity:{b['id']:[] for b in item['blocks']} for identity,item in stepped.items()}
        from .desktop_composition_rich_render import flatten_elements
        for e in flatten_elements(elements):
            if e['kind'] == 'space': continue
            if e.get('reveal_id') and not classroom: continue
            if e['kind'] == 'image':
                shape = slide.shapes.add_picture(BytesIO(image_data[e['asset_id']]), Mm(e['x']), Mm(e['y']),
                                         width=Mm(e['width']), height=Mm(e['height']))
                if 'transform' in e:
                    from .desktop_handout_paint_preview import validate_picture_transform
                    validate_picture_transform(e['transform'])
                    shape._element.spPr.get_or_add_xfrm().set('flipV', '1')
                if 'crop' in e:
                    for key,attribute in (('l','crop_left'),('t','crop_top'),('r','crop_right'),('b','crop_bottom')):
                        setattr(shape,attribute,e['crop'][key]/100000)
                if 'source_transparency' in e:
                    from .word_static_quote_render import add_ppt_transparency
                    add_ppt_transparency(shape, e['source_transparency'])
            elif e['kind'] in {'diagram_box', 'diagram_arrow'}:
                from pptx.enum.shapes import MSO_SHAPE
                kind=MSO_SHAPE.RECTANGLE if e['kind']=='diagram_box' else (MSO_SHAPE.DOWN_ARROW if e.get('vertical') else MSO_SHAPE.RIGHT_ARROW)
                shape=slide.shapes.add_shape(kind,
                    Mm(e['x']),Mm(e['y']),Mm(e['width']),Mm(e['height']))
                shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(e['fill'].lstrip('#'))
                if e['kind']=='diagram_box':
                    shape.line.color.rgb=RGBColor.from_string(e['stroke'].lstrip('#'));shape.line.width=Pt(1.4)
                else:shape.line.fill.background()
            else:
                shape = slide.shapes.add_textbox(Mm(e['x']), Mm(e['y']), Mm(e['width']), Mm(e['height']))
                if e.get('navigation_heading'):
                    shape.name = 'shchem-continuation-heading'
                frame = shape.text_frame; frame.clear(); frame.word_wrap = False
                frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
                paragraph = frame.paragraphs[0]; paragraph.space_before = paragraph.space_after = Pt(0)
                paragraph.line_spacing = 1.0
                run = paragraph.add_run(); run.text = e['text']; run.font.name = e['font_family']
                run.font.size = Pt(e['font_px']*72/96); run.font.bold = e['bold']; run.font.color.rgb = RGBColor(23,36,47)
            if e.get('reveal_id'):
                answer_shapes.setdefault(e['reveal_id'], []).append(shape)
            elif e['resource_id'] in reveal_shapes:
                if e['resource_id'] in stepped:
                    # The guiding question stays visible while complete steps
                    # and their incoming arrows appear on their own clicks.
                    key=e.get('block_id')
                    if key is not None:node_shapes[e['resource_id']][key].append(shape)
                else:reveal_shapes[e['resource_id']].append(shape)
        from .desktop_composition_animations import add_click_appear
        groups=[]
        for item in layout['composition']['resources']:
            identity=item['id']
            if identity in stepped:
                groups.extend((identity+'-'+b['id'],node_shapes[identity][b['id']]) for b in item['blocks'])
            elif identity in reveal_shapes:groups.append((identity,reveal_shapes[identity]))
            for reveal_id, shapes in answer_shapes.items():
                if reveal_id.startswith(identity+'-answer-'):
                    groups.append((reveal_id, shapes))
        add_click_appear(slide,groups)
    presentation.save(path)
