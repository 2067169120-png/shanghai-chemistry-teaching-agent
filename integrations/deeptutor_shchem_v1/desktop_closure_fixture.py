"""Small authored source for closure verification; never installed in user data."""
from pathlib import Path
from docx import Document


def seed_word(facade, directory):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    doc=Document();doc.add_heading('合成化学复练 · 非正式试卷',0)
    doc.add_paragraph('【例1】共同材料：恒温下比较可逆反应A(g)⇌B(g)的正逆反应速率。说明动态平衡的判断依据。')
    doc.add_paragraph('【答案】同一反应正逆反应速率相等且不为零；不是浓度必定相等。')
    doc.add_paragraph('【例2】沿用题内条件：某温度下c(A)=0.20 mol/L，c(B)=0.60 mol/L。写出K的表达式并计算。')
    doc.add_paragraph('【答案】K=c(B)/c(A)=3.0。')
    path=directory/'合成复练来源.docx';doc.save(path)
    facade.save_visual_import_batch(handout_files=(path,),source_type='教师讲义')
    rows=[r for r in facade.word_question_catalog()['items'] if r['source_name']==path.name]
    for row in rows:
        options=facade.word_question_attribute_options(row['key'],row['revision']);attrs=options['attributes']
        facade.word_question_save_attributes(row['key'],row['revision'],
            {'primary_knowledge':{'id':'K09','label':next(k['name'] for k in options['catalog']['knowledge_points'] if k['id']=='K09'),'status':'teacher_confirmed','evidence':[]},
             'teacher_note':'合成软件验收标签，不代表真实教材/原题核定。'},
            expected_attribute_revision=attrs['revision'],expected_stored_revision=options['stored_revision'])
    return path,[r for r in facade.word_question_catalog()['items'] if r['source_name']==path.name]
