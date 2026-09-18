"""Local, versioned output from one teacher-owned lesson design."""
from __future__ import annotations
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
from uuid import uuid4

from .desktop_lesson_design import candidate_from_design, content_fingerprint, coverage, validate_design
from .desktop_state import utc_now

FILES = ("lesson_presentation.pptx", "lesson_plan.docx", "student_worksheet.docx")


def _document(payload, student=False):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin = s.bottom_margin = Mm(18)
    s.left_margin = s.right_margin = Mm(20)
    for key, size in (("Normal", 11), ("Title", 20), ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[key]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = key != "Normal"
    doc.add_heading(payload["topic"] + (" · 学习单" if student else " · 教案"), 0)
    doc.add_paragraph((payload.get("audience", "") + "　" + payload.get("lesson_timing", "")).strip())
    return doc


def write_documents(payload, directory, image_data):
    from docx.shared import Mm
    from .desktop_preparation_images import verify_image_bytes
    plan = validate_design(payload["lesson_design"])
    images = {x["asset_id"]: x for x in payload.get("image_assets", [])}
    teacher, student = _document(payload), _document(payload, True)
    teacher.add_heading("教学目标与安排", 1)
    report = coverage(plan)
    for g in report["objectives"]:
        teacher.add_paragraph(g["text"] + "　[" + g["status"] + "]")
    teacher.add_paragraph(f"已估时合计 {report['known_minutes']} 分钟；"
                          f"{report['unestimated']} 个环节尚未估时。目标状态表示安排关系，不表示学生已掌握。")
    teacher.add_heading("教学过程", 1)
    for i, n in enumerate(plan["nodes"], 1):
        for doc, is_student in ((teacher, False), (student, True)):
            doc.add_heading(f"{i}. {n['title']}", 2 if not is_student else 1)
            if not is_student:
                goals = [g["text"] for g in plan["objectives"] if g["id"] in n["objective_ids"]]
                timing = "尚未估时" if n["minutes"] is None else f"{n['minutes']}分钟"
                doc.add_paragraph("目标：" + "；".join(goals) + "　用时：" + timing)
            if n["material_text"] and (not is_student or n["student_material"]):
                doc.add_paragraph(n["material_text"])
            if not is_student or n["student_material"]:
                for identity in n["image_ids"]:
                    a = images[identity]
                    data = verify_image_bytes(a, image_data[identity])
                    # Preserve proportions and original bytes; fit on one A4 page.
                    width = min(165, 200 * a["width"] / a["height"])
                    doc.add_picture(BytesIO(data), width=Mm(width))
                    doc.add_paragraph(a["caption"] + "　来源：" + a["source"])
            for key, label in (("teacher_action", "教师活动"), ("student_task", "学生任务"),
                               ("expected_output", "预期产出"), ("criteria", "评价依据"),
                               ("teacher_answer", "教师答案"), ("notes", "教师备注")):
                if n[key] and (not is_student or key == "student_task"):
                    doc.add_paragraph(label + "：" + n[key])
    teacher.add_heading("课后反思", 1)
    teacher.add_paragraph("授课后记录实际完成情况、学生证据及下一次调整；不预填未发生的课堂事实。")
    teacher.save(directory / "lesson_plan.docx")
    student.save(directory / "student_worksheet.docx")


def output_root(facade, record=None):
    root = Path(facade.paths.task_root) / "preparation-v1" / "node-exports"
    if record is None:
        return root
    # record has already been validated with the containing design.
    name = record["id"]
    import re
    if not re.fullmatch(r"OUT-[a-f0-9]{32}", name):
        raise ValueError("教学成品标识无效。")
    return root / name


def export_design(facade, payload, *, cancelled=lambda: False):
    from .desktop_preparation_renderer import NativePreparationRenderer
    plan = validate_design(payload["lesson_design"])
    c = candidate_from_design(payload)
    used = {im for n in plan["nodes"] for im in n["image_ids"]}
    assets = {x["asset_id"]: x for x in payload.get("image_assets", [])}
    image_data = {key: facade.preparation_image_bytes(assets[key]) for key in used}
    if cancelled():
        raise ValueError("已取消生成，教学设计保留。")
    root = output_root(facade)
    root.mkdir(parents=True, exist_ok=True)
    ident = "OUT-" + uuid4().hex
    target = root / ident
    with tempfile.TemporaryDirectory(prefix=".design-", dir=root) as staging:
        directory = Path(staging)
        result = NativePreparationRenderer().render(c, output_kind="ppt", output_dir=directory,
                                                    image_data=image_data, is_cancelled=cancelled)
        # Do not present known-overflow slide output as ready.
        deck = json.loads((directory / "deck.json").read_text(encoding="utf-8"))
        if any(e.get("overflow") for s in deck["slides"] for e in s["elements"]):
            raise ValueError("有页面内容超出可读范围，请拆分长环节后重试；原成品保留。")
        write_documents(payload, directory, image_data)
        if cancelled():
            raise ValueError("已取消生成，教学设计保留。")
        record = {"id": ident, "fingerprint": content_fingerprint(payload), "created_at": utc_now(),
                  "files": [{"name": name, "sha256": sha256((directory / name).read_bytes()).hexdigest()}
                            for name in FILES]}
        (directory / "lesson-design.json").write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
        (directory / "output.json").write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
        # Only finished bundles become discoverable. Existing versions are never overwritten.
        directory.rename(target)
    return record


def checked_file(facade, plan, export_id, name):
    plan = validate_design(plan)
    record = next((r for r in plan["exports"] if r["id"] == export_id), None)
    if record is None:
        raise ValueError("未找到这份成品记录。")
    file = next((f for f in record["files"] if f["name"] == name), None)
    if file is None:
        raise ValueError("该文件不属于已记录的教学成品。")
    path = output_root(facade, record) / name
    if not path.is_file() or sha256(path.read_bytes()).hexdigest() != file["sha256"]:
        raise ValueError("原成品缺失或已被外部修改，请重新生成或保留外部修订稿。")
    return path


def actual_ppt_preview(facade, plan, export_id):
    """Convert the *actual* PPTX, separately from the renderer's quick PNGs."""
    from .desktop_local_pagination import find_libreoffice
    source = checked_file(facade, plan, export_id, "lesson_presentation.pptx")
    office = find_libreoffice()
    if office is None:
        raise ValueError("实际PPTX预览需本机LibreOffice。PPTX仍可用PowerPoint打开；未删除或改写成品。")
    target = source.parent / "actual-pptx"
    target.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="shchem-pptx-") as profile:
        result = subprocess.run([str(office), "-env:UserInstallation=" + Path(profile).as_uri(),
                                "--headless", "--convert-to", "pdf:impress_pdf_Export",
                                "--outdir", str(target), str(source)],
                                capture_output=True, timeout=180,
                                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    pdf = target / (source.stem + ".pdf")
    if result.returncode or not pdf.is_file() or not pdf.read_bytes().startswith(b"%PDF-"):
        raise ValueError("实际PPTX转换未完成，请用PowerPoint核对原文件。")
    return {"path": str(pdf), "engine": "LibreOffice Impress",
            "source_sha256": sha256(source.read_bytes()).hexdigest(),
            "pdf_sha256": sha256(pdf.read_bytes()).hexdigest()}
