"""Source-closed image label candidates reviewed by the current root model.

This is an independent admission policy, not a relaxation of the text-only
route. It accepts selected paragraph ranges with original PNGs and validated
WMF previews, plus an explicit model reading record. An OLE equation's visible
preview is accepted only through its native same-object ShapeID relationship;
the binary is hashed but never interpreted or executed. Neither that technical
closure nor a model's reading declaration constitutes human confirmation.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from copy import deepcopy
from pathlib import Path

from docx import Document
from jsonschema import Draft202012Validator
from lxml import etree
from PIL import Image

from .desktop_preparation_sources import _block_text, _body_blocks, _word_images
from .desktop_word_question_attributes import validate_attributes
from .offline_word_label_review import (
    OUTPUT_SCHEMA, PROVENANCE, OfflineWordLabelReviewError,
    OfflineWordLabelReviewService, _SHA256, _TEXT, _digest,
    _has_substantive_answer_text, _object, _require,
)
from .word_handout_import import _validate_container
from .word_native_text import WordNativeTextReader

CANDIDATE_SCHEMA_VERSION = "shchem.offline-word-image-label-candidates.v1"
PLAN_SCHEMA_VERSION = "shchem.offline-word-image-label-plan.v1"
RECEIPT_SCHEMA_VERSION = "shchem.offline-word-image-label-receipt.v1"
SCOPE_SCHEMA_VERSION = "shchem.offline-word-image-reading-scope.v1"
ROOT_REVIEW_SCHEMA_VERSION = "shchem.root-image-label-review.v1"
READING_POLICY_REVISION = "complete-native-question-answer-image-root-model-v1"
RULE_REVISION = "codex-offline-image-reviewed-20260928-v1"
RECHECK_RULE_REVISION = "codex-offline-image-recheck-20260928-v1"
WHITE_RENDERER = "rgba-over-opaque-white-rgb-v1"

_FILE_REF = _object({
    "path": {"type": "string", "minLength": 1, "maxLength": 2048},
    "sha256": _SHA256, "bytes": {"type": "integer", "minimum": 1},
})
CANDIDATE_SCHEMA = _object({
    "schema_version": {"const": CANDIDATE_SCHEMA_VERSION},
    "provenance": {"const": PROVENANCE},
    "candidate_only": {"type": "boolean", "const": True},
    "human_review": {"type": "boolean", "const": False},
    "entries": {"type": "array", "minItems": 1, "maxItems": 100, "items": _object({
        "key": {**_TEXT, "maxLength": 160}, "question_revision": _TEXT,
        "source_sha256": _SHA256, "expected_stored_attribute_revision": _SHA256,
        "decoded": OUTPUT_SCHEMA, "root_review": _FILE_REF,
        "reading_scope_sha256": _SHA256,
        # Exact recursive equality against the reconstructed native scope below
        # validates all nested fields, rejecting omissions and invented fields.
        "reading_scope": _object({
            "schema_version": {"const": SCOPE_SCHEMA_VERSION},
            "identity": {"type": "object"}, "ranges": {"type": "object"},
            "blocks": {"type": "object"}, "objects": {"type": "array"},
            "images": {"type": "array", "minItems": 1},
            "document_xml_sha256": _SHA256,
        }),
    })},
})

_NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
    "v": "urn:schemas-microsoft-com:vml", "o": "urn:schemas-microsoft-com:office:office",
    "w10": "urn:schemas-microsoft-com:office:word",
}


def _q(name):
    prefix, local = name.split(":")
    return "{" + _NS[prefix] + "}" + local


_CONTAINERS = {
    _q("w:drawing"): ("word_drawing_or_shape", "图片或图形"),
    _q("w:pict"): ("vml_drawing_or_shape", "图片或旧式图形"),
    _q("w:object"): ("ole_object_reference", "嵌入对象或旧公式"),
}
_DRAWING_TAGS = {_q(name) for name in (
    "w:drawing", "wp:inline", "wp:extent", "wp:effectExtent", "wp:docPr",
    "wp:cNvGraphicFramePr", "a:graphicFrameLocks", "a:graphic", "a:graphicData",
    "pic:pic", "pic:nvPicPr", "pic:cNvPr", "pic:cNvPicPr", "a:picLocks",
    "pic:blipFill", "a:blip", "a:srcRect", "a:stretch", "a:fillRect",
    "pic:spPr", "a:xfrm", "a:off", "a:ext", "a:prstGeom", "a:avLst", "a:noFill", "a:ln",
)}
_IDENTITY = ("key", "source_sha256", "source_revision", "index_revision", "extraction_revision", "revision")
_RANGES = ("block_start", "question_end", "answer_start", "block_end", "context_start", "context_end", "origin_block_start")
_GROUPS = (("context_blocks", "context_start", "context_end"),
           ("question_blocks", "block_start", "question_end"),
           ("answer_blocks", "answer_start", "block_end"))


def _sha(raw):
    return hashlib.sha256(raw).hexdigest()


def _xml_sha(node):
    return _sha(etree.tostring(node, method="c14n"))


def validate_candidate_batch(value):
    _require(isinstance(value, dict) and not list(Draft202012Validator(CANDIDATE_SCHEMA).iter_errors(value)),
             "invalid_image_candidate_batch", "图文候选须包含完整原生范围、主代理读图记录及自动建议声明。")
    try:
        copied = deepcopy(value)
        _digest(copied)
    except (ValueError, TypeError, RecursionError) as exc:
        raise OfflineWordLabelReviewError("invalid_image_candidate_batch", "图文候选必须是完整JSON数据。") from exc
    keys = [entry["key"] for entry in copied["entries"]]
    _require(len(keys) == len(set(keys)), "duplicate_key", "图文候选不能重复包含同一道题。")
    return copied


def _unique_object(pairs):
    result = {}
    for key, value in pairs:
        _require(key not in result, "duplicate_evidence_key", "复核证据JSON含重复字段。")
        result[key] = value
    return result


def _bound_file(ref, *, max_bytes=8_000_000):
    _require(isinstance(ref, dict) and isinstance(ref.get("path"), str)
             and isinstance(ref.get("sha256"), str) and re.fullmatch(r"[0-9a-f]{64}", ref["sha256"])
             and type(ref.get("bytes")) is int and 0 < ref["bytes"] <= max_bytes,
             "invalid_evidence_file", "复核证据文件的路径、长度或摘要无效。")
    path = Path(ref["path"])
    _require(path.is_absolute() and not str(path).startswith(("\\\\", "//")),
             "invalid_evidence_file", "复核证据必须使用本机绝对路径。")
    try:
        _require(path.is_file() and all(not p.is_symlink() and not getattr(p, "is_junction", lambda: False)()
                                      for p in (path, *path.parents)),
                 "invalid_evidence_file", "复核证据必须是已有真实文件。")
        with path.open("rb") as stream:
            raw = stream.read(max_bytes + 1)
    except OSError as exc:
        raise OfflineWordLabelReviewError("evidence_file_unavailable", "复核证据文件不可读取。") from exc
    _require(len(raw) == ref["bytes"] and _sha(raw) == ref["sha256"],
             "evidence_file_changed", "复核证据文件与已绑定摘要不一致。")
    return raw


def _bound_json(ref):
    try:
        result = json.loads(_bound_file(ref).decode("utf-8-sig"), object_pairs_hook=_unique_object)
        _require(isinstance(result, dict), "invalid_review_evidence", "复核证据必须为JSON对象。")
        return result
    except (UnicodeError, json.JSONDecodeError) as exc:
        raise OfflineWordLabelReviewError("invalid_review_evidence", "复核证据不是完整JSON。") from exc


def _relation(document, rid, *, image):
    relation = document.part.rels.get(rid)
    expected = _NS["r"] + ("/image" if image else "/oleObject")
    _require(isinstance(rid, str) and relation is not None and not relation.is_external
             and relation.reltype == expected, "native_relationship_unresolved",
             "原生对象缺少可核对的内部图像或OLE关系。")
    part = relation.target_part
    _require(part.content_type in ({"image/png", "image/wmf", "image/x-wmf"} if image else
                                  {"application/vnd.openxmlformats-officedocument.oleObject"}),
             "native_object_not_supported", "该原生对象类型尚未纳入此图文候选流程。")
    return {
        "relationship_id": rid, "relationship_type": relation.reltype,
        "package_part": str(part.partname), "mime_type": part.content_type,
        "sha256": _sha(part.blob), "bytes_count": len(part.blob),
    }


def _objects(element, document, index, assets):
    result, covered = [], []
    containers = [node for node in element.iter() if node.tag in _CONTAINERS]
    for ordinal, node in enumerate(containers, 1):
        feature, label = _CONTAINERS[node.tag]
        _require(not any(ancestor.tag in _CONTAINERS for ancestor in node.iterancestors()),
                 "nested_native_object", "嵌套原生对象需另行复核。")
        ole, shape = None, None
        if node.tag == _q("w:drawing"):
            _require(all(child.tag in _DRAWING_TAGS for child in node.iter())
                     and len(list(node.iter(_q("wp:inline")))) == 1
                     and len(list(node.iter(_q("pic:pic")))) == 1
                     and all(child.get("uri") == _NS["pic"] for child in node.iter(_q("a:graphicData"))),
                     "drawing_layout_not_supported", "仅接受无未支持形状、浮动布局或扩展的内嵌原图。")
            for child in node.iter():
                if child.tag in {_q("a:srcRect"), _q("a:fillRect"), _q("a:xfrm")}:
                    _require(all(value in {"0", "false"} for value in child.attrib.values()),
                             "transformed_image_not_supported", "图片裁剪、翻转或旋转尚未纳入读图证据链。")
                if child.tag == _q("a:blip"):
                    _require(not list(child) and set(child.attrib) <= {_q("r:embed"), "cstate"},
                             "transformed_image_not_supported", "图片效果尚未纳入读图证据链。")
                if child.tag == _q("a:prstGeom"):
                    _require(dict(child.attrib) == {"prst": "rect"}, "transformed_image_not_supported", "非矩形图片需另行复核。")
                if child.tag in {_q("a:avLst"), _q("a:noFill")}:
                    _require(not child.attrib and not list(child), "transformed_image_not_supported", "图片几何调整尚未纳入读图证据链。")
                if child.tag == _q("a:ln"):
                    _require(not child.attrib and [part.tag for part in child] == [_q("a:noFill")],
                             "transformed_image_not_supported", "图片边框效果需另行复核。")
            images = list(node.iter(_q("a:blip")))
            rid = images[0].get(_q("r:embed")) if len(images) == 1 else None
        else:
            allowed = {node.tag, _q("v:shape"), _q("v:imagedata")}
            if node.tag == _q("w:object"):
                allowed.update({_q("o:OLEObject"), _q("v:stroke"), _q("o:lock"), _q("w10:anchorlock")})
            _require(all(child.tag in allowed for child in node.iter()),
                     "native_object_not_supported", "旧式图形包含尚未支持的对象或版式。")
            shapes, images = list(node.iter(_q("v:shape"))), list(node.iter(_q("v:imagedata")))
            _require(len(shapes) == len(images) == 1 and images[0].getparent() is shapes[0],
                     "native_preview_ambiguous", "旧式对象没有唯一的同对象图像预览。")
            shape = shapes[0]
            _require(shape.get("type") == "#_x0000_t75"
                     and set(shape.attrib) <= {"id", "type", "alt", "style", _q("o:ole"),
                                               "coordsize", _q("o:preferrelative"), "filled", "stroked"}
                     and shape.get("filled") in {None, "f", "false", "0"}
                     and shape.get("stroked") in {None, "f", "false", "0"}
                     and all(part.split(":", 1)[0].strip() in {"width", "height"}
                             for part in shape.get("style", "").split(";") if part.strip())
                     and set(images[0].attrib) <= {_q("r:id"), _q("o:title")},
                     "transformed_image_not_supported", "旧式预览含裁剪、旋转或外链，需另行复核。")
            for child in shape:
                # These are non-rendering edit locks, or an explicitly disabled
                # outline. Other VML effects remain unsupported.
                if child.tag == _q("v:stroke"):
                    _require(shape.get("stroked") in {"f", "false", "0"}
                             and dict(child.attrib) == {"joinstyle": "miter"} and not list(child),
                             "transformed_image_not_supported", "旧式图像边框尚未证明不影响预览。")
                if child.tag == _q("o:lock"):
                    _require(dict(child.attrib) == {_q("v:ext"): "edit", "aspectratio": "t"} and not list(child),
                             "native_object_not_supported", "旧式对象锁定结构不受支持。")
                if child.tag == _q("w10:anchorlock"):
                    _require(not child.attrib and not list(child), "native_object_not_supported", "旧式锚定结构不受支持。")
            rid = images[0].get(_q("r:id"))
            if node.tag == _q("w:object"):
                oles = list(node.iter(_q("o:OLEObject")))
                _require(len(oles) == 1 and oles[0].getparent() is node
                         and oles[0].get("Type") == "Embed" and oles[0].get("DrawAspect") == "Content"
                         and oles[0].get("ProgID") in {"Equation.DSMT4", "Equation.3"}
                         and bool(shape.get("id")) and oles[0].get("ShapeID") == shape.get("id"),
                         "ole_same_object_unproven", "OLE公式与显示预览的同对象关系尚未证明。")
                ole = _relation(document, oles[0].get(_q("r:id")), image=False)
                ole.update(shape_id=shape.get("id"), prog_id=oles[0].get("ProgID"), binary_interpreted=False)
        _require(len(images) == 1 and all(key != _q("r:link") for key in images[0].attrib),
                 "native_preview_ambiguous", "原生图片并非唯一内嵌图像。")
        relation = _relation(document, rid, image=True)
        matches = [asset for asset in assets if asset["sha256"] == relation["sha256"]]
        _require(len(matches) == 1, "native_asset_mismatch", "原生图像与区块素材索引不一致或存在歧义。")
        expected_rids = {rid} | ({ole["relationship_id"]} if ole else set())
        for child in node.iter():
            for key, value in child.attrib.items():
                if key.startswith("{" + _NS["r"] + "}"):
                    _require(value in expected_rids and key in {_q("r:embed"), _q("r:id")},
                             "native_relationship_unresolved", "原生对象含未归属关系。")
        result.append({
            "block_index": index, "ordinal": ordinal, "feature": feature,
            "placeholder": "【待查看原文：" + label + "】", "container_xml_sha256": _xml_sha(node),
            "asset_id": matches[0]["asset_id"], "image": relation, "ole": ole,
            "coverage": "same_object_visible_preview_only" if ole else "embedded_image",
        })
        covered.append(matches[0]["asset_id"])
    _require(set(covered) == {asset["asset_id"] for asset in assets},
             "native_object_inventory_incomplete", "存在未归属原生对象的图像。")
    all_refs = [node for node in element.iter() if any(key.startswith("{" + _NS["r"] + "}") for key in node.attrib)]
    _require(all(any(a.tag in _CONTAINERS for a in node.iterancestors()) for node in all_refs),
             "native_relationship_unresolved", "区块含图像对象以外的未支持关系。")
    return result


def _png(raw):
    with Image.open(io.BytesIO(raw)) as image:
        _require(image.format == "PNG" and image.width * image.height <= 50_000_000
                 and getattr(image, "n_frames", 1) == 1,
                 "invalid_readable_image", "候选显示依据必须是单帧且尺寸可读的PNG。")
        image.load()
        return image.copy()


def _native_scope(words, row, attributes, preview, source):
    _require(row.get("selection_ready") is True and row.get("export_ready") is True
             and row.get("boundary_status") in {"auto_detected", "manual_range"},
             "range_pending_review", "题答或公共材料范围尚未核对。")
    quality = row.get("content_quality")
    _require(isinstance(quality, dict) and quality.get("status") == "not_fully_reviewed" and quality.get("issues") == [],
             "source_issue", "来源内容问题尚未解决。")
    material = attributes["material_status"]
    _require(material["missing_context"] is False and material["missing_visual"] is False
             and row.get("shared_material_policy") not in {"unknown", "blocked_pending_review"},
             "missing_material", "缺失材料或图像尚未解决。")
    _require(bool(row.get("question_blocks")) and bool(row.get("answer_blocks")),
             "answer_range_required", "图文候选须完整读取题干、答案与公共材料。")
    _require(_sha(source.content) == row["source_sha256"] == source.source_sha256
             and preview.get("source_sha256") == row["source_sha256"]
             and preview.get("revision") == row["source_revision"]
             and preview.get("extraction_revision") == row["extraction_revision"],
             "source_identity_changed", "原文件或提取版本变化。")
    with zipfile.ZipFile(io.BytesIO(source.content)) as package:
        _validate_container(package)
        document_xml_sha = _sha(package.read("word/document.xml"))
    document = Document(io.BytesIO(source.content))
    elements = list(_body_blocks(document._element.body))
    reader = WordNativeTextReader(document.styles.element)
    native_assets = _word_images(elements, document)
    cached = {block["index"]: block for block in preview["blocks"]}
    groups, objects, images, payloads, seen = {}, [], [], {}, set()
    for group, start, end in _GROUPS:
        blocks = row.get(group)
        _require(isinstance(blocks, list), "range_missing", "题答区块范围不完整。")
        indices = [block.get("index") for block in blocks]
        _require((not blocks and group == "context_blocks" and row.get(start) is None and row.get(end) is None)
                 or (all(type(index) is int and 0 < index <= len(elements) for index in indices)
                     and type(row.get(start)) is int and type(row.get(end)) is int
                     and indices == list(range(row[start], row[end] + 1)) and not seen.intersection(indices)),
                 "range_gap", "完整题答范围存在缺段或重叠。")
        seen.update(indices)
        groups[group] = []
        for block in blocks:
            index, element = block["index"], elements[block["index"] - 1]
            _require(element.tag == _q("w:p"), "native_block_not_supported", "此图文入口当前只接受完整原生段落。")
            warnings = set()
            text = _block_text(element, document, warnings, reader)
            native = reader.read(element)
            assets = [asset for asset in native_assets if asset["block_index"] == index]
            origin = cached.get(index)
            _require(origin is not None and text == block.get("text") == origin.get("text")
                     and sorted(warnings) == block.get("warnings", []) == origin.get("warnings", [])
                     and block.get("assets", []) == assets
                     and [a for a in preview.get("assets", []) if a.get("block_index") == index] == assets
                     and all(not record.get(key) for record in (block, origin)
                             for key in ("display_only_split", "text_range", "unsupported_assets", "gaps")),
                     "range_content_mismatch", "原生全文、区块缓存与题答范围不一致。")
            _require(set(native.features) <= {feature for feature, _label in _CONTAINERS.values()},
                     "unsupported_source_content", "仍存在域代码、隐藏文字或其他未读原生内容。")
            block_objects = _objects(element, document, index, assets)
            placeholders = re.findall(r"【待查看原文：[^】]*】", text)
            expected_warnings = sorted({item["placeholder"][7:-1] + "未提取；需查看原文件，不得根据残句补写。"
                                        for item in block_objects})
            _require(Counter(placeholders) == Counter(item["placeholder"] for item in block_objects)
                     and sorted(warnings) == expected_warnings,
                     "uncovered_object_warning", "占位或对象警告尚未逐项获得原生图像证据。")
            groups[group].append({"index": index, "text": text, "warnings": sorted(warnings),
                                  "assets": assets, "xml_sha256": _xml_sha(element),
                                  "native_features": list(native.features), "native_math_count": native.native_math_count})
            objects.extend(block_objects)
            for asset in assets:
                try:
                    display = words.reader.word_asset_bytes(source.content, asset["asset_id"],
                                                           render_metafiles=True, expected_sha256=asset["sha256"])
                    raster = _png(display["bytes"])
                except Exception as exc:
                    if isinstance(exc, OfflineWordLabelReviewError):
                        raise
                    raise OfflineWordLabelReviewError("image_not_readable", "原图或可信WMF转换尚不可读。") from exc
                derived = asset["mime_type"] != "image/png"
                display_sha = _sha(display["bytes"])
                _require(not derived or (display.get("derived_preview") is True
                         and display.get("original_sha256") == asset["sha256"]
                         and display.get("preview_sha256") == display_sha and bool(display.get("renderer_revision"))),
                         "derived_image_unbound", "派生图未闭合到原图字节及转换器。")
                _require(derived or display_sha == asset["sha256"], "original_image_changed", "PNG原图字节变化。")
                images.append({**asset, "display": {
                    "sha256": display_sha, "bytes_count": len(display["bytes"]), "mime_type": "image/png",
                    "renderer_revision": display["renderer_revision"] if derived else "original-png-v1",
                    "derived_preview": derived, "width": raster.width, "height": raster.height,
                }})
                payloads[asset["asset_id"]] = display["bytes"]
    _require(row["answer_start"] == row["question_end"] + 1,
             "range_gap", "题目与答案之间存在未归属区块。")
    _require(_has_substantive_answer_text(groups["answer_blocks"]), "answer_range_required", "答案正文不可确认为已完整读取。")
    _require(images, "image_scope_required", "纯文字题请使用既有纯文字入口。")
    _require(material["question_image_count"] == sum(len(b["assets"]) for g in ("question_blocks", "context_blocks") for b in groups[g])
             and material["answer_image_count"] == sum(len(b["assets"]) for b in groups["answer_blocks"])
             and sorted(material["warnings"]) == sorted({w for g in groups.values() for b in g for w in b["warnings"]}),
             "material_inventory_mismatch", "属性材料清单与完整原生范围不一致。")
    return {
        "schema_version": SCOPE_SCHEMA_VERSION, "identity": {key: row[key] for key in _IDENTITY},
        "ranges": {key: row.get(key) for key in _RANGES}, "blocks": groups,
        "objects": objects, "images": images, "document_xml_sha256": document_xml_sha,
    }, payloads


def _root_review(entry, row, scope, payloads):
    root = _bound_json(entry["root_review"])
    _require(root.get("schema_version") == ROOT_REVIEW_SCHEMA_VERSION
             and root.get("candidate_only") is True and root.get("human_reviewed") is False,
             "root_model_review_required", "须绑定主代理实际读图的候选复核记录，不能使用人工确认入口。")
    matches = [item for item in root.get("items", []) if isinstance(item, dict) and item.get("key") == row["key"]]
    _require(len(matches) == 1, "root_review_selection_mismatch", "主代理复核记录没有唯一对应题目。")
    item = matches[0]
    _require(item.get("source_sha256") == row["source_sha256"] and item.get("question_revision") == row["revision"]
             and item.get("expected_stored_attribute_revision") == entry["expected_stored_attribute_revision"]
             and item.get("decision") == "candidate" and item.get("candidate_only") is True
             and item.get("human_reviewed") is False and item.get("teacher_confirmed") is False
             and item.get("automatic_source_gate_cleared") is False
             and item.get("full_extracted_question_answer_context_text_read") is True
             and item.get("neighbor_text_read") is True
             and item.get("block_indices_read") == {group: [b["index"] for b in scope["blocks"][group]] for group, _, _ in _GROUPS}
             and item.get("decoded_label_proposal_reviewed") == entry["decoded"],
             "root_review_binding_mismatch", "候选内容、完整题答范围或旧标签与主代理复核记录不一致。")
    evidence = _bound_json(item.get("source_evidence"))
    _bound_file(item.get("reading_file"))
    reviewed_row = evidence.get("question", {})
    _require(all(reviewed_row.get(key) == row.get(key) for key in (*_IDENTITY, *_RANGES))
             and all(reviewed_row.get(group) == row.get(group) for group, _, _ in _GROUPS)
             and evidence.get("document_xml_sha256") == scope["document_xml_sha256"],
             "reviewed_text_changed", "主代理已读的冻结文字与当前完整原生题答不一致。")
    # Verify every archived original object, including uninterpreted OLE bytes.
    original_occurrences = [(obj["block_index"], rel)
                            for obj in scope["objects"] for rel in (obj["image"], obj["ole"]) if rel is not None]
    originals = {(index, rel["relationship_id"]): rel for index, rel in original_occurrences}
    expected_occurrences = Counter((index, rel["relationship_id"]) for index, rel in original_occurrences)
    media = evidence.get("media")
    _require(isinstance(media, list) and len(media) == len(original_occurrences),
             "reviewed_object_inventory_changed", "复核证据未覆盖全部原生对象。")
    media_keys = []
    for media_item in media:
        key = (media_item.get("block_index"), media_item.get("relationship_id"))
        relation = originals.get(key)
        _require(relation is not None and all(media_item.get(name) == relation[name] for name in
                                            ("package_part", "mime_type", "sha256", "bytes_count")),
                 "reviewed_object_inventory_changed", "冻结媒体与原生对象关系不一致。")
        _bound_file({"path": media_item.get("local_original_path"), "sha256": relation["sha256"],
                     "bytes": relation["bytes_count"]}, max_bytes=32_000_000)
        media_keys.append(key)
    _require(Counter(media_keys) == expected_occurrences, "reviewed_object_inventory_changed", "媒体证据的原生出现次数不一致。")
    trusted = {}
    for image in scope["images"]:
        trusted.setdefault(image["display"]["sha256"], set()).add(image["asset_id"])
    display_bindings = []
    for composite in root.get("white_composites_verified", []):
        # Other questions in a batch need not be opened for this selection.
        original_ref = composite.get("original", {})
        original_sha = original_ref.get("sha256")
        matches = [image for image in scope["images"] if image["mime_type"] == "image/png" and image["sha256"] == original_sha]
        if not matches:
            continue
        _require(composite.get("pixel_equality") is True,
                 "white_display_unbound", "白底显示图缺少唯一原图或像素相等声明。")
        original = _bound_file(original_ref, max_bytes=32_000_000)
        _require(original == payloads[matches[0]["asset_id"]], "white_display_unbound", "白底显示图原件不一致。")
        display_ref = composite.get("display")
        displayed = _png(_bound_file(display_ref, max_bytes=32_000_000))
        rgba = _png(original).convert("RGBA")
        expected = Image.alpha_composite(Image.new("RGBA", rgba.size, "white"), rgba).convert("RGB")
        _require(displayed.mode == "RGB" and displayed.size == expected.size and displayed.tobytes() == expected.tobytes(),
                 "white_display_pixels_changed", "白底显示图与原图无损合成像素不符。")
        trusted.setdefault(display_ref["sha256"], set()).update(image["asset_id"] for image in matches)
        display_bindings.append({"original_sha256": original_sha, "display_sha256": display_ref["sha256"],
                                 "renderer_revision": WHITE_RENDERER})
    viewed = item.get("images_actually_viewed")
    _require(isinstance(viewed, list) and bool(viewed), "actual_image_read_required", "主代理尚未声明实际查看图像。")
    seen_assets, viewed_shas = set(), set()
    for record in viewed:
        _require(isinstance(record, dict) and record.get("viewed_by_root_model") is True
                 and record.get("human_reviewed") is False and bool(record.get("observation"))
                 and record.get("sha256") in trusted,
                 "actual_image_read_unbound", "实际读图声明未绑定可信原图或派生图。")
        _bound_file(record, max_bytes=32_000_000)
        seen_assets.update(trusted[record["sha256"]])
        viewed_shas.add(record["sha256"])
    _require(seen_assets == {image["asset_id"] for image in scope["images"]},
             "image_not_actually_read", "题干、答案或共用材料仍有未实际查看的图像。")
    allowed_by_block = {}
    for sha in viewed_shas:
        for asset_id in trusted[sha]:
            image = next(image for image in scope["images"] if image["asset_id"] == asset_id)
            allowed_by_block.setdefault(image["block_index"], []).append(sha)
    return allowed_by_block, {
        "root_review_sha256": entry["root_review"]["sha256"],
        "source_evidence_sha256": item["source_evidence"]["sha256"],
        "reading_file_sha256": item["reading_file"]["sha256"],
        "viewed_display_sha256s": sorted(viewed_shas), "white_display_bindings": display_bindings,
        "native_object_count": len(scope["objects"]),
        "uninterpreted_ole_with_same_object_preview_count": sum(obj["ole"] is not None for obj in scope["objects"]),
        "native_object_warnings_preserved": True,
    }


class OfflineWordImageLabelReviewService(OfflineWordLabelReviewService):
    """Independent graph/image admission, shared transactional label writer."""

    def _validate_candidates(self, candidates):
        return validate_candidate_batch(candidates)

    def _policy(self, mode):
        return (RECHECK_RULE_REVISION if mode == "recheck_automatic" else RULE_REVISION,
                READING_POLICY_REVISION, PLAN_SCHEMA_VERSION)

    def _receipt_schema(self):
        return RECEIPT_SCHEMA_VERSION

    def describe_entry(self, identity):
        """Read-only full evidence for authoring a candidate; grants no admission.

        Unlike public preview reports this explicit evidence API contains the
        complete local question/answer/context text and the native inventory.
        """
        with self.words._lock:
            self._prime_locations([identity])
            rows, inventory = self.words._resolve([
                {"key": identity["key"], "revision": identity["question_revision"]},
            ], read_only=True)
            row = rows[0]
            old = validate_attributes(self.words.attribute_store.get(row["key"]))
            _require(row["source_sha256"] == identity["source_sha256"]
                     and old["revision"] == identity["expected_stored_attribute_revision"],
                     "identity_changed", "来源或已有属性版本已变化。")
            source, _batch, preview, _index = inventory[row["source_id"]]
            scope, _payloads = _native_scope(self.words, row, old, preview, source)
            return {"reading_scope": scope, "reading_scope_sha256": _digest(scope)}

    def _reading_scope(self, entry, row, old, preview, source, *, mode):
        try:
            scope, payloads = _native_scope(self.words, row, old, preview, source)
            _require(_digest(entry["reading_scope"]) == entry["reading_scope_sha256"] == _digest(scope),
                     "image_reading_scope_changed", "完整题答文字、原生对象或图像转换与已读范围不一致。")
            allowed, metadata = _root_review(entry, row, scope, payloads)
            blocks = [{"index": block["index"], "kind": kind, "text": block["text"],
                       "image_sha256s": sorted(allowed.get(block["index"], []))}
                      for group, kind in (("context_blocks", "shared_context"), ("question_blocks", "question_text"))
                      for block in scope["blocks"][group]]
            return blocks, _digest(scope), metadata
        except OfflineWordLabelReviewError:
            raise
        except (OSError, ValueError, TypeError, KeyError, AttributeError) as exc:
            raise OfflineWordLabelReviewError("image_review_evidence_invalid", "图文复核证据不完整或不可复现，未应用候选。") from exc
