"""Activity-owned student worksheets, separate from teacher notes and answers."""

from collections.abc import Mapping
from typing import Any


def _object(properties: dict[str, Any]) -> dict[str, Any]:
    return {
        "type": "object",
        "properties": properties,
        "required": list(properties),
        "additionalProperties": False,
    }


def worksheet_schema() -> dict[str, Any]:
    text = {"type": "string"}
    strings = {"type": "array", "items": text}
    return {
        "anyOf": [
            {"type": "null"},
            _object(
                {
                    "title": text,
                    "instructions": strings,
                    "sections": {
                        "type": "array",
                        "items": _object(
                            {
                                "heading": text,
                                "prompt": text,
                                "response_kind": {
                                    "type": "string",
                                    "enum": ["lines", "table"],
                                },
                                "response_lines": {"type": "integer"},
                                "columns": strings,
                                "row_labels": strings,
                            }
                        ),
                    },
                }
            ),
        ]
    }


def _keys(value: Any, keys: set[str]) -> Mapping[str, Any]:
    if not isinstance(value, Mapping) or set(value) != keys:
        raise ValueError("invalid worksheet fields")
    return value


def _text(value: Any, limit: int) -> str:
    if not isinstance(value, str) or not 1 <= len(value.strip()) <= limit:
        raise ValueError("invalid worksheet text")
    return value.strip()


def _list(value: Any, minimum: int, maximum: int) -> list[Any]:
    if not isinstance(value, list) or not minimum <= len(value) <= maximum:
        raise ValueError("invalid worksheet item count")
    return value


def normalize_worksheet(value: Any) -> dict[str, Any] | None:
    """Normalize only supplied tasks; never derive questions from teacher prose."""
    if value is None:
        return None
    value = _keys(value, {"title", "instructions", "sections"})
    result: dict[str, Any] = {
        "title": _text(value["title"], 80),
        "instructions": [
            _text(item, 300) for item in _list(value["instructions"], 1, 4)
        ],
        "sections": [],
    }
    for section in _list(value["sections"], 1, 6):
        section = _keys(
            section,
            {
                "heading",
                "prompt",
                "response_kind",
                "response_lines",
                "columns",
                "row_labels",
            },
        )
        kind = section["response_kind"]
        lines = section["response_lines"]
        if type(lines) is not int:
            raise ValueError("worksheet response lines must be integer")
        if kind == "lines":
            if not 2 <= lines <= 10:
                raise ValueError("invalid worksheet writing space")
            _list(section["columns"], 0, 0)
            _list(section["row_labels"], 0, 0)
        elif kind == "table":
            if lines != 0:
                raise ValueError("table cannot also specify writing lines")
            _list(section["columns"], 2, 4)
            _list(section["row_labels"], 1, 6)
        else:
            raise ValueError("unknown worksheet response kind")
        result["sections"].append(
            {
                "heading": _text(section["heading"], 80),
                "prompt": _text(section["prompt"], 500),
                "response_kind": kind,
                "response_lines": lines,
                "columns": [_text(item, 40) for item in section["columns"]],
                "row_labels": [_text(item, 60) for item in section["row_labels"]],
            }
        )
    return result


def worksheet_records(
    candidate: Mapping[str, Any],
) -> list[tuple[int, Mapping[str, Any], dict[str, Any]]]:
    result = []
    for index, activity in enumerate(candidate.get("activities", []), 1):
        if not isinstance(activity, Mapping):
            raise TypeError("invalid worksheet activity")
        worksheet = normalize_worksheet(activity.get("worksheet"))
        if worksheet is not None:
            result.append((index, activity, worksheet))
    return result


def build_worksheet_document(candidate: Mapping[str, Any]) -> Any:
    """Build student-facing OOXML, leaving all response areas genuinely blank."""
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt, RGBColor, Twips

    records = worksheet_records(candidate)
    if not records:
        raise ValueError("no activity worksheets supplied")
    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Mm(210), Mm(297)
    page.top_margin = page.bottom_margin = Inches(0.7)
    page.left_margin = page.right_margin = Inches(0.8)
    # Word stores page and table widths in twips (1/20 point).  Compute the
    # usable width in that same integer unit before assigning columns, so
    # rounding cannot make a table one unit wider than the page.
    usable_width_twips = round(
        (page.page_width - page.left_margin - page.right_margin) / 635
    )
    for name, size in (
        ("Normal", 11),
        ("Title", 21),
        ("Heading 1", 14),
        ("Heading 2", 12),
    ):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.1
        properties = style._element.get_or_add_pPr()
        for border in list(properties.findall(qn("w:pBdr"))):
            properties.remove(border)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(10)

    for record_index, (activity_number, _activity, worksheet) in enumerate(records):
        title_paragraph = document.add_paragraph(worksheet["title"], style="Title")
        if record_index:
            # Attach the page break to the next worksheet's title.  A separate
            # add_page_break() paragraph can be pushed onto a fresh page when
            # the preceding worksheet already reaches the page bottom,
            # leaving an avoidable blank page in the printed document.
            title_paragraph.paragraph_format.page_break_before = True
        document.add_paragraph(f"活动 {activity_number}  课堂学习单")
        for instruction in worksheet["instructions"]:
            document.add_paragraph(instruction)
        for section in worksheet["sections"]:
            document.add_heading(section["heading"], level=1)
            paragraph = document.add_paragraph(section["prompt"])
            paragraph.paragraph_format.keep_with_next = True
            if section["response_kind"] == "lines":
                for line_index in range(section["response_lines"]):
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = Pt(24)
                    # A nonbreaking blank preserves real writing space. It is
                    # not a placeholder answer or an underscore text run.
                    paragraph.add_run("\u00a0")
                    border = OxmlElement("w:pBdr")
                    for edge in ("bottom", "between"):
                        rule = OxmlElement("w:" + edge)
                        for attr, value in (
                            ("val", "single"),
                            ("sz", "4"),
                            ("color", "D9D9D9"),
                        ):
                            rule.set(qn("w:" + attr), value)
                        border.append(rule)
                    paragraph._p.get_or_add_pPr().append(border)
                    paragraph.paragraph_format.keep_with_next = line_index < 1
            else:
                columns = section["columns"]
                table = document.add_table(rows=1, cols=len(columns))
                table.autofit = False
                # Preserve the former first-column proportion while deriving
                # the complete table width from the current A4 page margins.
                label_width = round(usable_width_twips * 1.35 / (1.35 + 5.55))
                column_base, extra = divmod(
                    usable_width_twips - label_width, len(columns) - 1
                )
                widths = [label_width] + [
                    column_base + (1 if index < extra else 0)
                    for index in range(len(columns) - 1)
                ]
                props = table._tbl.tblPr
                borders = OxmlElement("w:tblBorders")
                for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                    border = OxmlElement("w:" + edge)
                    for attr, value in (
                        ("val", "single"),
                        ("sz", "4"),
                        ("color", "D9D9D9"),
                    ):
                        border.set(qn("w:" + attr), value)
                    borders.append(border)
                props.append(borders)
                margins = OxmlElement("w:tblCellMar")
                for edge, value in (
                    ("top", "80"),
                    ("bottom", "80"),
                    ("left", "120"),
                    ("right", "120"),
                ):
                    margin = OxmlElement("w:" + edge)
                    margin.set(qn("w:w"), value)
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                props.append(margins)
                table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                for index, label in enumerate(columns):
                    cell = table.rows[0].cells[index]
                    cell.text = label
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "E8F0F6")
                    cell._tc.get_or_add_tcPr().append(shading)
                    for run in cell.paragraphs[0].runs:
                        run.font.bold = True
                for label in section["row_labels"]:
                    row = table.add_row()
                    row.height = Inches(0.5)
                    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
                    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                    row.cells[0].text = label
                for index, width in enumerate(widths):
                    column_width = Twips(width)
                    table.columns[index].width = column_width
                    for row in table.rows:
                        row.cells[index].width = column_width
                        row.cells[
                            index
                        ].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        for paragraph in row.cells[index].paragraphs:
                            paragraph.paragraph_format.space_after = Pt(0)
                spacer = document.add_paragraph()
                spacer.paragraph_format.space_after = Pt(0)
                spacer.paragraph_format.line_spacing = Pt(4)
    document.core_properties.title = "课堂学习单"
    document.core_properties.author = ""
    return document
