"""Synthetic teacher operations for source and shipping-EXE verification only."""
from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import sys


def exercise(window, settle, capture, output):
    from docx import Document
    from PySide6.QtCore import QPoint
    from PySide6.QtWidgets import QApplication, QMessageBox
    from .desktop_mixed_paper_service import MixedPaperService, REQUEST_SCHEMA
    output = Path(output); output.mkdir(parents=True, exist_ok=True)
    facade, page = window.facade, window.preparation_page
    window.resize(800,700);window.navigate('preparation');settle()
    page.materials.setPlainText('合成软件验收材料。\n' * 30)
    before = deepcopy(page._payload())
    page.workspace.requirements.click();page.workspace.materials.click()
    page.editor_scroll.verticalScrollBar().setValue(page.editor_scroll.verticalScrollBar().maximum())
    settle()
    for button in (page.save_button, page.generate_button, page.workspace.results):
        pos = button.mapTo(page, QPoint(0,0))
        assert button.isVisible() and 0 <= pos.y() and pos.y()+button.height() <= page.height()
    assert page._payload()==before
    capture(window,'preparation-fixed-actions.png')
    window.resize(1360,900);page.workspace.materials.click();settle()
    capture(window,'preparation-materials.png')
    app=QApplication.instance()
    assert app._workbench_chinese_translation_loaded
    message=QMessageBox(window)
    message.setWindowTitle('操作确认（合成演示）');message.setText('是否保留本次修改？')
    message.setStandardButtons(QMessageBox.StandardButton.Ok|QMessageBox.StandardButton.Cancel)
    assert '确定' in message.button(QMessageBox.StandardButton.Ok).text()
    assert '取消' in message.button(QMessageBox.StandardButton.Cancel).text()
    message.show();settle();capture(message,'chinese-confirmation.png');message.close();message.deleteLater()

    source=Document()
    source.add_heading('本卷编号验收（合成，非真实试题）',0)
    source.add_paragraph('【例27】标记甲：填写化学式，并保留答题位置 ______。')
    source.add_paragraph('【答案】27. 合成答案甲，只能出现在教师版。')
    source.add_paragraph('【例38】标记乙：参考第27题所给条件。数值12.5不应改变 ______。')
    source.add_paragraph('【答案】38. 合成答案乙，参考第27题。')
    path=output/'合成题号验收.docx';source.save(path)
    original=path.read_bytes()
    facade.save_visual_import_batch(handout_files=(path,), source_type='教师讲义')
    rows=[r for r in facade.word_question_catalog()['items'] if r['source_name']==path.name]
    assert len(rows)==2
    # Only this synthetic test basket is modified. Source records are untouched.
    for entry in list(facade.basket()):
        facade.remove_basket_item(entry['key'])
    facade.add_word_questions_to_basket([{'key':r['key'],'revision':r['revision']} for r in reversed(rows)])
    service=MixedPaperService(facade);projection=service.projection()
    request={'schema_version':REQUEST_SCHEMA,'title':'重新编号验收（合成）','subtitle':'仅验证题序与题答对应',
             'mode':'daily_practice','duration_minutes':40,'show_question_scores':False,
             'basket_sha256':projection['basket_sha256'],'section_order':[r['key'] for r in projection['items']],
             'settings_by_key':{}}
    preview=service.create_preview(request)
    folder,_record,snapshot=service._load(preview.preview_id)
    documents=service._build_docx(preview.preview_id,snapshot,folder)
    for audience in ('student','teacher'):
        raw=documents[audience+'_bytes'];(output/(audience+'.docx')).write_bytes(raw)
        text='\n'.join(p.text for p in Document(BytesIO(raw)).paragraphs)
        assert '1. 标记乙' in text and '2. 标记甲' in text
        assert '数值12.5' in text and '参考第2题' in text
        assert '27. ' not in text and '38. ' not in text
        assert ('合成答案乙' in text) == (audience=='teacher')
    assert path.read_bytes()==original
    page_count = None
    if getattr(sys, 'frozen', False):
        paginated=service.prepare_pagination(preview.preview_id,preview.preview_hash)
        model=paginated.preview_model
        page_count={role:model['pagination']['documents'][role]['page_count'] for role in ('student','teacher')}
        assert all(page_count.values())
        # Capture real rendered PDF pages, not the content-only preview.
        for role in ('student','teacher'):
            record=model['pagination']['documents'][role]['pages'][0]
            data=service.image(paginated.preview_id,record['image_id'])['data']
            from PySide6.QtWidgets import QLabel
            from PySide6.QtGui import QPixmap
            view=QLabel();pix=QPixmap();assert pix.loadFromData(data)
            view.setPixmap(pix.scaledToHeight(900));view.show();settle()
            capture(view,'paper-renumbered-'+role+'.png');view.close();view.deleteLater()
    return {'real_word_import':True,'reordered_student_and_teacher':True,'source_bytes_unchanged':True,
            'reference_numbers_updated':True,'fixed_actions_800x700':True,'chinese_standard_buttons':True,
            'actual_pdf_pages':page_count,'raster_numbers_auto_rewritten':False,'model_calls':0}
