"""Display sealed Word tables without exposing the extraction coordinates.

The work copy and the original coordinate text remain untouched. A table is
projected only while its complete work text still equals the sealed original.
All geometry comes from the source grid and merge records, never cell wording.
"""
import math
import re

_SCRIPT_START = re.compile(r'[_^]\{')
_SCRIPT_CHARACTERS = {
    'subscript': dict(zip('0123456789a', '₀₁₂₃₄₅₆₇₈₉ₐ')),
    'superscript': dict(zip('0123456789+-−＋－', '⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁻⁺⁻')),
}


def _cell_parts(text):
    """Project only complete, explicit markers with supported characters.

    Unknown, escaped, nested and unfinished markers remain literal, including
    their braces. A nested marker is never recovered from an unknown outer one.
    The same payload and script boundaries drive Unicode and editable Word.
    """
    offset = cursor = 0
    while match := _SCRIPT_START.search(text, cursor):
        depth, end = 1, match.end()
        while end < len(text) and depth:
            if text[end] == '{': depth += 1
            elif text[end] == '}': depth -= 1
            end += 1
        if depth:
            break
        cursor = end
        payload = text[match.end():end - 1]
        script = 'subscript' if text[match.start()] == '_' else 'superscript'
        if (not payload or any(c not in _SCRIPT_CHARACTERS[script] for c in payload)
                or (match.start() and text[match.start() - 1] == '\\')):
            continue
        if offset < match.start():
            yield text[offset:match.start()], None
        yield payload, script
        offset = end
    if offset < len(text):
        yield text[offset:], None


def classroom_cell_text(text):
    """Display supported explicit scripts without changing the sealed text."""
    return ''.join(''.join(_SCRIPT_CHARACTERS[script][c] for c in payload)
        if script else payload for payload, script in _cell_parts(text))


def source_table(item, row):
    """Compatibility helper for text-only consumers; never omit table images."""
    from .desktop_handout_table_projection import source_table_projection
    projection = source_table_projection(item, row)
    return projection['rows'] if projection is not None and not projection['occurrences'] else None


def _grid(rows):
    from .desktop_handout_knowledge import _validate_table_rows
    _validate_table_rows(rows)
    counts = [r['grid_before'] + r['grid_after'] + sum(c['column_span'] for c in r['cells']) for r in rows]
    if not counts[0] or len(set(counts)) != 1:
        raise ValueError('原表格网格宽度不一致，不能推测缺失的列。')
    cells, active = [], {}
    for ri, row in enumerate(rows):
        column, current = row['grid_before'], {}
        for cell in row['cells']:
            end = column + cell['column_span']
            projected = {'row': ri, 'column': column, 'end': end, 'cell': cell, 'merge_start': ri}
            key = (column, end)
            if cell['vertical_merge'] == 'continue':
                if key not in active:
                    raise ValueError('原表格纵向合并缺少对应起点，不能补猜合并关系。')
                projected['merge_start'] = active[key]
                current[key] = active[key]
            elif cell['vertical_merge'] == 'restart':
                current[key] = ri
            cells.append(projected); column = end
        active = current
    return counts[0], cells


def _column_widths(columns, width, weights):
    if weights is None: return [width / columns] * columns
    if (not isinstance(weights, (tuple, list)) or len(weights) != columns
            or any(type(w) not in (int, float) or not math.isfinite(w) or w <= 0 for w in weights)
            or not math.isfinite(sum(weights))):
        raise ValueError('表格显示列宽须对应完整原网格，并使用正数权重。')
    total = sum(weights)
    return [width * w / total for w in weights]


def _row_segments(ri, sources, measured, height, budget, padding, minimum):
    """Slice only between complete measured lines; native objects stay whole."""
    ranges = {s['column']: (0, len(measured[(ri, s['column'])])) for s in sources}
    if height <= budget + .001:
        return [dict(row=ri, height=height, ranges=ranges, fragment=(1, 1))]
    if any(s['cell']['vertical_merge'] != 'none' for s in sources):
        raise ValueError('完整表格行包含纵向合并，不能在行内推测合并续段。')
    usable = budget - 2 * padding
    if usable <= 0 or minimum > budget + .001:
        raise ValueError('完整表格行超过投影画幅，单元格没有完整行的可用高度。')
    for s in sources:
        for chunk in measured[(ri, s['column'])]:
            h = chunk.get('height')
            if (type(h) not in (int, float) or not math.isfinite(h) or h < 0
                    or h > usable + .001):
                raise ValueError('完整表格行中的单个文字行、图片或公式超过续页画幅。')
    cursors = {ci: 0 for ci in ranges}; segments = []
    while any(cursors[ci] < ranges[ci][1] for ci in ranges):
        selected, tallest, progressed = {}, 0.0, False
        for ci in ranges:
            chunks = measured[(ri, ci)]; start = end = cursors[ci]; used = 0.0
            while end < len(chunks) and used + chunks[end]['height'] <= usable + .001:
                used += chunks[end]['height']; end += 1
            selected[ci] = (start, end); cursors[ci] = end
            tallest = max(tallest, used); progressed |= end > start
        if not progressed:
            raise ValueError('完整表格行没有可安全续页的完整排版边界。')
        segments.append(dict(row=ri, height=max(minimum, tallest + 2 * padding), ranges=selected))
        if len(segments) > 30000:
            raise ValueError('完整表格超出续页排版容量。')
    for index, segment in enumerate(segments, 1): segment['fragment'] = (index, len(segments))
    return segments


def _continued_table_chunks(rows, cells, groups, measured, heights, column_positions,
        block_id, width, max_height, first_height, padding, minimum):
    """Continue oversized rows, and only proven empty vertical continuations.

    The ordinary whole-row path remains unchanged. Here the first source row is
    still a complete repeated guide. A merge chain may cross a source-row page
    boundary only when every continuation cell is empty. Its original origin,
    grid and span remain attached to each derived box; restart content is never
    repeated. Nonempty continuation cells and oversized merged rows are refused.
    """
    if groups[0] != [0] or len(rows) == 1:
        raise ValueError('表格首行纵向合并或单行表格不能推测续页列关系。')
    header = heights[0]; budget = max_height - header
    if budget < minimum - .001:
        raise ValueError('完整表格行超过投影画幅，完整列标题后没有可用高度。')
    by_row = {ri: [s for s in cells if s['row'] == ri] for ri in range(len(rows))}
    def complete(ri):
        return dict(row=ri, height=heights[ri], fragment=(1, 1),
            ranges={s['column']: (0, len(measured[(ri, s['column'])])) for s in by_row[ri]})
    units = []
    for group in groups[1:]:
        height = sum(heights[ri] for ri in group)
        if height <= budget + .001:
            units.append([complete(ri) for ri in group]); continue
        continuations = [s for ri in group for s in by_row[ri]
            if s['cell']['vertical_merge'] == 'continue']
        if any(s['cell']['text'] != '' for s in continuations):
            raise ValueError('完整表格行的纵向合并续接含内容，不能推测跨页分配。')
        # Empty source text is not enough if a projected token somehow carries
        # a picture or other object. Only genuinely empty native chunks qualify.
        if any(any(c['kind'] not in {'space', 'text'} or c.get('text', '') != ''
                for c in measured[(s['row'], s['column'])]) for s in continuations):
            raise ValueError('完整表格行的纵向合并续接包含显示对象，不能跨页。')
        for ri in group:
            units.extend([[s] for s in _row_segments(ri, by_row[ri], measured,
                heights[ri], budget, padding, minimum)])
    if not units:
        raise ValueError('完整表格没有可续页的数据行。')
    first_limit = min(max_height, max_height if first_height is None else first_height)
    restarted = header + sum(s['height'] for s in units[0]) > first_limit + .001
    if restarted: first_limit = max_height
    panels, pending, used, limit = [], [complete(0)], header, first_limit
    for unit in units:
        height = sum(s['height'] for s in unit)
        same_fragmented_row = (len(pending) > 1 and pending[-1]['row'] == unit[0]['row']
            and unit[0]['fragment'][1] > 1)
        if len(pending) > 1 and (used + height > limit + .001 or same_fragmented_row):
            panels.append(pending); pending, used, limit = [complete(0)], header, max_height
        if used + height > limit + .001:
            raise ValueError('完整表格行超过续页可用画幅。')
        pending.extend(unit); used += height
    if len(pending) > 1: panels.append(pending)
    result = []
    for pi, segments in enumerate(panels):
        offsets, cursor = [], 0.0
        for segment in segments:
            offsets.append(cursor); cursor += segment['height']
        parts, drawn_merges = [], set()
        for si, segment in enumerate(segments):
            ri = segment['row']; y = offsets[si]
            for source in by_row[ri]:
                ci, end = source['column'], source['end']; x = column_positions[ci]
                key = (ri + 1, ci + 1, end - ci)
                origin = (source['merge_start'] + 1, ci + 1, end - ci)
                metadata = dict(_table_cell=key, _table_merge_origin=origin,
                    _table_vertical_merge=source['cell']['vertical_merge'],
                    _table_row_fragment=segment['fragment'])
                merged = source['cell']['vertical_merge'] != 'none'
                if not merged or origin not in drawn_merges:
                    height = segment['height']
                    if merged:
                        height = sum(other['height'] for other in segments
                            if any(s['column'] == ci and s['end'] == end
                                and s['merge_start'] == source['merge_start']
                                for s in by_row[other['row']]))
                        drawn_merges.add(origin)
                    parts.append(dict(kind='diagram_box', offset_x=x, offset_y=y,
                        width=column_positions[end] - x, height=height,
                        fill='#edf5fb' if ri == 0 else '#ffffff', stroke='#5b7184',
                        block_id=block_id, **metadata))
                start, finish = segment['ranges'][ci]; offset = y + padding
                for index in range(start, finish):
                    chunk = measured[(ri, ci)][index]
                    parts.append({**chunk, 'offset_x': x + padding, 'offset_y': offset,
                        **metadata, '_table_source_chunk': index + 1,
                        '_table_source_chunk_count': len(measured[(ri, ci)])})
                    offset += chunk['height']
        if len(parts) > 30000:
            raise ValueError('完整表格超出排版容量。')
        result.append(dict(kind='diagram', width=width, height=cursor, block_id=block_id,
            _parts=parts, _table_source_rows=[s['row'] + 1 for s in segments],
            _table_repeated_first_row=pi > 0, _table_restarted_page=pi == 0 and restarted,
            _table_continued_cells=True,
            _table_row_fragments=[(s['row'] + 1, *s['fragment']) for s in segments]))
    return result


def table_chunks(rows, block_id, width, lines, body_px, *, max_height=130.0, first_height=None, column_weights=None,
        cell_tokens=None, assets=None, image_data=None, image_limits=None):
    """Measure native text and emit whole table panels at the same body size.

    Panels contain consecutive source rows. When a table needs more than one
    panel, its first source row repeats as a column guide. A vertical merge is
    kept together when it fits. Oversized rows continue only at complete native
    line boundaries; empty merge continuations may cross original row boundaries.
    Indivisible objects and unproven merge splits still raise. ``first_height`` is
    the measured space after preceding content, while later panels use the
    complete page budget. A first panel always contains at least the first
    source row and one complete data-row group. If that group cannot fit the
    remaining area, its first panel uses a complete new-page budget rather than
    carrying an unnecessarily small panel onto the fresh page. ``lines`` has the same call
    signature as desktop_composition_render._lines.
    ``column_weights`` changes display widths only; each merged cell spans the
    exact sum of its original grid columns.
    """
    if (type(width) not in (int, float) or not math.isfinite(width) or width <= 0
            or type(max_height) not in (int, float) or not math.isfinite(max_height) or max_height <= 0
            or type(body_px) is not int or body_px < 1):
        raise ValueError('表格画幅与正文字号无效。')
    if first_height is not None and (type(first_height) not in (int, float)
            or not math.isfinite(first_height) or first_height < 0):
        raise ValueError('表格首段剩余画幅无效。')
    columns, cells = _grid(rows)
    column_widths = _column_widths(columns, width, column_weights)
    column_positions = [0.0]
    for column_width in column_widths: column_positions.append(column_positions[-1] + column_width)
    padding = max(1.5, body_px / (96 / 25.4) * .22)
    minimum = body_px / (96 / 25.4) + 2 * padding
    heights = [minimum] * len(rows)
    measured = {}
    for source in cells:
        ri, ci = source['row'], source['column']
        cell_width = column_positions[source['end']] - column_positions[ci]
        if cell_width <= 2 * padding + 2:
            raise ValueError('表格列太窄，请加宽画幅；完整单元格文字保留。')
        if (ri, ci) in (cell_tokens or {}):
            from .desktop_composition_rich_render import rich_lines
            tokens = [{**t, 'text': classroom_cell_text(t['text'])} if t['kind'] == 'text' else t
                for t in cell_tokens[(ri, ci)]]
            chunks = rich_lines({'id': block_id}, {}, cell_width - 2 * padding,
                assets or {}, image_data or {}, max_height - 2 * padding,
                font_px=body_px, bold=ri == 0 or ci == 0, tokens=tokens, image_limits=image_limits)
        else:
            chunks = lines(classroom_cell_text(source['cell']['text']), cell_width - 2 * padding,
                block_id=block_id, font_px=body_px, bold=ri == 0 or ci == 0)
        measured[(ri, ci)] = chunks
        heights[ri] = max(heights[ri], sum(c['height'] for c in chunks) + 2 * padding)
    # Source merge chains determine which row boundaries may be used. No cell
    # labels or chemical meaning participate in this decision.
    cuts = set(range(1, len(rows)))
    for cell in cells:
        for boundary in range(cell['merge_start'] + 1, cell['row'] + 1):
            cuts.discard(boundary)
    groups, start = [], 0
    for end in sorted(cuts | {len(rows)}):
        groups.append(list(range(start, end))); start = end
    if (groups[0] == [0] and any(heights[0] + sum(heights[r] for r in group) > max_height + .001
            for group in groups[1:])):
        return _continued_table_chunks(rows, cells, groups, measured, heights,
            column_positions, block_id, width, max_height, first_height, padding, minimum)
    total = sum(heights)
    first_limit = min(max_height, max_height if first_height is None else first_height)
    minimum_group = heights[0] + (sum(heights[r] for r in groups[1]) if len(groups) > 1 else 0)
    restarted = first_limit + .001 < minimum_group
    if restarted: first_limit = max_height
    if total <= first_limit + .001 or (groups[0] != [0] and total <= max_height + .001):
        panels = [list(range(len(rows)))]
    else:
        if groups[0] != [0]:
            raise ValueError('表格首行纵向合并跨越续页，当前画幅不能完整重复列关系。')
        # Never emit a lone column guide, nor detach a vertical-merge chain.
        # An undersized remaining area advances a complete new-page panel,
        # avoiding a page containing only the minimum first data-row group.
        if len(groups) == 1:
            if total > max_height + .001:
                raise ValueError('完整表格行超过投影画幅，请加宽画幅或按课时调整；未缩小或删减文字。')
            panels = [list(range(len(rows)))]
        else:
            panels, pending, used = [], [0], heights[0]
            limit = first_limit
            for group in groups[1:]:
                height = sum(heights[r] for r in group)
                if heights[0] + height > max_height + .001:
                    raise ValueError('完整表格行超过投影画幅，请加宽画幅或按课时调整；未缩小或删减文字。')
                if used + height > limit + .001:
                    panels.append(pending); pending, used, limit = [0], heights[0], max_height
                pending.extend(group); used += height
            if len(pending) > 1:
                panels.append(pending)
    result = []
    for pi, indices in enumerate(panels):
        y_positions, cursor = {}, 0.0
        for ri in indices:
            y_positions[ri] = cursor; cursor += heights[ri]
        parts = []
        selected = [c for c in cells if c['row'] in y_positions]
        for source in selected:
            ri, ci = source['row'], source['column']
            x = column_positions[ci]; y = y_positions[ri]
            key = (ri + 1, ci + 1, source['end'] - ci)
            if source['cell']['vertical_merge'] != 'continue':
                chain = [c['row'] for c in selected if c['column'] == ci and c['end'] == source['end']
                    and c['merge_start'] == ri]
                height = sum(heights[r] for r in chain)
                parts.append({'kind': 'diagram_box', 'offset_x': x, 'offset_y': y,
                    'width': column_positions[source['end']] - column_positions[ci], 'height': height,
                    'fill': '#edf5fb' if ri == 0 else '#ffffff', 'stroke': '#5b7184',
                    'block_id': block_id, '_table_cell': key})
            offset = y + padding
            for chunk in measured[(ri, ci)]:
                parts.append({**chunk, 'offset_x': x + padding, 'offset_y': offset, '_table_cell': key})
                offset += chunk['height']
        if len(parts) > 30000:
            raise ValueError('完整表格超出排版容量，请按课时调整；未截断单元格。')
        result.append({'kind': 'diagram', 'width': width, 'height': cursor,
            'block_id': block_id, '_parts': parts, '_table_source_rows': [r + 1 for r in indices],
            '_table_repeated_first_row': pi > 0, '_table_restarted_page': pi == 0 and restarted})
    return result


def write_word_table(document, rows, *, cell_tokens=None, assets=None, image_data=None, image_limits=None):
    """Append an editable native table with the exact sealed grid and cells.

    Native merge properties are written directly: python-docx's merge() joins
    cell paragraphs, which would change the source text/occurrence sequence.
    Runs inherit the surrounding Word document's font. Original empty and
    omitted edge cells stay distinct; no placeholder or coordinate label is
    inserted into classroom text. The first source row repeats on Word pages.
    """
    columns, _cells = _grid(rows)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.table import _Cell
    table = document.add_table(rows=0, cols=columns)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        line = OxmlElement('w:' + side)
        for name, value in (('val', 'single'), ('sz', '6'), ('color', '5B7184')):
            line.set(qn('w:' + name), value)
        borders.append(line)
    table._tbl.tblPr.append(borders)
    for ri, source in enumerate(rows):
        row = table.add_row()
        for tc in list(row._tr.tc_lst):
            row._tr.remove(tc)
        properties = row._tr.get_or_add_trPr()
        for name in ('grid_before', 'grid_after'):
            if source[name]:
                prop = OxmlElement('w:gridBefore' if name == 'grid_before' else 'w:gridAfter')
                prop.set(qn('w:val'), str(source[name])); properties.append(prop)
        if ri == 0:
            properties.append(OxmlElement('w:tblHeader'))
        column = source['grid_before']
        for cell in source['cells']:
            tc = OxmlElement('w:tc'); props = OxmlElement('w:tcPr'); tc.append(props)
            width = OxmlElement('w:tcW'); width.set(qn('w:type'), 'dxa')
            width.set(qn('w:w'), str(sum(int(table._tbl.tblGrid.gridCol_lst[c].w.twips)
                for c in range(column, column + cell['column_span'])))); props.append(width)
            if cell['column_span'] != 1:
                span = OxmlElement('w:gridSpan'); span.set(qn('w:val'), str(cell['column_span'])); props.append(span)
            if cell['vertical_merge'] != 'none':
                merge = OxmlElement('w:vMerge'); merge.set(qn('w:val'), cell['vertical_merge']); props.append(merge)
            tc.append(OxmlElement('w:p')); row._tr.append(tc)
            target = _Cell(tc, table)
            target.text = cell['text']
            tokens = (cell_tokens or {}).get((ri, column))
            parts = list(_cell_parts(cell['text']))
            if tokens or any(script for _payload, script in parts):
                paragraph = target.paragraphs[0]; paragraph.clear()
                for token in tokens or [{'kind': 'text', 'text': cell['text']}]:
                    if token['kind'] == 'text':
                        for payload, script in _cell_parts(token['text']):
                            run = paragraph.add_run(payload)
                            if script == 'subscript': run.font.subscript = True
                            elif script == 'superscript': run.font.superscript = True
                    else:
                        from docx.shared import Mm
                        from .desktop_preparation_images import verify_image_bytes
                        from .desktop_word_picture import add_word_picture
                        from .desktop_composition_image_display import cap_image_size
                        identity = token['asset_id']
                        if identity not in (assets or {}) or identity not in (image_data or {}):
                            raise ValueError('表格单元格原图缺失，未生成省略图片的文档。')
                        asset = assets[identity]; data = verify_image_bytes(asset, image_data[identity])
                        w, h = token['width_mm'], token['height_mm']
                        if w == 'unknown' or h == 'unknown':
                            w, h = asset['width'] * 25.4 / 96, asset['height'] * 25.4 / 96
                        # Word's default cell margins are 108 twips each. The
                        # actual native grid determines usable width, including
                        # horizontal spans; image DPI never changes this fit.
                        available = max(.1, target.width.mm - 2 * 108 * 25.4 / 1440)
                        factor = min(1, available / w, 220 / h)
                        w, h = cap_image_size(w * factor, h * factor, (image_limits or {}).get(token['image_block_id']))
                        add_word_picture(paragraph.add_run(), data, asset, width=Mm(w), height=Mm(h))
            column += cell['column_span']
    return table
