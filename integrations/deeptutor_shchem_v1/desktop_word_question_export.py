"""Export selected, hash-bound Word body blocks without flattening chemistry.

This module is a pure bytes-to-bytes builder.  It never opens a source path,
follows a link, calls a model, or imports the original document's other parts
wholesale.  Its selection indices use the preparation preview's 1-based body
block traversal.  Unsupported dependencies fail before either result is returned.
"""

from __future__ import annotations

import re
from copy import deepcopy
from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from math import isfinite
from typing import Any
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .word_native_text import WordNativeTextReader

WORD_QUESTION_EXPORT_REVISION = "20260909-native-selected-blocks-v1"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
V = "urn:schemas-microsoft-com:vml"
O = "urn:schemas-microsoft-com:office:office"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
PIC = "http://schemas.openxmlformats.org/drawingml/2006/picture"
IMAGE_TYPES = {
    "image/png", "image/jpeg", "image/gif", "image/bmp", "image/tiff",
    "image/x-emf", "image/x-wmf", "image/emf", "image/wmf",
}
_UNSUPPORTED = {
    f"{{{W}}}{name}": label
    for name, label in {
        "altChunk": "外部正文块",
        "footnoteReference": "脚注", "endnoteReference": "尾注",
        "commentReference": "批注", "commentRangeStart": "批注",
        "commentRangeEnd": "批注", "fldSimple": "域代码",
        "fldChar": "域代码", "instrText": "域代码",
        "hyperlink": "超链接", "txbxContent": "文本框",
        "del": "修订", "ins": "修订", "moveFrom": "修订", "moveTo": "修订",
        "pPrChange": "修订", "rPrChange": "修订", "tblPrChange": "修订",
        "trPrChange": "修订", "tcPrChange": "修订",
    }.items()
}


class WordQuestionExportError(ValueError):
    def __init__(self, message: str):
        self.code = "word_question_export_invalid"
        self.message_zh = message
        super().__init__(message)


def _fail(message: str) -> None:
    raise WordQuestionExportError(message)


def _body_blocks(element):
    # Deliberately matches desktop_preparation_sources._body_blocks without
    # importing that service and its provider/import-stack dependencies.
    for child in element:
        name = child.tag.rsplit("}", 1)[-1]
        if name in {"p", "tbl", "altChunk"}:
            yield child
        elif name in {"sdt", "sdtContent", "customXml"}:
            yield from _body_blocks(child)
        elif name not in {"sectPr", "bookmarkStart", "bookmarkEnd", "proofErr"}:
            yield child


def _document(data: bytes):
    if not isinstance(data, bytes) or not data or len(data) > 256 * 1024 * 1024:
        _fail("Word 来源内容为空或过大，请重新导入。")
    try:
        with ZipFile(BytesIO(data)) as package:
            infos = package.infolist()
            if len(infos) > 12_000 or sum(i.file_size for i in infos) > 512 * 1024 * 1024:
                _fail("Word 解压内容过大，无法导出所选题目。")
            names = [i.filename for i in infos]
            if len(names) != len(set(names)) or "word/document.xml" not in names:
                _fail("Word 文件结构不完整，请重新导入。")
            if any(i.flag_bits & 1 for i in infos):
                _fail("Word 文件加密，无法导出所选题目。")
        return Document(BytesIO(data))
    except WordQuestionExportError:
        raise
    except (BadZipFile, KeyError, ValueError, TypeError, OSError) as exc:
        raise WordQuestionExportError("Word 原文件无法读取，请重新导入后再选题。") from exc


def _indices(value: Any, count: int, role: str) -> tuple[int, ...]:
    if not isinstance(value, list):
        _fail(f"{role}范围无效，请重新预览所选题目。")
    result = []
    for block in value:
        index = block.get("index") if isinstance(block, dict) else None
        if type(index) is not int or not 1 <= index <= count:
            _fail(f"{role}范围已失效，请重新预览所选题目。")
        if block.get("split_display") or block.get("display_only") or block.get("display_only_split"):
            _fail("题干和答案在同一原文段落中，请先修订范围，暂不能原生导出。")
        result.append(index)
    if result != sorted(set(result)):
        _fail(f"{role}范围重复或顺序不正确，请重新预览。")
    return tuple(result)


@dataclass
class _Source:
    digest: str
    document: Any
    blocks: list[Any]


@dataclass
class _Question:
    source: _Source
    source_name: str
    question: tuple[int, ...]
    answer: tuple[int, ...]
    context: tuple[int, ...]
    points: float
    nested_starts: tuple[int, ...] = ()


def _prepare(questions: list[dict]) -> list[_Question]:
    if not isinstance(questions, list) or not 1 <= len(questions) <= 200:
        _fail("请先选择 1 至 200 道题目。")
    sources: dict[str, _Source] = {}
    result = []
    identities = set()
    for item in questions:
        if not isinstance(item, dict) or item.get("export_ready") is not True:
            _fail("所选题目仍有题干答案边界问题，请先核对并修订范围。")
        identity = (item.get("key"), item.get("revision"))
        if not all(isinstance(v, str) and v for v in identity) or identity in identities:
            _fail("题目选择缺失版本或重复，请重新选择。")
        identities.add(identity)
        data, digest = item.get("source_bytes"), item.get("source_sha256")
        if not isinstance(data, bytes) or sha256(data).hexdigest() != digest:
            _fail("Word 原文件与预览版本不一致，请重新导入并选题。")
        if digest not in sources:
            document = _document(data)
            sources[digest] = _Source(digest, document, list(_body_blocks(document.element.body)))
        source = sources[digest]
        question = _indices(item.get("question_blocks"), len(source.blocks), "题干")
        answer = _indices(item.get("answer_blocks", []), len(source.blocks), "答案")
        context = _indices(item.get("context_blocks", []), len(source.blocks), "公共材料")
        if not question or set(question) & set(answer) or set(context) & (set(question) | set(answer)):
            _fail("题干、答案和公共材料的范围重叠或题干为空，请先修订范围。")
        # Inline answer labels must not be trusted merely because a caller set
        # export_ready.  The index/service normally stops this earlier.
        reader = WordNativeTextReader(source.document.styles.element)
        for index in (*context, *question):
            text = reader.read(source.blocks[index - 1]).text
            if re.search(r"【\s*(?:参考答案|答案|解析|分析|解答)\s*】|(?:^|\n)\s*(?:参考答案|答案|解析|解答)\s*[:：]", text):
                _fail(f"第 {index} 段题面仍含答案标记，请修订题干和答案范围后导出。")
        points = item.get("points", 2)
        if isinstance(points, bool) or not isinstance(points, (int, float)) or not isfinite(points) or not 0 < points <= 1000:
            _fail("每道题的本次分值须大于 0 且不超过 1000。")
        nested = item.get("nested_section_starts", [])
        if (not isinstance(nested, list) or any(type(v) is not int or v not in question for v in nested)
                or nested != sorted(set(nested))):
            _fail("题组的小题起点已失效，请重新核对选题范围。")
        result.append(_Question(source, str(item.get("source_name") or "Word 讲义"), question, answer, context, float(points), tuple(nested)))
    # A block recognized as an answer for any selected question may never
    # become another selected question's shared context/student content.
    answer_sets: dict[str, set[int]] = {}
    for item in result:
        answer_sets.setdefault(item.source.digest, set()).update(item.answer)
    for item in result:
        if (set(item.question) | set(item.context)) & answer_sets[item.source.digest]:
            _fail("所选题目的题面与另一道题的答案范围重叠，请先修订范围。")
    return result


def _new_document(title: str, teacher: bool):
    document = Document()
    document.core_properties.author = "沪上化学智研台"
    document.core_properties.title = title
    document.core_properties.subject = "逐题选编练习"
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.75)
    for name, size in (("Normal", 12), ("Title", 20), ("Subtitle", 12), ("Heading 1", 14)):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
        style.paragraph_format.space_after = Pt(6)
    document.styles["Normal"].paragraph_format.line_spacing = 1.15
    document.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    title_paragraphs = (
        document.add_paragraph(title, "Title"),
        document.add_paragraph("教师参考答案" if teacher else "学生练习", "Subtitle"),
    )
    # The bundled python-docx template can contain a blue bottom rule on
    # Title.  Remove decorations only from our generated heading block; source
    # question formatting (including colored printed numbers) remains intact.
    for name in ("Title", "Subtitle", "Title Char", "Subtitle Char"):
        if name not in document.styles:
            continue
        style = document.styles[name]
        for border in list(style.element.iter(qn("w:pBdr"))) + list(style.element.iter(qn("w:bdr"))):
            border.getparent().remove(border)
        style.font.underline = False
    for paragraph in title_paragraphs:
        for border in list(paragraph._p.iter(qn("w:pBdr"))) + list(paragraph._p.iter(qn("w:bdr"))):
            border.getparent().remove(border)
        for run in paragraph.runs:
            run.font.underline = False
    if teacher:
        document.add_paragraph("各题分值为本次练习设置。原文答案按所选范围附后，供备课和讲评使用。")
    return document


class _Writer:
    """One destination, with independent namespaces for every source package."""

    def __init__(self, document):
        self.document = document
        self.styles: dict[tuple[str, str], str] = {}
        self.nums: dict[tuple[str, str], str] = {}
        self.abstracts: dict[tuple[str, str], str] = {}
        self.sources: dict[str, str] = {}
        self.images: dict[tuple[str, str], str] = {}
        self.shape_counter = 0
        self.default_styles: dict[str, str] = {}
        self.warnings: list[str] = []

    def prefix(self, source: _Source) -> str:
        return self.sources.setdefault(source.digest, f"wq{len(self.sources) + 1}")

    def default_style(self, source: _Source, kind: str) -> str | None:
        for style in source.document.styles.element:
            if style.tag == qn("w:style") and style.get(qn("w:type")) == kind and style.get(qn("w:default")) in {"1", "true", "on"}:
                return style.get(qn("w:styleId"))
        return "Normal" if kind == "paragraph" else None

    def source_defaults(self, source: _Source) -> str:
        if source.digest in self.default_styles:
            return self.default_styles[source.digest]
        style_id = self.prefix(source) + "_Defaults"
        self.default_styles[source.digest] = style_id
        clone = OxmlElement("w:style")
        clone.set(qn("w:type"), "paragraph")
        clone.set(qn("w:styleId"), style_id)
        name = OxmlElement("w:name")
        name.set(qn("w:val"), style_id)
        clone.append(name)
        defaults = source.document.styles.element.find(qn("w:docDefaults"))
        if defaults is not None:
            for tag in ("pPr", "rPr"):
                holder = defaults.find(qn("w:" + tag + "Default"))
                properties = holder.find(qn("w:" + tag)) if holder is not None else None
                if properties is not None:
                    clone.append(deepcopy(properties))
        self._theme(clone, source)
        self.document.styles.element.append(clone)
        return style_id

    def style(self, source: _Source, original: str) -> str:
        key = (source.digest, original)
        if key in self.styles:
            return self.styles[key]
        found = next((s for s in source.document.styles.element if s.tag == qn("w:style") and s.get(qn("w:styleId")) == original), None)
        if found is None:
            _fail(f"原 Word 缺少样式 {original}，无法可靠保留题面。")
        mapped = f"{self.prefix(source)}_s{len(self.styles) + 1}"
        self.styles[key] = mapped
        clone = deepcopy(found)
        clone.set(qn("w:styleId"), mapped)
        clone.attrib.pop(qn("w:default"), None)
        name = clone.find(qn("w:name"))
        if name is not None:
            name.set(qn("w:val"), mapped)
        for node in list(clone):
            if node.tag in {qn("w:basedOn"), qn("w:next"), qn("w:link")}:
                node.set(qn("w:val"), self.style(source, node.get(qn("w:val"), "")))
        if clone.get(qn("w:type")) == "paragraph" and clone.find(qn("w:basedOn")) is None:
            base = OxmlElement("w:basedOn")
            base.set(qn("w:val"), self.source_defaults(source))
            clone.insert(1, base)
        self._references(clone, source)
        self._theme(clone, source)
        self.document.styles.element.append(clone)
        return mapped

    def numbering(self, source: _Source, original: str) -> str:
        if original == "0":
            return "0"
        key = (source.digest, original)
        if key in self.nums:
            return self.nums[key]
        try:
            origin = source.document.part.numbering_part.element
        except KeyError:
            _fail("原 Word 的列表编号定义缺失，请先在原文件中修复编号。")
        found = next((n for n in origin if n.tag == qn("w:num") and n.get(qn("w:numId")) == original), None)
        if found is None:
            _fail("原 Word 的列表编号定义缺失，请先在原文件中修复编号。")
        destination = self.document.part.numbering_part.element
        new_id = str(max((int(n.get(qn("w:numId"))) for n in destination if n.tag == qn("w:num")), default=0) + 1)
        self.nums[key] = new_id
        clone = deepcopy(found)
        clone.set(qn("w:numId"), new_id)
        destination.append(clone)
        ref = clone.find(qn("w:abstractNumId"))
        if ref is None:
            _fail("原 Word 的列表编号定义不完整。")
        old_abstract = ref.get(qn("w:val"), "")
        abstract_key = (source.digest, old_abstract)
        if abstract_key not in self.abstracts:
            abstract = next((n for n in origin if n.tag == qn("w:abstractNum") and n.get(qn("w:abstractNumId")) == old_abstract), None)
            if abstract is None or any(n.tag == qn("w:lvlPicBulletId") for n in abstract.iter()):
                _fail("题面使用缺失编号或图片项目符号，暂不能可靠导出。")
            mapped = str(max((int(n.get(qn("w:abstractNumId"))) for n in destination if n.tag == qn("w:abstractNum")), default=-1) + 1)
            self.abstracts[abstract_key] = mapped
            new_abstract = deepcopy(abstract)
            new_abstract.set(qn("w:abstractNumId"), mapped)
            # Reserve in the tree before following style links, which may
            # themselves resolve another numbering definition recursively.
            first_num = next((n for n in destination if n.tag == qn("w:num")), None)
            if first_num is None:
                destination.append(new_abstract)
            else:
                destination.insert(destination.index(first_num), new_abstract)
            self._references(new_abstract, source)
            self._theme(new_abstract, source)
        ref.set(qn("w:val"), self.abstracts[abstract_key])
        self._references(clone, source)
        return new_id

    def _references(self, element, source: _Source) -> None:
        for node in element.iter():
            if node.tag in {qn("w:pStyle"), qn("w:rStyle"), qn("w:tblStyle"), qn("w:styleLink"), qn("w:numStyleLink")}:
                node.set(qn("w:val"), self.style(source, node.get(qn("w:val"), "")))
            elif node.tag == qn("w:numId"):
                node.set(qn("w:val"), self.numbering(source, node.get(qn("w:val"), "")))

    def _theme(self, element, source: _Source) -> None:
        # Theme references point to a single package-level theme; resolve the
        # source's fonts/colors before mixing documents with different themes.
        from lxml import etree

        rel = next((r for r in source.document.part.rels.values() if r.reltype == RT.THEME and not r.is_external), None)
        if rel is None:
            return
        try:
            theme = etree.fromstring(rel.target_part.blob)
        except etree.XMLSyntaxError as exc:
            raise WordQuestionExportError("原 Word 主题定义损坏，无法保留字体颜色。") from exc
        scheme = theme.find(f".//{{{A}}}fontScheme")
        colors = theme.find(f".//{{{A}}}clrScheme")
        lang = source.document.settings.element.find(qn("w:themeFontLang"))
        east_language = lang.get(qn("w:eastAsia"), "") if lang is not None else ""
        script = {"zh-CN": "Hans", "zh-SG": "Hans", "zh-TW": "Hant", "zh-HK": "Hant", "ja-JP": "Jpan", "ko-KR": "Hang"}.get(east_language)
        for node in element.iter():
            if node.tag == qn("w:rFonts") and scheme is not None:
                for attr, target in (("asciiTheme", "ascii"), ("hAnsiTheme", "hAnsi"), ("eastAsiaTheme", "eastAsia"), ("cstheme", "cs")):
                    value = node.get(qn("w:" + attr))
                    if value is None:
                        continue
                    family = scheme.find(f"{{{A}}}{'majorFont' if value.startswith('major') else 'minorFont'}")
                    leaf = family.find(f"{{{A}}}{'ea' if target == 'eastAsia' else 'cs' if target == 'cs' else 'latin'}") if family is not None else None
                    font = leaf.get("typeface", "") if leaf is not None else ""
                    if not font and target == "eastAsia" and script and family is not None:
                        font = next((n.get("typeface", "") for n in family if n.get("script") == script), "")
                    if font:
                        node.set(qn("w:" + target), font)
                        node.attrib.pop(qn("w:" + attr), None)
                    elif node.get(qn("w:" + target)):
                        node.attrib.pop(qn("w:" + attr), None)
            if colors is not None:
                for attr, target in (("themeColor", "val"), ("themeFill", "fill")):
                    name = node.get(qn("w:" + attr))
                    if name is None:
                        continue
                    name = {"dark1": "dk1", "dark2": "dk2", "light1": "lt1", "light2": "lt2", "text1": "dk1", "text2": "dk2", "background1": "lt1", "background2": "lt2"}.get(name, name)
                    color = colors.find(f"{{{A}}}{name}")
                    if color is None or not len(color):
                        _fail("原 Word 颜色主题不完整，无法保留题面颜色。")
                    value = color[0].get("lastClr") or color[0].get("val")
                    try:
                        channels = [int(value[i:i + 2], 16) for i in (0, 2, 4)]
                        prefix = "theme" if attr == "themeColor" else "themeFill"
                        shade, tint = node.get(qn("w:" + prefix + "Shade")), node.get(qn("w:" + prefix + "Tint"))
                        if shade:
                            channels = [round(c * int(shade, 16) / 255) for c in channels]
                        if tint:
                            channels = [round(c + (255 - c) * (1 - int(tint, 16) / 255)) for c in channels]
                        node.set(qn("w:" + target), "".join(f"{c:02X}" for c in channels))
                    except (ValueError, TypeError) as exc:
                        raise WordQuestionExportError("原 Word 颜色主题不完整。") from exc
                    for suffix in ("", "Shade", "Tint"):
                        node.attrib.pop(qn("w:" + (attr if not suffix else prefix + suffix)), None)

    def _image(self, source: _Source, rid: str) -> str:
        relationship = source.document.part.rels.get(rid)
        if relationship is None or relationship.is_external or relationship.reltype != RT.IMAGE:
            _fail("题面含外部图片或未支持的对象关系，请改用内嵌图片后重新导入。")
        part = relationship.target_part
        if part.content_type not in IMAGE_TYPES or part.rels:
            _fail("题面含暂不支持的图片格式或关联对象，请先在原 Word 中转为内嵌 PNG 图片。")
        key = (sha256(part.blob).hexdigest(), part.content_type)
        if key not in self.images:
            extension = part.partname.ext
            if not re.fullmatch(r"[A-Za-z0-9]{1,8}", extension):
                _fail("题面图片名称无效，请重新插入图片后导入。")
            name = PackURI(f"/word/media/wq_image{len(self.images) + 1}.{extension}")
            image_part = Part(name, part.content_type, part.blob, self.document.part.package)
            self.images[key] = self.document.part.relate_to(image_part, RT.IMAGE)
        return self.images[key]

    def append(self, source: _Source, indices: tuple[int, ...], *, keep_question: bool = False, numbering=None, answer: bool = False, references=None) -> None:
        reader = WordNativeTextReader(source.document.styles.element)
        for offset, index in enumerate(indices):
            original = source.blocks[index - 1]
            if original.tag not in {qn("w:p"), qn("w:tbl")}:
                _fail(f"第 {index} 段使用暂不支持的正文容器，请调整选题范围。")
            features = reader.read(original).features
            if "hidden_text_omitted" in features or "tracked_changes" in features:
                _fail(f"第 {index} 段含隐藏文字或修订，请在原 Word 中核对后重新导入。")
            clone = deepcopy(original)
            # Section settings/header relationships belong to the original
            # full paper, not this selected exercise. Keep the output layout.
            for node in list(clone.iter()):
                if node.tag == qn("w:sectPr"):
                    node.getparent().remove(node)
            self._ole_fallback(clone, source, index)
            self._validate(clone, index)
            for node in clone.iter():
                for attr, value in list(node.attrib.items()):
                    if attr.startswith("{" + R + "}"):
                        node.set(attr, self._image(source, value))
                if node.tag == qn("w:p") and node.find("./" + qn("w:pPr") + "/" + qn("w:pStyle")) is None:
                    style = OxmlElement("w:pStyle")
                    style.set(qn("w:val"), self.default_style(source, "paragraph"))
                    node.get_or_add_pPr().insert(0, style)
                if node.tag == qn("w:tbl") and node.find("./" + qn("w:tblPr") + "/" + qn("w:tblStyle")) is None:
                    default = self.default_style(source, "table")
                    if default:
                        style = OxmlElement("w:tblStyle")
                        style.set(qn("w:val"), default)
                        node.tblPr.insert(0, style)
            if numbering is not None:
                from .desktop_paper_numbering import apply_word_numbering
                apply_word_numbering(clone, index, numbering, answer=answer, references=references)
            self._references(clone, source)
            self._theme(clone, source)
            self._shapes(clone, source)
            for node in list(clone.iter()):
                # Unused bookmarks are nonvisual and may be only half selected;
                # references were rejected above, so avoid duplicate identities.
                if node.tag in {qn("w:bookmarkStart"), qn("w:bookmarkEnd")}:
                    node.getparent().remove(node)
            if keep_question:
                self._keep_question_block(clone, last=offset == len(indices) - 1)
            self.document.element.body.insert(-1, clone)

    @staticmethod
    def _keep_question_block(block, *, last: bool) -> None:
        """Keep a short question's paragraphs/options together, not its answer.

        keepNext is a pagination preference, not a fixed-size box or manual
        page break: Word can still paginate questions longer than a page.
        Tables need every cell linked to subsequent rows, with each terminal
        cell's final paragraph released when the table ends the question.
        """
        paragraphs = list(block.iter(qn("w:p")))
        for paragraph in paragraphs:
            properties = paragraph.get_or_add_pPr()
            properties.get_or_add_keepNext().val = True
            properties.get_or_add_keepLines().val = True
        if not last:
            return
        if block.tag == qn("w:p"):
            block.get_or_add_pPr().get_or_add_keepNext().val = False
            return
        rows = list(block.iter(qn("w:tr")))
        if not rows:
            return
        # Only rows owned by this table determine its terminal row; a nested
        # table in an earlier cell must not terminate the outer-table chain.
        outer_rows = [row for row in rows if next((ancestor for ancestor in row.iterancestors() if ancestor.tag == qn("w:tbl")), None) is block]
        if not outer_rows:
            return
        for cell in outer_rows[-1]:
            if cell.tag != qn("w:tc"):
                continue
            cell_paragraphs = list(cell.iter(qn("w:p")))
            if cell_paragraphs:
                cell_paragraphs[-1].get_or_add_pPr().get_or_add_keepNext().val = False

    def _ole_fallback(self, clone, source: _Source, index: int) -> None:
        """Retain a linked, visible Content fallback, never an executable OLE part.

        An icon is not an equation/diagram preview.  Ambiguous or incomplete
        fallbacks remain a hard error, including OLE outside a Word object.
        """
        for obj in list(clone.iter(qn("w:object"))):
            objects = list(obj.iter(f"{{{O}}}OLEObject"))
            shapes = list(obj.iter(f"{{{V}}}shape"))
            if len(objects) != 1 or len(shapes) != 1:
                _fail(f"第 {index} 段旧版对象缺少完整静态预览，请先转为图片。")
            ole, shape = objects[0], shapes[0]
            images = list(shape.iter(f"{{{V}}}imagedata"))
            if (ole.get("DrawAspect") != "Content" or ole.get("ShapeID") != shape.get("id")
                    or len(images) != 1 or list(shape.iter(qn("w:txbxContent")))):
                _fail(f"第 {index} 段旧版对象只有图标或预览不完整，请先转为图片。")
            style = shape.get("style", "")
            if not all(re.search(rf"(?:^|;)\s*{dimension}\s*:\s*(?:[1-9]\d*(?:\.\d+)?|0\.\d*[1-9]\d*)(?:pt|in|cm|mm|px)(?:;|$)", style) for dimension in ("width", "height")):
                _fail(f"第 {index} 段旧版对象预览尺寸无效，请先转为图片。")
            rid = images[0].get(f"{{{R}}}id", "")
            # This validates type, bytes and internal relationship.  append()
            # subsequently remaps the source id in the surviving image node.
            self._image(source, rid)
            ole.getparent().remove(ole)
            obj.tag = qn("w:pict")
            obj.attrib.clear()
            warning = f"{source.digest[:12]} 第 {index} 段旧版对象已保留为原文件自带的静态预览图；不保留 OLE 编辑能力。"
            if warning not in self.warnings:
                self.warnings.append(warning)

    def _validate(self, clone, index: int) -> None:
        for node in clone.iter():
            label = _UNSUPPORTED.get(node.tag)
            if label or node.tag == f"{{{O}}}OLEObject":
                _fail(f"第 {index} 段含{label or '旧版 OLE 对象'}，暂不能可靠原生导出；请在原 Word 中转换或修订范围。")
            if node.tag == f"{{{A}}}graphicData" and node.get("uri") != PIC:
                _fail(f"第 {index} 段含非图片图表对象，暂不能可靠导出。")
            if node.tag == f"{{{A}}}schemeClr":
                _fail(f"第 {index} 段图片依赖主题颜色，请在原 Word 中转为静态图片后导入。")
            if node.tag in {qn("w:vanish"), qn("w:webHidden"), qn("w:specVanish")} and node.get(qn("w:val"), "1") not in {"0", "false", "off"}:
                _fail(f"第 {index} 段含隐藏文字，请先在原 Word 中核对。")
            if node.tag == f"{{{V}}}imagedata" and not node.get(f"{{{R}}}id"):
                _fail(f"第 {index} 段图片未内嵌到 Word，请重新插入图片。")

    def _shapes(self, clone, source: _Source) -> None:
        # Other renderers may append pictures between native Word blocks.
        # Synchronize on every append, not just when this writer is created.
        for node in self.document.element.body.iter():
            if node.tag in {f"{{{WP}}}docPr", f"{{{PIC}}}cNvPr"}:
                value = node.get("id", "")
                if value.isdecimal():
                    self.shape_counter = max(self.shape_counter, int(value))
        for node in clone.iter():
            if node.tag in {f"{{{WP}}}docPr", f"{{{PIC}}}cNvPr"}:
                self.shape_counter += 1
                node.set("id", str(self.shape_counter))
            if node.tag in {f"{{{V}}}{name}" for name in (
                "shape", "rect", "roundrect", "oval", "line", "polyline",
                "arc", "curve", "image", "group",
            )}:
                self.shape_counter += 1
                node.set("id", f"wq_shape_{self.shape_counter}")
                node.attrib.pop(f"{{{O}}}spid", None)
        # VML shapetype definitions may live in an unselected earlier block.
        # Carry only definitions referenced by selected static image shapes.
        definitions = {n.get("id"): n for n in source.document.element.body.iter(f"{{{V}}}shapetype")}
        for shape in list(clone.iter(f"{{{V}}}shape")):
            ref = shape.get("type", "")
            if not ref.startswith("#"):
                continue
            source_type = definitions.get(ref[1:])
            if source_type is None:
                # Some WPS/Word documents omit the standard picture-frame
                # shapetype 75 but keep the complete embedded image/geometry.
                # A VML rect with imagedata is an explicitly supported image
                # container; materialize that rectangle without fabricating
                # any image bytes or dropping crop, transform or dimensions.
                # https://learn.microsoft.com/en-us/windows/win32/vml/msdn-online-vml-imagedata-element
                images = list(shape.iter(f"{{{V}}}imagedata"))
                custom_geometry = shape.get("path") or any(
                    n.tag in {f"{{{V}}}formulas", qn("w:txbxContent")}
                    or (n.tag == f"{{{V}}}path" and n.get("v"))
                    for n in shape.iter()
                )
                if (ref == "#_x0000_t75" and shape.get(f"{{{O}}}spt", "75") == "75"
                        and len(images) == 1 and not custom_geometry):
                    shape.tag = f"{{{V}}}rect"
                    shape.attrib.pop("type", None)
                    shape.set("stroked", shape.get("stroked", "f"))
                    shape.set("filled", shape.get("filled", "f"))
                    continue
                _fail("原 Word 的 VML 图片形状定义缺失，请重新插入图片后导入。")
            self.shape_counter += 1
            target = f"wq_type_{self.shape_counter}"
            definition = deepcopy(source_type)
            definition.set("id", target)
            if any(attr.startswith("{" + R + "}") for n in definition.iter() for attr in n.attrib):
                _fail("VML 图片形状含关联资源，暂不能可靠导出。")
            shape.getparent().insert(shape.getparent().index(shape), definition)
            shape.set("type", "#" + target)
        for definition in list(clone.iter(f"{{{V}}}shapetype")):
            if not definition.get("id", "").startswith("wq_type_"):
                definition.getparent().remove(definition)


def export_word_questions(title: str, questions: list[dict], *, show_student_scores: bool = False) -> dict[str, Any]:
    """Return both native DOCX outputs, or fail without returning partial data.

    ``points`` means the teacher's allocation for this exercise, not a claim
    about the source paper. Original printed scores/answer space are preserved.
    """
    if not isinstance(title, str) or not title.strip() or len(title.strip()) > 160:
        _fail("请填写 1 至 160 字的练习名称。")
    if type(show_student_scores) is not bool:
        _fail("题面分值开关无效，请重新选择。")
    items = _prepare(questions)
    from .desktop_paper_numbering import plan_word, word_reference_maps
    plans, cursor = [], 1
    for item in items:
        plan = plan_word(item, cursor)
        plans.append(plan)
        cursor += plan.count
    refs = word_reference_maps(items, plans)
    outputs = {}
    warnings = []
    for role, teacher in (("student", False), ("teacher", True)):
        document = _new_document(title.strip(), teacher)
        writer = _Writer(document)
        included_context: set[tuple[str, int]] = set()
        for index, item in enumerate(items, 1):
            context = tuple(n for n in item.context if (item.source.digest, n) not in included_context)
            if context:
                document.add_paragraph("公共材料", "Heading 1")
                writer.append(item.source, context)
                included_context.update((item.source.digest, n) for n in context)
            score = f"  {item.points:g} 分" if teacher or show_student_scores else ""
            document.add_paragraph(f"练习 {index}{score}", "Heading 1")
            writer.append(item.source, item.question, keep_question=True, numbering=plans[index - 1], references=refs[item.source.digest])
            if teacher:
                reference = document.add_paragraph(f"来源：{item.source_name}")
                for run in reference.runs:
                    run.font.size = Pt(10)
                if item.answer:
                    document.add_paragraph("参考答案", "Heading 1")
                    writer.append(item.source, item.answer, numbering=plans[index - 1], answer=True, references=refs[item.source.digest])
                else:
                    document.add_paragraph("当前选定范围未识别到本题答案，请核对原教案与题答边界；这不表示原文没有答案。")
        output = BytesIO()
        document.save(output)
        outputs[role + "_bytes"] = output.getvalue()
        warnings.extend(writer.warnings)
    outputs["warnings"] = list(dict.fromkeys(warnings))
    return outputs
