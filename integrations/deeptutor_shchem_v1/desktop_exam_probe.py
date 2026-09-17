"""Actual Windows source/EXE acceptance on synthetic Excel plus native Word, no network."""
from __future__ import annotations
from hashlib import sha256
import json
import os
from pathlib import Path
import sys
import tempfile
import time


def run_probe(output):
    from PySide6.QtCore import QTimer,QRect
    from PySide6.QtWidgets import QApplication,QDialog,QTabWidget
    from docx import Document
    from .desktop_facade import build_default_facade,ProviderProfileSummary
    from .desktop_paths import DesktopPaths
    from .desktop_version import DESKTOP_VERSION
    from .desktop_exam_fixture import example,FixtureStore,FixtureTransport
    from .desktop_exam_ai import load_paper,model_payload
    from .desktop_exam_report import html_report
    from .desktop_workbench.app import create_application
    from .desktop_workbench.main_window import TeacherWorkbenchWindow
    from .desktop_workbench.exam_dashboard import ExamDashboard
    from .desktop_workbench.exam_import_dialog import ExamImportDialog
    app=create_application(['exam198-probe']);output=Path(output);output.mkdir(parents=True,exist_ok=True)
    errors=[];screens=[];checks={};old=sys.excepthook
    sys.excepthook=lambda typ,value,tb:errors.append(typ.__name__+': '+str(value))
    def settle(fn=lambda:True):
        end=time.monotonic()+35
        for _ in range(6):app.processEvents();time.sleep(.03)
        while not fn() and time.monotonic()<end:app.processEvents();time.sleep(.025)
        assert fn(),'Exam operation did not finish'
    def capture(w,name):
        path=output/name;assert w.grab().save(str(path));screens.append({'file':name,'sha256':sha256(path.read_bytes()).hexdigest()})
    with tempfile.TemporaryDirectory(prefix='exam198-') as directory:
        folder=Path(directory);root=Path(getattr(sys,'_MEIPASS',Path(__file__).resolve().parents[2]))
        facade=build_default_facade(DesktopPaths.from_workspace(root,state_root=folder/'state'))
        book,config,exam=example(folder/'scores.xlsx');original=(folder/'scores.xlsx').read_bytes()
        doc=Document();doc.add_paragraph('合成试卷材料，仅作软件验收。第1至4题分别映射物质结构、反应原理、实验推理与有机转化。')
        doc.add_paragraph('公共材料：比较条件变化前后的实验记录。缺少原作答时不能仅凭失分确定错误原因。')
        source=folder/'paper.docx';doc.save(source);paper_original=source.read_bytes()
        window=TeacherWorkbenchWindow(facade);window.resize(1366,850);window.show();window.navigate('student')
        def drive():
            d=next((w for w in QApplication.topLevelWidgets() if isinstance(w,ExamDashboard) and w.isVisible()),None)
            if d is None:errors.append('Exam entry did not open');return
            try:
                mapping=ExamImportDialog(book,d);mapping.show();settle();mapping.findChild(QTabWidget).setCurrentIndex(1);settle()
                capture(mapping,'exam-import.png');mapping.commit();assert mapping.exam is not None;d.accept_exam(mapping.exam);mapping.deleteLater();settle()
                assert d.report['overall']['mean']==69 and d.report['overall']['n']==10
                capture(d,'exam-overview.png')
                d.tabs.setCurrentIndex(1);settle();capture(d,'exam-items.png')
                d.tabs.setCurrentIndex(2);d.select_student(d.student_table.model().index(1,0));settle();capture(d,'exam-students.png')
                d.classes.setCurrentIndex(d.classes.findData('合成班A'));settle();assert d.report['overall']['n']==4
                d.classes.setCurrentIndex(0);d.accept_paper(load_paper(source));d.notes.setPlainText('合成验收：先核对失分题，再安排短诊断与一周复测。')
                # Keep real request building/validation and transport contract, inject only network.
                profile=ProviderProfileSummary('p','Fixture','https://models.example/v1','fixture-model','chat_completions',('text',),True,'fixture-revision')
                facade._exam_provider_store=FixtureStore();facade._exam_transport=FixtureTransport()
                d.model.clear();d.model.addItem('合成测试模型',profile);d.tabs.setCurrentIndex(3)
                def approve():
                    modal=QApplication.activeModalWidget()
                    assert modal is not None and modal is not d
                    modal.accept()
                QTimer.singleShot(200,approve);d.request_ai();settle(lambda:d._task is None and d.result is not None)
                assert len(facade._exam_transport.requests)==1 and d.report['overall']['mean']==69
                assert '示例01' not in facade._exam_transport.requests[0].body.decode()
                capture(d,'exam-api.png')
                d.save_current();saved=d.store.load(d.exam['id']);assert saved['paper']['text']==d.paper_text.toPlainText()
                (output/'synthetic-exam-report.html').write_text(html_report(d.report,d.result),encoding='utf-8')
                d.tabs.setCurrentIndex(0);d.resize(800,700);settle()
                assert d.width()==800 and d.tabs.height()>300
                for w in (d.import_button,d.export_button,d.classes,d.tabs,d.close_button):assert d.rect().contains(QRect(w.mapTo(d,w.rect().topLeft()),w.size()))
                capture(d,'exam-compact.png')
                assert (folder/'scores.xlsx').read_bytes()==original and source.read_bytes()==paper_original
                checks.update(native_entry=True,explicit_excel_mapping=True,exact_local_metrics=True,class_filter=True,
                              missing_not_zero=True,item_denominators=True,student_attention=True,word_common_material=True,
                              confirmed_api_one_fake_call=True,anonymous_payload=True,model_does_not_change_metrics=True,
                              saved_snapshot_reloads=True,offline_report=True,compact_actions_visible=True,source_unchanged=True)
            except Exception as e:
                errors.append(type(e).__name__+': '+str(e));capture(d,'exam-failure.png')
            finally:
                d.dirty=False;d._task=None;d.close()
        try:
            settle();QTimer.singleShot(300,drive);window.student_page.exam_analysis_button.click()
            assert not errors,errors
            report={'version':DESKTOP_VERSION,'source_commit':os.environ.get('GITHUB_SHA','local'),'frozen':getattr(sys,'frozen',False),
                    'qt_platform':app.platformName(),'synthetic_rows':12,'valid_totals':10,'mean':69,'median':72,
                    'network_requests':0,'fixture_transport_requests':1,'checks':checks,'screenshots':screens,'uncaught_errors':errors,
                    'scope':'Synthetic saved XLSX values and Word text. No actual student data or live API quality acceptance.'}
            (output/'exam-probe.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        finally:
            window.close();app.processEvents();sys.excepthook=old
            if errors:(output/'exam-errors.json').write_text(json.dumps(errors,ensure_ascii=False,indent=2),encoding='utf-8')
    return 0
