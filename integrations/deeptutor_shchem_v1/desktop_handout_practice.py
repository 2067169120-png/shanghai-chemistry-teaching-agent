"""Editable practice from hash-bound native Word paragraphs, not a mock paper."""

from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from .desktop_handout_candidates import HandoutCandidateError, HandoutCandidateService
from .desktop_state import utc_now

EXPORT_KIND = "native_handout_practice_export"
DRAFT_ID = "HANDOUT-PRACTICE-SELECTION"
FILES = {"student": "讲义练习-学生版.docx", "teacher": "讲义练习-教师参考答案.docx"}
SOURCE_ROOT = "sh-chem-db/.intake/2026-07-30-user-teaching-pack"
_LOCATOR = re.compile(r"word/document\.xml#/w:document/w:body/w:p\[([1-9][0-9]*)\]")
_TABLE_LOCATOR = re.compile(
    r"word/document\.xml#/w:document/w:body/w:tbl\[([1-9][0-9]*)\]/w:tr\[[1-9][0-9]*\]/w:tc\[[1-9][0-9]*\]/w:p\[[1-9][0-9]*\]"
)
_EXPORT_ID = re.compile(r"HANDOUT-PRACTICE-[a-f0-9]{32}")


def _error(code: str, message: str) -> HandoutCandidateError:
    return HandoutCandidateError("handout_practice_" + code, message)


def _normalized(text: str) -> str:
    return re.sub(r"\s+", "", text)


class NativeParagraphs:
    """Resolve only documented paragraph locators in the exact source DOCX bytes."""

    def __init__(self, workspace: Path):
        self.root = (workspace / SOURCE_ROOT / "expanded").resolve()
        self.documents: dict[str, Any] = {}

    def _document(self, sha: str, question_path: str) -> Any:
        if sha in self.documents:
            return self.documents[sha]
        source = (self.root.parent / question_path).resolve()
        source.relative_to(self.root)
        # The answer path is not guessed from its filename: match the recorded hash
        # among DOCXs in this question's package directory.
        candidates = [source, *sorted(source.parent.glob("*.docx"))]
        for path in dict.fromkeys(candidates):
            path.resolve().relative_to(self.root)
            data = path.read_bytes()
            if hashlib.sha256(data).hexdigest() == sha:
                document = Document(io.BytesIO(data))
                self.documents[sha] = document
                return document
        raise _error(
            "source_changed", "原始 Word 缺失或哈希已变化，请刷新并核对来源文件。"
        )

    def read(self, item: dict[str, Any], role: str) -> list[Any]:
        binding = item.get("editable_source") or {}
        locators = binding.get(role + "_locators")
        if not isinstance(locators, list) or not locators:
            raise _error(
                "locators_missing", "此题尚缺原生 Word 段落定位，暂不能保真导出。"
            )
        sha = (
            item["source_document"]["sha256"]
            if role == "question"
            else binding.get("answer_source_sha256")
        )
        try:
            document = self._document(sha, item["source_document"]["relative_path"])
            paragraphs = []
            positions = []
            blocks = []
            tables = {}
            ordering = {
                element: index
                for index, element in enumerate(document._element.iter(qn("w:p")))
            }
            for locator in locators:
                match = _LOCATOR.fullmatch(locator)
                table_match = _TABLE_LOCATOR.fullmatch(locator)
                if match is None and table_match is None:
                    raise ValueError("unsupported locator")
                xpath = locator.split("#/w:document", 1)[1]
                elements = document._element.xpath("." + xpath)
                if len(elements) != 1:
                    raise ValueError("missing paragraph")
                paragraph = Paragraph(elements[0], document._body)
                positions.append(ordering[elements[0]])
                # Do not flatten unknown objects, hyperlinks, hidden text or equations.
                for node in paragraph._p:
                    if node.tag not in {
                        qn("w:pPr"),
                        qn("w:r"),
                        qn("w:bookmarkStart"),
                        qn("w:bookmarkEnd"),
                        qn("w:proofErr"),
                    }:
                        raise ValueError("unsupported paragraph content")
                for run in paragraph.runs:
                    if run.font.hidden:
                        raise ValueError("hidden run")
                    for node in run._r:
                        if node.tag not in {
                            qn("w:rPr"),
                            qn("w:t"),
                            qn("w:tab"),
                            qn("w:br"),
                            qn("w:cr"),
                            qn("w:lastRenderedPageBreak"),
                        }:
                            raise ValueError("non-text object")
                paragraphs.append(paragraph)
                if table_match:
                    table_index = int(table_match[1])
                    if table_index not in tables:
                        element = document._element.xpath(
                            f"./w:body/w:tbl[{table_index}]"
                        )[0]
                        tables[table_index] = {"element": element, "selected": []}
                        blocks.append(Table(element, document._body))
                    tables[table_index]["selected"].append(elements[0])
                else:
                    blocks.append(paragraph)
            if not _LOCATOR.fullmatch(locators[0]):
                raise ValueError(
                    "candidate starts inside a table; question identity needs review"
                )
            for table in tables.values():
                if list(table["element"].iter(qn("w:p"))) != table["selected"]:
                    raise ValueError(
                        "partial table cannot be exported as a complete question"
                    )
            if positions != sorted(set(positions)):
                raise ValueError("duplicate or reordered locator")
            expected = (
                item["question_text"] if role == "question" else item["answer_text"]
            )
            if role == "answer" and expected.startswith("【非官方"):
                expected = expected.split("\n", 1)[1].strip()
            if _normalized("\n".join(p.text for p in paragraphs)) != _normalized(
                expected
            ):
                raise ValueError("native text differs")
            return blocks
        except HandoutCandidateError:
            raise
        except (OSError, ValueError, KeyError, IndexError, TypeError) as exc:
            raise _error(
                "native_mismatch",
                "原生段落与候选文字不一致，或含尚未支持的对象；请先查看原页。",
            ) from exc


def _font(style: Any, size: int) -> None:
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    for element in style.element.xpath("./w:pPr/w:pBdr"):
        element.getparent().remove(element)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + attr), None)


def _new_document(title: str, teacher: bool) -> Any:
    document = Document()
    document.core_properties.author = "沪上化学智研台"
    document.core_properties.title = title
    document.core_properties.subject = "讲义选题练习"
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.75)
    for name, size in (
        ("Normal", 12),
        ("Title", 20),
        ("Subtitle", 12),
        ("Heading 1", 14),
        ("Heading 2", 12),
    ):
        _font(document.styles[name], size)
    normal = document.styles["Normal"].paragraph_format
    normal.line_spacing = 1.15 if teacher else 1.25
    normal.space_after = Pt(6)
    for name, before, after in (
        ("Title", 0, 6),
        ("Subtitle", 0, 8),
        ("Heading 1", 14, 6),
        ("Heading 2", 12, 6),
    ):
        fmt = document.styles[name].paragraph_format
        fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    document.styles["Subtitle"].font.italic = False
    document.add_paragraph(title, "Title")
    document.add_paragraph("教师参考答案" if teacher else "学生练习", "Subtitle")
    document.add_paragraph(
        "按练习编号与原题号对应讲评。以下为原讲义配对的非官方参考答案，尚未独立验证化学正确性。"
        if teacher
        else "请按原题要求作答，必要时在题后空白处写出推理或计算过程。"
    )
    if not teacher:
        document.add_paragraph(
            "姓名 ____________________    班级 __________    日期 __________"
        )
    # A page number is useful when the teacher prints multiple sheets.
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return document


def _append_native(document: Any, paragraphs: list[Any]) -> None:
    for source in paragraphs:
        if isinstance(source, Table):
            _append_table(document, source)
            continue
        target = document.add_paragraph()
        target.paragraph_format.widow_control = True
        for run in source.runs:
            result = target.add_run(run.text)
            # Preserve chemistry's superscript/subscript and original emphasis;
            # use the practice's body font/size for readable, consistent printing.
            if run._r.rPr is not None:
                for tag in ("b", "i", "u", "vertAlign", "strike", "dstrike"):
                    element = run._r.rPr.find(qn("w:" + tag))
                    if element is not None:
                        result._r.get_or_add_rPr().append(deepcopy(element))


def _append_table(document: Any, source: Table) -> None:
    clone = deepcopy(source._tbl)
    document._body._body.insert(-1, clone)
    table = Table(clone, document._body)
    table.style = "Table Grid"
    table.autofit = False
    from docx.enum.table import WD_TABLE_ALIGNMENT

    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for node in clone.tblPr.findall(qn("w:tblInd")):
        clone.tblPr.remove(node)
    width = clone.tblPr.find(qn("w:tblW"))
    if width is not None:
        width.set(qn("w:w"), "10080")
        width.set(qn("w:type"), "dxa")
    grid = clone.tblGrid.gridCol_lst
    total = sum(int(col.w) for col in grid)
    scale = int(Inches(7)) / total if total else 1
    for col in grid:
        col.w = int(col.w * scale)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:" + side)
        for attr, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            edge.set(qn("w:" + attr), value)
        borders.append(edge)
    old = clone.tblPr.find(qn("w:tblBorders"))
    if old is not None:
        clone.tblPr.remove(old)
    clone.tblPr.append(borders)
    for index, row in enumerate(table.rows):
        row.height = None
        row.height_rule = None
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        seen_cells = set()
        for cell in row.cells:
            if cell._tc in seen_cells:
                continue
            seen_cells.add(cell._tc)
            if cell.width is not None:
                cell.width = int(cell.width * scale)
            from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            for tag in ("tcBorders", "tcMar", "shd"):
                for old in tc_pr.findall(qn("w:" + tag)):
                    tc_pr.remove(old)
            cell_borders = deepcopy(borders)
            cell_borders.tag = qn("w:tcBorders")
            tc_pr.append(cell_borders)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                edge = OxmlElement("w:" + side)
                edge.set(qn("w:w"), "100")
                edge.set(qn("w:type"), "dxa")
                margins.append(edge)
            cell._tc.get_or_add_tcPr().append(margins)
            if index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EDEDED")
                cell._tc.get_or_add_tcPr().append(shading)
            for paragraph in cell.paragraphs:
                paragraph.style = document.styles["Normal"]
                fmt = paragraph.paragraph_format
                fmt.space_before, fmt.space_after = Pt(3), Pt(3)
                fmt.line_spacing = 1.15
                fmt.left_indent = fmt.right_indent = fmt.first_line_indent = Pt(0)
                fmt.keep_with_next = False
                if index == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(12)
                    run.font.name = "Times New Roman"
                    run._r.get_or_add_rPr().get_or_add_rFonts().set(
                        qn("w:eastAsia"), "宋体"
                    )
                    run.font.color.rgb = RGBColor(0, 0, 0)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build_documents(
    folder: Path,
    title: str,
    items: list[dict[str, Any]],
    workspace: Path,
    answer_lines: int,
) -> dict[str, dict[str, Any]]:
    native = NativeParagraphs(workspace)
    # Validate both roles before authoring either output.
    content = [
        (native.read(item, "question"), native.read(item, "answer")) for item in items
    ]
    artifacts = {}
    for role, teacher in (("student", False), ("teacher", True)):
        document = _new_document(title, teacher)
        previous_group = None
        for index, (item, (question, answer)) in enumerate(
            zip(items, content, strict=True), 1
        ):
            group = (
                item["batch_id"],
                item["package_id"],
                item["editable_source"]["parent_id"],
            )
            if group != previous_group:
                document.add_paragraph(item["parent_title"], "Heading 1")
                previous_group = group
            document.add_paragraph(f"练习 {index}", "Heading 2")
            source = document.add_paragraph(
                f"来源：{item['source_name']}；原题号 {item['printed_number']}"
            )
            source.paragraph_format.keep_with_next = True
            for run in source.runs:
                run.font.size = Pt(10)
            _append_native(document, answer if teacher else question)
            if not teacher:
                for _ in range(answer_lines):
                    blank = document.add_paragraph(" ")
                    blank.paragraph_format.space_after = Pt(6)
        path = folder / FILES[role]
        document.save(path)
        artifacts[role] = {
            "filename": path.name,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "bytes": path.stat().st_size,
        }
    return artifacts


class HandoutPracticeService:
    def __init__(self, candidates: HandoutCandidateService):
        self.candidates = candidates
        self.state = candidates.state
        self.root = (
            candidates.workspace / "runtime/deeptutor_shchem/handout-practice-exports"
        )

    @staticmethod
    def _inputs(
        title: str, selections: list[dict[str, str]], answer_lines: int
    ) -> None:
        if (
            not isinstance(title, str)
            or not title.strip()
            or len(title) > 100
            or any(ord(c) < 32 for c in title)
        ):
            raise _error("title_invalid", "请填写 1 至 100 字的练习标题，不要换行。")
        if not isinstance(selections, list) or not 1 <= len(selections) <= 100:
            raise _error("selection_invalid", "请选择 1 至 100 道讲义题。")
        if any(
            not isinstance(s, dict)
            or not isinstance(s.get("key"), str)
            or not isinstance(s.get("revision"), str)
            for s in selections
        ):
            raise _error("selection_invalid", "选题信息不完整，请重新选题。")
        if len({s["key"] for s in selections}) != len(selections):
            raise _error("duplicate", "练习中存在重复题目，请移除重复项。")
        if type(answer_lines) is not int or not 0 <= answer_lines <= 8:
            raise _error("spacing_invalid", "每题作答留白应为 0 至 8 行。")

    def draft(self) -> dict[str, Any] | None:
        return self.state.snapshot()["drafts"].get(DRAFT_ID)

    def save_draft(
        self, title: str, selections: list[dict[str, str]], answer_lines: int
    ) -> dict[str, Any]:
        self._inputs(title, selections, answer_lines)
        result = {
            "kind": "native_handout_practice_selection",
            "title": title.strip(),
            "selections": deepcopy(selections),
            "answer_lines": answer_lines,
            "created_at": utc_now(),
        }
        self.state.save_draft(DRAFT_ID, result)
        return result

    def export(
        self, title: str, selections: list[dict[str, str]], answer_lines: int = 2
    ) -> dict[str, Any]:
        self._inputs(title, selections, answer_lines)
        catalog = {item["key"]: item for item in self.candidates.catalog()["items"]}
        items = []
        for selection in selections:
            item = catalog.get(selection["key"])
            if item is None or item["revision"] != selection["revision"]:
                raise _error(
                    "stale", "所选题目的来源已变化或不可用，请刷新后重新选题。"
                )
            if item["classification"] != "native_text_complete":
                raise _error(
                    "incomplete",
                    "尚有待补图或待补公式的题目，请移出后再导出可编辑练习。",
                )
            items.append(item)
        export_id = "HANDOUT-PRACTICE-" + uuid.uuid4().hex
        self.root.mkdir(parents=True, exist_ok=True)
        pending = self.root / (".incomplete-" + export_id)
        pending.mkdir()
        try:
            artifacts = build_documents(
                pending, title.strip(), items, self.candidates.workspace, answer_lines
            )
            record = {
                "kind": EXPORT_KIND,
                "export_id": export_id,
                "title": title.strip(),
                "created_at": utc_now(),
                "items": items,
                "answer_lines": answer_lines,
                "artifacts": artifacts,
                "candidate_only": True,
                "answer_correctness_verified": False,
                "paper_structure_claimed": False,
            }
            (pending / "来源与导出记录.json").write_text(
                json.dumps(record, ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )
            pending.rename(self.root / export_id)
        except HandoutCandidateError:
            raise
        except Exception as exc:
            raise _error(
                "write_failed",
                "练习导出未完成，请检查原始 Word 和输出目录；未登记不完整结果。",
            ) from exc
        self.state.save_draft(export_id, record)
        self.save_draft(title, selections, answer_lines)
        return record

    def history(self) -> list[dict[str, Any]]:
        return sorted(
            (
                item
                for item in self.state.snapshot()["drafts"].values()
                if isinstance(item, dict) and item.get("kind") == EXPORT_KIND
            ),
            key=lambda item: item["created_at"],
            reverse=True,
        )[:20]

    def artifact_path(self, export_id: str, role: str) -> Path:
        if not _EXPORT_ID.fullmatch(export_id) or role not in FILES:
            raise _error("artifact_invalid", "请选择有效的学生版或教师参考答案文件。")
        record = self.state.snapshot()["drafts"].get(export_id)
        if not record or record.get("kind") != EXPORT_KIND:
            raise _error("artifact_missing", "此导出记录不存在。")
        path = (self.root / export_id / FILES[role]).resolve()
        path.relative_to(self.root.resolve())
        if not path.is_file():
            raise _error("artifact_missing", "文件已被移动或删除，请重新导出。")
        return path
