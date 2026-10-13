"""Embed validated source pictures without changing their bytes.

python-docx recognizes JPEGs by JFIF/EXIF metadata and parses EXIF to find DPI.
Some fully decodable photos omit that optional metadata or contain an incomplete
EXIF field. On those parser failures, use the checked pixel dimensions and a
native picture part; source bytes and requested Word geometry stay intact.
"""
from __future__ import annotations

from io import BytesIO


def add_word_picture(run, data, asset, *, width=None, height=None):
    """Add a source image to an existing run, retaining its text and formatting."""
    from docx.image.exceptions import UnexpectedEndOfFileError, UnrecognizedImageError
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.shape import CT_Inline
    from docx.parts.image import ImagePart
    from docx.shape import InlineShape
    from docx.shared import Emu
    from .desktop_preparation_images import verify_image_bytes

    # Never treat a metadata-parser error as permission to embed an unverified
    # or partially decoded image, including on the normal insertion path.
    data = verify_image_bytes(asset, data)
    try:
        return run.add_picture(BytesIO(data), width=width, height=height)
    except (UnexpectedEndOfFileError, UnrecognizedImageError):
        if asset['content_type'] != 'image/jpeg':
            raise
        if width is None and height is None:
            raise ValueError('JPEG元数据无法读取，请指定图片的显示尺寸。') from None

    # Every production caller supplies at least one fitted size. Pixel aspect
    # ratio is independent of the optional DPI field that could not be parsed.
    if width is None:
        width = round(int(height) * asset['width'] / asset['height'])
    if height is None:
        height = round(int(width) * asset['height'] / asset['width'])

    package = run.part.package
    image_part = next((part for part in package.image_parts
                       if part.content_type == 'image/jpeg' and part.blob == data), None)
    if image_part is None:
        image_part = ImagePart(package.next_partname('/word/media/image%d.jpg'),
                               'image/jpeg', data)
        package.image_parts.append(image_part)
    relationship = run.part.relate_to(image_part, RT.IMAGE)
    inline = CT_Inline.new_pic_inline(run.part.next_id, relationship,
                                     image_part.filename, Emu(width), Emu(height))
    run._r.add_drawing(inline)
    return InlineShape(inline)


def add_word_document_picture(document, data, asset, *, width=None, height=None):
    """Match Document.add_picture's single paragraph/run on either path."""
    return add_word_picture(document.add_paragraph().add_run(), data, asset,
                            width=width, height=height)
