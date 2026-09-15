"""Synthetic source/EXE acceptance, deliberately isolated from teacher materials."""
from __future__ import annotations
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import time


def verify_answer_openings(path):
    """For this two-question fixture, each answer label must share its picture's page."""
    import re
    import pypdfium2 as pdfium
    from pypdfium2 import raw
    answers, images = {}, []
    with pdfium.PdfDocument(str(path)) as pdf:
        for index in range(len(pdf)):
            page = pdf[index]
            textpage = page.get_textpage()
            text = textpage.get_text_range()
            for match in re.finditer(r'【答案】\s*([12])[.．]', text):
                answers[int(match[1])] = index
                assert '参考答案' in text, 'Generated answer heading became orphaned'
            bounds = sorted((obj.get_bounds() for obj in page.get_objects(filter=[raw.FPDF_PAGEOBJ_IMAGE])),
                            key=lambda b: -b[3])
            images.extend(index for _ in bounds)
            textpage.close();page.close()
    assert len(images)==4 and set(answers)=={1,2}, 'Unexpected synthetic PDF content'
    assert answers[1]==images[1] and answers[2]==images[3], 'Answer label split from its picture'
    return True


def run_probe(output):
    from PIL import Image, ImageDraw, ImageFont, ImageChops
    from docx import Document
    from docx.shared import Inches
    from .desktop_facade import build_default_facade
    from .desktop_paths import DesktopPaths
    from .desktop_version import DESKTOP_VERSION
    from .desktop_raster_numbers import render_numbered_image
    from .desktop_raster_paper_service import RasterPaperService
    from .desktop_workbench.app import create_application
    from .desktop_workbench.main_window import TeacherWorkbenchWindow
    app=create_application(['scan194-probe'])
    output=Path(output);output.mkdir(parents=True,exist_ok=True)
    screenshots=[]
    def capture(widget,name):
        path=output/name;assert widget.grab().save(str(path))
        screenshots.append({'file':name,'sha256':sha256(path.read_bytes()).hexdigest()})
    def settle(predicate=lambda:True):
        end=time.monotonic()+90
        for _ in range(8):app.processEvents();time.sleep(.03)
        while not predicate() and time.monotonic()<end:app.processEvents();time.sleep(.03)
        assert predicate(), 'UI did not finish requested operation'
    with tempfile.TemporaryDirectory(prefix='scan194-') as temp:
        temp=Path(temp)
        root=Path(getattr(sys,'_MEIPASS',Path(__file__).resolve().parents[2]))
        facade=build_default_facade(DesktopPaths.from_workspace(root,state_root=temp/'state'))
        win=TeacherWorkbenchWindow(facade);win.resize(1280,900);win.show()
        try:
            doc=Document();raw_images={};new_numbers={}
            for number,new in ((27,2),(38,1)):
                scale=2 if number==38 else 1
                try:font=ImageFont.truetype('arial.ttf',38*scale)
                except OSError:font=ImageFont.load_default(size=38*scale)
                doc.add_paragraph(f'【例{number}】合成扫描题：看图作答，数据保持12.5。')
                for answer in (False,True):
                    if answer:doc.add_paragraph(f'【答案】{number}. 合成答案。')
                    image=Image.new('RGB',(800*scale,240*scale),'white');draw=ImageDraw.Draw(image)
                    draw.text((12*scale,12*scale),str(number)+'.',font=font,fill='black')
                    draw.text((110*scale,12*scale),'12.5 mol/L' + (' answer' if answer else ' question'),font=font,fill='black')
                    draw.line((100*scale,125*scale,700*scale,125*scale),width=2*scale,fill='black')
                    data=BytesIO();image.save(data,format='PNG');data=data.getvalue()
                    key=sha256(data).hexdigest();raw_images[key]=data;new_numbers[key]=new
                    doc.add_picture(BytesIO(data),width=Inches(5.5))
            source=temp/'合成扫描题.docx';doc.save(source);source_bytes=source.read_bytes()
            facade.save_visual_import_batch(handout_files=(source,),source_type='教师讲义')
            rows=[r for r in facade.word_question_catalog()['items'] if r['source_name']==source.name]
            assert len(rows)==2
            facade.add_word_questions_to_basket([{'key':r['key'],'revision':r['revision']} for r in reversed(rows)])
            original_state=deepcopy(facade.basket())
            win.navigate('paper');settle(lambda:win.paper_page._mixed_panel is not None and not win.paper_page._mixed_panel._busy)
            panel=win.paper_page._mixed_panel
            assert panel.number_button.isEnabled();panel.number_button.click()
            settle(lambda:panel._number_dialog is not None or not panel._busy)
            dialog=panel._number_dialog;assert dialog is not None, panel.status.text()
            assert len(dialog.images)==4
            for index,image in enumerate(dialog.images):
                dialog.picker.setCurrentIndex(index)
                dialog.canvas.selection=[int(v*image['width']/800) for v in [8,8,97,62]]
                assert image['suggested_number'] == new_numbers[image['image_sha256']]
                assert dialog.number.value() == image['suggested_number']
                dialog._add_region()
                dialog.remember.click()
                assert dialog.load_saved.isEnabled()
            assert len(dialog.edits)==4
            edits=deepcopy(dialog.edits)
            dialog.undo.click();assert len(dialog.edits)==3
            last=edits[-1];dialog.picker.setCurrentIndex(next(i for i,r in enumerate(dialog.images) if r['image_sha256']==last['image_sha256']))
            dialog.canvas.selection=list(last['box']);dialog.number.setValue(last['number']);dialog._add_region()
            assert dialog.edits==edits
            capture(dialog,'scan-number-editor.png')
            remembered = deepcopy(facade.state_store.snapshot()['scan_number_regions'])
            assert all('number' not in r for im in remembered['images'].values() for r in im['regions'])
            dialog.original.setChecked(True);capture(dialog,'scan-number-original.png');dialog.original.setChecked(False)
            dialog.apply.click()
            settle(lambda:panel._preview_dialog is not None or not panel._busy)
            assert panel._preview_dialog is not None, panel.status.text()
            preview=panel._preview
            assert preview.preview_model['scan_numbering']['edits']==edits
            service=RasterPaperService(facade)
            for role in ('student','teacher'):
                for page in preview.preview_model['pagination']['documents'][role]['pages']:
                    data=service.image(preview.preview_id,page['image_id'])['data']
                    path=output/f'scan-{role}-page-{page["page_number"]}.png';path.write_bytes(data)
                    screenshots.append({'file':path.name,'sha256':sha256(data).hexdigest()})
            service.approve(preview.preview_id,preview.preview_hash)
            exported=service.export(preview.preview_id,preview.preview_hash)
            for artifact in exported['artifacts']:
                shutil.copyfile(artifact['path'],output/Path(artifact['path']).name)
            answer_pdf = next(Path(a['path']) for a in exported['artifacts'] if a['artifact_id']=='teacher_pdf')
            answer_layout = verify_answer_openings(answer_pdf)
            # Re-create the state reader, and form a NEW paper order. The saved
            # boxes survive while the old paper's numbers are deliberately absent.
            from .desktop_number_regions import NumberRegionStore
            from .desktop_state import DesktopStateStore
            from .desktop_workbench.scan_number_dialog import ScanNumberDialog
            panel._preview_dialog.reject();settle()
            reordered = deepcopy(panel.request())
            reordered['section_order'] = list(reversed(reordered['section_order']))
            next_preview = facade.create_paper_preview(reordered)
            next_catalog = RasterPaperService(facade).inspect_images(next_preview.preview_id, next_preview.preview_hash)
            reusable = NumberRegionStore(DesktopStateStore(facade.paths.state_root))
            reopened = ScanNumberDialog(next_catalog, win, region_store=reusable)
            reopened.show();settle()
            for index, image in enumerate(reopened.images):
                reopened.picker.setCurrentIndex(index)
                assert image['suggested_number'] == 3-new_numbers[image['image_sha256']]
                assert reopened.number.value()==image['suggested_number']
                assert reopened.load_saved.isEnabled() and reopened.saved_selector.count()==1
                before_count=len(reopened.edits)
                reopened.load_saved.click()
                assert len(reopened.edits)==before_count
                reopened.add.click()
                assert reopened.edits[-1]['number']==image['suggested_number']
            reopened.resize(800,700);settle()
            for button in (reopened.remember,reopened.load_saved,reopened.apply,reopened.cancel):
                assert reopened.rect().contains(button.mapTo(reopened,button.rect().bottomRight()))
            capture(reopened,'scan-reused-regions.png')
            second_edits=deepcopy(reopened.edits)
            reopened.accept();reopened.deleteLater();settle()
            second=RasterPaperService(facade).prepare_corrected(next_preview.preview_id,next_preview.preview_hash,second_edits)
            service.approve(second.preview_id,second.preview_hash)
            second_export=service.export(second.preview_id,second.preview_hash)
            second_answer=next(Path(a['path']) for a in second_export['artifacts'] if a['artifact_id']=='teacher_pdf')
            reordered_layout=verify_answer_openings(second_answer)
            for artifact in second_export['artifacts']:
                shutil.copyfile(artifact['path'],output/('换序-'+Path(artifact['path']).name))
            for role in ('student','teacher'):
                for page in second.preview_model['pagination']['documents'][role]['pages']:
                    data=service.image(second.preview_id,page['image_id'])['data']
                    path=output/f'scan-reordered-{role}-page-{page["page_number"]}.png';path.write_bytes(data)
                    screenshots.append({'file':path.name,'sha256':sha256(data).hexdigest()})
            assert facade.state_store.snapshot()['scan_number_regions']==remembered
            assert source.read_bytes()==source_bytes and facade.basket()==original_state
            for edit in edits:
                raw=raw_images[edit['image_sha256']]
                changed=render_numbered_image(raw,[edit])
                a=Image.open(BytesIO(raw)).convert('RGB');b=Image.open(BytesIO(changed)).convert('RGB')
                delta=ImageChops.difference(a,b);assert delta.getbbox()
                delta.paste((0,0,0),tuple(edit['box']));assert delta.getbbox() is None
            report={'version':DESKTOP_VERSION,'source_commit':os.environ.get('GITHUB_SHA','local'),
                    'frozen':getattr(sys,'frozen',False),'qt_platform':app.platformName(),
                    'image_resolutions':[[800,240],[1600,480]],'edits':len(edits),'raw_word_import':True,
                    'native_entry_opened':True,'undo_exercised':True,'two_audiences_exported':True,
                    'outside_regions_unchanged':True,'source_unchanged':True,
                    'region_reuse':{'saved_positions_only':True,'reopened_state':True,'new_order_new_numbers':True,
                                   'load_requires_add':True,'actions_visible_800x700':True,
                                   'answer_openings_with_images':answer_layout,'reordered_answer_openings':reordered_layout},
                    'pages':{r:preview.preview_model['pagination']['documents'][r]['page_count'] for r in ('student','teacher')},
                    'screenshots':screenshots,'model_calls':0,'scope':'Synthetic Word-hosted raster images; manual pixel selection, no OCR or real-library acceptance.'}
            (output/'scan-probe.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        finally:
            if not (output/'scan-probe.json').is_file():
                for index,path in enumerate((temp/'state').rglob('*.docx')):
                    shutil.copyfile(path,output/f'diagnostic-{index}-{path.name}')
            win.close();app.processEvents()
    return 0
