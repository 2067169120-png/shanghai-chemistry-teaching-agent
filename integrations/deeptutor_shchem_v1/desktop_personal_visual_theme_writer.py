"""Image-first body blocks and DOCX append helper for frozen visual themes.

No provider, state mutation, document conversion or pagination certification.
The caller verifies current sources with the export service and owns the final
render/approval gate. Every required image is validated even for a student-only
request; missing originals never silently become OCR-only output.
"""

from __future__ import annotations

from collections.abc import Mapping
import hashlib
from io import BytesIO
import math
import unicodedata

from .desktop_personal_visual_questions import _digest, _question_text
from .desktop_personal_visual_theme_export import (
    ITEM_KIND,
    SCHEMA_VERSION,
    PersonalVisualThemeExportError,
)
from .desktop_preparation_images import image_info
from .reader_cancellation import check_read_cancelled


def _fail(message):
    raise PersonalVisualThemeExportError(message)


def _source_geometry(descriptor):
    """Return frozen pixel geometry, never infer a page size from a crop/DPI."""
    binding = descriptor["binding"]
    page, box = binding.get("page_size"), binding.get("pixel_xyxy")
    if (not isinstance(page, (list, tuple)) or len(page) != 2
            or any(type(value) is not int or value <= 0 for value in page)
            or not isinstance(box, (list, tuple)) or len(box) != 4
            or any(type(value) is not int for value in box)):
        _fail("主题原图缺少可追溯的原页像素尺寸或裁剪范围，请重新预览。")
    left, top, right, bottom = box
    if (not (0 <= left < right <= page[0] and 0 <= top < bottom <= page[1])
            or (right - left, bottom - top) != (descriptor["width"], descriptor["height"])):
        _fail("主题裁片尺寸与冻结原页几何不一致，请重新预览。")
    identity = (binding.get("source_file_id"), binding.get("page_number"),
                binding.get("page_sha256"))
    return identity, page, box


def _validate(item, assets):
    if (not isinstance(item, Mapping) or item.get("kind") != ITEM_KIND
            or item.get("settings") != {"use_source_scores": True}
            or not isinstance(item.get("content"), Mapping)
            or not isinstance(item.get("source_ref"), Mapping)
            or not isinstance(assets, Mapping)):
        _fail("完整主题导出参数无效；来源分值不可改成自设分值。")
    content, source = item["content"], item["source_ref"]
    if (content.get("schema_version") != SCHEMA_VERSION
            or source.get("content_sha256") != _digest(content)
            or content.get("candidate_only") is not True
            or content.get("teacher_reviewed") is not False
            or content.get("publication_allowed") is not False
            or not isinstance(content.get("images"), list)):
        _fail("完整主题内容与冻结摘要不一致，请重新预览。")
    images, pixel_roles, page_sizes = {}, {}, {}
    for descriptor in content["images"]:
        check_read_cancelled()
        if not isinstance(descriptor, Mapping):
            _fail("主题图片描述不完整。")
        key, binding = descriptor.get("asset_id"), descriptor.get("binding")
        if not isinstance(key, str) or key in images or not isinstance(binding, Mapping):
            _fail("主题图片身份缺失或重复。")
        role = descriptor.get("role")
        teacher = role == "answer"
        if (role not in {"question", "shared_material", "answer"}
                or descriptor.get("teacher_only") is not teacher
                or binding.get("source_role") != ("answer" if teacher else "question")):
            _fail("题面与教师答案图片角色冲突。")
        raw = assets.get(key)
        if not isinstance(raw, bytes) or hashlib.sha256(raw).hexdigest() != descriptor.get("sha256"):
            _fail("主题原图缺失或摘要已变化，不能用识别文字代替。")
        info = image_info(raw)
        if any(descriptor.get(field) != info[field] for field in ("width", "height", "content_type")):
            _fail("主题原图尺寸或格式与冻结记录不一致。")
        if key != "PVTHEMEIMG-" + _digest({"binding": binding, "role": role, "sha256": descriptor["sha256"]}):
            _fail("主题图片与来源定位不一致。")
        page_key, page_size, _ = _source_geometry(descriptor)
        if page_key in page_sizes and page_sizes[page_key] != tuple(page_size):
            _fail("同一冻结原页的像素尺寸不一致，请重新预览。")
        page_sizes[page_key] = tuple(page_size)
        sha = descriptor["sha256"]
        if sha in pixel_roles and pixel_roles[sha] != teacher:
            _fail("相同原图同时出现在题面与答案中，不能导出学生版。")
        pixel_roles[sha] = teacher
        images[key] = descriptor
    return content, images


def build_theme_blocks(item, assets, audience, *, show_scores=False):
    """Return mixed text/image blocks without duplicating image-backed OCR.

    Direct node evidence selects image-first presentation. A visual-object edge
    alone does not suppress text. The typed source remains intact in content;
    image-first does not assert that AI boxes have passed a human completeness
    check. Repeated exact source-page rectangles are displayed once per theme,
    while all explicit per-node relationships remain in the typed projection.
    """
    if audience not in {"student", "teacher"} or type(show_scores) is not bool:
        _fail("题目版本或来源分值显示选项无效。")
    content, images = _validate(item, assets)
    theme, blocks, shown = content["theme"], [], set()
    group_id, segment_id = "theme:context", "theme:context"

    def block_metadata():
        # Additive layout hints only: do not change the frozen source graph.
        return {"group_id": group_id, "segment_id": segment_id}

    def text(value, *, style="body"):
        if isinstance(value, str) and value.strip() and value.strip() != "unknown":
            blocks.append({"kind": "text", "text": value.strip(), "style": style,
                           **block_metadata()})

    def pictures(refs, expected_role):
        if not isinstance(refs, list):
            _fail("主题图文关系不完整。")
        for key in refs:
            check_read_cancelled()
            descriptor = images.get(key)
            if descriptor is None or descriptor["role"] != expected_role:
                _fail("主题必要原图未对应到正确内容，不能省略图形。")
            if descriptor["teacher_only"] and audience != "teacher":
                _fail("教师答案图片不能出现在学生版本。")
            binding = descriptor["binding"]
            identity = (binding["source_file_id"], binding["page_number"], binding["page_sha256"],
                        tuple(binding["pixel_xyxy"]), descriptor["sha256"], descriptor["teacher_only"])
            if identity not in shown:
                blocks.append({"kind": "image", "asset_id": key,
                               "caption_zh": descriptor["caption"], "teacher_only": descriptor["teacher_only"],
                               **block_metadata()})
                shown.add(identity)

    def has_body_image(node):
        direct = set(node.get("evidence_refs", []))
        return any(images.get(key, {}).get("binding", {}).get("evidence_id") in direct
                   for key in node.get("image_refs", []))

    # Context is a theme synopsis, not another shared-material source node.
    # Once shared source-body images are present, present those originals and
    # every text-only material instead of repeating the synopsis above them.
    # A figure-only visual edge does not establish body coverage.
    if not any(has_body_image(material) for material in theme["shared_materials"]):
        text(theme.get("context"))
    for index, material in enumerate(theme["shared_materials"]):
        group_id = segment_id = f"shared:{index}"
        if not has_body_image(material):
            text(material.get("content"))
            for expression in material.get("chemical_expressions", []):
                text(expression.get("raw"))
        pictures(material["image_refs"], "shared_material")
    for printed_index, printed in enumerate(theme["printed_questions"]):
        check_read_cancelled()
        group_id = f"question:{printed_index}"
        segment_id = group_id + ":stem"
        text(f"第{printed['question_number']}题", style="heading")
        if not has_body_image(printed):
            text(_question_text(printed))
        pictures(printed["image_refs"], "question")
        for atomic_index, atomic in enumerate(printed["atomic_parts"]):
            segment_id = group_id + f":atomic:{atomic_index}"
            if not has_body_image(atomic):
                value = _question_text(atomic)
                # Same typed parent text may be copied into a no-image atomic;
                # omit only literal equality, not semantic/chemistry rewriting.
                if value and value != _question_text(printed):
                    text(((atomic.get("part_label") or "") + " " + value).strip())
            pictures(atomic["image_refs"], "question")
        if audience == "teacher":
            group_id = f"answer:{printed_index}"
            segment_id = group_id + ":atomic:0"
            text("来源参考答案", style="heading")
            for atomic_index, atomic in enumerate(printed["atomic_parts"]):
                segment_id = group_id + f":atomic:{atomic_index}"
                answer = atomic["answer"]
                if answer["status"] == "missing":
                    text(f"{atomic.get('part_label') or ''} 来源答案待补充，未推导或借用其他题答案。")
                elif not has_body_image(answer):
                    text(answer.get("answer_body"))
                    text(answer.get("analysis"))
                    for expression in answer.get("chemical_expressions", []):
                        text(expression.get("raw"))
                    for point in answer.get("scoring_points", []):
                        text(f"{point['score']:g}分：{point['description']}")
                pictures(answer["image_refs"], "answer")
        if audience == "teacher" or show_scores:
            for atomic_index, atomic in enumerate(printed["atomic_parts"]):
                # Scores are source metadata, not a keep-chain into the next
                # question. Preserve original score values unchanged.
                group_id = segment_id = f"score:{printed_index}:{atomic_index}"
                answer = atomic["answer"]
                label = atomic.get("part_label") or ""
                if answer["status"] == "missing" or answer["max_score"] <= 0:
                    text(f"{label} 来源分值待核对", style="note")
                else:
                    text(f"{label} 来源参考分值：{answer['max_score']:g}分", style="note")
    if audience == "teacher":
        group_id = segment_id = "theme:notes"
        text(f"来源：{item['source_zh']}", style="note")
        text("来源参考答案及分值为 AI 识别候选，非官方评分认定，待对照原页；未作教师化学审核。", style="note")
        for warning in dict.fromkeys(content["warnings"]):
            text(warning, style="note")
    if not blocks:
        _fail("完整主题没有可编排的内容。")
    return blocks


def _image_widths(descriptors, available_width, max_picture_height):
    """One pixel-to-EMU scale per frozen page, independent of crop width.

    The full source page is mapped to the output content width. This deliberately
    does not invent a scanner DPI or enlarge a short answer to fill the line.
    If a tall crop needs reducing, reduce that page's other crops by the same
    factor. Both audiences use all descriptors, so shared question sizes agree.
    """
    scales = {}
    for descriptor in descriptors.values():
        page_key, page, _ = _source_geometry(descriptor)
        scale = min(available_width / page[0], max_picture_height / descriptor["height"])
        scales[page_key] = min(scales.get(page_key, scale), scale)
    return {key: max(1, int(descriptor["width"] * scales[_source_geometry(descriptor)[0]]))
            for key, descriptor in descriptors.items()}


def _text_height(paragraph, available_width):
    """Conservative layout budget, not a claim of measured Word pagination."""
    from docx.shared import Pt

    styles, current = [], paragraph.style
    while current is not None:
        styles.append(current)
        current = current.base_style
    sizes = [run.font.size for run in paragraph.runs if run.font.size is not None]
    sizes.extend(style.font.size for style in styles if style.font.size is not None)
    font = max(sizes or [Pt(11)])
    units_per_line = max(1.0, available_width / font)
    lines = sum(max(1, math.ceil(sum(1 if unicodedata.east_asian_width(char) in {"W", "F"}
                                   else .6 for char in line) / units_per_line))
                for line in paragraph.text.split("\n"))
    formats = [paragraph.paragraph_format, *(style.paragraph_format for style in styles)]

    def inherited(name):
        return next((getattr(fmt, name) for fmt in formats if getattr(fmt, name) is not None), None)

    spacing = inherited("line_spacing")
    # python-docx returns a float for multiples and a Length for exact spacing.
    line_height = (font * max(1.3, spacing) if isinstance(spacing, float)
                   else max(font * 1.3, spacing or 0))
    return int(lines * line_height + (inherited("space_before") or 0)
               + (inherited("space_after") or 0) + Pt(3))


def _keep_groups(layout, capacity):
    """Keep ordinary questions whole; bound oversized groups at source units.

    Each entry is (block, paragraph, estimated_height). Atomic/material units
    are preferred break boundaries. An oversized unit can break between image
    fragments, but no chain is made taller than the usable-page budget.
    """
    def chain(entries):
        for index, (_, paragraph, estimated) in enumerate(entries):
            paragraph.paragraph_format.keep_together = estimated <= capacity
            paragraph.paragraph_format.keep_with_next = index < len(entries) - 1

    start = 0
    while start < len(layout):
        end = start + 1
        while end < len(layout) and layout[end][0]["group_id"] == layout[start][0]["group_id"]:
            end += 1
        group = layout[start:end]
        if sum(entry[2] for entry in group) <= capacity:
            chain(group)
        else:
            units = []
            for entry in group:
                if not units or units[-1][-1][0]["segment_id"] != entry[0]["segment_id"]:
                    units.append([])
                units[-1].append(entry)
            # First split only units that are themselves too tall.
            bounded = []
            for unit in units:
                chunk, used = [], 0
                for entry in unit:
                    if chunk and used + entry[2] > capacity:
                        bounded.append(chunk)
                        chunk, used = [], 0
                    chunk.append(entry)
                    used += entry[2]
                bounded.append(chunk)
            chunk, used = [], 0
            for unit in bounded:
                size = sum(entry[2] for entry in unit)
                if chunk and used + size > capacity:
                    chain(chunk)
                    chunk, used = [], 0
                chunk.extend(unit)
                used += size
            chain(chunk)
        start = end


def append_personal_visual_theme(document, item, assets, audience, ordinal=1, show_scores=False):
    """Append the exact block projection to a configured python-docx document.

    Whole images retain a consistent source-page pixel density; no crop,
    reconstructed chemistry, new answer lines or source-score rewriting.
    Bounded keep chains protect question fragments and answer headings, not
    entire oversized themes. They are layout hints, not pagination approval.
    Physical pagination still requires the caller's real PDF render/preview.
    """
    from docx.shared import Emu, Pt

    if type(ordinal) is not int or ordinal < 1:
        _fail("完整主题顺序无效。")
    blocks = build_theme_blocks(item, assets, audience, show_scores=show_scores)

    def style(name):
        return name if name in document.styles else "Normal"

    heading = document.add_paragraph(f"题组 {ordinal} · {item['title_zh']}", style("ShChemThemeHeading"))
    heading.paragraph_format.keep_with_next = True
    section = document.sections[-1]
    width = int(section.page_width - section.left_margin - section.right_margin)
    # Reserve a line for the picture paragraph and conversion roundoff.
    height = int(section.page_height - section.top_margin - section.bottom_margin - Pt(24))
    if width <= 0 or height <= Pt(80):
        _fail("文档可用页面尺寸无效。")
    descriptors = {image["asset_id"]: image for image in item["content"]["images"]}
    widths = _image_widths(descriptors, width, height - int(Pt(60)))
    layout = []
    for block in blocks:
        check_read_cancelled()
        if block["kind"] == "text":
            style_name = {"heading": "ShChemQuestion", "note": "ShChemQuiet"}.get(block["style"], "ShChemQuestion")
            paragraph = document.add_paragraph(block["text"], style(style_name))
            estimated = _text_height(paragraph, width)
        else:
            descriptor = descriptors[block["asset_id"]]
            fitted_width = widths[block["asset_id"]]
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = False
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run()
            run.font.size = Pt(1)
            run.add_picture(BytesIO(assets[block["asset_id"]]), width=Emu(fitted_width))
            estimated = int(fitted_width * descriptor["height"] / descriptor["width"] + Pt(8))
        layout.append((block, paragraph, estimated))
    _keep_groups(layout, height - int(Pt(12)))
    return blocks
