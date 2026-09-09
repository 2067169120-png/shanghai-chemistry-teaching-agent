from __future__ import annotations

import argparse
import base64
import hashlib
import importlib.util
import io
import json
import os
import re
import shutil
import subprocess
import sys
from collections.abc import Mapping, Sequence
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, ClassVar
from zipfile import ZipFile

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import (
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image
from pypdf import PdfReader

if __package__:
    from .answer_diagrams import answer_diagram_png
    from .paper_format_presets import (
        PAPER_FORMAT_SCHEMA_VERSION,
        RENDER_RECEIPT_KIND,
        PaperFormatContractError,
        _chinese_number,
        build_assembly_blueprint,
        build_document_plans,
        build_render_request,
        default_shanghai_theme_preset,
        run_export_preflight,
        validate_render_receipts,
    )
else:  # Direct CLI execution avoids importing the gateway package and its optional HTTP deps.
    from answer_diagrams import answer_diagram_png
    from paper_format_presets import (  # type: ignore[no-redef]
        PAPER_FORMAT_SCHEMA_VERSION,
        RENDER_RECEIPT_KIND,
        PaperFormatContractError,
        _chinese_number,
        build_assembly_blueprint,
        build_document_plans,
        build_render_request,
        default_shanghai_theme_preset,
        run_export_preflight,
        validate_render_receipts,
    )


RENDERER_SCHEMA_VERSION = "shchem.paper-export-renderer.v1"
RENDER_BUNDLE_KIND = "paper_export_render_bundle"
QA_REPORT_KIND = "paper_export_render_qa_report"
HASH_MANIFEST_KIND = "paper_export_hash_manifest"
RENDERER_VERSION = "1.1.0"

_SOURCE_IMAGE_BLOCK_TYPES = {"image", "apparatus"}
_CHOICE_ITEM_MARKERS = ("choice",)
_ONE_LINE_ITEM_MARKERS = (
    "fill",
    "equation",
    "notation",
    "formula",
    "diagram_reading",
    "symbolic_response",
    "orbital_diagram",
)
_THREE_LINE_ITEM_MARKERS = (
    "calculation",
    "quantitative",
    "organic",
    "experiment",
    "graph",
)
_TWO_LINE_ITEM_MARKERS = (
    "explanation",
    "evaluation",
    "response",
    "observation",
    "interpretation",
    "evidence",
    "short_answer",
)

ARTIFACT_FILENAMES = {
    "student_docx": "上海化学主题练习-学生版.docx",
    "student_pdf": "上海化学主题练习-学生版.pdf",
    "teacher_docx": "上海化学主题练习-教师版.docx",
    "teacher_pdf": "上海化学主题练习-教师版.pdf",
}

_BUNDLE_KEYS = {
    "renderer_schema_version",
    "contract_kind",
    "bundle_id",
    "claim_boundary_zh",
    "preset",
    "blueprint",
    "student_plan",
    "teacher_plan",
    "preflight_report",
    "render_request",
    "publication_allowed",
}
_SHA256 = re.compile(r"^[0-9a-f]{64}$")
_SAFE_BUNDLE_ID = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.:-]{0,159}$")
_STUDENT_FORBIDDEN_MARKERS = (
    "参考答案（非官方",
    "建议答案",
    "教师解析",
    "AI候选解析",
    "易错点：",
    "来源标签：",
)
_SECRET_PATTERNS = (
    re.compile(r"\bsk-[A-Za-z0-9_-]{8,}"),
    re.compile(r"\bapi[_-]?key\b", re.IGNORECASE),
    re.compile(r"\bcredential[_-]?ref\b", re.IGNORECASE),
    re.compile(r"\bauthorization\s*:", re.IGNORECASE),
    re.compile(r"\bbearer\s+[A-Za-z0-9._~+/-]{8,}", re.IGNORECASE),
)


class PaperExportRendererError(RuntimeError):
    """Renderer failure with a stable code and teacher-facing Chinese message."""

    def __init__(self, code: str, message_zh: str) -> None:
        super().__init__(message_zh)
        self.code = code
        self.message_zh = message_zh


@dataclass(frozen=True)
class RendererToolchain:
    python_exe: Path
    render_docx_script: Path
    pdftoppm_exe: Path
    dpi: int = 300
    conversion_backend: str = "auto"

    def validated(self) -> RendererToolchain:
        for label, path in (
            ("Python", self.python_exe),
            ("render_docx.py", self.render_docx_script),
            ("Poppler", self.pdftoppm_exe),
        ):
            if not path.is_file():
                raise PaperExportRendererError(
                    "renderer_tool_missing", f"{label} 工具不存在，无法完成逐页渲染。"
                )
        if self.conversion_backend not in {"auto", "libreoffice", "word_com"}:
            raise PaperExportRendererError(
                "renderer_backend_invalid", "文档转换后端设置不正确。"
            )
        if not 150 <= self.dpi <= 2400:
            raise PaperExportRendererError(
                "renderer_dpi_invalid", "逐页检查清晰度必须在 150—2400 dpi 之间。"
            )
        return self


def _canonical_json_bytes(value: Any) -> bytes:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    ).encode("utf-8")


def _sha256_json(value: Any) -> str:
    return hashlib.sha256(_canonical_json_bytes(value)).hexdigest()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n",
        encoding="utf-8",
    )
    temporary.replace(path)


def _read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _relative(path: Path, root: Path) -> str:
    return path.resolve().relative_to(root.resolve()).as_posix()


def _require_exact_keys(value: Mapping[str, Any], expected: set[str], code: str) -> None:
    if set(value) != expected:
        raise PaperExportRendererError(code, "渲染包字段不完整或含未知字段。")


def validate_render_bundle(value: Mapping[str, Any]) -> dict[str, Any]:
    if not isinstance(value, Mapping):
        raise PaperExportRendererError("render_bundle_invalid", "渲染包格式不正确。")
    _require_exact_keys(value, _BUNDLE_KEYS, "render_bundle_invalid")
    record = deepcopy(dict(value))
    if (
        record.get("renderer_schema_version") != RENDERER_SCHEMA_VERSION
        or record.get("contract_kind") != RENDER_BUNDLE_KIND
        or record.get("publication_allowed") is not False
    ):
        raise PaperExportRendererError(
            "render_bundle_invalid", "渲染包版本或发布边界不正确。"
        )
    bundle_id = record.get("bundle_id")
    if not isinstance(bundle_id, str) or not _SAFE_BUNDLE_ID.fullmatch(bundle_id):
        raise PaperExportRendererError("render_bundle_invalid", "渲染包标识不正确。")
    claim = record.get("claim_boundary_zh")
    if (
        not isinstance(claim, str)
        or "不可发布" not in claim
        or not any(marker in claim for marker in ("非上海原题", "题目来源见卷内"))
    ):
        raise PaperExportRendererError(
            "render_bundle_claim_invalid",
            "渲染包必须说明题目来源边界并明确仅供本地备课、不可发布。",
        )

    try:
        recomputed_preflight = run_export_preflight(
            preset=record["preset"],
            blueprint=record["blueprint"],
            student_plan=record["student_plan"],
            teacher_plan=record["teacher_plan"],
        )
        if _canonical_json_bytes(recomputed_preflight) != _canonical_json_bytes(
            record["preflight_report"]
        ):
            raise PaperExportRendererError(
                "preflight_drift", "导出前检查结果与冻结计划不一致。"
            )
        recomputed_request = build_render_request(
            preset=record["preset"],
            student_plan=record["student_plan"],
            teacher_plan=record["teacher_plan"],
            preflight_report=record["preflight_report"],
        )
    except PaperFormatContractError as exc:
        raise PaperExportRendererError(exc.code, exc.message_zh) from exc
    if _canonical_json_bytes(recomputed_request) != _canonical_json_bytes(
        record["render_request"]
    ):
        raise PaperExportRendererError(
            "render_request_drift", "渲染请求与冻结文档计划不一致。"
        )
    blueprint = record["blueprint"]
    if (
        blueprint.get("top_level_unit") != "theme_big_question"
        or blueprint.get("standalone_choice_section_allowed") is not False
        or blueprint.get("status") != "ready_for_content_resolution"
    ):
        raise PaperExportRendererError(
            "theme_first_required", "只能渲染完整主题大题，不能生成独立选择题板块。"
        )
    return record


def _set_style_fonts(style: Any, *, cjk: str, latin: str, size_pt: float) -> None:
    style.font.name = latin
    style.font.size = Pt(size_pt)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), cjk)
    r_fonts.set(qn("w:cs"), latin)


def _content_width_dxa(preset: Mapping[str, Any]) -> int:
    margins = preset["page_layout"]["margins_mm"]
    width_mm = 210.0 - float(margins["left"]) - float(margins["right"])
    if width_mm <= 0:
        raise PaperExportRendererError("layout_width_invalid", "页边距导致正文宽度不可用。")
    return round(width_mm / 25.4 * 1440)


def _set_run_font(
    run: Any,
    *,
    cjk: str = "SimSun",
    latin: str = "Times New Roman",
    size_pt: float = 10.5,
    bold: bool | None = None,
    color: RGBColor | None = None,
) -> None:
    run.font.name = latin
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), cjk)
    r_fonts.set(qn("w:cs"), latin)


def _set_repeat_table_width(table: Any, width_dxa: int) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    indent = tbl_pr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), "0")
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for grid_col in grid.findall(qn("w:gridCol")):
        grid_col.set(qn("w:w"), str(width_dxa))
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width_dxa))
            tc_w.set(qn("w:type"), "dxa")


def _set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell: Any, *, top: int, start: int, bottom: int, end: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_bottom_border(cell: Any, color: str = "A6A6A6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "end", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), color)


def _set_paragraph_bottom_border(paragraph: Any, color: str = "A6A6A6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)


def _add_field(paragraph: Any, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, display, end])
    _set_run_font(run, cjk="SimSun", latin="Times New Roman", size_pt=9)


def _style_document(doc: DocumentType, preset: Mapping[str, Any]) -> None:
    layout = preset["page_layout"]
    body = layout["body"]
    normal = doc.styles["Normal"]
    _set_style_fonts(
        normal,
        cjk=body["cjk_font"],
        latin=body["latin_font"],
        size_pt=float(body["size_pt"]),
    )
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = float(body["line_spacing"])
    normal.paragraph_format.widow_control = True

    styles = doc.styles
    definitions = {
        "ShChemPaperTitle": (
            layout["title"]["cjk_font"],
            body["latin_font"],
            float(layout["title"]["size_pt"]),
            bool(layout["title"]["bold"]),
        ),
        "ShChemPaperMeta": (
            body["cjk_font"],
            body["latin_font"],
            float(body["size_pt"]),
            False,
        ),
        "ShChemThemeHeading": (
            layout["theme_heading"]["cjk_font"],
            body["latin_font"],
            float(layout["theme_heading"]["size_pt"]),
            bool(layout["theme_heading"]["bold"]),
        ),
        "ShChemQuestion": (
            body["cjk_font"],
            body["latin_font"],
            float(body["size_pt"]),
            False,
        ),
        "ShChemTeacherNote": (
            body["cjk_font"],
            body["latin_font"],
            max(8.0, float(body["size_pt"]) - 2.0),
            False,
        ),
        "ShChemQuiet": (
            body["cjk_font"],
            body["latin_font"],
            max(7.5, float(body["size_pt"]) - 2.5),
            False,
        ),
    }
    for name, (cjk, latin, size, bold) in definitions.items():
        if name in styles:
            style = styles[name]
        else:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        _set_style_fonts(style, cjk=cjk, latin=latin, size_pt=size)
        style.font.bold = bold
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = float(body["line_spacing"])
        style.paragraph_format.widow_control = True

    styles["ShChemPaperTitle"].paragraph_format.space_after = Pt(3)
    styles["ShChemThemeHeading"].paragraph_format.space_before = Pt(5)
    styles["ShChemThemeHeading"].paragraph_format.space_after = Pt(2)
    styles["ShChemThemeHeading"].paragraph_format.keep_with_next = True
    # Keep labels with their content locally, not every question in a theme.
    styles["ShChemQuestion"].paragraph_format.keep_with_next = False
    styles["ShChemTeacherNote"].paragraph_format.line_spacing = 1.08
    styles["ShChemTeacherNote"].paragraph_format.space_after = Pt(1)
    styles["ShChemQuiet"].paragraph_format.line_spacing = 1.05
    styles["ShChemQuiet"].paragraph_format.space_after = Pt(1)


def _populate_page_footer(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    first = paragraph.add_run("第 ")
    _set_run_font(first, cjk="SimSun", latin="Times New Roman", size_pt=9)
    _add_field(paragraph, "PAGE")
    middle = paragraph.add_run(" 页 / 共 ")
    _set_run_font(middle, cjk="SimSun", latin="Times New Roman", size_pt=9)
    _add_field(paragraph, "NUMPAGES")
    last = paragraph.add_run(" 页")
    _set_run_font(last, cjk="SimSun", latin="Times New Roman", size_pt=9)


def _configure_section(doc: DocumentType, preset: Mapping[str, Any], visible: Mapping[str, Any]) -> None:
    layout = preset["page_layout"]
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    margins = layout["margins_mm"]
    section.top_margin = Mm(float(margins["top"]))
    section.bottom_margin = Mm(float(margins["bottom"]))
    section.left_margin = Mm(float(margins["left"]))
    section.right_margin = Mm(float(margins["right"]))
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)
    section.different_first_page_header_footer = True

    header = section.header
    header.is_linked_to_previous = False
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_paragraph.paragraph_format.space_after = Pt(0)
    run = header_paragraph.add_run(visible["header"]["text_zh"])
    _set_run_font(run, cjk="SimSun", latin="Times New Roman", size_pt=9, color=RGBColor(89, 89, 89))

    first_header = section.first_page_header
    first_header.is_linked_to_previous = False
    first_header_paragraph = first_header.paragraphs[0]
    first_header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first_header_paragraph.paragraph_format.space_after = Pt(0)

    footer = section.footer
    footer.is_linked_to_previous = False
    _populate_page_footer(footer.paragraphs[0])
    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    _populate_page_footer(first_footer.paragraphs[0])


def _add_title_block(
    doc: DocumentType,
    visible: Mapping[str, Any],
    audience: str,
    *,
    preset: Mapping[str, Any],
) -> None:
    title = doc.add_paragraph(style="ShChemPaperTitle")
    title.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[preset["page_layout"]["title"]["alignment"]]
    title.paragraph_format.keep_with_next = True
    title.add_run(visible["paper_title_zh"])

    is_local_basket_export = "题目来源见教师版" in str(
        visible.get("version_label_zh") or ""
    )
    subtitle_candidates = (
        (visible.get("subtitle_zh"),)
        if is_local_basket_export
        else (visible.get("subtitle_zh"), visible.get("version_label_zh"))
    )
    subtitle_parts = [item for item in subtitle_candidates if item]
    if subtitle_parts:
        subtitle = doc.add_paragraph(style="ShChemPaperMeta")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.keep_with_next = True
        subtitle.add_run("　｜　".join(subtitle_parts))

    boundary_text = (
        "本地题篮导出｜题目来源见教师版｜不可发布"
        if is_local_basket_export
        else "合成布局演示｜非上海原题｜项目模板｜不可发布"
    )
    boundary = doc.add_paragraph(style="ShChemQuiet")
    boundary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    boundary.paragraph_format.keep_with_next = True
    boundary_run = boundary.add_run(boundary_text)
    _set_run_font(
        boundary_run,
        cjk="SimSun",
        latin="Times New Roman",
        size_pt=9,
        bold=True,
        color=(
            RGBColor(72, 104, 86)
            if is_local_basket_export
            else RGBColor(156, 35, 35)
        ),
    )

    exam = visible["exam_info"]
    info = doc.add_paragraph(style="ShChemPaperMeta")
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.keep_with_next = True
    info_text = f"建议时间：{exam['duration_minutes']} 分钟"
    if audience == "teacher" or preset["student_version"]["show_total_score"]:
        info_text += f"　　建议满分：{_format_score(exam['total_score'])} 分"
    info.add_run(info_text)

    if audience == "student":
        identity = doc.add_paragraph(style="ShChemPaperMeta")
        identity.paragraph_format.space_before = Pt(2)
        identity.paragraph_format.space_after = Pt(5)
        identity.paragraph_format.keep_with_next = True
        identity.add_run("姓名：________________　　班级：________________")

    rules = exam.get("scoring_rules") or {}
    rule_texts = [
        rules.get("selection_rule_zh"),
        rules.get("partial_credit_rule_zh"),
        rules.get("other_rule_zh"),
    ]
    rule_texts = [text for text in rule_texts if text]
    if rule_texts:
        rule = doc.add_paragraph(style="ShChemQuiet")
        rule.paragraph_format.space_after = Pt(5)
        rule.add_run("作答说明：" + " ".join(rule_texts))


def _format_score(value: Any) -> str:
    number = float(value)
    return str(int(number)) if number.is_integer() else f"{number:g}"


def _append_block_to_paragraph(paragraph: Any, block: Mapping[str, Any]) -> bool:
    block_type = block["block_type"]
    text = block.get("text_zh")
    if block_type in {"paragraph", "formula", "structure", "table"} and text:
        run = paragraph.add_run(text)
        if block_type in {"formula", "structure"}:
            inherited_size = (
                paragraph.style.font.size.pt
                if paragraph.style is not None and paragraph.style.font.size is not None
                else 10.5
            )
            _set_run_font(
                run,
                cjk="SimSun",
                latin="Cambria Math",
                size_pt=inherited_size,
            )
        return True
    return False


def _add_asset_block(
    container: Any,
    block: Mapping[str, Any],
    asset_root: Path | None,
    *,
    max_width_mm: float = 165.0,
    alt_prefix_zh: str | None = None,
    keep_with_next: bool = False,
) -> None:
    asset_ref = block.get("asset_ref")
    if not asset_ref or asset_root is None:
        raise PaperExportRendererError(
            "asset_unavailable", "题图或装置图缺少可读取的本地资源。"
        )
    asset_path = (asset_root / asset_ref).resolve()
    try:
        asset_path.relative_to(asset_root.resolve())
    except ValueError as exc:
        raise PaperExportRendererError("asset_path_invalid", "题图路径超出允许目录。") from exc
    if not asset_path.is_file():
        raise PaperExportRendererError("asset_unavailable", "题图或装置图文件不存在。")
    paragraph = container.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.left_indent = Mm(2)
    paragraph.paragraph_format.right_indent = Mm(2)
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run()
    # Question crops vary from one-line prompts to tall apparatus/flow-chart
    # panels.  Bound both dimensions so a legitimate tall crop cannot extend
    # beyond the printable A4 body while preserving its aspect ratio.
    try:
        with Image.open(asset_path) as source_image:
            width_px, height_px = source_image.size
    except (OSError, ValueError) as exc:
        raise PaperExportRendererError(
            "asset_image_invalid", "题图无法作为有效图片读取。"
        ) from exc
    if width_px <= 0 or height_px <= 0:
        raise PaperExportRendererError(
            "asset_image_invalid", "题图尺寸不正确。"
        )
    # A small source diagram is not a full-width question scan. Avoid blowing
    # a 130px apparatus thumbnail up to an entire page merely because space is
    # available. This changes only its physical layout size, never source bytes.
    width_mm = min(max_width_mm, width_px * 25.4 / 120.0)
    height_mm = width_mm * height_px / width_px
    if height_mm > 210.0:
        height_mm = 210.0
        width_mm = height_mm * width_px / height_px
    run.add_picture(str(asset_path), width=Mm(width_mm), height=Mm(height_mm))
    drawing = run._r.xpath(".//wp:docPr")
    if drawing:
        alt_text = block.get("alt_text_zh") or "化学题图"
        if alt_prefix_zh:
            alt_text = alt_prefix_zh + alt_text
        drawing[0].set("descr", alt_text)


def _is_source_image_question(atomic: Mapping[str, Any]) -> bool:
    for block in atomic.get("question_blocks") or []:
        if block.get("block_type") in _SOURCE_IMAGE_BLOCK_TYPES:
            return True
        if block.get("text_zh"):
            return False
    return False


def _source_asset_edge_review_notes(
    plan: Mapping[str, Any],
    asset_root: Path | None,
) -> list[dict[str, Any]]:
    if asset_root is None:
        return []
    resolved_root = asset_root.resolve()
    asset_refs: set[str] = set()
    for section in plan["visible"]["theme_sections"]:
        for material in section["shared_materials"]:
            for block in material["content_blocks"]:
                if block.get("block_type") in _SOURCE_IMAGE_BLOCK_TYPES and block.get(
                    "asset_ref"
                ):
                    asset_refs.add(str(block["asset_ref"]))
        for printed in section["printed_questions"]:
            for atomic in printed["atomic_parts"]:
                for block in atomic["question_blocks"]:
                    if block.get(
                        "block_type"
                    ) in _SOURCE_IMAGE_BLOCK_TYPES and block.get("asset_ref"):
                        asset_refs.add(str(block["asset_ref"]))

    risk_count = 0
    top_dark_total = 0
    bottom_dark_total = 0
    for asset_ref in sorted(asset_refs):
        asset_path = (resolved_root / asset_ref).resolve()
        try:
            asset_path.relative_to(resolved_root)
        except ValueError as exc:
            raise PaperExportRendererError(
                "asset_path_invalid", "题图路径超出允许目录。"
            ) from exc
        if not asset_path.is_file():
            continue
        try:
            with Image.open(asset_path) as source:
                rgba = source.convert("RGBA")
                white = Image.new("RGBA", rgba.size, "white")
                white.alpha_composite(rgba)
                gray = white.convert("L")
                band_height = min(2, gray.height)
                top_dark = sum(
                    pixel < 235
                    for pixel in gray.crop(
                        (0, 0, gray.width, band_height)
                    ).get_flattened_data()
                )
                bottom_dark = sum(
                    pixel < 235
                    for pixel in gray.crop(
                        (0, gray.height - band_height, gray.width, gray.height)
                    ).get_flattened_data()
                )
        except (OSError, ValueError) as exc:
            raise PaperExportRendererError(
                "asset_image_invalid", "题图无法作为有效图片读取。"
            ) from exc
        if top_dark or bottom_dark:
            risk_count += 1
            top_dark_total += top_dark
            bottom_dark_total += bottom_dark
    if risk_count == 0:
        return []
    return [
        {
            "reviewer_kind": "renderer_source_asset_edge_scan_warning",
            "note_zh": (
                f"题图边缘非阻断提醒：{len(asset_refs)} 张源裁片中 {risk_count} 张"
                f"的上或下边缘存在深色像素（上缘合计 {top_dark_total}，"
                f"下缘合计 {bottom_dark_total}）。导出器只在题图外围增加安全白边，"
                "未改动源像素；逐页目视时须确认上下边界未造成文字截断。"
            ),
            "scope_zh": "自动题图上下边缘安全扫描；只作视觉复核警示，不替代逐页目视或化学审校。",
            "human_reviewed": False,
            "chemistry_reviewed": False,
        }
    ]


def _resolved_asset_path(asset_root: Path, asset_ref: str) -> Path:
    resolved_root = asset_root.resolve()
    asset_path = (resolved_root / asset_ref).resolve()
    try:
        asset_path.relative_to(resolved_root)
    except ValueError as exc:
        raise PaperExportRendererError(
            "asset_path_invalid", "题图路径超出允许目录。"
        ) from exc
    if not asset_path.is_file():
        raise PaperExportRendererError(
            "asset_unavailable", "题图或装置图文件不存在。"
        )
    return asset_path


def _load_raster_rows(path: Path) -> dict[str, Any]:
    try:
        with Image.open(path) as source:
            rgba = source.convert("RGBA")
            white = Image.new("RGBA", rgba.size, "white")
            white.alpha_composite(rgba)
            rgb = white.convert("RGB")
            width, height = rgb.size
            raw = rgb.tobytes()
    except (OSError, ValueError) as exc:
        raise PaperExportRendererError(
            "asset_image_invalid", "题图无法作为有效图片读取。"
        ) from exc
    if width <= 0 or height <= 0:
        raise PaperExportRendererError("asset_image_invalid", "题图尺寸不正确。")
    stride = width * 3
    rows = tuple(raw[index : index + stride] for index in range(0, len(raw), stride))
    return {
        "width": width,
        "height": height,
        "rows": rows,
        "file_sha256": _sha256_file(path),
    }


def _longest_common_exact_run(first: Sequence[bytes], second: Sequence[bytes]) -> int:
    positions: dict[bytes, list[int]] = {}
    for index, row in enumerate(second):
        positions.setdefault(row, []).append(index)
    previous: dict[int, int] = {}
    longest = 0
    for row in first:
        current: dict[int, int] = {}
        for second_index in positions.get(row, []):
            length = previous.get(second_index - 1, 0) + 1
            current[second_index] = length
            longest = max(longest, length)
        previous = current
    return longest


def _column_signatures(rows: Sequence[bytes], width: int) -> tuple[bytes, ...]:
    return tuple(
        b"".join(row[column * 3 : column * 3 + 3] for row in rows)
        for column in range(width)
    )


def _is_exact_subimage(shared: Mapping[str, Any], question: Mapping[str, Any]) -> bool:
    shared_width = int(shared["width"])
    shared_height = int(shared["height"])
    question_width = int(question["width"])
    question_height = int(question["height"])
    if shared_width > question_width or shared_height > question_height:
        return False
    shared_rows = shared["rows"]
    question_rows = question["rows"]
    anchor_index = max(
        range(shared_height),
        key=lambda index: sum(value < 235 for value in shared_rows[index]),
    )
    anchor = shared_rows[anchor_index]
    for question_anchor_index, question_row in enumerate(question_rows):
        top = question_anchor_index - anchor_index
        if top < 0 or top + shared_height > question_height:
            continue
        start = 0
        while True:
            byte_offset = question_row.find(anchor, start)
            if byte_offset < 0:
                break
            if byte_offset % 3 == 0 and byte_offset // 3 + shared_width <= question_width:
                if all(
                    question_rows[top + row_index][
                        byte_offset : byte_offset + shared_width * 3
                    ]
                    == shared_rows[row_index]
                    for row_index in range(shared_height)
                ):
                    return True
            start = byte_offset + 3
    return False


def _shared_exact_overlap_ratio(
    shared: Mapping[str, Any],
    question: Mapping[str, Any],
) -> float:
    if shared["file_sha256"] == question["file_sha256"]:
        return 1.0
    if _is_exact_subimage(shared, question):
        return 1.0
    ratios = [0.0]
    if shared["width"] == question["width"]:
        ratios.append(
            _longest_common_exact_run(shared["rows"], question["rows"])
            / int(shared["height"])
        )
    if shared["height"] == question["height"]:
        shared_columns = _column_signatures(shared["rows"], int(shared["width"]))
        question_columns = _column_signatures(
            question["rows"], int(question["width"])
        )
        ratios.append(
            _longest_common_exact_run(shared_columns, question_columns)
            / int(shared["width"])
        )
    return max(ratios)


def _shared_question_visual_dedup(
    plan: Mapping[str, Any],
    asset_root: Path | None,
) -> dict[str, Any]:
    if asset_root is None:
        return {"suppressed_shared_keys": frozenset(), "records": []}
    cache: dict[str, dict[str, Any]] = {}

    def raster(asset_ref: str) -> dict[str, Any]:
        if asset_ref not in cache:
            cache[asset_ref] = _load_raster_rows(
                _resolved_asset_path(asset_root, asset_ref)
            )
        return cache[asset_ref]

    suppressed: set[str] = set()
    records: list[dict[str, Any]] = []
    for section in plan["visible"]["theme_sections"]:
        question_refs = [
            str(block["asset_ref"])
            for printed in section["printed_questions"]
            for atomic in printed["atomic_parts"]
            for block in atomic["question_blocks"]
            if block.get("block_type") in _SOURCE_IMAGE_BLOCK_TYPES
            and block.get("asset_ref")
        ]
        for material in section["shared_materials"]:
            blocks = list(material["content_blocks"])
            shared_refs = [
                str(block["asset_ref"])
                for block in blocks
                if block.get("block_type") in _SOURCE_IMAGE_BLOCK_TYPES
                and block.get("asset_ref")
            ]
            if not shared_refs or len(shared_refs) != len(blocks):
                continue
            material_matches: list[dict[str, Any]] = []
            for shared_ref in shared_refs:
                best: dict[str, Any] | None = None
                for question_ref in question_refs:
                    overlap = _shared_exact_overlap_ratio(
                        raster(shared_ref), raster(question_ref)
                    )
                    if overlap < 0.90:
                        continue
                    method = (
                        "same_sha256"
                        if raster(shared_ref)["file_sha256"]
                        == raster(question_ref)["file_sha256"]
                        else (
                            "exact_subimage"
                            if overlap == 1.0
                            else "exact_contiguous_pixel_overlap"
                        )
                    )
                    candidate = {
                        "shared_asset": Path(shared_ref).name,
                        "question_asset": Path(question_ref).name,
                        "method": method,
                        "exact_overlap_ratio": round(overlap, 6),
                    }
                    if best is None or candidate["exact_overlap_ratio"] > best[
                        "exact_overlap_ratio"
                    ]:
                        best = candidate
                if best is None:
                    material_matches = []
                    break
                material_matches.append(best)
            if material_matches:
                key = str(material["render_once_key"])
                suppressed.add(key)
                records.append(
                    {
                        "render_once_key": key,
                        "matches": material_matches,
                    }
                )
    return {
        "suppressed_shared_keys": frozenset(suppressed),
        "records": records,
    }


def _shared_question_dedup_review_notes(
    dedup: Mapping[str, Any],
) -> list[dict[str, Any]]:
    records = list(dedup.get("records") or [])
    if not records:
        return []
    method_counts: dict[str, int] = {}
    for record in records:
        for match in record["matches"]:
            method = str(match["method"])
            method_counts[method] = method_counts.get(method, 0) + 1
    method_summary = "、".join(
        f"{method}={count}" for method, count in sorted(method_counts.items())
    )
    return [
        {
            "reviewer_kind": "renderer_shared_question_exact_visual_dedup",
            "note_zh": (
                f"同主题 shared-vs-question 精确视觉去重：共抑制 "
                f"{len(records)} 个已被题面图包含的共享材料块（{method_summary}）。"
                "仅在同 SHA-256、精确子图或至少 90% 连续精确像素重叠时触发；"
                "题面图保留，源像素未改动。"
            ),
            "scope_zh": "渲染层同主题共享材料与题面图精确去重；不是化学内容或来源判定。",
            "human_reviewed": False,
            "chemistry_reviewed": False,
        }
    ]


def _add_content_blocks(
    doc: DocumentType,
    blocks: Sequence[Mapping[str, Any]],
    *,
    first_paragraph: Any | None = None,
    asset_root: Path | None = None,
    source_question_number: int | None = None,
    keep_last_with_next: bool = False,
) -> None:
    pending_first = first_paragraph
    source_number_anchor_pending = source_question_number is not None
    for block_index, block in enumerate(blocks):
        keep_next = block_index < len(blocks) - 1 or keep_last_with_next
        block_type = block["block_type"]
        if pending_first is not None and _append_block_to_paragraph(pending_first, block):
            pending_first.paragraph_format.keep_with_next = keep_next
            pending_first = None
            continue
        pending_first = None
        if block_type in _SOURCE_IMAGE_BLOCK_TYPES:
            alt_prefix = None
            if source_number_anchor_pending:
                alt_prefix = f"第{source_question_number}题题图："
                source_number_anchor_pending = False
            _add_asset_block(
                doc,
                block,
                asset_root,
                alt_prefix_zh=alt_prefix,
                keep_with_next=keep_next,
            )
            continue
        paragraph = doc.add_paragraph(style="ShChemQuestion")
        paragraph.paragraph_format.keep_with_next = keep_next
        if not _append_block_to_paragraph(paragraph, block):
            raise PaperExportRendererError(
                "content_block_unsupported", "题面包含当前渲染器无法处理的内容块。"
            )


def _add_shared_material(
    doc: DocumentType,
    material: Mapping[str, Any],
    *,
    audience: str,
    asset_root: Path | None,
    content_width_dxa: int,
    body_size_pt: float,
    keep_last_with_next: bool = True,
) -> None:
    # Consecutive one-cell tables can coalesce in Word, pushing a whole group
    # away from the title and splitting a material label from its image.
    # Ordinary inline paragraphs keep the figure and first question in flow.
    paragraph = doc.add_paragraph(style="ShChemQuestion")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_after = Pt(1)
    label = paragraph.add_run("【共同材料】")
    _set_run_font(
        label,
        cjk="SimHei",
        latin="Times New Roman",
        size_pt=body_size_pt,
        bold=True,
    )
    blocks = material["content_blocks"]
    has_source_label = audience == "teacher" and material.get("source_label_zh")
    for index, block in enumerate(blocks):
        keep_next = index < len(blocks) - 1 or bool(has_source_label) or keep_last_with_next
        if block["block_type"] in _SOURCE_IMAGE_BLOCK_TYPES:
            _add_asset_block(
                doc,
                block,
                asset_root,
                max_width_mm=min(162.0, content_width_dxa * 25.4 / 1440.0 - 4.0),
                keep_with_next=keep_next,
            )
            continue
        target = (
            paragraph
            if index == 0
            else doc.add_paragraph(style="ShChemQuestion")
        )
        target.paragraph_format.keep_with_next = keep_next
        if not _append_block_to_paragraph(target, block):
            raise PaperExportRendererError(
                "content_block_unsupported", "共同材料包含当前渲染器无法处理的内容块。"
            )
    if has_source_label:
        source = doc.add_paragraph(style="ShChemQuiet")
        source.paragraph_format.keep_with_next = keep_last_with_next
        source.paragraph_format.space_before = Pt(1)
        source.paragraph_format.space_after = Pt(0)
        source_run = source.add_run("材料来源标签：" + material["source_label_zh"])
        _set_run_font(
            source_run,
            cjk="SimSun",
            latin="Times New Roman",
            size_pt=8.5,
            color=RGBColor(89, 89, 89),
        )


def _effective_answer_line_count(requested_lines: int, item_type: str | None) -> int:
    requested = max(0, int(requested_lines))
    if item_type is None:
        return requested
    normalized = str(item_type).strip().lower()
    if any(marker in normalized for marker in _CHOICE_ITEM_MARKERS):
        cap = 0
    elif any(marker in normalized for marker in _ONE_LINE_ITEM_MARKERS):
        cap = 1
    elif any(marker in normalized for marker in _THREE_LINE_ITEM_MARKERS):
        cap = 3
    elif any(marker in normalized for marker in _TWO_LINE_ITEM_MARKERS):
        cap = 2
    else:
        cap = 2
    return min(requested, cap)


def _answer_line_width_dxa(content_width_dxa: int, lines: int) -> int:
    if lines <= 1:
        ratio = 0.62
    elif lines == 2:
        ratio = 0.84
    else:
        ratio = 0.90
    return max(1, round(content_width_dxa * ratio))


def _add_answer_space(
    doc: DocumentType,
    lines: int,
    *,
    content_width_dxa: int,
    body_size_pt: float,
) -> None:
    if lines <= 0:
        return
    table = doc.add_table(rows=lines, cols=1)
    _set_repeat_table_width(
        table,
        _answer_line_width_dxa(content_width_dxa, lines),
    )
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row_index, row in enumerate(table.rows):
        row.height = Pt(max(13.0, body_size_pt + 3.0))
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row_pr = row._tr.get_or_add_trPr()
        cannot_split = OxmlElement("w:cantSplit")
        row_pr.append(cannot_split)
        cell = row.cells[0]
        _set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        _set_cell_bottom_border(cell)
        paragraph = cell.paragraphs[0]
        paragraph.style = doc.styles["ShChemQuestion"]
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.keep_with_next = row_index < lines - 1
        run = paragraph.add_run("\u00a0")
        _set_run_font(run, size_pt=body_size_pt)


def _teacher_explanation_label(value: str) -> str:
    return {
        "source_analysis": "来源解析（非官方）",
        "teacher_analysis": "教师解析（非官方）",
        "ai_candidate": "AI候选解析（非官方，须复核）",
        "none": "解析",
    }[value]


def _compact_teacher_explanation(value: str) -> str:
    kept_lines = [
        line.strip()
        for line in str(value or "").splitlines()
        if line.strip()
        and not (
            "照录来源答案" in line
            and "不作独立正确性核验" in line
        )
    ]
    return "\n".join(kept_lines)


def _compact_answer_source_label(value: str | None) -> str:
    normalized = str(value or "").strip()
    if normalized == "题库已逐题对齐的非官方参考答案；按来源答案直接呈现。":
        return "非官方参考答案（题库逐题对齐）"
    return normalized


def _teacher_source_signature(notes: Mapping[str, Any]) -> tuple[str, str]:
    reference = notes["source_reference_answer"]
    return (
        str(notes.get("source_label_zh") or "").strip(),
        str(reference.get("source_label_zh") or "").strip(),
    )


def _add_answer_text(paragraph: Any, text: str) -> None:
    """Keep the Avogadro constant editable, with an actual subscript A."""
    for index, part in enumerate(text.split("N_A")):
        if index:
            paragraph.add_run("N").font.italic = True
            paragraph.add_run("A").font.subscript = True
        if part:
            paragraph.add_run(part)


def _add_teacher_notes(
    doc: DocumentType,
    notes: Mapping[str, Any],
    *,
    content_width_dxa: int,
    include_source_label: bool,
    scoring_label_zh: str | None = None,
    asset_root: Path | None = None,
) -> None:
    reference = notes["source_reference_answer"]
    answer_blocks = reference.get("content_blocks") or []
    table = doc.add_table(rows=1, cols=1)
    _set_repeat_table_width(table, content_width_dxa)
    # Source routes may be tall. Permit the note row to flow across pages;
    # each inline image remains whole instead of shrinking the entire answer.
    if not answer_blocks:
        row_pr = table.rows[0]._tr.get_or_add_trPr()
        cannot_split = OxmlElement("w:cantSplit")
        row_pr.append(cannot_split)
    cell = table.cell(0, 0)
    _set_cell_shading(cell, "FFF8E8")
    _set_cell_margins(cell, top=60, start=100, bottom=60, end=100)
    paragraph = cell.paragraphs[0]
    paragraph.style = doc.styles["ShChemTeacherNote"]
    paragraph.paragraph_format.space_after = Pt(1)
    note_size_pt = doc.styles["ShChemTeacherNote"].font.size.pt
    if scoring_label_zh:
        # Plain answers already have a non-splitting row. Adding keepNext to
        # its first cell paragraph makes Word chain following table rows and
        # can push the first theme off the title page.
        paragraph.paragraph_format.keep_with_next = bool(answer_blocks)
        scoring = paragraph.add_run(scoring_label_zh)
        _set_run_font(scoring, cjk="SimHei", latin="Times New Roman", size_pt=note_size_pt, bold=True)
        paragraph = cell.add_paragraph(style="ShChemTeacherNote")
        paragraph.paragraph_format.space_after = Pt(1)
    label = paragraph.add_run(f"{notes['answer_label_zh']}：")
    _set_run_font(
        label,
        cjk="SimHei",
        latin="Times New Roman",
        size_pt=note_size_pt,
        bold=True,
    )
    supplement = notes.get("supplemental_answer") or {}
    answer_text = supplement.get("text_zh") or reference.get("text_zh") or "本题暂无已对齐参考答案。"
    _add_answer_text(paragraph, answer_text)
    if answer_blocks:
        paragraph.paragraph_format.keep_with_next = True
        for block in answer_blocks:
            _add_asset_block(
                cell, block, asset_root=asset_root,
                max_width_mm=min(162.0, (content_width_dxa - 240) * 25.4 / 1440.0),
                alt_prefix_zh=None if str(block.get("alt_text_zh", "")).startswith("非官方参考答案图") else "非官方参考答案图：",
            )
    if supplement.get("diagram_key"):
        picture = cell.add_paragraph(style="ShChemTeacherNote")
        diagram_width = (
            (content_width_dxa - 240) / 20
            if supplement["diagram_key"] == "ferrate_single_electron_bridge"
            else min(130 if supplement["diagram_key"].endswith("atom_shells") else 160, (content_width_dxa - 240) / 20)
        )
        shape = picture.add_run().add_picture(
            io.BytesIO(answer_diagram_png(supplement["diagram_key"])),
            width=Pt(diagram_width),
        )
        shape._inline.docPr.set("descr", supplement["text_zh"])

    compact_explanation = _compact_teacher_explanation(
        str(notes.get("explanation_zh") or "")
    )
    if compact_explanation:
        explanation = cell.add_paragraph(style="ShChemTeacherNote")
        explanation.paragraph_format.space_after = Pt(1)
        heading = explanation.add_run(("解题过程" if supplement else _teacher_explanation_label(notes["explanation_label"])) + "：")
        _set_run_font(
            heading,
            cjk="SimHei",
            latin="Times New Roman",
            size_pt=note_size_pt,
            bold=True,
        )
        _add_answer_text(explanation, compact_explanation)
    if notes.get("pitfalls_zh"):
        pitfalls = cell.add_paragraph(style="ShChemTeacherNote")
        pitfalls.paragraph_format.space_after = Pt(1)
        heading = pitfalls.add_run("易错点：")
        _set_run_font(
            heading,
            cjk="SimHei",
            latin="Times New Roman",
            size_pt=note_size_pt,
            bold=True,
        )
        pitfalls.add_run("；".join(notes["pitfalls_zh"]))
    if include_source_label:
        source = cell.add_paragraph(style="ShChemQuiet")
        source.paragraph_format.space_after = Pt(0)
        source_label = str(notes.get("source_label_zh") or "").strip()
        answer_source = supplement.get("source_label_zh") or _compact_answer_source_label(reference.get("source_label_zh"))
        source_parts = []
        if source_label:
            source_parts.append("来源标签：" + source_label)
        if answer_source:
            source_parts.append("答案来源：" + answer_source)
        source.add_run("；".join(source_parts) or "来源标签：未提供")


def _blueprint_theme_sections(
    visible: Mapping[str, Any],
    blueprint: Mapping[str, Any] | None,
) -> list[Mapping[str, Any] | None]:
    visible_sections = visible["theme_sections"]
    if blueprint is None:
        return [None] * len(visible_sections)
    blueprint_sections = list(blueprint.get("theme_bundles") or [])
    if len(blueprint_sections) != len(visible_sections):
        raise PaperExportRendererError(
            "render_hint_alignment_invalid",
            "题型与共享材料的渲染提示无法与冻结文档计划对齐。",
        )
    return blueprint_sections


def _ordered_shared_materials(
    section: Mapping[str, Any],
    blueprint_section: Mapping[str, Any] | None,
) -> list[Mapping[str, Any]]:
    materials = list(section["shared_materials"])
    keys = [material["render_once_key"] for material in materials]
    if len(keys) != len(set(keys)):
        raise PaperExportRendererError(
            "shared_material_duplicate", "共同材料的冻结标识重复，不能确定唯一展示内容。"
        )
    if blueprint_section is None:
        return materials
    by_key = {material["render_once_key"]: material for material in materials}
    blueprint_materials = list(blueprint_section.get("shared_materials") or [])
    blueprint_keys = [row.get("render_once_key") for row in blueprint_materials]
    if (
        any(not isinstance(key, str) for key in blueprint_keys)
        or len(blueprint_keys) != len(set(blueprint_keys))
        or set(blueprint_keys) != set(by_key)
    ):
        raise PaperExportRendererError(
            "render_hint_alignment_invalid",
            "共享材料的原卷顺序无法与冻结文档计划对齐。",
        )

    def sort_key(index_and_row: tuple[int, Mapping[str, Any]]) -> tuple[int, int]:
        index, row = index_and_row
        page = row.get("page")
        return (int(page) if isinstance(page, int) and page > 0 else 10**9, index)

    ordered_rows = sorted(enumerate(blueprint_materials), key=sort_key)
    return [by_key[row["render_once_key"]] for _, row in ordered_rows]


def _shared_material_schedule(
    section: Mapping[str, Any],
    blueprint_section: Mapping[str, Any] | None,
) -> dict[int, list[Mapping[str, Any]]]:
    """Place each material at its first explicitly bound printed question.

    Older plans without member IDs retain theme-opening placement. Counts,
    filenames and neighboring material order are not membership evidence.
    """
    materials = _ordered_shared_materials(section, blueprint_section)
    references = {
        row["render_once_key"]: row
        for row in (blueprint_section or {}).get("shared_materials", [])
    }
    positions = {
        atomic["atomic_part_id"]: printed_index
        for printed_index, printed in enumerate(
            (blueprint_section or {}).get("printed_questions", [])
        )
        for atomic in printed.get("atomic_parts", [])
    }
    schedule: dict[int, list[Mapping[str, Any]]] = {}
    seen: set[str] = set()
    for material in materials:
        key = material["render_once_key"]
        if key in seen:
            raise PaperExportRendererError(
                "shared_material_duplicate", "共同材料的冻结标识重复，不能确定首次展示位置。"
            )
        seen.add(key)
        reference = references.get(key, {})
        first_index = 0
        if "used_by_atomic_ids" in reference:
            members = reference["used_by_atomic_ids"]
            if (
                not isinstance(members, list)
                or not members
                or any(not isinstance(value, str) or value not in positions for value in members)
                or len(set(members)) != len(members)
            ):
                raise PaperExportRendererError(
                    "shared_material_membership_invalid",
                    "共同材料的题目归属缺失、重复或不属于本次选题，请重新核对来源绑定。",
                )
            first_index = min(positions[value] for value in members)
            if first_index >= len(section["printed_questions"]):
                raise PaperExportRendererError(
                    "render_hint_alignment_invalid", "共同材料的首次使用位置与文档题目不一致。"
                )
        schedule.setdefault(first_index, []).append(material)
    return schedule


def _display_theme_heading(
    section: Mapping[str, Any], blueprint_section: Mapping[str, Any] | None
) -> str:
    """Strip a repeated source ordinal only when the source sequence supports it."""
    heading = section["heading_zh"]
    source = (blueprint_section or {}).get("source", {})
    title, sequence = source.get("theme_title"), source.get("source_theme_sequence")
    if not isinstance(title, str) or type(sequence) is not int or not 1 <= sequence <= 99:
        return heading
    prefix = re.escape(_chinese_number(sequence))
    cleaned = re.sub(r"^\s*" + prefix + r"(?:\s*[、．.]\s*|\s+)", "", title, count=1)
    expected_heading = f"{_chinese_number(section['theme_number'])}、{title}"
    if cleaned.strip() and cleaned != title and heading == expected_heading:
        return f"{_chinese_number(section['theme_number'])}、{cleaned}"
    return heading


def _source_question_number(
    printed: Mapping[str, Any], blueprint_printed: Mapping[str, Any] | None
) -> str | None:
    parts = printed["atomic_parts"]
    if not parts or not all(_is_source_image_question(part) for part in parts):
        return None
    number = (blueprint_printed or {}).get("source_number")
    if type(number) is int and number > 0:
        return str(number)
    # Preserve explicit source labels such as 2a or (3); never parse IDs or
    # turn a missing printed label into a guessed theme-local number.
    if isinstance(number, str) and number.strip() and number.strip() not in {
        "unknown", "unassigned", "待核验", "待补",
    } and len(number) <= 40 and not any(
        token in number for token in ("\n", "\r", "\t")
    ):
        return number.strip()
    return None


def _item_types_for_theme(
    section: Mapping[str, Any],
    blueprint_section: Mapping[str, Any] | None,
) -> list[str | None]:
    visible_count = sum(
        len(printed["atomic_parts"])
        for printed in section["printed_questions"]
    )
    if blueprint_section is None:
        return [None] * visible_count
    item_types = [
        str(atomic.get("item_type") or "unknown")
        for printed in blueprint_section.get("printed_questions") or []
        for atomic in printed.get("atomic_parts") or []
    ]
    if len(item_types) != visible_count:
        raise PaperExportRendererError(
            "render_hint_alignment_invalid",
            "题型数量无法与冻结文档计划对齐。",
        )
    return item_types


def _printed_question_blocks(
    atomic_parts: Sequence[Mapping[str, Any]],
) -> tuple[list[list[Mapping[str, Any]]], bool]:
    """Print an exact asset once per printed question, never merge answer units.

    Only references already rendered by an earlier atomic part are suppressed.
    Text, distinct assets and intentional repeats inside one part are preserved.
    The frozen plan and its asset bytes are not modified.
    """
    seen: set[str] = set()
    projected: list[list[Mapping[str, Any]]] = []
    suppressed = False
    for atomic in atomic_parts:
        blocks: list[Mapping[str, Any]] = []
        current: set[str] = set()
        for block in atomic["question_blocks"]:
            ref = block.get("asset_ref")
            is_image = block.get("block_type") in _SOURCE_IMAGE_BLOCK_TYPES
            if is_image and isinstance(ref, str) and ref:
                current.add(ref)
                if ref in seen:
                    suppressed = True
                    continue
            blocks.append(block)
        projected.append(blocks)
        seen.update(current)
    return projected, suppressed


def _teacher_scoring_label(
    printed: Mapping[str, Any], atomic_index: int, *, source_number: str | None = None
) -> str:
    parts = printed["atomic_parts"]
    atomic = parts[atomic_index]
    prefix = f"第{source_number or printed['question_number']}题"
    if len(parts) > 1:
        if atomic_index == 0:
            prefix += f"（本题共 {_format_score(sum(part['score'] for part in parts))} 分）"
        prefix += f"{atomic.get('part_label_zh') or f'（{atomic_index + 1}）'}"
    return prefix + f"评分（{_format_score(atomic['score'])} 分）"


def _add_theme_sections(
    doc: DocumentType,
    visible: Mapping[str, Any],
    *,
    audience: str,
    asset_root: Path | None,
    content_width_dxa: int,
    body_size_pt: float,
    blueprint: Mapping[str, Any] | None = None,
    suppressed_shared_keys: set[str] | frozenset[str] = frozenset(),
    show_question_scores: bool = True,
) -> None:
    seen_materials: set[str] = set()
    blueprint_sections = _blueprint_theme_sections(visible, blueprint)
    for section_index, section in enumerate(visible["theme_sections"]):
        blueprint_section = blueprint_sections[section_index]
        item_types = iter(_item_types_for_theme(section, blueprint_section))
        material_schedule = _shared_material_schedule(section, blueprint_section)
        blueprint_printed = list((blueprint_section or {}).get("printed_questions", []))
        if blueprint_section is not None and len(blueprint_printed) != len(section["printed_questions"]):
            raise PaperExportRendererError(
                "render_hint_alignment_invalid", "来源题号与文档题目数量不一致。"
            )
        seen_teacher_sources: set[tuple[str, str]] = set()
        heading = doc.add_paragraph(style="ShChemThemeHeading")
        heading.add_run(_display_theme_heading(section, blueprint_section) + (f"　（共 {_format_score(section['theme_score'])} 分）" if show_question_scores else ""))
        summary = str(section.get("context_summary_zh") or "")
        internal_placeholder = (
            summary.startswith("以卷面主题")
            and summary.endswith(("具体化学语义标签仍待入库。", "具体化学语义标签仍待人审。"))
        )
        if summary and not internal_placeholder:
            context = doc.add_paragraph(style="ShChemQuiet")
            context.paragraph_format.keep_with_next = True
            context.add_run("主题情境：" + summary)
        for printed_index, printed in enumerate(section["printed_questions"]):
            placed_materials = [
                material for material in material_schedule.get(printed_index, [])
                if material["render_once_key"] not in suppressed_shared_keys
            ]
            for material_index, material in enumerate(placed_materials):
                key = material["render_once_key"]
                if key in seen_materials:
                    raise PaperExportRendererError(
                        "shared_material_duplicate", "共同材料被重复渲染。"
                    )
                seen_materials.add(key)
                _add_shared_material(
                    doc, material, audience=audience, asset_root=asset_root,
                    content_width_dxa=content_width_dxa, body_size_pt=body_size_pt,
                    keep_last_with_next=material_index == len(placed_materials) - 1,
                )
            source_number = _source_question_number(
                printed, blueprint_printed[printed_index] if blueprint_printed else None
            )
            display_number = source_number or printed["question_number"]
            atomic_parts = printed["atomic_parts"]
            printed_blocks, _shared_question_image = _printed_question_blocks(atomic_parts)
            compact_source_scores = (
                show_question_scores
                and audience == "student"
                and len(atomic_parts) > 1
                and all(
                    _is_source_image_question(part)
                    and part["answer_space"]["lines"] == 0
                    for part in atomic_parts
                )
            )
            for atomic_index, atomic in enumerate(atomic_parts):
                question_blocks = printed_blocks[atomic_index]
                item_type = next(item_types)
                if compact_source_scores and atomic_index > 0:
                    # The original image contains the answer positions. Keep all
                    # scores together above it, not orphan labels on later pages.
                    # Distinct continuation images remain in their source order.
                    _add_content_blocks(doc, question_blocks, asset_root=asset_root)
                    continue
                answer_lines = (
                    atomic["answer_space"]["lines"]
                    if atomic["answer_space"]["mode"] == "ruled_lines_exact"
                    else _effective_answer_line_count(
                        atomic["answer_space"]["lines"], item_type,
                    )
                )
                rendered_answer_lines = answer_lines if audience == "student" else 0
                source_image_question = _is_source_image_question(atomic)
                question = None
                if show_question_scores or not source_image_question:
                    question = doc.add_paragraph(style="ShChemQuiet" if source_image_question else "ShChemQuestion")
                    question.paragraph_format.space_before = Pt(1 if source_image_question else 3)
                    question.paragraph_format.space_after = Pt(0)
                    question.paragraph_format.keep_with_next = True
                    if compact_source_scores:
                        question.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        prefix = f"第{display_number}题（共 {_format_score(sum(part['score'] for part in atomic_parts))} 分）"
                    elif source_image_question:
                        question.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        prefix = (
                            f"第{display_number}题{atomic.get('part_label_zh') or f'（{atomic_index + 1}）'}"
                            if len(atomic_parts) > 1 else ""
                        ) + f"（{_format_score(atomic['score'])} 分）"
                    else:
                        prefix = f"{printed['question_number']}. " if atomic_index == 0 else ""
                        if atomic.get("part_label_zh"):
                            prefix += atomic["part_label_zh"] + " "
                        if show_question_scores:
                            prefix += f"（{_format_score(atomic['score'])} 分）"
                    prefix_run = question.add_run(prefix)
                    _set_run_font(prefix_run, cjk="SimSun" if source_image_question else "SimHei", latin="Times New Roman", size_pt=max(8.0, body_size_pt - 1.5) if source_image_question else body_size_pt, bold=not source_image_question, color=RGBColor(89, 89, 89) if source_image_question else None)
                    if question_blocks:
                        question.add_run(" ")
                _add_content_blocks(
                    doc,
                    question_blocks,
                    first_paragraph=question,
                    asset_root=asset_root,
                    source_question_number=(
                        int(printed["question_number"])
                        if source_image_question and atomic_index == 0
                        else None
                    ),
                    keep_last_with_next=(
                        rendered_answer_lines > 0 or audience == "teacher"
                    ),
                )
                _add_answer_space(
                    doc,
                    rendered_answer_lines,
                    content_width_dxa=content_width_dxa,
                    body_size_pt=body_size_pt,
                )
                if audience == "teacher":
                    source_signature = _teacher_source_signature(
                        atomic["teacher_notes"]
                    )
                    include_source_label = source_signature not in seen_teacher_sources
                    seen_teacher_sources.add(source_signature)
                    _add_teacher_notes(
                        doc,
                        atomic["teacher_notes"],
                        content_width_dxa=content_width_dxa,
                        include_source_label=include_source_label,
                        asset_root=asset_root,
                        scoring_label_zh=_teacher_scoring_label(
                            printed, atomic_index, source_number=source_number
                        ),
                    )


def build_docx_from_plan(
    plan: Mapping[str, Any],
    *,
    preset: Mapping[str, Any],
    output_path: Path,
    asset_root: Path | None = None,
    blueprint: Mapping[str, Any] | None = None,
    suppressed_shared_keys: set[str] | frozenset[str] | None = None,
) -> None:
    audience = plan["audience"]
    if audience not in {"student", "teacher"}:
        raise PaperExportRendererError("audience_invalid", "文档受众不正确。")
    doc = Document()
    _style_document(doc, preset)
    visible = plan["visible"]
    _configure_section(doc, preset, visible)
    _add_title_block(doc, visible, audience, preset=preset)
    if suppressed_shared_keys is None:
        suppressed_shared_keys = set(
            _shared_question_visual_dedup(plan, asset_root)[
                "suppressed_shared_keys"
            ]
        )
    _add_theme_sections(
        doc,
        visible,
        audience=audience,
        asset_root=asset_root,
        content_width_dxa=_content_width_dxa(preset),
        body_size_pt=float(preset["page_layout"]["body"]["size_pt"]),
        blueprint=blueprint,
        suppressed_shared_keys=suppressed_shared_keys,
        show_question_scores=preset["student_version"]["show_item_scores"],
    )

    is_local_basket_export = "题目来源见教师版" in str(
        visible.get("version_label_zh") or ""
    )
    closing_text = (
        "—— 本地题篮导出结束｜题目来源见教师版｜不可发布 ——"
        if is_local_basket_export
        else "—— 合成布局演示结束｜非上海原题｜不可发布 ——"
    )
    closing = doc.add_paragraph(style="ShChemQuiet")
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    closing.paragraph_format.space_before = Pt(4)
    closing.add_run(closing_text)

    properties = doc.core_properties
    properties.title = visible["paper_title_zh"]
    properties.subject = (
        "上海高中化学主题式题篮导出"
        if is_local_basket_export
        else "上海高中化学主题式试卷导出布局演示"
    )
    properties.author = ""
    properties.last_modified_by = ""
    properties.keywords = ""
    properties.comments = (
        "本地题篮导出；题目来源见教师版；不可发布。"
        if is_local_basket_export
        else "合成布局演示；非上海原题；不可发布。"
    )
    properties.category = (
        "local-prep-basket-export" if is_local_basket_export else "internal-layout-smoke"
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def _iter_table_text(table: Any) -> list[str]:
    result: list[str] = []
    for row in table.rows:
        for cell in row.cells:
            result.extend(paragraph.text for paragraph in cell.paragraphs)
            for nested in cell.tables:
                result.extend(_iter_table_text(nested))
    return result


def _docx_body_text(path: Path) -> str:
    doc = Document(path)
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        parts.extend(_iter_table_text(table))
    return "\n".join(parts)


def _docx_header_text(path: Path) -> str:
    doc = Document(path)
    return "\n".join(
        paragraph.text
        for section in doc.sections
        for paragraph in section.header.paragraphs
    )


def _docx_first_header_text(path: Path) -> str:
    doc = Document(path)
    return "\n".join(
        paragraph.text
        for section in doc.sections
        for paragraph in section.first_page_header.paragraphs
    )


def _docx_footer_xml(path: Path) -> str:
    with ZipFile(path) as archive:
        return "\n".join(
            archive.read(name).decode("utf-8", errors="replace")
            for name in archive.namelist()
            if name.startswith("word/footer") and name.endswith(".xml")
        )


def _docx_figure_count(path: Path) -> int:
    with ZipFile(path) as archive:
        return len(
            {
                name
                for name in archive.namelist()
                if name.startswith("word/media/") and not name.endswith("/")
            }
        )


def _docx_image_alt_texts(path: Path) -> list[str]:
    with ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml").decode(
            "utf-8", errors="replace"
        )
    return re.findall(r'<wp:docPr\b[^>]*\bdescr="([^"]*)"', document_xml)


def _normalize_text(value: str) -> str:
    return re.sub(r"\s+", "", value).replace(" ", "")


def _secret_hits(text: str) -> list[str]:
    return [pattern.pattern for pattern in _SECRET_PATTERNS if pattern.search(text)]


def _expected_material_texts(plan: Mapping[str, Any]) -> list[str]:
    texts: list[str] = []
    for section in plan["visible"]["theme_sections"]:
        for material in section["shared_materials"]:
            for block in material["content_blocks"]:
                if block.get("text_zh"):
                    texts.append(block["text_zh"])
    return texts


def _expected_question_number_groups(
    plan: Mapping[str, Any],
) -> tuple[list[int], list[int]]:
    text_numbers: list[int] = []
    source_image_numbers: list[int] = []
    for section in plan["visible"]["theme_sections"]:
        for printed in section["printed_questions"]:
            number = int(printed["question_number"])
            atomic_parts = printed["atomic_parts"]
            if atomic_parts and _is_source_image_question(atomic_parts[0]):
                source_image_numbers.append(number)
            else:
                text_numbers.append(number)
    return text_numbers, source_image_numbers


def _teacher_answer_boundaries_present(text: str, plan: Mapping[str, Any]) -> bool:
    """Check every expected answer-state label, including absent answers.

    Requiring a nonofficial-answer label on an all-missing-answer theme falsely
    rejected honest teacher copies. The frozen plan determines which labels and
    how many are required; a single generic disclaimer cannot satisfy the check.
    """
    expected: dict[str, int] = {}
    for theme in plan["visible"]["theme_sections"]:
        for printed in theme["printed_questions"]:
            for atomic in printed["atomic_parts"]:
                label = _normalize_text(atomic["teacher_notes"]["answer_label_zh"] + "：")
                expected[label] = expected.get(label, 0) + 1
    normalized = _normalize_text(text)
    return bool(expected) and all(
        normalized.count(label) >= count for label, count in expected.items()
    )


def _score_labels_present(
    text: str, plan: Mapping[str, Any], show_question_scores: bool,
    blueprint: Mapping[str, Any] | None = None,
) -> bool:
    """Check requested score labels; teacher answers always retain every score."""
    audience = plan["audience"]
    if audience == "student" and not show_question_scores:
        return True  # No score label is requested in student question content.
    expected: dict[str, int] = {}
    blueprint_sections = _blueprint_theme_sections(plan["visible"], blueprint)
    for theme_index, theme in enumerate(plan["visible"]["theme_sections"]):
        blueprint_printed = list((blueprint_sections[theme_index] or {}).get("printed_questions", []))
        for printed_index, printed in enumerate(theme["printed_questions"]):
            parts = printed["atomic_parts"]
            if audience == "teacher":
                source_number = _source_question_number(
                    printed,
                    blueprint_printed[printed_index] if printed_index < len(blueprint_printed) else None,
                )
                labels = [_teacher_scoring_label(printed, index, source_number=source_number) for index in range(len(parts))]
            elif len(parts) > 1 and all(_is_source_image_question(part) and part["answer_space"]["lines"] == 0 for part in parts):
                labels = [f"（共 {_format_score(sum(part['score'] for part in parts))} 分）"]
            else:
                labels = [f"（{_format_score(part['score'])} 分）" for part in parts]
            for label in labels:
                normalized_label = _normalize_text(label)
                expected[normalized_label] = expected.get(normalized_label, 0) + 1
    normalized = _normalize_text(text)
    return bool(expected) and all(normalized.count(label) >= count for label, count in expected.items())


def _teacher_source_and_pitfalls_present(text: str, plan: Mapping[str, Any]) -> bool:
    """Require recorded pitfalls, not invented notes for an empty source field."""
    normalized = _normalize_text(text)
    if _normalize_text("来源标签：") not in normalized:
        return False
    expected: dict[str, int] = {}
    note_count = 0
    for theme in plan["visible"]["theme_sections"]:
        for printed in theme["printed_questions"]:
            for atomic in printed["atomic_parts"]:
                pitfalls = atomic["teacher_notes"].get("pitfalls_zh") or []
                if pitfalls:
                    note_count += 1
                for value in pitfalls:
                    note = _normalize_text(value)
                    expected[note] = expected.get(note, 0) + 1
    return (
        normalized.count(_normalize_text("易错点：")) >= note_count
        and all(normalized.count(note) >= count for note, count in expected.items())
    )


def audit_docx(
    path: Path,
    *,
    plan: Mapping[str, Any],
    preset: Mapping[str, Any],
    blueprint: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    audience = plan["audience"]
    show_scores = preset["student_version"]["show_item_scores"]
    doc = Document(path)
    section = doc.sections[0]
    body_text = _docx_body_text(path)
    normalized = _normalize_text(body_text)
    materials = _expected_material_texts(plan)
    expected_margins = preset["page_layout"]["margins_mm"]
    expected_title_alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[preset["page_layout"]["title"]["alignment"]]
    title_style = doc.styles["ShChemPaperTitle"]
    theme_style = doc.styles["ShChemThemeHeading"]
    question_style = doc.styles["ShChemQuestion"]
    expected_body_size = float(preset["page_layout"]["body"]["size_pt"])
    footer_xml = _docx_footer_xml(path)
    text_question_numbers, source_image_numbers = _expected_question_number_groups(
        plan
    )
    image_alt_texts = _docx_image_alt_texts(path)
    checks = {
        "a4_portrait": abs(section.page_width.mm - 210) < 0.2
        and abs(section.page_height.mm - 297) < 0.2,
        "margins_match_preset": all(
            abs(actual.mm - float(expected_margins[key])) < 0.2
            for key, actual in (
                ("top", section.top_margin),
                ("bottom", section.bottom_margin),
                ("left", section.left_margin),
                ("right", section.right_margin),
            )
        ),
        "title_style_matches_preset": abs(
            title_style.font.size.pt - float(preset["page_layout"]["title"]["size_pt"])
        )
        < 0.05
        and bool(title_style.font.bold) == bool(preset["page_layout"]["title"]["bold"])
        and doc.paragraphs[0].alignment == expected_title_alignment,
        "theme_style_matches_preset": abs(
            theme_style.font.size.pt
            - float(preset["page_layout"]["theme_heading"]["size_pt"])
        )
        < 0.05
        and bool(theme_style.font.bold)
        == bool(preset["page_layout"]["theme_heading"]["bold"]),
        "body_style_matches_preset": abs(question_style.font.size.pt - expected_body_size)
        < 0.05
        and abs(doc.styles["Normal"].font.size.pt - expected_body_size) < 0.05,
        "header_matches_contract": _docx_header_text(path).strip()
        == plan["visible"]["header"]["text_zh"],
        "first_page_header_blank": not _docx_first_header_text(path).strip(),
        "footer_page_fields_present": bool(
            re.search(r"<w:instrText[^>]*>\s*PAGE\s*</w:instrText>", footer_xml)
        )
        and bool(
            re.search(r"<w:instrText[^>]*>\s*NUMPAGES\s*</w:instrText>", footer_xml)
        ),
        "shared_material_once": all(
            normalized.count(_normalize_text(text)) == 1 for text in materials
        ),
        "question_numbers_present": all(
            (f"{number}.（" if show_scores else f"{number}.") in normalized for number in text_question_numbers
        )
        and all(
            any(f"第{number}题题图：" in alt for alt in image_alt_texts)
            for number in source_image_numbers
        ),
        "source_image_numbers_not_repeated_as_text_headings": all(
            f"{number}.（" not in normalized for number in source_image_numbers
        ),
        "scores_present": _score_labels_present(body_text, plan, show_scores, blueprint),
        "identity_scope_correct": (
            "姓名：" in body_text and "班级：" in body_text
            if audience == "student"
            else "姓名：" not in body_text and "班级：" not in body_text
        ),
        "student_answer_markers_absent": (
            all(marker not in body_text for marker in _STUDENT_FORBIDDEN_MARKERS)
            if audience == "student"
            else True
        ),
        "teacher_nonofficial_label_present": (
            _teacher_answer_boundaries_present(body_text, plan)
            if audience == "teacher"
            else True
        ),
        "teacher_source_and_pitfalls_present": (
            _teacher_source_and_pitfalls_present(body_text, plan)
            if audience == "teacher"
            else True
        ),
        "secret_scan_pass": not _secret_hits(body_text),
        "metadata_privacy_pass": not doc.core_properties.author
        and not doc.core_properties.last_modified_by,
        "publication_boundary_visible": (
            "不可发布" in body_text
            and any(marker in body_text for marker in ("非上海原题", "题目来源见教师版"))
        ),
    }
    return {
        "status": "pass" if all(checks.values()) else "blocked",
        "checks": checks,
        "figure_count": _docx_figure_count(path),
        "body_text_sha256": hashlib.sha256(body_text.encode("utf-8")).hexdigest(),
    }


def _fresh_directory(path: Path, output_root: Path) -> None:
    resolved = path.resolve()
    try:
        resolved.relative_to(output_root.resolve())
    except ValueError as exc:
        raise PaperExportRendererError("output_path_invalid", "渲染输出目录超出允许范围。") from exc
    if path.exists():
        shutil.rmtree(path)
    path.mkdir(parents=True, exist_ok=True)


def _run_command(command: Sequence[str], *, env: Mapping[str, str] | None = None, timeout: int = 180) -> subprocess.CompletedProcess[str]:
    completed = subprocess.run(
        list(command),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        env=dict(env) if env is not None else None,
        timeout=timeout,
        check=False,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    if completed.returncode != 0:
        safe_tail = (completed.stderr or completed.stdout)[-1200:]
        raise PaperExportRendererError(
            "renderer_command_failed", "文档渲染工具执行失败：" + safe_tail
        )
    return completed


def _word_com_pdf(docx_path: Path, pdf_path: Path) -> str:
    powershell = r"""
$ErrorActionPreference = 'Stop'
$word = $null
$doc = $null
try {
  $inputPath = [Environment]::GetEnvironmentVariable('SHCHEM_WORD_INPUT', 'Process')
  $outputPath = [Environment]::GetEnvironmentVariable('SHCHEM_WORD_OUTPUT', 'Process')
  $word = New-Object -ComObject Word.Application
  $word.Visible = $false
  $word.DisplayAlerts = 0
  $doc = $word.Documents.Open($inputPath, $false, $true)
  $doc.ExportAsFixedFormat($outputPath, 17)
} finally {
  if ($null -ne $doc) { $doc.Close(0) }
  if ($null -ne $word) { $word.Quit() }
  if ($null -ne $doc) { [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($doc) }
  if ($null -ne $word) { [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($word) }
  [GC]::Collect()
  [GC]::WaitForPendingFinalizers()
}
"""
    encoded = base64.b64encode(powershell.encode("utf-16le")).decode("ascii")
    env = os.environ.copy()
    env["SHCHEM_WORD_INPUT"] = str(docx_path.resolve())
    env["SHCHEM_WORD_OUTPUT"] = str(pdf_path.resolve())
    _run_command(
        [
            "powershell.exe",
            "-STA",
            "-NoProfile",
            "-NonInteractive",
            "-EncodedCommand",
            encoded,
        ],
        env=env,
        timeout=180,
    )
    if not pdf_path.is_file() or pdf_path.stat().st_size == 0:
        raise PaperExportRendererError(
            "word_pdf_missing", "Word 未生成可用的 PDF 文件。"
        )
    return "Word COM 只读打开 DOCX 并导出 PDF。"


def _load_render_docx_module(script_path: Path) -> Any:
    module_name = "_codex_render_docx_" + _sha256_file(script_path)[:12]
    spec = importlib.util.spec_from_file_location(module_name, script_path)
    if spec is None or spec.loader is None:
        raise PaperExportRendererError(
            "render_docx_import_failed", "无法加载标准 render_docx.py。"
        )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _render_with_canonical_docx_tool(
    docx_path: Path,
    output_dir: Path,
    *,
    toolchain: RendererToolchain,
) -> tuple[list[Path], Path, dict[str, Any]]:
    _fresh_directory(output_dir, output_dir.parent.parent)
    backend = toolchain.conversion_backend
    if backend == "auto":
        backend = "libreoffice" if shutil.which("soffice") else "word_com"
    script_sha = _sha256_file(toolchain.render_docx_script)
    if backend == "libreoffice":
        _run_command(
            [
                str(toolchain.python_exe),
                "-B",
                str(toolchain.render_docx_script),
                str(docx_path),
                "--output_dir",
                str(output_dir),
                "--dpi",
                str(toolchain.dpi),
                "--emit_pdf",
                "--verbose",
            ],
            timeout=240,
        )
        invocation_mode = "canonical_cli"
    else:
        # The Gateway process intentionally has a smaller dependency surface
        # than the bundled Documents runtime.  Export the PDF through local
        # Word COM here, then run the canonical render_docx.py inside its
        # bundled Python where pdf2image/Poppler are guaranteed to exist.
        preconverted_pdf = output_dir / f".{docx_path.stem}.word-com.pdf"
        _word_com_pdf(docx_path, preconverted_pdf)
        wrapper = r"""
import importlib.util
import os
import shutil
from pathlib import Path

script = Path(os.environ["SHCHEM_RENDER_SCRIPT"])
spec = importlib.util.spec_from_file_location("_shchem_canonical_render_docx", script)
if spec is None or spec.loader is None:
    raise RuntimeError("canonical render_docx.py import failed")
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)

def use_preconverted_pdf(input_path, user_profile, convert_tmp_dir, stem, verbose=False):
    del input_path, user_profile, verbose
    target = Path(convert_tmp_dir) / f"{stem}.pdf"
    shutil.copy2(Path(os.environ["SHCHEM_PRECONVERTED_PDF"]), target)
    return str(target), "Microsoft Word COM pre-conversion supplied by the local Gateway."

module.convert_to_pdf = use_preconverted_pdf
module.rasterize(
    os.environ["SHCHEM_DOCX_INPUT"],
    os.environ["SHCHEM_RENDER_OUTPUT"],
    int(os.environ["SHCHEM_RENDER_DPI"]),
    verbose=False,
    emit_pdf=True,
)
"""
        env = os.environ.copy()
        env.update(
            {
                "SHCHEM_RENDER_SCRIPT": str(toolchain.render_docx_script),
                "SHCHEM_PRECONVERTED_PDF": str(preconverted_pdf),
                "SHCHEM_DOCX_INPUT": str(docx_path),
                "SHCHEM_RENDER_OUTPUT": str(output_dir),
                "SHCHEM_RENDER_DPI": str(toolchain.dpi),
                "PATH": str(toolchain.pdftoppm_exe.parent)
                + os.pathsep
                + env.get("PATH", ""),
            }
        )
        try:
            _run_command(
                [str(toolchain.python_exe), "-B", "-X", "utf8", "-c", wrapper],
                env=env,
                timeout=240,
            )
        finally:
            if preconverted_pdf.is_file():
                preconverted_pdf.unlink()
        invocation_mode = "canonical_module_with_word_com_conversion"

    pages = sorted(output_dir.glob("page-*.png"), key=_page_number)
    emitted_pdf = output_dir / f"{docx_path.stem}.pdf"
    if not pages or not emitted_pdf.is_file():
        raise PaperExportRendererError(
            "render_docx_output_missing", "标准文档渲染器没有生成完整逐页图片和 PDF。"
        )
    return pages, emitted_pdf, {
        "tool": "render_docx.py",
        "script_sha256": script_sha,
        "invocation_mode": invocation_mode,
        "conversion_backend": backend,
        "dpi": toolchain.dpi,
    }


def _page_number(path: Path) -> int:
    match = re.search(r"(\d+)(?=\.png$)", path.name)
    return int(match.group(1)) if match else 0


def _render_pdf_with_poppler(
    pdf_path: Path,
    output_dir: Path,
    *,
    toolchain: RendererToolchain,
) -> tuple[list[Path], dict[str, Any]]:
    _fresh_directory(output_dir, output_dir.parent.parent)
    prefix = output_dir / "page"
    completed = _run_command(
        [
            str(toolchain.pdftoppm_exe),
            "-png",
            "-r",
            str(toolchain.dpi),
            str(pdf_path),
            str(prefix),
        ],
        timeout=180,
    )
    pages = sorted(output_dir.glob("page-*.png"), key=_page_number)
    if not pages:
        raise PaperExportRendererError("poppler_output_missing", "Poppler 没有生成逐页图片。")
    return pages, {
        "tool": "pdftoppm",
        "executable_sha256": _sha256_file(toolchain.pdftoppm_exe),
        "dpi": toolchain.dpi,
        "stderr_sha256": hashlib.sha256(completed.stderr.encode("utf-8")).hexdigest(),
    }


def _image_records(paths: Sequence[Path], output_root: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for path in paths:
        with Image.open(path) as image:
            width, height = image.size
        records.append(
            {
                "path": _relative(path, output_root),
                "page_number": _page_number(path),
                "width_px": width,
                "height_px": height,
                "sha256": _sha256_file(path),
            }
        )
    return records


def _pdf_text(path: Path) -> str:
    return "\n".join((page.extract_text() or "") for page in PdfReader(path).pages)


def _pdf_page_count(path: Path) -> int:
    return len(PdfReader(path).pages)


def _pdf_text_audit(
    path: Path, *, plan: Mapping[str, Any], preset: Mapping[str, Any] | None = None,
    blueprint: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    audience = plan["audience"]
    show_scores = preset["student_version"]["show_item_scores"] if preset else True
    text = _pdf_text(path)
    normalized = _normalize_text(text)
    materials = _expected_material_texts(plan)
    text_question_numbers, source_image_numbers = _expected_question_number_groups(
        plan
    )
    blueprint_sections = _blueprint_theme_sections(plan["visible"], blueprint)
    checks = {
        "title_present": _normalize_text(plan["visible"]["paper_title_zh"]) in normalized,
        "theme_headings_present": all(
            _normalize_text(_display_theme_heading(theme, blueprint_sections[index])) in normalized
            for index, theme in enumerate(plan["visible"]["theme_sections"])
        ),
        "shared_material_once": all(
            normalized.count(_normalize_text(material)) == 1 for material in materials
        ),
        "question_numbers_present": all(
            (f"{number}.（" if show_scores else f"{number}.") in normalized for number in text_question_numbers
        ),
        "source_image_numbers_not_repeated_as_text_headings": all(
            f"{number}.（" not in normalized for number in source_image_numbers
        ),
        "scores_present": _score_labels_present(text, plan, show_scores, blueprint),
        "identity_scope_correct": (
            "姓名：" in text and "班级：" in text
            if audience == "student"
            else "姓名：" not in text and "班级：" not in text
        ),
        "student_answer_markers_absent": (
            all(marker not in text for marker in _STUDENT_FORBIDDEN_MARKERS)
            if audience == "student"
            else True
        ),
        "teacher_nonofficial_label_present": (
            _teacher_answer_boundaries_present(text, plan)
            if audience == "teacher"
            else True
        ),
        "teacher_source_and_pitfalls_present": (
            _teacher_source_and_pitfalls_present(text, plan)
            if audience == "teacher"
            else True
        ),
        "secret_scan_pass": not _secret_hits(text),
        "publication_boundary_visible": (
            "不可发布" in text
            and any(marker in text for marker in ("非上海原题", "题目来源见教师版"))
        ),
    }
    return {
        "status": "pass" if all(checks.values()) else "blocked",
        "checks": checks,
        "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
    }


def _clean_known_outputs(output_root: Path) -> None:
    output_root.mkdir(parents=True, exist_ok=True)
    for name in ARTIFACT_FILENAMES.values():
        path = output_root / name
        if path.is_file():
            path.unlink()
    for name in ("render_bundle.json", "render_qa_report.json", "hash_manifest.json"):
        path = output_root / name
        if path.is_file():
            path.unlink()
    qa = output_root / "qa"
    if qa.is_dir():
        shutil.rmtree(qa)


def _machine_check(check_id: str, passed: bool, message_zh: str) -> dict[str, Any]:
    return {
        "check_id": check_id,
        "status": "pass" if passed else "blocked",
        "message_zh": message_zh,
    }


def _write_hash_manifest(output_root: Path) -> dict[str, Any]:
    manifest_path = output_root / "hash_manifest.json"
    files: list[dict[str, Any]] = []
    for path in sorted(output_root.rglob("*")):
        if not path.is_file() or path.resolve() == manifest_path.resolve():
            continue
        files.append(
            {
                "path": _relative(path, output_root),
                "sha256": _sha256_file(path),
                "size_bytes": path.stat().st_size,
            }
        )
    manifest = {
        "renderer_schema_version": RENDERER_SCHEMA_VERSION,
        "contract_kind": HASH_MANIFEST_KIND,
        "publication_allowed": False,
        "file_count": len(files),
        "files": files,
    }
    _write_json(manifest_path, manifest)
    return manifest


def render_export_bundle(
    bundle: Mapping[str, Any],
    *,
    output_dir: Path,
    toolchain: RendererToolchain,
    asset_root: Path | None = None,
) -> dict[str, Any]:
    frozen = validate_render_bundle(bundle)
    toolchain = toolchain.validated()
    minimum_dpi = int(frozen["preset"]["page_layout"]["minimum_raster_dpi"])
    if toolchain.dpi < minimum_dpi:
        raise PaperExportRendererError(
            "renderer_dpi_below_preset",
            f"逐页渲染清晰度低于当前模板要求的 {minimum_dpi} dpi。",
        )
    output_root = output_dir.resolve()
    _clean_known_outputs(output_root)
    _write_json(output_root / "render_bundle.json", frozen)
    shared_question_dedup = _shared_question_visual_dedup(
        frozen["student_plan"],
        asset_root,
    )
    automated_visual_notes = _source_asset_edge_review_notes(
        frozen["student_plan"],
        asset_root,
    )
    automated_visual_notes.extend(
        _shared_question_dedup_review_notes(shared_question_dedup)
    )

    artifact_records: list[dict[str, Any]] = []
    machine_checks: list[dict[str, Any]] = []
    qa_root = output_root / "qa"
    for audience in ("student", "teacher"):
        plan = frozen[f"{audience}_plan"]
        docx_id = f"{audience}_docx"
        pdf_id = f"{audience}_pdf"
        docx_path = output_root / ARTIFACT_FILENAMES[docx_id]
        pdf_path = output_root / ARTIFACT_FILENAMES[pdf_id]
        build_docx_from_plan(
            plan,
            preset=frozen["preset"],
            output_path=docx_path,
            asset_root=asset_root,
            blueprint=frozen["blueprint"],
            suppressed_shared_keys=set(
                shared_question_dedup["suppressed_shared_keys"]
            ),
        )
        docx_audit = audit_docx(
            docx_path, plan=plan, preset=frozen["preset"], blueprint=frozen["blueprint"]
        )
        docx_pages, emitted_pdf, docx_tool = _render_with_canonical_docx_tool(
            docx_path,
            qa_root / f"{audience}_docx_pages",
            toolchain=toolchain,
        )
        shutil.copy2(emitted_pdf, pdf_path)
        pdf_pages, pdf_tool = _render_pdf_with_poppler(
            pdf_path,
            qa_root / f"{audience}_pdf_pages",
            toolchain=toolchain,
        )
        pdf_count = _pdf_page_count(pdf_path)
        pdf_audit = _pdf_text_audit(
            pdf_path, plan=plan, preset=frozen["preset"], blueprint=frozen["blueprint"]
        )
        docx_page_records = _image_records(docx_pages, output_root)
        pdf_page_records = _image_records(pdf_pages, output_root)
        page_dimensions_match = [
            (item["width_px"], item["height_px"]) for item in docx_page_records
        ] == [(item["width_px"], item["height_px"]) for item in pdf_page_records]
        page_counts_match = len(docx_pages) == len(pdf_pages) == pdf_count
        machine_checks.extend(
            [
                _machine_check(
                    f"{audience}_docx_structure",
                    docx_audit["status"] == "pass",
                    f"{audience} 版 DOCX 结构、内容与隐私检查通过。",
                ),
                _machine_check(
                    f"{audience}_pdf_text",
                    pdf_audit["status"] == "pass",
                    f"{audience} 版 PDF 题面、共同材料与答案边界检查通过。",
                ),
                _machine_check(
                    f"{audience}_page_count_parity",
                    page_counts_match,
                    f"{audience} 版 DOCX 渲染页与 PDF 逐页渲染页数量一致。",
                ),
                _machine_check(
                    f"{audience}_page_geometry_parity",
                    page_dimensions_match,
                    f"{audience} 版两套逐页图片的尺寸一致。",
                ),
            ]
        )
        common = {
            "audience": audience,
            "page_count": pdf_count,
            "rendered_page_count": len(docx_pages),
            "figure_count": docx_audit["figure_count"],
            "docx_render_pages": docx_page_records,
            "pdf_render_pages": pdf_page_records,
            "docx_audit": docx_audit,
            "pdf_audit": pdf_audit,
            "render_docx_tool": docx_tool,
            "poppler_tool": pdf_tool,
        }
        artifact_records.append(
            {
                "artifact_id": docx_id,
                "format": "docx",
                "path": _relative(docx_path, output_root),
                "sha256": _sha256_file(docx_path),
                "size_bytes": docx_path.stat().st_size,
                **common,
            }
        )
        artifact_records.append(
            {
                "artifact_id": pdf_id,
                "format": "pdf",
                "path": _relative(pdf_path, output_root),
                "sha256": _sha256_file(pdf_path),
                "size_bytes": pdf_path.stat().st_size,
                **common,
            }
        )

    hashes = [item["sha256"] for item in artifact_records]
    machine_checks.append(
        _machine_check(
            "four_artifact_hashes_unique",
            len(hashes) == 4 and len(set(hashes)) == 4,
            "学生版与教师版 DOCX/PDF 四文件均有独立校验值。",
        )
    )
    machine_blockers = [item for item in machine_checks if item["status"] == "blocked"]
    layout = frozen["preset"]["page_layout"]
    margins = layout["margins_mm"]
    qa_report = {
        "renderer_schema_version": RENDERER_SCHEMA_VERSION,
        "contract_kind": QA_REPORT_KIND,
        "renderer_version": RENDERER_VERSION,
        "render_job_id": frozen["render_request"]["render_job_id"],
        "render_request_digest": frozen["render_request"]["render_request_digest"],
        "status": "awaiting_visual_review" if not machine_blockers else "blocked",
        "claim_boundary_zh": frozen["claim_boundary_zh"],
        "publication_allowed": False,
        "style_profile": {
            "profile_id": "shanghai_theme_exam_project_override_v1",
            "authority": "existing_paper_format_contract_and_exam_delivery_spec",
            "page": (
                "A4 portrait; margins mm "
                f"top={margins['top']}, bottom={margins['bottom']}, "
                f"left={margins['left']}, right={margins['right']}"
            ),
            "body": (
                f"{layout['body']['cjk_font']} {layout['body']['size_pt']} pt; "
                f"{layout['body']['latin_font']} Latin; "
                f"{layout['body']['line_spacing']} line spacing"
            ),
            "header_footer": "quiet centered header; 第 x 页 / 共 y 页 footer",
            "named_override_zh": "项目试卷合同覆盖通用文档预设；不声称为考试院官方印刷规范。",
        },
        "artifacts": artifact_records,
        "machine_checks": machine_checks,
        "machine_blocker_count": len(machine_blockers),
        "visual_review": {
            "status": "pending",
            "required_page_count": sum(
                len(item["docx_render_pages"]) + len(item["pdf_render_pages"])
                for item in artifact_records
                if item["format"] == "docx"
            ),
            "reviewed_pages": [],
            "layout_defect_count": None,
            "reviewer_kind": None,
            "human_reviewed": False,
            "chemistry_reviewed": False,
            "additional_review_notes": automated_visual_notes,
        },
        "receipts": [],
        "contract_validation_report": None,
    }
    _write_json(output_root / "render_qa_report.json", qa_report)
    _write_hash_manifest(output_root)
    return qa_report


def _artifact_by_id(report: Mapping[str, Any], artifact_id: str) -> Mapping[str, Any]:
    for artifact in report["artifacts"]:
        if artifact["artifact_id"] == artifact_id:
            return artifact
    raise PaperExportRendererError("artifact_missing", "逐页复核缺少导出文件。")


def _resolve_output_file(output_root: Path, relative_path: str) -> Path:
    if not isinstance(relative_path, str) or not relative_path:
        raise PaperExportRendererError("pending_report_drift", "待复核文件路径不正确。")
    candidate = (output_root / Path(relative_path)).resolve()
    try:
        candidate.relative_to(output_root.resolve())
    except ValueError as exc:
        raise PaperExportRendererError(
            "pending_report_drift", "待复核文件路径超出演示目录。"
        ) from exc
    if not candidate.is_file():
        raise PaperExportRendererError("pending_report_drift", "待复核文件已缺失。")
    return candidate


def _verify_pending_report_integrity(
    output_root: Path,
    report: Mapping[str, Any],
    bundle: Mapping[str, Any],
) -> None:
    request = bundle["render_request"]
    if (
        report.get("renderer_schema_version") != RENDERER_SCHEMA_VERSION
        or report.get("contract_kind") != QA_REPORT_KIND
        or report.get("status") != "awaiting_visual_review"
        or report.get("render_job_id") != request["render_job_id"]
        or report.get("render_request_digest") != request["render_request_digest"]
        or report.get("publication_allowed") is not False
        or report.get("machine_blocker_count") != 0
        or any(item.get("status") != "pass" for item in report.get("machine_checks", []))
    ):
        raise PaperExportRendererError(
            "pending_report_drift", "待复核报告已变化、仍有阻断或未绑定当前渲染请求。"
        )
    artifacts = report.get("artifacts")
    if not isinstance(artifacts, list) or len(artifacts) != 4:
        raise PaperExportRendererError("pending_report_drift", "待复核四文件记录不完整。")
    expected_ids = set(ARTIFACT_FILENAMES)
    actual_ids = {item.get("artifact_id") for item in artifacts if isinstance(item, Mapping)}
    if actual_ids != expected_ids:
        raise PaperExportRendererError("pending_report_drift", "待复核四文件标识不完整。")
    for artifact in artifacts:
        artifact_id = artifact["artifact_id"]
        expected_format = "docx" if artifact_id.endswith("_docx") else "pdf"
        expected_audience = artifact_id.split("_", 1)[0]
        if (
            artifact.get("format") != expected_format
            or artifact.get("audience") != expected_audience
            or artifact.get("path") != ARTIFACT_FILENAMES[artifact_id]
        ):
            raise PaperExportRendererError(
                "pending_report_drift", "待复核文件受众、格式或文件名已变化。"
            )
        artifact_path = _resolve_output_file(output_root, artifact["path"])
        if (
            _sha256_file(artifact_path) != artifact.get("sha256")
            or artifact_path.stat().st_size != artifact.get("size_bytes")
        ):
            raise PaperExportRendererError(
                "pending_artifact_hash_drift", "渲染后文件校验值已变化，必须重新渲染。"
            )
        if expected_format == "pdf" and _pdf_page_count(artifact_path) != artifact.get(
            "page_count"
        ):
            raise PaperExportRendererError(
                "pending_artifact_page_drift", "PDF 页数已变化，必须重新渲染。"
            )
        if artifact.get("docx_audit", {}).get("status") != "pass" or artifact.get(
            "pdf_audit", {}
        ).get("status") != "pass":
            raise PaperExportRendererError(
                "pending_report_drift", "四文件机器检查记录未全部通过。"
            )
        for audit_key in ("docx_audit", "pdf_audit"):
            checks = artifact[audit_key].get("checks")
            if not isinstance(checks, Mapping) or not checks or not all(
                value is True for value in checks.values()
            ):
                raise PaperExportRendererError(
                    "pending_report_drift", "四文件机器检查细项未全部通过。"
                )
        for page_key in ("docx_render_pages", "pdf_render_pages"):
            page_records = artifact.get(page_key)
            if not isinstance(page_records, list) or len(page_records) != artifact.get(
                "page_count"
            ):
                raise PaperExportRendererError(
                    "pending_artifact_page_drift", "逐页图片数量与文件页数不一致。"
                )
            for page_record in page_records:
                page_path = _resolve_output_file(output_root, page_record.get("path"))
                if _sha256_file(page_path) != page_record.get("sha256"):
                    raise PaperExportRendererError(
                        "pending_page_hash_drift", "逐页图片校验值已变化，必须重新渲染。"
                    )
                with Image.open(page_path) as image:
                    dimensions = image.size
                if dimensions != (
                    page_record.get("width_px"),
                    page_record.get("height_px"),
                ):
                    raise PaperExportRendererError(
                        "pending_page_geometry_drift", "逐页图片尺寸已变化，必须重新渲染。"
                    )


def finalize_visual_review(
    *,
    output_dir: Path,
    reviewed_pages: Sequence[Mapping[str, Any]],
    reviewed_at: str | None = None,
    reviewer_kind: str = "codex_visual_layout_review",
    additional_review_notes: Sequence[Mapping[str, Any]] = (),
) -> dict[str, Any]:
    output_root = output_dir.resolve()
    bundle = validate_render_bundle(_read_json(output_root / "render_bundle.json"))
    report = _read_json(output_root / "render_qa_report.json")
    _verify_pending_report_integrity(output_root, report, bundle)
    if report.get("machine_blocker_count") != 0:
        raise PaperExportRendererError(
            "machine_qa_blocked", "机器检查仍有阻断，不能确认逐页视觉通过。"
        )
    expected: set[str] = set()
    for artifact in report["artifacts"]:
        if artifact["format"] != "docx":
            continue
        expected.update(item["path"] for item in artifact["docx_render_pages"])
        expected.update(item["path"] for item in artifact["pdf_render_pages"])
    normalized_reviews: list[dict[str, Any]] = []
    seen: set[str] = set()
    for value in reviewed_pages:
        if set(value) != {"path", "status", "notes_zh"}:
            raise PaperExportRendererError(
                "visual_review_invalid", "逐页目视复核记录字段不正确。"
            )
        path = value["path"]
        if path not in expected or path in seen or value["status"] != "pass":
            raise PaperExportRendererError(
                "visual_review_invalid", "逐页目视复核缺页、重复或未通过。"
            )
        notes = value["notes_zh"]
        if not isinstance(notes, str) or not notes.strip():
            raise PaperExportRendererError(
                "visual_review_invalid", "每页必须记录简短的目视检查结论。"
            )
        seen.add(path)
        normalized_reviews.append(
            {"path": path, "status": "pass", "notes_zh": notes.strip()}
        )
    if seen != expected:
        raise PaperExportRendererError(
            "visual_review_incomplete", "尚未逐页查看全部 DOCX 与 PDF 渲染图片。"
        )
    normalized_reviews.sort(key=lambda item: item["path"])
    normalized_additional_notes: list[dict[str, Any]] = []
    note_keys = {
        "reviewer_kind",
        "note_zh",
        "scope_zh",
        "human_reviewed",
        "chemistry_reviewed",
    }
    pending_additional_notes = list(
        report.get("visual_review", {}).get("additional_review_notes") or []
    )
    pending_additional_notes.extend(additional_review_notes)
    for value in pending_additional_notes:
        if set(value) != note_keys:
            raise PaperExportRendererError(
                "additional_review_note_invalid", "额外目视复核备注字段不正确。"
            )
        if value["human_reviewed"] is not False or value["chemistry_reviewed"] is not False:
            raise PaperExportRendererError(
                "additional_review_note_invalid",
                "上游任务备注不能冒充化学人工审校或正式人工发布审核。",
            )
        for key in ("reviewer_kind", "note_zh", "scope_zh"):
            if not isinstance(value[key], str) or not value[key].strip():
                raise PaperExportRendererError(
                    "additional_review_note_invalid", "额外目视复核备注内容不完整。"
                )
        normalized_additional_notes.append(
            {
                "reviewer_kind": value["reviewer_kind"].strip(),
                "note_zh": value["note_zh"].strip(),
                "scope_zh": value["scope_zh"].strip(),
                "human_reviewed": False,
                "chemistry_reviewed": False,
            }
        )
    timestamp = reviewed_at or datetime.now(timezone.utc).replace(microsecond=0).isoformat()

    request = bundle["render_request"]
    receipts: list[dict[str, Any]] = []
    for artifact_id in (
        "student_docx",
        "student_pdf",
        "teacher_docx",
        "teacher_pdf",
    ):
        artifact = _artifact_by_id(report, artifact_id)
        page_key = "docx_render_pages" if artifact["format"] == "docx" else "pdf_render_pages"
        page_paths = {item["path"] for item in artifact[page_key]}
        page_reviews = [item for item in normalized_reviews if item["path"] in page_paths]
        page_receipt = {
            "artifact_id": artifact_id,
            "reviewed_at": timestamp,
            "reviewer_kind": reviewer_kind,
            "human_reviewed": False,
            "chemistry_reviewed": False,
            "pages": page_reviews,
            "layout_defect_count": 0,
        }
        receipts.append(
            {
                "schema_version": PAPER_FORMAT_SCHEMA_VERSION,
                "contract_kind": RENDER_RECEIPT_KIND,
                "artifact_id": artifact_id,
                "format": artifact["format"],
                "sha256": artifact["sha256"],
                "render_job_id": request["render_job_id"],
                "render_request_digest": request["render_request_digest"],
                "plan_digest": request["plan_digests"][artifact["audience"]],
                "preflight_digest": request["preflight_digest"],
                "page_count": artifact["page_count"],
                "rendered_page_count": artifact["rendered_page_count"],
                "page_review_receipt_sha256": _sha256_json(page_receipt),
                "all_pages_visual_pass": True,
                "metadata_privacy_pass": artifact["docx_audit"]["checks"][
                    "metadata_privacy_pass"
                ]
                and artifact["docx_audit"]["checks"]["secret_scan_pass"]
                and artifact["pdf_audit"]["checks"]["secret_scan_pass"],
                "content_version_id": request["content_version_id"],
                "blueprint_digest": request["blueprint_digest"],
                "question_anchor_digest": request["question_anchor_digest"],
                "answer_anchor_digest": request["answer_anchor_digest"],
                "figure_count": artifact["figure_count"],
            }
        )
    try:
        contract_report = validate_render_receipts(request, receipts)
    except PaperFormatContractError as exc:
        raise PaperExportRendererError(exc.code, exc.message_zh) from exc
    if contract_report["status"] != "local_delivery_candidate":
        raise PaperExportRendererError(
            "render_receipts_blocked", "四文件回执未通过既有渲染合同。"
        )

    report["status"] = "layout_smoke_passed"
    report["visual_review"] = {
        "status": "pass",
        "required_page_count": len(expected),
        "reviewed_pages": normalized_reviews,
        "layout_defect_count": 0,
        "reviewed_at": timestamp,
        "reviewer_kind": reviewer_kind,
        "human_reviewed": False,
        "chemistry_reviewed": False,
        "publication_allowed": False,
        "scope_zh": "仅检查截断、溢出、重叠、断题、字体替代和页眉页脚；不替代化学或人工发布审核。",
        "additional_review_notes": normalized_additional_notes,
    }
    report["receipts"] = receipts
    report["contract_validation_report"] = contract_report
    _write_json(output_root / "render_qa_report.json", report)
    _write_hash_manifest(output_root)
    return report


def _verification(value: Any) -> dict[str, Any]:
    return {
        "status": "project_template_value",
        "evidence_refs": [],
        "verified_for_paper_id": None,
        "note_zh": "本值只用于合成布局演示，未作为永久官方规则。",
    }


def _set_template_value(field: dict[str, Any], value: Any) -> None:
    field["value"] = value
    field["verification"] = _verification(value)


def _demo_dependency(kind: str, prior: Sequence[str]) -> dict[str, Any]:
    return {
        "kind": kind,
        "prior_atomic_part_ids": list(prior),
        "explicit_prior_edge_count": len(prior),
        "status": "validated_explicit",
    }


def _demo_atomic(
    atomic_id: str,
    printed_id: str,
    sequence: int,
    item_type: str,
    *,
    dependency_kind: str,
    prior: Sequence[str] = (),
) -> dict[str, Any]:
    return {
        "atomic_part_id": atomic_id,
        "printed_question_id": printed_id,
        "printed_question_number": str(sequence),
        "printed_sequence": sequence,
        "atomic_sequence_in_printed": 1,
        "item_type": item_type,
        "label_summary": {
            "primary_K": ["K-CO2-limewater"],
            "A": ["A-evidence-explanation"],
            "C": ["C-symbolic-representation"],
            "R": ["R-text-equation"],
            "RP": ["RP-text-to-symbol"],
        },
        "answer": {
            "availability": "present_part_aligned",
            "source_authority": "nonofficial",
            "has_quality_note": True,
        },
        "dependency": _demo_dependency(dependency_kind, prior),
        "detail_endpoint": f"/internal/synthetic-demo/{atomic_id}",
        "alias_units": [],
    }


def _build_demo_snapshot() -> dict[str, Any]:
    theme = {
        "theme": {
            "id": "SYNTH-CO2-T1",
            "title": "二氧化碳的检验与吸收（合成演示）",
            "sequence": 1,
            "sequence_status": "known_explicit",
            "page_span": {
                "start_page": 1,
                "end_page": 1,
                "page_numbers": [1],
                "status": "synthetic_layout_fixture",
            },
            "parent_chain_status": "complete",
        },
        "counts": {"atomic": 3},
        "shared_context": {
            "context_summary_zh": "同一实验现象贯穿选择、方程式书写与原因解释。",
            "context_status": "synthetic_layout_fixture",
            "material_count": 1,
            "materials": [
                {
                    "material_id": "SYNTH-CO2-M1",
                    "type": "text",
                    "page": 1,
                    "preview_allowed": True,
                    "used_by_atomic_count": 3,
                }
            ],
        },
        "dependencies": {
            "independent": 0,
            "shared_material_only": 2,
            "one_prior_part": 1,
            "multiple_prior_parts": 0,
            "per_alias_unit": 0,
            "explicit_prior_edge_count": 1,
            "blocked": 0,
        },
        "atomic_chain": [
            _demo_atomic(
                "SYNTH-CO2-A1",
                "SYNTH-CO2-P1",
                1,
                "choice_single",
                dependency_kind="shared_material_only",
            ),
            _demo_atomic(
                "SYNTH-CO2-A2",
                "SYNTH-CO2-P2",
                2,
                "fill_blank",
                dependency_kind="shared_material_only",
            ),
            _demo_atomic(
                "SYNTH-CO2-A3",
                "SYNTH-CO2-P3",
                3,
                "short_answer",
                dependency_kind="one_prior_part",
                prior=("SYNTH-CO2-A2",),
            ),
        ],
    }
    return {
        "data_snapshot_id": "snapshot_synthetic_co2_layout_v1",
        "scope": "synthetic_layout_demo",
        "catalog": {
            "schema_version": "shchem.theme-workbench.v1",
            "scope": "synthetic_layout_demo",
            "counts": {},
            "papers": [
                {
                    "paper": {
                        "id": "SYNTH-PAPER-CO2-V1",
                        "title": "合成主题式试卷导出演示",
                        "source_metadata": {
                            "claim_boundary": "synthetic_layout_fixture_not_a_shanghai_source_paper"
                        },
                    },
                    "theme_groups": [theme],
                }
            ],
            "unassigned_pending_review": {
                "status": "none",
                "count": 0,
                "reason_zh": None,
                "atomic_chain": [],
            },
            "authority": {"publication_allowed": False},
            "integrity": {
                "hash_verified_on_read": True,
                "semantic_invariants_verified_on_read": True,
                "complete_scope_coverage": True,
                "no_duplicate_atomic_parts": True,
                "explicit_order_only": True,
                "dependency_edges_validated": True,
                "answer_text_excluded": True,
                "pixel_reuse_allowed": False,
                "fail_closed": True,
            },
        },
    }


class _SyntheticDemoResolver:
    _scores: ClassVar[dict[str, int]] = {
        "SYNTH-CO2-A1": 2,
        "SYNTH-CO2-A2": 3,
        "SYNTH-CO2-A3": 5,
    }

    def resolve_shared_material(self, reference: Mapping[str, Any]) -> dict[str, Any]:
        if reference["material_id"] != "SYNTH-CO2-M1":
            raise PaperExportRendererError("demo_material_unknown", "合成演示材料标识不正确。")
        return {
            "content_blocks": [
                {
                    "block_type": "paragraph",
                    "text_zh": (
                        "某学习小组将二氧化碳通入澄清石灰水。通入少量二氧化碳时产生白色浑浊；"
                        "继续通入过量二氧化碳后，浑浊逐渐消失。"
                    ),
                    "asset_ref": None,
                    "alt_text_zh": None,
                }
            ],
            "source_label_zh": "合成布局演示材料；非上海原题；不可发布。",
        }

    def resolve_atomic_part(self, reference: Mapping[str, Any]) -> dict[str, Any]:
        atomic_id = reference["atomic_part_id"]
        payloads = {
            "SYNTH-CO2-A1": {
                "question": (
                    "通入少量 CO₂ 时，产生白色浑浊的主要物质是（　　）。\n"
                    "A. CaCO₃　　B. CaO　　C. CaCl₂　　D. Ca"
                ),
                "answer": "A（CaCO₃）",
                "explanation": "少量 CO₂ 与 Ca(OH)₂ 反应生成难溶的 CaCO₃。",
                "pitfalls": ["不要把后续过量 CO₂ 的反应产物提前作为本题答案。"],
                "lines": 1,
            },
            "SYNTH-CO2-A2": {
                "question": "填空：少量 CO₂ 与澄清石灰水反应的化学方程式为________________。",
                "answer": "CO₂ + Ca(OH)₂ → CaCO₃↓ + H₂O",
                "explanation": "该反应生成碳酸钙沉淀，对应共同材料中的白色浑浊。",
                "pitfalls": ["须保留 CaCO₃ 的沉淀符号“↓”。"],
                "lines": 2,
            },
            "SYNTH-CO2-A3": {
                "question": (
                    "结合第 2 题的产物，解释继续通入过量 CO₂ 后浑浊逐渐消失的原因，"
                    "并写出对应的化学方程式。"
                ),
                "answer": (
                    "CaCO₃ 与过量 CO₂ 和 H₂O 反应生成可溶的 Ca(HCO₃)₂："
                    "CaCO₃ + CO₂ + H₂O → Ca(HCO₃)₂。"
                ),
                "explanation": (
                    "本题显式依赖第 2 题生成的 CaCO₃；继续通入 CO₂ 后，固体转化为可溶的"
                    " Ca(HCO₃)₂，因此浑浊消失。"
                ),
                "pitfalls": ["原因解释与方程式应同时体现“过量 CO₂”和可溶产物。"],
                "lines": 5,
            },
        }
        if atomic_id not in payloads:
            raise PaperExportRendererError("demo_atomic_unknown", "合成演示小题标识不正确。")
        item = payloads[atomic_id]
        return {
            "question_blocks": [
                {
                    "block_type": "paragraph",
                    "text_zh": item["question"],
                    "asset_ref": None,
                    "alt_text_zh": None,
                }
            ],
            "score": self._scores[atomic_id],
            "answer_space_lines": item["lines"],
            "reference_answer": {
                "status": "aligned",
                "text_zh": item["answer"],
                "authority_label": "nonofficial",
                "independently_verified": False,
                "source_label_zh": "合成演示教师候选答案；非官方；未独立核验。",
            },
            "explanation_zh": item["explanation"],
            "explanation_label": "teacher_analysis",
            "pitfalls_zh": item["pitfalls"],
            "source_label_zh": "合成布局演示；非上海原题；非官方教师候选。",
        }


def build_synthetic_demo_bundle() -> dict[str, Any]:
    preset = default_shanghai_theme_preset()
    _set_template_value(preset["per_paper"]["theme_count"], 1)
    _set_template_value(preset["per_paper"]["total_score"], 10)
    _set_template_value(preset["per_paper"]["duration_minutes"], 15)
    _set_template_value(
        preset["per_paper"]["scoring_rules"],
        {
            "selection_rule_zh": "选择小问嵌入主题，按本卷标注分值计分。",
            "partial_credit_rule_zh": "简答题按完整因果链与方程式要点建议给分。",
            "other_rule_zh": "本卷仅为合成布局演示，不作正式评价使用。",
        },
    )
    _set_template_value(
        preset["structure"]["subquestion_numbering"], "continuous_across_paper"
    )
    snapshot = _build_demo_snapshot()

    def loader(scope: str, expected_data_snapshot_id: str) -> Mapping[str, Any]:
        if scope != snapshot["scope"] or expected_data_snapshot_id != snapshot["data_snapshot_id"]:
            raise PaperExportRendererError("demo_snapshot_mismatch", "合成演示快照不一致。")
        return deepcopy(snapshot)

    selection = {
        "scope": snapshot["scope"],
        "selection_unit": "theme",
        "theme_id": "SYNTH-CO2-T1",
        "target_atomic_id": None,
        "expected_data_snapshot_id": snapshot["data_snapshot_id"],
    }
    blueprint = build_assembly_blueprint(
        [selection],
        theme_loader=loader,
        expected_data_snapshot_id=snapshot["data_snapshot_id"],
        preset=preset,
    )
    plans = build_document_plans(
        blueprint,
        preset=preset,
        paper_metadata={
            "title_zh": "上海高中化学主题式试卷导出演示",
            "subtitle_zh": "完整主题小样：共同材料 + 混合作答形态 + 前序依赖",
            "version_label_zh": "学生/教师双版本 v1",
        },
        content_resolver=_SyntheticDemoResolver(),
    )
    preflight = run_export_preflight(
        preset=preset,
        blueprint=blueprint,
        student_plan=plans["student"],
        teacher_plan=plans["teacher"],
    )
    request = build_render_request(
        preset=preset,
        student_plan=plans["student"],
        teacher_plan=plans["teacher"],
        preflight_report=preflight,
    )
    bundle = {
        "renderer_schema_version": RENDERER_SCHEMA_VERSION,
        "contract_kind": RENDER_BUNDLE_KIND,
        "bundle_id": "synthetic_co2_theme_export_demo_v1",
        "claim_boundary_zh": "合成布局演示；非上海原题；未经过化学人工双检；不可发布。",
        "preset": preset,
        "blueprint": blueprint,
        "student_plan": plans["student"],
        "teacher_plan": plans["teacher"],
        "preflight_report": preflight,
        "render_request": request,
        "publication_allowed": False,
    }
    return validate_render_bundle(bundle)


def _default_toolchain_from_args(args: argparse.Namespace) -> RendererToolchain:
    return RendererToolchain(
        python_exe=Path(args.python_exe).resolve(),
        render_docx_script=Path(args.render_docx_script).resolve(),
        pdftoppm_exe=Path(args.pdftoppm_exe).resolve(),
        dpi=args.dpi,
        conversion_backend=args.conversion_backend,
    )


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="上海高中化学主题式试卷 DOCX/PDF 独立渲染器"
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    demo = subparsers.add_parser("demo", help="生成并渲染内置合成完整主题")
    demo.add_argument("--output-dir", required=True)
    demo.add_argument("--python-exe", required=True)
    demo.add_argument("--render-docx-script", required=True)
    demo.add_argument("--pdftoppm-exe", required=True)
    demo.add_argument("--dpi", type=int, default=300)
    demo.add_argument(
        "--conversion-backend",
        choices=("auto", "libreoffice", "word_com"),
        default="auto",
    )

    render = subparsers.add_parser("render", help="渲染一个冻结 bundle JSON")
    render.add_argument("--bundle", required=True)
    render.add_argument("--output-dir", required=True)
    render.add_argument("--asset-root")
    render.add_argument("--python-exe", required=True)
    render.add_argument("--render-docx-script", required=True)
    render.add_argument("--pdftoppm-exe", required=True)
    render.add_argument("--dpi", type=int, default=300)
    render.add_argument(
        "--conversion-backend",
        choices=("auto", "libreoffice", "word_com"),
        default="auto",
    )

    finalize = subparsers.add_parser(
        "finalize-qa", help="在逐页目视检查后形成合同回执"
    )
    finalize.add_argument("--output-dir", required=True)
    finalize.add_argument("--review-json", required=True)
    finalize.add_argument("--additional-review-notes-json")
    finalize.add_argument("--reviewed-at")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    try:
        if args.command == "demo":
            bundle = build_synthetic_demo_bundle()
            report = render_export_bundle(
                bundle,
                output_dir=Path(args.output_dir),
                toolchain=_default_toolchain_from_args(args),
            )
            print(json.dumps({"status": report["status"], "output_dir": args.output_dir}, ensure_ascii=False))
            return 0 if report["status"] == "awaiting_visual_review" else 2
        if args.command == "render":
            bundle = _read_json(Path(args.bundle))
            report = render_export_bundle(
                bundle,
                output_dir=Path(args.output_dir),
                toolchain=_default_toolchain_from_args(args),
                asset_root=Path(args.asset_root).resolve() if args.asset_root else None,
            )
            print(json.dumps({"status": report["status"], "output_dir": args.output_dir}, ensure_ascii=False))
            return 0 if report["status"] == "awaiting_visual_review" else 2
        reviews = _read_json(Path(args.review_json))
        if not isinstance(reviews, list):
            raise PaperExportRendererError(
                "visual_review_invalid", "逐页目视复核文件必须是列表。"
            )
        additional_notes = (
            _read_json(Path(args.additional_review_notes_json))
            if args.additional_review_notes_json
            else []
        )
        if not isinstance(additional_notes, list):
            raise PaperExportRendererError(
                "additional_review_note_invalid", "额外目视复核备注文件必须是列表。"
            )
        report = finalize_visual_review(
            output_dir=Path(args.output_dir),
            reviewed_pages=reviews,
            reviewed_at=args.reviewed_at,
            additional_review_notes=additional_notes,
        )
        print(json.dumps({"status": report["status"], "output_dir": args.output_dir}, ensure_ascii=False))
        return 0
    except (PaperExportRendererError, PaperFormatContractError) as exc:
        code = getattr(exc, "code", "paper_export_renderer_failed")
        message = getattr(exc, "message_zh", str(exc))
        print(json.dumps({"status": "blocked", "code": code, "message_zh": message}, ensure_ascii=False), file=sys.stderr)
        return 2


__all__ = [
    "ARTIFACT_FILENAMES",
    "HASH_MANIFEST_KIND",
    "QA_REPORT_KIND",
    "RENDERER_SCHEMA_VERSION",
    "RENDERER_VERSION",
    "RENDER_BUNDLE_KIND",
    "PaperExportRendererError",
    "RendererToolchain",
    "audit_docx",
    "build_docx_from_plan",
    "build_synthetic_demo_bundle",
    "finalize_visual_review",
    "render_export_bundle",
    "validate_render_bundle",
]


if __name__ == "__main__":
    raise SystemExit(main())
