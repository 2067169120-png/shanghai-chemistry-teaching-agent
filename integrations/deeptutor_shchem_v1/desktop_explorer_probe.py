"""Isolated source/EXE navigation and batch-image acceptance; no teacher data."""
from __future__ import annotations
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import sys
import tempfile
from time import perf_counter, monotonic, sleep


def run_probe(output):
    from PIL import Image, ImageDraw
    from docx import Document
    from docx.shared import Inches
    from PySide6.QtCore import QRect
    from PySide6.QtWidgets import QApplication
    from .desktop_facade import build_default_facade
    from .desktop_paths import DesktopPaths
    from .desktop_version import DESKTOP_VERSION
    from .desktop_workbench.app import create_application
    from .desktop_workbench.main_window import TeacherWorkbenchWindow
    from .desktop_workbench.question_explorer_page import QuestionExplorerPage
    from .desktop_workbench.dialogs import ImportDialog
    from .desktop_number_regions import NumberRegionStore
    from .desktop_state import DesktopStateStore
    from .desktop_backup import plan_backup, create_backup, inspect_backup, restore_backup, manifest_revision, summary_text
    app=create_application(['explorer195-probe'])
    output=Path(output);output.mkdir(parents=True,exist_ok=True)
    screenshots=[];errors=[]
    oldhook=sys.excepthook
    sys.excepthook=lambda typ,value,tb:errors.append(typ.__name__)
    def settle(predicate=lambda:True):
        end=monotonic()+45
        for _ in range(5):app.processEvents();sleep(.025)
        while not predicate() and monotonic()<end:app.processEvents();sleep(.02)
        assert predicate(), 'Question page did not complete requested operation'
    def capture(widget,name):
        path=output/name;assert widget.grab().save(str(path))
        screenshots.append({'file':name,'sha256':sha256(path.read_bytes()).hexdigest()})
    with tempfile.TemporaryDirectory(prefix='explorer195-') as temp:
        temp=Path(temp)
        root=Path(getattr(sys,'_MEIPASS',Path(__file__).resolve().parents[2]))
        facade=build_default_facade(DesktopPaths.from_workspace(root,state_root=temp/'state'))
        document=Document();source_images=[]
        for n in range(16):
            document.add_paragraph(f'【例{n+1}】合成软件验收题：观察图示，说明条件变化。')
            for i in range(8 if n==0 else 1):
                im=Image.new('RGB',(360,120),'white');draw=ImageDraw.Draw(im)
                draw.rectangle((20,25,330,95),outline='black',width=2)
                draw.text((35,40),f'SYNTHETIC  {n+1} / {i+1}   12.5 mol/L',fill='black')
                data=BytesIO();im.save(data,format='PNG');source_images.append(data.getvalue())
                document.add_picture(BytesIO(data.getvalue()),width=Inches(3.6))
            document.add_paragraph('【答案】合成参考说明，不是正式化学题或课堂评价。')
        path=temp/'多图验收讲义.docx';document.save(path);original=path.read_bytes()
        facade.save_visual_import_batch(handout_files=(path,),source_type='教师讲义')
        calls={'catalog':0,'resolve':0}
        original_catalog=facade.word_question_catalog
        def catalog():calls['catalog']+=1;return original_catalog()
        facade.word_question_catalog=catalog
        service=facade._word_questions();original_resolve=service._resolve
        def resolve(*a,**kw):calls['resolve']+=1;return original_resolve(*a,**kw)
        service._resolve=resolve
        win=TeacherWorkbenchWindow(facade);win.resize(1280,900);win.show()
        try:
            settle(lambda:not win.home_page._loading)
            page=win.library_page
            assert not page._started and calls['catalog']==0
            page.scope.setCurrentIndex(page.scope.findData('word_native'))
            start=perf_counter();win.navigate('library')
            settle(lambda:not page._loading and len(page.cards)==8 and page.cards[0].ready and bool(page._reader.load_metrics))
            first_complete=(perf_counter()-start)*1000
            assert len(page._reader.required)==8 and calls=={'catalog':1,'resolve':1},calls
            assert page._reader.loaded_tabs=={0}
            first=deepcopy(page.load_metrics);first['word_images']=deepcopy(page._reader.load_metrics)
            # Data-ready is not enough: the teacher must actually see the title,
            # source, reader and full action buttons inside the card.
            settle(lambda:page.cards[0].height() >= page._reader.height()+100)
            geometry=[]
            for card in page.cards:
                title=card.root.itemAt(1).widget()
                assert title.text() and title.height() >= title.fontMetrics().height()
                for widget in (title,card.open,card.add):
                    assert card.rect().contains(QRect(widget.mapTo(card,widget.rect().topLeft()),widget.size()))
                geometry.append({'height':card.height(),'minimum_height':card.minimumHeight(),
                                 'title_height':title.height(),'title_font_height':title.fontMetrics().height()})
            assert page.list_body.height()>page.scroll.viewport().height()
            (output/'card-geometry.json').write_text(json.dumps(geometry,indent=2),encoding='utf-8')
            capture(win,'explorer-loaded.png')
            current_index=page._sessions['word_native'].index
            page.next.click();settle(lambda:not page._loading and page._page==1 and page.cards[0].ready)
            assert page._sessions['word_native'].index is current_index
            assert page.load_metrics['query_cache_hit'] and calls['catalog']==1
            paging=deepcopy(page.load_metrics)
            page.previous.click();settle(lambda:not page._loading and page._page==0 and page.cards[0].ready)
            page.cards[0].add.click();settle(lambda:len(facade.basket())==1 and not page._adding)
            basket=deepcopy(facade.basket())
            page.query.setText('不存在的合成检索');page.search();settle(lambda:not page._loading)
            assert not page.cards and calls['catalog']==1 and facade.basket()==basket
            capture(win,'explorer-empty-guidance.png')
            page.copy_load_metrics();copied=json.loads(QApplication.clipboard().text())
            assert '不存在的合成检索' not in str(copied) and str(temp) not in str(copied)
            another=QuestionExplorerPage(facade,win.tasks)
            assert another.scope.currentData()=='word_native' and not another._started
            another.close();another.deleteLater()
            # A new explicit snapshot is required after import/edit or manual reload.
            page.query.clear();page.invalidate_catalogs()
            settle(lambda:not page._loading and len(page.cards)==8 and page.cards[0].ready)
            assert calls['catalog']==2 and page._sessions['word_native'].index is not current_index
            dialog=ImportDialog(facade,win.tasks,win);dialog.show();settle()
            assert dialog.question_files.mapTo(dialog,dialog.question_files.rect().topLeft()).y()<dialog.source_help_card.mapTo(dialog,dialog.source_help_card.rect().topLeft()).y()
            capture(dialog,'import-files-first.png');dialog.close();dialog.deleteLater()
            # Only geometry and original identity enter the backup; imported bank stays outside.
            image={'image_sha256':sha256(source_images[0]).hexdigest(),'width':360,'height':120}
            NumberRegionStore(facade.state_store).remember(image,[dict(image_sha256=image['image_sha256'],box=[4,4,50,30],number=8,punctuation='.')])
            plan=plan_backup(facade.paths.state_root);archive=temp/'saved-work.zip';create_backup(plan,archive)
            checked=inspect_backup(archive);target=temp/'restored';restore_backup(archive,target,expected_manifest=manifest_revision(checked))
            positions=NumberRegionStore(DesktopStateStore(target)).load(image)
            assert positions==[{'box':[4,4,50,30],'punctuation':'.'}]
            assert checked['summary']['number_regions']==1
            assert path.read_bytes()==original and facade.basket()==basket
            assert not errors,errors
            (output/'backup-summary.txt').write_text(summary_text(checked),encoding='utf-8')
            report={'version':DESKTOP_VERSION,'source_commit':os.environ.get('GITHUB_SHA','local'),
                'frozen':getattr(sys,'frozen',False),'qt_platform':app.platformName(),
                'synthetic_questions':16,'first_visible_images':8,'first_page_ready_ms':round(first_complete,2),
                'first_load':first,'next_page':paging,'diagnostic_sample':copied,
                'checks':{'hidden_page_no_catalog':True,'one_resolution_for_eight_images':True,
                  'answer_tab_lazy':True,'paging_reuses_index':True,'source_preference_retained':True,
                  'import_invalidates_index':True,'generic_files_before_old_pack':True,
                  'region_archive_roundtrip':True,'source_and_basket_unchanged':True,'diagnostic_no_content':True,'card_contents_visible':True},
                'screenshots':screenshots,'uncaught_errors':errors,'model_calls':0,
                'scope':'Synthetic Word data in isolated state. Event-loop readiness timings, not user hardware or actual classroom acceptance.'}
            (output/'explorer-probe.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        finally:
            win.close();app.processEvents();sys.excepthook=oldhook
    return 0
