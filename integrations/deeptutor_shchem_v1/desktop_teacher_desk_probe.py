"""Explicit synthetic teacher-desk acceptance; not loaded in normal operation."""
from __future__ import annotations
from copy import deepcopy
from pathlib import Path


def exercise(window, settle, capture, output):
    from docx import Document
    from PySide6.QtCore import QPoint
    from .desktop_workbench.main_window import ALL_ROUTES
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    prep, facade = window.preparation_page, window.facade
    for title in ("物质分类 · 新授课", "离子反应 · 课后讲评", "原电池 · 实验探究",
                  "化学反应速率 · 练习", "氧化还原反应 · 单元复习"):
        prep.topic.setText(title + "（合成演示）")
        prep.audience.setText("软件验收")
        prep.objective.setPlainText("保留材料依据，核对作答条件。仅用于软件运行验收。")
        prep.materials.setPlainText("合成演示材料：不作为真实试题或教师课程。")
        facade.create_preparation_draft(prep._payload())
    doc = Document()
    doc.add_heading("首页题篮演示（合成，非真实试卷）", 0)
    doc.add_paragraph("【例1】恒温密闭容器中，A(g) ⇌ B(g) 达到平衡时，正、逆反应速率的关系是______。")
    doc.add_paragraph("【答案】正、逆反应速率相等且不为零。")
    doc.add_paragraph("【例2】对于 A(g) ⇌ B(g)，某温度下平衡时 c(A)=0.20 mol/L，c(B)=0.60 mol/L。计算平衡常数。")
    doc.add_paragraph("【答案】K=c(B)/c(A)=3.0。")
    path = output / "首页合成题篮.docx"
    doc.save(path)
    facade.save_visual_import_batch(handout_files=(path,), source_type="教师讲义")
    rows = [r for r in facade.word_question_catalog()["items"] if r["source_name"] == path.name]
    assert len(rows) == 2
    facade.add_word_questions_to_basket([{"key": r["key"], "revision": r["revision"]} for r in rows])
    prep.topic.setText("化学平衡 · 第二课时（正在编辑演示）")
    before = deepcopy(prep._payload())
    saved = deepcopy(facade.state_store.snapshot())
    window.navigate("home"); settle(lambda: not window.home_page._loading)
    home = window.home_page
    assert len(home.records) == 5 and home.basket_count.text() == "已选 2 项"
    assert not home.diagnostics.toggle.isChecked()
    home.table.selectRow(0)
    capture(window, "home-current.png")
    home.resume_button.click()
    assert prep._payload() == before and window.stack.currentWidget() is prep
    window.navigate("home"); settle(lambda: not home._loading)
    assert facade.state_store.snapshot() == saved
    window.resize(800, 700); settle()
    capture(window, "home-compact.png")
    window.navigate("mywork")
    page = window.my_work_page
    page.query.clear(); page.shelves.setCurrentIndex(0); page.refresh()
    settle(lambda: not page._loading)
    page.results.setCurrentRow(0); page.results.scrollToBottom(); settle()
    for control in (page.open_button, page.rename_button, page.archive_button, page.backup_button):
        point = control.mapTo(page, QPoint(0, 0))
        assert control.isVisible() and 0 <= point.y() and point.y() + control.height() <= page.height()
    assert window.size().width() == 800 and window.size().height() == 700
    capture(window, "works-pinned-actions.png")
    window.resize(1360, 900); settle()
    # These are real page views, not evidence every underlying business is complete.
    for route in ALL_ROUTES:
        window.navigate(route)
        if route == "home": settle(lambda: not home._loading)
        elif route == "mywork": settle(lambda: not page._loading)
        else: settle()
        capture(window, "desk-page-" + route + ".png")
    return {"recent_current_five": True, "real_word_basket": True,
            "resume_preserved_editor": True, "home_did_not_write_state": True,
            "collapsed_diagnostics": True, "actions_visible_800x700": True,
            "model_calls": 0}
