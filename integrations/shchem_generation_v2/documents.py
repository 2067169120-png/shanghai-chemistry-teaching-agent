from __future__ import annotations

import re
from datetime import UTC, datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


BODY_CJK = "宋体"
BODY_LATIN = "Times New Roman"
HEADING_CJK = "黑体"


def _set_run_font(run, *, size: float = 10.5, bold: bool = False, cjk: str = BODY_CJK) -> None:
    run.font.name = BODY_LATIN
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cjk)
    run._element.rPr.rFonts.set(qn("w:ascii"), BODY_LATIN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_LATIN)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, *, top: int = 70, start: int = 90, bottom: int = 70, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _add_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))
    _set_run_font(run, size=8.5)


def _set_picture_alt_text(inline_shape, text: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", text)
    doc_pr.set("title", text)


def _configure_document(document: Document, *, title: str, version_id: str) -> None:
    section = document.sections[0]
    document.settings.odd_and_even_pages_header_footer = True
    section.different_first_page_header_footer = False
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.header_distance = Mm(7)
    section.footer_distance = Mm(8)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = BODY_LATIN
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_CJK)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(15)
    normal.paragraph_format.space_after = Pt(3)

    header_text = f"机器审核候选｜教师私域交付候选（非官方、非公开出版）｜{version_id}"
    for header in (section.header, section.even_page_header):
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = hp.add_run(header_text)
        _set_run_font(run, size=8.5, cjk=HEADING_CJK)
        run.font.color.rgb = RGBColor(80, 80, 80)

    for footer in (section.footer, section.even_page_footer):
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run("第 ")
        _set_run_font(run, size=8.5)
        _add_field(fp, "PAGE")
        run = fp.add_run(" 页 / 共 ")
        _set_run_font(run, size=8.5)
        _add_field(fp, "NUMPAGES")
        run = fp.add_run(" 页")
        _set_run_font(run, size=8.5)

    props = document.core_properties
    props.title = title
    props.subject = "上海高中化学主题化原创机器候选"
    props.author = "SHCHEM generation v2 machine-only pipeline"
    props.keywords = f"machine-only, suggested answers, {version_id}"
    props.comments = "非官方；无人类审核；不得声称官方采分点。"


def _add_masthead(document: Document, *, title: str, subtitle: str, metadata: Iterable[str]) -> None:
    table = document.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Mm(174)
    top = table.cell(0, 0)
    bottom = table.cell(1, 0)
    _set_cell_shading(top, "E8EDF2")
    for cell in (top, bottom):
        _set_cell_margins(cell, top=110, bottom=110, start=140, end=140)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = top.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    _set_run_font(r, size=16, bold=True, cjk=HEADING_CJK)
    p = bottom.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    _set_run_font(r, size=10.5, bold=True, cjk=HEADING_CJK)
    p = bottom.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("　｜　".join(metadata))
    _set_run_font(r, size=9)
    document.add_paragraph()


def _add_notice(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    _set_cell_shading(cell, "F2F2F2")
    _set_cell_margins(cell, top=90, bottom=90, start=120, end=120)
    p = cell.paragraphs[0]
    r = p.add_run(text)
    _set_run_font(r, size=9, cjk=HEADING_CJK)
    document.add_paragraph()


def _add_heading(document: Document, text: str, *, level: int = 1, page_break_before: bool = False) -> None:
    p = document.add_paragraph()
    p.paragraph_format.page_break_before = page_break_before
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    _set_run_font(r, size=12 if level == 1 else 11, bold=True, cjk=HEADING_CJK)


def _add_question(document: Document, part: dict, *, solutions: bool) -> None:
    p = document.add_paragraph()
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(3)
    r = p.add_run(f"{part['label']}．")
    _set_run_font(r, bold=True, cjk=HEADING_CJK)
    r = p.add_run(part["prompt"])
    _set_run_font(r)
    r = p.add_run(f"（{part['score']}分）")
    _set_run_font(r, size=9)
    for option in part.get("options", []):
        op = document.add_paragraph()
        op.paragraph_format.left_indent = Mm(7)
        op.paragraph_format.first_line_indent = Mm(-4)
        op.paragraph_format.space_after = Pt(1)
        r = op.add_run(option)
        _set_run_font(r)
    if not solutions:
        response_lines = 2 if part["response_R"] in {"short_explanation", "calculation"} else 1
        if part["response_R"] != "single_choice":
            for _ in range(response_lines):
                p = document.add_paragraph("________________________________________________________________________________")
                p.paragraph_format.space_after = Pt(1)
                for run in p.runs:
                    _set_run_font(run, size=8)
        return

    table = document.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Mm(168)
    answer_cell = table.cell(0, 0)
    point_cell = table.cell(1, 0)
    _set_cell_shading(answer_cell, "EEF4F8")
    for cell in (answer_cell, point_cell):
        _set_cell_margins(cell)
    p = answer_cell.paragraphs[0]
    r = p.add_run("建议答案：")
    _set_run_font(r, size=9.5, bold=True, cjk=HEADING_CJK)
    r = p.add_run(part["answer"]["value"])
    _set_run_font(r, size=9.5)
    p = point_cell.paragraphs[0]
    r = p.add_run("建议采分点（非官方）：")
    _set_run_font(r, size=9, bold=True, cjk=HEADING_CJK)
    r = p.add_run("；".join(part["answer"].get("score_points", [])))
    _set_run_font(r, size=9)


def build_exam_docx(
    paper: dict,
    figure_png: Path,
    output_path: Path,
    *,
    solutions: bool,
) -> None:
    document = Document()
    kind = "教师解析与建议采分点" if solutions else "学生卷"
    title = f"2026主题化化学时事模拟卷（{kind}）"
    _configure_document(document, title=title, version_id=paper["version_id"])
    task = paper["task_card"]
    _add_masthead(
        document,
        title=title,
        subtitle="原创机器候选｜无人工评审｜非官方",
        metadata=[
            f"时间 {task['duration_minutes']['value']} 分钟",
            f"满分 {task['total_score']['value']} 分",
            f"共 {task['theme_count']['value']} 个主题",
        ],
    )
    _add_notice(
        document,
        "纯机器审核候选｜非官方｜无人工审定。仅在随附manifest确认双审、对抗与parity全部通过后，"
        "可由教师本人管理并发给学生；不得自动外发。所有答案与采分点均为机器建议，"
        "不代表上海市教育考试院、学校、教师或专家意见，不可表述为官方采分点。",
    )
    _add_heading(document, "作答说明", level=2)
    p = document.add_paragraph(
        "全卷按主题大题组织，选择、填空、方程式、实验评价和计算均嵌入主题；题号连续。"
        "除注明外，计算结果保留题目所要求的有效数字。化学方程式须写明必要的物态、条件或物质名称。"
    )
    for run in p.runs:
        _set_run_font(run)
    if solutions:
        _add_notice(document, "解析标签说明：以下仅称“建议答案”“建议采分点”；易错点与证据边界另行说明，不存在人工审核或官方审定。")

    for theme_index, theme in enumerate(paper["themes"]):
        _add_heading(
            document,
            f"主题{theme['order']}　{theme['title']}（20分）",
            page_break_before=bool(theme_index),
        )
        p = document.add_paragraph(theme["shared_material"])
        p.paragraph_format.keep_with_next = True
        for run in p.runs:
            _set_run_font(run, size=9.5)
        for question in theme["printed_questions"]:
            part = question["atomic_parts"][0]
            if part["part_id"] == "P02":
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.keep_with_next = True
                shape = p.add_run().add_picture(str(figure_png), width=Mm(125))
                _set_picture_alt_text(shape, f"SH_CHEM_FIGURE:{paper['figure_refs'][0]}")
                cap = document.add_paragraph("图1　模拟含CuSO₄废液低压直流电解回收铜装置（独立SVG重绘）")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.keep_with_next = True
                for run in cap.runs:
                    _set_run_font(run, size=8.5)
            _add_question(document, part, solutions=solutions)

        if solutions:
            _add_heading(document, "本主题证据边界", level=2)
            boundary = "；".join(
                sorted(
                    {
                        "热点卡仅提供经核验事实或数值边界，不是正式题目",
                        "建议答案不代表官方评分细则",
                        "认知难度为预标，无真实学生作答数据",
                    }
                )
            )
            p = document.add_paragraph(boundary)
            for run in p.runs:
                _set_run_font(run, size=9)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def build_week_docx(
    paper: dict,
    plan: dict,
    output_path: Path,
    *,
    solutions: bool,
    content_bundle: dict | None = None,
) -> None:
    document = Document()
    kind = "教师解析版" if solutions else "学生版"
    title = f"完整示范周包（{kind}）"
    _configure_document(document, title=title, version_id=paper["version_id"])
    _add_masthead(
        document,
        title=title,
        subtitle="7日主题化训练＋第14日延时复测｜原创机器候选｜无人工评审",
        metadata=[f"周包编号 {plan['weekpack_id']}", f"关联试卷 {paper['paper_id']}"],
    )
    _add_notice(
        document,
        "纯机器审核候选｜非官方｜无人工审定。本周包未使用真实学生隐私或作答数据，不构成个体诊断。"
        "仅在随附manifest确认双审、对抗与parity全部通过后，可由教师本人管理并发给学生；不得自动外发。"
        "答案和采分点均为机器建议，非官方、非教师审核。",
    )
    table = document.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Mm(14), Mm(44), Mm(46), Mm(24), Mm(46))
    headers = ("日", "重点", "题目", "建议时长", "自检")
    for index, (cell, header, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = width
        _set_cell_shading(cell, "DCE6EF")
        _set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        _set_run_font(r, size=9, bold=True, cjk=HEADING_CJK)
    _set_repeat_table_header(table.rows[0])
    for day in plan["days"]:
        cells = table.add_row().cells
        values = (
            str(day["day"]),
            day["focus"],
            "、".join(day["question_refs"]),
            f"{day['target_minutes']}分钟",
            day["self_check"],
        )
        for cell, value, width in zip(cells, values, widths):
            cell.width = width
            _set_cell_margins(cell)
            p = cell.paragraphs[0]
            r = p.add_run(value)
            _set_run_font(r, size=8.7)
        _cant_split(table.rows[-1])

    _add_heading(document, "每日执行页")
    part_by_qid = {
        q["printed_question_id"]: q["atomic_parts"][0]
        for theme in paper["themes"]
        for q in theme["printed_questions"]
    }
    for day_index, day in enumerate(plan["days"]):
        if day_index:
            document.add_page_break()
        _add_heading(document, f"第{day['day']}日　{day['focus']}（建议{day['target_minutes']}分钟）")
        p = document.add_paragraph(f"当日自检：{day['self_check']}")
        for run in p.runs:
            _set_run_font(run, size=9.5, bold=True, cjk=HEADING_CJK)
        for qid in day["question_refs"]:
            _add_question(document, part_by_qid[qid], solutions=solutions)
        if solutions:
            p = document.add_paragraph("复盘建议：先独立反解计算与方程守恒，再核对条件、分母和证据边界；不要用答案命中代替过程判断。")
            for run in p.runs:
                _set_run_font(run, size=9)

    if content_bundle is not None:
        _add_heading(document, "合成学生画像内容示范（非真实学生效果）", page_break_before=True)
        _add_notice(
            document,
            "以下内容由synthetic_student_profile驱动，仅用于证明内容选择、治理绑定与导出接口；"
            "不含真实学生数据，不代表诊断结论或学习效果。题目状态为automated_verified_content。",
        )
        for section in content_bundle["demonstration_selection"]["sections"]:
            _add_heading(document, section["label"], level=2)
            for item in section["items"]:
                p = document.add_paragraph()
                r = p.add_run(f"{item['atomic_part_id']}　{item['prompt']}")
                _set_run_font(r, bold=True, cjk=HEADING_CJK)
                for option in item.get("options", []):
                    p = document.add_paragraph(option)
                    for run in p.runs:
                        _set_run_font(run, size=9.5)
                if solutions:
                    for label, text in (
                        ("建议答案（非官方）", item["suggested_answer"]),
                        ("详细解析", item["detailed_explanation"]),
                        ("建议采分点（非官方）", "；".join(item["suggested_score_points"])),
                    ):
                        p = document.add_paragraph()
                        r = p.add_run(f"{label}：")
                        _set_run_font(r, bold=True, cjk=HEADING_CJK)
                        r = p.add_run(str(text))
                        _set_run_font(r)
                else:
                    p = document.add_paragraph("作答：" + "_" * 58)
                    for run in p.runs:
                        _set_run_font(run)
                governance = item["governance"]
                p = document.add_paragraph(
                    f"治理证据：{governance['state']}｜{governance['child_chain_path']}｜"
                    f"sha256={governance['child_chain_sha256']}｜central_registry_exact_member=true"
                )
                for run in p.runs:
                    _set_run_font(run, size=7.5)
        else:
            p = document.add_paragraph("完成记录：用时 ______ 分钟　　确定性错误 ______ 个　　证据边界错误 ______ 个")
            for run in p.runs:
                _set_run_font(run, size=9)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def scrub_docx_metadata(path: Path) -> None:
    document = Document(path)
    props = document.core_properties
    props.last_modified_by = "SHCHEM generation v2"
    props.revision = 1
    fixed = datetime(2000, 1, 1, tzinfo=UTC)
    props.created = fixed
    props.modified = fixed
    props.last_printed = fixed
    document.save(path)


def docx_visible_text(path: Path) -> str:
    document = Document(path)
    chunks: list[str] = []
    for section in document.sections:
        chunks.extend(paragraph.text for paragraph in section.header.paragraphs if paragraph.text)
        chunks.extend(paragraph.text for paragraph in section.footer.paragraphs if paragraph.text)
    for paragraph in document.paragraphs:
        if paragraph.text:
            chunks.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells if cell.text)
    return re.sub(r"\s+", "", "\n".join(chunks))
